"""Build 7.Review_1_PPT.pptx from D:\\abstracts\\Templates\\Review_1_PPT.pptx (16 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import BASE_PAPER, BATCH, COLLEGE, DEPT, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, REFERENCES, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_1_PPT.pptx"
OUT = BATCH / "7.Review_1_PPT.pptx"


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 1", TITLE, "Batch A1", "health insurance, RAG, LLM, question answering"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A1"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), TITLE)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A1 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 1 covers problem identification and research for PolicyLens, an AI health-insurance assistant that combines retrieval-augmented generation with a large language model.")

    # 2 contents
    set_notes(S[1], "Nine parts: the abstract, the base paper and related work, the gap we found, our problem statement and objectives, the approach, expected outcomes and references.")

    # 3 abstract
    s = S[2]
    set_paras(shape(s, 2), "Abstract: policies explained from the policy itself")
    fill(s, 8, "Health insurance policies are long, technical documents, so users struggle to understand coverage, exclusions, deductibles and claim procedures. Agents give slow, generic answers, and general chatbots can state things the policy does not say.", 16)
    fill(s, 13, "PolicyLens combines retrieval-augmented generation (RAG) with an LLM. The uploaded policy is parsed into numbered clauses; a Policy Card lists the key terms; a chatbot answers in plain language, citing the clauses it used.", 16)
    fill(s, 18, "It also guides claims with eligibility pre-checks, an evidence checklist and step-by-step procedures, flags risky terms and compares two plans, in English, Hindi and Telugu.", 16)
    set_notes(s, "The abstract in one slide: the problem, the RAG approach, and the claim-guidance and comparison features.")

    # 4 base paper
    s = S[3]
    fill(s, 5, BASE_PAPER, 14)
    set_paras(shape(s, 7), "Method")
    fill(s, 8, "Hybrid pipeline for long Turkish court rulings: TF-IDF and TextRank first select the important sentences; LED, Long-T5, BART-large and GPT-3.5 Turbo then generate an abstractive summary from them.", 15)
    set_paras(shape(s, 10), "Data and reported results")
    fill(s, 11, "2,000 Turkish consumer-rights rulings. TF-IDF with keywords: content F1 0.47. Best model GPT-3.5 Turbo: ROUGE-1 55%, ROUGE-2 35%, ROUGE-L 42%, ROUGE-Sum 44%.", 15)
    set_paras(shape(s, 13), "What we adopt and what we change")
    fill(s, 14, "We adopt select-then-generate for long legal text. We change the task (cited question answering, claims, comparison), the domain (health-insurance wordings) and the evaluation (retrieval, citation, faithfulness).", 15)
    set_notes(s, "The base paper summarises Turkish court rulings and reports ROUGE. Its select-then-generate idea is the starting point of PolicyLens, which applies it to cited question answering over insurance wordings.")

    # 5 literature survey
    s = S[4]
    rows = [["Work", "Approach", "Data", "Reported result"],
            ["Albayati and Findik, 2025 (base) [1]", "TF-IDF/TextRank extraction, then LED, Long-T5, BART, GPT-3.5 summaries", "2,000 Turkish court rulings", "GPT-3.5 Turbo ROUGE-1 55%, ROUGE-2 35%, ROUGE-L 42%"],
            ["Beauchemin et al., 2024 [2]", "RAG with GPT-4o over a Quebec auto-insurance corpus", "Corpus + 82 expert answers", "Better on average; 5-13% of answers had a false statement"],
            ["Maharjan and Yadav, 2026 [6]", "Plain LLM vs basic RAG vs RAG with cross-encoder re-ranking", "Public-health policy documents", "Faithfulness 0.347 / 0.621 / 0.797"],
            ["Multi-strategy RAG for medical-insurance QA, 2026 [3]", "Multi-strategy RAG", "Not retrieved", "Not retrieved (listed in batch survey)"],
            ["AI chatbot, Iran Health Insurance Org., 2023 [4]", "Service chatbot", "Not retrieved", "Not retrieved (listed in batch survey)"],
            ["Knollmeyer et al., 2025 [5]", "Knowledge-graph-enhanced RAG for document QA", "Manufacturing documents", "Not read"]]
    sz = table_size(rows, 1164, 460, [0.28, 0.34, 0.17, 0.21], 13, 9)
    set_table(shape(s, 6), rows, sz, header_size=sz)
    set_notes(s, "Six works. [1], [2] and [6] were read in detail; [3] and [4] are listed in the batch survey but could not be retrieved, and the results of [5] were not read, so no results are claimed for them.")

    # 6 supporting evidence
    s = S[5]
    set_paras(shape(s, 2), "Supporting evidence: grounding raises faithfulness")
    ch = replace_chart(s, 4, ["Plain LLM", "Basic RAG", "RAG + cross-encoder"], [("Faithfulness [6]", [0.347, 0.621, 0.797])], number_format="0.000")
    style_chart(ch, "Faithfulness of policy-document answers (Maharjan and Yadav [6])", ["6D5BD0"], number_format="0.000", legend=False, label_pt=12, axis_title="faithfulness score")
    set_paras(shape(s, 8), "Takeaway")
    fill(s, 9, "On public-health policy documents, retrieval raised faithfulness from 0.347 to 0.621 and cross-encoder re-ranking to 0.797 [6]. Even so, 5-13% of retrieval-augmented answers about Quebec car insurance contained a false statement [2].", 15)
    set_paras(shape(s, 13), "Caveat")
    fill(s, 14, "Published figures from other corpora, models and metrics. They motivate re-ranking and an independent support check; they are not our results.", 15)
    set_notes(s, "These are published numbers, not ours. They show that retrieval and re-ranking help and that wrong statements remain, which is why PolicyLens adds citations and a faithfulness score.")

    # 7 existing system
    s = S[6]
    set_paras(shape(s, 2), "Existing system and current approaches")
    steps = ["Receive policy PDF", "Read it alone", "Ask agent / helpline", "Search comparison sites", "Ask a general chatbot", "Find out at claim time"]
    for sid, t in zip((5, 8, 11, 14, 17, 20), steps):
        fill(s, sid, t, 14)
    set_paras(shape(s, 24), "Human help")
    fill(s, 25, "Agents, call centres and FAQ pages answer in general terms. They are slow, inconsistent and depend on human availability.", 16)
    set_paras(shape(s, 29), "Tools without the policy")
    fill(s, 30, "Comparison sites compare premiums and headline benefits, not exclusions. General chatbots answer from training data and can hallucinate coverage.", 16)
    set_notes(s, "How a policyholder understands a policy today.")

    # 8 limitations
    s = S[7]
    set_paras(shape(s, 2), "Limitations of existing methods")
    for tid, bid, ttl, body in [
        (7, 8, "Technical and long", "Key exclusions, waiting periods and conditions are easily missed in a 20-60 page wording."),
        (12, 13, "Slow and inconsistent", "Answers from agents and helplines depend on who answers and when."),
        (17, 18, "Not grounded", "Generic chatbots can state coverage the policy does not contain, with no clause to check."),
        (22, 23, "No claim guidance", "No tool extracts deductibles and limits or pre-checks eligibility and documents before a claim.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Four limitations: the wording is hard to read, human help is slow, chatbots are not grounded, and claim filing is unguided.")

    # 9 research gap
    s = S[8]
    gaps = [("Structure", "Flat text chunks lose clause numbers and pages; segmentation is named a bottleneck in [6]."),
            ("Support", "Retrieval alone leaves false statements [2]; no reviewed system scores each answer's support."),
            ("Declining", "A policy may not cover a question; the assistant must say so instead of inventing an answer."),
            ("Guidance", "Question answering is not joined with claim pre-checks, checklists and estimates."),
            ("Evidence", "No repeatable evaluation on several real health-insurance wordings was found.")]
    for (lid, tid), (lab, txt) in zip([(5, 6), (8, 9), (11, 12), (14, 15), (17, 18)], gaps):
        set_paras(shape(s, lid), lab)
        fill(s, tid, txt, 17)
    set_notes(s, "Existing work summarises or answers questions over legal and general documents, but few systems combine attribute extraction, clause-grounded answers and claim guidance for health insurance.")

    # 10 problem statement
    s = S[9]
    fill(s, 5, "Given a health-insurance policy wording, extract its key attributes, answer questions in plain language using only that wording and citing the clauses used, say so when the wording does not cover a question, guide claims with eligibility pre-checks and a document checklist, and compare plans, in a secure application whose accuracy and faithfulness are measured.", 22)
    set_paras(shape(s, 8), "Policy PDF, questions, claim details")
    set_paras(shape(s, 11), "Policy Card, cited answers, claim plan, comparison")
    set_paras(shape(s, 14), "Faithfulness score + measured evaluation")
    set_notes(s, "The problem statement: inputs are the policy and the user's questions; outputs are cards, cited answers, claim plans and comparisons; extra property is a measured faithfulness.")

    # 11 objectives
    s = S[10]
    fill(s, 4, ["Extract key policy attributes automatically from uploaded health insurance documents.",
                "Build a RAG pipeline that retrieves the policy clauses relevant to each user question.",
                "Generate plain-language answers with an LLM, grounded in and citing the retrieved clauses.",
                "Provide guided claim support with eligibility pre-checks, evidence checklists and procedural steps.",
                "Summarise risks, exclusions and cost-coverage trade-offs to support plan comparison.",
                "Evaluate the assistant for answer accuracy, faithfulness and response time."], 20, bullet=True)
    set_notes(s, "The six objectives of the project.")

    # 12 proposed approach
    s = S[11]
    for tid, bid, ttl, body in [
        (7, 8, "Parse and segment", "PDF to numbered clauses with pages and highlight boxes."),
        (12, 13, "Hybrid retrieval", "BM25 + MiniLM embeddings, rank fusion, cross-encoder re-rank."),
        (17, 18, "Grounded generation", "Gemini answers from the retrieved clauses only, with citations."),
        (22, 23, "Check and guide", "NLI faithfulness score; calculated claim pre-checks and estimate.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 14)
    set_paras(shape(s, 27), "Web application")
    shape(s, 28).height = Emu(44 * 9525)
    fill(s, 28, "React + FastAPI + SQLite: Policy Card, risks, original-page viewer, chat, Claim Copilot, plan comparison, dashboard; English, Hindi, Telugu.", 16)
    set_notes(s, "The approach: parse and segment, retrieve, generate from the clauses only, then check and guide.")

    # 13 architecture
    s = S[12]
    put_picture(s, DIAG / "01_system_architecture.png", (215, 118, 790, 482), remove_ids=(4,), name="System architecture diagram")
    set_paras(shape(s, 5), "Proposed architecture: React web app, FastAPI backend with retrieval and answering services, local models, SQLite and ChromaDB, Gemini API")
    set_notes(s, "Browser to React to FastAPI. Retrieval, faithfulness and the rules run locally; only the language generation calls the Gemini API.")

    # 14 expected outcomes
    s = S[13]
    for tid, bid, ttl, body in [
        (7, 8, "Right clause found", "The clause that answers a question is among the top results (target)."),
        (12, 13, "Cited, grounded answers", "Every answer names clause and page; unanswerable questions are declined."),
        (17, 18, "Measured support", "A faithfulness score on each answer."),
        (22, 23, "Claims guided", "Verdict, checklist, steps and estimate before filing."),
        (27, 28, "Plans compared", "Two policies side by side with trade-offs."),
        (32, 33, "Honest evaluation", "Accuracy, faithfulness and speed measured; gaps reported plainly.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Expected outcomes are stated as targets at this stage and are measured in the later reviews.")

    # 15 references
    s = S[14]
    fill(s, 4, [f"[{i}] {r}" for i, r in enumerate(REFERENCES[:12], 1)], 14, min_pt=8.5, bullet=True, gap_pt=3)
    set_notes(s, "References [1] to [12]; the full list is in the report.")

    set_notes(S[15], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
