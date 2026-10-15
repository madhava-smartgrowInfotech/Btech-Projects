"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import BATCH, COLLEGE, DEPT, DIAG, EXP, GUIDE, GUIDE_DESIG, PRODUCT, SHOTS, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"


def px(v):
    return Emu(int(v * 9525))


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, "Batch A1", "health insurance, RAG, LLM, question answering"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A1"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), TITLE)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A1 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers design, methodology and progress of PolicyLens.")

    # 2 contents
    s = S[1]
    for sid, t in [(21, "Model and configuration"), (18, "Dataset and preparation"), (27, "Data flow, modules, database"), (36, "Preliminary results and screens")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts, from a recap of Review 1 through the design and the pipeline to progress, results, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Policies are long and technical; agents are slow; general chatbots are not grounded in the user's own policy; claim filing has no guided pre-check.", 16)
    fill(s, 13, "A RAG pipeline: clause-aware parsing, hybrid retrieval, Gemini answers from the retrieved clauses only, an independent faithfulness score, and calculated claim rules.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Upload policy PDF", "Parse + segment + index", "Policy Card + risks", "Retrieve + cited answer", "Claim plan, compare"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1: the problem and the proposed system as a five-step pipeline.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 18 + TypeScript + Vite (port 5101), TanStack Query, Recharts.",
                "Application: FastAPI (port 8101): JWT sign-in, owner-only access, 42 route handlers.",
                "Retrieval: BM25 + MiniLM/ChromaDB, rank fusion, cross-encoder re-rank.",
                "AI: Gemini (JSON schema, fallback chain) and a local NLI faithfulness scorer.",
                "Data: SQLite with 12 tables, vector index, PDFs and page images."], 16, bullet=True)
    set_notes(s, "The browser talks only to the Vite server, which proxies /api to FastAPI. Retrieval, faithfulness and the rules run locally; only language generation uses the Gemini API.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Upload + validate", "Parse + segment", "Index (BM25, vectors)", "Policy Card + risks", "Retrieve + re-rank", "Generate + cite", "Score + guide"]):
        fill(s, sid, t, 14)
    set_paras(shape(s, 25), "Clauses, not chunks")
    fill(s, 26, "Numbered clauses with pages and highlight boxes make every citation checkable on the original page.", 15)
    set_paras(shape(s, 28), "Grounded generation")
    fill(s, 29, "Gemini sees only the retrieved, tagged clauses; tags it was not given are removed.", 15)
    set_paras(shape(s, 31), "Independent checks")
    fill(s, 32, "An NLI model scores support; rules calculate waiting periods, co-pay and the estimate.", 15)
    set_notes(s, "End-to-end workflow in seven steps. The central idea: the language model only sees retrieved clauses and everything it returns is checked.")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    ch = replace_chart(s, 4, ["Star Health", "HDFC ERGO", "Niva Bupa", "Care Health"], [("Clauses", [219, 207, 166, 209])])
    style_chart(ch, "Clauses per sample policy (801 in total)", ["3F3FBF"], legend=False, label_pt=11, axis_title="clauses")
    rows = [["Dataset", "Size", "Use"],
            ["Sample policy wordings (public, insurers' sites)", "4 policies, 166 pages, 73,001 words", "Product samples, evaluation source"],
            ["Questions (hand-written)", "57: 51 answerable, 6 out-of-scope", "Retrieval and answer evaluation"],
            ["Hindi / Telugu questions", "5", "Multilingual check"],
            ["Claim scenarios", "14 (5 covered, 1 partly, 7 not, 1 needs info)", "Claim Copilot"],
            ["Policy Card checks", "52 hand-checked values", "Extraction accuracy"],
            ["Training data", "None: nothing is trained; whole set is test data", "-"]]
    sz = table_size(rows, 530, 400, [0.4, 0.38, 0.22], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.4, 0.38, 0.22], header_size=sz)
    set_notes(s, "The wordings are public documents that insurers publish; they are stored unmodified. The evaluation set was written by hand from them. There is no train/test split because no model is trained.")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation: parse, segment, index")
    set_paras(shape(s, 7), "Parsing")
    fill(s, 8, "PyMuPDF lines with page, position and font; running headers and footers removed; two-column reading order restored; letter-spaced headings joined; tables read by row; scans rejected.", 16)
    set_paras(shape(s, 12), "Segmentation")
    fill(s, 13, "Clause starts at 4.2.1, Def. 5, Excl03 and headings; exclusion-coded list items become clauses; heading path; 12-380 word chunks with highlight boxes. 801 clauses from four policies.", 16)
    set_paras(shape(s, 17), "Indexing")
    fill(s, 18, "BM25 Okapi (k1 1.4, b 0.7) with an insurance synonym list, and 384-d MiniLM embeddings in ChromaDB (cosine). One Gemini call then builds the Policy Card.", 16)
    set_notes(s, "Preprocessing here is turning a printed PDF into clean, numbered, indexed clauses.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Pipeline architecture (no model is trained)")
    put_picture(s, DIAG / "02_high_level_workflow.png", (75, 124, 693, 518), remove_ids=(4,), name="High-level workflow")
    rows = [["Component", "Detail"], ["Dense retrieval", "all-MiniLM-L6-v2, 384-d"], ["Re-ranker", "ms-marco-MiniLM-L6-v2"], ["Faithfulness", "nli-deberta-v3-xsmall"], ["Generator", "Gemini, JSON schema"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "All pre-trained; three small models run on the CPU, only generation uses the Gemini API.", 11)
    set_notes(s, "Every model is pre-trained and used as is. The pipeline is the design contribution.")

    # 9 configuration
    s = S[8]
    set_paras(shape(s, 2), "Pipeline configuration (nothing is trained)")
    rows = [["Setting", "Value"], ["Candidates per retriever", "30"], ["Fusion", "RRF, k = 60"], ["Re-rank pool", "fused top 20 + top 8 each"],
            ["Clauses to the LLM", "top 7 (+ up to 4 context)"], ["Abstain if", "re-rank < 0.004 and BM25 < 9.0"], ["Temperature", "0.2 answers, 0.1 card"],
            ["Gemini chain", "3.8-flash, 3.6-flash, 3.5-flash-lite"]]
    set_table(shape(s, 10), rows, 13, header_size=13)
    put_picture(s, EXP / "retrieval_ablation.png", (672, 232, 550, 330), remove_ids=(5,), name="Retrieval ablation chart")
    set_paras(shape(s, 7), "Retrieval ablation: hybrid + re-rank is best on hit@1, hit@3 and MRR", 11)
    set_notes(s, "Settings were chosen by measurement on the evaluation set: for example, adding each retriever's top 8 to the re-rank pool recovered questions where one retriever alone found the clause.")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (57, 124, 420, 520), remove_ids=(4,), name="Decision flowchart")
    for bid in (7, 10, 13):
        shape(s, bid).height = px(80)
    set_paras(shape(s, 6), "Abstain")
    fill(s, 7, "If the best re-rank score is below 0.004 and the best BM25 score below 9.0, answer \"not covered in this policy\" without calling the model.", 16)
    set_paras(shape(s, 9), "Cite")
    fill(s, 10, "Gemini cites clause tags; tags that were not retrieved are removed; each claim is scored by an NLI model.", 16)
    set_paras(shape(s, 12), "Claims")
    fill(s, 13, "Waiting periods, age/network co-payment, sub-limits and the estimate are calculated; a failed waiting period is never \"covered\".", 16)
    set_notes(s, "Decision logic: when to answer, how to cite and score, and which parts of a claim check are rules rather than model output.")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 330), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 330), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Answer a question' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(452)
    rows = [["Module", "Responsibility"], ["Ingestion", "validate, parse, segment, index, build card and risks in a background worker"],
            ["Retrieval + answering", "BM25 + dense, fusion, re-rank, abstention rule, cited answer, faithfulness"],
            ["Claim Copilot + compare", "calculated pre-checks and estimate, verdict, checklist; two-policy comparison"]]
    set_table(shape(s, 12), rows, 12, widths=[0.2, 0.8], header_size=12, row_h=34)
    shape(s, 12).top = px(485)
    set_notes(s, "Level 1 shows eight processes and nine data stores; level 2 opens the question-answering process.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 560, 340), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (632, 118, 589, 340), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 12 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(460), px(1164)
    rows = [["Table", "Key fields", "Relations"],
            ["documents", "sha256, page_count, insurer, uin, status, parse_stats", "1:N clauses, risk_flags; 1:1 policy_cards"],
            ["clauses", "ordinal, clause_ref, section_path, text, pages, bboxes", "N:1 documents"],
            ["messages", "role, content, citations, faithfulness, timings, model", "N:1 conversations"]]
    set_table(shape(s, 13), rows, 11, widths=[0.2, 0.44, 0.36], header_size=11, row_h=30, total_px=1164)
    t = shape(s, 13)
    t.left, t.top, t.width = px(57), px(492), px(1164)
    set_notes(s, "Twelve tables created by SQLAlchemy. Documents are shared by content hash. Structured results are stored as JSON columns.")

    # 13 technology
    s = S[12]
    cards = [(7, 8, "Language", "Python 3.11, TypeScript"), (12, 13, "Retrieval", "rank-bm25, ChromaDB, sentence-transformers, PyMuPDF"),
             (17, 18, "Frontend", "React 18, Vite, Tailwind CSS, shadcn/ui, TanStack Query, Recharts"), (22, 23, "Backend", "FastAPI, Pydantic v2, SQLAlchemy 2, PyJWT, bcrypt"),
             (27, 28, "AI", "Google Gemini (google-genai), three local CPU models, torch"), (32, 33, "Testing", "pytest (25 tests), strict TypeScript build, Playwright captures")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Everything runs on one laptop; the Gemini API key is the only credential.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Upload, parsing, clause segmentation, viewer", "Completed"], ["Policy Card with verified sources; risk highlights", "Completed"],
            ["Hybrid retrieval, cited answers, faithfulness score", "Completed"], ["Claim Copilot with calculated pre-checks", "Completed"],
            ["Plan comparison; English, Hindi, Telugu", "Completed"], ["Dashboard and Model performance page", "Completed"],
            ["Evaluation harness, 3 runs, 25 backend tests", "Completed"], ["User acceptance testing with real users", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.78, 0.22], header_size=13)
    for vid, lid, v, lab in [(6, 7, "42", "API route handlers"), (9, 10, "12", "database tables"), (12, 13, "25/25", "backend tests passing")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Every planned feature is implemented. The one open item is user-acceptance testing with real policyholders.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: retrieval and answers")
    ch = replace_chart(s, 4, ["BM25", "Dense", "Hybrid (RRF)", "Hybrid + re-rank"], [("hit@1", [0.647, 0.490, 0.686, 0.804]), ("hit@5", [1.000, 0.804, 0.961, 1.000])], number_format="0.000")
    style_chart(ch, "Relevant clause ranked first (hit@1) and in the top 5 (hit@5), 51 questions", ["3F3FBF", "2F9E6D"], number_format="0.000", legend=True, label_pt=9, axis_title="share of questions")
    ch.value_axis.maximum_scale = 1.05
    for vid, lid, v, lab in [(6, 7, "51 / 51", "answers correct and correctly cited"), (9, 10, "6 / 6", "out-of-scope questions declined, 0 false"), (12, 13, "87.0", "mean faithfulness; median answer 4.3 s")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Evaluation run eval-20260922-0027 on 51 answerable questions, 6 out-of-scope, 5 multilingual, 14 claim scenarios and 52 card values; all with the smallest Gemini model because of free-tier quota.")

    # 16 output screens
    s = S[15]
    put_picture(s, SHOTS / "11_chat_answer.png", (91, 128, 507, 476), remove_ids=(4,), name="Chat answer screenshot")
    put_picture(s, SHOTS / "16_claim_result_partly.png", (679, 128, 512, 476), remove_ids=(7,), name="Claim check screenshot")
    set_paras(shape(s, 6), "Cited answer: 24-month waiting period, two sources, faithfulness 88", 11)
    set_paras(shape(s, 9), "Claim Copilot: knee replacement is partly covered (20% co-payment)", 11)
    set_notes(s, "Real captures of the running application with the four sample policies and live Gemini answers.")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Authentication and accounts", "6", "Passed"], ["Parser and clause segmentation", "3", "Passed"], ["Faithfulness scoring", "3", "Passed"],
            ["Rules, fusion, risks, estimate", "6", "Passed"], ["Policy flow, search, chat, claims", "7", "Passed"], ["Total", "25", "25 passed"]]
    set_table(shape(s, 18), rows, 13, widths=[0.6, 0.15, 0.25], header_size=13)
    fill(s, 5, ["25 backend tests pass in about 15 s on a synthetic policy; Gemini is never called in tests.",
                "Web build: strict TypeScript and Vite production build succeed.",
                "Evaluation: relevant clause in the top 5 for 51/51; 51/51 answers correct and cited; 14/14 claim verdicts; 51/52 card values.",
                "Edge cases: scanned PDFs rejected; uncited tags removed; out-of-scope questions declined.",
                "Not done: user-acceptance, load and security testing."], 16, bullet=True)
    set_notes(s, "25 backend tests plus the measured evaluation and a scripted end-to-end demo run. The frontend is checked by a strict TypeScript compile and a production build.")

    # 18 challenges
    s = S[17]
    cards = [(7, 8, "Wordings are messy", "Two-column layouts, running headers, tables: a rule-based parser and segmenter restore order and clause numbers."),
             (12, 13, "Finding the clause", "BM25 and embeddings each miss cases; fusion plus a cross-encoder gives hit@5 of 100%."),
             (17, 18, "Unsupported answers", "Citations are filtered to retrieved clauses and an NLI score is shown; it is a heuristic, not a verdict."),
             (22, 23, "Knowing when to decline", "Neither score alone separates the cases; the pair (re-rank < 0.004, BM25 < 9.0) does."),
             (27, 28, "Quota and overload", "A fallback chain of three Gemini models, a cool-down and a cache keep answers coming."),
             (32, 33, "Translation quality", "One Telugu card line mixes scripts; recorded as a limitation.")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Six challenges and how each was handled.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["Evaluate on unseen policies", "User acceptance testing", "Final report", "Final presentation and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Write a fresh question set on policy wordings not used in development and re-measure (the current figures are partly tuned).",
                 "Collect feedback from policyholders on answers, citations and claim checks (not yet done).",
                 "Re-run the evaluation with the larger Gemini models; add plan-variant selection.",
                 "Document every diagram, test and result in the report."], 18, bullet=True)
    set_notes(s, "What remains: unseen-policy evaluation and real-user feedback, then the report and the final demo.")

    set_notes(S[19], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
