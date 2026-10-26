from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BATCH = Path(__file__).resolve().parents[2]
CODE = BATCH / "8.ProjectCode"
DIAG = CODE / "docs" / "diagrams"
CHART = CODE / "docs" / "charts"
SHOT = CODE / "docs" / "screenshots"

TITLE = "AI Powered Intelligent Hospital Patient Flow Management System"
PRODUCT = "CareFlow (implemented prototype: MedFlow)"
STUDENTS = [
    ("A. RAMU", "23K91A6715"),
    ("B. RUSHI", "23K91A6720"),
    ("B. RAJ KUMAR", "23K91A6725"),
    ("A. PRABHAS", "23K91A6703"),
]
GUIDE = "[GUIDE NAME]"

REFERENCES = [
    "S. A. N. Saritha, G. Kalnoor, P. R. Bhat, P. A. Rao, P. Prajwal, and P. T. Pradeep, ‘Alleviate: An AI-Powered Mobile Application for Personalized Healthcare Management and Doctor Recommendation,’ IEEE Access, vol. 13, pp. 215822–215832, 2025, doi: 10.1109/ACCESS.2025.3647089.",
    "S.-K. Tan, S.-C. Chong, K.-K. Wee, and L.-Y. Chong, ‘Personalized healthcare: A comprehensive approach for symptom diagnosis and hospital recommendations using AI and location services,’ Journal of Informatics and Web Engineering, vol. 3, no. 1, pp. 117–135, 2024.",
    "M. Abinaya and R. Ganesan, ‘Effective search mechanism for finding nearest healthcare facilities,’ in Proc. Global Conference on Communication Technologies, 2015, pp. 534–538.",
    "M. W. Munir, S. M. Omair, and M. Z. U. Haque, ‘An Android based application for determine a specialized hospital nearest to patient’s location,’ International Journal of Computer Applications, vol. 118, no. 9, pp. 43–46, 2015.",
    "R. A. Nimbalkar and R. A. Fadnavis, ‘Domain specific search of nearest hospital and healthcare management system,’ in Proc. Recent Advances in Engineering and Computational Sciences, 2014, pp. 1–5.",
    "A. Kumar, ‘Disease prediction and doctor recommendation system using machine learning approaches,’ International Journal for Research in Applied Science and Engineering Technology, vol. 9, no. 7, pp. 34–44, 2021.",
    "L. Bjorkdahl, O. Pauli, J. Ostman, C. Ceccobello, S. Lundell, and M. Kjellberg, ‘Towards holistic disease risk prediction using small language models,’ arXiv:2408.06943, 2024.",
    "R. Umbare, R. Patil, T. Mukund, A. Rathod, and P. Rajurkar, ‘A smart healthcare application using artificial intelligence and machine learning,’ in Proc. ICIPCN, 2024, pp. 204–210.",
    "D. Gujar, R. Biyani, T. Bramhane, S. Bhosale, and T. P. Vaidya, ‘Disease prediction and doctor recommendation system,’ International Research Journal of Engineering and Technology, vol. 5, pp. 3207–3209, 2018.",
    "D. Mogaveera, V. Mathur, and S. Waghela, ‘E-health monitoring system with diet and fitness recommendation using machine learning,’ in Proc. ICICT, 2021, pp. 694–700.",
    "P. Tanwar, K. Bansal, A. Sharma, and N. Dagar, ‘AI based chatbot for healthcare using machine learning,’ in Proc. ICAICCIT, 2023, pp. 751–755.",
    "R. Surve, T. Purohit, R. Joseph, and P. Shaikh, ‘Healthcare chatbot using machine learning and NLP,’ in Proc. ICAST, 2023, pp. 411–416.",
    "H. Mendapara, S. Digole, M. Thakur, and A. Dange, ‘AI based healthcare chatbot system using natural language processing,’ International Journal of Scientific Research and Engineering Development, vol. 4, no. 2, pp. 89–96, 2021.",
    "M. M. Rahman et al., ‘Web-based multiple disease prediction system with transfer learning and healthcare chatbot,’ in Proc. IEEE ICMI, 2024, pp. 1–5.",
    "J. Tseng et al., ‘An interactive healthcare system with personalized diet and exercise guideline recommendation,’ in Proc. TAAI, 2015, pp. 525–532.",
    "T. Bhanuteja, K. V. N. Kumar, K. S. Poornachand, C. C. V. N. Ashish, and P. Anudeep, ‘Symptoms based multiple disease prediction model using machine learning approach,’ IJITEE, vol. 10, no. 9, pp. 67–72, 2021.",
    "‘Waiting Time Prediction in Queue Management: Leveraging Machine Learning Approach,’ in Proc. International Conference on Recent Progresses in Science, Engineering and Technology, 2024.",
    "‘Prediction of outpatient waiting time using machine learning in a tertiary children’s hospital,’ Translational Pediatrics, 2023.",
    "‘Dissatisfaction-considered waiting time prediction for outpatients with interpretable machine learning,’ Health Care Management Science, 2024.",
]

METRIC_ROWS = [
    ["Wait-time regressor", "12,000 synthetic rows", "R-squared", "0.951041"],
    ["Wait-time regressor", "1,800-row test", "MAE", "5.982 min"],
    ["Wait-time regressor", "1,800-row test", "RMSE", "8.580 min"],
    ["Triage classifier", "2,250-row test", "Accuracy", "0.9680"],
    ["Triage classifier", "2,250-row test", "Macro F1", "0.898475"],
    ["Triage classifier", "100 critical cases", "Critical recall", "0.63"],
]

COMMON_LIMITS = [
    "Both model datasets are synthetic and generated from designed distributions; they are not real patient records.",
    "The measured model test partitions come from the same deterministic generators as training and are not independent clinical validation.",
    "No real hospital deployment, clinician study, user-acceptance study, security audit, accessibility audit, or concurrent-load test was performed.",
    "The 96.80% triage accuracy must be interpreted with the 63% critical-class recall; accuracy alone overstates safety-critical performance.",
    "SQLite, a development JWT secret fallback, an unauthenticated WebSocket, and in-process broadcasting are unsuitable for production without hardening.",
    "The landing-page claims of 42% wait reduction, 214 patients, and six departments are hard-coded marketing copy and are excluded from results.",
]


def set_cell_shading(cell, fill="DCE6F1"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def configure(doc: Document, letter=False):
    section = doc.sections[0]
    if letter:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.75)

    settings = doc.settings._element
    zoom = settings.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        settings.insert(0, zoom)
    zoom.set(qn("w:val"), "bestFit")
    zoom.set(qn("w:percent"), "100")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 18), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        s = styles[name]; s.font.name = "Times New Roman"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(20, 55, 90)
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Courier New"; code.font.size = Pt(8)
        code.paragraph_format.space_after = Pt(3)
        code.paragraph_format.line_spacing = 1.0
    header = section.header.paragraphs[0]
    header.text = TITLE
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = "Times New Roman"; header.runs[0].font.size = Pt(9); header.runs[0].font.italic = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    add_field(footer, "PAGE")
    return doc


def add_title(doc: Document, document_name: str, subtitle=""):
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(20); r.font.color.rgb = RGBColor(16, 70, 105)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(document_name); r.bold = True; r.font.size = Pt(17)
    if subtitle:
        p = doc.add_paragraph(subtitle); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
    table.rows[0].cells[0].text = "Student"
    table.rows[0].cells[1].text = "Hall Ticket No."
    for c in table.rows[0].cells: set_cell_shading(c, "1F4E78");
    for c in table.rows[0].cells:
        for run in c.paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
    for name, roll in STUDENTS:
        row = table.add_row().cells; row[0].text = name; row[1].text = roll
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Under the guidance of {GUIDE}").bold = True
    p = doc.add_paragraph("Department of Computer Science and Engineering (Data Science)\n[COLLEGE NAME]\n[ACADEMIC YEAR]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_para(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc: Document, items: Iterable[str], numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None, caption: str | None = None):
    if caption:
        p = doc.add_paragraph(caption); p.style = doc.styles["Caption"]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        cell=table.rows[0].cells[i]; cell.text=str(h); set_cell_shading(cell,"1F4E78")
        for run in cell.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(9)
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(2)
                for run in p.runs: run.font.size=Pt(9)
    return table


def add_figure(doc: Document, path: Path, caption: str, width=6.35):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c=doc.add_paragraph(caption); c.style=doc.styles["Caption"]; c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return c


def section_page(doc: Document, heading: str, paras: list[str], figure: Path | None=None, caption="", table=None, new_page=True):
    if new_page: doc.add_page_break()
    doc.add_heading(heading, level=2)
    if figure: add_figure(doc, figure, caption, width=6.15)
    if table: add_table(doc, *table)
    for text in paras: add_para(doc,text)


def save_doc(doc: Document, path: Path):
    props=doc.core_properties; props.title=TITLE; props.subject=path.stem; props.author="A. RAMU; B. RUSHI; B. RAJ KUMAR; A. PRABHAS"
    doc.save(path)
    print(f"Wrote {path.name} ({path.stat().st_size} bytes)")


def methodology_doc():
    doc=configure(Document()); add_title(doc,"METHODOLOGY, NOVELTY AND KEY CONTRIBUTIONS")
    doc.add_heading("1. Methodology Overview",1)
    add_para(doc,"CareFlow follows an evidence-driven patient-flow pipeline that begins with validated registration data and ends with a persisted operational state visible to each authorised role. The implementation combines deterministic workflow rules with two trained estimators: a triage classifier and a wait-time regressor. Predictions assist queue ordering and communication; they do not replace clinical judgement.")
    add_figure(doc,DIAG/"02_high_level_workflow.png","Figure 1. End-to-end patient-flow workflow.")
    section_page(doc,"2. Data Preparation",["The project does not use real patient records. The wait-time generator creates 12,000 rows from department mix, queue size, hour, priority, average service time and server count. The triage generator creates 15,000 vital-sign rows and assigns labels with a simplified NEWS2-style rule. Each dataset uses an 85/15 train/test split with random state 42.","Categorical department and priority values are label-encoded. Numerical features are kept on their original operational scales because tree ensembles do not require standardisation. Missing incoming vital values are replaced individually with documented defaults before inference."],CHART/"06_dataset_split.png","Figure 2. Synthetic dataset partitions.")
    section_page(doc,"3. Triage Classification",["The Gradient Boosting classifier consumes age, heart rate, systolic blood pressure, oxygen saturation, temperature and respiratory rate. It returns one of three priority classes—normal, urgent or critical—and the maximum class probability, which the current schema stores in the field named triage_score.","On the deterministic 2,250-row test partition the model achieved 0.9680 accuracy and 0.898475 macro F1. The critical-class recall was 0.63, so this model is an academic decision-support prototype and cannot be described as clinically validated."],CHART/"01_triage_confusion_matrix.png","Figure 3. Triage confusion matrix.")
    section_page(doc,"4. Wait-Time Regression",["The Random Forest regressor estimates minutes from department type, patients ranked ahead, UTC hour, priority, average service duration and a default assumption of two servers. Training targets follow a queue-inspired formula with rush-hour and priority factors plus Gaussian noise, clipped to 0–240 minutes.","The measured test results were R-squared 0.951041, MAE 5.982 minutes and RMSE 8.580 minutes. These values describe recovery of a synthetic generator, not proof of real-world waiting-time reduction."],CHART/"03_wait_error.png","Figure 4. Wait-time error measures.")
    section_page(doc,"5. Operational Integration",["During POST /api/visits, the service validates the department, finds or creates a patient, completes missing vitals, invokes both models, allocates the next department token, inserts a Visit and MovementLog, commits the transaction and broadcasts visit_created. Subsequent status or department changes create movement evidence and may recompute predicted wait.","Role restrictions are enforced with JWT decoding and require_roles dependencies. Doctors manage consultations and orders, laboratory technicians progress tests, nurses manage beds, reception creates visits and administrators view hospital-wide state."],DIAG/"10_core_algorithm_flowchart.png","Figure 5. Core prediction and persistence algorithm.")
    section_page(doc,"6. Novelty",["The novel contribution is not a new clinical scoring rule. It is the integration of priority-aware token ordering, explainable operational state, laboratory workflow, ward/ICU capacity and role-specific live dashboards in one compact prototype. Unlike the base paper, which focuses on patient-side personalised healthcare and recommendations, this system addresses coordination after a patient enters the hospital.","A second contribution is honest evidence separation: synthetic model evaluation, application smoke measurements and demonstration database outputs are reported as different evidence classes. Unsupported landing-page marketing numbers are excluded."],DIAG/"01_system_architecture.png","Figure 6. Integrated application architecture.")
    section_page(doc,"7. Key Contributions",[],table=(["No.","Contribution","Evidence"],[["1","Priority-aware digital token queue","Queue service and visit routes"],["2","AI-assisted triage and wait estimate","Saved models and measured test metrics"],["3","Cross-department movement trail","MovementLog records"],["4","Laboratory and bed workflow","Lab and bed routes with role checks"],["5","Real-time operational views","WebSocket broadcasts and dashboards"],["6","Reproducible evidence record","REPORT_FACTS.md and screenshots"]],None,"Table 1. Implemented contributions."))
    section_page(doc,"8. Methodological Limitations",COMMON_LIMITS)
    save_doc(doc,BATCH/"9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx")


def scope_doc():
    doc=configure(Document()); add_title(doc,"SCOPE AND LIMITATIONS")
    doc.add_heading("1. Scope",1)
    add_para(doc,"CareFlow covers the operational journey of a patient from registration through department queues, consultation, laboratory work, admission and discharge. It provides staff-facing role dashboards and a public waiting-room board. The prototype is intended for academic demonstration of coordinated flow, not autonomous diagnosis or production clinical use.")
    add_figure(doc,DIAG/"03_dfd_level_0.png","Figure 1. System boundary and external actors.")
    section_page(doc,"2. Functional Scope",[],table=(["Area","Included capability"],[["Reception","Create patient/visit, token, triage and wait estimate"],["Doctor","View queue, assign consultation, order lab, admit, discharge"],["Laboratory","View and progress test work"],["Nursing","View ward/ICU beds, update or release resources"],["Administration","KPIs, department load, bed and ICU occupancy"],["Public display","Now-serving token, upcoming tokens and expected wait"]],None,"Table 1. In-scope functions."))
    section_page(doc,"3. Non-Functional Scope",["The implementation demonstrates authenticated role separation, structured validation, responsive browser interfaces, local persistence, live update delivery and reproducible model loading. Performance evidence is limited to sequential loopback requests on one Windows workstation.","Maintainability is supported by route, schema, service, model and ML modules. Portability is demonstrated only for the documented local Python and Node.js setup; container or cloud deployment was not implemented."])
    section_page(doc,"4. Assumptions",["The two-server default used by wait inference is an operational simplification. Department average service minutes are configuration data. Staff use correct roles and keep states current. Synthetic demonstration visits are assumed to be sufficient for UI verification. Missing vitals are completed with neutral defaults, which can hide uncertainty and is unsuitable for clinical use without explicit warnings."])
    section_page(doc,"5. Out of Scope",[],table=(["Excluded item","Reason / boundary"],[["Diagnosis and treatment recommendation","No validated clinical model"],["Real EHR/HIS integration","No external hospital interface or standard implemented"],["Billing, insurance and pharmacy inventory","Outside patient-flow objective"],["Production identity management","Demo accounts and local JWT only"],["Mobile application","Implemented as responsive web application"],["Cloud scaling and disaster recovery","Single local process and SQLite"],["Real SMS, e-mail or push alerts","No external notification provider"]],None,"Table 2. Deliberate exclusions."))
    section_page(doc,"6. Evidence-Based Limitations",COMMON_LIMITS,CHART/"02_model_metrics.png","Figure 2. Measured model scores; values are not clinical validation.")
    section_page(doc,"7. Implementation Constraints",["The nurse dashboard requires an explicit tab click after first load because the uncontrolled default is calculated before asynchronous department data arrives. The WebSocket is unauthenticated and local to one process. Seeded occupied beds are display records and are not linked to admitted visits. Laboratory progress and visit status are not fully synchronised in every transition.","The checked-in application uses ports 8010 and 5173, although the documentation brief mentions 8110 and 5110. The documents follow the implemented ports."])
    save_doc(doc,BATCH/"10.SCOPE & LIMITATIONS.docx")


def diagrams_doc():
    doc=configure(Document()); add_title(doc,"BLOCK DIAGRAMS - ALL")
    doc.add_heading("Diagram Inventory",1)
    add_para(doc,"The following diagrams were generated from the implemented SQLAlchemy entities, FastAPI routes, services, model inference functions and React role flows. They describe the checked-in code rather than a proposed future architecture.")
    names=[
        ("01_system_architecture.png","Figure 1. System architecture","The React client calls FastAPI REST routes, which coordinate JWT role checks, SQLAlchemy persistence and the two saved estimators. WebSocket events refresh live clients."),
        ("02_high_level_workflow.png","Figure 2. High-level patient workflow","The operational path runs from registration and prediction through priority queueing, consultation, routing and closure."),
        ("03_dfd_level_0.png","Figure 3. DFD Level 0","The context view shows reception, clinical staff, administration and the waiting board exchanging information with CareFlow."),
        ("04_dfd_level_1.png","Figure 4. DFD Level 1","The system is decomposed into authentication, visit, prediction, resource, movement and reporting processes."),
        ("05_dfd_level_2.png","Figure 5. DFD Level 2","Visit creation validates inputs, resolves a patient, predicts priority and wait, persists the visit and publishes an event."),
        ("06_uml_class_diagram.png","Figure 6. UML class diagram","The class view records the persistent model types and their operational relationships."),
        ("07_sequence_diagram.png","Figure 7. Registration sequence","Reception, frontend, API, ML layer, database and WebSocket participate in one registration transaction."),
        ("08_use_case_diagram.png","Figure 8. Use-case diagram","Each role is linked only to the functions implemented by its dashboard and API permissions."),
        ("09_er_diagram.png","Figure 9. Entity-relationship diagram","Foreign-key paths connect departments, staff, patients, visits, beds, tests and movement logs."),
        ("10_core_algorithm_flowchart.png","Figure 10. Core algorithm flowchart","The prediction flow combines vital-sign classification and queue-context regression before persistence."),
    ]
    for i,(file,cap,explain) in enumerate(names):
        section_page(doc,cap.split(". ",1)[1],[explain],DIAG/file,cap,new_page=True)
    save_doc(doc,BATCH/"11. BLOCK DIAGRAMS ALL.docx")


def design_doc():
    doc=configure(Document()); add_title(doc,"DESIGN AND IMPLEMENTATION")
    doc.add_heading("1. Design Goals",1)
    add_para(doc,"The design prioritises a single source of operational state, role separation, visible queue ordering, traceable movement and small replaceable modules. Predictions are computed at the edge of visit creation and stored with the visit so every dashboard presents the same result.")
    add_figure(doc,DIAG/"01_system_architecture.png","Figure 1. Implemented architecture.")
    section_page(doc,"2. Module Breakdown",[],table=(["Module","Responsibility"],[["api/auth.py","Login and current-user endpoint"],["api/visits.py","Registration, queue, assignment, state and movement"],["api/beds.py","Capacity, admission and release"],["api/labs.py","Laboratory orders and progress"],["api/dashboard.py","KPIs, load and public queue"],["services/queue_service.py","Token and priority ordering"],["services/ws_manager.py","In-process WebSocket broadcast"],["ml/predictor.py","Model loading and inference"]],None,"Table 1. Backend modules."))
    section_page(doc,"3. Database Design",["Seven SQLAlchemy entities persist staff, departments, patients, visits, beds, laboratory tests and movement logs. String identifiers are generated from UUID values. Enumerations constrain workflow states. Relationships support role context, patient history, department capacity and visit-linked work."],DIAG/"09_er_diagram.png","Figure 2. Database entity relationships.")
    section_page(doc,"4. Visit State Design",["VisitStatusEnum includes waiting, in_consultation, lab_pending, lab_in_progress, pharmacy, admitted_ward, admitted_icu, discharged and cancelled. Priority is critical, urgent or normal. Queue sorting uses priority rank and token number; a critical new visit has zero patients counted ahead.","MovementLog is created at registration and when a status update includes a new department. This makes location transitions inspectable, although admission currently changes the visit department without adding a movement record."])
    section_page(doc,"5. API Design",[],table=(["Group","Principal methods","Protection"],[["Authentication","POST login; GET me","Public login / bearer token"],["Visits","POST, GET, PATCH status, PATCH doctor","Role checks"],["Beds","GET, PATCH, POST admit/release","Authenticated; write roles"],["Labs","POST, GET, PATCH","Doctor/admin or lab/admin"],["Dashboard","GET KPI/load/queue","Authenticated except queue board"],["Live","WebSocket /ws/live","No authentication"]],None,"Table 2. API design summary."))
    section_page(doc,"6. ML Integration",["The predictor lazily loads Joblib bundles and automatically trains only when either artefact is missing. Department and priority encoders are stored beside the wait model. The triage bundle stores its feature order. Inference constructs pandas DataFrames with the same feature names used in training."],DIAG/"10_core_algorithm_flowchart.png","Figure 3. Prediction integration.")
    section_page(doc,"7. Security and Error Handling",["Passwords are hashed with bcrypt through Passlib. JWTs carry subject, role and expiry and default to a 12-hour lifetime. OAuth2PasswordBearer extracts tokens and inactive users are rejected. Route-specific require_roles checks return HTTP 403 for forbidden writes; missing records return 404 and unavailable beds return 400.","Production work must replace the fallback secret, authenticate WebSockets, restrict CORS, add audit logging, enforce TLS and define privacy retention controls."])
    section_page(doc,"8. Frontend Design",["React Router separates public landing/login/board routes from protected role homes. Axios injects the bearer token and clears local session state after HTTP 401. Zustand stores session and toast state. Role dashboards fetch REST data and refresh when matching WebSocket events arrive.","The visual system uses dark high-contrast panels, priority/status badges, reusable cards and modals. Real 1440 x 900 captures confirm the principal screens, but no formal WCAG audit was performed."])
    section_page(doc,"9. Key Code Walkthrough",["Visit creation is the central orchestration path: resolve department; find or create patient; merge supplied values with defaults; classify priority; allocate token; count queue; predict wait; insert Visit and MovementLog; commit; publish VisitOut. The transaction intentionally stores the returned prediction so the public board and authenticated dashboards agree.","The dashboard aggregates counts and averages directly in SQLAlchemy. Bed occupancy divides occupied beds by total beds, with a denominator fallback of one to avoid division by zero."])
    section_page(doc,"10. Deployment and Execution",["The checked-in development configuration runs Uvicorn on port 8010 and Vite on 5173; Vite proxies /api and /ws to the backend. SQLite is created in the backend working directory. Models reside under app/ml/artifacts. The repository has no run.bat, Dockerfile or production process manager."])
    section_page(doc,"11. Design Limitations",COMMON_LIMITS)
    save_doc(doc,BATCH/"12.DESIGN &IMPLEMENTATION.docx")


def results_doc():
    doc=configure(Document()); add_title(doc,"RESULTS")
    doc.add_heading("1. Measurement Conditions",1)
    add_para(doc,"Measurements were taken on Windows 11 Pro 64-bit, Intel Core i5-12500 (6 cores, 12 logical processors), 15.7 GB RAM and Python 3.11.9. Models were evaluated on deterministic synthetic test partitions. API timings are 15 sequential loopback requests per read endpoint against Uvicorn, FastAPI and SQLite. The screenshots use a synthetic demonstration database and Google Chrome at 1440 x 900.")
    add_table(doc,["Model","Evidence","Metric","Value"],METRIC_ROWS,caption="Table 1. Measured model results.")
    section_page(doc,"2. Model Performance",["The wait model reconstructs its designed synthetic relationship with R-squared 0.951041. Its MAE of 5.982 minutes and RMSE of 8.580 minutes show that larger misses remain. The triage model reached 96.80% accuracy and 0.898475 macro F1, while its 63% critical recall is the decisive safety limitation."],CHART/"02_model_metrics.png","Figure 1. Measured model metrics.")
    section_page(doc,"3. Triage Error Analysis",["Thirty-seven of 100 actual critical examples were classified as urgent, while two urgent cases were classified as critical and 31 urgent cases as normal. The classifier therefore requires clinician override, threshold/cost-sensitive investigation and external validation before any clinical use."],CHART/"01_triage_confusion_matrix.png","Figure 2. Triage confusion matrix.")
    section_page(doc,"4. API and Build Results",["All 150 measured read requests returned HTTP 200. Mean endpoint times ranged from 3.985 ms for /api/auth/me to 14.231 ms for department load. One WebSocket connection completed in 46.657 ms. These are local sequential measurements and not a capacity claim.","Frontend lint exited zero with 14 warnings. The production build exited zero after 3,573 modules and 1.73 seconds, with warnings for two JavaScript chunks larger than 500 kB."],CHART/"04_api_latency.png","Figure 3. Mean local API latency.")
    screens=[
        ("03_waiting_board.png","Figure 4. Public waiting board","Live department cards show now-serving tokens, waiting counts, upcoming tokens and mean predicted waits."),
        ("04_admin_dashboard.png","Figure 5. Administrator command centre","The demonstration snapshot shows 23 patients today, 18 active visits, 11.1-minute predicted wait, two active critical cases and capacity."),
        ("05_reception_dashboard.png","Figure 6. Reception dashboard","Token, patient, complaint, priority, state and registration time are visible in one table."),
        ("06_patient_registration.png","Figure 7. Registration and triage input","The real form accepts patient identity, department, complaint and optional vitals."),
        ("07_doctor_dashboard.png","Figure 8. Doctor dashboard","Waiting and assigned patients are separated while retaining priority and predicted wait."),
        ("08_patient_detail.png","Figure 9. Visit detail","Vitals, complaint, priority, wait and movement history support the consultation decision."),
        ("09_order_lab_test.png","Figure 10. Laboratory order modal","Doctors may enter a named test or choose a frequent test."),
        ("10_admit_patient.png","Figure 11. Bed admission modal","Occupied beds are disabled and available ward or ICU beds are selectable."),
        ("11_nurse_bed_board.png","Figure 12. General Ward bed board","The synthetic snapshot contains 12 ward beds: seven occupied and five available."),
        ("12_icu_bed_board.png","Figure 13. ICU bed board","The synthetic snapshot contains six ICU beds: four occupied and two available."),
        ("13_laboratory_dashboard.png","Figure 14. Laboratory workflow","Five tests appear across scheduled and in-progress columns with valid actions."),
    ]
    for file,cap,text in screens: section_page(doc,cap.split(". ",1)[1],[text],SHOT/file,cap)
    section_page(doc,"16. Objectives versus Outcomes",[],table=(["Objective","Outcome","Evidence"],[["Digital tokens","Achieved","Token allocation and queue board"],["Movement tracking","Partially achieved","Registration and explicit department changes logged"],["Lab scheduling","Achieved as workflow state","Five populated tests; no optimisation study"],["Live bed/ICU monitoring","Achieved for state visibility","Ward and ICU boards"],["AI wait and priority","Prototype achieved","Measured synthetic-model results"],["Role dashboards","Achieved","Five staff roles plus public board"]],None,"Table 2. Objectives and outcomes."))
    section_page(doc,"17. Success and Failure Cases",["Successful cases include authenticated access, priority-first queue ordering, coherent token display, model inference, laboratory-state display, bed-capacity display and live dashboard refresh. The measured endpoints all returned HTTP 200 under the defined smoke conditions.","Observed limits include 37 critical-to-urgent triage errors in the synthetic test, the nurse page requiring an initial tab click, incomplete coupling between laboratory and visit states, seeded occupied beds without visit links, large frontend chunks and absence of automated tests. No unsupported wait-reduction claim is made."])
    section_page(doc,"18. Results Limitations",COMMON_LIMITS)
    save_doc(doc,BATCH/"14.RESULTS.docx")


def conclusion_doc():
    doc=configure(Document()); add_title(doc,"CONCLUSION, FUTURE SCOPE AND REFERENCES")
    doc.add_heading("1. Conclusion",1)
    add_para(doc,"CareFlow demonstrates an integrated, role-based hospital patient-flow prototype that connects registration, AI-assisted prioritisation, token queues, consultation, laboratory work, bed/ICU state and operational dashboards. The implementation is functional on the supplied local environment and all measured read endpoints passed the defined smoke run.")
    add_para(doc,"The strongest result is integration rather than clinical prediction. The wait regressor achieved R-squared 0.951041 on a generator-matched synthetic test and the triage classifier achieved 96.80% accuracy, but the latter recalled only 63% of critical cases. These results justify continued engineering investigation, not autonomous care decisions.")
    add_figure(doc,CHART/"02_model_metrics.png","Figure 1. Summary of measured model results.")
    section_page(doc,"2. Key Findings",["A single persisted Visit model lets role dashboards agree on token, status, priority and predicted wait. Priority-aware queue sorting is simple and auditable. Movement logs make explicit department changes inspectable. Bed and laboratory modules expose the operational dependencies that ordinary appointment applications omit.","The evidence also shows where design work remains: model safety is class-dependent, state transitions require tighter coupling, the initial nurse tab is faulty, and production security/scaling mechanisms are not present."])
    section_page(doc,"3. Future Scope",[],table=(["Priority","Future work","Expected evidence"],[["1","Validate on de-identified real patient-flow data with ethics and governance approval","External temporal test and subgroup metrics"],["2","Optimise critical recall with cost-sensitive learning and calibrated thresholds","Safety-focused confusion matrices"],["3","Integrate HL7 FHIR with hospital EHR, lab and admission systems","Interoperability conformance tests"],["4","Replace SQLite and in-process WebSocket with production services","Concurrent load, resilience and recovery tests"],["5","Authenticate live channels; rotate secrets; add audit and privacy controls","Security assessment"],["6","Fix dashboard state coupling and nurse default tab","Automated end-to-end tests"],["7","Forecast demand and recommend staffing/bed plans","Prospective operational study"]],None,"Table 1. Prioritised future scope."))
    section_page(doc,"4. Known Limitations",COMMON_LIMITS)
    doc.add_page_break(); doc.add_heading("REFERENCES",1)
    for i,ref in enumerate(REFERENCES,1):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.25); p.add_run(f"[{i}] {ref}")
    save_doc(doc,BATCH/"15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx")


def front_page(doc: Document, title: str, body: list[str]):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); r.bold=True; r.font.size=Pt(18); r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(20,55,90)
    for _ in range(2): doc.add_paragraph()
    for text in body:
        p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_page_break()


def report_page(doc: Document, chapter: str, subtitle: str, paras: list[str], figure: Path | None=None, caption="", table=None):
    doc.add_heading(chapter,1)
    doc.add_heading(subtitle,2)
    if figure: add_figure(doc,figure,caption,width=6.05)
    if table: add_table(doc,*table)
    for p in paras: add_para(doc,p)
    doc.add_page_break()


def final_report():
    doc=configure(Document())
    # 1 Cover
    for _ in range(1): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(TITLE.upper()); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(20,55,90)
    p=doc.add_paragraph("A PROJECT REPORT"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("submitted in partial fulfilment of the requirements for the award of the degree of\n[BACHELOR OF TECHNOLOGY]"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_table(doc,["Student","Hall Ticket No."],[[n,r] for n,r in STUDENTS])
    p=doc.add_paragraph(f"Under the guidance of\n{GUIDE}\n[DESIGNATION]\n\nDepartment of Computer Science and Engineering (Data Science)\n[COLLEGE NAME]\n[CITY]\n[ACADEMIC YEAR]"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    # 2-5
    front_page(doc,"CERTIFICATE",[f"This is to certify that the project report entitled “{TITLE}” is a bonafide record of the work carried out by {', '.join(n for n,_ in STUDENTS)} under the guidance of {GUIDE}, submitted to [COLLEGE NAME] in partial fulfilment of the requirements for the award of [DEGREE] during [ACADEMIC YEAR].","Project Guide: ____________________     Head of Department: ____________________     Principal: ____________________"])
    front_page(doc,"DECLARATION",[f"We declare that the work presented in this report, entitled “{TITLE},” is our original academic project work. Sources used for background, comparison and implementation have been acknowledged. The reported software measurements were obtained from this project’s artefacts and local execution; no clinical deployment or real patient study is claimed.","Signatures: A. RAMU __________  B. RUSHI __________  B. RAJ KUMAR __________  A. PRABHAS __________"])
    front_page(doc,"ACKNOWLEDGEMENT",[f"We express our sincere gratitude to our project guide, {GUIDE}, [DESIGNATION], for guidance and constructive review. We thank [HOD NAME], Head of the Department, [PRINCIPAL NAME], Principal, and the faculty of [COLLEGE NAME] for the academic environment and facilities used for this work.","We also acknowledge the authors of the Alleviate base paper and the open-source communities behind Python, FastAPI, scikit-learn, SQLAlchemy, React and their supporting libraries."])
    front_page(doc,"ABSTRACT",["Hospitals coordinate registration, queues, consultation, laboratory work, pharmacy, wards and intensive care through interdependent decisions. When these states are fragmented across calls, registers or isolated screens, staff cannot see one reliable operational picture. CareFlow is a full-stack academic prototype that issues digital tokens, tracks visits, manages laboratory and bed states, publishes live updates and provides role-specific dashboards.","Two machine-learning components assist the workflow. A Random Forest regressor estimates wait time from synthetic queue context, and a Gradient Boosting classifier estimates triage priority from synthetic vital-sign data labelled with a simplified early-warning rule. On deterministic test partitions the wait model achieved R-squared 0.951041, MAE 5.982 minutes and RMSE 8.580 minutes; triage achieved 96.80% accuracy and 0.898475 macro F1, but only 63% critical recall. The results therefore demonstrate software integration and reproducible synthetic evaluation, not clinical validity."])
    # 6 TOC
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("TABLE OF CONTENTS").bold=True
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()
    # 7 list figures
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("LIST OF FIGURES").bold=True
    fig_names=["System architecture","High-level workflow","DFD Level 0","DFD Level 1","DFD Level 2","Use-case diagram","UML class diagram","Registration sequence","Entity-relationship diagram","Core algorithm flowchart","Triage confusion matrix","Measured model metrics","API latency","Public waiting board","Administrator dashboard","Reception dashboard","Doctor dashboard","Ward and ICU boards","Laboratory dashboard"]
    for i,n in enumerate(fig_names,1): doc.add_paragraph(f"Figure {i}. {n}")
    doc.add_page_break()
    # 8 list tables
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("LIST OF TABLES").bold=True
    for i,n in enumerate(["Literature comparison","Feasibility summary","Functional requirements","Non-functional requirements","User roles","Database entities","API groups","Model results","Smoke-test results","Objectives versus outcomes","Future scope"],1): doc.add_paragraph(f"Table {i}. {n}")
    doc.add_page_break()
    # 9 abbreviations
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("LIST OF ABBREVIATIONS").bold=True
    add_table(doc,["Abbreviation","Meaning"],[["AI","Artificial Intelligence"],["API","Application Programming Interface"],["DFD","Data Flow Diagram"],["EHR","Electronic Health Record"],["ER","Entity Relationship"],["FHIR","Fast Healthcare Interoperability Resources"],["ICU","Intensive Care Unit"],["JWT","JSON Web Token"],["MAE","Mean Absolute Error"],["NEWS2","National Early Warning Score 2"],["OPD","Outpatient Department"],["REST","Representational State Transfer"],["RMSE","Root Mean Squared Error"],["UML","Unified Modeling Language"]])
    doc.add_page_break()

    # Chapter 1: 5 pages (10-14)
    report_page(doc,"CHAPTER 1 INTRODUCTION","1.1 Domain and Background",["Patient flow is the sequence of clinical and operational states through which a person moves after arriving at a hospital. Registration, queueing, consultation, tests, pharmacy, ward and intensive-care resources are coupled: a delay or stale state in one department propagates to others.","CareFlow treats that sequence as shared operational data. Its aim is not to diagnose disease but to make queues, resource state and inter-department transitions visible to the authorised people responsible for them."])
    report_page(doc,"CHAPTER 1 INTRODUCTION","1.2 Motivation and Current Practice",["The supplied project brief identifies manual name calling, telephone bed checks and hand-managed laboratory queues as sources of uncertainty. Conventional hospital management software may digitise billing and registration without coordinating the internal journey in real time.","A useful prototype therefore needs one visit identity, a visible token, explicit status values, department ownership and live resource views. Prediction can support communication only when its evidence and limitations are shown honestly."])
    report_page(doc,"CHAPTER 1 INTRODUCTION","1.3 Problem Statement and Need",["The problem is to coordinate a patient’s in-hospital flow across OPD, laboratory, pharmacy, ward, ICU and emergency services while preserving priority and current state. Staff need to know who is waiting, what action is pending and whether a resource is available. Patients need a clear token and expected wait.","The need is operational: reduce information fragmentation, not to claim a measured clinical outcome. This report distinguishes implemented functionality, synthetic-model performance and unmeasured impact."])
    report_page(doc,"CHAPTER 1 INTRODUCTION","1.4 Objectives and Contributions",["The objectives are to issue digital tokens, track visit movement, coordinate laboratory tasks, expose bed/ICU availability, assist priority and wait estimation, and provide role dashboards. The contributions are a unified Visit state, priority-aware queue ordering, two reproducible model artefacts, resource workflows, WebSocket refresh and a documented evidence boundary."],table=(["Objective","Implemented evidence"],[["Tokens","Department code plus daily sequence"],["Movement","MovementLog on registration/explicit transfer"],["Labs","Order and status routes"],["Beds","Availability, admit and release routes"],["AI assistance","Triage and wait predictors"],["Dashboards","Five roles plus public board"]],None,"Table 1.1 Objectives."))
    report_page(doc,"CHAPTER 1 INTRODUCTION","1.5 Scope, Limitations and Organisation",["Scope ends at an academic local prototype. Billing, real EHR integration, diagnosis, production identity, notification providers, cloud scaling and clinical decision authority are excluded. Synthetic data, 63% critical recall, missing automated tests and production-security gaps are material limitations.","Chapters 2–4 review prior work, feasibility and requirements; Chapters 5–6 describe design and implementation; Chapters 7–8 report testing and results; Chapters 9–10 conclude and identify future work. References and three appendices follow."])

    # Chapter 2: 6 pages
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.1 Overview",["Related work spans patient-facing healthcare applications, disease or specialist recommendation, location services, personalised diet/exercise planning and waiting-time prediction. These streams show the value of integrated access but often stop before in-hospital resource coordination.","The survey uses the supplied Alleviate paper as the base paper and the sources already cited in the batch literature document. Reported figures remain attributed to their original tasks."])
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.2 Base Paper Analysis",["Alleviate combines Flutter/Dart, Flask and Firebase with specialist recommendation, prescription management, clinic dashboards, maps, appointments, diet/exercise planning and doctor verification. It reports 97.24% precision for its logistic-regression recommendation task and 84.76% precision for a diet-plan Random Forest.","Its focus is patient-side access and personalisation. CareFlow adopts the integrated dashboard principle but changes the problem to internal queues, visit state, laboratory work and bed capacity. The model tasks and datasets differ, so numerical comparison would be invalid."],figure=DIAG/"01_system_architecture.png",caption="Figure 2.1 CareFlow architecture used for contrast with the base paper.")
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.3 Related Work: Access and Recommendation",["Tan et al. integrate symptom diagnosis, health records and hospital recommendation; Abinaya and Ganesan, Munir et al., and Nimbalkar and Fadnavis focus on locating appropriate or nearby facilities. Kumar, Gujar et al., and Bhanuteja et al. apply classification to disease or doctor recommendation.","These systems show that structured inputs and role-aware services can improve access. They do not by themselves represent the changing operational state of a patient already moving through a hospital."])
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.4 Related Work: Personalisation and Conversational Support",["Mogaveera et al. and Tseng et al. combine monitoring with diet and exercise guidance. Tanwar et al., Surve et al., Mendapara et al. and Rahman et al. investigate chatbots, NLP and conversational symptom capture. Umbare et al. combine prediction and location features.","CareFlow deliberately excludes general diagnosis and chatbot advice. Its human-computer interaction is role-specific workflow: a small set of auditable actions backed by explicit database state."])
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.5 Waiting-Time Research and Comparison",["The supplied survey identifies recent work on queue waiting-time prediction, outpatient waiting time and dissatisfaction-aware interpretable prediction. These studies motivate predictive communication but differ in institution, features and outcome definition.","CareFlow’s regressor is trained on a designed queue generator rather than observed hospital arrivals. Its R-squared and error values demonstrate consistency with that generator only."],table=(["Work/theme","Approach","Data","Reported result","Limitation for CareFlow"],[["Alleviate","LR/RF integrated app","Paper-specific datasets","97.24% recommendation precision","Different task and dataset"],["Waiting-time studies","ML regression","Observed outpatient context","Source-specific","No shared local dataset"],["CareFlow","RF wait + GB triage","Synthetic generators","R2 0.951; accuracy 0.968","No clinical validation"]],None,"Table 2.1 Comparison boundary."))
    report_page(doc,"CHAPTER 2 LITERATURE SURVEY","2.6 Research Gap",["The gap is a compact system that connects priority-aware tokens, department ownership, movement evidence, laboratory tasks, ward/ICU capacity and live staff views while making the evidence boundary explicit. CareFlow addresses that application gap.","The project does not claim novelty in NEWS2, Random Forests or Gradient Boosting. Its novelty is integrative and operational. Future research must replace synthetic generators with governed real data and evaluate impact prospectively."])

    # Chapter 3: 4 pages
    report_page(doc,"CHAPTER 3 FEASIBILITY STUDY","3.1 Technical and Economic Feasibility",["The prototype uses mature open-source libraries already available on a standard Windows workstation. FastAPI, SQLAlchemy, scikit-learn, React and SQLite reduce development cost. The saved model artefacts and local database fit comfortably within 15.7 GB RAM.","No paid cloud, messaging or data licence was required. This supports academic feasibility but does not estimate production hosting, support, compliance, integration or clinical-governance cost."],table=(["Dimension","Finding"],[["Hardware","Intel i5-12500 and 15.7 GB RAM sufficient"],["Software","Open-source local stack"],["Data","Synthetic generators available; real data absent"],["Cost","Low prototype cost; production cost unmeasured"]],None,"Table 3.1 Feasibility summary."))
    report_page(doc,"CHAPTER 3 FEASIBILITY STUDY","3.2 Operational and Schedule Feasibility",["Each role sees a focused dashboard consistent with its work: registration, consultation, laboratory, bed management or oversight. Seeded demonstration records make the workflow visible without real patient data.","Operational adoption would still require workflow mapping, staff training, downtime procedures, ownership of state updates and governance. The checked-in repository contains no schedule baseline; therefore schedule feasibility is qualitative rather than a measured variance claim."])
    report_page(doc,"CHAPTER 3 FEASIBILITY STUDY","3.3 Advantages and Constraints",["Advantages include one visit state, explicit queue ordering, traceable movement, visible capacity, reusable APIs and live refresh. Constraints include single-node persistence, development credentials, synthetic evidence, incomplete transition coupling and no production integration.","The system remains feasible as an academic proof of implementation. Its constraints are severe enough that deployment without a redesign would be irresponsible."])
    report_page(doc,"CHAPTER 3 FEASIBILITY STUDY","3.4 Risk Assessment",["Primary risks are clinical misclassification, stale state, unauthorised access, privacy exposure, service interruption and misleading interpretation of synthetic metrics. Mitigations include clinician override, validation warnings, authenticated channels, audit trails, monitoring, backups and explicit evidence labels.","The current build implements only part of these mitigations. Risk reduction is a future engineering requirement, not a completed outcome."])

    # Chapter 4: 6 pages
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.1 Existing and Proposed Systems",["The existing condition described in the project brief uses manual or fragmented coordination. CareFlow proposes a shared visit record, digital token, prediction assistance, laboratory and bed modules, role dashboards and public queue display.","The proposed advantages are visibility and consistency. They were demonstrated through populated screens and API smoke tests; no real hospital comparison was conducted."],figure=DIAG/"03_dfd_level_0.png",caption="Figure 4.1 System context.")
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.2 Functional Requirements",[],table=(["ID","Requirement","Implementation"],[["FR1","Authenticate staff","JWT login and role checks"],["FR2","Register visit","POST /api/visits"],["FR3","Issue token","Department code and sequence"],["FR4","Prioritise queue","Classifier plus priority sort"],["FR5","Manage laboratory work","Lab endpoints and board"],["FR6","Manage bed state","Bed endpoints and nurse board"],["FR7","Track movement","MovementLog"],["FR8","Report live state","KPIs, queue board, WebSocket"]],None,"Table 4.1 Functional requirements."))
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.3 Non-Functional Requirements",[],table=(["Category","Requirement","Current evidence / gap"],[["Performance","Responsive local reads","Mean 3.985–14.231 ms; no load test"],["Security","Authenticated staff writes","JWT and role checks; WebSocket public"],["Reliability","Consistent persisted state","SQLite transaction; no failover"],["Usability","Role-focused screens","Real captures; no UAT"],["Maintainability","Modular code","Routes/services/models separated"],["Portability","Documented local run","Windows verified only"]],None,"Table 4.2 Non-functional requirements."))
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.4 Hardware and Software Requirements",["Measured hardware was a 12th Gen Intel Core i5-12500 with 6 cores, 12 logical processors and 15.7 GB RAM. The development environment used Windows 11 Pro, Python 3.11.9, Node/Vite, Google Chrome and the local filesystem.","Backend packages include FastAPI 0.115.0, SQLAlchemy 2.0.35 and scikit-learn 1.5.2. Frontend packages include React 19, TypeScript, Vite, Axios, Zustand, Recharts and Tailwind CSS."])
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.5 User Roles and Constraints",[],table=(["Role","Principal permissions"],[["Reception","Register and view visits"],["Doctor","Queue, consultation, lab order, admission, discharge"],["Nurse","Visit state and bed management"],["Lab technician","Progress laboratory tests"],["Administrator","Hospital-wide read/write access"],["Public viewer","Queue board only"]],None,"Table 4.3 User roles."))
    report_page(doc,"CHAPTER 4 SYSTEM REQUIREMENTS AND ANALYSIS","4.6 Requirement Traceability",["Every implemented objective maps to routes, database entities and one or more output screens. The movement objective is partial because not every implicit transition creates a movement record. Laboratory scheduling is a status workflow, not an optimisation algorithm. Waiting-time reduction is not measured.","This traceability prevents a feature shown in marketing copy from being reported as an achieved research result."])

    # Chapter 5: 12 pages
    design_pages=[
        ("5.1 Overall System Architecture",DIAG/"01_system_architecture.png","Figure 5.1 System architecture","The client, API, prediction layer, persistence and live channel form five replaceable concerns."),
        ("5.2 High-Level Workflow",DIAG/"02_high_level_workflow.png","Figure 5.2 High-level workflow","Registration creates a predicted and prioritised visit that staff route to completion."),
        ("5.3 DFD Level 0",DIAG/"03_dfd_level_0.png","Figure 5.3 DFD Level 0","External actors exchange bounded operational data with the system."),
        ("5.4 DFD Level 1",DIAG/"04_dfd_level_1.png","Figure 5.4 DFD Level 1","Authentication, visit, prediction, resource, movement and reporting processes use four data stores."),
        ("5.5 DFD Level 2",DIAG/"05_dfd_level_2.png","Figure 5.5 DFD Level 2","The visit-creation process is decomposed into validation, identity, inference, queue calculation, persistence and publishing."),
        ("5.6 Use-Case Design",DIAG/"08_use_case_diagram.png","Figure 5.6 Use-case diagram","Use cases reflect the implemented role permissions and public board."),
        ("5.7 UML Class Design",DIAG/"06_uml_class_diagram.png","Figure 5.7 UML class diagram","Persistent domain classes and their relationships are shown without invented services."),
        ("5.8 Sequence Design",DIAG/"07_sequence_diagram.png","Figure 5.8 Registration sequence","A visit request crosses the UI, API, model, database and live-update channel."),
        ("5.9 Database Design",DIAG/"09_er_diagram.png","Figure 5.9 Entity-relationship diagram","Foreign keys connect the seven persistent entities."),
    ]
    for sub,fig,cap,text in design_pages: report_page(doc,"CHAPTER 5 SYSTEM DESIGN",sub,[text],figure=fig,caption=cap)
    report_page(doc,"CHAPTER 5 SYSTEM DESIGN","5.10 Database Schema",["Identifiers are 12-character hexadecimal strings. User email, department code and patient MRN are unique. Visits reference patients, departments, optional doctors and beds. Tests and movements are visit children. JSON stores vitals."],table=(["Entity","Key relationships"],[["User","department_id -> Department"],["Visit","patient, department, doctor, bed"],["Bed","department and optional current visit"],["LabTest","visit_id -> Visit"],["MovementLog","visit and from/to departments"]],None,"Table 5.1 Schema relationships."))
    report_page(doc,"CHAPTER 5 SYSTEM DESIGN","5.11 Module Design",["Backend modules separate configuration, database, security, models, schemas, route groups, queue rules, WebSocket handling and ML prediction. Frontend modules separate layout, common controls, API functions, session/toast stores, live-socket handling and role pages.","The central coupling remains Visit, which intentionally carries the state shared across modules."])
    report_page(doc,"CHAPTER 5 SYSTEM DESIGN","5.12 User Interface Design",["The interface uses a dark operational palette, cards, status badges, compact tables and workflow modals. Admin screens favour aggregate visibility; clinical screens favour actionable lists; the public board uses large tokens for distance viewing.","The captured nurse screen also exposed a real initial-tab defect, demonstrating why visual verification is necessary in addition to code review."],figure=SHOT/"04_admin_dashboard.png",caption="Figure 5.10 Administrator interface example.")

    # Chapter 6: 9 pages
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.1 Overall Methodology",["The method comprised fact extraction, code inspection, deterministic model evaluation, additive demonstration data, running-system capture and evidence-controlled documentation. The software pipeline performs validation, prediction, queue allocation, persistence and live publication."],figure=DIAG/"10_core_algorithm_flowchart.png",caption="Figure 6.1 Core algorithm flow.")
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.2 Dataset Source and Split",["Both datasets are generated in train.py with NumPy RNG seed 42. Wait data contains 12,000 rows and triage data 15,000 rows. Each is split 85/15, giving 10,200/1,800 wait rows and 12,750/2,250 triage rows.","No external licence applies to generated records and no real patient information is used."],figure=CHART/"06_dataset_split.png",caption="Figure 6.2 Dataset partitions.")
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.3 Preprocessing",["Department type and priority are encoded with stored LabelEncoder instances. Tree models consume numeric values directly. The inference path preserves training feature names in pandas DataFrames. Missing vital components are filled from DEFAULT_VITALS: heart rate 78, systolic 118, diastolic 76, SpO2 97, temperature 36.9 and respiratory rate 16.","The silent defaulting is convenient for a demonstration but should become an explicit uncertainty state in a clinical redesign."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.4 Triage Algorithm",["A simplified early-warning function assigns points for extreme respiratory rate, oxygen saturation, temperature, systolic pressure, heart rate and age. Total scores 0–3 are normal, 4–6 urgent and 7 or above critical. A 150-estimator Gradient Boosting classifier learns these labels.","The function is described as NEWS2-style rather than an exact clinical NEWS2 implementation. The saved classifier returns class probability as confidence."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.5 Wait-Time Algorithm",["Synthetic target wait is based on queue ahead divided by servers, multiplied by average service duration, a 1.35 rush factor during defined hours, a priority factor and Gaussian noise. Results are clipped to 0–240 minutes. The Random Forest uses 200 trees and maximum depth 10.","Runtime inference defaults to two servers and current UTC hour. Department average service time is stored in the database."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.6 Training and Evaluation",["The test partitions are deterministic and the existing Joblib models were measured without retraining. Wait R-squared was 0.951041, MAE 5.982 minutes and RMSE 8.580 minutes. Triage accuracy was 0.9680, macro precision 0.962435, macro recall 0.858574 and macro F1 0.898475.","These figures are reproducibility evidence within the generator, not clinical generalisation."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.7 Application Implementation",["FastAPI creates database tables on import and verifies model artefacts at startup. Route functions use SQLAlchemy sessions through dependency injection. Pydantic models validate request ranges and response shape. React pages call a shared Axios client and refresh on selected WebSocket events.","The seed script is destructive by design, but documentation data was inserted additively because existing content had to be preserved."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.8 API, Security and Error Handling",["JWT authentication, bcrypt passwords and route-level role checks provide the implemented security boundary. CORS accepts configured origins and a localhost regular expression for development. HTTP exceptions represent invalid credentials, permissions, missing records and unavailable beds.","Production controls—TLS termination, secret rotation, authenticated WebSockets, privacy logging, rate limits and least-privilege database accounts—are not implemented."])
    report_page(doc,"CHAPTER 6 METHODOLOGY AND SOFTWARE IMPLEMENTATION","6.9 Installation and Execution",["Create or activate the backend virtual environment, install requirements, ensure data exist, then run `python -m uvicorn app.main:app --port 8010`. In frontend, run `npm install` and `npm run dev`; Vite serves port 5173 and proxies API and WebSocket requests to 8010.","The repository has no run.bat. Detailed user steps are in Appendix B."])

    # Chapter 7: 6 pages
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.1 Strategy and Limits",["The repository contains no automated unit or integration tests. Verification therefore combined deterministic model evaluation, frontend lint/build, API smoke/latency requests, WebSocket connection and real-browser workflow capture.","This strategy verifies availability and presentation but cannot substitute for automated regression, concurrency, penetration, clinical or user-acceptance testing."])
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.2 API Smoke Results",["Login returned HTTP 200. Ten read endpoints received 15 sequential requests each, for 150 successful HTTP 200 responses. The WebSocket connection also succeeded."],table=(["Evidence","Result"],[["Login","HTTP 200"],["Read requests","150/150 HTTP 200"],["Mean endpoint range","3.985–14.231 ms"],["WebSocket connect","46.657 ms"]],None,"Table 7.1 Smoke results."))
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.3 Frontend Static and Build Verification",["Oxlint exited zero with 14 warnings concerning ref access, state updates in effects, missing dependencies and render-time randomness. TypeScript and Vite built successfully, transforming 3,573 modules in 1.73 seconds.","Two JavaScript chunks exceeded 500 kB after minification, indicating a code-splitting and performance opportunity."])
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.4 Functional and Edge Observations",["Authenticated dashboards loaded populated data; tokens, priorities, predicted waits, lab states and bed states were visible. Invalid or missing records have explicit HTTP errors in route code. Bed admission rejects non-available beds.","Observed edges include a blank initial nurse content panel until a tab click and incomplete coupling between laboratory test and visit status transitions."])
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.5 Security and Performance Boundaries",["JWT and role checks were exercised indirectly through successful role logins and protected data loads. No unauthorised access matrix, penetration test, dependency scan or privacy review was executed. The unauthenticated WebSocket is a known gap.","Latency values come from one sequential loopback client; no throughput, tail latency under load, memory saturation or multi-process behaviour was measured."],figure=CHART/"04_api_latency.png",caption="Figure 7.1 Local mean read latency.")
    report_page(doc,"CHAPTER 7 SYSTEM TESTING","7.6 User Acceptance Limit",["The screenshots show real application output with synthetic demonstration data, but no hospital staff or independent user group evaluated usefulness, clarity or workflow fit. Therefore user acceptance is unmeasured and is listed as future work."])

    # Chapter 8: 15 pages
    result_pages=[
        ("8.1 Public Waiting Board",SHOT/"03_waiting_board.png","Figure 8.1 Public waiting board","Large token cards make current and upcoming queue state visible without authentication."),
        ("8.2 Administrator Overview",SHOT/"04_admin_dashboard.png","Figure 8.2 Administrator dashboard","The populated snapshot combines KPI, department and capacity evidence."),
        ("8.3 Reception Output",SHOT/"05_reception_dashboard.png","Figure 8.3 Reception dashboard","The registration list shows the operational identity of each visit."),
        ("8.4 Registration Input",SHOT/"06_patient_registration.png","Figure 8.4 Registration form","Vitals are optional inputs to triage prediction."),
        ("8.5 Doctor Queue",SHOT/"07_doctor_dashboard.png","Figure 8.5 Doctor dashboard","Waiting and active visits are separated."),
        ("8.6 Patient Detail",SHOT/"08_patient_detail.png","Figure 8.6 Patient detail","The modal shows complaint, vital context and movement."),
        ("8.7 Laboratory Order",SHOT/"09_order_lab_test.png","Figure 8.7 Laboratory order","A named test can be ordered from the consultation."),
        ("8.8 Admission Selection",SHOT/"10_admit_patient.png","Figure 8.8 Admission modal","Availability is visible before selection."),
        ("8.9 Ward Capacity",SHOT/"11_nurse_bed_board.png","Figure 8.9 Ward bed board","Seven of twelve demonstration ward beds are occupied."),
        ("8.10 ICU Capacity",SHOT/"12_icu_bed_board.png","Figure 8.10 ICU bed board","Four of six demonstration ICU beds are occupied."),
        ("8.11 Laboratory Workflow",SHOT/"13_laboratory_dashboard.png","Figure 8.11 Laboratory dashboard","Five synthetic tests populate scheduled and in-progress work."),
        ("8.12 Model Results",CHART/"02_model_metrics.png","Figure 8.12 Model metrics","Scores summarise generator-matched test performance."),
        ("8.13 Confusion and Failure Analysis",CHART/"01_triage_confusion_matrix.png","Figure 8.13 Triage confusion matrix","Critical recall of 63% is the main model failure finding."),
        ("8.14 API and Capacity Results",CHART/"05_bed_capacity.png","Figure 8.14 Demonstration bed capacity","Capacity is a synthetic application snapshot, not hospital utilisation evidence."),
    ]
    for sub,fig,cap,text in result_pages: report_page(doc,"CHAPTER 8 RESULTS AND OUTPUT SCREENS",sub,[text],figure=fig,caption=cap)
    report_page(doc,"CHAPTER 8 RESULTS AND OUTPUT SCREENS","8.15 Objectives versus Outcomes",["Tokens, role dashboards, model integration, laboratory state and bed visibility were achieved as prototype functions. Movement tracking is partial, laboratory scheduling is not optimised and waiting-time reduction was not measured. The correct conclusion is functional integration with explicit gaps."],table=(["Objective","Outcome"],[["Tokens","Achieved"],["Movement","Partial"],["Lab workflow","Achieved; no optimisation evidence"],["Bed visibility","Achieved"],["AI assistance","Prototype achieved on synthetic data"],["Role dashboards","Achieved"]],None,"Table 8.1 Objectives versus outcomes."))

    # Chapters 9-10, 2 pages each
    report_page(doc,"CHAPTER 9 CONCLUSION","9.1 Summary and Objectives Achieved",["CareFlow implements one coherent patient-flow prototype across registration, queue, consultation, laboratory and bed state. The application executed successfully with populated data, all measured reads returned HTTP 200 and the frontend built successfully.","The model figures are reproducible but synthetic. The project’s valid contribution is an integrated operational design and measured prototype, not a clinical or deployment claim."])
    report_page(doc,"CHAPTER 9 CONCLUSION","9.2 Findings, Applications and Limitations",["Priority-aware sorting, shared visit state and live role views are practical foundations for a future hospital operations tool. The same architecture could support outpatient centres, diagnostic facilities or teaching simulations.","Critical recall, absent real data, missing tests, local-only persistence, security gaps and incomplete transition coupling define the present boundary."])
    report_page(doc,"CHAPTER 10 FUTURE SCOPE","10.1 Improvements and Further Research",["Next work should obtain governed, de-identified operational data; use temporal and external validation; optimise critical recall; calibrate prediction uncertainty; model service capacity explicitly; and conduct clinician-led usability and safety assessment.","Automated unit, integration and end-to-end tests should cover every role and state transition."])
    report_page(doc,"CHAPTER 10 FUTURE SCOPE","10.2 Scalability, Deployment and Integration",["A production redesign should use PostgreSQL, a managed message broker, authenticated multi-instance live channels, TLS, secret management, audit trails, backups and observability. HL7 FHIR interfaces could connect registration, laboratory, admission and EHR systems.","Demand forecasting and staffing recommendations should be introduced only after prospective evaluation with real operational outcomes."])

    # References: 3 pages
    for page_idx in range(3):
        doc.add_heading("REFERENCES",1)
        start=page_idx*7; end=min(len(REFERENCES),(page_idx+1)*7)
        for i in range(start,end):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.28); p.paragraph_format.line_spacing=1.0
            p.add_run(f"[{i+1}] {REFERENCES[i]}")
        doc.add_page_break()

    # Appendix A: 4 pages
    snippets=[
        ("A.1 Queue ordering",'''PRIORITY_RANK = {critical: 0, urgent: 1, normal: 2}\ndef sort_key(visit):\n    return (PRIORITY_RANK.get(visit.priority, 3), visit.token_number)'''),
        ("A.2 Wait prediction features",'''X = DataFrame([{\n  "department_type": dept_code, "queue_ahead": queue_ahead,\n  "hour": hour, "priority": priority_code,\n  "avg_service_minutes": avg_service_minutes, "servers": servers\n}])\nprediction = model.predict(X)[0]'''),
        ("A.3 Visit creation orchestration",'''priority_label, confidence = predict_triage(...)\nahead = queue_ahead_count(db, department.id, priority)\npredicted_wait = predict_wait_minutes(...)\ndb.add(visit); db.add(MovementLog(...)); db.commit()\nawait manager.broadcast("visit_created", visit_out.model_dump())'''),
        ("A.4 Role enforcement",'''def require_roles(*roles):\n    def checker(user = Depends(get_current_user)):\n        if user.role not in roles:\n            raise HTTPException(status_code=403)\n        return user\n    return checker'''),
    ]
    for title,code in snippets:
        doc.add_heading("APPENDIX A: IMPORTANT CODE",1); doc.add_heading(title,2)
        add_para(doc,"The excerpt below is condensed from the checked-in implementation; ellipses indicate omitted arguments rather than new logic.")
        for line in code.splitlines(): doc.add_paragraph(line,style="Code Block")
        doc.add_page_break()

    # Appendix B: 3 pages
    report_page(doc,"APPENDIX B: INSTALLATION AND USER GUIDE","B.1 Local Installation",["Backend: enter 8.ProjectCode/backend, activate the supplied virtual environment or create one, install requirements.txt and ensure medflow.db contains the desired data. Frontend: enter 8.ProjectCode/frontend and run npm install if dependencies are absent.","The project stores models under backend/app/ml/artifacts. Running `python -m app.ml.train` regenerates them and changes the artefact files; do so only when intentional."])
    report_page(doc,"APPENDIX B: INSTALLATION AND USER GUIDE","B.2 Start and Sign In",["Start the backend with `python -m uvicorn app.main:app --port 8010`. Start the frontend with `npm run dev`; open http://127.0.0.1:5173. Use the demo role account shown on the login screen. The public waiting board is /board.","The code contains no run.bat. Stop both development servers with Ctrl+C."])
    report_page(doc,"APPENDIX B: INSTALLATION AND USER GUIDE","B.3 Role Workflows",["Reception creates a registration; doctors open a queue card and start consultation; doctors may order a test or choose an available bed; laboratory staff advance test work; nurses select a ward or ICU tab to inspect beds; administrators monitor KPIs.","Use only synthetic demonstration information unless appropriate privacy, consent and institutional controls are established."])

    # Appendix C: 4 pages, last page has no trailing break removed later by Word naturally
    extras=[
        ("C.1 Staff Sign-In",SHOT/"02_login.png","Figure C.1 Staff sign-in screen."),
        ("C.2 Registration Detail",SHOT/"06_patient_registration.png","Figure C.2 Populated registration form."),
        ("C.3 Patient Detail",SHOT/"08_patient_detail.png","Figure C.3 Patient detail and movement."),
        ("C.4 Admission Detail",SHOT/"10_admit_patient.png","Figure C.4 Available-bed selection."),
    ]
    for i,(sub,fig,cap) in enumerate(extras):
        doc.add_heading("APPENDIX C: ADDITIONAL SCREENS",1); doc.add_heading(sub,2); add_figure(doc,fig,cap,width=6.05)
        add_para(doc,"This is an actual capture of the running prototype using synthetic demonstration data.")
        if i < len(extras)-1: doc.add_page_break()
    save_doc(doc,BATCH/"17.Final Document.docx")


def set_columns(section, num=2, space_twips="360"):
    sectPr=section._sectPr
    cols=sectPr.xpath('./w:cols')
    if cols: cols=cols[0]
    else: cols=OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'),str(num)); cols.set(qn('w:space'),space_twips)


def ieee_par(doc: Document, text: str, indent=True):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
    if indent: p.paragraph_format.first_line_indent=Inches(.2)
    for r in p.runs: r.font.size=Pt(10)
    p.add_run(text)
    for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(10)
    return p


def ieee_heading(doc: Document, text: str):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(text); r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(10)


def ieee_paper():
    doc=Document(); configure(doc,letter=True)
    s=doc.sections[0]; s.top_margin=Inches(.75); s.bottom_margin=Inches(1.0); s.left_margin=Inches(.69); s.right_margin=Inches(.69)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("CareFlow: AI-Assisted Hospital Patient Flow Management with Live Queue and Resource Coordination"); r.font.name="Times New Roman"; r.font.size=Pt(24); r.bold=True
    table=doc.add_table(rows=2,cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    authors=[("A. Ramu, B. Rushi","23K91A6715, 23K91A6720"),("B. Raj Kumar, A. Prabhas","23K91A6725, 23K91A6703"),(GUIDE,"[DESIGNATION]")]
    for i,(name,detail) in enumerate(authors):
        cell=table.cell(0,i); cell.text=name
        cell2=table.cell(1,i); cell2.text=f"{detail}\nDepartment of CSE (Data Science)\n[COLLEGE NAME], [CITY]\n[EMAIL]"
        for c in [cell,cell2]:
            for pp in c.paragraphs:
                pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for rr in pp.runs: rr.font.name="Times New Roman"; rr.font.size=Pt(9)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(3)
    a=p.add_run("Abstract—"); a.bold=True; a.italic=True; a.font.size=Pt(9)
    b=p.add_run("Hospitals depend on coordinated transitions among registration, outpatient queues, consultation, laboratories, wards and intensive care. This paper presents CareFlow, a role-based web prototype that creates digital tokens, records visit movement, exposes laboratory and bed state, publishes live updates and applies two machine-learning models for triage priority and waiting-time estimation. The implementation uses React, FastAPI, SQLAlchemy, SQLite, JWT and WebSockets. The models are evaluated only on deterministic synthetic generators: the wait-time Random Forest obtained R-squared 0.951041, MAE 5.982 minutes and RMSE 8.580 minutes, while the Gradient Boosting triage classifier obtained 96.80% accuracy and 0.898475 macro F1. Critical-class recall was 63%, which prevents a clinical-validity claim. In a local sequential smoke run, 150 of 150 read requests returned HTTP 200. The contribution is an integrated and measurable patient-flow prototype with explicit evidence boundaries, not an autonomous clinical system."); b.font.size=Pt(9)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
    a=p.add_run("Index Terms—"); a.bold=True; a.italic=True; a.font.size=Pt(9)
    b=p.add_run("bed management, hospital operations, patient flow, queue prediction, real-time dashboard, triage, WebSocket"); b.font.size=Pt(9)
    sec=doc.add_section(WD_SECTION.CONTINUOUS); sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.75); sec.bottom_margin=Inches(1); sec.left_margin=Inches(.69); sec.right_margin=Inches(.69); set_columns(sec,2,"360")
    # Roman sections
    ieee_heading(doc,"I. INTRODUCTION")
    ieee_par(doc,"Patient flow is an operational chain rather than a single appointment. A visit may begin at reception, wait in an outpatient or emergency queue, enter consultation, require laboratory work, move to a ward or intensive-care bed and finish at discharge. When those transitions are maintained in telephone calls, registers or disconnected screens, staff cannot see a reliable current state and patients receive uncertain queue information.")
    ieee_par(doc,"CareFlow addresses this coordination problem with one persisted Visit state and role-specific interfaces. Reception creates the visit and token; doctors manage consultation and routing; laboratory staff progress tests; nurses manage bed state; administrators see aggregate load; and a public display exposes now-serving and upcoming tokens. Two estimators provide decision support but do not replace staff judgement.")
    ieee_par(doc,"The contributions are: (1) priority-aware digital tokens linked to one visit identity; (2) integrated laboratory, ward and ICU workflow; (3) live role views using REST and WebSocket events; and (4) a reproducible evidence boundary that separates synthetic-model metrics, local software measurements and unmeasured clinical impact. Section II reviews related work, Section III presents the method, Section IV reports results and Section V concludes.")
    ieee_par(doc,"The scope is deliberately operational. CareFlow does not diagnose disease, recommend treatment, replace a validated emergency scoring protocol or claim interoperability with an installed hospital information system. Instead, it demonstrates how a shared visit identifier can connect actions that are frequently represented in separate lists. The distinction matters because a visually complete dashboard can otherwise be mistaken for clinical evidence. Every quantitative statement in this paper is therefore tagged by its evidence source: deterministic synthetic model evaluation, a populated demonstration database, or a local software smoke measurement.")
    p=doc.add_paragraph(); r=p.add_run("A. System Hardware"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"Measurements used Windows 11 Pro 64-bit on a 12th Gen Intel Core i5-12500 with six physical cores, twelve logical processors and 15.7 GB RAM. Python 3.11.9 executed the model evaluation; Chrome at 1440 x 900 produced the screen evidence.")
    p=doc.add_paragraph(); r=p.add_run("B. Software Requirements"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The backend uses FastAPI 0.115.0, SQLAlchemy 2.0.35, scikit-learn 1.5.2, Joblib and SQLite. The frontend uses React 19, TypeScript, Vite, Axios, Zustand, Recharts and Tailwind CSS. JWT protects staff routes and an in-process WebSocket manager publishes changes.")

    ieee_heading(doc,"II. LITERATURE SURVEY")
    ieee_par(doc,"The Alleviate base paper integrates appointments, prescriptions, clinic dashboards, doctor verification, maps and personalised diet/exercise support [1]. Its logistic-regression recommendation task reports 97.24% precision and its diet model 84.76% precision. Those tasks differ from hospital flow, so their figures are contextual rather than numerical baselines.")
    ieee_par(doc,"Access-oriented systems combine symptom assessment, records and hospital recommendation [2], locate nearby facilities [3]–[5], or predict disease and suggest doctors [6], [9], [16]. Personalisation studies add monitoring, diet or exercise [10], [15], while conversational systems use NLP and chatbots [11]–[14]. These systems improve access or advice but do not necessarily represent the operational state of a person moving among queues, tests and beds.")
    ieee_par(doc,"Recent waiting-time studies motivate regression and interpretable prediction [17]–[19]. Their institutions, features and outcomes differ from this project. CareFlow therefore reports its wait model only as recovery of a designed synthetic process. The research gap is an integrated, auditable application that links priority, token order, department movement, laboratory work, capacity and live role views while exposing the limits of its evidence.")
    ieee_par(doc,"The literature also highlights a recurring evaluation hazard. A high score on a recommendation, diagnosis or waiting-time task cannot be transferred to another hospital, target variable or patient population. CareFlow consequently avoids a numerical league table. The cited work establishes design motivation—integrated access, personalisation, conversational support and delay prediction—whereas the prototype contribution is the explicit linkage of operational states and the reproducible reporting of what has and has not been measured.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run("TABLE I\nRELATED WORK AND RESEARCH GAP"); rr.bold=True; rr.font.size=Pt(8)
    add_table(doc,["Work","Approach","Data/result","Limitation"],[["Alleviate [1]","Integrated mobile AI","97.24% recommendation precision","Different task"],["Access systems [2]–[6]","Diagnosis/location","Source-specific","No internal flow"],["Personalisation [10],[15]","Diet/exercise","User metrics","No capacity state"],["Waiting studies [17]–[19]","ML prediction","Observed outpatient settings","Not shared locally"],["CareFlow","RF + GB + workflow","Synthetic tests","No clinical validation"]])

    ieee_heading(doc,"III. METHODOLOGY")
    ieee_par(doc,"The implemented architecture has five concerns: React interfaces, FastAPI routes, SQLAlchemy persistence, a Joblib model layer and WebSocket publication. JWT decoding resolves the current user and role dependencies guard write routes. Seven entities—User, Department, Patient, Visit, Bed, LabTest and MovementLog—represent staff, operational state and traceability.")
    add_figure(doc,DIAG/"01_system_architecture.png","Fig. 1. Implemented CareFlow architecture.",width=3.2)
    ieee_par(doc,"At registration, the API validates the department, matches a patient by phone or creates an MRN, completes missing individual vital values with documented defaults, predicts priority, allocates the next token, counts higher-ranked waiting visits, predicts expected wait, persists Visit and MovementLog, commits and broadcasts visit_created. Queue order is critical, urgent and normal, then token number.")
    p=doc.add_paragraph(); r=p.add_run("A. State and Role Design"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"Visit is the central operational record. It contains the patient and department links, token, priority, predicted wait, current status and timing fields. MovementLog records transitions; LabTest records an order independently of its display state; and Bed represents General Ward or ICU capacity. Reception may register visits, doctors may start or complete consultation and route the patient, laboratory staff may progress tests, nurses may allocate or release beds, and administrators may view system-wide summaries. This separation reduces accidental cross-role writes while retaining a common read model.")
    ieee_par(doc,"The frontend requests REST resources for stable lists and subscribes to WebSocket messages for change notification. On receipt, the relevant view refreshes its data rather than trusting the event as the sole source of truth. The approach is simple and suitable for a single-process prototype, but it does not provide durable delivery, replay or multi-instance fan-out. Those properties would require a message broker, authenticated channel subscriptions and an event contract with versioning.")
    add_figure(doc,DIAG/"10_core_algorithm_flowchart.png","Fig. 2. Core triage and wait-time flow.",width=3.2)
    add_figure(doc,DIAG/"02_high_level_workflow.png","Fig. 3. Cross-role patient-flow state transitions.",width=3.2)
    p=doc.add_paragraph(); r=p.add_run("B. Synthetic Data and Triage Model"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The triage generator creates 15,000 rows of age, heart rate, systolic pressure, SpO2, temperature and respiratory rate. A simplified early-warning function assigns normal for scores 0–3, urgent for 4–6 and critical for 7 or more. An 85/15 stratified split produces 12,750 training and 2,250 test rows. GradientBoostingClassifier uses 150 estimators, maximum depth three and random state 42.")
    p=doc.add_paragraph(); r=p.add_run("C. Wait-Time Model"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The wait generator creates 12,000 rows with department type, queue ahead, UTC hour, priority, average service minutes and server count. The target approximates (queue/servers) times service duration, multiplied by rush and priority factors, plus Gaussian noise, clipped to 0–240 minutes. An 85/15 split gives 10,200 training and 1,800 test rows. RandomForestRegressor uses 200 trees and maximum depth ten.")
    ieee_par(doc,"Both generators use NumPy seed 42. Encoders and feature order are stored with the estimators. The datasets contain no real patient data and have no external dataset licence. This design supports reproducibility, not clinical generalisation.")
    p=doc.add_paragraph(); r=p.add_run("D. Measurement Protocol"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The evaluation regenerated each dataset from the recorded seed and applied the stored 85/15 split policy. Regression was assessed with coefficient of determination, mean absolute error and root mean squared error. Classification was assessed with accuracy and macro-averaged precision, recall and F1, followed by the complete three-class confusion matrix. Macro measures were included because normal, urgent and critical classes are not equally frequent. Critical recall was treated as a safety-relevant diagnostic rather than hidden behind aggregate accuracy.")
    ieee_par(doc,"Software observations were collected separately. Ten read endpoints were requested 15 times each over loopback, yielding 150 responses, and WebSocket connection establishment was timed once. The frontend was checked with its configured lint and production-build commands. These are engineering smoke checks, not concurrent load tests: they exclude network distance, browser concurrency, write contention, warm-up analysis, confidence intervals, database growth and production infrastructure.")
    p=doc.add_paragraph(); r=p.add_run("E. Reproducibility and Operational Invariants"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"Reproducibility rests on four recorded elements: fixed random seeds, explicit generator formulas, fixed estimator hyperparameters and saved preprocessing metadata. The model artefacts retain the label encoders and feature order required at inference time. The application facts file records the evaluated environment, sample counts, class order and exact metrics. This permits another evaluator to distinguish a regenerated experiment from a screenshot value and prevents a model score from being silently relabelled as a hospital outcome.")
    ieee_par(doc,"Several invariants are required for operational correctness. A token must be unique within its allocation scope; priority ordering must precede token ordering; a bed cannot be allocated to two active visits; discharge must release the associated resource; laboratory completion must not imply visit completion unless a separate transition occurs; and movement history must preserve who changed a state and when. The prototype implements portions of these rules through route logic and relational fields, but it does not yet enforce all of them with database constraints and automated concurrent tests.")
    ieee_par(doc,"Security was reviewed as implemented rather than assumed from framework choice. Staff HTTP routes use JWT-derived identity and role dependencies, passwords are represented through the authentication layer, and the frontend separates role views. Remaining production gaps include a development secret that must be replaced and rotated, missing authenticated WebSocket subscription, incomplete audit coverage, no documented rate limiting, and no formal privacy or retention policy. Consequently, the demonstration must not contain identifiable real patient records.")
    add_figure(doc,DIAG/"09_er_diagram.png","Fig. 4. Core data entities and relationships.",width=3.2)

    ieee_heading(doc,"IV. RESULTS")
    ieee_par(doc,"The wait model achieved R-squared 0.951041, MAE 5.982 minutes and RMSE 8.580 minutes. The triage classifier achieved 0.9680 accuracy, 0.962435 macro precision, 0.858574 macro recall and 0.898475 macro F1. Of 100 actual critical test examples, 63 were classified critical and 37 urgent. This 63% critical recall is more decision-relevant than the overall accuracy.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run("TABLE II\nMEASURED MODEL AND SOFTWARE RESULTS"); rr.bold=True; rr.font.size=Pt(8)
    add_table(doc,["Measure","Value","Condition"],[["Wait R-squared","0.951041","1,800 synthetic test rows"],["Wait MAE / RMSE","5.982 / 8.580 min","Generator-matched test"],["Triage accuracy / macro F1","0.9680 / 0.898475","2,250 synthetic test rows"],["Critical recall","0.63","100 critical test cases"],["HTTP smoke","150/150 status 200","Sequential loopback reads"],["Mean endpoint range","3.985–14.231 ms","15 requests per endpoint"]])
    add_figure(doc,CHART/"02_model_metrics.png","Fig. 5. Measured model scores.",width=3.2)
    ieee_par(doc,"A running synthetic demonstration database contained seven users, seven departments, 23 patients and visits, 18 beds and five laboratory tests. The administrator snapshot showed 23 patients for the day, 18 active visits, 11.1 minutes average predicted wait, two active critical cases, 61% overall bed occupancy and 67% ICU occupancy. These are demonstration values, not hospital-study results.")
    add_figure(doc,SHOT/"04_admin_dashboard.png","Fig. 6. Populated administrator dashboard.",width=3.2)
    add_figure(doc,SHOT/"07_doctor_dashboard.png","Fig. 7. Doctor queue and active patient view.",width=3.2)
    add_figure(doc,CHART/"04_api_latency.png","Fig. 8. Sequential loopback endpoint latency.",width=3.2)
    ieee_par(doc,"Frontend lint exited zero with 14 warnings and the production build exited zero after transforming 3,573 modules. The browser capture reported no console errors. Observed failures include a nurse tab that requires explicit selection on first load, incomplete coupling of laboratory and visit states, seeded occupied beds without linked visits, and two large JavaScript chunks.")
    ieee_par(doc,"The endpoint means ranged from 3.985 to 14.231 ms in the sequential loopback run, and the measured WebSocket connection took 46.657 ms. The variation reflects endpoint work and local scheduling, but the protocol is too small to support a service-level objective. Likewise, 150 successful reads establish only that the selected routes responded in this environment; they do not establish authorization completeness, transaction safety under competing writes or resilience after process failure.")
    p=doc.add_paragraph(); r=p.add_run("A. Error Analysis and Validity"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The triage confusion matrix, ordered critical, normal and urgent, was [[63, 0, 37], [0, 1525, 2], [2, 31, 590]]. Most errors therefore involved adjacent urgency boundaries, but 37 critical examples were reduced to urgent. Even though no critical example was predicted normal, the missed critical decisions are sufficient to reject autonomous use. Future evaluation must specify the clinical cost of each transition, establish thresholds with clinicians and test calibration and subgroup performance on governed real data.")
    ieee_par(doc,"Internal validity is limited by formula-derived labels and generator-matched train and test distributions. External validity is absent because the records are synthetic and originate from no hospital. Construct validity is also limited: the target represents a simplified waiting process and the early-warning function is not a validated triage protocol. Software validity is limited by one workstation, one browser viewport, sequential requests and no automated end-to-end test suite. These limitations are part of the result, not footnotes to be removed from it.")
    p=doc.add_paragraph(); r=p.add_run("B. Capacity and Workflow Observations"); r.italic=True; r.font.size=Pt(10)
    ieee_par(doc,"The demonstration bed inventory comprised 18 resources across ward and ICU categories. Overall occupancy was 61%, while ICU occupancy was 67% in the captured state. These figures verify that the dashboard can aggregate seeded resource records and distinguish categories; they do not estimate utilisation for a hospital. Five laboratory tests exercised pending and progressed views, and 23 visits provided non-empty queue, consultation and dashboard states. The records were intentionally synthetic and remain unsuitable for epidemiological or operational inference.")
    ieee_par(doc,"Visual walkthroughs covered the public landing page, staff login, public waiting board, administrator summary, reception registration, doctor queue, patient detail, laboratory ordering, admission, ward, ICU and laboratory views. This coverage establishes that the major screens could be reached and populated in the recorded environment. It is not a substitute for scripted acceptance testing: transition assertions, permission-denial cases, refresh consistency, keyboard accessibility, responsive layouts and recovery after interrupted writes remain to be tested.")
    ieee_par(doc,"The observed nurse first-load issue illustrates the value of testing workflow state rather than only route availability. A page may load successfully while its default tab does not expose the expected data until the user selects a ward explicitly. Similarly, a laboratory record can progress without proving that every corresponding visit transition is synchronised. These findings motivate end-to-end scenarios that begin with registration and assert the final patient, visit, test, movement and bed records together.")
    ieee_par(doc,"Comparison with Alleviate is qualitative because datasets and tasks differ. Alleviate emphasises patient access and recommendation; CareFlow emphasises operational flow and shared resource state. No measured reduction in real waiting time, clinical accuracy, deployment availability or user acceptance is claimed.")

    ieee_heading(doc,"V. CONCLUSION")
    ieee_par(doc,"CareFlow demonstrates that a compact full-stack system can unify digital tokens, priority-aware queues, movement, laboratory work, bed/ICU visibility and live role dashboards. The local software checks succeeded and the saved models produced reproducible synthetic-test results. However, critical recall, synthetic data and absent clinical evaluation prevent clinical use.")
    ieee_par(doc,"The practical lesson is that integration and evidence should advance together. A system may correctly expose a queue and still require transition invariants, authorization tests and failure recovery; a model may score highly on generated samples and still be unsafe for a patient population. Retaining these distinctions makes the prototype a useful engineering baseline for supervised research rather than an overstated clinical product.")
    ieee_par(doc,"A staged validation path is proposed. First, automated unit and API tests should lock down token allocation, state transitions, authorization and resource release. Second, concurrent synthetic scenarios should measure database and event consistency. Third, clinicians should review the feature definitions, error costs and user interface using only non-identifiable cases. Only after governance approval should a de-identified retrospective dataset be used for temporal and external model evaluation, calibration and subgroup analysis. Prospective deployment would then require monitoring, rollback criteria and human override.")
    ieee_par(doc,"Future work should obtain governed de-identified flow data, perform temporal and external validation, optimise critical recall, calibrate uncertainty, add automated state-transition tests and conduct clinician-led usability assessment. Production engineering should replace SQLite and in-process broadcasts, authenticate live channels, rotate secrets, add audit and privacy controls, and integrate hospital systems through HL7 FHIR.")
    ieee_heading(doc,"REFERENCES")
    for i,ref in enumerate(REFERENCES,1):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
        r=p.add_run(f"[{i}] {ref}"); r.font.name="Times New Roman"; r.font.size=Pt(8)
    save_doc(doc,BATCH/"19.IEEE_Paper.docx")


def main():
    methodology_doc(); scope_doc(); diagrams_doc(); design_doc(); results_doc(); conclusion_doc(); final_report(); ieee_paper()


if __name__ == "__main__":
    main()
