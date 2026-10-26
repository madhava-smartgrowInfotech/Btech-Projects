from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


BATCH = Path(__file__).resolve().parents[2]
CODE = BATCH / "8.ProjectCode"

DELIVERABLES = [
    "7.Review_1_PPT.pptx",
    "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx",
    "10.SCOPE & LIMITATIONS.docx",
    "11. BLOCK DIAGRAMS ALL.docx",
    "12.DESIGN &IMPLEMENTATION.docx",
    "13.Review_2_PPT.pptx",
    "14.RESULTS.docx",
    "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    "16.Final_PPT.pptx",
    "17.Final Document.docx",
    "19.IEEE_Paper.docx",
]

DECK_COUNTS = {
    "7.Review_1_PPT.pptx": 16,
    "13.Review_2_PPT.pptx": 20,
    "16.Final_PPT.pptx": 28,
}

STUDENTS = {
    "A. RAMU": "23K91A6715",
    "B. RUSHI": "23K91A6720",
    "B. RAJ KUMAR": "23K91A6725",
    "A. PRABHAS": "23K91A6703",
}
GUIDE = "[GUIDE NAME]"

BANNED = [
    "[STUDENT NAME",
    "[ROLL NO",
    "Fakeddit",
    "PHEME",
    "Snopes",
    "MCNN",
    "bank transaction",
    "social media manipulation",
    "skincare",
    "Dayananda",
    "Sindhu",
    "recommender system",
]

EXPECTED_SCREENSHOTS = [
    "01_landing.png",
    "02_login.png",
    "03_waiting_board.png",
    "04_admin_dashboard.png",
    "05_reception_dashboard.png",
    "06_patient_registration.png",
    "07_doctor_dashboard.png",
    "08_patient_detail.png",
    "09_order_lab_test.png",
    "10_admit_patient.png",
    "11_nurse_bed_board.png",
    "12_icu_bed_board.png",
    "13_laboratory_dashboard.png",
]

failures: list[str] = []


def check(condition: bool, message: str):
    if condition:
        print(f"PASS: {message}")
    else:
        failures.append(message)
        print(f"FAIL: {message}")


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def pptx_text(path: Path) -> str:
    prs = Presentation(path)
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                parts.append(shape.text)
    return "\n".join(parts)


def package_xml_text(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as package:
        check(package.testzip() is None, f"{path.name} package opens cleanly")
        for name in package.namelist():
            if name.lower().endswith((".xml", ".rels")):
                chunks.append(package.read(name).decode("utf-8", errors="ignore"))
    return "\n".join(chunks)


def first_slide_text(path: Path) -> str:
    slide = Presentation(path).slides[0]
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))


def first_doc_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs[:20]]
    for table in doc.tables[:2]:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def embedded_images(path: Path) -> int:
    with zipfile.ZipFile(path) as package:
        return sum(1 for name in package.namelist() if "/media/" in name and not name.endswith("/"))


def has_two_column_body(path: Path) -> bool:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    sections = root.findall(".//w:sectPr", ns)
    column_counts = []
    for section in sections:
        cols = section.find("w:cols", ns)
        column_counts.append(1 if cols is None else int(cols.get(f"{{{ns['w']}}}num", "1")))
    return len(column_counts) >= 2 and column_counts[0] == 1 and 2 in column_counts[1:]


def verify_deliverables():
    for name in DELIVERABLES:
        path = BATCH / name
        check(path.exists(), f"exact deliverable exists: {name}")
        if path.exists():
            check(path.stat().st_size >= 40 * 1024, f"{name} is at least 40 KB")


def verify_clean_text():
    for name in DELIVERABLES:
        path = BATCH / name
        if not path.exists():
            continue
        visible = docx_text(path) if path.suffix.lower() == ".docx" else pptx_text(path)
        package = package_xml_text(path)
        combined = visible + "\n" + package
        hits = [term for term in BANNED if term.casefold() in combined.casefold()]
        check(not hits, f"{name} has no forbidden template/reference leftovers" + (f" ({', '.join(hits)})" if hits else ""))


def verify_decks():
    for name, expected in DECK_COUNTS.items():
        path = BATCH / name
        if not path.exists():
            continue
        prs = Presentation(path)
        check(len(prs.slides) == expected, f"{name} has {expected} slides")
        title = first_slide_text(path)
        for student, roll in STUDENTS.items():
            check(student in title and roll in title, f"{name} title slide contains {student} and {roll}")
        check(GUIDE in title, f"{name} title slide contains guide value {GUIDE}")


def verify_report():
    path = BATCH / "17.Final Document.docx"
    if not path.exists():
        return
    text = docx_text(path)
    for number in range(1, 11):
        check(re.search(rf"CHAPTER\s+{number}\b", text, re.I) is not None, f"final report contains CHAPTER {number}")
    check(re.search(r"\bREFERENCES\b", text, re.I) is not None, "final report contains REFERENCES")
    for letter, label in [("A", "IMPORTANT CODE"), ("B", "INSTALLATION AND USER GUIDE"), ("C", "ADDITIONAL SCREENS")]:
        check(re.search(rf"APPENDIX\s+{letter}.*{label}", text, re.I | re.S) is not None, f"final report contains Appendix {letter}: {label.title()}")
    check(embedded_images(path) >= 20, "final report embeds at least 20 figures/screens")
    title = first_doc_text(path)
    for student, roll in STUDENTS.items():
        check(student in title and roll in title, f"final report title page contains {student} and {roll}")
    check(GUIDE in title, f"final report title page contains guide value {GUIDE}")


def verify_chapter_docs_and_screens():
    screenshot_dir = CODE / "docs" / "screenshots"
    missing = [name for name in EXPECTED_SCREENSHOTS if not (screenshot_dir / name).exists()]
    check(not missing, "all 13 referenced application screenshots exist" + (f" (missing: {missing})" if missing else ""))
    check(embedded_images(BATCH / "11. BLOCK DIAGRAMS ALL.docx") >= 10, "block-diagram document embeds all ten required diagrams")
    check(embedded_images(BATCH / "14.RESULTS.docx") >= 13, "results document embeds the captured output screens")
    chapter_docs = [name for name in DELIVERABLES if name.endswith(".docx") and name != "19.IEEE_Paper.docx"]
    for name in chapter_docs:
        title = first_doc_text(BATCH / name)
        for student, roll in STUDENTS.items():
            check(student in title and roll in title, f"{name} title page contains {student} and {roll}")
        check(GUIDE in title, f"{name} title page contains guide value {GUIDE}")


def verify_ieee():
    path = BATCH / "19.IEEE_Paper.docx"
    pdf = BATCH / "19.IEEE_Paper.pdf"
    if not path.exists():
        return
    text = docx_text(path)
    check(has_two_column_body(path), "IEEE paper has a single-column title block and two-column body")
    for section in ["I. INTRODUCTION", "II. LITERATURE SURVEY", "III. METHODOLOGY", "IV. RESULTS", "V. CONCLUSION"]:
        check(section in text, f"IEEE paper contains {section}")
    check("TABLE I" in text and "TABLE II" in text, "IEEE paper contains TABLE I and TABLE II")
    figure_captions = re.findall(r"\bFig\.\s*\d+\.", text)
    check(len(figure_captions) >= 4, "IEEE paper contains at least four captioned figures")
    references = re.findall(r"(?m)^\[(\d{1,2})\]\s", text)
    check(15 <= len(references) <= 20, f"IEEE paper contains 15–20 references (found {len(references)})")
    check(pdf.exists(), "IEEE PDF exists alongside the Word paper")
    if pdf.exists():
        check(len(PdfReader(pdf).pages) == 6, "IEEE PDF is exactly six pages")
    title = first_doc_text(path)
    for student, roll in STUDENTS.items():
        check(student.casefold() in title.casefold() and roll in title, f"IEEE author block contains {student} and {roll}")
    check(GUIDE in title, f"IEEE author block contains guide value {GUIDE}")


def main() -> int:
    print(f"Verifying documentation in {BATCH}")
    verify_deliverables()
    verify_clean_text()
    verify_decks()
    verify_report()
    verify_chapter_docs_and_screens()
    verify_ieee()
    if failures:
        print(f"\nFAILED: {len(failures)} check(s)")
        for item in failures:
            print(f" - {item}")
        return 1
    print("\nALL DOCUMENTATION CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
