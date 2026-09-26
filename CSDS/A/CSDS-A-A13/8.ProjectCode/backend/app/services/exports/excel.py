"""Excel exports: hall-wise lists (with a roll-order door list) and the attendance report."""
from __future__ import annotations

import io
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.engine.rolls import roll_sort_key
from app.services.exports.data import PlanExport, paper_hex

HEADER_FILL = PatternFill("solid", fgColor="493EE5")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MUTED_FONT = Font(italic=True, color="666666")
THIN = Side(style="thin", color="D9DBE6")


def _light(hex_colour: str, white: float = 0.7) -> str:
    """A pale tint of a paper colour for cell backgrounds (dark text stays readable)."""
    r, g, b = (int(hex_colour.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(v + (255 - v) * white):02X}" for v in (r, g, b))


def _sheet_name(text: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "-", text)[:31] or "Sheet"
    base, n = name, 2
    while name in used:
        name = f"{base[:28]}-{n}"
        n += 1
    used.add(name)
    return name


def _table(ws, start_row: int, headers: list[str], rows: list[list], widths: list[int]) -> None:
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=i, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = Border(bottom=THIN)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + max(len(rows), 1)}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{start_row}:{start_row}"


def _heading(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = MUTED_FONT


def hall_lists_xlsx(export: PlanExport) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    used = {"Summary"}
    _heading(summary, f"Hall lists - {export.title}", export.stamp)
    rows = [[h.hall.code, h.hall.name, h.hall.building, len(h.seats), h.hall.capacity,
             ", ".join(f"{p['paper']} ({p['count']})" for p in h.legend), ", ".join(h.invigilators)] for h in export.halls]
    rows.append(["Total", "", "", sum(len(h.seats) for h in export.halls), sum(h.hall.capacity for h in export.halls), "", ""])
    _table(summary, 4, ["Hall", "Name", "Building", "Candidates", "Seats", "Papers", "Invigilators"], rows,
           [11, 26, 18, 12, 9, 60, 30])
    for cell in summary[4 + len(rows)]:
        cell.font = Font(bold=True)

    for item in export.halls:
        ws = wb.create_sheet(_sheet_name(item.hall.code, used))
        _heading(ws, f"{item.hall.code} · {item.hall.name}", f"{export.title} · invigilator: {', '.join(item.invigilators) or '-'}")
        data = [[s.label, s.roll_no, s.full_name, s.department, s.course_code, s.course_name,
                 "yes" if s.needs_accessible else "", ""] for s in item.seats]
        _table(ws, 4, ["Seat", "Roll no.", "Name", "Dept.", "Course", "Course name", "Accessible", "Signature"], data,
               [7, 12, 26, 8, 10, 30, 11, 22])
        for r, s in enumerate(item.seats, start=5):
            ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=_light(paper_hex(s.colour)))
            ws.cell(row=r, column=5).font = Font(bold=True)

    door = wb.create_sheet(_sheet_name("Door list", used))
    _heading(door, "Where do I sit?", f"{export.title} · sorted by roll number - post at the entrance")
    everyone = sorted(((s, h.hall) for h in export.halls for s in h.seats), key=lambda t: roll_sort_key(t[0].roll_no))
    _table(door, 4, ["Roll no.", "Name", "Course", "Hall", "Seat", "Building"],
           [[s.roll_no, s.full_name, s.course_code, hall.code, s.label, hall.building] for s, hall in everyone],
           [12, 28, 10, 10, 7, 20])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def attendance_xlsx(export: PlanExport) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    used = {"Summary"}
    _heading(summary, f"Attendance - {export.title}", export.stamp)
    rows = []
    totals = Counter()
    for h in export.halls:
        c = Counter(s.attendance or "not marked" for s in h.seats)
        totals.update(c)
        rate = c["present"] / len(h.seats) if h.seats else 0
        rows.append([h.hall.code, len(h.seats), c["present"], c["absent"], c["not marked"], round(rate * 100, 1)])
    everyone = sum(len(h.seats) for h in export.halls)
    rows.append(["Total", everyone, totals["present"], totals["absent"], totals["not marked"],
                 round(totals["present"] / everyone * 100, 1) if everyone else 0])
    _table(summary, 4, ["Hall", "Candidates", "Present", "Absent", "Not marked", "Present %"], rows, [11, 12, 10, 10, 12, 11])
    for cell in summary[4 + len(rows)]:
        cell.font = Font(bold=True)

    by_paper: dict[str, Counter] = {}
    for h in export.halls:
        for s in h.seats:
            by_paper.setdefault(s.course_code, Counter())[s.attendance or "not marked"] += 1
    start = 7 + len(rows)
    summary.cell(row=start - 1, column=1, value="By course").font = Font(bold=True)
    for i, h in enumerate(["Course", "Candidates", "Present", "Absent", "Not marked"], start=1):
        cell = summary.cell(row=start, column=i, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r, (code, c) in enumerate(sorted(by_paper.items()), start=start + 1):
        for col, value in enumerate([code, sum(c.values()), c["present"], c["absent"], c["not marked"]], start=1):
            summary.cell(row=r, column=col, value=value)

    status_fill = {"present": "D1FAE5", "absent": "FEE2E2"}
    for item in export.halls:
        ws = wb.create_sheet(_sheet_name(item.hall.code, used))
        _heading(ws, f"{item.hall.code} · {item.hall.name}", export.title)
        data = [[s.label, s.roll_no, s.full_name, s.course_code, s.attendance or "not marked",
                 s.marked_at.strftime("%Y-%m-%d %H:%M") if s.marked_at else "", s.marked_by or ""] for s in item.seats]
        _table(ws, 4, ["Seat", "Roll no.", "Name", "Course", "Status", "Marked at (UTC)", "Marked by"], data,
               [7, 12, 26, 10, 12, 17, 20])
        for r, s in enumerate(item.seats, start=5):
            if s.attendance in status_fill:
                ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=status_fill[s.attendance])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
