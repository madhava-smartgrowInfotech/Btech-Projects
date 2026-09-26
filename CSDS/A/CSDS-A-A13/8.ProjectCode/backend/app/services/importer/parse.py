"""Read CSV / Excel uploads into rows keyed by canonical column names."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

from openpyxl import load_workbook

from app.core.errors import AppError
from app.services.importer.spec import KINDS, SPECS, TITLES, header_map, normalise_header

MAX_ROWS = 50_000
MAX_BYTES = 15 * 1024 * 1024


class Table:
    """Rows from one sheet or CSV file, with the spreadsheet row number of each."""

    def __init__(self, kind: str, rows: list[dict[str, str | None]], row_numbers: list[int],
                 unknown_columns: list[str], missing_columns: list[str]):
        self.kind = kind
        self.rows = rows
        self.row_numbers = row_numbers
        self.unknown_columns = unknown_columns
        self.missing_columns = missing_columns


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _to_table(kind: str, header: list, body: list[list], first_row_number: int) -> Table:
    mapping = header_map(kind)
    columns: list[str | None] = []
    unknown: list[str] = []
    for raw in header:
        name = mapping.get(normalise_header(raw)) if raw not in (None, "") else None
        if raw not in (None, "") and name is None:
            unknown.append(str(raw))
        columns.append(name)
    present = {c for c in columns if c}
    missing = [c.name for c in SPECS[kind] if c.required and c.name not in present]
    rows, numbers = [], []
    for offset, raw_row in enumerate(body):
        values = {}
        for idx, name in enumerate(columns):
            if name and idx < len(raw_row):
                text = _cell_text(raw_row[idx])
                if text is not None or name not in values:
                    values[name] = text
        if not any(v is not None for v in values.values()):
            continue  # skip blank lines
        rows.append({c.name: values.get(c.name) for c in SPECS[kind]})
        numbers.append(first_row_number + offset)
    if len(rows) > MAX_ROWS:
        raise AppError(f"The {TITLES[kind].lower()} file has more than {MAX_ROWS:,} rows. Split it into smaller files.")
    return Table(kind, rows, numbers, unknown, missing)


def _read_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise AppError("The CSV file is not in a readable text encoding. Save it as 'CSV UTF-8' and try again.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _sheet_rows(ws) -> list[list]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def read_upload(kind: str, filename: str, data: bytes) -> dict[str, Table]:
    """Return {kind: Table}. A workbook upload returns one table per recognised sheet."""
    if not data:
        raise AppError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise AppError("The file is larger than 15 MB.")
    name = filename.lower()
    if name.endswith(".csv"):
        if kind == "workbook":
            raise AppError("A combined import needs an Excel workbook (.xlsx) with one sheet per kind of data.")
        rows = _read_csv(data)
        if not rows:
            raise AppError("The file has no header row.")
        return {kind: _to_table(kind, rows[0], rows[1:], 2)}
    if not name.endswith((".xlsx", ".xlsm")):
        raise AppError("Upload a .csv or .xlsx file. Older .xls files must be saved as .xlsx first.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for damaged files
        raise AppError("The Excel file could not be opened. Check that it is a valid .xlsx file.") from exc

    tables: dict[str, Table] = {}
    if kind == "workbook":
        by_title = {normalise_header(t): k for k, t in TITLES.items()} | {k: k for k in KINDS}
        for ws in wb.worksheets:
            target = by_title.get(normalise_header(ws.title))
            if target:
                rows = _sheet_rows(ws)
                if rows:
                    tables[target] = _to_table(target, rows[0], rows[1:], 2)
        if not tables:
            raise AppError("No sheet named Courses, Candidates, Halls or Timetable was found in the workbook.")
        return tables

    # Single kind: prefer a sheet named after the kind, else the first sheet.
    ws = next((w for w in wb.worksheets if normalise_header(w.title) in {kind, normalise_header(TITLES[kind])}),
              wb.worksheets[0])
    rows = _sheet_rows(ws)
    if not rows:
        raise AppError("The sheet is empty.")
    return {kind: _to_table(kind, rows[0], rows[1:], 2)}
