"""Excel and CSV import templates, written from the column specification."""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.importer.spec import KINDS, SPECS, TITLES

HEADER_FILL = PatternFill("solid", fgColor="493EE5")
OPTIONAL_FILL = PatternFill("solid", fgColor="8B84F0")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FONT = Font(italic=True, color="666666")


def _fill_sheet(ws, kind: str, rows: Iterable[Mapping[str, object]] | None, example: bool) -> None:
    columns = SPECS[kind]
    ws.append([c.name for c in columns])
    for idx, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL if column.required else OPTIONAL_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        width = max(len(column.name), len(column.example)) + 4
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 12), 48)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    if rows is not None:
        for row in rows:
            ws.append(["" if row.get(c.name) is None else row.get(c.name) for c in columns])
    elif example:
        ws.append([c.example for c in columns])
    for idx, column in enumerate(columns, start=1):
        if column.name == "needs_accessible_seat":
            letter = get_column_letter(idx)
            dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}5000")


def _instructions(ws, kinds: Iterable[str], about: str | None) -> None:
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 30
    ws.append(["SeatWise import"])
    ws["A1"].font = Font(bold=True, size=14)
    if about:
        ws.append([about])
        ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
    ws.append(["Dark headers are required columns; light headers are optional. Keep the header row as it is."])
    ws.cell(row=ws.max_row, column=1).font = NOTE_FONT
    for kind in kinds:
        ws.append([])
        ws.append([f"Sheet: {TITLES[kind]}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(["Column", "Required", "What to enter", "Example"])
        for idx in range(1, 5):
            ws.cell(row=ws.max_row, column=idx).font = Font(bold=True)
        for column in SPECS[kind]:
            ws.append([column.name, "yes" if column.required else "no", column.description, column.example])
            ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")


def template_xlsx(kind: str) -> bytes:
    """A single-sheet template with one example row, plus a 'How to fill' sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = TITLES[kind]
    _fill_sheet(ws, kind, None, example=True)
    _instructions(wb.create_sheet("How to fill"), [kind], None)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def workbook_xlsx(data: Mapping[str, list[Mapping[str, object]]] | None = None, about: str | None = None) -> bytes:
    """A combined workbook with one sheet per kind (courses, candidates, halls, timetable)."""
    wb = Workbook()
    wb.remove(wb.active)
    _instructions(wb.create_sheet("How to fill"), KINDS, about)
    for kind in KINDS:
        ws = wb.create_sheet(TITLES[kind])
        _fill_sheet(ws, kind, data.get(kind) if data else None, example=data is None)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_csv(kind: str, rows: Iterable[Mapping[str, object]] | None = None) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([c.name for c in SPECS[kind]])
    for row in rows if rows is not None else [{c.name: c.example for c in SPECS[kind]}]:
        writer.writerow(["" if row.get(c.name) is None else row.get(c.name) for c in SPECS[kind]])
    return buffer.getvalue().encode("utf-8-sig")
