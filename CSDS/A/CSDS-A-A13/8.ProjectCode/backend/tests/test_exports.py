"""F5 - PDF and Excel outputs."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from tests.helpers import import_sample, reset

RULES = {"adjacency": 8, "roll_gap": 5, "department_mix": True, "fill_strategy": "compact", "accessible_per_hall": 2}


@pytest.fixture(scope="module")
def plan(client, admin_headers):
    reset(client, admin_headers)
    import_sample(client, admin_headers)
    session = client.get("/api/sessions", headers=admin_headers).json()[0]
    created = client.post(f"/api/sessions/{session['id']}/plans", headers=admin_headers, json={"rules": RULES, "seed": 7})
    assert created.status_code == 201
    return created.json()


def _pages(pdf: bytes) -> int:
    return pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page\r") + pdf.count(b"/Type /Page ") + pdf.count(b"/Type /Page>")


@pytest.mark.parametrize("name", ["seating-charts.pdf", "invigilator-sheets.pdf", "qr-slips.pdf"])
def test_pdf_exports(client, admin_headers, plan, name):
    response = client.get(f"/api/plans/{plan['id']}/exports/{name}", headers=admin_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert "attachment" in response.headers["content-disposition"]
    if name == "seating-charts.pdf":
        assert _pages(response.content) == plan["halls_used"]


def test_single_hall_chart(client, admin_headers, plan):
    hall = plan["halls"][0]
    response = client.get(f"/api/plans/{plan['id']}/exports/seating-charts.pdf", params={"hall_id": hall["hall_id"]},
                          headers=admin_headers)
    assert response.status_code == 200 and _pages(response.content) == 1
    assert hall["code"] in response.headers["content-disposition"]


def test_hall_lists_workbook(client, admin_headers, plan):
    response = client.get(f"/api/plans/{plan['id']}/exports/hall-lists.xlsx", headers=admin_headers)
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Door list"
    assert len(wb.sheetnames) == plan["halls_used"] + 2
    door = wb["Door list"]
    assert door.max_row - 4 == plan["candidates"]           # header on row 4, one row per candidate
    rolls = [door.cell(row=r, column=1).value for r in range(5, door.max_row + 1)]
    assert rolls == sorted(rolls)


def test_attendance_workbook(client, admin_headers, plan):
    response = client.get(f"/api/plans/{plan['id']}/exports/attendance.xlsx", headers=admin_headers)
    wb = load_workbook(io.BytesIO(response.content))
    assert wb["Summary"]["A1"].value.startswith("Attendance")


def test_invigilators_cannot_export_everything(client, invigilator_headers, plan):
    response = client.get(f"/api/plans/{plan['id']}/exports/hall-lists.xlsx", headers=invigilator_headers)
    assert response.status_code == 403
