"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import (BATCH, DIAG, EXP, GUIDE, GUIDE_DESIG, COLLEGE, YEAR, PRODUCT, REFERENCES, SHOTS, STUDENTS, TITLE)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE (Data Science)  |  Batch A13  |  Final Review"


def px(v):
    return Emu(int(v * 9525))


def table(s, sid, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 9)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l)
    set_paras(shape(s, cap_ids[1]), cap_r)
    for cid in cap_ids:
        shape(s, cid).top = px(554)


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Final Review", TITLE, "Batch A13", "exam seating, CP-SAT, optimisation"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A13"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    # 1 title
    s = S[0]
    set_paras(shape(s, 4), TITLE)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A13")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of SeatWise: AI-enabled digital examination seating management.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Engine",
                "Design & Implementation", "Results, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["Exam seating is prepared by hand or in spreadsheets before every cycle: slow, error-prone and predictable, so candidates of the same paper can end up side by side.",
                "SeatWise treats seating as constraint optimisation: OR-Tools CP-SAT seats everyone so that no two neighbours (diagonals included) write the same paper, close roll numbers are apart and accessible seats are honoured.",
                "Every plan is a random valid choice from a recorded seed, scored by an independent validator and reproducible with one click (Verify).",
                "A FastAPI + React application adds validated import, a drag-to-swap seat map, PDF/Excel outputs, QR seat slips, tablet attendance, public lookup, analytics and an audit trail.",
                "Benchmark (8 scenarios x 3 seeds): 24 of 24 plans met every hard rule; 2.4 s for 600 and 4.8 s for 2,000 candidates; roll-order seating leaves 1,253 same-paper pairs at 600 candidates.",
                "Keywords: exam seating, constraint programming, CP-SAT, malpractice prevention, web application."], 19, bullet=True)
    set_notes(s, "Abstract: problem, approach, application, and the benchmark headline.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [f"{REFERENCES[0]}",
                "Method: Flask + SQLite web application; students uploaded as CSV or PDF; a zipper algorithm (itertools.zip_longest) interleaves two subject lists; optional branch constraint; blocked seats skipped.",
                "Output: grid map, ReportLab PDF chart, hall-wise reports, public search by USN.",
                "Reported results: none measured. The only quantitative statement is a requirement: 100 students allocated in under 3 seconds.",
                "Limitations (our reading): two subjects only; no roll-number or accessible-seat rule; fixed pattern; no attendance; no optimisation. Authors' future work: mobile app, e-mail/SMS, face recognition, cloud, drag-and-drop swap."], 19, bullet=True)
    set_notes(s, "The base paper is a practical Flask system with no measured results, so we compare rules and behaviour.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Yukth et al., 2025 (base) [1]", "Flask app; zipper interleaving of two subjects", "Institutional lists", "No measured results; 100 students < 3 s target", "Two subjects; no roll or accessibility rule"],
            ["Adetona and Akintoye, 2023 [4]", "Web app; greedy graph colouring; SMS reminders", "University test group", "Limited testing reported in abstract", "Same-course separation only; figures not read"],
            ["Kashyap et al., 2021 [6]", "Genetic multi-objective search", "Not reviewed", "Not reviewed", "Not reviewed in detail"],
            ["Jitha and Fathima, 2022 [7]", "Web portal, CSV upload, static seat table", "Student CSV lists", "Per [1]: reduces notice-board crowding", "Per [1]: no interleaving"],
            ["Patel, 2021 [8]", "Biometric attendance + QR seat lookup", "University", "Per [1]: seat lookup by QR", "Per [1]: costly, needs internet"],
            ["LCG allocation, 2021 [9]", "Linear congruential generator shuffle", "Student lists", "Per [1]: fully random seats", "Per [1]: confuses invigilators"]]
    table(s, 6, rows, [0.2, 0.25, 0.14, 0.22, 0.19], 14, 440)
    set_notes(s, "Works [2], [3] and [5] from the batch survey could not be retrieved, so they are not summarised here.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Interleaves two subjects only; no guarantee for diagonals, roll-number gaps or accessible seats; no measured results"],
            ["Adetona and Akintoye [4]", "Colouring separates same-course students; no roll-number, accessibility or randomness control reported"],
            ["Kashyap et al. [6]", "Search-based and not guaranteed to satisfy every rule; no reproducible seed or audit"],
            ["Fixed-pattern tools [7], [9]", "Predictable patterns or full randomness that breaks rules; no hybrid"],
            ["All reviewed work", "No complete application joining validated import, optimisation, live swap checks, QR slips, attendance and audit"]]
    table(s, 7, rows, [0.26, 0.74], 17, 440)
    set_notes(s, "The gap is rule-respecting randomness with measured evidence in a complete application.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Seating plans are slow to prepare and error-prone, and predictable patterns allow malpractice.",
                "Many rules apply together: capacity, papers, roll-number spacing, accessible seats and department mix.",
                "Candidates and staff have no reliable digital way to find or track seat allocations.",
                "Problem: given candidates, courses, halls and a timetable, produce a fair, unpredictable seating plan in which no two neighbours write the same paper, close roll numbers are apart and accessible seats are honoured, and deliver it through a secure web application whose correctness is measured."], 20, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["Automate seat allocation from uploaded student, course, hall and timetable data, with row-level validation.",
                "Satisfy capacity, subject-combination and roll-number constraints with OR-Tools CP-SAT.",
                "Enforce anti-cheating spacing and randomise seating from a recorded, reproducible seed.",
                "Generate seating charts, hall lists, invigilator sheets and QR seat slips.",
                "Provide real-time seat lookup and tablet attendance tracking.",
                "Reduce the time and errors of manual preparation, and benchmark the engine against manual methods.",
                "Verify with automated tests and a reproducible benchmark."], 20, bullet=True)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["1. Import - Excel/CSV for courses, halls, candidates and timetable, validated row by row before saving.",
                "2. Pre-check - seats, accessible seats and the largest paper are checked; impossible requests are explained.",
                "3. Stage A - CP-SAT shares candidates between halls and gives each paper one colour class of the seat grid.",
                "4. Stage B - CP-SAT lays out seats per hall: roll gap, accessible seats, department mix, seeded random choice.",
                "5. Validate - an independent checker scores the plan; seed and hashes are stored.",
                "6. Review and publish - drag-to-swap with live checks, versions, audit trail.",
                "7. Use - charts, QR slips, public lookup, tablet attendance, analytics."], 19, bullet=False)

    # 10 architecture
    s = S[9]
    put_picture(s, DIAG / "01_system_architecture.png", (240, 178, 800, 424), remove_ids=(7,), name="System architecture")
    set_paras(shape(s, 8), "Layered architecture: browsers, React frontend, FastAPI backend with engine and services, SQLite storage")

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "End-to-end workflow: import, validate, set rules, generate, validate, review, publish, outputs, on the day, analyse")

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Rules & Model")
    fill(s, 7, ["Synthetic sample (seed 2026): 49 courses, 7 departments, 908 candidates, 2,724 registrations, 14 halls (871 seats), 6 sittings; no real person or record.",
                "Benchmark scenarios: 150 to 2,000 candidates, 4 to 14 papers, 3 to 39 halls; rebuilt from their seeds.",
                "Rules (the model's parameters): 8- or 4-neighbour adjacency, roll gap 5, department mix, compact or balanced halls, 2 accessible seats by default.",
                "Model: a two-stage CP-SAT constraint model. Nothing is trained, so there is no train/validation/test split.",
                "Settings: per-hall budget 0.15 (tuned), Stage A budget 0.3, one worker per hall, engine 1.0.0, OR-Tools 9.15.6755.",
                "Groups: candidates are grouped by (paper, department, accessible need); courses sharing a question paper count as one paper."], 19, bullet=True)
    set_notes(s, "The model is a constraint model, not a trained model: the honest evaluation is a benchmark against manual methods.")

    # 13 decision workflow
    s = S[12]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (300, 178, 680, 424), remove_ids=(7,), name="Algorithm flowchart")
    set_paras(shape(s, 8), "Core algorithm: pre-checks, Stage A (2 passes), seeded split, Stage B per hall, retry, validate and store")

    # 14 user-system interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    put_picture(s, DIAG / "09_sequence_diagram.png", (270, 178, 740, 424), remove_ids=(7,), name="Sequence diagram")
    set_paras(shape(s, 8), "Sequence: generate a plan, publish it, and a candidate finds the seat")

    # 15 database
    s = S[14]
    put_picture(s, DIAG / "06_er_diagram.png", (250, 178, 780, 424), remove_ids=(7,), name="ER diagram")
    set_paras(shape(s, 8), "Entity\u2013relationship diagram (16 SQLite tables)")

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["Import", "CSV/XLSX parsing, row checks incl. timetable clashes, preview, re-validation at commit, a batch applies once"],
            ["Engine", "CP-SAT Stage A (two passes) and Stage B per hall in parallel threads; seeded; at most 5 retries"],
            ["Reproducibility", "SHA-256 sub-seeds; data_fingerprint, solver_hash, assignment_hash; Verify re-runs the engine"],
            ["Seat map", "dnd-kit drag or tap; swap-check marks every seat green or red with the reason"],
            ["Outputs", "ReportLab PDFs (charts, sheets, QR slips), openpyxl hall lists and attendance report"],
            ["Security and tests", "bcrypt, JWT 12 h, roles, rate-limited public lookup, audit trail; 58 backend tests"]]
    table(s, 17, rows, [0.2, 0.8], 17, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Engine Benchmark (600 candidates)")
    rows = [["Measure (mean of 3 seeds)", "SeatWise", "Sequential (roll order)"],
            ["Same-paper neighbour pairs", "0", "1,253"], ["Roll-gap violations", "0", "453"], ["Accessible seats missed", "0", "6"],
            ["Same-department neighbour pairs", "109", "1,648"], ["Solve time", "2.36 s", "under 0.001 s"], ["Plans meeting every hard rule", "100% (24 of 24 overall)", "0%"]]
    table(s, 18, rows, [0.46, 0.27, 0.27], 19, 420)
    set_paras(shape(s, 8), "Synthetic 600-candidate sitting (8 papers, 6 departments, 12 halls); run engine-benchmark-20260921; not an estimate of real-world performance.")
    set_notes(s, "There is no accuracy or F1: nothing is trained. These are the measured rule violations and solve time.")

    # 18 violations & predictability
    s = S[17]
    set_paras(shape(s, 4), "Rule Violations & Predictability by Method")
    two_pics(s, EXP / "conflicts_by_method.png", EXP / "neighbour_overlap.png", "Same-paper neighbour pairs by method and size (SeatWise: 0)", "Neighbour pairs repeated across seeds: SeatWise 1.0%, sequential 100%")
    set_notes(s, "This is the equivalent of a confusion matrix for an optimisation engine: how many rule violations each method leaves, and how predictable each method is.")

    # 19 parameter analysis
    s = S[18]
    set_paras(shape(s, 4), "Parameter Analysis & Scaling")
    two_pics(s, EXP / "budget_tuning.png", EXP / "solve_time_by_size.png", "Per-hall budget: 0.15 chosen (2.63 s, 123 same-department pairs)", "Solve time grows slowly with sitting size: 4.8 s at 2,000 candidates")
    set_notes(s, "The budget is the one tunable setting; the scaling chart shows near-linear growth with the number of halls.")

    # 20 performance & testing
    s = S[19]
    rows = [["Measurement", "Result"],
            ["Solve time", "2.4 s (600 candidates), 4.8 s (2,000); slowest benchmark run 6.1 s (4-neighbour mode)"],
            ["API latency (local)", "median 2-11 ms for data endpoints; seat map 7 ms; seating-chart PDF 96 ms; Excel lists 314 ms; QR slips PDF 3.2 s"],
            ["Verify a plan", "reproduced exactly (hashes equal) in about 2.2 s"],
            ["Automated tests", "58 passed, 0 failed (engine 14, import 9, plans 6, exports 7, attendance 5, public 5, analytics 4, auth 8)"],
            ["Demo scenario", "end to end in 11.4 s: import, three plans (1.8-2.2 s, 0 same-paper pairs), exports, lookup, attendance"]]
    table(s, 21, rows, [0.22, 0.78], 18, 430)
    set_notes(s, "Not measured: concurrent-user load, memory use, other hardware; user-acceptance testing was not performed.")

    # 21 comparison
    s = S[20]
    rows = [["System", "Data", "Method", "Rules enforced", "Reported / measured", "Extras"],
            ["Base paper [1]", "Institutional lists; none reported", "Zipper interleave of two subjects", "Subject alternation, optional branch rule, blocked seats", "No results; 100 students in < 3 s target", "PDF chart, grid map, public search"],
            ["Round-robin baseline (ours)", "Synthetic, 600 candidates", "Papers interleaved cyclically", "None guaranteed", "281 same-paper pairs, 198 roll-gap violations", "-"],
            ["SeatWise", "Synthetic, 150-2,000 candidates", "CP-SAT, seeded, validated", "No same-paper neighbours (8 sides), roll gap, accessible seats, capacity", "0 violations in 24 of 24 plans; 2.4 s at 600", "Swap map, QR slips, attendance, analytics, audit"]]
    table(s, 22, rows, [0.13, 0.15, 0.16, 0.22, 0.19, 0.15], 13, 400)
    set_paras(shape(s, 8), "Different data: the base paper reports no measured results, so no numeric head-to-head is possible. Round-robin is our closest proxy for its interleaving idea, not a re-implementation of the base paper.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Plan and Seat Map")
    two_pics(s, SHOTS / "09_plan.png", SHOTS / "11_seatmap_targets.png", "Plan page: scorecard, versions, halls, Verify", "Seat map: a picked-up candidate and the seats that keep every rule")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Seat Lookup and Attendance")
    two_pics(s, SHOTS / "18_lookup_result.png", SHOTS / "21_attendance_list.png", "Public seat lookup with hall map and QR slip", "Tablet attendance: present, absent, to mark; submit and lock")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Dashboard and Analytics")
    two_pics(s, SHOTS / "04_dashboard.png", SHOTS / "14_analytics_engine.png", "Dashboard after three published plans and attendance", "Analytics: engine performance from the benchmark")

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built SeatWise: constraint-optimised, reproducible exam seating with a complete web application.",
                "24 of 24 benchmark plans met every hard rule: 0 same-paper neighbours against 1,253 for roll-order seating (600 candidates).",
                "2.4 s for 600 and 4.8 s for 2,000 candidates; roll-to-seat correlation 0.065; 1% of neighbour pairs repeat across seeds.",
                "58 backend tests pass; the demo scenario runs import, planning, exports, lookup and attendance end to end.",
                "Limits: synthetic data only; no user-acceptance test; front-row bias 4.7 points; 4-neighbour mode slower (6.0 s)."], 20, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Pilot on a real, anonymised sitting and gather feedback from invigilators and the examination cell.",
                "E-mail or SMS seat notifications and an installable mobile app (also the base paper's future work).",
                "Reduce front-row bias (for example by rotating colour classes) and speed up 4-neighbour mode.",
                "Invigilator rostering and objectives such as seat visibility and distance.",
                "Multi-institution deployment with a server database; face-recognition attendance only after a privacy review."], 20, bullet=True)

    # 27 references
    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], 14, min_pt=8, bullet=True, gap_pt=2)

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
