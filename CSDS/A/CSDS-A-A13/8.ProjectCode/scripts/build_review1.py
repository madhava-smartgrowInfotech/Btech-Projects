"""Build 7.Review_1_PPT.pptx from D:\\abstracts\\Templates\\Review_1_PPT.pptx (16 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (style_chart, fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, table_size)  # noqa: E402
from doc_common import (BATCH, BENCH, DIAG, GUIDE, GUIDE_DESIG, COLLEGE, YEAR, DEPT, PRODUCT, REFERENCES, STUDENTS, TITLE)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_1_PPT.pptx"
OUT = BATCH / "7.Review_1_PPT.pptx"


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 1", TITLE, "Batch A13", "exam seating, CP-SAT, optimisation"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A13"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), TITLE)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A13 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 1 covers problem identification and research for SeatWise, an AI-enabled digital examination seating management system.")

    # 2 contents: unchanged tiles
    set_notes(s := S[1], "Nine parts: the abstract, the base paper and related work, the gap we found, our problem statement and objectives, the approach, expected outcomes and references.")

    # 3 abstract
    s = S[2]
    set_paras(shape(s, 2), "Abstract: seating as a constraint problem")
    fill(s, 8, "Exam seating is still prepared by hand or in spreadsheets before every cycle. It is slow and error-prone, and predictable patterns let candidates of the same paper sit side by side, which invites malpractice. Capacity, papers, roll numbers and accessible seats are hard to satisfy together by hand.", 16)
    fill(s, 13, "Treat seating as a constraint-optimisation problem. SeatWise uses Google OR-Tools CP-SAT to seat every candidate so that no two neighbours, diagonals included, write the same paper and close roll numbers are kept apart. It then picks one valid plan at random from a recorded seed.", 16)
    fill(s, 18, "A web application with validated Excel/CSV import, plans in seconds, a seat map with live rule checks, printable charts and QR seat slips, tablet attendance, public seat lookup, analytics and an audit trail.", 16)
    set_notes(s, "The abstract in one slide: the manual process is slow and predictable; we model it as constraint optimisation and wrap it in a usable web application.")

    # 4 base paper
    s = S[3]
    fill(s, 5, BASE_PAPER_TEXT, 14)
    set_paras(shape(s, 7), "Method")
    fill(s, 8, "A Flask and SQLite web application. Student lists are uploaded as CSV or PDF. A \"zipper\" algorithm (itertools.zip_longest) interleaves two subject lists so subjects alternate, with an optional branch constraint and blocked seats skipped. Output: grid map, PDF chart, public search by USN.", 15)
    set_paras(shape(s, 10), "Data and reported results")
    fill(s, 11, "No dataset and no measured results are reported. The only quantitative statement is a requirement: allocate a batch of 100 students in under 3 seconds. The system is evidenced by screenshots of the working application.", 15)
    set_paras(shape(s, 13), "Limits and future work")
    fill(s, 14, "Our reading: only two subjects are interleaved; no roll-number or accessible-seat rule; the pattern is fixed. The authors' future work: mobile app, e-mail/SMS, face-recognition attendance, cloud hosting, drag-and-drop seat swapping.", 15)
    set_notes(s, "The base paper is a practical Flask system. It reports no numbers, so we compare behaviour and rules rather than metrics. Its own future work names drag-and-drop swapping, which SeatWise already provides.")

    # 5 literature survey
    s = S[4]
    rows = [["Work", "Approach", "Data", "Reported result"],
            ["Yukth et al., 2025 (base) [1]", "Flask web app; zipper interleaving of two subject lists; PDF chart and grid map", "Institutional student, hall and exam lists", "No measured results; target of 100 students in under 3 s"],
            ["Adetona and Akintoye, 2023 [4]", "Web app; greedy graph colouring keeps same-course students apart; SMS reminders", "University test group", "Abstract reports successful limited testing; no figures read"],
            ["Kashyap et al., 2021 [6]", "Genetic multi-objective search (subject similarity, distance, seat visibility)", "Not reviewed in detail", "Not reviewed in detail"],
            ["Jitha and Fathima, 2022 [7]", "Central web portal; CSV upload; static seat table", "Student CSV lists", "Per [1]: no interleaving of subjects"],
            ["Patel, 2021 [8]", "Smart seating with biometric attendance and QR lookup", "University setting", "Per [1]: costly hardware, needs internet"],
            ["LCG allocation, 2021 [9]", "Linear congruential generator shuffles seats", "Student lists", "Per [1]: random order confuses invigilators"],
            ["Batch survey [2], [3], [5]", "Automated hall allocation and seating systems, 2025-2026", "Not independently retrieved", "Listed by title only; not verified"]]
    sz = table_size(rows, 1164, 460, [0.26, 0.34, 0.18, 0.22], 13, 9)
    set_table(shape(s, 6), rows, sz, header_size=sz)
    set_notes(s, "Seven works. Only Adetona and Akintoye could be verified online beyond the base paper; the last row is listed by title because the abstracts could not be retrieved. We do not claim their methods or results.")

    # 6 supporting evidence
    s = S[5]
    set_paras(shape(s, 2), "Supporting evidence: manual-style seating breaks the rules")
    cats = ["150", "300", "600", "1,000", "2,000"]
    keys = ["xs-150", "s-300", "m-600", "l-1000", "xl-2000"]
    ser = [(m, [BENCH[k][m][1] for k in keys]) for m in ("Sequential", "Round-robin", "Random shuffle", "SeatWise")]
    ch = replace_chart(s, 4, cats, ser)
    style_chart(ch, "Same-paper neighbour pairs by seating method and sitting size (candidates)", ["E4572E", "F2A541", "7A8BA8", "2F9E6D"], axis_title="neighbour pairs", label_pt=7)
    set_paras(shape(s, 8), "Takeaway")
    fill(s, 9, "Roll-order seating of a 600-candidate sitting leaves 1,253 same-paper neighbour pairs. Round-robin still leaves 281 and a random shuffle 265. The same-paper pairs are plotted by sitting size (candidates).", 15)
    set_paras(shape(s, 13), "Caveat")
    fill(s, 14, "Measured by us on synthetic sittings (benchmark of 8 scenarios). The base paper reports no comparable figures, and real examination data may differ.", 15)
    set_notes(s, "This is our own measurement, from the engine benchmark run on 21 September 2026: same-paper neighbour pairs by method and sitting size. SeatWise has zero in every run.")

    # 7 existing system
    s = S[6]
    set_paras(shape(s, 2), "Existing system and current approaches")
    steps = ["Collect lists", "Sort by roll number", "Fill halls in order", "Print notice-board chart", "Fix clashes by hand", "Redo on any change"]
    for sid, t in zip((5, 8, 11, 14, 17, 20), steps):
        fill(s, sid, t, 14)
    set_paras(shape(s, 24), "Manual and spreadsheet")
    fill(s, 25, "The examination cell sorts candidates by roll number and fills halls row by row. It takes days for thousands of candidates, and every late change means redoing the plan.", 16)
    set_paras(shape(s, 29), "Software with fixed rules")
    fill(s, 30, "Basic tools, and the base paper's zipper method, alternate two subjects or branches in a fixed pattern. Faster, but predictable and limited to the rules they encode.", 16)
    set_notes(s, "How seating is done today: by hand or with software that applies a fixed pattern.")

    # 8 limitations
    s = S[7]
    set_paras(shape(s, 2), "Limitations of existing methods")
    for tid, bid, ttl, body in [
        (7, 8, "Slow and error-prone", "Days of effort; duplicate or missing seats; any late change forces a full redo."),
        (12, 13, "Predictable", "Sequential seating repeats 100% of neighbour pairs when regenerated (measured); fixed patterns can be guessed."),
        (17, 18, "Few rules", "Subject clashes, roll-number spacing, accessible seats and department mix are hard by hand; the base paper interleaves two subjects only."),
        (22, 23, "No digital access", "Printed charts only: candidates cannot look up a seat, and staff cannot track attendance or audit changes.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Four limitations: slow, predictable, few rules, and no digital access for candidates and staff.")

    # 9 research gap
    s = S[8]
    gaps = [("Rules", "Existing tools follow fixed sort or interleave rules; few guarantee that no neighbour, diagonals included, writes the same paper [1], [4]."),
            ("Spacing", "Roll-number adjacency and accessible-seat needs are rarely enforced together with paper separation."),
            ("Randomness", "Pure shuffling breaks the rules and fixed patterns are predictable: rule-respecting randomness is missing [9]."),
            ("Evidence", "Reviewed work reports no measured conflict counts or solve times; the base paper states only a 3-second target."),
            ("Application", "No single tool joins validated import, optimisation, live swap checks, QR slips, attendance and an audit trail.")]
    for (lid, tid), (lab, txt) in zip([(5, 6), (8, 9), (11, 12), (14, 15), (17, 18)], gaps):
        set_paras(shape(s, lid), lab)
        fill(s, tid, txt, 17)
    set_notes(s, "The gap is the combination: guaranteed anti-cheating spacing, controlled randomness, measured evidence and a complete application.")

    # 10 problem statement
    s = S[9]
    fill(s, 5, "Given candidate lists, course registrations, hall layouts and a timetable, seat every candidate so that no two neighbours write the same paper, close roll numbers are kept apart and accessible seats are honoured; choose one such plan at random from a recorded seed; and deliver it, with seat lookup, attendance and reports, through a secure web application whose correctness is measured on benchmark sittings.", 22)
    set_paras(shape(s, 8), "Candidates, courses, halls, timetable")
    set_paras(shape(s, 11), "Seating plan, scorecard, charts, QR slips")
    set_paras(shape(s, 14), "Reproducible seed + audit trail")
    set_notes(s, "The problem statement: inputs are the four data files; outputs are plans, charts and slips; extra properties are reproducibility and auditability.")

    # 11 objectives
    s = S[10]
    fill(s, 4, ["Automate seat allocation from uploaded student, course, hall and timetable data, with row-level validation.",
                "Satisfy capacity, subject-combination and roll-number constraints with constraint optimisation (OR-Tools CP-SAT).",
                "Enforce anti-cheating spacing (no same-paper neighbours, diagonals included) and randomise seating from a recorded seed.",
                "Generate exportable seating charts, hall lists, invigilator sheets and QR seat slips.",
                "Provide real-time seat lookup and tablet attendance tracking.",
                "Reduce the time and errors of manual preparation, and measure the engine against manual methods.",
                "Verify the system with automated tests and a reproducible benchmark."], 18, bullet=True)
    set_notes(s, "Six objectives from the project definition plus a seventh on verification.")

    # 12 proposed approach
    s = S[11]
    for tid, bid, ttl, body in [
        (7, 8, "Validated import", "Excel/CSV templates; row-by-row checks including timetable clashes."),
        (12, 13, "Stage A: halls", "CP-SAT shares candidates between halls and confines each paper to one colour class of the seat grid."),
        (17, 18, "Stage B: seats", "Per-hall CP-SAT: roll gap, accessible seats, department mix, seeded random choice."),
        (22, 23, "Validate + verify", "An independent checker scores every plan; seed and hashes make it reproducible.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 14)
    set_paras(shape(s, 27), "Web application")
    shape(s, 28).height = Emu(44 * 9525)
    fill(s, 28, "React + FastAPI + SQLite: seat map with drag-to-swap, PDF and Excel outputs, QR slips, tablet attendance, public lookup, analytics, audit trail.", 16)
    set_notes(s, "The approach: validated import, a two-stage CP-SAT engine, an independent validator and a complete web application around it.")

    # 13 architecture
    s = S[12]
    put_picture(s, DIAG / "01_system_architecture.png", (215, 118, 790, 482), remove_ids=(4,), name="System architecture diagram")
    set_paras(shape(s, 5), "Proposed architecture: React frontend, FastAPI backend with services and engine, SQLite storage")
    set_notes(s, "Browser to React to FastAPI. The engine and the independent validator sit in the service layer; everything is stored in SQLite.")

    # 14 expected outcomes
    s = S[13]
    for tid, bid, ttl, body in [
        (7, 8, "Conflict-free plans", "Zero same-paper neighbours and roll-gap violations (target)."),
        (12, 13, "Seconds, not days", "A plan for hundreds of candidates in a few seconds."),
        (17, 18, "Unpredictable, fair", "Seeded random choice; roll number says nothing about the seat."),
        (22, 23, "Reproducible", "Same data and seed give the same plan; Verify proves it."),
        (27, 28, "Usable system", "Charts, QR slips, attendance, lookup, analytics."),
        (32, 33, "Honest evaluation", "Benchmark against manual methods; gaps reported plainly.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Expected outcomes are stated as targets at this stage and are measured in the later reviews.")

    # 15 references
    s = S[14]
    fill(s, 4, [f"[{i}] {r}" for i, r in enumerate(REFERENCES[:13], 1)], 14, min_pt=8.5, bullet=True, gap_pt=3)
    set_notes(s, "References [1] to [13]; the full list is in the report.")

    # 16 thank you
    set_notes(S[15], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


BASE_PAPER_TEXT = ('P. Yukth, I. S. Mahmood, K. H. Shashank, S. G. Thejus and M. S. Darshini, "Exam Seating Arrangement System," '
                   'International Journal of Research Publication and Reviews (IJRPR), vol. 6, no. 12, pp. 2093-2105, Dec. 2025.')

if __name__ == "__main__":
    build()
