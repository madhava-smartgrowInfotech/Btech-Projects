"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import (BATCH, BENCH, DIAG, EXP, GUIDE, GUIDE_DESIG, COLLEGE, YEAR, DEPT, PRODUCT, SHOTS, STUDENTS, TITLE)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"


def px(v):
    return Emu(int(v * 9525))


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, "Batch A13", "exam seating, CP-SAT, optimisation"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A13"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), TITLE)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A13 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers design, methodology and progress of SeatWise.")

    # 2 contents
    s = S[1]
    for sid, t in [(21, "Engine and configuration"), (18, "Data preparation"), (27, "Data flow, modules, database"), (36, "Preliminary results and screens")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts, from a recap of Review 1 through the design and the engine to progress, results, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Seating by hand or in spreadsheets is slow, error-prone and predictable: neighbours write the same paper, several rules clash, and candidates cannot look up a seat.", 16)
    fill(s, 13, "A two-stage OR-Tools CP-SAT engine and an independent validator, served by a FastAPI backend and a React web application.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Excel / CSV data", "Stage A: halls + colour classes", "Stage B: seats per hall", "Validator + scorecard", "Publish, charts, QR slips"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1: the problem and the proposed system as a five-step pipeline.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 18 + TypeScript + Vite (port 5113), role guards, seat map with dnd-kit, TanStack Query.",
                "Application: FastAPI (port 8113): JWT sign-in, administrator and invigilator roles, rate-limited public endpoints.",
                "Engine: OR-Tools CP-SAT Stage A and Stage B, an independent validator and a swap checker.",
                "Data: SQLite with 16 tables and an append-only audit trail.",
                "Outputs: ReportLab PDFs, openpyxl Excel files, qrcode slips."], 16, bullet=True)
    set_notes(s, "The browser talks only to the Vite server, which proxies /api to FastAPI. The backend owns every rule; the frontend never decides whether a move is legal.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Import files", "Validate rows", "Pre-check capacity", "Stage A: halls + classes", "Stage B: seats per hall", "Validate + scorecard", "Publish + outputs"]):
        fill(s, sid, t, 14)
    set_paras(shape(s, 25), "Two separate stages")
    fill(s, 26, "Stage A decides each paper's hall and colour class; Stage B only places seats inside that class. Each stage can be tested alone.", 15)
    set_paras(shape(s, 28), "One rule set")
    fill(s, 29, "The engine, the swap checker and the validator share the neighbour-graph and roll-gap code, so map, exports and benchmark agree.", 15)
    set_paras(shape(s, 31), "Reproducible")
    fill(s, 32, "Sub-seeds derive from the plan seed with SHA-256; inputs, result and seating are fingerprinted; Verify re-runs the engine.", 15)
    set_notes(s, "End-to-end workflow in seven steps. The key idea: confine each paper to one colour class of the seat grid, so same-paper neighbours are impossible by construction.")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    ch = replace_chart(s, 4, ["Mon 12 Oct AM", "Mon 12 Oct PM", "Tue 13 Oct AM", "Tue 13 Oct PM", "Wed 14 Oct AM", "Wed 14 Oct PM"], [("Candidates", [451, 457, 451, 457, 511, 397])])
    style_chart(ch, "Sample dataset: candidates per sitting", ["3F3FBF"], legend=False, label_pt=10, axis_title="candidates")
    rows = [["Dataset", "Size", "Use"],
            ["Sample workbook (synthetic, seed 2026)", "49 courses, 908 candidates, 14 halls (871 seats), 6 sittings", "Demo, import test"],
            ["Registrations", "2,724; 10 candidates need an accessible seat", "Subject and access rules"],
            ["Benchmark scenarios (seeded)", "150 to 2,000 candidates, 8 scenarios", "Engine evaluation"],
            ["Import templates", "4 empty files + 1 workbook", "Real data entry"],
            ["External data", "None; no downloads, no API keys", "Fully offline"]]
    sz = table_size(rows, 530, 400, [0.4, 0.38, 0.22], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.4, 0.38, 0.22], header_size=sz)
    set_notes(s, "All data is synthetic and generated by a seeded script, so the same seed always gives the same files. There is no external dataset and no train/test split because nothing is trained.")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation and validation")
    set_paras(shape(s, 7), "Import parsing")
    fill(s, 8, "CSV or XLSX with flexible headers (Roll No, roll_number and candidate_id are all accepted); lists split on ; , or |; dates in two formats; combined workbook or single files.", 16)
    set_paras(shape(s, 12), "Validation")
    fill(s, 13, "Row-by-row checks: required columns, duplicates, unknown departments or courses, seat labels outside the layout, aisle positions, time order, split paper groups and timetable clashes. Nothing is saved until the administrator imports.", 16)
    set_paras(shape(s, 17), "Engine inputs")
    fill(s, 18, "Candidates become (paper, department, accessible need) groups; courses that share a question paper count as one paper. Each hall becomes a seat graph; blocked seats and aisles remove neighbours. Halls without accessible seats get 2 by default.", 16)
    set_notes(s, "Preprocessing in this project means parsing and validating the four data files, then turning them into the engine's plain data types.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Model architecture: the seating engine")
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (75, 124, 693, 518), remove_ids=(4,), name="Engine algorithm flowchart")
    rows = [["Component", "Detail"], ["Stage A (CP-SAT)", "groups to halls and colour classes, 2 passes"], ["Stage B (CP-SAT)", "x[candidate, seat] per hall, in parallel"],
            ["Colour classes", "4 (8 neighbours) or 2 (4 neighbours)"], ["Validator", "plain Python, no solver; scorecard"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "No neural network or trained parameters: the model is a constraint model solved by OR-Tools.", 11)
    set_notes(s, "The flowchart is the real algorithm. Stage A confines each paper to one colour class per hall; Stage B picks seats inside the class with a seeded random objective.")

    # 9 configuration
    s = S[8]
    set_paras(shape(s, 2), "Engine configuration (nothing is trained)")
    rows = [["Setting", "Value"], ["Solver", "OR-Tools CP-SAT 9.15.6755"], ["Per-hall budget", "0.15 deterministic seconds (tuned)"],
            ["Stage A budget", "0.3 (pass 2: half)"], ["Workers per hall", "1 (deterministic); halls in parallel"], ["Retries", "up to 5 re-balancing rounds"],
            ["Random weights", "0-19 per seat, 0-999 per class"], ["Hardware", "12-thread Intel desktop, Python 3.11.9"]]
    set_table(shape(s, 10), rows, 13, header_size=13)
    put_picture(s, EXP / "budget_tuning.png", (672, 232, 550, 330), remove_ids=(5,), name="Budget tuning chart")
    set_paras(shape(s, 7), "Budget tuning: 0.15 is the fastest budget within 10% of the best department mix", 11)
    set_notes(s, "The only setting that matters for speed is the per-hall budget. The benchmark tried five values on 600 and 1,000 candidates and chose 0.15: 2.63 seconds with 123 same-department pairs; 0.6 gives 114 pairs but takes 6.58 seconds.")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (57, 124, 420, 520), remove_ids=(4,), name="Decision flowchart")
    for bid in (7, 10, 13):
        shape(s, bid).height = px(80)
    set_paras(shape(s, 6), "Valid plan")
    fill(s, 7, "A plan is hard-valid when same-paper pairs, roll-gap violations, accessible misses and capacity violations are all 0 and everyone is seated.", 16)
    set_paras(shape(s, 9), "Impossible input")
    fill(s, 10, "The pre-check names the cause and suggests a fix. A failing hall is retried at 4x budget, then re-balanced (max 5).", 16)
    set_paras(shape(s, 12), "Saving")
    fill(s, 13, "Each plan stores seed, rules, engine version and three SHA-256 fingerprints; publishing archives the previous version; every action is audited.", 16)
    set_notes(s, "Decision logic: the engine either returns a plan that satisfies every hard rule or explains why none exists.")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 330), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 330), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Generate seating plan' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(452)
    rows = [["Module", "Responsibility"], ["Importer", "parse CSV/XLSX, validate row by row, preview, commit; re-validates against current data"],
            ["Engine and validator", "pre-check, Stage A, Stage B in parallel, independent scorecard, roll-order baseline"],
            ["Outputs and attendance", "PDF charts and sheets, Excel lists, QR slips, tablet register with lock"]]
    set_table(shape(s, 12), rows, 12, widths=[0.2, 0.8], header_size=12, row_h=34)
    shape(s, 12).top = px(485)
    set_notes(s, "Level 1 shows eight processes and five data stores; level 2 opens the plan-generation process.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 560, 340), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (632, 118, 589, 340), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 16 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(460), px(1164)
    rows = [["Table", "Key fields", "Relations"],
            ["plans", "session, version, status, seed, rules, three hashes, scorecard", "1:N seat_assignments, attendance"],
            ["seat_assignments", "plan, candidate, hall, row, col, seat_label", "unique per plan+candidate and plan+seat"],
            ["candidates", "roll_no, department, needs_accessible_seat", "N:M courses through registrations"]]
    set_table(shape(s, 13), rows, 11, widths=[0.2, 0.44, 0.36], header_size=11, row_h=30, total_px=1164)
    t = shape(s, 13)
    t.left, t.top, t.width = px(57), px(492), px(1164)
    set_notes(s, "Sixteen tables created by SQLAlchemy. The UML diagram shows the domain entities and the engine's data types and services.")

    # 13 technology
    s = S[12]
    cards = [(7, 8, "Language", "Python 3.11, TypeScript 5.8"), (12, 13, "Optimisation", "Google OR-Tools CP-SAT 9.15, NumPy, pandas"),
             (17, 18, "Frontend", "React 18, Vite 6, Tailwind CSS, shadcn/ui, Recharts, dnd-kit"), (22, 23, "Backend", "FastAPI, Pydantic v2, SQLAlchemy 2, PyJWT, bcrypt"),
             (27, 28, "Storage and outputs", "SQLite, ReportLab (PDF), openpyxl (Excel), qrcode"), (32, 33, "Testing", "pytest (58 tests), strict TypeScript check, Playwright captures")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Everything runs offline: no API keys and no external services.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Validated import (CSV/XLSX, four kinds)", "Completed"], ["Two-stage CP-SAT engine + independent validator", "Completed"],
            ["Plans: generate, verify, publish, versions", "Completed"], ["Seat map with drag-to-swap and live rule checks", "Completed"],
            ["Exports: charts, hall lists, sheets, QR slips, attendance report", "Completed"], ["Tablet attendance, public lookup, analytics, audit trail", "Completed"],
            ["Automated tests (58 passed) and engine benchmark", "Completed"], ["User acceptance testing with real users", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.78, 0.22], header_size=13)
    for vid, lid, v, lab in [(6, 7, "57", "API operations"), (9, 10, "16", "database tables"), (12, 13, "58/58", "backend tests passing")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Every planned feature is implemented. The one open item is user-acceptance testing with real invigilators and examination staff.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: engine benchmark")
    keys = ["xs-150", "s-300", "m-600", "l-1000", "xl-2000"]
    ch = replace_chart(s, 4, ["150", "300", "600", "1,000", "2,000"], [("SeatWise solve time (s)", [BENCH[k]["SeatWise"][0] for k in keys])], number_format="0.00")
    style_chart(ch, "Solve time by sitting size (seconds, mean of 3 seeds)", ["3F3FBF"], number_format="0.00", legend=False, label_pt=11, axis_title="seconds")
    for vid, lid, v, lab in [(6, 7, "24 / 24", "plans meeting every hard rule"), (9, 10, "2.4 s", "to seat 600 candidates (4.8 s for 2,000)"), (12, 13, "0", "same-paper neighbours (roll order: 1,253)")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Benchmark run engine-benchmark-20260921: 8 scenarios, 3 seeds each. Every SeatWise plan met every hard rule; the maximum solve time was 6.1 seconds (4-neighbour mode).")

    # 16 output screens
    s = S[15]
    put_picture(s, SHOTS / "09_plan.png", (91, 128, 507, 476), remove_ids=(4,), name="Plan page screenshot")
    put_picture(s, SHOTS / "11b_swap_refused.png", (679, 128, 512, 476), remove_ids=(7,), name="Seat map screenshot")
    set_paras(shape(s, 6), "Plan page: 0 same-paper neighbours, 0 roll-gap violations, 1,099 conflicts avoided", 11)
    set_paras(shape(s, 9), "Seat map: an illegal move is refused with the reason", 11)
    set_notes(s, "Real captures of the running application with the sample data imported and three plans published.")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Authentication, roles, team", "8", "Passed"], ["Engine and rules", "14", "Passed"], ["Import and validation", "9", "Passed"],
            ["Plans, swap, verify", "6", "Passed"], ["Exports, attendance, public, analytics", "21", "Passed"], ["Total", "58", "58 passed"]]
    set_table(shape(s, 18), rows, 13, widths=[0.6, 0.15, 0.25], header_size=13)
    fill(s, 5, ["58 backend tests passed on a throw-away database; the slowest takes 7.1 s.",
                "The same seed reproduces the plan exactly (hashes equal); different seeds give different plans.",
                "Illegal swaps are refused with reasons; impossible inputs are explained; the public lookup returns 429 after 30 requests a minute.",
                "Demo scenario: import, three plans (1.8-2.2 s, 0 same-paper pairs), exports, lookup and attendance in 11.4 s.",
                "Not done: user-acceptance testing and load testing."], 16, bullet=True)
    set_notes(s, "58 backend tests plus a scripted end-to-end demo. The frontend is checked by a strict TypeScript compile, a production build and walkthroughs in Chrome.")

    # 18 challenges
    s = S[17]
    cards = [(7, 8, "Rules by construction", "Each paper is confined to one colour class per hall, so same-paper neighbours are impossible before seats are chosen."),
             (12, 13, "Speed", "Halls solve in parallel; the 0.15 budget was tuned. 2.4 s at 600 candidates."),
             (17, 18, "Reproducibility", "Deterministic budget, one worker per hall and SHA-256 sub-seeds give the same plan from the same seed."),
             (22, 23, "Impossible inputs", "A pre-check explains the cause; a failing hall is retried at 4x budget, then re-balanced."),
             (27, 28, "Front-row bias", "The construction gives a 4.7-point front-row bias (random shuffle 1.6); recorded as a limitation."),
             (32, 33, "Public data", "Lookup shows published seats only, a shortened name, rate-limited; no e-mail or date of birth.")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Six challenges and how each was handled.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["User acceptance testing", "Anonymised real-sitting pilot", "Final report", "Final presentation and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Collect feedback from invigilators and the examination cell on the seat map, attendance screen and lookup page (not yet done).",
                 "Run the engine on a real, anonymised sitting and compare it with the cell's manual plan.",
                 "Reduce the front-row bias and speed up 4-neighbour mode (6.0 s at 600 candidates).",
                 "Document every diagram, test and result in the report."], 18, bullet=True)
    set_notes(s, "What remains: real-user feedback and a pilot with real data, then the report and the final demo.")

    set_notes(S[19], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
