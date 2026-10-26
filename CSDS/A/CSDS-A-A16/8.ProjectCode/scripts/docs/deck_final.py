"""16.Final_PPT.pptx: 28 slides. Built from D:\\Abstracts\\Templates\\Final_PPT.pptx (rule, title and footer chrome inherited;
the template's college letterhead picture is replaced by a plain '[COLLEGE NAME]' placeholder)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from . import facts as F
from . import lit
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT) and sh.shape_type != MSO_SHAPE_TYPE.PICTURE


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "a16_deck_crops")


def crop_shot(name, box=(256, 22, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def p1(x):
    return f"{x * 100:.1f}%"


def p2(x):
    return f"{x * 100:.2f}%"


def mstd(vals, d=2):
    import statistics
    return f"{statistics.mean(vals) * 100:.{d}f} ± {statistics.stdev(vals) * 100:.{d}f}%"


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE (Data Science)"
    cp.comments = ""
    cp.last_modified_by = cp.author
    cp.subject = F.PRODUCT
    cp.keywords = "facial emotion recognition, compound emotions, cEmoGAN, multi-label CNN, FER2013"
    notes = {}

    N, B, G, SY, DT = F.FIN, F.BASE, F.GAN, F.SYN, F.DET
    cmp_ = N["compound_on_real_test"]
    ra, rf = F.abl_vals("real_only", "accuracy"), F.abl_vals("real_only", "macro_f1")
    sa, sf = F.abl_vals("synth", "accuracy"), F.abl_vals("synth", "macro_f1")
    s05 = SY["final_model"]["0.5"]
    ctrl = SY.get("real_only_control", {})
    lat = {l["label"]: l for l in F.APIM["latency"]}
    an_anon = lat["POST /api/analyze anonymous (face)"]
    an_auth = lat["POST /api/analyze authenticated (face, persisted)"]
    login = lat["POST /api/auth/login (bcrypt)"]
    T = F.TESTS

    # ---- chrome: drop the other college's letterhead picture everywhere; footers
    for sl in S:
        for sh_ in list(sl.shapes):
            if sh_.shape_type == MSO_SHAPE_TYPE.PICTURE and sh_.top == TOP_BANNER:
                K.drop_shape(sl, sh_)
            elif sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE (Data Science)  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 26)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE (Data Science)")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and Mosaic: basic and compound emotion estimation from a face, built on a CNN multi-label classifier and a conditional GAN, with an honest evaluation."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Model\nDecision Workflow, Interaction & Database Design\nImplementation Highlights\nResults: Model Performance, Confusion Matrix, Ablation\nParameter Analysis: Compound Threshold & GAN Condition\nPerformance, Testing & Comparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Most facial-emotion systems name one basic emotion, yet faces often show two at once (happily surprised, sadly fearful) and labelled mixed-emotion data is scarce. Mosaic detects a face, runs a 64x64 CNN with 18 sigmoid outputs (7 basic + 11 compound emotions) and serves the result through a FastAPI service and a Next.js console with accounts and history. A conditional GAN synthesises compound-expression faces to train the compound outputs.", {"bullet": True, "space_after": 7}),
        (f"Measured on the real FER2013 test set ({F.DATA['n_test']:,} faces): the single-label baseline scores {p2(B['accuracy'])}, the shipped multi-label model {p2(N['accuracy'])} (macro-F1 {N['macro_f1']:.3f}); a real-data-only control trained the same number of epochs scores {mstd(ra, 1)} versus {mstd(sa, 1)} with the synthetic faces, so the gain comes from the extra training, not from the GAN.", {"bullet": True, "space_after": 7}),
        (f"Honest finding: the compound outputs stay near zero on real faces (mean top compound probability {cmp_['mean_top_compound_confidence']:.4f}) and rank synthetic compound classes at chance (ROC-AUC {s05['compound_roc_auc_macro']:.2f}); the 10-epoch GAN reads as its intended emotion on only {p1(G['mean_basic_conditional_accuracy'])} of faces (chance 14.3%). No compound ground truth exists in FER2013.", {"bullet": True, "space_after": 7}),
        (f"{T['tests']} pytest tests pass; the API answers an analysis in about {an_anon['p50_ms']:.0f} ms on a CPU.", {"bullet": True}),
    ], size=15)
    notes[3] = "Product first, then the headline honest results: the multi-label model beats the baseline only through extra epochs; the compound head is not working; the GAN is weak."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: complex emotion estimation using analysis-by-synthesis"],
            ["Method", "cEmoGANs: an encoder maps a face to an identity-preserving latent, a conditional GAN with a 6-d multi-hot label (anger, disgust, fear, happiness, sadness, surprise) generates mixed expressions; a CNN with sigmoid outputs is trained on the generated images"],
            ["Data", "CK+, JAFFE and MUG; mixed-emotion ground truth made by morphing two peak-expression images; 15,000 generated images split 7:3"],
            ["Reported", "f-beta 0.99 / 0.99 / 0.80 and exact-match ratio 100% / 99% / 68.7% (CK+ / JAFFE / MUG), against EfficientNet-B0 0.45 / 0.60 / 0.51; identity cosine similarity above 0.7 for 98% / 94% / 99% of generated faces"],
            ["Scored on", "Test split of the generated images; real unconstrained photos only shown qualitatively (Fig. 11)"],
            ["Stated future work", "Demographic diversity, better backbones and interpretability, multimodal cues"]]
    ftable(s, 0.55, 2.6, [1.9, 10.3], rows, size=12, row_h=0.7)
    notes[4] = "Khine et al., IEEE Access 2025. The recogniser is scored on images from its own generator, which matters when comparing numbers."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    idx = [0, 1, 3, 5, 6, 7, 8, 9]
    rows = [["Work", "Year", "Approach", "Finding / limitation"]] + [list(lit.LIT_TABLE[i]) for i in idx]
    ftable(s, 0.55, 1.75, [2.7, 0.7, 4.6, 4.2], rows, size=10, row_h=0.55, align=["l", "c", "l", "l"])
    ftext(s, 0.55, 6.75, 12.2, 0.3, "Two further batch-survey items (Ni et al. 2024; Zhou et al. 2024) are known by title and venue only; no result is quoted from them.", size=10, color=K.FNAVY)
    notes[5] = "Only facts we could read are quoted: the base paper and archived PDFs, and abstract-level statements recorded in the reference notes."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [(g_, {"bullet": True, "space_after": 18}) for g_ in lit.GAPS], size=21)
    notes[6] = "The batch's gap, plus two sharper points: evaluation on generated data only, and the missing no-synthesis control."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.5)
    ftext(s, 0.85, 1.9, 11.6, 2.4, lit.PROBLEM_STATEMENT, size=17, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "One face photo (upload or webcam frame), optionally with a signed-in account"), ("Output", "Scores for 7 basic and 11 compound emotions, the dominant emotion, a face thumbnail and a stored history"), ("Need", lit.NEED)]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.6, 3.9, 2.0, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.68, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.15, 3.5, 1.4, b_, size=14)
    notes[7] = "The batch's problem statement and need, then the input and output of the system."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 16}) for i, o in enumerate(lit.OBJECTIVES, 1)], size=22)
    notes[8] = "The five objectives from the batch document; slide 25 gives the status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Per-image pipeline", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Upload / webcam", K.FNAVY), ("Haar face detection", K.FBLUE), ("64x64 grayscale", K.GREEN), ("CNN, 18 sigmoids", K.ORANGE), ("Rank + dominant", K.PURPLE), ("Result + history", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.35, 6.0, 0.4, "Training (offline)", size=15, color=K.FNAVY, bold=True)
    ftext(s, 0.55, 3.8, 6.0, 3.0, [
        ("Stage 1: 7-class backbone CNN on FER2013 (28,709 faces)", {"bullet": True}),
        ("Stage 2: conditional GAN on 2,800 faces, 10 epochs", {"bullet": True}),
        ("Stage 3: 2,200 synthetic compound faces (11 pairs x 200) added; 18-output classifier with BCE", {"bullet": True}),
    ], size=14, space_after=6)
    ftext(s, 6.9, 3.35, 5.9, 0.4, "Application", size=15, color=K.FNAVY, bold=True)
    ftext(s, 6.9, 3.8, 5.9, 3.0, [
        ("FastAPI: register, login, analyze, history, health (JWT, bcrypt)", {"bullet": True}),
        ("SQLite stores each analysis as JSON; only the cropped face is saved", {"bullet": True}),
        ("Next.js console: upload or webcam, bar chart, compound bars, history", {"bullet": True}),
    ], size=14, space_after=6)
    notes[9] = "Two flows: the offline training in three stages and the online per-image pipeline."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 3.0, 1.6, 7.3, 4.8)
    ftext(s, 0.6, 6.45, 12.1, 0.4, "Four tiers: browser, FastAPI service, ML modules (PyTorch + OpenCV), local files; no external service at run time", size=14, color=K.FNAVY, align="c")
    notes[10] = "One process on one CPU machine; the generator is used only at training time."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("02_high_level_workflow.png"), 1.6, 1.6, 10.1, 4.8)
    ftext(s, 0.6, 6.45, 12.1, 0.4, "Three offline stages, then one online pass per image", size=14, color=K.FNAVY, align="c")
    notes[11] = "Analysis by synthesis: generate compound faces, train a multi-label recogniser on real plus generated data."

    # ---- 12 dataset, features and model
    s = S[11]
    set_title(s, "Dataset, Preprocessing & Model")
    wipe(s)
    ftext(s, 0.55, 1.75, 6.3, 5.1, [
        (f"FER2013 (Hugging Face mirror, apache-2.0 tag): {F.DATA['n_train']:,} train / {F.DATA['n_test']:,} test faces, 48x48 grayscale, 7 basic labels; no compound labels", {"bullet": True, "space_after": 6}),
        ("Disgust is 1.5% of the training set; the baseline uses class-weighted cross-entropy", {"bullet": True, "space_after": 6}),
        ("Preprocessing: Haar crop, resize 64x64, scale to [-1, 1]", {"bullet": True, "space_after": 6}),
        (f"EmotionCNN: 4 conv blocks (32-64-128-128, BatchNorm, ReLU, MaxPool), FC 2048-256, dropout 0.4, out 18; {N['parameters']:,} parameters", {"bullet": True, "space_after": 6}),
        (f"Generator: 5 transposed convs, noise 100 + condition 7, {G['generator_parameters']:,} parameters; compound condition = 0.5 / 0.5 on the two emotions", {"bullet": True, "space_after": 6}),
        ("Adam, batch 256, 5 epochs per CNN; CPU only", {"bullet": True}),
    ], size=14)
    K.picture(s, F.fig("fig_class_distribution.png"), 7.0, 1.85, 5.9, 3.0)
    ftext(s, 7.0, 5.0, 5.8, 1.6, "Compound taxonomy: 11 two-emotion categories from the compound-expression work of Du, Tao and Martinez, each labelled with its two basic emotions plus the compound bit.", size=13, color=K.FNAVY)
    notes[12] = "The dataset has only basic labels; compound targets exist only on the synthetic faces."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_core_algorithm.png"), 3.9, 1.6, 5.5, 4.8)
    ftext(s, 0.6, 6.45, 12.1, 0.4, "analyze_image_bytes(): detect face, crop, 64x64, CNN, sigmoid, rank; dominant = argmax of all 18", size=14, color=K.FNAVY, align="c")
    notes[13] = "No face found returns 200 with face_detected false; the dominant emotion is the argmax over all 18 outputs, so it is basic unless a compound output beats every basic one."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User–System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_main_flow.png"), 3.2, 1.6, 6.9, 4.8)
    ftext(s, 0.6, 6.45, 12.1, 0.4, "Sequence of one analysis through the router, Haar detector, CNN and SQLite", size=14, color=K.FNAVY, align="c")
    notes[14] = "The alternative flows at the bottom of the diagram list anonymous calls, no face, missing model and bad uploads."

    # ---- 15 database
    s = S[14]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 1.2, 1.6, 10.9, 4.8)
    ftext(s, 0.6, 6.45, 12.1, 0.4, "Entity-relationship diagram: users 1 : N analyses; emotion lists stored as JSON", size=14, color=K.FNAVY, align="c")
    notes[15] = "Two tables. The uploaded photo is never stored; only the cropped grayscale face is written to the thumbnails folder."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    rows = [["Area", "What was built"],
            ["Backend", "FastAPI, SQLAlchemy/SQLite, JWT (HS256, 7 days), bcrypt, pydantic schemas, CORS for localhost:3000; 8 endpoints"],
            ["ML", "PyTorch CPU: EmotionCNN (7 or 18 outputs), conditional GAN, Haar detector, 3-stage pipeline (run_pipeline.py)"],
            ["Frontend", "Next.js 16 / React 19 / TypeScript: landing page, login, register, analyze (upload or webcam), history; Recharts"],
            ["Data", "FER2013 pickles (28,709 + 7,178); 2,200 synthetic compound faces generated per training run"],
            ["Quality", f"{T['tests']} pytest tests (units, API, privacy); 24 abuse probes; production build used for the screens"],
            ["Added in this run", "classifier_final.pt trained (it was missing); experiments, tests and API measurements added under scripts/ and backend/tests/"]]
    ftable(s, 0.55, 1.85, [2.4, 9.8], rows, size=13, row_h=0.62)
    notes[16] = "Be clear that the final checkpoint did not exist when the batch was received."

    # ---- 17 model performance
    s = S[16]
    set_title(s, "Results – Model Performance (FER2013 test, 7,178 faces)")
    wipe(s)
    rows = [["Model", "Accuracy", "Macro-F1"],
            ["Single-label baseline (shipped stage-1 CNN, 5 epochs)", p2(B["accuracy"]), f"{B['macro_f1']:.3f}"],
            ["Final multi-label model (shipped, 5 + 5 epochs, argmax of 7 basic outputs)", p2(N["accuracy"]), f"{N['macro_f1']:.3f}"],
            [f"Control: real data only, same recipe (mean of {len(ra)} seeds)", mstd(ra), mstd(rf)],
            [f"Real + 200 synthetic per compound class (mean of {len(sa)} seeds)", mstd(sa), mstd(sf)]]
    ftable(s, 0.55, 1.8, [7.3, 2.5, 2.4], rows, size=13, row_h=0.46, align=["l", "c", "c"])
    K.picture(s, F.fig("fig_ablation.png"), 0.55, 4.25, 6.4, 2.7)
    ftext(s, 7.2, 4.3, 5.6, 2.6, [
        ("The 7.6-point gain over the baseline is explained by five extra epochs: the real-only control reaches the same level", {"bullet": True, "space_after": 6}),
        ("Synthetic faces change accuracy by less than the seed-to-seed spread and lower macro-F1", {"bullet": True, "space_after": 6}),
        (f"Always predicting 'happy' scores {p1(F.DATA['test_majority_baseline'])}", {"bullet": True}),
    ], size=14)
    notes[17] = "Mean and sample standard deviation over three seeds; the shipped checkpoint is a single run."

    # ---- 18 confusion matrix
    s = S[17]
    set_title(s, "Confusion Matrix & Per-Class F1")
    wipe(s)
    K.picture(s, F.fig("fig_confusion_final.png"), 0.55, 1.8, 4.9, 4.4)
    K.picture(s, F.fig("fig_per_class_f1.png"), 5.7, 1.8, 7.1, 3.6)
    ftext(s, 0.55, 6.3, 4.9, 0.6, "Final model, rows = true class, columns = predicted", size=12, color=K.FNAVY)
    pc_ = N["per_class"]
    ftext(s, 5.7, 5.5, 7.1, 1.4, [
        (f"Strongest: happy F1 {pc_['happy']['f1']:.2f}, surprise {pc_['surprise']['f1']:.2f}; weakest: disgust {pc_['disgust']['f1']:.2f} (precision 1.00, recall {pc_['disgust']['recall']:.2f}), fear {pc_['fear']['f1']:.2f}", {"bullet": True, "space_after": 4}),
        ("Fear, sad and neutral are confused with each other, the usual FER2013 pattern", {"bullet": True}),
    ], size=13)
    notes[18] = "The disgust class has only 111 test faces and 436 training faces."

    # ---- 19 parameter analysis
    s = S[18]
    set_title(s, "Parameter Analysis – Compound Threshold & GAN Condition")
    wipe(s)
    K.picture(s, F.fig("fig_compound_threshold.png"), 0.45, 1.75, 6.2, 3.1)
    K.picture(s, F.fig("fig_gan_fidelity.png"), 6.75, 1.75, 6.2, 3.1)
    ftext(s, 0.55, 5.0, 12.2, 1.95, [
        (f"Threshold sweep 0.05 to 0.5 on 1,100 fresh synthetic compound faces: compound recall {SY['final_model']['0.05']['compound_recall'] * 100:.0f}% at 0.05 (everything fires, ROC-AUC {s05['compound_roc_auc_macro']:.2f}) and {s05['compound_recall'] * 100:.0f}% from 0.2 up; no threshold separates the compound classes", {"bullet": True, "space_after": 5}),
        (f"Real-only control scores the same AUC ({ctrl.get('compound_roc_auc_macro', 0.5):.2f}); on real faces no compound output reaches 0.3 (mean top {cmp_['mean_top_compound_confidence']:.4f})", {"bullet": True, "space_after": 5}),
        (f"Generator test: a classifier reads the intended emotion on {p1(G['mean_basic_conditional_accuracy'])} of generated faces (chance 14.3%); compound pairs: top-2 equals the pair {p1(G['mean_compound_top2_equals_pair'])} (chance 4.8%)", {"bullet": True}),
    ], size=13)
    notes[19] = "The decision rule is argmax over all 18 outputs, so the compound head never wins; a lower threshold only makes every compound fire together."

    # ---- 20 performance and testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    rows = [["Measure", "Result"],
            ["API latency (127.0.0.1, CPU)", f"Analyze p50 {an_anon['p50_ms']:.0f} ms, p95 {an_anon['p95_ms']:.0f} ms (signed-in with storage p50 {an_auth['p50_ms']:.0f} ms); login (bcrypt) p50 {login['p50_ms']:.0f} ms"],
            ["Automated tests", f"{T['tests']} pytest tests pass in {T['seconds']:.1f} s: taxonomy, tokens, shapes, preprocessing, synthesis labels, auth, analyze, history privacy"],
            ["Abuse probes", f"{F.APIM['probes_ok']}/{F.APIM['probes_total']} behave as expected; text, empty and truncated uploads end in HTTP 500 (no handler); unlisted CORS origins refused"],
            ["Face detection", f"Haar finds 0% of native 48x48 faces and {p1(DT['haar_detection_rate_scale_3x'])} after 3x upscaling; on those faces accuracy is {p1(DT['accuracy_via_haar_crop_on_detected'])} through the product path vs {p1(DT['accuracy_direct_64px_on_same_images'])} with a direct resize"],
            ["Training cost", "Final classifier: 10 min 37 s for 5 epochs on a shared CPU (about 2 min per epoch); no GPU used"]]
    ftable(s, 0.55, 1.85, [2.8, 9.4], rows, size=13, row_h=0.72)
    notes[20] = "Latency numbers come from scripts/measure_api.py; the Haar detection gap is a training/serving mismatch."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different data and scoring)")
    wipe(s)
    rows = [["Aspect", "Khine et al. (base paper)", "Mosaic (this project)"],
            ["Training data", "CK+, JAFFE, MUG; cEmoGAN-generated compound faces (15,000)", f"FER2013 real faces ({F.DATA['n_train']:,}) + 2,200 synthetic compound faces"],
            ["Generator", "Encoder (identity-preserving latent) + conditional GAN", "Conditional GAN only; no identity encoder"],
            ["Scored on", "Split of its own generated images", f"Real FER2013 test faces, basic labels only ({F.DATA['n_test']:,})"],
            ["Headline", "f-beta 0.99 / 0.99 / 0.80; exact match 100% / 99% / 68.7%", f"Accuracy {p2(N['accuracy'])}, macro-F1 {N['macro_f1']:.3f}; compound ROC-AUC {s05['compound_roc_auc_macro']:.2f} (chance)"],
            ["Compound result on real photos", "Qualitative examples only", f"Compound outputs near zero (mean top {cmp_['mean_top_compound_confidence']:.4f})"]]
    ftable(s, 0.55, 1.75, [2.4, 4.9, 4.9], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 5.65, 12.2, 1.3, [
        ("Not comparable: different datasets, different test images (generated vs real), different label sets. The base paper's scores say its recogniser reads its own generator's output; ours say a small CNN reaches about 58% on real FER2013.", {"bullet": True, "space_after": 4}),
        (f"Closest independent reference: a lightweight CNN on FER-2013 reaches 51.6% [{R['galkaduwa']}]; our baseline {p1(B['accuracy'])}, final {p1(N['accuracy'])} (splits not guaranteed identical).", {"bullet": True}),
    ], size=13, color=K.FNAVY)
    notes[21] = "State the 'not comparable' point first; the numbers are context."

    # ---- 22-24 output screens
    cap = "real FER2013 test faces, production build"
    s = S[21]
    set_title(s, "Output – Analyze a Face")
    wipe(s)
    K.picture(s, crop_shot("08_result_happy.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("12_result_neutral_misread.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, f"Left: a happy face is read as Happy (47%) with Surprise close behind; the compound bars all show 0% ({cap}). Right: a neutral face read as Fear (24%), one of the misses behind the 58% accuracy. The thumbnail beside the dominant emotion shows a broken image (known defect).", size=12, color=K.FNAVY)
    notes[22] = "Real captures of the running application."
    s = S[22]
    set_title(s, "Output – Generated Face and No-Face Case")
    wipe(s)
    K.picture(s, crop_shot("14_result_gan_happily_surprised.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("16_result_no_face.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: a face from the project's own cEmoGAN conditioned on 'happily surprised': a blurry blob, read as Sad (53%), compound scores about 1%. Right: a real 48x48 FER2013 face as uploaded is not found by the Haar detector: the API answers with face_detected false.", size=12, color=K.FNAVY)
    notes[23] = "Shows the generator quality and the detection limit honestly."
    s = S[23]
    set_title(s, "Output – History")
    wipe(s)
    K.picture(s, crop_shot("17_history.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("18_history_detail.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: stored analyses of the signed-in user, newest first, with dominant emotion and confidence (delete on hover). Right: a saved result opened by id; the route exists but nothing in the UI links to it, so it was opened by URL.", size=12, color=K.FNAVY)
    notes[24] = "History works per user; the detail page is unreachable from the list."

    # ---- 25 conclusion and objectives
    s = S[24]
    set_title(s, "Conclusion & Objectives vs Outcomes")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 1.1, [
        (f"Built and tested end to end: Haar detection, an 18-output CNN, accounts, history and a console; {p2(N['accuracy'])} on real FER2013 against a {p2(B['accuracy'])} baseline.", {"bullet": True, "space_after": 3}),
        ("The compound capability is not demonstrated: its outputs are inert on real faces and at chance on synthetic ones, and the GAN does not follow its condition.", {"bullet": True}),
    ], size=13)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Identity-preserving compound faces with cEmoGANs", "Not achieved: no identity encoder; generator ignores its condition"],
            ["2", "Enlarge the complex-emotion dataset", f"Done mechanically (2,200 faces); no accuracy or compound benefit measured"],
            ["3", "CNN multi-label classifier, basic + compound", f"Basic: {p2(N['accuracy'])}; compound outputs not functional"],
            ["4", "Evaluate against single-label baselines", f"Done: baseline {p2(B['accuracy'])}; real-only control {mstd(ra, 1)}"],
            ["5", "Interface for uploaded images", f"Achieved; {T['tests']} tests; known defects listed in the report"]]
    ftable(s, 0.55, 3.0, [0.5, 4.6, 7.1], rows, size=12, row_h=0.55, align=["c", "l", "l"])
    notes[25] = "Be explicit about what was and was not achieved."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("Train the GAN to convergence on a larger, identity-labelled set and add the base paper's encoder so that identity is preserved", {"bullet": True, "space_after": 10}),
        ("Obtain real compound labels (for example RAF-DB compound subset, licence permitting) to evaluate and fine-tune the compound outputs", {"bullet": True, "space_after": 10}),
        ("Weight or oversample synthetic faces, add a dedicated compound head and report exact-match ratio and per-class F1", {"bullet": True, "space_after": 10}),
        ("Train on Haar-style crops, or use a stronger detector, to close the training/serving gap", {"bullet": True, "space_after": 10}),
        ("Fix the product defects: error handling for bad uploads, thumbnail URL, link to the detail page, session after reload, e-mail validation", {"bullet": True}),
    ], size=18)
    notes[26] = "Each item answers a measured weakness."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "cdgan", "nstgan", "galkaduwa", "abaw7", "du2014", "fer2013", "gan", "cgan", "viola"]
    ftext(s, 0.55, 1.75, 12.2, 5.2, [(f"[{R[k_]}]  {clean_ref(REFS[R[k_] - 1])}", {"space_after": 5}) for k_ in keys], size=12)
    notes[27] = "Numbering follows the report's reference list."

    # ---- 28 thank you
    s = S[27]
    for sh_ in s.shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, "THANK YOU")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        if i in notes:
            K.set_notes(sl, notes[i])

    # ---- college placeholder in the letterhead position (added last so wipe() cannot remove it)
    for sl in S:
        ftext(sl, 2.07, 0.22, 9.2, 0.5, "[COLLEGE NAME]", size=26, color=K.FNAVY, bold=True, align="c")
        ftext(sl, 2.07, 0.68, 9.2, 0.3, f"{F.DEPT}", size=12, color=K.FNAVY, align="c")

    prs.save(out_path)
    K.rename_theme(out_path, "Mosaic")
    return out_path


if __name__ == "__main__":
    import sys

    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(F.BATCH, "16.Final_PPT.pptx")
    print(build(out))
