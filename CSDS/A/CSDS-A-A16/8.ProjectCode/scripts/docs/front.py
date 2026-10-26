"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations (Mosaic, CSDS-A-A16)."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures); None if unavailable."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    par = d.doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.line_spacing = 1.0
    if b:
        par.add_run().add_picture(b, width=Inches(width))
    else:  # plain text banner when the letterhead image cannot be extracted
        add_runs(par, F.COLLEGE.upper(), 14, True)


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p(F.DEPT.upper(), align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering (Data Science), "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering (Data Science), is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the related work cited in this report, the creators and maintainers of the FER2013 dataset and of its Hugging Face "
        "mirror used for training and evaluation, and the maintainers of the open-source software used in this project (Python, PyTorch, OpenCV, FastAPI, SQLAlchemy, scikit-learn, "
        "NumPy, matplotlib, Next.js, React, Recharts, pytest and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    N, B, S, G = F.FIN, F.BASE, F.SYN["final_model"]["0.5"], F.GAN
    c = N["compound_on_real_test"]
    d.h1("ABSTRACT")
    d.p("Most facial emotion recognition systems assign one basic emotion (happy, sad, angry, fear, disgust, surprise or neutral) to a face, although real expressions often mix "
        "several emotions, such as happiness with surprise. Labelled data for such compound expressions is scarce, and expression-editing methods that could create it tend to distort the person's identity. "
        f"This project builds {F.PRODUCT}, a working web application that follows the analysis-by-synthesis approach of Khine et al. (IEEE Access, 2025): a conditional GAN (cEmoGAN) is trained to synthesise "
        "faces, eleven compound-emotion classes are synthesised by conditioning the generator on two emotions at once, and a convolutional neural network with 18 sigmoid outputs (7 basic and 11 compound labels) "
        "is trained on real FER2013 faces plus the synthetic ones. A FastAPI service detects the largest face with an OpenCV Haar cascade, runs the network, stores history for signed-in users in SQLite, "
        "and serves a Next.js console that shows a basic-emotion chart, compound-emotion scores and the dominant emotion.")
    d.p(f"Measured on the {F.DATA['n_test']:,}-image FER2013 test set, the multi-label network reaches {N['accuracy'] * 100:.2f} % basic-emotion accuracy (macro-F1 {N['macro_f1']:.3f}) against "
        f"{B['accuracy'] * 100:.2f} % (macro-F1 {B['macro_f1']:.3f}) for the single-label backbone it was initialised from; a seeded ablation in Chapter 8 shows how much of that difference comes from the synthetic images. "
        f"The compound outputs, however, are not functional in the shipped model: on real test faces their largest probability averages {c['mean_top_compound_confidence']:.4f}, so the dominant emotion is never a compound label, "
        f"and on {F.SYN['n']:,} freshly generated compound faces their ROC-AUC is {S['compound_roc_auc_macro']:.2f}, which is chance. The generator itself does not obey its condition (a classifier recovers the intended basic emotion on only "
        f"{G['mean_basic_conditional_accuracy'] * 100:.1f} % of its faces, chance being 14.3 %), and FER2013 contains no compound ground truth, so compound recognition could not be validated on real data. "
        f"The base paper's scores were measured on its own generated images and are not comparable. {F.TESTS['tests']} automated tests pass and the API answers an analysis in about {F.APIM['latency'][3]['p50_ms']:.0f} ms on a CPU-only computer. "
        "The report documents these gaps and the product defects found, and proposes concrete next steps.")
    d.p("**Keywords:** complex emotion recognition, compound facial expressions, conditional GAN, multi-label CNN, FER2013, FastAPI, Next.js.")


ABBR = [
    ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("BCE", "Binary Cross-Entropy"), ("BGR", "Blue-Green-Red colour order (OpenCV)"),
    ("CK+", "Extended Cohn-Kanade facial expression dataset"), ("CNN", "Convolutional Neural Network"), ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"),
    ("cEmoGAN", "Conditioned Emotion Generative Adversarial Network"), ("CSS", "Cascading Style Sheets"), ("DFD", "Data Flow Diagram"), ("EMR", "Exact Match Ratio"),
    ("ER", "Entity-Relationship"), ("FER", "Facial Expression Recognition"), ("FER2013", "Facial Expression Recognition 2013 dataset"), ("FID", "Fréchet Inception Distance"),
    ("GAN", "Generative Adversarial Network"), ("GPU", "Graphics Processing Unit"), ("HOG", "Histogram of Oriented Gradients"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JAFFE", "Japanese Female Facial Expression dataset"), ("JPEG", "Joint Photographic Experts Group image format"), ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"),
    ("MUG", "Multimedia Understanding Group facial expression dataset"), ("ORM", "Object-Relational Mapping"), ("RAF-DB", "Real-world Affective Faces Database"),
    ("ReLU", "Rectified Linear Unit"), ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("SQL", "Structured Query Language"),
    ("SVM", "Support Vector Machine"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("ViT", "Vision Transformer"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
