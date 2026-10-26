"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides) for CropSight.

    python scripts/build_final.py
Every number comes from docs/REPORT_FACTS.md and experiments/*/metrics.json.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import BATCH, CHARTS, COLLEGE, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, REFERENCES, SHOTS, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE  |  Batch A17  |  Final Review"


def px(v):
    return Emu(int(v * 9525))


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 8)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182, cap_top=546):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l, 13)
    set_paras(shape(s, cap_ids[1]), cap_r, 13)
    for cid in cap_ids:
        c = shape(s, cid)
        c.top = px(cap_top)
        c.height = px(100)


def diagram_with_panel(s, img, points, pic_box, panel_box, size=16):
    """Large diagram on the left, key points (the template's caption shape, moved and left-aligned) on the right."""
    put_picture(s, img, pic_box, remove_ids=(7,), name=Path(img).stem)
    cap = shape(s, 8)
    cap.left, cap.top, cap.width, cap.height = (px(v) for v in panel_box)
    fill(s, 8, ["\u2022  " + p for p in points], size, min_pt=11, gap_pt=6)
    for p in cap.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.italic = False
            r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    cap.text_frame.vertical_anchor = MSO_ANCHOR.TOP


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author = f"{PRODUCT}: Final Review", TITLE, "Batch A17"
    cp.keywords = "crop quality grading, price forecasting, delivery planning, IoT, explainable AI"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A17"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    # 1 title
    s = S[0]
    set_paras(shape(s, 4), TITLE)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A17")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of CropSight: a farmer-to-buyer platform that grades crop quality from a photo, forecasts the price, plans delivery and monitors the shipment.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Pipeline",
                "Design & Implementation", "Results, Real-Photo Check, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["Farmers face unpredictable prices, manual and inconsistent quality grading, unmonitored storage and transport, and little decision support after harvest.",
                "CropSight is a web platform for farmers, buyers, distributors and admins: a crop photo is graded A/B/C/Reject from 11 computer-vision features by a Random Forest and explained with SHAP.",
                "A gradient-boosted model forecasts a 31-day price with a 15-85 % band and a best day to sell; a weighted score ranks buyers and routes; a simulated IoT stream monitors temperature, humidity and shock.",
                "Measured: grader 92.3 % on its synthetic test set; price model R\u00b2 0.992 (synthetic series); 36 tests pass; upload + grade 174 ms median.",
                "Honest finding: both models are trained on synthetic data. On 3,200 real fruit photos the shipped grader is at chance (50.4 %), while the same features retrained on real labels reach 90.6 %.",
                "Keywords: crop quality grading, price forecasting, delivery planning, IoT, explainable AI."], 17, bullet=True)
    set_notes(s, "Abstract: problem, approach, headline measurements and the honest real-photo result.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [REFERENCES[0],
                "Task: pre-harvest smart agriculture - irrigation advice, soil analysis, crop and fertiliser recommendation from soil, weather and crop data, with IoT sensors, a weather API and a chatbot.",
                "Method: Transformer-based Tabular Learning (irrigation), TabNet regressor and classifier (soil, fertiliser), Sparse Weighted Fusion Transformer (crop); XAI explains soil and fertiliser predictions.",
                "Reported results: irrigation 99.13 %; soil RMSE 1.51, R\u00b2 98.7 %; crop 98.75 %; fertiliser 99.3 %.",
                "Our reading: strong tabular results, but everything happens before harvest. We adopt its IoT + deep learning + XAI idea and apply it after harvest to quality grading, price and delivery."], 17, bullet=True)
    set_notes(s, "The base paper is read from the PDF in Documentary/. CropSight differs in task, data and models, so the numbers cannot be compared directly.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Khaliq et al., 2025 (base) [1]", "TTL, TabNet, SwiFT + XAI, IoT, weather API", "Soil, weather, fertiliser, crop data (tabular)", "Irrigation 99.13 %; soil R\u00b2 98.7 %; crop 98.75 %; fertiliser 99.3 %", "Pre-harvest only; no grading, price or delivery"],
            ["Wang and Zhang, 2026 [6]", "17 forecasting approaches incl. zero-shot time-series foundation models", "USDA ERS monthly commodity prices, 1997-2025", "Beat USDA futures-based forecasts on 3 of 4 commodities; Time-MoE +54.9 % on wheat", "US commodity prices; no farmer-level decisions"],
            ["Feng and Zong, 2026 [7]", "ConvLSTM + coordinate attention + ASPP; LCMST-Net", "Large-scale crop images (not examined in detail)", "Accuracy 0.979 (ConvLSTM), 0.982 (LCMST-Net)", "Recognises crop types; no quality grade, price or delivery"],
            ["Leaf-disease multi-task [2]; quality-inspection grading [3]", "Image-based diagnosis / grading", "Not retrieved", "Not retrieved", "Listed in batch survey only"],
            ["Crop price prediction, 2023 [4] and 2024 [5]", "Machine-learning price prediction", "Not retrieved", "Not retrieved", "Listed in batch survey only"]]
    table(s, rows, [0.18, 0.23, 0.19, 0.23, 0.17], 16, 440)
    set_notes(s, "[1], [6] and [7] were read from their PDFs. [2]-[5] could not be retrieved, so no results are claimed for them.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Pre-harvest decisions only: soil, irrigation, crop and fertiliser. No post-harvest quality grading, price forecast or delivery planning"],
            ["Price forecasting [4], [5], [6]", "Predict price only; not linked to the crop's quality grade, to buyers or to transport and spoilage"],
            ["Image-based work [2], [3], [7]", "Recognise crops, diseases or grade from images; no price forecast, buyer ranking or cold-chain monitoring"],
            ["Existing practice (batch survey)", "Manual grading is subjective; middlemen set mandi prices; portals such as e-NAM publish prices without forecasts; storage and transport are rarely monitored"],
            ["All reviewed work", "No system joins grading, price forecast, buyer and route recommendation, shipment monitoring and a marketplace with explanations - and none is shown to work on real farm data. CropSight prototypes the joined system"]]
    table(s, rows, [0.26, 0.74], 18, 440)
    set_notes(s, "The gap is the joined post-harvest workflow; the prototype does not yet close the evidence gap because its models use synthetic data.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Crop quality is graded manually and inconsistently, which affects the price a farmer receives.",
                "Farmers cannot predict prices or choose the best time to sell.",
                "Storage and transport conditions are not monitored, causing post-harvest losses.",
                "Middlemen and limited market access reduce farmer income.",
                "Need: an integrated platform that grades quality, forecasts prices, recommends buyers and routes, monitors the supply chain and connects farmers to buyers - with explanations for every decision."], 20, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["Grade crop quality automatically from images - done with classical ML (Random Forest on 11 CV features), not deep learning.",
                "Forecast crop prices from market, weather and demand data - done with gradient boosting on a synthetic series.",
                "Monitor storage and transport with IoT sensors - done as a simulated WebSocket stream (no hardware).",
                "Explain grading and pricing decisions with XAI - SHAP top-6 factors on every grade and forecast.",
                "Recommend buyers, routes and the best time to sell - weighted score and best day within a 14-day window.",
                "Connect farmers, distributors and buyers - role-gated marketplace with an order state machine."], 19, bullet=True)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["Upload - farmer adds a crop photo (JPEG/PNG/WEBP, max 8 MB), crop, quantity and region.",
                "Grade - OpenCV + scikit-image extract 11 features; Random Forest returns A/B/C/Reject with probabilities and SHAP.",
                "Forecast - gradient boosting gives today's price, a 15-85 % band, 31 days and the best day to sell.",
                "Plan delivery - buyers ranked by price, distance, quality fit, reliability and spoilage risk.",
                "Publish and order - listing goes to the marketplace; buyer orders; farmer confirms.",
                "Transit and monitor - distributor starts transit; a shipment is created and a sensor stream is shown live.",
                "Oversee - admin KPIs and a model-insights page that discloses the synthetic training data."], 18, bullet=False)

    # 10 architecture
    s = S[9]
    diagram_with_panel(s, DIAG / "01_system_architecture.png",
                       ["React 19 + Vite web app: farmer, buyer, distributor and admin screens",
                        "FastAPI backend: 7 routers, 21 endpoints (20 REST + 1 WebSocket), JWT + bcrypt, role guards",
                        "Models: Random Forest grader + SHAP; Gradient Boosting price model with quantile bands; weighted delivery score",
                        "Storage: SQLite (9 tables), uploaded photos, .joblib model files",
                        "Sensor stream: simulated, one reading every 1.4 s over WebSocket"],
                       (43, 178, 640, 494), (700, 200, 540, 440), 17)

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "End-to-end workflow: upload, grade, forecast, plan, publish, order, confirm, transit, monitor, deliver", 14)

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Features & Models")
    fill(s, 7, ["Grader data (synthetic): 6,000 feature rows, 1,500 per grade, 11 features, 6 % of labels moved to a neighbouring grade; 80/20 stratified split (4,800 / 1,200).",
                "Price data (synthetic): 23,360 rows = 8 crops x 8 regions x every 3rd day, 2022-2024; INR 402-19,176 per quintal; 85/15 split (19,856 / 3,504).",
                "Real photos used only for checking: 3,200 fresh/rotten fruit photos (CC BY 4.0) and tomato, onion and potato photos for the demo screens.",
                "Features: hue, saturation, value, uniformity, GLCM texture (3), edge density, blemish ratio, size score; price: season, rainfall, temperature, demand, supply, fuel, days to harvest, grade, region, crop.",
                "Models: Random Forest (250 trees, depth 9); Gradient Boosting (300 trees, depth 4) + 2 quantile models (alpha 0.15 / 0.85); SHAP TreeExplainer.",
                "Delivery score = 0.38 price + 0.22 distance + 0.18 quality fit + 0.14 reliability + 0.08 (1 - spoilage risk); not trained."], 16, bullet=True)
    set_notes(s, "No real crop-image or market-price dataset was available; both models are trained on synthetic data and the app says so in its admin page.")

    # 13 decision workflow
    s = S[12]
    diagram_with_panel(s, DIAG / "10_algorithm_flowchart.png",
                       ["Grading: validate type and size, centre on a 256 px canvas, foreground mask (whole image if mask < 3 % or > 98 %), 11 features, Random Forest, SHAP",
                        "Delivery: for every buyer compute haversine distance, offered price, five normalised terms and a weighted score x100; return the top 5",
                        "Best sell day = highest forecast price within days 0-14 (spoilage window)"],
                       (43, 178, 560, 494), (630, 210, 610, 420), 17)

    # 14 user-system interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    diagram_with_panel(s, DIAG / "09_sequence_diagram.png",
                       ["Farmer uploads a photo; the API checks the JWT and role, saves the file and a draft listing, then grades and stores the result",
                        "Forecast and delivery plan are requested from the listing; publishing sets the status to listed",
                        "Buyer orders (listing becomes sold); farmer confirms; distributor starts transit and a shipment with a WebSocket stream is created"],
                       (43, 178, 560, 494), (630, 210, 610, 420), 17)

    # 15 database
    s = S[14]
    diagram_with_panel(s, DIAG / "06_er_diagram.png",
                       ["9 SQLite tables: users, crop_listings, quality_grades, towns, price_predictions, shipments, sensor_readings, orders, delivery_recommendations",
                        "One grade per listing; orders link a listing, a buyer and a distributor",
                        "Shipments hold the sensor readings; recommendations are stored per listing",
                        "price_predictions exists in the schema but no route writes to it"],
                       (43, 178, 460, 494), (540, 210, 700, 420), 17)

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["Grading", "OpenCV HSV mask (Otsu), GLCM texture, Canny edges, blemish and size features; Random Forest 250 trees; SHAP top 6"],
            ["Price", "Gradient Boosting 300 trees + two quantile models; 31-day context walk; best day = argmax over days 0-14"],
            ["Delivery", "Weighted score of price, distance (haversine), quality fit, reliability, spoilage; top 5 buyers and a route"],
            ["IoT", "Seeded random-walk sensor simulator with injected events; WebSocket every 1.4 s; readings stored per shipment"],
            ["Marketplace", "Order state machine placed - confirmed - in transit - delivered; role checks on every transition"],
            ["Web app + security", "React 19, TypeScript, Tailwind, Recharts; bcrypt, JWT 7 days, role guards; 21 endpoints; 36 backend tests"]]
    table(s, rows, [0.18, 0.82], 18, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Measured Performance")
    rows = [["Measure", "Result", "Note"],
            ["Grader accuracy, synthetic test (1,200 rows)", "92.3 %  (macro F1 0.924)", "5 seeds 92.9 \u00b1 0.3 %; every error is between neighbouring grades"],
            ["Price model MAE / R\u00b2 / MAPE (3,504 rows)", "INR 164.6 / 0.992 / 5.08 %", "noise floor MAPE 2.8 %; 2024 hold-out R\u00b2 0.988"],
            ["15-85 % price band coverage", "69.3 %", "nominal 70 %; mean width 23.1 % of price"],
            ["Shipped grader on 3,200 real fruit photos", "50.4 % (AUC 0.57)", "chance = 50 %; 99 % of photos graded B"],
            ["Same 11 features retrained on real labels", "90.6 % (AUC 0.97)", "5-fold CV; unseen fruit types 66.9 %"],
            ["Delivery top-1 stability", "97.7 %", "weights x0.5-1.5, 256 cases, worst case 74 %"],
            ["IoT simulator in alarm", "91.6 % of steps", "all 50 shipments by step 400 (9.3 min)"],
            ["Automated tests", "36 passed, 3 expected-fail", "the 3 document real gaps"]]
    table(s, rows, [0.38, 0.24, 0.38], 17, 410)
    set_paras(shape(s, 8), "Grader and price figures are on synthetic test sets (they measure recovery of the generator, not field performance); real-photo rows are the independent check.", 13)
    set_notes(s, "Both models are synthetic-trained. The real-photo rows show the grader does not transfer; the feature extractor itself does carry signal.")

    # 18 confusion matrix + real photo check
    s = S[17]
    set_paras(shape(s, 4), "Confusion Matrix & Real-Photo Check")
    two_pics(s, CHARTS / "01_quality_confusion_matrix.png", CHARTS / "05_real_photos_grade_share.png",
             "Synthetic test set: 1,108 of 1,200 correct (92.3 %); all 92 errors are between neighbouring grades",
             "3,200 real photos: 99.8 % of fresh and 99.1 % of rotten fruit are graded B")
    set_notes(s, "Left: confusion matrix on the shipped model's synthetic test split. Right: the shipped model on real photographs of fresh and rotten fruit.")

    # 19 parameter / ablation analysis
    s = S[18]
    set_paras(shape(s, 4), "Baselines, Seeds & Real-Label Retraining")
    two_pics(s, CHARTS / "03_quality_baselines_seeds.png", CHARTS / "08_real_label_retrain.png",
             "Same split: logistic regression 93.1 % and k-NN 93.0 % match the forest (92.3 %); five seeds 92.9 ± 0.3 %",
             "Fresh vs rotten: shipped model 50.4 %, same features trained on real labels 90.6 %, unseen fruits 66.9 %",
             box_h=330)
    set_notes(s, "Baselines show the synthetic classes are close to linearly separable (dropping any one feature group costs at most 1.2 points). Retraining on real labels shows the features carry signal while the shipped classifier does not transfer.")

    # 20 performance & testing
    s = S[19]
    rows = [["Measurement", "Result"],
            ["API response, median / 95th percentile", "health 1.6 / 4.7 ms; marketplace 17.7 / 36.5; price predict 54.9 / 119.4; delivery 99.8 / 147.9; upload + grade 173.9 / 207.3; login (bcrypt) 325.3 / 349.0 ms"],
            ["In-process model time", "price forecast 13.1 ms median (p95 29.9); delivery ranking 64.7 ms (p95 107.1)"],
            ["Automated tests", "36 passed, 3 expected-fail, 0 failed in 30 s: auth, role guards, grading, price, delivery, order lifecycle, IoT, admin"],
            ["Defects found", "corrupt image -> HTTP 500 and an orphan draft listing; unvalidated crop, region and order quantity/price; admin volume overstated 100x (INR per quintal x kg); unauthenticated WebSocket"],
            ["Machine", "Intel Core i5-12500, 15.7 GB RAM, no GPU; both models train in about 20 s"],
            ["Not measured", "concurrent load, memory, other hardware or browsers, user acceptance, real sensors"]]
    table(s, rows, [0.22, 0.78], 17, 430)
    set_notes(s, "Latencies are from a live server on 127.0.0.1 while other jobs ran on the same machine.")

    # 21 comparison
    s = S[20]
    rows = [["System", "Task and data", "Method", "Reported / measured", "Explainability and extras", "Comparable?"],
            ["Base paper [1]", "Pre-harvest irrigation, soil, crop, fertiliser advice; tabular soil/weather/crop data", "TTL, TabNet, SwiFT; IoT sensors; weather API", "Irrigation 99.13 %; soil RMSE 1.51, R\u00b2 98.7 %; crop 98.75 %; fertiliser 99.3 %", "XAI, chatbot", "-"],
            ["Wang and Zhang [6]", "Monthly US commodity prices 1997-2025", "Zero-shot time-series foundation models", "Beat USDA futures-based forecasts on 3 of 4 commodities", "-", "No: price only"],
            ["CropSight", "Post-harvest grading, price, delivery; synthetic features and price series; 3,200 real photos for checking", "Random Forest (11 CV features), Gradient Boosting + quantiles, weighted score, simulated sensors", "Grader 92.3 % synthetic, 50.4 % on real photos; price R\u00b2 0.992 synthetic", "SHAP on every grade and forecast; marketplace; live monitor", "No: different task and data"]]
    table(s, rows, [0.11, 0.2, 0.2, 0.22, 0.14, 0.13], 15, 400)
    set_paras(shape(s, 8), "Different tasks, data and metrics: no numeric head-to-head is possible. The base paper's figures are on tabular pre-harvest data; ours are partly on synthetic data.", 13)
    set_notes(s, "State plainly: no claim that CropSight beats the base paper; the comparison is of scope and approach.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Grading and Price Forecast")
    two_pics(s, SHOTS / "07_grading_result.png", SHOTS / "08_price_forecast.png",
             "Real tomato photo graded B (70.1 %): radar of features and SHAP drivers", "Forecast INR 2,272 per quintal with band, 30-day curve and drivers")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Delivery Planner and Marketplace")
    two_pics(s, SHOTS / "09_delivery_planner.png", SHOTS / "13_marketplace.png",
             "Four buyers ranked (top: 67.7 % match) with a suggested route", "Buyer marketplace: graded listings with photo, grade and price")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Live Monitoring and Model Insights")
    two_pics(s, SHOTS / "22c_live_alert_later.png", SHOTS / "24_model_insights.png",
             "Simulated shipment stream: alert banner when readings leave the safe band", "Admin model insights: 92.3 % accuracy, R\u00b2 0.992 and the synthetic-data disclosure")

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built CropSight: a role-based platform that grades a crop photo with explanations, forecasts price, ranks buyers and routes, streams shipment sensors and runs an order lifecycle.",
                "Measured: grader 92.3 % and price R\u00b2 0.992 on synthetic test sets; 36 tests pass; upload + grade 174 ms median.",
                "Key finding: both models are trained on synthetic data; on 3,200 real fruit photos the shipped grader is at chance (50.4 %), while the same features retrained on real labels reach 90.6 %.",
                "Limits: classical ML rather than deep learning; simulated sensors whose alarm drifts into permanent alarm; hash-based market context; defects found (100x admin volume, unvalidated inputs).",
                "Objectives are met as a working prototype, not as field-validated performance."], 20, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Collect real labelled crop images and retrain the grader (the features already reach 90.6 % on real labels); consider a CNN.",
                "Replace the synthetic price series by real mandi data and a true time-series forecaster; validate on a time-based split.",
                "Connect real IoT sensors; add mean reversion and event-based alarms; authenticate the WebSocket.",
                "Fix the defects: unit error in admin volume, input validation, corrupt-image handling, authorisation gaps, secrets.",
                "Real buyer quotes and road routing; payments; deployment with a hosted database, HTTPS and load testing; a user study."], 20, bullet=True)

    # 27 references
    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], 11, min_pt=6.5, bullet=True, gap_pt=1.5)

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
