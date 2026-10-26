"""Build 19.IEEE_Paper.docx: a six-page, two-column IEEE conference paper on CropSight.

    python scripts/make_paper_figures.py ; python scripts/build_paper.py ; python scripts/paper_finish.py

Every number comes from docs/REPORT_FACTS.md and experiments/*/metrics.json. References are numbered in order of first citation
from doc_common.REFERENCES; the four batch-survey papers that could not be retrieved (entries 2-5) are not cited here.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paper_kit as K  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from doc_common import BATCH, GUIDE, REFERENCES, ROOT, STUDENTS  # noqa: E402

OUT = BATCH / "19.IEEE_Paper.docx"
PAPER = ROOT / "docs" / "paper"

KEYS = {"base": 0, "wang": 5, "feng": 6, "sultana": 7, "vegds": 8, "breiman": 9, "friedman": 10, "koenker": 11, "shap": 12, "shaptree": 13,
        "haralick": 14, "canny": 15, "otsu": 16, "sklearn": 17, "skimage": 18, "opencv": 19, "lime": 20, "tabnet": 21}
ORDER: list[str] = []


def C(text: str) -> str:
    """Replace [@a,b] by [n], [m] numbered in order of first citation."""
    def sub(m):
        out = []
        for k in m.group(1).split(","):
            if k not in ORDER:
                ORDER.append(k)
            out.append(f"[{ORDER.index(k) + 1}]")
        return ", ".join(out)
    return re.sub(r"\[@([a-z,]+)\]", sub, text)


class Paper(K.Paper):
    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(K.TITLE)
        K.font(r, 24)
        p.paragraph_format.space_after = Pt(6)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [(GUIDE, "[DESIGNATION], Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Dept. of CSE (Data Science)", False), ("[COLLEGE NAME]", False), ("[CITY]", False), ("[EMAIL]", False)):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                K.font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        sectpr = sec._sectPr
        cols = sectpr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectpr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def references(self):
        self.h1("References")
        for i, key in enumerate(ORDER, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]\t{REFERENCES[KEYS[key]]}")
            K.font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


K.TITLE = "CropSight: AI-Based Crop Quality Grading, Price Forecasting and Delivery Planning on a Farmer-Buyer Marketplace"
K.ABSTRACT = (
    "Farmers sell perishable crops with little objective grading, no price outlook and no view of what happens to the produce in transit. "
    "This paper presents CropSight, a web platform that connects farmers, buyers and distributors. A photograph is graded A, B, C or "
    "Reject by a Random Forest on 11 colour, texture, edge, blemish and size features, with SHAP explanations; a gradient-boosting "
    "model gives a price, a 70% band and a 30-day curve with a best day to sell; a weighted score ranks buyers; and a simulated "
    "sensor stream raises cold-chain alerts. Both models are trained on synthetic data because no real labelled set was available. "
    "On their own test sets the grader reaches 92.3% accuracy and the price model an R2 of 0.992, but a check on 3,200 real fruit "
    "photographs shows chance-level transfer (50.4% accuracy, AUC 0.57), whereas the same features retrained on real labels reach "
    "90.6%. The synthetic scores therefore measure the data generator, not the farm. Thirty-six automated tests pass and three known "
    "gaps are documented."
)
K.KEYWORDS = "crop quality grading, delivery planning, explainable AI, IoT monitoring, price forecasting, smart agriculture, synthetic data"


def build():
    P = Paper()
    P.front()
    P.abstract()
    para = lambda t, **k: P.para(C(t), **k)

    # ---------------------------------------------------------------- I
    P.h1("I. Introduction")
    para("Agriculture supports millions of farmers whose income depends on what a buyer will pay on the day the crop is sold. Three "
         "things decide that price: the quality of the produce, the state of the market, and whether the produce reaches the buyer "
         "in good condition. Today quality is judged by eye by a trader, so two traders may grade the same lot differently; farmers "
         "usually sell through middlemen at the prevailing market price, and government portals publish prices without a forecast; "
         "and storage and transport conditions are rarely monitored, so damage is discovered only on arrival. The farmer cannot "
         "show objective evidence of quality, cannot decide when to sell, and carries the loss.")
    para("Smart-agriculture research has mostly addressed the period before harvest. The base paper [@base] integrates IoT sensors, "
         "deep learning and explainable AI for soil analysis, irrigation advice and crop-fertiliser recommendation, and reports "
         "accuracies near 99%, but it stops at the farm gate. We found no system that combines post-harvest quality grading, price "
         "forecasting, delivery planning and a marketplace that connects the farmer to buyers and distributors.")
    para("This paper presents CropSight, a full-stack platform for exactly that gap, and reports honestly what it does and does not "
         "show. Its contributions are:", after=2)
    P.numbered([
        "a photograph-to-grade pipeline of real computer-vision features (colour, texture, edges, blemishes, size) and a Random Forest grader with per-grade SHAP explanations;",
        "a price model with a point estimate, a 15-85% quantile band, a 30-day curve and a best-day-to-sell rule, and a transparent weighted score that ranks buyers by price, distance, quality fit, reliability and spoilage risk;",
        "a role-gated farmer-buyer-distributor marketplace with an order state machine and a live, simulated cold-chain sensor stream with alerts;",
        "an evaluation that separates what the synthetic training data can show from what real photographs show, with 36 passing tests and the defects it exposed.",
    ])
    para("Section II reviews related work, Section III describes the method, Section IV reports the results and Section V concludes.", before=2)
    P.h2("A. System Hardware")
    para("All measurements were taken on one desktop with a 12th-generation Intel Core i5-12500 (6 cores, 12 threads), 15.7 GB of memory "
         "and no discrete graphics card (Intel UHD 770), running Windows 11 Pro. Nothing was trained on a GPU: both models fit in "
         "about twenty seconds on the CPU.")
    P.h2("B. Software Requirements")
    para("The back end uses Python 3.11.9, FastAPI 0.115.6, SQLAlchemy 2.0.36 over SQLite, JSON Web Tokens with bcrypt hashes, "
         "scikit-learn 1.5.2 [@sklearn], SHAP 0.46.0, OpenCV 4.10 [@opencv] and scikit-image 0.24 [@skimage]. The front end is React 19 "
         "with TypeScript, Vite 8, Tailwind CSS 4, Recharts and Three.js, run on Node 24.19.0. Tests use pytest 9.1.1; screenshots were "
         "captured with Playwright driving Chrome.")

    # ---------------------------------------------------------------- II
    P.h1("II. Literature Survey")
    para("**Base paper.** Khaliq et al. [@base] build an AI-driven smart-agriculture system for soil analysis, irrigation advice, "
         "crop recommendation and fertiliser recommendation. IoT sensors supply soil nutrient, moisture and temperature readings that are "
         "combined with historical weather and crop-yield records, cleaned, normalised and encoded. A Transformer-based Tabular "
         "Learning (TTL) model advises irrigation (accuracy 99.13%), a TabNet [@tabnet] regressor analyses the soil (RMSE 1.51, R2 98.7%), "
         "a Sparse Weighted Fusion Transformer (SwiFT) recommends crops (98.75%) and a TabNet classifier recommends fertiliser "
         "(99.3%). Explainable-AI methods, LIME [@lime] among them, explain the soil and fertiliser predictions; a weather API supports "
         "irrigation timing and a chatbot eases use by farmers. The models were developed in Python on a Windows 10 machine with an "
         "Intel Core i5 and 8 GB of memory. The paper contributes an integrated, explainable, sensor-fed pipeline, and names real-time "
         "adaptive control and generalisation across regions as future work. It does not address produce after harvest: there is no "
         "quality grading of harvested crops, no market price, no buyer or route recommendation and no transport monitoring. CropSight "
         "adopts its IoT, learning and explanation architecture and applies it to that later stage.")
    para("**Price forecasting.** Wang and Zhang [@wang] evaluate 17 approaches in four classes, including five time-series foundation "
         "models, on monthly USDA ERS commodity prices from 1997 to 2025. Zero-shot foundation models beat traditional, machine-learning "
         "and deep-learning models trained from scratch, and Time-MoE improved on the USDA benchmark by 54.9% for wheat and 18.5% for "
         "corn on 2017-2024 data. Their series are national monthly prices; they do not address grade, region or delivery.")
    para("**Image-based recognition and data.** Feng and Zong [@feng] improve a ConvLSTM with coordinate attention and an ASPP module "
         "and classify the extracted spatiotemporal features with LCMST-Net, reporting a mean classification accuracy of 0.982 for crop "
         "recognition; the task is identifying the crop, not grading its quality. Sultana et al. [@sultana] published a dataset of fresh "
         "and rotten fruit photographs in eight fruit types, which we use to test our grader on real images.")
    para("**Methods used.** The grader is a Random Forest [@breiman] on grey-level co-occurrence texture statistics [@haralick], Canny "
         "edges [@canny] and an Otsu-threshold foreground mask [@otsu]; the price model is gradient boosting [@friedman] with quantile "
         "loss [@koenker]; explanations use Shapley values for trees [@shap,shaptree].")
    P.table([["Work", "Approach", "Data", "Reported result", "Limitation"],
             [C("[@base]"), "TTL, TabNet, SwiFT + IoT + XAI", "Soil, weather, crop, fertiliser data", "99.13% irrigation; 98.75% crop; 99.3% fertiliser", "Pre-harvest only"],
             [C("[@wang]"), "17 methods incl. 5 foundation models", "USDA ERS monthly prices, 1997-2025", "Time-MoE +54.9% wheat, +18.5% corn vs USDA", "National series; no grade or delivery"],
             [C("[@feng]"), "ConvLSTM + attention + LCMST-Net", "Crop images", "Accuracy 0.982", "Crop type, not quality"],
             [C("[@sultana]"), "Dataset paper", "Fresh and rotten fruit photos, 8 fruits", "Dataset", "Binary freshness only"],
             [C("[@breiman]"), "Random Forest", "-", "Method reference", "-"],
             [C("[@friedman,koenker]"), "Gradient boosting; quantile loss", "-", "Method reference", "-"],
             [C("[@shaptree]"), "Tree SHAP explanations", "-", "Method reference", "-"]],
            "Related work", widths=[0.13, 0.26, 0.22, 0.24, 0.15], size=6.3)
    para("**Reading across the work.** Three points shape the design. First, the base paper [@base] shows that tabular learning, "
         "sensing and explanation can be combined for farm decisions, but its tasks are decisions about inputs, not about produce "
         "already harvested. Second, the foundation-model study [@wang] shows that price forecasting depends on long, real, "
         "monthly histories; a synthetic series cannot test that claim, so we state our price results as recovery of a generator "
         "and not as forecasting skill. Third, recognition accuracy such as 0.982 [@feng] concerns which crop an image shows; grading "
         "needs labels for quality, which is why a labelled freshness dataset [@sultana] is the right test of a grader and why we use it.")
    para("**Gap.** Further grading and price-prediction papers appear in our batch survey, but we could not retrieve their text and "
         "state no result for them. Among the work we read, none joins post-harvest photo grading, a price outlook with a "
         "best day to sell, buyer and route ranking, transport monitoring and a marketplace in one explainable system, and none "
         "reports how a model trained on synthetic data behaves on real photographs. CropSight addresses this combination and "
         "measures the transfer problem.")

    # ---------------------------------------------------------------- III
    P.h1("III. Methodology")
    P.h2("A. Architecture and Main Flow")
    para("CropSight has a React web app and a FastAPI back end with seven routers (authentication, crops, prices, delivery, IoT, "
         "marketplace, admin), 21 endpoints (20 REST and one WebSocket), nine SQLite tables and local model files (Fig. 1). There is no "
         "external service. Four roles exist: farmer, buyer, distributor and admin; the admin role cannot self-register. Fig. 2 shows "
         "the flow: a farmer uploads a photo and receives a grade, forecasts the price, plans delivery and publishes the listing; a buyer "
         "orders; the farmer confirms; a distributor starts transit and watches the sensor stream; and the order is delivered.")
    P.figure(PAPER / "fig1_architecture.png", "System architecture of CropSight.", width=3.4)
    P.figure(PAPER / "fig2_flow.png", "Main flow from photo to delivery.", width=3.4)
    P.h2("B. Data")
    para("No labelled crop-image set and no market-price history were available, so both models are trained on synthetic data, "
         "and the application states this on its admin screen. The grading data are 6,000 feature vectors (1,500 per grade) drawn from "
         "per-grade normal distributions in the 11-feature space, with 6% of labels moved to a neighbouring grade (80/20 split). The price "
         "data are 23,360 rows (8 crops, 8 regions, every third day from 2022 to 2024), generated by Eq. (1) and split 85/15. To test "
         "the grader on reality we used 3,200 photographs from the fresh-and-rotten fruit dataset [@sultana] (CC BY 4.0), and tomato, onion "
         "and potato photographs from a vegetable dataset [@vegds] (CC BY 4.0) for the demonstration screens.")
    P.h2("C. Quality Grading")
    para("A photograph is centred on a 256 x 256 canvas; a foreground mask comes from Otsu thresholding [@otsu] of saturation with "
         "morphological clean-up (the whole image is used if the mask covers under 3% or over 98%). Eleven features are computed on "
         "the foreground: circular mean hue, mean saturation, mean value, saturation spread, colour uniformity, three grey-level "
         "co-occurrence statistics [@haralick] (contrast, homogeneity, energy; 16 levels, four angles), Canny edge density [@canny], "
         "blemish ratio (pixels darker than the mean value by 1.15 standard deviations) and a size score from the mask area. A Random "
         "Forest [@breiman] of 250 trees (depth 9, minimum leaf 4) outputs probabilities, the grade is their arg-max, and tree SHAP "
         "[@shap,shaptree] returns the six strongest feature contributions for the predicted grade.")
    para("Training and configuration. Both models are fitted with fixed random states (42 for the split and the forest, 7 for the price "
         "generator), the grader on 4,800 rows and the price models on 19,856 rows; the three price models (point, 0.15 and 0.85 quantiles) "
         "and the grader take about twenty seconds in all. No hyper-parameter search was run: the settings above were fixed in advance "
         "and not tuned on the test rows. Model files are stored with joblib and loaded once per process; if a file is missing the "
         "server retrains from the generator, which makes a fresh clone runnable.")
    para("Why hand-crafted features and not a convolutional network? The batch abstract speaks of deep learning for grading, but a "
         "network needs labelled photographs of graded crops, and none existed for this project; the only labels available were synthetic "
         "feature vectors. Classical features also give explanations a farmer can read (colour uniformity, blemish ratio), train in "
         "seconds without a GPU, and let the same extractor be used for training and for inference. The cost is that the representation "
         "is fixed by us, and Section IV-C shows where that representation and the synthetic labels fail. A network trained on a labelled "
         "photograph set is the natural next step.")
    P.h2("D. Price Intelligence")
    para("The price model takes 19 inputs: month (sine and cosine), rainfall, temperature, demand and supply indices, fuel index, days "
         "since harvest, a trend, grade premium, region factor and eight crop indicators. The generator used for training and "
         "evaluation is")
    P.eq("p = b_{c} · π_{g} · ρ_{r} · δ(D/S) · s(m) · w(R) · (1 − 0.004 d) · f^0.15 · (1 + T) · (1 + ε)", 1)
    para("with crop base price b, grade premium π (A 1.22, B 1.05, C 0.85, Reject 0.45), region factor ρ, a clipped demand-supply "
         "term δ, a seasonal term s, a rain penalty w, fuel index f, trend T and noise ε ~ N(0, 0.035²). A gradient-boosting regressor "
         "[@friedman] (300 trees, depth 4) gives the point estimate and two quantile models [@koenker] (alpha 0.15 and 0.85) give the "
         "band, trained on the pinball loss", indent=False)
    P.eq("L_{α} = Σ max( α (y − q), (α − 1)(y − q) )", 2)
    para("The 30-day curve evaluates the point model on 31 successive daily market contexts (smoothed 0.8/0.2), and the best day to sell "
         "within the 14-day spoilage window is", indent=False)
    P.eq("d* = argmax_{0 ≤ d ≤ 14} p̂(d)", 3)
    P.h2("E. Delivery, Orders and Monitoring")
    para("Each buyer b receives a score from five normalised terms, with distance by the haversine formula between town coordinates, "
         "transit time of distance/380 + 0.5 days and a spoilage span of 5 days for tomato, mango, onion and potato and 21 days otherwise:", indent=False)
    P.eq("S_{b} = 100 (0.38 x_{price} + 0.22 x_{dist} + 0.18 x_{fit} + 0.14 x_{rel} + 0.08 (1 − x_{spoil}))", 4)
    para("The top five buyers and a route are returned. Orders follow placed, confirmed, in transit, delivered (or cancelled), "
         "each step restricted to the owning farmer, a distributor or the admin. Starting transit creates a shipment whose sensor "
         "stream is a seeded random walk (temperature, humidity, shock) with an injected event every 18 to 55 steps, sent every 1.4 s over a "
         "WebSocket; a reading is flagged if it is an event, outside 0.5-9.5 C or 47-83% humidity, or at least 2.4 g. Authentication "
         "uses signed tokens and bcrypt, uploads are limited to JPEG, PNG or WEBP of at most 8 MB, and every endpoint checks the role.")

    P.h2("F. Implementation Notes")
    para("The grading route first stores the upload under a random file name and creates a draft listing, then extracts features, "
         "classifies, explains and stores the grade with its probabilities, features and SHAP contributions as JSON, so that the "
         "buyer later sees exactly what the farmer saw. The price route is stateless and the delivery route stores each recommendation. "
         "Listing status moves from draft to listed to sold; the buyer's order moves the listing to sold at once. Tokens expire after "
         "seven days. The front end keeps the token in a persisted store, redirects each role to its own area and draws the radar, "
         "SHAP, forecast and live-sensor charts from the same JSON.")

    # ---------------------------------------------------------------- IV
    P.h1("IV. Results")
    P.h2("A. Test Conditions")
    para("Quality results use the 1,200 held-out synthetic rows; price results the 3,504 held-out rows (and a 2024 time hold-out); the "
         "real-photo check uses 3,200 photographs (400 per fruit, half fresh), mapped fresh to grades A or B and rotten to C or Reject. "
         "That mapping is our assumption, because fresh and rotten is a binary label. The machine is the one in Section I-A; five seeds "
         "were run for each model.")
    P.table([["Measure", "Result"],
             ["Grader, synthetic test (1,200): accuracy / macro-F1", "0.923 / 0.924"],
             ["Grader, five seeds: accuracy", "0.929 ± 0.003"],
             ["Logistic regression, same split", "0.931"],
             ["Shipped grader on 3,200 real photos: accuracy / AUC", "0.504 / 0.571"],
             ["Same 11 features, real labels, 5-fold CV: accuracy / AUC", "0.906 / 0.970"],
             ["Real labels, unseen fruit types: accuracy", "0.669"],
             ["Price GBR (3,504): MAE / R2 / MAPE", "164.6 INR / 0.992 / 5.1%"],
             ["Price noise floor MAPE; time hold-out 2024 MAE", "2.8%; 191.8 INR"],
             ["70% price band, empirical coverage", "0.693"],
             ["Delivery top buyer unchanged under weight changes", "97.7%"],
             ["IoT simulator: steps flagged", "91.6%"],
             ["Automated tests", "36 passed, 3 expected failures"],
             ["Median API time: upload and grade / price forecast", "174 ms / 55 ms"]],
            "Measured results", widths=[0.66, 0.34], size=6.8)
    P.h2("B. Quality Grading on Synthetic Data")
    para("The grader reaches 92.3% on its own test set and 92.9% ± 0.3% over five seeds; every one of the 92 errors is between neighbouring "
         "grades. Because 6% of labels are noisy, the same model scores 98.1% against the clean labels, so the ceiling on noisy labels is "
         "about 94%. A linear model scores 0.931, k-nearest neighbours 0.930, histogram boosting 0.922, a depth-9 tree 0.885 and the "
         "majority class 0.271: the synthetic classes are close to linearly separable, and the forest is no better than a logistic "
         "regression. Colour alone gives 0.85, texture alone 0.83 and blemish ratio alone 0.75, and colour uniformity and blemish ratio "
         "carry the most importance (0.239 and 0.241), as the generator was built to make them.")
    para("Per grade, precision, recall and F1 are 0.928, 0.964 and 0.946 for A; 0.924, 0.862 and 0.892 for B; 0.909, 0.894 and 0.901 for C; and "
         "0.933, 0.986 and 0.959 for Reject. Grade B is the hardest because it sits between two neighbours (21 B lots were graded A and 24 "
         "graded C). The mean top-class probability is 0.869 on correct and 0.788 on wrong predictions, so the reported confidence "
         "separates them only weakly.")
    P.h2("C. Check on Real Photographs")
    para("On 3,200 real photographs the shipped grader assigns grade B to 99.8% of fresh and 99.1% of rotten fruit (Fig. 3a); accuracy is "
         "0.504, which is chance, and the AUC of P(A)+P(B) is 0.571 and falls below 0.5 for guava, jujube and orange. The reason is the "
         "domain gap: 99.8% of the real images lie outside the three-standard-deviation envelope of every synthetic grade profile on at "
         "least one feature, for example a mean hue of 0.14 against about 0.27, a texture contrast of 0.02 against 0.18 to 0.64, and "
         "a blemish ratio of 0.14 for both fresh and rotten fruit. The extractor itself is sound: a Random Forest trained on the same 11 features "
         "with real labels reaches 0.906 in five-fold cross-validation (AUC 0.970), although only 0.669 on fruit types it has not seen "
         "(Fig. 3b). Real labels reverse the importances, with mean hue first (0.29) and blemish ratio last (0.03).")
    P.figure(PAPER / "fig3_results.png", "Real-photo check: predicted grades (a) and accuracy in four settings (b).", width=3.4)
    P.h2("D. Price, Delivery and Monitoring")
    para("The price model reaches an MAE of 164.6 INR per quintal, R2 0.992 and MAPE 5.1% on random held-out rows; the noise floor of the "
         "generator is a MAPE of 2.8%, the 2024 time hold-out gives MAE 191.8 and R2 0.988, and a crop-mean baseline gives MAE 990.0, "
         "linear regression 527.7 and a random forest 245.3. The 70% band covers 69.3% of test prices, but 119 test rows have a "
         "low bound above the high bound. Because the target is the formula of Eq. (1), these figures show recovery of the generator: "
         "crop type carries 71.5% of importance and rainfall and temperature almost none (0.0001), and the served 30-day curve moves "
         "by a mean of 33.5% of price because its daily context is hash-generated noise, so a recommended best day gains a mean "
         "14.6% that no real market supports (Fig. 6). The delivery score keeps the same top buyer in 97.7% of 25,600 weight perturbations "
         "(weights multiplied by 0.5 to 1.5), but buyer offers are simulated. In the sensor simulator 91.6% of steps are flagged "
         "and 97.3% of flags are persistent drift, not events, because the walks do not return to range.")
    para("Across crops the error is largest where prices are large (MAE 97 INR for tomato, 250 for mango, 383 for cotton) and across grades "
         "the relative error is largest for Reject (MAPE 11.4%, against 4.2% for A and B), where prices are lowest. Five seeds give an "
         "MAE of 162.8 ± 5.5 INR. The served recommendation puts the best day at today for 24.2% of the 256 crop, region and grade "
         "combinations and a mean of 6.2 days ahead, and a better grade is never priced lower (64 of 64 crop and region pairs). The "
         "delivery engine returns in a median 65 ms and the quality grader in about 174 ms including the upload and the database "
         "write, so both are fast on a CPU.")
    P.h2("E. Output Screens and Testing")
    P.figure(PAPER / "fig4_grade.png", "Grading screen: a real tomato photograph graded B (70.1%).", width=3.4)
    P.figure(PAPER / "fig5_price.png", "Price screen: forecast, band and SHAP drivers.", width=3.4)
    para("Figs. 4 and 5 are captures of the running application for a tomato lot of 620 kg: grade B at 70.1% confidence, a current price "
         "of 2,272 INR per quintal and best day 9 November; the top-ranked buyer matched 67.7%. The grading screen explains the grade with colour uniformity and blemish ratio as the largest SHAP drivers, followed by texture homogeneity and edge density; the price screen shows regional factor and grade pushing the price up and crop type, supply, demand and season pulling it down relative to the average. Thirty-six pytest tests pass in 30 s, "
         "covering authentication, role guards, grading, price, delivery, the order lifecycle, the WebSocket stream and the admin "
         "screens; three expected-failure tests document gaps: a corrupt image returns an error and leaves a gradeless draft listing, "
         "catalogue values are not validated, and order quantities are not validated. The admin volume figure multiplies a price per "
         "quintal by kilograms and so overstates volume 100 times. Latency medians are 1.6 ms for health, 174 ms for upload and grading "
         "and 325 ms for login (bcrypt).")
    para("For the delivery engine the replica of the score reproduces the served top buyer in 256 of 256 crop, region and grade "
         "cases. The top buyer varies (Harvest & Co. 32.8% of cases, Coastal Fresh 28.1%, GreenMart 21.9%, Riverstone 17.2%), the mean "
         "top score is 76.0 and the mean gap between best and worst buyer is 18.1 points. Distance separates buyers most (standard "
         "deviation of its weighted term 0.052), then price (0.021), while the quality-fit term is the same for every buyer of a "
         "listing and so never changes the ranking; because offers are a hash-derived factor of 0.90 to 1.18 times the model price, "
         "the price term carries no market information.")
    P.figure(PAPER / "fig6_price_iot.png", "Price baselines (a) and simulator alarm over time (b).", width=3.4)
    para("Failure and edge cases were exercised on purpose. A wrong password is rejected with HTTP 401 and shown on the log-in page; a "
         "file that is not an image, or one over 8 MB, is refused with 400; skipping an order stage returns 400 and a wrong role returns "
         "403. The gaps are these: a corrupt file that claims to be a JPEG raises an error inside grading (HTTP 500) after the "
         "draft listing has been written; an unknown crop or region is priced without complaint, from an all-zero crop indicator; a "
         "negative quantity or a zero price is accepted in an order; the sensor WebSocket needs no token; any signed-in user can read "
         "any listing; and reopening a shipment replays the same sensor trajectory from its seed and writes duplicate history rows.")
    P.h2("F. Comparison with the Base Paper")
    para("The base paper [@base] reports 99.13% to 99.3% on soil, irrigation, crop and fertiliser tasks from sensor and tabular data, "
         "whereas CropSight grades photographs, prices produce and plans delivery after harvest on synthetic data; the datasets, tasks "
         "and conditions differ, so no head-to-head comparison is valid and none is claimed. What the two share is the combination of "
         "sensing, learning and explanation; what differs is the stage of the value chain, and our evidence is explicitly weaker on "
         "real data.")

    P.h2("G. Threats to Validity")
    para("The strongest threat is circularity: each synthetic score is measured against data produced by a formula that the model "
         "can learn, so it cannot reveal how a model behaves in the field. The real-photo check addresses this for the grader only, "
         "uses fruits that are not in the application's crop list, and rests on a binary freshness label mapped to grades by our "
         "assumption. The price model has no real-data check at all. The seed studies vary the split and the forest but not the "
         "generator, and the delivery results depend on four seeded buyers and eight towns. Latency was measured with other jobs "
         "running on the same machine. Finally, the demonstration photographs were chosen from the first 40 images of each vegetable "
         "so that more than one grade appears, although all 80 scanned onions and tomatoes were graded B.")

    # ---------------------------------------------------------------- V
    P.h1("V. Conclusion")
    para("CropSight delivers a working, tested platform that grades a crop from a photograph with SHAP explanations, forecasts a "
         "price with a band and a best day to sell, ranks buyers, runs a role-gated order lifecycle and streams simulated cold-chain "
         "readings. Against the six objectives, grading, forecasting, monitoring, explanation, recommendation and a marketplace all "
         "exist, but the evidence differs. The key finding is negative and useful: the grader's 92.3% on synthetic data does not transfer, "
         "since on 3,200 real photographs it is at chance, while the same features retrained with real labels reach 90.6%. The honest "
         "limits are synthetic training data, a grader that is classical machine learning on hand-crafted features and not deep "
         "learning, a price forecast that is not a time-series forecast, simulated sensors and buyer offers, and no user study.")
    para("Taking the six objectives in turn: grading from images is met with classical machine learning, not deep learning, and is not "
         "validated on real crop photographs; price forecasting is implemented and validated only against synthetic prices; storage "
         "and transport monitoring runs on simulated sensors; explanations are delivered by SHAP on the grading and price screens; "
         "buyer, route and best-day recommendations are implemented, with simulated offers; and the marketplace with its order "
         "lifecycle is implemented and covered by the tests.")
    para("Future work: train the grader on labelled photographs of the target crops and report accuracy on a held-out real set; fit the "
         "price model to real mandi price series and evaluate it with a rolling forecast against time-series baselines; read real "
         "sensors and make the alarm return to normal; validate inputs and close the authorisation gaps found; and test with farmers.")

    P.references()
    P.d.core_properties.title = K.TITLE
    P.d.core_properties.author = "Batch A17"
    P.d.save(OUT)
    print("saved", OUT, "refs:", len(ORDER), ORDER)


if __name__ == "__main__":
    build()
