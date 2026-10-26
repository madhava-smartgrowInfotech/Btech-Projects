"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides). Every number is from docs/REPORT_FACTS.md."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import BATCH, CHARTS, COLLEGE, DEPT, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, SHOTS, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"


def px(v):
    return Emu(int(v * 9525))


def cards(s, items, size=15):
    for tid, bid, ttl, body in items:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, size, min_pt=10)


def frame_to(pic, frame):
    frame.left, frame.top, frame.width, frame.height = pic.left, pic.top, pic.width, pic.height


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, "Batch A17", "crop quality grading, price forecasting, smart delivery, SHAP, IoT"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A17"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), "Smart Crop Delivery with AI-Based Price and Quality Intelligence")
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A17 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers the design, methodology and progress of CropSight.")

    # 2 contents (template labels already generic)
    s = S[1]
    for sid, t in [(21, "Model and configuration"), (18, "Data preparation"), (27, "Data flow, modules, database"), (36, "Preliminary results and screens")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts: recap, architecture, methodology, data, preparation, models, decision logic, data flow and database, technology, progress, results and screens, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Farmers see fluctuating prices, inconsistent manual quality grading, unmonitored storage and transport, and little decision support or direct access to buyers.", 16)
    fill(s, 13, "A farmer-buyer-distributor platform: photo-based quality grade with SHAP, price forecast with best day to sell, buyer ranking, a marketplace and live sensor monitoring.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Crop photo upload", "11 vision features, Random Forest grade", "Price forecast, 31 days", "Buyer ranking, best day", "Order, transit, live sensors"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1. The base paper [1] is pre-harvest (soil, irrigation, crop and fertiliser advice); CropSight applies the IoT + ML + explainable-AI idea after harvest.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 19 + TypeScript + Vite (port 5190); the dev server proxies /api, /uploads and /ws.",
                "API: FastAPI (port 8000), 21 endpoints (20 REST + 1 WebSocket); JWT sign-in and role guards.",
                "ML services: OpenCV / scikit-image features, Random Forest grader, Gradient Boosting price model, SHAP.",
                "Rule services: buyer scoring, order state machine, simulated sensor stream.",
                "Data: SQLite (9 tables), uploaded photos, .joblib model files."], 16, bullet=True)
    set_notes(s, "Everything runs on one machine. The browser talks only to the Vite server, which proxies to FastAPI. No external AI service is used.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Generate synthetic data", "Extract 11 vision features", "Train Random Forest grader", "Train price GBR + quantile bands", "Evaluate: seeds, baselines", "Real-photo check", "Serve: FastAPI + React"]):
        fill(s, sid, t, 14, min_pt=9)
    set_paras(shape(s, 25), "Real vision, synthetic labels")
    fill(s, 26, "Features are computed from the real pixels of the uploaded photo; the grade classes were learned from generated feature data.", 15)
    set_paras(shape(s, 28), "Explained and disclosed")
    fill(s, 29, "SHAP shows what drove each grade and price; the admin screen states that the training data is synthetic.", 15)
    set_paras(shape(s, 31), "Reproducible")
    fill(s, 32, "Fixed seeds; scripts/exp_*.py regenerate every figure into experiments/*/metrics.json.", 15)
    set_notes(s, "Seven-step workflow from data generation to serving, then a real-photo check that was added to test the synthetic-trained grader.")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    ch = replace_chart(s, 4, ["Synthetic quality features", "Synthetic price series", "Fresh/rotten fruit photos (real)", "Vegetable photos scanned (real)"], [("Rows or images", [6000, 23360, 3200, 120])])
    style_chart(ch, "Rows or images per dataset", ["1B6B3A"], legend=False, label_pt=11, axis_title="rows / images")
    rows = [["Dataset", "Size", "Use"],
            ["Synthetic quality features", "6,000 rows, 11 features, 4 grades", "Train and test the grader"],
            ["Synthetic market series", "23,360 rows, 2022-2024, 8 crops x 8 regions", "Train and test the price model"],
            ["Fresh and Rotten Fruits [8] (CC BY 4.0)", "3,200 photos, 8 fruits", "Real-photo check"],
            ["Vegetable photos [9] (CC BY 4.0)", "120 scanned of 4,319; 4 used", "Demo screens"],
            ["Demo database", "8 users, 8 towns, 1 seeded listing", "Marketplace demo"]]
    sz = table_size(rows, 530, 400, [0.38, 0.38, 0.24], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.38, 0.38, 0.24], header_size=sz)
    set_notes(s, "No real labelled crop-quality set and no real price series were available, so both models use synthetic data. Two public real-photo sets were used only to test the grader and to produce the demo screens.")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation and features")
    set_paras(shape(s, 7), "Image features")
    fill(s, 8, "Photo centred on a 256 x 256 white canvas; HSV foreground mask (Otsu on saturation; whole image if the mask is under 3 % or over 98 %). 11 features: colour (5), GLCM texture (3), edge density, blemish ratio, size score.", 15)
    set_paras(shape(s, 12), "Market features")
    fill(s, 13, "19 inputs: month sin/cos, rainfall, temperature, demand, supply, fuel index, days to harvest, year trend, grade premium, region factor and 8 crop one-hots. The daily context is hash-seeded per crop, region and date, not a live feed.", 15)
    set_paras(shape(s, 17), "Synthetic data")
    fill(s, 18, "Quality: a normal distribution per grade in feature space; 360 of 6,000 labels (6 %) moved to a neighbouring grade. Price: base x grade x region x demand/supply x season x weather x freshness x fuel^0.15 x trend, plus N(0, 3.5 %) noise.", 15)
    set_notes(s, "Preparation = turning a photo into 11 numbers, building the 19 market inputs, and generating the two training sets.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Models and explanation")
    put_picture(s, DIAG / "02_high_level_workflow.png", (75, 124, 693, 518), remove_ids=(4,), name="High-level workflow")
    rows = [["Component", "Detail"], ["Grader", "Random Forest, 250 trees"], ["Price", "Gradient Boosting, 300 trees"], ["Band", "Quantile GBR, 15 % / 85 %"], ["Explanation", "SHAP TreeExplainer"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "Classical machine learning on the CPU; no deep network and no GPU.", 11)
    set_notes(s, "The abstract speaks of deep learning for grading; the implemented grader is a Random Forest on hand-crafted features.")

    # 9 training configuration
    s = S[8]
    set_paras(shape(s, 2), "Training configuration")
    rows = [["Setting", "Value"], ["Grader data", "6,000 synthetic rows, 6 % label noise"], ["Grader split", "80/20 stratified (4,800 / 1,200), seed 42"],
            ["Random Forest", "250 trees, depth 9, min leaf 4"], ["Price data", "23,360 synthetic rows, 19 features"], ["Price split", "85/15 random (19,856 / 3,504)"],
            ["Gradient Boosting", "300 trees, depth 4, learning rate 0.06"], ["Quantile bands", "alpha 0.15 / 0.85, 250 trees, depth 3"]]
    set_table(shape(s, 10), rows, 13, header_size=13)
    pic = put_picture(s, CHARTS / "03_quality_baselines_seeds.png", (672, 230, 550, 270), remove_ids=(5,), name="Baselines and seeds chart")
    frame_to(pic, shape(s, 6))
    set_paras(shape(s, 7), "Baselines on the same split (left) and five seeds: 0.929 +- 0.003 (right)", 11)
    shape(s, 7).top = Emu(pic.top + pic.height + px(10))
    set_notes(s, "A logistic-regression baseline scores 0.931 and k-NN 0.930, so the synthetic classes are close to linearly separable; the shipped forest scores 0.923.")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (45, 124, 620, 520), remove_ids=(4,), name="Decision flowchart")
    items = [(5, 6, 7, "Grade", "Class with the highest forest probability; confidence = that probability; top-6 SHAP features explain it."),
             (8, 9, 10, "Best day to sell", "Price model on 31 successive days; best day = highest price in days 0-14; band from the 15 % / 85 % models."),
             (11, 12, 13, "Buyer ranking", "Score x100 = 0.38 price + 0.22 distance + 0.18 quality fit + 0.14 reliability + 0.08 (1 - spoilage); top 5 shown.")]
    for card, tid, bid, ttl, body in items:
        c = shape(s, card)
        t, b = shape(s, tid), shape(s, bid)
        c.left, c.width = px(690), px(531)
        t.left, t.width = px(714), px(490)
        b.left, b.width = px(714), px(490)
        set_paras(t, ttl)
        t.top, t.height = c.top + px(10), px(40)
        b.top, b.height = c.top + px(50), px(92)
        fill(s, bid, body, 15, min_pt=10)
    set_notes(s, "Decision logic: how a grade is chosen, how the best selling day is picked and how buyers are scored. The buyer offers are simulated from the reference price; the engine is a transparent formula, not a trained model.")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 340), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 340), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Grade crop' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(462)
    rows = [["Module", "Responsibility"], ["Grading", "validate photo, 11 vision features, Random Forest, SHAP, store grade"],
            ["Price and delivery", "31-day forecast with band and best day; weighted buyer score and route"],
            ["Marketplace and monitoring", "listings, order state machine, role guards, simulated sensor stream over WebSocket"]]
    set_table(shape(s, 12), rows, 12, widths=[0.22, 0.78], header_size=12, row_h=34)
    shape(s, 12).top = px(492)
    set_notes(s, "Level 1 splits the system into seven processes and six data stores; level 2 opens the grading process.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 385, 400), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (455, 118, 370, 400), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 9 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(524), px(768)
    rows = [["Table", "Key fields", "Relations"],
            ["users", "email, role, region, reliability", "1:N listings, orders"],
            ["crop_listings", "crop, quantity, status, price", "1:1 quality_grades"],
            ["quality_grades", "grade, confidence, SHAP (JSON)", "N:1 listing"],
            ["orders, shipments", "status, buyer, distributor", "1:N sensor_readings"]]
    set_table(shape(s, 13), rows, 10, widths=[0.28, 0.42, 0.30], header_size=10, row_h=30, total_px=371)
    t = shape(s, 13)
    t.left, t.top, t.width = px(850), px(150), px(371)
    set_notes(s, "Nine SQLite tables created by SQLAlchemy. Structured results (probabilities, features, SHAP, forecasts) are stored as JSON columns. The price_predictions table exists but no route writes to it.")

    # 13 technology
    s = S[12]
    cards(s, [(7, 8, "Language", "Python 3.11, TypeScript"), (12, 13, "Machine learning", "scikit-learn 1.5.2 (Random Forest, Gradient Boosting), SHAP 0.46, NumPy, pandas"),
              (17, 18, "Vision", "OpenCV 4.10, scikit-image 0.24 (GLCM texture), Pillow"),
              (22, 23, "Backend", "FastAPI 0.115, SQLAlchemy 2.0, Pydantic 2, JWT (python-jose), bcrypt, WebSocket"),
              (27, 28, "Frontend", "React 19, Vite 8, Tailwind CSS 4, Recharts, Framer Motion, Three.js, Zustand"),
              (32, 33, "Storage and tests", "SQLite, joblib model files, pytest (36 pass, 3 expected fail), Playwright captures")], 15)
    set_notes(s, "All versions are the ones pinned in requirements.txt and package.json. There is no external AI service.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Photo upload, 11 vision features, Random Forest grade, SHAP", "Completed"], ["Price model, 31-day forecast, 15-85 % band, best day", "Completed"],
            ["Delivery planner: weighted buyer score and route", "Completed"], ["Marketplace, order state machine, role guards", "Completed"],
            ["Live monitor: sensor simulator over WebSocket", "Completed (simulated)"], ["Admin overview and model insights", "Completed"],
            ["Experiments, 39 tests, latency measurements", "Completed"], ["Real crop-image and price data; real sensors; user testing", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.7, 0.3], header_size=13)
    for vid, lid, v, lab in [(6, 7, "21", "API endpoints (20 REST + 1 WebSocket)"), (9, 10, "9", "database tables"), (12, 13, "36", "tests pass; 3 expected failures")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "All planned features are implemented. The open items are real training data, real sensors and user-acceptance testing.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: grading accuracy")
    ch = replace_chart(s, 4, ["Synthetic test (shipped RF)", "Synthetic, 5-seed mean", "Real photos, shipped RF", "Real photos, retrained (CV)"], [("Accuracy", [0.923, 0.929, 0.504, 0.906])], number_format="0.000")
    style_chart(ch, "Accuracy of the 11-feature classifier: synthetic data versus 3,200 real fruit photos", ["1B6B3A"], number_format="0.000", legend=False, label_pt=11, axis_title="accuracy")
    ch.value_axis.maximum_scale = 1.05
    for vid, lid, v, lab in [(6, 7, "92.3 %", "grader test accuracy, synthetic data"), (9, 10, "0.992", "price model R2 (MAE 164.6, MAPE 5.1 %)"), (12, 13, "50.4 %", "shipped grader on real photos (chance 50 %)")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Synthetic scores recover the generator, not field performance. On 3,200 real fruit photos (fresh versus rotten, mapped to grades A/B versus C/Reject) the shipped grader gives grade B to about 99 % of photos and sits at chance (AUC 0.57); the same 11 features retrained on real labels reach 90.6 % in 5-fold cross-validation (66.9 % on unseen fruit types). The price noise floor is MAPE 2.8 %.")

    # 16 output screens
    s = S[15]
    p1 = put_picture(s, SHOTS / "07_grading_result.png", (45, 128, 590, 476), remove_ids=(4,), name="Grading result screenshot")
    p2 = put_picture(s, SHOTS / "08_price_forecast.png", (645, 128, 590, 476), remove_ids=(7,), name="Price forecast screenshot")
    frame_to(p1, shape(s, 5))
    frame_to(p2, shape(s, 8))
    set_paras(shape(s, 6), "Real tomato photo graded B (70.1 %): radar of features and SHAP drivers", 11)
    set_paras(shape(s, 9), "Price forecast: INR 2,272 per quintal, 30-day curve and drivers", 11)
    for cap_id, p in ((6, p1), (9, p2)):
        shape(s, cap_id).top = Emu(max(p.top + p.height + px(12), px(500)))
    set_notes(s, "Real captures of the running application with real photographs (tomato from a public vegetable dataset, CC BY 4.0).")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Health, authentication, role guards", "9", "9 passed"], ["Listings, grading, vision features", "12", "11 passed, 1 expected fail"],
            ["Price and delivery", "8", "7 passed, 1 expected fail"], ["Marketplace and orders", "5", "4 passed, 1 expected fail"], ["IoT, simulator, admin", "5", "5 passed"], ["Total", "39", "36 passed, 3 expected fail"]]
    set_table(shape(s, 18), rows, 12, widths=[0.5, 0.12, 0.38], header_size=12)
    fill(s, 5, ["36 tests pass in about 30 s; 3 tests assert hardened behaviour and fail today: a corrupt image gives HTTP 500 and leaves a gradeless draft, unknown crops are priced, order quantity and price are not validated.",
                "API median times: health 1.6 ms, upload and grade 174 ms, price 55 ms, delivery plan 100 ms, login 325 ms.",
                "Delivery ranking: top buyer unchanged in 97.7 % of weight perturbations.",
                "Found while testing: admin volume is 100x too high (INR per quintal x kg).",
                "Not done: user-acceptance, load and security testing."], 15, bullet=True)
    set_notes(s, "Expected-fail tests are reported as such rather than hidden. Latencies were measured on a live server over 127.0.0.1 on an i5-12500 with other jobs running.")

    # 18 challenges
    s = S[17]
    cards(s, [(7, 8, "No real training data", "No labelled crop photos or price series: both models use synthetic data, disclosed in the app."),
              (12, 13, "Synthetic-to-real gap", "On 3,200 real photos the grader is at chance (50.4 %); the same features retrained on real labels reach 90.6 %."),
              (17, 18, "Circular price score", "R2 0.992 recovers the generator (noise floor MAPE 2.8 %); weather features carry about zero importance."),
              (22, 23, "Simulated sensors", "A random walk without mean reversion: 97 % of alarm steps are drift. Real sensors are needed."),
              (27, 28, "Unvalidated inputs", "Corrupt image: HTTP 500 and a gradeless draft; order quantity and price unchecked. Logged as expected-fail tests."),
              (32, 33, "Unit error in a KPI", "Admin volume multiplies INR per quintal by kg, so it reads 100x too high; found while capturing screens.")], 14)
    set_notes(s, "Six challenges and findings, each backed by a measurement in REPORT_FACTS.md.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["Collect real labelled photos", "Retrain and validate", "Harden and fix", "Final report and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Collect a labelled set of real crop photos with grades and retrain the grader; test on unseen photos.",
                 "Replace the hash-generated market context with a real mandi price series and validate over time.",
                 "Fix the admin unit error, validate inputs, clean up failed uploads, make the sensor alarm clear, authenticate the WebSocket.",
                 "Collect feedback from farmers and buyers (not yet done); complete the report and final demo."], 17, bullet=True)
    set_notes(s, "What remains: real data first, then the hardening items found in testing, then the report and final demonstration.")

    set_notes(S[19], "Thank you. Questions are welcome.")
    # template leftovers in shape names and picture descriptions (slide 18 cards)
    for n, sid in enumerate((4, 9, 14, 19, 24, 29), 1):
        shape(S[17], sid).name = f"Challenge card {n}"
    for sh in S[17].shapes:
        if sh.shape_type == 13:
            sh._element.nvPicPr.cNvPr.set("descr", "icon")
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    for m in prs.slide_masters:
        th = m.part.part_related_by(RT.THEME)
        th._blob = th.blob.replace(b"Hidden Risk", b"CropSight")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
