"""Excel (.xlsx) export of attendance for a session."""
from __future__ import annotations

import io
from core.timeutil import to_local

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def records_dataframe(session, records, expected_rolls: list[str] | None = None) -> pd.DataFrame:
    """Records as a DataFrame; if an explicit roll list is known, absentees are added."""
    rows = []
    for r in sorted(records, key=lambda x: x.roll_number):
        rows.append({
            "Roll Number": r.roll_number,
            "Student Name": r.student_name,
            "Status": "Present",
            "Marked At": to_local(r.marked_at).strftime("%Y-%m-%d %H:%M:%S"),
            "Distance (m)": round(r.distance_m, 1),
            "GPS Accuracy (m)": round(r.gps_accuracy_m, 1) if r.gps_accuracy_m is not None else "",
            "Challenge": r.challenge,
            "Liveness Score": round(r.liveness_score, 2),
            "Identity Similarity": round(r.identity_similarity, 3) if r.identity_similarity is not None else "",
            "Max Duplicate Similarity": round(r.max_duplicate_similarity, 3) if r.max_duplicate_similarity is not None else "",
            "Device": r.device_hash[:12],
        })
    present = {r.roll_number for r in records}
    for roll in expected_rolls or []:
        if roll not in present:
            rows.append({"Roll Number": roll, "Student Name": "", "Status": "Absent"})
    df = pd.DataFrame(rows)
    if not df.empty:
        df = df.sort_values("Roll Number").reset_index(drop=True)
    return df


def to_excel_bytes(session, records, attempts=None, expected_rolls: list[str] | None = None) -> bytes:
    df = records_dataframe(session, records, expected_rolls)
    present = int((df["Status"] == "Present").sum()) if not df.empty else 0
    absent = int((df["Status"] == "Absent").sum()) if not df.empty else 0

    summary = pd.DataFrame({
        "Field": ["Course", "Section", "Faculty", "Session Start", "Session End",
                  "Geo-fence centre", "Radius (m)", "Present", "Absent", "Blocked attempts"],
        "Value": [
            f"{session.course_code} – {session.course_name}", session.section, session.faculty.full_name,
            to_local(session.starts_at).strftime("%Y-%m-%d %H:%M"),
            to_local(session.ends_at).strftime("%Y-%m-%d %H:%M"),
            f"{session.latitude:.6f}, {session.longitude:.6f}", session.radius_m, present, absent,
            sum(1 for a in (attempts or []) if not a.passed),
        ],
    })

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        (df if not df.empty else pd.DataFrame({"Roll Number": [], "Status": []})).to_excel(
            xw, sheet_name="Attendance", index=False)
        if attempts:
            adf = pd.DataFrame([{
                "Time": to_local(a.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                "Roll Number": a.roll_number, "Stage": a.stage,
                "Result": "PASS" if a.passed else "BLOCKED", "Reason": a.reason,
                "Device": a.device_hash[:12],
            } for a in attempts])
            adf.to_excel(xw, sheet_name="Verification Log", index=False)

        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 10), 60)
            ws.freeze_panes = "A2"
    return buf.getvalue()
