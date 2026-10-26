"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import (BASE_PAPER, BATCH, CHARTS, DIAG, FF, FV, GUIDE, GUIDE_DESIG, LT, LV, PB, PRODUCT, REFERENCES, ROOT, SHOTS, STUDENTS, YEAR)  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE (Data Science)  |  Batch A2  |  Final Review"
SUB = "Attendance Magic: A Smart Geo-Fenced Attendance System with Live Challenge Verification"
CROPS = ROOT / "docs" / "deck_crops"
CROPS.mkdir(parents=True, exist_ok=True)


def px(v):
    return Emu(int(v * 9525))


def crop(name, box, out):
    p = CROPS / out
    Image.open(SHOTS / name).crop(box).save(p)
    return p


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 9)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l)
    set_paras(shape(s, cap_ids[1]), cap_r)
    for cid in cap_ids:
        shape(s, cid).top = px(554)


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Final Review", SUB, "Batch A2", "geo-fencing, liveness, face verification, attendance"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A2"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    s = S[0]
    set_paras(shape(s, 4), SUB)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A2")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of Attendance Magic: geo-fenced attendance with live challenge verification.")

    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Verification Chain",
                "Design & Implementation", "Results, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    s = S[2]
    fill(s, 7, ["Attendance by roll call or shared link is slow and easy to fake; face recognition confirms identity but a photo can still be shown.",
                "Attendance Magic combines, in one time-bound session: JWT and device-bound sign-in, roll restrictions, a GPS geo-fence, a random live facial challenge (MediaPipe Face Mesh) and ArcFace face matching (InsightFace) with a duplicate-face check.",
                "Faculty create sessions and share a link or QR code; students check in from a phone; records, a map, a verification log and an Excel export follow.",
                "No model is trained. Measured on public LFW pairs: AUC 0.999; at the deployed threshold 98.7% of same-person pairs pass and 0 of 2,166 different-person pairs do.",
                "A single stolen photo fails 7 of 8 challenges; a rotated photo or several photos of the victim can still pass, and this is reported as a limitation.",
                "Keywords: attendance, geo-fencing, liveness, face verification, ArcFace, MediaPipe."], 19, bullet=True)
    set_notes(s, "Abstract: problem, combined checks, application, and the measured headline with its limits.")

    s = S[3]
    fill(s, 7, [BASE_PAPER,
                "Method: Python desktop application with a Tkinter GUI; webcam captures 25-30 images per student; Haar cascade face detection and LBPH recognition; subject-wise CSV logs; manual entry and text-to-speech feedback.",
                "Reported results: about 92-95% recognition under good frontal light, 5-8% false negatives, 0% false positives, under 2 s per face; basic hardware, no GPU.",
                "Limitations (our reading): identity only; no liveness, location, session window or duplicate-identity check; accuracy depends on light and pose; no dataset size or protocol reported.",
                "Surveyed in the paper: Haar+LBPH, CNN+SVM/KNN, VGG16/ResNet transfer learning, FaceNet+MTCNN."], 19, bullet=True)
    set_notes(s, "The base paper is a classroom face-recognition attendance system with LBPH. It is silent about spoofing and location.")

    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Dharmaraj et al., 2025 (base) [1]", "Haar cascade + LBPH; Tkinter; CSV logs", "Classroom simulations", "~92-95% recognition; 0% false positives", "No liveness, location or duplicate check"],
            ["Shandilya and Shah, 2026 [3]", "ArcFace (DeepFace), LBP, EAR blink, HSV dress code; FastAPI + Streamlit", "Not read beyond abstract", "96.2% recognition; 97.1% liveness; 92.4% dress code", "No geo-fence reported (abstract)"],
            ["Salam et al., 2025 [5]", "Silent-face anti-spoofing: CNN + Fourier transform; Raspberry Pi 4", "Varied light and distance", "Figures not read", "Spoof detection only (abstract)"],
            ["Classroom system, 2026 [4]", "ResNet101 anti-spoofing + ArcFace (title)", "Not retrieved", "Not retrieved", "Not reviewed"],
            ["Review, 2026 [2]", "Review: recognition, liveness, dress code (title)", "Not retrieved", "Not retrieved", "Not reviewed"],
            ["Face-recognition works in [1]", "LBPH; CNN+SVM/KNN; VGG16/ResNet; FaceNet+MTCNN", "Various", ">99% in controlled light", "Poor light, occlusion, look-alikes"]]
    table(s, rows, [0.2, 0.28, 0.14, 0.2, 0.18], 14, 440)
    set_notes(s, "Only the abstracts of [3] and [5] were read; [2] and [4] could not be retrieved, so we do not claim more.")

    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Identity only: no liveness, no location proof, no session window, no duplicate-identity check; small informal tests"],
            ["Shandilya and Shah [3]", "Liveness by LBP and blink; no location or per-session duplicate check reported in the abstract"],
            ["Salam et al. [5]", "Anti-spoofing module for a low-cost device; no attendance-session, location or roll-restriction logic reported"],
            ["Link / QR systems", "Easy to forward to absent students; no identity or presence check"],
            ["All reviewed work", "No single tool joins time-bound sessions, geo-fence, device binding, random challenge, duplicate check and Excel reports, with measured thresholds"]]
    table(s, rows, [0.26, 0.74], 17, 440)
    set_notes(s, "The gap is the combination of checks in one time-bound session, with evidence.")

    s = S[6]
    fill(s, 7, ["Manual and link-based attendance is slow, error-prone and shareable with absent students.",
                "Face recognition without liveness can be fooled by a photo, and does not prove the student is in the classroom.",
                "One person can mark several identities if faces in a session are not compared.",
                "Problem: given a faculty-defined session, accept a student's attendance only when the session is live, the student and device are eligible, the position is inside the geo-fence, a random live challenge is performed, and the face matches the enrolled face and no other face already recorded in the session."], 20, bullet=True)

    s = S[7]
    fill(s, 7, ["Enable faculty to create time-bound attendance sessions with configurable geo-fences.",
                "Authenticate students with JWT, device identification and roll-number restrictions.",
                "Verify physical presence by comparing GPS location with the session boundary.",
                "Detect liveness through a random live facial challenge using MediaPipe.",
                "Prevent duplicate identities by matching InsightFace embeddings within each session.",
                "Provide attendance summaries and Excel export for faculty.",
                "Verify with automated tests and measure the checks on a public benchmark."], 20, bullet=True)

    s = S[8]
    fill(s, 7, ["1. Session - faculty set course, window, centre, radius (60 m default) and allowed roll patterns; a unique link and QR code are shared.",
                "2. Sign-in - PBKDF2 password check; JWT valid 120 minutes and bound to a device hash.",
                "3. Eligibility - session live, roll allowed, not already marked, one device per student.",
                "4. Geo-fence - GPS inside radius + 15 m tolerance; accuracy worse than 150 m rejected.",
                "5. Live challenge - random head, mouth or eye action on two photos; three tries.",
                "6. Face match - same person in both photos, matches enrolment, not a duplicate in the session.",
                "7. Record and report - verification log, map, analytics, Excel export."], 19, bullet=False)

    s = S[9]
    put_picture(s, DIAG / "01_system_architecture.png", (240, 178, 800, 424), remove_ids=(7,), name="System architecture")
    set_paras(shape(s, 8), "Layered architecture: browsers, React frontend, FastAPI backend with verification services, SQLite and local AI models")

    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "End-to-end workflow: create session, share, sign in, eligibility, geo-fence, challenge, face match, record, monitor, export")

    s = S[11]
    set_paras(shape(s, 4), "Data, Features & Model")
    fill(s, 7, ["No model is trained: MediaPipe Face Mesh (468 landmarks) and InsightFace buffalo_l (SCRFD + ArcFace R50, 512-d) are used as published, on CPU.",
                "Features: yaw, pitch, roll, mouth aspect ratio and eye aspect ratio from landmarks; 512-d face embeddings; haversine distance.",
                "Challenge thresholds: yaw 0.12, pitch 0.08, roll 12 deg, MAR 0.35, EAR 0.18; neutral frame |yaw| <= 0.10, MAR <= 0.25.",
                "Face thresholds: frame consistency 0.45, identity 0.40, duplicate 0.45 (cosine similarity).",
                "Evaluation data: LFW protocol (6,000 pairs; 4,307 scored), 1,000 LFW photos for liveness, 157 identities with 10+ photos for attack tests. No train/test split of our own.",
                "Demo data: 5 accounts and 2 sessions; LFW photos stand in for student photographs."], 19, bullet=True)
    set_notes(s, "Nothing is trained, so the evaluation is of thresholds and attacks, not of accuracy curves.")

    s = S[12]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (300, 178, 680, 424), remove_ids=(7,), name="Algorithm flowchart")
    set_paras(shape(s, 8), "Decision flow of one check-in: eligibility, location, liveness, consistency and identity, duplicate scan, record")

    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    put_picture(s, DIAG / "09_sequence_diagram.png", (270, 178, 740, 424), remove_ids=(7,), name="Sequence diagram")
    set_paras(shape(s, 8), "Sequence of a student check-in through the React app, FastAPI, core services and SQLite")

    s = S[14]
    put_picture(s, DIAG / "06_er_diagram.png", (250, 178, 780, 424), remove_ids=(7,), name="ER diagram")
    set_paras(shape(s, 8), "Entity\u2013relationship diagram (5 SQLite tables)")

    s = S[15]
    rows = [["Module", "Implementation"],
            ["Sessions and roll patterns", "secrets.token_urlsafe(24) link tokens; fnmatch roll patterns; explicit lists give absentees"],
            ["Authentication", "PBKDF2-HMAC-SHA256 (200,000 iterations); JWT HS256 bound to a SHA-256 device hash; role guards"],
            ["Geo-fence", "haversine on a 6,371 km sphere (0.47% from WGS-84); radius + 15 m; accuracy gate 150 m"],
            ["Liveness", "Face Mesh landmarks to yaw, pitch, roll, MAR, EAR; 8 random challenges; 3 tries then lock"],
            ["Face match", "ArcFace embeddings, cosine similarity; one face per frame; duplicate scan against session records"],
            ["API and UI", "27 FastAPI endpoints; React 19 wizard; Excel export (Summary, Attendance, Verification Log)"]]
    table(s, rows, [0.22, 0.78], 17, 440)

    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Verification Chain on Public Pairs")
    rows = [["Measure", "Attendance Magic", "Base paper [1]"],
            ["Same person accepted (identity 0.40)", "98.7% (LFW, 2,141 pairs)", "~92-95% (own classroom data)"],
            ["Different person accepted", "0 of 2,166 pairs (max 0.246)", "0% false positives (reported)"],
            ["Duplicate face caught (0.45)", "97.2% of same-person pairs", "not provided"],
            ["Single stolen photo vs live challenge", "0% pass for 7 of 8 actions", "no liveness check"],
            ["Check-in server time (median)", "0.56 s on a CPU", "< 2 s per face (reported)"],
            ["Location proof", "geo-fence 75 m limit", "none"]]
    table(s, rows, [0.38, 0.34, 0.28], 17, 420)
    set_paras(shape(s, 8), "Different data: LFW press photos (ours) vs the base paper's own classroom data; not a head-to-head comparison.")
    set_notes(s, "There is no accuracy curve because nothing is trained. These are threshold measurements on 4,307 LFW pairs plus attack and timing experiments.")

    s = S[17]
    set_paras(shape(s, 4), "Similarity Distributions & Threshold Behaviour")
    two_pics(s, CHARTS / "similarity_hist.png", CHARTS / "threshold_sweep.png", "Same-person and different-person similarity (LFW)", "Error rates by threshold: 0.40 gives 1.3% false rejects, 0 false accepts")
    set_notes(s, "This is the equivalent of a confusion matrix for a threshold decision: the two distributions barely overlap.")

    s = S[18]
    set_paras(shape(s, 4), "Liveness Attacks & Parameter Analysis")
    two_pics(s, CHARTS / "attack_success.png", CHARTS / "rotation_sweep.png", "Attack success by challenge (stills; no live person)", "Rotating one photo by 15 deg or more passes both tilt challenges")
    set_notes(s, "Our own attack experiments show where liveness is weak: rotated photos and photo banks.")

    s = S[19]
    B, C = LT["B_stage_ms"], LT["C_http_ms"]
    rows = [["Measurement", "Result"],
            ["Model load (cold start)", f"MediaPipe {LT['A_cold_start_ms']['mediapipe_face_mesh'] / 1000:.1f} s; InsightFace {LT['A_cold_start_ms']['insightface_buffalo_l'] / 1000:.1f} s"],
            ["Per image (CPU)", f"landmarks {B['mediapipe_landmarks']['p50']:.1f} ms; ArcFace embedding {B['insightface_embed']['p50']:.0f} ms (median)"],
            ["Whole check-in (HTTP, in-process)", f"median {C['checkin_total_after_login']['p50']:.0f} ms; 95th percentile {C['checkin_total_after_login']['p95']:.0f} ms (40 runs)"],
            ["Duplicate scan", f"{LT['D_duplicate_scan_ms']['1000']['p50']:.0f} ms for 1,000 faces; {LT['D_duplicate_scan_ms']['5000']['p50']:.0f} ms for 5,000"],
            ["Automated tests", "23 pytest tests and 41 API smoke checks pass; 8 attack scenarios blocked at the right stage"],
            ["Not measured", "concurrent load, phone GPS error, camera trials with volunteers"]]
    table(s, rows, [0.28, 0.72], 18, 430)
    set_notes(s, "Measured on an i5-12500 laptop-class desktop with other applications running; network time excluded.")

    s = S[20]
    rows = [["System", "Data", "Method", "Checks enforced", "Reported / measured", "Extras"],
            ["Base paper [1]", "Own classroom data; size not reported", "Haar + LBPH", "Identity only", "~92-95% recognition; 0% false positives", "Tkinter GUI, CSV, speech"],
            ["Shandilya and Shah [3]", "Not read", "ArcFace + LBP + blink + HSV", "Identity, liveness, dress code (abstract)", "96.2% / 97.1% / 92.4% (abstract)", "FastAPI, Streamlit, parent alerts"],
            ["Attendance Magic", "LFW pairs and stills (public)", "ArcFace + Face Mesh challenge + haversine", "Session, device, GPS, challenge, identity, duplicate", "98.7% same-person pass at 0.40; 0 of 2,166 impostors", "Web app, map, log, Excel"]]
    table(s, rows, [0.13, 0.15, 0.16, 0.2, 0.22, 0.14], 13, 400)
    set_paras(shape(s, 8), "Different data and protocols: no numeric head-to-head is claimed. Reported figures for [1] and [3] are from their own papers (abstract only for [3]).")

    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Check-in Challenge and Result")
    a = crop("23_checkin_photos_ready.png", (110, 470, 1330, 1284), "challenge_photos.png")
    b = crop("24_checkin_marked.png", (110, 380, 1340, 1123), "marked.png")
    two_pics(s, a, b, "Random challenge with two photos loaded", "Marked present: 8 m, liveness 1.00, identity 0.86")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Blocked Attempts")
    a = crop("25_checkin_outside_geofence.png", (110, 380, 1340, 1138), "outside.png")
    b = crop("29_checkin_duplicate_face_blocked.png", (110, 280, 1340, 915), "duplicate.png")
    two_pics(s, a, b, "Blocked at location: outside the geo-fence", "Blocked at face check: duplicate face in the session")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Faculty Records and Dashboard")
    a = crop("13_session_records.png", (250, 640, 1440, 1470), "records.png")
    b = crop("11_faculty_dashboard.png", (240, 90, 1440, 960), "dashboard.png")
    two_pics(s, a, b, "Live session: share link, QR and records (present, absent)", "Dashboard: check-ins per session, checks by stage, latest attempts")

    s = S[24]
    fill(s, 7, ["Built Attendance Magic: a time-bound, geo-fenced attendance system with a random live challenge and duplicate-face check, with a React and FastAPI application.",
                "On 4,307 LFW pairs the deployed thresholds accept 98.7% of same-person pairs and none of 2,166 different-person pairs (AUC 0.999).",
                "A single photo is rejected by 7 of 8 challenges; rotated photos and photo banks still pass, so liveness is the weakest link.",
                "23 pytest tests and 41 API checks pass; a check-in takes about 0.56 s of server time on a CPU.",
                "Limits: no live-camera or phone trials, public-benchmark photos, GPS drift not measured, 15.7% of LFW photos rejected by the single-face rule."], 20, bullet=True)

    s = S[25]
    fill(s, 7, ["Trial with volunteers: genuine pass rate of each challenge, retake rate, usability.",
                "Stronger liveness: reject rotated photos, passive texture or depth cues (for example silent-face anti-spoofing [5]), time-limited capture.",
                "Phone deployment over HTTPS with measured GPS error and indoor drift; optional mock-location detection.",
                "Server database (PostgreSQL) and multi-faculty roles for institution-wide use.",
                "Notifications and LMS integration; privacy review for storing face embeddings."], 20, bullet=True)

    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], 14, min_pt=8, bullet=True, gap_pt=2)
    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
