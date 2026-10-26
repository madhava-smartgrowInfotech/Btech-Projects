"""Build 7.Review_1_PPT.pptx from D:\\abstracts\\Templates\\Review_1_PPT.pptx (16 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size  # noqa: E402
from doc_common import (BASE_PAPER, BATCH, CH_LABEL, CHALLENGES, COLLEGE, DEPT, DIAG, FV, GUIDE, GUIDE_DESIG, LV, PRODUCT, REFERENCES, STUDENTS, TITLE, YEAR)  # noqa: E402
from pptx import Presentation  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_1_PPT.pptx"
OUT = BATCH / "7.Review_1_PPT.pptx"


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 1", TITLE, "Batch A2", "geo-fencing, liveness, face verification, attendance"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A2"

    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), "A Smart Geo-Fenced Attendance System with Live Challenge Verification")
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A2 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 1 covers problem identification and research for Attendance Magic, a geo-fenced attendance system with live challenge verification.")
    set_notes(S[1], "Nine parts: abstract, base paper, literature, existing approaches, the gap, problem statement and objectives, the approach, expected outcomes and references.")

    s = S[2]
    set_paras(shape(s, 2), "Abstract: proving presence, not just identity")
    fill(s, 8, "Attendance is recorded every class, yet roll calls cost class time and are error-prone, and shared QR codes or links can be forwarded to students who are absent. Face recognition confirms who a person is, but a photo of that person can still be shown to the camera.", 16)
    fill(s, 13, "Check three things inside one time-bound session: the student is inside a geo-fence (GPS), answers a random live facial challenge (MediaPipe landmarks), and the face matches the enrolled one and no other record in the session (InsightFace embeddings).", 16)
    fill(s, 18, "A web application: faculty create sessions and share a link or QR code; students sign in with a device-bound JWT and check in from a phone; faculty see records, a map and a verification log, and export to Excel.", 16)
    set_notes(s, "The abstract in one slide: the weakness of existing attendance, the three combined checks, and the application.")

    s = S[3]
    fill(s, 5, BASE_PAPER, 13)
    set_paras(shape(s, 7), "Method")
    fill(s, 8, "A Python desktop application with a Tkinter interface. A webcam captures about 25 to 30 images per student; a Haar cascade detects faces and an LBPH recogniser matches them. Attendance goes to subject-wise CSV files, with manual entry and spoken feedback.", 15)
    set_paras(shape(s, 10), "Data and reported results")
    fill(s, 11, "Tested in classroom simulations with a small group. Reported: about 92 to 95% recognition under good frontal light, 5 to 8% false negatives, 0% false positives, under 2 s per face. No dataset size or protocol is given.", 15)
    set_paras(shape(s, 13), "Limits and future work")
    fill(s, 14, "Our reading: identity only, in a classroom where the camera is. No liveness, no location, no duplicate-identity check, no session window, and accuracy falls with poor light or angled faces. Future work is not stated in the paper.", 15)
    set_notes(s, "The base paper automates attendance with Haar plus LBPH face recognition. It is a useful baseline for identity, and silent about spoofing and location.")

    s = S[4]
    rows = [["Work", "Approach", "Data", "Reported result"],
            ["Dharmaraj et al., 2025 (base) [1]", "Haar cascade + LBPH face recognition, Tkinter GUI, CSV logs", "Classroom simulations", "About 92-95% recognition; 0% false positives"],
            ["Shandilya and Shah, 2026 [3]", "ArcFace (DeepFace), LBP texture, EAR blink, HSV dress-code check; FastAPI + Streamlit", "Not read beyond abstract", "96.2% recognition; 97.1% liveness; 92.4% dress code"],
            ["Salam et al., 2025 [5]", "Silent-face anti-spoofing: CNN + Fourier transform; Raspberry Pi 4", "Different light and distance", "Figures not read (abstract only)"],
            ["Classroom system, 2026 [4]", "ResNet101 anti-spoofing + ArcFace (title)", "Not retrieved", "Not retrieved"],
            ["Review, 2026 [2]", "Review of recognition, liveness and dress-code techniques (title)", "Not retrieved", "Not retrieved"],
            ["Works surveyed in [1]", "LBPH, CNN+SVM/KNN, VGG16/ResNet transfer, FaceNet+MTCNN", "Various", "Over 99% (controlled light); falls with poor light, masks"]]
    sz = table_size(rows, 1164, 460, [0.22, 0.34, 0.18, 0.26], 13, 9)
    set_table(shape(s, 6), rows, sz, header_size=sz)
    set_notes(s, "Only the abstracts of [3] and [5] could be read, and [2] and [4] could not be retrieved, so we do not claim their methods or results beyond what is shown.")

    s = S[5]
    set_paras(shape(s, 2), "Supporting evidence: a matching photo passes face recognition alone")
    cats = [CH_LABEL[c] for c in CHALLENGES]
    face_alone = [round(100 * FV["at_identity_threshold"]["tar"], 1)] * 8
    live = [round(100 * LV["C_static_replay"][c]["rate"], 1) for c in CHALLENGES]
    ch = replace_chart(s, 4, cats, [("A photo of the right person, face match only (%)", face_alone), ("Same photo, with a live challenge (%)", live)], number_format="0.0")
    style_chart(ch, "Acceptance of a still photo of the victim, by check (%)", ["E4572E", "2F9E6D"], number_format="0", axis_title="% accepted", label_pt=7)
    set_paras(shape(s, 8), "Takeaway")
    fill(s, 9, "At the deployed threshold (0.40) a different photo of the same person is accepted 98.7% of the time (LFW pairs). A live challenge rejects a single still photo for seven of eight actions; only 'close eyes' lets 8.2% through.", 15)
    set_paras(shape(s, 13), "Caveat")
    fill(s, 14, "Measured by us on public LFW still images with our own pipeline, not on the base paper's system; no live volunteers. A rotated photo or several photos can still pass (see the later reviews).", 15)
    set_notes(s, "This is our own measurement: face matching alone would admit a stolen photo; a random challenge changes that for a single photo.")

    s = S[6]
    set_paras(shape(s, 2), "Existing system and current approaches")
    steps = ["Call roll / pass sheet", "Mark on paper", "Type into register", "Share QR code or link", "Friend marks for absentee", "Reconcile by hand"]
    for sid, t in zip((5, 8, 11, 14, 17, 20), steps):
        fill(s, sid, t, 14)
    set_paras(shape(s, 24), "Manual and shared-link methods")
    fill(s, 25, "Roll calls and paper registers use class time and invite errors and proxy marking. QR codes and attendance links are quick, but anyone who receives the link can use it from anywhere.", 16)
    set_paras(shape(s, 29), "Biometric and face-recognition methods")
    fill(s, 30, "Fingerprint devices need hardware and queues. Camera-based face recognition confirms identity, but without a liveness check a photo or video can still be shown.", 16)
    set_notes(s, "How attendance is done today, and where each method breaks.")

    s = S[7]
    set_paras(shape(s, 2), "Limitations of existing methods")
    for tid, bid, ttl, body in [
        (7, 8, "Slow and error-prone", "Manual roll calls waste class time and allow proxy marking and typing errors."),
        (12, 13, "Shareable", "QR codes and links can be forwarded to students who are not in the room."),
        (17, 18, "Spoofable", "Face recognition without liveness accepts a photo; one person can mark several identities if duplicates are not checked."),
        (22, 23, "No presence proof", "Existing systems do not check that the student is physically inside the classroom area.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Four limitations taken from the project's literature survey.")

    s = S[8]
    gaps = [("Presence", "Face-recognition attendance systems verify identity but rarely check that the student is inside the classroom area [1], [3]."),
            ("Liveness", "A random facial challenge before the face match is uncommon in classroom attendance; photos can otherwise pass."),
            ("Duplicates", "Few systems compare each new face with the faces already recorded in the same session."),
            ("Evidence", "The base paper reports small, informal tests; reviewed work gives few measurements of spoof resistance."),
            ("Application", "No single tool joins time-bound sessions, geo-fence, device binding, liveness, duplicate check and Excel reports.")]
    for (lid, tid), (lab, txt) in zip([(5, 6), (8, 9), (11, 12), (14, 15), (17, 18)], gaps):
        set_paras(shape(s, lid), lab)
        fill(s, tid, txt, 17)
    set_notes(s, "The gap is the combination of location, liveness and duplicate-identity checks within one time-bound session.")

    s = S[9]
    fill(s, 5, "Given a faculty-defined session (course, time window, geo-fence and roll-number patterns) and students who enrolled a reference face, accept a student's attendance only if the session is live, the student and device are eligible, the GPS position is inside the fence, a random live challenge is performed, and the face matches the enrolled face and no other face already recorded in the session; and give faculty the records, a verification log and an Excel export.", 21)
    set_paras(shape(s, 8), "Session, GPS fix, two photos, credentials")
    set_paras(shape(s, 11), "Attendance record, reason for every block, Excel")
    set_paras(shape(s, 14), "Time-bound, device-bound, fully logged")
    set_notes(s, "The problem statement: inputs, outputs and the extra properties.")

    s = S[10]
    fill(s, 4, ["Enable faculty to create time-bound attendance sessions with configurable geo-fences.",
                "Authenticate students securely with JWT, device identification and roll-number restrictions.",
                "Verify physical presence by checking GPS location against the session boundary.",
                "Detect liveness through a random live facial challenge using MediaPipe.",
                "Prevent duplicate identities by matching InsightFace embeddings within each session.",
                "Provide attendance summaries and Excel export for faculty.",
                "Verify the system with automated tests and measure its checks on a public benchmark."], 18, bullet=True)
    set_notes(s, "Six objectives from the project definition plus a seventh on verification.")

    s = S[11]
    for tid, bid, ttl, body in [
        (7, 8, "Session and link", "Faculty set course, window, centre, radius and allowed rolls; a unique link and QR code are shared."),
        (12, 13, "Eligibility and location", "Session live, roll allowed, not already marked, device bound; then GPS inside radius + 15 m."),
        (17, 18, "Live challenge", "A random head, mouth or eye action checked on two photos with Face Mesh landmark metrics."),
        (22, 23, "Face verification", "ArcFace embeddings: same person in both photos, matches enrolment, not a duplicate in the session.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 14)
    set_paras(shape(s, 27), "Web application")
    fill(s, 28, "React front end, FastAPI back end, SQLite; records, map, verification log, analytics and Excel export.", 16)
    set_notes(s, "The approach: four verification stages in sequence, inside a web application.")

    s = S[12]
    put_picture(s, DIAG / "01_system_architecture.png", (215, 118, 790, 482), remove_ids=(4,), name="System architecture diagram")
    set_paras(shape(s, 5), "Proposed architecture: React front end, FastAPI back end with verification services, SQLite and local AI models")
    set_notes(s, "Browser to React to FastAPI. The verification logic sits in the service layer; the AI models run locally on CPU.")

    s = S[13]
    for tid, bid, ttl, body in [
        (7, 8, "Proxy attendance blocked", "A forwarded link, a photo or a second identity should fail a check (to be measured)."),
        (12, 13, "Presence proven", "Only students inside the geo-fence can complete a check-in."),
        (17, 18, "Fast check-in", "A student completes the checks in seconds on a phone."),
        (22, 23, "Faculty control", "Time window, allowed rolls, device reset, close session."),
        (27, 28, "Usable records", "Live list, map, verification log and Excel with absentees."),
        (32, 33, "Honest evaluation", "Thresholds measured on a public benchmark; limits reported plainly.")]:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 16)
    set_notes(s, "Expected outcomes are stated as targets here and are measured in the later reviews.")

    s = S[14]
    fill(s, 4, [f"[{i}] {r}" for i, r in enumerate(REFERENCES[:8], 1)], 13, min_pt=8.5, bullet=True, gap_pt=4)
    set_notes(s, "References [1] to [8]; the full list is in the report.")
    set_notes(S[15], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
