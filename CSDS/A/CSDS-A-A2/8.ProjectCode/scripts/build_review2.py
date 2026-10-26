"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size  # noqa: E402
from doc_common import (BATCH, CHARTS, COLLEGE, DEPT, DIAG, FV, GUIDE, GUIDE_DESIG, LT, PRODUCT, SHOTS, STUDENTS, YEAR)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"
SUB = "A Smart Geo-Fenced Attendance System with Live Challenge Verification"


def px(v):
    return Emu(int(v * 9525))


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", SUB, "Batch A2", "geo-fencing, liveness, face verification, attendance"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A2"

    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), SUB)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A2 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers design, methodology and progress of Attendance Magic.")

    s = S[1]
    for sid, t in [(21, "Verification components"), (18, "Inputs and preparation"), (27, "Data flow, modules, database"), (36, "Preliminary results and screens"), (15, "Evaluation data")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts: recap, design, methodology, data, components, decision logic, data flow, stack, progress, results, testing and next steps.")

    s = S[2]
    fill(s, 8, "Attendance by roll call or a shared link is slow, error-prone and easy to fake: a forwarded link, a photo or a friend can mark an absent student present.", 16)
    fill(s, 13, "Time-bound sessions with a geo-fence, a random live facial challenge and duplicate-face checks, delivered by a FastAPI and React web application.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Session link + sign-in (JWT, device)", "Eligibility + GPS geo-fence", "Random live challenge (Face Mesh)", "Face match + duplicate check", "Record, log, Excel export"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1: the problem and the proposed system as a five-step pipeline.")

    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 19 + Vite 8 (port 5102): 12 routes, camera or upload, Geolocation, Leaflet map.",
                "API: FastAPI (port 8102): 27 endpoints; JWT bound to the device; server-side check-in wizard state.",
                "Core services: geofence, liveness (MediaPipe Face Mesh), face_match (InsightFace ArcFace), eligibility, export.",
                "Data: SQLite with 5 tables and a verification-attempt log.",
                "Original Streamlit UI kept on the same core."], 16, bullet=True)
    set_notes(s, "The browser talks only to the Vite server, which proxies /api to FastAPI. The API is thin: every rule lives in core/, shared with the Streamlit interface.")

    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Create session", "Share link / QR", "Sign in + eligibility", "GPS geo-fence", "Random live challenge", "Face match + duplicate", "Record + export"]):
        fill(s, sid, t, 14)
    set_paras(shape(s, 25), "Staged and logged")
    fill(s, 26, "Each stage can stop the check-in with a plain reason, and every attempt (pass or block) is stored with its stage and reason.", 15)
    set_paras(shape(s, 28), "Rules in one place")
    fill(s, 29, "The REST API and the Streamlit UI both call core/, so there is a single implementation of every check.", 15)
    set_paras(shape(s, 31), "Configurable")
    fill(s, 32, "Radius, tolerance, GPS accuracy, thresholds and device limits are environment settings with documented defaults.", 15)
    set_notes(s, "End-to-end workflow in seven steps. A failed stage never writes an attendance record.")

    s = S[5]
    set_paras(shape(s, 2), "Evaluation data and input sources")
    ch = replace_chart(s, 4, ["Same-person pairs scored", "Different-person pairs scored", "Pairs dropped (no single face)"], [("LFW verification pairs", [FV["genuine_pairs"], FV["impostor_pairs"], FV["pairs_dropped_no_single_face"]])])
    style_chart(ch, "LFW protocol pairs (6,000 in total)", ["3F3FBF"], legend=False, label_pt=11, axis_title="pairs")
    rows = [["Data", "Size", "Use"],
            ["LFW benchmark, pairs.txt protocol", "13,233 images; 6,000 pairs (3,000 + 3,000); 7,701 unique images", "Threshold evaluation"],
            ["LFW stills, random sample", "1,000 photos; 157 identities with 10+ photos", "Liveness and attack tests"],
            ["Demo data", "5 accounts, 2 sessions; LFW photos as stand-ins", "Screens, smoke test"],
            ["Student photos / camera trials", "None collected", "Not done"],
            ["Training data", "None: models are pre-trained", "No train/test split"]]
    sz = table_size(rows, 530, 400, [0.34, 0.42, 0.24], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.34, 0.42, 0.24], header_size=sz)
    set_notes(s, "No model is trained in this project, so there is no training set. LFW is a public benchmark used only to evaluate the thresholds of the verification chain.")

    s = S[6]
    set_paras(shape(s, 2), "Inputs and preparation")
    set_paras(shape(s, 7), "Photos")
    fill(s, 8, "JPEG or PNG from the camera or an upload; decoded with OpenCV and shrunk so the longest side is at most 1280 px. Exactly one face is required in each frame; the first (neutral) frame must be frontal (|yaw| <= 0.10) with the mouth closed.", 16)
    set_paras(shape(s, 12), "Location")
    fill(s, 13, "Browser GPS latitude, longitude and accuracy (or a manual coordinate in test mode). A fix worse than 150 m is rejected; the distance is a haversine on a 6,371 km sphere.", 16)
    set_paras(shape(s, 17), "Identity and roll data")
    fill(s, 18, "Each student enrols one reference photo, stored as a 512-dimensional ArcFace embedding. Roll numbers are matched with wildcard patterns (for example 23K91A67*) and an explicit list also gives the absentees.", 16)
    set_notes(s, "Preparation here means decoding and validating the photos, the GPS fix and the roll numbers before the checks run.")

    s = S[7]
    set_paras(shape(s, 2), "Verification components (pre-trained, not trained here)")
    put_picture(s, DIAG / "05_dfd_level2.png", (58, 128, 745, 500), remove_ids=(4,), name="Face verification data flow")
    rows = [["Component", "Detail"], ["MediaPipe Face Mesh", "468 landmarks; yaw, pitch, roll, MAR, EAR"], ["InsightFace buffalo_l", "SCRFD detector + ArcFace R50, 512-d, CPU"],
            ["Thresholds", "consistency 0.45, identity 0.40, duplicate 0.45"], ["Liveness", "8 random challenges, 3 tries"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "No network is trained: thresholds on landmarks and cosine similarity decide.", 11)
    set_notes(s, "The data-flow diagram of the face-verification process. The two AI components are used as published; the project contribution is the chain, the thresholds and their measurement.")

    s = S[8]
    set_paras(shape(s, 2), "Configuration and threshold analysis (nothing is trained)")
    rows = [["Setting", "Value"], ["Geo-fence", "radius 60 m + 15 m tolerance; GPS accuracy <= 150 m"], ["Challenges", "yaw 0.12, pitch 0.08, roll 12 deg, MAR 0.35, EAR 0.18"],
            ["Neutral frame", "|yaw| <= 0.10, MAR <= 0.25, one face"], ["Face thresholds", "consistency 0.45, identity 0.40, duplicate 0.45"],
            ["Session / token", "JWT 120 min, device-bound; 1 device per student"], ["Liveness tries", "3, then the wizard locks"], ["Hardware", "Intel i5-12500, 15.7 GB, CPU only"]]
    set_table(shape(s, 10), rows, 12, header_size=12)
    put_picture(s, CHARTS / "threshold_sweep.png", (672, 232, 550, 330), remove_ids=(5,), name="Threshold sweep chart")
    set_paras(shape(s, 7), "Error rates against threshold on LFW: 0.40 gives 1.3% false rejects and 0 false accepts", 11)
    set_notes(s, "The settings are defaults from config.py. The chart is our measurement on 4,307 LFW pairs: no different-person pair reached 0.40 (highest 0.246).")

    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (57, 124, 420, 520), remove_ids=(4,), name="Decision flowchart")
    for bid in (7, 10, 13):
        shape(s, bid).height = px(80)
    set_paras(shape(s, 6), "Marked present")
    fill(s, 7, "Only when eligibility, geo-fence, challenge, frame consistency, identity and duplicate checks all pass is a record written (one per student per session).", 16)
    set_paras(shape(s, 9), "Blocked")
    fill(s, 10, "The first failing stage stops the check-in with a reason; three failed challenges lock the wizard until the student starts over.", 16)
    set_paras(shape(s, 12), "Logged")
    fill(s, 13, "Every attempt is stored with stage, result, reason and device, and shown in the faculty verification log.", 16)
    set_notes(s, "Decision logic: a strict sequence of checks; the order is the order of the wizard.")

    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 530, 330), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "08_use_case_diagram.png", (605, 118, 616, 330), name="Use-case diagram")
    set_paras(shape(s, 5), "DFD level 1 (left) and use-case diagram (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(452)
    rows = [["Module", "Responsibility"], ["auth, device", "PBKDF2 passwords, JWT bound to a SHA-256 device hash"],
            ["geofence, liveness, face_match", "haversine check, Face Mesh challenge metrics, ArcFace consistency / identity / duplicate"],
            ["services, export, api", "eligibility, pipeline and records; 3-sheet Excel; 27 REST endpoints"]]
    set_table(shape(s, 12), rows, 12, widths=[0.28, 0.72], header_size=12, row_h=34)
    shape(s, 12).top = px(485)
    set_notes(s, "Level 1 shows eight processes and four data stores; the use-case diagram shows what faculty and students can do.")

    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 450, 340), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (520, 118, 701, 340), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 5 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(460), px(1164)
    rows = [["Table", "Key fields", "Constraints"],
            ["attendance_records", "session, student, distance_m, challenge, liveness_score, identity_similarity, max_duplicate_similarity, face_embedding", "UNIQUE(session_id, student_id)"],
            ["device_bindings", "user_id, device_hash, label, first/last seen", "UNIQUE(user_id, device_hash)"],
            ["verification_attempts", "session, roll, stage, passed, reason, device", "one row per stage attempt"]]
    set_table(shape(s, 13), rows, 11, widths=[0.2, 0.52, 0.28], header_size=11, row_h=30, total_px=1164)
    t = shape(s, 13)
    t.left, t.top, t.width = px(57), px(492), px(1164)
    set_notes(s, "Five tables created by SQLAlchemy. The UML diagram shows the entities and the decision classes of core/.")

    s = S[12]
    cards = [(7, 8, "Language", "Python 3.11, JavaScript (ES modules, JSX)"), (12, 13, "Computer vision", "MediaPipe 0.10.21, InsightFace 2.0 (ONNX Runtime), OpenCV 4.10, NumPy"),
             (17, 18, "Frontend", "React 19, Vite 8, React Router, Leaflet, Recharts, Framer Motion, GSAP"), (22, 23, "Backend", "FastAPI, Pydantic, SQLAlchemy 2, PyJWT, Streamlit (original UI)"),
             (27, 28, "Storage and reports", "SQLite, pandas + openpyxl (Excel), qrcode"), (32, 33, "Testing", "pytest (23 tests), 41-check API smoke test, Playwright captures")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Everything runs offline on a laptop CPU; only the map tiles need the internet.")

    s = S[13]
    rows = [["Component", "Status"], ["Sessions with geo-fence, window, roll patterns, link and QR", "Completed"], ["JWT sign-in bound to device; device bindings", "Completed"],
            ["GPS geo-fence check", "Completed"], ["Random live challenge (8 actions)", "Completed"], ["Face consistency, identity and duplicate checks", "Completed"],
            ["Faculty records, map, log, analytics, Excel export", "Completed"], ["Automated tests (23 + 41 checks) and LFW evaluation", "Completed"], ["Camera trials with volunteers, HTTPS phone test", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.78, 0.22], header_size=13)
    for vid, lid, v, lab in [(6, 7, "27", "REST endpoints"), (9, 10, "5", "database tables"), (12, 13, "23 + 41", "pytest tests + smoke checks, all passing")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Every planned feature is implemented. Open items are volunteer camera trials and a phone test over HTTPS.")

    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: threshold behaviour on LFW pairs")
    ths = [0.30, 0.40, 0.45, 0.50, 0.60]
    sw = {r["threshold"]: r for r in FV["sweep"]}
    ch = replace_chart(s, 4, [f"{t:.2f}" for t in ths], [("Same person accepted (%)", [round(100 * sw[t]["tar"], 1) for t in ths])], number_format="0.0")
    style_chart(ch, "Same-person pairs accepted by similarity threshold (%)", ["3F3FBF"], number_format="0.0", legend=False, label_pt=11, axis_title="% accepted")
    for vid, lid, v, lab in [(6, 7, "98.7%", "same-person pairs accepted at 0.40"), (9, 10, "0 of 2,166", "different-person pairs accepted (max 0.246)"), (12, 13, "0.999", "ROC AUC; equal error rate 0.30%")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Measured on the standard LFW verification protocol with our own pipeline: 4,307 scored pairs; 1,693 pairs were dropped because a photo failed the single-face rule.")

    s = S[15]
    put_picture(s, SHOTS / "24_checkin_marked.png", (91, 128, 507, 476), remove_ids=(4,), name="Check-in result screenshot")
    put_picture(s, SHOTS / "13_session_records.png", (679, 128, 512, 476), remove_ids=(7,), name="Session records screenshot")
    set_paras(shape(s, 6), "Student result: marked present, 8 m from the centre, liveness 1.00, identity 0.86", 11)
    set_paras(shape(s, 9), "Faculty view: live session, 3 present, 7 blocked attempts", 11)
    set_notes(s, "Real captures of the running application with demo data (LFW photos as stand-ins for student photographs).")

    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Passwords, JWT bound to device", "2", "Passed"], ["Haversine and geo-fence", "2", "Passed"], ["Liveness maths, 11 challenge cases, neutral frame", "13", "Passed"],
            ["Face logic, roll patterns, services, model test", "4", "Passed"], ["Streamlit end-to-end flows", "2", "Passed"], ["API smoke test (real models)", "41", "41 / 41 pass"]]
    set_table(shape(s, 18), rows, 12, widths=[0.58, 0.14, 0.28], header_size=12)
    fill(s, 5, ["A forwarded link, an outside position, a static photo, another person's face and a duplicate face were each blocked at their stage.",
                "A rotated still photo passes the tilt challenges and several photos of the victim can pass all challenges: a limitation to address.",
                "One check-in takes about 0.56 s of server time on a CPU (median).",
                "15.7% of LFW photos fail the single-face rule, so retakes are expected."], 16, bullet=True)
    set_notes(s, "Testing so far: automated tests and scripted attacks. The weaknesses found by our own attack experiments are reported openly.")

    s = S[17]
    cards = [(7, 8, "No trained model", "Both AI models are pre-trained; evaluation uses a public benchmark and attack experiments."),
             (12, 13, "Photo spoofing", "A random challenge defeats a single photo (0% pass for 7 of 8 actions) but not rotated or multiple photos."),
             (17, 18, "Two-photo capture", "Two still photos replace a video stream; simple, but weaker than continuous liveness."),
             (22, 23, "Camera and GPS need HTTPS", "On phones they work only over HTTPS; localhost is fine on a laptop."),
             (27, 28, "Single-face rule", "A second face above 50% of the main face width rejects the photo; 15.7% of LFW photos failed."),
             (32, 33, "Duplicate enrolment", "A proxy enrolled with the same face passes identity, and is caught by the duplicate check.")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 14)
    set_notes(s, "Six challenges and how each was handled or recorded.")

    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["Camera trials with volunteers", "Phone test over HTTPS", "Final report", "Final presentation and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Collect real selfie captures to measure the genuine pass rate of each challenge and the retake rate.",
                 "Test on phones over HTTPS with real GPS and measure drift indoors.",
                 "Strengthen liveness: reject rotated photos, add texture or depth cues, time-limit the two photos.",
                 "Document every diagram, test and result in the report."], 18, bullet=True)
    set_notes(s, "What remains: real-user trials and a stronger liveness check, then the report and the final demo.")
    set_notes(S[19], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
