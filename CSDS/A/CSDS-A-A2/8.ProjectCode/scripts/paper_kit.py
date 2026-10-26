"""Two-column IEEE-style page helpers (python-docx): fonts, headings, figures, tables, author block, references."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from doc_common import REFERENCES, STUDENTS  # noqa: E402

FONT = "Times New Roman"
COL_W = 3.43  # inches
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")
TITLE = ""
ABSTRACT = ""
KEYWORDS = ""


def font(run, size=10, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


class Paper:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        s = self.d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
        s.left_margin = s.right_margin = Inches(0.69)
        s.header_distance = s.footer_distance = Inches(0.4)
        self.fig = 0
        self.tab = 0

    # ---------------------------------------------------------------- text
    def runs(self, p, text, size=10, italic=False, bold=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2]); font(r, size, bold=True, italic=italic)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); font(r, size, bold=bold, italic=not italic)
            else:
                r = p.add_run(part); font(r, size, bold=bold, italic=italic)

    def para(self, text, indent=True, size=10, align="j", after=0, before=0, keep=False):
        p = self.d.add_paragraph()
        self.runs(p, text, size)
        p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.keep_with_next = keep
        return p

    def h1(self, text):
        p = self.d.add_paragraph()
        r = p.add_run(text.upper()); font(r, 10, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True

    def h2(self, text):
        p = self.d.add_paragraph()
        r = p.add_run(text); font(r, 10, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True

    def eq(self, text, num):
        p = self.d.add_paragraph()
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Inches(COL_W / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Inches(COL_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t"); font(r, 10)
        for part in re.split(r"(_\{[^}]*\})", text):
            if part.startswith("_{"):
                sub = p.add_run(part[2:-1]); font(sub, 10, italic=True); sub.font.subscript = True
            elif part:
                run = p.add_run(part); font(run, 10, italic=True)
        r2 = p.add_run(f"\t({num})"); font(r2, 10)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        for i, t in enumerate(items, 1):
            p = self.d.add_paragraph()
            self.runs(p, f"{i})\t{t}")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))

    # ---------------------------------------------------------------- floats
    def figure(self, path, caption, width=3.4):
        self.fig += 1
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(str(path), width=Inches(width))
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(5)
        r = c.add_run(f"Fig. {self.fig}. {caption}"); font(r, 8)
        return self.fig

    def table(self, rows, caption, widths, size=7):
        self.tab += 1
        roman = ["I", "II", "III", "IV"][self.tab - 1]
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_before = Pt(4)
        r = c.add_run(f"TABLE {roman}"); font(r, 8)
        r.add_break()
        r2 = c.add_run(caption.upper()); font(r2, 8)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for ci, txt in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(COL_W * widths[ci])
                tcpr = cell._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for side in ("left", "right"):
                    m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), "40"); m.set(qn("w:type"), "dxa"); mar.append(m)
                tcpr.append(mar)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = ri < len(rows) - 1
                run = p.add_run(txt); font(run, size, bold=(ri == 0))
                if ri == 0:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E7E9EF"); tcpr.append(shd)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        for r_ in sp.runs:
            r_.font.size = Pt(2)
        sp.paragraph_format.line_spacing = Pt(3)
        return self.tab

    # ---------------------------------------------------------------- front block and columns
    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(TITLE); font(r, 24)
        p.paragraph_format.space_after = Pt(8)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [("[GUIDE NAME]", "[DESIGNATION], Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Department of Computer Science and Engineering (Data Science)", False),
                               ("[COLLEGE NAME]", False), ("[CITY]", False), ("[EMAIL]", False)):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line); font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        # continuous section with two columns
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        sectpr = sec._sectPr
        cols = sectpr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectpr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def abstract(self):
        p = self.d.add_paragraph()
        r = p.add_run("Abstract—"); font(r, 9, bold=True, italic=True)
        r = p.add_run(ABSTRACT); font(r, 9)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        q = self.d.add_paragraph()
        r = q.add_run("Index Terms—"); font(r, 9, bold=True, italic=True)
        r = q.add_run(KEYWORDS + "."); font(r, 9)
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        q.paragraph_format.space_after = Pt(2)

    def references(self):
        self.h1("References")
        for i, ref in enumerate(REFERENCES, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]\t{ref}"); font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


