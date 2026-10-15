"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import (BATCH, CHARTS, COLLEGE, CV, DIAG, GUIDE, GUIDE_DESIG, LAT, LR, M, PRODUCT, R, REFERENCES, RF, SHOTS, STUDENTS, T, TITLE, X, YEAR, p)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE (Data Science)  |  Batch A3  |  Final Review"


def px(v):
    return Emu(int(v * 9525))


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 9)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l)
    set_paras(shape(s, cap_ids[1]), cap_r)
    for cid in cap_ids:
        shape(s, cid).top = px(554)


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Final Review", TITLE, "Batch A3", "phishing detection, URL classification, crowdsourced threat intelligence"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A3"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    # 1 title
    s = S[0]
    set_paras(shape(s, 4), TITLE, 24)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A3")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of PhishGuard: a dynamic phishing URL detection system with machine learning and crowdsourced threat intelligence.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Decision Logic",
                "Design & Implementation", "Results, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["Phishing sites steal credentials by imitating trusted pages; new domains appear daily and live only hours. Blacklists miss zero-day URLs, fixed rules are easy to evade and models trained once go stale.",
                "PhishGuard implements the base paper's three stages: 47 offline URL features (23 lexical, 9 structural, 15 domain), classification with Random Forest, XGBoost and Logistic Regression, and refinement through crowdsourced threat intelligence.",
                "A confirmed community verdict, phishing or safe, overrides the model at once; otherwise the best model scores the URL with thresholds 0.35 / 0.60. Two net votes or a trusted feed confirm a URL; confirmed URLs retrain the model, which is hot-reloaded.",
                f"XGBoost reaches accuracy {p(X['accuracy'], 1)}, F1 {p(X['f1'], 1)} and ROC-AUC {X['roc_auc']:.3f} on 8,000 held-out URLs (5-fold CV F1 {p(CV['xgboost']['f1']['mean'], 1)}); a check takes {LAT['api_check_ms_mean']:.1f} ms through the API.",
                "Flask API, CLI and a seven-page React interface; 95 automated checks pass. Limits: one 40,000-URL dataset, URL-only features, anonymous votes.",
                "Keywords: phishing detection, URL classification, crowdsourced threat intelligence, XGBoost, adaptive security."], 19, bullet=True)
    set_notes(s, "Abstract: problem, approach, system and the headline results, with the limits stated.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [f"{REFERENCES[0]}",
                "Design: Extractor, Detection, Updating by crowdsourcing. Novel URLs are verified by a human, added to the training set, and the model is retrained.",
                "Data and model: PhiUSIIL, 235,804 URLs, 54 features (34 kept); LSTM with attention plus a dense branch.",
                "Reported: accuracy 99.8 %, F1 0.9988, 5-fold cross-validation, learning-curve simulation of crowd data growth, about 0.20 ms per URL.",
                "Limits: one dataset, crowd loop simulated, real-world validation, scalability, adversarial misuse and privacy left as future work."], 19, bullet=True)
    set_notes(s, "The base paper is a framework paper with a strong benchmark number; its crowd loop is evaluated by simulation.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Aldhuwayhi and Alromih, 2026 (base) [1]", "LSTM-attention + dense; crowd feedback loop", "PhiUSIIL, 235,804 URLs", "Acc. 99.8 %, F1 0.9988", "Crowd loop simulated only"],
            ["Vajrobol et al., 2024 [6]", "Mutual information + logistic regression", "PhiUSIIL", "99.97 % with 5 features", "No real-time test"],
            ["Etem and Teke, 2024 [7]", "t-SNE reduction + decision tree", "51-feature set", "99.7 %", "Single dataset"],
            ["Alnemari and Alshammari, 2023 [8]", "Classical classifiers", "UCI phishing domains", "97.3 % (RF)", "Static data"],
            ["Priya et al., 2024 [11]", "RF / LR browser extension", "11,055 URLs", "97 %", "Latency not discussed"],
            ["Assefa and Katarya, 2022 [14]", "Autoencoder", "Small URL set", "91.24 %", "Below supervised models"]]
    table(s, rows, [0.2, 0.24, 0.16, 0.17, 0.23], 14, 440)
    set_notes(s, "Rows [6] to [14] are as summarised in the base paper. Works [2] to [5] from the batch survey were recorded by title only.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Crowd loop only simulated; no live validation; no explanation of verdicts"],
            ["Supervised detectors [6], [7], [8]", "High accuracy on static datasets; no continuous retraining from new threat data"],
            ["Deployed tools [11], [12]", "Latency, load and retraining not discussed"],
            ["Unsupervised detectors [14]", "Suit novel sites but score lower and use small datasets"],
            ["Blacklists and manual reports", "Miss zero-day URLs; slow, manual updates"]]
    table(s, rows, [0.28, 0.72], 17, 440)
    set_notes(s, "The gap: a live, explainable loop that joins ML with community verdicts and measured evidence.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Phishing websites keep stealing credentials and financial data at scale.",
                "Blacklists cannot detect newly created, zero-day phishing URLs.",
                "Static rules and one-time-trained models fail as attack patterns evolve; false positives block legitimate sites.",
                "Problem: given a URL, give a verdict (phishing, suspicious, safe) with probability and explanation, using a trained classifier and community-confirmed verdicts that override it and feed its retraining, in real time through a web interface."], 20, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["Collect and preprocess a dataset of phishing and legitimate URLs.", "Extract lexical, structural and domain-based URL features.",
                "Train and compare Random Forest, XGBoost and Logistic Regression.", "Integrate crowdsourced threat intelligence for newly reported phishing URLs.",
                "Refine the model continuously as new threat data arrives.", "Provide real-time URL classification through a simple web interface."], 21, bullet=True)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["1. Extract - 47 offline features from the URL string; red flags explain them.",
                "2. Classify - Random Forest, XGBoost and Logistic Regression on one split; best F1 is deployed.",
                "3. Decide - a confirmed community verdict wins; else model score, reputation x0.4, pending votes +-0.15, thresholds 0.35 / 0.60.",
                "4. Report - users vote phishing or safe; net 2 confirms; OpenPhish and URLhaus confirm at once.",
                "5. Refine - confirmed URLs join the training data (x3); retrain; the detector hot-reloads.",
                "6. Use - Flask API, CLI and React interface with batch check, activity log and model comparison."], 19, bullet=False)

    # 10 architecture
    s = S[9]
    put_picture(s, DIAG / "01_system_architecture.png", (240, 178, 800, 424), remove_ids=(7,), name="System architecture")
    set_paras(shape(s, 8), "Layered architecture: users, React interface, Flask API with detection modules, local storage")

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "End-to-end workflow: prepare the models once, then check, collect community reports, retrain and reload")

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Features & Models")
    fill(s, 7, ["Data: 40,000 URLs from a public set (20,000 phishing, 20,000 legitimate), cleaned and split 80/20 (32,000 / 8,000), stratified, seed 42.",
                "47 offline features: 23 lexical, 9 structural, 15 domain; no DNS, WHOIS or page download.",
                "Random Forest: 300 trees. XGBoost: 400 trees, depth 8, learning rate 0.08. Logistic Regression: scaled, C 1.0.",
                "Selection: best F1 on the test set (ROC-AUC breaks ties) saved as best_model.joblib.",
                "Community store: SQLite (reports, intel, checks); consensus = 2 net votes; trusted feeds confirm at once.",
                f"Retraining: base data + confirmed URLs x3, same split; training takes {RF['train_seconds'] + X['train_seconds'] + LR['train_seconds']:.1f} s for the three models on this PC."], 19, bullet=True)
    set_notes(s, "Everything is computed from the URL string, so training and real-time checking use identical code.")

    # 13 decision workflow
    s = S[12]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (300, 178, 680, 424), remove_ids=(7,), name="Algorithm flowchart")
    set_paras(shape(s, 8), "Decision order: community verdict first, then model score with adjustments and thresholds")

    # 14 interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    put_picture(s, DIAG / "09_sequence_diagram.png", (330, 178, 620, 424), remove_ids=(7,), name="Sequence diagram")
    set_paras(shape(s, 8), "Sequence: check a URL, report it, and refine the model")

    # 15 database
    s = S[14]
    put_picture(s, DIAG / "06_er_diagram.png", (250, 178, 780, 424), remove_ids=(7,), name="ER diagram")
    set_paras(shape(s, 8), "Entity\u2013relationship diagram: 3 SQLite tables and 3 file stores")

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["Features", "47 offline features, red flags; robust to malformed ports and brackets"],
            ["Training", "Three models on one split; metrics, confusion matrix, importances in metrics.json"],
            ["Threat intelligence", "SQLite reports and intel; consensus 2; trusted feeds; host-level lookup; check log"],
            ["Detector", "Community first; model; reputation x0.4; pending votes; thresholds; hot reload under a lock"],
            ["API and UI", "Flask JSON API (16 routes), CLI, React interface with 7 pages"],
            ["Tests", "3 + 20 + 72 = 95 automated checks; Playwright captures"]]
    table(s, rows, [0.2, 0.8], 17, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Model Comparison (8,000 test URLs)")
    rows = [["Model", "Accuracy  |  Precision  |  Recall", "F1  |  ROC-AUC"]]
    for n, m in (("XGBoost (selected)", X), ("Random Forest", RF), ("Logistic Regression", LR)):
        rows.append([n, f"{p(m['accuracy'])}  |  {p(m['precision'])}  |  {p(m['recall'])}", f"{p(m['f1'])}  |  {m['roc_auc']:.4f}"])
    c = CV["xgboost"]
    rows.append(["XGBoost, 5-fold CV", f"{p(c['accuracy']['mean'])}  |  {p(c['precision']['mean'])}  |  {p(c['recall']['mean'])}", f"{p(c['f1']['mean'])}  |  {c['roc_auc']['mean']:.4f}"])
    table(s, rows, [0.26, 0.46, 0.28], 19, 420)
    set_paras(shape(s, 8), "Held-out 20 % of a 40,000-URL sample; the shipped numbers were re-checked from the saved models. Not comparable with the base paper's dataset.")
    set_notes(s, "XGBoost is best by F1. The 5-fold cross-validation result is slightly lower and very stable (std 0.18 points).")

    # 18 confusion matrix + ROC
    s = S[17]
    set_paras(shape(s, 4), "Confusion Matrices & ROC Curves")
    two_pics(s, CHARTS / "confusion_matrices.png", CHARTS / "roc_curves.png", "XGBoost: 3,745 TN, 255 FP, 311 FN, 3,689 TP", "ROC-AUC: XGBoost 0.981, Random Forest 0.978, LR 0.922")
    set_notes(s, "The error is balanced: 6.4 % of legitimate URLs are flagged and 7.8 % of phishing URLs are missed.")

    # 19 parameter analysis
    s = S[18]
    set_paras(shape(s, 4), "Threshold & Feature Analysis")
    two_pics(s, CHARTS / "threshold_sweep.png", CHARTS / "feature_ablation.png", "F1 is flat between thresholds 0.40 and 0.60; 0.35 marks 'suspicious'", "Lexical features give F1 89.2 %; all 47 give 92.8 %")
    set_notes(s, "The threshold sweep justifies 0.35 / 0.60; the ablation shows lexical features carry most of the signal and the model does not depend on its two strongest features alone.")

    # 20 performance & testing
    s = S[19]
    rows = [["Measurement", "Result"],
            ["Response time", f"/api/check mean {LAT['api_check_ms_mean']:.1f} ms, p95 {LAT['api_check_ms_p95']:.1f} ms; feature extraction {LAT['feature_extraction_ms_per_url']:.3f} ms per URL"],
            ["Training time", f"RF {RF['train_seconds']} s, XGBoost {X['train_seconds']} s, LR {LR['train_seconds']} s on 32,000 URLs"],
            ["Automated tests", "95 passed, 0 failed (3 original, 20 smoke, 72 extended: unit 23, integration 8, functional 14, validation 17, security 8, performance 2)"],
            ["Crowd loop (live run)", "Zero-day page 9.5 % -> phishing 97 % after 2 reports; feed sync +800 URLs; retrain v1 -> v2"],
            ["Retraining effect", "Adding 250 to 2,000 confirmed URLs changes F1 by under 0.1 point on unseen URLs"]]
    table(s, rows, [0.22, 0.78], 18, 430)
    set_notes(s, "Not measured: concurrent-user load, memory use, other hardware, user acceptance, authentication.")

    # 21 comparison
    s = S[20]
    rows = [["System", "Data", "Method", "Crowd loop", "Reported / measured", "Extras"],
            ["Base paper [1]", "PhiUSIIL, 235,804 URLs", "LSTM-attention + dense, 34 features", "Human-verified, retrained periodically; simulated", "Accuracy 99.8 %, F1 0.9988; 0.20 ms per URL", "-"],
            ["PhishGuard", "40,000-URL sample", "XGBoost on 47 offline features (RF, LR compared)", "Vote consensus + trusted feeds; applied at once; retrain + hot reload", f"Accuracy {p(X['accuracy'], 1)}, F1 {p(X['f1'], 1)}; {LAT['api_check_ms_mean']:.1f} ms per check", "Red flags, React UI, CLI, activity log"]]
    table(s, rows, [0.1, 0.15, 0.2, 0.22, 0.19, 0.14], 13, 400)
    set_paras(shape(s, 8), "Different data, features and models: the numbers are not a head-to-head benchmark. The gap of 6.9 points is not explained experimentally.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Check URL")
    two_pics(s, SHOTS / "03_check_phishing.png", SHOTS / "05_check_community_override.png", "Phishing verdict with gauge and risk scale", "Community overrides the model: 9.5 % -> 97.0 %")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Report and Threat Intelligence")
    two_pics(s, SHOTS / "08_report_form.png", SHOTS / "09_intel.png", "Report form with live vote tally and consensus rules", "Threat-intelligence board: statuses and votes")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Models and Retraining")
    two_pics(s, SHOTS / "11_models.png", SHOTS / "20_models_after_retrain.png", "Model comparison, version 1", "After retraining on 807 confirmed URLs: version 2")

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built PhishGuard: a working, explainable implementation of the base paper's three-stage design.",
                f"XGBoost reaches F1 {p(X['f1'], 1)} and ROC-AUC {X['roc_auc']:.3f} on held-out URLs; 5-fold CV F1 {p(CV['xgboost']['f1']['mean'], 1)}; {LAT['api_check_ms_mean']:.0f} ms per check.",
                "A confirmed community verdict acts at once in both directions; the retraining loop with hot reload runs end to end on real feed data.",
                "Retraining on random in-distribution URLs does not change the average F1; the loop's value is the override.",
                "Limits: one 40,000-URL dataset (not comparable with 99.8 %), URL-only features, anonymous votes, no sign-in, no user-acceptance test."], 20, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Evaluate on a larger, more recent dataset and on URLs collected after training.",
                "Add sign-in, rate limiting, reporter reputation and an administrator review queue.",
                "Better reputation data than a hand-written list; smarter host-level override.",
                "Optional sandboxed page features for uncertain URLs; browser extension and mail plug-in.",
                "Study the crowd loop under noisy reports and new campaigns; reproduce the base paper's network for a like-for-like comparison."], 20, bullet=True)

    # 27 references
    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES[:14], 1)], 13, min_pt=7.5, bullet=True, gap_pt=2)

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
