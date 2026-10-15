"""Build 19.IEEE_Paper.docx: a six-page, two-column IEEE conference paper on PhishGuard.

    python scripts/build_paper.py        then   python scripts/paper_finish.py   (page count + PDF through Word)

Every number comes from docs/REPORT_FACTS.md (via doc_common.py).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from doc_common import BATCH, CV, GUIDE, GUIDE_DESIG, LAT, M, R, REFERENCES, ROOT, SHOTS, STUDENTS  # noqa: E402

OUT = BATCH / "19.IEEE_Paper.docx"
PAPER = ROOT / "docs" / "paper"
FONT = "Times New Roman"
COL_W = 3.43  # inches
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")

TITLE = "PhishGuard: A Dynamic Phishing URL Detection System with Machine Learning and Crowdsourced Threat Intelligence"
ABSTRACT = (
    "Phishing remains a leading route to credential theft, and blacklists and one-time-trained classifiers miss newly created URLs. "
    "This paper presents PhishGuard, a phishing URL detection system that follows the three-stage design of a recent base paper: feature "
    "extraction, machine-learning classification and crowdsourced threat-intelligence refinement. Forty-seven offline lexical, structural and "
    "domain features are extracted from the URL string alone, and Random Forest, XGBoost and Logistic Regression are compared on a balanced "
    "40,000-URL sample. A community store confirms a URL after two net votes or one trusted-feed report; a confirmed verdict overrides the "
    "model at once and feeds a retraining loop with hot reload. On 8,000 held-out URLs XGBoost reaches accuracy "
    f"{M['models']['xgboost']['accuracy'] * 100:.2f} %, F1 {M['models']['xgboost']['f1'] * 100:.2f} % and ROC-AUC {M['models']['xgboost']['roc_auc']:.3f}, with five-fold "
    f"cross-validated F1 of {CV['xgboost']['f1']['mean'] * 100:.2f} %; a check takes {LAT['api_check_ms_mean']:.1f} ms. In a live run two reports turned a page the model had scored at "
    "9.5 % into a phishing verdict, whereas retraining with random in-distribution URLs did not change the average score. The dataset is small "
    "and from one source, so the results are not comparable with the base paper's 99.8 %, and no user study was carried out."
)
KEYWORDS = "adaptive security, crowdsourced threat intelligence, machine learning, phishing detection, URL classification, XGBoost"
KEYS = {"base": 0, "ijsr": 2, "vajrobol": 5, "etem": 6, "alnemari": 7, "priya": 10, "sawant": 11, "fan": 12, "assefa": 13, "huang": 14, "lazaros": 15,
        "prasad": 16, "ma": 17, "sahingoz": 18, "krombholz": 19, "breiman": 20, "chen": 21, "pedregosa": 22, "dataset": 24, "flask": 26, "react": 27}


def font(run, size=10, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)


class Paper:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        s = self.d.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
        s.left_margin = s.right_margin = Inches(0.69)
        s.header_distance = s.footer_distance = Inches(0.4)
        self.fig = 0
        self.tab = 0

    # ---------------------------------------------------------------- text
    def runs(self, p, text, size=10, italic=False, bold=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2]); font(r, size, bold=True, italic=italic)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); font(r, size, bold=bold, italic=not italic)
            else:
                r = p.add_run(part); font(r, size, bold=bold, italic=italic)

    def para(self, text, indent=True, size=10, align="j", after=0, before=0, keep=False):
        p = self.d.add_paragraph()
        self.runs(p, text, size)
        p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.keep_with_next = keep
        return p

    def h1(self, text):
        p = self.d.add_paragraph()
        r = p.add_run(text.upper()); font(r, 10, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True

    def h2(self, text):
        p = self.d.add_paragraph()
        r = p.add_run(text); font(r, 10, italic=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True

    def eq(self, text, num):
        p = self.d.add_paragraph()
        ts = p.paragraph_format.tab_stops
        ts.add_tab_stop(Inches(COL_W / 2), WD_TAB_ALIGNMENT.CENTER)
        ts.add_tab_stop(Inches(COL_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t"); font(r, 10)
        for part in re.split(r"(_\{[^}]*\})", text):
            if part.startswith("_{"):
                sub = p.add_run(part[2:-1]); font(sub, 10, italic=True); sub.font.subscript = True
            elif part:
                run = p.add_run(part); font(run, 10, italic=True)
        r2 = p.add_run(f"\t({num})"); font(r2, 10)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        for i, t in enumerate(items, 1):
            p = self.d.add_paragraph()
            self.runs(p, f"{i})\t{t}")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))

    # ---------------------------------------------------------------- floats
    def figure(self, path, caption, width=3.4):
        self.fig += 1
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(str(path), width=Inches(width))
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(5)
        r = c.add_run(f"Fig. {self.fig}. {caption}"); font(r, 8)
        return self.fig

    def table(self, rows, caption, widths, size=7):
        self.tab += 1
        roman = ["I", "II", "III", "IV"][self.tab - 1]
        c = self.d.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_before = Pt(4)
        r = c.add_run(f"TABLE {roman}"); font(r, 8)
        r.add_break()
        r2 = c.add_run(caption.upper()); font(r2, 8)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for ci, txt in enumerate(row):
                cell = tr.cells[ci]
                cell.width = Inches(COL_W * widths[ci])
                tcpr = cell._tc.get_or_add_tcPr()
                mar = OxmlElement("w:tcMar")
                for side in ("left", "right"):
                    m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), "40"); m.set(qn("w:type"), "dxa"); mar.append(m)
                tcpr.append(mar)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = ri < len(rows) - 1
                run = p.add_run(txt); font(run, size, bold=(ri == 0))
                if ri == 0:
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E7E9EF"); tcpr.append(shd)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        for r_ in sp.runs:
            r_.font.size = Pt(2)
        sp.paragraph_format.line_spacing = Pt(3)
        return self.tab

    # ---------------------------------------------------------------- front block and columns
    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(TITLE); font(r, 24)
        p.paragraph_format.space_after = Pt(8)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [(GUIDE, GUIDE_DESIG + ", Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Department of Computer Science and Engineering (Data Science)", False),
                               ("[COLLEGE NAME]", False), ("[CITY]", False), ("[EMAIL]", False)):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line); font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        # continuous section with two columns
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        sectpr = sec._sectPr
        cols = sectpr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectpr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def abstract(self):
        p = self.d.add_paragraph()
        r = p.add_run("Abstract—"); font(r, 9, bold=True, italic=True)
        r = p.add_run(ABSTRACT); font(r, 9)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        q = self.d.add_paragraph()
        r = q.add_run("Index Terms—"); font(r, 9, bold=True, italic=True)
        r = q.add_run(KEYWORDS + "."); font(r, 9)
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        q.paragraph_format.space_after = Pt(2)

    def references(self, order):
        self.h1("References")
        for i, key in enumerate(order, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]	{REFERENCES[KEYS[key]]}"); font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


class Cites:
    """Numbers references in order of first citation so that every entry is cited and every citation is listed."""

    def __init__(self):
        self.order = ["base"]

    def __call__(self, *keys):
        nums = []
        for k in keys:
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        return ", ".join(f"[{n}]" for n in nums)


def build():
    P = Paper()
    c = Cites()
    P.front()
    P.abstract()
    X, RF, LR = M["models"]["xgboost"], M["models"]["random_forest"], M["models"]["logistic_regression"]
    ds, B = R["dataset"], R["verdict_bands"]

    # ---------------------------------------------------------------- I
    P.h1("I. Introduction")
    P.para("Phishing remains one of the most common routes to stolen credentials and financial loss: an attacker imitates a trusted "
           f"service and lures the victim to a look-alike page through a link {c('krombholz')}. The link is the one artefact every campaign needs, so "
           "judging it before the page is opened is attractive: a URL is short, always available and cheap to inspect. Today links are "
           "checked mostly by blacklists, fixed rules and classifiers trained once on a fixed dataset. Blacklists act only after a URL has been "
           "reported, so they miss zero-day sites whose domains may live only hours; fixed rules are easy to evade; and static models age as attack "
           f"patterns change {c('base')}.")
    P.para("The base paper argues that these weaknesses are complementary: a classifier can score a URL it has never seen, while a crowd of users "
           f"and feeds recognises a new campaign before a model could learn it {c('base')}. It proposes a three-phase framework, extractor, detection and "
           "crowdsourced updating, but evaluates the crowd loop by simulation on one benchmark. This paper builds and measures a small, complete "
           "version of that idea. PhishGuard is a phishing URL detection system that combines offline URL features, three classifiers and a community "
           "threat-intelligence store whose confirmed verdicts override the model and retrain it. Our contributions are:", after=2)
    P.numbered([
        f"a 47-feature offline URL extractor (23 lexical, 9 structural, 15 domain) with human-readable red flags, and a fair comparison of Random Forest, XGBoost and Logistic Regression;",
        "a vote-based community store with trusted public feeds in which a confirmed phishing or safe verdict overrides the model at once, and a retraining loop with hot reload;",
        "a live end-to-end evaluation on real feed data, plus ablation, threshold and crowd-feedback experiments that state plainly what retraining does and does not change;",
        "a Flask API and a seven-page React interface, with 95 automated checks and documented limits.",
    ])
    P.para("Section II reviews the literature, Section III describes the method, Section IV reports the results and Section V concludes.", before=2)
    P.h2("A. System Hardware")
    P.para("All measurements were taken on one desktop with an Intel Core i5-12500 (6 cores, 12 threads), 15.7 GB of memory and Windows 11 Pro. No GPU "
           "was used. Training, evaluation and timing ran on this machine only; other hardware was not tested.")
    P.h2("B. Software Requirements")
    P.para(f"The system uses Python 3.11.9, scikit-learn 1.4.2 {c('pedregosa')}, XGBoost 2.0.3 {c('chen')}, pandas 2.2.1, NumPy 1.26.4 and joblib for the models, Flask 3.1 "
           f"{c('flask')} with SQLite for the API and storage, and React 18.3 with Vite 5.4 {c('react')} for the interface. Node.js 24 served the frontend. "
           "Everything runs offline except the optional feed synchronisation.")

    # ---------------------------------------------------------------- II
    P.h1("II. Literature Survey")
    P.h2("A. Base Paper")
    P.para(f"Aldhuwayhi and Alromih {c('base')} describe a dynamic framework with three phases. The *extractor* derives features from the URL and answers "
           "directly if the URL matches a known pattern. The *detection* phase classifies the remaining URLs with a trained model and flags inputs that "
           "match no known pattern as novel. In the *updating* phase novel URLs are queued for human verification, which filters noisy reports; "
           f"verified URLs are added to the training set and the model is retrained periodically. They evaluate on PhiUSIIL {c('prasad')}, 235,804 URLs with "
           "54 URL and page-content features, balance the classes by undersampling, keep 34 features after ANOVA and chi-square selection, and train a "
           "hybrid network: an embedding and an LSTM with attention over the URL text, concatenated with a dense branch over the numeric features. They "
           "report accuracy 0.9980, F1 0.9988, five-fold cross-validation, a learning-curve simulation of growing crowd data that plateaus near 99.7 % "
           "accuracy, and about 0.20 ms per URL.")
    P.para("The paper is clear about the gap and about its own limits: results come from one dataset, real-world validation is still needed, and "
           "the crowd channel can be abused. Two points limit what can be taken from it. The crowd loop is not run with real reports, and the "
           "page-content features it inherits from the dataset would require fetching every page in a live system, which the paper does not discuss.")
    P.h2("B. Related Work")
    P.para(f"*Supervised URL classifiers.* As summarised in {c('base')}, Vajrobol et al. {c('vajrobol')} rank features by mutual information and use logistic regression "
           f"(99.97 % with five features on PhiUSIIL); Etem and Teke {c('etem')} reduce 51 features to two with t-SNE and report 99.7 % for a decision tree; Alnemari "
           f"and Alshammari {c('alnemari')} find Random Forest best among classical models (97.3 %). Ma et al. {c('ma')} and Sahingoz et al. {c('sahingoz')} are early and "
           "widely cited studies of detecting malicious sites from URL features. All of these train once on a fixed dataset.")
    P.para(f"*Combined and deployed systems.* Fan {c('fan')} combines SVM and Naive Bayes over text, visual and URL features (99.16 %). Priya et al. {c('priya')} "
           f"build a Random Forest and logistic regression browser extension (97 % on 11,055 URLs) and Sawant et al. {c('sawant')} integrate ML, DL and NLP in an "
           f"extension; the base paper notes that both neglect latency and continuous retraining {c('base')}. *Unsupervised detection.* Assefa and Katarya {c('assefa')} use an "
           "autoencoder over 16 features (91.24 %), which suits novel sites but scores lower.")
    P.para(f"*Crowds and humans in the loop.* Huang et al. {c('huang')} show that crowdsourced cyber-threat intelligence can contain misleading information, and "
           f"Lazaros et al. {c('lazaros')} review human-in-the-loop AI for high-stakes settings; both motivate verification before crowd data reaches a model. "
           "TABLE I compares the works discussed.")
    P.table([
        ["Work", "Approach", "Data", "Reported result", "Limitation"],
        [c("base"), "LSTM-attention + dense; crowd loop", "PhiUSIIL, 235,804", "Acc. 99.8 %", "Crowd loop simulated"],
        [c("vajrobol"), "Mutual info + LR", "PhiUSIIL", "99.97 %", "Static data"],
        [c("etem"), "t-SNE + decision tree", "51-feature set", "99.7 %", "One dataset"],
        [c("alnemari"), "Classical models", "UCI domains", "97.3 % (RF)", "Static data"],
        [c("priya"), "RF / LR extension", "11,055 URLs", "97 %", "Latency unstated"],
        [c("assefa"), "Autoencoder", "Small set", "91.24 %", "Lower accuracy"],
        [c("fan"), "SVM + NB ensemble", "Text, visual, URL", "99.16 %", "Predefined features"],
    ], "Related work (results as reported in [1])", [0.12, 0.28, 0.2, 0.18, 0.22], size=6.5)
    P.para("Three observations follow from the survey. First, accuracy figures between 91 % and 99.97 % are reported on different, static datasets, so they cannot "
           "be ranked against each other, and none of them says how a model behaves on a campaign that starts after training. Second, the studies that "
           "use only the URL string trade some accuracy for speed and safety, because no page has to be fetched, which is the trade this project also makes. "
           "Third, the human-in-the-loop idea is well motivated but rarely built: verification protects the training data, yet it adds delay and needs "
           "staff, and a purely vote-based alternative is faster but easier to game. PhishGuard implements the vote-based variant and measures what it does, "
           "so that the trade-off can be discussed with evidence.")
    P.para("The gap is therefore a system that keeps its detector current from live community evidence, explains its verdicts, and reports "
           "measured behaviour of the loop. Blacklists fail on zero-day URLs, and few systems join lexical, structural and domain features with "
           f"continuously updated threat feeds {c('base')}.")

    # ---------------------------------------------------------------- III
    P.h1("III. Methodology")
    P.h2("A. Pipeline and Data")
    P.para("PhishGuard implements the three phases as one pipeline (Fig. 1). The training data is a balanced sample of 40,000 URLs (20,000 phishing, "
           f"20,000 legitimate, mean length {ds['length_mean']} characters) of a public labelled set of about 420,000 URLs {c('dataset')}; its licence was not "
           "verified. Rows are cleaned (missing values, duplicates, labels mapped to 0/1) and split 80/20, stratified, with a fixed seed: 32,000 for "
           f"training and 8,000 for testing. Only {ds['with_scheme']} URLs carry a scheme, so each URL is normalised first (scheme added, host lower-cased, "
           "www. removed).")
    P.para("Cleaning follows a fixed recipe: rows with a missing URL or label are dropped, text labels (bad, phishing, malicious; good, benign, legitimate) are "
           "mapped to 1 and 0, whitespace is trimmed, URLs shorter than four characters are removed and duplicates are dropped. Because normalisation is "
           "deterministic and never raises on malformed input (for example an invalid port or an unmatched bracket), the same function serves training "
           "rows and live requests. Phishing URLs in the sample are slightly longer than legitimate ones "
           f"({ds['length_mean_phishing']} against {ds['length_mean_legitimate']} characters on average), but the length distributions overlap strongly, so no single "
           "simple feature separates the classes.")
    P.figure(PAPER / "fig1_architecture.png", "System architecture of PhishGuard.", 3.4)
    P.h2("B. Features")
    P.para("Forty-seven features are computed from the URL string alone, with no DNS, WHOIS or page download, so training and run time share one code path. "
           "Twenty-three lexical features count lengths, separators, digits and special characters, the Shannon entropy of the URL, credential-related "
           "words and encoded characters. Nine structural features describe the build of the URL: IP host, port, sub-domain count, path depth, shorteners and "
           "executable extensions. Fifteen domain features describe the registered domain and its TLD: entropy, digit and hyphen counts, vowel ratio, TLD "
           "class, brand names inside or outside the domain, and membership in a built-in list of 145 popular domains used as a reputation signal. "
           "An explainer turns the features into red flags such as “host is a raw IP address”.")
    P.eq("H(s) = -Σ_{c} p(c) log_{2} p(c)", 1)
    P.h2("C. Classifiers and Selection")
    P.para(f"Three models are fitted on the same features and split: Logistic Regression with scaling (C = 1), Random Forest {c('breiman')} with 300 trees and "
           f"XGBoost {c('chen')} with 400 trees, depth 8 and learning rate 0.08. They represent a linear model, bagged trees and boosted trees, and train in seconds. "
           "The model with the highest F1 on the test set is saved and loaded by the detector. No hyper-parameter search was run. We did not reproduce "
           "the base paper's LSTM network: it relies on text and page features of a much larger dataset, whereas tree ensembles on engineered features are "
           "light and can explain their decisions.")
    P.h2("D. Community Intelligence and Decision Rule")
    P.para("Users report a URL as phishing or safe with a reporter name and a note. With a phishing votes and b safe votes, n = a − b; a URL is "
           "confirmed phishing if n ≥ 2 or if it comes from a trusted public feed (OpenPhish, URLhaus), confirmed safe if n ≤ −2, and pending otherwise. "
           "At check time a confirmed entry for the URL or its host decides at once. Otherwise the model probability p is adjusted and thresholded:")
    P.eq("p′ = clip(r·p + 0.15·n, 0, 1)", 2)
    P.para("where r = 0.4 if the registered domain is popular and its TLD is common or longer than two letters, and 1 otherwise. The verdict is "
           "phishing if p′ ≥ 0.60, suspicious if p′ ≥ 0.35 and safe otherwise (Fig. 2). Confirmed URLs are repeated three times in the training matrix and "
           "all models are retrained; the detector then reloads the best model under a lock, without a restart.")
    P.figure(PAPER / "fig2_flow.png", "Decision order of the detector for one URL.", 3.4)
    P.h2("E. Implementation")
    P.para("A Flask application exposes the JSON API (check, report, stats, retrain, sync and five read-only endpoints) on port 8103; the React interface "
           "on port 5103 proxies /api to it. State is a SQLite file with three tables (reports, intel, checks) and the model files. The interface has "
           "seven pages: landing, check (single and batch), report, threat intelligence, models and activity. SQL statements use bound parameters and "
           "the URL is never fetched, so a malicious link cannot run anything on the server.")

    P.h2("F. Robustness and Security")
    P.para("Input is untrusted, so the system never opens the URL; it only reads the string. An empty or whitespace-only URL, an unknown verdict and an unknown "
           "status filter return HTTP 400 with a JSON error, and unusual strings (a bare domain, a 5,000-character URL, a unicode host, a percent-encoded "
           "traversal, javascript: and data: URLs) are classified as text without error. Reports and notes are stored through bound parameters, and the "
           "web pages escape their output, so SQL and script injection are neutralised. A failing public feed is reported per feed and never stops the "
           "application. Known weaknesses are listed in Section IV-E.")

    # ---------------------------------------------------------------- IV
    P.h1("IV. Results")
    P.h2("A. Test Conditions")
    P.para("Models were evaluated on the 8,000 held-out URLs (4,000 per class) that no model saw in training; re-scoring the saved models reproduced the "
           "stored metrics exactly. We also ran five-fold stratified cross-validation over all 40,000 URLs. Timings were measured in process with Flask's "
           "test client on the machine above. Results are therefore from one sampled dataset and are not comparable with other datasets.")
    P.table([
        ["Model", "Acc. %", "Prec. %", "Rec. %", "F1 %", "AUC"],
        ["XGBoost", f"{X['accuracy'] * 100:.2f}", f"{X['precision'] * 100:.2f}", f"{X['recall'] * 100:.2f}", f"{X['f1'] * 100:.2f}", f"{X['roc_auc']:.4f}"],
        ["Random Forest", f"{RF['accuracy'] * 100:.2f}", f"{RF['precision'] * 100:.2f}", f"{RF['recall'] * 100:.2f}", f"{RF['f1'] * 100:.2f}", f"{RF['roc_auc']:.4f}"],
        ["Logistic Regr.", f"{LR['accuracy'] * 100:.2f}", f"{LR['precision'] * 100:.2f}", f"{LR['recall'] * 100:.2f}", f"{LR['f1'] * 100:.2f}", f"{LR['roc_auc']:.4f}"],
        ["XGBoost, 5-fold CV", f"{CV['xgboost']['accuracy']['mean'] * 100:.2f}", f"{CV['xgboost']['precision']['mean'] * 100:.2f}", f"{CV['xgboost']['recall']['mean'] * 100:.2f}", f"{CV['xgboost']['f1']['mean'] * 100:.2f}", f"{CV['xgboost']['roc_auc']['mean']:.4f}"],
    ], "Held-out test results", [0.28, 0.14, 0.15, 0.14, 0.14, 0.15], size=7)
    P.h2("B. Model Performance")
    P.para(f"XGBoost is selected: accuracy {X['accuracy'] * 100:.2f} %, F1 {X['f1'] * 100:.2f} % and ROC-AUC {X['roc_auc']:.3f} (TABLE II, Fig. 3). Random Forest is within 0.5 "
           f"points; Logistic Regression is clearly behind (recall {LR['recall'] * 100:.1f} %), so the classes are not linearly separable in these features. "
           f"Cross-validated F1 is {CV['xgboost']['f1']['mean'] * 100:.2f} % with a standard deviation of {CV['xgboost']['f1']['std'] * 100:.2f} points. XGBoost makes "
           "255 false positives and 311 false negatives on the test set. Lexical features alone reach F1 89.2 %, domain features alone 77.7 % and structural features "
           "alone 66.0 %; all 47 give 92.8 %, and removing the two strongest features (is_common_tld, is_popular_domain) costs only 0.2 points. F1 is flat between "
           f"thresholds 0.40 and 0.60 (92.9 % to 92.4 %), while lowering the threshold to 0.35 raises recall to 94.4 % at a false-positive rate of 9.6 %. With the deployed thresholds, {B['phishing_urls']['phishing']:,} of 4,000 phishing URLs are called phishing and {B['phishing_urls']['suspicious']} suspicious, while "
           f"{B['legitimate_urls']['phishing']} of 4,000 legitimate URLs are called phishing.")
    P.figure(PAPER / "fig3_results.png", "Test metrics (left) and precision, recall and F1 against the threshold (right).", 3.4)
    A = R["ablation"]
    P.table([
        ["Features used (count)", "Acc. %", "F1 %", "AUC"],
        ["Lexical only (23)", f"{A['lexical only']['accuracy'] * 100:.1f}", f"{A['lexical only']['f1'] * 100:.1f}", f"{A['lexical only']['roc_auc']:.3f}"],
        ["Structural only (9)", f"{A['structural only']['accuracy'] * 100:.1f}", f"{A['structural only']['f1'] * 100:.1f}", f"{A['structural only']['roc_auc']:.3f}"],
        ["Domain only (15)", f"{A['domain only']['accuracy'] * 100:.1f}", f"{A['domain only']['f1'] * 100:.1f}", f"{A['domain only']['roc_auc']:.3f}"],
        ["Lexical + domain (38)", f"{A['lexical + domain']['accuracy'] * 100:.1f}", f"{A['lexical + domain']['f1'] * 100:.1f}", f"{A['lexical + domain']['roc_auc']:.3f}"],
        ["All features (47)", f"{A['all 47']['accuracy'] * 100:.1f}", f"{A['all 47']['f1'] * 100:.1f}", f"{A['all 47']['roc_auc']:.3f}"],
        ["Without two strongest (45)", f"{A['without both']['accuracy'] * 100:.1f}", f"{A['without both']['f1'] * 100:.1f}", f"{A['without both']['roc_auc']:.3f}"],
    ], "Feature ablation (XGBoost refitted, threshold 0.5)", [0.46, 0.17, 0.17, 0.2], size=7)
    P.para("The ablation (TABLE III) shows where the signal comes from. The most important features of the selected model are is_common_tld (0.195), "
           "is_popular_domain (0.165) and suspicious_word_count (0.121), followed by brand_outside_domain and has_executable_extension (Fig. 4, left). Because "
           "two of the top three are domain-reputation signals, they were removed together: F1 fell by only 0.2 points, so the model does not depend on them "
           "alone. The learning curve (Fig. 4, right) rises from F1 89.5 % with 3,200 training URLs to 93.2 % with 32,000 and is still rising, which suggests "
           "that more data would help; the final point differs from the shipped model by about 0.3 points because the rows were permuted for this experiment.")
    P.figure(PAPER / "fig6_features.png", "XGBoost feature importance (left) and learning curve (right).", 3.4)
    P.h2("C. Crowd Loop and Output Screens")
    P.para("The crowd loop was exercised on the running application (Fig. 5). A look-alike document-share URL that the model scored at 9.5 % became "
           "phishing at 97.0 % after two community reports, decided by the community, with no retraining; two legitimate pages the model flagged "
           "(docs.python.org and amazon.in) were cleared by two safe reports each. Synchronising the OpenPhish and URLhaus feeds added 800 URLs, and "
           "a retrain from the Models page (Fig. 6) produced model version 2 from 40,000 base rows plus 807 confirmed URLs (XGBoost F1 92.99 % on a "
           "different test set of 8,162 URLs, so not a measured improvement).")
    P.figure(PAPER / "fig4_check.png", "Check screen: community verdict overrides the model.", 3.4)
    P.figure(PAPER / "fig5_models.png", "Models screen after retraining (version 2).", 3.4)
    P.para("To measure whether retraining improves the model, 250 to 2,000 correctly labelled test URLs were added to the training data (×3) and the "
           f"model was scored on the remaining URLs: F1 changed by less than 0.1 point at every size ({R['crowd'][0]['f1_before'] * 100:.2f} % to {R['crowd'][0]['f1_after'] * 100:.2f} % with 250; "
           f"{R['crowd'][3]['f1_before'] * 100:.2f} % to {R['crowd'][3]['f1_after'] * 100:.2f} % with 2,000). Random in-distribution URLs therefore do not change the average score; the value of the loop "
           "is the immediate override and the reversal of known mistakes. Its benefit on a new campaign could not be measured with a static dataset.")
    P.table([
        ["Confirmed URLs added", "Held-out URLs", "F1 before %", "F1 after %", "Missed before", "Missed after"],
    ] + [[f"{d['confirmed_urls']:,}", f"{d['heldout_urls']:,}", f"{d['f1_before'] * 100:.2f}", f"{d['f1_after'] * 100:.2f}", f"{d['missed_before']:.0f}", f"{d['missed_after']:.0f}"] for d in R["crowd"]],
        "Effect of adding confirmed URLs to training (mean of 3 draws)", [0.22, 0.17, 0.15, 0.15, 0.16, 0.15], size=6.5)
    P.para("TABLE IV gives the full experiment. The number of phishing URLs missed on the unseen part moves by two or three at most, and in one case "
           "it rises slightly, which is within the variation of a boosted-tree fit. This is the expected outcome when the model already saw 32,000 similar "
           "examples: correctly labelled URLs from the same distribution add little. The result should not be read as showing that crowd data is useless; "
           "it shows that, on this static dataset, the benefit of a crowd must come from information the model lacks, such as a new campaign or a mistake "
           "that a human can correct. That is what the override provides, and measuring the adaptation to a genuinely new distribution is the main open "
           "experiment of this work.")
    P.h2("D. Speed, Tests and Comparison")
    P.para(f"Feature extraction takes {LAT['feature_extraction_ms_per_url']:.2f} ms per URL and a check through the API takes {LAT['api_check_ms_mean']:.1f} ms on average (p95 {LAT['api_check_ms_p95']:.1f} ms). "
           "Scoring 8,000 feature rows in one batch takes a few milliseconds, and featurising all 40,000 training URLs takes about four seconds. Training the three models takes 15.1 s in total (Random Forest 12.2 s, XGBoost 2.7 s, Logistic Regression 0.2 s), so a retrain on the base data is quick; the retrain in the live run is slower because it also featurises and splits the enlarged data and ran while a browser polled the page. All 95 automated checks pass: 3 original tests, a 20-step API smoke test and 72 extended checks covering unit, integration, functional, "
           f"validation, security and performance cases, including SQL-injection and script-tag inputs. The base paper reports accuracy 99.80 % and F1 0.9988 {c('base')}. "
           f"Our {X['accuracy'] * 100:.2f} % is 6.9 points lower, but the datasets, features and models differ: it used 235,804 URLs, page-content features and an LSTM network, "
           "whereas we use a 40,000-URL sample, URL-only features and tree ensembles. No claim of matching it is made.")
    P.h2("E. Failure Cases and Limits")
    P.para("The most confident false negatives are ordinary-looking pages (a blog post, a form-sharing page, a PDF on a research site): a phishing page on a "
           "clean host has no suspicious URL features. False positives are unusual pages on small sites and, in the live set, docs.python.org (0.60) and "
           "amazon.in (0.79), because the hand-written popular-domain list lacks python and gives no discount to two-letter country TLDs. A confirmed "
           "phishing URL marks every URL on its host as phishing, which would over-block shared hosting. The application has no sign-in, votes are anonymous, "
           "and no user-acceptance, load or independent security test was done.")

    # ---------------------------------------------------------------- V
    P.h1("V. Conclusion and Future Scope")
    P.para(f"PhishGuard implements the three-stage design of the base paper end to end: 47 offline URL features, XGBoost as the best of three classifiers "
           f"(F1 {X['f1'] * 100:.1f} %, ROC-AUC {X['roc_auc']:.3f}, five-fold F1 {CV['xgboost']['f1']['mean'] * 100:.1f} %), a community store whose confirmed verdicts override the "
           "model in both directions, and a retraining loop with hot reload that ran on real feed data. Against the project objectives all six were met, "
           "except that crowd confirmation is a vote count and not a human review. The key finding is that the loop's value is the immediate override, "
           "not a better average score. Limits are the small single-source dataset, URL-only features, anonymous votes and the absence of user testing.")
    P.para("Several practical lessons come out of the work. Interpretable URL features and a reason sentence make a verdict easy to question, and the "
           "community channel turned a false negative and two false positives into correct verdicts within one check. The weakest points are all on "
           "the trust side: anonymous votes, no sign-in, host-wide override and a hand-written reputation list. They are cheap to describe and moderate "
           "to fix, and the next version should treat them before any claim about deployment is made.")
    P.para("Future work: evaluate on larger and more recent data and on URLs collected after training; add sign-in, rate limiting and an administrator "
           "review queue; replace the hand-written reputation list; study the loop under noisy reports and new campaigns; and reproduce the base paper's "
           "network for a like-for-like comparison.")

    P.references(c.order)
    P.d.core_properties.title = TITLE
    P.d.core_properties.author = "Batch A3"
    P.d.save(OUT)
    print("saved", OUT, "references:", len(c.order))


if __name__ == "__main__":
    build()
