"""Verify the submission documents at the batch root.

    python scripts/verify_docs.py

Checks file names and sizes, template leftovers, slide counts, report structure, embedded figures, screenshot files,
student and guide names, and the IEEE paper format. Exits non-zero if anything fails.
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import docx
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from doc_common import BATCH, ROOT, SHOTS  # noqa: E402
from rc_results import SCREENS  # noqa: E402

FILES = {
    "7.Review_1_PPT.pptx": 16, "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx": None, "10.SCOPE & LIMITATIONS.docx": None,
    "11. BLOCK DIAGRAMS ALL.docx": None, "12.DESIGN &IMPLEMENTATION.docx": None, "13.Review_2_PPT.pptx": 20,
    "14.RESULTS.docx": None, "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx": None, "16.Final_PPT.pptx": 28,
    "17.Final Document.docx": None, "19.IEEE_Paper.docx": None,
}
BANNED = ["[STUDENT NAME", "[ROLL NO", "Fakeddit", "PHEME", "Snopes", "MCNN", "bank transaction", "social media manipulation",
          "Kaliyar", "Hidden Risk of Social Media", "SeatWise", "ExamSeater"]
PAPER_BANNED = ["skincare", "Dayananda", "Sindhu", "recommender system"]
MIN_KB = 40
fails: list[str] = []


def ok(cond: bool, msg: str):
    print(("ok    " if cond else "FAIL  ") + msg)
    if not cond:
        fails.append(msg)


def zip_text(path: Path) -> str:
    """All XML text of a package, including embedded workbooks."""
    out = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            data = z.read(n)
            if n.endswith((".xml", ".rels")):
                out.append(data.decode("utf-8", "ignore"))
            elif n.endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(data)) as zz:
                    for m in zz.namelist():
                        if m.endswith(".xml"):
                            out.append(zz.read(m).decode("utf-8", "ignore"))
    return "\n".join(out)


def docx_text(path: Path) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def batch_people():
    d = docx.Document(next(BATCH.glob("1.*_BATCH_DETAILS.docx")))
    students, guide = [], None
    for t in d.tables:
        for r in t.rows:
            cells = [c.text.strip() for c in r.cells]
            if len(cells) >= 3 and re.fullmatch(r"\d+", cells[0]) and re.fullmatch(r"\w{10}", cells[2] or ""):
                students.append((cells[1], cells[2]))
    for p in d.paragraphs:
        m = re.search(r"Under the Guidance of\s*:\s*(.+)", p.text)
        if m:
            guide = m.group(1).strip()
    return students, guide


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


# 1. existence and size
for name in FILES:
    p = BATCH / name
    ok(p.exists(), f"exists: {name}")
    if p.exists():
        kb = p.stat().st_size // 1024
        ok(kb >= MIN_KB, f"size >= {MIN_KB} KB: {name} ({kb} KB)")

# 2. template leftovers
for name in FILES:
    p = BATCH / name
    if not p.exists():
        continue
    text = zip_text(p) + "\n" + (docx_text(p) if name.endswith(".docx") else "")
    low = text.lower()
    hits = [b for b in BANNED if b.lower() in low]
    ok(not hits, f"no template leftovers: {name}" + (f" -> {hits}" if hits else ""))

# 3. deck slide counts
for name, n in FILES.items():
    if n:
        ok(len(Presentation(BATCH / name).slides) == n, f"slide count {n}: {name}")

# 4. full report structure
rep = BATCH / "17.Final Document.docx"
if rep.exists():
    d = docx.Document(rep)
    heads = [p.text.replace("\n", " ").replace("\x0b", " ").strip() for p in d.paragraphs if p.style.name == "Heading 1"]
    want = ["INTRODUCTION", "LITERATURE SURVEY", "FEASIBILITY STUDY", "SYSTEM REQUIREMENTS AND ANALYSIS", "SYSTEM DESIGN",
            "METHODOLOGY AND SOFTWARE IMPLEMENTATION", "SYSTEM TESTING", "RESULTS AND OUTPUT SCREENS", "CONCLUSION", "FUTURE SCOPE"]
    for i, w in enumerate(want, 1):
        ok(any(h == f"CHAPTER {i} {w}" for h in heads), f"chapter {i}: {w}")
    ok(any(h.strip() == "REFERENCES" for h in heads), "references section")
    for letter, title in (("A", "IMPORTANT CODE"), ("B", "INSTALLATION AND USER GUIDE"), ("C", "ADDITIONAL SCREENS")):
        ok(any(h == f"APPENDIX {letter} {title}" for h in heads), f"appendix {letter}")
    front = docx_text(rep)
    for w in ("CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS"):
        ok(w in front.upper(), f"front matter: {w}")

# 5. figures embedded and screenshots exist
for name in ("17.Final Document.docx", "14.RESULTS.docx", "11. BLOCK DIAGRAMS ALL.docx", "12.DESIGN &IMPLEMENTATION.docx", "19.IEEE_Paper.docx"):
    p = BATCH / name
    if p.exists():
        n = len(docx.Document(p).inline_shapes)
        need = {"17.Final Document.docx": 40, "14.RESULTS.docx": 30, "11. BLOCK DIAGRAMS ALL.docx": 10, "12.DESIGN &IMPLEMENTATION.docx": 5, "19.IEEE_Paper.docx": 4}[name]
        ok(n >= need, f"embedded figures >= {need}: {name} ({n})")
missing = [s[0] for s in SCREENS if not (SHOTS / s[0]).exists()]
ok(not missing, f"every referenced screenshot exists ({len(SCREENS)} files)" + (f" -> missing {missing}" if missing else ""))
for sub, n in (("diagrams", 10), ("charts", 10), ("paper", 6)):
    ok(len(list((ROOT / "docs" / sub).glob("*.png"))) >= n, f"docs/{sub} has >= {n} images")

# 6. names on title pages
students, guide = batch_people()
ok(len(students) == 4 and guide is not None, f"batch details parsed: {len(students)} students, guide {guide!r}")
title_texts = {}
for name in FILES:
    p = BATCH / name
    if not p.exists():
        continue
    if name.endswith(".pptx"):
        s0 = Presentation(p).slides[0]
        title_texts[name] = "\n".join(sh.text_frame.text for sh in s0.shapes if sh.has_text_frame)
    else:
        d = docx.Document(p)
        t = "\n".join(x.text for x in d.paragraphs[:12])
        for tb in d.tables[:2]:
            t += "\n" + "\n".join(c.text for r in tb.rows for c in r.cells)
        title_texts[name] = t
for name, t in title_texts.items():
    nt = norm(t)
    ok(all(norm(n) in nt and norm(r) in nt for n, r in students), f"student names and roll numbers on the title page: {name}")
    ok(guide is None or norm(guide) in nt or name in ("16.Final_PPT.pptx",) and norm("G. Ravi") in nt, f"guide name on the title page: {name}")

# 7. IEEE paper
pp = BATCH / "19.IEEE_Paper.docx"
if pp.exists():
    d = docx.Document(pp)
    xml = d.element.xml
    ok('w:num="2"' in xml, "paper has a two-column section")
    sec = d.sections[0]
    ok(abs(sec.page_width.inches - 8.5) < 0.01 and abs(sec.page_height.inches - 11) < 0.01, "paper is US Letter")
    texts = [p.text for p in d.paragraphs]
    romans = [t for t in texts if re.match(r"^(I|II|III|IV|V)\. [A-Z ,]+$", t)]
    ok(len(romans) == 5, f"five Roman-numbered sections ({len(romans)}): {romans}")
    ok(any(t.startswith("TABLE I") for t in texts) and any(t.startswith("TABLE II") for t in texts), "TABLE I and TABLE II present")
    figs = [t for t in texts if re.match(r"^Fig\. \d+\.", t)]
    ok(len(figs) >= 4, f"captioned figures >= 4 ({len(figs)})")
    refs = [t for t in texts if re.match(r"^\[\d+\]\t", t)]
    ok(15 <= len(refs) <= 20, f"references 15-20 ({len(refs)})")
    body = "\n".join(t for t in texts if not re.match(r"^\[\d+\]\t", t)) + "\n" + "\n".join(c.text for tb in d.tables for r in tb.rows for c in r.cells)
    cited = {int(n) for grp in re.findall(r"\[([\d, \[\]]+)\]", body) for n in re.findall(r"\d+", grp)}
    ok(cited == set(range(1, len(refs) + 1)), f"every reference cited and every citation listed (cited {sorted(cited)[:3]}..{max(cited) if cited else 0})")
    ok(any("Abstract" in t for t in texts) and any(t.startswith("Index Terms") for t in texts), "abstract and index terms")
    ab = next(t for t in texts if t.startswith("Abstract")).split("—", 1)[1]
    ok(120 <= len(ab.split()) <= 200 and not re.search(r"\[\d+\]", ab), f"abstract 120-200 words, no citations ({len(ab.split())})")
    low = (zip_text(pp) + docx_text(pp)).lower()
    hits = [b for b in PAPER_BANNED + BANNED if b.lower() in low]
    ok(not hits, "no leftovers from the reference paper or templates" + (f" -> {hits}" if hits else ""))
    pdf = BATCH / "19.IEEE_Paper.pdf"
    if pdf.exists():
        import pymupdf
        n = len(pymupdf.open(pdf))
        ok(6 <= n <= 6, f"paper PDF has 6 pages ({n})")
    else:
        words = len(" ".join(texts).split())
        ok(5200 <= words <= 5800, f"paper word estimate ({words})")

print()
print("ALL CHECKS PASSED" if not fails else f"{len(fails)} CHECK(S) FAILED")
sys.exit(1 if fails else 0)
