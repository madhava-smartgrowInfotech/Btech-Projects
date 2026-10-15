"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import (BATCH, CHARTS, COLLEGE, DEM, DEPT, DIAG, GUIDE, GUIDE_DESIG, M, P, PRODUCT, REFERENCES, RF, B1, B3, ROL, SCAL, SHOTS, STUDENTS, TITLE, YEAR)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE (Data Science)  |  Batch A4  |  Final Review"
AD, ST = ROL["adaptive_avg"], ROL["static_avg"]


def px(v):
    return Emu(int(v * 9525))


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 9)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l)
    set_paras(shape(s, cap_ids[1]), cap_r)
    for cid in cap_ids:
        shape(s, cid).top = px(554)


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Final Review", TITLE, "Batch A4", "hub location, matheuristic, random forest, e-commerce logistics"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A4"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    # 1 title
    s = S[0]
    set_paras(shape(s, 4), TITLE, 24)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A4")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of HubOptima: forecast-driven hub location for e-commerce logistics.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Model",
                "Design & Implementation", "Results, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["E-commerce networks route parcels through hubs, so hub placement drives cost and delivery time. Classical models assume a fixed demand, and exact solvers slow down on large networks.",
                "HubOptima forecasts regional demand with a Random Forest and feeds the forecast to a capacitated hub-location model solved by a matheuristic (greedy construction, MIP-priced local search, kernel search), re-planning every month.",
                "Plans are scored on the demand that actually occurred; five traditional methods run on the same instance; a FastAPI + React application exposes everything.",
                "Synthetic data (90 regions x 36 months): forecast MAPE 7.67% vs 17.6-19.7% for naive forecasters; Oct 2025 plan INR 62.2 M, equal to the exact MIP, 3.6-16.1% cheaper than the alternatives; re-planning 98.3% vs 83.7% fulfilment.",
                "Limits: synthetic data, no robustness analysis, peak under-forecast by 13.2% overloads 3 of 6 hubs.",
                "Keywords: hub location, matheuristic, Random Forest, demand forecasting, e-commerce logistics."], 18, bullet=True)
    set_notes(s, "Abstract: problem, approach, evaluation and the honest limits.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [REFERENCES[0],
                "Method: Random Forest demand forecast + matheuristic (relaxed MILP start, Variable Neighbourhood Search); forecasts re-optimise the hubs; validated on peak, spike and disruption scenarios.",
                "Data: proprietary data of a global transport company, two years of demand, thousands of customer locations.",
                "Reported: delivery-time compliance 92.3%, fulfilment 97.6%, robustness 0.87; +7.8% / +9.4% / +16%; cost reduction 15-20%; networks of 50-500 nodes.",
                "Limitations (authors): sector-specific data, scalability not extensively tested, offline evaluation, other ML models not explored; RF gain given only as 'X %'."], 18, bullet=True)
    set_notes(s, "The base paper is a framework paper evaluated on private data; its numbers are not reproduced here.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Danach et al., 2025 [1] (base)", "RF forecast + matheuristic (MILP, VNS)", "Proprietary logistics data", "92.3% compliance, 97.6% fulfilment", "Data not public; X % gain unstated"],
            ["Zheng et al., 2024 [2]", "Diversity-based hybrid PSO", "Village drone delivery", "42.58% efficiency gain (abstract)", "No forecasting; no optimality bound"],
            ["Yang et al., 2023 [3]", "Maximal-detour p-hub centre; MILP + Benders", "Liner shipping", "Faster exact solve", "Fixed demand"],
            ["Dehghan Chenary et al., 2024 [4]", "Pareto clustering + metaheuristic", "Australian Post", "Better Pareto sets than MOPSO, NSGA-II", "Demand given"],
            ["Zhang et al., 2025 [5]", "Two-stage adaptive robust (C&CG + Benders)", "Not stated in abstract", "Better out-of-sample than stochastic", "Assumed uncertainty set"],
            ["Danach et al., 2019 [9]", "Lagrangian relaxation + hyper-heuristic", "Not reviewed", "Described in [1]", "No forecasting"]]
    table(s, rows, [0.2, 0.25, 0.14, 0.22, 0.19], 14, 440)
    set_notes(s, "Only what the abstracts and the base paper state is summarised; full texts of [2]-[5] were not read.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Data, code and matheuristic parameters unavailable; RF gain stated only as 'X %'; effect of forecast error on the plan not shown"],
            ["Zheng et al. [2]; Dehghan Chenary et al. [4]", "Metaheuristic search with demand taken as given; no forecast"],
            ["Yang et al. [3]; Ernst and Krishnamoorthy [8]", "Exact methods and decomposition for fixed demand; no adaptation"],
            ["Zhang et al. [5]", "Robust under an assumed uncertainty set, not learned from data"],
            ["All reviewed work", "No measured value of monthly re-planning; no exact-versus-matheuristic measurement as networks grow"]]
    table(s, rows, [0.28, 0.72], 17, 440)
    set_notes(s, "The gap is forecast-driven, reproducible and measured hub planning, including the effect of forecast error and of re-planning.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Poor hub placement increases transportation cost and delays deliveries.",
                "Traditional hub location models assume fixed demand and ignore fluctuations.",
                "Exact optimisation is too slow for large, changing logistics networks.",
                "Forecasting and optimisation are done separately, which weakens hub decisions.",
                "Problem: given order history, candidate hubs and cost and service parameters, choose hubs and an allocation that minimise cost within capacity and keep the plan good as demand changes."], 20, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["Collect and preprocess historical e-commerce logistics data.", "Predict regional demand using a Random Forest model.",
                "Formulate hub location as an optimisation problem with cost, time and capacity constraints.", "Solve the problem efficiently with a matheuristic approach.",
                "Re-optimise hub configurations as demand changes.", "Compare performance with traditional hub location methods."], 22, bullet=True)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["1. Data - 36 months of orders for 90 regions, cleaned and turned into 19 features.",
                "2. Forecast - a Random Forest predicts next-month demand per region.",
                "3. Model - a capacitated single-allocation MIP: fixed hub cost, transport cost, delivery-time penalty, capacity, p hubs.",
                "4. Solve - matheuristic: greedy construction, MIP-priced local search, kernel search; exact MIP as reference.",
                "5. Evaluate - plans scored on realised demand: cost, delivery time, fulfilment, utilisation.",
                "6. Compare and adapt - five baselines, a scalability benchmark and monthly re-optimisation.",
                "7. Use - FastAPI + React: forecast, plan with a map, compare, re-optimise."], 19, bullet=False)

    # 10 architecture
    s = S[9]
    put_picture(s, DIAG / "01_system_architecture.png", (240, 178, 800, 424), remove_ids=(7,), name="System architecture")
    set_paras(shape(s, 8), "Layered architecture: browser/CLI, React frontend, FastAPI layer, hublocator core, libraries and files")

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "Predict-then-optimise workflow: load, preprocess, features, train, forecast, model, solve, evaluate, compare and re-plan, present")

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Features & Model")
    fill(s, 7, ["Synthetic dataset (seed 42): 90 regions (30 cities + 60 satellite districts) x 36 months = 3,240 rows, 48.3 M orders; peak 2.38 M in Oct 2025; no real data.",
                "Split by time: train 2023-04 to 2025-06 (2,430 rows), test 2025-07 to 2025-12 (540 rows).",
                "19 features: lags 1-12, rolling mean/std, trend, calendar, promotion, region attributes, order value, return rate.",
                "Random Forest: 400 trees, min leaf 2, 60% features, seed 42; fit 0.8 s; no tuning, no cross-validation.",
                "Hub model: p = 6, capacity 1.25 x demand / p, transport INR 0.04/order-km, time penalty INR 1.5/order-h, fixed INR 2.5 M per hub (by tier), SLA 48 h soft.",
                "Matheuristic: greedy/k-medoids start, swap/drop/add moves screened within 6%, exact assignment MIPs, kernel of 24 hubs."], 18, bullet=True)
    set_notes(s, "All numbers are read from the project's own files; the dataset is synthetic, so absolute values are model outputs.")

    # 13 decision workflow
    s = S[12]
    put_picture(s, DIAG / "11_algorithm_flowchart_wide.png", (80, 178, 1120, 424), remove_ids=(7,), name="Algorithm flowchart")
    set_paras(shape(s, 8), "Core algorithm: construct, price, screen, price moves, accept first improvement, kernel search, return")

    # 14 user-system interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    put_picture(s, DIAG / "09_sequence_diagram.png", (270, 178, 740, 424), remove_ids=(7,), name="Sequence diagram")
    set_paras(shape(s, 8), "Sequence: planning hubs from the browser, through FastAPI, the forecaster and the matheuristic")

    # 15 database
    s = S[14]
    put_picture(s, DIAG / "06_er_diagram.png", (250, 178, 780, 424), remove_ids=(7,), name="ER diagram")
    set_paras(shape(s, 8), "Entity\u2013relationship diagram of the file-based data model (CSV, joblib, JSON; no database)")

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["data.py / forecast.py", "seeded generation, cleaning, 19 shifted features; Random Forest with time split; forecast per region"],
            ["model.py", "Instance (distance x1.25, time, cost, capacity); exact MIP via PuLP/CBC; KPI evaluation on realised demand"],
            ["matheuristic.py", "greedy/k-medoids start; screened swap/drop/add priced by assignment MIPs; kernel search; cache of priced hub sets"],
            ["baselines.py / rolling.py", "greedy, k-means, largest markets, no-forecast; comparison; rolling horizon; scalability benchmark"],
            ["backend/main.py", "FastAPI, 14 endpoints, pydantic validation, background jobs in threads, NaN-safe JSON"],
            ["frontend/", "React + Vite: 7 pages, SVG network map, Recharts, light/dark theme, job polling"]]
    table(s, rows, [0.2, 0.8], 17, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Methods Compared (Oct 2025, p = 6)")
    rows = [["Method", "Operating cost (INR M per month)", "Avg delivery / fulfilment"]]
    for k, lab in [("proposed", "RF + matheuristic (proposed)"), ("exact_mip", "RF + exact MIP"), ("greedy", "RF + greedy"), ("kmeans", "RF + k-means"), ("largest", "Largest-markets rule"), ("fixed_avg", "Historical avg, no forecast")]:
        m = M[k]
        cost = f"{m['total_cost'] / 1e6:.1f}" + ("  (reference)" if k == "proposed" else f"  ({m['cost_vs_proposed_pct']:+.1f}%)")
        rows.append([lab, cost, f"{m['avg_delivery_hours']:.2f} h / {m['fulfilment_rate'] * 100:.1f}%"])
    table(s, rows, [0.4, 0.3, 0.3], 18, 420)
    set_paras(shape(s, 8), "Synthetic data; planned on the forecast, scored on realised demand. Fulfilment is simulated from hub capacity; 3 of 6 hubs overloaded after a 13.2% under-forecast.")
    set_notes(s, "The matheuristic equals the exact MIP on this instance; the heuristics cost 3.6% to 16.1% more.")

    # 18 forecast accuracy (equivalent of a confusion matrix)
    s = S[17]
    set_paras(shape(s, 4), "Forecast Accuracy & Error Distribution")
    two_pics(s, CHARTS / "forecast_comparison.png", CHARTS / "error_histogram.png", "Held-out accuracy: Random Forest vs naive forecasters (MAPE 7.67%)", "Per-region error, Dec 2025: median 7.0%, 8 regions above 20%")
    set_notes(s, "This is the regression counterpart of a confusion matrix: error against two naive benchmarks and the spread of per-region error.")

    # 19 parameter analysis
    s = S[18]
    set_paras(shape(s, 4), "Feature Importance & Scalability")
    two_pics(s, CHARTS / "feature_importance.png", CHARTS / "scalability_time.png", "Top features: 3-month mean 0.379, last month 0.242, population 0.155", "Exact MIP vs matheuristic: 270 nodes, exact not solved; matheuristic +3.67%")
    set_notes(s, "Feature importance shows the level of a region dominates; the scalability chart shows where the matheuristic helps.")

    # 20 performance & testing
    s = S[19]
    rows = [["Measurement", "Result"],
            ["Plan latency (median, end to end)", "3.4 s (4 hubs), 10.7-12.2 s (6 hubs, matheuristic), 15.3 s (6 hubs, exact)"],
            ["Read endpoints", "2-30 ms; forecast endpoint 264 ms"],
            ["Solver time (saved runs)", "matheuristic 4.89 s, exact 7.99 s (90 nodes); 3.13 s vs 12.74 s; 8.31 s vs 42.94 s"],
            ["Automated tests", "tests.py 4 of 4 passed; API smoke test 18 of 18 passed"],
            ["Re-optimisation", "98.3% vs 83.7% fulfilment, -4.9% cost, 1 hub change in 6 months"],
            ["Not done", "user acceptance, load and authentication testing"]]
    table(s, rows, [0.28, 0.72], 17, 430)
    set_notes(s, "Measured on an i5-12500 with 15.7 GB RAM; run-to-run variation of several seconds comes from CBC and processor load.")

    # 21 comparison
    s = S[20]
    rows = [["Indicator", "Base paper [1]", "This project", "Where measured", "Comparable?", "Reading"],
            ["Delivery-time compliance", "92.3%", "92.3% (SLA reliability, Oct 2025 peak)", "synthetic, simulated", "No", "different definition"],
            ["Fulfilment", "97.6%", "92.3% (peak); 98.3% average with re-planning", "synthetic, simulated", "No", "simulated vs observed"],
            ["Cost vs alternatives", "15-20% lower", "3.6-16.1% lower; 0% vs exact MIP", "synthetic, six methods", "Direction only", "baselines differ"],
            ["Gain from adapting", "+7.8% / +9.4% (abstract)", "+14.6 points fulfilment vs a static plan", "synthetic, Jul-Dec 2025", "Different comparison", "control differs"],
            ["Robustness", "0.87", "Not implemented", "-", "-", "gap in this project"],
            ["Scalability", "50-500 nodes", "90-270 nodes; +3.67% at 270", "synthetic benchmark", "Smaller range", "not tested beyond 270"]]
    table(s, rows, [0.16, 0.15, 0.27, 0.16, 0.13, 0.13], 13, 400)
    set_paras(shape(s, 8), "Datasets and metric definitions differ, so this is context only: no head-to-head claim is made.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Plan Hubs")
    two_pics(s, SHOTS / "09_plan_kpis.png", SHOTS / "12_plan_searchlog.png", "Plan Oct 2025: 6 hubs, INR 62.2 M, 92.3% fulfilment", "Search log and allocation of regions to hubs")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Compare and Scalability")
    two_pics(s, SHOTS / "15_compare.png", SHOTS / "17_compare_scalability.png", "Six methods scored on realised demand", "Exact MIP vs matheuristic at 90, 180, 270 nodes")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Forecast and Re-optimisation")
    two_pics(s, SHOTS / "05_forecast.png", SHOTS / "19_rolling.png", "Forecast page: Random Forest vs naive forecasters", "Re-optimise page: 98.3% vs 83.7% fulfilment")

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built HubOptima: forecast-driven hub location with a matheuristic, a comparison suite and a web application.",
                "Forecast MAPE 7.67% (naive 17.6-19.7%); the matheuristic matches the exact MIP at 90 and 180 nodes and is cheapest of six methods on Oct 2025.",
                "Re-planning each month: 98.3% vs 83.7% simulated fulfilment, 4.9% lower cost.",
                "4 tests and 18 API checks pass; 22 screens captured with real data.",
                "Limits: synthetic data; 13.2% peak under-forecast overloads 3 hubs; +3.67% at 270 nodes; no robustness analysis or user study."], 20, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Pilot on real order data and gather feedback from logistics planners.",
                "Robust plans: a safety margin or scenario objective against forecast error.",
                "Retrain the forecaster monthly; try gradient boosting and LSTM; handle extrapolation beyond past peaks.",
                "Add Variable Neighbourhood Search, multiple capacity levels and split allocation.",
                "Database, sign-in and a job queue for multi-user deployment."], 20, bullet=True)

    # 27 references
    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], 13, min_pt=7.5, bullet=True, gap_pt=2)

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
