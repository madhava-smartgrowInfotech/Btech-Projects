"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides)."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import (BATCH, CHARTS, COLLEGE, DEM, DEPT, DIAG, FM, GUIDE, GUIDE_DESIG, M, P, PRODUCT, RF, ROL, SCAL, SHOTS, STUDENTS, TITLE, YEAR)  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"


def px(v):
    return Emu(int(v * 9525))


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, "Batch A4", "hub location, matheuristic, random forest, e-commerce logistics"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A4"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), TITLE, 22)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A4 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers design, methodology and progress of HubOptima.")

    # 2 contents
    s = S[1]
    for sid, t in [(21, "Forecaster and optimiser settings"), (18, "Data preparation"), (27, "Data flow, modules, data model"), (36, "Preliminary results and screens")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts, from a recap of Review 1 through design and methodology to progress, results, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Hub networks are planned on a fixed demand with slow exact solvers or simple rules, so plans go stale as e-commerce demand moves between regions and seasons.", 16)
    fill(s, 13, "A Random Forest forecast feeds a capacitated hub-location model solved by a matheuristic, re-planned every month, behind a FastAPI + React application.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Order history, 90 regions", "Random Forest forecast", "Hub-location model", "Matheuristic solve", "KPIs, compare, re-plan"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1: the problem and the proposed system as a five-step pipeline.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 18 + Vite (port 5104): 7 pages, SVG network map, Recharts, light/dark theme.",
                "API: FastAPI (port 8104), 14 endpoints, pydantic validation, background jobs in threads.",
                "Core: hublocator package: data, forecast, model, matheuristic, baselines, rolling.",
                "Solvers: scikit-learn Random Forest; PuLP with CBC for the MIPs.",
                "Storage: CSV inputs, a joblib model and JSON results; no database."], 16, bullet=True)
    set_notes(s, "The browser talks only to the Vite server, which proxies /api to FastAPI. All science lives in the core package; the API only validates, runs jobs and serialises.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Load + clean data", "19 features", "Train Random Forest", "Forecast next month", "Build hub model", "Matheuristic solve", "Evaluate on realised demand"]):
        fill(s, sid, t, 14)
    set_paras(shape(s, 25), "Plan, then judge")
    fill(s, 26, "Plans are built from the forecast and scored on the demand that actually occurred, so forecast error shows up as overloaded hubs.", 15)
    set_paras(shape(s, 28), "No look-ahead")
    fill(s, 29, "For a past month the forecast uses only earlier data, and all rolling features are shifted by one month.", 15)
    set_paras(shape(s, 31), "Compare and adapt")
    fill(s, 32, "Five baselines run on the same instance; a rolling horizon re-forecasts and re-plans every month against a static plan.", 15)
    set_notes(s, "End-to-end workflow in seven steps. The key idea is to separate planning demand (forecast) from realised demand (what happened).")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    m25 = DEM[DEM.period.str.startswith("2025")].groupby("period").orders.sum()
    ch = replace_chart(s, 4, [p[5:] for p in m25.index], [("Orders (2025)", [int(v) for v in m25.values])])
    style_chart(ch, "Network orders per month, 2025 (synthetic data)", ["4B46C8"], legend=False, label_pt=7, axis_title="orders")
    rows = [["Item", "Facts", "Use"],
            ["Source", "Synthetic, seed 42; no real data available", "Training, planning"],
            ["Regions", "90 = 30 cities + 60 satellite districts", "Demand points and hub sites"],
            ["History", "36 months, 2023-01 to 2025-12; 3,240 rows", "Features, back-tests"],
            ["Orders", "48.3 M in total; peak 2.38 M (Oct 2025)", "Festival peak test"],
            ["Split", "Train 2023-04..2025-06 (2,430 rows); test 2025-07..12 (540)", "Time-based evaluation"],
            ["External data", "None; generated by the project, no licence", "Fully offline"]]
    sz = table_size(rows, 530, 400, [0.2, 0.5, 0.3], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.2, 0.5, 0.3], header_size=sz)
    set_notes(s, "All data is synthetic and generated by a seeded script, so the same seed gives the same files. Festival peaks in Oct-Nov and a growth trend were built in; the model therefore learns a process the authors wrote down.")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation and features")
    set_paras(shape(s, 7), "Cleaning")
    fill(s, 8, "De-duplicate (region, month); coerce orders to numbers; sort; interpolate gaps inside each region; derive year and month. The generated data has no gaps: the guards exist for real files.", 16)
    set_paras(shape(s, 12), "Features (19)")
    fill(s, 13, "Lags 1, 2, 3, 6, 12; rolling mean (3, 6) and std (3); trend; month, quarter, festival; promo and previous promo; population, tier, penetration; order value, return rate.", 16)
    set_paras(shape(s, 17), "No leakage")
    fill(s, 18, "Every rolling window is computed on values shifted by one month; a past month is forecast only from earlier data; the split is by time, never random.", 16)
    set_notes(s, "Preprocessing is cleaning plus 19 engineered features; the shift(1) before every rolling window keeps the current month out of its own features.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Model: forecaster and hub-location solver")
    put_picture(s, DIAG / "11_algorithm_flowchart_wide.png", (60, 130, 710, 450), remove_ids=(4,), name="Matheuristic flowchart")
    rows = [["Component", "Detail"], ["Random Forest", "400 trees, leaf 2, 60% features, seed 42"], ["Hub MIP", "90 regions, p = 6, capacity 1.25"], ["Matheuristic", "construct, swap/drop/add, kernel (24)"], ["Exact MIP", "CBC via PuLP, time-limited"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "The 'model' is a Random Forest plus a mixed-integer model solved by CBC; no neural network is used.", 11)
    set_notes(s, "The flowchart is the real matheuristic. Fixing the hub set leaves a small assignment MIP that CBC solves exactly; the local search moves between hub sets.")

    # 9 training configuration
    s = S[8]
    set_paras(shape(s, 2), "Training and solver configuration")
    rows = [["Setting", "Value"], ["Forecaster", "RandomForestRegressor, 400 trees, depth None"], ["Tuning", "None (one fixed setting, no CV)"], ["Train / test", "2,430 rows / 540 rows (time split)"],
            ["Fit time", "0.8 s"], ["Plan defaults", "p 6, capacity 1.25, SLA 48 h (soft)"], ["MIP limit", "20 s per call (sub-MIP 5 s)"], ["Hardware", "i5-12500, 15.7 GB, Python 3.11.9"]]
    set_table(shape(s, 10), rows, 13, header_size=13)
    put_picture(s, CHARTS / "feature_importance.png", (672, 232, 550, 330), remove_ids=(5,), name="Feature importance chart")
    set_paras(shape(s, 7), "Top features: 3-month rolling mean (0.379), last month (0.242), population (0.155)", 11)
    set_notes(s, "Nothing is tuned: the hyper-parameters were set once. The importances show that the level of a region dominates; calendar features carry little weight.")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "12_decision_logic.png", (57, 124, 420, 520), remove_ids=(4,), name="Decision logic")
    for bid in (7, 10, 13):
        shape(s, bid).height = px(80)
    set_paras(shape(s, 6), "Accepting a move")
    fill(s, 7, "A move is screened by a greedy estimate (within 6%), priced by an exact MIP and accepted only if cheaper.", 15)
    set_paras(shape(s, 9), "Stopping")
    fill(s, 10, "Stops when no promising move remains or 70% of the time is used; a kernel MIP over about 24 hubs then polishes the plan.", 15)
    set_paras(shape(s, 12), "Infeasible plans")
    fill(s, 13, "A start set that breaks capacity or a hard SLA gets more hubs; if still infeasible the API explains why.", 15)
    set_notes(s, "Decision logic: the search accepts only strictly better plans, restarts after each improvement and finishes with a restricted exact MIP.")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 330), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 330), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Solve' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(452)
    rows = [["Module", "Responsibility"], ["data.py, forecast.py", "generate, clean, features, Random Forest, forecasts"],
            ["model.py, matheuristic.py", "Instance, exact MIP, KPI evaluation, matheuristic"],
            ["baselines.py, rolling.py", "greedy, k-means, largest markets; comparison; rolling horizon; scalability"]]
    set_table(shape(s, 12), rows, 12, widths=[0.25, 0.75], header_size=12, row_h=34)
    shape(s, 12).top = px(485)
    set_notes(s, "Level 1 shows nine processes and four data stores; level 2 opens the solve process.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Data model and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 560, 340), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (632, 118, 589, 340), name="UML class diagram")
    set_paras(shape(s, 6), "File-based data model (left, no database) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(460), px(1164)
    rows = [["File", "Content", "Size"],
            ["regions.csv, demand_history.csv", "90 regions; region x month orders with promo, value, return rate", "90 / 3,240 rows"],
            ["rf_demand.joblib, forecast_metrics.json", "fitted forest + 19 features; metrics, importances", "38.5 MB / 12 KB"],
            ["comparison / rolling / scalability .json", "method KPIs and 90-row allocations; monthly adaptive vs static; benchmark runs", "118 KB / 6 KB / 1 KB"]]
    set_table(shape(s, 13), rows, 11, widths=[0.3, 0.5, 0.2], header_size=11, row_h=30, total_px=1164)
    t = shape(s, 13)
    t.left, t.top, t.width = px(57), px(492), px(1164)
    set_notes(s, "There is no database: data and results are readable files, which suits the small data and keeps everything inspectable.")

    # 13 technology
    s = S[12]
    cards = [(7, 8, "Language", "Python 3.11, JavaScript (JSX)"), (12, 13, "Data and ML", "pandas 2.2, NumPy 1.26, scikit-learn 1.4, joblib"),
             (17, 18, "Optimisation", "PuLP 3.3 with the CBC solver"), (22, 23, "Backend", "FastAPI 0.110, uvicorn, pydantic; Flask 3 (classic UI)"),
             (27, 28, "Frontend", "React 18, Vite 5, recharts, framer-motion, gsap"), (32, 33, "Testing", "tests.py (4), API smoke test (18), Playwright captures")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "All components are open source and free; no GPU or commercial solver is needed.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Data generation, cleaning and 19 features", "Completed"], ["Random Forest training and evaluation", "Completed"],
            ["Capacitated hub MIP and exact solver", "Completed"], ["Matheuristic (construct, local search, kernel)", "Completed"], ["Baselines, comparison and scalability benchmark", "Completed"],
            ["Rolling re-optimisation", "Completed"], ["FastAPI layer and React application", "Completed"], ["User acceptance testing with planners", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.78, 0.22], header_size=13)
    for vid, lid, v, lab in [(6, 7, "14", "API endpoints"), (9, 10, "7", "web pages"), (12, 13, "4 + 18", "tests and API checks passing")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Every planned feature is implemented. The one open item is user-acceptance testing with real logistics planners.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: forecast and plan, Oct 2025")
    keys = ["proposed", "exact_mip", "greedy", "kmeans", "largest", "fixed_avg"]
    labs = ["RF + matheuristic", "RF + exact MIP", "RF + greedy", "RF + k-means", "Largest markets", "No forecast"]
    ch = replace_chart(s, 4, labs, [("Cost (INR M)", [round(M[k]["total_cost"] / 1e6, 1) for k in keys])], number_format="0.0")
    style_chart(ch, "Operating cost by method, Oct 2025 (INR M per month)", ["4B46C8"], number_format="0.0", legend=False, label_pt=10, axis_title="INR million")
    for vid, lid, v, lab in [(6, 7, "7.67%", "forecast MAPE (naive: 17.6-19.7%)"), (9, 10, "62.2 M", "monthly cost, same as exact MIP"), (12, 13, "92.3%", "simulated fulfilment (3 hubs overloaded)")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Results on synthetic data: the proposed plan equals the exact optimum and is the cheapest; the others cost 3.6% to 16.1% more. The peak was under-forecast by 13.2%, which overloads three hubs.")

    # 16 output screens
    s = S[15]
    put_picture(s, SHOTS / "09_plan_kpis.png", (91, 128, 507, 476), remove_ids=(4,), name="Plan result screenshot")
    put_picture(s, SHOTS / "15_compare.png", (679, 128, 512, 476), remove_ids=(7,), name="Compare screenshot")
    set_paras(shape(s, 6), "Plan hubs: 6 hubs, INR 62.2 M, 92.3% fulfilment (Oct 2025)", 11)
    set_paras(shape(s, 9), "Compare methods: six methods on realised demand", 11)
    set_notes(s, "Real captures of the running application with the project's data.")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Count", "Result"], ["Data, forecast, optimisation, web app (tests.py)", "4", "Passed"], ["API smoke checks (endpoints, validation)", "18", "Passed"],
            ["Browser walkthrough screens", "22", "Captured"], ["Edge cases (n_hubs 0, unknown month, hard SLA)", "3", "Handled"], ["Total automated", "22", "All passed"]]
    set_table(shape(s, 18), rows, 13, widths=[0.6, 0.15, 0.25], header_size=13)
    fill(s, 5, ["Matheuristic and exact MIP agree on the 90-region plan; the test requires them to be within 2%.",
                "Planning takes about 3 s (4 hubs) to 12 s (6 hubs) end to end; read endpoints answer in 2-30 ms.",
                "A 13.2% under-forecast of the peak overloads 3 of 6 hubs: found, reported, not hidden.",
                "Not done: user-acceptance testing, load testing, authentication testing."], 16, bullet=True)
    set_notes(s, "Four unit tests, 18 API checks and a scripted browser walkthrough. The frontend has no automated unit tests.")

    # 18 challenges
    s = S[17]
    cards = [(7, 8, "Synthetic data", "No real order history was available; stated in every result and limitation."),
             (12, 13, "Peak under-forecast", "13.2% low in Oct 2025; plans overload 3 hubs. Next: a safety margin."),
             (17, 18, "Forests cannot extrapolate", "Oct 2025 exceeded every earlier month; likely cause of the miss."),
             (22, 23, "Exact MIP slows down", "270 nodes not solved in 67 s: the case for the matheuristic."),
             (27, 28, "Matheuristic gap", "3.67% above the exact incumbent at 270 nodes; kernel size is a lever."),
             (32, 33, "No authentication", "Local planning tool only; recorded as a limitation.")]
    for tid, bid, ttl, body in cards:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, 15)
    set_notes(s, "Six challenges and how each was handled or recorded.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["User acceptance testing", "Robust plan with a safety margin", "Final report", "Final presentation and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Collect feedback from logistics planners on the plan, compare and re-optimise screens (not yet done).",
                 "Add a safety margin or scenario objective so forecast error does not overload the biggest hubs.",
                 "Retrain the forecaster monthly inside the rolling horizon; test other learners.",
                 "Document every diagram, test and result in the report."], 18, bullet=True)
    set_notes(s, "What remains: real-user feedback and robustness, then the report and the final demo.")

    set_notes(S[19], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
