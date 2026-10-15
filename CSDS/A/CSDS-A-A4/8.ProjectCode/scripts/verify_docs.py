"""Verify the submission documents at the batch root.

    python scripts/verify_docs.py

Checks: (1) all files exist with the exact names and are not under 40 KB (decks excepted when legitimately smaller); (2) no template
leftovers in any output; (3) deck slide counts 16 / 20 / 28; (4) the full report has all ten chapter headings, the references and three
appendices; (5) figures are embedded and every screenshot referenced exists; (6) student names, roll numbers and the guide name from
1.*_BATCH_DETAILS.docx appear on the title pages; (7) the IEEE paper is two-column with five Roman-numbered sections, both tables, at least
four captioned figures, 15-20 references (each cited, each citation listed) and six pages; (8) no leftovers from the reference paper.
Needs python-docx, python-pptx and PyMuPDF. Exit code 0 only when every check passes.
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from doc_common import BATCH, DIAG, SCREENS, SHOTS  # noqa: E402

FILES = ["7.Review_1_PPT.pptx", "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "10.SCOPE & LIMITATIONS.docx", "11. BLOCK DIAGRAMS ALL.docx",
         "12.DESIGN &IMPLEMENTATION.docx", "13.Review_2_PPT.pptx", "14.RESULTS.docx", "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
         "16.Final_PPT.pptx", "17.Final Document.docx"]
PAPER = "19.IEEE_Paper.docx"
DECKS = {"7.Review_1_PPT.pptx": 16, "13.Review_2_PPT.pptx": 20, "16.Final_PPT.pptx": 28}
MIN_KB = 40
EXACT = ["[STUDENT NAME", "[ROLL NO"]
PHRASES = ["Fakeddit", "PHEME", "Snopes", "MCNN", "bank transaction", "social media manipulation", "Hidden Risk of Social Media", "FraudShield",
           "Kafka", "Batch D1", "Batch A13", "Royyala", "GenImage", "ResNet50", "Kaliyar", "Sarwosri", "Chandrakala", "SeatWise"]
PAPER_PHRASES = ["skincare", "Dayananda", "Sindhu", "recommender system"]
CHAPTERS = ["INTRODUCTION", "LITERATURE SURVEY", "FEASIBILITY STUDY", "SYSTEM REQUIREMENTS AND ANALYSIS", "SYSTEM DESIGN",
            "METHODOLOGY AND SOFTWARE IMPLEMENTATION", "SYSTEM TESTING", "RESULTS AND OUTPUT SCREENS", "CONCLUSION", "FUTURE SCOPE"]

failures: list[str] = []


def check(ok: bool, what: str) -> None:
    print(("PASS  " if ok else "FAIL  ") + what)
    if not ok:
        failures.append(what)


def zip_texts(path: Path) -> dict[str, str]:
    """All XML text of a docx/pptx package (headers, notes, layouts, charts and embedded workbooks included)."""
    out: dict[str, str] = {}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            data = z.read(name)
            if name.endswith((".xml", ".rels")):
                out[name] = re.sub(r"<[^>]+>", " ", data.decode("utf-8", "ignore"))
            elif name.endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(data)) as inner:
                    for n2 in inner.namelist():
                        if n2.endswith(".xml"):
                            out[f"{name}!{n2}"] = re.sub(r"<[^>]+>", " ", inner.read(n2).decode("utf-8", "ignore"))
    return out


def batch_details():
    path = next(BATCH.glob("1.*_BATCH_DETAILS.docx"))
    doc = Document(path)
    tokens: list[str] = []
    for t in doc.tables:
        header = [c.text.strip().lower() for c in t.rows[0].cells]
        if any("student" in h for h in header) and any("hall" in h or "roll" in h for h in header):
            for row in t.rows[1:]:
                tokens += [c.text.strip() for c in row.cells[1:] if c.text.strip() and c.text.strip().upper() not in ("STUDENT NAME", "GUIDE NAME")]
    guide = ""
    for p in doc.paragraphs:
        if "guidance" in p.text.lower() and ":" in p.text:
            guide = p.text.split(":", 1)[1].strip()
    return tokens, guide


def docx_first_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs[:40]]
    for t in doc.tables[:3]:
        for row in t.rows:
            chunks += [c.text for c in row.cells]
    return "\n".join(chunks)


def paper_checks(path: Path) -> None:
    import pymupdf
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    sect = doc.sections[-1]
    cols = sect._sectPr.find(qn("w:cols"))
    check(cols is not None and cols.get(qn("w:num")) == "2", "paper: body section is two-column")
    check(abs(sect.page_width.inches - 8.5) < 0.01 and abs(sect.page_height.inches - 11) < 0.01, "paper: US Letter")
    check(abs(sect.top_margin.inches - 0.75) < 0.01 and abs(sect.bottom_margin.inches - 1.0) < 0.01 and abs(sect.left_margin.inches - 0.69) < 0.01, "paper: margins 0.75/1.0/0.69 in")
    romans = [t for t in paras if re.match(r"^(I|II|III|IV|V)\. [A-Z ]+$", t.strip())]
    check(len(romans) == 5, f"paper: five Roman-numbered sections ({romans})")
    check(any(t.strip() == "REFERENCES" for t in paras), "paper: REFERENCES heading")
    full = "\n".join(paras)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                full += "\n" + c.text
    check("TABLE I" in full and "TABLE II" in full, "paper: TABLE I and TABLE II present")
    figs = [t for t in paras if re.match(r"^Fig\. \d+\. ", t.strip())]
    check(len(figs) >= 4, f"paper: {len(figs)} captioned figures")
    check(len(doc.inline_shapes) >= 4, f"paper: {len(doc.inline_shapes)} embedded images")
    refs = [t for t in paras if re.match(r"^\[\d+\]\t", t)]
    check(15 <= len(refs) <= 20, f"paper: {len(refs)} references (15-20)")
    body_text = "\n".join(t for t in paras if not re.match(r"^\[\d+\]\t", t))
    cited = set()
    for m in re.finditer(r"\[([\d,\s\-–]+)\]", body_text):
        for part in re.split(r",\s*", m.group(1)):
            part = part.strip()
            if "-" in part or "–" in part:
                a, b = re.split(r"[-–]", part)
                cited.update(range(int(a), int(b) + 1))
            elif part.isdigit():
                cited.add(int(part))
    listed = set(range(1, len(refs) + 1))
    check(cited <= listed, f"paper: every citation is listed (unlisted: {sorted(cited - listed)})")
    check(listed <= cited, f"paper: every reference is cited (uncited: {sorted(listed - cited)})")
    abstract = next((t for t in paras if t.startswith("Abstract")), "")
    n_abs = len(abstract.replace("Abstract—", "").split())
    check(120 <= n_abs <= 200, f"paper: abstract {n_abs} words (120-200)")
    check(not re.search(r"\[\d+\]", abstract), "paper: abstract has no citations")
    idx = next((t for t in paras if t.startswith("Index Terms")), "")
    terms = [x.strip() for x in idx.replace("Index Terms—", "").rstrip(".").split(",")]
    check(5 <= len(terms) <= 7 and terms == sorted(terms, key=str.lower), f"paper: {len(terms)} index terms, alphabetical")
    pdf = BATCH / "19.IEEE_Paper.pdf"
    check(pdf.exists(), "paper: PDF exported")
    if pdf.exists():
        n = len(pymupdf.open(pdf))
        check(n == 6, f"paper: {n} pages (target 6)")
    low = full.lower()
    hits = [p for p in PAPER_PHRASES if p.lower() in low]
    check(not hits, f"paper: no reference-paper leftovers {hits}")
    texts = zip_texts(path)
    hits = [(n, ph) for n, tx in texts.items() for ph in PHRASES + PAPER_PHRASES + EXACT if ph.lower() in tx.lower()]
    check(not hits, f"paper: no template leftovers anywhere in the package {hits[:3]}")


def main() -> int:
    paths = {f: BATCH / f for f in FILES}
    print("== 1. files")
    for f, p in paths.items():
        check(p.exists(), f"exists: {f}")
        if p.exists():
            kb = p.stat().st_size / 1024
            check(kb >= MIN_KB or f in DECKS and kb >= 20, f"size {kb:,.0f} KB: {f}")
    check((BATCH / PAPER).exists(), f"exists: {PAPER}")
    if failures:
        return 1

    print("== 2. template leftovers")
    for f, p in paths.items():
        texts = zip_texts(p)
        hits = []
        for name, txt in texts.items():
            for tok in EXACT:
                if tok in txt:
                    hits.append((name, tok))
            low = txt.lower()
            for ph in PHRASES + PAPER_PHRASES:
                if ph.lower() in low:
                    hits.append((name, ph))
        check(not hits, f"no template leftovers in {f}" + (f"  -> {hits[:4]}" if hits else ""))

    print("== 3. deck slide counts")
    for f, n in DECKS.items():
        prs = Presentation(paths[f])
        check(len(prs.slides) == n, f"{f}: {len(prs.slides)} slides (expected {n})")

    print("== 4. full report structure")
    doc = Document(paths["17.Final Document.docx"])
    heads = [p.text.upper().replace("\n", " ").replace("\x0b", " ") for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    for i, ch in enumerate(CHAPTERS, 1):
        check(any(h.startswith(f"CHAPTER {i}") and ch in h for h in heads), f"chapter {i} heading: {ch}")
    check(any(h.strip() == "REFERENCES" for h in heads), "references section")
    for letter in "ABC":
        check(any(h.startswith(f"APPENDIX {letter}") for h in heads), f"appendix {letter}")
    for front in ("CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENT", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS"):
        check(any(p.text.strip().upper() == front for p in doc.paragraphs), f"front matter: {front}")

    print("== 5. figures and screenshots")
    for f in FILES:
        with zipfile.ZipFile(paths[f]) as z:
            n = len([m for m in z.namelist() if m.startswith(("word/media/", "ppt/media/"))])
        need = 3 if f.endswith(".docx") and any(k in f for k in ("17.", "11.", "14.")) else (1 if f.endswith(".docx") else 5)
        check(n >= need, f"{f}: {n} embedded images / media files")
    check(len(doc.inline_shapes) >= 40, f"full report: {len(doc.inline_shapes)} inline figures (need at least 40)")
    missing = [s[0] for s in SCREENS if not (SHOTS / s[0]).exists()]
    check(not missing, f"every screenshot referenced exists ({len(SCREENS)} listed)" + (f"  -> missing {missing}" if missing else ""))
    check(len(sorted(DIAG.glob("*.png"))) >= 10, f"{len(sorted(DIAG.glob('*.png')))} diagram files in docs/diagrams")

    print("== 6. names, roll numbers and guide on the title pages")
    tokens, guide = batch_details()
    for f in FILES + [PAPER]:
        p = BATCH / f
        if f.endswith(".pptx"):
            prs = Presentation(p)
            text = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
        else:
            text = docx_first_text(p)
        missing = [t for t in tokens if t not in text]
        check(not missing, f"{f}: students/rolls on the title page" + (f"  -> missing {missing}" if missing else ""))
        check(guide != "" and guide.lower().strip("[]") in text.lower(), f"{f}: guide name ({guide})")

    print("== 7-8. IEEE paper")
    paper_checks(BATCH / PAPER)

    print()
    if failures:
        print(f"{len(failures)} CHECK(S) FAILED")
        return 1
    print("ALL CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
