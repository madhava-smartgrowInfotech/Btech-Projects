"""Build 7.Review_1_PPT.pptx, 13.Review_2_PPT.pptx and 16.Final_PPT.pptx from the templates in D:\\abstracts\\Templates.

Usage (from 8.ProjectCode):  python scripts/docs_build/build_decks.py [r1|r2|final|all]
Numbers are read from experiments/*.json so the decks agree with the documents.
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation  # noqa: E402

import deck_lib as L  # noqa: E402
from deck_lib import Deck, ACCENTS, BLUE, GREEN, ORANGE, PURPLE, RED  # noqa: E402
from common import (ROOT, BATCH, STUDENTS, GUIDE, GUIDE_DESIG, COLLEGE, PROJECT, ARCH, FLOW, DFD1, DFD2, UML, SEQ, UC, ER, ALGO, C, S, SCREENS, REFS)  # noqa: E402

TPL = r"D:\abstracts\Templates"
EX = os.path.join(ROOT, "experiments")


def J(*p):
    return json.load(open(os.path.join(EX, *p), encoding="utf-8"))


EV = J("eval", "metrics.json")
DET = EV["detection"]
TP = DET["taxpayer"]
M, B = TP["model"], TP["rule_baseline"]
AB = TP["ablation"]
TR = J("metrics.json")
LAT = J("latency.json")


def _hu(x, nd):
    from decimal import ROUND_HALF_UP, Decimal
    return str(Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP))


def f3(x):
    return _hu(x, 3)


def f2(x):
    return _hu(x, 2)


def load_seed():
    p = os.path.join(EX, "seed_study.json")
    return json.load(open(p)) if os.path.exists(p) else None


def load_edge():
    return J("edge_tests.json")


def load_err():
    p = os.path.join(EX, "error_analysis.json")
    return json.load(open(p)) if os.path.exists(p) else None


YEARS = "2026-2027"
TITLE_SHORT = "Self-Supervised, Explainable GST ITC Fraud Detection"
SUB = "JEPA-LLM Tuning Framework for Combating GST and Sales Tax Input Tax Credit Fraud Pattern Detection"
SCR = lambda k: S(SCREENS[k][0])  # noqa: E731


def review_keep(shp):
    if shp.is_placeholder:
        t = str(shp.placeholder_format.type)
        return "TITLE" in t or "SLIDE_NUMBER" in t
    return False


def fix_title_slide(D, s, label):
    for shp in s.shapes:
        if shp.name == "Text 0" and "TITLE" in str(shp.placeholder_format.type):
            D.set_runs(shp, [TITLE_SHORT])
            for r in shp.text_frame.paragraphs[0].runs:
                r.font.size = L.Pt(36)
        elif shp.name == "Text 0":
            D.set_runs(shp, [SUB])
        elif shp.name == "review label":
            D.set_runs(shp, [label])
        elif shp.name == "team":
            D.set_runs(shp, [f"{n}   {r}" for n, r in STUDENTS])
        elif shp.name == "Text 6":
            D.set_runs(shp, [f"{GUIDE}, {GUIDE_DESIG}"])
        elif shp.name == "Text 7":
            D.set_runs(shp, ["Department of Computer Science and Engineering (Data Science)", f"{COLLEGE}  ·  Batch A5 (CSE-DS)  ·  {YEARS}"])
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = L.Pt(13)


def finish(prs, path, label):
    for lay in prs.slide_layouts:
        for shp in lay.shapes:
            if shp.has_text_frame and "Hidden Risk" in shp.text_frame.text:
                for p in shp.text_frame.paragraphs:
                    for r in p.runs:
                        r.text = r.text.replace("The Hidden Risk of Social Media", "TaxSentinel: GST ITC fraud detection")
    for sl_ in prs.slides:
        if sl_.has_notes_slide:
            for shp in sl_.notes_slide.shapes:
                if shp.has_text_frame and not (shp.is_placeholder and 'SLIDE_IMAGE' in str(shp.placeholder_format.type)):
                    shp.text_frame.text = ''
    cp = prs.core_properties
    cp.title = f"{PROJECT} - {label}"
    cp.author = "Batch A5: " + ", ".join(n for n, _ in STUDENTS)
    cp.subject = "TaxSentinel"
    cp.comments = ""
    cp.keywords = "GST, ITC fraud, JEPA, explainable AI"
    cp.last_modified_by = "Batch A5"
    prs.save(path)
    print("saved", os.path.basename(path), os.path.getsize(path) // 1024, "KB")


# =====================================================================================================  REVIEW 1
def build_r1():
    prs = Presentation(os.path.join(TPL, "Review_1_PPT.pptx"))
    D = Deck(prs, "review")
    sl = list(prs.slides)
    seed = load_seed()
    fix_title_slide(D, sl[0], "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    # 2 contents: labels are generic, keep as in the template
    # 3 abstract
    s = sl[2]
    D.wipe(s, review_keep)
    D.title(s, "Abstract: learning what normal GST filing looks like", 30)
    D.card(s, 0.6, 1.5, 3.95, 4.7, "The problem", [
        "Input Tax Credit (ITC) fraud through fake invoices, circular trading and shell companies causes large revenue losses.",
        "Audits and fixed rules are slow and easy to evade, fraud labels are scarce, and flags rarely come with reasons.",
        "DGGI detected evasion of Rs 2.01 lakh crore in FY 2023-24 [1]."], ORANGE, "1", body_size=15)
    D.card(s, 4.7, 1.5, 3.95, 4.7, "Our idea", [
        "Learn normal GST behaviour without labels with a JEPA encoder.",
        "Combine its deviation score with circular-trading, shell-cluster and invoice evidence into one risk.",
        "Let Gemini write a case note in which every claim cites the evidence."], BLUE, "2", body_size=15)
    D.card(s, 8.8, 1.5, 3.95, 4.7, "The deliverable", [
        "TaxSentinel: a secured FastAPI + React web application.",
        "Ranked taxpayers and invoice chains, an interactive fraud graph, numbered evidence and cited case notes.",
        "An evaluation page against a rule baseline. Data: a seeded synthetic GST ecosystem."], GREEN, "3", body_size=15)
    # 4 base paper
    s = sl[3]
    D.wipe(s, review_keep)
    D.title(s, "Base paper", 32)
    D.rect(s, 0.6, 1.4, 12.1, 1.15)
    D.text(s, 0.8, 1.5, 11.7, 0.95, "S. Datta, H. Thakkar and S. Dabhade, \"Combating GST Input Tax Credit Fraud in India: Forensic Accounting, Data Analytics, and Legal Enforcement Strategies,\" Int. J. Management, Public Policy and Research, vol. 5, no. 1, Jan-Mar 2026, doi 10.55829/ykz3ea41.", size=14, anchor="m")
    D.plain_card(s, 0.6, 2.8, 3.95, 3.7, "Method", "A qualitative synthesis of public DGGI and CBIC data and case studies. It explains how fake invoicing, circular trading and shell networks work and recommends data analytics, network analysis, AI-based risk modelling and forensic capacity.", ORANGE, body_size=14)
    D.plain_card(s, 4.7, 2.8, 3.95, 3.7, "Data and reported results", "FY 2022-23 to FY 2023-24: detected evasion Rs 1.01 to Rs 2.01 lakh crore; cases booked 4,872 to 6,084; arrests 97 to 140. It reports no detection model and no accuracy figures.", BLUE, body_size=14)
    D.plain_card(s, 8.8, 2.8, 3.95, 3.7, "Limitations stated by the authors", "Public sources only, no primary or empirical data, qualitative. No detection model is implemented or evaluated; quantitative, AI-supported detection is left as future work.", GREEN, body_size=14)
    # 5 literature survey
    s = sl[4]
    D.wipe(s, review_keep)
    D.title(s, "Literature survey", 32)
    rows = [["Work", "Approach", "Data", "Reported result", "Limitation for this problem"],
            ["Datta et al. [1] (base paper)", "Forensic accounting, analytics, legal strategy", "DGGI and CBIC statistics, cases", "Evasion Rs 1.01 to 2.01 lakh crore (FY23 to FY24)", "Qualitative; no model"],
            ["AI/ML in tax compliance [2]; income-tax fraud ML [3], [4]; boosting [5]", "ML and AI for tax fraud (per titles; full texts not available to the team)", "Not available", "Not quoted", "Supervised on labelled records; per-taxpayer; no network or explanation (per titles)"],
            ["Mehta et al. [6]", "node2vec + DBSCAN circular-trading communities", "Telangana GST: 77.8 M invoices, 388 k taxpayers", "Communities found; no accuracy metric", "One pattern; no explanation"],
            ["Xu et al. [7] (TED)", "Heterogeneous graph neural network", "Two real tax datasets", "F1 76.35 to 94.96 %", "Supervised; needs labels"],
            ["Assran et al. [8] (I-JEPA)", "Predict masked-block embeddings", "Images", "Strong transfer", "Not applied to tax behaviour"]]
    D.table(s, 0.6, 1.4, 12.1, [2.3, 2.5, 2.2, 2.4, 2.7], rows, size=12, max_h=5.3)
    # 6 supporting evidence
    s = sl[5]
    D.wipe(s, review_keep)
    D.title(s, "Supporting evidence: the scale of GST evasion (base paper data)", 28)
    D.bar_chart(s, 0.6, 1.4, 7.4, 5.1, ["FY 2022-23", "FY 2023-24"], [("Evasion detected (Rs lakh crore)", [1.01, 2.01])], colors=[ORANGE], number_format="0.00", title="DGGI: evasion detected (Rs lakh crore)", legend=False, maxv=2.5, font_size=13)
    D.card(s, 8.3, 1.5, 4.4, 2.3, "Takeaway", "Evasion detected doubled in a year, while cases booked rose from 4,872 to 6,084 (about 25 %) [1].", BLUE, "1", body_size=14)
    D.card(s, 8.3, 4.0, 4.4, 2.5, "Caveat", "These are government figures quoted in the base paper. They show the scale of the problem, not the accuracy of any method; our own results are on synthetic data.", ORANGE, "!", body_size=14)
    # 7 existing system
    s = sl[6]
    D.wipe(s, review_keep)
    D.title(s, "Existing system and current approaches", 32)
    D.flow(s, 0.6, 1.8, 12.1, 1.2, ["Returns filed (GSTR-1, 2B, 3B)", "Mismatch rules and notices", "Manual scrutiny", "Audit and investigation", "Recovery and prosecution"], size=13)
    D.card(s, 0.6, 3.6, 5.95, 2.9, "Rule-based checks", "Fixed thresholds such as ITC claimed above GSTR-2B. Transparent, but easy to stay under. In our baseline test (rules R1-R5) they flagged 236 taxpayers, 168 of them false alarms.", BLUE, "R", body_size=15)
    D.card(s, 6.75, 3.6, 5.95, 2.9, "Manual audit and investigation", "Experts examine a taxpayer after suspicion arises: reliable, but slow, and it starts after revenue has been lost.", ORANGE, "M", body_size=15)
    # 8 limitations
    s = sl[7]
    D.wipe(s, review_keep)
    D.title(s, "Limitations of existing methods", 32)
    cards = [("Slow and reactive", "Audits detect fraud only after revenue is lost.", ORANGE), ("Easily evaded", "Fixed rules can be avoided by staying under the thresholds.", BLUE),
             ("Scarce labels", "Confirmed fraud cases are few and late, which limits supervised learning.", GREEN), ("Blind to networks, silent on reasons", "Per-entity checks miss loops and shell clusters, and scores rarely explain themselves.", PURPLE)]
    for i, (h, b, c) in enumerate(cards):
        D.card(s, 0.6 + (i % 2) * 6.2, 1.5 + (i // 2) * 2.5, 5.95, 2.3, h, b, c, str(i + 1), body_size=15)
    # 9 research gap
    s = sl[8]
    D.wipe(s, review_keep)
    D.title(s, "Research gap", 32)
    gaps = [("Labels", "Supervised work needs confirmed fraud; little learns normal GST behaviour without labels [2]-[5], [7]."), ("Combined evidence", "Invoice, return-behaviour and network evidence are rarely scored together."),
            ("Group fraud", "Circular trading and shell clusters need graph analysis, not per-entity scores [6]."), ("Explanation", "Surveyed GST work returns scores or communities, not cited reasons an officer can check."),
            ("Application", "The base paper calls for AI risk modelling [1] but proposes no working, evaluable system.")]
    for i, (h, b) in enumerate(gaps):
        y = 1.4 + i * 1.05
        D.rect(s, 0.6, y, 2.6, 0.85, fill=ACCENTS[i], line=None)
        D.text(s, 0.6, y, 2.6, 0.85, h, size=16, bold=True, color="FFFFFF", align="c", anchor="m")
        D.rect(s, 3.35, y, 9.35, 0.85)
        D.text(s, 3.5, y + 0.03, 9.1, 0.8, b, size=14, anchor="m")
    # 10 problem statement
    s = sl[9]
    D.wipe(s, review_keep)
    D.title(s, "Problem statement", 32)
    D.rect(s, 0.6, 1.5, 12.1, 3.0)
    D.text(s, 1.0, 1.65, 11.3, 2.7, "Input Tax Credit fraud through fake invoices, circular trading and shell companies causes large revenue losses. Manual audits and rule-based checks are slow and reactive, confirmed fraud labels are scarce, and existing systems rarely explain why a claim is flagged. The problem is to design a self-supervised, explainable framework that learns normal GST behaviour, ranks taxpayers and invoice chains by deviation from it, and explains each flagged claim in language an investigator can verify.", size=18, anchor="m")
    D.plain_card(s, 0.6, 4.8, 3.95, 1.75, "Input", "Returns, invoices and registration data (a seeded synthetic ecosystem)", ORANGE, body_size=13)
    D.plain_card(s, 4.7, 4.8, 3.95, 1.75, "Output", "Risk ranking, rings, invoice chains and cited case notes", BLUE, body_size=13)
    D.plain_card(s, 8.8, 4.8, 3.95, 1.75, "Boundary", "A risk indication for review, not a legal finding", GREEN, body_size=13)
    # 11 objectives
    s = sl[10]
    D.wipe(s, review_keep)
    D.title(s, "Objectives", 32)
    D.text(s, 0.6, 1.4, 7.4, 5.3, [
        "1.  Build a multi-layer GST dataset covering invoices, taxpayer behaviour and transaction networks.",
        "2.  Learn normal GST behaviour with a self-supervised JEPA encoder.",
        "3.  Detect fraud patterns such as fake invoices, circular trading and shell-entity activity.",
        "4.  Score and rank suspicious taxpayers and invoice chains by deviation from learned behaviour.",
        "5.  Generate natural-language explanations for flagged claims using an LLM (prompt-tuned and checked, not fine-tuned).",
        "6.  Evaluate detection accuracy and explanation quality."], size=18, spacing=10)
    D.picture(s, UC, 8.2, 1.7, 4.6, 4.2, caption="Use cases the system serves (analyst view)")
    # 12 proposed approach
    s = sl[11]
    D.wipe(s, review_keep)
    D.title(s, "Proposed approach", 32)
    cs = [("JEPA encoder", "19 features per taxpayer-month; predicts embeddings from history and with one layer hidden; error = deviation.", ORANGE), ("Network analysis", "Circular-trading loops (bounded cycles) and shared-address / contact shell clusters.", BLUE),
          ("Invoice and risk", "Eight-signal invoice score; noisy-OR of JEPA, behaviour, network and invoice scores; flag at 0.5.", GREEN), ("Cited explanation", "Gemini restates numbered evidence only; citations and numbers are checked automatically.", PURPLE)]
    for i, (h, b, c) in enumerate(cs):
        D.card(s, 0.6 + i * 3.05, 1.5, 2.9, 3.3, h, b, c, str(i + 1), body_size=13, head_size=15)
    D.card(s, 0.6, 4.95, 12.1, 1.8, "Web application", "FastAPI with JWT, SQLite and background runs; React interface with ranking, profile, fraud graph, chains, case notes, performance page and run history.", BLUE, "W", body_size=14)
    # 13 architecture
    s = sl[12]
    D.wipe(s, review_keep)
    D.title(s, "High-level proposed architecture", 32)
    D.picture(s, ARCH, 0.6, 1.3, 6.2, 4.9, caption="Four-tier architecture; Gemini is the only external service")
    D.picture(s, FLOW, 7.0, 1.3, 5.7, 4.9, caption="Eight-step workflow from GST data to a cited case note")
    # 14 expected outcomes
    s = sl[13]
    D.wipe(s, review_keep)
    D.title(s, "Expected outcomes and contributions", 30)
    oc = [("Label-free behaviour model", "A small JEPA trained without labels in seconds on a CPU.", ORANGE), ("Multi-layer risk", "Ablations to show each layer's role; JEPA is expected to complement, not replace, network evidence.", BLUE),
          ("Ring and cluster detection", "Scored loops and shell clusters drawn as an interactive graph.", GREEN), ("Cited notes", "Case notes whose citations and numbers are checked.", PURPLE),
          ("Honest evaluation", "Rule baseline, per-pattern recall, cross-seed study; synthetic data stated plainly.", RED), ("Usable system", "Secured multi-user application tested end to end.", "0E8F9B")]
    for i, (h, b, c) in enumerate(oc):
        D.card(s, 0.6 + (i % 3) * 4.1, 1.5 + (i // 3) * 2.6, 3.9, 2.4, h, b, c, str(i + 1), body_size=14, head_size=15)
    # 15 references
    s = sl[14]
    D.wipe(s, review_keep)
    D.title(s, "References", 32)
    D.text(s, 0.6, 1.3, 12.1, 5.4, [
        "[1] S. Datta, H. Thakkar and S. Dabhade, Int. J. Management, Public Policy and Research, 5(1), 2026 (base paper).",
        "[2]-[5] Four tax-fraud papers listed in the batch survey: AI/ML in tax compliance in India (2025); income-tax fraud detection with ML (ICOEI 2025) and with AI and ML (COMPUTINGCON 2025); boosting for income-tax fraud (2026).",
        "[6] P. Mehta et al., \"Representation learning on graphs to identifying circular trading in GST,\" arXiv:2208.07660, 2022.",
        "[7] Y. Xu et al., \"TED: Related party transaction guided tax evasion detection on heterogeneous graph,\" arXiv:2605.26984, 2026.",
        "[8] M. Assran et al., \"Self-supervised learning from images with a joint-embedding predictive architecture,\" CVPR 2023.",
        "[9] Y. LeCun, \"A path towards autonomous machine intelligence,\" 2022.",
        "[10] J. Pearl, Probabilistic Reasoning in Intelligent Systems, 1988 (noisy-OR)."], size=14, spacing=8)
    # 16 thank you: keep template text
    finish(prs, os.path.join(BATCH, "7.Review_1_PPT.pptx"), "Review 1")


# =====================================================================================================  REVIEW 2
def build_r2():
    prs = Presentation(os.path.join(TPL, "Review_2_PPT.pptx"))
    D = Deck(prs, "review")
    sl = list(prs.slides)
    seed = load_seed()
    edge = load_edge()
    fix_title_slide(D, sl[0], "REVIEW 2  ·  DESIGN, METHODOLOGY AND PROGRESS")
    # 3 recap
    s = sl[2]
    D.wipe(s, review_keep)
    D.title(s, "Recap: problem and proposed system", 32)
    D.card(s, 0.6, 1.5, 5.95, 2.4, "Problem", "ITC fraud (fake invoices, circular trading, shell companies) is hard to catch with audits and fixed rules; labels are scarce; flags are rarely explained.", ORANGE, "1", body_size=14)
    D.card(s, 6.75, 1.5, 5.95, 2.4, "Proposed system", "TaxSentinel: JEPA behaviour encoder + network and invoice analysis + noisy-OR risk + cited Gemini case notes, served by a secured FastAPI and React application.", BLUE, "2", body_size=14)
    D.flow(s, 0.6, 4.4, 12.1, 1.5, ["Returns, invoices, registrations", "19 features in 3 layers", "JEPA + network + invoice scores", "Risk, evidence, chains", "Cited case note"], size=13)
    D.text(s, 0.6, 6.15, 12.1, 0.5, "Data: seeded synthetic GST ecosystem with injected fraud labels (no real GST data was available).", size=13, italic=True, align="c")
    # 4 architecture
    s = sl[3]
    D.wipe(s, review_keep)
    D.title(s, "System architecture", 32)
    D.picture(s, ARCH, 0.6, 1.3, 7.4, 5.3)
    D.text(s, 8.3, 1.5, 4.4, 5.1, [
        ("Presentation: React 18, Vite, Tailwind; ten screens, route-level code splitting.", {}),
        ("Application: FastAPI, JWT guard, five routers, background generation thread.", {}),
        ("Analytics: ml/ package with features, JEPA, graph and pipeline.", {}),
        ("Stores: SQLite (users, runs, notes), pickled pipeline state, CSV ecosystem, model weights.", {}),
        ("External: Google Gemini for case notes.", {}),
        ("Read endpoints are lookups of a state computed once per ecosystem.", {"bold": True})], size=14, bullet=True, spacing=8)
    # 5 methodology
    s = sl[4]
    D.wipe(s, review_keep)
    D.title(s, "Methodology: end-to-end workflow", 32)
    D.flow(s, 0.6, 1.6, 12.1, 1.7, ["Generate ecosystem", "Build 19 features", "Train JEPA (no labels)", "Network + invoice scores", "Combine risk", "Evidence + notes", "Evaluate vs labels"], size=12, gap=0.3)
    D.plain_card(s, 0.6, 3.8, 3.95, 2.7, "Self-supervised", "Labels are never used for training; they are used only to score the result.", ORANGE, body_size=14)
    D.plain_card(s, 4.7, 3.8, 3.95, 2.7, "Layer attribution", "Four mask variants (whole month; invoice, behaviour or network layer hidden) show which layer deviates.", BLUE, body_size=14)
    D.plain_card(s, 8.8, 3.8, 3.95, 2.7, "Reproducible", "Seeded generator, torch and numpy seeds; seed 42 reproduces every reported figure.", GREEN, body_size=14)
    # 6 dataset
    s = sl[5]
    D.wipe(s, review_keep)
    D.title(s, "Dataset and input sources", 32)
    D.bar_chart(s, 0.6, 1.4, 6.3, 4.9, ["Shell entity", "Circular trading", "Fake-invoice buyer", "ITC spike"], [("Injected fraud taxpayers", [27, 21, 19, 12])], colors=[ORANGE], number_format="0", title="Injected fraud taxpayers (79 of 600)", legend=False, font_size=12)
    rows = [["Item", "Bundled ecosystem (seed 42)"], ["Period", "Apr 2025 - Mar 2026 (12 months)"], ["Taxpayers", "600 (8 states, 5 sectors)"], ["Invoices", "64,474 (530 fraud)"], ["Monthly returns", "6,977"], ["Fraud taxpayers", "79 (13.2 %)"],
            ["Hard negatives", "Festival seasons, capital purchases, new firms, group companies, late filers"], ["Transfer check", "PaySim, 100k rows (CC BY-SA 4.0)"]]
    D.table(s, 7.2, 1.5, 5.5, [1.7, 3.8], rows, size=12, max_h=4.8)
    D.text(s, 7.2, 6.2, 5.5, 0.5, "All GST data is synthetic sample data.", size=12, italic=True)
    # 7 preprocessing
    s = sl[6]
    D.wipe(s, review_keep)
    D.title(s, "Data preprocessing and feature layers", 30)
    D.card(s, 0.6, 1.5, 3.95, 4.8, "Invoice layer (6)", ["Purchase and sales invoice counts (log)", "Share of purchases not reported by suppliers", "Goods invoices of Rs 50,000+ without e-way bill", "Round-value share", "Share bought from suppliers under 6 months old"], ORANGE, "I", body_size=13, bullet=True)
    D.card(s, 4.7, 1.5, 3.95, 4.8, "Behaviour layer (7)", ["Outward turnover and ITC claimed (log)", "ITC to output tax", "ITC claimed above GSTR-2B", "Share of tax paid in cash", "Filing delay; registration age"], BLUE, "B", body_size=13, bullet=True)
    D.card(s, 8.8, 1.5, 3.95, 4.8, "Network layer (6)", ["Suppliers and buyers (log)", "Share of new partners", "Reciprocal sales share", "Largest buyer's share", "Value added"], GREEN, "N", body_size=13, bullet=True)
    D.text(s, 0.6, 6.4, 12.1, 0.4, "Robust scaling (median / IQR, clipped at 10); months without a return are marked absent, not imputed.", size=13, italic=True, align="c")
    # 8 model architecture
    s = sl[7]
    D.wipe(s, review_keep)
    D.title(s, "Model architecture and risk model", 32)
    D.picture(s, ALGO, 0.6, 1.2, 5.6, 5.5)
    rows = [["Component", "Specification"], ["Online encoder", "19 > 64 > 64 > 32 (LayerNorm, GELU)"], ["Target encoder", "EMA copy, momentum 0.99, no gradient"], ["Predictor", "135 > 128 > 128 > 32"], ["Mask variants", "4 (month, invoice, behaviour, network)"],
            ["Saved values", "53,420 (45,696 trainable)"], ["Risk", "1 - (1-0.75J)(1-0.6B)(1-0.7N)(1-0.6I)"]]
    D.table(s, 6.6, 1.5, 6.1, [1.8, 4.3], rows, size=12, max_h=4.0)
    D.text(s, 6.6, 5.3, 6.1, 1.2, "J = JEPA deviation, B = behaviour-layer deviation, N = network score, I = invoice score. Weights were set on the seed-42 ecosystem.", size=12, italic=True)
    # 9 training configuration
    s = sl[8]
    D.wipe(s, review_keep)
    D.title(s, "Training configuration", 32)
    rows = [["Setting", "Value"], ["Samples", "6,977 taxpayer-months, no labels"], ["Window", "3 previous months"], ["Optimiser", "AdamW, lr 1e-3, weight decay 1e-4"], ["Batch / epochs", "256 / 100"], ["Loss", "MSE in embedding space + variance hinge (weight 1)"],
            ["Hardware / time", "CPU only, 11.7 s"], ["Final loss", f"{TR['final_loss']:.4f}"]]
    D.table(s, 0.6, 1.5, 6.2, [2.0, 4.2], rows, size=13, max_h=4.8)
    D.picture(s, C("c04_jepa_training_loss.png"), 7.0, 1.6, 5.7, 3.6, caption="Loss per epoch: peaks at 0.261 (epoch 11) as the EMA target moves, then settles near 0.11")
    # 10 decision logic
    s = sl[9]
    D.wipe(s, review_keep)
    D.title(s, "Decision logic", 32)
    pic = C("c10_risk_distribution.png")
    if not os.path.exists(pic):
        pic = C("c02_method_comparison.png")
    if os.path.exists(pic):
        D.picture(s, pic, 0.6, 1.4, 5.3, 4.2, caption="Risk score of legitimate (grey) and injected-fraud (red) taxpayers, seed 42")
    D.plain_card(s, 6.2, 1.5, 6.5, 1.6, "Flag", "Risk R >= 0.5 flags a taxpayer; an invoice with anomaly score >= 0.5 is anomalous.", ORANGE, body_size=14)
    D.plain_card(s, 6.2, 3.3, 6.5, 1.6, "Pattern", "Named from the evidence: circular trading, shell entity, fake invoices, ITC spike or unusual behaviour.", BLUE, body_size=14)
    D.plain_card(s, 6.2, 5.1, 6.5, 1.5, "Evidence and note", "Evidence items E1..En are built per taxpayer; the note may restate only those, citing each.", GREEN, body_size=14)
    # 11 data flow and modules
    s = sl[10]
    D.wipe(s, review_keep)
    D.title(s, "Data flow and modules", 32)
    D.picture(s, DFD1, 0.6, 1.3, 6.3, 5.2, caption="DFD level 1 (levels 0 and 2 are in the report)")
    rows = [["Module", "File", "Lines"], ["Data generator", "scripts/generate_gst_data.py", "483"], ["Features", "ml/features.py", "169"], ["JEPA", "ml/jepa.py", "149"], ["Graph analysis", "ml/graph.py", "110"], ["Pipeline", "ml/pipeline.py", "530"],
            ["Evaluation", "ml/eval.py", "133"], ["API and services", "backend/app/", "787"], ["Web pages", "frontend/src/pages/", "1,415"]]
    D.table(s, 7.2, 1.5, 5.5, [1.7, 2.9, 0.9], rows, size=12, max_h=4.9)
    # 12 database and UML
    s = sl[11]
    D.wipe(s, review_keep)
    D.title(s, "Database design and UML", 32)
    rows = [["Table", "Columns (keys)"], ["users", "id PK, email unique, name, password_hash, created_at"], ["pipeline_runs", "id PK, seed, n_taxpayers, status, message, stats JSON, user_id FK, started_at, finished_at"], ["explanations", "id PK, run_key, gstin, model, text, evidence JSON, quality JSON, user_id FK, created_at"]]
    D.table(s, 0.6, 1.5, 6.0, [1.5, 4.5], rows, size=12, max_h=3.5)
    D.text(s, 0.6, 5.0, 6.0, 1.6, "Analytical data (taxpayers, invoices, returns) stays in CSV files joined by GSTIN; the three SQL tables hold users, runs and stored notes.", size=13, italic=True)
    D.picture(s, UML, 6.8, 1.4, 5.9, 5.0, caption="UML class diagram")
    # 13 technology stack
    s = sl[12]
    D.wipe(s, review_keep)
    D.title(s, "Technology stack", 32)
    items = [("Language", "Python 3.11.9, JavaScript (React 18)", ORANGE), ("Machine learning", "PyTorch 2.14 (CPU), NumPy, NetworkX 3.4", BLUE), ("Generative AI", "Google Gemini through google-genai; prompt + grounding checks", GREEN),
             ("Backend", "FastAPI 0.115, SQLAlchemy 2.0, PyJWT, bcrypt, SQLite", PURPLE), ("Frontend", "Vite 6, Tailwind 4, recharts, react-force-graph-2d, framer-motion", RED), ("Data and tools", "kagglehub (PaySim), pytest-free smoke test, setup.bat / run.bat", "0E8F9B")]
    for i, (h, b, c) in enumerate(items):
        D.card(s, 0.6 + (i % 3) * 4.1, 1.5 + (i // 3) * 2.6, 3.9, 2.4, h, b, c, str(i + 1), body_size=14, head_size=16)
    # 14 implementation progress
    s = sl[13]
    D.wipe(s, review_keep)
    D.title(s, "Implementation progress", 32)
    rows = [["Component", "Status", "Evidence"], ["Synthetic generator", "Done", "600 taxpayers, 64,474 invoices, 4 patterns"], ["Features and JEPA", "Done", "19 features; 11.7 s CPU training"], ["Graph, invoice and risk", "Done", "9 rings, 7 suspicious clusters, 16 chains"],
            ["Evidence and Gemini notes", "Done", "Cited notes with grounding checks"], ["API and web app", "Done", "10 screens, JWT, background runs"], ["Evaluation", "Done", "Baseline, ablations, per-pattern, cross-seed"], ["Unit tests", "Not done", "Smoke and black-box tests only"]]
    D.table(s, 0.6, 1.5, 8.2, [2.4, 1.1, 4.7], rows, size=13, max_h=5.0)
    D.tile(s, 9.1, 1.5, 3.6, 1.6, "20/20", "smoke-test checks passed", GREEN)
    D.tile(s, 9.1, 3.3, 3.6, 1.6, "11.7 s", "JEPA training on a CPU", BLUE)
    D.tile(s, 9.1, 5.1, 3.6, 1.6, "2,710", "lines of Python (ml, backend, scripts)", ORANGE)
    # 15 preliminary results
    s = sl[14]
    D.wipe(s, review_keep)
    D.title(s, "Preliminary results: seed-42 ecosystem", 32)
    ks = ["10", "25", "50", "79"]
    D.bar_chart(s, 0.6, 1.4, 8.3, 4.8, [f"Top {k}" for k in ks], [("TaxSentinel", [M["precision_at_k"][k] for k in ks]), ("Rule baseline", [B["precision_at_k"][k] for k in ks])], title="Precision@k", number_format="0.00", maxv=1.2, font_size=12)
    D.tile(s, 9.2, 1.5, 3.5, 1.5, f3(M["at_threshold"]["f1"]), f"F1 at threshold (rules {f3(B['at_threshold']['f1'])})", GREEN)
    D.tile(s, 9.2, 3.15, 3.5, 1.5, f3(M["roc_auc"]), f"ROC-AUC (rules {f3(B['roc_auc'])})", BLUE)
    D.tile(s, 9.2, 4.8, 3.5, 1.5, f3(M["precision_at_k"]["79"]), f"precision@79 (rules {f3(B['precision_at_k']['79'])})", ORANGE)
    note = "Synthetic data; weights tuned on this seed."
    if seed:
        note += f" Mean F1 on five other seeds: {f3(seed['mean']['f1'])} (range {f3(seed['min']['f1'])} to {f3(seed['max']['f1'])})."
    D.text(s, 0.6, 6.3, 12.1, 0.5, note, size=13, italic=True, align="c")
    # 16 output screens
    s = sl[15]
    D.wipe(s, review_keep)
    D.title(s, "Output screens", 32)
    D.picture(s, SCR("ov_done"), 0.6, 1.4, 6.0, 4.6, caption="Overview after a run: 78 flagged taxpayers, ITC at risk, precision against rules")
    D.picture(s, SCR("tp_sig"), 6.8, 1.4, 6.0, 4.6, caption="Top-ranked taxpayer: layer scores and monthly GSTR-2B vs GSTR-3B")
    # 17 testing
    s = sl[16]
    D.wipe(s, review_keep)
    D.title(s, "Initial testing and observations", 32)
    ec = edge
    rows = [["Test", "Result"], ["End-to-end smoke test", "20 / 20 checks passed (39.4 s)"], ["Edge and security checks", f"{ec['passed']} / {ec['total']} passed"], ["Median API latency", "1.7 to 12.1 ms (read endpoints)"], ["Full generate + analyse", "18.6 to 29.0 s typical (600 taxpayers)"], ["Unit tests", "None; stated as a gap"]]
    D.table(s, 0.6, 1.5, 5.8, [2.4, 3.4], rows, size=13, max_h=4.0)
    D.text(s, 6.8, 1.5, 5.9, 5.0, [
        "A wrong password, an altered token, a missing token, bad input and a second concurrent run are all rejected with the expected status codes.",
        "SQL-injection-style and script-tag search text is inert.",
        "Timings were taken on a machine at times heavily loaded by other jobs.",
        "Gemini quota and overload errors (429, 503, 504) stopped 4 of 10 case notes in the evaluation."], size=14, bullet=True, spacing=8)
    # 18 challenges
    s = sl[17]
    D.wipe(s, review_keep)
    D.title(s, "Challenges and corrective actions", 32)
    ch = [("No real GST data", "Seeded generator with hard negatives; every figure labelled synthetic.", ORANGE), ("JEPA misses shell firms", "Similar shells form their own normal; the network layer catches them (JEPA-only recall 0.20).", BLUE),
          ("Weights tuned on one seed", "Cross-seed study run; headline figures treated as a best case.", GREEN), ("Gemini quota and overload", "Retries, fallback models, notes on demand; failures reported, not hidden.", PURPLE),
          ("Legitimate loops look like fraud", "2 of 9 detected rings are legitimate reciprocal trade; notes suggest the checks that clear them.", RED), ("Loaded machine", "Timings varied widely; typical and worst runs are both reported.", "0E8F9B")]
    for i, (h, b, c) in enumerate(ch):
        D.card(s, 0.6 + (i % 3) * 4.1, 1.5 + (i // 3) * 2.6, 3.9, 2.4, h, b, c, str(i + 1), body_size=13, head_size=15)
    # 19 next steps
    s = sl[18]
    D.wipe(s, review_keep)
    D.title(s, "Next steps", 32)
    D.flow(s, 0.6, 1.8, 12.1, 1.5, ["Final report and IEEE paper", "Head-to-head baselines", "Unit-test suite", "Final presentation and demo"], size=14, gap=0.4)
    D.text(s, 0.6, 3.9, 12.1, 2.7, [
        "Train an isolation forest and an autoencoder on the same features and compare them with the JEPA.",
        "Re-fit or learn the risk weights and threshold so that cross-seed F1 is more stable.",
        "Add unit tests and automated UI tests; run a user-acceptance session with analysts.",
        "Evaluate the case notes on all ten planned cases (paid or retried Gemini quota)."], size=16, bullet=True, spacing=10)
    finish(prs, os.path.join(BATCH, "13.Review_2_PPT.pptx"), "Review 2")


# =====================================================================================================  FINAL
def final_keep(shp):
    if shp.name in ("Shape 0", "Text 1", "Text 2", "Text 3"):
        return True
    if shp.shape_type == 13 and shp.name == "Image 0" and shp.top < 1000000:
        return True
    return False


def build_final():
    prs = Presentation(os.path.join(TPL, "Final_PPT.pptx"))
    D = Deck(prs, "final")
    sl = list(prs.slides)
    seed = load_seed()
    edge = load_edge()
    err = load_err()
    foot = "Dept. of CSE (Data Science)  |  Batch A5  |  Final Review"
    # slide 1 title
    s = sl[0]
    for shp in s.shapes:
        if shp.name == "Text 1":
            D.set_runs(shp, [PROJECT])
            for r in shp.text_frame.paragraphs[0].runs:
                r.font.size = L.Pt(28)
        elif shp.name == "Text 2":
            D.set_runs(shp, ["Final Review  ·  TaxSentinel"])
        elif shp.name == "Text 3":
            D.set_runs(shp, ["Team Details : Batch A5 (CSE-DS)"])
        elif shp.name == "Text 4":
            D.set_runs(shp, [f"{r}   {n}" for n, r in STUDENTS])
        elif shp.name == "Text 5":
            D.set_runs(shp, ["Under the Esteemed Guidance of", GUIDE, GUIDE_DESIG])
        elif shp.name == "Text 6":
            D.set_runs(shp, [f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  ·  Academic Year {YEARS.replace('-', ' – ')}"])
    # regular slides: set footer, number, wipe content
    def setup(i, title, size=None):
        s = sl[i]
        D.wipe(s, final_keep)
        D.title(s, title, size or 30)
        for shp in s.shapes:
            if shp.name == "Text 2":
                D.set_runs(shp, [foot])
            if shp.name == "Text 3":
                D.set_runs(shp, [str(i + 1)])
        return s

    def body(s, paras, size=18, y=1.9, h=4.6, spacing=8, bullet=True):
        D.text(s, 0.6, y, 12.1, h, paras, size=max(size, 22), bullet=bullet, spacing=spacing + 2)

    s = setup(1, "Table of Contents", 32)
    D.text(s, 0.9, 1.9, 5.6, 4.8, ["1. Abstract", "2. Base paper and literature survey", "3. Research gap and problem statement", "4. Objectives", "5. Proposed system and architecture", "6. Methodology, dataset and model", "7. Decision workflow and design"], size=20, spacing=10)
    D.text(s, 6.9, 1.9, 5.8, 4.8, ["8. Implementation highlights", "9. Results and threshold analysis", "10. Performance and testing", "11. Comparison with the base paper", "12. Output screens", "13. Conclusion and future scope", "14. References"], size=20, spacing=10)
    s = setup(2, "Abstract", 32)
    body(s, ["ITC fraud under GST (fake invoices, circular trading, shell companies) causes large revenue losses; audits and fixed rules are slow and easy to evade, labels are scarce and flags are rarely explained. The base paper describes the problem but proposes no model.",
             "TaxSentinel learns normal GST behaviour without labels with a JEPA encoder (53,420 values, 11.7 s on a CPU), combines its deviation with circular-trading, shell-cluster and invoice evidence in a noisy-OR risk, and has Gemini write case notes that cite numbered evidence.",
             f"On a synthetic ecosystem (600 taxpayers, 64,474 invoices, 79 injected fraud taxpayers) it reaches precision@79 {f3(M['precision_at_k']['79'])}, F1 {f3(M['at_threshold']['f1'])} and ROC-AUC {f3(M['roc_auc'])} against {f3(B['precision_at_k']['79'])}, {f3(B['at_threshold']['f1'])} and {f3(B['roc_auc'])} for rules; all 5 injected rings were recovered.",
             "Limits: data is synthetic; the LLM is prompted, not fine-tuned; JEPA alone is the weakest layer; F1 on other generated ecosystems is lower."], size=16, bullet=False, spacing=10)
    s = setup(3, "Base Paper", 32)
    body(s, ["S. Datta, H. Thakkar and S. Dabhade, \"Combating GST Input Tax Credit Fraud in India: Forensic Accounting, Data Analytics, and Legal Enforcement Strategies,\" Int. J. Management, Public Policy and Research, vol. 5, no. 1, Jan-Mar 2026, doi 10.55829/ykz3ea41.",
             "Qualitative synthesis of public DGGI and CBIC data: detected evasion Rs 1.01 to Rs 2.01 lakh crore (FY23 to FY24), cases 4,872 to 6,084, arrests 97 to 140.",
             "Recommends analytics, network analysis, AI-based risk modelling at registration, forensic capacity and faster recovery.",
             "Limitations: public sources only; no detection model and no accuracy figures. This project builds and evaluates the model it calls for."], size=17, spacing=10)
    s = setup(4, "Literature Survey", 32)
    rows = [["Work", "Approach", "Data", "Reported result", "Limitation"],
            ["Datta et al. [1]", "Forensic analytics review", "DGGI, CBIC data", "Evasion Rs 1.01 to 2.01 lakh crore", "No model"],
            ["[2]-[5] (titles only)", "ML / AI for tax fraud", "Not available", "Not quoted", "Supervised; no network"],
            ["Mehta et al. [6]", "node2vec + DBSCAN loops", "Telangana GST, 77.8 M invoices", "Communities found", "One pattern; no accuracy"],
            ["Xu et al. [7] (TED)", "Heterogeneous GNN", "Two real tax datasets", "F1 76.35 to 94.96 %", "Needs labels"],
            ["Assran et al. [8]", "I-JEPA", "Images", "Strong transfer", "Not tax data"]]
    D.table(s, 0.5, 1.9, 12.3, [2.2, 2.6, 2.6, 2.6, 2.3], rows, size=13, max_h=4.6)
    s = setup(5, "Research Gap", 32)
    rows = [["Gap", "Existing work", "TaxSentinel"], ["Labels", "Supervised models need confirmed fraud", "Learns normal behaviour without labels"], ["Evidence", "One view (returns or graph)", "Invoice, behaviour and network layers in one risk"],
            ["Group fraud", "Per-entity scores miss loops and shell clusters", "Scored loops (5 of 5 found) and clusters"], ["Explanation", "Scores or communities only", "Numbered evidence and cited notes, checked automatically"], ["Evaluation", "Rarely against a rule baseline", "Baseline, ablations, per-pattern and cross-seed results"]]
    D.table(s, 0.5, 1.9, 12.3, [1.6, 4.3, 5.4], rows, size=15, max_h=4.6)
    s = setup(6, "Problem Statement", 32)
    body(s, ["ITC fraud through fake invoices, circular trading and shell companies causes large revenue losses.", "Manual audits and rule-based checks are slow and reactive.", "Confirmed fraud labels are scarce, which limits supervised learning.", "Existing systems rarely explain why a claim is flagged.",
             ("Need: a self-supervised, explainable framework that learns normal GST behaviour and flags and explains suspicious ITC claims.", {"bold": True})], size=20, spacing=12)
    s = setup(7, "Objectives", 32)
    body(s, ["Build a multi-layer GST dataset (invoices, behaviour, network).", "Learn normal GST behaviour with a self-supervised JEPA encoder.", "Detect fake invoices, circular trading and shell-entity activity.", "Score and rank taxpayers and invoice chains by deviation.",
             "Explain flagged claims with an LLM (prompt-tuned and checked, not fine-tuned).", "Evaluate detection accuracy and explanation quality."], size=20, spacing=10)
    s = setup(8, "Proposed System", 32)
    body(s, ["1. Generate or load a GST ecosystem (taxpayers, invoices, GSTR returns).", "2. Build 19 features per taxpayer-month in invoice, behaviour and network layers.", "3. JEPA encoder scores deviation (4 mask variants).", "4. Network analysis finds loops and shared-identity clusters.",
             "5. Eight-signal invoice score; noisy-OR risk; flag at 0.5.", "6. Evidence list, patterns and ranked invoice chains.", "7. Gemini writes a cited note; grounding checks run; analyst reviews."], size=18, bullet=False, spacing=8)
    s = setup(9, "System Architecture", 32)
    D.picture(s, ARCH, 1.6, 1.9, 10.1, 4.8, caption="Four-tier architecture: React, FastAPI, analytics layer, stores; Gemini external")
    s = setup(10, "Methodology", 32)
    D.picture(s, FLOW, 1.6, 1.9, 10.1, 4.8, caption="End-to-end processing flow")
    s = setup(11, "Dataset, Features & Model", 32)
    body(s, ["Synthetic seeded ecosystem (seed 42): 600 taxpayers, 64,474 invoices, 6,977 returns; 79 fraud taxpayers (27 shell, 21 circular, 19 fake-invoice buyers, 12 ITC spikes); hard negatives included.",
             "19 features per taxpayer-month: invoice (6), behaviour (7), network (6); robust-scaled.",
             "JEPA: encoder 19-64-64-32, EMA target, predictor 135-128-128-32, window 3, 4 mask variants; AdamW, 100 epochs, 53,420 values, 11.7 s on a CPU.",
             "Network: bounded cycles scored by balance, recurrence, concentration; address and contact clusters; Louvain communities.",
             "Risk = noisy-OR(0.75 J, 0.6 B, 0.7 N, 0.6 I), flag at 0.5. LLM: Gemini, prompted and grounded."], size=16, spacing=9)
    s = setup(12, "Decision Workflow", 32)
    D.picture(s, ALGO, 3.4, 1.9, 6.5, 4.8, caption="Risk scoring: combine layers, flag at 0.5, evidence, note")
    s = setup(13, "User–System Interaction", 32)
    D.picture(s, SEQ, 2.4, 1.9, 8.5, 4.8, caption="Sequence: sign in, generate and analyse (background thread), open taxpayer, write explanation")
    s = setup(14, "Database Design", 32)
    D.picture(s, ER, 2.2, 1.9, 8.9, 4.8, caption="Entity-relationship diagram: 3 SQL tables and the CSV entities joined by GSTIN")
    s = setup(15, "Implementation Highlights", 32)
    rows = [["Component", "Implementation", "Evidence"], ["Generator", "Seeded; GSTIN checksum; 4 patterns; hard negatives", "64,474 invoices, 79 fraud taxpayers"], ["JEPA", "PyTorch, EMA target, 4 mask variants", "53,420 values; 11.7 s on a CPU"], ["Graph", "NetworkX bounded cycles; identity clusters", "5 of 5 rings; 9 rings, 7 suspicious clusters"],
            ["Explainer", "Gemini; evidence-only prompt; grounding checks", "0 invalid citations in 6 notes"], ["API", "FastAPI, JWT, bcrypt, SQLite, background runs", "20/20 smoke-test checks"], ["Web app", "React, Vite, Tailwind, force graph", "10 screens, dark mode, command palette"]]
    D.table(s, 0.5, 1.9, 12.3, [1.7, 5.4, 5.2], rows, size=14, max_h=4.6)
    # 16 results
    s = setup(16, "Results – Model Performance (seed 42, threshold 0.5)", 28)
    rows = [["Method", "P@79", "Precision", "Recall", "F1", "ROC-AUC"],
            ["TaxSentinel (all layers)", f3(M["precision_at_k"]["79"]), f3(M["at_threshold"]["precision"]), f3(M["at_threshold"]["recall"]), f3(M["at_threshold"]["f1"]), f3(M["roc_auc"])],
            ["Rule baseline R1-R5", f3(B["precision_at_k"]["79"]), f3(B["at_threshold"]["precision"]), f3(B["at_threshold"]["recall"]), f3(B["at_threshold"]["f1"]), f3(B["roc_auc"])],
            ["JEPA + behaviour only", f3(AB["jepa_only"]["precision_at_k"]["79"]), f3(AB["jepa_only"]["at_threshold"]["precision"]), f3(AB["jepa_only"]["at_threshold"]["recall"]), f3(AB["jepa_only"]["at_threshold"]["f1"]), f3(AB["jepa_only"]["roc_auc"])],
            ["Graph + invoice only", f3(AB["graph_and_invoice_only"]["precision_at_k"]["79"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["precision"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["recall"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["f1"]), f3(AB["graph_and_invoice_only"]["roc_auc"])]]
    if seed:
        mm = seed["mean"]
        rows.append([f"Five other seeds (mean), TaxSentinel", f3(mm["p_at_k"]), f3(mm["precision"]), f3(mm["recall"]), f3(mm["f1"]), f3(mm["roc_auc"])])
        rows.append([f"Five other seeds (mean), rules", f3(mm["base_p_at_k"]), "-", "-", f3(mm["base_f1"]), f3(mm["base_roc_auc"])])
    _, h = D.table(s, 0.5, 1.9, 12.3, [3.8, 1.5, 1.8, 1.6, 1.6, 1.9], rows, size=14, max_h=4.2)
    D.text(s, 0.5, 1.9 + h + 0.15, 12.3, 0.9, "Synthetic ecosystem with injected labels; weights and threshold were tuned on seed 42, so the first rows are a best case. JEPA alone is the weakest layer (F1 0.327); the layers are complementary. Invoice level: F1 0.924 vs 0.160 for the GSTR-1 mismatch check.", size=13, italic=True)
    s = setup(17, "Confusion Matrix & Recall by Pattern", 30)
    D.picture(s, C("c08_confusion_matrices.png"), 0.5, 1.9, 6.4, 4.3, caption="Test: TP 75, FN 4, FP 3, TN 518 (rules: TP 68, FP 168)")
    D.picture(s, C("c03_recall_by_pattern.png"), 7.0, 1.9, 5.8, 4.3, caption="Recall by pattern: ITC spikes 0.83, below the rule baseline's 1.00")
    s = setup(18, "Threshold Analysis & Layer Scores", 30)
    p9, p12 = C("c09_threshold_sweep.png"), C("c12_layer_scores.png")
    if os.path.exists(p9):
        D.picture(s, p9, 0.5, 1.9, 6.0, 4.2, caption="Flag threshold sweep (seed 42): F1 is highest around 0.4 to 0.6; 0.5 chosen")
    if os.path.exists(p12):
        D.picture(s, p12, 6.8, 1.9, 6.0, 4.2, caption="Network and invoice layers separate fraud best; JEPA adds recall")
    s = setup(19, "Performance & Testing", 32)
    rows = [["Measure", "Result"], ["JEPA training (CPU)", "11.7 s, final loss 0.1115"], ["Full generate + analyse, 600 taxpayers", "18.6 to 29.0 s typical (analysis stage 16.0 s)"], ["Read-endpoint latency (median of 20)", "1.7 ms (health) to 12.1 ms (whole graph)"],
            ["End-to-end smoke test", "20 / 20 checks passed (39.4 s)"], ["Edge and security checks", f"{edge['passed']} / {edge['total']} passed (wrong password, forged token, bad input, concurrent run)"], ["Unit tests / UI tests / load tests", "Not done (stated as gaps); single-client timings on a loaded machine"]]
    D.table(s, 0.5, 1.9, 12.3, [4.5, 7.8], rows, size=14, max_h=4.6)
    s = setup(20, "Comparison with Base Paper (context only)", 28)
    rows = [["Aspect", "Base paper [1]", "TaxSentinel"], ["Type", "Qualitative review", "Implemented and evaluated system"], ["Detection model", "None", "JEPA + network + invoice noisy-OR"], ["Reported accuracy", "None", f"F1 {f3(M['at_threshold']['f1'])} (seed 42, synthetic); lower on other seeds"],
            ["Explanation", "Not addressed", "Cited, checked case notes"], ["Registration-time scoring", "Recommended", "Not built"], ["Context only: TED [7]", "-", "F1 76 to 95 % on real data, supervised: not comparable"]]
    D.table(s, 0.5, 1.9, 12.3, [3.0, 4.0, 5.3], rows, size=14, max_h=4.2)
    D.text(s, 0.5, 6.1, 12.3, 0.6, "The base paper reports no model metrics, so there is no like-for-like number; our comparison is with the recommendations and with a rule baseline.", size=13, italic=True)
    # 21-23 output screens
    pairs = [(20, "Output – Generate & Analyse", ("ov_run", "Run in progress: JEPA training epoch 10/100"), ("ov_done", "Overview: 78 flagged taxpayers, ITC at risk, precision vs rules")),
             (21, "Output – Taxpayer Investigation", ("tp_sig", "Profile: four layer scores, GSTR-2B vs 3B chart"), ("tp_net", "Network tab: circular-trading loop in red")),
             (22, "Output – Case Note & Invoice Chain", ("case1", "Gemini case note with [E#] citations beside evidence"), ("chain_c", "Chain CH-R01: 5-party loop, 98 invoices, Rs 44.65 Cr"))]
    for idx, t, a, b in pairs:
        s = setup(idx + 1, t, 30)
        D.picture(s, SCR(a[0]), 0.5, 1.9, 6.1, 4.3, caption=a[1])
        D.picture(s, SCR(b[0]), 6.8, 1.9, 6.1, 4.3, caption=b[1])
    s = setup(24, "Conclusion", 32)
    body(s, ["Built TaxSentinel: label-free behaviour learning, network and invoice evidence and cited case notes in a secured web application.",
             f"Seed 42: precision@79 {f3(M['precision_at_k']['79'])}, F1 {f3(M['at_threshold']['f1'])}, ROC-AUC {f3(M['roc_auc'])} vs 0.620 / 0.432 / 0.876 for rules; 5 of 5 rings; invoice F1 0.924."]
         + ([f"Honest limits: mean F1 {f3(seed['mean']['f1'])} on five other seeds (range {f3(seed['min']['f1'])} to {f3(seed['max']['f1'])}); data is synthetic; JEPA alone F1 0.327; LLM prompted, not fine-tuned; 6 of 10 notes evaluated; no unit-test suite."] if seed else ["Limits: synthetic data; JEPA alone is weakest; LLM prompted, not fine-tuned."]), size=18, spacing=10)
    s = setup(25, "Future Scope", 32)
    body(s, ["Validate on real, anonymised returns; recalibrate or learn the risk weights and threshold.", "Compare the JEPA with an isolation forest and an autoencoder on the same features.", "Add e-invoice, e-way-bill and payment data; scale the graph analysis to state size.", "Role-based access, case management and feedback loop; HTTPS, rate limiting, audit logging.",
             "Fine-tune or self-host a language model and run a human evaluation of notes; add unit and UI tests."], size=18, spacing=10)
    s = setup(26, "References", 32)
    short = ["[1] S. Datta, H. Thakkar, S. Dabhade, Int. J. Manage., Public Policy Res., 5(1), 2026.", "[2]-[5] Tax-fraud ML and AI papers listed in the batch survey (2025-2026).", "[6] P. Mehta et al., arXiv:2208.07660, 2022 (circular trading in GST).", "[7] Y. Xu et al., TED, arXiv:2605.26984, 2026.",
             "[8] M. Assran et al., I-JEPA, CVPR 2023.", "[9] Y. LeCun, A path towards autonomous machine intelligence, 2022.", "[10] J.-B. Grill et al., BYOL, NeurIPS 2020.", "[11] A. Bardes et al., VICReg, ICLR 2022.", "[12] J. Pearl, Probabilistic Reasoning in Intelligent Systems, 1988.", "[13] Gemini Team, arXiv:2312.11805, 2023."]
    body(s, short, size=15, spacing=6, bullet=False)
    s = sl[27]
    for shp in s.shapes:
        if shp.name == "Text 3":
            D.set_runs(shp, [foot])
    finish(prs, os.path.join(BATCH, "16.Final_PPT.pptx"), "Final Review")


if __name__ == "__main__":
    what = sys.argv[1] if len(sys.argv) > 1 else "all"
    if what in ("r1", "all"):
        build_r1()
    if what in ("r2", "all"):
        build_r2()
    if what in ("final", "all"):
        build_final()
    for w in L.WARN:
        print("WARN", w)
