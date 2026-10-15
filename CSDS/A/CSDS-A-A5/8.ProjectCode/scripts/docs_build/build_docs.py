"""Build the six chapter documents (9, 10, 11, 12, 14, 15) and the final report (17) at the batch root.

Usage (from 8.ProjectCode, system python):  python scripts/docs_build/build_docs.py [chapters|final|all]
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))

from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402

from common import *  # noqa: E402,F401,F403
from docbuild import *  # noqa: E402,F401,F403
from docbuild import _runs, _set_font  # noqa: E402
import content_a as A  # noqa: E402
import content_b as B  # noqa: E402
import content_c as Cc  # noqa: E402
import content_d as Dd  # noqa: E402
import content_e as E  # noqa: E402
import content_f as Ff  # noqa: E402
import content_g as G  # noqa: E402
import content_h as H  # noqa: E402

OUTDIR = BATCH
NOTE_NUMBERING = "Section, table and figure numbers in this document follow the Final Document (17.Final Document.docx) where a section is cited; tables and figures are numbered within this document."


def refs_block(heading="REFERENCES"):
    def raw(doc):
        for i, r in enumerate(REFS, 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.45)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _runs(p, f"[{i}]\t" + r, 11)
    return [("h1", heading), ("raw", raw)]


def front_header(doc, subtitle):
    center(doc, DEPT, 12, True, after=2)
    center(doc, "CSDS-A  |  Batch A5", 12, after=8)
    center(doc, PROJECT, 14, True, after=6)
    center(doc, subtitle, 14, True, after=8, caps=True)
    student_table(doc, STUDENTS)
    center(doc, f"Project guide: {GUIDE}, {GUIDE_DESIG}", 12, before=6, after=2)
    center(doc, f"{COLLEGE}  |  A.Y. {YEAR}", 11, after=10)


def intro_note(text):
    return [("note", text), ("note", NOTE_NUMBERING)]


def write_doc(filename, subtitle, blocks, chap=False):
    doc = new_document()
    front_header(doc, subtitle)
    render(doc, blocks, Numbering(chap))
    path = os.path.join(OUTDIR, filename)
    doc.save(path)
    print("saved", filename, os.path.getsize(path) // 1024, "KB")
    return path


USED = []


def all_blocks():
    """Build every content section once so that all documents are consistent."""
    R = G.results_blocks()
    T = Dd.testing_blocks()
    W = Ff.code_walkthroughs()
    return R, T, W


def doc9(R, T, W):
    b = intro_note("This document describes how TaxSentinel works end to end, the algorithms used and why they were chosen, what is new against the base paper and the contributions of the project.")
    b += B.DES["flow"]
    for k in ("overall", "data", "prep", "choice", "math", "train", "pipeline", "llm"):
        b += Cc.METH[k]
    b += Cc.NOV["novelty"] + Cc.NOV["contrib"]
    return write_doc("9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "Methodology, Novelty and Key Contributions", b)


def doc10():
    b = intro_note("This document states what the project covers, what it deliberately does not, the assumptions it rests on and its honest limitations.")
    b += A.INTRO["scope"] + E.SCOPE["glance"] + E.SCOPE["in"] + E.SCOPE["out"] + E.SCOPE["boundaries"] + E.SCOPE["assume"] + E.SCOPE["deviation"] + E.SCOPE["limits"] + A.INTRO["limits"]
    return write_doc("10.SCOPE & LIMITATIONS.docx", "Scope and Limitations", b)


def doc11():
    b = intro_note("Every diagram of the system is collected here with a caption and an explanation. The diagrams were drawn from the code (models, routes and pipeline) and each matches what the implementation does.")
    for k in ("arch", "flow", "dfd", "uml", "db"):
        b += B.DES[k]
    b += Cc.METH["overall"]
    return write_doc("11. BLOCK DIAGRAMS ALL.docx", "Block Diagrams (All)", b)


def doc12(W):
    b = intro_note("This document records the design of TaxSentinel, how it was implemented, the code that matters most and the reasons behind the main implementation decisions.")
    for k in ("arch", "modules", "db", "uml"):
        b += B.DES[k]
    b += Cc.METH["app"] + Cc.METH["api"] + Cc.METH["sec"] + B.DES["ui"]
    b += [("h2", "Key Code Walkthroughs")]
    for k in ("jepa", "score", "cycles", "risk", "explain", "auth", "bg"):
        b += W[k]
    b += [("h2", "Implementation Decisions and Reasons"),
          ("p", "The table below lists the main decisions, what else was possible and why the choice was made. Where an alternative was not tested, the table says so; the choice is then a design argument and not a measured result.")]
    b += H.DECISIONS
    b += Cc.METH["install"]
    return write_doc("12.DESIGN &IMPLEMENTATION.docx", "Design and Implementation", b)


def doc14(R, T):
    b = intro_note("This document presents the measured results of the project: model performance with tables and charts, every output screen of the application with a caption, success and failure cases, a comparison with the base paper and the objectives measured against the outcomes. All numbers come from the project's own experiment files.")
    for k in ("screens", "perf", "expl", "sample", "sf", "cmp", "limits", "obj"):
        b += R[k]
    # remaining screens (every output screen is shown)
    rest = [k for k in SCREENS if k not in G.BODY_SCREENS]
    b += [("h2", "8.9 Further Output Screens")]
    for k in rest:
        b.append(SCR(k, 5.8, label=f"r14_{k}"))
    return write_doc("14.RESULTS.docx", "Results", b)


def doc15():
    b = intro_note("This document gives the conclusions of the project, its future scope and the numbered reference list (base paper first, then the literature-survey sources).")
    b += G.conclusion_blocks()["all"]
    b += [("h1", "FUTURE SCOPE")] + E.FUT["all"]
    b += refs_block()
    return write_doc("15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx", "Conclusion, Future Scope and References", b)


# ------------------------------------------------------------------ final report
def chapter(n, title, blocks):
    return [("h1", f"CHAPTER {n}: {title}")] + blocks


def final_blocks(R, T, W):
    a = A.INTRO
    blocks = chapter(1, "INTRODUCTION", sum((a[k] for k in ("domain", "background", "problem", "need", "objectives", "scope", "contrib", "limits", "org")), []))
    blocks += chapter(2, "LITERATURE SURVEY", sum((A.LIT[k] for k in ("overview", "base", "related", "tables", "gap")), []))
    blocks += chapter(3, "FEASIBILITY STUDY", A.FEAS["all"])
    blocks += chapter(4, "SYSTEM REQUIREMENTS AND ANALYSIS", sum((B.REQ[k] for k in ("existing", "proposed", "func", "nonfunc", "hw_sw", "users")), []))
    blocks += chapter(5, "SYSTEM DESIGN", sum((B.DES[k] for k in ("arch", "flow", "dfd", "uml", "db", "modules", "ui")), []))
    blocks += chapter(6, "METHODOLOGY AND SOFTWARE IMPLEMENTATION", sum((Cc.METH[k] for k in ("overall", "data", "prep", "choice", "math", "train", "pipeline", "llm", "app", "api", "sec", "install")), []))
    blocks += chapter(7, "SYSTEM TESTING", sum((T[k] for k in ("strategy", "unit", "integration", "functional", "security", "validation", "perf", "uat", "limits")), []))
    blocks += chapter(8, "RESULTS AND OUTPUT SCREENS", sum((R[k] for k in ("screens", "perf", "expl", "sample", "sf", "cmp", "limits", "obj")), []))
    blocks += chapter(9, "CONCLUSION", G.conclusion_blocks()["all"])
    blocks += chapter(10, "FUTURE SCOPE", E.FUT["all"])
    blocks += refs_block("REFERENCES")
    blocks += Ff.appendix_a() + Ff.appendix_b() + Ff.appendix_c(set(G.BODY_SCREENS))
    return blocks


def set_section_numbering(section, fmt, start=1):
    sectPr = section._sectPr
    for e in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(e)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)


def toc_lines(doc, entries, pages):
    """entries: list of (indent_level, text, key). pages: dict key->page or None."""
    for lvl, text, key in entries:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(1 if lvl else 2)
        pf.space_before = Pt(0 if lvl else 3)
        pf.left_indent = Inches(0.0 if lvl == 0 else 0.3)
        pf.first_line_indent = Inches(0)
        pf.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        pg = pages.get(key, "0") if pages else "0"
        run = p.add_run(f"{text}\t{pg}")
        _set_font(run, 12 if lvl == 0 else 11, True if lvl == 0 else False)


def build_final(pages=None):
    R, T, W = all_blocks()
    blocks = final_blocks(R, T, W)
    num = Numbering(chap=True)
    num.scan(blocks)
    # TOC entries (h1 and h2 in order), figure and table lists
    heads = [(1 if b[0] == "h1" else 2, b[1]) for b in blocks if b[0] in ("h1", "h2")]
    toc_entries = [(0 if lv == 1 else 1, t, ("h", i)) for i, (lv, t) in enumerate(heads)]
    fig_entries = [(0, f"Fig. {lab}: {cap}", ("f", i)) for i, (lab, cap, _k) in enumerate(num.fig_caps)]
    tab_entries = [(0, f"Table {lab}: {cap}", ("t", i)) for i, (lab, cap, _k) in enumerate(num.tab_caps)]
    # resolve tokens in captions of lists
    fig_entries = [(a, resolve(t, num), k) for a, t, k in fig_entries]
    tab_entries = [(a, resolve(t, num), k) for a, t, k in tab_entries]

    doc = new_document()
    clear_footer(doc.sections[0])
    H.cover(doc)
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    clear_footer(s2)
    add_page_number_footer(s2)
    set_section_numbering(s2, "lowerRoman", 1)
    H.certificate(doc)
    doc.add_page_break()
    H.declaration(doc)
    doc.add_page_break()
    H.acknowledgement(doc)
    doc.add_page_break()
    H.abstract(doc)
    doc.add_page_break()
    center(doc, "TABLE OF CONTENTS", 14, True, before=0, after=10)
    toc_lines(doc, toc_entries, pages)
    doc.add_page_break()
    center(doc, "LIST OF FIGURES", 14, True, before=0, after=10)
    toc_lines(doc, fig_entries, pages)
    doc.add_page_break()
    center(doc, "LIST OF TABLES", 14, True, before=0, after=10)
    toc_lines(doc, tab_entries, pages)
    doc.add_page_break()
    H.abbreviations(doc)
    s3 = doc.add_section(WD_SECTION.NEW_PAGE)
    clear_footer(s3)
    add_page_number_footer(s3)
    set_section_numbering(s3, "decimal", 1)
    # body: render with page break before each chapter / appendix heading
    body = []
    first = True
    for b in blocks:
        if b[0] == "h1" and not first:
            body.append(("pb",))
        if b[0] == "h1":
            first = False
        body.append(b)
    render(doc, body, num)
    return doc, heads, num


def final(passes=2):
    import wordtools
    path = os.path.join(OUTDIR, "17.Final Document.docx")
    pdf = os.path.join(ROOT, "docs", "_final_report.pdf")
    pages = None
    info = {}
    for i in range(passes):
        doc, heads, num = build_final(pages)
        doc.save(path)
        total = wordtools.export_pdf(path, pdf)
        head_txt = [resolve(t, num) for _lv, t in heads]
        fig_txt = [f"Figure {lab}: {resolve(cap, num)}" for lab, cap, _k in num.fig_caps]
        tab_txt = [f"Table {lab}: {resolve(cap, num)}" for lab, cap, _k in num.tab_caps]
        pages, body_start, npdf = wordtools.locate_pages(pdf, head_txt, fig_txt, tab_txt)
        missing = [k for k, v in pages.items() if v == "?"]
        print(f"pass {i + 1}: {total} pages in total, front matter {body_start} pages, body {npdf - body_start} pages, {len(fig_txt)} figures, {len(tab_txt)} tables, unlocated {len(missing)}")
        info = dict(total=total, body=npdf - body_start, front=body_start, figs=len(fig_txt), tabs=len(tab_txt))
    print("saved", os.path.basename(path), os.path.getsize(path) // 1024, "KB")
    return info


if __name__ == "__main__":
    what = sys.argv[1] if len(sys.argv) > 1 else "all"
    R, T, W = all_blocks()
    if what in ("chapters", "all"):
        doc9(R, T, W)
        doc10()
        doc11()
        doc12(W)
        doc14(R, T)
        doc15()
    if what in ("final", "all"):
        final()
