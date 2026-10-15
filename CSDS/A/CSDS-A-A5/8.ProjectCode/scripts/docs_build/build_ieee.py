"""Build 19.IEEE_Paper.docx (6-page, two-column IEEE conference format) from REPORT_FACTS.md / experiment files.

Usage (from 8.ProjectCode, system python):  python scripts/docs_build/build_ieee.py [--no-pdf]
Reads experiments/*.json at build time; if seed_study.json is missing the cross-seed sentences say so.
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from common import ROOT, BATCH, REFS, STUDENTS, GUIDE, GUIDE_DESIG, PROJECT, SS, DIAG, CH

EXP = os.path.join(ROOT, "experiments")
FIGDIR = os.path.join(ROOT, "docs", "paper_figs")
os.makedirs(FIGDIR, exist_ok=True)
OUT = os.path.join(BATCH, "19.IEEE_Paper.docx")
FONT = "Times New Roman"


def jl(*p, default=None):
    f = os.path.join(EXP, *p)
    return json.load(open(f, encoding="utf-8")) if os.path.exists(f) else default


EV = jl("eval", "metrics.json")["detection"]
TRAIN = jl("metrics.json")
LAT = jl("latency.json")
EDGE = jl("edge_tests.json")
SD = jl("seed_study.json")
ERR = jl("error_analysis.json")
PAY = jl("eval", "metrics.json")["paysim"]
EXPL = jl("eval", "metrics.json")["explanations"]["summary"]
TP = EV["taxpayer"]
M, B, AB = TP["model"], TP["rule_baseline"], TP["ablation"]


def f3(x):
    from decimal import Decimal, ROUND_HALF_UP
    return str(Decimal(str(x)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


# ---------------------------------------------------------------------------------------------- references
# keys -> index into common.REFS (1-based numbers there)
KEYS = {"base": 1, "s2": 2, "s3": 3, "s4": 4, "s5": 5, "mehta": 6, "ted": 7, "ijepa": 8, "lecun": 9, "byol": 11, "vicreg": 12,
        "anom": 13, "johnson": 17, "nx": 16, "pearl": 22, "rag": 23, "halluc": 24, "gemini": 25, "paysim": 21, "torch": 20}
PRE = ["base", "s2", "s3", "s4", "s5"]  # base paper first, then the literature-survey sources


def number_citations(texts):
    order = list(PRE)
    for t in texts:
        for grp in re.findall(r"\[\[([a-z0-9,]+)\]\]", t):
            for k in grp.split(","):
                if k not in order:
                    order.append(k)
    return {k: i + 1 for i, k in enumerate(order)}, order


def fmt_cites(t, num):
    def rep(m):
        ks = m.group(1).split(",")
        return ", ".join(f"[{num[k]}]" for k in ks)
    return re.sub(r"\[\[([a-z0-9,]+)\]\]", rep, t)


# ---------------------------------------------------------------------------------------------- paper figures
plt.rcParams.update({"font.family": "Times New Roman", "font.size": 6.5, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.linewidth": 0.6, "xtick.major.width": 0.5, "ytick.major.width": 0.5})
GRN, DK, LT = "#15803d", "#14532d", "#dcfce7"
BLU, BLU_L = "#0369a1", "#e0f2fe"
AMB, AMB_L = "#b45309", "#fef3c7"
RED, RED_L = "#b91c1c", "#fee2e2"
PUR, PUR_L = "#6d28d9", "#ede9fe"
GRY, GRY_L = "#475569", "#f1f5f9"


def _box(ax, x, y, w, h, text, fc, ec, fs=6.3, bold=False):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=0.04", fc=fc, ec=ec, lw=0.8))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, fontweight="bold" if bold else "normal", linespacing=1.15)


def _arrow(ax, p, q, color=GRY, ls="-", style="-|>", lw=0.8):
    ax.add_patch(FancyArrowPatch(p, q, arrowstyle=style, mutation_scale=6, color=color, lw=lw, ls=ls, shrinkA=0, shrinkB=0))


def fig_arch():
    fig, ax = plt.subplots(figsize=(3.4, 2.45), dpi=300)
    ax.set_xlim(0, 3.4); ax.set_ylim(0, 2.45); ax.axis("off")
    _box(ax, 0.05, 2.0, 3.3, 0.4, "React SPA: ranking, profile, fraud graph,\ncase note, performance, run history", BLU_L, BLU, bold=True)
    _box(ax, 0.05, 1.38, 2.05, 0.45, "FastAPI service: auth, pipeline, investigate,\nexplanations  (bcrypt, JWT guard, CORS)", LT, GRN, bold=True)
    _box(ax, 2.3, 1.38, 1.05, 0.45, "Gemini API\n(evidence-only)", RED_L, RED, bold=True)
    _box(ax, 0.05, 0.62, 2.05, 0.55, "Analytics (ml/): 19 features, JEPA,\ncycles and clusters, invoice scores,\nrisk, evidence, chains, evaluation", PUR_L, PUR, bold=True)
    _box(ax, 2.3, 0.62, 1.05, 0.55, "Stores: SQLite,\nstate.pkl, CSV,\njepa.pt", AMB_L, AMB, bold=True)
    _box(ax, 0.05, 0.05, 2.05, 0.35, "Seeded GST generator (labelled fraud)", GRY_L, GRY, bold=True)
    _arrow(ax, (1.7, 2.0), (1.7, 1.83), BLU)
    _arrow(ax, (1.1, 1.38), (1.1, 1.17), GRN)
    _arrow(ax, (2.1, 1.6), (2.3, 1.6), RED, "--", "<|-|>")
    _arrow(ax, (2.1, 0.9), (2.3, 0.9), AMB, "-", "<|-|>")
    _arrow(ax, (1.95, 1.38), (2.75, 1.17), AMB)
    _arrow(ax, (1.1, 0.4), (1.1, 0.62), GRY)
    fig.savefig(os.path.join(FIGDIR, "fig1_architecture.png"), bbox_inches="tight", pad_inches=0.02, facecolor="white")
    plt.close(fig)


def fig_flow():
    fig, ax = plt.subplots(figsize=(3.4, 2.15), dpi=300)
    ax.set_xlim(0, 3.4); ax.set_ylim(0, 2.15); ax.axis("off")
    r1 = [("1 Data", "taxpayers,\ninvoices,\nreturns", GRY_L, GRY), ("2 Features", "19 per\ntaxpayer-\nmonth", BLU_L, BLU), ("3 JEPA", "embedding\nerror =\ndeviation", PUR_L, PUR), ("4 Network", "loops, shared-\nidentity\nclusters", LT, GRN)]
    r2 = [("5 Invoices", "noisy-OR of\n8 flags\n(>= 0.5)", AMB_L, AMB), ("6 Risk", "noisy-OR of\nJ, B, N, I;\nflag >= 0.5", RED_L, RED), ("7 Evidence", "E1..En,\nrings,\nchains", LT, GRN), ("8 Note", "Gemini,\ncited,\nchecked", BLU_L, BLU)]
    w, gap = 0.75, 0.11
    for i, (t, d, fc, ec) in enumerate(r1):
        x = 0.03 + i * (w + gap)
        _box(ax, x, 1.2, w, 0.82, "", fc, ec)
        ax.text(x + w / 2, 1.9, t, ha="center", va="center", fontsize=6.6, fontweight="bold", color=ec)
        ax.text(x + w / 2, 1.5, d, ha="center", va="center", fontsize=6.0, linespacing=1.15)
        if i < 3:
            _arrow(ax, (x + w, 1.6), (x + w + gap, 1.6))
    for i, (t, d, fc, ec) in enumerate(r2):
        x = 0.03 + (3 - i) * (w + gap)
        _box(ax, x, 0.3, w, 0.82, "", fc, ec)
        ax.text(x + w / 2, 1.0, t, ha="center", va="center", fontsize=6.6, fontweight="bold", color=ec)
        ax.text(x + w / 2, 0.6, d, ha="center", va="center", fontsize=6.0, linespacing=1.15)
        if i < 3:
            _arrow(ax, (x, 0.7), (x - gap, 0.7))
    _arrow(ax, (3.37 - w / 2, 1.2), (3.37 - w / 2, 1.12))
    ax.text(1.7, 0.1, "Evaluation against injected labels and a rule baseline runs on the same scores", ha="center", fontsize=5.8, style="italic", color=GRY)
    fig.savefig(os.path.join(FIGDIR, "fig2_flow.png"), bbox_inches="tight", pad_inches=0.02, facecolor="white")
    plt.close(fig)


def fig_results():
    fig, axs = plt.subplots(1, 2, figsize=(3.4, 1.95), dpi=300, gridspec_kw={"width_ratios": [1.1, 1]})
    ks = ["10", "25", "50", "79"]
    series = [("TaxSentinel", M, GRN), ("Graph+invoice", AB["graph_and_invoice_only"], BLU), ("JEPA+behav.", AB["jepa_only"], PUR), ("Rules R1-R5", B, GRY)]
    ax = axs[0]
    w = 0.2
    for i, (n, d, c) in enumerate(series):
        ax.bar([j + (i - 1.5) * w for j in range(4)], [d["precision_at_k"][k] for k in ks], w, color=c, label=n)
    ax.set_xticks(range(4)); ax.set_xticklabels([f"@{k}" for k in ks], fontsize=6)
    ax.set_ylim(0, 1.32); ax.set_yticks([0, .5, 1]); ax.set_ylabel("Precision@k", fontsize=6.5)
    ax.legend(fontsize=5.2, frameon=False, ncol=2, loc="upper center", handlelength=0.8, columnspacing=0.6, borderpad=0.1)
    ax.tick_params(labelsize=6)
    ax = axs[1]
    names = ["Full", "Graph+\ninv.", "JEPA+\nbeh.", "Rules"]
    f1 = [M["at_threshold"]["f1"], AB["graph_and_invoice_only"]["at_threshold"]["f1"], AB["jepa_only"]["at_threshold"]["f1"], B["at_threshold"]["f1"]]
    cols = [GRN, BLU, PUR, GRY]
    if SD:
        names.insert(1, "Mean\n5 seeds"); f1.insert(1, SD["mean"]["f1"]); cols.insert(1, "#86efac")
    ax.bar(range(len(f1)), f1, color=cols, edgecolor=[c if c == "#86efac" else "none" for c in cols])
    for i, v in enumerate(f1):
        ax.text(i, v + 0.02, f"{v:.2f}", ha="center", fontsize=5.8)
    ax.set_xticks(range(len(f1))); ax.set_xticklabels(names, fontsize=5.0, rotation=0)
    ax.set_ylim(0, 1.15); ax.set_yticks([0, .5, 1]); ax.set_ylabel("F1 at threshold 0.5", fontsize=6.5); ax.tick_params(labelsize=6)
    fig.tight_layout(pad=0.3, w_pad=0.6)
    fig.savefig(os.path.join(FIGDIR, "fig3_results.png"), bbox_inches="tight", pad_inches=0.02, facecolor="white")
    plt.close(fig)


def crop(src, box, dst):
    im = Image.open(os.path.join(SS, src)).convert("RGB").crop(box)
    im.save(os.path.join(FIGDIR, dst), dpi=(300, 300))


def make_figs():
    fig_arch(); fig_flow(); fig_results()
    crop("14_fraud_graph.png", (305, 250, 1110, 815), "fig4_graph.png")
    crop("19_case_explanation.png", (288, 322, 952, 900), "fig5_case.png")


# ---------------------------------------------------------------------------------------------- docx helpers
def set_font(run, size, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def add_runs(par, text, size, base_bold=False, base_italic=False):
    for part in re.split(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); set_font(r, size, True, base_italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, size, base_bold, True)
        else:
            r = par.add_run(part); set_font(r, size, base_bold, base_italic)


def pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=2, first=0.0, left=0.0, keep=False, line=1.0):
    f = p.paragraph_format
    p.alignment = align
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.first_line_indent = Inches(first) if first else None
    f.left_indent = Inches(left) if left else None
    f.line_spacing = line
    f.keep_with_next = keep
    f.widow_control = True


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for n in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{n}"); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)


def cell_margins(table, top=10, bottom=10, left=40, right=40):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for n, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{n}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def set_cols(section, num, space_twips=360):
    sp = section._sectPr
    for c in sp.findall(qn("w:cols")):
        sp.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num)); cols.set(qn("w:space"), str(space_twips))
    sp.append(cols)


def borders_cell(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge, val in kw.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), val); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e)
    tcPr.append(b)


# ---------------------------------------------------------------------------------------------- content
def cross_seed_sentence():
    if not SD:
        return "[cross-seed results pending]"
    m, s, lo, hi = SD["mean"], SD["std"], SD["min"], SD["max"]
    seeds = ", ".join(SD["seeds"].keys())
    pp = {}
    for k in ("circular_trading", "shell_entity", "fake_invoice_buyer", "itc_spike"):
        v = [x["per_pattern"][k] for x in SD["seeds"].values()]
        pp[k] = (min(v), max(v))
    rings = all(x["rings_recovered"] == x["rings_injected"] for x in SD["seeds"].values())
    return (f"Re-running the whole pipeline on five other generated ecosystems of 600 taxpayers (seeds {seeds}) with the same weights and threshold "
            f"gives a mean F1 of {f3(m['f1'])} (std {f3(s['f1'])}, range {f3(lo['f1'])}-{f3(hi['f1'])}), a mean ROC-AUC of {f3(m['roc_auc'])} "
            f"(range {f3(lo['roc_auc'])}-{f3(hi['roc_auc'])}) and a mean precision@k of {f3(m['p_at_k'])}, against a rule-baseline mean F1 of {f3(m['base_f1'])} "
            f"(range {f3(lo['base_f1'])}-{f3(hi['base_f1'])}). Ranking quality is therefore stable and the advantage over rules holds on every seed, but the F1 at the fixed "
            f"0.5 threshold is lower and less stable than the seed-42 figure, so the headline numbers are a best case. The loss is concentrated in two patterns: recall ranges "
            f"from {pp['fake_invoice_buyer'][0]:.2f} to {pp['fake_invoice_buyer'][1]:.2f} for fake-invoice buyers and from {pp['itc_spike'][0]:.2f} to {pp['itc_spike'][1]:.2f} for ITC spikes, "
            f"while circular trading stays at {pp['circular_trading'][0]:.2f}-{pp['circular_trading'][1]:.2f}, shell entities at {pp['shell_entity'][0]:.2f}-{pp['shell_entity'][1]:.2f}"
            + (" and every injected ring is recovered." if rings else "."))


def sweep_sentence():
    if not SD:
        return ""
    sw = {r["threshold"]: r for r in SD["threshold_sweep_seed42"]}
    return (f" Sweeping the threshold on the seed-42 scores gives F1 of {f3(sw[0.4]['f1'])}, {f3(sw[0.5]['f1'])} and {f3(sw[0.6]['f1'])} at 0.4, 0.5 and 0.6, but only "
            f"{f3(sw[0.3]['f1'])} at 0.3 ({sw[0.3]['flagged']} flagged) and {f3(sw[0.7]['f1'])} at 0.7 ({sw[0.7]['flagged']} flagged): the choice of threshold matters.")


def failure_text():
    t = ("The 3 false alarms are not loops: they are legitimate wholesalers and manufacturers flagged on behavioural deviation alone (JEPA scores 0.37-0.80, network score 0, no ring or cluster), "
         "so the JEPA's own false positives are where a single signal misleads. The 4 misses all lie below the threshold but within the top 21 % of the ranking (ranks 80 to 122 of 600): "
         "two ITC-spike cases whose single-month spikes are diluted in the deviation score (risk 0.41 and 0.47) and two fake-invoice buyers with weak network (0.31) and invoice (0.11-0.12) evidence (risk 0.29 and 0.32). "
         "ITC-spike recall (0.833) is below the rule baseline's (1.00). Separately, two of the nine detected rings, the two with the lowest scores (0.30 and 0.27), contain no injected-fraud taxpayer: "
         "they are legitimate reciprocal trade loops that the generator includes on purpose, so a loop is a lead, not proof. ")
    if ERR:
        fp = ERR["false_positives"]
        fn = ERR["false_negatives"]
        assert len(fp) == 3 and len(fn) == 4, "error analysis changed; update failure_text"
    return t + ("Other limits are the synthetic data, weights and threshold tuned on one seed, the absence of a held-out split for the encoder (the cross-seed study is the nearest substitute), "
                "and the dependence on an external language-model quota.")


def edge_sentence():
    if not EDGE:
        return ""
    return f"{EDGE['passed']} of {EDGE['total']} black-box edge and security checks passed"


def lat_range():
    meds = [v[0] for v in LAT.values()]
    return f"{min(meds):.1f}-{max(meds):.1f} ms"


ABSTRACT = ("Input tax credit (ITC) fraud under India's Goods and Services Tax, carried out through fake invoices, circular trading and shell companies, "
            "causes large revenue losses, yet confirmed fraud labels are scarce and most systems cannot say why a claim is flagged. This paper presents TaxSentinel, "
            "a self-supervised and explainable framework. A Joint Embedding Predictive Architecture (JEPA) learns normal taxpayer behaviour without labels; its prediction "
            "error is combined, through a noisy-OR, with circular-trading and shared-identity cluster detection and with invoice-level anomaly scores. Flagged taxpayers receive "
            "a numbered evidence list, and a large language model writes a case note in which every claim cites that evidence. On a seeded synthetic ecosystem of 600 taxpayers "
            "and 64,474 invoices, the system reaches precision@79 of 0.949, F1 0.955 and ROC-AUC 0.997, against 0.620, 0.432 and 0.876 for a rule baseline, and recovers all five "
            "injected rings. Performance at the fixed threshold is lower on other generated ecosystems, the language model is prompted rather than fine-tuned, and no real tax data was used.")
INDEX = "circular trading, explainable AI, GST fraud, JEPA, large language model, self-supervised learning, shell entities"

# body: list of ("h", text) ("sh", text) ("p", text) ("eq", text, num) ("fig", file, caption, width) ("tbl", caption, header, rows, widths) ("pnoind", text)
BODY = []


def H(t): BODY.append(("h", t))
def SH(t): BODY.append(("sh", t))
def P(t): BODY.append(("p", t))
def EQ(t, n): BODY.append(("eq", t, n))
def FIG(f, c, w=3.38): BODY.append(("fig", f, c, w))
def TBL(c, h, r, w, fs=6.6): BODY.append(("tbl", c, h, r, w, fs))
def LST(items): BODY.append(("list", items))


H("I. INTRODUCTION")
P("The Goods and Services Tax (GST), introduced in India in 2017, lets a business set the tax it paid on purchases, its input tax credit (ITC), against the tax it collects on sales. "
  "The mechanism removes cascading taxes but turns an invoice into something that carries money, and because returns are self-assessed, an invoice for a supply that never happened can still hand a buyer a credit [[base]]. "
  "Three patterns dominate: fake invoicing by networks of paper firms, circular trading in which a few firms pass the same value around a loop, and shell entities that share an address, trade among themselves, pay almost no tax in cash and stop filing. "
  "Official figures cited in the base paper show detected evasion rising from Rs 1.01 lakh crore in FY 2022-23 to Rs 2.01 lakh crore in FY 2023-24, with cases booked rising from 4,872 to 6,084 [[base]].")
P("Current practice combines audits after suspicion, fixed mismatch rules and, increasingly, supervised machine learning on labelled taxpayer records [[s2]], [[s3]], [[s4]], [[s5]]. Each has a cost. Audits are slow and find fraud after the revenue is lost; fixed rules are transparent and therefore easy to stay under; supervised models need confirmed cases, which are scarce, late and biased towards what investigators happened to pursue; and checks on one return or invoice cannot see a loop of five firms that each look ordinary. Few systems explain a flag in terms an officer can verify.")
P("This paper presents **TaxSentinel**, a self-supervised, explainable framework that learns normal GST behaviour from unlabelled data, combines invoice, return-behaviour and network evidence, and explains each flag with cited evidence. Its contributions are:")
LST(["a three-layer, 19-feature representation of taxpayer-months and a small JEPA behaviour encoder trained without labels in 11.7 s on a CPU;",
     "scored circular-trading and shared-identity detectors and a noisy-OR risk that combines them with the JEPA deviation and invoice-level anomaly scores;",
     "evidence-gated case notes written by a large language model, with automatic citation and grounding checks; and",
     "an evaluation against a rule baseline with ablations, a cross-seed study and a transfer check, and a secured full-stack application."])
P("Section II reviews the base paper and related work, Section III the methodology, Section IV the results and Section V concludes.")
SH("A. System Hardware")
P("All development and measurement used one desktop: an Intel Core i5-12500 (6 cores, 12 threads), 15.7 GB of RAM and integrated Intel UHD 770 graphics under Windows 11 Pro. No GPU was used; the machine was at times heavily loaded by other jobs, so timings are indicative.")
SH("B. Software Requirements")
P("Python 3.11.9 with PyTorch 2.14.0 (CPU build) [[torch]], NumPy 2.4.6 and NetworkX 3.4.2 [[nx]]; FastAPI 0.115.6, SQLAlchemy 2.0.36 with SQLite, PyJWT and bcrypt; React 18.3 with Vite 6 and Node.js v24.19.0; and the Google Gemini API through the google-genai client [[gemini]].")

H("II. LITERATURE SURVEY")
P("**Base paper.** Datta, Thakkar and Dabhade [[base]] give a qualitative, evidence-synthesis study of GST ITC fraud. It explains why the self-assessed design invites fraud, drawing on asymmetric information, moral hazard and the Fraud Triangle, and describes how fake invoicing, circular trading and shell-company networks are carried out. Using DGGI and CBIC statistics it reports detected evasion of Rs 1.01 and Rs 2.01 lakh crore in FY 2022-23 and FY 2023-24, cases booked rising from 4,872 to 6,084, arrests from 97 to 140, and special drives against fake registrations (21,791 fake GSTINs in May-July 2023). It recommends data analytics and network analysis, dynamic AI-based risk modelling at registration, officer capacity building, faster recovery and an inter-agency intelligence grid. By its own account it relies on public sources and contains no empirical model: it names the patterns and calls for the tool, but builds and evaluates none. TaxSentinel adopts its three patterns as detection targets and implements the call for analytics, network analysis and AI-based risk modelling.")
P("**Machine learning for tax fraud.** Four further papers in the batch survey apply AI and machine learning to tax compliance and income-tax fraud: an overview of AI and ML in Indian tax compliance [[s2]], two conference papers on income-tax fraud detection [[s3]], [[s4]] and a boosting-based study [[s5]]. Their full texts were not available to the authors, so, judging only by their titles and venues, they use supervised learning on labelled taxpayer records; no results are quoted from them.")
P("**Graph methods.** Mehta et al. [[mehta]] study circular trading on real Telangana GST data (77.8 million invoices, 388,448 taxpayers): they learn node2vec embeddings of the sales-flow graph and cluster them to find communities of circular traders, reporting communities found rather than accuracy. Xu et al. [[ted]] model taxation as a heterogeneous graph and propose TED, a supervised graph neural network with related-party groups, reporting F1 between 76.35 % and 94.96 % on two real tax datasets. Both target specific patterns or need labels, and neither explains its output.")
P("**Self-supervised prediction.** LeCun's joint-embedding predictive architecture [[lecun]] argues for predicting in representation space; I-JEPA [[ijepa]] predicts the embeddings of masked image blocks from a context block with an exponential-moving-average (EMA) target encoder, an idea from BYOL [[byol]], and a variance term as in VICReg [[vicreg]] prevents collapse. We transfer this recipe to monthly tax behaviour vectors. Anomaly detection in general is surveyed in [[anom]].")
P("**Structure and language.** Short directed cycles can be enumerated exactly with Johnson-style algorithms [[johnson]], available with a length bound in NetworkX [[nx]]; a noisy-OR [[pearl]] combines independent causes into one probability. Language models write fluent but sometimes unsupported text [[halluc]]; retrieval-augmented generation [[rag]] reduces this by supplying source passages. TaxSentinel supplies only computed evidence and checks the answer.")
TBL("SURVEY OF THE MOST RELEVANT WORKS", ["Work", "Approach", "Data", "Reported result", "Limitation"], [
    ["[[base]]", "Forensic analytics review", "DGGI, CBIC statistics", "Evasion Rs 1.01 to 2.01 lakh cr", "No model"],
    ["[[s2]]-[[s5]]", "ML for tax fraud (per titles)", "Not available", "Not quoted", "Supervised; full text unseen"],
    ["[[mehta]]", "node2vec + clustering", "Telangana GST, 77.8 M invoices", "Circular-trader communities", "One pattern; no explanation"],
    ["[[ted]]", "Heterogeneous GNN", "Two real tax sets", "F1 76.35-94.96 %", "Supervised; no explanation"],
    ["[[ijepa]]", "I-JEPA, masked-block prediction", "Images", "Strong transfer", "Not applied to tax"],
    ["This work", "JEPA + graph + invoice + LLM", "Synthetic GST", "F1 0.955 (seed 42)", "Synthetic; prompted LLM"],
], [0.42, 0.78, 0.7, 0.8, 0.7])
P("**Research gap.** Existing GST fraud work is largely descriptive or rule-based or needs labels. Few systems learn behaviour from unlabelled tax data, combine invoice, return and network evidence in one score, recover loops and shell clusters, and explain every flag with checkable evidence. This work addresses that combination; it makes no claim to outperform supervised graph methods [[ted]] on real data, against which it was not tested.")

H("III. METHODOLOGY")
SH("A. Data")
P("Real GST returns are not public, so a seeded generator produces one financial year (April 2025 to March 2026) of GST activity: 600 taxpayers in eight states and five sectors with valid-format GSTINs, 64,474 business-to-business invoices on a tiered supply network and 6,977 monthly returns (outward supplies, GSTR-2B credit available, GSTR-3B credit claimed, cash paid, filing delay). Four patterns are injected with labels: 27 shell entities in seven clusters, 19 fake-invoice buyers, 21 circular-trading firms in five rings of three to five, and 12 ITC spikes, 79 fraud taxpayers (13.2 %) and 530 fraud invoices in all. Legitimate hard negatives are included on purpose: festival seasonality, one-off capital purchases, new registrations, group companies sharing an address, late filers, credit catch-ups and reciprocal service-manufacturer trade. Labels are used only to evaluate. The generator is seeded, so seed 42 reproduces every figure reported here, and other ecosystems of 150 to 1,500 taxpayers can be generated from the application. In the bundled ecosystem 39,582 invoices carry an e-way bill and 1,071 were not reported by the seller in GSTR-1. A 100,000-row sample of PaySim [[paysim]] (134 fraud rows) serves as an out-of-domain check.")
SH("B. Features")
P("Every active taxpayer-month becomes a 19-dimensional vector in three layers: invoice (counts, share of purchases not reported by suppliers, goods invoices without an e-way bill, round values, young suppliers), behaviour (turnover, ITC claimed, ITC-to-output-tax ratio, claim above GSTR-2B, cash share, filing delay, registration age) and network (suppliers, buyers, new partners, reciprocity, largest-buyer share, value added). Heavy-tailed values are log-transformed and features are robust-scaled (median and IQR, clipped at ten).")
SH("C. JEPA Behaviour Encoder")
P("Let f be an online encoder (19-64-64-32 perceptron with layer normalisation), f' an EMA copy of it without gradient, and g a predictor. For month t of taxpayer i with scaled vector z(t), previous months z(t-1..t-3) and a mask variant v, the model predicts the embedding of the whole month from the three previous embeddings (with presence flags) and from z(t) with one block hidden. Variant 0 hides everything; variants 1-3 hide the invoice, behaviour or network block, so a deviation can be attributed to a layer:")
EQ("ŝ = g( f(zₜ₋₁), f(zₜ₋₂), f(zₜ₋₃), m, f(kᵥ ⊙ zₜ), onehot(v) ),   s = f′(zₜ)", "1")
EQ("L = mean (ŝ − s)² + λ · meanⱼ ReLU(1 − stdⱼ f(z))", "2")
EQ("ξ ← 0.99 ξ + 0.01 θ   (after every optimiser step)", "3")
P("Training uses AdamW [lr 1e-3, batch 256, 100 epochs, λ = 1] on all 6,977 taxpayer-months without labels; the saved model holds 53,420 values, trains in 11.7 s on the CPU and ends at a loss of 0.1115. At scoring time each variant's error is divided by the median error of that month, the four variants are combined by a geometric mean, the taxpayer's score is the mean of its two worst months, and a robust z-score and a logistic squash give a deviation score J in (0, 1). The behaviour-masked variant alone gives a second score B that responds to one-month ITC spikes. The training loss is not monotone: it rises from 0.012 to a peak of 0.261 at epoch 11 while the online embeddings spread out under the variance term and the EMA target follows, then falls and settles between 0.11 and 0.12 from about epoch 60.")
P("**Why JEPA.** A predictive model asks whether a month is what the history and the other layers would lead one to expect. Predicting embeddings, not raw values, means the model need not reproduce the heavy-tailed noise of amounts and counts; the EMA target and the variance term prevent collapse without negative pairs; and the four masks correspond to meaningful hypotheses (history, invoice layer, behaviour layer, network layer). We did not train an isolation forest or an autoencoder on the same features, so the choice rests on design arguments, not on a measured comparison; the ablation in Section IV shows that JEPA alone is not sufficient, which is why explicit network and invoice evidence is added.")
SH("D. Network and Invoice Analysis")
P("A directed seller-to-buyer graph is built from the invoices. Bounded cycle search (length ≤ 5) on recurring (at least two months), above-median edges finds loops, and each loop is scored as")
EQ("score(C) = √balance · recurrence · √concentration", "4")
P("where balance is the smallest over the largest edge value, recurrence is the share of months (capped at three) in which the whole loop trades, and concentration is the loop's share of every member's purchases; overlapping loops with a score of at least 0.25 are merged into rings. Registrations sharing an address or contact number form clusters, scored from mean deviation, youth, intra-cluster trade and stopped filing. Each invoice receives a noisy-OR of eight weighted flags (unreported by the supplier, no e-way bill, round value, young supplier, supplier in a suspicious cluster, supplier later silent, loop edge, value spike) and counts as anomalous at 0.5 or more; the taxpayer's invoice score I is the peak monthly anomalous share, and the network score N combines ring, cluster and exposure evidence. A flagged taxpayer is given a named pattern from its evidence: circular trading if it belongs to a ring; shell entity if it is in a suspicious cluster and is young, has stopped filing or pays almost no cash tax; fake invoices if it buys from suspicious clusters (10 % of purchases in some month) or has at least 15 % of its purchase value flagged; ITC spike if a month's claim is at least 2.5 times its median and 1.5 times the credit available; otherwise unusual behaviour. On the bundled ecosystem 79 candidate cycles were scanned, giving 9 rings, 12 identity clusters (7 suspicious) and 16 invoice chains.")
SH("E. Risk, Evidence and Explanation")
P("The four layer scores combine in a weighted noisy-OR, so that one strong layer suffices and agreeing layers reinforce each other; a taxpayer is flagged at 0.5:")
EQ("R = 1 − (1 − 0.75 J)(1 − 0.6 B)(1 − 0.7 N)(1 − 0.6 I)", "5")
P("The weights were set on the seed-42 ecosystem. For each taxpayer the pipeline builds a numbered evidence list E1..En (layer, severity, concrete figures), names the suspected pattern and ranks taxpayers and invoice chains (each loop and each suspicious cluster with its buyers). On request a Gemini model [[gemini]] receives only the profile, scores and evidence list under a fixed system prompt: use only the evidence, cite [E#] after every claim, give no legal conclusion and write three sections (summary, why flagged, next steps) in under 320 words, at temperature 0.2. The answer is checked automatically for invalid citations, evidence and high-severity coverage and the share of its numbers found in the evidence, and is stored with the evidence it was written from. The model is therefore prompted and verified, not fine-tuned.")
FIG("fig1_architecture.png", "Fig. 1. Architecture of TaxSentinel: React front end, FastAPI service, analytics package, stores and the Gemini API.")
SH("F. Implementation")
P("The backend exposes 11 endpoint families behind a JWT bearer guard (HS256, 12 h, bcrypt password hashes, pydantic validation). Generation runs on a background thread with a polled status, and the finished analytical state is held in memory and pickled, so every read endpoint is a lookup. Users, runs and stored notes are kept in SQLite. The React interface offers a ranked list, taxpayer profiles with evidence, an interactive fraud graph, invoice-chain drill-down, a case page with citation links, a performance page and a run history, in light and dark themes with keyboard shortcuts and a command palette. While a generation runs, a stage stepper and an activity log show each pipeline stage (generate, features, train, score, rings, invoices, risk, evidence) with the training epoch and loss.")
FIG("fig2_flow.png", "Fig. 2. Main flow from GST data to a cited case note.")

H("IV. RESULTS")
SH("A. Test Conditions")
P("All results are on the bundled seed-42 ecosystem unless stated: 600 taxpayers (79 fraud), 64,474 invoices (530 fraud), scored without labels and then compared with the injected labels. The comparison baseline is a rule set built from standard scrutiny checks: ITC claimed over GSTR-2B by 20 % and Rs 50,000 (R1), young registration with large turnover (R2), under 1 % cash tax (R3), ITC more than tripling (R4) and stopped filing (R5). Precision@k is the share of fraud among the k highest-risk taxpayers (k = 10, 25, 50 and 79, the number of fraud taxpayers); precision, recall and F1 are taken at the flag threshold of 0.5; ROC-AUC is rank-based. An injected ring counts as recovered when one detected ring holds at least 80 % of its members. Hardware is described in Section I-A. **The data is synthetic, and the weights and threshold were chosen on this ecosystem, so these are best-case figures.**")
SH("B. Detection Performance")
P("TaxSentinel flags 78 taxpayers, 75 correctly (3 false alarms, 4 missed): P@79 is 0.949, F1 0.955 and ROC-AUC 0.997 (TABLE II). The rule baseline flags 236, only 68 of them fraud (168 false alarms), for an F1 of 0.432. The first 50 ranked taxpayers contain no false alarm and the 10 top-ranked invoice chains are all majority-fraud. The top-ranked taxpayer (risk 0.965) belongs to ring R01, a five-party loop of 98 invoices worth Rs 44.65 Cr that passes Rs 9.78 Cr of ITC (Figs. 4 and 5).")
rows = [["TaxSentinel (seed 42)", f3(M["precision_at_k"]["79"]), f3(M["at_threshold"]["precision"]), f3(M["at_threshold"]["recall"]), f3(M["at_threshold"]["f1"]), f3(M["roc_auc"])],
        ["Rule baseline R1-R5", f3(B["precision_at_k"]["79"]), f3(B["at_threshold"]["precision"]), f3(B["at_threshold"]["recall"]), f3(B["at_threshold"]["f1"]), f3(B["roc_auc"])],
        ["JEPA + behaviour only", f3(AB["jepa_only"]["precision_at_k"]["79"]), f3(AB["jepa_only"]["at_threshold"]["precision"]), f3(AB["jepa_only"]["at_threshold"]["recall"]), f3(AB["jepa_only"]["at_threshold"]["f1"]), f3(AB["jepa_only"]["roc_auc"])],
        ["Graph + invoice only", f3(AB["graph_and_invoice_only"]["precision_at_k"]["79"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["precision"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["recall"]), f3(AB["graph_and_invoice_only"]["at_threshold"]["f1"]), f3(AB["graph_and_invoice_only"]["roc_auc"])]]
if SD:
    m = SD["mean"]
    rows += [["TaxSentinel, 5 other seeds (mean)", f3(m["p_at_k"]), f3(m["precision"]), f3(m["recall"]), f3(m["f1"]), f3(m["roc_auc"])],
             ["Rule baseline, 5 other seeds (mean)", f3(m["base_p_at_k"]), "-", "-", f3(m["base_f1"]), f3(m["base_roc_auc"])]]
inv, invb = EV["invoice"]["model"], EV["invoice"]["rule_baseline"]["at_threshold"]
rows += [["Invoice score >= 0.5", "-", f3(inv["at_threshold"]["precision"]), f3(inv["at_threshold"]["recall"]), f3(inv["at_threshold"]["f1"]), f"{inv['roc_auc']:.4f}"],
         ["Invoice not in supplier GSTR-1", "-", f3(invb["precision"]), f3(invb["recall"]), f3(invb["f1"]), "-"]]
TBL("DETECTION RESULTS (TAXPAYER LEVEL, THEN INVOICE LEVEL)", ["Method", "P@79", "Prec.", "Recall", "F1", "AUC"], rows, [1.42, 0.4, 0.4, 0.4, 0.4, 0.4], fs=6.8)
SH("C. Ablations and Patterns")
P("The ablations are the most informative result (Fig. 3). The JEPA with the behaviour variant alone reaches an F1 of only 0.327: precision 0.842 but recall 0.203, because the 27 look-alike shell firms form their own \"normal\". Graph plus invoice evidence reaches 0.887 with precision 1.000, and the full model 0.955; removing the invoice layer leaves P@79 unchanged at 0.949. The JEPA is thus a useful but partial signal that mainly adds recall, and most detection power comes from the network and invoice layers. Recall at the threshold is 1.00 for circular trading (rules 0.86), 1.00 for shell entities (rules 1.00), 0.89 for fake-invoice buyers (rules 0.58) and 0.83 for ITC spikes, where **the rule baseline is better** (1.00): a claim that triples in one month is exactly what rule R4 tests, while the model's score dilutes a single month. All five injected rings were recovered.")
FIG("fig3_results.png", "Fig. 3. Left: precision@k of the full model, two ablations and the rule baseline. Right: F1 at the 0.5 threshold (seed 42, and the mean over five other seeds).")
PAT = {"circular_trading": "Circular trading", "shell_entity": "Shell entity", "fake_invoice_buyer": "Fake-invoice buyer", "itc_spike": "ITC spike"}
TBL("RECALL BY INJECTED PATTERN AT THE FLAG THRESHOLD (SEED 42)", ["Pattern", "Count", "TaxSentinel", "Rules", "Named"],
    [[PAT[k], str(v["count"]), f"{v['model_recall']:.2f}", f"{v['baseline_recall']:.2f}", f"{v['pattern_named']:.2f}"] for k, v in EV["per_pattern"].items()], [1.2, 0.5, 0.7, 0.5, 0.5], fs=6.8)
SH("D. Generalisation, Invoices and Transfer")
P(cross_seed_sentence() + " At invoice level, the invoice score reaches F1 0.924 (precision 0.883, recall 0.968, 581 flagged) against 0.160 for the classic check that a purchase is missing from the supplier's GSTR-1 (1,071 flagged)." + sweep_sentence() + " As a transfer check the same JEPA module, trained without labels on 100,000 PaySim transactions for 6.2 s, reaches a ROC-AUC of 0.935 against 0.790 for ranking by amount (precision@100 0.31 against 0.02).")
SH("E. Explanations, Testing and Speed")
def stored_note_range():
    smp = jl("samples.json")
    if not smp:
        return ""
    g = smp["taxpayer"]["gstin"]
    qs = [q["quality"] for q in smp["stored_note_quality"] if q["gstin"] == g]
    if len(qs) < 2:
        return ""
    lo, hi = min(q["numeric_grounding"] for q in qs), max(q["numeric_grounding"] for q in qs)
    return f" The same case written {len(qs)} times during development gave numeric grounding between {lo:.2f} and {hi:.2f}, because the model is not deterministic even at temperature 0.2."


def edge_detail():
    if not EDGE:
        return ""
    return (f"Black-box checks sent invalid, forged and hostile requests (wrong password, unknown user, duplicate registration, short password, malformed e-mail or JSON, a token with an altered signature, every protected endpoint without a token, SQL-injection-style and script-tag search text, out-of-range paging and generation parameters, unknown identifiers and a second generation while one was running); the expected 401, 404, 409 or 422 status was returned in {EDGE['passed']} of {EDGE['total']} cases." + ((" The " + str(EDGE['total'] - EDGE['passed']) + " exceptions are the final checks that the bundled ecosystem was restored after the concurrency test; that run timed out while the machine was heavily loaded and was re-run separately.") if EDGE['passed'] < EDGE['total'] else ""))


P(f"Of 10 planned case notes, 6 completed before the free Gemini quota and availability limits (HTTP 429, 503 and 504) stopped the rest. Over these 6 notes there were {EXPL['invalid_citation_rate']*100:.0f} % invalid citations, {EXPL['mean_citations']:.1f} citations per note, {EXPL['evidence_coverage']*100:.0f} % evidence coverage, {EXPL['high_severity_coverage']*100:.0f} % high-severity coverage and {EXPL['numeric_grounding']*100:.1f} % numeric grounding; the ungrounded numbers are mostly derived figures (a total of two evidence amounts, a rounded share) that the simple check cannot match." + stored_note_range() + " These checks measure grounding, not the correctness of the advice, and no human judged the notes.")
P("An end-to-end smoke test (registration, sign-in, two regenerations, every endpoint family and one real Gemini call) passed 20 of 20 checks in 39.4 s. " + edge_detail() + f" Median read-endpoint latency is {lat_range()} (20 sequential calls each, single client) and a full generate-train-analyse run takes 18.6 to 29.0 s for 600 taxpayers (longer when other jobs load the machine); the API process held about 0.5 GB with the 600-taxpayer state loaded. There is no unit-test suite, load test or user study.")
FIG("fig4_graph.png", "Fig. 4. Fraud network of the seed-42 ecosystem: red edges are detected circular-trading loops, ringed nodes are shared-identity members.")
FIG("fig5_case.png", "Fig. 5. Case note for the top-ranked taxpayer; each claim ends in an [E#] chip linked to its evidence.")
SH("F. Comparison with the Base Paper")
P("The base paper [[base]] reports no model, accuracy or other measured performance, so no like-for-like numerical comparison is possible. Qualitatively, TaxSentinel implements what it recommends: the three fraud patterns as detection targets, analytics and network analysis, and AI-based risk modelling (applied to existing taxpayers, not at registration). Against the rule checks that the paper's setting implies, the gain on the same data is F1 0.955 against 0.432 on seed 42" + (f" and {f3(SD['mean']['f1'])} against {f3(SD['mean']['base_f1'])} on average over five other seeds" if SD else "") + ". For context only, TED [[ted]] reports F1 of 76.35-94.96 % on two real datasets with supervised training; the data, task and supervision differ, so the figures are not comparable.")
SH("G. Failure Cases and Observed Limits")
P(failure_text())

H("V. CONCLUSION")
P("TaxSentinel shows that GST ITC-fraud detection can be built without fraud labels and still explain itself. A 53,420-value JEPA learns normal taxpayer behaviour in 11.7 s on a CPU; combined with loop and cluster detection and invoice scoring, it ranks the synthetic ecosystem with precision@79 of 0.949 and ROC-AUC 0.997, recovers all five injected rings, scores invoices with F1 0.924, and supports case notes with no invalid citations in the six evaluated. The key finding is that several layers beat one: the JEPA alone reaches only F1 0.327, and its value lies in adding recall to network and invoice evidence. The honest limits are the synthetic data, a fixed threshold whose F1 is lower on other generated ecosystems, a prompted rather than fine-tuned language model evaluated on six notes, and no unit, load or user testing. The output is a risk indication for review, never a finding of fraud.")
P("Future work is, in order of importance: calibrating weights on a small set of confirmed real cases, validation on anonymised real returns with a tax administration, head-to-head baselines against an isolation forest and an autoencoder, role-based access with case management, scaling the graph analysis to state-sized data, and a human evaluation of the explanations.")


# ---------------------------------------------------------------------------------------------- build
def build():
    make_figs()
    texts = [ABSTRACT] + [x[1] for x in BODY if x[0] in ("h", "sh", "p")] + [i for x in BODY if x[0] == "list" for i in x[1]] + \
            [c for x in BODY if x[0] == "tbl" for c in (x[1], *x[2], *[cell for row in x[3] for cell in row])] + [x[2] for x in BODY if x[0] == "fig"]
    num, order = number_citations(texts)
    # citation order must follow the rule: numbers non-decreasing at first appearance
    seen, last = set(), 0
    for t in texts:
        for grp in re.findall(r"\[\[([a-z0-9,]+)\]\]", t):
            for k in grp.split(","):
                if k not in seen:
                    seen.add(k)
                    assert num[k] >= last, f"citation order violated at {k}"
                    last = num[k]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Inches(0.75), Inches(1.0), Inches(0.69), Inches(0.69)
    sec.header_distance = sec.footer_distance = Inches(0.4)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    rf = st.element.rPr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    # ---- title block (single column)
    p = doc.add_paragraph()
    pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
    add_runs(p, PROJECT, 24)
    # authors
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t); cell_margins(t, 10, 10, 30, 30)
    people = [(n, r, "Student") for n, r in STUDENTS] + [(GUIDE, GUIDE_DESIG, "Guide")]
    for i, (n, extra, kind) in enumerate(people):
        c = t.rows[0].cells[i]
        c.width = Inches(7.12 / 5)
        c.text = ""
        first = True
        lines = [(n, True, False), (extra if kind == "Guide" else f"Roll No. {extra}", False, True),
                 ("Dept. of CSE (Data Science)", False, False), ("TKR College of Engineering and Technology (Autonomous)", False, False),
                 ("[City, State, Country]", False, False), ("[email address]", False, False)]
        for txt, b, it in lines:
            par = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            pfmt(par, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
            r = par.add_run(txt); set_font(r, 9, b, it)
    sp = doc.add_paragraph(); pfmt(sp, WD_ALIGN_PARAGRAPH.CENTER, 0, 2)
    r = sp.add_run(""); set_font(r, 4)
    # ---- continuous break -> two columns
    s2 = doc.add_section(WD_SECTION.CONTINUOUS)
    s2.top_margin, s2.bottom_margin, s2.left_margin, s2.right_margin = Inches(0.75), Inches(1.0), Inches(0.69), Inches(0.69)
    s2.page_width, s2.page_height = Inches(8.5), Inches(11)
    set_cols(s2, 2, 360)
    # abstract and index terms
    p = doc.add_paragraph(); pfmt(p, first=0.2, after=3)
    r = p.add_run("Abstract—"); set_font(r, 9, True, True)
    add_runs(p, ABSTRACT, 9, base_bold=True)
    p = doc.add_paragraph(); pfmt(p, first=0.2, after=6)
    r = p.add_run("Index Terms—"); set_font(r, 9, True, True)
    add_runs(p, INDEX, 9, base_bold=True)

    for item in BODY:
        k = item[0]
        if k == "h":
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 8, 3, keep=True)
            add_runs(p, item[1], 10, base_bold=True)
        elif k == "sh":
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.LEFT, 4, 2, keep=True)
            add_runs(p, item[1], 10, base_italic=True)
        elif k == "p":
            p = doc.add_paragraph(); pfmt(p, first=0.2, after=3)
            add_runs(p, fmt_cites(item[1], num), 10)
        elif k == "list":
            for i, it in enumerate(item[1], 1):
                p = doc.add_paragraph(); pfmt(p, left=0.3, after=1)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                add_runs(p, f"{i}) " + fmt_cites(it, num), 10)
        elif k == "eq":
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1, 3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(1.69), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.37), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run("\t" + item[1]); set_font(r, 8.5, None, True)
            r = p.add_run(f"\t({item[2]})"); set_font(r, 10)
        elif k == "fig":
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 3, 1, keep=True)
            p.add_run().add_picture(os.path.join(FIGDIR, item[1]), width=Inches(item[3]))
            c = doc.add_paragraph(); pfmt(c, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6)
            add_runs(c, item[2], 8)
        elif k == "tbl":
            _, cap, hdr, rws, widths, fs = item
            roman = "I" if not any(isinstance(b, str) and b == "x" for b in []) else ""
            n_tbl = sum(1 for x in BODY[:BODY.index(item) + 1] if x[0] == "tbl")
            lab = ["I", "II", "III", "IV"][n_tbl - 1]
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 0, keep=True)
            add_runs(p, f"TABLE {lab}", 8)
            p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 2, keep=True)
            add_runs(p, cap, 8)
            tb = doc.add_table(rows=1, cols=len(hdr))
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            tb.autofit = False
            tblPr = tb._tbl.tblPr
            cell_margins(tb, 5, 5, 25, 25)
            for i, h in enumerate(hdr):
                c = tb.rows[0].cells[i]; c.text = ""
                pp = c.paragraphs[0]; pfmt(pp, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
                add_runs(pp, h, fs, base_bold=True)
                borders_cell(c, top="single", bottom="single")
            for row in rws:
                cells = tb.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = ""
                    pp = cells[i].paragraphs[0]
                    pfmt(pp, WD_ALIGN_PARAGRAPH.LEFT if (i == 0 or not re.fullmatch(r"[-\d.]+", v)) else WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
                    add_runs(pp, fmt_cites(v, num), fs)
                trpr = tb.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
            for c in tb.rows[-1].cells:
                borders_cell(c, bottom="single")
            for row in tb.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
            sp = doc.add_paragraph(); pfmt(sp, after=3)
            r = sp.add_run(""); set_font(r, 3)
    # references
    p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.CENTER, 8, 3, keep=True)
    add_runs(p, "REFERENCES", 10, base_bold=True)
    for k in order:
        i = num[k]
        p = doc.add_paragraph(); pfmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1, left=0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        add_runs(p, f"[{i}] " + REFS[KEYS[k] - 1], 8)
    doc.save(OUT)
    print("saved", OUT, os.path.getsize(OUT) // 1024, "KB; references:", len(order))
    words = sum(len(re.findall(r"\w+", fmt_cites(t, num))) for t in texts) + len(re.findall(r"\w+", " ".join(REFS[KEYS[k] - 1] for k in order)))
    print("approx words:", words)
    return num


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        import wordtools
        pdf = os.path.join(BATCH, "19.IEEE_Paper.pdf")
        n = wordtools.export_pdf(OUT, pdf)
        print("pages:", n)
