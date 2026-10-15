"""Renderer: turns a list of content blocks into a formatted .docx (python-docx).

Block grammar (tuples):
  ("h1", text)  ("h2", text)  ("h3", text)          headings (TNR 14 bold / 14 bold / 12 bold-italic)
  ("p", text)                                       body paragraph (justified, 1.5 spacing); supports **bold** and {fig:key}/{tab:key}
  ("bul", [text, ...])   ("num", [text, ...])       lists
  ("fig", key, path, caption, width_in)             picture + numbered caption below
  ("tbl", key, caption, header, rows, col_widths)   table + numbered caption above
  ("code", text)                                    monospace block
  ("eq", text, number)                              centred equation line
  ("note", text)                                    small italic note
  ("pb",)                                           page break
  ("raw", callable(doc))                            escape hatch for cover pages etc.
"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
GREEN = "DCFCE7"


def _set_font(style_or_run, size=None, bold=None, italic=None, name=FONT, color=None):
    f = style_or_run.font
    f.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") and hasattr(style_or_run.element, "get_or_add_rPr") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), name)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)


def new_document(margins=(1.0, 1.0, 1.0, 1.0)):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = (Inches(m) for m in margins)
    st = doc.styles
    n = st["Normal"]
    _set_font(n, 12)
    pf = n.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, bold, italic, align, before, after in [("Heading 1", 14, True, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 12),
                                                            ("Heading 2", 14, True, False, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
                                                            ("Heading 3", 12, True, True, WD_ALIGN_PARAGRAPH.LEFT, 8, 4)]:
        h = st[name]
        _set_font(h, size, bold, italic, color="000000")
        h.paragraph_format.alignment = align
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.line_spacing = 1.15
    c = st["Caption"]
    _set_font(c, 11, True, False, color="000000")
    c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.line_spacing = 1.15
    c.paragraph_format.space_before = Pt(4)
    c.paragraph_format.space_after = Pt(10)
    for nm in ("List Bullet", "List Number"):
        _set_font(st[nm], 12)
        st[nm].paragraph_format.line_spacing = 1.5
        st[nm].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_page_number_footer(sec)
    return doc


def add_field(par, instr, cached="1", bold=None):
    def r(kind=None, text=None, instr_text=None):
        run = par.add_run()
        _set_font(run, None, bold)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        if instr_text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr_text
            run._r.append(it)
        if text is not None:
            run.text = text
        return run
    r("begin")
    r(instr_text=f" {instr} ")
    r("separate")
    r(text=cached)
    r("end")


def add_page_number_footer(section, prefix=""):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if prefix:
        run = p.add_run(prefix)
        _set_font(run, 11)
    add_field(p, "PAGE", "1")
    for r in p.runs:
        _set_font(r, 11)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _runs(par, text, size=None, italic=None, bold_all=False):
    """Add text with **bold** and *italic* markup to a paragraph."""
    for part in re.split(r"(`[^`]+`|\*\*.+?\*\*|\*[^*\s][^*]*?\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = par.add_run(part[1:-1])
            _set_font(run, (size or 12) - 1.5, None, italic, name="Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            _set_font(run, size, True, italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            _set_font(run, size, bold_all or None, True)
        else:
            run = par.add_run(part)
            _set_font(run, size, True if bold_all else None, italic)


def _cell_border_off(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{b}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tblPr.append(borders)


class Numbering:
    """Figure/table numbers. chap=True gives chapter-based labels (5.1, 5.2, ..., A.1) and records captions."""

    def __init__(self, chap=False):
        self.chap = chap
        self.fig = {}
        self.tab = {}
        self.fig_caps = []   # (label, text, key)
        self.tab_caps = []

    def scan(self, blocks):
        self.fig, self.tab, self.fig_caps, self.tab_caps = {}, {}, [], []
        prefix, cf, ct = "", {}, {}
        f = t = 0
        for b in blocks:
            if b[0] == "h1":
                m = re.match(r"CHAPTER (\d+)", b[1]) or re.match(r"APPENDIX ([A-Z])", b[1])
                prefix = m.group(1) if m else ""
            elif b[0] == "fig":
                f += 1
                cf[prefix] = cf.get(prefix, 0) + 1
                lab = f"{prefix}.{cf[prefix]}" if (self.chap and prefix) else str(f)
                self.fig[b[1]] = lab
                self.fig_caps.append((lab, b[3], b[1]))
            elif b[0] == "tbl":
                t += 1
                ct[prefix] = ct.get(prefix, 0) + 1
                lab = f"{prefix}.{ct[prefix]}" if (self.chap and prefix) else str(t)
                self.tab[b[1]] = lab
                self.tab_caps.append((lab, b[2], b[1]))


def resolve(text, num):
    def fig(m):
        return f"Fig. {num.fig[m.group(1)]}" if m.group(1) in num.fig else "the corresponding figure"

    def tab(m):
        return f"Table {num.tab[m.group(1)]}" if m.group(1) in num.tab else "the corresponding table"
    text = re.sub(r"\{fig:([a-z0-9_]+)\}", fig, text)
    return re.sub(r"\{tab:([a-z0-9_]+)\}", tab, text)


def caption(doc, label, number, text, seq=True):
    p = doc.add_paragraph(style="Caption")
    if seq:
        run = p.add_run(f"{label} ")
        _set_font(run, 11, True)
        add_field(p, "SEQ " + label + " " + chr(92) + "* ARABIC", str(number), bold=True)
    else:
        run = p.add_run(f"{label} {number}")
        _set_font(run, 11, True)
    run = p.add_run(f": {text}")
    _set_font(run, 11, True)
    return p


def _new_list_instance(doc):
    """A fresh numbering instance of the 'List Number' style so that every numbered list restarts at 1."""
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"]
    numpr = style.element.pPr.find(qn("w:numPr")) if style.element.pPr is not None else None
    base_id = numpr.find(qn("w:numId")).get(qn("w:val"))
    abstract = None
    for n in numbering.findall(qn("w:num")):
        if n.get(qn("w:numId")) == base_id:
            abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    an = OxmlElement("w:abstractNumId")
    an.set(qn("w:val"), abstract)
    num.append(an)
    ov = OxmlElement("w:lvlOverride")
    ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride")
    so.set(qn("w:val"), "1")
    ov.append(so)
    num.append(ov)
    numbering.append(num)
    return new_id


def render(doc, blocks, num=None):
    num = num or Numbering()
    num.scan(blocks)
    for b in blocks:
        k = b[0]
        if k in ("h1", "h2", "h3"):
            doc.add_paragraph(resolve(b[1], num), style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[k])
        elif k == "p":
            p = doc.add_paragraph()
            _runs(p, resolve(b[1], num))
        elif k in ("bul", "num"):
            nid = _new_list_instance(doc) if k == "num" else None
            for it in b[1]:
                p = doc.add_paragraph(style="List Bullet" if k == "bul" else "List Number")
                if nid is not None:
                    ppr = p._p.get_or_add_pPr()
                    numpr = OxmlElement("w:numPr")
                    il = OxmlElement("w:ilvl")
                    il.set(qn("w:val"), "0")
                    ni = OxmlElement("w:numId")
                    ni.set(qn("w:val"), str(nid))
                    numpr.append(il)
                    numpr.append(ni)
                    ppr.append(numpr)
                _runs(p, resolve(it, num))
        elif k == "fig":
            _, key, path, cap, width = b
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            p.add_run().add_picture(path, width=Inches(width))
            caption(doc, "Figure", num.fig[key], resolve(cap, num), seq=not num.chap)
        elif k == "tbl":
            _, key, cap, header, rows, widths = b
            cp = caption(doc, "Table", num.tab[key], resolve(cap, num), seq=not num.chap)
            cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_after = Pt(4)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for i, h in enumerate(header):
                c = table.rows[0].cells[i]
                c.text = ""
                par = c.paragraphs[0]
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(2)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.keep_with_next = True
                _runs(par, resolve(str(h), num), 10, bold_all=True)
                shade(c, GREEN)
            trpr = table.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trpr.append(th)
            for row in rows:
                cells = table.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = ""
                    par = cells[i].paragraphs[0]
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(2)
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 or not re.fullmatch(r"[-+\d.,%/ ()x]+", str(v))) else WD_ALIGN_PARAGRAPH.CENTER
                    _runs(par, resolve(str(v), num), 10)
                    if len(table.rows) == 2:
                        par.paragraph_format.keep_with_next = True
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                cs = OxmlElement("w:cantSplit")
                trpr.append(cs)
            if widths:
                total = sum(widths)
                scale = min(1.0, 6.25 / total)
                for row in table.rows:
                    for i, w in enumerate(widths):
                        row.cells[i].width = Inches(w * scale)
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(6)
            sp.paragraph_format.line_spacing = 1.0
        elif k == "code":
            for line in b[1].rstrip("\n").split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Inches(0.15)
                run = p.add_run(line if line else " ")
                _set_font(run, 8.5, name="Consolas")
                ppr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "F1F5F9")
                ppr.append(shd)
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)
            sp.paragraph_format.line_spacing = 1.0
        elif k == "eq":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            run = p.add_run(resolve(b[1], num))
            _set_font(run, 12, italic=True)
            if len(b) > 2 and b[2]:
                r2 = p.add_run(f"        ({b[2]})")
                _set_font(r2, 12)
        elif k == "note":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(resolve(b[1], num))
            _set_font(run, 10.5, italic=True)
        elif k == "pb":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "raw":
            b[1](doc)
        else:
            raise ValueError(f"unknown block {k}")
    return num


def center(doc, text, size=12, bold=False, italic=False, before=0, after=6, color=None, caps=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text.upper() if caps else text)
    _set_font(run, size, bold, italic, color=color)
    return p


def student_table(doc, students, guide_line=None):
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("S.No.", "Name of the student", "Hall ticket no.")):
        c = t.rows[0].cells[i]
        c.text = ""
        _runs(c.paragraphs[0], h, 11, bold_all=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
        shade(c, GREEN)
    for i, (n, r) in enumerate(students, 1):
        cells = t.add_row().cells
        for j, v in enumerate((str(i), n, r)):
            cells[j].text = ""
            _runs(cells[j].paragraphs[0], v, 11)
            cells[j].paragraphs[0].paragraph_format.line_spacing = 1.0
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for i, w in enumerate((0.8, 3.0, 1.8)):
            row.cells[i].width = Inches(w)
    return t


def title_block(doc, project, subtitle=None):
    """Header used at the top of each chapter document: department, batch, project title, document title."""
    center(doc, "Department of Computer Science and Engineering (Data Science)", 12, True, after=2)
    center(doc, "CSDS-A  |  Batch A5", 12, False, after=8)
    center(doc, project, 14, True, after=8)
    if subtitle:
        center(doc, subtitle, 14, True, after=10, caps=True)


def toc_field(doc, instr, placeholder):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_field(p, instr, placeholder)
    return p


def set_update_fields_on_open(doc):
    settings = doc.settings.element
    e = OxmlElement("w:updateFields")
    e.set(qn("w:val"), "true")
    settings.append(e)
