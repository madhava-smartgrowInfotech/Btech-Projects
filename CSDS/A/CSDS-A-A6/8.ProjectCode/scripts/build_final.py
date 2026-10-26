"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides) for ClauseGuard.

    python scripts/build_final.py
Every number comes from docs/REPORT_FACTS.md and experiments/**/metrics.json.
"""
from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import drop_slide_notes_text, fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import BATCH, CHARTS, COLLEGE, DEPT, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, REFERENCES, SHOTS, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE  |  Batch A6  |  Final Review"


def px(v):
    return Emu(int(v * 9525))


def scrub_theme(path):
    """The template's theme is named after another project: rename it."""
    tmp = Path(str(path) + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/theme/") and item.filename.endswith(".xml"):
                text = data.decode("utf-8")
                for old in ("Hidden Risk", "FraudShield"):
                    text = text.replace(old, PRODUCT)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 8)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182, cap_top=546, cap_pt=13):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l, cap_pt)
    set_paras(shape(s, cap_ids[1]), cap_r, cap_pt)
    for cid in cap_ids:
        c = shape(s, cid)
        c.top = px(cap_top)
        c.height = px(100)


def diagram_with_panel(s, img, points, pic_box, panel_box, size=16):
    """Large diagram on the left, key points (the template's caption shape, moved and left-aligned) on the right."""
    put_picture(s, img, pic_box, remove_ids=(7,), name=Path(img).stem)
    cap = shape(s, 8)
    cap.left, cap.top, cap.width, cap.height = (px(v) for v in panel_box)
    fill(s, 8, ["\u2022  " + p for p in points], size, min_pt=11, gap_pt=6)
    for p in cap.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.italic = False
            r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    cap.text_frame.vertical_anchor = MSO_ANCHOR.TOP


def build():
    prs = Presentation(SRC)
    drop_slide_notes_text(prs)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author = f"{PRODUCT}: Final Review", TITLE, "Batch A6"
    cp.keywords = "legal contracts, predatory clauses, clause classification, complexity score, Indian contract law"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A6"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)

    # 1 title
    s = S[0]
    set_paras(shape(s, 4), TITLE)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A6")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE (Data Science)"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING (DATA SCIENCE)  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of ClauseGuard: it reads a PDF or DOCX contract, classifies its clauses, flags predatory ones under eight Indian-law rules and explains the risk in plain language.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives and what was built", "Proposed System & Architecture", "Methodology & Pipeline",
                "Design & Implementation", "Results, Threshold Analysis, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["Individuals, freelancers and small businesses sign contracts with unfair termination, penalty, liability and arbitration terms and cannot afford a lawyer to read them.",
                "ClauseGuard takes a PDF or DOCX contract, splits it into clauses, classifies each into one of 20 CUAD categories with Gemini embeddings, flags 8 predatory patterns tied to Indian statutes, and finds 3 risky clause combinations.",
                "It adds an A-F complexity grade from readability, jargon, cross-references and sentence length, and Gemini writes a plain-language reason, the quoted provision and safer wording for each flag.",
                "Measured: classification top-1 79.2 %, top-3 97.9 % on 240 held-out CUAD spans; rules P 1.00 / R 0.94 on their own 40-clause set but P 0.79 / R 0.46 on an independent probe set; 68 tests pass.",
                "Differences from the abstract: Gemini replaces Claude, SQLite + NumPy replaces pgvector, an in-memory NetworkX graph replaces Neo4j.",
                "Keywords: contract analysis, predatory clauses, CUAD, embeddings, Indian contract law."], 16, bullet=True)
    set_notes(s, "Abstract: problem, approach, headline measurements and the honest differences between the abstract and the build.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [REFERENCES[0],
                "Task: a Saudi legal assistant with two functions - a legal chatbot and a contract analysis module - for Arabic text under Saudi law.",
                "Method: Gemini 1.5 Flash with RAG over a custom Arabic corpus (GATE-AraBERT-v1 embeddings, Chroma, cross-encoder reranking, JSON-constrained output at temperature 0.2); contract clauses checked against an essential-clause checklist.",
                "Reported: 150 legal questions over 10 domains and 50 contract drafts; Gemini 1.5 Flash 98.67 % (148 of 150), GPT-3.5 98.00 %, GPT-4 97.33 %; contract analysis judged clause by clause on sample outputs.",
                "Our reading: a strong RAG-first, jurisdiction-aware design, but Arabic and Saudi-specific, with a small expert-reviewed set and no clause-combination reasoning. We adopt the jurisdiction-aware idea for Indian contracts and add combinations and a complexity score."], 15, bullet=True)
    set_notes(s, "Read from the PDF in Documentary/. ClauseGuard differs in task, language, data and metrics, so the numbers are not comparable.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Almutairy et al., 2026 (base) [1]", "Gemini 1.5 Flash + RAG; checklist-based contract analysis", "150 questions, 50 contract drafts (Saudi law, Arabic)", "98.67 % Q&A accuracy; contract analysis judged by inspection", "Saudi/Arabic only; small set; no clause combinations or complexity score"],
            ["Liu et al., 2025 [2]", "19 LLMs asked to find clause-level legal risks", "CUAD test data, 41 categories", "Proprietary models lead; top two F2 0.672 and 0.678", "Benchmark of extraction, not an Indian-law tool"],
            ["Lincoln et al., 2026 [3]", "Domain-trained small model vs 5 frontier LLMs", "24 held-out contracts, 508 field instances", "Macro F1 0.812, micro F1 0.842; cost down 78-97 %", "Field extraction, not predatory-clause detection"],
            ["Charan R et al., 2026 [4]", "AWS Lambda + GenAI, 10 clause types, risk score 0-100", "20 synthetic contracts; 30 for expert correlation", "94.2 % extraction accuracy; risk vs expert r = 0.87", "Synthetic documents; template-gap logic"],
            ["Hendrycks et al., 2021 [8]", "CUAD: expert-annotated contract review dataset", "510 contracts, 41 categories", "Dataset used here for classification and the corpus run", "US SEC filings, not Indian contracts"],
            ["ClauseMiner, Legal-BERT, construction-contract LLM [5]-[7]", "Clause extraction and risk detection", "Not retrieved", "Not retrieved", "Listed in the batch survey only"]]
    table(s, rows, [0.17, 0.25, 0.19, 0.22, 0.17], 15, 500)
    set_notes(s, "[1]-[4] and [8] were read or counted from local copies. [5]-[7] could not be retrieved, so no results are claimed for them.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Saudi law and Arabic text; consultation plus a checklist of missing clauses. No graph of clause combinations, no complexity score, and no separate scoring of predatory-clause detection"],
            ["Benchmarks [2], [3]", "Measure how well LLMs extract clauses from US contracts; none flags predatory terms or targets Indian law"],
            ["Risk-score system [4]", "Scores missing clauses against a template on synthetic documents; does not judge whether the wording of a present clause is one-sided"],
            ["Existing practice (batch survey)", "Lawyer review is costly and slow; commercial tools target US/EU enterprises; general chatbots summarise without clause-level references"],
            ["All reviewed work", "Jurisdiction-aware detection of predatory clause combinations for Indian contracts, explained to non-experts, is underrepresented. ClauseGuard prototypes it; the evidence gap remains (no Indian contract corpus)"]]
    table(s, rows, [0.24, 0.76], 17, 440)
    set_notes(s, "The gap is the combination of clause-level fairness, Indian law and plain-language explanation. The prototype is evaluated on US CUAD spans and short written clauses only.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Non-experts cannot easily spot predatory or ambiguous clauses in contracts they are asked to sign.",
                "Legal review is expensive and inaccessible for freelancers, gig workers and small businesses.",
                "Existing tools mostly ignore Indian contract law and analyse each clause in isolation, missing combinations that are harmless alone but dangerous together.",
                "Users sign without understanding hidden liabilities, penalties, one-sided arbitration or auto-renewal.",
                "Need: an affordable, jurisdiction-aware tool that flags predatory clauses and combinations and explains the risk in plain language before signing."], 20, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["1. Segment contracts and classify clause types - done: heading splitter; 79.2 % top-1, 97.9 % top-3 on 240 CUAD spans.",
                "2. RAG with Claude API + pgvector - not as stated: Gemini with SQLite + NumPy; no retrieval over a legal corpus (prompts carry the rule's legal basis).",
                "3. Neo4j knowledge graph of clauses, concepts, precedents - partly: NetworkX graph with 3 combination patterns; no concept or precedent nodes.",
                "4. Flag predatory clauses and combinations under Indian law - done in a limited form: 8 regex rules + 3 combinations; own set P 1.00 / R 0.94, independent probe P 0.79 / R 0.46.",
                "5. Complexity score and plain-language summary - done: A-F grade and Gemini summary (wording quality not expert-verified).",
                "6. Evaluate classification and predatory detection separately - done: separate metrics, plus a probe set and a 510-contract run."], 17, bullet=False)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["Upload - a signed-in user sends a PDF or DOCX; the file is read, analysed and deleted.",
                "Parse and segment - PDF per page, DOCX as page 1; numbered headings or paragraphs become clauses with page numbers.",
                "Classify - Gemini embedding (3,072-d) compared with 20 CUAD category centroids; optional Gemini tie-break.",
                "Detect - 8 regex rules mapped to Indian statutes flag predatory patterns; Gemini writes the reason, quote and safer wording.",
                "Combine and score - 3 risky combinations across clauses; complexity score 0-100 and grade A-F.",
                "Summarise - Gemini plain-language summary and key obligations.",
                "Review - report, clause detail, clause graph and an evaluation page, all behind a JWT login."], 17, bullet=False)

    # 10 architecture
    s = S[9]
    diagram_with_panel(s, DIAG / "01_system_architecture.png",
                       ["React 19 + Vite web app: sign-up, upload, report, clause detail, graph, evaluation",
                        "FastAPI backend: 9 endpoints, JWT + bcrypt, ownership check on every contract route",
                        "Services: parsing, classification, rules, graph, complexity, explain, vector store",
                        "Storage: SQLite (4 tables, embeddings as BLOBs), cached centroids, metrics file",
                        "Gemini API: embeddings, tie-break, explanations, summary - the only network dependency"],
                       (43, 178, 720, 494), (780, 200, 460, 440), 16)

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "One contract from upload to report: ten steps; pink steps call Gemini", 14)

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Rules & Models")
    fill(s, 7, ["CUAD v1 (Atticus Project, CC BY 4.0): 510 SEC-filed contracts, 41 categories. ClauseGuard uses a 20-category substantive subset: 25 training spans per category (500) and 12 held-out spans (240), seed 42.",
                "Split caveat: the split is by span text, not by contract; 192 of 240 held-out spans share a contract with training spans and 4 are exact duplicates.",
                "Rule tests: 40 clauses written with the rules, plus a 40-clause probe set (24 predatory paraphrases, 16 benign look-alikes) written in this run.",
                "Classifier: gemini-embedding-001 vectors, cosine to 20 category centroids; tie-break only if top-1 minus top-2 < 0.03 and top-1 < 0.60; Unclassified below 0.35.",
                "Rules: restraint of trade, penalty, unilateral termination, unlimited liability, one-sided arbitration, auto-renewal, IP overreach, no-refund; tagged to ICA 1872, Arbitration Act 1996, Copyright Act 1957, Consumer Protection Act 2019.",
                "Complexity = 0.35 (100 - Flesch) + 0.30 legalese + 0.20 cross-references + 0.15 sentence length; weights chosen by us, not validated."], 15, bullet=True)
    set_notes(s, "No Indian contract corpus was available; CUAD is US data. No model is trained: classification uses centroids of embeddings and the rules are regular expressions.")

    # 13 decision workflow
    s = S[12]
    diagram_with_panel(s, DIAG / "10_algorithm_flowchart.png",
                       ["Algorithm 1 (process_upload): validate type, extract pages, split clauses, classify and flag each clause, combine, score, summarise, commit once",
                        "Algorithm 2 (classify_clause): embed, rank 20 centroids, tie-break only when close and below 0.60, floor at 0.35",
                        "A Gemini failure after 3 retries returns HTTP 502 and nothing is stored"],
                       (43, 178, 560, 494), (630, 210, 610, 420), 17)

    # 14 user-system interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    diagram_with_panel(s, DIAG / "09_sequence_diagram.png",
                       ["The browser posts the file with the JWT; the route checks the user and writes a temporary file",
                        "For every clause the pipeline embeds, applies the rules and asks Gemini to explain each hit; rows are flushed",
                        "Combinations, complexity and the Gemini summary are added, the contract is committed and the temporary file deleted",
                        "Clause detail and graph calls check ownership and answer 404 for other users' contracts"],
                       (43, 178, 560, 494), (630, 210, 610, 420), 17)

    # 15 database
    s = S[14]
    diagram_with_panel(s, DIAG / "06_er_diagram.png",
                       ["4 SQLite tables: users, contracts, clauses, flags; 3 foreign keys",
                        "A clause stores its page, type, confidence, method and a 3,072-float embedding BLOB",
                        "Risky combinations are JSON on the contract; the clause graph is rebuilt per request",
                        "contract_type_guess is never set (always 'general')"],
                       (43, 178, 720, 494), (780, 210, 460, 420), 16)

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["Parsing", "PyMuPDF per page, python-docx as one page; heading regex split, paragraph fallback up to ~400 characters; fragments under 25 characters dropped"],
            ["Classification", "Gemini embeddings + NumPy cosine over 20 centroids cached in an .npz file; Gemini JSON tie-break; clause vectors kept in SQLite"],
            ["Rules + graph", "8 case-insensitive regex rules with legal basis; 3 combinations (NetworkX); one Flag row per hit"],
            ["Explanations", "One Gemini call per flag (reason, quoted provision, safer wording) and one for the summary; 3 retries; JSON parsing"],
            ["Complexity", "Flesch reading ease, 34-term jargon list, cross-reference regex, sentence length -> 0-100 -> A-F"],
            ["Web app + security", "React 19 + TanStack Query; FastAPI, bcrypt, JWT 1,440 min, ownership checks; 68 pytest tests with Gemini stubbed"]]
    table(s, rows, [0.18, 0.82], 17, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Measured Performance")
    rows = [["Measure", "Result", "Note"],
            ["Classification, 240 held-out CUAD spans, 20 classes", "top-1 79.2 %, top-3 97.9 %", "chance 5 %; contract-disjoint top-1 80.0 %; embedding + logistic regression 80.0 %"],
            ["Rules, shipped 40-clause set", "P 1.00, R 0.94, F1 0.97", "clauses written with the rules"],
            ["Rules, independent probe set (24 + 16)", "P 0.79, R 0.46, F1 0.58", "13 paraphrases missed, 3 false alarms"],
            ["Rules on 5,032 unrelated CUAD spans", "fired on 0.74 %", "recall proxies: Non-Compete 8.2 %, Liquidated Damages 56.2 %, Uncapped 0 %"],
            ["Same contract as DOCX vs PDF", "8 vs 6 of 8 rules", "line breaks defeat two rules; collapsed text gives 8"],
            ["All 510 CUAD contracts, no API", "56 ms per contract", "23,869 clauses, 2.2 % flagged; grades A4 B169 C268 D68 F1"],
            ["Uploads with real Gemini calls", "15.5 / 54.6 / 113.6 s", "0, 1 and 8 flags; about 12 s per text call"],
            ["Automated tests", "68 passed, 5 xfailed", "Gemini stubbed; xfails document known gaps"]]
    table(s, rows, [0.30, 0.22, 0.48], 17, 460)
    set_paras(shape(s, 8), "Classification is measured on US CUAD spans and the rules on short written clauses; no Indian contract corpus exists here and no lawyer reviewed the Gemini text.", 13)
    set_notes(s, "The shipped 40-clause set was written together with the rules, so its 100 % precision is optimistic; the probe set is the fairer estimate.")

    # 18 confusion matrix + rule test sets
    s = S[17]
    set_paras(shape(s, 4), "Confusion Matrix & Rule Test Sets")
    two_pics(s, CHARTS / "05_confusion_matrix.png", CHARTS / "03_rules_two_test_sets.png",
             "240 held-out spans: Uncapped vs Cap On Liability and Non-Compete vs No-Solicit are the main confusions",
             "Rules: own 40-clause set P 1.00 / R 0.94; independent probe set P 0.79 / R 0.46", box_h=405, cap_top=592, cap_pt=12)
    set_notes(s, "Left: confusion matrix of the shipped nearest-centroid classifier. Right: the same rules scored on two test sets.")

    # 19 threshold / baseline analysis
    s = S[18]
    set_paras(shape(s, 4), "Threshold Analysis & Baselines")
    two_pics(s, CHARTS / "09_confidence_thresholds.png", CHARTS / "02_classification_baselines.png",
             "Top-1 similarity is 0.657-0.849: the 0.60 and 0.35 tests never fire (152 of 240 spans pass the margin test alone)",
             "Same split: nearest centroid 0.792; TF-IDF + logistic regression 0.675; embedding + logistic regression 0.800",
             box_h=405, cap_top=592, cap_pt=12)
    set_notes(s, "The tie-break needs both a margin below 0.03 and a top-1 below 0.60, and no span is below 0.657, so the Gemini tie-break and the Unclassified label are dead code with this embedding model.")

    # 20 performance & testing
    s = S[19]
    rows = [["Measurement", "Result"],
            ["API response, median / p95 (no Gemini)", "health 2.5 / 4.0 ms; report 9.3 / 11.2; clause detail 5.0 / 6.1; graph 9.2 / 11.8; bcrypt login 254.7 / 260.8 ms"],
            ["Upload to report (real Gemini calls)", "15.5 s (0 flags), 54.6 s (1 flag), 113.6 s (8 flags); the frontend times out at 180 s"],
            ["Automated tests", "68 passed, 5 xfailed in about 10 s: rules, graph, complexity, parsing, classification, auth, upload, ownership, errors"],
            ["Defects found", "PDF line breaks lose 2 of 8 rules; text before the first heading is dropped; tie-break and Unclassified never fire; Evaluation screen calls a self-written set 'held-out'"],
            ["Machine and quota", "i5-12500, 15.7 GB, no GPU; the default Gemini text model allowed 20 requests per day, so a lighter model was used for the captures"],
            ["Not measured", "load and concurrency, lawyer review of explanations, Indian contracts, browser tests"]]
    table(s, rows, [0.24, 0.76], 17, 470)
    set_notes(s, "Latencies are from a live server on 127.0.0.1 with other jobs running on the same machine.")

    # 21 comparison
    s = S[20]
    set_paras(shape(s, 4), "Comparison with Base Paper (context only \u2013 different task and data)")
    rows = [["System", "Task and data", "Method", "Reported / measured", "Extras", "Comparable?"],
            ["Base paper [1]", "Saudi legal Q&A (150 questions) and contract analysis (50 drafts); Arabic", "Gemini 1.5 Flash + RAG (Chroma), reranking", "98.67 % Q&A accuracy (148 of 150); contract analysis judged by inspection", "Chatbot, checklist of missing clauses", "-"],
            ["ContractEval [2]", "CUAD test data, 41 categories", "19 LLMs prompted per category", "Top two F2 0.672 and 0.678", "Benchmark only", "No: span extraction, other metrics"],
            ["ClauseGuard", "Contract fairness for Indian-law use; CUAD (US) spans for classification, 40 + 40 written clauses for rules", "Gemini embeddings + centroids, 8 regex rules, 3 combinations, complexity score", "Top-1 79.2 %, top-3 97.9 %; rules P 1.00 / R 0.94 (own set), P 0.79 / R 0.46 (probe)", "Safer wording, clause graph, summary", "No: different task, data, metrics"]]
    table(s, rows, [0.11, 0.21, 0.19, 0.23, 0.13, 0.13], 15, 470)
    set_paras(shape(s, 8), "Different tasks, languages, data and metrics: no numeric head-to-head is possible and no superiority is claimed.", 13)
    set_notes(s, "State plainly: the base paper reports Q&A accuracy on Saudi-law questions; ClauseGuard reports clause classification and rule scores on other data.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Contract Report")
    two_pics(s, SHOTS / "08_report_freelance_top.png", SHOTS / "09_report_freelance_clauses.png",
             "One-sided freelance DOCX: grade B (23.3), 8 of 11 clauses flagged, Gemini summary and severity chart",
             "Key obligations, three risky combinations and the clause table with type, flags and risk")
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Clause Detail and Clause Graph")
    two_pics(s, SHOTS / "11_clause_detail_termination.png", SHOTS / "14_clause_graph.png",
             "Clause 2: Unilateral Termination (high) with reason, quoted provision and safer wording",
             "Clause graph: red nodes triggered a rule; lines join clauses in a risky combination")
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Fair Contract and Evaluation")
    two_pics(s, SHOTS / "15_report_balanced.png", SHOTS / "21_evaluation_lower.png",
             "Fair consulting DOCX: grade A (13.7), no flagged clause",
             "Evaluation screen: per-rule counts and CUAD classification 79 % top-1, 98 % top-3")

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built ClauseGuard: upload a PDF or DOCX, get clauses classified, 8 predatory patterns and 3 combinations flagged, an A-F complexity grade and plain-language reasons with safer wording.",
                "Measured: classification top-1 79.2 % / top-3 97.9 % on 240 CUAD spans; rules P 1.00 / R 0.94 on their own set but P 0.79 / R 0.46 on an independent probe set; 68 tests pass.",
                "Findings: PDF line breaks hide 2 of 8 rules; the Gemini tie-break and Unclassified label never fire; the complexity grade measures wording, not fairness (the one-sided contract scored B).",
                "Differences from the abstract: Gemini, SQLite + NumPy and NetworkX replace Claude, pgvector and Neo4j; there is no legal-corpus retrieval.",
                "Objectives are met as a working prototype for Indian-law-oriented screening, not as validated legal advice."], 19, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Replace regexes with a trained or LLM-assisted detector and test it on independently labelled clauses, ideally from Indian contracts reviewed by lawyers.",
                "Fix input handling: join PDF lines before matching, keep the preamble, handle 'Article I' headings and clauses that span pages, and store partial results when Gemini fails.",
                "Re-tune the tie-break and Unclassified thresholds to the embedding model; split train and test by contract.",
                "Build the promised retrieval layer over statutes and precedents (vector database) and a real knowledge graph with legal concepts.",
                "Add a lawyer-reviewed evaluation of explanations, a paid Gemini tier or a local model, rate limiting, and a user study."], 20, bullet=True)

    # 27 references
    s = S[26]
    set_paras(shape(s, 7), [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], 9)
    for sp in shape(s, 7).text_frame._txBody.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}spcPts"):
        sp.set("val", "150")

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    scrub_theme(OUT)
    print("saved", OUT, len(Presentation(OUT).slides), "slides")


if __name__ == "__main__":
    build()
