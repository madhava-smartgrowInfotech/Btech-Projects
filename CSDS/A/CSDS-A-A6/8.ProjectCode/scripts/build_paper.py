"""Build 19.IEEE_Paper.docx: a six-page, two-column IEEE conference paper on ClauseGuard.

    python scripts/make_paper_figures.py ; python scripts/build_paper.py ; python scripts/paper_finish.py

Every number comes from docs/REPORT_FACTS.md and experiments/**/metrics.json. The paper has its own reference list (20 entries,
base paper first) chosen from doc_common.REFERENCES and numbered in order of first citation; the three batch-survey papers that could
not be retrieved (entries 5-7 of that list) are not cited.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paper_kit as K  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from doc_common import BATCH, GUIDE, REFERENCES, ROOT, STUDENTS  # noqa: E402

OUT = BATCH / "19.IEEE_Paper.docx"
PAPER = ROOT / "docs" / "paper"

# key -> index in doc_common.REFERENCES
KEYS = {"base": 0, "cont": 1, "fgc": 2, "ldas": 3, "cuad": 7, "lbert": 8, "bert": 9, "rag": 10, "sbert": 11, "gemb": 12, "tfidf": 13,
        "flesch": 14, "nx": 15, "sk": 16, "bcrypt": 17, "jwt": 18, "ica": 19, "arb": 20, "cpa": 21, "cr": 22}
ORDER: list[str] = []


def C(text: str) -> str:
    """Replace [@a,b] by [n], [m] numbered in order of first citation."""
    def sub(m):
        out = []
        for k in m.group(1).split(","):
            if k not in ORDER:
                ORDER.append(k)
            out.append(f"[{ORDER.index(k) + 1}]")
        return ", ".join(out)
    return re.sub(r"\[@([a-z,]+)\]", sub, text)


class Paper(K.Paper):
    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(K.TITLE)
        K.font(r, 24)
        p.paragraph_format.space_after = Pt(6)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [(GUIDE, "[DESIGNATION], Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Dept. of CSE (Data Science)", False), ("[COLLEGE NAME]", False), ("[CITY]", False), ("[EMAIL]", False)):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                K.font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        sectpr = sec._sectPr
        cols = sectpr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectpr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def references(self):
        self.h1("References")
        for i, key in enumerate(ORDER, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]\t{REFERENCES[KEYS[key]]}")
            K.font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


K.TITLE = "ClauseGuard: Legal Contract Complexity and Predatory Clause Detector"
K.ABSTRACT = (
    "Freelancers, individuals and small businesses often sign contracts they cannot assess, and professional review is costly. "
    "This paper presents ClauseGuard, a web system that reads a PDF or DOCX contract, splits it into clauses, labels each clause with one of "
    "20 CUAD categories using Gemini embeddings and a nearest-centroid rule, flags eight predatory patterns tied to Indian contract law "
    "with regular expressions, detects three risky clause combinations in a graph, grades textual complexity from A to F and asks a language "
    "model for plain-language reasons and safer wording. Classification and predatory detection are evaluated separately. On 240 held-out CUAD "
    "spans, top-1 accuracy is 0.792 and top-3 accuracy 0.979. The rules score precision 1.00 and recall 0.94 on the 40 clauses written with them, "
    "but precision 0.79 and recall 0.46 on an independent probe set. PDF line breaks hide two of eight rules on identical text, and two "
    "classifier thresholds never trigger. Sixty-eight automated tests pass. The limits are reported plainly."
)
K.KEYWORDS = "clause classification, complexity scoring, contract analysis, Indian contract law, legal NLP, predatory clause detection, text embeddings"


def build():
    P = Paper()
    P.front()
    P.abstract()
    para = lambda t, **k: P.para(C(t), **k)

    # ---------------------------------------------------------------- I
    P.h1("I. Introduction")
    para("Individuals, freelancers, gig workers and small businesses now sign digital contracts without a lawyer. Such contracts can hold "
         "unfair termination rights, penalties that bear no relation to loss, unlimited liability, one-sided arbitration and automatic "
         "renewals, and the weaker party seldom notices them. Professional review is slow and expensive, so the information gap stays. "
         "Large language models can read a contract and explain it in plain words, but a general chatbot is not grounded in the "
         "contract's clauses and gives no clause-level reference. Commercial analysis tools and research systems mostly target US or EU law, "
         "classify one clause at a time, and often mix up two different questions: what type a clause is, and whether it is predatory.")
    para("The base paper [@base] builds a retrieval-augmented legal assistant and contract checker for Saudi law. We adopt its jurisdiction-"
         "aware, language-model-explained design for Indian contracts, but our build differs from the original plan in three ways: Gemini replaces the "
         "Claude API, vectors are kept in SQLite and compared with NumPy instead of pgvector, and a per-contract NetworkX graph replaces Neo4j. "
         "This paper therefore reports what was built and measured, not what was planned. Its contributions are:", after=2)
    P.numbered([
        "an end-to-end pipeline for PDF and DOCX contracts: segmentation, 20-category CUAD classification, eight rules tagged with Indian statutes, three risky combinations, an A-F complexity grade, and LLM-written reasons, safer wording and summaries;",
        "a separate evaluation of classification (240 CUAD spans) and predatory detection (a 40-clause set, an independent 40-clause probe set, a CUAD cross-check and a 510-contract run);",
        "measured findings about the design: where the rules and the classifier fail, what PDF input does to detection, and which thresholds never act;",
        "a tested React and FastAPI application with 68 passing tests and real captured output screens.",
    ])
    para("Section II reviews related work, Section III describes the method, Section IV reports results and Section V concludes.", before=2)
    P.h2("A. System Hardware")
    para("All measurements were taken on one desktop with a 12th-generation Intel Core i5-12500 (12 threads), 15.7 GB of memory and no GPU "
         "in use, running Windows 11. Embeddings and text generation are remote Gemini API calls, so no model was trained or run locally.")
    P.h2("B. Software Requirements")
    para("The back end uses Python 3.11.9, FastAPI 0.115, SQLAlchemy 2 over SQLite, PyMuPDF and python-docx for parsing, NumPy, NetworkX 3.3 "
         "[@nx], the google-generativeai client, JSON Web Tokens [@jwt] and bcrypt [@bcrypt]. The front end is React 19 with Vite and Tailwind "
         "on Node 24.19.0. Baselines use scikit-learn [@sk]; tests use pytest 9.1.1.")

    # ---------------------------------------------------------------- II
    P.h1("II. Literature Survey")
    para("**Base paper.** Almutairy et al. [@base] present SmartLaw Advisor, a Saudi-law assistant with two functions: a legal chatbot and a "
         "contract analyser. The chatbot uses Gemini 1.5 Flash with a standard retrieval-augmented generation (RAG) pipeline [@rag] over a "
         "custom Arabic corpus of Saudi regulations: semantic chunking, GATE-AraBERT-v1 embeddings in a Chroma store (a FAISS index of about 18,400 "
         "vectors is also described), cross-encoder re-ranking, OCR-driven normalisation and JSON-constrained clause extraction at temperature 0.2. "
         "The contract analyser checks each contract against a checklist of essential clauses validated by Saudi practitioners and applies an "
         "LLM-assisted semantic check. The system was evaluated on 150 legal questions in ten domains and on 50 real contract drafts. "
         "Gemini 1.5 Flash answered 148 of 150 questions correctly (98.67%) in 1.63 s on average, ahead of GPT-3.5 (98.00%, 147), GPT-4 "
         "(97.33%, 146) and ALLaM (96.00%, 144). Contract-analysis results are qualitative: the authors inspect the clauses returned by each "
         "model. Reported failures are retrieval across chunk boundaries and clause misclassification. The work is limited to Arabic and Saudi "
         "law, uses a small expert-reviewed dataset, and has neither graph reasoning, clause-combination analysis nor a complexity score.")
    para("**Related work.** ContractEval [@cont] benchmarks 19 LLMs (four proprietary, 15 open-source) on the CUAD test data [@cuad] over 41 "
         "risk categories, using F1, F2, Jaccard overlap and the rate of wrongly answering \"no related clause\"; proprietary models lead and "
         "the top two reach F2 of 0.672 and 0.678. Lincoln et al. [@fgc] compare a domain-trained small language model with five frontier models "
         "on 24 held-out contracts (508 field instances, 26 targets) and report a macro F1 of 0.812 and a micro F1 of 0.842, with extraction "
         "cost lower by 78% to 97%. The legal document analyser of [@ldas] extracts ten clause types with cloud services and a language model "
         "and scores missing or risky clauses against a template on a 0-100 scale; it reports 94.2% mean extraction accuracy on 20 synthetic "
         "contracts and a Pearson correlation of 0.87 between its score and expert ratings on 30 contracts. CUAD [@cuad] supplies 510 expert-"
         "annotated contracts with 41 categories. LEGAL-BERT [@lbert] adapts BERT [@bert] to legal text by pre-training on legal corpora, and "
         "Sentence-BERT [@sbert] and Gemini Embedding [@gemb] provide the dense sentence vectors on which nearest-centroid and retrieval "
         "methods rest. TF-IDF weighting [@tfidf] is the classical sparse baseline, and the Flesch score [@flesch] the classical readability measure.")
    P.table([["Work", "Approach", "Data", "Reported result", "Limitation"],
             [C("[@base]"), "Gemini 1.5 Flash, RAG, clause checklist", "Saudi corpus; 150 questions; 50 contracts", "98.67% (148/150) on questions", "Saudi law; contract part qualitative"],
             [C("[@cont]"), "19 LLMs, clause-level risk", "CUAD test data, 41 categories", "Top two F2 0.672, 0.678", "Extraction; US contracts"],
             [C("[@fgc]"), "Small model vs 5 frontier LLMs", "24 contracts, 508 fields", "Macro F1 0.812; micro 0.842", "Field extraction, not risk"],
             [C("[@ldas]"), "Cloud NLP, template risk score", "20 synthetic, 30 sample contracts", "94.2% accuracy; r = 0.87", "Synthetic; missing clauses only"],
             [C("[@cuad]"), "Annotated corpus", "510 contracts, 41 categories", "Dataset", "US SEC filings"],
             [C("[@lbert]"), "Legal-domain BERT", "Legal corpora", "No result quoted", "No contract-risk output"]],
            "Related work", widths=[0.1, 0.25, 0.24, 0.22, 0.19], size=6.0)
    para("**Reading across the work.** Three points shape our design. First, the base paper [@base] shows that a jurisdiction-specific "
         "language-model assistant can reach high answer accuracy, but its contract part is judged by reading outputs, with no count of "
         "missed or false flags; we therefore score flags against labelled clauses. Second, the CUAD-based studies [@cont,cuad] treat "
         "\"which clause type is this\" and \"find the clause that answers a risk question\" as extraction tasks over US filings; none says "
         "whether a clause is unfair under a given law, so a predatory label cannot be read off a CUAD category and must come from rules "
         "or from new annotation. Third, the template-based scoring of [@ldas] reports a strong correlation with expert ratings, but on "
         "synthetic and sample contracts and for missing clauses, not for one-sided wording; we keep clause-level evidence and test "
         "the rules on phrasing they were not written for.")
    para("**Gap.** The tools and studies above concentrate on US, EU or Saudi law, treat clauses singly, and score extraction or "
         "question answering. We found none that combines Indian-law predatory rules, risky clause combinations and a complexity grade, "
         "or that reports clause classification and predatory detection as two separate measurements.")

    # ---------------------------------------------------------------- III
    P.h1("III. Methodology")
    P.h2("A. Architecture and Main Flow")
    para("ClauseGuard is a React app and a FastAPI back end with three routers (authentication, contracts, evaluation), nine endpoints and "
         "four SQLite tables (Fig. 1). A signed-in user uploads a file; the pipeline of Fig. 2 parses, segments, classifies, flags, "
         "combines, scores and summarises it and stores the result. Only embeddings, tie-breaks, reasons, safer wording and summaries "
         "call Gemini; the decision whether a clause is risky is taken by local rules.")
    P.figure(PAPER / "fig1_architecture.png", "System architecture of ClauseGuard.", width=3.35)
    P.figure(PAPER / "fig2_flow.png", "Main flow from upload to report.", width=3.3)
    P.h2("B. Data")
    para("Clause types come from CUAD [@cuad] (510 contracts, CC BY 4.0). We use 20 substantive categories: 500 training spans (25 per "
         "category) build the centroids and 240 held-out spans (12 per category) test them; the split is seeded but made by span text, not by "
         "contract, so 192 of 240 test spans share a contract with a training span. Predatory detection uses the 40 labelled clauses "
         "shipped with the project, which were written by the author of the rules, and a probe set written for this study: 24 predatory "
         "paraphrases (three per rule, mostly without the rule's keywords) and 16 benign look-alikes. Three sample contracts drive the "
         "screens: an invented one-sided freelance agreement, an invented fair consulting agreement and a real CUAD supply agreement.")
    para("Why these choices. A nearest-centroid rule needs no training run, can be rebuilt in seconds from 500 labelled spans, and is easy "
         "to audit: the label of a clause is the category whose average example it most resembles. Training spans are embedded as "
         "documents and test clauses as queries, as the Gemini embedding interface requires, and each clause vector is stored as a "
         "float32 BLOB beside the clause. Deterministic rules were preferred for the risk decision because they are reproducible and "
         "inspectable: the same clause always gets the same flag, and the language model only explains a flag that already exists. "
         "The cost is that a rule sees words, not meaning, which Section IV measures.")
    P.h2("C. Segmentation and Classification")
    para("A PDF is read page by page, a DOCX as one page. A page is split at numbered headings (\"1.\", \"Section 2.1\"); with fewer than two "
         "headings it is split into blank-line paragraphs merged to about 400 characters, and fragments under 25 characters are dropped. Each "
         "clause is embedded with gemini-embedding-001 (3,072 dimensions). With the centroid of category k taken as the mean vector of its 25 training spans,", indent=True)
    P.eq("s_{k} = cos(e, centroid_{k}),   k = 1..20", 1)
    para("the label is the category with the highest s. If the gap between the best two is below 0.03 and the best is below 0.60, the "
         "language model chooses among the top three:", indent=False)
    P.eq("tie-break if (s_{1} − s_{2} < 0.03) ∧ (s_{1} < 0.60)", 2)
    para("and a clause whose final confidence is below 0.35 is labelled \"Other / Unclassified\".", indent=False)
    P.h2("D. Rules, Combinations and Complexity")
    para("Eight case-insensitive regular expressions detect restraint of trade (Indian Contract Act s. 27 [@ica]), penalty against liquidated "
         "damages (s. 74), unilateral termination, unlimited liability, one-sided arbitration (Arbitration and Conciliation Act [@arb]), automatic "
         "renewal and no-refund terms (Consumer Protection Act [@cpa]) and intellectual-property overreach (Copyright Act [@cr]). Each hit becomes one "
         "flag; the model then writes a reason, quotes the triggering provision and drafts safer wording. Three combinations (termination + "
         "no-refund + penalty; restraint + unlimited liability; auto-renewal + termination) are reported when all their rules occur "
         "anywhere in the contract; the graph joins the clauses involved. With Flesch reading ease F [@flesch], legalese terms L per 100 words, "
         "cross-references X over n clauses and mean sentence length S, the complexity score is")
    P.eq("Q = 0.35(100 − F) + 0.30 min(100, 8L) + 0.20 min(100, 40X/n) + 0.15 clip(3(S − 15), 0, 100)", 3)
    para("with grades A below 20, B below 40, C below 60, D below 80 and F above. The weights and cut-offs were chosen by the authors and are "
         "not validated against human ratings. The four rules of high severity are restraint of trade, unilateral termination, unlimited "
         "liability and IP overreach; penalty, one-sided arbitration, automatic renewal and no-refund are of medium severity. A combination is "
         "matched on rule identifiers anywhere in the contract, not on proximity or on which party is bound, so it is a coarse warning. The "
         "explanation prompt asks the model for a two- to three-sentence reason that quotes the clause, the exact triggering phrase and a "
         "fairer rewrite that keeps the clause's legitimate purpose; the summary prompt returns four to six sentences and four to eight key "
         "obligations from the first 40 clauses and the flag list.", indent=False)
    P.h2("E. Implementation")
    para("Passwords are bcrypt-hashed; a JWT (HS256, 1,440 minutes) guards every contract route, and a contract of another user returns 404. "
         "Uploads are limited to .pdf and .docx, written to a temporary file and deleted after analysis. Each contract costs one Gemini "
         "embedding per clause, one text call per flag and one for the summary, with three retries; a Gemini failure returns HTTP 502 and "
         "stores nothing. The default text model resolved to a model with a free quota of 20 requests a day and timed out when this study "
         "began, so all screens were made with gemini-3.1-flash-lite set on the server process only.")
    P.h2("F. Data Model and Interfaces")
    para("Four tables hold the data: users; contracts, with the grade, the four component scores, the summary, the key obligations and the "
         "combinations as JSON; clauses, with text, page, type, confidence, method and embedding; and flags, with rule, severity, reason, "
         "provision and safer wording. Deleting a user cascades to all of them, and the clause graph is not stored but rebuilt from clauses "
         "and flags on each request. The nine endpoints are health, register, login, upload, list, report, clause detail, graph and the "
         "evaluation metrics; a missing token gives 403, a bad or expired token 401, a wrong file type 400 and an unknown or foreign "
         "contract 404. Cross-origin access is open to all origins, which suits a local demonstration but not deployment.")
    P.h2("G. Evaluation Procedure")
    para("The unmodified evaluation script was run on cached embeddings, so classification numbers come from the shipped code. Extra checks "
         "add contract-disjoint scoring (centroids rebuilt without training spans from the test span's contract), four baselines on the "
         "same split, exact-duplicate counting (four test spans also occur in training; excluding them changes top-1 from 0.792 to 0.797), "
         "and the confidence analysis. Rules were scored per clause and per rule as true and false positives and false negatives. "
         "The corpus run applies the non-API pipeline to every CUAD contract treated as one page. Output screens were captured with "
         "a browser driven by script, on three uploads made through the interface.")

    # ---------------------------------------------------------------- IV
    P.h1("IV. Results")
    P.h2("A. Test Conditions")
    para("Classification and rules were scored on the data of Section III-B on the machine of Section I-A. Embeddings were cached so that "
         "the unmodified evaluation script and extra checks used identical vectors. Latencies are medians of 30 calls to 127.0.0.1 while other "
         "jobs ran on the machine. Table II lists the results.")
    P.table([["Measure", "Result"],
             ["Classification: top-1 / top-3 / top-5 (240 spans)", "0.792 / 0.979 / 0.983"],
             ["Contract-disjoint top-1; chance", "0.800; 0.05"],
             ["Baselines, top-1: TF-IDF + LR; 5-NN; embedding + LR", "0.675; 0.746; 0.800"],
             ["Tie-break triggers; Unclassified labels (of 240)", "0; 0"],
             ["Rules, shipped 40 clauses: P / R / F1", "1.000 / 0.941 / 0.970"],
             ["Rules, probe set 24+16: P / R / F1", "0.786 / 0.458 / 0.579"],
             ["Rules on 5,032 unrelated CUAD spans: any rule fires", "37 (0.74%)"],
             ["Same contract: DOCX / PDF / PDF, spaces collapsed", "8 / 6 / 8 rules of 8"],
             ["510 CUAD contracts, no API: time per contract", "56 ms"],
             ["Contracts with fewer than 5 clauses; flagged clauses", "84 of 510; 2.2%"],
             ["Login (bcrypt) / report / clause detail, median", "254.7 / 9.3 / 5.0 ms"],
             ["Upload to report: fair 8 clauses / CUAD PDF / one-sided", "15.5 / 54.6 / 113.6 s"],
             ["Automated tests", "68 passed, 5 xfailed"]],
            "Measured results", widths=[0.68, 0.32], size=6.6)
    P.h2("B. Clause Classification")
    para("The nearest-centroid classifier reaches 0.792 top-1 and 0.979 top-3 on 20 categories (Fig. 3a), far above chance (0.05) and above "
         "TF-IDF with logistic regression (0.675); logistic regression on the same embeddings scores 0.800, so the shipped rule is as good "
         "as a trained classifier. Rebuilding the centroids without any training span from the same contract gives 0.800, so the split "
         "does not inflate the score. Governing Law, Renewal Term and IP Ownership Assignment are perfect; Non-Compete and Cap On "
         "Liability score 0.50 recall. The main confusions are Uncapped to Capped Liability (4), Non-Compete to No-Solicit of Customers (4) "
         "and the two no-solicit classes with each other (3 and 3).")
    para("The weak categories are the ones whose wording overlaps: No-Solicit Of Customers has recall 0.58 but precision only 0.44 because "
         "Non-Compete and employee no-solicit spans fall into it, Exclusivity, Liquidated Damages and Uncapped Liability score 0.67 recall, "
         "and Cap On Liability, whose spans often mention warranties or damages, is sent to Warranty Duration three times and to Uncapped "
         "Liability three times. Among the simple methods, 1-NN (0.604) and TF-IDF centroids (0.621) are weakest, and 5-NN reaches 0.746; "
         "the gain from averaging 25 vectors per class is therefore real. Classification is deterministic given the cached vectors, so "
         "no seed study applies; the open uncertainty is the 240-span sample, in which one span is 0.4 points.")
    para("Two parts of the rule never act. Top-1 similarity lies between 0.657 and 0.849 (median 0.767), so the 0.60 tie-break "
         "condition and the 0.35 floor are never met (Fig. 4): 152 of 240 spans pass the margin test but none passes the 0.60 test. Every "
         "clause therefore receives one of the 20 labels, including clauses that belong to none. On whole clauses of the three sample "
         "contracts (19 clauses, our manual reading, indicative only) about half of the labels were plausible: a payment clause was labelled "
         "Liquidated Damages and an unlimited-liability clause Cap On Liability, because the 20 categories describe annotated spans, not "
         "whole clauses with headings.")
    P.figure(PAPER / "fig3_results.png", "(a) Top-1 accuracy of the shipped classifier and baselines. (b) Rule scores on the two test sets.", width=3.35)
    P.figure(PAPER / "fig6_thresholds.png", "Similarity of held-out spans to their nearest centroid and margin to the second.", width=3.3)
    P.h2("C. Predatory-Clause Detection")
    para("On the 40 clauses shipped with the rules the detector scores precision 1.00, recall 0.94 and F1 0.97, with two misses. On the "
         "independent probe set the same rules reach precision 0.79 and recall 0.46 (11 of 24 predatory clauses found, 3 false alarms; Fig. 3b). "
         "The misses are paraphrases without the keyword: \"agrees not to accept any engagement from another software house\", \"liability "
         "shall be unlimited\", \"fees once paid cannot be returned\" and \"shall not be entitled to any refund\" (the pattern expects "
         "\"not entitled\"). The false alarms are a fair, capped liquidated-damages clause, a conflict-of-interest bar flagged as restraint of "
         "trade, and mutual liability for gross negligence flagged as unlimited liability. The headline 97% F1 is therefore a measurement on "
         "the author's own phrasing; the Evaluation screen's wording \"held-out\" overstates it.")
    para("Per rule, recall on the probe set is 0.67 for auto-renewal, IP overreach and unilateral termination and 0.33 for the other five "
         "rules, with precision 0.5 for the penalty, restraint and unlimited-liability rules, each of which produced one false alarm. "
         "Because the probe set was written by us after reading the patterns, it is itself biased towards finding their weaknesses, and "
         "three examples per rule is small; the numbers show direction, not a precise rate. Still, the gap between 0.94 and 0.46 recall "
         "is too large to be sampling noise on clauses of this kind.")
    para("A cross-check on CUAD is only a proxy, since its categories are not predatory labels. The restraint rule fires on 21 of 257 "
         "Non-Compete spans (8.2%), the penalty rule on 68 of 121 Liquidated Damages spans (56.2%), the auto-renewal rule on 69 of 208 "
         "Renewal Term spans (33.2%), and the unlimited-liability and IP rules on none of 167 and 317 spans. On 5,032 spans of unrelated "
         "categories any rule fired 37 times (0.74%). Over all 510 contracts the non-API pipeline needs 56 ms per contract, flags "
         "2.2% of 23,869 clauses, finds a risky combination in 11 contracts, and grades them A 4, B 169, C 268, D 68, F 1; 84 contracts "
         "give fewer than five clauses because their headings (for example \"ARTICLE I\") are not recognised.")
    P.h2("D. Input Format and Output Screens")
    para("The same invented freelance contract gives all eight rules as DOCX but six as PDF: PDF text keeps visual line breaks, which the "
         "patterns' bounded gaps of the form .{0,N} cannot cross, and collapsing whitespace first restores all eight. Text before the first numbered heading is "
         "also dropped. Fig. 5 shows the report for the one-sided contract (grade B, score 23.3: wording, not fairness, is graded; 8 of 11 "
         "clauses flagged; three combinations), and Fig. 6 a flagged clause with its reason, quoted provision and safer wording. The fair "
         "contract gives no flag and grade A (13.7); the real CUAD supply agreement grade C (53.8) with one flag. Upload to report took "
         "15.5 s, 54.6 s and 113.6 s, about 12 s per Gemini text call; the front end gives up at 180 s, so a contract with many flags would "
         "fail in the browser although the server would finish.")
    P.figure(PAPER / "fig4_report.png", "Report screen for the one-sided freelance contract.", width=3.35)
    P.figure(PAPER / "fig5_clause.png", "Clause detail: reason, provision and safer wording.", width=3.35)
    para("The complexity score does not track fairness. The one-sided contract is graded B (23.3), the fair one A (13.7), and the real "
         "supply agreement C (53.8) with 43 cross-references. Across the 510 CUAD contracts the median Flesch score is 25.7, legalese density "
         "1.13 per 100 words, cross-references 26 and mean sentence length 28.9 words; the score correlates with contract length at 0.30. "
         "The summaries were fluent but were not checked by a lawyer; for example the summary of the freelance contract says the assignment "
         "covers all of the signer's past and future intellectual property, a paraphrase that goes slightly beyond the clause text.")
    P.h2("E. Testing and Cost")
    para("Sixty-eight tests pass and five are expected failures that document known limits (three keyword-free paraphrases, one fair "
         "capped liquidated-damages clause and malformed Gemini JSON answered with HTTP 400). They stub Gemini with deterministic vectors and JSON and run every API test on a temporary "
         "database. They cover each rule on a textbook clause and on benign clauses, the shipped-set score, the three combinations, the "
         "eight grade boundaries, segmentation, DOCX and PDF reading, the tie-break and its failure fallback, registration and login, "
         "missing, forged, expired and deleted-user tokens, the full upload, cross-user access, wrong file types, a missing key (503) and an "
         "upstream failure (502, with the temporary file removed). There are no browser tests, no load tests and no assertions on the "
         "quality of generated text. Local endpoints answer in 2 to 15 ms except bcrypt login and registration at about 255 to 265 ms; "
         "all measurable cost is in Gemini calls (about 12 s per text call here) and in quota, which on the default model allowed 20 "
         "text requests a day.")
    P.h2("F. Threats to Validity")
    para("First, the shipped test set was written with the rules, so its scores are optimistic; the probe set was written by us and is "
         "small. Second, CUAD contains US filings, not Indian agreements, and its annotations are spans; no Indian corpus and no lawyer "
         "labels were available, so no claim is made about Indian practice beyond the statutes named in the rules. Third, the manual "
         "reading of 19 whole clauses is one reader's judgement. Fourth, latencies were measured on a shared machine and with a remote "
         "service whose speed changed during the study. Fifth, Gemini output is non-deterministic and was not repeated.")
    P.h2("G. Comparison with the Base Paper")
    para("The base paper [@base] reports 98.67% correct answers on 150 legal questions and a qualitative review of 50 contract drafts under "
         "Saudi law. ClauseGuard classifies clauses and flags predatory terms in Indian-law-oriented contracts: the datasets, tasks and metrics "
         "differ, so no head-to-head comparison is valid and none is claimed. We share its Gemini-based, jurisdiction-aware design, but we do "
         "not retrieve from a legal corpus, we add combination logic and a complexity grade, and we report negative results.")

    # ---------------------------------------------------------------- V
    P.h1("V. Conclusion")
    para("ClauseGuard works end to end: it reads contracts, classifies clauses at 0.792 top-1 on CUAD spans, flags eight Indian-law "
         "patterns, reports three risky combinations and an A-F grade, and explains each flag in plain language. "
         "The key finding is that the encouraging 0.97 F1 does not survive independent phrasing (0.58), and that PDF input "
         "and unreachable thresholds weaken the system further. Other limits are that no Indian contract corpus or lawyer review was "
         "available, the explanations are unverified, and free-tier quotas and 12 s calls constrain use.")
    para("Taking the objectives in turn: segmenting and classifying clauses is met, with a heading-based segmenter that fails on some styles "
         "and a classifier at 0.792 top-1; the RAG pipeline with the Claude API and pgvector is not met as stated, since Gemini and SQLite "
         "are used and nothing is retrieved from a legal corpus; the Neo4j knowledge graph is met only in part, by a per-contract graph "
         "of three combination patterns without legal concepts or precedents; predatory flags and combinations under Indian criteria are met "
         "by eight rules and three combinations of limited recall; the complexity score and plain-language summary are met, with unverified "
         "summaries; and the separate evaluation of classification and predatory detection is met and extended.")
    para("Future work: widen the rules with a legal-text classifier trained on labelled predatory clauses and score it on a held-out Indian "
         "set; normalise PDF text and recognise more heading styles; set the thresholds from the new embeddings and add an \"unknown\" class; "
         "retrieve statute text for the explanations and stream partial results.")

    P.references()
    P.d.core_properties.title = K.TITLE
    P.d.core_properties.author = "Batch A6"
    P.d.save(OUT)
    print("saved", OUT, "refs:", len(ORDER), ORDER)
    missing = [k for k in KEYS if k not in ORDER]
    print("uncited keys:", missing)


if __name__ == "__main__":
    build()
