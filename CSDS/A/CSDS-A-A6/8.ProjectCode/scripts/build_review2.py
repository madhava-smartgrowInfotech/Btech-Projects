"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides) for ClauseGuard.

    python scripts/build_review2.py
Every number comes from docs/REPORT_FACTS.md and experiments/**/metrics.json.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import BATCH, CHARTS, COLLEGE, DEPT, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, SHOTS, STUDENTS, SUBTITLE, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"


def px(v):
    return Emu(int(v * 9525))


def cards(s, items, size=15):
    for tid, bid, ttl, body in items:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, size, min_pt=10)


def frame_to(pic, frame):
    frame.left, frame.top, frame.width, frame.height = pic.left, pic.top, pic.width, pic.height


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, "Batch A6", "contract analysis, predatory clauses, CUAD, Gemini embeddings, Indian contract law"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A6"

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), SUBTITLE)
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  Batch A6 (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers the design, methodology and progress of ClauseGuard, a contract complexity and predatory-clause detector for Indian agreements.")

    # 2 contents
    s = S[1]
    set_paras(shape(s, 21), "Model and configuration")
    set_notes(s, "Twelve parts: recap, architecture, methodology, data, preparation, model and configuration, decision logic, data flow and database, technology, progress, results and screens, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Freelancers and small businesses sign contracts they cannot assess. Legal review is costly, and existing tools target US or EU law and read each clause alone.", 16)
    fill(s, 13, "Upload a PDF or DOCX: clauses are classified, 8 Indian-law rules flag predatory terms, 3 combinations are checked, and a complexity grade A-F comes with a plain-language report.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Upload PDF or DOCX", "Segment, embed, classify into 20 CUAD types", "8 rules + 3 risky combinations", "Gemini: reason and safer wording", "Grade A-F, report, clause graph"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1. The base paper [1], SmartLaw Advisor, is a Gemini RAG legal chatbot and contract checker for Saudi law; ClauseGuard keeps the Gemini-based idea for Indian contracts and adds clause-level rules, combinations and a complexity score. Note: the abstract mentions the Claude API, pgvector and Neo4j; the build uses Gemini, SQLite with NumPy, and an in-memory NetworkX graph.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 19 + Vite (port 5106); the dev server proxies /api to the backend.",
                "API: FastAPI (port 8106), 9 REST endpoints; JWT sign-in, bcrypt, owner check on every contract route.",
                "Services: PyMuPDF / python-docx parsing, clause classifier, 8 rules, NetworkX combinations, complexity score.",
                "Data: SQLite (4 tables), cached centroid file, evaluation metrics file.",
                "Only external dependency: the Google Gemini API."], 16, bullet=True)
    set_notes(s, "Everything runs on one machine. The browser talks only to the Vite server, which proxies to FastAPI. Gemini supplies embeddings, tie-break answers, explanations and summaries; there is no vector or graph database.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["Parse PDF / DOCX", "Split into clauses", "Embed + classify", "Apply 8 rules, 3 combos", "Gemini explains flags", "Score complexity A-F", "Store + show report"]):
        fill(s, sid, t, 14, min_pt=9)
    set_paras(shape(s, 25), "Rules decide, Gemini explains")
    fill(s, 26, "Whether a clause is flagged is decided by auditable regex rules; Gemini only writes the reason, the quote and the safer wording.", 15)
    set_paras(shape(s, 28), "Two scores, kept apart")
    fill(s, 29, "Clause classification (CUAD) and predatory detection (curated clauses) are evaluated separately, as objective 6 requires.", 15)
    set_paras(shape(s, 31), "Reproducible")
    fill(s, 32, "Scripts scripts/exp_*.py regenerate every figure into experiments/*/metrics.json.", 15)
    set_notes(s, "Seven-step workflow from upload to report. The two evaluations are separate by design; an extra independent probe set was added because the shipped test set was written together with the rules.")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    ch = replace_chart(s, 4, ["CUAD contracts", "Training spans", "Held-out spans", "Shipped test clauses", "Probe clauses"], [("Items", [510, 500, 240, 40, 40])])
    style_chart(ch, "Items per dataset", ["3F3FBF"], legend=False, label_pt=11, axis_title="contracts / spans / clauses")
    rows = [["Dataset", "Size", "Use"],
            ["CUAD v1 [8] (CC BY 4.0)", "510 contracts, 41 categories", "Corpus run, centroids"],
            ["Training spans", "500 (20 categories x 25)", "Build 20 centroids"],
            ["Held-out spans", "240 (12 per category)", "Classification test"],
            ["Predatory test set", "40 clauses, written by the team", "Rule test (shipped)"],
            ["Probe set", "24 predatory + 16 benign", "Independent rule test"]]
    sz = table_size(rows, 530, 400, [0.34, 0.38, 0.28], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.34, 0.38, 0.28], header_size=sz)
    set_notes(s, "CUAD contracts are US SEC filings; there is no public Indian contract corpus, so the rules (Indian statutes) are tested only on clauses written by us. The train and held-out spans are split by text, not by contract: 192 of 240 held-out spans share a contract with a training span (accuracy is 0.800 when that is removed).")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation")
    set_paras(shape(s, 7), "Parsing")
    fill(s, 8, "PDF text is read page by page with PyMuPDF (real page numbers); a DOCX is read as one page. Scanned images cannot be read.", 15)
    set_paras(shape(s, 12), "Segmentation")
    fill(s, 13, "A page with 2 or more numbered headings (1., 2.1, Section 3) is split at them; otherwise paragraphs are merged to about 400 characters. Fragments under 25 characters are dropped; text before the first heading is lost.", 14)
    set_paras(shape(s, 17), "Features")
    fill(s, 18, "Each clause becomes a 3,072-number Gemini embedding. The complexity score uses Flesch reading ease, legal-jargon density, cross-references and average sentence length.", 15)
    set_notes(s, "Preparation = turning a file into clauses and each clause into a vector. Measured on all 510 CUAD contracts: 444 split on headings and 84 give fewer than 5 clauses because their headings are in other styles.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Models and components")
    put_picture(s, DIAG / "02_high_level_workflow.png", (75, 124, 693, 518), remove_ids=(4,), name="High-level workflow")
    rows = [["Component", "Detail"], ["Embedding", "gemini-embedding-001, 3,072-d"], ["Classifier", "nearest of 20 centroids"], ["Tie-break", "Gemini text model, top 3"], ["Rules", "8 regex rules, Indian statutes"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "No model is trained: the centroids are averages and the text model is called through the API.", 11)
    set_notes(s, "The abstract speaks of a RAG pipeline with the Claude API and pgvector. The implemented pipeline classifies by cosine distance to 20 centroids and prompts Gemini with the matched rule's legal basis; there is no retrieval from a legal corpus.")

    # 9 training configuration
    s = S[8]
    set_paras(shape(s, 2), "Configuration")
    rows = [["Setting", "Value"], ["Training spans", "500 = 20 categories x 25, seed 42"], ["Centroid", "mean of 25 embeddings per category"], ["Held-out set", "240 spans, 12 per category"],
            ["Tie-break trigger", "top-1 minus top-2 < 0.03 and top-1 < 0.60"], ["Unclassified floor", "confidence < 0.35"], ["Complexity weights", "0.35 / 0.30 / 0.20 / 0.15"], ["Text model", "gemini-flash-latest (demo: 3.1-flash-lite)"]]
    set_table(shape(s, 10), rows, 12, header_size=12)
    pic = put_picture(s, CHARTS / "02_classification_baselines.png", (672, 215, 550, 330), remove_ids=(5,), name="Baselines chart")
    set_paras(shape(s, 7), "Top-1 accuracy of the shipped centroid classifier (0.792) against simple baselines on the same 240 spans", 11)
    shape(s, 7).top = Emu(pic.top + pic.height + px(8))
    set_notes(s, "Embedding logistic regression scores 0.800 and 5-NN 0.746; TF-IDF with logistic regression 0.675; chance is 0.05. The shipped centroid method is simple and nearly as good as the best baseline. The text model was changed for the demo run only, because the default model hit its free-tier limit of 20 requests per day.")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (45, 124, 620, 520), remove_ids=(4,), name="Decision flowchart")
    items = [(5, 6, 7, "Classification", "Nearest of 20 centroids by cosine. The tie-break needs top-1 < 0.60, but similarities range 0.657-0.849, so it never fires: every clause gets a label."),
             (8, 9, 10, "Predatory flag", "A regex rule decides (8 rules). Gemini then writes the reason, quotes the provision and drafts safer wording."),
             (11, 12, 13, "Risk and grade", "A combination fires when all its rule ids appear anywhere in the contract (3 patterns). Score = 0.35 readability + 0.30 jargon + 0.20 references + 0.15 sentence length; grade cut-offs 20 / 40 / 60 / 80.")]
    for card, tid, bid, ttl, body in items:
        c = shape(s, card)
        t, b = shape(s, tid), shape(s, bid)
        c.left, c.width = px(690), px(531)
        t.left, t.width = px(714), px(490)
        b.left, b.width = px(714), px(490)
        set_paras(t, ttl)
        t.top, t.height = c.top + px(10), px(40)
        b.top, b.height = c.top + px(50), px(92)
        fill(s, bid, body, 14, min_pt=10)
    set_notes(s, "Left: Algorithm 1 (process_upload) and Algorithm 2 (classify_clause). Right: the three decisions. Complexity measures wording, not fairness: the one-sided freelance sample scored B (23.3) and the fair one A (13.7).")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 340), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 340), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Classify clauses' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(462)
    rows = [["Module", "Responsibility"], ["Parsing and segmentation", "PDF / DOCX text, page numbers, numbered-heading split"],
            ["Classification and rules", "Gemini embedding, 20 centroids, 8 regex rules, 3 combinations"],
            ["Explanation and report", "Gemini reasons, safer wording, summary; complexity grade; JSON report"]]
    set_table(shape(s, 12), rows, 12, widths=[0.25, 0.75], header_size=12, row_h=34)
    shape(s, 12).top = px(492)
    set_notes(s, "Level 1 has seven processes and five data stores; level 2 opens the classification process and shows where the 0.60 and 0.35 tests sit.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 480, 330), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (560, 118, 661, 500), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 4 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(622), px(1164)
    rows = [["Table", "Key fields", "Relations"],
            ["users", "email, password hash", "1:N contracts"],
            ["contracts", "grade, score, summary, combos", "1:N clauses"],
            ["clauses", "text, page, type, embedding", "1:N flags"],
            ["flags", "rule, severity, reason, wording", "N:1 clause"]]
    set_table(shape(s, 13), rows, 11, widths=[0.2, 0.5, 0.3], header_size=11, row_h=30, total_px=480)
    t = shape(s, 13)
    t.left, t.top, t.width = px(57), px(462), px(480)
    set_notes(s, "Four SQLite tables created by SQLAlchemy. The clause graph is not stored: it is rebuilt from clauses and flags on each request. contract_type_guess is never set by the code.")

    # 13 technology
    s = S[12]
    cards(s, [(7, 8, "Language", "Python 3.11, JavaScript and TypeScript"), (12, 13, "AI services", "Google Gemini API (embedding-001 and a flash text model), NumPy cosine search"),
              (17, 18, "Documents and graph", "PyMuPDF 1.24, python-docx 1.1, NetworkX 3.3"),
              (22, 23, "Backend", "FastAPI 0.115, SQLAlchemy 2.0, Pydantic, PyJWT 2.9, bcrypt 4.2"),
              (27, 28, "Frontend", "React 19, Vite, Tailwind CSS 4, TanStack Query, Recharts, axios"),
              (32, 33, "Storage and tests", "SQLite, npz centroid cache, pytest (68 pass, 5 expected fail), Playwright captures")], 15)
    set_notes(s, "Versions are the ones pinned in backend/requirements.txt and frontend/package.json. scikit-learn was used only for the baseline experiments.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Upload, parsing, clause segmentation", "Completed"], ["Classification into 20 CUAD types (79.2 % top-1)", "Completed"],
            ["8 Indian-law rules and 3 risky combinations", "Completed"], ["Gemini reasons, safer wording, summary", "Completed"],
            ["Complexity score and A-F grade", "Completed"], ["Report, clause, graph, evaluation screens", "Completed"],
            ["Experiments, 73 tests, latency measurements", "Completed"], ["Claude / pgvector / Neo4j, Indian corpus, lawyer review", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.7, 0.3], header_size=13)
    for vid, lid, v, lab in [(6, 7, "9", "REST endpoints"), (9, 10, "4", "database tables"), (12, 13, "68", "tests pass; 5 expected failures")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Every feature of the build is implemented and tested. The open items are the technologies named in the abstract that were replaced (Claude, pgvector, Neo4j), an Indian contract corpus, and expert review of the explanations.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: classification and rules")
    ch = replace_chart(s, 4, ["CUAD top-1", "CUAD top-3", "Rules precision (shipped)", "Rules recall (shipped)", "Rules precision (probe)", "Rules recall (probe)"],
                       [("Score", [0.792, 0.979, 1.000, 0.941, 0.786, 0.458])], number_format="0.000")
    style_chart(ch, "ClauseGuard scores: 240 held-out CUAD spans and two sets of rule test clauses", ["3F3FBF"], number_format="0.000", legend=False, label_pt=11, axis_title="score")
    ch.value_axis.maximum_scale = 1.1
    for vid, lid, v, lab in [(6, 7, "79.2 %", "CUAD top-1 (top-3 97.9 %), 240 spans"), (9, 10, "1.00 / 0.94", "rule precision / recall, shipped 40 clauses"), (12, 13, "0.79 / 0.46", "rule precision / recall, independent probe")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Classification: top-1 0.792 (190 of 240), top-3 0.979; weakest categories Non-Compete and Cap On Liability (recall 0.50). Rules: the shipped 40 clauses were written with the rules (P 1.000, R 0.941, F1 0.970). On the independent probe set (24 predatory paraphrases, 16 benign look-alikes) precision is 0.786, recall 0.458, F1 0.579: the regexes miss paraphrases and flag some fair clauses. These are measured values from experiments/*/metrics.json.")

    # 16 output screens
    s = S[15]
    p1 = put_picture(s, SHOTS / "08_report_freelance_top.png", (45, 128, 590, 476), remove_ids=(4,), name="Contract report screenshot")
    p2 = put_picture(s, SHOTS / "11_clause_detail_termination.png", (645, 128, 590, 476), remove_ids=(7,), name="Clause detail screenshot")
    frame_to(p1, shape(s, 5))
    frame_to(p2, shape(s, 8))
    set_paras(shape(s, 6), "Report of the one-sided freelance contract: grade B (23.3), 8 flagged clauses, 3 combinations", 11)
    set_paras(shape(s, 9), "Clause 2: Unilateral Termination, reason, quoted provision, safer wording", 11)
    for cap_id, p in ((6, p1), (9, p2)):
        shape(s, cap_id).top = Emu(max(p.top + p.height + px(12), px(500)))
    set_notes(s, "Real captures of the running application. The contract is an invented sample written for this review; the explanation text is produced by the Gemini text model.")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Rules, combinations, graph", "26", "22 passed, 4 expected fail"], ["Complexity, parsing, vector store", "21", "21 passed"],
            ["Classification (Gemini stubbed)", "6", "6 passed"], ["API: auth, upload, ownership, errors", "20", "19 passed, 1 expected fail"], ["Total", "73", "68 passed, 5 expected fail"]]
    set_table(shape(s, 18), rows, 12, widths=[0.5, 0.12, 0.38], header_size=12)
    fill(s, 5, ["68 tests pass in about 10 s; the 5 expected failures record known limits (paraphrases the rules miss, a fair clause that is flagged, malformed Gemini JSON answered 400).",
                "Independent probe set: rule precision 0.79, recall 0.46, against 1.00 / 0.94 on the set written with the rules.",
                "The same contract as DOCX fires 8 of 8 rules, as PDF 6 of 8: line breaks defeat the patterns.",
                "API medians: report 9.3 ms, clause 5.0 ms, login 254.7 ms; upload to report 15.5 s (0 flags) to 113.6 s (8 flags).",
                "Not done: UI automation, load and security tests, lawyer review."], 14, bullet=True)
    put_picture(s, CHARTS / "03_rules_two_test_sets.png", (57, 480, 300, 190), name="Rules on two test sets")
    put_picture(s, CHARTS / "10_format_parity.png", (385, 480, 270, 190), name="DOCX versus PDF parity")
    set_notes(s, "Expected-fail tests are reported as such rather than hidden. Latencies were measured on 127.0.0.1 on an i5-12500 with other jobs running; upload time is dominated by about 12 s per Gemini text call.")

    # 18 challenges
    s = S[17]
    cards(s, [(7, 8, "Abstract versus build", "Claude, pgvector and Neo4j became Gemini, SQLite with NumPy and in-memory NetworkX; no legal-corpus retrieval."),
              (12, 13, "Rules miss paraphrases", "Independent recall 0.46 and precision 0.79, against 1.00 / 0.94 on our own set."),
              (17, 18, "Thresholds never fire", "Similarities run 0.657-0.849, so tie-break (< 0.60) and 'Unclassified' (< 0.35) never act."),
              (22, 23, "PDF line breaks", "8 of 8 rules on DOCX, 6 of 8 on the same PDF; text before the first heading is dropped."),
              (27, 28, "Gemini quota and speed", "Default model allows 20 requests a day; 113.6 s for 8 flags; the UI times out at 180 s."),
              (32, 33, "No Indian contracts", "CUAD holds US filings; Indian-law rules are tested on our clauses only; no lawyer review.")], 14)
    set_notes(s, "Six challenges and findings, each backed by a measurement in docs/REPORT_FACTS.md.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["Fix rules and parsing", "Recalibrate classifier", "Indian data and review", "Final report and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Normalise whitespace before matching, keep text before the first heading, widen the rules, and test them on clauses written by someone other than the rule author.",
                 "Set the tie-break and 'Unclassified' thresholds for the embedding model's similarity range; split train and test by contract.",
                 "Collect Indian contracts with lawyer-reviewed labels; show the honest scores on the Evaluation screen.",
                 "Run user testing with signatories (not yet done); complete the report and final demo."], 16, bullet=True)
    set_notes(s, "What remains: repair the weaknesses found in testing, collect Indian-law data and expert review, then the report and the final demonstration.")

    set_notes(S[19], "Thank you. Questions are welcome.")

    # template leftovers: icon alt text, shape names, theme name
    for idx in (2, 6, 12, 17):
        for sh in S[idx].shapes:
            if sh.shape_type == 13:
                sh._element.nvPicPr.cNvPr.set("descr", "icon")
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    for m in prs.slide_masters:
        th = m.part.part_related_by(RT.THEME)
        th._blob = th.blob.replace(b"Hidden Risk", PRODUCT.encode())
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
