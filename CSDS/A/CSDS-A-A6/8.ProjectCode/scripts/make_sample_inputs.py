"""Writes the sample contracts used for the documentation screenshots into docs/sample_inputs/.

    venv\\Scripts\\python.exe scripts\\make_sample_inputs.py

* freelance_services_agreement.docx / .pdf - an invented one-sided freelance contract (the PDF is 2 pages) written
                                       to contain all eight predatory patterns; the two files hold identical text
* balanced_consulting_agreement.docx - an invented fair consulting agreement (should raise no, or almost no, flags)
* supply_agreement_cuad.pdf          - a real contract from CUAD v1 (Fusion Pharmaceuticals supply agreement,
                                       a public SEC exhibit), typeset to PDF unchanged, so the screenshots
                                       also show the tool on a real-world document
All company names in the two invented contracts are fictitious.
"""
from __future__ import annotations

import json
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "sample_inputs"

FREELANCE = [
    ("FREELANCE SERVICES AGREEMENT", None),
    ("", "This Agreement is made at Hyderabad on 4 May 2026 between Northbridge Digital Solutions Private Limited (the \"Company\") and Ananya Rao, an independent software developer (the \"Consultant\")."),
    ("1. Scope of Services", "The Consultant shall design, build and maintain the Company's customer portal and any related software as described in Schedule A and as the Company may reasonably request from time to time."),
    ("2. Term and Termination", "This Agreement begins on 4 May 2026 for a period of twelve months. The Company may terminate this Agreement at any time, without cause or notice, at its sole discretion, and the Consultant shall have no right to object or to claim compensation for unfinished work."),
    ("3. Fees and Payment", "The Company shall pay the Consultant a monthly fee of Rs. 60,000 against invoices raised on the last working day of each month. All payments made under this Agreement are non-refundable under any circumstances, including early termination by the Company, and the Consultant shall not be entitled to any refund of expenses incurred."),
    ("4. Delay and Penalty", "If the Consultant delays delivery of any milestone, the Consultant shall pay a penalty of Rs. 25,000 per day of delay, irrespective of the actual loss suffered by the Company, and the Company may deduct the penalty from any amount due."),
    ("5. Exclusivity and Non-Compete", "During the term of this Agreement and for a period of three years thereafter, the Consultant shall not, directly or indirectly, engage in any competing business or provide services to any competitor of the Company anywhere in India."),
    ("6. Liability", "The Consultant shall be liable for any and all losses, damages, claims and expenses arising in connection with this Agreement, with no limitation on liability whatsoever."),
    ("7. Intellectual Property", "All intellectual property created by the Consultant, whether related to this engagement or otherwise, including prior and pre-existing works, shall be assigned to the Company irrevocably and worldwide, and the Consultant waives all moral rights."),
    ("8. Renewal", "This Agreement shall automatically renew for successive terms of twelve months each unless the Consultant gives written notice of non-renewal at least ninety days before the end of the then-current term."),
    ("9. Dispute Resolution", "Any dispute arising out of this Agreement shall be referred to arbitration at Hyderabad before a sole arbitrator solely appointed by the Company, and the Consultant waives the right to approach any court except to enforce the award."),
    ("10. Governing Law", "This Agreement shall be governed by and construed in accordance with the laws of India, and the courts at Hyderabad shall have exclusive jurisdiction."),
    ("11. Confidentiality", "Each party shall keep confidential all non-public information of the other party disclosed in connection with this Agreement and shall not use it for any purpose other than performing this Agreement."),
]

BALANCED = [
    ("CONSULTING AGREEMENT", None),
    ("", "This Agreement is made on 12 June 2026 between Lakeview Analytics LLP (the \"Client\") and Rohan Mehta (the \"Consultant\")."),
    ("1. Services", "The Consultant will prepare a quarterly sales report and a dashboard for the Client, as listed in Schedule A. Changes to the list need the written agreement of both parties."),
    ("2. Fees", "The Client will pay Rs. 45,000 per month within fifteen days of receiving a valid invoice. Fees for work that has not been carried out will be returned if this Agreement ends early."),
    ("3. Term and Termination", "This Agreement runs for six months. Either party may end it earlier by giving the other thirty days' written notice. If either party breaches a material term and does not fix the breach within fifteen days of written notice, the other party may end the Agreement at once."),
    ("4. Ownership of Work", "The Client owns the reports and dashboards made for it under this Agreement once they are paid for. The Consultant keeps the ownership of tools, templates and methods that existed before this Agreement and gives the Client a licence to use them with the deliverables."),
    ("5. Liability", "Each party's total liability under this Agreement is limited to the fees paid or payable in the six months before the claim. Neither party is liable for indirect loss."),
    ("6. Confidentiality", "Each party will keep the other's confidential information private and use it only for this Agreement. This duty continues for two years after the Agreement ends."),
    ("7. Disputes", "The parties will first try to settle any dispute by discussion within thirty days. If that fails, the dispute will go to a single arbitrator chosen jointly by both parties under the Arbitration and Conciliation Act, 1996, with the seat at Hyderabad."),
    ("8. Governing Law", "This Agreement is governed by the laws of India."),
]


def write_pdf(path: Path, blocks, font="tiro", size=10.5, per_page_lines=None):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    y, left, right, bottom = 70, 62, 533, 780
    width = right - left

    def need(h):
        nonlocal page, y
        if y + h > bottom:
            page = doc.new_page(width=595, height=842)
            y = 70

    for head, body in blocks:
        if body is None:  # title
            rect = fitz.Rect(left, y, right, y + 30)
            page.insert_textbox(rect, head, fontsize=15, fontname="tibo", align=1)
            y += 38
            continue
        lines = max(1, int(len(body) * (size * 0.47) / width) + 1)
        h = lines * (size * 1.32) + 8
        if head:
            need(24 + h)
            page.insert_text((left, y + 11), head, fontsize=size + 1, fontname="tibo")
            y += 17
        need(h)
        page.insert_textbox(fitz.Rect(left, y, right, y + h), body, fontsize=size, fontname=font, align=3)
        y += h + 6
    for i, p in enumerate(doc, 1):
        p.insert_textbox(fitz.Rect(0, 805, 595, 825), f"Page {i} of {len(doc)}", fontsize=8.5, fontname="tiro", align=1)
    doc.save(str(path))
    doc.close()


def write_docx(path: Path, blocks):
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(11)
    for head, body in blocks:
        if body is None:
            d.add_heading(head, level=1)
            continue
        if head:
            p = d.add_paragraph()
            p.add_run(head).bold = True
        d.add_paragraph(body)
    d.save(str(path))


def write_cuad_pdf(path: Path):
    data = json.loads((ROOT / "data/cuad/CUADv1.json").read_text(encoding="utf-8"))["data"]
    doc = next(x for x in data if x["title"].startswith("FUSIONPHARMACEUTICALSINC_06_05_2020-EX-10.17"))
    text = doc["paragraphs"][0]["context"]
    pdf = fitz.open()
    paras = [p.strip() for p in text.replace("\r", "").split("\n\n") if p.strip()]
    page = pdf.new_page(width=595, height=842)
    y, left, right, bottom, size = 62, 56, 539, 790, 9.5
    for p in paras:
        p = " ".join(p.split())
        lines = int(len(p) * (size * 0.455) / (right - left)) + 1
        h = lines * size * 1.3 + 4
        if y + h > bottom:
            page = pdf.new_page(width=595, height=842)
            y = 62
        page.insert_textbox(fitz.Rect(left, y, right, y + h + 4), p, fontsize=size, fontname="tiro", align=0)
        y += h + 5
    for i, pg in enumerate(pdf, 1):
        pg.insert_textbox(fitz.Rect(0, 810, 595, 830), f"- {i} -", fontsize=8.5, fontname="tiro", align=1)
    pdf.save(str(path))
    n = len(pdf)
    pdf.close()
    return doc["title"], n


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    write_pdf(OUT / "freelance_services_agreement.pdf", FREELANCE)
    write_docx(OUT / "freelance_services_agreement.docx", FREELANCE)
    write_docx(OUT / "balanced_consulting_agreement.docx", BALANCED)
    title, n = write_cuad_pdf(OUT / "supply_agreement_cuad.pdf")
    for f in sorted(OUT.iterdir()):
        print(f.name, f.stat().st_size)
    print("CUAD source:", title, "pages:", n)


if __name__ == "__main__":
    main()
