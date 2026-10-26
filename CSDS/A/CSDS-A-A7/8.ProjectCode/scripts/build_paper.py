"""Build 19.IEEE_Paper.docx: a six-page, two-column IEEE conference paper on SHEGUARD.

    python scripts/make_paper_figures.py ; python scripts/build_paper.py ; python scripts/paper_finish.py

Every number comes from docs/REPORT_FACTS.md and experiments/study/*.json. References are numbered in order of first citation
from doc_common.REFERENCES; the three batch-survey papers that could not be retrieved (entries 3-5) are not cited here.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paper_kit as K  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from doc_common import BATCH, GUIDE, GUIDE_DESIG, REFERENCES, ROOT, STUDENTS  # noqa: E402

OUT = BATCH / "19.IEEE_Paper.docx"
FIG = ROOT / "docs" / "charts"

# key -> index into doc_common.REFERENCES
KEYS = {"base": 0, "farooq": 1, "thawale": 5, "shinde": 6, "extranom": 7, "macqueen": 8, "rousseeuw": 10, "db": 11, "ch": 12, "em": 13,
        "ward": 14, "sklearn": 15, "osrm": 16, "osm": 17, "jwt": 18, "websocket": 19, "sha": 21, "kaggle": 22, "speech": 24, "gemini": 25}
ORDER: list[str] = []


def C(text: str) -> str:
    """Replace [@a,b] by [n], [m] numbered in order of first citation."""
    def sub(m):
        out = []
        for k in m.group(1).split(","):
            if k not in ORDER:
                ORDER.append(k)
            out.append(f"[{ORDER.index(k) + 1}]")
        return ", ".join(out)
    return re.sub(r"\[@([a-z,]+)\]", sub, text)


class Paper(K.Paper):
    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(K.TITLE)
        K.font(r, 24)
        p.paragraph_format.space_after = Pt(6)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [(GUIDE, f"{GUIDE_DESIG}, Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Dept. of CSE (Data Science)", False), ("[COLLEGE NAME]", False), ("[CITY]", False), ("[EMAIL]", False)):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line)
                K.font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        sectpr = sec._sectPr
        cols = sectpr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sectpr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def references(self):
        self.h1("References")
        for i, key in enumerate(ORDER, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]\t{REFERENCES[KEYS[key]]}")
            K.font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


K.TITLE = "SHEGUARD: AI Powered Women's Safety Platform"
K.ABSTRACT = (
    "Most women's safety applications wait for the user to press an SOS button, choose routes for speed rather than safety and keep no "
    "trustworthy record of an incident. This paper presents SHEGUARD, a web application that combines a district-level area-risk model, "
    "risk-scored route alternatives, a spoken SafePhrase and one-tap SOS, live location streaming to guardians over a WebSocket, "
    "SHA-256-hashed photo and audio evidence, and a Gemini-based assistant. Districts are tiered by K-Means on one feature, the average "
    "crimes against women per reporting year (467 districts, 2001-2012), giving a silhouette of 0.580 and a Davies-Bouldin index of "
    "0.531. The tiers reduce to two cut-offs, the night factor never changes which route is recommended, and 7 of 40 city centres lie "
    "more than 100 km from any district in the table. A guardian sees a new position about 15 ms after the push, and 51 automated "
    "tests pass while six documented defects fail as expected. Alert delivery and the assistant could not be demonstrated in this run."
)
K.KEYWORDS = "area-risk clustering, evidence integrity, live location tracking, safe route, SafePhrase, SOS, women's safety"


def build():
    P = Paper()
    P.front()
    P.abstract()
    para = lambda t, **k: P.para(C(t), **k)

    # ---------------------------------------------------------------- I
    P.h1("I. Introduction")
    para("Harassment, stalking and assault can happen on an isolated street as easily as in a crowd, and a frightened user often cannot "
         "open an app, find a button and make several choices. Today most safety software sends one alert when a button is pressed, "
         "map applications choose the fastest route, and wearable devices need extra hardware and a trigger [@farooq,thawale]. The cost "
         "is late or no warning, no preventive advice before a journey, guardians who receive a single message instead of a moving "
         "position, and no evidence that can later be shown to be unchanged. Work on safe routes does exist [@base,shinde], but it "
         "stops at the map and does not cover what happens when the user is in danger.")
    para("This paper reports SHEGUARD, a web application that covers the sequence before, during and after an emergency. Its contributions are:")
    P.numbered([
        "an area-risk model that tiers 467 Indian districts by K-Means, and a route scorer that ranks OSRM alternatives against it;",
        "two quick triggers, a spoken SafePhrase matched by word share and a one-tap SOS with a hold-to-cancel control;",
        "live location for guardians through an unguessable link and a WebSocket, with SHA-256-hashed photo and audio evidence;",
        "an honest evaluation: clustering indices, 16 real routes, a 57-test suite, latency and WebSocket delay, and the defects found.",
    ])
    para("SHEGUARD is built as a browser application, installable as a progressive web app, rather than a native phone app. This keeps "
         "one code base, lets guardians open a plain link with no installation or account, and allows evaluation on a PC; the price is "
         "that voice activation depends on the speech recogniser of an open Chrome or Edge tab and that the page cannot listen with the "
         "screen off. The platform is therefore a working prototype whose limits are stated alongside its measurements.")
    para("Section II reviews related work, Section III describes the method, Section IV reports the results and Section V concludes.")
    P.h2("A. System Hardware")
    para("All measurements were taken on one Windows 11 PC with an Intel Core i5-12500 (6 cores, 12 threads), 15.7 GB of RAM and no GPU; "
         "the model trains in well under a second on the CPU.")
    P.h2("B. Software Requirements")
    para("Python 3.11.9 with FastAPI 0.115.6, SQLAlchemy 2.0.36, scikit-learn 1.5.2 [@sklearn], pandas 2.2.3, PyJWT 2.10.1 and "
         "bcrypt 4.2.1 on SQLite; Node 24.19.0 with React 19, Vite 6, Tailwind 4 and Leaflet 1.9.4 in Chrome. External services: the public "
         "OSRM server [@osrm], OpenStreetMap tiles and Nominatim [@osm], and the Gemini API [@gemini].")

    # ---------------------------------------------------------------- II
    P.h1("II. Literature Survey")
    para("The base paper, SafeRoutes [@base], clusters regions by women's safety using data from data.gov.in: district population (2001 "
         "and 2011), IPC crimes against women (2013), literacy, per-capita expenditure, police strength and forest cover. After "
         "imputation, scaling and Pearson and Spearman correlation, it applies K-Means [@macqueen] and Gaussian mixtures [@em] and "
         "finds that three to four clusters suit the data; a cluster is scored against dummy data from cities regarded as safest, and "
         "a ride-hailing integration with a 1.96x safety buffer raises alerts on deviation. Its ablation tables report, for k = 3, a "
         "silhouette [@rousseeuw] of 0.416, a Calinski-Harabasz index [@ch] of 221.0 and a Davies-Bouldin index [@db] of 0.970, "
         "with the best Davies-Bouldin value among k = 2 to 4. The results section is stated as anticipated outcomes: no accuracy is "
         "given and the cab integration is designed, not evaluated.")
    para("Hardware-oriented work is surveyed in [@farooq], which reviews IoT devices from 2016 to 2022 and finds pulse and pressure "
         "sensors, GPS, GSM and Raspberry Pi most common and calls for alerts that need less human action. An Android application with an "
         "ESP32 camera reports 93% accuracy on a risk-zone classification of predefined locations and SOS delivery in 4-6 s [@thawale], "
         "on its own test set. A route recommender built on OpenStreetMap and OSRM avoids crime hotspots and adds a one-touch SOS, and "
         "reports only qualitative results [@shinde]. Women-centred video anomaly detection is addressed by a benchmark of 1,001 videos in "
         "five crime categories [@extranom]. Spoken triggers rely on the browser's speech recogniser [@speech]. Table I compares the works.")
    T1 = ([["Work", "Approach", "Data", "Reported result", "Limitation"],
             ["[@base]", "K-Means, GMM clusters; cab alerts", "data.gov.in, 2013 IPC", "Cluster tables; no accuracy", "Cab part not evaluated"],
             ["[@farooq]", "Review of IoT devices", "Studies 2016-2022", "Taxonomy; sensors, GPS/GSM", "No system built"],
             ["[@thawale]", "ESP32-CAM app, crime zones", "Own test locations", "93% zone accuracy; SOS 4-6 s", "Own test set; hardware"],
             ["[@shinde]", "OSM/OSRM safe routes, SOS", "Crime reports, user input", "Qualitative only", "No measurements"],
             ["[@extranom]", "Video anomaly benchmark", "1,001 videos", "Benchmark, 5 categories", "Video, not routes"],
             ["This work", "K-Means tiers, scorer, SOS", "467 districts, 2001-2012", "Section IV", "One feature; old data"]])
    T1 = [[C(x) for x in row] for row in T1]
    P.table(T1, "Comparison of related work", [0.12, 0.26, 0.21, 0.24, 0.17], size=6.5)
    para("The methods used here are standard. K-Means [@macqueen] partitions points by squared distance; cluster quality is judged "
         "with the silhouette [@rousseeuw] and Davies-Bouldin [@db] indices, which the base paper also uses; hierarchical clustering "
         "with Ward linkage [@ward] and mixture models fitted by expectation maximisation [@em] are the usual alternatives and are "
         "compared with K-Means in Section IV. Routing relies on OSRM [@osrm] over OpenStreetMap [@osm]. None of the surveyed route "
         "works reports a check of how the risk data cover the places where people travel, which is one of the questions asked below.")
    para("Read capability by capability, the surveyed works split the problem: [@base] and [@shinde] address where to travel, [@thawale] "
         "and [@farooq] address how to raise an alert with extra hardware, and [@extranom] addresses detecting violence in video. None "
         "of them combines a spoken trigger, a live link that a guardian can open without an account, and evidence that can be "
         "checked afterwards, and only [@thawale] reports timings for the alert itself.")
    para("The gap is a platform that joins route prevention, hands-free activation, live guardian tracking and tamper-evident evidence "
         "in one running system and measures it, while saying where the evidence is weak.")

    # ---------------------------------------------------------------- III
    P.h1("III. Methodology")
    P.h2("A. Architecture and Main Flow")
    para("A React progressive web app talks to a FastAPI server over REST and one WebSocket route. Ten routers handle "
         "authentication, guardians, risk areas, route planning, place search, SOS sessions, evidence, the public tracking view and the "
         "assistant; six SQLite tables store users, guardians, sessions, location pings, evidence and the district table (Fig. 1). "
         "Every query is filtered by the user id in a JWT [@jwt]; guardians have no account and see a session only through a random "
         "32-character share token. The main flow is shown in Fig. 2.")
    para("The API is small: registration and login return a bearer token; guardians have create, read, update and delete calls; "
         "one public call lists the 467 risk areas; the route call takes an origin and a destination; and the SOS family has start, "
         "location, cancel, read, evidence upload and evidence list calls. Missing or invalid credentials return 401, a session or "
         "guardian that belongs to someone else returns 404, an unavailable routing or assistant service returns 502, and "
         "malformed input returns 422.")
    P.figure(FIG / "p_arch.png", "System architecture of SHEGUARD.", width=3.4)
    P.figure(FIG / "p_flow.png", "Main flow before, during and after an emergency.", width=3.4)
    P.h2("B. Data and Area-Risk Model")
    para("The district file [@kaggle] has 9,017 rows (8,597 after dropping TOTAL roll-ups), 35 states and territories, 827 district "
         "pairs and seven counted categories for 2001-2012. A district's feature is its crimes per reporting year,")
    P.eq("r_{d} = (1 / Y_{d}) Σ_{y} Σ_{c=1..7} n_{d,y,c}", 1)
    para("where the number of years reported is at most twelve (638 of the 827 districts reported all twelve). Nominatim geocoded 467 of the 827 pairs (56.5%); the "
         "rest were dropped, not guessed. The feature is standardised and K-Means with k = 3, random_state 42 and ten initialisations "
         "minimises the within-cluster sum of squares. Clusters are named low, medium and high by their mean rate, and a 0-100 risk score is the "
         "min-max scaled rate. There is no train/test split because there is no ground truth. With one feature, the clusters are "
         "intervals: the tiers are exactly the rules r < 213.75, r < 529.83, and the model is a data-driven pair of cut-offs.")
    para("The tier boundaries found by K-Means were 0-212.0 crimes a year for low, 215.5-522.0 for medium and 537.7-1191.6 for high. "
         "Because the model is fitted on all 467 points and no label exists, it cannot be scored for accuracy; instead the experiments in "
         "Section IV vary k, the random initialisation and the clustering method and report how much the tiers move.")
    P.h2("C. Route Scoring")
    para("For a trip, OSRM [@osrm] returns up to a few alternatives. Every eighth polyline point is mapped to the nearest district by haversine "
         "distance and given a tier weight w (1, 2.5 or 5); the route score is the mean,")
    P.eq("S = m · (1 / N) Σ_{i=1..N} w(tier(nearest(p_{i})))", 2)
    para("where m = 1.6 between about 20:30 and 05:30 IST and 1 otherwise. The routes are sorted by S and the lowest is recommended; "
         "the interface labels a route HIGH when at least 20% of its samples are in high-risk districts.")
    para("Two properties follow from Eq. (2). The score depends only on the share of samples in each tier, so a long route and a short "
         "one through the same districts score alike, and distance enters only through the number of samples. Because m is the same "
         "for every route in a request, the ranking is determined by the tier weights alone. The nearest-district lookup compares each "
         "sample with all 467 districts and has no distance limit.")
    P.h2("D. SafePhrase and SOS")
    para("The Web Speech API runs in continuous mode. For a stored phrase with words T, the matcher slides a window of |T| spoken words over the transcript "
         "and fires when")
    P.eq("max_{window} |T ∩ window| / |T| ≥ θ", 3)
    para("with θ = 1, 0.85, 0.75 or 0.6 for the Low, Medium (default), High and Ultra settings. An exact substring always fires. Phrases stay in "
         "the browser. SOS starts from the phrase or one tap: the server creates a session with a share token and tells each guardian by "
         "Telegram or e-mail; the phone then pushes its position about every 4 s and the server forwards it to guardians "
         "over the WebSocket [@websocket]. Photos and audio are uploaded, hashed server-side with SHA-256 [@sha], stored with time and "
         "location, and listed with their hash. Cancelling needs a three-second hold.")

    P.h2("E. Alerts, Evidence and Access Control")
    para("When a session starts, the server loops over the user's guardians and, for each, sends a Telegram message if a chat "
         "identifier is stored and an e-mail through Gmail SMTP if an address is stored; both carry the tracking link, and the reply "
         "reports per guardian which channel succeeded so that the screen can say so. Delivery is best effort and its failure never "
         "blocks the session. The public tracking endpoint returns the owner's name, the status, the last position and the whole path; "
         "the WebSocket group is keyed by the share token. Passwords are hashed with bcrypt and tokens expire after seven days. "
         "Evidence is stored under a file name made from the first 16 hex digits of its hash, with the full hash, kind, time and "
         "position in the database; a hash is recomputed only if someone chooses to check it, because the product has no verification endpoint.")
    P.h2("F. Assistant and Implementation Notes")
    para("The assistant sends a fixed system prompt (safety scope; the first line of an emergency answer must point to SOS or the "
         "emergency number; fewer than 150 words), the question and, if permitted, the user's coordinates to the Gemini API [@gemini] "
         "at temperature 0.4. It keeps no conversation memory. Provider errors map to HTTP 502 and the interface then shows a plain "
         "unavailable message. The application is installable as a PWA with a service worker that caches the shell (six precached "
         "files, 481 KiB); the production bundle is 450 kB of JavaScript (136 kB compressed). The whole Python side is about 2,300 lines "
         "and the TypeScript side about 2,950 lines in 26 files.")

    para("The word-share rule was chosen because it needs no training data, runs in the browser in about 1.5 microseconds per "
         "transcript and tolerates the dropped and substituted words that continuous recognition produces, at the price of an adjustable "
         "trade-off between missed phrases and accidental alerts. The hold-to-cancel control is the counterpart for the other direction: "
         "a stray tap cannot end an alert, because the button must be held for three seconds.")

    # ---------------------------------------------------------------- IV
    P.h1("IV. Results")
    P.h2("A. Test Conditions")
    para("Everything ran on the machine of Section I-A, with the API on 127.0.0.1 (localhost added about 2 s per call here), one sequential "
         "client, 5 to 30 calls per measurement, and the public OSRM and Nominatim servers. Telegram and Gmail credentials were empty, "
         "so no alert was sent, and the Gemini key's quota was exhausted during the run. Table II summarises the measurements.")
    P.table([["Measure", "Result"],
             ["Silhouette / Davies-Bouldin / Calinski-Harabasz, k = 3", "0.580 / 0.531 / 1155.4"],
             ["Tiers low / medium / high (districts)", "278 / 148 / 41"],
             ["K-Means vs two thresholds (213.75, 529.83)", "100% agreement"],
             ["k = 2 and k = 4: silhouette; Davies-Bouldin", "0.662; 0.551 and 0.596; 0.493"],
             ["30 random initialisations: ARI min / mean", "0.659 / 0.790"],
             ["Ward / Gaussian mixture agreement with K-Means", "83.3% / 72.6%"],
             ["Of 16 OSRM pairs: with 2 alternatives; day order = night order", "10; 16"],
             ["Nearest district, 40 cities: median / max (km)", "8.1 / 231.3"],
             ["SafePhrase Medium: recall (24) / false triggers (80)", "54.2% / 0%"],
             ["SafePhrase Ultra: recall / false triggers", "100% / 10.0%"],
             ["Median latency: location push; track page; login", "14.7; 8.3; 279.0 ms"],
             ["Live route plan (OSRM + scoring), mean of 5", "1.63 s"],
             ["WebSocket push, mean: 1 viewer; 5 viewers", "15.4 ms; 14.6 ms"],
             ["Automated tests", "51 pass; 6 expected fail"],
             ["Assistant: probe prompts answered", "0 of 10 (HTTP 429)"]],
            "Measured results", [0.64, 0.36], size=6.8)
    para("The test suite has 45 API-level tests, of which 39 pass and six are marked as expected failures, and 12 unit tests of the "
         "scorer and the model files, all of which pass. The API tests run against a throw-away database with the routing, assistant, "
         "Telegram and Gmail calls replaced by stubs, so they are offline and repeatable; the live measurements in this section "
         "use the real server and the real routing service.")
    P.h2("B. Area-Risk Model")
    para("The re-run reproduced the shipped metrics exactly. The three tiers hold 278, 148 and 41 districts with mean rates of 97.8, "
         "327.9 and 720.9 crimes a year. k = 3 is a design choice for three interface tiers and not the optimum: k = 2 has the highest "
         "silhouette and k = 4 beats k = 3 on both indices (Fig. 3a). Thirty single-initialisation runs agree with the shipped "
         "assignment with an adjusted Rand index of at least 0.659 (mean 0.790), and other tiering methods (Ward [@ward], a Gaussian mixture, fixed cut-offs, terciles) agree on only 49% to 83% of "
         "districts, so the 41-district high tier is method dependent. The score counts raw cases, not rates per population, and 40.0% of "
         "the counted cases are cruelty by husband or relatives. The highest rates are West Godavari, Vijayawada, Guntur and Jaipur. "
         "The indices are far from the base paper's 0.970 and 221.0, but the data, feature count and number of districts differ, so no comparison "
         "is claimed and none is valid.")
    P.figure(FIG / "p_results.png", "Cluster-quality indices against k (a), SafePhrase recall and false triggers by sensitivity (b), and distance from city centres to the nearest district (c).", width=3.4)
    para("The 0-100 risk score that accompanies each tier is the district rate scaled between the lowest and highest rate in the table "
         "(0 to 1191.6 crimes a year), so a district at 600 crimes a year scores about 50 whichever tier it falls in; the score is "
         "returned by the API but used neither by the map, which colours districts by tier, nor by the route scorer.")
    P.h2("C. Route Scoring")
    para("Of 16 real city pairs, 10 returned two alternatives and six only one, so for more than a third of trips there was nothing to choose. The "
         "recommended route was also the fastest in 11 pairs; where alternatives existed the safest pick cost 3.7 min more on average. "
         "The night factor multiplies every route by 1.6 (a day score of 2.5 becomes 4.0), so the order was identical by day and "
         "night in all 16 pairs: it changes the displayed number and label only. Changing the tier weights to 1/2/3 or 1/3/9 did "
         "not change the recommendation in any multi-route pair. Coverage is the main problem: no district in Uttar Pradesh, West "
         "Bengal, Delhi or eight other states and territories was geocoded, so the nearest district is 149 km from Hyderabad, 231 km from Lucknow and "
         "191 km from Kolkata, and a route is then scored by a district far away. Scoring one route takes about 305 ms in pure Python.")
    para("Two runs through the interface illustrate this. From Delhi to Jaipur OSRM returned a 274.1 km route of 204 min with "
         "33% low, 25% medium and 42% high samples and an alternative of 299.5 km and 205 min with 19%, 5% and 76%; both were labelled "
         "HIGH risk, so the first was recommended on its lower score alone, with one minute less travel time. From Hyderabad to Vijayawada "
         "(269.1 km, 213 min) there was a single route, 20% low and 80% high, scored against the districts nearest its sample points "
         "while the starting point itself has no district within 149 km.")
    P.h2("D. SafePhrase, Tracking, Evidence and Tests")
    para("On written transcripts, exact and embedded phrases fire at every setting. The default Medium setting behaves like Low for a "
         "three-word phrase (two of three words is 0.67, below 0.85): recall was 54.2% with no false triggers, High 75.0%, and Ultra 100% "
         "with 10.0% false triggers on ordinary sentences such as \"tell me about sheguard\" (Fig. 3b). No audio was used, so speech "
         "recognition accuracy was not measured. A guardian sees a new position 15.4 ms after the push (p95 17.2 ms; 5 viewers 14.6 ms). "
         "Evidence upload took 15-68 ms on average for 0.1-5 MB, and a one-bit change alters the stored hash. Access-control probes behaved as designed: "
         "all unauthenticated calls returned 401 and every cross-user call returned 404. The 57 tests give 51 passes and six expected failures "
         "that record real defects: an assistant timeout returns HTTP 500, latitude 999 and any trigger string are accepted, the WebSocket "
         "accepts any token, uploads have no size limit, and an uploaded .html file is served back as text/html. Evidence files are also "
         "retrievable without a token if the name is known, and 25 wrong passwords were never throttled.")
    P.table([["Transcript type (n)", "Low", "Medium", "High", "Ultra"],
             ["Exact phrase (5)", "5", "5", "5", "5"],
             ["Phrase inside a sentence (5)", "5", "5", "5", "5"],
             ["Words in another order (3)", "3", "3", "3", "3"],
             ["One word dropped (6)", "0", "0", "3", "6"],
             ["One word misheard (5)", "0", "0", "2", "5"],
             ["False triggers, 80 checks", "0", "0", "0", "8"]],
            "SafePhrase matcher: transcripts that fired", [0.43, 0.14, 0.15, 0.14, 0.14], size=6.8)
    para("Table III shows where the sensitivity levels differ: only dropped or misheard words are affected, and the setting that "
         "catches them is the one that mistakes ordinary speech for a phrase. Two phrases of different length were tested "
         "(\"help me sheguard\" and \"please call my brother now\"), and the four-word-plus phrase tolerates one missing word at High "
         "(4 of 5 words is 0.8) while the three-word default does not.")
    para("Latency was low everywhere except where an outside service or a deliberately slow hash is involved. Listing the 467 risk "
         "areas took 31 ms at the median, a guardian's location push 15 ms, the public tracking page 8 ms and a session start with 30 "
         "stored guardians and no channel configured 34 ms. Evidence upload grew roughly linearly with size (median 13 ms at 100 KB, 18 ms at "
         "1 MB, 61 ms at 5 MB), and hashing 1 MB takes about 1 ms. Login costs 279 ms because of bcrypt, which is intended. "
         "These numbers exclude any real Telegram or SMTP call, which would add network time to every session start.")
    P.h2("E. Comparison with the Base Paper")
    para("SafeRoutes [@base] and SHEGUARD share the idea of clustering regions by crime and acting on route deviations, but nothing in their "
         "evidence can be set side by side. The base paper uses seven public datasets, several features per district, its own "
         "district count and 2013 crime data, reports clustering tables and describes the cab-vendor alerts as future integration. "
         "SHEGUARD uses one feature from 2001-2012, 467 districts and a working route scorer, and adds activation, tracking and "
         "evidence that the base paper does not attempt. Conversely, SHEGUARD lacks population and per-capita data, police strength, "
         "socio-economic features, a mixture-model comparison inside the product and any vehicle integration. Its silhouette of 0.580 "
         "on a single count is not evidence of a better model; it shows only that districts separate along one axis.")
    P.h2("F. Failure Cases and Limits")
    para("Four observed failures matter most. First, the assistant: the first two calls returned HTTP 500 after the 20 s read timeout "
         "because that exception is not caught, one returned 503 from the provider, and the ten probe prompts all returned 429 because "
         "the key's quota was exhausted, so answer quality and the emergency first line were not measured. Second, alerts: with "
         "no Telegram or Gmail credentials, every guardian showed \"not sent\"; the code also calls a blocking SMTP routine inside an "
         "asynchronous route, which would stall the server while it runs. Third, coverage: eleven of 35 states and territories "
         "have no geocoded district. Fourth, the user interface prints \"Call 911\" on the emergency screen although India's number is 112. "
         "Not tested: real device GPS, battery and speech accuracy, load beyond five WebSocket viewers, iOS Safari, a screen reader "
         "and a user study.")
    P.h2("G. Threats to Validity")
    para("The risk tiers have no ground truth: silhouette and Davies-Bouldin measure separation along one count, not whether a "
         "district is unsafe. Counts of reported cases mix reporting behaviour with crime and are dominated by domestic-violence "
         "cases, which say little about street safety, and they are not divided by population. The 2001-2012 window is old, "
         "geocoding silently dropped 43.5% of district pairs, and three of the ten highest-rate districts (Vijayawada, Jaipur, Mumbai) reported only five years. "
         "The route study uses 16 pairs between large cities, the latency figures come from one idle machine with a single client, and "
         "the SafePhrase transcripts were written by the authors; the matcher result may differ with real speech.")
    P.h2("H. Output Screens")
    P.figure(FIG / "p_route.png", "Safe Routes screen: Delhi to Jaipur, two OSRM alternatives, both labelled HIGH risk.", width=3.4)
    P.figure(FIG / "p_track.png", "Guardian's tracking page, opened without a login during a live emergency.", width=3.0)
    para("The screens come from a registered account with three guardians. The emergency screen reported 0 of 3 guardians reached because "
         "no channel was configured, and the assistant returned its unavailable message, so neither is shown as working.")

    # ---------------------------------------------------------------- V
    P.h1("V. Conclusion")
    para("SHEGUARD delivers a running prevention-to-evidence platform: tiered district risk, scored routes, SafePhrase and one-tap SOS, "
         "live guardian tracking in tens of milliseconds, hashed evidence and an assistant, with owner-scoped access verified by probes. "
         "Against the six objectives, tracking and evidence are met, route guidance, quick activation and data protection are partly met, "
         "and the assistant is implemented but was not demonstrated. The key finding is that the area-risk model is two cut-offs on one "
         "count from 2001-2012, with 56.5% geocoding and no district near several large cities, and that its night factor cannot "
         "change a recommendation.")
    para("Other limits are untested alert channels, SafePhrase checked only on written text, no field trial, "
         "and the defects above. Future work: population-normalised, current data with a distance limit, a night factor applied per "
         "segment, verified Telegram and e-mail delivery, authenticated evidence download with size limits, real speech tests and a user study.")

    para("Taken one by one, the objectives stand as follows. Safe-route guidance works but rests on old, coarse and incomplete data. "
         "SafePhrase and SOS activate quickly, but the phrase matcher was checked on text only. Live guardian tracking and evidence "
         "hashing work as specified. The assistant and the alert channels are implemented but unverified here, and data protection "
         "rests on owner-scoped queries instead of database-level row security.")
    P.references()
    P.d.core_properties.title = K.TITLE
    P.d.core_properties.author = "Batch A7"
    P.d.save(OUT)
    print("saved", OUT, "refs:", len(ORDER), ORDER)


if __name__ == "__main__":
    build()
