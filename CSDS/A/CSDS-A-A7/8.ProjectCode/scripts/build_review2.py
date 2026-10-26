"""Build 13.Review_2_PPT.pptx from D:\\abstracts\\Templates\\Review_2_PPT.pptx (20 slides) for SHEGUARD.

Every number is from docs/REPORT_FACTS.md and experiments/study/*.json.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import (fill, put_picture, replace_chart, replace_everywhere, set_notes, set_paras, set_table, shape, style_chart, table_size)  # noqa: E402
from doc_common import BATCH, CHARTS, COLLEGE, DEPT, DIAG, GUIDE, GUIDE_DESIG, PRODUCT, SHOTS, STUDENTS, TITLE, YEAR  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OUT = BATCH / "13.Review_2_PPT.pptx"
BATCH_TAG = "Batch A7"
if os.environ.get("SG_DIAG"):  # test hook: build with stand-in diagram files
    DIAG = Path(os.environ["SG_DIAG"])


def px(v):
    return Emu(int(v * 9525))


def cards(s, items, size=15):
    for tid, bid, ttl, body in items:
        set_paras(shape(s, tid), ttl)
        fill(s, bid, body, size, min_pt=10)


def frame_to(pic, frame):
    frame.left, frame.top, frame.width, frame.height = pic.left, pic.top, pic.width, pic.height


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    replace_everywhere(prs, {"The Hidden Risk of Social Media": PRODUCT})
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Review 2", TITLE, BATCH_TAG, "women's safety, safe route, K-Means, SafePhrase, SOS, live tracking"
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = BATCH_TAG

    # 1 title
    s = S[0]
    set_paras(shape(s, 2), PRODUCT)
    set_paras(shape(s, 3), "AI Powered Women's Safety Platform")
    set_paras(shape(s, 6), [f"{n}  {r}" for n, r in STUDENTS], 15)
    set_paras(shape(s, 8), f"{GUIDE}, {GUIDE_DESIG}")
    set_paras(shape(s, 9), [DEPT, f"{COLLEGE}  \u00b7  {BATCH_TAG} (CSE-DS, Section A)  \u00b7  {YEAR}"], 13)
    set_notes(s, "Review 2 covers the design, methodology and progress of SHEGUARD.")

    # 2 contents
    s = S[1]
    for sid, t in [(21, "Model and configuration"), (18, "Data preparation"), (27, "Data flow, modules, database"), (36, "Preliminary results and screens")]:
        set_paras(shape(s, sid), t)
    set_notes(s, "Twelve parts: recap, architecture, methodology, data, preparation, model, decision logic, data flow and database, technology, progress, results and screens, testing and next steps.")

    # 3 recap
    s = S[2]
    fill(s, 8, "Women face harassment on isolated and crowded streets. Safety apps need a manual SOS press, mostly send one alert, choose routes for speed and rarely keep evidence.", 16)
    fill(s, 13, "A safety platform: risk-scored routes, SafePhrase or one-tap SOS, a live guardian link, hashed evidence and a Gemini assistant. Built as a React PWA with FastAPI and SQLite.", 16)
    for sid, t in zip((15, 18, 21, 24, 27), ["Trip, phrase or tap", "District risk tiers, OSRM routes", "SOS session + alerts", "Live location, evidence", "Guardian link, cancel"]):
        fill(s, sid, t, 15)
    set_notes(s, "Recap of Review 1. The base paper [1] (SafeRoutes) clusters regions for safe paths; SHEGUARD adopts that idea and adds the emergency features. The batch abstract names Flutter and Supabase; the implemented stack is React, FastAPI and SQLite.")

    # 4 architecture
    s = S[3]
    put_picture(s, DIAG / "01_system_architecture.png", (57, 124, 720, 518), remove_ids=(4,), name="System architecture diagram")
    fill(s, 5, ["Client: React 19 + TypeScript + Vite (port 5107), installable PWA; Leaflet map, Web Speech API, geolocation, MediaRecorder.",
                "API: FastAPI (port 8107), 20 endpoints (19 HTTP + 1 WebSocket); JWT login, per-user queries.",
                "Services: route scorer, OSRM and Nominatim clients, SHA-256 hashing, Gemini client, Telegram and Gmail senders.",
                "Data: SQLite (6 tables), evidence folder, risk_model.joblib and district_risk.csv.",
                "Alert channels are coded but not configured in this run."], 16, bullet=True)
    set_notes(s, "The browser talks to the Vite server, which proxies /api and /ws to FastAPI. Outside services: OSRM demo server, Nominatim, Gemini, Telegram Bot API and Gmail SMTP.")

    # 5 methodology
    s = S[4]
    for sid, t in zip((5, 8, 11, 14, 17, 20, 23), ["District crime data 2001-2012", "Sum 7 crimes per reporting year", "Geocode: 467 of 827 districts", "K-Means k=3 on crime rate", "Evaluate: indices, k sweep", "Score OSRM routes by tier", "Serve: FastAPI + React"]):
        fill(s, sid, t, 14, min_pt=9)
    set_paras(shape(s, 25), "Cluster, not classify")
    fill(s, 26, "There are no safe/unsafe labels, so districts are clustered on one number and named low, medium, high.", 15)
    set_paras(shape(s, 28), "Tested around the model")
    fill(s, 29, "Alternative tierings, seed stability, route sensitivity and the SafePhrase matcher were measured separately.", 15)
    set_paras(shape(s, 31), "Reproducible")
    fill(s, 32, "Fixed seed 42; scripts/run_experiments.py regenerates experiments/study/metrics.json.", 15)
    set_notes(s, "Seven-step workflow from crime tables to a served application. The model is trained once, in under a second on the CPU, and its output table is loaded into the database at start-up.")

    # 6 dataset
    s = S[5]
    set_paras(shape(s, 2), "Dataset and input sources")
    ch = replace_chart(s, 4, ["Crime rows (2001-2012)", "State-year rows (not used)", "District pairs", "Geocoded and used"], [("Rows or districts", [8597, 736, 827, 467])])
    style_chart(ch, "Rows or districts per input", ["2F6FB0"], legend=False, label_pt=11, axis_title="rows / districts")
    rows = [["Source", "Size", "Use"],
            ["Crime in India district file [23] (CC0)", "8,597 rows, 827 districts, 12 years", "Risk feature"],
            ["Nominatim geocodes", "467 of 827 districts (56.5 %)", "Map position"],
            ["State-year file 2001-2021 [24] (CC0)", "736 rows", "Not used by any feature"],
            ["OSRM, Nominatim (live)", "Public demo servers", "Routes, place search"],
            ["Demo database", "Demo account + registered users", "Walk-through"]]
    sz = table_size(rows, 530, 400, [0.38, 0.38, 0.24], 12, 9)
    set_table(shape(s, 7), rows, sz, widths=[0.38, 0.38, 0.24], header_size=sz)
    set_notes(s, "All inputs are real public data; there is no synthetic data in this project. They are old (2001-2012), district-level counts and only 56.5 % of districts could be geocoded.")

    # 7 preprocessing
    s = S[6]
    set_paras(shape(s, 2), "Data preparation and features")
    set_paras(shape(s, 7), "Crime table")
    fill(s, 8, "Drop TOTAL roll-up rows (9,017 to 8,597); sum the 7 crime columns; divide by the years a district reported (638 of 827 reported all 12). Feature: crimes per reporting year.", 15)
    set_paras(shape(s, 12), "Geocoding and join")
    fill(s, 13, "Nominatim once, 1 request per second; 467 of 827 kept, the rest dropped not guessed. No geocoded district for Uttar Pradesh, West Bengal, Delhi and 8 other states/UTs. StandardScaler on the feature; score 0-100 by min-max.", 15)
    set_paras(shape(s, 17), "Run-time inputs")
    fill(s, 18, "Route polyline sampled at every 8th point, nearest district by haversine (no distance limit). SafePhrase transcript lowercased, punctuation removed, matched in a word window.", 15)
    set_notes(s, "Preparation = turning the crime table into one number per district, giving each district a position, and turning a route or a transcript into inputs for the scorer or matcher.")

    # 8 model
    s = S[7]
    set_paras(shape(s, 2), "Models and decision rules")
    put_picture(s, DIAG / "02_high_level_workflow.png", (75, 124, 693, 518), remove_ids=(4,), name="High-level workflow")
    rows = [["Component", "Detail"], ["Risk model", "K-Means, k=3, one feature"], ["Tier weights", "low 1.0, medium 2.5, high 5.0"], ["Night factor", "x 1.6 (about 20:30-05:30 IST)"], ["SafePhrase", "word-window match, 4 levels"]]
    set_table(shape(s, 9), rows, 12, header_size=12)
    set_paras(shape(s, 6), "Classical clustering on the CPU; no deep network and no GPU.", 11)
    set_notes(s, "The abstract speaks of an AI-powered platform; the only trained model is a three-cluster K-Means. The route scorer and the SafePhrase matcher are transparent rules.")

    # 9 training configuration
    s = S[8]
    set_paras(shape(s, 2), "Training configuration")
    rows = [["Setting", "Value"], ["Data", "467 districts, 1 feature"], ["Algorithm", "K-Means, k=3, seed 42, n_init 10"], ["Scaling", "StandardScaler"],
            ["Split", "none (unsupervised, all 467)"], ["Fit time", "under 1 s on CPU"], ["Quality", "silhouette 0.580, DB 0.531"], ["Tiers", "278 low, 148 medium, 41 high"], ["Cut-offs", "213.75 and 529.83 crimes/year"]]
    set_table(shape(s, 10), rows, 13, header_size=13)
    pic = put_picture(s, CHARTS / "01_k_sweep.png", (672, 230, 550, 270), remove_ids=(5,), name="k sweep chart")
    frame_to(pic, shape(s, 6))
    set_paras(shape(s, 7), "Silhouette and Davies-Bouldin for k = 2 to 8: k=3 is a UI choice, not the optimum", 11)
    shape(s, 7).top = Emu(pic.top + pic.height + px(10))
    set_notes(s, "Silhouette is 0.662 at k=2, 0.580 at k=3 and 0.596 at k=4; Davies-Bouldin is 0.531 at k=3 and 0.493 at k=4. Three tiers were chosen for the interface. Thirty single-initialisation runs agree with the shipped tiers with an adjusted Rand index of at least 0.66 (mean 0.79).")

    # 10 decision logic
    s = S[9]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (45, 124, 620, 520), remove_ids=(4,), name="Decision flowchart")
    items = [(5, 6, 7, "Route score", "Every 8th point: nearest district tier weight; mean x 1.6 at night; lowest score is recommended."),
             (8, 9, 10, "SOS trigger", "One tap, or a SafePhrase match; start session, alert guardians, stream location; 3-second hold cancels."),
             (11, 12, 13, "SafePhrase level", "Medium (default) 54 % recall, 0 of 80 false triggers; Ultra 100 % recall, 10 % false triggers.")]
    for card, tid, bid, ttl, body in items:
        c = shape(s, card)
        t, b = shape(s, tid), shape(s, bid)
        c.left, c.width = px(690), px(531)
        t.left, t.width = px(714), px(490)
        b.left, b.width = px(714), px(490)
        set_paras(t, ttl)
        t.top, t.height = c.top + px(10), px(40)
        b.top, b.height = c.top + px(50), px(92)
        fill(s, bid, body, 15, min_pt=10)
    set_notes(s, "Decision logic. The night factor multiplies every route by the same 1.6, so it changes the displayed score but never the recommended route (16 of 16 tested city pairs keep their order). SafePhrase figures are from written transcripts only; no audio or speech recogniser was measured.")

    # 11 data flow
    s = S[10]
    set_paras(shape(s, 2), "Data flow and modules")
    put_picture(s, DIAG / "04_dfd_level1.png", (57, 118, 470, 340), remove_ids=(4,), name="DFD level 1")
    put_picture(s, DIAG / "05_dfd_level2.png", (545, 118, 676, 340), name="DFD level 2")
    set_paras(shape(s, 5), "DFD level 1 (left) and level 2 expansion of 'Plan safe route' (right)", 11)
    shape(s, 5).left, shape(s, 5).top = px(57), px(462)
    rows = [["Module", "Responsibility"], ["Routes and risk", "OSRM alternatives, nearest-district scoring, risk areas for the map layer"],
            ["SOS and tracking", "session, location pings, WebSocket push, public share-token page, cancel"],
            ["Evidence, assistant, access", "photo/audio with SHA-256; Gemini question endpoint; JWT and owner scoping"]]
    set_table(shape(s, 12), rows, 12, widths=[0.22, 0.78], header_size=12, row_h=34)
    shape(s, 12).top = px(492)
    set_notes(s, "Level 1 splits the system into processes and data stores; level 2 opens the route-planning process.")

    # 12 database and UML
    s = S[11]
    set_paras(shape(s, 2), "Database design and UML")
    put_picture(s, DIAG / "06_er_diagram.png", (57, 118, 385, 400), remove_ids=(5,), name="ER diagram")
    put_picture(s, DIAG / "07_uml_class_diagram.png", (455, 118, 370, 400), name="UML class diagram")
    set_paras(shape(s, 6), "ER diagram, 6 tables (left) and UML class diagram (right)", 11)
    shape(s, 6).left, shape(s, 6).top, shape(s, 6).width = px(57), px(524), px(768)
    rows = [["Table", "Key fields", "Relations"],
            ["users, guardians", "email, password hash; owner_id", "1:N guardians, sessions"],
            ["emergency_sessions", "share_token, status, trigger", "1:N pings, evidence"],
            ["location_pings, evidence", "lat, lng, ts; kind, sha256", "N:1 session"],
            ["district_risk", "district, lat, lng, tier", "seeded from CSV (467)"]]
    set_table(shape(s, 13), rows, 10, widths=[0.36, 0.38, 0.26], header_size=10, row_h=30, total_px=371)
    t = shape(s, 13)
    t.left, t.top, t.width = px(850), px(150), px(371)
    set_notes(s, "Six SQLite tables created by SQLAlchemy. Guardian rows belong to a user; a guardian has no login and sees only a session's share-token page. SafePhrase settings live in the browser, not in the database.")

    # 13 technology
    s = S[12]
    cards(s, [(7, 8, "Language", "Python 3.11, TypeScript 5.8"), (12, 13, "Machine learning", "scikit-learn 1.5.2 (K-Means), pandas 2.2.3, NumPy 2.1.3, joblib"),
              (17, 18, "Maps and services", "OSRM, Nominatim, OpenStreetMap tiles, Leaflet 1.9, Gemini API"),
              (22, 23, "Backend", "FastAPI 0.115, SQLAlchemy 2.0, PyJWT, bcrypt, WebSocket, Telegram and Gmail senders"),
              (27, 28, "Frontend", "React 19, Vite 6, Tailwind CSS 4, Web Speech API, PWA plugin"),
              (32, 33, "Storage and tests", "SQLite, evidence folder, pytest (51 pass, 6 expected fail), Playwright captures")], 15)
    set_notes(s, "Versions are those pinned in requirements.txt and package.json or reported by the installed packages. The abstract's Flutter and Supabase were not used.")

    # 14 progress
    s = S[13]
    rows = [["Component", "Status"], ["Accounts, guardians, per-user data scoping", "Completed"], ["K-Means district risk and route scoring", "Completed"],
            ["SOS, live location, WebSocket tracking page", "Completed"], ["Evidence capture with SHA-256", "Completed"],
            ["SafePhrase: speech + matcher", "Completed; recogniser not measured"], ["Alerts: Telegram, Gmail", "Coded, never delivered"],
            ["Gemini assistant", "Coded; quota error in test"], ["Real-device, field and user testing", "Not done"]]
    set_table(shape(s, 15), rows, 13, widths=[0.6, 0.4], header_size=13)
    for vid, lid, v, lab in [(6, 7, "20", "API endpoints (19 HTTP + 1 WebSocket)"), (9, 10, "6", "database tables"), (12, 13, "51", "tests pass; 6 expected failures")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "The core flow works end to end. Guardian alerts could not be verified because the Telegram and Gmail credentials are empty, and the Gemini key's quota was exhausted during testing.")

    # 15 preliminary results
    s = S[14]
    set_paras(shape(s, 2), "Preliminary results: SafePhrase matcher and risk model")
    ch = replace_chart(s, 4, ["Low (100 %)", "Medium (85 %)", "High (75 %)", "Ultra (60 %)"], [("Recall on 24 phrase transcripts (%)", [54.2, 54.2, 75.0, 100.0]), ("False triggers on 80 checks (%)", [0.0, 0.0, 0.0, 10.0])], number_format="0.0")
    style_chart(ch, "SafePhrase matcher on written transcripts, by sensitivity (word share needed)", ["2E8B57", "C8453C"], number_format="0.0", legend=True, label_pt=11, axis_title="percent")
    ch.value_axis.maximum_scale = 110
    for vid, lid, v, lab in [(6, 7, "0.580", "silhouette of the 3 tiers (Davies-Bouldin 0.531)"), (9, 10, "149 km", "Hyderabad to its nearest district in the table"), (12, 13, "16 / 16", "city pairs where the night factor keeps the route order")]:
        set_paras(shape(s, vid), v)
        set_paras(shape(s, lid), lab)
    set_notes(s, "Silhouette measures cluster separation on one number, not real safety. Coverage: 40 city centres were checked; the nearest district is a median of 8.1 km away but up to 231 km (Lucknow); 7 cities are more than 100 km from any district. SafePhrase numbers come from hand-written transcripts, not speech audio.")

    # 16 output screens
    s = S[15]
    p1 = put_picture(s, SHOTS / "06_route_delhi_jaipur.png", (45, 128, 590, 476), remove_ids=(4,), name="Route screenshot")
    p2 = put_picture(s, SHOTS / "13_emergency_evidence.png", (645, 128, 590, 476), remove_ids=(7,), name="Emergency screenshot")
    frame_to(p1, shape(s, 5))
    frame_to(p2, shape(s, 8))
    set_paras(shape(s, 6), "Delhi to Jaipur: two real OSRM routes, both HIGH risk", 11)
    set_paras(shape(s, 9), "Emergency: photo and audio saved with SHA-256; 0 of 3 guardians reached (no channel set)", 11)
    for cap_id, p in ((6, p1), (9, p2)):
        shape(s, cap_id).top = Emu(max(p.top + p.height + px(12), px(500)))
    set_notes(s, "Real captures of the running application. Location was set by the test browser to a Delhi and a Hyderabad city centre; the AI tips panel shows the unavailable message because the assistant's quota was exhausted.")

    # 17 testing
    s = S[16]
    set_paras(shape(s, 2), "Testing and observations")
    rows = [["Test group", "Tests", "Result"], ["Authentication, guardians", "14", "14 passed"], ["Risk data, route planning", "6", "6 passed"], ["SOS, tracking, WebSocket", "7", "7 passed"],
            ["Evidence", "6", "6 passed"], ["AI endpoint, places search", "6", "6 passed"], ["Scorer and model files", "12", "12 passed"], ["Known defects", "6", "6 expected fail"]]
    set_table(shape(s, 18), rows, 12, widths=[0.5, 0.12, 0.38], header_size=12)
    fill(s, 5, ["51 pass in about 14 s; 6 tests assert hardened behaviour and fail today: Gemini timeout gives 500, latitude 999 accepted, any trigger string, WebSocket accepts any token, no upload size limit, .html evidence served as HTML.",
                "API median times: health 2.1 ms, SOS start 34 ms, location push 15 ms, login 279 ms (bcrypt), live route plan 1.6 s; WebSocket push 13 ms.",
                "Stored evidence files are readable without login by URL.",
                "Not done: real alert delivery, real device, user and load testing."], 15, bullet=True)
    set_notes(s, "Expected-fail tests are reported as such rather than hidden. Latencies were measured on 127.0.0.1 on an i5-12500, one client at a time.")

    # 18 challenges
    s = S[17]
    cards(s, [(7, 8, "Old, coarse data", "2001-2012 district counts, not per population; 56.5 % of districts geocoded."),
              (12, 13, "One-feature model", "Tiers are two cut-offs (213.75, 529.83); k=4 scores better than the chosen k=3."),
              (17, 18, "Coverage gaps", "No district for Uttar Pradesh, West Bengal, Delhi; Hyderabad is 149 km from the nearest."),
              (22, 23, "Night factor is a no-op", "x 1.6 on every route: ranking unchanged in 16 of 16 pairs."),
              (27, 28, "Alerts and AI unverified", "Telegram and Gmail not configured; Gemini quota 429 and a timeout returns 500."),
              (32, 33, "Open evidence URLs", "Files readable without login; no size limit; found by the security probes.")], 14)
    set_notes(s, "Six challenges and findings, each backed by a measurement in REPORT_FACTS.md.")

    # 19 next steps
    s = S[18]
    for sid, t in zip((5, 8, 11, 14), ["Better risk data", "Fix the findings", "Real-device test", "Final report and demo"]):
        fill(s, sid, t, 16)
    fill(s, 15, ["Use population-normalised, newer data and cover the missing states; cap the distance to the nearest district.",
                 "Handle provider timeouts, validate coordinates, authenticate the WebSocket, protect evidence files, cap uploads.",
                 "Configure Telegram and Gmail and test delivery on a phone; measure speech recognition; run a small user trial with real users.",
                 "Complete the report and the final demonstration."], 17, bullet=True)
    set_notes(s, "What remains: better data and the hardening items found in testing, then real-device checks, then the report and final demonstration.")

    set_notes(S[19], "Thank you. Questions are welcome.")
    # template leftovers in shape names and picture descriptions (slide 18 cards)
    for n, sid in enumerate((4, 9, 14, 19, 24, 29), 1):
        shape(S[17], sid).name = f"Challenge card {n}"
    for sh in S[17].shapes:
        if sh.shape_type == 13:
            sh._element.nvPicPr.cNvPr.set("descr", "icon")
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    for m in prs.slide_masters:
        th = m.part.part_related_by(RT.THEME)
        th._blob = th.blob.replace(b"Hidden Risk", PRODUCT.encode())
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
