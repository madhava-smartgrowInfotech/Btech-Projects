"""python-docx helpers for the SeatWise submission documents.

Times New Roman 12 body, 14 bold headings, 1.5 line spacing, justified text, numbered figures and tables with
captions (real Word fields for the table of contents and the lists of figures and tables), page numbers.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

FONT = "Times New Roman"
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")


def _set_font(style_or_run, name=FONT, size=None, bold=None, italic=None, color=None):
    f = style_or_run.font
    f.name = name
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    el = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(att), name)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(att)) is not None:
                del rfonts.attrib[qn(att)]


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _field(paragraph, instr, placeholder="", bold=False, size=None):
    """Insert a complex field (TOC, PAGE, ...) whose cached text is `placeholder`."""
    def run_with(child):
        r = paragraph.add_run()
        if size:
            r.font.size = Pt(size)
        r.bold = bold
        r._r.append(child)
        return r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run_with(b); run_with(i); run_with(s)
    t = paragraph.add_run(placeholder)
    t.bold = bold
    if size:
        t.font.size = Pt(size)
    run_with(e)


class Doc:
    def __init__(self, header_text: str, numbering: str = "flat"):
        """numbering: 'chapter' (final report: Fig. 5.1, 5.2) or 'flat' (chapter documents: Fig. 1, 2)."""
        self.d = Document()
        self.numbering = numbering
        self.header_text = header_text
        self.ch = 0
        self.sec = 0
        self.sub = 0
        self.fig_n = 0
        self.tab_n = 0
        self.appendix = None
        self._style()
        self._page(self.d.sections[0])

    # ------------------------------------------------------------ setup
    def _style(self):
        st = self.d.styles
        n = st["Normal"]
        _set_font(n, size=12)
        n.paragraph_format.line_spacing = 1.5
        n.paragraph_format.space_after = Pt(2)
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size, align, before in (("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER, 0), ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 12),
                                          ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 8)):
            h = st[name]
            _set_font(h, size=size, bold=True, italic=False, color="000000")
            h.paragraph_format.alignment = align
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.15
        for name in ("Figure Caption", "Table Caption"):
            c = st.add_style(name, 1)
            c.base_style = st["Normal"]
            _set_font(c, size=11, bold=True)
            c.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.line_spacing = 1.0
            c.paragraph_format.space_before = Pt(4)
            c.paragraph_format.space_after = Pt(10)
        st["Figure Caption"].paragraph_format.keep_together = True
        st["Table Caption"].paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            _set_font(st[name], size=12)
            st[name].paragraph_format.line_spacing = 1.5
            st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fh = st.add_style("Front Heading", 1)
        fh.base_style = st["Normal"]
        _set_font(fh, size=14, bold=True)
        fh.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fh.paragraph_format.space_after = Pt(12)
        fh.paragraph_format.keep_with_next = True
        code = st.add_style("Code Block", 1)
        code.base_style = st["Normal"]
        _set_font(code, name="Courier New", size=7.5)
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        toc = st["Normal"]

    def _page(self, section):
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin, section.right_margin = Inches(1.0), Inches(0.8)
        section.header_distance = section.footer_distance = Inches(0.45)

    @property
    def text_width(self):
        s = self.d.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin) / 914400

    def header_footer(self, section=None):
        sec = section or self.d.sections[-1]
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.text = ""
        r = hp.add_run(self.header_text)
        _set_font(r, size=9, italic=True, color="555555")
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.line_spacing = 1.0
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp, "PAGE", "1", size=10)
        fp.paragraph_format.line_spacing = 1.0

    def number_format(self, fmt="decimal", start=None, section=None):
        sec = section or self.d.sections[-1]
        ppr = sec._sectPr
        pg = ppr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            ppr.append(pg)
        pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))

    def new_section(self, landscape=False):
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        pg = sec._sectPr.find(qn("w:pgNumType"))
        if pg is not None and pg.get(qn("w:start")) is not None:
            del pg.attrib[qn("w:start")]  # continue numbering across the new section
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
            sec.left_margin = sec.right_margin = Inches(1)
            sec.top_margin = sec.bottom_margin = Inches(0.9)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            self._page(sec)
        return sec

    # ------------------------------------------------------------ text
    def _runs(self, p, text, size=None, color=None, base_bold=False, base_italic=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2]); r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1]); _set_font(r, name="Courier New", size=(size or 12) - 1.5)
                continue
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.italic = True
            else:
                r = p.add_run(part)
                if base_bold:
                    r.bold = True
            if base_italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def p(self, text, align=None, size=None, bold=False, italic=False, space_after=None, keep_next=False, indent=None, line=None):
        para = self.d.add_paragraph()
        self._runs(para, text, size=size, base_bold=bold, base_italic=italic)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            para.paragraph_format.keep_with_next = True
        if indent is not None:
            para.paragraph_format.left_indent = Inches(indent)
        if line is not None:
            para.paragraph_format.line_spacing = line
        return para

    def paras(self, *texts):
        for t in texts:
            self.p(t)

    def bullets(self, items, size=None):
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            self._runs(para, it, size=size)
            para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            para = self.d.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.4)
            para.paragraph_format.first_line_indent = Inches(-0.3)
            para.paragraph_format.space_after = Pt(2)
            self._runs(para, f"{i}.\t{it}")
            para.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ headings
    def h1(self, text, chapter=True, page_break=True):
        """A chapter (numbered), or an unnumbered heading when chapter=False."""
        if page_break and len(self.d.paragraphs) > 0:
            self.d.paragraphs[-1].paragraph_format.page_break_before = False
        if chapter:
            if self.appendix:
                self.appendix = chr(ord(self.appendix) + 1)
            self.ch += 1
            self.sec = self.sub = 0
            label = f"CHAPTER {self.ch}" if not self.appendix else f"APPENDIX {self.appendix}"
            para = self.d.add_paragraph(style="Heading 1")
            para.add_run(f"{label}")
            para.add_run().add_break()
            para.add_run(text.upper())
        else:
            para = self.d.add_paragraph(style="Heading 1")
            para.add_run(text.upper())
        if page_break:
            para.paragraph_format.page_break_before = True
        if self.numbering == "chapter":
            self.fig_n = self.tab_n = 0
        return para

    def appendix_h1(self, letter, text):
        self.appendix = letter
        para = self.d.add_paragraph(style="Heading 1")
        para.add_run(f"APPENDIX {letter}")
        para.add_run().add_break()
        para.add_run(text.upper())
        para.paragraph_format.page_break_before = True
        self.sec = self.sub = 0
        self.fig_n = self.tab_n = 0
        self.ch_label = letter

    def _prefix(self):
        return getattr(self, "ch_label", None) if self.appendix else (self.ch if self.numbering == "chapter" else None)

    def h2(self, text):
        self.sec += 1
        self.sub = 0
        pre = self._prefix()
        num = f"{pre}.{self.sec}" if pre is not None else f"{self.sec}"
        para = self.d.add_paragraph(style="Heading 2")
        para.add_run(f"{num}  {text}")
        return para

    def h3(self, text):
        self.sub += 1
        pre = self._prefix()
        num = f"{pre}.{self.sec}.{self.sub}" if pre is not None else f"{self.sec}.{self.sub}"
        para = self.d.add_paragraph(style="Heading 3")
        para.add_run(f"{num}  {text}")
        return para

    def front_heading(self, text, page_break=True):
        para = self.d.add_paragraph(style="Front Heading")
        para.add_run(text.upper())
        if page_break:
            para.paragraph_format.page_break_before = True
        return para

    def simple_table(self, rows, widths=None, font=11, bold_header=False, borders=True, align=None):
        """A table without caption (title blocks, signature blocks)."""
        ncols = max(len(r) for r in rows)
        t = self.d.add_table(rows=len(rows), cols=ncols)
        if borders:
            t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = self.text_width
        widths = widths or [1 / ncols] * ncols
        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = t.rows[ri].cells[ci]
                cell.width = Inches(total * widths[ci])
                para = cell.paragraphs[0]
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.space_after = Pt(2)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if align and align[ci] == "c":
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._runs(para, str(row[ci]) if ci < len(row) else "", size=font, base_bold=(bold_header and ri == 0))
                if borders and bold_header and ri == 0:
                    _shade(cell, "D9E2F3")
        return t

    def h_plain(self, text, level=1):
        para = self.d.add_paragraph(style=f"Heading {level}")
        para.add_run(text.upper() if level == 1 else text)
        return para

    # ------------------------------------------------------------ figures and tables
    def _label(self, kind):
        if kind == "Fig.":
            self.fig_n += 1
            n = self.fig_n
        else:
            self.tab_n += 1
            n = self.tab_n
        pre = self._prefix()
        return f"{kind} {pre}.{n}" if pre is not None and self.numbering == "chapter" else f"{kind} {n}"

    def figure(self, path, caption, width=None, landscape=False, max_h=None):
        if landscape:
            self.new_section(landscape=True)
        path = Path(path)
        with Image.open(path) as im:
            iw, ih = im.size
        avail = self.text_width
        w = min(width or avail * 0.85, avail)
        h = w * ih / iw
        limit = max_h or (4.0 if not landscape else 5.4)
        if getattr(self, 'report', False) and not landscape:
            limit = min(limit, 3.3)
            w = min(w, avail * 0.8)
            h = w * ih / iw
        if h > limit:
            w = w * limit / h
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.line_spacing = 1.0
        para.add_run().add_picture(str(path), width=Inches(w))
        cap = self.d.add_paragraph(style="Figure Caption")
        cap.add_run(f"{self._label('Fig.')}: {caption}")
        if landscape:
            self.new_section(landscape=False)


    def figure_pair(self, items, width=2.85, max_h=3.0):
        """Two figures side by side (borderless table), each with its own numbered caption. items: [(path, caption), ...]."""
        t = self.d.add_table(rows=1, cols=len(items))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cw = self.text_width / len(items)
        for ci, (path, caption) in enumerate(items):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(cw)
            with Image.open(path) as im:
                iw, ih = im.size
            w = min(width, cw - 0.1)
            h = w * ih / iw
            if h > max_h:
                w = w * max_h / h
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(path), width=Inches(w))
            cap = cell.add_paragraph(style="Figure Caption")
            cap.add_run(f"{self._label('Fig.')}: {caption}")
            cap.paragraph_format.space_after = Pt(4)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sp.paragraph_format.line_spacing = 0.6

    def table(self, rows, caption, widths=None, font=10, align=None, header=True, first_col_bold=False):
        cap = self.d.add_paragraph(style="Table Caption")
        cap.add_run(f"{self._label('Table')}: {caption}")
        ncols = max(len(r) for r in rows)
        t = self.d.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = self.text_width
        widths = widths or [1 / ncols] * ncols
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trpr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trpr.append(cs)
            if ri == 0 and header:
                th = OxmlElement("w:tblHeader")
                trpr.append(th)
            for ci in range(ncols):
                cell = tr.cells[ci]
                cell.width = Inches(total * widths[ci])
                txt = str(row[ci]) if ci < len(row) else ""
                para = cell.paragraphs[0]
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if align and ci < len(align) and align[ci] == "c":
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align and ci < len(align) and align[ci] == "r":
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                head = ri == 0 and header
                self._runs(para, txt, size=font, base_bold=head or (first_col_bold and ci == 0))
                if head:
                    _shade(cell, "D9E2F3")
        after = self.d.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = 0.6
        return t

    def code(self, text, caption=None):
        if caption:
            c = self.d.add_paragraph(style="Table Caption")
            c.add_run(f"Listing: {caption}")
            c.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            para = self.d.add_paragraph(style="Code Block")
            para.add_run(line if line else " ")
            ppr = para._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F4F7")
            ppr.append(shd)
            para.paragraph_format.left_indent = Inches(0.3)
            para.paragraph_format.first_line_indent = Inches(-0.25)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.6

    # ------------------------------------------------------------ fields
    def toc(self, title="TABLE OF CONTENTS", instr='TOC \\o "1-2" \\h \\z \\u', placeholder="Right-click and choose Update Field to build the table of contents."):
        self.front_heading(title)
        para = self.d.add_paragraph()
        _field(para, instr, placeholder)

    def lof(self):
        self.front_heading("LIST OF FIGURES")
        para = self.d.add_paragraph()
        _field(para, 'TOC \\h \\z \\t "Figure Caption,1"', "Update field to build the list of figures.")

    def lot(self):
        self.front_heading("LIST OF TABLES")
        para = self.d.add_paragraph()
        _field(para, 'TOC \\h \\z \\t "Table Caption,1"', "Update field to build the list of tables.")

    def save(self, path):
        self.d.core_properties.author = "Batch A17"
        self.d.save(str(path))
