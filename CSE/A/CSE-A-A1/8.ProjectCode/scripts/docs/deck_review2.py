"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress. Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx."""

import os
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"

BANNED = ["social media", "fakeddit", "pheme", "snopes", "mcnn", "bank transaction", "hidden risk", "recommender system", "skincare",
          "dayananda", "sindhu", "[student name", "[roll no", "genimage", "resnet", "glove", "manipulat", "sign language", "churn", "fraud",
          "placement", "xgboost"]


def scan_deck(path):
    """Banned strings found anywhere in the saved package (slides, notes, layouts, masters, theme, properties)."""
    hits = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                t = re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore")).lower()
                for b in BANNED:
                    if b in t:
                        hits.append((n, b))
    return hits


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "adaptive learning, retrieval-augmented generation, TF-IDF, learning analytics, educational document analysis"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 30)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}
    SH = F.SHIPPED
    h = lambda name, k: F.hit(name, k)          # noqa: E731
    C4, C1 = F.CMP["4"], F.CMP["1"]
    A = F.API
    DA = F.DOCAN

    # ---- 1 title
    title_slide(S[0], "Design, methodology and progress: a TF-IDF retrieval-augmented study assistant with quizzes and an adaptive roadmap",
                "REVIEW 2  ·  DESIGN, METHODOLOGY AND PROGRESS")
    notes[1] = "This review shows how the platform is designed and built: data, algorithms, what is implemented, the first measured results and the open items."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap and architecture"), ("02", "Methodology and dataset"), ("03", "Preprocessing and algorithms"), ("04", "Experiment configuration and decision logic"),
               ("05", "Data flow, database and UML"), ("06", "Technology stack"), ("07", "Implementation progress and results"),
               ("08", "Output screens and testing"), ("09", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Students keep syllabi, notes and question banks as scattered files. Portals store them without reading them, and a general chatbot answers from the internet, not from the student's own syllabus.", K.ORANGE, "!"),
        ("Proposed system", "Upload PDF, DOCX or TXT, extract subjects and topics, answer questions only from the uploaded text with the source passage shown, generate quizzes, and turn quiz accuracy into a ranked study roadmap.", K.BLUE, "i"),
        ("Base paper", "SAC-RAG [1] routes questions by type, retrieves with BM25 plus dense vectors and compresses the context with an LLM (38 to 58 percent fewer tokens). We build the retrieve-then-answer core and measure how far it gets.", K.GREEN, "B"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, 4.7, t, b_, col, gl, body_size=19)
    notes[3] = "Recap of Review 1: the gap between file storage and understanding, our working system, and what the base paper adds beyond it."

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.6, 5.5)
    K.text(s, 8.5, 1.5, 4.2, 5.2, [
        ("**Layers**", {"space_after": 8}),
        ("Presentation: React 19 app (port 5301) and the original Flask pages (port 5000).", {"bullet": True}),
        ("Application: FastAPI (port 8301), 14 endpoints, signed session cookie, CORS for the front end only.", {"bullet": True}),
        ("Core: text extraction, a per-user TF-IDF index (RAGEngine) and the quiz generator, shared by both interfaces.", {"bullet": True}),
        ("Storage: SQLite with 5 tables and the uploaded files; Claude is optional and used only if a key is set.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[4] = "The FastAPI layer imports the original Flask project's modules, so both interfaces use the same retrieval and quiz code."

    # ---- 5 methodology workflow
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.2, [("Upload PDF / DOCX / TXT", K.NAVY), ("Extract and clean text", K.NAVY), ("Chunk 180 / 30 words", K.GREEN), ("Subject, unit and topic hints", K.PURPLE),
                                      ("Store in SQLite", K.ORANGE), ("TF-IDF index per user", K.BLUE), ("Retrieve top 4, answer with sources", K.NAVY)], size=12, gap=0.22)
    for i, (t, b_) in enumerate([("One shared core", "The Flask pages and the React front end call the same functions, so the answers, quizzes and roadmap behave identically in both."),
                                 ("Grounded by construction", "The extractive answer is made of sentences copied from the retrieved chunks, and every answer lists the file, passage and match score."),
                                 ("Measured, not assumed", f"{F.CORPUS['questions']} questions over {F.CORPUS['chunks']} chunks compare seven retrievers; scripts regenerate every number.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.1, 3.85, 3.4)
        K.text(s, x + 0.25, 3.25, 3.4, 0.5, t, size=20, color=K.INK, bold=True)
        K.text(s, x + 0.25, 3.9, 3.4, 2.5, b_, size=18, color=K.TXT)
    notes[5] = "The workflow is also drawn as a diagram in the report. Nothing is trained: the retriever fits a TF-IDF vocabulary on the user's own chunks after every upload."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Evaluation data")
    wipe(s)
    for i, (v, l, c) in enumerate([("36", "Wikipedia articles in 6 subjects", K.BLUE), (f"{F.CORPUS['words']:,}", "words of notes", K.GREEN),
                                   (str(F.CORPUS["chunks"]), "chunks of 180 words (30 overlap)", K.ORANGE), (str(F.CORPUS["questions"]), "hand-written questions (+12 off-topic)", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=12)
    docs = F.CORPUS["documents"]
    K.bar_chart(s, 0.6, 2.95, 5.6, 3.9, [d["subject"].replace(" Systems", "").replace("Computer ", "") for d in docs], [("Chunks per subject", [d["chunks"] for d in docs])],
                colors=[K.BLUE], number_format="0", font_size=11)
    K.text(s, 6.5, 3.0, 6.2, 3.9, [
        ("**No real student documents were available**, so each subject file is built from six English Wikipedia articles (CC BY-SA 4.0, pinned revisions) in the Subject / Unit layout the analyser reads.", {"bullet": True, "size": 15}),
        ("72 questions, 24 each of fact, definition and reasoning (the question types of [1]). The gold evidence is a sentence copied from the text; a chunk is a hit if it contains it.", {"bullet": True, "size": 15}),
        ("12 off-topic questions test whether the system refuses.", {"bullet": True, "size": 15}),
        ("The bundled 619-word demo notes were used for the screenshots.", {"bullet": True, "size": 15}),
    ], size=14, color=K.INK)
    notes[6] = "Be upfront: this is a notes-style corpus built for measurement, not data from real students. The gold evidence strings are verified to occur exactly once in the corpus."

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Document analysis and preprocessing")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Extract and clean", "PDF pages with pypdf, DOCX paragraphs with python-docx, TXT as UTF-8. Spaces and blank lines are collapsed. Table text in DOCX and scanned PDFs without a text layer give nothing.", K.BLUE, "1"),
        ("Chunk", "Word windows of 180 words with a 30-word overlap, so a sentence cut at a boundary survives in the next chunk. The 619-word demo file gives 4 chunks; the corpus gives 672.", K.GREEN, "2"),
        ("Subjects, units and topics", f"Regular expressions read Subject, Course, Unit n and Chapter n lines (36 of 36 unit headings found in the corpus); TF-IDF keywords fill up to 12 topics. UNIT-I, Module 1 and numbered headings are missed (0 of 5).", K.ORANGE, "3")]):
        K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 4.6, t, b_, col, gl, body_size=17)
    K.text(s, 0.6, 6.25, 12.13, 0.6, "Result of the analysis: a document row with subject and topics, and the chunk rows that form the knowledge repository.", size=14, color=K.TXT, align="c", italic=True)
    notes[7] = "Heading detection is regex based. On a five-unit test syllabus, formats Unit 1:, Unit 1 - and Chapter 1: found 5 of 5; UNIT-I, Unit I:, Module 1: and 1. found 0 of 5."

    # ---- 8 algorithms
    s = S[7]
    set_title(s, "Algorithms")
    wipe(s)
    rows = [["Component", "Algorithm and parameters"],
            ["Retrieval", "TF-IDF with unigrams and bigrams, English stop words, cosine similarity, top 4 chunks with score above 0 [7]"],
            ["Extractive answer", "split chunks into sentences, TF-IDF rank against the question, keep the top 4 with score above 0, join"],
            ["Optional LLM answer", "Claude receives the question and chunks, told to answer only from them; any error falls back to the extractive answer"],
            ["Quiz", "top 40 TF-IDF terms; blank the first whole-word occurrence in a shuffled sentence; 3 other terms as options; 5 questions"],
            ["Roadmap", "accuracy per topic: High below 50 percent, Medium below 80, otherwise Low; sorted by priority then accuracy"]]
    K.table(s, 0.6, 1.4, [2.3, 6.6], rows, font_size=13, row_h=0.8, first_bold=True)
    K.tile(s, 9.9, 1.4, 2.83, 1.45, "0", "trained parameters in the shipped code", K.BLUE, vsize=28, lsize=11)
    K.tile(s, 9.9, 3.0, 2.83, 1.45, "1-2 grams", "TF-IDF vocabulary fitted per user", K.GREEN, vsize=22, lsize=11)
    K.tile(s, 9.9, 4.6, 2.83, 1.45, "CPU only", f"{F.SCALE[0]['index_build_ms']:.0f} ms to index {F.SCALE[0]['chunks']} chunks", K.ORANGE, vsize=22, lsize=11)
    notes[8] = "SAC-RAG classifies the question with an LLM and uses BM25 plus dense retrieval and an LLM compressor [1]. We use none of these: retrieval is lexical and compression is sentence selection."

    # ---- 9 experiment configuration
    s = S[8]
    set_title(s, "Experiment configuration (nothing is trained)")
    wipe(s)
    rows = [["Setting", "Value"],
            ["Corpus", f"6 subject files, 36 articles, {F.CORPUS['words']:,} words, {F.CORPUS['chunks']} chunks (180 / 30)"],
            ["Questions", "72 hand-written: 24 fact, 24 definition, 24 reasoning; 12 off-topic"],
            ["Metrics", "Hit@k (evidence inside top-k chunks), MRR@10, answer-key recall, words removed"],
            ["Retrievers", "7: shipped TF-IDF, unigram, BM25 [6], LSA [8], BM25 + LSA, dense BGE-small [9], BM25 + dense"],
            ["Quiz study", "30 seeds x 6 files x 5 questions = 900; 30 reviewed by hand"],
            ["Hardware", f"{F.ENV['cpu']}, {F.ENV['ram']}, no GPU; Python {F.ENV['python']}, scikit-learn {F.ENV['sklearn']} [16]"]]
    K.table(s, 0.6, 1.4, [1.9, 6.0], rows, font_size=12, row_h=0.72, first_bold=True)
    K.picture(s, F.fig("fig_chunk_size.png"), 8.7, 1.4, 4.0, 2.5)
    K.text(s, 8.7, 4.0, 4.0, 2.9, f"Chunk size trade-off for the shipped retriever: Hit@4 rises from {F.CHUNKS['60/15']['hit4']:.3f} (60 words) to {F.CHUNKS['180/30']['hit4']:.3f} (180, shipped) and {F.CHUNKS['300/50']['hit4']:.3f} (300), but larger chunks carry more words per hit.", size=13, color=K.TXT)
    notes[9] = "The product has no training stage. Experiments fix seeds (quiz seeds 0 to 29) and the evidence strings, so every table can be regenerated with scripts/run_experiments.py."

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic")
    wipe(s)
    K.picture(s, F.diag("12_flowchart_rag.png"), 0.6, 1.3, 5.2, 5.6)
    K.text(s, 6.2, 1.5, 6.5, 5.3, [
        ("**Refuse only if nothing overlaps.** An empty question is rejected; a question with no shared term gets 'No relevant content found'.", {"bullet": True, "space_after": 9}),
        ("**No relevance threshold.** 9 of 12 off-topic questions still received an answer; a threshold near 0.12 to 0.15 would refuse 75 to 83 percent of them.", {"bullet": True, "space_after": 9}),
        ("**LLM is optional.** With a key the answer comes from Claude, with a context-only instruction; otherwise (or on any error) the extractive answer is used.", {"bullet": True, "space_after": 9}),
        ("**Roadmap priority.** Below 50 percent High, below 80 percent Medium, else Low.", {"bullet": True}),
    ], size=17, color=K.INK)
    notes[10] = "This is the flowchart of RAGEngine.generate_answer. The abstention gap is a measured weakness reported in the results."

    # ---- 11 data flow and modules
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 6.2, 5.6)
    rows = [["Module", "Role"], ["utils/text_extraction.py", "extract, clean, chunk, topics"], ["utils/rag_engine.py", "TF-IDF index, retrieve, answer"], ["utils/quiz_generator.py", "cloze quiz questions"],
            ["db.py", "SQLite schema, 5 tables"], ["backend/main.py", "14 FastAPI routes"], ["app.py + templates/", "original Flask pages"], ["frontend/src", "React pages and charts"]]
    K.table(s, 7.1, 1.5, [2.6, 3.03], rows, font_size=12, row_h=0.55)
    notes[11] = "Level-1 data flow. The Python modules are the original project code; the FastAPI layer and the React front end are the new layers."

    # ---- 12 database and UML
    s = S[11]
    set_title(s, "Database design and UML")
    wipe(s)
    K.picture(s, F.diag("11_er_diagram.png"), 0.6, 1.35, 6.0, 5.1)
    K.picture(s, F.diag("07_uml_class_diagram.png"), 6.8, 1.35, 5.9, 5.1)
    K.text(s, 0.6, 6.55, 12.13, 0.4, "SQLite with five tables (users, documents, chunks, queries, quiz_results), foreign keys with cascade delete; the uploaded files live on disk under uploads/<user>/.", size=12, color=K.TXT, align="c", italic=True)
    notes[12] = "quiz_results keeps subject and topic as text, so the roadmap survives deleting a document."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    E = F.ENV
    rows = [["Layer", "Technology"], ["Retrieval and analysis", f"scikit-learn {E['sklearn']} (TF-IDF, cosine), NumPy {E['numpy']}, pypdf {E['pypdf']}, python-docx {E['python-docx']}"],
            ["API", f"FastAPI {E['fastapi']}, Uvicorn {E['uvicorn']}, Pydantic, Starlette sessions (port 8301)"],
            ["Original UI", f"Flask {E['flask']}, Jinja templates, Bootstrap (port 5000)"],
            ["Front end", f"React {E['react']}, TypeScript, Vite {E['vite']}, Tailwind {E['tailwind']}, TanStack Query, Recharts, framer-motion (port 5301)"],
            ["Storage", "SQLite (5 tables) and uploaded files on disk"],
            ["Optional generation", "Anthropic API (Claude), only when ANTHROPIC_API_KEY is set"],
            ["Testing", "pytest (84 tests), API smoke test, Playwright for screenshots"]]
    K.table(s, 0.6, 1.5, [3.0, 9.13], rows, font_size=15, row_h=0.62, first_bold=True)
    notes[13] = f"Python {E['python']} on Windows 11. run.bat creates the virtual environment, installs packages and starts both servers."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Upload, extraction, chunking, subject and topic hints", "Done", "36 of 36 unit headings found"],
            ["TF-IDF retrieval, extractive answer with sources", "Done", f"Hit@4 {h(SH, 4):.3f}"],
            ["Quiz, grading, roadmap, dashboard", "Done", f"{F.QUIZ['generated']} of {F.QUIZ['requested']} questions generated"],
            ["FastAPI (14 endpoints) and React interface", "Done", f"{F.SMOKE_TOTAL} of {F.SMOKE_TOTAL} smoke checks"],
            ["Claude answer mode", "Coded", "not measured, no API key"],
            ["Dense embeddings, question-type routing, LLM compression", "Not done", "SAC-RAG features absent"],
            ["Relevance threshold, revision schedule", "Not done", "9 of 12 off-topic answered"]]
    K.table(s, 0.6, 1.5, [5.6, 1.4, 5.13], rows, font_size=14, row_h=0.62)
    notes[14] = "State of the code base. The honest gaps are the missing semantic retrieval and compression stages of the base paper, and the missing relevance threshold."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: retrieval and compression")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{h(SH, 4):.3f}", "Hit@4, shipped TF-IDF", K.BLUE), (f"{h('bm25', 4):.3f}", "Hit@4, BM25", K.GREEN),
                                   (f"{h('hybrid_bm25_dense', 4):.3f}", "Hit@4, BM25 + dense (best)", K.PURPLE), (f"{C4['token_reduction_pct']:.1f}%", "words removed by sentence selection", K.ORANGE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=12)
    K.picture(s, F.fig("fig_retrieval_methods.png"), 0.6, 2.95, 7.6, 3.9)
    K.card(s, 8.5, 2.95, 4.23, 3.9, "Reading",
           f"72 questions, 672 chunks. The shipped retriever finds the evidence at rank 1 in {F.pct(h(SH, 1), 0)} of questions. Sentence selection cuts 720 words to about 90 but answer-key recall falls from {C4['key_in_ctx']:.3f} to {C4['key_in_ans']:.3f}. SAC-RAG reports 38 to 58 percent fewer tokens for 1 to 2 points of EM on other data [1].",
           K.BLUE, "i", body_size=13)
    notes[15] = "Bootstrap 95 percent intervals are wide with 72 questions (Hit@4 0.69 to 0.88), so only the larger gaps matter: BM25 over TF-IDF at Hit@1 and the hybrid being best. The comparison with SAC-RAG is of the trade-off, not of the numbers: different data and no LLM."

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    K.picture(s, F.shot("09_ask.png"), 0.6, 1.3, 7.4, 4.65, border=True)
    K.picture(s, F.shot("07_dashboard.png"), 8.3, 1.3, 4.43, 2.77, border=True)
    K.picture(s, F.shot("13_roadmap.png"), 8.3, 4.2, 4.43, 2.77, border=True)
    K.text(s, 0.6, 6.05, 7.4, 0.9, "Ask your notes: 'What is a deadlock?' answered from 02_Operating_Systems.txt, passage 8 (match 0.46), offline extractive mode. Right: dashboard with subject accuracy and the study roadmap. Real application, five uploaded documents.", size=12, color=K.TXT)
    notes[16] = "Real screenshots from the running application captured with Playwright. The quiz answers behind the dashboard were scripted to simulate a student with different mastery per subject."

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{F.TESTS['tests'] - F.TESTS['failures'] - F.TESTS['errors']} / {F.TESTS['tests']}", "pytest tests passed", K.GREEN), (f"{F.SMOKE_TOTAL} / {F.SMOKE_TOTAL}", "API smoke checks", K.BLUE),
                                   (f"{A['ask']['median_ms']:.0f} ms", f"median POST /api/ask ({F.CORPUS['chunks']} chunks)", K.ORANGE), (f"{A['quiz_generate']['median_ms']:.0f} ms", "median quiz generation", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    rows = [["Observation", "Action"],
            ["9 of 12 off-topic questions still get an answer", "Add a relevance threshold (0.12 to 0.15)"],
            [f"Definition questions: key in the chunks {F.CMP['by_type_k4']['definition']['key_in_ctx']:.3f}, in the answer {F.CMP['by_type_k4']['definition']['key_in_ans']:.3f}", "Rank sentences for definitions, not word overlap"],
            [f"Quiz answer still visible in the question stem in {F.pct(F.QUIZ['answer_repeated_in_stem_rate'])} of questions", "Blank every occurrence of the term"],
            ["Two uploads with one file name share one file on disk", "Store files under unique names"]]
    K.table(s, 0.6, 3.05, [6.0, 6.13], rows, font_size=14, row_h=0.72)
    notes[17] = "84 tests: 45 API, 15 text extraction, 11 security, 9 retrieval, 4 quiz. Latency measured on 127.0.0.1 with one user; quiz submit showed 1 to 2 second outliers on 1 or 2 of 30 calls, cause not isolated."

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["No real student documents or ground truth to measure against", "Built a 36-article corpus and 72 questions with verifiable evidence strings"],
            [f"Lexical retrieval misses paraphrases (Hit@1 {h(SH, 1):.3f})", "Measured BM25, LSA, dense and hybrid retrievers as reference points"],
            ["Math markup in the Wikipedia text polluted topics and chunks", "Stripped it and regenerated the corpus from pinned revisions"],
            ["Roadmap topics are single keywords, some generic", "Reported; unit headings as topics listed as future work"],
            ["Timings fluctuate on a shared PC", "Medians of repeated runs, two API runs reported"]]
    K.table(s, 0.6, 1.5, [5.6, 6.53], rows, font_size=15, row_h=0.85)
    notes[18] = "Each challenge led to a concrete change in the measurement setup or in how results are reported."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Retrieve better", "Add BM25 and a dense model in a hybrid (Hit@4 0.889 in our test) and route questions by type as in [1][11].", K.BLUE, "1"),
        ("Stay grounded", "Add a relevance threshold, measure the Claude answer mode and improve sentence choice for definitions.", K.GREEN, "2"),
        ("Better practice", "Blank all term occurrences, use unit headings as roadmap topics, add knowledge tracing [17].", K.ORANGE, "3"),
        ("Validate", "Test with real student syllabi in several heading styles and run a small user study.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.1, 1.6, 2.9, 4.6, t, b_, col, gl, body_size=18)
    notes[19] = "Plan to the final review and beyond; items 1 and 2 close the largest gaps to the base paper and to the stated objectives."

    # ---- 20 thank you
    thank_you(S[19])
    notes[20] = "Likely questions: why Wikipedia text; why not embeddings in the product; how answers stay grounded; what is missing against SAC-RAG."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "Adaptive Learning")
    return out_path


if __name__ == "__main__":
    p = build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx"))
    print(p, scan_deck(p))
