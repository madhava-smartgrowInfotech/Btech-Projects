"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p("A document-grounded study platform: syllabus and notes analysis, source-cited retrieval-augmented answers, auto-generated quizzes and a priority-ranked study roadmap", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"The undersigned hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"The project team expresses its sincere gratitude to the project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"The team thanks {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given.")
    d.p("Thanks are due to the project coordinators and all faculty and staff members of the department for their help, and to families and friends for their constant support.")
    d.p("The team also acknowledges the authors of the base paper and the supporting papers, the contributors to the English Wikipedia articles used to build the evaluation corpus, and the developers of the "
        "open-source software used in this project (scikit-learn, NumPy, pypdf, python-docx, Flask, FastAPI, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    sh = F.SHIPPED
    c4 = F.CMP["4"]
    ab = F.ABS
    d.p("Students handle scattered PDFs, syllabi, lecture notes and question banks across six to eight subjects every semester, yet learning portals store these files without understanding what is "
        "inside them, and general chatbots answer from public data rather than the prescribed syllabus. This project presents the Adaptive Learning Intelligence Platform, a web application in "
        "which a student uploads study material, the system extracts subjects, units and topics into a searchable repository, answers questions from that repository with the source passage shown, "
        "generates fill-in-the-blank quizzes, and turns quiz accuracy per topic into a priority-ranked roadmap and a dashboard. It follows the retrieve-then-answer idea of the SAC-RAG base paper, "
        "but retrieval is lexical (TF-IDF cosine similarity), answers are extractive and run offline, and an optional Claude mode exists that was not measured.")
    d.p(f"The platform was measured on a corpus of {F.CORPUS['words']:,} words from 36 Wikipedia articles arranged as six subjects, with {F.CORPUS['questions']} hand-written questions. The shipped retriever places the "
        f"answer sentence in the top four chunks for {F.pct(F.hit(sh, 4))} of questions (Hit@1 {F.pct(F.hit(sh, 1))}); BM25 reaches {F.pct(F.hit('bm25', 4))} and a BM25 plus dense hybrid {F.pct(F.hit('hybrid_bm25_dense', 4))}. "
        f"Extractive sentence selection removes {c4['token_reduction_pct']:.1f}% of the retrieved words but lowers answer-key recall from {c4['key_in_ctx']:.3f} to {c4['key_in_ans']:.3f}, and "
        f"{ab['off_topic_answered_by_shipped']} of {ab['n_off_topic']} off-topic questions still receive an answer because no relevance threshold exists. A question is answered in about "
        f"{F.API['ask']['median_ms']:.0f} ms. {F.TESTS['tests']} automated tests and a {F.SMOKE_TOTAL}-check smoke test pass. Learned embeddings, question-type routing, LLM compression, revision scheduling "
        "and coverage analytics are not implemented, and are the main future steps.")
    d.p("**Keywords:** adaptive learning, educational document analysis, retrieval-augmented generation (RAG), learning analytics, AI tutoring, personalised roadmaps, academic intelligence.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("BM25", "Best Matching 25 (probabilistic ranking function)"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CPU", "Central Processing Unit"), ("DFD", "Data Flow Diagram"), ("DOCX", "Office Open XML Word document"), ("EM", "Exact Match"),
    ("ER", "Entity-Relationship"), ("FK / PK", "Foreign Key / Primary Key"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"),
    ("LLM", "Large Language Model"), ("LSA", "Latent Semantic Analysis"), ("MRR", "Mean Reciprocal Rank"), ("PDF", "Portable Document Format"), ("RAG", "Retrieval-Augmented Generation"),
    ("REST", "Representational State Transfer"), ("SAC-RAG", "Semantic Adaptive Compress RAG"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"),
    ("SVD", "Singular Value Decomposition"), ("TF-IDF", "Term Frequency - Inverse Document Frequency"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=10)
