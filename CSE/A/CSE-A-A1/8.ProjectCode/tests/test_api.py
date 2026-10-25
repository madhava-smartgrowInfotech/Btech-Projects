import uuid

import pytest
from fastapi.testclient import TestClient

from backend.main import app as api_app
from conftest import upload
from db import get_db, now


def uname():
    return "u_" + uuid.uuid4().hex[:10]


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json() == {"status": "ok", "llm_enabled": False}


# ------------------------------------------------------------------ auth
def test_register_login_me_logout(client):
    n = uname()
    assert client.get("/api/auth/me").json() is None
    assert client.post("/api/auth/register", json={"username": n, "password": "x1", "interests": "AI"}).status_code == 201
    r = client.post("/api/auth/login", json={"username": n, "password": "x1"})
    assert r.status_code == 200 and r.json()["username"] == n and r.json()["interests"] == "AI"
    assert client.get("/api/auth/me").json()["username"] == n
    assert client.post("/api/auth/logout").status_code == 200
    assert client.get("/api/auth/me").json() is None


def test_register_validation(client):
    assert client.post("/api/auth/register", json={"username": "", "password": "x"}).status_code == 400
    assert client.post("/api/auth/register", json={"username": "abc", "password": ""}).status_code == 400
    assert client.post("/api/auth/register", json={"username": "   ", "password": "x"}).status_code == 400


def test_duplicate_username_rejected(client):
    n = uname()
    assert client.post("/api/auth/register", json={"username": n, "password": "x"}).status_code == 201
    assert client.post("/api/auth/register", json={"username": n, "password": "y"}).status_code == 409


def test_wrong_password_and_unknown_user(client):
    n = uname()
    client.post("/api/auth/register", json={"username": n, "password": "right"})
    assert client.post("/api/auth/login", json={"username": n, "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/login", json={"username": "nobody_" + n, "password": "right"}).status_code == 401


PROTECTED = [("get", "/api/dashboard"), ("get", "/api/documents"), ("get", "/api/ask/history"), ("get", "/api/roadmap"),
             ("post", "/api/ask"), ("post", "/api/quiz/generate"), ("post", "/api/quiz/submit"),
             ("delete", "/api/documents/1"), ("post", "/api/documents")]


@pytest.mark.parametrize("method,path", PROTECTED)
def test_protected_routes_need_login(method, path):
    with TestClient(api_app) as anon:
        r = getattr(anon, method)(path, **({"json": {}} if method == "post" and path != "/api/documents" else {}))
        assert r.status_code == 401


# ------------------------------------------------------------------ documents
def test_upload_txt_creates_document_topics_and_chunks(user, sample_bytes):
    r = upload(user["client"], sample_bytes, "sample.txt", "Data Structures")
    assert r.status_code == 201
    body = r.json()
    assert body["document"]["subject"] == "Data Structures" and body["chunks_count"] == 4
    assert body["topics_count"] == 12 and body["warning"] is None
    docs = user["client"].get("/api/documents").json()
    assert len(docs) == 1 and docs[0]["chunk_count"] == 4 and "filepath" not in docs[0]


def test_upload_default_subject_is_general(user):
    r = upload(user["client"], b"A short paragraph about graphs and trees and their traversal order.", "g.txt", "")
    assert r.json()["document"]["subject"] == "General"


def test_upload_rejects_unsupported_type(user):
    assert upload(user["client"], b"MZ...", "virus.exe").status_code == 400


def test_upload_rejects_files_over_16_mb(user):
    r = upload(user["client"], b"a " * (9 * 1024 * 1024), "big.txt")
    assert r.status_code == 413


def test_upload_empty_text_gives_warning(user):
    r = upload(user["client"], b"", "empty.txt")
    assert r.status_code == 201 and r.json()["chunks_count"] == 0 and "Could not extract" in r.json()["warning"]


def test_upload_docx(user, tmp_path):
    import docx
    p = tmp_path / "n.docx"
    d = docx.Document()
    d.add_paragraph("Subject: Networks")
    d.add_paragraph("TCP is a connection oriented protocol that guarantees ordered delivery of bytes. " * 5)
    d.save(str(p))
    r = upload(user["client"], p.read_bytes(), "n.docx", "Networks")
    assert r.status_code == 201 and r.json()["chunks_count"] >= 1


def test_delete_document_removes_chunks_and_file(user, sample_bytes):
    r = upload(user["client"], sample_bytes, "del.txt")
    doc_id = r.json()["document"]["id"]
    assert user["client"].delete(f"/api/documents/{doc_id}").status_code == 200
    assert user["client"].get("/api/documents").json() == []
    assert user["client"].delete(f"/api/documents/{doc_id}").status_code == 404
    conn = get_db()
    assert conn.execute("SELECT COUNT(*) c FROM chunks WHERE document_id=?", (doc_id,)).fetchone()["c"] == 0
    conn.close()


def test_users_cannot_see_or_delete_each_others_documents(client, sample_bytes):
    a, b = uname(), uname()
    for n in (a, b):
        client.post("/api/auth/register", json={"username": n, "password": "p"})
    client.post("/api/auth/login", json={"username": a, "password": "p"})
    doc_id = upload(client, sample_bytes, "a.txt").json()["document"]["id"]
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": b, "password": "p"})
    assert client.get("/api/documents").json() == []
    assert client.delete(f"/api/documents/{doc_id}").status_code == 404
    assert client.post("/api/quiz/generate", json={"doc_id": doc_id}).status_code == 404
    assert client.post("/api/ask", json={"question": "What is the LIFO principle?"}).json()["sources"] == []


# ------------------------------------------------------------------ ask
def test_ask_requires_text(user):
    assert user["client"].post("/api/ask", json={"question": "   "}).status_code == 400


def test_ask_without_documents_says_so(user):
    r = user["client"].post("/api/ask", json={"question": "What is the LIFO principle?"})
    assert r.status_code == 200 and r.json()["answer"].startswith("No relevant content") and r.json()["sources"] == []


def test_ask_answers_from_the_upload_with_sources_and_history(user, sample_bytes):
    c = user["client"]
    upload(c, sample_bytes, "sample_notes.txt")
    r = c.post("/api/ask", json={"question": "What is the LIFO principle?"}).json()
    assert "Last In First Out" in r["answer"] and r["llm_enabled"] is False
    assert r["sources"] and r["sources"][0]["doc_name"] == "sample_notes.txt"
    hist = c.get("/api/ask/history").json()
    assert hist[0]["question"] == "What is the LIFO principle?" and hist[0]["sources"] == ["sample_notes.txt"] * len(r["sources"])


def test_history_is_capped_at_ten(user, sample_bytes):
    c = user["client"]
    upload(c, sample_bytes)
    for i in range(12):
        c.post("/api/ask", json={"question": f"What is a stack {i}?"})
    assert len(c.get("/api/ask/history").json()) == 10


def test_deleting_a_document_removes_it_from_answers(user, sample_bytes):
    c = user["client"]
    doc_id = upload(c, sample_bytes).json()["document"]["id"]
    c.delete(f"/api/documents/{doc_id}")
    assert c.post("/api/ask", json={"question": "What is a queue?"}).json()["sources"] == []


# ------------------------------------------------------------------ quiz + roadmap + dashboard
def test_quiz_flow_grades_and_feeds_roadmap_and_dashboard(user, sample_bytes):
    c = user["client"]
    doc_id = upload(c, sample_bytes, subject="Data Structures").json()["document"]["id"]
    q = c.post("/api/quiz/generate", json={"doc_id": doc_id}).json()
    assert len(q["questions"]) == 5 and all("answer_index" not in x for x in q["questions"])
    res = c.post("/api/quiz/submit", json={"answers": [None] * 5}).json()
    assert res["correct"] == 0 and res["total"] == 5 and res["score_pct"] == 0
    assert all(r["is_correct"] is False for r in res["results"])
    road = c.get("/api/roadmap").json()
    assert len(road) == 5 and {r["priority"] for r in road} == {"High"}
    dash = c.get("/api/dashboard").json()
    assert dash["doc_count"] == 1 and dash["quiz_attempts"] == 5
    assert dash["performance"] == [{"subject": "Data Structures", "accuracy": 0.0, "attempts": 5}]


def test_quiz_submit_with_correct_answers_scores_100(user, sample_bytes):
    c = user["client"]
    doc_id = upload(c, sample_bytes).json()["document"]["id"]
    c.post("/api/quiz/generate", json={"doc_id": doc_id})
    from backend.main import _active_quizzes
    answers = [q["answer_index"] for q in _active_quizzes[user["id"]]["questions"]]
    res = c.post("/api/quiz/submit", json={"answers": answers}).json()
    assert res["correct"] == 5 and res["score_pct"] == 100
    assert {r["priority"] for r in c.get("/api/roadmap").json()} == {"Low"}


def test_quiz_submit_without_active_quiz_conflicts(user):
    assert user["client"].post("/api/quiz/submit", json={"answers": []}).status_code == 409


def test_quiz_cannot_be_submitted_twice(user, sample_bytes):
    c = user["client"]
    doc_id = upload(c, sample_bytes).json()["document"]["id"]
    c.post("/api/quiz/generate", json={"doc_id": doc_id})
    assert c.post("/api/quiz/submit", json={"answers": [0] * 5}).status_code == 200
    assert c.post("/api/quiz/submit", json={"answers": [0] * 5}).status_code == 409


def test_quiz_on_tiny_document_is_refused(user):
    c = user["client"]
    doc_id = upload(c, b"Too short to quiz.", "tiny.txt").json()["document"]["id"]
    assert c.post("/api/quiz/generate", json={"doc_id": doc_id}).status_code == 422


def test_partial_answers_are_graded_as_wrong(user, sample_bytes):
    c = user["client"]
    doc_id = upload(c, sample_bytes).json()["document"]["id"]
    c.post("/api/quiz/generate", json={"doc_id": doc_id})
    res = c.post("/api/quiz/submit", json={"answers": [0]}).json()
    assert res["total"] == 5 and res["correct"] <= 1


@pytest.mark.parametrize("correct,total,priority", [
    (0, 1, "High"), (1, 3, "High"), (49, 100, "High"), (1, 2, "Medium"), (2, 4, "Medium"), (7, 9, "Medium"),
    (79, 100, "Medium"), (4, 5, "Low"), (80, 100, "Low"), (5, 5, "Low"),
])
def test_roadmap_priority_thresholds(user, correct, total, priority):
    conn = get_db()
    for i in range(total):
        conn.execute("INSERT INTO quiz_results (user_id, subject, topic, correct, created_at) VALUES (?,?,?,?,?)",
                     (user["id"], "S", "topic", 1 if i < correct else 0, now()))
    conn.commit()
    conn.close()
    item = user["client"].get("/api/roadmap").json()[0]
    assert item["priority"] == priority and item["attempts"] == total


def test_roadmap_is_sorted_high_first_then_by_accuracy(user):
    conn = get_db()
    rows = [("a", 1, 1), ("b", 0, 1), ("c", 1, 2), ("d", 4, 5), ("e", 0, 4)]
    for topic, ok, n in rows:
        for i in range(n):
            conn.execute("INSERT INTO quiz_results (user_id, subject, topic, correct, created_at) VALUES (?,?,?,?,?)",
                         (user["id"], "S", topic, 1 if i < ok else 0, now()))
    conn.commit()
    conn.close()
    order = [r["topic"] for r in user["client"].get("/api/roadmap").json()]
    assert order == ["b", "e", "c", "d", "a"]


def test_dashboard_for_new_user_is_zero(user):
    d = user["client"].get("/api/dashboard").json()
    assert (d["doc_count"], d["query_count"], d["quiz_attempts"], d["performance"], d["documents"]) == (0, 0, 0, [], [])
