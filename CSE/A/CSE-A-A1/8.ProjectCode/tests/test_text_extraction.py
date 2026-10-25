import docx

from conftest import minimal_pdf
from utils.text_extraction import (allowed_file, chunk_text, clean_text, detect_subject_hints, extract_text,
                                   extract_topics)


def test_extract_txt(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("Subject: Maths\nHello world", encoding="utf-8")
    assert "Hello world" in extract_text(str(p))


def test_extract_docx_paragraphs(tmp_path):
    p = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("Unit 1: Sets")
    d.add_paragraph("A set is a collection of objects.")
    d.save(str(p))
    assert "A set is a collection" in extract_text(str(p))


def test_extract_docx_ignores_table_cells(tmp_path):
    """Known limitation: only paragraphs are read, text inside Word tables is lost."""
    p = tmp_path / "t.docx"
    d = docx.Document()
    d.add_paragraph("Outside the table")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "Only inside the table"
    d.save(str(p))
    text = extract_text(str(p))
    assert "Outside the table" in text and "Only inside the table" not in text


def test_extract_pdf_text_layer(tmp_path):
    p = tmp_path / "a.pdf"
    p.write_bytes(minimal_pdf("Stacks follow the LIFO principle"))
    assert "LIFO" in extract_text(str(p))


def test_extract_pdf_without_text_layer_is_empty(tmp_path):
    p = tmp_path / "scan.pdf"
    p.write_bytes(minimal_pdf(""))
    assert extract_text(str(p)).strip() == ""


def test_extract_corrupt_pdf_does_not_raise(tmp_path):
    p = tmp_path / "bad.pdf"
    p.write_bytes(b"this is not a pdf")
    assert extract_text(str(p)).startswith("[Could not extract text")


def test_clean_text_collapses_whitespace():
    assert clean_text("a   b\t c\n\n\n\nd") == "a b c\n\nd"


def test_chunk_text_size_and_overlap():
    words = ["w%d" % i for i in range(500)]
    chunks = chunk_text(" ".join(words), chunk_size=180, overlap=30)
    assert [len(c.split()) for c in chunks] == [180, 180, 180, 50]
    assert chunks[0].split()[-30:] == chunks[1].split()[:30]


def test_chunk_text_edge_cases():
    assert chunk_text("") == []
    assert chunk_text("   ") == []
    assert chunk_text("one two three") == ["one two three"]
    assert len(chunk_text(" ".join(["x"] * 180))) == 1


def test_detect_subject_hints_standard_headings():
    text = "Subject: Data Structures\nUnit 1: Arrays\nUnit 2 - Stacks\nChapter 3: Trees\nCourse: CS201"
    assert set(detect_subject_hints(text)) == {"Data Structures", "Arrays", "Stacks", "Trees", "CS201"}


def test_detect_subject_hints_deduplicates_and_caps():
    text = "\n".join("Unit %d: Topic %d" % (i, i) for i in range(1, 40)) + "\nUnit 1: Topic 1"
    hints = detect_subject_hints(text)
    assert len(hints) == 15 and len({h.lower() for h in hints}) == 15


def test_detect_subject_hints_roman_numerals_not_recognised():
    """Known limitation: 'UNIT-I' / 'Module 1' style headings are not matched."""
    assert detect_subject_hints("UNIT-I Arrays\nUnit I: Stacks\nModule 1: Trees") == []


def test_extract_topics_prefers_hints_then_keywords():
    sample = open("sample_data/sample_notes_data_structures.txt", encoding="utf-8").read()
    topics = extract_topics(sample)
    assert topics[0] == "Data Structures" and len(topics) == 12
    assert "Introduction to Data Structures" in topics


def test_extract_topics_short_text_falls_back():
    assert extract_topics("tiny") == []


def test_allowed_file():
    ok = {"pdf", "docx", "txt"}
    assert allowed_file("a.PDF", ok) and allowed_file("x.y.txt", ok)
    assert not allowed_file("a.exe", ok) and not allowed_file("noext", ok) and not allowed_file("a.txt.exe", ok)
