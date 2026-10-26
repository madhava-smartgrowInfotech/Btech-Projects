"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (letterhead, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "a10_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def lat(prefix):
    for r in F.API["results"]:
        if r["endpoint"].startswith(prefix):
            return r
    raise KeyError(prefix)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "smart traffic, YOLOv8, adaptive signal control, emergency priority, queue forecasting"
    notes = {}
    bp = F.BASE_PAPER_RESULTS
    d2, d3 = F.derived("2x2"), F.derived("3x3")
    fx2, ad2, em2, av2 = (F.gm("2x2", k, "Avg delay / vehicle (s)") for k in F.CTRL_KEYS)
    fx3, ad3, em3, av3 = (F.gm("3x3", k, "Avg delay / vehicle (s)") for k in F.CTRL_KEYS)
    efx2, eav2 = F.gm("2x2", "fixed", "Emergency delay (s)"), F.gm("2x2", "advanced", "Emergency delay (s)")
    efx3, eav3 = F.gm("3x3", "fixed", "Emergency delay (s)"), F.gm("3x3", "advanced", "Emergency delay (s)")
    mt, amb = F.VR["mixed_traffic"], F.VR["ambulance_classification"]
    pc = mt["per_class"]
    lb, lby = amb["lightbar_cluster_any"], amb["ambulance_flag_yolo_plus_lightbar"]
    vc = F.VS["results"]["clean"]
    pr_rand, pr_loso, sh = F.PR["random_split_5x"], F.PR["leave_one_simulation_out"], F.PR["shipped_fit"]["metrics"]
    ab = F.AB["results"]
    ab_full, ab_noml = ab["advanced (full)"], ab["no ML forecast"]

    for sl in S:  # footers on every slide except the title
        for sh_ in sl.shapes:
            if sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 28)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and FlowSense: one pipeline that detects vehicles, forecasts queues, adapts signal timing, gives ambulances priority and coordinates neighbouring junctions."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Data & Controller Logic\nInteraction & Data Design\nImplementation Highlights\nResults: Controllers, Emergency Priority, Vision, Forecaster\nParameter Analysis, Performance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Limitations, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Fixed-time signals cannot see queues or ambulances. FlowSense puts vehicle detection, queue forecasting, adaptive timing, emergency priority and inter-junction coordination into one working web product (FastAPI + React).", {"bullet": True, "space_after": 8}),
        (f"A seeded simulator runs four controllers on identical traffic. The advanced controller cuts average delay {d2['delay_cut_advanced_vs_fixed_pct']:.1f}% (2x2) and {d3['delay_cut_advanced_vs_fixed_pct']:.1f}% (3x3) against fixed-time and ambulance delay from {efx2[0]:.1f} s to {eav2[0]:.1f} s (2x2).", {"bullet": True, "space_after": 8}),
        (f"Honestly: it beats plain adaptive control by only {d2['delay_cut_advanced_vs_adaptive_pct']:.1f}% (within one standard deviation), and a base-paper-style pre-emption controller is {-d2['delay_cut_basepaper_vs_adaptive_pct']:.1f}% worse than adaptive on mean delay.", {"bullet": True, "space_after": 8}),
        (f"The 30-second queue forecaster scores R\u00b2 {pr_rand['Gradient Boosting (used)']['r2_mean']:.3f} on a random split but fails on an unseen heavily congested run (MAE 14.66 vs 5.49 for persistence).", {"bullet": True, "space_after": 8}),
        (f"Vision works on synthetic scenes (recall {vc['recall'] * 100:.1f}%) but not on real photos: light-bar recall {lb['recall'] * 100:.1f}% on 600 real images; YOLOv8n 4-wheeler recall {pc['4-wheeler']['recall'] * 100:.1f}%.", {"bullet": True, "space_after": 8}),
        (f"{F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks pass. No SUMO, no trained light-bar detector, no real cameras: the limits are stated on every results slide.", {"bullet": True}),
    ], size=17)
    notes[3] = "Give the product, the headline gains, then the honest limits. Say that the simulator is not SUMO and the vision gap on real photos."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: emergency-aware adaptive signal control with YOLO light-bar detection and SUMO validation"],
            ["Detector", f"YOLO26m, one class (roof light-bar, lights ON or OFF), 3,130 images (2,599 / 266 / 265), 100 epochs; precision {bp['precision']:.2f}, recall {bp['recall']:.2f}, mAP50 {bp['map50']:.2f}, mAP50-95 {bp['map5095']:.2f}; 11.45-13.89 ms on a T4"],
            ["Control", "Pre-optimised GA-SGD signal plan; a priority arbitrator (confidence 0.5, 3 consecutive frames) forces green for the emergency approach, then restores the plan"],
            ["Validation", "SUMO model of the Souissi roundabout, Rabat; Fixed-time vs GA-SGD vs GA-SGD + vision override; selected scenario at demand scale 12"],
            ["Reported", "Ambulance travel 56 s vs 70 s (GA-SGD) and 64 s (fixed); ambulance wait 15 s vs 21 s; total waiting 19,887 s vs 22,827 s and 33,719 s"],
            ["Stated limits", "Moderate dataset; simulation only; one roundabout; light-bar presence only (no flashing-state check); edge hardware not tested"]]
    ftable(s, 0.55, 2.6, [1.7, 10.5], rows, size=12, row_h=0.7)
    notes[4] = "Lamrabet et al., IEEE Access 2026. We keep its idea (a vision cue triggering emergency priority over an adaptive controller) but extend it to forecasting and corridor coordination; its trained detector and SUMO case are not reproduced."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Finding / limitation"],
            [f"Lamrabet et al. [{R['base']}] (base)", "2026", "YOLO26m light-bar detector + priority override on GA-SGD", "SUMO roundabout: ambulance travel 56 s vs 70 s; one network"],
            [f"Goenawan, ASTM [{R['astm']}]", "2024", "YOLOv5 counts + RNN-LSTM 12-hour forecast + cycle-length rule", "CARLA: 21 vs ~15 vehicles/min, 5 s vs ~12 s pass delay; no emergency priority"],
            [f"Gupta et al., NeCDM [{R['necdm']}]", "2026", "Neighbour-embedded GNN congestion prediction for crowd delivery", "Predicts flow at stations; no signal control or vision"],
            [f"Sasikala et al. [{R['sasikala']}]", "2025", "IoT + GPS + RFID emergency priority and dynamic routes", "35 % faster ambulance response (as cited in [1]); needs infrastructure"],
            [f"Zlatkovic et al. [{R['zlatkovic']}]", "2023", "Pre-emption in connected-vehicle environments (VISSIM)", "Up to 36 % emergency delay cut (as cited in [1]); needs connected vehicles"],
            [f"Jayakumar et al. [{R['jayakumar']}]", "2024", "Siren audio features with SVM and ensembles", "99.5 % accuracy (as cited in [1]); noise, no localisation"],
            [f"YOLO emergency / vehicle detectors [{R['phuong']}], [{R['suhana']}]", "2024-25", "YOLOv5-v9 vehicle and emergency-vehicle classes", "Appearance-based; small objects and occlusion are hard (as cited in [1])"],
            [f"Batch survey [{R['s_signaliq']}]-[{R['s_incip']}]", "2024-26", "Signal-IQ, density estimation, accelerated EV detection, client-server control", "Title-level only; authors and results not verified"]]
    ftable(s, 0.55, 1.8, [3.1, 0.9, 4.2, 4.0], rows, size=11, row_h=0.58, align=["l", "c", "l", "l"])
    notes[5] = "Rows marked 'as cited in [1]' are second-hand from the base paper's related work. The four batch-survey papers could not be verified, so no result is quoted for them."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in [
        "The surveyed systems do adaptive signal control or vehicle detection separately; no single pipeline detects, adapts timing, prioritises emergency vehicles and coordinates junctions (batch survey).",
        "The base paper validates one roundabout in simulation and gives priority at one junction; corridor coordination and proactive queue forecasting are outside its scope.",
        "ASTM forecasts vehicle counts but has no emergency priority; GNN congestion work does not control signals.",
        "Few papers report seeds, standard deviations or ablations, so the value of each added component is unclear.",
        "Synthetic-scene vision accuracy is rarely tested on real photos, so reported detection numbers may not transfer."]], size=19)
    notes[6] = "Five gaps. FlowSense addresses the first, second and fourth directly; the fifth we test with two public real-image sets."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.6)
    ftext(s, 0.85, 1.9, 11.6, 2.5, [
        "Fixed signal timings waste green time on empty approaches while queues build on others.",
        "Manual traffic control cannot scale to the number of junctions in a modern city.",
        "Emergency vehicles lose critical minutes waiting at signals with no priority mechanism.",
        "Junctions operate independently, so clearing one simply moves the queue to the next.",
        "Need: a vision-based controller that measures real traffic density, adapts signal timing and gives emergency vehicles priority."], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Camera frames (image upload), demand and controller settings, simulated vehicle arrivals"), ("Output", "Signal phases and green times, ambulance priority, queue forecasts, density levels, violation log"), ("Record", "Benchmark tables, metrics JSON, screenshots of the running console")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.7, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.8, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.3, 3.5, 1.0, b_, size=14)
    notes[7] = "Four problem points from the batch statement, then input, output and record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate([
        "Ingest junction-camera video and detect and count vehicles with a YOLO-based model.",
        "Estimate traffic density per approach lane from the detection output.",
        "Compute adaptive signal timings that allocate green time in proportion to measured density.",
        "Detect emergency vehicles and grant priority passage through the junction.",
        "Enable coordination between adjacent junctions to smooth traffic along a corridor.",
        "Evaluate detection accuracy, average waiting time and emergency vehicle clearance time."], 1)], size=22)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the real status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Perception and forecast", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Camera frame", K.FNAVY), ("YOLOv8n vehicles", K.FBLUE), ("Light-bar rule", K.GREEN), ("PCU density", K.ORANGE), ("30 s queue forecast", K.PURPLE), ("Green-time advice", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Control (advanced controller, per junction)", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Queues + PCU", K.FNAVY), ("Forecast excess", K.FBLUE), ("Inbound platoon", K.GREEN), ("Spill-back damping", K.ORANGE), ("Emergency green wave", K.PURPLE), ("Phase decision", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, [
        ("The same traffic is replayed under four controllers (fixed, adaptive, base-paper style, advanced) so every gain is like-for-like.", {"bullet": True, "space_after": 7}),
        ("Ambulances are cleared by pre-emption plus a 15-second look-ahead that turns neighbouring junctions green ahead of them.", {"bullet": True, "space_after": 7}),
        ("A violation log, benchmark page and live map make the behaviour visible; vision is a separate demo module, not wired into the simulator's loop.", {"bullet": True}),
    ], size=16)
    notes[9] = "The proposed system as two chains. Be explicit that vision is not closed-loop with the controllers and that violations are simulated."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.2, 5.1)
    ftext(s, 7.95, 1.95, 4.9, 4.9, [
        (f"Console: React 19 + Vite, port {F.WEB_PORT}; 8 pages", {"bullet": True, "space_after": 8}),
        (f"API: FastAPI + Uvicorn, port {F.API_PORT}; 20 JSON routes, demo token login", {"bullet": True, "space_after": 8}),
        ("Engine: smart_traffic (simulator, forecaster, vision) imported unchanged by the API service layer", {"bullet": True, "space_after": 8}),
        ("State: one shared live simulation, benchmark cache and violation log in memory; no database", {"bullet": True, "space_after": 8}),
        ("Files: yolov8n.pt, metrics JSON, screenshots", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. No database: state lives in process memory. The Streamlit app remains as the original front end."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, [
        ("Simulator: 1 s ticks, Poisson arrivals with rush-hour profile, saturation 0.5 veh/s, 12 s links, spill-back at 30 vehicles.", {"bullet": True, "space_after": 8}),
        ("Vehicles carry PCU weights (car 1, bike 0.5, truck 2, bus 3).", {"bullet": True, "space_after": 8}),
        ("Forecaster: gradient boosting on 11 features predicts each approach's queue 30 s ahead.", {"bullet": True, "space_after": 8}),
        ("Benchmarks: 5 seeds, 1800 s, demand x1.2; demand sweep, ablation and emergency-heavy tests.", {"bullet": True, "space_after": 8}),
        ("Vision checked on synthetic scenes and two public real-image sets.", {"bullet": True}),
    ], size=14)
    notes[11] = "Method in one picture: perceive, forecast, decide, measure. Demand is controller-independent, so controllers face identical vehicles and ambulances."

    # ---- 12 dataset
    s = S[11]
    set_title(s, "Dataset & Models")
    wipe(s)
    ftile(s, 0.55, 1.8, 3.9, 1.0, f"{F.PR['n_samples']:,}", "forecaster samples x 11 features", K.FBLUE, 24)
    ftile(s, 4.7, 1.8, 3.9, 1.0, f"{mt['images']}", f"real frames, {mt['gt_boxes_excluding_pedestrians']:,} vehicle boxes", K.GREEN, 24)
    ftile(s, 8.85, 1.8, 3.9, 1.0, f"{amb['images']}", "real photos: 300 ambulance, 300 other", K.ORANGE, 24)
    rows = [["Data", "Size", "Used for"],
            ["Simulator demand (generated)", "seeds 40-44", "controller benchmarks; no external traffic dataset"],
            ["Forecaster training runs", "6 x 1800 s", "2x2 / 3x3, demand 0.7-1.6, adaptive controller"],
            [f"Visaitech mixed traffic [{R['ds_traffic']}]", "293 frames", "YOLOv8n on real dashcam traffic (CC BY 4.0)"],
            [f"Ambulance photos [{R['ds_ambulance']}]", "600 images", "light-bar rule on real photos (CC BY 4.0)"],
            ["Synthetic overhead scenes", "300 scenes", "classical detector + light-bar rule"]]
    ftable(s, 0.55, 3.0, [2.7, 1.2, 3.5], rows, size=11, row_h=0.5)
    K.picture(s, F.fig("f13_real_examples.png"), 8.2, 3.0, 4.6, 3.95)
    ftext(s, 0.55, 6.15, 7.4, 0.8, "Models: YOLOv8n (COCO weights, not fine-tuned); gradient boosting, random forest and persistence for the forecaster. Right: real photos, a hit and a miss for each test.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "No traffic dataset is needed to drive the simulator. The two public image sets are used only to test the vision layer honestly; nothing is trained on them."

    # ---- 13 control workflow
    s = S[12]
    set_title(s, "Controller Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_control.png"), 0.55, 1.8, 5.7, 5.1)
    ftext(s, 6.5, 1.95, 6.3, 4.9, [
        ("Pressure per approach: PCU queue + 0.7 x max(forecast - queue, 0) + inbound PCU within 8 s + head wait / 12.", {"bullet": True, "space_after": 9}),
        ("Spill-back: pressure x 0.3 if the straight-ahead downstream queue is already at capacity - 4.", {"bullet": True, "space_after": 9}),
        ("Switch when the current phase is empty, when the other phase exceeds 2 x current + 4, or at 60 s; minimum green 13 s.", {"bullet": True, "space_after": 9}),
        ("Emergency: the nearest ambulance (queued, or inbound within 15 s) forces its phase green after 2 s.", {"bullet": True, "space_after": 9}),
        ("Compared with the base-paper style, which looks 5 s ahead at one junction only.", {"bullet": True}),
    ], size=15)
    notes[13] = "The core algorithm in smart_traffic/sim.py (_pressure, _decide, _emergency_target) as a flowchart."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_flow.png"), 0.55, 1.8, 6.4, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[14] = "Sequence of the main flow and the use cases of the operator."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design (no database)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.7, 4.6)
    rows = [["Store", "Content"], ["LIVE simulation (memory)", "one shared Sim, reset on grid change"], ["PREVIEW simulation", "landing-page pulse, separate seed"], ["benchmark cache", "keyed by grid, seconds, demand, seeds, ML"],
            ["last-benchmark violations", "feeds the Violations page"], ["predictor (memory)", "trained in a background thread, ~6 s"], ["yolov8n.pt, metrics JSON", "files on disk"]]
    ftable(s, 7.5, 1.9, [2.4, 2.9], rows, size=11, row_h=0.5)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "No foreign keys and no persistence: simulation state is lost when the server restarts.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "FlowSense has no database. The ER diagram shows logical entities."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("20", "JSON routes", K.FBLUE), (f"{F.PYTEST_TOTAL} + {F.SMOKE_TOTAL}", "pytest tests + smoke checks", K.GREEN), ("8", "React pages", K.ORANGE), (f"~{round(F.TH['2x2/advanced']['speedup'], -2):,.0f}x", "simulator speed vs real time (2x2)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        ("Service layer wraps the unchanged engine; HS256 token auth with a demo operator account; readable JSON errors (message + hint) for 401, 413, 415, 422, 503.", {"bullet": True, "space_after": 9}),
        ("The predictor trains in a background thread at start-up; /api/health reports readiness; benchmarks are cached by their settings.", {"bullet": True, "space_after": 9}),
        ("Same traffic for every controller: arrivals and ambulances come from a controller-independent random stream.", {"bullet": True, "space_after": 9}),
        ("Console: live map, benchmark, vision, prediction and violations pages; run.bat starts both servers; Windows needs torch imported before OpenCV.", {"bullet": True, "space_after": 9}),
        ("Not present: database, TLS, multi-user accounts, real camera streams, hardware control.", {"bullet": True}),
    ], size=16)
    notes[16] = "Highlights and honest absences."

    # ---- 17 results: controllers
    s = S[16]
    set_title(s, "Results – Controller Benchmark")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"-{d2['delay_cut_advanced_vs_fixed_pct']:.1f}%", f"delay vs fixed-time, 2x2 ({F.mstd(av2)} vs {F.mstd(fx2)} s)", K.FBLUE),
                                   (f"-{d3['delay_cut_advanced_vs_fixed_pct']:.1f}%", f"delay vs fixed-time, 3x3 ({F.mstd(av3)} vs {F.mstd(fx3)} s)", K.GREEN),
                                   (f"-{d2['delay_cut_advanced_vs_adaptive_pct']:.1f}%", "vs plain adaptive, 2x2: within one std", K.ORANGE),
                                   (f"{eav2[0]:.1f} s", f"ambulance delay vs {efx2[0]:.1f} s fixed (-{d2['emergency_cut_advanced_vs_fixed_pct']:.1f}%)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.1, v, l, c, 26, lsize=11)
    K.picture(s, F.fig("f01_delay_by_controller.png"), 0.55, 3.0, 6.0, 3.1)
    K.picture(s, F.fig("f02_emergency_delay.png"), 6.8, 3.0, 6.0, 3.1)
    ftext(s, 0.55, 6.2, 12.2, 0.8, f"Five seeds, 1800 s, demand x1.2. The base-paper-style controller is {-d2['delay_cut_basepaper_vs_adaptive_pct']:.1f}% worse than plain adaptive on mean delay (2x2) but cuts ambulance delay to {F.gm('2x2', 'emergency', 'Emergency delay (s)')[0]:.1f} s; advanced vs base-paper style: delay -{d2['delay_cut_advanced_vs_basepaper_pct']:.1f}%, ambulance -{d2['emergency_cut_advanced_vs_basepaper_pct']:.1f}% (2x2), -{d3['emergency_cut_advanced_vs_basepaper_pct']:.1f}% (3x3). Most of the gain over fixed-time is adaptivity itself.", size=12, color=K.FNAVY, italic=True)
    notes[17] = "Mean ± std over five seeds. The advanced controller's edge over plain adaptive control is small; its clearest benefit is ambulance delay."

    # ---- 18 confusion matrix & vision
    s = S[17]
    set_title(s, "Confusion Matrix & Vision on Real Photos")
    wipe(s)
    K.picture(s, F.fig("f08_ambulance_real_confusion.png"), 0.55, 1.75, 6.0, 3.2)
    K.picture(s, F.fig("f12_real_vision_per_class.png"), 6.8, 1.75, 6.0, 3.2)
    ftext(s, 0.55, 5.1, 12.2, 1.85, [
        (f"Light-bar rule on 600 real photos: any red+blue cluster TP {lb['tp']}, FP {lb['fp']}, FN {lb['fn']}, TN {lb['tn']} (recall {lb['recall'] * 100:.1f}%, false alarms {lb['false_alarm_rate'] * 100:.1f}%); YOLO box + light bar recall {lby['recall'] * 100:.1f}%.", {"bullet": True, "space_after": 4}),
        (f"YOLOv8n on 293 real dashcam frames (conf 0.3, IoU 0.5): 4-wheeler P {pc['4-wheeler']['precision']:.2f} / R {pc['4-wheeler']['recall']:.2f}; 2-wheeler R {pc['2-wheeler']['recall']:.2f}; 3-wheeler R {pc['3-wheeler']['recall']:.2f}; all vehicles R {pc['any vehicle class']['recall']:.2f}.", {"bullet": True, "space_after": 4}),
        (f"Synthetic scenes: recall {vc['recall'] * 100:.1f}%, ambulance recall {vc['ambulance_recall'] * 100:.0f}% - but noise sigma 10 gives zero detections and darkening removes every light bar. Base paper's trained detector: P 0.73, R 0.69.", {"bullet": True}),
    ], size=13)
    notes[18] = "Equivalent of the confusion matrix for this project: the real-photo test of the emergency rule. In-domain synthetic accuracy does not transfer."

    # ---- 19 parameter analysis
    s = S[18]
    set_title(s, "Parameter Analysis – Demand, Ablation, Forecaster")
    wipe(s)
    K.picture(s, F.fig("f03_demand_sweep.png"), 0.55, 1.75, 6.0, 3.0)
    K.picture(s, F.fig("f05_forecaster_validation.png"), 6.8, 1.75, 6.0, 3.0)
    rows = [["Ablation of the advanced controller", "Avg delay (s)", "Ambulance delay (s)"]]
    for k, m in ab.items():
        rows.append([k, F.mstd((m["Avg delay / vehicle (s)"]["mean"], m["Avg delay / vehicle (s)"]["std"])), f"{m['Emergency delay (s)']['mean']:.1f}"])
    ftable(s, 0.55, 4.85, [3.1, 1.5, 1.4], rows, size=11, row_h=0.34, align=["l", "c", "c"])
    ftext(s, 6.8, 4.85, 6.0, 2.1, [
        ("Demand sweep (2x2): advanced is worse than adaptive at x0.6 (26.4 vs 20.6 s), better from x1.2 (41.1 vs 42.2 s) to x1.8 (127.6 vs 135.3 s).", {"bullet": True, "space_after": 4}),
        ("No pre-emption gives the lowest network delay but the slowest ambulance: a real efficiency / safety trade-off.", {"bullet": True, "space_after": 4}),
        (f"Forecaster: random split MAE {pr_rand['Gradient Boosting (used)']['mae_mean']:.2f} (persistence {pr_rand['Persistence']['mae_mean']:.2f}); held-out R² {pr_loso['Gradient Boosting (used)']['r2_mean']:.2f}; beats persistence in 5 of 6 runs, fails on the congested one. Random forest is as good.", {"bullet": True}),
    ], size=12)
    notes[19] = "Parameter analysis: demand level, which component matters, and how far the forecaster generalises."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{lat('live step 2 s')['mean_ms']:.1f} ms", "API: advance live sim 2 s", K.FBLUE), (f"{lat('live step 300 s')['mean_ms']:.0f} ms", "API: skip 5 simulated minutes", K.GREEN),
                                   (f"{lat('vision synthetic')['mean_ms']:.0f} ms", "API: synthetic-scene detection", K.ORANGE), (f"{F.VY['sample_photo']['latency_ms_median']:.0f} ms", "YOLOv8n per image, CPU (median)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f10_api_latency.png"), 0.55, 2.9, 6.1, 2.4)
    K.picture(s, F.fig("f11_queue_timeseries.png"), 6.8, 2.9, 6.0, 2.4)
    rows = [["Check", "Result"],
            [f"pytest (engine {F.PYTEST_ENGINE}, API {F.PYTEST_API})", f"{F.PYTEST_PASSED} of {F.PYTEST_TOTAL} passed in ~{F.PYTEST_SECONDS} s; smoke test {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} checks"],
            ["First benchmark (4 controllers x 3 seeds x 1800 s)", f"{lat('POST /api/benchmark (4 ctrl')['mean_ms'] / 1000:.1f} s; repeats from cache in ~{lat('POST /api/benchmark (cached)')['mean_ms']:.0f} ms"],
            ["Not tested", "concurrent users, load, security penetration, real cameras or signal hardware"]]
    ftable(s, 0.55, 5.4, [4.7, 7.5], rows, size=11, row_h=0.38)
    notes[20] = "Latencies measured on 127.0.0.1 on an i5-12500, CPU only. YOLOv8n takes ~48 ms per image on CPU; the first call also loads the model."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "FlowSense (this project)"],
            ["Detector", f"YOLO26m light-bar, trained on 3,130 images: P {bp['precision']:.2f}, R {bp['recall']:.2f}, mAP50 {bp['map50']:.2f}", f"YOLOv8n (not fine-tuned) + HSV light-bar rule; real photos: light-bar recall {lb['recall'] * 100:.1f}%, 4-wheeler recall {pc['4-wheeler']['recall'] * 100:.0f}%"],
            ["Controller", "GA-SGD plan + supervisory override", "PCU pressure + ML forecast + look-ahead + green wave"],
            ["Testbed", "SUMO, one Rabat roundabout, selected scenario", "In-house simulator, 2x2 and 3x3 grids, 5 seeds"],
            ["Ambulance", "Travel 56 s vs 70 s GA-SGD / 64 s fixed (-20 % / -12.5 %)", f"Delay {eav2[0]:.1f} s vs {efx2[0]:.1f} s fixed (-{d2['emergency_cut_advanced_vs_fixed_pct']:.0f}%); -{d2['emergency_cut_advanced_vs_basepaper_pct']:.0f}% vs base-paper style"],
            ["Network", "Total waiting 19,887 s vs 22,827 s / 33,719 s (-12.9 % / -41 %)", f"Avg delay {av2[0]:.1f} s vs {ad2[0]:.1f} s adaptive (-{d2['delay_cut_advanced_vs_adaptive_pct']:.1f}%) / {fx2[0]:.1f} s fixed (-{d2['delay_cut_advanced_vs_fixed_pct']:.1f}%)"]]
    ftable(s, 0.55, 1.8, [1.35, 3.35, 3.5], rows, size=12, row_h=0.85)
    K.picture(s, F.fig("f09_base_paper_comparison.png"), 8.95, 1.8, 3.9, 2.6)
    ftext(s, 8.95, 4.5, 3.9, 2.4, [
        ("Different simulators, networks and metrics: these figures are indicative, not a head-to-head.", {"bullet": True, "space_after": 6}),
        ("The base paper's trained detector is the standard we did not reach.", {"bullet": True, "space_after": 6}),
        ("No superiority is claimed.", {"bullet": True}),
    ], size=12)
    notes[21] = "State plainly that testbeds differ. The relative reductions are of the same kind but not comparable in absolute terms."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Live Control and Benchmark")
    wipe(s)
    K.picture(s, crop_shot("06_live_map_and_status.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("07_benchmark.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: live 2x2 network after 360 simulated seconds - signals, queued vehicles, junction table with +30 s forecasts. Right: benchmark with UI defaults (3 seeds): advanced -38% vs fixed-time, -11% vs emergency-aware, ambulance delay 5.1 vs 18.7 s.", size=13, color=K.FNAVY)
    notes[22] = "Real captures of the running console. The report tables use 5 seeds, so their numbers differ slightly from this 3-seed screen."
    s = S[22]
    set_title(s, "Output – Vision: Synthetic Scene and YOLOv8n")
    wipe(s)
    K.picture(s, crop_shot("10_vision_synthetic.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("12_vision_yolo.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: synthetic overhead scene - boxes by class, a magenta ring on the red+blue light-bar cluster and the ambulance label. Right: YOLOv8n on the bundled sample photo detects the bus (0.87). Real-photo accuracy is much lower than on synthetic scenes (see results).", size=13, color=K.FNAVY)
    notes[23] = "Vision page: classical detector on a synthetic scene and YOLOv8n on a real photo."
    s = S[23]
    set_title(s, "Output – Queue Forecast and Violation Log")
    wipe(s)
    K.picture(s, crop_shot("14_prediction_forecast.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("16_violations_log.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: model comparison (persistence 2.535, random forest 1.605, gradient boosting 1.644 MAE) and the live per-approach 30 s forecast. Right: simulated red-light and wrong-way events with plate and confidence - generated from ground-truth events, not detected by a model.", size=13, color=K.FNAVY)
    notes[24] = "Forecast and violation screens. Say clearly that violations are simulated events."

    # ---- 25 conclusion and objectives
    s = S[24]
    set_title(s, "Conclusion & Objectives vs Outcomes")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 1.45, [
        (f"A working simulator-validated platform: adaptive control cuts delay about a third vs fixed-time ({d2['delay_cut_advanced_vs_fixed_pct']:.0f}% / {d3['delay_cut_advanced_vs_fixed_pct']:.0f}%); pre-emption cuts ambulance delay {d2['emergency_cut_advanced_vs_fixed_pct']:.0f}-{d3['emergency_cut_advanced_vs_fixed_pct']:.0f}%; {F.PYTEST_PASSED} tests and {F.SMOKE_PASSED} smoke checks pass.", {"bullet": True, "space_after": 4}),
        (f"Findings: advanced vs adaptive only -{d2['delay_cut_advanced_vs_adaptive_pct']:.1f}%; pre-emption trades a little network delay for ambulance speed; the forecaster does not extrapolate; real-photo vision is weak (light-bar recall {lb['recall'] * 100:.1f}%).", {"bullet": True}),
    ], size=13)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Detect and count vehicles with YOLO", f"Partly: YOLOv8n on images only (no video stream); real recall {pc['any vehicle class']['recall'] * 100:.0f}%, 4-wheelers {pc['4-wheeler']['recall'] * 100:.0f}%"],
            ["2", "Density per approach", "Partly: PCU density per frame; per-lane density only in the simulator"],
            ["3", "Adaptive green time by density", f"Met in simulation: -{d2['delay_cut_advanced_vs_fixed_pct']:.1f}% delay vs fixed-time"],
            ["4", "Emergency vehicle priority", f"Met in simulation (ground-truth ambulances): {eav2[0]:.1f} s vs {efx2[0]:.1f} s; real-photo detection not achieved"],
            ["5", "Coordination between junctions", f"Partly: simulated look-ahead; only -{d2['delay_cut_advanced_vs_adaptive_pct']:.1f}% over adaptive"],
            ["6", "Evaluate accuracy, waiting, clearance", "Met: measured on simulator and two real-image sets; accuracy on real data is low"]]
    ftable(s, 0.55, 3.4, [0.5, 3.6, 8.1], rows, size=12, row_h=0.5, align=["c", "l", "l"])
    notes[25] = "Be explicit about which objectives were met, partly met or not met."

    # ---- 26 limitations and future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 5.1, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("In-house simulator, not SUMO; absolute numbers are not comparable with the base paper", {"bullet": True, "space_after": 7}),
        ("Ambulances known by ground truth; vision not closed-loop", {"bullet": True, "space_after": 7}),
        ("No trained light-bar detector: real-photo recall 1.3%", {"bullet": True, "space_after": 7}),
        ("Violations simulated, not detected; no tracking or plate reading", {"bullet": True, "space_after": 7}),
        ("Forecaster fails outside its training range (R\u00b2 ~0.04 held-out)", {"bullet": True, "space_after": 7}),
        ("Advanced controller worse at low demand; gain over adaptive within noise", {"bullet": True, "space_after": 7}),
        ("No database, no cameras or hardware, single demo user, CPU only", {"bullet": True}),
    ], size=15)
    fbox(s, 6.8, 1.8, 6.0, 5.1, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Validate in SUMO or CARLA on real junction layouts", {"bullet": True, "space_after": 7}),
        ("Train a light-bar / emergency-vehicle detector on labelled data; check flashing state", {"bullet": True, "space_after": 7}),
        ("Wire vision into the control loop with temporal confirmation", {"bullet": True, "space_after": 7}),
        ("Train the forecaster on wider congestion; try sequence or graph models", {"bullet": True, "space_after": 7}),
        ("Fine-tune YOLO on local mixed traffic incl. three-wheelers", {"bullet": True, "space_after": 7}),
        ("Real camera streams, persistent storage, multi-user roles", {"bullet": True, "space_after": 7}),
        ("Pilot with edge hardware and a signal controller", {"bullet": True}),
    ], size=15)
    notes[26] = "Limitations first, then concrete next steps."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [R[k] for k in ("base", "astm", "necdm", "gasgd", "sasikala", "zlatkovic", "jayakumar", "yolo", "ultralytics", "sumo", "varaiya", "intellilight", "gbm", "ds_traffic")]
    short = [f"[{n}] {clean_ref(REFS[n - 1])}" for n in ids]
    short = [t if len(t) < 400 else t[:397] + "..." for t in short]
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[:7]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[7:]] + [("The full numbered list of 33 references, including the batch survey items [2]-[5], is in the report.", {"italic": True})], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh_ in S[27].shapes:
        if sh_.name == "Text 2":
            K.replace_text(sh_, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "FlowSense")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
