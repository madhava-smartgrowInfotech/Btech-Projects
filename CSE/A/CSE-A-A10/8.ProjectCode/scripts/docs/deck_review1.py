"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "smart traffic management, YOLOv8, adaptive signal control, emergency vehicle priority, queue forecasting"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


# ---------------------------------------------------------------- slide text (batch documents + verified facts only)
PROBLEM_STATEMENT = [
    "Fixed signal timings waste green time on empty approaches while queues build on others",
    "Manual traffic control cannot scale to the number of junctions in a modern city",
    "Emergency vehicles lose critical minutes waiting at signals with no priority mechanism",
    "Junctions operate independently, so clearing one simply moves the queue to the next",
    "Need: a vision-based controller that measures real traffic density, adapts signal timing and gives emergency vehicles priority",
]
OBJECTIVES = [
    "Ingest live video from junction cameras and detect and count vehicles using a YOLO-based model",
    "Estimate traffic density per approach lane from the detection output",
    "Compute adaptive signal timings that allocate green time in proportion to measured density",
    "Detect emergency vehicles and grant priority passage through the junction",
    "Enable coordination between adjacent junctions to smooth traffic along a corridor",
    "Evaluate detection accuracy, average waiting time and emergency vehicle clearance time",
]

# [work, year, approach, reported result / limitation]  (cells kept short)
LIT_ROWS = [
    [f"Lamrabet et al. [{R['base']}] (base)", "2026", "Trained YOLO26m roof light-bar detector as a priority layer over a GA-SGD controller; SUMO roundabout",
     "P 0.73, R 0.69, mAP50 0.73; ambulance trip 56 s vs 70 s (GA-SGD), 64 s (fixed); one roundabout, simulation only"],
    [f"Goenawan (ASTM) [{R['astm']}]", "2024", "YOLOv5 vehicle detection + RNN-LSTM 12-hour count forecast + cycle-length analysis (CARLA)",
     "Forecast MSE 4.521 / RMSE 2.232 vehicles; 21 vs about 15 vehicles/min, 5 s vs about 12 s delay"],
    [f"Gupta et al. [{R['necdm']}]", "2026", "Neighbour-embedded graph neural network predicts traffic flow at delivery stations",
     "L1 loss -4.03 %, L2 loss -16.66 %; crowd-delivery focus, no signal control"],
    [f"Lamrabet et al. [{R['gasgd']}]", "2025", "Hybrid GA-SGD optimisation of signal timing, Moroccan case study",
     "Per the base paper: no emergency-vehicle detection or priority"],
    [f"Sasikala et al. [{R['sasikala']}]", "2025", "IoT, GPS and RFID emergency prioritisation with dynamic signal control",
     "35 % shorter ambulance response (as reported in [1]); needs dedicated infrastructure"],
    [f"Zlatkovic et al. [{R['zlatkovic']}]", "2023", "Pre-emption concepts for connected-vehicle environments (VISSIM)",
     "Up to 36 % emergency delay reduction (as reported in [1]); needs connected vehicles"],
    [f"Jayakumar et al. [{R['jayakumar']}]", "2024", "Siren recognition from temporal and spectral audio features, SVM and ensembles",
     "99.5 % accuracy (as reported in [1]); audio has no spatial localisation"],
    [f"Suhana; Satyanarayana; Sharma [{R['suhana']}], [{R['satyanarayana']}], [{R['sharma']}]", "2025", "YOLOv5 / YOLOv8 recognise ambulances, fire trucks and police cars from CCTV",
     "Rely on vehicle appearance, which varies by region and fleet [1]"],
    [f"Signal-IQ [{R['s_signaliq']}]", "2026", "Low-cost adaptive control with YOLOv8 vehicle and emergency-vehicle detection (title only)",
     "Hardware-validated per its title; authors and results not verified here"],
    [f"Batch survey [{R['s_cisct']}]-[{R['s_incip']}]", "2024-25", "Density-based adaptive control; accelerated emergency detection; client-server dynamic control (titles only)",
     "Full texts not accessed; no figures quoted"],
]

GAPS = [
    "The surveyed systems perform adaptive signal control or vehicle detection separately: one pipeline that detects, adapts timing, prioritises emergency vehicles and coordinates junctions is missing",
    f"Emergency priority relies on dedicated infrastructure [{R['sasikala']}], [{R['zlatkovic']}], on audio without localisation [{R['jayakumar']}], or on vehicle appearance [{R['suhana']}], [{R['satyanarayana']}], [{R['sharma']}]",
    f"The base paper validates one modelled roundabout; it has no queue forecast and no coordination between neighbouring junctions [{R['base']}]",
    "Controllers are rarely compared on identical traffic with the price of pre-emption (delay to other vehicles) measured alongside the ambulance gain",
    f"Validation stays in simulation for want of authorisation and real-camera trials [{R['base']}]; transfer of detectors to street cameras is untested",
]


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    bp = F.BASE_PAPER_RESULTS

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: FlowSense puts vehicle detection, queue forecasting, adaptive signal timing, emergency-vehicle priority and "
                "junction-to-junction coordination into one working system, and measures it against fixed-time and emergency-aware baselines.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design details and measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: see the traffic, adapt the signals")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Traffic congestion raises travel time, fuel use and delay for ambulances. Fixed-time signals cannot react to real traffic, and most systems either detect vehicles or adapt timing, rarely both together with emergency priority.", K.ORANGE, "!"),
        ("Our idea", "One pipeline: YOLOv8 vehicle detection and density, a machine-learning forecast of queues 30 s ahead, adaptive green times, emergency-vehicle priority, coordination between neighbouring junctions and a violation log.", K.BLUE, "i"),
        ("The deliverable", "FlowSense: a simulator with four controllers on identical traffic, a vision module, a queue forecaster and a FastAPI + React control console, with delay and ambulance clearance time measured.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: the weakness of fixed-time and detect-only systems, the idea of one pipeline, and what is delivered. Everything runs locally on a CPU, without a GPU."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "A YOLO26m detector, trained on 3,130 annotated images, finds the roof-mounted siren / light-bar (single class). A priority arbitrator (confidence 0.5, 3 consecutive frames) overrides a GA-SGD optimised signal plan and restores it after the vehicle passes.", K.BLUE, "M"),
        ("Reported results", f"Detector: precision {bp['precision']}, recall {bp['recall']}, mAP50 {bp['map50']}, {bp['latency_ms'][0]}-{bp['latency_ms'][1]} ms per frame on a T4. SUMO, one Rabat roundabout: ambulance travel 56 s vs 70 s (GA-SGD) and 64 s (fixed-time); total waiting 19,887 s vs 22,827 s and 33,719 s.", K.GREEN, "R"),
        ("Relevance and limits", "Same core idea: detect the emergency cue and adapt the signal. The paper itself lists a moderate dataset, simulation-only validation, a single roundabout and no flashing-state check. FlowSense adds forecasting and junction coordination.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The paper adds a trained light-bar detector to an already optimised controller and shows, in SUMO, shorter ambulance trips and lower total waiting. "
                "It is validated on one modelled roundabout in simulation, which is the main limit we try to widen: several junctions, a forecast, coordination.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    rows += LIT_ROWS
    K.table(s, 0.6, 1.25, [2.5, 0.65, 4.55, 4.43], rows, font_size=10, row_h=0.5, header_size=11, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey is the base paper plus four recent works; the ASTM paper is the companion reference. Figures are quoted only where they were read in the paper itself or, marked 'as reported in [1]', "
                "from the base paper's own related-work section. For the four survey titles only the title is available, so no results are quoted.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: adaptive and priority control pay off")
    wipe(s)
    tw = bp["total_wait_s"]
    gf = K.bar_chart(s, 0.6, 1.4, 6.1, 4.9, ["Fixed-time", "GA-SGD", "GA-SGD +\nlight-bar priority"],
                     [("Total network waiting time (s)", [tw["fixed"], tw["ga_sgd"], tw["proposed"]])], colors=[K.BLUE], number_format="#,##0",
                     font_size=12, max_value=40000)
    pt = gf.chart.plots[0].series[0].points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 0.6, 6.35, 6.1, 0.6, "Base paper [1]: SUMO, modelled Souissi roundabout, selected scenario (demand scale 12).", size=12, color=K.TXT, italic=True)
    K.text(s, 7.0, 1.45, 5.75, 5.0, [
        ("**What the literature shows**", {"space_after": 8}),
        ("Adaptive control cuts waiting: 41.0 % lower total waiting than fixed-time in the base paper's scenario [1].", {"bullet": True, "space_after": 8}),
        (f"A forecast plus cycle-length control raised throughput from about 15 to 21 vehicles/min in CARLA [{R['astm']}].", {"bullet": True, "space_after": 8}),
        ("Pre-emption reduces ambulance delay: 20.0 % shorter travel than GA-SGD in the base paper [1]; up to 36 % in a VISSIM study reported there.", {"bullet": True, "space_after": 8}),
        ("Reported detector quality is moderate for small light-bars: mAP50 0.73 [1].", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("These numbers are the published evidence that motivates the project, taken from the papers we read. They come from different simulators and datasets, so they are not combined or compared with each other. "
                "Our own measurements start in Review 2.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Cameras / loop sensors", K.NAVY), ("Fixed-time or manual plan", K.NAVY), ("Signal phase per junction", K.NAVY), ("Each junction in isolation", K.NAVY)], size=15, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Fixed-time and manual control",
           "A fixed cycle gives the same green to an empty approach and a full one.\nManual control by police or operators does not scale to a city of junctions.\nAn ambulance waits in the queue like any other vehicle.", K.BLUE, "1", body_size=18)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Adaptive and priority systems",
           "Adaptive timing or vehicle detection is built separately from emergency priority.\nPriority uses IoT / GPS hardware, audio sirens, vehicle class or a light-bar detector.\nNeighbouring junctions rarely share what they see.", K.ORANGE, "2", body_size=18)
    notes[7] = "Today most junctions run a fixed plan. Where adaptive or priority systems exist, detection, timing, emergency priority and coordination are separate pieces."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Wasted green time", "Does not scale", "Emergency delay", "No coordination"]
    bodies = [
        "Fixed timings keep green on empty approaches while queues grow elsewhere, raising delay, fuel use and emissions.",
        "Manual control and hand-tuned plans cannot follow demand across many junctions through the day.",
        "Without a priority mechanism, ambulances wait in the queue; IoT and acoustic fixes need extra infrastructure or lose spatial precision.",
        "Junctions act independently, so clearing one moves the queue to the next; no forecast anticipates the build-up.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=16)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"The base paper adds: moderate dataset size, light-bar presence only (not flashing state), one roundabout, simulation-only validation [{R['base']}].", size=14, color=K.INK, anchor="m")
    notes[8] = "The first four limitations are the batch problem statement; the last line is the base paper's own limitations section."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=16, color=K.INK, anchor="m")
    notes[9] = ("Gap 1 is the batch survey's own statement. Gaps 2, 3 and 5 come from the base paper's related-work and limitations sections. "
                "Gap 4 is our reading: the cost that pre-emption imposes on other traffic is seldom reported next to the ambulance gain.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Junction camera frames or an uploaded image; simulated demand and ambulance arrivals; road network of 2x2 or 3x3 junctions"),
                                 ("Output", "Vehicle counts, density level and green-time advice; signal phases; 30 s queue forecast; ambulance pre-emption; violation log"),
                                 ("Measures", "Average delay per vehicle, queue length, ambulance delay, throughput, fuel wasted idling, detection precision and recall")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=13, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a vision-based controller that measures traffic, adapts timing and gives ambulances priority."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Validation is in a seeded simulator rather than SUMO; vision is checked on synthetic scenes and on public real photographs.", size=14, color=K.TXT, italic=True)
    notes[11] = ("The six objectives from the proposal. Later reviews report which were met. We validate on our own simulator because SUMO was not used, and we say so; "
                 "detection accuracy is reported on real public images as well as synthetic ones.")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Detect", "YOLOv8n finds cars, bikes, buses and trucks; a red+blue light-bar rule flags an emergency vehicle; density is counted in passenger-car units (PCU).", K.BLUE, "1"),
        ("Forecast", "A gradient-boosting model predicts each approach's queue 30 s ahead from 11 features (queue, lags, inflow, signal state, time of day).", K.GREEN, "2"),
        ("Control", "Green time follows PCU-weighted pressure plus the forecast and arriving platoons; ambulances get a green wave across junctions.", K.ORANGE, "3"),
        ("Evaluate", "Four controllers on identical seeded traffic: fixed-time, adaptive, emergency-aware (base-paper style) and FlowSense advanced.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React console: live map, benchmark, vision, prediction and violation screens", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "Violations (red-light, wrong-way) are logged with plate and confidence from simulated events. No database and no GPU are needed.", size=14, color=K.TXT)
    notes[12] = ("Four stages and the console that wraps them. The base-paper-style controller is our emulation of the paper's idea inside the simulator, not a re-implementation of its GA-SGD or its trained detector.")

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    if os.path.exists(F.diag("01_system_architecture.png")):
        K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.2, 7.5, 5.9)
    else:
        print("WARNING: diagram 01 missing")
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite console (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}) with token sign-in and input checks.", {"bullet": True}),
        ("Core: traffic simulator with four controllers, queue forecaster, YOLOv8n vision and density.", {"bullet": True}),
        ("State: in memory; the model is trained at start-up; no database.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the console calls the API, the API calls the simulator, forecaster and vision module. There is no database."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working platform", "Controller comparison", "Emergency priority", "Honest vision check", "Stated limits"]
    bodies = [
        "A console that runs the simulator live, benchmarks controllers, analyses images, forecasts queues and logs violations.",
        "Delay, queue and fuel compared across four controllers, 2x2 and 3x3 networks, several seeds and demand levels.",
        "Ambulance delay with and without pre-emption, and what pre-emption costs the other traffic.",
        "Vehicle and light-bar detection tested on synthetic scenes and on real public photographs, with failures reported.",
        "Simulator not SUMO, no real camera stream, simulated violations: each limit documented, not hidden.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=16)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=17)
    notes[14] = ("These are the outcomes we commit to. A first prototype run already shows the adaptive controllers well below fixed-time delay and ambulance delay cut sharply, but those results, "
                 "and where they are weaker than hoped, are presented in Review 2 and the final review.")

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [R[k] for k in ("base", "s_signaliq", "astm", "necdm", "gasgd", "sasikala", "zlatkovic", "jayakumar", "suhana", "yolo", "ultralytics")]
    K.text(s, 0.6, 1.25, 12.13, 5.8, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 3}) for n in idx], size=13, color=K.INK)
    notes[15] = "Selected references; the complete numbered list, with notes on how each entry was checked, is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why a simulator instead of SUMO (self-contained, seeded, identical demand for every controller); why YOLOv8n and an HSV light-bar rule instead of a trained light-bar model "
                 "(no light-bar dataset was trained; real-photo results are reported honestly); how the base paper differs (trained detector, GA-SGD, one roundabout).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
