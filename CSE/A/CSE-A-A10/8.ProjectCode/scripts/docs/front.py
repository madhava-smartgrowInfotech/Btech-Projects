"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the publishers of the two public image sets (Visaitech mixed-traffic frames and the ambulance photo set) "
        "used for the real-image evaluation, and the maintainers of the open-source software used in this project (Ultralytics YOLOv8, PyTorch, OpenCV, scikit-learn, "
        "FastAPI, React, Vite, Recharts, pandas and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    a2, a3 = F.gm("2x2", "advanced", "Avg delay / vehicle (s)"), F.gm("3x3", "advanced", "Avg delay / vehicle (s)")
    f2_, f3_ = F.gm("2x2", "fixed", "Avg delay / vehicle (s)"), F.gm("3x3", "fixed", "Avg delay / vehicle (s)")
    e2a, e2f = F.gm("2x2", "advanced", "Emergency delay (s)"), F.gm("2x2", "fixed", "Emergency delay (s)")
    d2, d3 = F.derived("2x2"), F.derived("3x3")
    rs = F.PR["random_split_5x"]
    gb, pers = {"mae": rs["Gradient Boosting (used)"]["mae_mean"]}, {"mae": rs["Persistence"]["mae_mean"]}
    loso = F.PR["leave_one_simulation_out"]["Gradient Boosting (used)"]
    mt = F.VR["mixed_traffic"]["per_class"]
    amb = F.VR["ambulance_classification"]
    vsc = F.VS["results"]["clean"]
    d.p("Fixed-time traffic signals cannot see queues, ambulances or the next junction, so green time is wasted while vehicles wait and emergency vehicles lose minutes. "
        f"This project builds {F.PRODUCT}, one pipeline that detects vehicles with YOLOv8, estimates density in passenger-car units, forecasts each approach's queue 30 seconds ahead, "
        "adapts signal timing, clears the way for ambulances and lets neighbouring junctions look ahead at each other, with a violation log and a FastAPI and React console. "
        "Four controllers are compared on identical seeded traffic in a purpose-built simulator of 2x2 and 3x3 junction grids; it is not SUMO and its absolute numbers are not "
        "comparable with the base paper's.")
    d.p(f"Over five seeds on a 2x2 grid the advanced controller averages {a2[0]:.1f} s of delay per vehicle against {f2_[0]:.1f} s for fixed-time signals "
        f"({d2['delay_cut_advanced_vs_fixed_pct']:.1f} % lower; {d3['delay_cut_advanced_vs_fixed_pct']:.1f} % on 3x3, {a3[0]:.1f} s against {f3_[0]:.1f} s) and cuts ambulance delay "
        f"from {e2f[0]:.1f} s to {e2a[0]:.1f} s. Its edge over plain queue-actuated control is only {d2['delay_cut_advanced_vs_adaptive_pct']:.1f} %, within one standard deviation, and "
        f"the base-paper-style pre-emption controller is {-d2['delay_cut_basepaper_vs_adaptive_pct']:.1f} % slower than plain adaptive control. The queue forecaster reaches MAE {gb['mae']:.2f} against "
        f"{pers['mae']:.2f} for persistence on a random split, but its leave-one-simulation-out R2 averages {loso['r2_mean']:.2f}. Vision works on synthetic scenes (recall {vsc['recall'] * 100:.1f} %) "
        f"but not on real photographs: YOLOv8n recalls {mt['4-wheeler']['recall'] * 100:.0f} % of four-wheelers and {mt['2-wheeler']['recall'] * 100:.0f} % of two-wheelers on real dashcam frames, and "
        f"the light-bar rule flags {amb['ambulance_flag_yolo_plus_lightbar']['tp']} of {amb['ambulance_images']} real ambulance photos. Violations and emergency knowledge are simulated, "
        f"and {F.PYTEST_PASSED} unit and API tests pass.")
    d.p("**Keywords:** smart traffic management, adaptive signal control, YOLOv8, emergency vehicle priority, queue forecasting, inter-junction coordination, traffic simulation, computer vision.")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("CCTV", "Closed-Circuit Television"), ("CNN", "Convolutional Neural Network"),
    ("CO2", "Carbon Dioxide"), ("COCO", "Common Objects in Context (dataset)"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"),
    ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"), ("FPS", "Frames Per Second"), ("GA-SGD", "Genetic Algorithm - Stochastic Gradient Descent"),
    ("GB", "Gradient Boosting"), ("GNN", "Graph Neural Network"), ("HSV", "Hue, Saturation, Value (colour space)"), ("HTTP", "Hypertext Transfer Protocol"),
    ("IoT", "Internet of Things"), ("IoU", "Intersection over Union"), ("ITS", "Intelligent Transportation Systems"), ("JSON", "JavaScript Object Notation"),
    ("JWT", "JSON Web Token"), ("KPI", "Key Performance Indicator"), ("LSTM", "Long Short-Term Memory"), ("MAE", "Mean Absolute Error"),
    ("ML", "Machine Learning"), ("PCU", "Passenger-Car Unit"), ("R2", "Coefficient of determination"), ("RF", "Random Forest"),
    ("SUMO", "Simulation of Urban Mobility"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("YOLO", "You Only Look Once"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
