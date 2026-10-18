"""Build 16.Final_PPT.pptx from D:\\abstracts\\Templates\\Final_PPT.pptx (28 slides)."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deckkit import fill, put_picture, set_notes, set_paras, set_table, shape, table_size  # noqa: E402
from doc_common import (BASE_PAPER, BATCH, BEST, CHARTS, COLLEGE, DEMO, DIAG, EV, GUIDE, GUIDE_DESIG, LAT, MODELS, PRODUCT, REFERENCES, REP, RET, SEEDS,  # noqa: E402
                        SHOTS, STUDENTS, TITLE, YEAR, f3)
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

SRC = r"D:\abstracts\Templates\Final_PPT.pptx"
OUT = BATCH / "16.Final_PPT.pptx"
FOOT = "Dept. of CSE  |  Batch A11  |  Final Review"
CACHE = Path(tempfile.gettempdir()) / "autosecai_deck"
CACHE.mkdir(exist_ok=True)
KEYWORDS = "pull request review, vulnerability detection, source-code metrics, multi-agent, code quality"


def px(v):
    return Emu(int(v * 9525))


def crop(name, top=0, bottom=0, left=0):
    """Crop the sticky top bar and the sidebar off a screenshot; returns the cached file."""
    out = CACHE / f"{Path(name).stem}_{top}_{bottom}_{left}.png"
    with Image.open(SHOTS / name) as im:
        im.crop((left, top, im.width, im.height - bottom)).save(out)
    return out


def table(s, rows, widths, max_pt=16, h=430, w=1180):
    sz = table_size(rows, w, h, widths, max_pt, 9)
    tshape = [sh for sh in s.shapes if sh.has_table][0]
    set_table(tshape, rows, sz, widths=widths, header_size=sz)


def two_pics(s, left, right, cap_l, cap_r, ids_remove=(7, 9), cap_ids=(8, 10), box_h=352, top=182):
    put_picture(s, left, (43, top, 576, box_h), remove_ids=(ids_remove[0],), name=cap_l)
    put_picture(s, right, (657, top, 576, box_h), remove_ids=(ids_remove[1],), name=cap_r)
    set_paras(shape(s, cap_ids[0]), cap_l)
    set_paras(shape(s, cap_ids[1]), cap_r)
    for cid in cap_ids:
        shape(s, cid).top = px(top + box_h + 6)


def build():
    prs = Presentation(SRC)
    S = list(prs.slides)
    cp = prs.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = f"{PRODUCT}: Final Review", TITLE, "Batch A11", KEYWORDS
    cp.category = cp.comments = cp.content_status = ""
    cp.last_modified_by = "Batch A11"
    for s in S[1:]:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("Dept. of CSE"):
                set_paras(sh, FOOT)
    best, rf = MODELS[BEST], MODELS["Random Forest"]
    others = [m["f1"] for n, m in MODELS.items() if n not in (BEST, "Random Forest")]
    sm = SEEDS["summary"]

    # 1 title
    s = S[0]
    fill(s, 4, TITLE, 27)
    set_paras(shape(s, 5), "Final Review")
    set_paras(shape(s, 6), "Team Details : Batch A11 (CSE, Section A)")
    set_paras(shape(s, 7), [f"{r}   {n}" for n, r in STUDENTS], 16)
    set_paras(shape(s, 8), ["Under the Esteemed Guidance of", GUIDE, f"{GUIDE_DESIG}, Department of CSE"])
    set_paras(shape(s, 9), f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  \u00b7  Academic Year {YEAR}", 12)
    set_notes(s, "Final review of AutoSecAI: an autonomous pull-request reviewer for security, vulnerability detection and code quality.")

    # 2 contents
    s = S[1]
    fill(s, 7, ["Abstract", "Base Paper & Literature Survey", "Research Gap & Problem Statement", "Objectives", "Proposed System & Architecture", "Methodology & Pipeline",
                "Design & Implementation", "Results, Testing & Comparison", "Output Screens", "Conclusion & Future Scope", "References"], 20, bullet=True)

    # 3 abstract
    s = S[2]
    fill(s, 7, ["Pull requests arrive faster than they can be reviewed by hand; security flaws such as SQL injection and XSS slip through, static analysers add many low-value warnings, and newcomers struggle to find their way around a large codebase.",
                "AutoSecAI extracts the changed files, runs AST and taint static analysis, scores each changed function with the base paper's 14 source-code metrics, and uses seven cooperating agents to explain each finding, propose a patch that is verified by re-scanning, and issue a verdict.",
                "A codebase assistant answers natural-language questions about a repository; results export as Markdown, SARIF, JSON and CSV through a React + FastAPI web application.",
                f"Measured: XGBoost F1 {f3(best['f1'])} (precision {f3(best['precision'])}, recall {f3(best['recall'])}) on 1,500 held-out functions of a generated Python corpus; rule engine F1 {f3(EV['rule_engine']['f1'])}; either detector {f3(EV['combined_or']['f1'])}; a full demo review takes {LAT['median_s']} s.",
                "Keywords: pull request review, vulnerability detection, source-code metrics, multi-agent systems, code quality."], 17, bullet=True)
    set_notes(s, "Abstract: problem, approach, application, and the headline measurements. The ML scores are on a generated corpus, not on real pull requests.")

    # 4 base paper
    s = S[3]
    fill(s, 7, [BASE_PAPER,
                "Task: predict vulnerable C/C++ functions, and their type, from 14 source-code metrics taken from a code property graph; 21,759 functions from 10 open-source projects, 275 vulnerable.",
                "Method: seven feature-selection methods keep the five most useful metrics; SmoteTomek balances the classes; Min-Max scaling; nine ML models, tuned by grid search.",
                "Reported results: Random Forest accuracy 0.96, precision 0.948, recall 0.974, F1 0.961; top metrics M5, M6, M1, M9, M11.",
                "Our reading: the signal is computed offline and per function; no explanation, patch or pull-request context. The paper lists other languages than C and C++ as future work. We adopt the pipeline, reproduce it for Python and place it inside a pull-request review."], 17, bullet=True)
    set_notes(s, "The base paper predicts vulnerable functions offline; AutoSecAI applies the same signal inside the pull-request workflow and adds explanation and remediation.")

    # 5 literature survey
    s = S[4]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Abbasi and Shameli-Sendi, 2026 (base) [1]", "14 code metrics, feature selection, SmoteTomek, 9 ML models", "21,759 C/C++ functions, 10 projects", "Random Forest F1 0.961, accuracy 0.96", "C/C++ only; offline; no explanation"],
            ["Zhang et al., 2026 [2]", "Sphinx: structured data generation, checklist benchmark, checklist-reward training for LLM PR review", "Pull-request review comments", "Up to 40% higher checklist coverage", "Review text, not security rules or fixes"],
            ["CNN-LSTM hybrid, 2025 [3]; AI code security auditing, 2025 [4]", "Deep hybrid detection; automated detection and remediation", "Not retrieved", "Not retrieved", "Listed in batch survey only"],
            ["Multi-modal + meta-learning, 2025 [5]; DcodeBERT, 2024 [6]", "Code simplification and meta-learning; BERT-style detection and repair", "Not retrieved", "Not retrieved", "Listed in batch survey only"]]
    table(s, rows, [0.2, 0.28, 0.17, 0.18, 0.17], 15, 440)
    set_notes(s, "[1] and [2] were read; [3] to [6] are listed in the batch survey but could not be retrieved, so no results are claimed for them.")

    # 6 research gap
    s = S[5]
    rows = [["Author / Paper", "Research Gap"],
            ["Base paper [1]", "Predicts vulnerable functions offline from code metrics; no per-pull-request use, explanation, patch or repository context; C/C++ only"],
            ["Sphinx [2]", "Trains an LLM to write review comments; does not combine static security rules with a metric model or verify fixes"],
            ["Static analysers", "Report known patterns as raw warnings, many of low value, without impact, flow or a checked fix"],
            ["Manual review", "Slow, uneven on security, and blind to the wider codebase when judging a small diff"],
            ["All of the above", "No single reviewer joins rules, a metric model, explained findings, verified patches, a verdict and a repository assistant, with measured results"]]
    table(s, rows, [0.26, 0.74], 17, 440)
    set_notes(s, "The batch survey states the gap: the signal is predicted offline, and it must be applied inside the live pull-request workflow together with agentic review that explains each finding.")

    # 7 problem statement
    s = S[6]
    fill(s, 7, ["Manual pull-request review does not scale with the pace of modern development.",
                "Security review depends on whether the reviewer happens to have security expertise.",
                "Static analysers overwhelm developers with warnings that are mostly false positives.",
                "Reviewers lack the context of a large codebase when judging a small diff.",
                "Need: an autonomous reviewer that analyses each pull request for security and quality using codebase-aware multi-agent reasoning, and whose behaviour is measured."], 21, bullet=True)

    # 8 objectives
    s = S[7]
    fill(s, 7, ["Extract changed files and diff context automatically when a pull request is opened.",
                "Compute source-code metrics for changed files and use them to rank vulnerability likelihood.",
                "Perform AST-based static analysis to detect injection, access-control and validation flaws.",
                "Deploy multiple agents covering security, code quality and maintainability concerns.",
                "Retrieve related code from the repository so review comments reflect the wider codebase.",
                "Post consolidated review comments and evaluate precision against manually reviewed pull requests (not yet measured: no labelled set)."], 21, bullet=True)

    # 9 proposed system
    s = S[8]
    fill(s, 7, ["1. Sign in and choose the input - the built-in demo, a pasted diff, two file versions, a GitHub PR link, or files / zip / folder.",
                "2. Ingestion Agent - changed files, hunks, line ranges and languages.",
                "3. Static Analysis Agent - AST + taint rules for Python (30 security rules), regex rules for JS, PHP, Java, and quality checks.",
                "4. ML Risk Agent - 14 metrics per changed function -> vulnerability probability, likely type, top drivers.",
                "5. Security Reasoner and Remediation Agent - CWE / OWASP explanation, impact, flow; a patch kept only if a re-scan confirms it.",
                "6. Quality Reviewer and Judge Agent - de-duplicate, risk score 0-100, verdict BLOCK / REQUEST_CHANGES / APPROVE.",
                "7. Review, export (Markdown, SARIF, JSON, CSV), history; and a codebase assistant that answers where a feature lives."], 18, bullet=False)

    # 10 architecture
    s = S[9]
    put_picture(s, DIAG / "01_system_architecture.png", (240, 178, 800, 424), remove_ids=(7,), name="System architecture")
    set_paras(shape(s, 8), "Layered architecture: browser, React app, FastAPI backend with the seven-agent engine, SQLite and model artefacts, optional GitHub and LLM calls")

    # 11 methodology
    s = S[10]
    put_picture(s, DIAG / "02_high_level_workflow.png", (57, 178, 1161, 424), remove_ids=(7,), name="High-level workflow")
    set_paras(shape(s, 8), "End-to-end workflow: sign in, choose input, seven agents, review, export and history")

    # 12 dataset, features, model
    s = S[11]
    set_paras(shape(s, 4), "Dataset, Features & Models")
    fill(s, 7, [f"Data: a generated Python corpus of {REP['n_functions']:,} functions, {REP['n_vulnerable']} vulnerable ({REP['class_ratio'] * 100:.2f}%): benign filler, injected flaws of eight types, look-alike safe patterns and 1% label noise. The base paper's real C/C++ data is not available.",
                f"Split 75 / 25 stratified: {REP['train_size']:,} train, {REP['test_size']:,} test (125 vulnerable). SMOTE (k = 5) + Tomek links on the training part only: {REP['balanced_train_size']:,} rows.",
                f"Features: 14 AST metrics M1-M14 (M12-M14 added for taint-style flaws); seven selectors vote; k swept 5 to 14, k = {REP['best_k']} chosen; Min-Max scaling.",
                "Models: Random Forest, Decision Tree, Logistic Regression, SVM, KNN, Naive Bayes, Gradient Boosting, XGBoost, soft-voting ensemble; a Random Forest classifies the vulnerability type.",
                f"Stored model: {BEST} (best F1). Training takes {REP['seconds']} s on a CPU.",
                "Because the corpus is synthetic, scores show the pipeline works; they are not evidence about real pull requests."], 16, bullet=True)
    set_notes(s, "The corpus is generated to resemble the base paper's imbalance; the metrics are measured from the real generated source code.")

    # 13 decision workflow
    s = S[12]
    put_picture(s, DIAG / "10_algorithm_flowchart.png", (200, 178, 880, 424), remove_ids=(7,), name="Algorithm flowchart")
    set_paras(shape(s, 8), "Core algorithm review_changes(): ingest, static analysis, ML risk, reasoner, remediation, quality, judge, score and verdict")

    # 14 user-system interaction
    s = S[13]
    set_paras(shape(s, 4), "User\u2013System Interaction")
    put_picture(s, DIAG / "09_sequence_diagram.png", (270, 178, 740, 424), remove_ids=(7,), name="Sequence diagram")
    set_paras(shape(s, 8), "Sequence: start a review, poll the job with live agent messages, load the review, export a report")

    # 15 database
    s = S[14]
    put_picture(s, DIAG / "06_er_diagram.png", (200, 178, 880, 424), remove_ids=(7,), name="ER diagram")
    set_paras(shape(s, 8), "Entity\u2013relationship diagram: users (app.db) and reviews with a JSON payload (history.db)")

    # 16 implementation highlights
    s = S[15]
    rows = [["Module", "Implementation"],
            ["Ingestion", "Unified diffs, before/after versions, GitHub PR and repo zip, uploaded files, zip, local folder -> changed files with line ranges"],
            ["Static analysis", "Python AST visitor with taint tracking; 30 security rules mapped to CWE and OWASP; regex rules for JS, PHP, Java; 12 quality rules"],
            ["ML risk", "14 metrics per function; XGBoost probability; Random Forest type; top drivers from the metric z-scores against safe code"],
            ["Remediation", "Parameterised SQL, list-form subprocess and other rewrites applied to a copy and re-scanned; kept only when the finding is gone"],
            ["Judge", "Score = 100 (1 - exp(-W / 18)); in-diff findings drive BLOCK, REQUEST_CHANGES or APPROVE"],
            ["Codebase assistant", "TF-IDF + concept scores + import / call graph with personalised PageRank; cited chunks, map and onboarding guide"],
            ["Web app + API", "React 19 + FastAPI, 25 routes, bcrypt + JWT, background jobs; 17 pytest tests"]]
    table(s, rows, [0.18, 0.82], 16, 440)

    # 17 results
    s = S[16]
    set_paras(shape(s, 4), "Results \u2013 Model and Detector Performance")
    rows = [["Measure", "Result", "Note"],
            [f"{BEST} (stored model) F1", f3(best["f1"]), f"precision {f3(best['precision'])}, recall {f3(best['recall'])}, accuracy {f3(best['accuracy'])}, AUC {f3(best['auc'])}"],
            ["Random Forest (base paper's choice) F1", f3(rf["f1"]), f"precision {f3(rf['precision'])}, recall {f3(rf['recall'])}, AUC {f3(rf['auc'])}"],
            ["Other seven models, F1 range", f"{f3(min(others))} - {f3(max(others))}", "Voting ensemble best of them; Naive Bayes lowest"],
            ["Five corpus seeds, mean F1", f"XGB {sm['XGBoost|f1']['mean']:.3f}, RF {sm['Random Forest|f1']['mean']:.3f}", f"sd {sm['XGBoost|f1']['sd']:.3f} / {sm['Random Forest|f1']['sd']:.3f}; headline is the best seed"],
            ["Rule engine alone, F1", f3(EV["rule_engine"]["f1"]), f"precision {f3(EV['rule_engine']['precision'])}, recall {f3(EV['rule_engine']['recall'])}"],
            ["Either rules or ML flags, F1", f3(EV["combined_or"]["f1"]), f"precision {f3(EV['combined_or']['precision'])}, recall {f3(EV['combined_or']['recall'])}"],
            ["Vulnerability-type accuracy", f3(REP["type_accuracy"]), "8 types; deserialization F1 0.207, SQL injection 0.875"]]
    table(s, rows, [0.34, 0.2, 0.46], 16, 420)
    set_paras(shape(s, 8), "1,500 held-out functions (125 vulnerable) of a generated Python corpus; not comparable with real C/C++ code; rule scores are optimistic because the same team wrote the generator and the rules.")
    set_notes(s, "All numbers come from autosecai/artifacts/report.json and docs/measured_facts.json. Seed 7 reproduces the stored report exactly.")

    # 18 confusion matrix + ROC
    s = S[17]
    set_paras(shape(s, 4), "Confusion Matrix & ROC Curves")
    two_pics(s, CHARTS / "confusion.png", CHARTS / "roc.png", "XGBoost: 6 false alarms, 28 missed, 97 caught of 125", f"ROC: AUC {f3(best['auc'])} XGBoost, {f3(rf['auc'])} Random Forest")
    set_notes(s, "Rows are the truth. Recall is 0.776 for both models: about a fifth of the vulnerable functions are missed by the ML model alone, which is why the rule engine runs beside it.")

    # 19 threshold and parameter analysis
    s = S[18]
    set_paras(shape(s, 4), "Parameter & Ablation Analysis")
    two_pics(s, CHARTS / "k_sweep.png", CHARTS / "ablation.png", "Feature count k: validation F1 0.566 at k = 5, 0.786 at k = 13", "Random Forest: the paper's top-5 setting scores F1 0.645, all 14 metrics 0.847")
    set_notes(s, "The ablation changes the metric count and the balancing together, so their effects are confounded; the sweep shows the extra taint metrics matter. Balancing did not raise recall on this corpus.")

    # 20 performance & testing
    s = S[19]
    pa = LAT["per_agent_median_s"]
    rows = [["Measurement", "Result"],
            ["Full review of the demo PR (20 runs)", f"median {LAT['median_s']} s, p95 {LAT['p95_s']} s; ML Risk Agent {pa['ML Risk Agent']} s, Remediation {pa['Remediation Agent']} s, Static {pa['Static Analysis Agent']} s"],
            ["Rule engine speed", f"{EV['rule_engine_seconds_per_function_ms']} ms per function"],
            ["HTTP latency (localhost)", "health 1.2 ms, review list 3.2 ms, lab analyse 5.2 ms, login 60 ms (bcrypt), demo review job to done 436 ms"],
            ["Codebase assistant", f"{RET['top1']} of {RET['questions']} questions put the right file first, {RET['top3']} of {RET['questions']} in the top 3; index 14 ms, answer 2 ms (small demo repo)"],
            ["Automated tests", "17 pytest tests pass (10 engine, 7 API); 16 of 16 API smoke checks; 22 of 22 browser flows, 0 console errors"],
            ["Build and layout", "strict TypeScript + Vite build passes; no horizontal overflow at 360 and 768 px, light and dark"],
            ["Not measured", "precision on real reviewed PRs, other scanners, LLM text mode, real GitHub PR fetch, load, user study"]]
    table(s, rows, [0.24, 0.76], 16, 430)
    set_notes(s, "Timings are on an i5-12500 with 15.7 GB RAM and no GPU, measured when the machine was idle.")

    # 21 comparison
    s = S[20]
    rows = [["System", "Data", "Method", "Output", "Reported / measured", "Extras"],
            ["Base paper [1]", "21,759 C/C++ functions, 10 real projects, 275 vulnerable", "14 metrics, selectors, SmoteTomek, 9 models", "Vulnerable or not; type", "Random Forest F1 0.961, accuracy 0.96", "Metric importance"],
            ["Sphinx [2]", "Pull-request review comments", "LLM trained with checklist rewards", "Review comments", "Up to 40% higher checklist coverage", "Benchmark"],
            ["AutoSecAI", f"Generated Python corpus: {REP['n_functions']:,} functions, {REP['n_vulnerable']} vulnerable", "AST / taint rules + 14-metric XGBoost + 7 agents", "Verdict, explained findings, verified patches", f"ML F1 {f3(best['f1'])}; rules {f3(EV['rule_engine']['f1'])}; either {f3(EV['combined_or']['f1'])}", "Codebase assistant, exports, web app"]]
    table(s, rows, [0.12, 0.2, 0.2, 0.17, 0.17, 0.14], 14, 400)
    set_paras(shape(s, 8), "Different language, data origin and class balance: the scores are not comparable and we do not claim to beat the paper's F1 of 0.961. The comparison is of approach.")

    # 22-24 outputs
    s = S[21]
    set_paras(shape(s, 4), "Output \u2013 Review Verdict and Findings")
    two_pics(s, crop("26_review_verdict.png", 0, 0, 270), crop("07_review_findings.png", 80, 0, 270), "BLOCK: risk 99, 15 security findings, 7 verified auto-fixes", "Findings with CWE, explanation, impact and fix", box_h=405, top=176)
    s = S[22]
    set_paras(shape(s, 4), "Output \u2013 Function Risk and Agent Trace")
    two_pics(s, crop("08_review_function_risk.png", 80, 0, 270), crop("10_review_agent_trace.png", 80, 0, 270), "ML risk per changed function with the metrics that drive it", "Audit trail: what each of the seven agents did", box_h=405, top=176)
    s = S[23]
    set_paras(shape(s, 4), "Output \u2013 Codebase Assistant and Model Lab")
    two_pics(s, crop("12_codebase_ask.png", 110, 0, 270), crop("18_lab_roc_confusion.png", 0, 0, 270), "Where is the authentication mechanism implemented?", "Model Lab: base-paper reproduction, ROC and confusion matrix", box_h=405, top=176)

    # 25 conclusion
    s = S[24]
    fill(s, 7, ["Built AutoSecAI: a seven-agent pull-request reviewer that explains findings, offers verified patches and gives a verdict, plus a repository assistant and a web application.",
                f"Stored XGBoost model F1 {f3(best['f1'])} on 1,500 held-out functions of a generated corpus (five-seed mean {sm['XGBoost|f1']['mean']:.3f}); rule engine {f3(EV['rule_engine']['f1'])}; either detector {f3(EV['combined_or']['f1'])}.",
                f"A full demo review runs in {LAT['median_s']} s; the demo PR scores {DEMO['score']}/100 and BLOCK with 7 of 7 patches verified; 17 pytest tests, 16 API checks and 22 browser flows pass.",
                "Limits: synthetic training data; Python-only deep analysis; no labelled real-PR benchmark; no comparison with other scanners; LLM text mode and real GitHub fetch untested; reviews are not scoped to a user."], 22, bullet=True)

    # 26 future scope
    s = S[25]
    fill(s, 7, ["Collect manually reviewed pull requests and measure the precision and recall of each agent against them.",
                "Train on real labelled vulnerable Python code; compare with Bandit, Semgrep or CodeQL on the same code.",
                "Extend the AST and taint analysis beyond Python; post comments back to GitHub automatically.",
                "Evaluate the optional LLM explanation mode; add a vector-embedding retriever to the codebase assistant.",
                "Per-user review history, a hosted database, rate limiting, load testing and a user study."], 24, bullet=True)

    # 27 references
    s = S[26]
    fill(s, 7, [f"[{i}] {r}" for i, r in enumerate(REFERENCES[:12], 1)], 15, min_pt=7.5, bullet=True, gap_pt=2)

    set_notes(S[27], "Thank you. Questions are welcome.")
    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
