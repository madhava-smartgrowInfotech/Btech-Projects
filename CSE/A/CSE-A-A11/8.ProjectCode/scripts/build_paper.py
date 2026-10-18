"""Build 19.IEEE_Paper.docx: a six-page, two-column IEEE conference paper on AutoSecAI.

    python scripts/build_paper.py        then   python scripts/paper_finish.py   (page count + PDF through Word)

Every number comes from docs/REPORT_FACTS.md (via doc_common.py: report.json, measured_facts.json, API latency).
References are numbered by first appearance: the document is built twice, the first pass only collects the order.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paper_kit as K  # noqa: E402
from doc_common import (API, BATCH, BEST, DEMO, EV, LAT, MODELS, REFERENCES, REP, RET, SEEDS, STUDENTS, f3)  # noqa: E402

OUT = BATCH / "19.IEEE_Paper.docx"
PAPER = Path(__file__).resolve().parents[1] / "docs" / "paper"

# citation key -> index into doc_common.REFERENCES
KEYS = dict(base=0, sphinx=1, cnnlstm=2, realtime=3, multimodal=4, dcode=5, yamaguchi=9, sklearn=19, vuldee=11, vudenc=12,
            smote=14, tomek=15, batista=16, rf=17, xgb=18, mccabe=20, tfidf=21, pagerank=22, cwe=24, sarif=25)

K.TITLE = "AutoSecAI: Multi-Agent Pull Request Analysis for Security, Vulnerability Risk and Code Quality"
K.KEYWORDS = ("automated code review, codebase intelligence, multi-agent systems, pull requests, source-code metrics, static analysis, "
              "vulnerability prediction")
xg, rfm = MODELS[BEST], MODELS["Random Forest"]
sx = SEEDS["summary"]
K.ABSTRACT = (
    "Pull requests are the main entry point for security flaws, yet manual review is slow, depends on the reviewer's security "
    "expertise, and static analysers bury developers in warnings. This paper presents AutoSecAI, a pull-request analyser in which "
    "seven cooperating agents process a change: ingestion, AST and taint-based static analysis, a machine-learning risk agent that "
    "scores every changed function from 14 source-code metrics, a security reasoner, a remediation agent that verifies each patch by "
    "re-scanning it, a quality reviewer and a judge that issues a risk score and a merge verdict. A codebase assistant answers "
    "questions about a repository with cited code. The risk model reproduces a published metrics-based pipeline on Python; because "
    "the original C datasets are unavailable it is trained on a generated corpus of 6,000 functions. On 1,500 held-out functions "
    f"XGBoost reaches F1 {f3(xg['f1'])} (mean {sx['XGBoost|f1']['mean']:.3f} over five corpora), the rule engine {f3(EV['rule_engine']['f1'])}, "
    f"and flagging by either detector {f3(EV['combined_or']['f1'])}. The built-in demo pull request is reviewed in a median of "
    f"{LAT['median_s']:.3f} s with every patch verified. The data are generated, no labelled real-PR benchmark exists, and the optional "
    "language-model layer was not measured."
)


class Paper(K.Paper):
    """Paper with this project's author block and a reference list numbered by first appearance."""

    def __init__(self, order):
        super().__init__()
        self.order = order          # citation keys in numbering order (empty on the first pass)
        self.seen: list[str] = []

    def c(self, *keys):
        nums = []
        for k in keys:
            if k not in self.seen:
                self.seen.append(k)
            n = (self.order.index(k) + 1) if self.order else self.seen.index(k) + 1
            nums.append(f"[{n}]")
        return ", ".join(nums)

    def front(self):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(K.TITLE); K.font(r, 24)
        p.paragraph_format.space_after = Pt(8)
        authors = [(n, f"Roll No. {rn}") for n, rn in STUDENTS] + [("P. KISHORE", "[DESIGNATION], Project Guide")]
        t = self.d.add_table(rows=1, cols=len(authors))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = 7.12 / len(authors)
        for ci, (name, sub) in enumerate(authors):
            cell = t.rows[0].cells[ci]
            cell.width = Inches(w)
            first = True
            for line, bold in ((name, True), (sub, False), ("Dept. of Computer Science and Engineering", False),
                               ("TKR College of Engineering & Technology (Autonomous)", False), ("[CITY]", False), ("[EMAIL]", False)):
                par = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = par.add_run(line); K.font(run, 9, bold=bold)
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        from docx.enum.section import WD_SECTION
        from docx.oxml import OxmlElement
        sec = self.d.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sec._sectPr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    def references(self):
        self.h1("References")
        for i, key in enumerate(self.seen, 1):
            p = self.d.add_paragraph()
            r = p.add_run(f"[{i}]\t{REFERENCES[KEYS[key]]}"); K.font(r, 8)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            p.paragraph_format.space_after = Pt(1.5)


def content(P):
    c = P.c
    P.front()
    P.abstract()

    # ---------------------------------------------------------------- I
    P.h1("I. Introduction")
    P.para("Every change that reaches a shared code base arrives as a pull request, so each one is a possible entry point for an "
           "injection flaw, a leaked secret or a weak cryptographic call. Review is done by people: it is slow, it varies with the "
           "reviewer's security knowledge, and it lacks the context of a large repository when judging a small diff. Static "
           "analysers catch known patterns but return many low-value warnings with little explanation, and source-code metrics, "
           "which correlate with vulnerability, are used offline on whole projects rather than inside the review. Language models "
           "can explain a change in context, but a single general reviewer is hard to audit.")
    P.para("AutoSecAI answers this with specialised agents, each with one narrow job and an auditable trace. Static rules supply "
           "precise, explainable findings; a metric model ranks functions by how much their structure resembles that of "
           "vulnerable code and corroborates or questions the rules; a remediation agent proposes a fix and proves that the rule no "
           "longer fires; and a judge turns the evidence into one decision a reviewer can act on. Nothing is merged, posted or "
           "changed automatically: the system advises and the reviewer decides.")
    P.para(f"The base paper {c('base')} predicts vulnerable C/C++ functions from 14 code metrics, and recent work moves towards "
           f"model-based pull-request review {c('sphinx')}. The gap is applying the metric signal inside the live pull-request "
           "workflow together with rule-based analysis, explanation and a checked fix. This paper presents AutoSecAI. Its "
           "contributions are:", after=2)
    P.numbered([
        "a seven-agent pull-request pipeline that combines an AST and taint rule engine (30 security rules mapped to CWE and OWASP) with a metric-based risk model and issues a risk score and a merge verdict;",
        f"a Python adaptation of the 14 metrics of the base paper {c('base')} and a reproduction of its selection, balancing and nine-model pipeline, with three added taint metrics;",
        "patches that are verified by re-scanning, exports (Markdown, SARIF, JSON, CSV), and a codebase assistant that answers repository questions with cited code;",
        "an honest evaluation: rule engine against ML on the same held-out functions, five corpus seeds, latency and automated tests.",
    ])
    P.para("Section II reviews related work, Section III describes the method, Section IV reports the results and Section V concludes.", before=2)
    P.h2("A. System Hardware")
    P.para("All measurements were taken on one desktop: Intel Core i5-12500 (12 threads), 15.7 GB of memory, no discrete graphics "
           "card, Windows 11. Models are trained and run on the CPU.")
    P.h2("B. Software Requirements")
    P.para(f"Python 3.11.9 with scikit-learn 1.5.0 {c('sklearn')}, XGBoost 2.0.3 {c('xgb')}, pandas, NetworkX and the standard "
           "ast module; FastAPI with PyJWT, bcrypt and SQLite for the service; React 19, TypeScript, Vite and Tailwind CSS for the "
           "web interface; pytest and Playwright for tests. An Anthropic key is optional and unused in every measurement.")

    # ---------------------------------------------------------------- II
    P.h1("II. Literature Survey")
    P.para(f"**Base paper.** Abbasi and Shameli-Sendi {c('base')} study which source-code factors predict vulnerabilities and their "
           "types. Their data are 21,759 functions from ten C/C++ projects (SQLite, FreeType, LibTIFF, Libxslt, Binutils, FFmpeg, "
           "ImageMagick, OpenSC, rdesktop and OpenSSL), of which 275 (1.26 %) are vulnerable, with types taken from CVE and NVD. A "
           "fuzzy parser builds a code property graph from which 14 metrics are computed, among them cyclomatic complexity "
           f"{c('mccabe')}, caller-argument count, pointer arithmetic, nesting depth and control dependency. Seven feature-selection "
           f"methods vote for the five most important metrics (M5, M6, M1, M9 and M11); the imbalance is treated with SmoteTomek "
           f"{c('smote', 'tomek', 'batista')}, which yields 21,474 balanced samples; values are Min-Max normalised; and nine models "
           f"are tuned by grid search and repeated k-fold validation. Random Forest {c('rf')} performs best: accuracy 0.96, "
           "precision 0.948, recall 0.974 and F1 0.961, and 0.972 accuracy for one vulnerability type (denial of service). The work is "
           "an offline function classifier for C/C++: it produces no explanation, no fix and no integration with review.")
    P.para(f"**Pull-request review.** Sphinx {c('sphinx')} builds context-rich review comments by comparing pseudo-modified and "
           "merged code, evaluates them with a checklist of actionable verification points instead of BLEU or ROUGE, and trains a "
           "model with rule-based checklist rewards; it reports up to 40 % higher checklist coverage than baselines. Four further "
           f"works are listed in our batch survey by title only: a CNN-LSTM hybrid for vulnerability detection {c('cnnlstm')}, real-time AI "
           f"code security auditing with remediation {c('realtime')}, a multi-modal framework with code simplification and meta-learning "
           f"{c('multimodal')}, and DcodeBERT for detection and repair {c('dcode')}. We could not retrieve their text and state no "
           "result for them.")
    P.para(f"**Learning-based vulnerability detection.** Earlier work learns from abstract syntax trees {c('yamaguchi')}, deep "
           f"learning on code {c('vuldee')} and, for Python specifically, on a natural code base "
           f"{c('vudenc')}. We read only their titles here. Most are trained on C/C++ or on token sequences and none returns a "
           "reviewer-facing explanation and patch.")
    P.para(f"**Techniques reused.** SMOTE {c('smote')} creates minority samples by interpolating between a sample and its nearest "
           f"minority neighbours, Tomek links {c('tomek')} are pairs of nearest neighbours from different classes whose majority member "
           f"is removed, and Batista et al. {c('batista')} compare their combination with other balancing methods. Random Forest "
           f"{c('rf')} and gradient-boosted trees {c('xgb')} are the standard learners for tabular metric data; they are cheap to run "
           f"on a CPU, which matters when a model must score every function of every pull request. Repository retrieval builds on "
           f"TF-IDF weighting {c('tfidf')} and on link analysis over a graph {c('pagerank')}.")
    P.table([["Work", "Approach", "Data", "Reported result", "Limitation"],
             ["[1] base paper", "14 CPG metrics, 7 selectors, SmoteTomek, 9 models", "21,759 C/C++ functions, 275 vulnerable", "RF acc 0.96, F1 0.961", "C/C++; offline; no explanation"],
             ["Sphinx", "Checklist benchmark, rule-reward training", "PR review data", "Up to 40 % more checklist coverage", "Comments only; data not yet released"],
             ["CNN-LSTM; AI audit", "Hybrid network; detect and remediate", "Not retrieved", "Not retrieved", "Titles only"],
             ["Multi-modal; DcodeBERT", "Meta-learning; detect and repair", "Not retrieved", "Not retrieved", "Titles only"],
             ["Deep / AST learners", "AST, deep, Python code learning", "Not read", "Not read", "Not read here"]],
            "Related work", widths=[0.17, 0.26, 0.21, 0.19, 0.17], size=6.5)
    P.para("**Gap.** The metric signal has been used offline, review models produce comments without a security engine or code "
           "metrics, and we found no system that joins metric-based risk, taint rules, verified patches and repository-aware "
           "answers in the pull-request workflow, evaluated openly including its weaknesses. AutoSecAI addresses this combination.")

    # ---------------------------------------------------------------- III
    P.h1("III. Methodology")
    P.h2("A. Architecture and Pipeline")
    P.para("A React single-page application talks to a FastAPI service that wraps the analysis engine (Fig. 1). Long operations run "
           "as background jobs whose agent messages the interface polls, so the user watches each agent finish. Accounts (bcrypt, "
           "JWT) live in one SQLite file and saved reviews, stored as JSON, in another. A review passes through seven agents (Fig. 2).")
    P.figure(PAPER / "fig1_architecture.png", "System architecture of AutoSecAI.", width=3.4)
    P.h2("B. Metrics, Corpus and Risk Model")
    P.para(f"The Python ast module replaces the code property graph. Metrics M1-M11 keep the families of the base paper (cyclomatic "
           "complexity, size, variables, calls, parameters, index arithmetic in place of pointer arithmetic, I/O calls, nesting, "
           "loops, control dependency). Because Python flaws are injections rather than memory errors, three taint-oriented metrics "
           "are added: M12 risky API calls, M13 string-built sinks and M14 untrusted or secret-like sources. The C datasets are not "
           f"distributable, so a corpus of {REP['n_functions']:,} Python functions is generated: random benign logic of varying "
           "complexity, injected patterns of eight types (SQL and command injection, code execution, deserialisation, path "
           "traversal, XSS, hard-coded secret, weak cryptography), look-alike safe code as hard negatives, and 1 % label noise. "
           f"{REP['n_vulnerable']} functions ({REP['class_ratio'] * 100:.2f} %) are vulnerable. Metrics are measured on the real generated source.")
    P.para("The metric extractor visits each function once. Cyclomatic complexity counts decision points and boolean operators plus "
           "one; nesting and control dependency are the deepest block depth and the largest number of guarding conditions on a "
           "statement; M12 counts calls from a list of dangerous APIs (eval, exec, os.system, subprocess, pickle and yaml loaders, "
           "weak hashes and random generators); M13 counts call arguments built by f-string, concatenation, percent or format; and "
           "M14 counts request-like sources and secret-named literals. Each metric is a plain count, so a reviewer can see why a "
           "function scored high.")
    P.para(f"Seven selectors (Pearson, linear SVM, Lasso, chi-squared, two RFE variants, variance) vote on the training split. The paper "
           f"keeps five metrics; here the number k is validated from 5 to 14 and k = {REP['best_k']} gives the best validation F1. "
           f"Values are Min-Max scaled and the minority class is balanced by SMOTE {c('smote')} followed by Tomek-link removal "
           f"{c('tomek')}, implemented in house: a synthetic sample is given by Eq. (1), with λ drawn uniformly between 0 and 1 and x_nn "
           f"one of the five nearest minority neighbours, and a Tomek link is removed by dropping its majority member. Nine models are trained on {REP['balanced_train_size']:,} balanced rows and "
           f"tested on 1,500 untouched functions: Random Forest (300 trees), decision tree, logistic regression, RBF SVM, KNN, "
           f"naive Bayes, gradient boosting, XGBoost and a soft-voting ensemble. The best F1 model, XGBoost, is stored. A second "
           "Random Forest classifies the vulnerability type from all 14 metrics.")
    P.eq("x_{new} = x + λ (x_{nn} − x),   λ ~ U(0, 1)", 1)
    P.para("Quality is reported as precision P = TP / (TP + FP), recall R = TP / (TP + FN) and their harmonic mean F1 = 2PR / (P + R), "
           "because with 8.35 % positives accuracy alone would reward a model that flags nothing.", indent=False)
    P.h2("C. Rule Engine and Agents")
    P.para(f"The static agent walks the AST with taint tracking: values from request objects, input, sys.argv and the environment are "
           "sources, calls such as cursor.execute, os.system, eval, pickle.loads and file opens are sinks, and sanitisers "
           f"(parameter binding, shlex, escaping) stop the flow. It holds 30 security rules, each with a CWE {c('cwe')} and OWASP "
           "class, severity, impact, remediation and a safe pattern, plus 12 quality rules; JavaScript, PHP and Java get regex rules. "
           "The ML agent scores each changed function and reports its three largest drivers. The reasoner merges both signals and "
           "marks a finding as corroborated when the model gives p ≥ 0.5. The remediation agent rewrites the code (for example "
           "parameterising a query), re-parses it and re-runs the rules, and keeps the patch only if the finding is gone.")
    P.para("Every agent appends an entry to a trace (name, role, seconds, summary, details) that is stored with the review, so the "
           "interface can show what each agent did. Findings carry the rule, CWE and OWASP class, severity, confidence, the code line "
           "and a source-to-sink description; findings on lines the pull request did not touch are kept but marked as pre-existing "
           "and weighted down, so a reviewer is not blamed for old code. A minimum-severity filter and an ML on/off switch are "
           "options of the review.")
    P.figure(PAPER / "fig2_flow.png", "Review flow of the seven agents and the verdict.", width=3.4)
    P.para("The judge removes duplicates with the same file, line and CWE, and gives a finding the weight w = s·κ·γ, where s is the "
           "severity weight (critical 10, high 6, medium 3, low 1, info 0.2), κ the confidence weight (1.0, 0.7, 0.4) and γ = 1 inside "
           "the changed lines and 0.15 outside. With W the sum of the weights, the risk score is", indent=False)
    P.eq("score = round(100 · (1 − exp(−W / 18)))", 2)
    P.para("The verdict is BLOCK if a changed high or critical security finding has medium or high confidence, or the score is at "
           "least 70; REQUEST_CHANGES if a changed medium finding exists or the score is at least 30; otherwise APPROVE.", indent=False)
    P.h2("D. Codebase Assistant")
    P.para(f"A repository is cut into function, class and window chunks. Retrieval scores a chunk with TF-IDF {c('tfidf')} over "
           "split identifiers (unigrams and bigrams), adds 0.30 times a concept affinity for the detected topic (authentication, "
           "routing, database, user interface, configuration, tests, security), boosts exact identifier and path matches, and adds "
           f"a personalised PageRank {c('pagerank')} boost (α = 0.6, weight 0.08) over the import and call graph. The answer lists "
           "the evidence chunks with the reasons they matched; an LLM may rewrite the text only if a key is given.")
    P.h2("E. Implementation")
    P.para("The service has 25 route handlers (7 public) with HS256 tokens valid for 12 h and bcrypt with 10 rounds. Reviews accept a "
           "diff, a before and after pair, a GitHub pull-request URL, or files, a zip or a folder, and can be exported as Markdown, "
           f"SARIF 2.1.0 {c('sarif')}, JSON or CSV. The engine has 2,881 non-blank lines, the service 502 and the React interface 6,525.")
    P.para("Inputs are validated before work starts: an empty or malformed diff, an unreadable folder or an unreachable GitHub "
           "address is rejected with a readable message, uploads are limited to 300 files per review and 800 KB per file, and "
           "passwords need eight characters. Jobs are held in memory and pruned after an hour; at most six repository indexes are "
           "kept. The tool never posts a comment or changes a repository: every output is a report that a person reads, and the "
           "history of reviews can be deleted from the interface.")

    # ---------------------------------------------------------------- IV
    P.h1("IV. Results")
    P.h2("A. Test Conditions")
    P.para(f"Models are evaluated on a stratified 25 % split (1,500 functions, 125 vulnerable) of the generated corpus, with the training "
           f"part balanced and the test part untouched. To check stability the pipeline is repeated on five independently generated "
           f"corpora (k fixed at {SEEDS['k']}). The rule engine is scored on the same 1,500 functions: a function counts as flagged if it "
           "has a security finding of medium severity or higher. Latency is the median of 20 runs on the demo pull request (one file, "
           f"+67/−1 lines, {DEMO['stats']['functions_scored']} functions) after a warm-up. Codebase retrieval is checked with 12 "
           "hand-written questions on the built-in demo repository (10 files). All numbers are measured on the machine of Section I-A.")
    P.table([["Model", "Acc.", "Prec.", "Rec.", "F1", "AUC"]] +
            [[n.replace(" (RBF)", ""), f3(m["accuracy"]), f3(m["precision"]), f3(m["recall"]), f3(m["f1"]), f3(m["auc"])] for n, m in MODELS.items()],
            "Nine models on 1,500 held-out functions", widths=[0.33, 0.13, 0.14, 0.13, 0.13, 0.14], size=7)
    P.h2("B. Risk Model")
    P.para(f"XGBoost gives the best F1 ({f3(xg['f1'])}, precision {f3(xg['precision'])}, recall {f3(xg['recall'])}) and Random Forest the "
           f"highest AUC of the single models ({f3(rfm['auc'])}); only 6 safe functions are flagged and 28 vulnerable ones missed. Over five "
           f"corpora XGBoost reaches F1 {sx['XGBoost|f1']['mean']:.3f} ± {sx['XGBoost|f1']['sd']:.3f} (range {sx['XGBoost|f1']['min']:.3f}-"
           f"{sx['XGBoost|f1']['max']:.3f}) and Random Forest {sx['Random Forest|f1']['mean']:.3f} ± {sx['Random Forest|f1']['sd']:.3f}, so the "
           "headline value is the best of five. With the paper's five metrics and SMOTE-Tomek the Random Forest F1 falls to "
           f"{f3(REP['ablation']['RF - top-5 metrics + SmoteTomek (paper setting)']['f1'])}, against "
           f"{f3(REP['ablation']['RF - all 14 metrics, no balancing']['f1'])} with all 14 metrics and no balancing and {f3(rfm['f1'])} with "
           f"{REP['best_k']} metrics and balancing. The settings differ in both metric count and balancing, so the effect of balancing "
           "alone is not isolated; recall is identical (0.776) with and without it. The type classifier reaches "
           f"{f3(REP['type_accuracy'])} accuracy over eight types: SQL injection (F1 0.875) and XSS (0.872) are recognised, deserialisation "
           "(0.207) and path traversal (0.333) are not.")
    P.para("The metrics that matter most are those that were added for Python: in the Random Forest of the Model Lab (the shipped scorer is XGBoost) the largest importances are "
           f"risky API calls ({f3(REP['importance']['M12 Risky API calls'])}), untrusted sources "
           f"({f3(REP['importance']['M14 Untrusted/sensitive sources'])}) and string-built sinks "
           f"({f3(REP['importance']['M13 String-built sinks'])}); the three together carry about 46 % of the total, whereas cyclomatic "
           f"complexity carries {f3(REP['importance']['M1 Cyclomatic complexity'])}. The base paper's complexity-led ranking therefore does not "
           "transfer to injection flaws in Python.")
    P.figure(PAPER / "fig3_results.png", "F1 of nine models (a); rule engine, ML model and their combinations (b).", width=3.4)
    P.h2("C. Rule Engine, ML Model and Their Combination")
    P.para(f"On the same functions the rule engine reaches precision {f3(EV['rule_engine']['precision'])}, recall {f3(EV['rule_engine']['recall'])} and "
           f"F1 {f3(EV['rule_engine']['f1'])}; the XGBoost model {f3(EV['ml_model']['f1'])}. Flagging a function when either detector "
           f"does reaches recall {f3(EV['combined_or']['recall'])} and F1 {f3(EV['combined_or']['f1'])} (Fig. 3b), while requiring both gives "
           f"precision {f3(EV['combined_and']['precision'])} at recall {f3(EV['combined_and']['recall'])}. The detectors are "
           "complementary, which supports using both in one pipeline. The rule figures are optimistic: the generator and the rules "
           "were written by the same team around the same eight patterns, and 1 % label noise creates positives without vulnerable code.")
    P.para(f"In counts, XGBoost misses {EV['ml_model']['fn']} of 125 vulnerable functions and flags {EV['ml_model']['fp']} of 1,375 safe ones; the "
           f"rule engine finds {EV['rule_engine']['tp']}, misses {EV['rule_engine']['fn']} and flags {EV['rule_engine']['fp']}; the union finds "
           f"{EV['combined_or']['tp']}, misses {EV['combined_or']['fn']} and flags {EV['combined_or']['fp']}. For a reviewer the trade-off is visible: the union "
           f"catches {EV['combined_or']['tp'] - EV['ml_model']['tp']} more vulnerable functions than XGBoost for {EV['combined_or']['fp'] - EV['ml_model']['fp']} more false alarms, while requiring agreement is the setting for "
           "a gate that must almost never block a safe change. In the product both signals reach the judge, which weighs the "
           "findings rather than trusting either detector alone.")
    P.h2("D. Demonstration, Speed and Tests")
    P.para(f"The built-in demo pull request (67 added lines) receives score {DEMO['score']} and verdict {DEMO['verdict']}: {DEMO['findings']} findings "
           f"(1 critical, 10 high, 4 medium, 2 info), all 7 functions flagged and all 7 generated patches verified by re-scan (Figs. 4 and 5). "
           f"A review takes a median of {LAT['median_s']:.3f} s (95th percentile {LAT['p95_s']:.3f} s); the ML agent needs {LAT['per_agent_median_s']['ML Risk Agent']:.2f} s of it. "
           f"Over HTTP a polled review job finishes in a median of {API['POST /api/reviews demo-diff -> done (polled 50 ms)']['median_ms']:.0f} ms and a "
           f"login, dominated by bcrypt, takes {API['POST /api/auth/login (bcrypt)']['median_ms']:.0f} ms. The assistant ranks the expected file first for "
           f"{RET['top1']} of {RET['questions']} questions and in the top three for {RET['top3']}; an index of 10 files builds in {RET['index_build_ms']:.0f} ms. "
           "Seventeen pytest tests, 16 API smoke checks and 22 browser flows pass with no console errors.")
    P.para("The 15 changed-line security findings span 12 weakness classes: two each for CWE-798 (hard-coded credential), CWE-89 "
           "(SQL injection) and CWE-502 (unsafe deserialisation), and one each for CWE-22, 78, 79, 95, 295, 327, 338, 489 and 918. "
           f"The static agent needs {LAT['per_agent_median_s']['Static Analysis Agent'] * 1000:.0f} ms, remediation "
           f"{LAT['per_agent_median_s']['Remediation Agent'] * 1000:.0f} ms and the ML agent {LAT['per_agent_median_s']['ML Risk Agent'] * 1000:.0f} ms, "
           "so scoring the functions through the stored model dominates a review. The question the assistant missed asked how a user is "
           "looked up by name; it returned the authentication module and a test file instead of the database module that defines the lookup. "
           "These measurements come from one idle machine and one small demonstration; they show that the pipeline is fast enough to run on "
           "every pull request, not how it scales to large changes.")
    P.figure(PAPER / "fig4_review.png", "Verdict and risk score of the demo pull request.", width=3.4)
    P.figure(PAPER / "fig5_finding.png", "A finding with its verified patch.", width=3.4)
    P.h2("E. Comparison and Limits")
    P.para(f"The base paper {c('base')} reports F1 0.961 for Random Forest on real C/C++ functions; this work reports {f3(rfm['f1'])} for its Random Forest and {f3(xg['f1'])} for the XGBoost model it ships, on generated "
           "Python functions. The language, data, class balance and parser differ, so the numbers are not comparable and no superiority is "
           "claimed; the contribution is the integration. Failure cases and limits: the corpus is synthetic, so scores are a "
           "prioritisation signal rather than proof of a flaw; the rule engine misses 37 of 125 positives; the type classifier is weak "
           "for three types; retrieval ranked the right file first for only 8 of 12 questions (a lookup question missed it entirely); the "
           "whole-file fix needs full-file input; saved reviews are not separated by account; no comparison with Bandit, Semgrep or "
           "CodeQL was run; precision against manually reviewed real pull requests, objective 6 of the project, was not measured; and "
           "the optional LLM explanations were not evaluated.")
    P.para("All figures can be regenerated: the corpus, models and report come from one training call of about 32 s, the "
           "measurement script writes the engine, seed, latency and retrieval numbers, and the screens in Figs. 4 and 5 were "
           "captured from the running application with real data. The 17 automated tests check the engine (SQL injection found and "
           "safe code not flagged, taint from a request, sanitiser stopping taint, auto-fix that re-parses, cyclomatic metric, diff "
           "parsing, pre-existing findings, ML ranking, SMOTE-Tomek balance, repository retrieval) and the service (token checks, "
           "demo login, registration rules, a full review with SARIF export and deletion, input validation, codebase questions, "
           "the Model Lab scorer).")

    # ---------------------------------------------------------------- V
    P.h1("V. Conclusion")
    P.para(f"AutoSecAI brings metric-based risk, taint rules, verified patches and repository-aware answers into one pull-request review. "
           f"On generated Python data XGBoost reaches F1 {f3(xg['f1'])} (mean {sx['XGBoost|f1']['mean']:.3f} over five corpora), the rule engine "
           f"{f3(EV['rule_engine']['f1'])}, and their union {f3(EV['combined_or']['f1'])}; the demo review runs in {LAT['median_s']:.2f} s. "
           "Against the objectives, changed files and diff context are extracted, metrics rank functions, AST rules detect injection "
           "flaws, agents cover security and quality, retrieval supplies repository context, and a consolidated verdict is produced; the "
           "precision test on real reviewed pull requests remains open. The main finding is that the two detectors are complementary. "
           "The practical value is a reviewer-facing report that is explained, patched and checked in under half a second for a "
           "small change. Future work: a labelled set of real pull requests and a comparison with existing analysers such as "
           "Bandit, Semgrep and CodeQL; real Python vulnerability data for the model, with a per-language corpus; a measured "
           "evaluation of the language-model layer and of embedding-based retrieval; and per-account storage with CI and webhook "
           "integration so that a verdict can gate a merge.")

    P.references()


def build():
    first = Paper([])
    content(first)                       # pass 1: collect the citation order
    order = list(first.seen)
    P = Paper(order)
    content(P)                           # pass 2: final numbers
    assert P.seen == order
    P.d.core_properties.title = K.TITLE
    P.d.core_properties.author = "Batch A11"
    P.d.save(OUT)
    print("saved", OUT, "| references:", len(order))


if __name__ == "__main__":
    build()
