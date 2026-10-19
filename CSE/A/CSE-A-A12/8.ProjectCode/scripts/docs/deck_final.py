"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .ch1 import OBJECTIVES
from .ch2 import LIT_TABLE
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    t = sh.top
    return t in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.subject = f"{F.TITLE}"
    notes = {}

    M = F.DEPLOYED or {}
    T = M.get("test", {})
    TR = M.get("training", {})
    DS = M.get("dataset", {})
    LAT = M.get("latency_ms_mean_cpu", {})
    B = T.get("binary_normal_vs_pneumonia", {})
    PC = T.get("per_class", {})
    E2, E3, E4 = F.E5, F.E4, F.E3   # no-CLAHE, MobileNetV3Small, patient-level
    E1 = F.E1
    py = (F.TESTS or {}).get("pytest", {})
    sm = (F.TESTS or {}).get("smoke", {})
    ui = (F.TESTS or {}).get("ui", {})
    api = F.API_LAT or {}
    mi = F.MODEL_INFO or {}
    LD = F.LEAK or {}

    # footers and numbers on every slide except the title
    for i, sl in enumerate(S, 1):
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        t = sh.text_frame.text if sh.has_text_frame else ""
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE)
        elif sh.name == "Text 2":
            K.replace_text(sh, "Final Review")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_.title()}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE}  ·  Academic Year 2026 – 2027")
    notes[1] = "Introduce the team and the project: an explainable chest X-ray classifier with a complete web application."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap & Problem Statement\nObjectives\nProposed System & Architecture\nMethodology, Dataset & Model\nDecision Workflow, Interaction & Database\nImplementation Highlights\nResults, Threshold Analysis, Performance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References")
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    acc = F.pct(T.get("accuracy", 0)) if T else "n/a"
    ftext(s, 0.55, 1.9, 12.2, 4.9, [
        ("Pneumonia is a leading infectious cause of death and chest X-ray readers are scarce where it is most dangerous; manual reading is slow and varies between readers.", {"bullet": True, "space_after": 10}),
        ("We built an end-to-end system: CLAHE preprocessing, a compact CNN that classifies Normal, Bacterial or Viral pneumonia, and a Grad-CAM heat map for every prediction.", {"bullet": True, "space_after": 10}),
        (f"The CNN ({mi.get('total_params', 0):,} parameters, CPU-trained) reaches **{acc}** three-class accuracy on the official 624-image test split of the public Kermany data; separating pneumonia from normal gives **{F.pct(B.get('sensitivity', 0))}** sensitivity and **{F.pct(B.get('specificity', 0))}** specificity. Normal films mistaken for pneumonia are the weak point.", {"bullet": True, "space_after": 10}),
        ("A FastAPI service and a React interface add login, history, dashboard, CSV export and printing; 73 automated tests pass.", {"bullet": True, "space_after": 10}),
        ("U-Net segmentation, calibrated uncertainty and external validation are not done and are stated as limits.", {"bullet": True}),
    ], size=19)
    notes[3] = "Read the first two bullets, then give the headline numbers with their protocol (official test split). End on the honest limits."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper"],
            ["Method", "MobileNet family (V3Small chosen) with FixMatch, FlexMatch, FreeMatch; Grad-CAM; Monte Carlo Dropout, test-time augmentation, ECE"],
            ["Data", "Mendeley paediatric chest X-rays (5,863; Normal / Pneumonia) and a merged 4-class set (7,135); balanced subsets, seed 42"],
            ["Reported results", "Up to 98 % (binary) and 97 % (4-class) supervised; FreeMatch 97 % with 20 % (binary) / 30 % (4-class) labels; latency 0.11 s, 116.9 M FLOPs"],
            ["Stated limits", "Benchmark data, no external validation, confusion among similar diseases, benefit of extra classes not isolated"]]
    ftable(s, 0.55, 2.9, [2.0, 10.2], rows, size=14, row_h=0.78)
    notes[4] = "A research framework on public benchmarks. We take the input size, Grad-CAM and the lightweight-model idea; we do not reproduce its semi-supervised or uncertainty methods."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"]]
    for i in (1, 2, 3, 4, 7, 8):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (CheXNet)", ""), r[1], r[2], r[3], r[4]])
    ftable(s, 0.55, 1.85, [1.9, 2.5, 2.4, 3.2, 2.2], rows, size=11, row_h=0.7)
    notes[5] = "Representative work on chest X-ray classification. Figures marked 'as summarised in [1]' come from the base paper's own review."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, [
        ("**Application:** few works go from an uploaded image to an explained, stored, reviewable prediction.", {"bullet": True, "space_after": 12}),
        ("**Pathogen type:** bacterial against viral pneumonia is far less studied than pneumonia against normal.", {"bullet": True, "space_after": 12}),
        ("**Protocol:** results are rarely stated with their split and balance, so numbers on the same dataset (80 % to 98 %) cannot be compared fairly.", {"bullet": True, "space_after": 12}),
        ("**Explanation and uncertainty:** Grad-CAM, calibration and uncertainty are usually separate post-hoc add-ons.", {"bullet": True, "space_after": 12}),
        ("**Segmentation:** restricting the classifier to the lungs is discussed but rarely combined with the rest.", {"bullet": True}),
    ], size=20)
    notes[6] = "Five gaps drawn from the survey. We address the first three directly and the fourth partly; segmentation is future work."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.95, 12.2, 2.0)
    ftext(s, 0.85, 2.0, 11.6, 1.9, "Given one frontal chest X-ray, classify it as **Normal**, **Bacterial Pneumonia** or **Viral Pneumonia**, report the probability of each class, show the image regions that supported the decision, and keep a reviewable record, all through a web interface a non-expert can use.", size=20, anchor="m")
    for i, (t, b_) in enumerate([("Input", "PNG or JPEG chest X-ray, up to 8 MB"), ("Output", "Class, confidence, three probabilities, Grad-CAM overlay"), ("Record", "History, dashboard, CSV export, printable report")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.3, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.4, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 4.9, 3.5, 1.0, b_, size=15)
    notes[7] = "The statement has three parts: classify, explain, record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion reports for each whether it was met; objective 2 (U-Net) was not."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 2.0, 12.2, 1.1, [("Upload X-ray", K.FNAVY), ("Preprocess (CLAHE)", K.FBLUE), ("CNN classify", K.GREEN), ("Grad-CAM explain", K.ORANGE), ("Store + review", K.PURPLE)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Compact three-class CNN written for the project; the same preprocessing at training and at inference.", {"bullet": True, "space_after": 10}),
        ("Grad-CAM overlay and three probabilities for every scan, stored with the scan in SQLite.", {"bullet": True, "space_after": 10}),
        ("FastAPI service reusing the original Flask project's model code; React interface with demo login, dashboard, history, CSV export and print.", {"bullet": True, "space_after": 10}),
        ("Demo Mode warning whenever the network is untrained; clinical disclaimer on every screen.", {"bullet": True}),
    ], size=18)
    notes[9] = "The proposed system in one line of boxes, then the four properties that matter."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.85, 8.3, 5.0)
    ftext(s, 9.1, 2.0, 3.7, 4.7, [
        ("Presentation: React SPA, original Flask UI", {"bullet": True, "space_after": 8}),
        ("Application: FastAPI, signed-cookie login, validation", {"bullet": True, "space_after": 8}),
        ("Inference: preprocess, CNN, Grad-CAM", {"bullet": True, "space_after": 8}),
        ("Data: SQLite, image files, weights", {"bullet": True}),
    ], size=16)
    notes[10] = "Four tiers; the model code is shared by both interfaces."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Public data", K.FNAVY), ("Label from file names", K.FNAVY), ("Preprocess + cache", K.PURPLE), ("Patient-level split", K.GREEN), ("Train CNN", K.ORANGE), ("Test once", K.FBLUE), ("metrics.json", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Shared preprocessing: the training script calls the same function as the web app.", {"bullet": True, "space_after": 10}),
        ("Validation cut by patient; official test split used once; seed 42; balanced class weights.", {"bullet": True, "space_after": 10}),
        ("Additional views: patient-disjoint part of the test split, a fully patient-level re-partition, a CLAHE ablation and a MobileNetV3Small comparison.", {"bullet": True, "space_after": 10}),
        ("Every metric, prediction and figure is written to disk and regenerated by scripts.", {"bullet": True}),
    ], size=18)
    notes[11] = "The workflow, and the additional protocols added because the official test split shares ids with training."

    # ---- 12 dataset, features, model
    s = S[11]
    set_title(s, "Dataset, Features & Model")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, "5,856", "Kermany chest X-rays (CC BY 4.0)", K.FBLUE, 26)
    ftile(s, 0.55, 3.15, 2.9, 1.1, "224 x 224", "input, CLAHE + blur, [0, 1]", K.GREEN, 24)
    ftile(s, 0.55, 4.4, 2.9, 1.1, f"{mi.get('total_params', 0):,}", "CNN parameters", K.ORANGE, 24)
    K.picture(s, F.fig("dataset_distribution.png"), 3.7, 1.85, 4.7, 2.8)
    rows = [["Layer", "Output"], ["4 x (Conv3x3 + BN + MaxPool)", "14x14x256"], ["GAP, Dense 512, Dropout 0.35", "512"], ["Dense 128, Dropout 0.25", "128"], ["Dense 3, softmax", "3 classes"]]
    ftable(s, 8.6, 1.9, [2.9, 1.3], rows, size=11, row_h=0.5)
    ftext(s, 3.7, 4.85, 9.0, 1.9, [
        (f"Training {DS.get('train', 0):,} / validation {DS.get('validation', 0):,} / test {DS.get('test', 0)} images; pathogen labels from file names.", {"bullet": True}),
        (f"Adam {M.get('hyperparameters', {}).get('learning_rate', 3e-4)}, batch 32, early stopping, class weights, augmentation (rotation, shift, shear, zoom, flip).", {"bullet": True}),
    ], size=14)
    notes[12] = "Dataset and model in one view. Pathogen labels are derived from file names, as the dataset authors encoded them."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 5.0, 5.1)
    ftext(s, 5.9, 2.0, 6.9, 4.8, [
        ("Validate: session, name, extension, 8 MB, decodable. Errors return 401 / 400 / 413 and store nothing.", {"bullet": True, "space_after": 12}),
        ("Decide: class = arg-max of the three probabilities; confidence = that probability.", {"bullet": True, "space_after": 12}),
        ("Explain: Grad-CAM for the predicted class, overlay alpha 0.4.", {"bullet": True, "space_after": 12}),
        ("Flag: accuracy against a minimum-confidence threshold is reported so low-confidence scans can be sent for review.", {"bullet": True}),
    ], size=18)
    notes[13] = "The prediction route in decision form. The confidence threshold is analysed on slide 19."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_diagram.png"), 0.55, 1.8, 7.9, 5.1)
    K.picture(s, F.diag("08_use_case_diagram.png"), 8.6, 1.8, 4.2, 5.1)
    notes[14] = "Sequence of one analysis and the use cases of the user and the developer."

    # ---- 15 database design
    s = S[14]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.9)
    rows = [["Column", "Type"], ["id", "INTEGER PK"], ["filename", "TEXT"], ["original_path", "TEXT"], ["heatmap_path", "TEXT"], ["prediction", "TEXT"], ["confidence", "REAL"], ["prob_normal / _bacterial / _viral", "REAL"], ["created_at", "TEXT (UTC)"]]
    ftable(s, 7.8, 1.9, [3.1, 1.9], rows, size=12, row_h=0.46)
    notes[15] = "One table; each row owns one original image and one overlay image on disk."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("10", "REST endpoints", K.FBLUE), (str(py.get("passed", 73)), "pytest tests", K.GREEN), ("8", "React screens", K.ORANGE), ("1", "command to start (run.bat)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Signed HttpOnly session cookie (HMAC-SHA256, 12 h), 401 on protected routes, path-traversal guard on image files, bound-parameter SQL.", {"bullet": True, "space_after": 9}),
        ("One TensorFlow inference at a time under a lock; model loads in a background thread at start-up.", {"bullet": True, "space_after": 9}),
        ("Upload progress bar from a real XMLHttpRequest; side-by-side and blended Grad-CAM views; CSV export; print stylesheet.", {"bullet": True, "space_after": 9}),
        ("Training script reuses the inference preprocessing and records every metric to JSON.", {"bullet": True}),
    ], size=18)
    notes[16] = "Highlights of the implementation: security controls, concurrency, interface features and the reproducible training script."

    # ---- 17 results: model performance
    s = S[16]
    set_title(s, "Results – Model Performance (official test split)")
    wipe(s)
    if T:
        rows = [["Class", "Precision", "Recall", "F1", "Support"]]
        for c in F.CLASSES:
            p = PC[c]
            rows.append([c, F.f3(p["precision"]), F.f3(p["recall"]), F.f3(p["f1"]), str(p["support"])])
        rows.append(["Macro average / accuracy", F.f3(T["macro_precision"]), F.f3(T["macro_recall"]), F.f3(T["macro_f1"]), f"acc {F.f3(T['accuracy'])}"])
        ftable(s, 0.55, 1.9, [3.0, 1.2, 1.2, 1.2, 1.3], rows, size=14, row_h=0.55, align=["l", "c", "c", "c", "c"])
        for i, (v, l, c) in enumerate([(F.pct(B["sensitivity"]), "pneumonia sensitivity", K.GREEN), (F.pct(B["specificity"]), "normal specificity", K.FBLUE), (F.f3(B["auc"]), "binary AUC", K.ORANGE), (F.f3(T["expected_calibration_error"]), "calibration error (ECE)", K.PURPLE)]):
            ftile(s, 0.55 + (i % 2) * 3.05, 5.4 + (i // 2) * 0.0, 2.95, 1.2, v, l, c, 24) if i < 2 else None
        K.bar_chart(s, 8.6, 1.85, 4.2, 4.9, ["Normal", "Bacterial", "Viral"], [("F1", [PC[c]["f1"] for c in F.CLASSES])], colors=[K.FBLUE], number_format="0.00", font_size=12, max_value=1.1)
    notes[17] = ("Protocol: train on 4,688 patient-split training images, validate on 544, test once on the official 624-image test split. Class weights balanced.")

    # ---- 18 confusion matrix & per-class
    s = S[17]
    set_title(s, "Confusion Matrix & Per-class Metrics")
    wipe(s)
    cm_path = os.path.join(F.EXP, F.DEPLOYED_NAME, "confusion_matrix.png")
    if os.path.exists(cm_path):
        K.picture(s, cm_path, 0.55, 1.85, 5.2, 4.7)
    K.picture(s, F.fig("per_class_metrics.png"), 6.0, 1.85, 6.8, 4.7)
    if T:
        cm = T["confusion_matrix"]
        ftext(s, 0.55, 6.5, 12.2, 0.4, f"Test n = {T['n']}: Normal row {cm[0]}, Bacterial row {cm[1]}, Viral row {cm[2]} (columns Normal, Bacterial, Viral).", size=12, color=K.FNAVY, align="c")
    notes[18] = "Rows are the true class, columns the prediction. The largest confusion is normal films predicted as viral pneumonia, then viral predicted as bacterial."

    # ---- 19 threshold / parameter analysis
    s = S[18]
    set_title(s, "Threshold & Parameter Analysis")
    wipe(s)
    K.picture(s, F.fig("confidence_threshold.png"), 0.55, 1.7, 6.1, 2.9)
    K.picture(s, F.fig("decision_threshold_sweep.png"), 6.7, 1.7, 6.1, 2.9)
    if E1 or E2:
        rows = [["Experiment", "3-class acc.", "Macro F1", "Sensitivity", "Specificity"]]
        for lab, e in (("Baseline: project defaults", E1), ("Deployed: tuned recipe", F.DEPLOYED), ("Ablation: no CLAHE", E2), ("MobileNetV3Small (base-paper backbone)", E3)):
            if e:
                t = e["test"]
                rows.append([lab, F.pct(t["accuracy"]), F.f3(t["macro_f1"]), F.pct(t["binary_normal_vs_pneumonia"]["sensitivity"]), F.pct(t["binary_normal_vs_pneumonia"]["specificity"])])
        ftable(s, 0.55, 4.75, [4.6, 1.8, 1.8, 2.0, 2.0], rows, size=12, row_h=0.38, align=["l", "c", "c", "c", "c"])
    notes[19] = "Left: accuracy against the minimum confidence required to auto-report. Right: moving the pneumonia decision threshold trades sensitivity for specificity. Bottom: the CLAHE ablation and the MobileNetV3Small comparison on the same test images."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    if LAT:
        for i, (v, l, c) in enumerate([(f"{LAT['preprocess']:.0f} ms", "preprocessing", K.FBLUE), (f"{LAT['predict']:.0f} ms", "network forward pass", K.GREEN), (f"{LAT['gradcam_overlay']:.0f} ms", "Grad-CAM + overlay", K.ORANGE), (f"{api.get('mean_ms', 0):.0f} ms", "end-to-end POST /api/predict (mean)", K.PURPLE)]):
            ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.15, v, l, c, 26)
    rows = [["Test layer", "What it covers", "Result"],
            ["Unit (pytest)", "preprocessing, network, Grad-CAM, database, tokens", f"{py.get('per_file', {}).get('tests/test_preprocess.py', 0) + py.get('per_file', {}).get('tests/test_model_gradcam.py', 0) + py.get('per_file', {}).get('tests/test_database.py', 0) + py.get('per_file', {}).get('tests/test_auth_tokens.py', 0)} passed"],
            ["Integration (pytest)", "login, validation, prediction, history, stats, files", f"{py.get('per_file', {}).get('tests/test_api.py', 0)} passed"],
            ["Functional (Flask)", "upload form, result, history, errors", f"{py.get('per_file', {}).get('tests/test_flask_app.py', 0)} passed"],
            ["Live smoke test", "running servers through the proxy", f"{sm.get('checks', 0)} checks passed"],
            ["UI verification", "360 / 768 / 1440 px, light / dark, keyboard, motion", f"{ui.get('configurations', 0)} configurations, no overflow or console errors"]]
    ftable(s, 0.55, 3.35, [2.6, 6.0, 3.6], rows, size=13, row_h=0.5)
    notes[20] = "Latency was measured on a CPU-only desktop: about half a second per scan through the API. The test table shows every layer passed."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different datasets)")
    wipe(s)
    acc_b = F.pct(B.get("accuracy", 0)) if B else "n/a"
    rows = [["Aspect", "Base paper [1]", "This project"],
            ["Task", "Binary (Normal / Pneumonia) and 4-class (COVID-19, Normal, Pneumonia, TB)", "3-class: Normal / Bacterial / Viral pneumonia; Normal vs pneumonia view derived"],
            ["Data", "Mendeley paediatric set, balanced subset of 1,300 per class (2,080 / 260 / 260)", f"Same image source, full official data: {DS.get('train', 0):,} / {DS.get('validation', 0)} / {DS.get('test', 0)}, imbalanced, class weights"],
            ["Model", "MobileNet family, ImageNet-pretrained; FixMatch / FlexMatch / FreeMatch", f"Custom 4-block CNN, {mi.get('total_params', 0):,} parameters, trained from scratch on a CPU"],
            ["Accuracy", "Up to 98 % (binary), 97 % (4-class); 97 % with 20 % labels", f"{F.pct(T.get('accuracy', 0))} 3-class; {acc_b} Normal vs pneumonia (official test split)"],
            ["Explanation / reliability", "Grad-CAM, MC Dropout (entropy 0.30), TTA, ECE", f"Grad-CAM, ECE {F.f3(T.get('expected_calibration_error', 0))}, confidence-threshold analysis"],
            ["Deliverable", "Research framework", "Working web application with history, dashboard and tests"]]
    ftable(s, 0.55, 1.85, [2.0, 5.0, 5.2], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 6.35, 12.2, 0.5, "Not a like-for-like comparison: different class set, balance, split and model family. Where the protocol differs the numbers are context, not a ranking.", size=13, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that datasets and protocols differ: the base paper balances and subsamples, we keep the official imbalanced data and add a third class."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Upload and Analysis")
    wipe(s)
    K.picture(s, F.shot("05_detect_preview.png"), 0.55, 1.85, 6.0, 4.95, border=True)
    K.picture(s, F.shot("06_result.png"), 6.8, 1.85, 6.0, 4.95, border=True)
    notes[22] = "Left: the detect screen with a preview before analysis. Right: the result page with the class badge, probabilities and the Grad-CAM overlay."
    s = S[22]
    set_title(s, "Output – Grad-CAM Result (blend and dark theme)")
    wipe(s)
    K.picture(s, F.shot("07_result_blend.png"), 0.55, 1.85, 6.0, 4.95, border=True)
    K.picture(s, F.shot("08_result_dark.png"), 6.8, 1.85, 6.0, 4.95, border=True)
    notes[23] = "The heat-map strength slider blends the overlay over the X-ray; the dark theme is remembered."
    s = S[23]
    set_title(s, "Output – History and Dashboard")
    wipe(s)
    K.picture(s, F.shot("11_history.png"), 0.55, 1.85, 6.0, 4.95, border=True)
    K.picture(s, F.shot("09_dashboard.png"), 6.8, 1.85, 6.0, 4.95, border=True)
    notes[24] = "History with search, filter, sort and CSV export; dashboard with totals, class distribution and scans per day."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("**Delivered:** an end-to-end pneumonia detection system: preprocessing, a trained three-class CNN, Grad-CAM, storage, history, dashboard, tests and documentation.", {"bullet": True, "space_after": 10}),
        (f"**Measured:** {F.pct(T.get('accuracy', 0))} three-class accuracy and {F.pct(B.get('sensitivity', 0))} / {F.pct(B.get('specificity', 0))} pneumonia sensitivity / normal specificity on the official test split, with the protocol stated.", {"bullet": True, "space_after": 10}),
        ("**Finding:** normal versus pneumonia is learnable by a compact CPU-trained CNN; specificity (normal films called pneumonia) is the weak point, and the public test split shares pneumonia ids with training, so patient-disjoint results are reported too.", {"bullet": True, "space_after": 10}),
        ("**Not achieved:** U-Net lung segmentation (objective 2), calibrated uncertainty, external validation. Not a medical device.", {"bullet": True}),
    ], size=18)
    notes[25] = "Five objectives were met in full or in part; objective 2 was not. Findings are stated with their protocol."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("U-Net lung segmentation to restrict the classifier to the lungs and make Grad-CAM maps more specific.", {"bullet": True, "space_after": 12}),
        ("Monte Carlo Dropout or temperature scaling for calibrated uncertainty and a built-in review flag.", {"bullet": True, "space_after": 12}),
        ("Transfer-learning backbones (MobileNetV3Small, EfficientNet) with a Grad-CAM that can reach nested layers.", {"bullet": True, "space_after": 12}),
        ("External validation on data from other hospitals; adult cohorts; prospective clinical evaluation.", {"bullet": True, "space_after": 12}),
        ("Real accounts and roles, audit log, DICOM input, container deployment.", {"bullet": True}),
    ], size=20)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    refs = [1, 3, 5, 6, 7, 9, 10, 11, 12, 13, 14, 15, 16, 20, 22, 29]
    left = refs[:8]
    right = refs[8:]
    ftext(s, 0.55, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1].split(', doi')[0]}", {"space_after": 4}) for n in left], size=9)
    ftext(s, 6.8, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1].split(', doi')[0]}", {"space_after": 4}) for n in right], size=9)
    notes[27] = "Selected references; the complete list is in the report."

    # ---- 28 thank you
    s = S[27]
    for sh in s.shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    from .pptkit import scrub_theme
    scrub_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
