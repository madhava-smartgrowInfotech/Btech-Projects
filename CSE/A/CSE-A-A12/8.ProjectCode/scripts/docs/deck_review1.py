"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .ch2 import LIT_TABLE
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


KEEP = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY)


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.TITLE}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "pneumonia, chest X-ray, CNN, Grad-CAM"
    cp.comments = ""
    cp.last_modified_by = ""


def content_slides(prs):
    return list(prs.slides)


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, "An explainable chest X-ray classifier: Normal, Bacterial or Viral Pneumonia, with a Grad-CAM heat map for every prediction")
        elif nm == "review label":
            K.replace_text(sh, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")
    notes[1] = ("Introduce the team and the project title. One sentence: we build a web tool that reads a chest X-ray, says Normal, Bacterial or Viral "
                "pneumonia, and shows the image regions behind its answer.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. The review is about why the problem matters and what is missing in current work, not yet about results."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: a second opinion that shows its reasoning")
    wipe(s)
    W, H = 3.85, 4.5
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Pneumonia is a leading infectious cause of death, yet chest X-ray readers are scarce exactly where it is most dangerous. Manual reading is slow and varies between readers, and most classifiers return a bare label with no visual justification.", K.ORANGE, "!"),
        ("Our idea", "Enhance the radiograph with CLAHE, classify it with a compact CNN into Normal, Bacterial or Viral pneumonia, and overlay a Grad-CAM heat map so every prediction can be checked against the image.", K.BLUE, "i"),
        ("The deliverable", "A web application: upload an X-ray, receive the class, the three probabilities and the heat map; every scan is stored with a history and a dashboard.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three sentences: the problem, our idea, what we deliver. Emphasise explanation: a heat map next to every prediction."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "MobileNet, MobileNetV2, MobileNetV3Small and V3Large with ImageNet weights; FixMatch, FlexMatch and FreeMatch pseudo-labelling; Grad-CAM for explanation; Monte Carlo Dropout, test-time augmentation and calibration error for reliability.", K.BLUE, "M"),
        ("Data and reported results", "Mendeley paediatric chest X-rays (5,863) for Normal vs Pneumonia and a merged four-class set (7,135). Supervised accuracy up to 98 % (binary) and 97 % (multiclass); FreeMatch reaches 97 % with 20 % (binary) or 30 % (multiclass) labels.", K.GREEN, "R"),
        ("Limits stated by the authors", "Benchmark datasets may not reflect clinical variability; the benefit of adding COVID-19 and TB classes is a hypothesis, not isolated; errors cluster in visually similar classes and high-entropy samples.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = "Read the citation. The paper is a research framework on public benchmarks; it does not deliver an application. We reuse its input size, Grad-CAM idea and caution about orientation-changing augmentation."

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result"]]
    pick = [0, 1, 2, 3, 4, 7, 8]
    for i in pick:
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.4, [2.6, 3.3, 3.0, 3.23], rows, font_size=10.5, row_h=0.62, header_size=12)
    notes[5] = "The pattern: classifiers on public chest X-ray data, hybrid and multi-disease variants, and recent comparisons on the same paediatric images. Figures marked as summarised in [1] are quoted from the base paper's own survey."

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: protocol changes the number")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 7.4, 5.2, ["Kermany 2018\n(official test)", "EfficientNet-B0\nKhadidos 2026", "DenseNet121\nKhadidos 2026", "Best model\nShahi 2025"],
                [("Reported accuracy", [0.928, 0.846, 0.797, 0.98])], colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.1)
    K.card(s, 8.3, 1.5, 4.43, 2.45, "Takeaway", "Reported accuracy on the paediatric dataset ranges from about 80 % to 98 % depending on the split, class balance and protocol.", K.GREEN, "↑", body_size=14)
    K.card(s, 8.3, 4.15, 4.43, 2.45, "Caveat", "Shahi and Bagale report 96-98 % (98 % plotted) after a patient-wise re-split; we plot the top of that range. Numbers are only comparable when the protocol is stated.", K.ORANGE, "!", body_size=14)
    notes[6] = ("Chart values come from the papers: 92.8 % [3], 84.6 % and 79.7 % [5], and the 96-98 % range of [6] (98 % plotted). The point is protocol sensitivity; "
                "our own results in later reviews state the protocol beside every number.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Chest X-ray taken", K.NAVY), ("Clinician reads the film", K.NAVY), ("Opinion / report", K.NAVY), ("Treatment decision", K.NAVY)], size=17, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Manual reading", "A trained reader interprets each radiograph by eye. It is accurate in expert hands but slow when staff are scarce, tiring, and variable between readers.", K.BLUE, "1", body_size=18)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Research classifiers", "Papers train CNNs on public data and report a label or a score. Code is usually a notebook: no interface, no record of past cases, and rarely any visual explanation.", K.ORANGE, "2", body_size=18)
    notes[7] = "Describe today's typical path. Research models stop at a number in a notebook; they are not something a clinician or student can open and try."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("No explanation", "A bare label cannot be checked against the image, so it is hard to trust in a clinical workflow.", K.ORANGE, "1"),
        ("Protocol-dependent numbers", "The same images give anywhere from about 80 % to 98 % accuracy depending on how they are split.", K.BLUE, "2"),
        ("Pathogen type is hard", "Separating bacterial from viral pneumonia is less studied than Normal vs Pneumonia and is where results fall.", K.GREEN, "3"),
        ("Not usable as a tool", "No interface, history, search or export: research code, not an application.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=17)
    notes[8] = "Four limitations our project targets. Each maps to a design decision shown later: heat map, stated protocol, three-class output, web application."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Application", "Few works go from an uploaded image to an explained, stored, reviewable prediction."),
            ("Pathogen type", "Bacterial against viral pneumonia is far less studied than pneumonia against normal."),
            ("Protocol", "Results are rarely stated with their split and balance, so published numbers cannot be compared fairly."),
            ("Explanation + uncertainty", "Grad-CAM, calibration and uncertainty are usually separate post-hoc add-ons, not part of a usable tool."),
            ("Segmentation", "Restricting the classifier to the lung region (U-Net) is discussed but rarely combined with the rest.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=18, color=K.INK, anchor="m")
    notes[9] = "Each gap is backed by the literature on the previous slides. The batch literature survey phrases it as: combine segmentation, calibrated uncertainty and visual explanation into one deployable aid."

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, "Given one frontal chest X-ray, classify it as **Normal**, **Bacterial Pneumonia** or **Viral Pneumonia**, report the probability of each class, show the image regions that supported the decision, and keep a reviewable record, all through a web interface a non-expert can use.", size=20, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "A PNG or JPEG chest X-ray, up to 8 MB"), ("Output", "Class, confidence, three probabilities, Grad-CAM overlay"), ("Record", "History, dashboard, CSV export")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=15, color=K.TXT)
    notes[10] = "Read the statement slowly. Three parts: classify into three classes, explain with a heat map, keep a record."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    from .ch1 import OBJECTIVES

    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 9}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "Six objectives from the project proposal: data and preprocessing, U-Net segmentation, CNN classifier, evaluation, heat maps, deployment."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Preprocess", "Resize to 224 x 224, CLAHE contrast enhancement, light blur, scale to [0, 1]. Planned: U-Net lung segmentation.", K.BLUE, "1"),
        ("Classify", "Compact CNN with four convolution blocks; softmax over Normal, Bacterial and Viral.", K.GREEN, "2"),
        ("Explain", "Grad-CAM on the last convolution layer, overlaid on the image.", K.ORANGE, "3"),
        ("Evaluate", "Accuracy, sensitivity, specificity, F1, confusion matrix, calibration, on a held-out test split.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "W", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Web application", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "FastAPI service over the existing Flask project's model code, React interface, demo login, SQLite history, dashboard and CSV export.", size=15, color=K.TXT)
    notes[12] = "Four pipeline stages and the application around them. U-Net segmentation is part of the proposal and is marked as planned."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 2.2, 1.35, 8.9, 5.3)
    K.text(s, 0.6, 6.62, 12.13, 0.3, "Four tiers: presentation, application (FastAPI), inference (preprocess, CNN, Grad-CAM) and data (SQLite, images, weights)", size=12, color=K.TXT, align="c", italic=True)
    notes[13] = "Follow the diagram left to right: the user interface, the API, the three inference modules, and the data stores."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working classifier", "A trained three-class CNN evaluated on the public test split.", K.BLUE, "1"),
            ("Visible reasoning", "A Grad-CAM overlay and three probabilities for every scan.", K.GREEN, "2"),
            ("Honest evaluation", "Metrics stated with their protocol, including where the model fails.", K.ORANGE, "3"),
            ("Usable system", "Web application with login, history, dashboard and export.", K.PURPLE, "4"),
            ("Reproducibility", "One script trains and evaluates; every figure is regenerated from JSON.", K.BLUE, "5"),
            ("Quality", "Automated tests for the API, the model code and the interface.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=16)
    notes[14] = "These are expected outcomes, not results. We will report whatever the held-out test split shows, including weaknesses."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    refs = [(1, REFS[0].split(", doi")[0] + "."), (3, REFS[2]), (5, REFS[4]), (6, REFS[5]), (7, REFS[6]), (9, REFS[8]), (10, REFS[9]), (11, REFS[10])]
    K.text(s, 0.6, 1.35, 12.13, 5.5, [(f"[{n_}] {t}", {"space_after": 5}) for n_, t in refs], size=13, color=K.INK)
    notes[15] = "Full references are in the project report."

    # ---- 16 thank you
    s = S[15]
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")
    notes[16] = "Likely questions: why a compact CNN (trains on a CPU, explainable); why Grad-CAM; why three classes; what about segmentation (planned, see later reviews)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    from .pptkit import scrub_theme
    scrub_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
