"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress. Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .deck_review1 import agenda, is_ph, prepare, set_title, wipe

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"


def metrics():
    return F.DEPLOYED or {}


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}
    M = metrics()
    T = M.get("test", {})
    TR = M.get("training", {})
    DS = M.get("dataset", {})
    py = (F.TESTS or {}).get("pytest", {})
    sm = (F.TESTS or {}).get("smoke", {})

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, "Design, methodology and progress: a trained three-class CNN with Grad-CAM, behind a FastAPI service and a React interface")
        elif nm == "review label":
            K.replace_text(sh, "REVIEW 2  ·  DESIGN, METHODOLOGY AND PROGRESS")
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")
    notes[1] = "This review shows how the system is designed and built: the data and model, what has been implemented, the first measured results and the open items."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap and architecture"), ("02", "Methodology and dataset"), ("03", "Preprocessing and model"), ("04", "Training and decision logic"),
               ("05", "Data flow, database and UML"), ("06", "Technology stack"), ("07", "Implementation progress and results"),
               ("08", "Output screens and testing"), ("09", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Pneumonia needs a fast, checkable reading of the chest X-ray. Readers are scarce, readings vary, and most classifiers give a label with no visual justification.", K.ORANGE, "!"),
        ("Proposed system", "Preprocess with CLAHE, classify into Normal, Bacterial or Viral pneumonia with a compact CNN, explain with Grad-CAM, and keep every scan in a reviewable history.", K.BLUE, "i"),
        ("Base paper link", "Reuses 224 x 224 input and Grad-CAM; replaces its semi-supervised research framework by a complete application. Segmentation and Monte Carlo Dropout are not reproduced.", K.GREEN, "B"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, 4.6, t, b_, col, gl, body_size=17)
    notes[3] = "One-minute recap of Review 1: the base paper is a research framework; we deliver the end-to-end tool and measure it honestly."

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.6, 5.45)
    K.text(s, 8.5, 1.5, 4.2, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        ("Presentation: React SPA (port 5312) and the original Flask UI.", {"bullet": True}),
        ("Application: FastAPI service (port 8312) with signed-cookie login and validation.", {"bullet": True}),
        ("Inference: preprocess, CNN, Grad-CAM modules shared with the Flask app.", {"bullet": True}),
        ("Data: SQLite history, image files, saved model weights.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[4] = "Four tiers. The FastAPI layer imports the original model, Grad-CAM and database modules, so nothing was rewritten."

    # ---- 5 methodology workflow
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.35, [("Public chest X-ray data", K.NAVY), ("Label from file names", K.NAVY), ("Preprocess + cache", K.PURPLE), ("Patient-level split", K.GREEN),
                                      ("Train CNN", K.ORANGE), ("Evaluate on test split", K.BLUE), ("metrics.json + figures", K.NAVY)], size=14, gap=0.22)
    for i, (t, b_) in enumerate([("Shared preprocessing", "`load_and_preprocess_image` is used by training, Flask and FastAPI: no train/serve mismatch."),
                                 ("Patient-level split", "Validation is cut by patient so one patient never sits on both sides; the test split is used once."),
                                 ("Reproducible", "Seed 42; every metric, prediction and figure is written to disk and regenerated by scripts.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.3, 3.85, 3.0)
        K.text(s, x + 0.25, 3.45, 3.4, 0.5, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 4.05, 3.4, 2.1, b_.replace("`", ""), size=16, color=K.TXT)
    notes[5] = "The training script is one command and writes every number to JSON. Preprocessing is the same function the web app uses."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Dataset")
    wipe(s)
    for i, (v, l, c) in enumerate([("5,856", "chest X-ray images, paediatric", K.BLUE), ("3", "classes: Normal, Bacterial, Viral", K.GREEN), ("624", "official test images, held out", K.ORANGE), ("CC BY 4.0", "Kermany et al., Mendeley Data v3", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    K.picture(s, F.fig("dataset_distribution.png"), 0.6, 2.95, 6.6, 3.85)
    K.text(s, 7.5, 3.0, 5.2, 3.8, [
        ("Pathogen labels come from file names (person12_bacteria_34, person7_virus_21).", {"bullet": True}),
        ("Official val folder has only 16 images, so train and val are merged and a 10 % validation set is cut by patient.", {"bullet": True}),
        ("Many pneumonia ids in the test folder also occur in training; results are also reported on a patient-disjoint view.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[6] = "Kermany et al. data, CC BY 4.0. Pathogen labels are derived from the file names. The overlap of ids between the official folders is analysed in the results."

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Data preprocessing")
    wipe(s)
    K.picture(s, F.fig("preprocessing_steps.png"), 0.6, 1.4, 7.9, 4.3)
    K.text(s, 8.8, 1.5, 3.95, 5.2, [
        ("1  Decode to grayscale", {}), ("2  Resize to 224 x 224", {}), ("3  CLAHE, clip 2.5, 8 x 8 tiles", {}), ("4  Gaussian blur 3 x 3", {}), ("5  Three channels, divide by 255", {}),
        ("Training adds rotation, shift, shear, zoom and horizontal flip.", {"italic": True, "size": 14, "color": K.TXT}),
    ], size=17, color=K.INK, space_after=10)
    notes[7] = "The same five steps run at training and at inference. Augmentation is training-only."

    # ---- 8 model architecture
    s = S[7]
    set_title(s, "Model architecture")
    wipe(s)
    rows = [["Block", "Layers", "Output"], ["1", "Conv 3x3, 32 · BN · MaxPool", "112 x 112 x 32"], ["2", "Conv 3x3, 64 · BN · MaxPool", "56 x 56 x 64"], ["3", "Conv 3x3, 128 · BN · MaxPool", "28 x 28 x 128"],
            ["4", "Conv 3x3, 256 (last_conv) · BN · MaxPool", "14 x 14 x 256"], ["Head", "GAP · Dense 512 · BN · Dropout 0.35 · Dense 128 · Dropout 0.25", "256 -> 512 -> 128"], ["Output", "Dense 3, softmax", "3 class probabilities"]]
    K.table(s, 0.6, 1.5, [1.1, 5.0, 2.3], rows, font_size=14, row_h=0.62, align=["c", "l", "l"])
    mi = F.MODEL_INFO or {}
    K.tile(s, 9.4, 1.5, 3.33, 1.45, f"{mi.get('total_params', 0):,}", "parameters", K.BLUE, vsize=28)
    K.tile(s, 9.4, 3.15, 3.33, 1.45, f"{mi.get('float32_mb', 0)} MB", "weights in float32", K.GREEN, vsize=28)
    K.tile(s, 9.4, 4.8, 3.33, 1.45, "No GPU", "trains on a desktop CPU", K.ORANGE, vsize=26)
    notes[8] = "A compact CNN written for the project. The convolution layer named last_conv is where Grad-CAM reads its feature maps."

    # ---- 9 training configuration
    s = S[8]
    set_title(s, "Training configuration")
    wipe(s)
    hp = M.get("hyperparameters", {})
    rows = [["Setting", "Value"],
            ["Optimiser / loss", f"Adam, learning rate {hp.get('learning_rate', 1e-4)} / categorical cross-entropy"],
            ["Class weights", "balanced: " + ", ".join(f"{v:.2f}" for v in (M.get('class_weights') or {}).values())],
            ["Batch / max epochs", f"{hp.get('batch_size', 32)} / {hp.get('epochs_max', 25)}"],
            ["Early stopping", f"validation loss, patience {hp.get('early_stopping_patience', 10)}, best weights restored"],
            ["Learning-rate schedule", f"halved after {hp.get('reduce_lr_patience', 3)} epochs without improvement"],
            ["Augmentation", "rotation 12, shift 0.08, shear 0.08, zoom 0.12, horizontal flip"],
            ["Seed / hardware", f"{hp.get('seed', 42)} / Intel Core i5-12500, CPU only, about {TR.get('seconds_per_epoch', 240):.0f} s per epoch"]]
    K.table(s, 0.6, 1.5, [3.2, 8.93], rows, font_size=16, row_h=0.64, first_bold=True)
    notes[9] = "The base configuration of the project's custom CNN, with class weights for the Bacterial-heavy balance."

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.6, 1.3, 5.2, 5.5)
    K.text(s, 6.2, 1.5, 6.5, 5.3, [
        ("**Validation first.** Valid session, named file, extension jpg / jpeg / png, size up to 8 MB, decodable image; each failure returns a clear 401, 400 or 413 and stores nothing.", {"bullet": True, "space_after": 10}),
        ("**Decision.** The class is the arg-max of the three softmax probabilities; the confidence is that probability.", {"bullet": True, "space_after": 10}),
        ("**Explanation.** Grad-CAM for the predicted class, blended over the display image at alpha 0.4.", {"bullet": True, "space_after": 10}),
        ("**Caution.** Confidence is a raw softmax value; Results report accuracy against a confidence threshold so that low-confidence scans can be flagged for review.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[10] = "The route validates in the same order as the original Flask route. The confidence threshold is analysed on the results slides, not applied automatically."

    # ---- 11 data flow and modules
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 7.3, 5.5)
    rows = [["Module", "Role"], ["model/preprocess.py", "decode, resize, CLAHE, blur"], ["model/model_builder.py", "CNN, last-conv finder"], ["model/gradcam.py", "heat map and overlay"], ["database/db.py", "SQLite access"], ["backend/main.py", "FastAPI routes, login"], ["frontend/src", "React screens"]]
    K.table(s, 8.2, 1.5, [2.45, 2.08], rows, font_size=12, row_h=0.5)
    notes[11] = "Level-1 data flow: five processes. Level 0 and 2 are in the report."

    # ---- 12 database and UML
    s = S[11]
    set_title(s, "Database design and UML")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.6, 1.35, 6.0, 5.1)
    K.picture(s, F.diag("06_uml_class_diagram.png"), 6.8, 1.35, 5.9, 5.1)
    K.text(s, 0.6, 6.55, 12.13, 0.35, "One table (predictions, 10 columns) with bound-parameter SQL; the code base is function-oriented, so the UML draws modules as classes.", size=12, color=K.TXT, align="c", italic=True)
    notes[12] = "A single predictions table; each row owns one original image and one overlay. The UML shows modules as classes because the code is function-oriented."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    rows = [["Layer", "Technology"], ["Deep learning", "TensorFlow / Keras, NumPy, scikit-learn, matplotlib"], ["Image processing", "OpenCV (CLAHE, resize, colour maps), Pillow"],
            ["API", "FastAPI, Uvicorn, python-multipart; Flask (original UI)"], ["Database", "SQLite"], ["Front end", "React 19, Vite, React Router, Framer Motion, Recharts, Neo Brutal Grid kit"],
            ["Testing", "pytest, FastAPI TestClient, Playwright"], ["Run", "run.bat (Windows), ports 8312 / 5312"]]
    K.table(s, 0.6, 1.5, [3.0, 9.13], rows, font_size=16, row_h=0.62, first_bold=True)
    notes[13] = "Everything is open source; one script starts the stack."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Preprocessing, CNN, Grad-CAM", "Done", "model/, 19 unit tests"],
            ["Training and evaluation script", "Done", "scripts/train_experiment.py, metrics.json"],
            ["FastAPI service and demo login", "Done", "10 routes, 34 integration tests"],
            ["React interface (8 screens)", "Done", "13 screenshots, UI checks at 360-1440 px"],
            ["History, dashboard, CSV export", "Done", "smoke test, screens 09 to 11"],
            ["U-Net lung segmentation", "Not implemented", "listed as future work"],
            ["Monte Carlo Dropout, semi-supervised learning", "Not implemented", "out of scope; confidence-threshold analysis instead"]]
    K.table(s, 0.6, 1.5, [4.4, 1.9, 5.83], rows, font_size=14, row_h=0.62)
    notes[14] = "State of the code base. The two honest gaps are segmentation and the base paper's uncertainty and semi-supervised methods."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: official test split")
    wipe(s)
    if T:
        b = T["binary_normal_vs_pneumonia"]
        for i, (v, l, c) in enumerate([(F.pct(T["accuracy"]), "3-class accuracy", K.BLUE), (F.f3(T["macro_f1"]), "macro F1", K.GREEN), (F.pct(b["sensitivity"]), "pneumonia sensitivity", K.ORANGE), (F.pct(b["specificity"]), "normal specificity", K.PURPLE)]):
            K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28)
        pc = T["per_class"]
        K.bar_chart(s, 0.6, 2.9, 7.4, 3.95, ["Normal", "Bacterial", "Viral"], [("Precision", [pc[c]["precision"] for c in F.CLASSES]), ("Recall", [pc[c]["recall"] for c in F.CLASSES]), ("F1", [pc[c]["f1"] for c in F.CLASSES])],
                    colors=[K.BLUE, K.GREEN, K.ORANGE], number_format="0.00", font_size=12, max_value=1.1)
        K.card(s, 8.3, 2.95, 4.43, 3.75, "Reading", f"n = {T['n']} test images. Protocol: official Kermany test split, model trained on the patient-split train pool. Details and caveats in the next review.", K.BLUE, "i", body_size=15)
    notes[15] = "First measured numbers on the official test split. The protocol is stated beside them; the final review adds the patient-disjoint analysis and the comparison with the base paper."

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    K.picture(s, F.shot("06_result.png"), 0.6, 1.3, 4.8, 5.55, border=True)
    K.picture(s, F.shot("09_dashboard.png"), 5.6, 1.3, 3.6, 5.55, border=True)
    K.picture(s, F.shot("11_history.png"), 9.4, 1.3, 3.33, 5.55, border=True)
    notes[16] = "Real screenshots from the running application with the trained model: a result with the Grad-CAM overlay, the dashboard and the history."

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([(str(py.get("passed", 0)), "pytest tests passed", K.GREEN), (str(sm.get("checks", 0)), "live smoke-test checks", K.BLUE), ("0", "console errors on the screenshot walk", K.ORANGE), ("360-1440 px", "light and dark, no overflow", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    rows = [["Observation", "Action"],
            ["Fresh checkout had no trained model; predictions were near-equal", "Trained the CNN on the public data; Demo Mode banner kept for the untrained case"],
            ["Official test split shares pneumonia ids with training", "Added a patient-disjoint analysis and a patient-level re-split"],
            ["Transfer-learning backbones break the shipped Grad-CAM", "Deployed the flat custom CNN; documented the limitation"]]
    K.table(s, 0.6, 3.1, [5.6, 6.53], rows, font_size=14, row_h=0.8)
    notes[17] = "The suite runs offline with the real model and a throw-away database. Observations on the right drove three of the design decisions."

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["No trained weights or dataset in the repository", "Downloaded the public Kermany data; wrote a reproducible training script"],
            ["Original train.py trains without CLAHE but the app uses it", "Experiment script reuses the inference preprocessing"],
            ["Official validation folder has 16 images", "Merged train and val; patient-level 10 % validation"],
            ["Bacterial images about twice as frequent", "Balanced class weights in the loss"],
            ["CPU-only training, about 4 minutes per epoch", "Compact model; few, well-chosen experiments"],
            ["TensorFlow failed to load after other scientific imports", "Import TensorFlow first in the scripts"]]
    K.table(s, 0.6, 1.5, [5.6, 6.53], rows, font_size=14, row_h=0.68)
    notes[18] = "Each challenge was met by a concrete change in the code or the evaluation protocol."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Evaluate", "Finish the patient-disjoint evaluation, the CLAHE ablation and a MobileNetV3Small comparison.", K.BLUE, "1"),
        ("Compare", "Set the results against the base paper, stating where datasets and protocols differ.", K.GREEN, "2"),
        ("Document", "Final report, paper and deck with every figure regenerated from JSON.", K.ORANGE, "3"),
        ("Extend (future)", "U-Net segmentation, calibration or Monte Carlo Dropout, external hospital data.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.1, 1.6, 2.9, 4.4, t, b_, col, gl, body_size=16)
    notes[19] = "Plan to the final review. The extensions at the right are out of scope for this submission."

    # ---- 20 thank you
    s = S[19]
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")
    notes[20] = "Likely questions: why the custom CNN; why patient-level splitting; what about segmentation."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    from .pptkit import scrub_theme
    scrub_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx")))
