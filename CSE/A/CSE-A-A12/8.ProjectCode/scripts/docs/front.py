"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=22, bold=True, space_after=8)
    d.p("An explainable three-class chest X-ray classifier with a web application", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p("Medbowli, Meerpet, Balapur, Hyderabad, Telangana - 500 097", align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the creators of the public chest X-ray dataset released under CC BY 4.0, and the developers of the open-source software "
        "used in this project (TensorFlow, OpenCV, FastAPI, Flask, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    T = (F.DEPLOYED or {}).get("test", {})
    B = T.get("binary_normal_vs_pneumonia", {})
    mi = F.MODEL_INFO or {}
    py = ((F.TESTS or {}).get("pytest") or {}).get("passed", 0)
    d.p("Pneumonia is a leading infectious cause of death, and chest X-ray readers are scarce where the disease is most dangerous. Manual reading is slow and varies between readers, "
        "and most published classifiers return a label with no visual justification. This project builds an end-to-end, explainable pneumonia detection system. An uploaded chest X-ray "
        "is validated, enhanced with contrast-limited adaptive histogram equalisation, classified by a compact convolutional neural network into Normal, Bacterial Pneumonia or Viral Pneumonia "
        "with a probability for each class, and explained with a Grad-CAM heat map; every scan is stored in a history that a dashboard summarises.")
    if T:
        d.p(f"The network ({mi.get('total_params', 0):,} parameters, trained on a desktop CPU) was trained on the public Kermany paediatric chest X-ray data with a patient-level validation split. "
            f"On the official 624-image test split it reaches {F.pct(T['accuracy'])} three-class accuracy and a macro F1 of {F.f3(T['macro_f1'])}; viewed as pneumonia against normal it has "
            f"{F.pct(B['sensitivity'])} sensitivity and {F.pct(B['specificity'])} specificity. Normal radiographs are the weak point: many are called pneumonia, mostly viral. Because many pneumonia ids in the official test "
            "folder also occur in the training folders, results on the patient-disjoint part of the test split and on a fully patient-level re-partition are reported as well, and every figure is "
            "stated together with its protocol. Datasets and protocols differ from those of the base paper, so the comparison is given as context, not as a ranking.")
    d.p("The model is served by a FastAPI layer over the original Flask project's code, behind a React interface with a demo login, history search, CSV export and a printable report. "
        f"{py} automated tests, a live smoke test and browser checks at 360, 768 and 1440 pixels pass. U-Net lung segmentation, calibrated uncertainty and external validation were not implemented "
        "and are listed as limitations and future work. The system is a research prototype and not a medical device.")
    d.p("**Keywords:** pneumonia detection, chest X-ray, deep learning, convolutional neural network, CLAHE, Grad-CAM, explainable AI, medical image analysis, FastAPI, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the ROC Curve"), ("BN", "Batch Normalisation"),
    ("CAM", "Class Activation Mapping"), ("CLAHE", "Contrast Limited Adaptive Histogram Equalisation"), ("CNN", "Convolutional Neural Network"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("CXR", "Chest X-Ray"), ("DFD", "Data Flow Diagram"), ("ECE", "Expected Calibration Error"),
    ("ER", "Entity-Relationship"), ("FN / FP", "False Negative / False Positive"), ("GAP", "Global Average Pooling"), ("GPU", "Graphics Processing Unit"),
    ("Grad-CAM", "Gradient-weighted Class Activation Mapping"), ("HMAC", "Hash-based Message Authentication Code"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JPEG / PNG", "Image file formats"), ("JSON", "JavaScript Object Notation"), ("ReLU", "Rectified Linear Unit"), ("REST", "Representational State Transfer"),
    ("ROC", "Receiver Operating Characteristic"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("TN / TP", "True Negative / True Positive"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("WHO", "World Health Organization"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
