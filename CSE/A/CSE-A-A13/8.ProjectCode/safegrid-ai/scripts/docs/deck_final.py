"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .content_shared import BASE_COMPARE, GAPS, LIT_TABLE, OBJECTIVES, PROBLEM_POINTS, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520
FOOTER = f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review"


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


def caption(slide, x, y, w, text, size=12):
    ftext(slide, x, y, w, 0.3, text, size=size, color=K.FNAVY, align="c", italic=True)


def m(arm, key):
    return F.mean(arm, key)


def nice(sc):
    return sc.replace("_", " ")


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.comments = ""
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.keywords = "agentic AI, multi-agent, smart grid, safety shield, explainability"
    cp.last_modified_by = ""
    notes = {}

    # footer text lives on every slide (and possibly in layouts / master)
    for sl in S:
        for sh in sl.shapes:
            if sh.has_text_frame and "Final Review" in sh.text_frame.text and sh.top == TOP_FOOT:
                K.replace_text(sh, FOOTER)
    for lay in list(prs.slide_layouts) + [prs.slide_master]:
        for sh in lay.shapes:
            if sh.has_text_frame and "Batch D1" in sh.text_frame.text:
                K.replace_text(sh, FOOTER)

    OV = F.OVERALL
    NOSHIELD = (m('agents_unshielded', 'avg_step_time_ms') + m('baseline_rule', 'avg_step_time_ms')) / 2
    interventions = F.PROJECTED + F.REJECTED

    # ---- 1 title
    for sh in S[0].shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE)
        elif sh.name == "Text 2":
            K.replace_text(sh, "Final Review")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_.title()}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE.title()}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE.upper()}\nAcademic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: a multi-agent power-routing system in which AI agents propose actions and a deterministic safety shield decides what may run."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap & Problem Statement\nObjectives\nProposed System & Architecture\nMethodology, Scenarios & Agents\nDecision Workflow, Interaction & Database\nImplementation Highlights\nResults, Shield Decisions, Stress & Cost of Safety\nPerformance, Testing & Comparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References")
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("A language model that controls physical infrastructure is non-deterministic and can propose actions that break electrical limits; operators also need to see why each action was taken.", {"bullet": True, "space_after": 8}),
        ("SafeGrid-AI is neuro-symbolic: three agents (prosumer, EV, renewable) on a simulated 3-node low-voltage feeder each propose one action per step, and a deterministic **Symbolic Safety Shield** approves, scales down (bisection) or rejects it before anything is executed. Every decision is explained and every intervention is written to a hash-chained log.", {"bullet": True, "space_after": 8}),
        (f"In 63 simulation runs ({F.PROPOSALS:,} proposals) the shield approved **{F.pct(F.APPROVED / F.PROPOSALS)}**, projected **{F.pct(F.PROJECTED / F.PROPOSALS)}** and rejected **{F.pct(F.REJECTED / F.PROPOSALS)}**; executed actions caused **0 violations in all 504 shielded steps**, against {F.UNSHIELDED_CAUSED} violations in {F.UNSHIELDED_CAUSED_STEPS} steps for the same agents without it.", {"bullet": True, "space_after": 8}),
        (f"Safety costs time: mean step time {m('safegrid_shielded', 'avg_step_time_ms'):.0f} ms against about {NOSHIELD:.0f} ms; renewable use {m('safegrid_shielded', 'renewable_utilization'):.2f} against {m('agents_unshielded', 'renewable_utilization'):.2f}.", {"bullet": True, "space_after": 8}),
        ("Limits: the language model is a deterministic mock, the grid is synthetic, and the guarantee holds in the simulated model only.", {"bullet": True}),
    ], size=18)
    notes[3] = "Read the problem, the idea, then the headline numbers (63 runs, 0 violations caused by executed actions). End with the limits: mock LLM, synthetic feeder, model-relative guarantee."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.85, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper"],
            ["System", "DC microgrid with PV, wind, synchronous generator, battery storage, grid, critical and non-critical loads; MATLAB/Simulink, 800 V bus, 20 microsecond sampling"],
            ["Method", "Four rule-based agents (DER, energy storage, load, grid) coordinate generation, storage and load shedding"],
            ["Reported result", "Bus voltage regulated at 800 V and critical loads continuously supplied in all eight operating scenarios"],
            ["Stated limits", "Scalability, software integration, cybersecurity; no experimental validation"],
            ["What we take", "The multi-agent decomposition (generation, storage, load); we add an LLM-driven proposer, an independent safety shield and explanations"]]
    ftable(s, 0.55, 2.85, [2.0, 10.2], rows, size=14, row_h=0.62)
    notes[4] = "The base paper uses rule-based agents in a DC microgrid and reports regulation in eight scenarios. It has no language model, no formal safety layer and no explanation interface. We keep its agent decomposition and add those three things."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Approach", "Setting", "Reported result / limit"]]
    short = {
        0: ("Eyimaya & Altin 2026 [1] (base)", "Rule-based agents: DER, storage, load, grid", "DC microgrid, Simulink, 800 V", "Voltage held in 8 scenarios; scalability, cybersecurity, validation open"),
        1: ("Zhang et al. 2025 [2]", "LLM planner + validation agent with sandbox and rollback", "IEEE 69-bus, CIGRE MV, IEEE 30-bus", "Best LLMs 100 % success; weaker 63 % and 44 %; no rule layer, no operator explanations"),
        2: ("Picault & Goubet 2026 [3]", "Position paper: MCP server exposes a simulator to AI agents", "TSO grid studies (pypowsybl)", "No results; asks for traceability and human validation"),
        7: ("McArthur et al. 2007 [8]", "Concepts and challenges of multi-agent systems in power engineering", "Survey, Part I", "Agents fit distributed grids; verification and trust remain open"),
    }
    for i in (0, 1, 2, 7):
        rows.append(list(short[i]))
    rows.append(["[4]-[7] Four further surveyed works", "Multi-agent deep RL, fault reconstruction, EV scheduling, MAS review (classified from titles)", "Smart-grid energy management", "Full text not reviewed for this report"])
    ftable(s, 0.55, 1.85, [2.7, 3.7, 2.6, 3.2], rows, size=12, row_h=0.8)
    notes[5] = "Three papers were read in full (base paper, Grid-Agent, the MCP position paper). The four survey entries in the batch record were classified from title and venue only; the slide says so."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, [
        ("**Planning layer:** the base paper's agents are rule-based; no language-model proposer is used.", {"bullet": True, "space_after": 14}),
        ("**Safety path:** Grid-Agent checks actions by sandboxed trial and rollback; an independent rule layer that scales an unsafe request into the permitted region is not part of the surveyed work.", {"bullet": True, "space_after": 11}),
        ("**Explanation:** operator-facing, per-decision explanations are not reported; the MCP position paper asks for traceability but gives no results.", {"bullet": True, "space_after": 11}),
        ("**Audit:** tamper-evident records of every blocked or scaled action are not addressed.", {"bullet": True, "space_after": 11}),
        ("**Evaluation:** few comparisons separate violations caused by an action from stress already on the grid.", {"bullet": True}),
    ], size=21)
    notes[6] = "Gaps are stated against the surveyed work only, not against all literature. Our contribution is the combination: agent proposals, an independent shield, explanations, an audit trail and a strategy comparison with identical seeds."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.9, 12.2, 1.95)
    ftext(s, 0.85, 1.95, 11.6, 1.85, PROBLEM_STATEMENT, size=17, anchor="m")
    ftext(s, 0.55, 4.05, 6.6, 2.8, [(p, {"bullet": True, "space_after": 7}) for p in PROBLEM_POINTS], size=15)
    for i, (t, b_) in enumerate([("Input", "Feeder state and one proposed action per agent per 15-minute step"), ("Output", "Validated action plus explanation and safety result"), ("Record", "Hash-chained safety log, dashboard, run history")]):
        y = 4.05 + i * 0.95
        fbox(s, 7.5, y, 5.25, 0.85, fill=K.FZEBRA)
        ftext(s, 7.7, y + 0.04, 1.3, 0.77, t, size=15, color=K.FNAVY, bold=True, anchor="m")
        ftext(s, 8.9, y + 0.04, 3.75, 0.77, b_, size=12, anchor="m")
    notes[7] = "The problem is four points from the batch proposal; the statement turns them into an engineering requirement: execute only what cannot create or worsen a violation, and explain every decision."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 16}) for i, o in enumerate(OBJECTIVES, 1)], size=21)
    notes[8] = "The six objectives from the proposal. The conclusion reports each outcome; objective 2 ran with a mock model and objective 6 is only partly measured."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.1, [("Observe feeder", K.FNAVY), ("Agents propose", K.FBLUE), ("Coordinate", K.PURPLE), ("Safety Shield", K.RGBColor(0xC0, 0x39, 0x2B)), ("Execute + log", K.GREEN)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.4, 12.2, 3.4, [
        ("Three agents (prosumer, EV, renewable) get a structured context and return one schema-validated EnergyAction through a pluggable provider (mock, API, rule baseline, failing).", {"bullet": True, "space_after": 9}),
        ("A coordinator detects peer mismatch, target contention and shared-feeder conflicts and orders relief actions first.", {"bullet": True, "space_after": 9}),
        ("The shield checks ten ordered conditions on a trial AC power flow; only approved or scaled actions are committed.", {"bullet": True, "space_after": 9}),
        ("Each decision has an explanation, each intervention a hash-chained event; a React dashboard shows it live.", {"bullet": True}),
    ], size=19)
    notes[9] = "The proposed system in five boxes. The agent is untrusted; the coordinator, shield and simulator are trusted deterministic code."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.75, 8.0, 5.25)
    ftext(s, 8.85, 2.0, 3.95, 4.7, [
        ("Presentation: React dashboard (Vite, port 5313)", {"bullet": True, "space_after": 9}),
        ("Application: FastAPI REST and WebSocket (port 8313), SimulationService", {"bullet": True, "space_after": 9}),
        ("Core: engine, agents, coordinator, shield, pandapower simulator", {"bullet": True, "space_after": 9}),
        ("Data: SQLite store, hash-chained safety log, seeded profiles", {"bullet": True, "space_after": 9}),
        ("Trust boundary: only the LLM provider is untrusted", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. The shield imports nothing from the llm, agents, simulation or grid packages; a test enforces that."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Seeded inputs", K.FNAVY), ("No-action power flow", K.FNAVY), ("Agents propose", K.FBLUE), ("Coordinate", K.PURPLE), ("Shield validates", K.RGBColor(0xC0, 0x39, 0x2B)), ("Commit power flow", K.GREEN), ("Log + explain", K.FNAVY)], size=12, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.45, 12.2, 3.4, [
        ("Experiment: 7 scenarios x 3 strategies x 3 seeds x 24 steps = 63 runs (15-minute step). Strategies: rule baseline, agents without shield, SafeGrid-AI.", {"bullet": True, "space_after": 9}),
        ("A violation is blamed on an action only if it is new or worse than the grid state with every agent idle; pre-existing stress is counted separately.", {"bullet": True, "space_after": 9}),
        ("Projection: bisection for the largest scale factor in [0, 1], then up to four re-verifications of the exact rounded action; below 0.05 kW is rejected.", {"bullet": True, "space_after": 9}),
        ("Everything is reproducible: results come from run_experiment and are stored as JSON.", {"bullet": True}),
    ], size=18)
    notes[11] = "The method in one line of boxes and the three design decisions that matter: the three strategies, the attribution rule and the projection."

    # ---- 12 scenarios, inputs and agents
    s = S[11]
    set_title(s, "Scenarios, Inputs & Agents")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, "No dataset", "inputs are synthetic and seeded; no trained model", K.FBLUE, 22, 11)
    ftile(s, 0.55, 3.15, 2.9, 1.1, "7", "seeded scenarios, 12 to 32 steps", K.GREEN, 26)
    ftile(s, 0.55, 4.4, 2.9, 1.1, "3", "agents on 3 nodes, 0.4 kV feeder", K.ORANGE, 26)
    K.picture(s, F.diag("11_feeder_topology.png"), 3.7, 1.85, 9.1, 3.7)
    ftext(s, 0.55, 5.7, 12.25, 1.2, [
        ("Scenarios: normal, high_solar, peak_demand, battery_stress, line_congestion, voltage_violation, unsafe_ai_action (injects oversized and malformed proposals).", {"bullet": True, "space_after": 5}),
        ("Profiles: solar sin^1.2 with a seeded cloud factor, household load shape with 5 % noise, illustrative time-of-use tariff. The mock LLM is deterministic; a real-LLM provider exists but was not run.", {"bullet": True}),
    ], size=13)
    notes[12] = "There is no dataset and no trained model, so no accuracy, F1 or confusion matrix exists for this project. Inputs are generated from a seed. The three agents are a 6 kWp prosumer with a 10 kWh battery, a 40 kWh EV charging at up to 7.4 kW, and a 20 kWp renewable prosumer with a 30 kWh battery."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.75, 5.0, 5.2)
    ftext(s, 5.9, 1.95, 6.9, 4.9, [
        ("Structural checks first: unknown agent or node, limits, reason codes, magnitude; failure means REJECT.", {"bullet": True, "space_after": 11}),
        ("Idle or near-zero actions are approved and still recorded.", {"bullet": True, "space_after": 11}),
        ("Otherwise a trial AC power flow of the request; only new or worse violations count against the action.", {"bullet": True, "space_after": 11}),
        ("Violation found: bisection for the largest safe scale (40 local, 14 power-flow iterations), re-verify, then PROJECT or REJECT.", {"bullet": True, "space_after": 11}),
        ("Every non-approved action becomes a hash-chained safety event.", {"bullet": True}),
    ], size=16)
    notes[13] = "This is the core algorithm, evaluate() in the shield. Ten ordered checks: structure, power limits, battery, node limits, power balance, power flow, voltage, line loading, transformer, grid import/export."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_diagram.png"), 0.45, 1.75, 6.5, 5.2)
    K.picture(s, F.diag("08_use_case_diagram.png"), 7.1, 1.75, 5.8, 4.7)
    notes[14] = "Left: one simulation step from the dashboard through the API, engine, agents, coordinator, shield and simulator. Right: what the operator can do."

    # ---- 15 database design
    s = S[14]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.2, 4.5)
    rows = [["Table", "Holds"], ["runs", "configuration, resolved settings, summary, chain fingerprint"], ["events", "live event stream per run (typed payload)"], ["safety_events", "hash-chained interventions"], ["steps", "one full record per simulated step"], ["experiments", "specification and results of a comparison"]]
    ftable(s, 8.0, 1.9, [1.7, 3.1], rows, size=12, row_h=0.72)
    notes[15] = "Five SQLite tables from the standard library. Payloads are JSON text. The safety log carries prev_hash and hash so any change is detectable."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("23 + 1", "REST endpoints + WebSocket", K.FBLUE), (f"{F.LOC['backend']:,}", "lines of backend Python", K.GREEN), (f"{F.LOC['frontend']:,}", "lines of dashboard code", K.ORANGE), (str(F.TESTS["total"]), "test functions", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 26)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("The shield imports nothing from the llm, agents, simulation or grid packages; a test enforces it.", {"bullet": True, "space_after": 9}),
        ("Any provider failure, timeout or invalid output becomes a labelled deterministic fallback decision that is still shielded.", {"bullet": True, "space_after": 9}),
        ("EnergyAction forbids extra fields and non-finite power; model output is never executed as code, SQL or shell.", {"bullet": True, "space_after": 9}),
        ("Explanations are structured (inputs, factors, action, safety result); no chain-of-thought is requested or stored.", {"bullet": True, "space_after": 9}),
        ("React dashboard with live WebSocket updates; run.bat starts backend (8313) and frontend (5313).", {"bullet": True}),
    ], size=18)
    notes[16] = "Highlights: isolation of the shield, robust fallbacks, strict input schema, explanations without chain-of-thought, and a one-command start."

    # ---- 17 results: overall
    s = S[16]
    set_title(s, "Results – Strategy Comparison (63 runs)")
    wipe(s)
    A = F.ARMS

    def row(label, key, fmt):
        return [label] + [fmt.format(m(a, key)) for a in A]

    rows = [["Metric (mean of 21 runs)", "Rule baseline", "Agents, no shield", "SafeGrid-AI"],
            row("Steps with violations caused by actions", "violation_steps_caused_by_actions", "{:.2f}"),
            row("Violations caused by actions", "violations_caused_by_actions", "{:.2f}"),
            row("Steps with any violation after execution", "violation_steps_executed", "{:.2f}"),
            row("Voltage violation steps", "voltage_violation_steps", "{:.2f}"),
            row("Line overload events", "line_overload_events", "{:.2f}"),
            row("Energy cost (INR, illustrative)", "total_cost", "{:.1f}"),
            row("Renewable utilisation", "renewable_utilization", "{:.3f}"),
            row("Mean step time (ms)", "avg_step_time_ms", "{:.1f}")]
    ftable(s, 0.55, 1.85, [3.5, 1.15, 1.3, 1.15], rows, size=12, row_h=0.46, align=["l", "c", "c", "c"])
    cats = [nice(sc).replace(" ", "\n") for sc in F.SCENARIOS]
    ser = [(F.ARM_LABEL[a], [round(F.sc_mean(sc, a, "violation_steps_caused_by_actions"), 2) for sc in F.SCENARIOS]) for a in A]
    K.bar_chart(s, 7.85, 1.8, 5.0, 4.2, cats, ser, colors=[K.RGBColor(0x8A, 0x97, 0xAB), K.ORANGE, K.GREEN], number_format="0.0", font_size=9, label_size=7, gap=40)
    caption(s, 7.85, 6.02, 5.0, "Steps with violations caused by actions, per scenario (of 24)", 11)
    ftext(s, 0.55, 6.15, 7.1, 0.7, "Mean of 7 scenarios x 3 seeds, 24 steps each. Mock LLM, synthetic 3-node feeder: a simulation result, not field data.", size=12, color=K.FNAVY, italic=True)
    notes[17] = ("Table: means over 21 runs per strategy. The key row is the first: 2.19 steps (rule baseline), 4.71 (agents without shield) and 0.00 (SafeGrid-AI). "
                 "Chart: line congestion and peak demand also hurt the rule baseline (12.0 and 3.3 steps), so the problem is not only an AI problem.")

    # ---- 18 shield decisions (replaces confusion matrix)
    s = S[17]
    set_title(s, "Shield Decisions")
    wipe(s)
    K.picture(s, F.fig("fig_shield_decisions.png"), 0.55, 1.8, 7.3, 4.0)
    for i, (v, l, c) in enumerate([(F.pct(F.APPROVED / F.PROPOSALS), f"approved ({F.APPROVED:,} of {F.PROPOSALS:,})", K.GREEN), (F.pct(F.PROJECTED / F.PROPOSALS), f"projected, scaled down ({F.PROJECTED})", K.GOLD), (F.pct(F.REJECTED / F.PROPOSALS), f"rejected ({F.REJECTED})", K.RGBColor(0xC0, 0x39, 0x2B))]):
        ftile(s, 8.2, 1.9 + i * 1.3, 4.6, 1.15, v, l, c, 26, 13)
    ftext(s, 0.55, 5.9, 12.2, 1.0, [
        (f"{F.SAFETY_EVENTS} safety events logged; hash chain valid in 21 of 21 runs; {F.CONFLICTS} coordinator conflicts. Raw shielded-arm proposals would have caused {F.RAW_WOULD} violations in {F.RAW_WOULD_STEPS} steps; none reached the grid.", {"bullet": True, "space_after": 4}),
        ("No classifier was trained, so a confusion matrix, accuracy or F1 do not apply; the decision counts above take their place.", {"bullet": True}),
    ], size=13)
    notes[18] = "Per scenario the intervention load varies: line congestion has the most projections (23 per run), voltage violation and battery stress the most rejections (about 8 per run). Totals are over 21 shielded runs."

    # ---- 19 pre-existing stress and cost of safety
    s = S[18]
    set_title(s, "Pre-existing Stress & Cost of Safety")
    wipe(s)
    K.picture(s, F.fig("fig_preexisting_stress.png"), 0.55, 1.8, 5.5, 2.8)
    K.picture(s, F.fig("fig_timeseries.png"), 6.2, 1.8, 6.65, 2.5)
    caption(s, 6.2, 4.35, 6.65, "Seed 42: worst loading / voltage with and without the shield (dashed = limit)", 10)
    rows = [["Metric (mean of 21 runs)", "Rule baseline", "Agents, no shield", "SafeGrid-AI"],
            row("Steps with any violation after execution", "violation_steps_executed", "{:.2f}"),
            row("Energy cost (INR, illustrative)", "total_cost", "{:.1f}"),
            row("Renewable utilisation", "renewable_utilization", "{:.3f}"),
            row("Mean step time (ms)", "avg_step_time_ms", "{:.1f}")]
    ftable(s, 0.55, 4.75, [4.4, 2.2, 2.9, 2.7], rows, size=12, row_h=0.4, align=["l", "c", "c", "c"])
    ftext(s, 0.55, 6.75, 12.2, 0.3, "The shield removes what actions add; stress already on the grid (6.33 steps per run) is not removed by design.", size=11, color=K.FNAVY, italic=True)
    notes[19] = ("Left: with every agent idle the grid is already outside its limits in 6.33 of 24 steps on average (20 of 24 in high_solar). The shield does not promise to repair that. "
                 "Right: in the single seed-42 runs the shield holds line loading at exactly 100 % instead of 137.7 %, the highest voltage at 1.056 pu instead of 1.117 pu, and the lowest voltage at 0.932 pu instead of 0.861 pu. "
                 "Cost of safety: step time 348 ms against about 107 ms, renewable use 0.860 against 0.884, cost 225.0 INR against 380.8 INR.")

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    lat = {r[0]: r for r in F.API_LAT}
    for i, (v, l, c) in enumerate([(f"{m('safegrid_shielded', 'avg_step_time_ms'):.0f} ms", "mean step, shield on (about 107 ms off)", K.FBLUE),
                                   (f"{lat['POST /api/simulation/step'][2]:.0f} ms", f"API step, median (p95 {lat['POST /api/simulation/step'][3]:.0f} ms)", K.GREEN),
                                   (f"{lat['POST /api/agents/{id}/decision (agent proposal, dry run)'][2]:.0f} ms", "agent dry run, median", K.ORANGE),
                                   (f"{lat['GET /api/grid/state'][2]:.1f} ms", "grid state read, median", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26, 11)
    rows = [["Check", "Result"],
            [f"pytest ({F.TESTS['total']} functions)", f"{F.TESTS['passed']} passed, {F.TESTS['failed']} failed, {F.TESTS['hung']} hangs (excluded)"],
            ["The 2 failures", "numeric tolerances with pandapower 3.5.6"],
            ["The hanging test", "WebSocket test client; live endpoint works"],
            ["Live smoke test", "20 of 20 checks pass"],
            ["Browser walk", "360 / 768 / 1440 px, light and dark: no console errors"]]
    ftable(s, 0.55, 3.2, [2.2, 4.0], rows, size=12, row_h=0.52)
    lrows = [["Live API call (localhost)", "median ms", "p95 ms"]]
    for name, n_, med, p95 in F.API_LAT:
        lrows.append([name.replace(" (agent proposal, dry run)", " (dry run)").replace("POST dry run, ", "dry run: "), f"{med:.1f}", f"{p95:.1f}"])
    ftable(s, 7.0, 3.2, [3.9, 1.0, 0.9], lrows, size=10.5, row_h=0.37, align=["l", "c", "c"])
    ftext(s, 0.55, 6.6, 12.2, 0.4, f"Measured on {F.MACHINE}. Mock-provider decision latency is 0.1-0.2 ms and says nothing about a real LLM.", size=12, color=K.FNAVY, italic=True)
    notes[20] = ("Latency was measured against a live uvicorn server. Step time rises because every proposal triggers trial power flows. The two failing tests are tolerance checks written for an older pandapower; "
                 "they were left as delivered. No real-LLM latency was measured.")

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (qualitative)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "SafeGrid-AI"]] + [list(r) for r in BASE_COMPARE]
    ftable(s, 0.55, 1.85, [2.0, 4.6, 5.6], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 6.25, 12.2, 0.65, "No numeric comparison exists: the systems, benchmarks and metrics differ. The difference is in what is added (LLM proposer, independent shield, explanations, audit), not in a better number.", size=13, color=K.FNAVY, italic=True)
    notes[21] = "Be explicit that a like-for-like numeric comparison is impossible: the base paper is a DC microgrid in Simulink with rule-based agents; ours is an AC feeder with agent proposals under a shield."

    # ---- 22-24 output screens
    def screens(idx, title, a, ca, b, cb, note):
        sl = S[idx]
        set_title(sl, title)
        wipe(sl)
        K.picture(sl, F.shot(a), 0.4, 1.85, 6.35, 4.0, border=True)
        K.picture(sl, F.shot(b), 6.85, 1.85, 6.35, 4.0, border=True)
        caption(sl, 0.4, 5.98, 6.35, ca)
        caption(sl, 6.85, 5.98, 6.35, cb)
        notes[idx + 1] = note

    screens(21, "Output – Dashboard and Grid View", "04_dashboard.png", "Dashboard after the unsafe_ai_action run: 12 interventions of 36 proposals, 0 violations caused",
            "05_grid_view.png", "Single-line diagram with bus voltage and line loading colours",
            "Real captures from the running application: scenario unsafe_ai_action, mock provider, seed 42, 12 of 12 steps.")
    screens(22, "Output – Safety Shield and Agent Monitor", "07_safety_shield.png", "Safety log, dry-run rejection of a 1000 kW export, active constraints",
            "06_agent_monitor.png", "Agent cards: proposal, shield result and source",
            "The dry run submits a custom action to the shield; nothing is applied. The 1000 kW export is rejected with the failed checks listed.")
    screens(23, "Output – Scenarios and Explainability", "09_scenarios.png", "Scenario runner and strategy comparison",
            "10_explainability.png", "Decision factors, reason codes and safety result per decision",
            "Left: scenario runner with the strategy comparison. Right: the explanation of one decision (inputs, factors, action and safety result); no chain-of-thought is stored.")

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("**Delivered:** a multi-agent power-routing system with a deterministic Symbolic Safety Shield, a hash-chained audit log, per-decision explanations, a React dashboard and a 63-run evaluation.", {"bullet": True, "space_after": 8}),
        (f"**Measured:** 0 violations caused by executed actions in 504 shielded steps ({F.PROPOSALS:,} proposals; {F.pct(F.PROJECTED / F.PROPOSALS)} projected, {F.pct(F.REJECTED / F.PROPOSALS)} rejected); the same agents without the shield caused {F.UNSHIELDED_CAUSED} violations in {F.UNSHIELDED_CAUSED_STEPS} steps.", {"bullet": True, "space_after": 8}),
        ("**Cost:** step time about 3x higher (348 ms against 107 ms); some PV is curtailed in high_solar and voltage_violation.", {"bullet": True, "space_after": 8}),
        ("**Honest gaps:** mock LLM only (no real-LLM result or latency); violation-resolution rate not measured; three seeds on a synthetic 3-node feeder; the guarantee is model-relative, not a proof; 2 tests fail on numeric tolerances and 1 hangs; no user study.", {"bullet": True}),
    ], size=19)
    notes[25] = "Objectives 1, 3, 4 and 5 are met, objective 2 with a mock model, objective 6 partly. Read the gaps slowly: they are part of the result."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("Run the implemented API provider against a live language model; measure decision latency, fallback rate and reasoning quality against the mock.", {"bullet": True, "space_after": 11}),
        ("Measure a violation-resolution rate so pre-existing stress can be repaired, not only prevented.", {"bullet": True, "space_after": 11}),
        ("Larger and three-phase unbalanced networks; real load and PV data; more seeds.", {"bullet": True, "space_after": 11}),
        ("Property-based tests or formal verification of the shield; fix the two tolerance tests and the WebSocket test.", {"bullet": True, "space_after": 11}),
        ("An MCP tool interface, operator roles and authentication, and a user study with grid operators.", {"bullet": True}),
    ], size=21)
    notes[26] = "Ordered by value: the real-LLM evaluation closes the largest gap in this work."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    idx = list(range(1, len(REFS) + 1))
    half = (len(idx) + 1) // 2
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(f"[{n}] {REFS[n - 1]}", {"space_after": 5}) for n in idx[:half]], size=10)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(f"[{n}] {REFS[n - 1]}", {"space_after": 5}) for n in idx[half:]], size=10)
    notes[27] = "References [4] to [7] come from the batch literature survey; their author lists were not recorded and are not guessed."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
