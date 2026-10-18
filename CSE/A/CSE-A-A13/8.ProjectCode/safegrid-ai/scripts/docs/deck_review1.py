"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import content_shared as C
from . import facts as F
from . import pptkit as K
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.SHORT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = C.KEYWORDS
    cp.comments = ""
    cp.last_modified_by = ""


def patch_theme(path, old="Hidden Risk", new="SafeGrid-AI"):
    """The template's theme, colour scheme and font scheme are named after its deck title: rename them in the saved package."""
    import shutil
    import zipfile

    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/theme/") and item.filename.endswith(".xml"):
                data = data.replace(old.encode(), new.encode())
            zout.writestr(item, data)
    shutil.move(tmp, path)


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}\nBatch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s, line):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, line)


# short table rows for the literature survey slide (text condensed from content_shared.LIT_TABLE)
LIT_ROWS = [
    ("Eyimaya and Altin, 2026 [1] (base)", "Four rule-based agents (DER, storage, load, grid)", "DC microgrid, MATLAB/Simulink, 800 V bus", "Voltage held at 800 V in 8 scenarios; scalability, security, no experimental validation left open"),
    ("Zhang et al., 2025 [2]", "LLM planner + validation agent with sandbox and rollback", "IEEE 69-bus, CIGRE MV, IEEE 30-bus", "Best LLM set-ups 100 % success; weaker 63 % and 44 %"),
    ("Picault, Goubet, 2026 [3]", "MCP server exposes a grid simulator to AI agents, human in the loop", "TSO studies, pypowsybl-mcp", "Position paper: requirements and an evaluation plan, no results"),
    ("Multi-agent deep RL, 2025 [4]", "Multi-agent deep reinforcement learning", "Smart-grid energy management", "Listed in the batch survey; full text not reviewed"),
    ("Fault reconstruction, 2024 [5]", "Multi-agent deep reinforcement learning", "Energy-system fault handling", "Listed in the batch survey; full text not reviewed"),
    ("Distributed EV scheduling, 2025 [6]", "Distributed scheduling by vehicle, station and grid agents", "Electric-vehicle charging", "Listed in the batch survey; full text not reviewed"),
    ("MAS in power systems, 2025 [7]", "Review of multi-agent energy management", "Power-grid energy management", "Listed in the batch survey; full text not reviewed"),
    ("McArthur et al., 2007 [8]", "Concepts and challenges of agents for power engineering", "Survey, Part I", "Agents suit distributed grids; verification and trust remain open"),
]


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}

    # ---- 1 title
    title_slide(S[0], "A neuro-symbolic multi-agent controller: language-model agents propose power set-points, a deterministic Symbolic Safety Shield approves, scales down or rejects them, and every decision is explained",
                "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: AI agents propose how power should be routed on a small feeder, and a rule-based shield "
                "checks every proposal against the physics before anything is executed.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work, not yet about results."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: AI that proposes, a shield that decides")
    wipe(s)
    W, H = 3.85, 4.6
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Grids must balance supply and demand continuously while solar, batteries and electric vehicles add variability. Language-model agents can plan, but a non-deterministic model that issues commands to physical equipment can issue an unsafe one.", K.ORANGE, "!"),
        ("Our idea", "Agents see a structured picture of their node and propose one action. A deterministic Symbolic Safety Shield tests it with a power flow and either approves it, scales it down to the largest safe value, or rejects it.", K.BLUE, "i"),
        ("The deliverable", "A simulated three-node feeder, three agents, the shield, an audit log that cannot be edited unnoticed, a REST and WebSocket service and a dashboard that explains each decision.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=16)
    notes[3] = "Three sentences: the problem, the idea, the deliverable. Stress that the shield is ordinary code, not a model."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=14, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("What it does", "Four rule-based agents (distributed energy, energy storage, load, grid) coordinate a DC microgrid with solar, wind, a generator, a battery and critical and non-critical loads, modelled in MATLAB/Simulink.", K.BLUE, "M"),
        ("Reported result", "The bus voltage is held at 800 V and critical loads stay supplied in all eight operating scenarios of the study.", K.GREEN, "R"),
        ("Limits stated by the authors", "Scalability, software integration and cybersecurity are open issues, and the work has no experimental validation. It uses no language model, no separate safety layer and no explanation interface.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. We keep the idea of role-based agents and scenario-based evaluation, and add what the paper lists as open: a language-model proposer, "
                "an independent safety layer, explanations and a tamper-evident record. The last sentence on the slide is our reading of the paper.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Setting", "Reported result or status"]] + [list(r) for r in LIT_ROWS]
    K.table(s, 0.6, 1.35, [2.55, 3.2, 2.75, 3.63], rows, font_size=10.5, row_h=0.6, header_size=12)
    K.text(s, 0.6, 6.85, 12.13, 0.3, "Rows 4-7 are taken from the batch literature survey and classified from title and venue only; their authors were not recorded.", size=11, color=K.TXT, italic=True)
    notes[5] = ("Three papers were read in full: the base paper, Grid-Agent and the MCP position paper. The four conference and review items come from the batch survey; "
                "we did not read their full text, so the table says so instead of guessing their results.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: agent quality varies widely")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 7.4, 5.0, ["Best LLM\nconfigurations", "Weaker\nconfiguration A", "Weaker\nconfiguration B"],
                [("Success rate", [1.0, 0.63, 0.44])], colors=[K.BLUE], number_format="0%", font_size=12, max_value=1.15)
    K.text(s, 0.6, 6.5, 7.4, 0.5, "Source: Zhang et al. [2], violation-resolution success on IEEE and CIGRE test networks. Their numbers, their systems, not ours.", size=11, color=K.TXT, italic=True)
    K.card(s, 8.3, 1.5, 4.43, 2.55, "Takeaway", "When the proposer is a language model, outcomes depend heavily on which model is used. A safety rule layer that does not depend on the proposer is therefore attractive.", K.GREEN, "↑", body_size=14)
    K.card(s, 8.3, 4.25, 4.43, 2.55, "Caveat", "These figures measure how often a planner fixed a violation on other networks. We ran no such comparison, so they are context, not a benchmark for our results.", K.ORANGE, "!", body_size=14)
    notes[6] = ("The chart reports the success rates given in the Grid-Agent paper: the best configurations reach 100 %, weaker ones 63 % and 44 %. They are another group's results on other "
                "networks. We use them only to motivate a model-independent safety layer.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Grid state measured", K.NAVY), ("Controller decides", K.NAVY), ("Set-points sent", K.NAVY), ("Grid responds", K.NAVY)], size=17, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Rule-based agents", "Fixed logic per role, as in the base paper: predictable in the scenarios it was designed for, but each new topology or source needs new rules and nothing explains why a set-point was chosen.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "LLM agents with trial and rollback", "Grid-Agent plans with a language model and checks the plan in a sandbox, rolling back on failure. Safety is a second agent inside the same framework, not an independent rule layer.", K.ORANGE, "2", body_size=17)
    notes[7] = "Two families of existing work: fixed-rule agents and language-model agents with a sandbox. Neither separates the safety rules from the decision maker."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Fixed rules", "Rule sets do not adapt to new sources or topologies; the base paper itself names scalability as open.", K.BLUE, "1"),
        ("Unconstrained AI actions", "A language model can return a plausible but unsafe or malformed command; trial and rollback acts after the fact.", K.ORANGE, "2"),
        ("Opaque decisions", "Operators will not accept set-points they cannot interrogate; little of the reviewed work offers an explanation interface.", K.GREEN, "3"),
        ("Thin verification", "The base paper lacks experimental validation and the MCP paper is a position paper without results.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=16)
    notes[8] = "Four limitations that map to design decisions later: a separate shield, validation before execution, a structured explanation and an evaluation with a baseline."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Independent safety layer", "A rule-based gate whose limits are configuration, separate from whichever model proposes the action."),
            ("Operator explanation", "Per-decision inputs, factors, action and safety result, shown next to the grid, without relying on model reasoning text."),
            ("Prosumer-level coordination", "Several agents sharing one feeder: conflicting proposals must be detected and ordered, not just validated one by one."),
            ("Tamper-evident record", "An audit trail in which an edited or removed safety event is detectable."),
            ("Measured with and without", "The same agents and scenarios run with the shield off and on, so the effect of the layer is a number, not a claim.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.9, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.9, 0.9, k, size=15, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.75, y, 8.95, 0.9, v, size=16, color=K.INK, anchor="m")
    notes[9] = "The batch survey states the gap as combining agentic planning with formal safety shielding and operator-facing explanations. The five rows break that into testable pieces."

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, C.PROBLEM_STATEMENT, size=19, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Seeded scenario, grid state, one proposed action per agent"), ("Output", "Approve, project or reject, with an explanation per decision"), ("Record", "Hash-chained safety log of every intervention")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=14, color=K.TXT)
    notes[10] = "Read the statement slowly. Three parts: agents propose, only safe set-points execute, and everything is explained and recorded."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 10}) for i, o in enumerate(C.OBJECTIVES, 1)], size=21, color=K.INK)
    notes[11] = "Six objectives from the batch proposal. Outcomes against each are reported in the final review."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Perceive", "Each agent receives load, solar, battery state, voltage, price, neighbours and the limits.", K.BLUE, "1"),
        ("Propose", "A language-model provider returns one schema-checked action; failures become a labelled fallback.", K.GREEN, "2"),
        ("Validate", "The shield runs ten ordered checks with trial power flows and projects unsafe requests by bisection.", K.ORANGE, "3"),
        ("Explain and log", "Inputs, factors and the safety result are stored and chained with SHA-256.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "S", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Simulation, service and dashboard", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "pandapower AC power flow on a three-node feeder, FastAPI REST and WebSocket service, SQLite store and a React dashboard for runs, grid, agents, shield and power flow.", size=15, color=K.TXT)
    notes[12] = "Four stages plus the platform around them. The language-model provider is pluggable; a deterministic mock is used for reproducible measurements."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 2.2, 1.35, 8.9, 5.3)
    K.text(s, 0.6, 6.62, 12.13, 0.3, "Presentation, application, simulation and safety core, and data tiers; the LLM provider is the only untrusted component", size=12, color=K.TXT, align="c", italic=True)
    notes[13] = "Follow the diagram left to right: dashboard, API, simulation and safety core, data. The red box is the shield; the orange box is the untrusted proposer."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working multi-agent loop", "Three agents proposing on a simulated feeder every 15 simulated minutes.", K.BLUE, "1"),
            ("Safe execution", "No executed action that creates or worsens a limit violation, to be tested over many runs.", K.GREEN, "2"),
            ("Visible reasoning", "A structured explanation and a safety verdict for every decision.", K.ORANGE, "3"),
            ("Honest evaluation", "Same scenarios with and without the shield, with the limits of the test stated.", K.PURPLE, "4"),
            ("Auditability", "A hash-chained log that reveals edits to the record.", K.BLUE, "5"),
            ("Usable system", "A service and dashboard an operator can run, step and inspect.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=15)
    notes[14] = "These are expectations at this stage, not results. The final review reports the measured values and the gaps."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    pick = [1, 2, 3, 8, 11, 12, 13, 15, 16, 17]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(f"[{n_}] {REFS[n_ - 1]}", {"space_after": 7}) for n_ in pick], size=13, color=K.INK)
    notes[15] = "Selected references; the full list is in the project report."

    # ---- 16 thank you
    thank_you(S[15], "Questions and suggestions are welcome.")
    notes[16] = "Likely questions: why a mock model (reproducible, no key needed, provider is pluggable); is the guarantee a proof (no, it is relative to the simulated model); why a three-node feeder."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    patch_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
