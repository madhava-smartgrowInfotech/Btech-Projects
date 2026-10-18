"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .content_shared import HEADLINE, KEYWORDS
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")
ADDRESS = "[COLLEGE ADDRESS / AFFILIATING UNIVERSITY]"


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                os.makedirs(os.path.dirname(BANNER), exist_ok=True)
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=20, bold=True, space_after=8)
    d.p("A neuro-symbolic multi-agent controller with a deterministic safety shield and operator-facing explanations", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.CITY}", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The software, the experiments and the figures reported here were produced by running the project's own code, and "
        f"the sources we rely on are cited in the references. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.CITY}", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the supporting papers on LLM-based grid control, and the developers of the open-source software "
        "on which the project stands: pandapower, FastAPI, pydantic, numpy, scipy, SQLite, React, Vite, three.js, matplotlib, pytest and Playwright.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    pct = lambda x: F.pct(x / F.PROPOSALS)  # noqa: E731
    d.p("Power networks that host rooftop solar, batteries and electric vehicles must keep voltages, cable currents and transformer loading inside their limits while "
        "generation and demand change every few minutes. Fixed rules do not adapt, and language-model agents that do adapt can, in principle, emit a wrong or "
        "unauthorised command to physical equipment. This project, SafeGrid-AI, builds an explainable multi-agent controller for a simulated three-node low-voltage feeder "
        "in which every action proposed by an agent must pass a deterministic Symbolic Safety Shield before it is executed. Three agents perceive a structured context and obtain one schema-validated action from a pluggable provider; a coordinator detects conflicting "
        "proposals; the shield approves a proposal, scales it down to the largest safe value by bisection on trial AC power flows, or rejects it; only validated "
        "set-points are committed to the pandapower model. Every decision is explained from structured facts, without chain-of-thought text, and non-approved actions are "
        "written to a SHA-256 hash-chained log.")
    d.p(f"The system was evaluated on seven seeded scenarios, three control strategies (a rule baseline, the agents without the shield, and SafeGrid-AI) and three seeds, in "
        f"{F.N_RUNS} simulation runs. " + HEADLINE + " The shield does not repair stress that is already present on the grid when every agent is idle, so steps with any violation fall from "
        f"{F.f2(F.mean('baseline_rule', 'violation_steps_executed'))} (rule baseline) and {F.f2(F.mean('agents_unshielded', 'violation_steps_executed'))} (agents without the shield) to "
        f"{F.f2(F.mean('safegrid_shielded', 'violation_steps_executed'))} (SafeGrid-AI) per 24-step run, not to zero. Safety has a price: the mean step time rises from about "
        f"{F.mean('baseline_rule', 'avg_step_time_ms'):.0f} ms to {F.mean('safegrid_shielded', 'avg_step_time_ms'):.0f} ms because of the trial power flows, and renewable utilisation is "
        f"{F.f2(F.mean('safegrid_shielded', 'renewable_utilization'))} against 1.00 for the baseline.")
    d.p("The results have clear limits. Every agent result comes from a deterministic mock provider; a real-LLM provider is implemented but was not run against a live model, so language-model "
        "reasoning quality and real inference latency are unmeasured. The feeder is synthetic, only three seeds were used, no violation-resolution rate was measured, the guarantee is "
        f"relative to the simulated model and not a formal proof, and the comparison with the base paper (a rule-based DC microgrid study) is qualitative. {F.TESTS['passed']} of {F.TESTS['total']} "
        "automated tests pass; two numeric-tolerance tests fail and one WebSocket test hangs. The system is a research prototype, not a control product.")
    d.p(f"**Keywords:** {KEYWORDS}.", align="left")


ABBR = [
    ("AC", "Alternating Current"), ("API", "Application Programming Interface"), ("ASGI", "Asynchronous Server Gateway Interface"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CSV", "Comma-Separated Values"), ("DC", "Direct Current"), ("DER", "Distributed Energy Resource"), ("DFD", "Data Flow Diagram"), ("EMS", "Energy Management System"),
    ("ER", "Entity-Relationship"), ("EV", "Electric Vehicle"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"),
    ("LLM", "Large Language Model"), ("LV", "Low Voltage"), ("MAS", "Multi-Agent System"), ("MCP", "Model Context Protocol"), ("MV", "Medium Voltage"),
    ("pu", "per unit (fraction of nominal voltage)"), ("PV", "Photovoltaic"), ("REST", "Representational State Transfer"), ("SHA-256", "Secure Hash Algorithm, 256-bit"),
    ("SOC", "State of Charge"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("TSO", "Transmission System Operator"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("WebSocket", "Full-duplex communication protocol over a single TCP connection"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
