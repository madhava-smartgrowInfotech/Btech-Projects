"""
Verifies the documentation set at the batch root.

    python scripts/verify_docs.py [--no-pdf]

Checks (DOCS_PROMPT.md section 8):
 1. every required file exists with its exact name and is not tiny
 2. no template leftovers in any part of any file (including nested packages)
 3. deck slide counts 16 / 20 / 28
 4. the full report has the ten chapters, the references and the three appendices
 5. figures are embedded and every screenshot referenced by the builders exists
 6. student names, roll numbers and the guide name are on the title pages
 7. the IEEE paper is two-column, has five Roman sections, both tables, 4+ captioned figures, 15-20 references and 6 pages
 8. nothing from the reference paper the paper format was taken from
"""

import argparse
import hashlib
import io
import os
import re
import sys
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BATCH = os.path.dirname(os.path.dirname(ROOT))
sys.path.insert(0, ROOT)

REQUIRED = [
    "7.Review_1_PPT.pptx", "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "10.SCOPE & LIMITATIONS.docx", "11. BLOCK DIAGRAMS ALL.docx",
    "12.DESIGN &IMPLEMENTATION.docx", "13.Review_2_PPT.pptx", "14.RESULTS.docx", "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    "16.Final_PPT.pptx", "17.Final Document.docx", "19.IEEE_Paper.docx",
]
DECKS = {"7.Review_1_PPT.pptx": 16, "13.Review_2_PPT.pptx": 20, "16.Final_PPT.pptx": 28}
TEMPLATE_WORDS = ["[STUDENT NAME", "[ROLL NO", "Fakeddit", "PHEME", "Snopes", "MCNN", "bank transaction", "social media manipulation",
                  "Hidden Risk", "Kaliyar", "GloVe", "pneumonia", "Batch D1", "Batch B17"]
PAPER_WORDS = ["skincare", "Dayananda", "Sindhu", "recommender system"]
CHAPTERS = ["CHAPTER 1:", "CHAPTER 2:", "CHAPTER 3:", "CHAPTER 4:", "CHAPTER 5:", "CHAPTER 6:", "CHAPTER 7:", "CHAPTER 8:", "CHAPTER 9:", "CHAPTER 10:"]
STUDENTS = [("DISHA KAUSHAL", "23K91A0563"), ("BOLLAM RAKESH", "23K91A0542"), ("CHINNOLLA RAVINDHAR", "23K91A0554"), ("AFFAN MOHAMMED RAMY", "23K91A0505")]
GUIDE = "CHANDRASHEKAR"

problems, notes = [], []


def fail(msg):
    problems.append(msg)
    print("  FAIL", msg)


def ok(msg):
    print("  ok  ", msg)


def path(n):
    return os.path.join(BATCH, n)


def all_text_blobs(p):
    """Yield (member, bytes) for every part of a package, descending into nested zips (embedded workbooks)."""
    with zipfile.ZipFile(p) as z:
        for m in z.namelist():
            data = z.read(m)
            yield m, data
            if m.lower().endswith((".xlsx", ".docx", ".pptx", ".zip")):
                try:
                    with zipfile.ZipFile(io.BytesIO(data)) as zz:
                        for mm in zz.namelist():
                            yield f"{m}!{mm}", zz.read(mm)
                except zipfile.BadZipFile:
                    pass


def docx_text(p):
    import docx

    d = docx.Document(p)
    out = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out += [c.text for c in r.cells]
    for s in d.sections:
        out += [x.text for x in s.header.paragraphs] + [x.text for x in s.footer.paragraphs]
    return "\n".join(out)


def pdf_pages(p):
    import pypdf

    return len(pypdf.PdfReader(p).pages)


def to_pdf(p):
    out = os.path.splitext(p)[0] + ".pdf"
    if os.path.exists(out) and os.path.getmtime(out) >= os.path.getmtime(p):
        return out
    cached = os.path.join(ROOT, "docs", "_pdf", os.path.basename(out))
    if os.path.exists(cached) and os.path.getmtime(cached) >= os.path.getmtime(p):
        return cached
    try:
        import win32com.client

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(os.path.abspath(p), ReadOnly=True, AddToRecentFiles=False)
            doc.Repaginate()
            doc.ExportAsFixedFormat(os.path.abspath(out), 17)
            doc.Close(False)
        finally:
            word.Quit()
        return out if os.path.exists(out) else None
    except Exception as e:  # noqa: BLE001
        notes.append(f"PDF conversion failed for {os.path.basename(p)}: {e}")
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()

    print("1. files exist")
    for n in REQUIRED:
        p = path(n)
        if not os.path.exists(p):
            fail(f"missing {n}")
            continue
        size = os.path.getsize(p)
        small = size < 40 * 1024
        if small:
            fail(f"{n} is only {size} bytes")
        else:
            ok(f"{n} ({size // 1024} KB)")

    print("2. template leftovers")
    for n in REQUIRED:
        p = path(n)
        if not os.path.exists(p):
            continue
        hits = set()
        for m, data in all_text_blobs(p):
            low = data.lower()
            for w in TEMPLATE_WORDS:
                if w.lower().encode() in low:
                    hits.add((w, m))
        if hits:
            for w, m in sorted(hits)[:6]:
                fail(f"{n}: '{w}' in {m}")
        else:
            ok(f"{n}: clean")

    print("3. deck slide counts")
    from pptx import Presentation

    for n, cnt in DECKS.items():
        if os.path.exists(path(n)):
            got = len(Presentation(path(n)).slides)
            (ok if got == cnt else fail)(f"{n}: {got} slides (expected {cnt})")

    print("4. full report structure")
    p17 = path("17.Final Document.docx")
    if os.path.exists(p17):
        t = docx_text(p17)
        for c in CHAPTERS:
            (ok if c in t else fail)(f"heading {c}")
        for w in ("REFERENCES", "APPENDIX A", "APPENDIX B", "APPENDIX C"):
            (ok if w in t else fail)(f"heading {w}")

    print("5. figures and screenshots")
    for dname in ("01_system_architecture", "02_high_level_workflow", "03_dfd_level0", "04_dfd_level1", "05_dfd_level2", "06_uml_class_diagram",
                  "07_sequence_diagram", "08_use_case_diagram", "09_er_diagram", "10_core_algorithm_flowchart"):
        (ok if os.path.exists(os.path.join(ROOT, "docs", "diagrams", dname + ".png")) else fail)(f"diagram {dname}.png")
    (ok if os.path.exists(os.path.join(ROOT, "docs", "REPORT_FACTS.md")) else fail)("docs/REPORT_FACTS.md")
    for n in ("14.RESULTS.docx", "17.Final Document.docx", "11. BLOCK DIAGRAMS ALL.docx"):
        if os.path.exists(path(n)):
            with zipfile.ZipFile(path(n)) as z:
                imgs = [m for m in z.namelist() if m.startswith("word/media/")]
            ok(f"{n}: {len(imgs)} embedded images")
            if n.startswith("17") and len(imgs) < 30:
                fail(f"{n}: only {len(imgs)} images (need 30+)")
    referenced = set()
    for f in os.listdir(os.path.join(ROOT, "scripts", "docs")):
        if f.endswith(".py"):
            referenced |= set(re.findall(r'shot\("([^"]+\.png)"\)', open(os.path.join(ROOT, "scripts", "docs", f), encoding="utf-8").read()))
    for s in sorted(referenced):
        (ok if os.path.exists(os.path.join(ROOT, "docs", "screenshots", s)) else fail)(f"screenshot {s}")

    print("6. title pages carry names, rolls and guide")
    if os.path.exists(p17):
        t = docx_text(p17)[:6000]
        for nme, roll in STUDENTS:
            (ok if (nme in t and roll in t) else fail)(f"report front matter: {nme} {roll}")
        (ok if GUIDE in t else fail)("report front matter: guide")
    for n in DECKS:
        if os.path.exists(path(n)):
            prs = Presentation(path(n))
            txt = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
            for nme, roll in STUDENTS:
                if roll not in txt or nme.split()[0].lower() not in txt.lower():
                    fail(f"{n}: title slide lacks {nme} {roll}")
            if GUIDE.lower() not in txt.lower():
                fail(f"{n}: title slide lacks the guide")
            else:
                ok(f"{n}: title slide has the team and the guide")
    for n in ("9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "10.SCOPE & LIMITATIONS.docx", "11. BLOCK DIAGRAMS ALL.docx", "12.DESIGN &IMPLEMENTATION.docx", "14.RESULTS.docx", "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx"):
        if os.path.exists(path(n)):
            import docx as _d

            _doc = _d.Document(path(n))
            t = chr(10).join([x.text for x in _doc.paragraphs[:25]] + [c.text for r_ in _doc.tables[0].rows for c in r_.cells])
            good = all(r in t for _, r in STUDENTS) and GUIDE in t
            (ok if good else fail)(f"{n}: team and guide on the first page")

    print("7. IEEE paper")
    pp = path("19.IEEE_Paper.docx")
    if os.path.exists(pp):
        import docx
        from docx.oxml.ns import qn

        d = docx.Document(pp)
        sect = d.sections[-1]
        cols = sect._sectPr.xpath("./w:cols")
        two = bool(cols) and cols[0].get(qn("w:num")) == "2"
        (ok if two else fail)("body section is two-column")
        full = docx_text(pp)
        romans = re.findall(r"^(I|II|III|IV|V)\. [A-Z ]+$", "\n".join(p.text for p in d.paragraphs), re.M)
        (ok if len(set(romans)) == 5 else fail)(f"Roman-numbered sections: {sorted(set(romans))}")
        for t in ("TABLE I", "TABLE II"):
            (ok if t in full else fail)(t)
        figs = re.findall(r"Fig\. \d+\.", full)
        (ok if len(figs) >= 4 else fail)(f"{len(figs)} figure captions")
        refs = re.findall(r"^\[(\d+)\]\s", "\n".join(p.text for p in d.paragraphs), re.M)
        n_refs = max(map(int, refs)) if refs else 0
        (ok if 15 <= n_refs <= 20 else fail)(f"{n_refs} references")
        body = "\n".join(p.text for p in d.paragraphs)
        cited = {int(x) for grp in re.findall(r"\[([\d\], \[]+)\]", body.split("REFERENCES")[0]) for x in re.findall(r"\d+", grp)}
        (ok if cited == set(range(1, n_refs + 1)) else fail)(f"every reference cited and every citation listed (cited {len(cited)} of {n_refs})")
        for w in PAPER_WORDS:
            if w.lower() in full.lower():
                fail(f"paper contains '{w}'")
        words = len(full.split())
        notes.append(f"IEEE paper word count (incl. captions and references): {words}")
        if not a.no_pdf:
            pdf = to_pdf(pp)
            if pdf:
                pg = pdf_pages(pdf)
                (ok if 5.5 <= pg <= 6.5 else fail)(f"paper is {pg} pages")
                notes.append(f"IEEE paper pages: {pg} (PDF: {os.path.basename(pdf)})")
    else:
        fail("19.IEEE_Paper.docx missing")

    print("8. no leftovers from the reference paper or other documents")
    for n in REQUIRED:
        p = path(n)
        if os.path.exists(p) and n.endswith(".docx"):
            hits = [w for w in PAPER_WORDS if w.lower() in docx_text(p).lower()]
            (fail if hits else ok)(f"{n}: {hits or 'clean'}")

    if not a.no_pdf and os.path.exists(p17):
        pdf = to_pdf(p17)
        if pdf:
            pg = pdf_pages(pdf)
            notes.append(f"Full report pages: {pg}")
            (ok if 80 <= pg <= 95 else fail)(f"17.Final Document is {pg} pages (target 80-95)")

    print()
    for n in notes:
        print("note:", n)
    if problems:
        print(f"\n{len(problems)} problem(s)")
        return 1
    print("\nAll documentation checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
