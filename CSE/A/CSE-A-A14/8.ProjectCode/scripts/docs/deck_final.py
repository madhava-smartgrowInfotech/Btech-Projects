"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import json
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .refs import NUM, REF_BY_KEY
from .shared import GAP, LIT_HEADER, LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520
OLD_WORDS = ("Bank Transaction", "Fraud", "Batch D1")


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def tiles(s, items, y=1.9, h=1.15, vsize=26):
    for i, (v, l, c) in enumerate(items):
        ftile(s, 0.55 + i * 3.08, y, 2.88, h, v, l, c, vsize)


def bullets(items, gap=10):
    out = []
    for i, t in enumerate(items):
        out.append((t, {"bullet": True, "space_after": gap if i < len(items) - 1 else 0}))
    return out


def screens(s, shots):
    for i, (f, cap) in enumerate(shots):
        r, c = divmod(i, 3)
        x, y = 0.55 + c * 4.1, 1.85 + r * 2.6
        K.picture(s, F.shot(f), x, y, 3.95, 2.25, border=True)
        ftext(s, x, y + 2.27, 3.95, 0.3, cap, size=12, color=K.FNAVY, align="c")


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE"
    cp.comments = ""
    cp.subject = F.TITLE
    cp.keywords = "customer churn, telecom, machine learning, retention"
    cp.last_modified_by = ""
    notes = {}

    lr, rf, dt = F.LR, F.RF, F.DT
    ds, sp = F.DS, F.SP
    ex = json.load(open(os.path.join(F.EXP, "extra_analysis.json"), encoding="utf8"))
    once = (F.API or {}).get("once_ms", {})
    rep = (F.API or {}).get("repeated_ms", {})
    cv, spl = F.CV, F.SPL

    # footers on every slide
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")
    for lay in list(prs.slide_layouts) + [prs.slide_master]:
        for sh in lay.shapes:
            if sh.has_text_frame and any(w in sh.text_frame.text for w in OLD_WORDS):
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE)
        elif sh.name == "Text 2":
            K.replace_text(sh, "Final Review")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_.title()}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: a machine learning retention system that predicts churn, explains it and suggests actions."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap & Problem Statement\nObjectives\nProposed System & Architecture\nMethodology, Dataset & Model\nDecision Workflow, Interaction & Data Model\nImplementation Highlights\nResults: Metrics, Stability, Threshold & Calibration\nPerformance, Testing & Comparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References")
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, bullets([
        "Customers who leave a telecom provider cost far more to replace than to keep, yet churn is noticed only after they have gone, and a churn score with no reason gives the retention team nothing to act on.",
        "We built ChurnIQ: cleaning and encoding, five classifiers (Logistic Regression, Decision Tree, Random Forest, Gradient Boosting, Hist Gradient Boosting), imbalance handling, selection on F1, risk bands, per-customer drivers, what-if and rule-based retention actions.",
        f"On a 1,000-customer synthetic dataset with the Telco schema, Logistic Regression is best: F1 {F.f3(lr['f1_score'])}, ROC-AUC {F.f3(lr['roc_auc'])}, recall {F.f3(lr['recall'])} on a 200-row test split (53 churners). Results vary by 0.03-0.10 between splits and are moderate, not strong.",
        "A FastAPI service and a React interface add batch scoring, drift monitoring and a demo login; 44 pytest tests and a 63-check smoke test pass.",
        "Data are synthetic, SHAP is not installed (fallback explanations), probabilities are uncalibrated and the retention rules are unmeasured heuristics: all stated as limits.",
    ], 8), size=17)
    notes[3] = "Read the first two bullets, then the headline numbers with their protocol. End on the honest limits."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, F.BASE_PAPER, size=14, color=K.FNAVY)
    rows = [["Aspect", "Base paper"],
            ["Method", "LR, SVC, RF, balanced RF, CatBoost, XGBoost, LightGBM; unified perturbation-based explainer (RNP, FA, IM, AP) compared with SHAP and LIME"],
            ["Data", "Real Telco churn data, 7,043 customers, 21 columns, 26.5% churn; 70/30 split (2,113 test rows)"],
            ["Reported results", "Recall is the objective: up to 0.9077 (CatBoost); Spearman 0.8125 with SHAP; runtime SHAP 289.4 s, LIME 35.1 s, perturbation 18.5 s"],
            ["Stated limits", "Recall as the primary objective may change rankings under other metrics; one structured dataset; CRM integration proposed, not built"]]
    ftable(s, 0.55, 2.9, [2.0, 10.2], rows, size=14, row_h=0.8)
    notes[4] = "A research framework on the real Telco data. We take its motivation; we do not reproduce its explainer or its model set."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [LIT_HEADER] + [list(r) for r in LIT_TABLE]
    ftable(s, 0.45, 1.8, [1.95, 3.1, 2.2, 2.95, 2.1], rows, size=10, row_h=0.55)
    ftext(s, 0.45, 6.62, 12.4, 0.35, "Details read from each paper's own page. Three papers named in the batch survey could not be located online by title and were not analysed.", size=10.5, color=K.FNAVY, italic=True)
    notes[5] = "Results are quoted as the authors report them; the datasets and splits differ, so the numbers are not comparable with each other or with ours."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, bullets([
        "**Prediction + explanation:** most surveyed work compares classifiers; per-customer explanation is a separate step or a research method.",
        "**Explanation to action:** even with SHAP, mapping drivers to a concrete retention step is left to the reader.",
        "**Cost:** SHAP took 289.4 s in the base paper's experiment; lighter ways to rank drivers are useful.",
        "**Imbalance and metrics:** accuracy hides the minority class; recall, precision and F1 must be shown together.",
        "**Usability:** few works deliver an application an analyst can run end to end.",
    ], 14), size=21)
    notes[6] = GAP

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.95, 12.2, 2.3)
    ftext(s, 0.85, 2.0, 11.6, 2.2, PROBLEM_STATEMENT, size=18, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Customer records as CSV (sample: 1,000 rows, 19 features)"), ("Output", "Churn probability, risk band, drivers, retention actions"), ("Record", "Metrics, saved model, batch-scored CSV")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.55, 3.9, 1.8, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.65, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.15, 3.5, 1.1, b_, size=15)
    notes[7] = "The statement has three parts: predict and band, explain, act."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "Six objectives of the proposal. The conclusion states how each was met."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 2.0, 12.2, 1.1, [("Load data", K.FNAVY), ("Prepare + balance", K.FBLUE), ("Train + select", K.GREEN), ("Score + band", K.ORANGE), ("Explain + act", K.PURPLE)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, bullets([
        "Five scikit-learn classifiers behind one registry; best model chosen on F1 (selectable), saved with joblib and restored at start-up.",
        "Risk bands Low / Medium / High / Critical at 30 / 60 / 80, editable and validated on the server; Customer 360 view for one customer.",
        "Local drivers (coefficient x value) and global importance (built-in, permutation); what-if scenarios; ten rule-based retention actions.",
        "FastAPI service and React interface (ChurnIQ) over the same core as the original Streamlit app; demo sign-in; batch scoring; drift check.",
    ], 10), size=18)
    notes[9] = "The proposed system in one line of boxes, then the four properties that matter."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.85, 6.0, 5.1)
    ftext(s, 6.9, 2.0, 5.9, 4.7, bullets([
        "**Presentation:** React SPA (port 5314); original Streamlit UI (8501)",
        "**Application:** FastAPI (8314), 31 routes, HMAC-signed bearer token, in-memory AppState",
        "**Core:** 12 pure-Python modules in src/",
        "**Data and state:** CSV files, joblib model files, memory; no SQL database",
    ], 12), size=17)
    notes[10] = "Four tiers; the core is shared by both interfaces."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Load CSV", K.FNAVY), ("Profile + clean", K.FNAVY), ("Split 80/20", K.PURPLE), ("Fit preprocessor", K.GREEN), ("Balance + train", K.ORANGE), ("Evaluate", K.FBLUE), ("Save + score", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, bullets([
        "Stratified split, test size 20% (seed 42). Median imputation + StandardScaler for 4 numeric columns; mode imputation + one-hot for 15 categorical columns (45 model inputs). Fitted on the training split only.",
        "Imbalance: class_weight=balanced (default), SMOTE on the transformed training split only, or none. Selection on F1 by default.",
        "Beyond one split: 5-fold stratified cross-validation and five random splits, because the 200-row test set holds only 53 churners.",
        "Every number is written to JSON by scripts/run_experiments.py and regenerated into the documents.",
    ], 10), size=17)
    notes[11] = "The workflow is what the Model Training page runs, plus the extra protocols added because the test split is small."

    # ---- 12 dataset, features and model drivers
    s = S[11]
    set_title(s, "Dataset, Features & Drivers")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, f"{ds['rows']:,}", "customers, 21 columns (synthetic)", K.FBLUE, 26)
    ftile(s, 0.55, 3.15, 2.9, 1.1, f"{ds['churn_rate_pct']}%", "churn: 264 churned, 736 retained", K.GREEN, 26)
    ftile(s, 0.55, 4.4, 2.9, 1.1, "19 -> 45", "raw features -> encoded inputs", K.ORANGE, 24)
    K.picture(s, F.fig("feature_importance.png"), 3.7, 1.85, 9.1, 3.55)
    ftext(s, 3.7, 5.5, 9.1, 1.4, bullets([
        "Logistic Regression drivers: contract type, monthly charges, tenure, payment method; permutation importance agrees on monthly charges and month-to-month contracts.",
        "The data were generated with these relationships, so recovering them checks the pipeline: it is not evidence about real customers.",
    ], 6), size=13)
    notes[12] = "Dataset and drivers. Be explicit that the rows are generated with the Telco schema and that the drivers are a check on the pipeline."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 4.1, 5.15)
    ftext(s, 5.0, 2.0, 7.8, 4.8, bullets([
        "Validate: two classes, rows and a target; failures return HTTP 422 and train nothing.",
        "Decide: probability = predict_proba x 100; band from the 30 / 60 / 80 thresholds; batch cut-off 50%.",
        "Explain: coefficient x value (top 5 up, top 5 down); SHAP optional, not installed.",
        "Act: ten IF/THEN rules fire at probability >= 60 or Critical; two fallbacks otherwise. Heuristics: effect not measured.",
    ], 12), size=17)
    notes[13] = "The training and scoring path in decision form."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_diagram.png"), 0.55, 1.8, 5.4, 5.15)
    K.picture(s, F.diag("08_use_case_diagram.png"), 6.3, 1.9, 6.5, 4.8)
    notes[14] = "Sequence of one train-save-explain session and the use cases of the analyst."

    # ---- 15 data model
    s = S[14]
    set_title(s, "Data Model (logical)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.6, 4.65)
    rows = [["Entity", "Where it lives"], ["DATASET, CUSTOMER", "CSV / DataFrame"], ["TRIAL_RUN, RISK_RECORD", "in memory (AppState)"], ["MODEL_ARTEFACT", "models/*.pkl (joblib)"], ["RISK_THRESHOLDS", "in memory, editable"], ["RECOMMENDATION", "computed per customer"]]
    ftable(s, 7.5, 1.9, [2.7, 2.6], rows, size=13, row_h=0.55)
    ftext(s, 7.5, 5.4, 5.3, 1.3, "There is no SQL database: only the active model survives a restart. This is the logical data model of what the code handles.", size=14, color=K.FNAVY, italic=True)
    notes[15] = "State plainly that there is no database; the diagram is a logical model."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    tiles(s, [(str(F.N_ROUTES), "REST route handlers", K.FBLUE), ("44", "pytest tests", K.GREEN), ("16", "React page components", K.ORANGE), ("1", "command to start (run.bat)", K.PURPLE)], 1.9, 1.1, 28)
    ftext(s, 0.55, 3.3, 6.6, 3.6, bullets([
        "HMAC-SHA256 bearer token (12 h), 401 on protected routes; server-side validation (422) of settings and uploads.",
        "In-memory session state; the active model is restored from models/*.pkl at start-up.",
        "Batch scoring validates columns, scores every row and can add top-3 drivers; drift check compares new data with the training data.",
        "Same src/ core serves the Streamlit app, the API and the experiment scripts.",
    ], 8), size=14)
    K.picture(s, F.fig("risk_bands.png"), 7.4, 3.25, 5.4, 3.4)
    notes[16] = ("Highlights: security basics, session state, batch and drift, shared core. Risk bands: Low 253, Medium 481, High 213, Critical 53 of the 1,000 customers at the default thresholds "
                 "(scored on the training data too, so this is the portfolio view, not generalisation).")

    # ---- 17 results: model performance
    s = S[16]
    set_title(s, "Results – Model Performance (80/20 split, seed 42)")
    wipe(s)
    rows = [["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]]
    for n in F.MODEL_ORDER:
        m = F.CW[n]
        rows.append([n, f"{m['accuracy']:.3f}", f"{m['precision']:.3f}", f"{m['recall']:.3f}", f"{m['f1_score']:.3f}", f"{m['roc_auc']:.3f}"])
    ftable(s, 0.55, 1.9, [2.9, 1.0, 1.0, 0.9, 0.8, 1.0], rows, size=13, row_h=0.5, align=["l", "c", "c", "c", "c", "c"])
    K.bar_chart(s, 8.4, 1.8, 4.5, 3.3, [F.SHORT[n] for n in F.MODEL_ORDER], [("F1", [F.CW[n]["f1_score"] for n in F.MODEL_ORDER]), ("ROC-AUC", [F.CW[n]["roc_auc"] for n in F.MODEL_ORDER])],
                colors=[K.FBLUE, K.GREEN], number_format="0.00", font_size=10, max_value=1.0)
    ftext(s, 0.55, 5.35, 12.2, 1.5, bullets([
        f"Logistic Regression is best on F1 ({lr['f1_score']:.3f}), ROC-AUC and recall; Random Forest is within one standard deviation (F1 {rf['f1_score']:.3f}); Decision Tree is weakest.",
        "Gradient boosting has the highest accuracy (0.745) but finds only 15 of 53 churners (recall 0.283): why selection is on F1. Class weights were used; test set 200 rows.",
    ], 8), size=14)
    notes[17] = "Protocol: 800 training rows, 200 test rows (53 churners), class weights, seed 42. ROC-AUC about 0.72 is moderate."

    # ---- 18 confusion matrices and ROC
    s = S[17]
    set_title(s, "Confusion Matrices & ROC / Precision-Recall")
    wipe(s)
    K.picture(s, F.fig("confusion_matrices.png"), 0.55, 1.85, 6.2, 2.2)
    K.picture(s, F.fig("roc_pr_curves.png"), 7.0, 1.85, 5.8, 2.46)
    rows = [["Model", "TN", "FP", "FN", "TP"]]
    for n in ("Logistic Regression", "Decision Tree", "Random Forest", "Gradient Boosting", "Hist Gradient Boosting"):
        cm = F.CW[n]["confusion_matrix"]
        rows.append([n, cm[0][0], cm[0][1], cm[1][0], cm[1][1]])
    ftable(s, 0.55, 4.45, [3.0, 0.9, 0.9, 0.9, 0.9], rows, size=12, row_h=0.4, align=["l", "c", "c", "c", "c"])
    ftext(s, 7.3, 4.5, 5.5, 2.3, bullets([
        "Logistic Regression: 35 of 53 churners found, 56 of 147 loyal customers flagged (precision 0.385).",
        "The curves for LR and RF are close; the Decision Tree is lower (AUC 0.608).",
        "Test split: 200 rows, 53 churners.",
    ], 8), size=14)
    notes[18] = "Rows are the true class and columns the prediction. Cost: a false alarm is a wasted offer; a miss is a lost customer. The business costs were not measured."

    # ---- 19 stability
    s = S[18]
    set_title(s, "Stability: Cross-Validation and Five Splits")
    wipe(s)
    K.picture(s, F.fig("stability.png"), 0.55, 1.8, 7.0, 2.9)
    rows = [["Model", "F1 (5-fold CV)", "F1 (5 splits)", "ROC-AUC (5-fold CV)", "ROC-AUC (5 splits)"]]
    for n in ("Logistic Regression", "Decision Tree", "Random Forest"):
        rows.append([n, f"{cv[n]['f1_score']['mean']:.3f} ± {cv[n]['f1_score']['std']:.3f}", f"{spl[n]['f1_score']['mean']:.3f} ± {spl[n]['f1_score']['std']:.3f}",
                     f"{cv[n]['roc_auc']['mean']:.3f} ± {cv[n]['roc_auc']['std']:.3f}", f"{spl[n]['roc_auc']['mean']:.3f} ± {spl[n]['roc_auc']['std']:.3f}"])
    ftable(s, 0.55, 4.9, [2.7, 2.3, 2.3, 2.5, 2.4], rows, size=13, row_h=0.5, align=["l", "c", "c", "c", "c"])
    ftext(s, 7.8, 1.9, 5.0, 2.9, bullets([
        "Mean ± standard deviation over folds or seeds; each fold holds 200 rows.",
        "LR and RF are within one standard deviation of each other on F1 and ROC-AUC.",
        "The single split is not unusual, but the spread (0.03-0.10) shows how noisy 53 churners are.",
    ], 8), size=14)
    notes[19] = "Cross-validation and five random splits, class weights. The ranking LR ~ RF > DT is stable; the gap between LR and RF is not."

    # ---- 20 threshold, calibration and imbalance
    s = S[19]
    set_title(s, "Threshold, Calibration & Imbalance")
    wipe(s)
    K.picture(s, F.fig("threshold_sweep.png"), 0.55, 1.8, 3.5, 2.42)
    K.picture(s, F.fig("calibration.png"), 4.2, 1.8, 3.1, 2.42)
    sw = {str(x["threshold"]): x for x in ex["threshold_sweep"]}
    lf = {str(x["top_fraction"]): x for x in ex["targeting_lift"]}
    ftext(s, 7.6, 1.85, 5.2, 2.5, bullets([
        f"Threshold 0.50: precision {sw['0.5']['precision']:.3f}, recall {sw['0.5']['recall']:.3f}; 0.30: recall {sw['0.3']['recall']:.3f}, precision {sw['0.3']['precision']:.3f}. F1 is flat from 0.40 to 0.60.",
        f"Calibration: Brier {ex['brier_score']:.3f} against {ex['brier_no_skill_constant_base_rate']:.3f} for a constant predictor; mean prediction {ex['mean_predicted_probability']:.2f} vs {ex['observed_churn_rate_test']:.2f} observed. Scores are rankings, not probabilities.",
        f"Top 20% by score (40 customers) holds {lf['0.2']['churners_captured']} of 53 churners: lift {lf['0.2']['lift']:.1f}x.",
    ], 6), size=12)
    K.picture(s, F.fig("imbalance_strategy.png"), 0.55, 4.45, 5.9, 2.44)
    ftext(s, 6.8, 4.5, 6.0, 2.4, bullets([
        "Without imbalance handling accuracy looks fine (LR 0.745, RF 0.755) while recall collapses (LR 0.189, RF 0.094).",
        "Class weights (default) and SMOTE restore recall; SMOTE gave LR its best single-split F1 (0.526) but was tried on one split only.",
    ], 8), size=14)
    notes[20] = "Threshold, calibration and imbalance on the Logistic Regression test predictions (200 rows). With 20-60 contacts the lift counts are small."

    # ---- 21 performance and testing
    s = S[20]
    set_title(s, "Performance & Testing")
    wipe(s)
    tiles(s, [(f"{F.LAT['single_customer_median']:.1f} ms", "score one customer (median, 200 runs)", K.FBLUE), (f"{F.LAT['full_dataset_1000_rows_ms']:.1f} ms", "score all 1,000 customers", K.GREEN),
              (f"{F.M['app_default_split']['class_weight']['wall_seconds']} s", "train and evaluate five models", K.ORANGE), (f"{rep['dashboard']['median']:.0f} ms", "API dashboard call (median of 10)", K.PURPLE)], 1.9, 1.15, 26)
    rows = [["Test layer", "What it covers", "Result"],
            ["pytest (src/ core, token auth)", "sample data, preprocessing, SMOTE, training, risk bands, batch, rules, persistence", "44 passed (pandas 3 and pandas 2)"],
            ["API smoke test", "63 checks: auth, data, train, evaluation, risk, explain, what-if, retention, batch, drift", "63 PASS, 0 FAIL"],
            ["Browser walk (Playwright, Chrome)", "16 screens with real data at 1440 x 900; 360-1440 px, light and dark", "console clean, no overflow"],
            ["Not done", "load test, penetration test, user study, other browsers, SHAP path", "stated as limits"]]
    ftable(s, 0.55, 3.3, [3.1, 6.3, 2.8], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 6.72, 12.2, 0.4, "Over 127.0.0.1 on an i5-12500 CPU. Via the name localhost every request costs an extra ~2 s on this Windows machine (IPv6 tried first); not a code fault.", size=11, color=K.FNAVY, italic=True)
    notes[21] = "Latency on a CPU-only desktop. Single-customer scoring was 2.85 ms under scikit-learn 1.5.0 / pandas 2.2.1 and 19 ms in the pandas 3 environment of record."

    # ---- 22 comparison with the base paper
    s = S[21]
    set_title(s, "Comparison with Base Paper (context only – different data)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "This project"],
            ["Data", "Real Telco churn, 7,043 customers, 26.5% churn", "Synthetic, 1,000 customers, same schema, 26.4% churn"],
            ["Split", "70/30 (2,113 test rows)", "80/20 stratified (200 test rows, 53 churners)"],
            ["Models", "LR, SVC, RF, balanced RF, CatBoost, XGBoost, LightGBM", "LR, DT, RF (required) + GB, HGB (scikit-learn only)"],
            ["Objective", "Recall is the primary objective", "F1 by default, user-selectable"],
            ["Reported", "Recall 0.4112 (LR) to 0.9077 (CatBoost); LR accuracy 0.7120", f"LR on synthetic data: recall {lr['recall']:.3f}, accuracy {lr['accuracy']:.3f}, ROC-AUC {lr['roc_auc']:.3f}"],
            ["Explanation", "SHAP, LIME, unified perturbation explainer; runtime comparison", "Built-in + permutation importance, coefficient x value; SHAP code present, not run"],
            ["Beyond the paper", "—", "Risk bands, Customer 360, what-if, retention rules, batch, drift, web UI"]]
    ftable(s, 0.55, 1.85, [1.9, 5.0, 5.3], rows, size=12, row_h=0.56)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "Not like-for-like: different data, size, split and model set. The numbers are context, not a ranking; we do not claim to match or beat the paper.", size=13, color=K.FNAVY, italic=True)
    notes[22] = "State plainly that the base paper uses real Telco data and a different protocol. Our data are synthetic; the drivers we recover were built into the generator."

    # ---- 23-24 output screens
    s = S[22]
    set_title(s, "Output Screens – Analysis")
    wipe(s)
    screens(s, [("04_dashboard.png", "Dashboard"), ("09_risk.png", "Customer risk"), ("10_customer360.png", "Customer 360"),
                ("11_explain.png", "Explainable AI"), ("12_whatif.png", "What-if"), ("13_retention.png", "Retention actions")])
    notes[23] = "Real screenshots with 1,000 sample customers and three trained models: dashboard, risk segmentation, Customer 360, explanations, what-if and retention actions."
    s = S[23]
    set_title(s, "Output Screens – Data, Training and Operations")
    wipe(s)
    screens(s, [("05_data.png", "Data"), ("06_explore.png", "Explore"), ("07_train.png", "Model training"),
                ("08_evaluation.png", "Evaluation"), ("14_batch.png", "Batch prediction"), ("15_monitoring.png", "Monitoring")])
    notes[24] = "Data upload and quality, exploration, training, evaluation, batch scoring and monitoring."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, bullets([
        "**Delivered:** an end-to-end churn retention system: data profiling, five models, risk bands, drivers, what-if, retention rules, batch scoring, drift check, tests and documentation.",
        f"**Measured:** Logistic Regression F1 {lr['f1_score']:.3f}, ROC-AUC {lr['roc_auc']:.3f}, recall {lr['recall']:.3f} on a 200-row split; 5-fold CV F1 {cv['Logistic Regression']['f1_score']['mean']:.3f} ± {cv['Logistic Regression']['f1_score']['std']:.3f}. Moderate, and noisy.",
        "**Finding:** contract type, monthly charges and tenure drive risk, as built into the synthetic data; imbalance handling is essential (recall 0.19 without it).",
        "**Not achieved or not measured:** real operator data, SHAP results, calibration (Brier no better than a constant), business effect of the rules, user study.",
    ], 12), size=20)
    notes[25] = "All six objectives were addressed; the evaluation is on synthetic data and is stated that way."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, bullets([
        "Validate on real operator data (for example the public Telco data used in the base paper) and report the comparison on a matched protocol.",
        "Install SHAP and add the base paper's perturbation explainer and LIME; compare their rankings and runtimes with the current fallbacks.",
        "Calibrate probabilities (isotonic or Platt) and choose the threshold by campaign cost and offer value.",
        "Measure the retention rules in a controlled campaign; add gradient-boosting libraries and a drift schedule.",
        "Real accounts and roles, persistent storage, containerised deployment.",
    ], 14), size=21)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = ["oprea", "sikri", "elattar", "rasool", "lundberg", "ribeiro", "slack", "smote", "hegarcia", "breiman", "sklearn", "telco", "friedman"]
    left, right = keys[:7], keys[7:]
    ftext(s, 0.55, 1.85, 6.1, 5.1, [(f"[{NUM[k]}] {REF_BY_KEY[k]}", {"space_after": 5}) for k in left], size=12)
    ftext(s, 6.8, 1.85, 6.0, 5.1, [(f"[{NUM[k]}] {REF_BY_KEY[k]}", {"space_after": 5}) for k in right], size=12)
    notes[27] = "Selected references; the complete list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
