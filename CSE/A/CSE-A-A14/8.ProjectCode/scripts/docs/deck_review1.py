"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import NUM, REF_BY_KEY
from .shared import GAP, LIT_HEADER, LIT_TABLE, NEED, OBJECTIVES, PROBLEM_POINTS, PROBLEM_STATEMENT

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
OLD_LAYOUT_WORDS = ("Hidden Risk", "Social Media", "Multimodal")


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout / master text carries the template's deck title
    holders = list(prs.slide_layouts) + [prs.slide_master]
    for lay in holders:
        for sh in lay.shapes:
            if sh.has_text_frame and any(w in sh.text_frame.text for w in OLD_LAYOUT_WORDS):
                K.replace_text(sh, f"{F.TITLE}  ·  {review_label}")
    # the template's theme carries its own deck name
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            part._blob = part.blob.replace(b"Hidden Risk", b"ChurnIQ")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "customer churn, telecom, machine learning, retention"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, review_label, subtitle):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle)
        elif nm == "review label":
            K.replace_text(sh, review_label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}

    # ---- 1 title
    title_slide(S[0], "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH",
                "A machine learning retention system: predict who will leave, show why, and suggest what to do")
    notes[1] = ("Introduce the team and the title. One sentence: we build a web tool that scores every telecom customer for churn risk, "
                "shows the drivers behind each score and maps them to retention actions.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. Review 1 is about why the problem matters and what is missing in current work; results come in later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: predict churn early, explain it, act on it")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Customers who leave a telecom provider cost far more to replace than to keep, yet churn is usually noticed only after the customer has gone. "
                        "A churn score with no reason attached gives the retention team nothing to act on.", K.ORANGE, "!"),
        ("Our idea", "Clean and encode customer records, compare Logistic Regression, Decision Tree and Random Forest (plus two boosted models) on churn, "
                     "handle the minority-class imbalance, and pick the best model on F1 rather than accuracy.", K.BLUE, "i"),
        ("The deliverable", "A web application, ChurnIQ: risk bands for every customer, the drivers behind each score, what-if scenarios, rule-based retention "
                            "actions and batch scoring.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=19)
    notes[3] = "Three sentences: the problem, our idea, what we deliver. Emphasise that every score comes with its drivers."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=14, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Seven classifiers (LR, SVC, RF, balanced RF, CatBoost, XGBoost, LightGBM) and a unified perturbation-based explainer (RNP, FA, IM, AP) compared with SHAP and LIME.", K.BLUE, "M"),
        ("Data and reported results", "Real Telco churn data, 7,043 customers, 30% test split. Recall is the objective: up to 0.9077 (CatBoost). Explainer runtime 18.5 s against 289.4 s for SHAP.", K.GREEN, "R"),
        ("Stated limits", "Recall is chosen as the primary objective, so rankings may change under other metrics; one structured dataset; CRM integration is proposed as future use, not built.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=17)
    notes[4] = ("Read the citation. The paper is a research framework on a public benchmark. Our project takes its motivation (explain churn, "
                "tie it to retention) and builds an application; it does not reproduce the perturbation explainer.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result"]]
    for r in LIT_TABLE:
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.5, 1.3, [2.45, 3.9, 2.75, 3.23], rows, font_size=10.5, row_h=0.62, header_size=12)
    K.text(s, 0.5, 6.55, 12.3, 0.4,
           "Details were read from each paper's own page. Three papers named in the batch survey could not be located online by title and were not analysed.",
           size=10.5, color=K.TXT, italic=True)
    notes[5] = ("The pattern: classifiers compared on Telco-type data, SHAP or LIME added for explanation, and the base paper focused on explanation cost. "
                "Figures are quoted as reported by the authors; they use different data and splits, so they are not comparable with each other.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: the objective changes the ranking")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 7.6, 5.2, ["LR", "SVC", "RF", "Bal. RF", "CatBoost", "XGBoost", "LightGBM"],
                [("Recall reported in the base paper", [0.4112, 0.5369, 0.5634, 0.8223, 0.9077, 0.7300, 0.8293])], colors=[K.BLUE],
                number_format="0.00", font_size=12, max_value=1.1)
    K.card(s, 8.5, 1.5, 4.23, 2.45, "Takeaway", "On the real Telco data the same churners are caught 41% of the time by Logistic Regression and 91% by CatBoost: the metric and the model both matter.", K.GREEN, "↑", body_size=13)
    K.card(s, 8.5, 4.15, 4.23, 2.45, "Why it matters here", "Accuracy alone hides the minority class, so our app selects the model on F1 by default and shows recall and precision beside it.", K.ORANGE, "!", body_size=13)
    notes[6] = ("Chart values are the recall column of Table 2 of the base paper [1]. We plot them to show how much the model and the objective matter, not as our own results.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Customer calls to cancel", K.NAVY), ("Retention team notified", K.NAVY), ("Generic offer sent", K.NAVY), ("Customer already lost?", K.NAVY)], size=17, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Reactive retention", "Churn is detected after the customer has left or has already asked to cancel; the same discount is offered to everyone, whether or not they were at risk.", K.BLUE, "1", body_size=20)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Research classifiers", "Papers train models on churn tables and report a score. The output is a notebook or a table: no per-customer drivers, no risk bands, no suggested action.", K.ORANGE, "2", body_size=20)
    notes[7] = "Describe the usual path: reactive and uniform. Research models stop at a metric table and are not something a retention analyst can open and use."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Late detection", "Churn is detected after the customer has left, when retention is no longer possible.", K.ORANGE, "1"),
        ("Untargeted offers", "Retention offers are applied uniformly rather than targeted at genuinely at-risk customers.", K.BLUE, "2"),
        ("Scores without reasons", "A churn score without a reason gives the retention team nothing to act upon.", K.GREEN, "3"),
        ("Imbalanced data", "Churners are a minority, so models under-predict them and accuracy looks better than it is.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=20)
    notes[8] = "These four limitations are the problem statement of the batch. Each maps to a design decision: risk bands, targeting, drivers, imbalance handling."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Prediction + explanation", "Most surveyed work compares classifiers; per-customer explanation is usually a separate step or a research method."),
            ("Explanation to action", "Even where SHAP is used, mapping drivers to a concrete retention step is left to the reader."),
            ("Explanation cost", "SHAP is heavy (289.4 s in the base paper's experiment); lighter ways to rank drivers are needed."),
            ("Imbalance and metrics", "Accuracy alone hides the minority class; recall, precision and F1 must be shown together."),
            ("Usability", "Few works deliver an application an analyst can run end to end.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.9, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.9, 0.9, k, size=15, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.75, y, 8.95, 0.9, v, size=18, color=K.INK, anchor="m")
    notes[9] = ("The batch literature survey words the gap as: pairing accurate prediction with per-customer explanation so that retention action can be targeted rather than generic. "
                "The rows above spell that out.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.55, 11.45, 2.6, PROBLEM_STATEMENT, size=19, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "A CSV of customer records (sample: 1,000 rows, 19 features)"), ("Output", "Churn probability, risk band, drivers, retention actions"), ("Record", "Metrics, saved model, batch-scored CSV")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=15, color=K.TXT)
    K.text(s, 0.6, 6.2, 12.13, 0.5, "Need: " + NEED, size=16, color=K.NAVY, bold=True)
    notes[10] = "Read the statement slowly: predict, group into bands, show drivers, turn into actions, all without writing code."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 11}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "Six objectives from the batch proposal: data and EDA, cleaning and imbalance, three models, five metrics, drivers, dashboard."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "Detect target and ID, drop duplicates, impute, scale numeric fields, one-hot encode categorical fields.", K.BLUE, "1"),
        ("Balance + train", "Stratified 80/20 split; class weights, SMOTE or none; LR, DT, RF (plus GB and HGB).", K.GREEN, "2"),
        ("Evaluate", "Accuracy, precision, recall, F1, ROC-AUC, confusion matrix; best model chosen on F1.", K.ORANGE, "3"),
        ("Explain + act", "Risk bands 30 / 60 / 80, drivers per customer, what-if, rule-based retention actions.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.4, cw, 2.55, t, b_, col, gl, body_size=14)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.6, 4.1, 7.2, 2.75)
    K.circle(s, 8.1, 4.3, 0.6, K.NAVY, "W", 16)
    K.text(s, 8.85, 4.3, 3.9, 0.6, "Web application: ChurnIQ", size=16, color=K.INK, bold=True, anchor="m")
    K.text(s, 8.1, 5.05, 4.6, 1.6, "React interface and FastAPI service over a pure-Python scikit-learn core; demo sign-in; runs locally with one command. The diagram shows the eight steps of one analysis.", size=14, color=K.TXT)
    notes[12] = "Four stages and the application around them. No database and no LLM: data are CSV, the saved model is a joblib file."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 6.8, 5.2)
    K.text(s, 7.7, 1.5, 5.0, 5.0, [
        ("**Presentation:** React single-page app (ChurnIQ).", {"bullet": True, "space_after": 8}),
        ("**Application:** FastAPI, bearer-token demo login, 31 routes.", {"bullet": True, "space_after": 8}),
        ("**Core:** pure-Python modules for data, preprocessing, training, evaluation, explanations, risk, retention.", {"bullet": True, "space_after": 8}),
        ("**Data and state:** CSV files, joblib model files, in-memory session state.", {"bullet": True}),
    ], size=18, color=K.INK)
    notes[13] = "Follow the diagram left to right: the interface, the API, the core modules and the data stores. The Streamlit app is the original front end over the same core."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Model comparison", "Three required classifiers and two boosted ones evaluated on five metrics.", K.BLUE, "1"),
            ("Visible reasoning", "A ranked list of drivers for every customer and for the portfolio.", K.GREEN, "2"),
            ("Targeted action", "Risk bands and rule-based retention steps in place of one offer for all.", K.ORANGE, "3"),
            ("Honest evaluation", "Metrics with their split, spread across seeds, and known limits stated.", K.PURPLE, "4"),
            ("Usable system", "Dashboard, what-if, batch scoring and a drift check in one application.", K.BLUE, "5"),
            ("Reproducibility", "One script reproduces every reported number from the code.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=18)
    notes[14] = "These are expected outcomes, not results. We will report what the held-out test split shows, including weaknesses."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = ["oprea", "sikri", "elattar", "rasool", "lundberg", "ribeiro", "slack", "smote"]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(f"[{NUM[k]}] {REF_BY_KEY[k]}", {"space_after": 5}) for k in keys], size=13, color=K.INK)
    notes[15] = "Selected references; the full list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = "Likely questions: why Logistic Regression, Decision Tree and Random Forest; why F1; why no SHAP result; what data we use (a synthetic sample with the Telco schema)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
