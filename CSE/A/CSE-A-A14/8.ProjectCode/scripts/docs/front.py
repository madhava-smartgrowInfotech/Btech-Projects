"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.FIGS, "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                os.makedirs(F.FIGS, exist_ok=True)
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=22, bold=True, space_after=4)
    d.p(f"({F.SUBTITLE})", align="center", size=14, bold=True, space_after=8)
    d.p("An explainable churn-risk and retention application: ChurnIQ", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: [PLACE]", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: [PLACE]", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the related work we studied, the providers of the public Telco churn data set whose schema this project follows, and the developers of the open-source software "
        "used in this project (Python, scikit-learn, pandas, imbalanced-learn, FastAPI, Streamlit, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    d.p("Customer churn, the loss of subscribers to a competitor, is costly for telecom operators because acquiring a new customer costs several times more than keeping an existing one, "
        "and because churn is usually noticed only after the customer has left. Customer records contain early signals such as tenure, contract type, monthly charges, internet service and payment method. "
        "This project builds a machine-learning customer-retention system around those signals. The data are loaded and profiled, cleaned, explored, and prepared with a pipeline that imputes, scales and one-hot encodes the fields. "
        "Logistic Regression, Decision Tree and Random Forest classifiers (and optionally Gradient Boosting) are trained with class weights or SMOTE to counter the minority churn class, compared on accuracy, precision, recall, F1-score and ROC-AUC, and the best model is saved.")
    d.p("Each customer receives a churn probability and a Low, Medium, High or Critical risk band with user-editable thresholds, the factors pushing the score up or down, a what-if analysis, and rule-based retention suggestions. "
        "A new customer file can be scored in batch and checked for drift. The system is delivered as the original Streamlit application and as a React web interface (ChurnIQ) over a FastAPI service.")
    d.p(f"Because no real operator data were available, the evaluation uses a synthetic 1,000-customer data set with the public Telco schema and {F.DS['churn_rate_pct']}% churn; the results describe the pipeline and not real customers. "
        f"On the 200-row test split the best model by F1, Logistic Regression, reaches accuracy {F.LR['accuracy']:.3f}, recall {F.LR['recall']:.3f}, F1 {F.LR['f1_score']:.3f} and ROC-AUC {F.LR['roc_auc']:.3f}. "
        f"Cross-validation and repeated splits show spreads of 0.03 to 0.10 and that Logistic Regression and Random Forest cannot be separated reliably; without imbalance handling recall falls to {F.NONE['Logistic Regression']['recall']:.3f} although accuracy looks good. "
        "The probabilities are not calibrated and are used as a ranking. SHAP was not installed, so explanations use coefficients and permutation importance. "
        "The base paper's unified explainer is not reproduced, and the metrics are not comparable with its real-data results. "
        f"{F.TESTS['pytest']} unit tests and a {F.TESTS['smoke']}-check end-to-end API test pass.")
    d.p("**Keywords:** " + F.KEYWORDS + ".", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CRM", "Customer Relationship Management"), ("CSV", "Comma-Separated Values"), ("DFD", "Data Flow Diagram"), ("DT", "Decision Tree"),
    ("EDA", "Exploratory Data Analysis"), ("ER", "Entity-Relationship"), ("FN / FP", "False Negative / False Positive"), ("GB / HGB", "Gradient Boosting / Histogram Gradient Boosting"),
    ("HMAC", "Hash-based Message Authentication Code"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"), ("LIME", "Local Interpretable Model-agnostic Explanations"),
    ("LR", "Logistic Regression"), ("ML", "Machine Learning"), ("PBE", "Perturbation-Based Explainer"), ("PR", "Precision-Recall"), ("RF", "Random Forest"),
    ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("SHAP", "SHapley Additive exPlanations"), ("SMOTE", "Synthetic Minority Over-sampling Technique"),
    ("SPA", "Single-Page Application"), ("TN / TP", "True Negative / True Positive"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("XAI", "Explainable Artificial Intelligence"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.6], size=11)
