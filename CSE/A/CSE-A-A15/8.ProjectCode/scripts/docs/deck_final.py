"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (letterhead, rule, title and footer chrome inherited).

    python -m scripts.docs.deck_final            # build (waits up to 25 min for the diagrams/figures to exist)
"""

import os
import sys
import tempfile
import time

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520

NEEDED = ["01_system_architecture.png", "02_high_level_workflow.png", "06_uml_class.png", "09_er_diagram.png",
          "10_flowchart_algorithm.png", "11_bundle_layout.png"]
NEEDED_FIGS = ["f01_psnr_vs_fill.png", "f02_ssim_vs_fill.png", "f04_stage_timings.png", "f05_rsa_key_size.png", "f06_robustness_matrix.png",
               "f08_steganalysis.png", "f10_api_latency.png", "f11_reported_psnr_comparison.png", "f13_psnr_by_payload.png"]


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "a15_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "hybrid encryption, AES-GCM, RSA-OAEP, LSB steganography, PSNR, SSIM, data protection"
    notes = {}

    EV, EX, API, SM = F.EV, F.EX, F.API, F.S
    qs = [r for r in EV["quality_and_speed"] if r["fits"]]
    fs = EX["fill_study"]
    rsa = {r["rsa_bits"]: r for r in EX["rsa_study"]}
    sa = EX["steganalysis"]
    tam = EX["tamper_study"]
    fresh = EX["freshness"]
    pe = EX["plain_vs_encrypted"]["rows"]
    core_rob = {r["attack"]: r for r in EV["robustness"]}
    extra = EX["extra_attacks"]
    n_core_surv = sum(1 for r in EV["robustness"] if r["recovered"] == 6)
    n_extra_surv = sum(1 for r in extra if r["recovered"] == r["of"])
    n_att = len(EV["robustness"]) + len(extra)
    fill100 = [r for r in fs["rows"] if r["fill"] == 1.0]
    floor_lo, floor_hi = min(r["psnr_mean"] for r in fill100), max(r["psnr_mean"] for r in fill100)
    ssim_min_fill = min(r["ssim_mean"] for r in fs["rows"])
    psnr_min_all = min(r["psnr_mean"] for r in fs["rows"])
    chi_n = sa["chi2_total"]
    chi_stego_min = min(r["chi2_p_stego_prefix"] for r in sa["rows"])
    chi_cover_lo = sum(1 for r in sa["rows"] if r["chi2_p_cover_prefix"] < 0.5)
    ws_base = sorted(r["ws_estimate_cover"] for r in {x["cover"]: x for x in sa["rows"]}.values())

    def lat(k):
        return API[k]["median_ms"]

    for sl in S:  # footers on every slide except the title
        for sh_ in sl.shapes:
            if sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 30)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = ("Introduce the team and CipherVeil: a file is encrypted with AES-256-GCM, the AES key is wrapped with the receiver's RSA key, "
                "and the bundle is hidden in the least significant bits of a PNG.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System, Architecture & Methodology\nBundle Format, Formulation & Design\nResults: Quality, Robustness, Parameters\nSteganalysis Finding & Performance\nComparison with Base Paper & Output Screens\nConclusion, Limitations & Future Scope\nReferences")
    notes[2] = "The agenda follows the report: problem, design, method, results, the steganalysis finding, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Encryption alone shows that something secret is moving; hiding alone protects nothing once found. CipherVeil gives a file two layers: AES-256-GCM for the data, RSA-OAEP for the key, and LSB steganography to hide the result in an ordinary-looking PNG (FastAPI + React).", {"bullet": True, "space_after": 8}),
        (f"Measured on 6 covers and payloads of 1-128 KiB: {SM['runs']} of {SM['runs']} round trips byte-exact; PSNR {SM['psnr_db_min']:.1f}-74.6 dB (mean {SM['psnr_db_mean']:.1f}); SSIM at least {SM['ssim_min']:.4f}; capacity {SM['max_capacity_bits_per_pixel']:.1f} bits per pixel.", {"bullet": True, "space_after": 8}),
        (f"A closed-form PSNR model matches measurement within {fs['max_abs_dev_vs_theory_db']:.3f} dB; PSNR never drops below {floor_lo:.1f} dB, even with the cover full.", {"bullet": True, "space_after": 8}),
        (f"Tampering is detected: {tam['trials']} of {tam['trials']} single-bit flips were rejected, none released wrong data.", {"bullet": True, "space_after": 8}),
        (f"Honestly: the hidden data lives only in lossless PNGs (JPEG, noise, resize and rotation destroy it) and a chi-square test flags the embedded region in {chi_n} of {chi_n} cases, so it is hidden from the eye, not from statistics.", {"bullet": True, "space_after": 8}),
        (f"{F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks pass; no dataset is needed, and the six covers and four payloads are small.", {"bullet": True}),
    ], size=17)
    notes[3] = "Give the product, the headline numbers, then the honest limits: lossless only, and statistically detectable."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.6, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "ProSt: Provably Secure Homomorphic Steganography"],
            ["Idea", "Secret bits are hidden at a secret position among random decoy LSBs of cover images; an untrusted cloud then evaluates a Boolean circuit of Fredkin gates on every pixel without learning where the secrets are"],
            ["Scheme", "Four algorithms: LocGen (secret position), Emb (LSB substitution), Comp (gate evaluation on all pixels), Ext (extract the result at the secret position)"],
            ["Security", "Honest-but-curious adversary; advantage of order 1/n in the number n of LSB positions (Theorem 1); n of about 10^6 or more is recommended for advantage below 10^-6"],
            ["Evaluation", "Python/NumPy, 128x128x3 image, 17 gates, 61 wires: 8.38 s on an i5-1135G7 (the paper's own projection: about 33 s for 256x256). No PSNR, SSIM or capacity is reported"],
            ["Stated limits", "Circuit privacy not proven; one image per wire is costly; scalability and GPU acceleration left as future work"]]
    ftable(s, 0.55, 2.45, [1.6, 10.6], rows, size=12, row_h=0.72)
    notes[4] = ("Ahmad and Rass, IEEE Access 2026. CipherVeil adopts the LSB carrier and the key-exchange framing. It does not implement homomorphic computation, "
                "decoy-bit location hiding or a provable-security argument; it moves a whole file instead.")

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Relevance / limitation"],
            [f"Ahmad & Rass [{R['base']}] (base)", "2026", "Homomorphic computation on LSB-hidden bits with Fredkin gates", "No PSNR, SSIM or capacity; hides location, not existence"],
            [f"Raiyan & Kabir, SCReedSolo [{R['screed']}]", "2025", "Shuffle + Fernet + Reed-Solomon + LSB, 3 bpp", "PSNR 51.8-57.7 dB; closest to our pipeline; adds error correction"],
            [f"Fatemieh et al. [{R['memr']}]", "2026", "Trivium / Grain-128a on memristive in-memory logic, steganography demo", "Hardware-level; energy and steps down by up to 42 % / 44 %"],
            [f"Batch survey [{R['s_triple']}]-[{R['s_dwt']}]", "2024-25", "Visual cryptography, image-object hiding, chaotic encryption + CNN, AES + DWT-DCT", "Listed by title only; not verified, no result quoted"],
            [f"Chan & Cheng [{R['chan']}]", "2004", "Simple LSB substitution with pixel adjustment", "Establishes LSB quality; no key protection"],
            [f"Westfeld & Pfitzmann [{R['westfeld']}]; Fridrich et al. [{R['rs']}]", "1999-2001", "Chi-square and RS attacks on LSB replacement", "Show sequential LSB is statistically detectable"],
            [f"Fridrich & Goljan [{R['ws']}]; Ker & Bohme [{R['wsrev']}]", "2004-08", "Weighted-stego payload estimation", "Used here as a blind check of our stego images"]]
    ftable(s, 0.55, 1.8, [3.3, 0.9, 4.0, 4.0], rows, size=11, row_h=0.6, align=["l", "c", "l", "l"])
    notes[5] = "The four batch-survey papers could not be verified beyond title and year, so no result is quoted for them. Steganalysis work is why we test detectability."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in [
        "The surveyed work advances steganography and encryption independently, including provably secure constructions (batch survey).",
        "The base paper reports no PSNR, SSIM or capacity: it targets computation on hidden bits, not file transfer.",
        "What is missing is a practical hybrid AES-RSA plus LSB pipeline, evaluated end to end on capacity, imperceptibility and speed.",
        "Robustness and detectability of such a pipeline are rarely measured with real decryption and statistical tests, so claims of invisibility go untested.",
        "Usable tooling (key handling, capacity checks, analysis views) is seldom provided with the algorithm."]], size=19)
    notes[6] = "The gap is the third bullet; the fourth and fifth explain why we measure robustness and detectability and ship a working tool."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.6)
    ftext(s, 0.85, 1.9, 11.6, 2.5, [
        "Encryption alone signals that sensitive data is being transmitted, inviting attack.",
        "Steganography alone offers no protection once the hidden payload is discovered.",
        "Symmetric encryption is fast, but its key exchange remains the weak point.",
        "Naive LSB embedding degrades the cover image and is detectable by steganalysis.",
        "Need: a two-layer scheme that encrypts data with hybrid AES-RSA and hides the result imperceptibly inside a cover image."], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "A file up to 25 MB, a receiver's RSA public key and a PNG/image cover"), ("Output", "A stego PNG, PSNR/SSIM/capacity measures, and on receipt the original file with a hash check"), ("Record", "Operation log with sizes and timings (no keys), metrics JSON, screenshots")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.7, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.8, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.3, 3.5, 1.0, b_, size=14)
    notes[7] = "Four problem points from the batch statement, then the system's input, output and record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate([
        "Implement AES encryption for the secret payload and RSA encryption for the AES key.",
        "Build a secure key generation and exchange mechanism between sender and receiver.",
        "Embed the encrypted payload into a cover image using LSB steganography.",
        "Implement the extraction and decryption path to recover the original data.",
        "Measure imperceptibility using PSNR and SSIM against the original cover image.",
        "Evaluate embedding capacity, processing time and robustness of the combined scheme."], 1)], size=22)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Send", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("File + SHA-256 header", K.FNAVY), ("AES-256-GCM (fresh key, nonce)", K.FBLUE), ("RSA-OAEP wraps the key", K.GREEN), ("CVL1 bundle", K.ORANGE), ("Capacity check", K.PURPLE), ("LSB embed, save PNG", K.FNAVY)], size=12, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Receive", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Stego PNG", K.FNAVY), ("Read length + bits", K.FBLUE), ("RSA unwrap with private key", K.GREEN), ("AES-GCM decrypt, tag check", K.ORANGE), ("SHA-256 compare", K.PURPLE), ("Original file", K.FNAVY)], size=12, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, [
        ("A fresh 256-bit AES key and 96-bit nonce for every message; the magic CVL1 is bound in as associated data, so a modified bundle never decrypts.", {"bullet": True, "space_after": 7}),
        ("Analysis, robustness and benchmark views measure how visible the hiding is and what survives: PSNR, SSIM, histograms, bit planes, nine attacks.", {"bullet": True, "space_after": 7}),
        ("The operation log records sizes, timings and metrics, never keys or file contents.", {"bullet": True}),
    ], size=16)
    notes[9] = "Two chains: send and receive. The chain is standard; the contribution is the authenticated bundle format, the measurement and the honest evaluation."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.2, 5.1)
    ftext(s, 7.95, 1.95, 4.9, 4.9, [
        (f"Console: React 18 + Vite + Tailwind, port {F.WEB_PORT}; 8 pages", {"bullet": True, "space_after": 8}),
        (f"API: FastAPI, port {F.API_PORT}; 24 endpoints, JWT (12 h), bcrypt", {"bullet": True, "space_after": 8}),
        ("Services: crypto, lsb, metrics, pipeline, plots, robustness, benchmark - about 1,040 backend lines", {"bullet": True, "space_after": 8}),
        ("Storage: SQLite with 3 tables (users, keypairs, oplog); covers, payloads and outputs as files", {"bullet": True, "space_after": 8}),
        ("No external services; runs offline", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. The same pipeline module is used by the API, the benchmark and ml/eval.py, so measured numbers are the numbers of the product."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, [
        ("No dataset: 6 committed cover images (200x200 to 1280x720) and 4 synthetic payloads (99 B to 21.5 KiB), regenerated by a seeded script.", {"bullet": True, "space_after": 8}),
        ("Core evaluation: payloads of 1-128 KiB, RSA-3072, median of 3 repeats; 22 of 30 runs fit, 8 blocked by the capacity check.", {"bullet": True, "space_after": 8}),
        ("Extra experiments: RSA key size, 6 covers x 7 fill levels x 5 seeds, plain vs encrypted, steganalysis, 240 tamper flips, extra attacks.", {"bullet": True, "space_after": 8}),
        ("Robustness: every attack ends in a real decryption and a hash comparison.", {"bullet": True}),
    ], size=14)
    notes[11] = "Method in one picture, then the evaluation design. Everything is seeded and runs offline."

    # ---- 12 bundle format and algorithm
    s = S[11]
    set_title(s, "Bundle Format & Core Algorithm")
    wipe(s)
    K.picture(s, F.diag("11_bundle_layout.png"), 0.55, 1.75, 6.5, 4.55)
    K.picture(s, F.diag("10_flowchart_algorithm.png"), 7.3, 1.75, 5.5, 5.2)
    ftext(s, 0.55, 6.35, 6.5, 0.7, "Inside AES-GCM: header length | header JSON {name, size, sha256} | file. Hidden stream: 32-bit length, then the bundle, one bit per R, G, B value, row by row from the top-left pixel.", size=11.5, color=K.FNAVY, italic=True)
    notes[12] = "Left: byte layout. Right: the embed and extract flow with the capacity check and the GCM tag check. A wrong key, a flipped bit or a truncated bundle fails the tag and releases nothing."

    # ---- 13 formulation
    s = S[12]
    set_title(s, "Formulation & Test Material")
    wipe(s)
    fbox(s, 0.55, 1.8, 7.3, 4.0, fill=K.FZEBRA)
    ftext(s, 0.8, 1.9, 6.9, 3.9, [
        ("Capacity of a W x H cover:  C = (W*H*3 - 32) / 8 bytes  (3.0 bits per pixel)", {"space_after": 9}),
        ("Imperceptibility:  MSE = mean (x - x')^2,   PSNR = 10 log10(255^2 / MSE),  SSIM per colour channel, averaged", {"space_after": 9}),
        ("Random payload bits flip half of the touched values, so with b embedded bits in n values:  E[MSE] = 0.5 b / n", {"space_after": 9}),
        ("Closed form:  PSNR = 10 log10(255^2 n / (0.5 b));  with the cover full (b = n) the floor is 51.14 dB", {"space_after": 9}),
        ("Authenticity: GCM tag over header + file with CVL1 as associated data; RSA-OAEP with SHA-256 and MGF1-SHA-256 wraps the 256-bit key", {"space_after": 0}),
    ], size=14)
    ftile(s, 8.1, 1.8, 4.65, 1.05, f"{fs['max_abs_dev_vs_theory_db']:.3f} dB", "largest gap between measured PSNR and the closed form", K.FBLUE, 24, 12)
    ftile(s, 8.1, 3.0, 4.65, 1.05, "6 covers", "200x200 to 1280x720; public domain / CC0 / generated", K.GREEN, 24, 12)
    ftile(s, 8.1, 4.2, 4.65, 1.05, "4 payloads", "memo PDF 21.5 KiB, CSV 2.7 KiB, notes, JSON (synthetic)", K.ORANGE, 24, 12)
    ftext(s, 0.55, 6.0, 12.2, 0.9, "Why these choices: AES-GCM gives secrecy and integrity in one pass; OAEP is the standard randomised RSA padding; 1-bit sequential LSB is the simplest carrier and the one the base paper uses. LSB matching, adaptive embedding and error-correcting codes were left out of scope.", size=13, color=K.FNAVY, italic=True)
    notes[13] = "The formulation is short on purpose: capacity, PSNR and the random-bit model. The model was validated over 6 covers x 7 fill levels x 5 payload seeds."

    # ---- 14 design
    s = S[13]
    set_title(s, "Design – Classes and Data")
    wipe(s)
    K.picture(s, F.diag("06_uml_class.png"), 0.55, 1.8, 6.5, 4.5)
    K.picture(s, F.diag("09_er_diagram.png"), 7.2, 1.8, 5.6, 3.3)
    rows = [["Table", "Holds"], ["users", "username, bcrypt hash"], ["keypairs", "name, bits, fingerprint, public PEM, private PEM (empty = public only)"], ["oplog", "operation, status, JSON details (sizes, timings, metrics; no secrets)"]]
    ftable(s, 7.2, 5.2, [1.3, 4.3], rows, size=10.5, row_h=0.4)
    ftext(s, 0.55, 6.4, 6.5, 0.6, "Private keys are stored unencrypted in the local SQLite file: treat data/app.db as sensitive.", size=12, color=K.FNAVY, italic=True)
    notes[14] = "Class model of the services and routes, and the three tables. State the key-storage limitation here."

    # ---- 15 implementation highlights
    s = S[14]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("24", "API endpoints", K.FBLUE), (f"{F.PYTEST_TOTAL} + {F.SMOKE_TOTAL}", "pytest tests + smoke checks", K.GREEN), ("3", "database tables", K.ORANGE), ("8", "React pages", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        ("Python 3.11, FastAPI, SQLAlchemy, `cryptography` 44, Pillow, NumPy, scikit-image (metrics), matplotlib (plots); React 18, Vite 6, Tailwind 4, Recharts.".replace("`", ""), {"bullet": True, "space_after": 9}),
        ("One pipeline module (hide / reveal) serves the API, the benchmark and the evaluation script; samples are resolved from a directory listing, never joined from user input.", {"bullet": True, "space_after": 9}),
        ("RSA keys below 2048 bits are rejected; uploads are capped at 25 MB; the oversize check runs before any embedding.", {"bullet": True, "space_after": 9}),
        ("Known gaps: private keys unencrypted in the database, placeholder JWT secret in .env, no login rate limit, RSA only, no sender signatures.", {"bullet": True, "space_after": 9}),
        ("Defect found while documenting: the Send page crashes for an account with no keys (a missing icon import); not fixed in the product code.", {"bullet": True}),
    ], size=16)
    notes[15] = "Highlights and honest gaps, including the Send-page defect found while documenting."

    # ---- 16 results: headline
    s = S[15]
    set_title(s, "Results – Quality, Capacity and Speed")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{SM['runs']} / {SM['runs']}", "round trips byte-exact (8 blocked by the capacity check)", K.FBLUE),
                                   (f"{SM['psnr_db_min']:.1f} dB", f"minimum PSNR; mean {SM['psnr_db_mean']:.1f} dB", K.GREEN),
                                   (f"{SM['ssim_min']:.4f}", f"minimum SSIM; mean {SM['ssim_mean']:.4f}", K.ORANGE),
                                   (f"{SM['median_decrypt_ms']:.0f} ms", f"RSA-3072 unwrap + decrypt; AES encrypt {SM['median_encrypt_ms']:.2f} ms", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.1, v, l, c, 26, lsize=11)
    rows = [["Cover", "Pixels", "Capacity (KiB)", "PSNR in the grid (dB)", "Min SSIM"]]
    for c in EV["capacity"]:
        rr = [r for r in qs if r["cover"] == c["cover"]]
        rows.append([c["cover"], f"{c['width']}x{c['height']}", f"{c['max_bundle_kib']}", f"{min(r['psnr_db'] for r in rr):.1f} - {max(r['psnr_db'] for r in rr):.1f}", f"{min(r['ssim'] for r in rr):.4f}"])
    ftable(s, 0.55, 3.05, [3.0, 1.15, 1.25, 1.85, 1.0], rows, size=10.5, row_h=0.4, align=["l", "c", "c", "c", "c"])
    K.picture(s, F.fig("f13_psnr_by_payload.png"), 9.0, 3.05, 3.8, 2.85)
    ftext(s, 0.55, 6.05, 12.2, 0.9, f"Grid: payloads of 1, 4, 16, 64, 128 KiB with RSA-3072. The lowest PSNR ({SM['psnr_db_min']:.1f} dB) is the 600x400 cover filled to 73 %; PSNR depends on how much of the cover is used, not on which file is hidden. Median embed {SM['median_embed_ms']:.2f} ms, extract {SM['median_extract_ms']:.2f} ms.", size=12, color=K.FNAVY)
    notes[16] = "Headline numbers from ml/eval.py. Capacity is 3.0 bits per pixel on every cover. The PSNR range reflects fill ratio."

    # ---- 17 robustness matrix
    s = S[16]
    set_title(s, "Robustness Matrix – What Survives")
    wipe(s)
    K.picture(s, F.fig("f06_robustness_matrix.png"), 0.55, 1.75, 7.6, 5.2)
    surv = [r["label"] for r in EV["robustness"] if r["recovered"] == 6] + [r["attack"] for r in extra if r["recovered"] == r["of"]]
    ftext(s, 8.35, 1.85, 4.5, 5.1, [
        (f"{n_att} attacks on 6 covers (accounts.csv payload); each ends in a real decrypt and a hash check.", {"bullet": True, "space_after": 8}),
        (f"Survive ({n_core_surv + n_extra_surv} of {n_att}, 6 of 6 covers): " + "; ".join(surv) + ".", {"bullet": True, "space_after": 8}),
        (f"Lost: JPEG at q95, q75, q50 and even q100; lossy WebP; Gaussian noise sigma 1 and 3; resize 90 %; crop left 10 %; grayscale; rotate; flip; 256-colour palette. LSB error rate about 50 %.", {"bullet": True, "space_after": 8}),
        ("Crop bottom 10 % survives only because bits run from the top row downwards and the payload still fits in the remaining rows.", {"bullet": True, "space_after": 8}),
        ("Consequence: send the stego image as a PNG file; a chat app or social site that re-encodes it destroys the message.", {"bullet": True}),
    ], size=13)
    notes[17] = "This is the confusion-matrix equivalent: attack by outcome, counted over six covers. There is no error correction, so one flipped LSB region is fatal."

    # ---- 18 parameter analysis
    s = S[17]
    set_title(s, "Parameter Analysis – Fill Ratio and RSA Key Size")
    wipe(s)
    K.picture(s, F.fig("f01_psnr_vs_fill.png"), 0.55, 1.75, 4.2, 3.0)
    K.picture(s, F.fig("f02_ssim_vs_fill.png"), 4.85, 1.75, 4.2, 3.0)
    K.picture(s, F.fig("f05_rsa_key_size.png"), 9.15, 1.75, 3.7, 3.0)
    rows = [["RSA bits", "Keygen (ms)", "Wrap (ms)", "Unwrap + decrypt (ms)", "Wrapped key", "Overhead"]]
    for b in (2048, 3072, 4096):
        r = rsa[b]
        rows.append([str(b), f"{r['keygen_ms_median']:.0f}", f"{r['encrypt_wrap_ms_median']:.2f}", f"{r['unwrap_decrypt_ms_median']:.0f}", f"{r['wrapped_key_bytes']} B", f"{r['overhead_bytes']} B"])
    ftable(s, 0.55, 4.95, [1.0, 1.2, 1.1, 1.9, 1.2, 1.1], rows, size=11, row_h=0.4, align=["c"] * 6)
    ftext(s, 8.3, 4.9, 4.5, 2.1, [
        (f"6 covers x 7 fill levels x 5 seeds: measured PSNR within {fs['max_abs_dev_vs_theory_db']:.3f} dB of the closed form; seed-to-seed std at most {fs['max_psnr_std_db']:.2f} dB.", {"bullet": True, "space_after": 5}),
        (f"Full cover: PSNR {floor_lo:.2f}-{floor_hi:.2f} dB, SSIM at least {ssim_min_fill:.4f}.", {"bullet": True, "space_after": 5}),
        ("Larger RSA keys add bytes and unwrap time, not image distortion beyond the bit count.", {"bullet": True}),
    ], size=12)
    notes[18] = "The parameters: how much of the cover is used (drives PSNR) and the RSA key size (drives time and overhead). Key generation times vary widely between runs."

    # ---- 19 steganalysis
    s = S[18]
    set_title(s, "Steganalysis – Hidden from the Eye, Not from Statistics")
    wipe(s)
    K.picture(s, F.fig("f08_steganalysis.png"), 0.55, 1.75, 7.4, 5.2)
    ftext(s, 8.15, 1.85, 4.7, 5.1, [
        (f"Chi-square (Westfeld-Pfitzmann) on the embedded prefix: stego p-value above 0.5 in {chi_n} of {chi_n} cases (lowest {chi_stego_min:.2f}); the clean cover prefix is below 0.5 in {chi_cover_lo} of them.", {"bullet": True, "space_after": 8}),
        ("In 3 cases (cat 10 % and 25 %, small coffee 10 %) the clean cover already looks equalised, so the test cannot separate them. It is an oracle test: it knows the used length.", {"bullet": True, "space_after": 8}),
        (f"Weighted-stego (blind) estimates the embedded fraction well: true 0.5 gave 0.49-0.62. Clean covers read {ws_base[0]:.2f} to {ws_base[-1]:.2f} (natural-image bias).", {"bullet": True, "space_after": 8}),
        ("Encryption does not cost quality (about 50 % of embedded bits flip a value either way) and makes the stream balanced, but sequential, key-independent embedding stays detectable at medium and high fill.", {"bullet": True}),
    ], size=12.5)
    notes[19] = "The key honest finding. PSNR above 50 dB means invisible to the eye; it does not mean undetectable. Low fill is safer, not safe. Stronger detectors were not run."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{lat('send (21 KiB memo, 512x512 cover, RSA-3072)'):.0f} ms", "API send, 512x512 cover (median)", K.FBLUE),
                                   (f"{lat('send (21 KiB memo, 1280x720 cover)') / 1000:.1f} s", "API send, 1280x720 cover", K.GREEN),
                                   (f"{lat('receive (RSA-3072)'):.0f} ms", "API receive, RSA-3072", K.ORANGE),
                                   (f"{tam['trials'] - tam['accepted_with_wrong_data']} / {tam['trials']}", "tampered bundles rejected", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f04_stage_timings.png"), 0.55, 2.85, 6.1, 2.35)
    K.picture(s, F.fig("f10_api_latency.png"), 6.8, 2.85, 6.0, 2.35)
    rows = [["Check", "Result"],
            [f"pytest (crypto {F.PYTEST_CRYPTO}, LSB {F.PYTEST_LSB}, metrics {F.PYTEST_METRICS}, API {F.PYTEST_API})", f"{F.PYTEST_PASSED} of {F.PYTEST_TOTAL} passed; smoke test {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} checks"],
            ["Cost split", f"AES encrypt {SM['median_encrypt_ms']:.2f} ms is negligible; PSNR/SSIM and PNG encoding dominate send time"],
            ["Not tested", "load or concurrent users, camera photographs, external security audit, other machines"]]
    ftable(s, 0.55, 5.3, [4.7, 7.5], rows, size=10.5, row_h=0.34)
    notes[20] = "Latencies measured on 127.0.0.1 on an i5-12500, CPU only, on the same PC that ran other jobs, so milliseconds vary by roughly 10-30 % between runs."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper and Supporting Work")
    wipe(s)
    rows = [["Aspect", "ProSt [1] (base)", "SCReedSolo [3]", "CipherVeil"],
            ["Goal", "Compute on hidden bits at an untrusted cloud", "Robust LSB hiding of text, audio, image", "Send a file securely and invisibly"],
            ["Crypto", "None for the data; secret position as key", "Fernet + Reed-Solomon, shuffling", "AES-256-GCM + RSA-OAEP-SHA256"],
            ["Reported quality", "No PSNR, SSIM or capacity", "PSNR 51.8 / 57.7 / 52.9 dB; SSIM 0.996 / 0.936 / 0.996; 3 bpp", f"PSNR {SM['psnr_db_min']:.1f}-74.6 dB (up to 73 % fill), {floor_lo:.1f} dB full; SSIM >= {ssim_min_fill:.4f}; 3.0 bpp"],
            ["Reported time", "8.38 s: 17-gate circuit, 128x128x3", "not compared here", f"embed {SM['median_embed_ms']:.2f} ms, unwrap {SM['median_decrypt_ms']:.0f} ms (RSA-3072)"],
            ["Robustness", "Not evaluated", "Reed-Solomon survives bit corruption", f"No error correction; {n_core_surv + n_extra_surv} of {n_att} attacks survive (lossless only)"]]
    ftable(s, 0.55, 1.8, [1.5, 2.9, 3.4, 4.4], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 5.75, 12.2, 1.2, [
        ("Different goals, covers and payloads: the figures are indicative, not a head-to-head. SCReedSolo hides far larger payloads in far larger covers. No superiority over either paper is claimed.", {"bullet": True, "space_after": 5}),
        ("CipherVeil takes the LSB carrier and key-exchange framing from ProSt and adds authenticated hybrid encryption plus the quality, robustness and detectability measurements ProSt does not report.", {"bullet": True}),
    ], size=13)
    notes[21] = "State plainly that the datasets and goals differ. SCReedSolo's figures are for much larger covers and payloads. ProSt reports no PSNR or SSIM, so no quality comparison with it exists."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Send and Analysis")
    wipe(s)
    K.picture(s, crop_shot("06_send_result.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("08_analysis.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: the 21 KiB memo hidden in the astronaut photo: PSNR 57.64 dB, SSIM 0.9991, 22.4 % of the cover used, SHA-256 of the file. Right: analysis of the same message: MSE 0.112, largest change 1 of 255, 11.2 % of values changed, cover and stego side by side.", size=13, color=K.FNAVY)
    notes[22] = "Real captures of the running console. The change is invisible to the eye; the statistics tell the truth, as the steganalysis slide shows."
    s = S[22]
    set_title(s, "Output – Receive and Wrong Key")
    wipe(s)
    K.picture(s, crop_shot("11_receive_result.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("12_receive_wrong_key.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: the right key restores the memo byte for byte; the AES-GCM tag is verified and both SHA-256 values are identical. Right: another keypair of the same size fails at the RSA unwrap and nothing is released.", size=13, color=K.FNAVY)
    notes[23] = "Success and failure of the receive path."
    s = S[23]
    set_title(s, "Output – Robustness and Operation Log")
    wipe(s)
    K.picture(s, crop_shot("13_robustness_results.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("17_log.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: nine attacks on one stego image: the payload survives 2 of 9, the rest give an LSB error rate near 50 %. Right: the operation log: keys, sends (one blocked by the capacity check), receives (one failed), robustness and benchmark, with no key material.", size=13, color=K.FNAVY)
    notes[24] = "Robustness page and operation log from the same session."

    # ---- 25 contributions + conclusion
    s = S[24]
    set_title(s, "Contributions, Conclusion & Objectives")
    wipe(s)
    ftext(s, 0.55, 1.75, 6.3, 5.2, [
        ("Contributions", {"bold": True, "color": K.FNAVY, "size": 15, "space_after": 5}),
        ("1. Authenticated CVL1 bundle: name, size and SHA-256 inside the AES-GCM layer, magic as associated data.", {"space_after": 4}),
        (f"2. Hybrid AES-256-GCM + RSA-OAEP-SHA256 with a fresh key and nonce per message: byte-exact in {SM['runs']}/{SM['runs']} runs, {tam['trials']}/{tam['trials']} tampered bundles rejected.", {"space_after": 4}),
        (f"3. Capacity-aware 1-bit LSB embedder, 3.0 bpp; PSNR model within {fs['max_abs_dev_vs_theory_db']:.3f} dB of measurement; PSNR >= {floor_lo:.1f} dB at any fill.", {"space_after": 4}),
        (f"4. Honest evaluation: only lossless handling survives ({n_core_surv + n_extra_surv} of {n_att} attacks); chi-square flags {chi_n}/{chi_n}; weighted-stego tracks the fill.", {"space_after": 4}),
        (f"5. Full-stack tool (FastAPI + React) with key management, analysis, robustness and a secret-free log; {F.PYTEST_PASSED} tests + {F.SMOKE_PASSED} smoke checks.", {"space_after": 8}),
        ("Novelty is integration and evaluation, not new primitives: AES, RSA-OAEP and LSB are standard.", {"italic": True, "color": K.FNAVY}),
    ], size=13.5)
    rows = [["#", "Objective", "Outcome"],
            ["1", "AES for data, RSA for key", "Met: AES-256-GCM + RSA-OAEP"],
            ["2", "Key generation and exchange", "Met for RSA 2048-4096, PEM import/export; no identity check"],
            ["3", "LSB embedding", "Met: 3.0 bpp, capacity check"],
            ["4", "Extract and decrypt", "Met: byte-exact, tag and hash"],
            ["5", "PSNR and SSIM", f">= {floor_lo:.1f} dB at any fill; SSIM >= {ssim_min_fill:.3f}"],
            ["6", "Capacity, time, robustness", "Measured; lossless-only, detectable"]]
    ftable(s, 7.0, 1.8, [0.4, 2.3, 3.1], rows, size=11, row_h=0.62, align=["c", "l", "l"])
    notes[25] = "Five contributions, then the six objectives against outcomes. Objectives 1-5 are met; objective 6 is met as a measurement, and the measurement is unflattering on robustness and detectability."

    # ---- 26 limitations and future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 5.1, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("Data survives only lossless PNG handling; no error correction", {"bullet": True, "space_after": 7}),
        ("Sequential, key-independent embedding: detected by chi-square and weighted-stego at medium and high fill", {"bullet": True, "space_after": 7}),
        ("Private keys unencrypted in the database; placeholder JWT secret; no login rate limit", {"bullet": True, "space_after": 7}),
        ("RSA only, no sender signature: the receiver cannot tell who sent it", {"bullet": True, "space_after": 7}),
        ("Six covers, four small payloads, one machine; no camera photographs", {"bullet": True, "space_after": 7}),
        ("ProSt's homomorphic computation and provable security are out of scope; Send page crashes for accounts with no keys", {"bullet": True}),
    ], size=14)
    fbox(s, 6.8, 1.8, 6.0, 5.1, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Key-dependent pixel permutation, LSB matching or adaptive (edge) embedding", {"bullet": True, "space_after": 7}),
        ("Reed-Solomon coding for resilience to light corruption", {"bullet": True, "space_after": 7}),
        ("Sender signatures (RSA-PSS or Ed25519) and an encrypted key store", {"bullet": True, "space_after": 7}),
        ("Post-quantum key encapsulation in place of RSA", {"bullet": True, "space_after": 7}),
        ("Test against stronger steganalysis (RS, sample-pair, learned detectors)", {"bullet": True, "space_after": 7}),
        ("Audio and video carriers; rate limiting; a formal usability study", {"bullet": True}),
    ], size=14)
    notes[26] = "Limitations first, then concrete next steps. The first future item answers the steganalysis finding directly."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [R[k] for k in ("base", "memr", "screed", "fridrich", "chan", "westfeld", "rs", "ws", "aes", "gcm", "rfc8017", "ssim")]
    short = [f"[{n}] {clean_ref(REFS[n - 1])}" for n in ids]
    short = [t if len(t) < 330 else t[:327] + "..." for t in short]
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[:6]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[6:]] + [(f"The full numbered list of {len(REFS)} references, including the four batch survey items [4]-[7] (title level only), is in the report.", {"italic": True})], size=11)
    notes[27] = "Selected references; the complete numbered list is in the project report."

    # ---- 28 thank you
    for sh_ in S[27].shapes:
        if sh_.name == "Text 2":
            K.replace_text(sh_, "Questions & Suggestions")
    notes[28] = "Invite questions. Likely ones: why LSB, how detectable it is, what survives, how this differs from the base paper."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


def wait_for_assets(minutes=25):
    need = [F.diag(n) for n in NEEDED] + [F.fig(n) for n in NEEDED_FIGS]
    t0 = time.time()
    while True:
        missing = [os.path.basename(p) for p in need if not os.path.exists(p)]
        if not missing:
            return
        if time.time() - t0 > minutes * 60:
            raise SystemExit(f"missing after {minutes} min: {missing}")
        time.sleep(20)


if __name__ == "__main__":
    wait_for_assets(25 if "--nowait" not in sys.argv else 0)
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
