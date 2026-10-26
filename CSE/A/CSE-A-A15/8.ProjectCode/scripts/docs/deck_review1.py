"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "hybrid encryption, AES, RSA, steganography, LSB, data security, information hiding"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


# ---------------------------------------------------------------- slide text (batch documents + verified facts only)
PROBLEM_STATEMENT = [
    "Encryption alone signals that sensitive data is being transmitted, inviting attack",
    "Steganography alone offers no protection once the hidden payload is discovered",
    "Symmetric encryption is fast, but its key exchange remains the weak point",
    "Naive LSB embedding degrades the cover image and is detectable by steganalysis",
    "Need: a two-layer scheme that encrypts data with hybrid AES-RSA and hides the result imperceptibly inside a cover image",
]
OBJECTIVES = [
    "Implement AES encryption for the secret payload and RSA encryption for the AES key",
    "Build a secure key generation and exchange mechanism between sender and receiver",
    "Embed the encrypted payload into a cover image using LSB steganography",
    "Implement the extraction and decryption path to recover the original data",
    "Measure imperceptibility using PSNR and SSIM against the original cover image",
    "Evaluate embedding capacity, processing time and robustness of the combined scheme",
]

# [work, year, approach, reported result / limitation]  (cells kept short; only what the papers themselves state)
LIT_ROWS = [
    [f"Ahmad and Rass [{R['base']}] (base)", "2026", "Secret bits hidden by LSB at a secret position among random decoy bits; an untrusted cloud computes on them with Fredkin gates",
     "Adversary advantage of order 1/n; 8.38 s for 17 gates on a 128x128x3 image; no PSNR, SSIM or capacity reported"],
    [f"Fatemieh et al. [{R['memr']}]", "2026", "Trivium and Grain-128a stream ciphers in memristive IMPLY logic, tried in an LSB steganography application",
     "Up to 42 % and 44 % fewer steps/energy; hardware level, no end-to-end software system"],
    [f"Raiyan and Kabir [{R['screed']}]", "2025", "Random shuffle + Fernet symmetric cipher + Reed-Solomon codes, then LSB embedding at 3 bits per pixel",
     "PSNR 51.774 / 57.653 / 52.888 dB on three cases; symmetric key shared in advance"],
    [f"Triple-layer visual cryptography [{R['s_triple']}]", "2024", "Image data hiding in three layers using visual cryptography (title only)", "Full text not accessed; no figures quoted"],
    [f"Data hiding using image object [{R['s_object']}]", "2024", "Gurukul Int. Multidisciplinary Research Journal (title only)", "Full text not accessed; no figures quoted"],
    [f"Chaotic encryption and CNN [{R['s_chaos']}]", "2024", "7th ICMRA: hybrid steganography with chaotic encryption and a CNN (title only)", "Full text not accessed; no figures quoted"],
    [f"AES with hybrid DWT-DCT [{R['s_dwt']}]", "2025", "Int. J. Advanced Research: AES encryption with a DWT-DCT transform-domain embedding (title only)", "Full text not accessed; no figures quoted"],
]

GAPS = [
    "The surveyed work advances steganography and encryption independently, including provably secure constructions; a practical hybrid AES-RSA plus LSB pipeline evaluated end to end on capacity, imperceptibility and speed is missing",
    f"The base paper targets computation on hidden bits by an untrusted party; it reports no PSNR, SSIM or capacity for a sender-to-receiver file transfer [{R['base']}]",
    f"Closest practical work uses a symmetric cipher with a pre-shared secret [{R['screed']}]: the key-exchange weakness named in the problem statement stays open",
    f"Visual quality is usually reported alone; robustness to re-encoding and detectability by statistics [{R['westfeld']}], [{R['ws']}] are rarely shown next to it",
    f"Lightweight-cipher steganography is studied at hardware level [{R['memr']}]; a working software tool with key handling and measurements is not provided",
]


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: CipherVeil protects a confidential file twice - it is encrypted with AES-256-GCM, its key is wrapped with the receiver's RSA key, "
                "and the result is hidden in the least significant bits of an ordinary-looking PNG - and the tool measures how well that works.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design details and measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: encrypt it, then hide it")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Sensitive data crosses the Internet all the time. Encryption protects the content but shows that a secret exists; hiding alone protects nothing once found; symmetric keys are hard to share.", K.ORANGE, "!"),
        ("Our idea", "A double layer: AES encrypts the data efficiently, RSA protects the AES key, and the encrypted bundle is hidden in the least significant bits of an image. The receiver extracts, decrypts and recovers the original.", K.BLUE, "i"),
        ("The deliverable", "CipherVeil: key management, send and receive screens, PSNR and SSIM analysis, capacity and speed benchmark, robustness tests and an operation log, built as a FastAPI + React tool.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: the weakness of encryption-only and hiding-only protection, the idea of the two layers, and what is delivered. Everything runs locally on a CPU; there is no dataset and no GPU."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "LSB steganography is used as a computing medium. A secret bit sits at a secret position among random decoy LSBs; the cloud applies a Fredkin gate pixel-wise to every bit and never learns which bits are real. Extraction needs the secret position.", K.ORANGE, "1"),
        ("Reported results", "Security: adversary advantage of order 1/n in the number n of positions; n of 2^20 (about 10^6) is recommended for advantage below 10^-6. Runtime: 8.38 s for a 17-gate, 61-wire circuit on a 128x128x3 image (i5-1135G7, 16 GB). No PSNR, SSIM or capacity is reported.", K.BLUE, "2"),
        ("Relevance and limits", "We adopt the LSB carrier and the idea of a shared secret exchanged by standard cryptography. The paper itself notes that circuit privacy is only heuristic and that scaling is a challenge. It does not move a file from sender to receiver, so its numbers are not like-for-like with ours.", K.GREEN, "3"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=14)
    notes[4] = ("Read the citation. ProSt is a provably secure steganography scheme for outsourcing computation: the data are hidden among decoys and processed with a universal reversible gate. "
                "CipherVeil does not compute on hidden bits; it uses the same carrier (image LSBs) to deliver an AES-GCM-encrypted file. The 8.38 s figure is for a different task and must not be compared with our millisecond timings.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    rows += LIT_ROWS
    K.table(s, 0.6, 1.25, [2.75, 0.65, 4.7, 4.03], rows, font_size=11, row_h=0.66, header_size=12, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey is the base paper plus four recent works; we also read two further papers from the batch folder. Figures are quoted only where we read them in the paper itself. "
                "For the four survey titles only the title, year and source are available, so no results are quoted.")

    # ---- 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 6.3, 4.7, ["Text into\n14589x8423", "Audio into\n21188x30328", "Image into\n512x512"],
                     [("PSNR reported for SCReedSolo (dB)", [51.774, 57.653, 52.888])], colors=[K.BLUE], number_format="0.0", font_size=12, max_value=70)
    K.text(s, 0.6, 6.2, 6.3, 0.8, f"Reported in [{R['screed']}] for three cases at 3 bits per pixel; covers and payloads differ from ours, so this is context, not a target.", size=12, color=K.TXT, italic=True)
    K.text(s, 7.2, 1.45, 5.55, 5.5, [
        ("**What the literature and standards give us**", {"space_after": 8}),
        (f"LSB substitution is simple, and the change per value is at most 1, so images stay visually intact [{R['provos']}], [{R['chan']}].", {"bullet": True, "space_after": 8}),
        (f"AES (FIPS 197 [{R['aes']}]) in Galois/Counter Mode [{R['gcm']}] gives confidentiality and an integrity tag.", {"bullet": True, "space_after": 8}),
        (f"RSA [{R['rsa']}] with OAEP padding [{R['rfc8017']}] wraps the session key for the receiver.", {"bullet": True, "space_after": 8}),
        (f"PSNR and SSIM [{R['ssim']}] are the usual imperceptibility measures; chi-square and weighted-stego analysis [{R['westfeld']}], [{R['ws']}] show that LSB changes can still be detected.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("The chart shows what a closely related hybrid scheme reported, so we know the PSNR range to expect when the whole capacity is used. The right-hand side lists the standards and steganalysis papers "
                "that our design and evaluation follow. Our own measurements start in Review 2.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Sensitive file", K.NAVY), ("Encrypt (or not)", K.NAVY), ("Send over the network", K.NAVY), ("Receiver decrypts", K.NAVY)], size=15, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Encryption only, or hiding only",
           "Encryption alone makes the transfer look like what it is: a secret on the move.\nHiding alone gives no protection once the payload is found.\nBoth are used separately in most tools.", K.BLUE, "1", body_size=18)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Symmetric keys and plain LSB",
           "Symmetric encryption is fast, but sharing its key over the same channel is the weak point.\nNaive LSB embedding with no care for capacity or measurement can degrade the cover.", K.ORANGE, "2", body_size=18)
    notes[7] = "Today a confidential file is usually protected by encryption alone, or an image is used to hide plain data. The two protections and the key-sharing step are handled separately."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Visible secrecy", "Single layer", "Key exchange", "Detectable LSB"]
    bodies = [
        "Encryption alone signals that sensitive data is being transmitted, which invites attack.",
        "Steganography alone offers no protection once the hidden payload is discovered.",
        "Symmetric encryption is fast but its key exchange remains the weak point; a shared secret must travel first.",
        "Naive LSB embedding degrades the cover image and is detectable by steganalysis, especially when much of the capacity is used.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=16)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"The base paper adds: circuit privacy is only heuristic, cost grows with image size, and no imperceptibility or capacity figures are given [{R['base']}].", size=14, color=K.INK, anchor="m")
    notes[8] = "The four limitations are the batch problem statement. The last line is taken from the base paper's own discussion of scalability and circuit privacy."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=15, color=K.INK, anchor="m")
    notes[9] = ("Gap 1 is the batch survey's own statement. Gaps 2, 3 and 5 follow from the papers we read. Gap 4 is our reading: visual quality is reported by almost everyone, "
                "but robustness and detectability are seldom shown next to it, so we measure all of them.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "A confidential file, the receiver's RSA public key and a PNG or other RGB cover image"),
                                 ("Output", "A stego PNG that looks like the cover; at the receiver the original file, with matching SHA-256 hashes"),
                                 ("Measures", "PSNR, SSIM, bits per pixel of capacity, time per stage and which attacks the payload survives")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=13, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a two-layer scheme: hybrid AES-RSA encryption, then imperceptible hiding in a cover image."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "No dataset is needed: evaluation uses six public-domain or generated cover images and four synthetic payload files.", size=14, color=K.TXT, italic=True)
    notes[11] = "The six objectives from the proposal. Later reviews report which were met and where the results are weaker than hoped."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Encrypt", "AES-256-GCM with a fresh key and nonce for every message; a small header carries the file name, size and SHA-256.", K.BLUE, "1"),
        ("Wrap the key", "The AES key is wrapped with the receiver's RSA public key using OAEP (SHA-256); keys of 2048, 3072 or 4096 bits.", K.GREEN, "2"),
        ("Hide", "The bundle is written into the lowest bit of each R, G, B value after a capacity check; the output is a PNG.", K.ORANGE, "3"),
        ("Recover and verify", "Extract, unwrap, check the GCM tag and compare hashes: the file comes back byte for byte or not at all.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React console: keys, send, receive, analysis, robustness and log screens", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "PSNR, SSIM, histograms and bit planes show how visible the change is; nine attack tests show what the hidden data survives. No GPU is needed.", size=14, color=K.TXT)
    notes[12] = ("Four stages and the console around them. Encryption and key wrapping use the standard cryptography library; the LSB embedder is our own code. "
                 "The console also measures: it is how we can say how invisible, how large and how robust the hiding is.")

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    pic = F.diag("01_system_architecture.png")
    if os.path.exists(pic):
        K.picture(s, pic, 0.6, 1.2, 7.5, 5.9)
    else:
        print("WARNING: diagram 01 missing")
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite console (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}) with token sign-in and input checks.", {"bullet": True}),
        ("Services: AES-GCM and RSA-OAEP, LSB embedder, PSNR and SSIM, robustness attacks and benchmark.", {"bullet": True}),
        ("State: SQLite for users, keypairs and the operation log, plus image files on disk.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the console calls the API, the API calls the services, and results are stored as database rows and image files. The log never records keys or file contents."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working double layer", "Byte-exact recovery", "Imperceptible hiding", "Capacity and speed", "Honest robustness"]
    bodies = [
        "A tool that encrypts, hides, extracts and decrypts a file, with key management for sender and receiver.",
        "Original file restored exactly, and any tampering or wrong key rejected instead of releasing wrong data.",
        "PSNR above 50 dB and SSIM near 1 against the cover; histograms and bit planes as evidence.",
        "Bits per pixel, bytes per cover and time per stage measured on several covers and payload sizes.",
        "Which attacks the payload survives, which destroy it, and how detectable it is by statistics; limits stated, not hidden.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=16)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=17)
    notes[14] = ("These are the outcomes we commit to. A first prototype run already round-trips files byte for byte (22 of 22 runs), with PSNR between 52.5 and 74.6 dB and 3.0 bits per pixel of capacity. "
                 "Those are prototype figures; the full evaluation, including where it is weaker than hoped, is presented in Review 2 and the final review.")

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [R[k] for k in ("base", "memr", "screed", "s_triple", "s_object", "s_chaos", "s_dwt", "fridrich", "westfeld", "ws", "aes", "gcm", "rsa", "rfc8017", "ssim")]
    K.text(s, 0.6, 1.2, 12.13, 5.9, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 2}) for n in idx], size=11.5, color=K.INK)
    notes[15] = "Selected references; the complete numbered list, with notes on which entries could only be listed by title, is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why LSB and not an adaptive method (simple, measurable, matches the base paper's carrier; adaptive embedding is future scope); is the result undetectable (no: it is invisible to the eye "
                 "but statistically detectable at higher fill, and we measure this); how it differs from the base paper (ProSt computes on hidden bits, CipherVeil delivers a file).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
