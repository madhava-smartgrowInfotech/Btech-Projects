"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the supporting papers, the maintainers of the open-source software used in this project "
        "(Python, FastAPI, SQLAlchemy, the `cryptography` library, Pillow, NumPy, scikit-image, matplotlib, React, Vite, Tailwind CSS, Recharts, pytest and Playwright), "
        "and the providers of the public-domain and CC0 sample images shipped with scikit-image.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    S = F.S
    fs = F.EX["fill_study"]
    d.p("Encryption protects the content of a message but shows that a secret is moving; hiding conceals the message but protects nothing once it is found. "
        f"This project builds {F.PRODUCT}, which combines both: a file is encrypted with a fresh AES-256-GCM key, the key is wrapped with the receiver's RSA public key "
        "(OAEP, SHA-256), and the resulting bundle is hidden in the least significant bits of a PNG image. The receiver extracts the bundle, unwraps the key, verifies the "
        "authentication tag and recovers the file byte for byte. A FastAPI and React application provides key management, analysis, robustness and benchmark views.")
    d.p(f"On six cover images the file was restored exactly in {S['runs']} of {S['runs']} runs, with PSNR from {S['psnr_db_min']:.1f} to 74.6 dB, SSIM of at least {S['ssim_min']:.4f} and a capacity of "
        f"{S['max_capacity_bits_per_pixel']:.1f} bits per pixel; a closed-form model predicts PSNR within {fs['max_abs_dev_vs_theory_db']} dB, and all 240 single-bit tamper trials were rejected. "
        "The payload survives only lossless handling (5 of 18 attacks), and chi-square steganalysis flagged all 30 embedded prefixes, so the hiding is imperceptible but not undetectable. "
        f"{F.PYTEST_PASSED} tests pass. The base paper's homomorphic computation and provable security are not implemented.")
    d.p("**Keywords:** hybrid encryption, AES-GCM, RSA-OAEP, steganography, LSB, data security, information hiding, PSNR, SSIM, secure communication.")


ABBR = [
    ("AES", "Advanced Encryption Standard"), ("API", "Application Programming Interface"), ("bpp", "Bits Per Pixel"), ("CSV", "Comma-Separated Values"),
    ("DCT", "Discrete Cosine Transform"), ("DFD", "Data Flow Diagram"), ("DWT", "Discrete Wavelet Transform"), ("ER", "Entity-Relationship"),
    ("GCM", "Galois/Counter Mode"), ("HTTP", "Hypertext Transfer Protocol"), ("JPEG", "Joint Photographic Experts Group (lossy image format)"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("KEM", "Key Encapsulation Mechanism"), ("LSB", "Least Significant Bit"),
    ("MSE", "Mean Squared Error"), ("OAEP", "Optimal Asymmetric Encryption Padding"), ("PEM", "Privacy-Enhanced Mail (text format for keys)"),
    ("PNG", "Portable Network Graphics"), ("PSNR", "Peak Signal-to-Noise Ratio"), ("REST", "Representational State Transfer"), ("RGB", "Red, Green, Blue"),
    ("RS", "Reed-Solomon (error-correcting code)"), ("RSA", "Rivest-Shamir-Adleman"), ("SHA", "Secure Hash Algorithm"), ("SQL", "Structured Query Language"),
    ("SSIM", "Structural Similarity Index Measure"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("WS", "Weighted Stego (steganalysis estimator)"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
