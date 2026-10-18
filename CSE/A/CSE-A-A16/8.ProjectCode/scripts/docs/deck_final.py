"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .lit import LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "sign language recognition, CNN, ASL"
    notes = {}
    T, DS, LAT = F.T, F.DS, F.LAT

    # footers on every slide except the title
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE, 28)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: SignSpeak, a CNN-based recogniser for the ASL alphabet with a complete web application."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Model\nDecision Workflow, Interaction & Storage\nImplementation Highlights\nResults, Confusion Matrix & Threshold Analysis\nPerformance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [
        ("Sign language is the main way many deaf and hard-of-hearing people communicate, but few hearing people understand it, and interpreters are scarce.", {"bullet": True, "space_after": 10}),
        ("SignSpeak recognises the ASL alphabet from a webcam: the hand region is cropped, segmented (grayscale, blur, Otsu), resized to 64 x 64 and classified by a CNN into 29 classes; recognised letters build a sentence that is shown and spoken.", {"bullet": True, "space_after": 10}),
        (f"The CNN ({F.MODEL['parameters']:,} parameters, CPU-trained on 7,250 images) reaches **{F.pct(T['accuracy'], 2)}** accuracy on a held-out 1,088-image split (macro F1 {F.f3(T['macro_f1'])}); on the server a frame takes about {LAT['server_side_total_per_frame']['median_ms']:.0f} ms, so the live loop runs at roughly 17 to 19 frames per second.", {"bullet": True, "space_after": 10}),
        ("A FastAPI service and a React interface add collection, training and a model report; 52 pytest tests and a 26-check smoke test pass.", {"bullet": True, "space_after": 10}),
        ("Not done: words and motion signs, and evaluation on unseen signers; accuracy drops sharply when the hand is rotated by 10 degrees.", {"bullet": True}),
    ], size=18)
    notes[3] = "Read the first two bullets, then give the headline number with its protocol (held-out split). End on the honest limits."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper (MAG-KD)"],
            ["Method", "Skeleton-based GCN student (STA-GCN) distilled from three teachers (SL-GCN, CTR-GCN, TD-GCN); hard correctness mask plus adaptive gating network"],
            ["Data", "AUTSL (38,336 videos, 43 signers, 226 signs) and WLASL-100 (1,442 train / 596 validation videos); seeds 1, 42, 100"],
            ["Reported results", "94.25 +/- 0.19 % Top-1 on AUTSL with 0.48 M parameters; 66.33 +/- 0.54 % on WLASL-100 (baseline 54.92 %); Jetson Nano 33.13 FPS at 6.23 W (TensorRT FP16)"],
            ["Stated limits", "Mask needs ground-truth labels; accuracy bounded by the upstream pose extractor in difficult scenes"]]
    ftable(s, 0.55, 2.9, [2.0, 10.2], rows, size=14, row_h=0.78)
    notes[4] = "A skeleton and distillation method for video word signs on an edge device. We share the goal of accessible real-time recognition but not the method or the task."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"]]
    for i in (0, 1, 2, 3, 4, 7):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", ""), r[1], r[2], r[3], r[4]])
    ftable(s, 0.55, 1.85, [1.9, 2.7, 2.2, 3.0, 2.4], rows, size=10, row_h=0.65)
    notes[5] = "Representative work. Where a work's figures were not checked they are not quoted. Four further 2024-2025 papers named in the project survey are in the report's reference list with that caveat."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, [
        ("**End-to-end system:** few works go from a live webcam to letters, a sentence and speech in one usable application.", {"bullet": True, "space_after": 12}),
        ("**Accessible hardware:** skeleton and depth pipelines need extra models or sensors; an ordinary webcam and CPU should suffice for letters.", {"bullet": True, "space_after": 12}),
        ("**Honest evaluation:** results are rarely given with split, per-class scores, confidence behaviour and latency together.", {"bullet": True, "space_after": 12}),
        ("**Robustness:** effects of clutter, lighting and hand orientation are seldom measured.", {"bullet": True, "space_after": 12}),
        ("**Unseen signers:** public alphabet data has no signer identity, so signer-independent accuracy is hard to test.", {"bullet": True}),
    ], size=19)
    notes[6] = "Five gaps. We address the first three directly, measure the fourth, and state the fifth as not done."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.95, 12.2, 2.0)
    ftext(s, 0.85, 2.0, 11.6, 1.9, PROBLEM_STATEMENT, size=19, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Webcam frames, 960 x 540, hand inside a 300 x 300 region"), ("Output", "Letter and confidence, sentence on screen, spoken text"), ("Record", "Collected images, trained model, metrics, confusion matrix")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.3, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.4, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 4.9, 3.5, 1.0, b_, size=15)
    notes[7] = "The statement has three parts: recognise, speak, and keep an inspectable record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion reports for each whether it was met; objectives 1 (words) and 6 (unseen signers) were not fully met."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 2.0, 12.2, 1.1, [("Webcam frame", K.FNAVY), ("ROI + segmentation", K.FBLUE), ("CNN classify", K.GREEN), ("Vote + stability", K.ORANGE), ("Sentence + speech", K.PURPLE)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Compact three-block CNN, 2.2 million parameters, trained from scratch on the CPU; the same preprocessing for collection, training and recognition.", {"bullet": True, "space_after": 10}),
        ("Confidence threshold 0.75, majority vote over 15 frames and 18 stable frames before a letter is typed.", {"bullet": True, "space_after": 10}),
        ("Sentence builder with offline text-to-speech (browser voice fallback); in-app collection, background training and a model report.", {"bullet": True, "space_after": 10}),
        ("FastAPI service reusing the original desktop code; React interface in light and dark themes, usable at phone width.", {"bullet": True}),
    ], size=18)
    notes[9] = "The proposed system in one line of boxes, then four properties that matter."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.85, 8.3, 5.0)
    ftext(s, 9.1, 2.0, 3.7, 4.7, [
        ("Presentation: React app owns the webcam", {"bullet": True, "space_after": 8}),
        ("Application: FastAPI, 16 endpoints, in-memory sessions", {"bullet": True, "space_after": 8}),
        ("Domain: preprocessing, CNN, recogniser, speech", {"bullet": True, "space_after": 8}),
        ("Storage: image folders and model files", {"bullet": True}),
    ], size=16)
    notes[10] = "Three tiers plus storage; the domain code is shared by the desktop and web interfaces."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Kaggle ASL Alphabet", K.FNAVY), ("250 per class", K.FNAVY), ("ROI + Otsu mask", K.PURPLE), ("Stratified split", K.GREEN), ("Train CNN", K.ORANGE), ("Test once", K.FBLUE), ("metrics.json", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("One preprocessing function for collection, training and live recognition: no train/serve mismatch.", {"bullet": True, "space_after": 10}),
        ("Stratified 70 / 15 / 15 split with seed 42; test part used once; augmentation only in training.", {"bullet": True, "space_after": 10}),
        ("Extra measurements: confidence threshold sweep, robustness to rotation, brightness, noise and blur, latency per stage.", {"bullet": True, "space_after": 10}),
        ("Every number is written to docs/metrics/metrics.json by one script and re-derived from the saved model.", {"bullet": True}),
    ], size=18)
    notes[11] = "Workflow and measurements. The test accuracy was independently reproduced from the saved model and the same seed."

    # ---- 12 dataset, features, model
    s = S[11]
    set_title(s, "Dataset, Features & Model")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, "7,250", "images, 250 x 29 classes (of 87,000)", K.FBLUE, 26)
    ftile(s, 0.55, 3.15, 2.9, 1.1, "64 x 64 x 1", "segmented mask, scaled to [0, 1]", K.GREEN, 24)
    ftile(s, 0.55, 4.4, 2.9, 1.1, f"{F.MODEL['parameters']:,}", "CNN parameters", K.ORANGE, 24)
    K.picture(s, F.fig("f08_dataset_split.png"), 3.7, 1.85, 4.8, 2.0)
    rows = [["Layer", "Output"], ["3 x (Conv3x3 + BN + MaxPool)", "8x8x128"], ["Flatten, Dense 256, Dropout 0.4", "256"], ["Dense 29, softmax", "29 classes"]]
    ftable(s, 8.7, 1.9, [2.9, 1.2], rows, size=11, row_h=0.5)
    K.picture(s, F.fig("f01_training_curves.png"), 3.7, 4.0, 4.8, 1.8)
    ftext(s, 8.7, 4.2, 4.1, 2.6, [
        (f"Train {DS['train']:,} / validation {DS['validation']:,} / test {DS['test']:,}.", {"bullet": True}),
        ("Adam 1e-3, batch 32, 30 epochs, augmentation, early stopping on validation accuracy.", {"bullet": True}),
        ("Frames come from video: near-duplicates, no signer id.", {"bullet": True}),
    ], size=13)
    ftext(s, 0.55, 5.8, 8.0, 0.9, f"Saved run: best validation accuracy {F.pct(F.RUN['best_val_accuracy'])} at epoch {F.RUN['best_epoch']} of 30 (the cap), final training accuracy {F.pct(F.RUN['final_train_accuracy'])}.", size=12, color=K.FNAVY)
    notes[12] = "Dataset, input and model in one view. The training curves show validation above training accuracy because dropout and augmentation act only in training."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 5.0, 5.1)
    ftext(s, 5.9, 2.0, 6.9, 4.8, [
        ("Gate: top softmax probability below 0.75 counts as 'nothing'.", {"bullet": True, "space_after": 12}),
        ("Smooth: majority vote over the last 15 frame labels.", {"bullet": True, "space_after": 12}),
        ("Commit: 18 stable frames, label not 'nothing' and different from the last committed letter.", {"bullet": True, "space_after": 12}),
        ("Controls: 'space' and 'del' edit the sentence; 'nothing' resets the repeat guard.", {"bullet": True}),
    ], size=18)
    notes[13] = "The decision rules in SignRecognizer.predict_frame. The threshold behaviour is analysed on slide 19."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_recognition.png"), 0.55, 1.8, 6.6, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.4, 1.8, 5.4, 5.1)
    notes[14] = "Sequence of one recognition request and the use cases of the signer and the operator."

    # ---- 15 storage design
    s = S[14]
    set_title(s, "Database Design (file store)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.5)
    rows = [["Store", "Content"], ["data/<CLASS>/", "training images, 29 folders"], ["models/asl_cnn_model.keras", "trained network"], ["models/labels.json", "class order"],
            ["models/training_history.json", "per-epoch curves"], ["models/classification_report.txt", "test report"], ["models/confusion_matrix.json", "29 x 29 counts"], ["memory: sessions", "at most 20 recognisers"]]
    ftable(s, 7.8, 1.9, [2.9, 2.1], rows, size=11, row_h=0.5)
    ftext(s, 0.55, 6.35, 12.2, 0.5, "There is no SQL database: the ER diagram models the file-system entities. Recognition sessions are not persisted.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Honest point: the project stores files, not rows."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("16", "REST endpoints (+ /docs)", K.FBLUE), (str(F.PYTEST_TOTAL), "pytest tests", K.GREEN), ("7", "React screens", K.ORANGE), ("1", "command to start (run.bat)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Browser captures frames with getUserMedia, cover-crops to 16:9, mirrors, and posts a 960 x 540 JPEG (quality 0.8, about 27 KB); one request in flight, at least 60 ms apart.", {"bullet": True, "space_after": 9}),
        ("Validation: epochs 1-500, batch 1-1024, learning rate 1e-6 to 1, at least two classes with images, 2,000-character speech limit, image-name pattern.", {"bullet": True, "space_after": 9}),
        ("Training runs in a background thread with a stop flag and live progress; speech is offline (pyttsx3) with a browser-voice fallback.", {"bullet": True, "space_after": 9}),
        ("The original desktop app is unchanged; the API imports its modules.", {"bullet": True}),
    ], size=17)
    notes[16] = "Highlights: the capture loop, validation, background training and reuse of the original code."

    # ---- 17 results: model performance
    s = S[16]
    set_title(s, "Results – Model Performance (held-out test split)")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.pct(T["accuracy"], 2), "accuracy, n = 1,088", K.FBLUE), (F.f3(T["macro_f1"]), "macro F1", K.GREEN),
                                   (F.pct(T["top3_accuracy"]), "top-3 accuracy", K.ORANGE), (F.f4(T["loss"]), "test loss", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26)
    K.picture(s, F.fig("f02_per_class.png"), 0.55, 3.15, 7.7, 3.2)
    rows = [["Best F1", "Weakest F1"], [", ".join(f"{k} {v:.2f}" for k, v in sorted(T["per_class_f1"].items(), key=lambda kv: -kv[1])[:4]), ", ".join(f"{k} {v:.2f}" for k, v in F.sorted_f1()[:4])]]
    ftable(s, 8.5, 3.2, [2.15, 2.15], rows, size=12, row_h=0.8, first_bold=False)
    ftext(s, 8.5, 4.95, 4.3, 1.8, [
        (f"Other archive photos (n = {F.UNSEEN['n']:,}): {F.pct(F.UNSEEN['accuracy'], 2)}, optimistic: near-duplicate frames of trained images.", {"bullet": True, "size": 12}),
        (f"Official Kaggle test: {F.OFFICIAL['correct']} of {F.OFFICIAL['n']} (A read as S, I read as D).", {"bullet": True, "size": 12}),
    ], size=12)
    ftext(s, 0.55, 6.4, 7.7, 0.5, "Test split = 15 % never used for training or model selection; single run, seed 42.", size=12, color=K.FNAVY, italic=True)
    notes[17] = "Headline 82.81 %. The weakest letters are Y, V, X, W, E, U: hand shapes that differ by a finger or an angle. The 87.24 % figure is not used as the headline."

    # ---- 18 confusion matrix & per-class
    s = S[17]
    set_title(s, "Confusion Matrix & Frequent Confusions")
    wipe(s)
    K.picture(s, F.fig("f03_confusion_matrix.png"), 0.55, 1.85, 5.6, 5.0)
    K.picture(s, F.fig("f04_top_confusions.png"), 6.4, 1.85, 6.4, 3.1)
    top = ", ".join("%s read as %s (%d)" % (c["true"], c["predicted"], c["count"]) for c in F.CONFUSIONS[:4])
    ftext(s, 6.4, 5.1, 6.4, 1.7, [
        ("Most frequent: " + top + ".", {"bullet": True, "size": 14}),
        ("V, U, W and R, T, S share similar finger layouts; E is confused with B and A.", {"bullet": True, "size": 14}),
    ], size=14)
    notes[18] = "Rows are the true letter, columns the prediction. Errors concentrate on letters whose hand shapes differ by a finger position."

    # ---- 19 threshold / parameter analysis
    s = S[18]
    set_title(s, "Threshold & Robustness Analysis")
    wipe(s)
    K.picture(s, F.fig("f05_threshold.png"), 0.55, 1.85, 6.1, 3.3)
    K.picture(s, F.fig("f06_robustness.png"), 6.7, 1.85, 6.1, 3.3)
    ths = [0.5, 0.6, 0.7, 0.75, 0.8, 0.9, 0.95, 0.99]
    rows = [["Confidence threshold"] + [f"{t:.2f}" for t in ths],
            ["Frames accepted"] + [F.pct(F.thr_row(t)["coverage"]) for t in ths],
            ["Accuracy on accepted"] + [F.pct(F.thr_row(t)["accuracy_on_accepted"]) for t in ths]]
    ftable(s, 0.55, 5.3, [2.6] + [1.2] * 8, rows, size=12, row_h=0.42, align=["l"] + ["c"] * 8)
    ftext(s, 0.55, 6.62, 12.2, 0.4, "Deployed threshold 0.75: 70.3 % of frames accepted at 97.1 % accuracy. Rotation of +/- 10 degrees drops accuracy from 85.0 % to 67.6 % / 68.6 % (580-photo robustness set).", size=12, color=K.FNAVY)
    notes[19] = "Raising the confidence threshold trades coverage for accuracy. At 0.75, seven in ten frames are used and 97 % of them are right. Rotation is the weak spot."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{LAT['jpeg_decode_960x540']['median_ms']:.1f} ms", "JPEG decode", K.FBLUE), (f"{LAT['roi_crop_and_segmentation_and_normalisation']['median_ms']:.1f} ms", "crop + segmentation", K.GREEN),
                                   (f"{LAT['cnn_forward_batch1']['median_ms']:.1f} ms", "CNN forward (model.predict)", K.ORANGE), (f"{LAT['http_recognize_frame_round_trip']['median_ms']:.1f} ms", "HTTP round trip (about 17.6 fps)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.15, v, l, c, 26)
    rows = [["Test layer", "What it covers", "Result"],
            ["Unit (pytest)", "preprocessing, recogniser logic, dataset and model", f"{F.PYTEST_PRE + F.PYTEST_REC + F.PYTEST_DS} passed"],
            ["API (pytest)", "all endpoints, validation, errors, sessions", f"{F.PYTEST_API} passed"],
            ["Smoke test", "running server: collect, train, recognise, errors", f"{F.SMOKE_TOTAL} / {F.SMOKE_TOTAL} checks"],
            ["Browser walk", "7 screens, 360 / 768 / 1440 px, light and dark, fake camera", "no console or network errors"],
            ["Build", "npm run build (tsc + vite)", "passes"]]
    ftable(s, 0.55, 3.3, [2.4, 6.2, 3.6], rows, size=13, row_h=0.5)
    ftext(s, 0.55, 6.4, 12.2, 0.5, "Median of 300 frames on an Intel Core i5-12500 (CPU only, 15.7 GB). Not tested: a physical webcam, load, a user study, React component tests.", size=12, color=K.FNAVY, italic=True)
    notes[20] = "The CNN call dominates the 53 ms server time. Latency was measured on an otherwise idle desktop; the live loop is limited to about 17-19 fps."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different task and data)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "This project (SignSpeak)"],
            ["Task", "Isolated word signs from video (AUTSL 226 signs; WLASL-100)", "Static ASL fingerspelling: 26 letters + space, del, nothing from single frames"],
            ["Input", "2D skeleton sequences (27 keypoints) from a pose estimator", "One 64 x 64 segmented hand mask cropped from a webcam frame"],
            ["Model", "0.48 M-parameter GCN student distilled from three GCN teachers", f"{F.MODEL['parameters'] / 1e6:.2f} M-parameter CNN trained from scratch; no distillation"],
            ["Accuracy", "94.25 +/- 0.19 % (AUTSL), 66.33 +/- 0.54 % (WLASL-100); 3 seeds", f"{F.pct(T['accuracy'], 2)} on a held-out split (one run); not signer-independent"],
            ["Hardware / speed", "Jetson Nano, TensorRT FP16: 33.13 FPS at 6.23 W", "Desktop CPU, no optimisation: about 17.6 fps end to end; power not measured"],
            ["Deliverable", "Research method and measurements", "Application: collect, train, recognise, speak"]]
    ftable(s, 0.55, 1.85, [1.9, 5.0, 5.3], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 6.4, 12.2, 0.55, "Different task, data, model family and hardware: the numbers are context, not a ranking, and no superiority is claimed.", size=13, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that the tasks differ: video word signs versus static letters. Even the accuracy numbers measure different things."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Live Recognition and Collection")
    wipe(s)
    K.picture(s, F.shot("04_recognize.png"), 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, F.shot("05_collect.png"), 6.8, 1.85, 6.0, 4.2, border=True)
    ftext(s, 0.55, 6.2, 12.2, 0.7, "Left: recognition (letter H at 87 %, mask preview, sentence). Right: collection with the ROI box and recent captures. Camera = Chromium fake camera fed with real photos.", size=12, color=K.FNAVY)
    notes[22] = "Real screenshots. The recognition clip was assembled from real held-out photos because the build machine had no physical webcam."
    s = S[22]
    set_title(s, "Output – Training and Model Performance")
    wipe(s)
    K.picture(s, F.shot("06_train.png"), 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, F.shot("07b_model_confusion.png"), 6.8, 1.85, 6.0, 4.2, border=True)
    ftext(s, 0.55, 6.2, 12.2, 0.7, "Left: training screen (coverage 7,250 images, hyper-parameters, finished run 82.8 %). Right: model screen with the confusion matrix.", size=12, color=K.FNAVY)
    notes[23] = "Training runs in the background with live curves; the model screen shows true per-class results."
    s = S[23]
    set_title(s, "Output – Dashboard and Phone Layout")
    wipe(s)
    K.picture(s, F.shot("03_dashboard.png"), 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, F.shot("12_mobile_recognize.png"), 7.8, 1.75, 2.0, 4.3, border=True)
    K.picture(s, F.shot("11_mobile_dashboard.png"), 10.4, 1.75, 2.0, 4.3, border=True)
    ftext(s, 0.55, 6.2, 12.2, 0.7, "Left: dashboard (7,250 images, 29 of 29 classes, 82.8 %). Right: the same app at 390 x 844 px (recognition and dashboard).", size=12, color=K.FNAVY)
    notes[24] = "Dashboard with real counts and the phone-width layout."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("**Delivered:** a working ASL-alphabet recogniser: segmentation, a trained 2.2 M-parameter CNN, decision logic, sentence builder with speech, collection, training and a model report, with tests.", {"bullet": True, "space_after": 10}),
        (f"**Measured:** {F.pct(T['accuracy'], 2)} accuracy on a held-out 1,088-image split; {F.pct(F.thr_row(0.75)['accuracy_on_accepted'])} on the frames passing the 0.75 confidence gate; about 53 ms per frame on a CPU.", {"bullet": True, "space_after": 10}),
        ("**Objectives:** 2, 3, 5 met; 4 partly (17-19 fps without model optimisation); 1 and 6 not fully met (no words; no unseen-signer test).", {"bullet": True, "space_after": 10}),
        ("**Limits:** same-recording data, Otsu masks keep background, rotation sensitivity, letters only, no physical-webcam test. Not a certified communication aid.", {"bullet": True}),
    ], size=18)
    notes[25] = "Be explicit about which objectives were met and which were not."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("Collect data from many signers and backgrounds with signer ids, and report signer-independent accuracy.", {"bullet": True, "space_after": 12}),
        ("Rotation, brightness and background augmentation; compare skin-colour segmentation or hand landmarks with the Otsu mask.", {"bullet": True, "space_after": 12}),
        ("Faster inference: quantisation or TFLite export in place of model.predict; measure power as the base paper does.", {"bullet": True, "space_after": 12}),
        ("Words and motion signs: sequence models over several frames; two-hand signs and sentence-level translation.", {"bullet": True, "space_after": 12}),
        ("Packaging: installer, calibration screen for the ROI, user study with deaf signers.", {"bullet": True}),
    ], size=20)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [1, 2, 3, 4, 5, 6, 7, 11, 14, 15, 17, 19, 20, 21, 23, 25]
    ftext(s, 0.55, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in ids[:8]], size=11)
    ftext(s, 6.8, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in ids[8:]], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
