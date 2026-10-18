"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.dml.color import RGBColor

from . import facts as F
from . import pptkit as K
from .lit import LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "sign language recognition, CNN, ASL, computer vision"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 32)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    T = F.T

    # ---- 1 title
    title_slide(S[0], "Real-time ASL fingerspelling recognition: a webcam, a CNN, and the signed letters turned into text and speech",
                "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: SignSpeak watches a signer's hand through a webcam, recognises the American Sign Language "
                "alphabet with a convolutional neural network, and shows the sentence on screen and speaks it.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: from a webcam to text and speech")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Sign language is the main medium of communication for people with hearing and speech impairment, yet few hearing people understand it. Without an interpreter, everyday conversation breaks down.", K.ORANGE, "!"),
        ("Our idea", "Capture the hand through a webcam, extract the hand region, segment it, and classify the letter with a convolutional neural network trained on ASL alphabet images.", K.BLUE, "i"),
        ("The deliverable", "A working application: live recognition of 26 letters plus space and delete, a sentence builder with speech output, and in-app data collection, training and model reports.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=19)
    notes[3] = "Three points: the communication gap, the CNN-based idea, and what is delivered: a complete application, not only a model."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Skeleton-based sign recognition: a small graph convolutional network (student) learns from three larger GCN teachers. A hard correctness mask removes wrong teacher guidance; an adaptive gate weighs the rest.", K.BLUE, "M"),
        ("Data and reported results", "AUTSL (38,336 videos, 43 signers, 226 signs): 94.25 +/- 0.19 % Top-1 with 0.48 M parameters. WLASL-100: 66.33 +/- 0.54 %. Jetson Nano: 33.13 FPS at 6.23 W.", K.GREEN, "R"),
        ("Relevance to SignSpeak", "Same motivation: accessible, real-time sign recognition on modest hardware. Different route: SignSpeak classifies hand images with a CNN and does not use skeletons or distillation.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The paper is about compressing skeleton-based video models for edge devices. SignSpeak shares its goal (accessible, real-time recognition) "
                "but is an image-based CNN for fingerspelling; the two sets of numbers are not comparable.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result / role"]]
    for i in (0, 1, 2, 3, 4, 7):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.4, [2.2, 3.7, 2.9, 3.33], rows, font_size=12, row_h=0.7, header_size=13)
    notes[5] = ("The pattern: early camera-based systems, depth-sensor fingerspelling, deep CNNs on static ASL images, video translation, and surveys. "
                "Where a work's figures were not checked they are not quoted. The project literature survey also lists four 2024-2025 papers whose authors were not recorded; they are in the report's reference list with that caveat.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: task and protocol set the number")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 7.6, 5.3,
                     ["Base paper\nAUTSL\n(226 word signs)", "Base paper\nWLASL-100 baseline\n(STA-GCN)", "Base paper\nWLASL-100\nMAG-KD", "SignSpeak\n29 static classes\n(own test split)"],
                     [("Top-1 accuracy", [0.9425, 0.5492, 0.6633, T["accuracy"]])], colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.1)
    pt = gf.chart.plots[0].series[0].points[3]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = K.ORANGE
    K.card(s, 8.5, 1.45, 4.23, 2.6, "Takeaway", "Accuracy ranges from 55 % to 94 % in the base paper alone, depending on the dataset and method. A number is meaningful only with its task and protocol.", K.GREEN, "↑", body_size=14)
    K.card(s, 8.5, 4.2, 4.23, 2.5, "Caveat", "The first three bars are video word signs from the base paper; the orange bar is our own measurement on letters. They are not comparable and no ranking is implied.", K.ORANGE, "!", body_size=14)
    notes[6] = ("Chart values: 94.25 % (AUTSL), 54.92 % and 66.33 % (WLASL-100 baseline and MAG-KD) are from the base paper; 82.81 % is SignSpeak's accuracy on its own held-out 1,088-image test split. "
                "Different tasks, data and models, so the point is only that protocol and task decide the number.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Signer signs", K.NAVY), ("Human interpreter or written note", K.NAVY), ("Hearing person understands", K.NAVY), ("Reply (spoken)", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Human interpretation", "Accurate, but interpreters are scarce and costly, and are not available for everyday spontaneous exchanges such as a shop counter, a bus stop or a family meal.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Research recognisers", "Studies use depth cameras, video and skeleton pipelines or CNNs on static images. Many stop at a model and an accuracy figure, with no live camera loop, sentence or speech.", K.ORANGE, "2", body_size=17)
    notes[7] = "Today's path is a human interpreter or writing. Research systems exist, but most are not something a person can open, sign to and hear back."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Scarce, costly interpreters", "Not available for spontaneous, everyday conversation.", K.ORANGE, "1"),
        ("Special hardware or heavy pipelines", "Depth sensors, pose estimators or large video models are hard to run on an ordinary laptop.", K.BLUE, "2"),
        ("Static images, no live loop", "Many CNN studies classify single images and do not turn a camera stream into a sentence with speech.", K.GREEN, "3"),
        ("Fragile in the real world", "Accuracy degrades with background clutter, lighting changes and unseen signers.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=17)
    notes[8] = "Four limitations, taken from the project's problem statement and the survey. Each maps to a design decision later: webcam only, on-device CNN, sentence builder with speech, measured robustness."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("End-to-end system", "Few works go from a live webcam to letters, a sentence and speech inside one usable application."),
            ("Accessible hardware", "Skeleton and depth pipelines need extra models or sensors; an ordinary webcam and CPU should be enough for letters."),
            ("Honest evaluation", "Results are rarely given with their split, per-class scores, confidence behaviour and latency in one place."),
            ("Robustness", "Background clutter, lighting and hand orientation effects are seldom measured, so real-use accuracy is unknown."),
            ("Unseen signers", "Most public alphabet data carries no signer identity, so signer-independent accuracy is hard to test.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=18, color=K.INK, anchor="m")
    notes[9] = ("The batch literature survey states the gap as: an accessible end-to-end system that recognises signs live and renders them as text and speech. "
                "We address the first three directly; the last two we measure or state honestly.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, PROBLEM_STATEMENT, size=20, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Webcam frames (960 x 540 JPEG) with the hand inside a fixed region"), ("Output", "Letter and confidence, a sentence on screen, spoken text"), ("Record", "Collected images, trained model, per-class metrics")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=15, color=K.TXT)
    notes[10] = "Read the statement slowly. Input is the camera stream, output is letters, a sentence and speech, and the system keeps data and a trained model that can be inspected."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "The six objectives from the proposal. The later reviews report which were met; words (objective 1) and unseen signers (objective 6) were not."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Capture", "Webcam frame, cropped to a fixed 300 x 300 region of interest that contains the signing hand.", K.BLUE, "1"),
        ("Segment", "Grayscale, Gaussian blur and Otsu threshold with morphology give a clean hand mask; resize to 64 x 64.", K.GREEN, "2"),
        ("Classify", "Three-block CNN with 2.2 million parameters; softmax over 29 classes (A-Z, space, del, nothing).", K.ORANGE, "3"),
        ("Decide", "Confidence threshold, majority vote and a stability counter commit one letter at a time.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=17)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "W", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Application", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "Sentence builder with text-to-speech, in-app image collection, background training with live curves, a model report page; FastAPI service and React interface over the original desktop code.", size=15, color=K.TXT)
    notes[12] = "Four stages from camera to committed letter, and the application around them. The CNN, preprocessing and decision rules are the same ones used in the desktop app."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 2.8, 1.3, 7.7, 5.3)
    K.text(s, 0.6, 6.62, 12.13, 0.3, "Browser owns the webcam; the API runs the recognition pipeline; the CNN model and image store sit behind it", size=12, color=K.TXT, align="c", italic=True)
    notes[13] = "Follow the diagram: the browser captures frames, the FastAPI service runs preprocessing and the CNN, and the files on disk hold the images and the trained model."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working recogniser", "A trained CNN that reads the ASL alphabet from a live webcam.", K.BLUE, "1"),
            ("Letters to speech", "A sentence builder with stability logic and spoken output.", K.GREEN, "2"),
            ("Honest evaluation", "Accuracy, per-class scores, confidence behaviour, robustness and latency, with gaps stated.", K.ORANGE, "3"),
            ("Complete tool", "Collect, train and inspect from one application, not a notebook.", K.PURPLE, "4"),
            ("Reproducibility", "One script writes every metric to JSON; figures are regenerated from it.", K.BLUE, "5"),
            ("Quality", "Automated tests for the API, preprocessing, recogniser and model.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=16)
    notes[14] = "Expected outcomes and contributions. We will report whatever the held-out test split shows, including the weak letters."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 2, 3, 4, 5, 6, 14, 15]
    K.text(s, 0.6, 1.35, 12.13, 5.5, [(f"[{n}] {REFS[n - 1]}", {"space_after": 5}) for n in idx], size=14, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = "Likely questions: why a CNN on segmented images rather than hand landmarks; why only letters; how the system handles different signers (not yet tested; stated as a limitation)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
