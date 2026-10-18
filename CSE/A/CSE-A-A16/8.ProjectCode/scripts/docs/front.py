"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p("SignSpeak: a live webcam recogniser for the ASL alphabet with text and speech output", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the creator of the public ASL Alphabet image collection on Kaggle, and the developers of the open-source software "
        "used in this project (TensorFlow, Keras, OpenCV, NumPy, scikit-learn, FastAPI, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    t = F.T
    thr = F.thr_row(0.75)
    d.p("Sign language is the main means of communication for many people with hearing or speech impairment, yet few hearing people understand it, and human interpreters are scarce, costly and rarely "
        "available for everyday exchanges. This project presents SignSpeak, a real-time recognition system for the American Sign Language manual alphabet that turns hand signs shown to a webcam into text and "
        "synthesised speech. A browser camera sends frames to a FastAPI service; a fixed region of each frame is converted to a binary hand mask by grayscale conversion, Gaussian blur, Otsu thresholding and "
        "morphology, resized to 64x64 and classified by a three-block convolutional neural network of 2,199,453 parameters into 29 classes (26 letters and the control signs space, delete and nothing). A temporal rule "
        "(confidence of at least 0.75, a majority vote over 15 frames and 18 stable frames) decides when a letter is typed; the letters form a sentence that can be spoken. The application also collects images, trains "
        "models and reports their performance, and the original desktop application is retained.")
    d.p(f"Trained on 7,250 images (250 per class) from the public Kaggle ASL Alphabet collection on a CPU-only PC, the network reaches {F.pct(t['accuracy'], 2)} accuracy (macro F1 {F.f3(t['macro_f1'])}) on a held-out split of "
        f"{t['n']:,} images. At the 0.75 confidence threshold the system accepts {F.pct(thr['coverage'])} of frames and is right on {F.pct(thr['accuracy_on_accepted'])} of them. A recognition request takes a median of "
        f"{F.LAT['http_recognize_frame_round_trip']['median_ms']:.0f} ms (about {F.LAT['http_round_trip_fps_sequential']} frames per second). {F.PYTEST_TOTAL} automated tests and a {F.SMOKE_TOTAL}-check smoke test pass. "
        "The held-out split comes from the same recordings as the training images, so accuracy on unseen signers was not measured and is expected to be lower; recognition of words, evaluation on a physical webcam "
        "and a user study were not done. The base paper (MAG-KD, skeleton-based video recognition on edge devices) addresses a different task, so the comparison is given as context only.")
    d.p("**Keywords:** sign language recognition, American Sign Language, convolutional neural network, fingerspelling, hand segmentation, computer vision, deep learning, assistive technology, FastAPI, React.", align="left")


ABBR = [
    ("API", "Application Programming Interface"), ("ASL", "American Sign Language"), ("BN", "Batch Normalisation"), ("CNN", "Convolutional Neural Network"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"),
    ("FPS", "Frames Per Second"), ("GCN", "Graph Convolutional Network"), ("GPU", "Graphics Processing Unit"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JPEG / PNG", "Image file formats"), ("JSON", "JavaScript Object Notation"), ("KD", "Knowledge Distillation"), ("MAG-KD", "Mask-Guided Adaptive Gating Knowledge Distillation"),
    ("ReLU", "Rectified Linear Unit"), ("REST", "Representational State Transfer"), ("ROI", "Region of Interest"), ("SLR", "Sign Language Recognition"),
    ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("TTS", "Text-to-Speech"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
