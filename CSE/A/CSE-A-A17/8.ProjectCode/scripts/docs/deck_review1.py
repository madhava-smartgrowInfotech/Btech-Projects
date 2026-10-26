"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\Abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "fashion demand forecasting, XGBoost, domain adaptation, H&M, retail analytics"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


PROBLEM_STATEMENT = [
    "Forecasting models trained on one market lose accuracy when applied to another",
    "New products have no sales history, so conventional forecasting cannot cover them",
    "Fashion trends shift within a season, invalidating models trained on older data",
    "Forecasts delivered without explanation are not trusted by inventory planners",
    "Need: a demand forecasting system that adapts across markets and categories and explains the drivers behind each forecast",
]
OBJECTIVES = [
    "Collect historical fashion retail sales data and engineer seasonal, pricing and category features",
    "Train gradient-boosting models such as XGBoost to forecast demand at product and category level",
    "Apply domain adaptation to transfer the model across markets and product categories",
    "Compare adapted and non-adapted models on a target market with a different data distribution",
    "Explain forecast drivers so planners understand why demand is predicted to move",
    "Evaluate using MAE, RMSE and MAPE, and translate error into inventory impact",
]

# [work, year, approach, reported result / limitation]
LIT_ROWS = [
    [f"Mishra et al. {N('base')} (base paper)", "2026", "LightGBM plus four explainability levels (PFI, PDP, SHAP, ICE); 41 features; Optuna tuning",
     "UCI Online Retail II, 1.07 M rows, 4,631 SKUs: RMSE 33.71, R2 0.243; one dataset, one-day-ahead"],
    [f"Chen and Guestrin {N('xgb')}", "2016", "XGBoost: regularised tree boosting that accepts per-row sample weights", "Default tabular learner; assumes train and test share one distribution"],
    [f"Makridakis et al. {N('m5')}", "2022", "M5 retail forecasting competition (as summarised in the base paper)", "Boosted trees with hand-built features matched or beat deep learning"],
    [f"Shimodaira {N('shim')}", "2000", "Weight the training loss by the target/source density ratio under covariate shift", "Theory; needs a density-ratio estimate, weights can be extreme"],
    [f"Bickel et al. {N('bickel')}", "2009", "Learn the shift weights with a logistic source-versus-target classifier", "The idea TrendCast implements; depends on overlap between domains"],
    [f"Thomassey {N('thom')}", "2010", "Clothing sales forecasting with fuzzy, neural and data-mining models", "Volatile, seasonal, short-lived items; pre-dates boosted trees"],
    [f"Ekambaram et al. {N('ekam')}", "2020", "Attention-based multi-modal model for new fashion products (images, attributes, external factors)", "Beats k-NN baselines; needs images and large history"],
    [f"Singh et al. {N('singh')}", "2019", "Attribute-based demand forecasting for new fashion items", "Finds which attributes drive demand; no transfer across markets"],
    [f"Zhou et al. {N('cdlf')}", "2026", "Conditional diffusion model for cold-start product life cycles", "More accurate than classical baselines; not fashion-specific"],
    [f"Batch survey {N('s2')}-{N('s5')}", "2026", "AI-driven demand forecasting and inventory planning (titles only)", "Full texts not accessed; no figures quoted"],
]

GAPS = [
    f"The base paper validates one dataset from one UK wholesaler and does not test transfer to a differently distributed market {N('base')}",
    f"Importance weighting for covariate shift is well established {N('shim', 'bickel')}, but is not reported together with gradient-boosted demand models and honest controls (time overlap, weight scale)",
    f"Fashion cold-start work relies on product images and external signals {N('ekam', 'skend')}; a plain tabular route is seldom compared with simple baselines",
    f"Planners override forecasts they cannot interpret {N('base')}: explanations must accompany any adapted model",
    "Gains are rarely checked against unweighted pooling or a repeat-yesterday forecast, so reported improvements may not survive a fair control",
]

BASE_RMSE = {"Naive lag-1": 33.71 + 14.06, "Moving avg (7d)": 33.71 + 2.86, "Random forest": 33.71 + 0.45, "CatBoost": 33.83, "XGBoost": 33.89, "LightGBM": 33.71}


def pic_or_note(s, name, x, y, w, h):
    p = F.diag(name)
    if os.path.exists(p):
        K.picture(s, p, x, y, w, h)
        return True
    print("WARNING: missing", p)
    return False


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    xg = F.fm(F.XGB)
    ma = F.fm(F.MA)
    nv = F.fm("Naive (lag 1)")
    dm = F.DOM["summary"]

    # 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: TrendCast turns retail transactions into daily product demand, forecasts it with XGBoost and "
                "tests whether importance-weighted domain adaptation carries a model from one sales domain to another.")

    # 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Outcomes measured so far"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. The review covers why the problem matters and what is missing in current work, then what we built and what we measured."

    # 3 abstract
    s = S[2]
    set_title(s, "Abstract: forecast fashion demand across markets")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Fashion demand moves with short product life cycles, seasons and promotions. A model trained in one market or channel loses accuracy when the data distribution changes, and planners distrust forecasts they cannot explain.", K.ORANGE, "!"),
        ("Our idea", "Daily product demand from retail transactions, forecast with XGBoost on calendar, price, attribute, lag and rolling features. A logistic domain classifier gives importance weights p/(1-p) that move a source-trained model to a shifted target domain.", K.BLUE, "i"),
        ("The deliverable", "TrendCast: a FastAPI service (22 routes) and a 13-page React console on real H&M data, measured with RMSE, MAE, MAPE and R2, with controls that test whether adaptation really helps.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=17)
    notes[3] = "Three points: the weakness of single-market models, the idea of weighting source rows by how target-like they look, and what is delivered and measured on real H&M data."

    # 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "LightGBM on 41 engineered features (temporal, lag, rolling, price), tuned with Optuna (30 trials, 3-fold TimeSeriesSplit), then four explanation levels: permutation importance, partial dependence, SHAP and ICE curves.", K.BLUE, "M"),
        ("Data and reported results", "UCI Online Retail II: 1.07 M transactions, 4,631 SKUs, daily one-step forecasts. LightGBM RMSE 33.71, R2 0.243; CatBoost 33.83; XGBoost 33.89; MASE 0.799. Rolling statistics are the dominant feature group.", K.GREEN, "R"),
        ("Limits stated by the authors", "Single dataset (a UK wholesaler), stationary-relationship assumption, no promotion or holiday data, one-day horizon, a slightly leaking price feature. TrendCast adds transfer across domains.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The base paper pairs LightGBM with layered explanations on UCI online-retail data. Its own limitations section names single-dataset validation and the "
                "stationarity assumption, which is exactly where domain adaptation is meant to help.")

    # 5 literature table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]] + LIT_ROWS
    K.table(s, 0.6, 1.25, [2.35, 0.6, 4.95, 4.23], rows, font_size=10, row_h=0.5, header_size=11, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey is the base paper plus four 2026 titles that we could not locate online, so no results are quoted from them. "
                "The other rows were checked against the papers or their abstract pages.")

    # 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence: boosted trees beat simple baselines")
    wipe(s)
    cats = list(BASE_RMSE.keys())
    gf = K.bar_chart(s, 0.6, 1.4, 6.6, 4.9, cats, [("RMSE (lower is better)", [round(v, 2) for v in BASE_RMSE.values()])], colors=[K.BLUE],
                     number_format="0.00", font_size=11, max_value=55, legend=False)
    K.text(s, 0.6, 6.35, 6.6, 0.7, "Base paper, UCI Online Retail II. Naive, moving-average and random-forest bars are derived from the RMSE gaps the paper states (14.06, 2.86, 0.45).",
           size=11, color=K.TXT, italic=True)
    K.text(s, 7.5, 1.45, 5.25, 5.0, [
        ("**What the literature shows**", {"space_after": 8}),
        ("LightGBM, CatBoost and XGBoost end within 0.18 RMSE of each other; the choice of boosting library matters little " + N('base') + ".", {"bullet": True, "space_after": 8}),
        ("MASE 0.799: a 20.1% gain over naive forecasting; rolling statistics are the dominant feature group " + N('base') + ".", {"bullet": True, "space_after": 8}),
        ("In the M5 competition boosted trees with hand-built features matched deep learning " + N('m5') + " (as summarised in the base paper).", {"bullet": True, "space_after": 8}),
        ("Caveat: R2 is only 0.243 on SKU-level daily demand, so absolute accuracy is modest.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("These numbers are published evidence from the base paper, on different data from ours, so they are not put on one scale with our own results. "
                "Our own measurements are in the last slides.")

    # 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Sales history", K.NAVY), ("Classical model (moving average, ARIMA-type)", K.NAVY), ("Planner adjusts by judgement", K.NAVY), ("Order quantity", K.NAVY), ("One model per market", K.NAVY)], size=14, gap=0.45)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Statistical and judgement-based planning",
           "Classical methods give parameters planners can read, but assume linear, stationary demand.\nPromotions, seasons and short life cycles break those assumptions.\nPlanners override forecasts they do not understand " + N('base') + ".", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Machine-learning forecasters",
           "Gradient boosting predicts retail demand well " + N('base', 'm5') + ".\nA model is usually trained and tested inside one market.\nMoving to a new market or a new product category normally means retraining.", K.ORANGE, "2", body_size=17)
    notes[7] = "Today most planning uses classical models plus judgement; where machine learning is used it is trained and tested in a single market."

    # 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Accuracy drops across markets", "No history for new products", "Trends shift in-season", "Unexplained forecasts"]
    bodies = [
        "A model trained on one market loses accuracy where the data distribution differs: other stores, channels, regions or categories.",
        "A new product has no sales history, so lag-based forecasting cannot cover it.",
        "Fashion trends move within a season and invalidate models trained on older data.",
        "Forecasts delivered without a reason are not trusted by inventory planners, who override them.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=16)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"The base paper adds: one dataset, a stationarity assumption, no promotion or holiday signals and a one-day horizon {N('base')}.", size=14, color=K.INK, anchor="m")
    notes[8] = "The four limitations are the batch problem statement; the last line is the base paper's own limitations section."

    # 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=16, color=K.INK, anchor="m")
    notes[9] = ("Gap 1 and 4 come from the base paper's limitations and introduction. Gap 2 and 3 are our reading of the covariate-shift and fashion cold-start literature. "
                "Gap 5 is why we added unweighted pooling, shuffled weights and a time-clean protocol as controls.")

    # 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Transactions, articles and optional customers CSV (H&M schema); source and target domain choice"),
                                 ("Output", "Daily demand forecast per article (7, 14, 30 days), before/after adaptation report, feature importance"),
                                 ("Measures", "RMSE, MAE, safe MAPE, R2 (and WAPE offline) against a moving-average baseline")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=13, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a demand forecaster that adapts across markets and explains its drivers."

    # 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Later reviews report which objectives were met: product-level forecasting, adaptation and error metrics are built; inventory impact and per-forecast explanation are partial.",
           size=13, color=K.TXT, italic=True)
    notes[11] = "The six objectives from the proposal. The final slides state plainly which are met, partly met or not met."

    # 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "Validate, clean and merge the three CSV tables; count units per article per day; fill missing days with zero.", K.BLUE, "1"),
        ("Forecast", "31 features: calendar, price, channel, six attributes, shifted lags and rolling means. Chronological 80/20 split; XGBoost versus moving average.", K.GREEN, "2"),
        ("Adapt", "Logistic domain classifier gives weight p/(1-p) clipped at 20; XGBoost is retrained with those sample weights and tested on the target domain.", K.ORANGE, "3"),
        ("Evaluate", "RMSE, MAE, MAPE and R2 before and after, feature importance and recursive 7/14/30-day forecasts.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React console for the whole pipeline", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "Upload, explore, features, training, domain adaptation, evaluation, forecast and AI insights; every step is gated until its inputs exist. No database and no GPU are needed; Gemini is optional and sees only aggregated numbers.", size=14, color=K.TXT)
    notes[12] = "Four stages and the console that wraps them. Domain adaptation here means importance weighting for covariate shift, not adversarial or deep adaptation."

    # 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    pic_or_note(s, "01_system_architecture.png", 0.6, 1.2, 7.6, 5.9)
    K.text(s, 8.4, 1.5, 4.35, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"React console on port {F.WEB_PORT}; Vite proxies /api.", {"bullet": True}),
        (f"FastAPI on port {F.API_PORT}: 22 routes, one gated pipeline state.", {"bullet": True}),
        ("src/ package: cleaning, daily demand, features, XGBoost, domain adaptation, metrics.", {"bullet": True}),
        ("State in memory; trained models saved as joblib files; no database.", {"bullet": True}),
        ("Optional Gemini explanation over aggregated numbers.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the console calls the API, the API calls the src package. There is no database."

    # 14 outcomes
    s = S[13]
    set_title(s, "Outcomes measured so far")
    wipe(s)
    r = F.pair("Online")
    heads = ["Working platform", "Forecast accuracy", "Domain adaptation", "Multi-day forecast", "Stated limits"]
    bodies = [
        f"22 routes, 13 console pages, {F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks pass; real H&M subset ({F.DS['raw_rows']:,} rows).",
        f"Test RMSE: XGBoost {xg:.2f}, moving average {ma:.2f}, but repeat-yesterday {nv:.2f}: boosting beats the baseline, not the strongest simple rule.",
        f"Helped on {dm['pairs_where_adaptation_helped']} of {dm['pairs_total']} source-target pairs; mean RMSE change {dm['mean_rmse_improvement_pct']:+.1f}%. The one clear gain shrinks to {r['time_clean']['rmse_improvement_pct_mean']:+.1f}% without time overlap.",
        "Recursive forecast error grows with horizon (WAPE 41% over days 1-7 to 124% over days 15-30) and trails a flat 7-day mean.",
        "107 days, 300 articles, one data shard; no cold-start mode in the product; SHAP and inventory impact computed offline only.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=14)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=15)
    notes[14] = ("Because the system is already built, this slide reports measured outcomes, not promises. The results are mixed and we state them that way; details are in the later reviews and the report.")

    # 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = sorted(["base", "xgb", "m5", "shap", "shim", "sugi", "bickel", "pan", "thom", "ekam", "singh", "skend", "cdlf", "hm"], key=lambda k: _IDX[k])
    K.text(s, 0.6, 1.25, 12.13, 5.8, [(f"[{_IDX[k]}] {REFS[_IDX[k] - 1]}", {"space_after": 4}) for k in keys], size=12.5, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why only 300 articles (best-selling, enough history for lag features, fast to repeat 5 seeds); what 'domain' means here (sales channel, product group or department); "
                 "why adaptation did not help more (time overlap, degenerate weights, unweighted pooling is as good).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
