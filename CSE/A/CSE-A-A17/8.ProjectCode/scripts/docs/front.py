"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        os.makedirs(os.path.dirname(BANNER), exist_ok=True)
        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, L. N. Mishra, B. Senapati, S. K. Nayak and A. J. Rangari, whose explainable retail demand forecasting work "
        "framed this project; H&M Group, whose released articles, customers and transactions tables made an honest test on real retail data possible; "
        "and the maintainers of the open-source software used here: XGBoost, scikit-learn, pandas, NumPy, FastAPI, React, Vite, Recharts and Playwright, "
        "and LightGBM, which we used only as a comparison model.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    fc, dom = F.FC["models"], F.DOM
    on, sh = F.pair("Online"), F.pair("Store ->")
    p = on
    xgb, ma, lag1 = fc[F.XGB]["RMSE"], fc[F.MA]["RMSE"], fc["Naive (lag 1)"]["RMSE"]
    ds = F.DS
    rec = F.REC["results"]
    d.p("Fashion retailers must decide how much of each article to stock before demand is known, and a forecasting model fitted on one sales context loses accuracy when "
        f"it is moved to another. This project builds {F.PRODUCT}, a demand-forecasting application that turns H&M-style retail transactions into daily units sold per article, "
        "forecasts them with XGBoost on calendar, price, product-attribute, lag and rolling features, and adds importance-weighted domain adaptation, in which a logistic domain "
        "classifier gives each source row the weight p/(1-p), so that a model trained on one sales domain can be tested on a shifted one. "
        "It is delivered as a FastAPI service with 22 routes and a 13-page React console, and is covered by 39 unit and API tests and a 26-check smoke test.")
    d.p(f"We measured it on a real subset of the H&M data: the {ds['articles']} best-selling articles, {ds['raw_rows']:,} transaction rows, {ds['date_min']} to {ds['date_max']}. "
        f"On the last 20 % of the dates (the Christmas peak and the January sales), XGBoost reaches an RMSE of {xgb:.2f} against {ma:.2f} for the product's moving-average baseline, "
        f"but not against simply repeating yesterday's demand ({lag1:.2f}). Domain adaptation as shipped improved RMSE on {dom['summary']['pairs_where_adaptation_helped']} of {dom['summary']['pairs_total']} "
        f"source-target pairs; the one clear gain ({p['app']['rmse_improvement_pct_mean']:+.1f} %) disappears ({p['time_clean']['rmse_improvement_pct_mean']:+.1f} %) once the source model may not train on dates "
        "inside the target test window, and badly scaled weights can make XGBoost predict a constant. Recursive 7 to 30-day forecasts were worse than a flat 7-day mean. "
        "We report these findings as measured, with the controls that expose them.")
    d.p("**Keywords:** fashion demand forecasting, XGBoost, domain adaptation, covariate shift, importance weighting, retail analytics, inventory planning, explainability, H&M dataset.")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("CORS", "Cross-Origin Resource Sharing"), ("CSV", "Comma-Separated Values"),
    ("DA", "Domain Adaptation"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"), ("FR", "Functional Requirement"),
    ("GBM", "Gradient Boosting Machine"), ("HTTP", "Hypertext Transfer Protocol"), ("ICE", "Individual Conditional Expectation"), ("JSON", "JavaScript Object Notation"),
    ("KPI", "Key Performance Indicator"), ("MA", "Moving Average"), ("MAE", "Mean Absolute Error"), ("MAPE", "Mean Absolute Percentage Error"),
    ("MASE", "Mean Absolute Scaled Error"), ("ML", "Machine Learning"), ("NFR", "Non-Functional Requirement"), ("PDP", "Partial Dependence Plot"),
    ("PFI", "Permutation Feature Importance"), ("R2", "Coefficient of determination"), ("REST", "Representational State Transfer"), ("RF", "Random Forest"),
    ("RMSE", "Root Mean Square Error"), ("SHAP", "SHapley Additive exPlanations"), ("SKU", "Stock-Keeping Unit"), ("UI", "User Interface"),
    ("UML", "Unified Modeling Language"), ("WAPE", "Weighted Absolute Percentage Error"), ("XAI", "Explainable Artificial Intelligence"), ("XGBoost", "eXtreme Gradient Boosting"),
]


def abbreviations(d):
    """Two abbreviation/expansion pairs per row so the list takes half the space."""
    half = (len(ABBR) + 1) // 2
    rows = [list(ABBR[i]) + (list(ABBR[i + half]) if i + half < len(ABBR) else ["", ""]) for i in range(half)]
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion", "Abbreviation", "Expansion"], rows, widths=[0.95, 2.25, 0.95, 2.25], size=10)
