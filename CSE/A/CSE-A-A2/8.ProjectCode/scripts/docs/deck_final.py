"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .deck_review1 import GAPS, OBJECTIVES, PROBLEM_STATEMENT, ref_line
from .refs import CITE

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520
CROPS = os.path.join(tempfile.gettempdir(), "a2_deck_crops")
MAIN = (248, 62, 1440, 900)   # screenshot without the sidebar and top bar


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def crop(name, box):
    """Crop a region (left, top, right, bottom in pixels) from a real screenshot into a temp PNG and return its path."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_{box[1]}_{box[3]}.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "accident detection, severity prediction, smartphone sensors, random forest, GPS, emergency response"
    cp.last_modified_by = ""
    notes = {}
    DET, SEVM, ALERT, SEEDS, INF = F.DET, F.SEVM, F.ALERT, F.SEEDS, F.INF
    EP = F.ENDPOINTS
    sev_acc, sev_f1 = f"{SEVM['accuracy']:.5f}", f"{SEVM['macro_f1']:.5f}"
    cb = CITE

    # footers on every slide except the title
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.FULL_TITLE, 24)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: ResQAI, a web application that detects a crash from sensor readings, grades its severity, asks for confirmation and drafts a located alert."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Models\nDetection Workflow, Interaction & Storage\nImplementation Highlights\nResults: Confusion Matrices, Baselines, Ablation\nThreshold, Noise & Unseen Events\nPerformance & Testing\nComparison with Base Paper\nOutput Screens, Conclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("Road accident victims often cannot call for help, and crash detection that only reports an impact gives responders no severity to prioritise by.", {"bullet": True, "space_after": 9}),
        ("ResQAI scales eleven motion-sensor features, decides with a random forest whether an event is a crash, grades it None, Minor, Moderate or Severe with a second forest, asks the occupant to confirm within 15 seconds, and then locates the crash and drafts a dispatch message with the three nearest hospitals.", {"bullet": True, "space_after": 9}),
        (f"On 2,400 held-out synthetic events the detector scores {DET['accuracy']:.3f} and the severity model {sev_acc} accuracy (macro-F1 {sev_f1}); over five seeds severity accuracy is {SEEDS['mean']['sev_accuracy']:.5f}. The classes are separable by construction: one acc_magnitude threshold matches the detector.", {"bullet": True, "space_after": 9}),
        (f"Alert-path compute is {F.ALERT_PATH['mean']:.0f} ms on loopback, plus the deliberate 15 s confirmation. {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} smoke checks and {F.TEST_PASSED} of {F.TEST_TOTAL} further checks pass; the 8 failures are input-validation findings.", {"bullet": True, "space_after": 9}),
        ("Not done: real phone sensors, real SMS or call delivery, and potholes and speed breakers (a probe flags about half of them as accidents, never as Moderate or Severe).", {"bullet": True}),
    ], size=16)
    notes[3] = "Read the problem, then the system in one breath. Give each number with its data and end on the honest limits: synthetic data, simulated sensors, nothing transmitted."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.8, F.BASE_PAPER.replace("*", ""), size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: Multi-Task FT-Transformer for traffic accident severity prediction"],
            ["Method", "Feature-token transformer for tabular data with three task heads (injury, fatal, property loss), uncertainty-weighted loss, SHAP explanations"],
            ["Data", "Police-reported crashes of Dongguan, China, 2019-2023: 30,442 records (19,483 train, 4,871 validation, 6,088 test), proprietary"],
            ["Task", "Predict how severe a crash was from driver, vehicle, road and environment attributes, after the crash"],
            ["Reported results", "Accuracy 70.31 % (injury), 90.38 % (fatal), 80.79 % (property loss); F1 64.48 / 86.71 / 79.18 %; beats XGBoost, LightGBM, CatBoost and MT-CNN"],
            ["Stated limits", "One city, injury proxied by number of injured, no ordinal modelling, no training or inference cost comparison"]]
    ftable(s, 0.55, 2.7, [2.0, 10.2], rows, size=13, row_h=0.62)
    notes[4] = "Wang, Liang and He, IEEE Access 2026. It analyses recorded crashes; it does not sense, detect, locate or dispatch. We borrow severity grading and per-class metrics only."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Data / finding / limitation"],
            [f"Wang et al. (base) [{cb['base']}]", "2026", "Multi-task FT-Transformer, SHAP", "30,442 Dongguan crashes; 70.31 / 90.38 / 80.79 % accuracy; after the event"],
            [f"GreenBox [{cb['greenbox']}]", "2026", "Arduino, GSM/GPS, impact switches, MPU-6050 rollover; SMS", "Alert 'in seconds'; no accuracy or latency numbers; no machine learning"],
            [f"CIRS [{cb['cirs']}]", "2025", "YOLOv11 on camera frames + video-language model", "5,200 + 4,800 frames; top-1 86.5 %, F1 0.861; fixed cameras"],
            [f"Karthik et al. [{cb['karthik']}]", "2023", "Arduino, GSM, vibration sensor, accelerometer; SMS (as described in [6])", "No machine learning; results not quoted"],
            [f"Sharma and Sebastian [{cb['sharma']}]", "2019", "Arduino, accelerometer, vibration, seat-belt heart-rate sensor, GPS/GSM", "Sensor rules; results not quoted"],
            [f"Sameen and Pradhan [{cb['sameen']}]", "2017", "Recurrent neural network on recorded crashes", "73.76 % accuracy (as reported in [1])"],
            [f"Chen and Chen [{cb['chen']}]", "2020", "Logistic regression, decision tree, random forest; 9,472 Taiwan crashes", "Random forest most effective (as cited in [1])"],
            [f"Hamdan and Sipos [{cb['hamdan']}]", "2025", "Review of ML severity models (RF, XGBoost, SVM)", "No numeric result in the abstract"]]
    ftable(s, 0.55, 1.8, [2.7, 0.7, 4.6, 4.2], rows, size=11, row_h=0.52, align=["l", "c", "l", "l"])
    notes[5] = "Severity models use recorded data, IoT prototypes alert without grading, camera systems need infrastructure. Four more survey titles from the batch file could not be located and are not used for any claim."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in GAPS], size=19)
    notes[6] = "The gap in one sentence: nothing surveyed closes the loop from live sensing through severity grading to a located, prioritised alert."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.9, 12.2, 2.4)
    ftext(s, 0.85, 1.95, 11.6, 2.3, PROBLEM_STATEMENT, size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Eleven sensor features per event: acceleration, rotation, speed, jerk"), ("Output", "Crash or not, severity grade, located alert with the nearest hospitals"), ("Safeguard", "A 15 s confirmation so a false alarm can be cancelled")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.6, 3.9, 1.6, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.7, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.2, 3.5, 0.95, b_, size=14)
    notes[7] = "The need is a smartphone-based system that detects, grades and dispatches, with a way to cancel false alarms."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion states the real status of each: 1 and 5 partly, 2 partly, 3, 4 and 6 achieved at demonstration level."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.1, [("Sensor reading", K.FNAVY), ("Scale", K.FBLUE), ("Detect", K.GREEN), ("Grade severity", K.ORANGE), ("Confirm 15 s", K.PURPLE), ("Locate + dispatch", K.FNAVY)], size=15, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.35, 12.2, 3.5, [
        ("Two random forests of the same family: a detector (200 trees) and a severity grader (250 trees), both with balanced class weights.", {"bullet": True, "space_after": 10}),
        ("The grader never overrides the detector: no accident means severity None.", {"bullet": True, "space_after": 10}),
        ("Only Moderate or Severe starts the confirmation countdown; cancelling is logged as a false alarm.", {"bullet": True, "space_after": 10}),
        ("After confirmation: browser GPS, haversine ranking of a 10-hospital directory, a drafted message with a map link, 112 and 108. A demonstration: nothing is transmitted.", {"bullet": True}),
    ], size=17)
    notes[9] = "The proposed system as a line of boxes and the design rules behind it."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 8.4, 5.1)
    ftext(s, 9.2, 2.0, 3.6, 4.8, [
        (f"Console: React 19 + Vite 8, port {F.WEB_PORT}", {"bullet": True, "space_after": 8}),
        (f"API: Flask 3, port {F.API_PORT}, 10 routes", {"bullet": True, "space_after": 8}),
        ("ML: scaler + two random forests, loaded once", {"bullet": True, "space_after": 8}),
        ("Data: users.json, hospitals.json, four pickles; no database", {"bullet": True, "space_after": 8}),
        ("Outside: OpenStreetMap tiles and the browser's location", {"bullet": True}),
    ], size=15)
    notes[10] = "Layers and storage. Readings, predictions and alerts are not stored."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.9, 12.2, 1.2, [("Generate", K.FNAVY), ("Split 80/20", K.FBLUE), ("Scale", K.GREEN), ("Train 2 forests", K.ORANGE), ("Evaluate", K.PURPLE), ("Baselines", K.FBLUE), ("Stress tests", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.4, 12.2, 3.4, [
        (f"{F.N_TOTAL:,} synthetic events from train_models.py (seed 42): {F.N_TRAIN:,} for training, {F.N_TEST:,} held out, stratified on the accident label.", {"bullet": True, "space_after": 9}),
        ("The shipped pickles were checked against a fresh retrain: 100 % agreement on all 2,400 test events.", {"bullet": True, "space_after": 9}),
        ("Added measurements: a five-seed repeat, four baselines, feature importance and ablation, a noise sweep, a threshold sweep and pothole / speed-breaker probes.", {"bullet": True, "space_after": 9}),
        ("API latency measured on 127.0.0.1; 55 unit, edge and security checks plus the shipped 33-check smoke test.", {"bullet": True, "space_after": 9}),
        ("Every number is written by scripts/measure_metrics.py, measure_api.py or run_tests.py to experiments/resqai_eval.", {"bullet": True}),
    ], size=16)
    notes[11] = "Workflow and the measurements added on top of the shipped training script, which has no validation split and no cross-validation."

    # ---- 12 dataset and models
    s = S[11]
    set_title(s, "Dataset & Models")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.05, f"{F.N_TOTAL:,}", "synthetic events, seed 42", K.FBLUE, 26)
    ftile(s, 0.55, 3.1, 2.9, 1.05, "11", "features per event", K.GREEN, 26)
    ftile(s, 0.55, 4.3, 2.9, 1.05, "19.9 %", "crashes: minor 1,215, moderate 725, severe 447", K.ORANGE, 24)
    K.picture(s, F.fig("f01_feature_separation.png"), 3.7, 1.85, 5.4, 2.4)
    rows = [["Component", "Choice"], ["Scaler", "StandardScaler"], ["Detector", "RF, 200 trees, depth 10"], ["Severity", "RF, 250 trees, depth 12"], ["Fallback", "acc_magnitude 15 / 28 / 45"]]
    ftable(s, 9.3, 1.9, [1.4, 2.1], rows, size=11, row_h=0.6, first_bold=True)
    ftext(s, 3.7, 4.4, 5.5, 2.5, [
        ("Crashes start at 15.3 m/s² and harsh braking never exceeds 13.2: one gap separates the classes.", {"bullet": True}),
        ("The generator has no potholes, speed breakers, phone drops or real sensor noise.", {"bullet": True}),
        ("The same generator parameters drive the live simulator in the console.", {"bullet": True}),
    ], size=13)
    ftext(s, 0.55, 5.55, 3.0, 1.3, "Test split: 1,463 normal, 460 braking, 233 minor, 156 moderate, 88 severe events.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "Be frank about the data: synthetic. It makes detection easy and the metrics optimistic, and it contains nothing like a pothole."

    # ---- 13 detection workflow
    s = S[12]
    set_title(s, "Detection Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_core_algorithm.png"), 0.55, 1.75, 5.0, 5.15)
    ftext(s, 5.9, 1.95, 6.9, 4.9, [
        ("The payload is completed (acc_z defaults to 9.81) and magnitudes are derived.", {"bullet": True, "space_after": 10}),
        ("The scaler uses the training mean and deviation; the detector returns the class and the probability of an accident.", {"bullet": True, "space_after": 10}),
        ("The severity forest returns 0 to 3; an accident flag of 0 forces 0.", {"bullet": True, "space_after": 10}),
        ("If the models cannot be loaded, three thresholds on acc_magnitude answer and the response says rule_based_fallback.", {"bullet": True, "space_after": 10}),
        ("In the browser only Moderate or Severe opens the 15 s confirmation and pauses live monitoring.", {"bullet": True}),
    ], size=16)
    notes[13] = "The core algorithm is predict_from_features in app.py, plus the confirmation rule in MonitorContext.jsx."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_main_flow.png"), 0.55, 1.75, 6.6, 5.2)
    K.picture(s, F.diag("08_use_case.png"), 7.3, 1.8, 5.5, 5.1)
    notes[14] = "Sequence of one event, with the cancel and the send branches, and the use cases of the occupant."

    # ---- 15 database design
    s = S[14]
    set_title(s, "Storage & Data Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.5)
    rows = [["Store", "Content"], ["users.json", "id, name, unique e-mail, scrypt hash, created"], ["hospitals.json", "10 Hyderabad hospitals + default location"], ["accident_detector.pkl", "385 KB, 200-tree forest"],
            ["severity_predictor.pkl", "2.6 MB, 250-tree forest"], ["feature_scaler.pkl", "StandardScaler, 11 features"], ["feature_names.pkl", "feature order"]]
    ftable(s, 7.8, 1.9, [2.4, 2.6], rows, size=11, row_h=0.55)
    ftext(s, 0.55, 6.4, 12.2, 0.5, "No relational database: readings, predictions and alerts are transient, and an alert carries no user id.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Two JSON files and four pickles; an honest simplification for a demonstration."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("10", "REST routes", K.FBLUE), (f"{F.SMOKE_TOTAL} + {F.TEST_TOTAL}", "smoke + unit / edge / security checks", K.GREEN), ("6", "React screens, light and dark", K.ORANGE), ("672", "lines of Python (app, auth, training)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Accounts: scrypt hashes, signed 12-hour tokens, e-mail case-insensitive, five failed logins lock an account for five minutes (429); REQUIRE_AUTH=1 guards every API route except /api/health.", {"bullet": True, "space_after": 9}),
        ("Confirmation: 15-second countdown in the browser, cancel by button or Escape, live monitoring paused while it is open.", {"bullet": True, "space_after": 9}),
        ("Location and dispatch: browser geolocation with a 3 s timeout, haversine ranking, message with a map link, alert id, 112 and 108 tap-to-call.", {"bullet": True, "space_after": 9}),
        ("Known gaps: bad input can return HTTP 500, two alerts in one second share an id, requirements are not pinned.", {"bullet": True}),
    ], size=16)
    notes[16] = "Highlights and, in the last line, the defects the tests found. The product code was not changed during documentation."

    # ---- 17 results: detector and severity
    s = S[16]
    set_title(s, "Results – Detector & Severity")
    wipe(s)
    for i, (v, l, c) in enumerate([("1.000", "detector accuracy, precision, recall, F1", K.FBLUE), (sev_acc, "severity accuracy after the gate", K.GREEN), (sev_f1, "severity macro-F1", K.ORANGE),
                                   (f"{ALERT['precision']:.3f} / {ALERT['recall']:.3f}", "alert-trigger precision / recall", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 24)
    K.picture(s, F.fig("f02_confusion_matrices.png"), 0.55, 3.1, 7.2, 3.3)
    rows = [["Class", "Precision", "Recall", "F1", "n"]]
    for k, v in SEVM["per_class"].items():
        rows.append([k, f"{v['precision']:.3f}", f"{v['recall']:.4f}", f"{v['f1']:.4f}", str(v["support"])])
    ftable(s, 8.0, 3.2, [1.3, 1.0, 0.9, 0.9, 0.7], rows, size=12, row_h=0.45, align=["l", "c", "c", "c", "c"])
    ftext(s, 8.0, 5.55, 4.8, 1.4, f"Five-seed mean: severity accuracy {SEEDS['mean']['sev_accuracy']:.5f} (SD {SEEDS['std']['sev_accuracy']:.5f}), macro-F1 {SEEDS['mean']['sev_macro_f1']:.5f} (SD {SEEDS['std']['sev_macro_f1']:.5f}); detector 1.000 in every run. One error in 2,400: a Minor crash graded Moderate.", size=12)
    ftext(s, 0.55, 6.5, 7.2, 0.45, "2,400 held-out synthetic events; 1,923 non-accidents and 477 accidents.", size=12, color=K.FNAVY, italic=True)
    notes[17] = "The detector made no error; the severity model made one. These are synthetic-data results: they show the pipeline works, not that it would score this on real roads."

    # ---- 18 baselines and ablation
    s = S[17]
    set_title(s, "Results – Baselines & Ablation")
    wipe(s)
    K.picture(s, F.fig("f03_baselines.png"), 0.55, 1.85, 6.1, 3.6)
    K.picture(s, F.fig("f10_ablation.png"), 6.8, 1.85, 6.0, 3.6)
    ftext(s, 0.55, 5.6, 12.2, 1.4, [
        ("A one-line acc_magnitude threshold detects crashes as well as the forest (1.000); the forest earns its place on severity: 0.9996 against 0.975 accuracy and 0.9987 against 0.912 macro-F1. HistGradientBoosting ties it.", {"bullet": True, "size": 14}),
        ("acc_magnitude carries the detector (permutation F1 drop 0.80; every other feature 0.00); the gyroscope adds nothing measurable here; speed change and speed before matter for severity.", {"bullet": True, "size": 14}),
    ])
    notes[18] = "Honest reading of the comparison: on this data detection does not need machine learning; grading severity does."

    # ---- 19 threshold, noise and unseen events
    s = S[18]
    set_title(s, "Threshold, Noise & Unseen Events")
    wipe(s)
    K.picture(s, F.fig("f06_threshold_sweep.png"), 0.55, 1.85, 4.0, 2.7)
    K.picture(s, F.fig("f05_noise_robustness.png"), 4.65, 1.85, 4.0, 2.7)
    K.picture(s, F.fig("f07_ood_probes.png"), 8.75, 1.85, 4.05, 2.7)
    ftext(s, 0.55, 4.7, 4.0, 2.2, "Threshold sweep: precision and recall stay at 1.000 from 0.2 to 0.9; 5 false alarms at 0.05, 2 missed crashes at 0.95. Not sensitive on this data.", size=13)
    ftext(s, 4.65, 4.7, 4.0, 2.2, f"Noise (never trained on): stable to 10 % noise; at 25 % the detector false-alarm rate is {F.NOISE[3]['det_false_alarm'] * 100:.1f} %, at 50 % it is {F.NOISE[4]['det_false_alarm'] * 100:.1f} %, at 100 % {F.NOISE[5]['det_false_alarm'] * 100:.0f} %. Recall stays at 0.97 or more.", size=13)
    ftext(s, 8.75, 4.7, 4.05, 2.2, f"Unseen events (constructed probes, not real data): {F.OOD['pothole']['predicted_accident_rate'] * 100:.1f} % of pothole-like and {F.OOD['speed_breaker']['predicted_accident_rate'] * 100:.1f} % of speed-breaker-like events flagged as accidents, none graded Moderate or Severe. Harsh braking: 0 of 460.", size=13)
    notes[19] = "Three stress views. The unseen-event probes are our own construction, so they show a weakness rather than a rate: objective 2 is met for hard braking and not demonstrated for potholes and speed breakers."

    # ---- 20 performance and testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{INF['mean']:.0f} ms", "in-process prediction, mean (p95 %.1f)" % INF["p95"], K.FBLUE), (f"{EP['POST /api/predict']['mean']:.0f} ms", "POST /api/predict, mean (p95 %.0f)" % EP["POST /api/predict"]["p95"], K.GREEN),
                                   (f"{F.ALERT_PATH['mean']:.0f} ms", "alert path: simulate + predict + emergency", K.ORANGE), (f"{F.LAT['predict_throughput_req_per_s_single_client']:.0f} / s", "predict requests, one client", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.15, v, l, c, 26)
    K.picture(s, F.fig("f08_api_latency.png"), 0.55, 3.15, 6.0, 3.0)
    ftext(s, 6.8, 3.2, 6.0, 3.6, [
        (f"Smoke test: {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} checks pass.", {"bullet": True, "space_after": 6}),
        (f"Unit, edge and security: {F.TEST_PASSED} of {F.TEST_TOTAL}; unit 17 of 17, security 18 of 18, edge 12 of 20.", {"bullet": True, "space_after": 6}),
        ("The 8 failures are input validation: non-numeric values give HTTP 500, k = -1 returns 9 hospitals, coordinates are not range-checked, same-second alert ids collide.", {"bullet": True, "space_after": 6}),
        ("Plus the deliberate 15 s confirmation; real SMS or call delivery time was not measured.", {"bullet": True}),
    ], size=14)
    ftext(s, 0.55, 6.3, 6.0, 0.6, "Flask development server, loopback 127.0.0.1, i5-12500, 15.7 GB RAM under heavy memory pressure: conservative figures. No load test.", size=11, color=K.FNAVY, italic=True)
    notes[20] = "Compute is tens of milliseconds; the 15 second confirmation dominates the alert time by design. The failed edge checks are findings, left unfixed on purpose."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different data and task)")
    wipe(s)
    rows = [["Aspect", f"Base paper [{cb['base']}]", "This project (ResQAI)"],
            ["Task", "Severity of a past crash from recorded attributes", "Detect a crash live, grade it, confirm, locate, dispatch"],
            ["Data", "30,442 police-reported crashes, Dongguan, 2019-2023", f"{F.N_TOTAL:,} synthetic sensor events; no real recordings"],
            ["Models", "Multi-task FT-Transformer; baselines RF, XGBoost, SVM ...", "Two random forests; baselines threshold, LR, tree, boosting"],
            ["Reported", "Accuracy 70.31 / 90.38 / 80.79 % (three severity tasks)", f"Detector 1.000; severity {sev_acc}; macro-F1 {sev_f1}"],
            ["Explainability", "SHAP at sub-class level", "Gini and permutation importance, ablation"],
            ["Live sensing, location, dispatch", "None", "Simulated feed, GPS, hospitals, drafted message (not sent)"]]
    ftable(s, 0.55, 1.85, [2.4, 4.7, 5.1], rows, size=12, row_h=0.6)
    ftext(s, 0.55, 6.15, 12.2, 0.8, "Datasets, labels and tasks differ, and ours is synthetic and easily separable: the figures cannot be compared as like for like, and no superiority is claimed.", size=13, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that the numbers are not comparable. The base paper's task is harder (real, noisy, five-class labels from police reports)."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Live Monitor")
    wipe(s)
    K.picture(s, crop("08_monitor_live.png", MAIN), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("16_monitor_minor_logged.png", MAIN), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: a normal-driving reading, 0 % accident probability, sensor trace and scenario buttons. Right: a Minor crash: detected, graded Minor at 100 % probability and only logged in the event log, with no countdown.", size=12, color=K.FNAVY)
    notes[22] = "Real screenshots from the running console. Sensor readings are simulated by the backend and pass through the real models."
    s = S[22]
    set_title(s, "Output – Confirmation and Dispatch")
    wipe(s)
    K.picture(s, crop("09_alert_countdown.png", MAIN), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("11_response_dispatched.png", MAIN), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: 'Possible accident detected' for a Severe crash, 11 seconds left, with 'I'm fine, cancel' and 'Send alert now'. Right: the dispatched alert with a map, the nearest hospital, the alert id and tap-to-call. The message is displayed; nothing is sent.", size=12, color=K.FNAVY)
    notes[23] = "The confirmation is the only protection against false alarms in the live app. The dispatch summary shows the location, the nearest hospital and the alert id."
    s = S[23]
    set_title(s, "Output – Manual Override and Cancelled Alert")
    wipe(s)
    K.picture(s, crop("15_manual_form.png", MAIN), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("17_monitor_moderate_cancelled.png", MAIN), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: manual override: type accelerometer, gyroscope, speed and jerk values and run detection. Right: a Moderate crash that the occupant cancelled; the event log records 'User confirmed they are OK. Alert cancelled.'", size=12, color=K.FNAVY)
    notes[24] = "Manual values let a reader try any reading; the right screenshot shows the cancel branch of the confirmation."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 5.0, [
        (f"**Delivered:** a working console and API from sensor reading to detection, severity, 15 s confirmation, location, nearest hospitals and a drafted alert; {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} smoke checks pass.", {"bullet": True, "space_after": 9}),
        (f"**Measured (synthetic data):** detector 1.000; severity accuracy {sev_acc}, macro-F1 {sev_f1}; five-seed means {SEEDS['mean']['sev_accuracy']:.5f} / {SEEDS['mean']['sev_macro_f1']:.5f}; alert path {F.ALERT_PATH['mean']:.0f} ms plus the 15 s confirmation.", {"bullet": True, "space_after": 9}),
        ("**Objectives:** 3, 4 and 6 achieved at demonstration level; 2 partly (hard braking yes, potholes and speed breakers not shown); 1 and 5 partly (no real recordings; alert drafted, not sent).", {"bullet": True, "space_after": 9}),
        ("**Findings:** a threshold matches the detector; noise robustness ends near 25 %; 8 input-validation failures; pickles written with a newer scikit-learn than pinned.", {"bullet": True, "space_after": 9}),
        ("**Position:** a demonstration of the pipeline, not a validated safety system.", {"bullet": True}),
    ], size=18)
    notes[25] = "Be explicit about which objectives were met and which were not, and what was never measured."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 5.0, [
        ("Collect real accelerometer and gyroscope recordings on phones, including potholes, speed breakers, hard braking and low-speed impacts, and retrain with a validation split.", {"bullet": True, "space_after": 10}),
        ("Add road-surface classes and report the false-alarm rate per event kind; add noise to training.", {"bullet": True, "space_after": 10}),
        ("Connect an SMS or call gateway, store alerts per user and measure end-to-end delivery time.", {"bullet": True, "space_after": 10}),
        ("Validate every input, use unique alert ids, pin library versions, and use a live hospital service with capacity and traffic.", {"bullet": True, "space_after": 10}),
        ("Run on the phone itself, test battery cost, and evaluate with users and responders.", {"bullet": True}),
    ], size=19)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = sorted(["base", "greenbox", "cirs", "karthik", "sharma", "hamdan", "chen", "sameen", "zheng", "gorishniy", "kendall", "breiman", "sklearn", "haversine"], key=lambda k: CITE[k])
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(ref_line(k), {"space_after": 5}) for k in keys[:7]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(ref_line(k), {"space_after": 5}) for k in keys[7:]], size=11)
    notes[27] = "Selected references; the complete numbered list, with four marked not verified, is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
