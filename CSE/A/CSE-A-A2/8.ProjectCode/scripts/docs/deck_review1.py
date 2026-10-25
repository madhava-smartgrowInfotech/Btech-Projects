"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import CITE, REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"
KEYWORDS = "accident detection, severity prediction, smartphone sensors, random forest, GPS, emergency response"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = KEYWORDS
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def ref_line(key):
    return f"[{CITE[key]}] " + REFS[CITE[key] - 1].replace("*", "")


OBJECTIVES = [
    "Collect and preprocess smartphone accelerometer, gyroscope, speed and GPS data for normal driving and collision events.",
    "Train a classifier that distinguishes genuine accidents from potholes, speed breakers and hard braking.",
    "Build a severity model that grades an impact as minor, moderate or severe.",
    "Use GPS to locate the crash and identify a suitable nearby hospital.",
    "Dispatch an automated alert carrying location and severity to emergency contacts and medical services.",
    "Evaluate detection accuracy, false alarm rate and end-to-end alert latency.",
]
PROBLEM_STATEMENT = [
    "Accidents on remote stretches often go unreported until a passer-by intervenes.",
    "Existing crash detection reports only that an impact occurred, with no indication of how serious it is.",
    "Without severity information, responders cannot prioritise between a minor scrape and a critical collision.",
    "Dedicated in-vehicle crash hardware is costly and absent from most two-wheelers and older cars.",
]
GAPS = [
    "Severity prediction in the literature works on recorded crash reports after the event, not on live sensor data.",
    "Low-cost IoT prototypes alert contacts by SMS but add no severity grade and, in the work surveyed, report no accuracy figures.",
    "Camera-based detection needs fixed cameras and still says nothing about how serious the crash is.",
    "No surveyed work closes the loop: live sensing, then severity grading, then a located, prioritised alert with a hospital.",
    "False alarms from braking, potholes and speed breakers are rarely quantified; a confirmation step is seldom evaluated.",
]


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: ResQAI watches motion-sensor readings, decides whether an event is a crash, grades how severe it is, "
                "asks the occupant to confirm, and then prepares a located alert with the nearest hospitals.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: from a motion sensor to a located alert")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Road accidents can be fatal when victims cannot call for help; the first hour decides survival. Crash detection that only says an impact happened gives responders no way to prioritise.", K.ORANGE, "!"),
        ("Our idea", "Use the sensors every smartphone already has. Two random forests decide whether a reading is a crash and grade it None, Minor, Moderate or Severe; only Moderate or Severe starts a 15-second confirmation.", K.BLUE, "i"),
        ("The deliverable", "A Flask API and a React console: simulated sensor feed, detection, countdown, GPS location, the three nearest hospitals and a drafted dispatch message. A demonstration: no real phone is attached and nothing is transmitted.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: why delay costs lives, the two-model idea with a confirmation step, and what is delivered, stressing that it is a demonstration on simulated sensor data."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER.replace("*", ""), size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "A multi-task FT-Transformer predicts injury, fatal and property-loss severity together, balancing the three losses by learned uncertainty; SHAP explains the features. Tabular crash records, SMOTE on the training split.", K.BLUE, "M"),
        ("Data and reported results", "30,442 police-reported crashes from Dongguan, China (2019-2023): 19,483 train, 4,871 validation, 6,088 test. Accuracy 70.31 % (injury), 90.38 % (fatal), 80.79 % (property loss); F1 64.48 / 86.71 / 79.18 %.", K.GREEN, "R"),
        ("Relevance to this project", "It shows that severity can be graded as a classification task and reported with accuracy, precision, recall and F1. It does not sense, detect, locate or dispatch: it analyses a crash after it happened. Its numbers are not comparable with ours.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Wang, Liang and He, IEEE Access 2026. A different task: records of past crashes, not live sensors. We borrow the idea of grading severity and of reporting per-class metrics; "
                "we do not use the FT-Transformer, the data, SMOTE or SHAP, and we do not claim to beat these numbers.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    cb = CITE
    rows = [["Work", "Year", "Approach", "Reported result / limitation"],
            [f"Wang et al. (base) [{cb['base']}]", "2026", "Multi-task FT-Transformer, uncertainty-weighted loss, SHAP; tabular crash records", "70.31 / 90.38 / 80.79 % accuracy (injury / fatal / property); one city, after the event"],
            [f"GreenBox [{cb['greenbox']}]", "2026", "Arduino + GSM/GPS shield; push-button impact and MPU-6050 rollover sensing; SMS to contacts", "Alerts 'in seconds'; no accuracy or latency figures; no machine learning"],
            [f"CIRS [{cb['cirs']}]", "2025", "YOLOv11 on camera frames plus a video-language model; multi-agent dispatch", "Top-1 86.5 %, F1 0.861 on 10,000 frames; fixed cameras, not phones"],
            [f"Karthik et al. [{cb['karthik']}]", "2023", "Arduino, GSM, vibration sensor and accelerometer; SMS alert (as described in [6])", "No machine learning; results not quoted here"],
            [f"Sharma and Sebastian [{cb['sharma']}]", "2019", "Arduino with accelerometer, vibration and seat-belt heart-rate sensor; GPS and GSM alert", "Sensor rules; results not quoted here"],
            [f"Sameen and Pradhan [{cb['sameen']}]", "2017", "Recurrent neural network on recorded crash data", "73.76 % accuracy (as reported in [1]); data from past crashes"],
            [f"Chen and Chen [{cb['chen']}]", "2020", "Logistic regression, decision tree, random forest on 9,472 Taiwan crashes (as cited in [1])", "Random forest most effective; recorded crashes, no live sensing"],
            [f"Hamdan and Sipos [{cb['hamdan']}]", "2025", "Review of machine-learning severity models (random forest, XGBoost, SVM)", "No numeric result in the abstract; severity from records"]]
    K.table(s, 0.6, 1.35, [2.5, 0.7, 4.6, 4.33], rows, font_size=12, row_h=0.6, header_size=12, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: severity models work on recorded data; IoT prototypes alert but do not grade; camera systems need infrastructure. "
                "Four more papers in the batch survey (IJRPR 2025, ITSC 2024, INDICON 2024, a 2026 review) could not be located, so the report lists them as not verified and uses none of them for a claim.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 5.7, 5.0, ["Injury\nseverity", "Fatal\nseverity", "Property\nloss"], [("Base-paper accuracy", [0.7031, 0.9038, 0.8079])],
                colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.15)
    ev = [f"Severity can be learned as a classification task: the base paper reaches 70.31 / 90.38 / 80.79 % on 6,088 test records [{cb['base']}].",
          f"Random forests were the most effective of three models on 9,472 Taiwan crashes, as the base paper reports [{cb['chen']}].",
          f"Low-cost accelerometer, GPS and GSM hardware can send an SMS alert in seconds, and the owner can cancel a false detection [{cb['greenbox']}].",
          f"Camera detection reaches F1 0.861 but needs fixed cameras [{cb['cirs']}]; a phone is already inside the vehicle.",
          "Smartphones already carry an accelerometer, a gyroscope and GPS (batch introduction)."]
    K.text(s, 6.6, 1.45, 6.1, 4.8, [(e, {"bullet": True, "space_after": 10}) for e in ev], size=15, color=K.INK)
    K.text(s, 0.6, 6.5, 12.1, 0.5, "Different tasks and data from ours: shown as context for the approach, not as our result.", size=12, color=K.TXT, italic=True)
    notes[6] = "Evidence that each ingredient works on its own. None of these numbers is ours, and they come from different tasks; this project measures its own pipeline on its own data in later reviews."

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Crash happens", K.NAVY), ("Witness or victim calls", K.NAVY), ("Responders told 'accident'", K.NAVY), ("Severity learned on arrival", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Manual reporting and dedicated hardware", "Accidents on remote stretches go unreported until a passer-by intervenes. Dedicated in-vehicle crash units exist but are costly and absent from most two-wheelers and older cars. Where detection exists, it reports only that an impact occurred.", K.BLUE, "1", body_size=16)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Recorded-data models and IoT prototypes", "Machine-learning severity models are trained on past crash reports and used for analysis and planning, not for a live alert. Arduino and GSM prototypes raise SMS alerts from impact switches or accelerometers but add no severity grade.", K.ORANGE, "2", body_size=16)
    notes[7] = "Today the severity of a crash is learned only when someone arrives. Existing research either analyses past crashes or sends an unprioritised alert."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["No severity in the alert", "Analysis after the event", "Hardware cost and reach", "Unquantified false alarms"]
    bodies = ["Responders cannot tell a minor scrape from a critical collision, so every alert looks the same.",
              "Severity models use recorded crash reports; they cannot raise an alert at the moment of impact.",
              "In-vehicle crash units and fixed cameras need installation; many vehicles have neither, while almost every driver has a phone.",
              "Braking, potholes and speed breakers shake a sensor too; the surveyed work gives few or no false-alarm figures."]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=18)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: GreenBox reports alert delivery only as 'a matter of seconds', with no measured latency or accuracy.", size=15, color=K.INK, anchor="m")
    notes[8] = "Four limitations, each answered by a design decision: a severity grade, live sensing, a phone-based design and a measured false-alarm rate with a confirmation step."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 1.05
        K.circle(s, 0.6, y + 0.12, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.95, g, size=18, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as the missing closed loop from live sensing through severity grading to automated hospital dispatch. "
                "We address it end to end on simulated data; gaps 1-4 are addressed by design, gap 5 is measured later and the answer for potholes and speed breakers is only partly positive.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.9, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.65, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=20, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Accelerometer, gyroscope, speed and jerk readings (11 features) from the phone"), ("Output", "Accident or not, severity None / Minor / Moderate / Severe, and a located alert with the nearest hospitals"), ("Safeguard", "A 15-second confirmation so the occupant can cancel a false alarm")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.65, 3.85, 1.9)
        K.text(s, x + 0.25, 4.77, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.25, 3.4, 1.25, b_, size=15, color=K.TXT)
    notes[10] = "Need: a smartphone-based system that detects a crash, grades its severity, and dispatches a located, prioritised alert automatically, with a way to cancel false alarms."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 18}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "The six objectives from the proposal. The final review reports honestly which were met: objectives 1 and 5 are only partly met (no real recordings; the alert is drafted, not sent)."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Sense", "Eleven features per event: three acceleration axes and magnitude, three gyroscope axes and magnitude, speed before, speed change and jerk.", K.BLUE, "1"),
        ("Detect", "A random forest (200 trees) says whether the event is a crash and with what probability; hard braking is a trained negative class.", K.GREEN, "2"),
        ("Grade", "A second random forest (250 trees) grades None, Minor, Moderate or Severe; no accident forces None.", K.ORANGE, "3"),
        ("Respond", "Moderate or Severe starts a 15 s countdown; if not cancelled, GPS and the 3 nearest hospitals build the dispatch message.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "R", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "Alert rule: dispatch only if severity is Moderate or Severe and the occupant has not cancelled within 15 seconds", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.85, 10.7, 0.8, "Flask API and React console, scikit-learn models, haversine ranking over a 10-hospital directory. The sensor feed is simulated and the dispatch is a demonstration: nothing is sent.", size=14, color=K.TXT)
    notes[12] = "Four stages from sensor reading to dispatch message. The honest limits are on the slide: simulated sensors and a displayed, not transmitted, alert."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.4, 5.5)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"Presentation: React and Vite console (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: Flask API (port {F.API_PORT}) with simulate, predict, hospitals, emergency and account routes.", {"bullet": True}),
        ("Machine learning: scaler, detector forest, severity forest; rule fallback if the models are missing.", {"bullet": True}),
        ("Data: two JSON files (users, hospitals) and four model files; no database.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser sends a reading, the API scales it and calls the two forests, and the hospital directory and accounts are local JSON files."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working pipeline", "Measured metrics", "Fewer false alerts", "Honest limits", "Reusable and open"]
    bodies = ["Sensor reading to detection, severity, confirmation, location, nearest hospitals and a dispatch message in one console.",
              "Accuracy, precision, recall, F1, false-alarm rate and latency, with baselines and a noise study, reported with the data they come from.",
              "A 15-second confirmation lets the occupant cancel before any alert is raised.",
              "Data are synthetic, sensors simulated, nothing transmitted; potholes and speed breakers are not modelled.",
              "Code, scripts and tests are part of the project, so every figure can be reproduced."]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.55, h, b_, C[i % 4], str(i + 1), body_size=15)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=17)
    notes[14] = "Expected outcomes. The later reviews report what was actually measured, including what could not be measured."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = sorted(["base", "greenbox", "cirs", "karthik", "sharma", "sameen", "chen", "hamdan", "breiman", "sklearn"], key=lambda k: CITE[k])
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(ref_line(k), {"space_after": 4}) for k in keys], size=12, color=K.INK)
    notes[15] = f"Selected references; the complete numbered list of {len(REFS)} entries, with four marked as not verified, is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why two models instead of one; why a random forest; what the 100 percent detector score means (the synthetic classes are easily separable); "
                 "what happens with potholes (not modelled); and whether an alert is really sent (no, it is drafted and displayed).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
