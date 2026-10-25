"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                os.makedirs(os.path.dirname(BANNER), exist_ok=True)
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p("A system that detects a crash from smartphone-style motion readings, grades its severity, asks the occupant to confirm, and drafts a located emergency alert with the nearest hospitals",
        align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, Y. Wang, Z. Liang and Y. He, whose work on accident severity prediction gave this project its starting point, "
        "the authors of the GreenBox and CIRS papers, whose systems frame the problem of automatic detection and alerting, and the maintainers of the open-source software "
        "the project depends on (Flask, scikit-learn, NumPy, pandas, React, Vite, Leaflet, Recharts and OpenStreetMap).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    seeds = F.SEEDS
    pot, sb = F.OOD["pothole"], F.OOD["speed_breaker"]
    rule = F.BASE["rule_based_fallback (acc_magnitude < 15 / 28 / 45)"]
    d.p("Road accidents cause serious injury and death, and the harm grows when victims cannot call for help. Existing automatic crash detection typically reports only that an impact occurred, "
        "with no indication of how serious it was, and dedicated in-vehicle hardware is costly and missing from most two-wheelers and older cars. This project presents ResQAI, a system that detects "
        "a crash from the readings a smartphone provides, grades its severity, gives the occupant a short window to cancel a false alarm, and then drafts a located emergency alert that names the nearest hospital. "
        "A reading of eleven features (accelerometer, gyroscope, speed change and jerk) is scaled and passed to an accident detector, a random forest of 200 trees, and to a severity model, a random forest of 250 trees, "
        "which grades the event as None, Minor, Moderate or Severe; the detector gates the grade. Only a Moderate or Severe result starts a 15-second confirmation. If the occupant does not cancel, "
        "the browser location is read, the three nearest hospitals are ranked by great-circle distance, and a dispatch message with a map link is composed. The system is a Flask API with a React console and was built, tested and documented by the team.")
    d.p(f"No real crash recordings were available, so the models are trained on {F.N_TOTAL:,} synthetic sensor events and every figure below describes that data, not real roads. On the {F.N_TEST:,} held-out events the detector scores accuracy "
        f"{F.f4(F.DET['accuracy'])} and the severity model {F.f5(F.SEVM['accuracy'])} (macro-F1 {F.f5(F.SEVM['macro_f1'])}; five-seed mean accuracy {F.f5(seeds['mean']['sev_accuracy'])}). A threshold on acceleration magnitude detects crashes just as well, "
        f"so only the severity grade shows an advantage over a simple rule ({F.pct(rule['sev_accuracy'])} accuracy for the rule). Hard braking is never mistaken for a crash in the test set, but two constructed probes of pothole-like and speed-breaker-like events were flagged as accidents "
        f"{F.pct(pot['predicted_accident_rate'])} and {F.pct(sb['predicted_accident_rate'])} of the time, although none reached Moderate or Severe. The alert path takes about {F.ALERT_PATH['mean']:.0f} ms on a local machine, plus the 15-second confirmation window. "
        f"The interface passes all {F.SMOKE_PASSED} smoke checks and {F.TEST_PASSED} of {F.TEST_TOTAL} checks of a wider suite, whose {F.TEST_TOTAL - F.TEST_PASSED} failures, mostly missing input validation, are reported. "
        "Nothing is sent by SMS or call in this build. The base paper predicts severity from recorded Chinese crash reports, so its results cannot be compared with these.")
    d.p("**Keywords:** accident detection, severity assessment, smartphone sensors, random forest, GPS, hospital recommendation, emergency response, false-alarm confirmation, Flask, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"),
    ("CPU", "Central Processing Unit"), ("CSS", "Cascading Style Sheets"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"),
    ("FN", "False Negative"), ("FP", "False Positive"), ("FT-Transformer", "Feature Tokenizer Transformer"), ("GPS", "Global Positioning System"),
    ("GSM", "Global System for Mobile Communications"), ("HTTP", "Hypertext Transfer Protocol"), ("IMU", "Inertial Measurement Unit"),
    ("IoT", "Internet of Things"), ("JSON", "JavaScript Object Notation"), ("ML", "Machine Learning"), ("OOD", "Out of Distribution"),
    ("OSM", "OpenStreetMap"), ("REST", "Representational State Transfer"), ("RF", "Random Forest"), ("ROC", "Receiver Operating Characteristic"),
    ("SHAP", "SHapley Additive exPlanations"), ("SMOTE", "Synthetic Minority Over-sampling Technique"), ("SMS", "Short Message Service"),
    ("SPA", "Single-Page Application"), ("TN / TP", "True Negative / True Positive"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
    ("WHO", "World Health Organization"), ("WSGI", "Web Server Gateway Interface"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
