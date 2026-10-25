"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .shared import NEED, OBJECTIVES, PROBLEM_POINTS, PROBLEM_STATEMENT

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
OLD_LAYOUT_WORDS = ("Hidden Risk", "Social Media", "Multimodal")


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout / master text carries the template's deck title
    holders = list(prs.slide_layouts) + [prs.slide_master]
    for lay in holders:
        for sh in lay.shapes:
            if sh.has_text_frame and any(w in sh.text_frame.text for w in OLD_LAYOUT_WORDS):
                K.replace_text(sh, f"{F.TITLE}  ·  {review_label}")
    # the template's theme carries its own deck name
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            part._blob = part.blob.replace(b"Hidden Risk", b"Multi-Disease Prediction")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "multi-disease prediction, machine learning, diabetes, heart disease, liver disease, kidney disease"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, review_label, subtitle):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle)
        elif nm == "review label":
            K.replace_text(sh, review_label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def shorten(t, n):
    t = t.replace("\n", " ")
    return t if len(t) <= n else t[: n - 1].rsplit(" ", 1)[0] + "…"


def build(out_path):
    from .lit import LIT_HEADER, LIT_TABLE
    from .refs import NUM, REFS

    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    ho = F.HO

    # ---- 1 title
    title_slide(S[0], "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH",
                "One web interface that screens four chronic conditions from routine clinical values")
    notes[1] = ("Introduce the team and the title. One sentence: we compare four classical classifiers for diabetes, heart, liver and kidney risk "
                "and serve the best one per disease through a single web application.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. Review 1 is about why the problem matters and what is missing in current work; the measured results are summarised later."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: screen several diseases from one record")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Chronic diseases are often found late, and prediction tools handle one disease at a time. Patients visit several specialists and take separate tests "
                        "before any risk is identified; rural centres lack specialists to read routine reports.", K.ORANGE, "!"),
        ("Our idea", "Train Logistic Regression, Decision Tree, Random Forest and SVM for diabetes, heart, liver and kidney disease on standardised clinical values, "
                     "compare them on a held-out split and serve the best model per disease.", K.BLUE, "i"),
        ("The deliverable", "A web application, the Intelligent Health Assistant: one form per disease, a risk class with the model's confidence, a model-comparison page "
                            "and a dataset page, over a FastAPI service.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = ("Three sentences: the problem, our idea, what we deliver. Be upfront that the training data are synthetic because public datasets could not be downloaded when the project was first built; "
                "a supplementary check on real public data is part of later reviews.")

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.15, fill=K.CARD)
    K.text(s, 0.82, 1.47, 11.7, 1.0, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Thyroid status (Normal / Hypothyroid / Hyperthyroid) classified with SVC, Random Forest, Extra Trees, AdaBoost and CatBoost under nested cross-validation (outer 5-fold, inner 3-fold grid search); SHAP and counterfactual explanations.", K.BLUE, "M"),
        ("Data and reported results", "UCI thyroid data, 7,200 samples, 21 features. Random Forest and CatBoost reach 0.996-0.997 accuracy and AUC 0.999-1.000; SVC 0.965-0.967. Ten consensus features keep the accuracy.", K.GREEN, "R"),
        ("Stated limits", "Strong class imbalance, no external validation and one disease only. SMOTE was tried and rejected. Our project borrows the idea of comparing classical classifiers and extends it to four diseases.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.1, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The base paper is a rigorous single-disease study with explanations. We do not reproduce it: we use different diseases, different (synthetic) data "
                "and a single hold-out split, so our numbers are not comparable with its numbers.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result"]]
    for r in LIT_TABLE[:9]:
        rows.append([r[0], r[1], r[2], r[3]])
    K.table(s, 0.5, 1.25, [2.5, 4.0, 2.6, 3.2], rows, font_size=9.5, row_h=0.52, header_size=11)
    K.text(s, 0.5, 6.85, 12.3, 0.4,
           "Details were read from each paper itself; results use different data and splits and are not comparable with one another.",
           size=10.5, color=K.TXT, italic=True)
    notes[5] = ("The pattern: single-disease studies on public benchmarks, a few symptom-based multi-disease systems, and explainability added in the most recent work. "
                "Figures are quoted as reported by the authors.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: the algorithm matters")
    wipe(s)
    models = F.MODEL_ORDER
    K.bar_chart(s, 0.6, 1.4, 7.9, 5.2, [F.NAME[d] for d in F.DISEASES],
                [(m, [ho[d][m]["f1"] for d in F.DISEASES]) for m in models],
                colors=[K.BLUE, K.GOLD, K.GREEN, K.PURPLE], number_format="0.00", font_size=11, max_value=1.1, gap=50, label_size=9)
    K.card(s, 8.75, 1.5, 3.98, 2.55, "Takeaway", "Early measurements: on our hold-out split the F1-score of the same four algorithms differs by up to 0.12 within one disease, so comparing them is worthwhile.", K.GREEN, "↑", body_size=13)
    K.card(s, 8.75, 4.25, 3.98, 2.55, "Caveat", "The data are synthetic with 8% flipped labels, so no model can exceed about 92% accuracy; these bars show the pipeline works, not clinical accuracy.", K.ORANGE, "!", body_size=13)
    notes[6] = ("Chart: hold-out F1 of the four algorithms for each disease (our own early measurements from models/training_report.json, 400-row test sets, seed 42). "
                "The synthetic-data caveat must be said aloud.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Symptoms noticed late", K.NAVY), ("Visit a specialist per condition", K.NAVY), ("Separate tests per disease", K.NAVY), ("Risk identified late", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Fragmented screening", "Patients see several specialists and take separate tests before any risk is found; rural and semi-urban centres often have no specialist to read a routine report early.", K.BLUE, "1", body_size=19)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "One tool per disease", "Published prediction models and web tools usually cover a single condition, trained and packaged separately, so they cannot be reused across conditions or reached from one interface.", K.ORANGE, "2", body_size=19)
    notes[7] = "Describe the usual path: separate specialists, separate tests, and prediction tools that each cover one disease."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Several visits and tests", "One disease at a time", "No specialist locally", "Late diagnosis"]
    cols = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    for i, (h_, b_) in enumerate(zip(heads, PROBLEM_POINTS)):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, h_, b_, cols[i], str(i + 1), body_size=20)
    notes[8] = "These four limitations are the problem statement of the batch. Each maps to a design decision: one interface, four disease modules, routine inputs only, an early risk class."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    gaps = [("Single pipeline", "The surveyed papers each target one disease in isolation; a single pipeline for several conditions is missing."),
            ("One record, many conditions", "Few works let one interface screen several conditions from routine clinical values."),
            ("Explanation of drivers", "The drivers of each prediction are rarely shown to the user next to the risk class."),
            ("Honest evaluation", "Results on small benchmark sets are reported from one split; model choice needs cross-validation and checks on real data."),
            ("Usability", "Few works deliver an application a non-programmer can run end to end.")]
    wipe(s)
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.9, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.9, 0.9, k, size=15, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.75, y, 8.95, 0.9, v, size=17, color=K.INK, anchor="m")
    notes[9] = ("The batch literature survey words the gap as: a single pipeline that screens several conditions from one patient record and explains which parameters drove each prediction. "
                "The rows above spell that out. In this project the first two are addressed; explanation of drivers is only an offline analysis.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.55, 11.45, 2.6, PROBLEM_STATEMENT, size=17, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Routine clinical values of one patient (8 to 13 fields per disease)"),
                                 ("Output", "High-risk or low-risk class and the model's confidence"),
                                 ("Record", "Four models, a training report and model-comparison tables")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=15, color=K.TXT)
    K.text(s, 0.6, 6.2, 12.13, 0.7, "Need: " + NEED, size=16, color=K.NAVY, bold=True)
    notes[10] = "Read the statement slowly: four conditions, one record, a risk class with confidence, one interface, models chosen by comparison on held-out data."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 11}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    notes[11] = "Six objectives from the batch proposal: datasets, preprocessing, four algorithms, selection by metrics, a single interface, and risk with its drivers."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "Four 2,000-row clinical tables (8 to 13 features); StandardScaler fitted on the training part only.", K.BLUE, "1"),
        ("Train + compare", "Stratified 80/20 split; Logistic Regression, Decision Tree, Random Forest and RBF SVM per disease.", K.GREEN, "2"),
        ("Select", "Accuracy, precision, recall and F1 on the test set; the highest-F1 model is saved for each disease.", K.ORANGE, "3"),
        ("Serve", "FastAPI with 7 routes and a React interface with one form per disease and a result card.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.35, cw, 2.35, t, b_, col, gl, body_size=14)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.6, 3.85, 7.4, 3.3)
    K.circle(s, 8.3, 4.3, 0.6, K.NAVY, "W", 16)
    K.text(s, 9.05, 4.3, 3.7, 0.6, "Web application", size=16, color=K.INK, bold=True, anchor="m")
    K.text(s, 8.3, 5.05, 4.45, 1.9, "React interface on port 5303 over a FastAPI service on port 8303; no database and no stored patient data. The diagram shows the eight steps from data generation to the result card.", size=14, color=K.TXT)
    notes[12] = "Four stages and the application around them. Models are pickle files, data are CSV files; a prediction is computed from the request and returned."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.5, 1.15, 7.1, 5.85)
    K.text(s, 7.7, 1.5, 5.0, 5.0, [
        ("**Offline:** generate_data.py writes four CSVs; train_models.py trains 16 models and saves the best per disease.", {"bullet": True, "space_after": 8}),
        ("**Application:** FastAPI imports the Flask module's model bundles; 7 routes, CORS limited to the web app.", {"bullet": True, "space_after": 8}),
        ("**Presentation:** React single-page app: landing, overview, four prediction forms, models, datasets, about.", {"bullet": True, "space_after": 8}),
        ("**Data and state:** CSV files, pickle model files and a JSON report; nothing is stored per prediction.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the offline pipeline that creates the models, then the online application. The original Flask pages use the same bundles."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Model comparison", "Four algorithms evaluated for each of four diseases with the same split and settings.", K.BLUE, "1"),
            ("One interface", "A single application where a user picks a disease, enters values and gets a risk class.", K.GREEN, "2"),
            ("Confidence shown", "The model used, its confidence and its hold-out accuracy shown beside every result.", K.ORANGE, "3"),
            ("Honest evaluation", "Metrics with their split, a check on stability across splits and known limits stated.", K.PURPLE, "4"),
            ("Reproducible code", "Scripts regenerate the data, retrain the models and rerun every measurement.", K.BLUE, "5"),
            ("Screening aid, not diagnosis", "Positioned as a teaching and screening demonstration with a visible disclaimer.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=17)
    notes[14] = "These are expected outcomes and contributions, not clinical claims. Later reviews report what the held-out data show, including weaknesses."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = [k for k in ("base", "bangla", "scikit", "rf", "svm", "shap", "smote", "varma", "pima", "cleveland", "ilpd", "ckd") if k in NUM]
    keys.sort(key=lambda k: NUM[k])
    K.text(s, 0.6, 1.25, 12.13, 5.8, [(f"[{NUM[k]}] {REFS[NUM[k] - 1]}", {"space_after": 4}) for k in keys], size=11.5, color=K.INK)
    notes[15] = "Selected references; the full list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why these four algorithms; why F1; why the data are synthetic and what happens on real data; why no SHAP explanation in the app yet.")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
