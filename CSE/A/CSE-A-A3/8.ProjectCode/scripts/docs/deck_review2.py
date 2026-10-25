"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress. Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx.

    python -m scripts.docs.deck_review2        (from 8.ProjectCode)
"""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .shared import NEED, PROBLEM_POINTS

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"
OLD_LAYOUT_WORDS = ("Hidden Risk", "Social Media", "Multimodal")


# ------------------------------------------------------------------ helpers (self-contained)
def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    holders = list(prs.slide_layouts) + [prs.slide_master]
    for lay in holders:
        for sh in lay.shapes:
            if sh.has_text_frame and any(w in sh.text_frame.text for w in OLD_LAYOUT_WORDS):
                K.replace_text(sh, f"{F.TITLE}  ·  {review_label}")
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            part._blob = part.blob.replace(b"Hidden Risk", b"Health Assistant")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "multi-disease prediction, machine learning, diabetes, heart, liver, kidney"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 4
    w, h, gx, gy = 2.88, 1.5, 0.2, 0.25
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.22, y + 0.12, 1.0, 0.5, n_, size=24, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.22, y + 0.7, w - 0.4, 0.75, label, size=15, color=K.INK, bold=True)


def title_slide(s, review_label, subtitle):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle)
        elif nm == "review label":
            K.replace_text(sh, review_label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def pc(x, d=1):
    return F.pct(x + 1e-9, d)


# ------------------------------------------------------------------ build
def build(out_path):
    X, R, A = F.X, F.REAL, F.API
    D = F.DISEASES
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}
    lat = A["latency"]
    lat_p50 = lat["POST /api/predict/diabetes"]["p50_ms"]
    rps = lat["POST /api/predict/diabetes (4 concurrent clients)"]["throughput_rps"]
    fit_total = sum(X[d]["timing"][m]["fit_s"] for d in D for m in F.MODEL_ORDER)
    lr_wins = all(X[d]["seed_winners"]["Logistic Regression"] == 10 for d in D)
    assert lr_wins
    best_acc = {d: X[d]["holdout"][F.BEST[d]]["accuracy"] for d in D}

    # ---- 1 title
    title_slide(S[0], "REVIEW 2  ·  DESIGN, METHODOLOGY AND PROGRESS",
                "Design, methodology and progress: four classifiers for each of four diseases, behind one web interface")
    notes[1] = ("This review shows how the Intelligent Health Assistant is designed and built: data, models, what is implemented, the first measured results, "
                "and the open items. The training data are synthetic and we say so on every relevant slide.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap"), ("02", "Architecture"), ("03", "Methodology"), ("04", "Dataset"), ("05", "Preprocessing"), ("06", "Models and training"),
               ("07", "Decision logic"), ("08", "Data flow, data model, UML"), ("09", "Technology stack"), ("10", "Progress and results"),
               ("11", "Output screens and testing"), ("12", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    base = ("Alam et al. (IEEE Access 2026) classify thyroid status on 7,200 UCI records with five classifiers, nested cross-validation, SHAP and counterfactuals. "
            "We take the idea of comparing classifiers under held-out evaluation and widen it to four diseases; the explanation part is not yet built.")
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Patients see several specialists for separate tests, tools cover one disease each, and rural centres cannot interpret routine reports early. Late diagnosis raises cost and mortality.", K.ORANGE, "!"),
        ("Proposed system", "One web application that takes routine clinical values and returns a high / low risk class with the model's confidence for diabetes, heart, liver and kidney disease.", K.BLUE, "i"),
        ("Base paper link", base, K.GREEN, "B"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, 4.9, t, b_, col, gl, body_size=18)
    notes[3] = ("One-minute recap of Review 1. Need: " + NEED + " The base paper studies one disease with real data and explanations; this project covers four diseases "
                "on synthetic data and does not yet explain predictions.")

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 6.3, 5.6)
    K.text(s, 7.2, 1.5, 5.5, 5.3, [
        ("**Data:** generate_data.py writes four synthetic CSVs; train_models.py trains and saves models/*.pkl and training_report.json.", {"bullet": True, "space_after": 8}),
        ("**Flask layer:** app.py loads the model bundles (model, scaler, features) and field metadata; original server-side pages on port 5000.", {"bullet": True, "space_after": 8}),
        ("**API:** FastAPI (port 8303) imports app.py, so predictions are identical; 7 routes.", {"bullet": True, "space_after": 8}),
        ("**UI:** React + Vite (port 5303): prediction forms, model comparison, datasets.", {"bullet": True, "space_after": 8}),
        ("No database, no accounts, no stored patient data.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[4] = "Four layers. The FastAPI layer imports the Flask module, so no model code was rewritten for the React interface. There is no database."

    # ---- 5 methodology
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.35, [("Generate / load CSV", K.NAVY), ("Split 80/20", K.PURPLE), ("StandardScaler", K.GREEN), ("Train 4 models", K.ORANGE),
                                      ("Metrics on test", K.BLUE), ("Best by F1", K.NAVY), ("Save + serve", K.PURPLE)], size=14, gap=0.22)
    for i, (t, b_) in enumerate([("Same recipe per disease", "Four diseases, four algorithms, one split rule (stratified, seed 42) and the same settings: 16 models in total."),
                                 ("Selection on F1", "The model with the highest F1 on the 400-row test set is saved; this reuses the test set for selection, so the reported best is slightly optimistic."),
                                 ("Served unchanged", "The saved model, scaler and feature list are loaded by app.py and used by the API for every prediction.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.3, 3.85, 3.2)
        K.text(s, x + 0.25, 3.45, 3.4, 0.5, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 4.05, 3.4, 2.3, b_, size=18, color=K.TXT)
    notes[5] = "The workflow is exactly what train_models.py runs. The scaler is fitted on the training part only, so the test rows do not leak into scaling."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Dataset: four synthetic tables")
    wipe(s)
    for i, (v, l, c) in enumerate([("8,000", "rows: 2,000 for each disease", K.BLUE), ("43", "input features: 8 + 13 + 10 + 12", K.GREEN),
                                   ("8%", "of labels flipped at random", K.ORANGE), ("~92%", "accuracy ceiling set by the noise", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.08, 1.4, 2.88, 1.2, v, l, c, vsize=26, lsize=11)
    gen = {"diabetes": "glucose, BMI, age, pedigree, blood pressure", "heart": "age, sex, chest pain, BP, cholesterol, angina, ST depression, vessels, max HR",
           "liver": "bilirubin, ALP, ALT, AST, albumin, age", "kidney": "albumin, sugar, creatinine, urea, haemoglobin, glucose, gravity, BP, age"}
    rows = [["Disease", "Features", "High risk", "Risk score is built from"]]
    for d in D:
        r = X[d]
        rows.append([F.NAME[d], r["n_features"], pc(r["positive"] / r["rows"]), gen[d]])
    K.table(s, 0.6, 2.85, [1.9, 1.2, 1.3, 7.73], rows, font_size=13, row_h=0.55, first_bold=True)
    K.text(s, 0.6, 5.75, 12.13, 1.2, [
        ("**Synthetic:** generate_data.py (seed 42) draws feature values from clipped normal / exponential distributions, scores them with a hand-written linear rule, applies a sigmoid at 0.5 and flips 8% of labels.", {"bullet": True, "space_after": 5}),
        ("Split: 1,600 training and 400 test rows per disease; 0 missing cells, 0 duplicates. Results show how well the pipeline recovers a known rule, not clinical performance.", {"bullet": True}),
    ], size=14, color=K.INK)
    notes[6] = ("Be explicit: the build environment had no access to public data, so the CSVs are generated. Because labels are a linear rule plus 8% flips, "
                "no model can exceed about 92% accuracy. The README and the Datasets page of the app say the same. Feature counts per disease: 8, 13, 10 and 12.")

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Data preprocessing")
    wipe(s)
    rows = [["Step", "What the code does"],
            ["Split", "train_test_split, test size 20%, stratified on the label, random_state 42"],
            ["Scale", "StandardScaler fitted on the training rows only, applied to the test rows"],
            ["Persist", "scaler saved as models/<disease>_scaler.pkl and reused for every online prediction"],
            ["Validate (online)", "UI checks required / numeric; API rejects missing or non-finite values with HTTP 422"],
            ["Range hint", "UI warns when a value is outside the training range of that feature"]]
    K.table(s, 0.6, 1.45, [2.2, 6.0], rows, font_size=14, row_h=0.75, first_bold=True)
    K.card(s, 9.1, 1.45, 3.63, 4.5, "Not implemented", "No imputation, outlier handling, feature selection, SMOTE or class weights: the synthetic data has no missing cells or real outliers. Objective 2 is only partly met.", K.ORANGE, "!", body_size=15)
    notes[7] = ("Preprocessing is a single StandardScaler. The objective to handle missing values and outliers has nothing to act on in the synthetic data, "
                "and no imputation code was written; the supplementary real-data check imputes with the median.")

    # ---- 8 models
    s = S[7]
    set_title(s, "Models compared")
    wipe(s)
    rr = lambda m, key: (min(X[d]["holdout"][m][key] for d in D), max(X[d]["holdout"][m][key] for d in D))
    dt_lo, dt_hi = rr("Decision Tree", "accuracy")
    rf_lo, rf_hi = rr("Random Forest", "accuracy")
    lr_lo, lr_hi = rr("Logistic Regression", "accuracy")
    sv_lo, sv_hi = rr("SVM", "accuracy")
    rows = [["Model", "Settings (train_models.py)", "Test accuracy, 4 diseases"],
            ["Logistic Regression", "max_iter 1000, random_state 42", f"{pc(lr_lo)} - {pc(lr_hi)}  (served for 3 diseases)"],
            ["Decision Tree", "max_depth 6", f"{pc(dt_lo)} - {pc(dt_hi)}  (fits label noise)"],
            ["Random Forest", "200 trees, max_depth 8", f"{pc(rf_lo)} - {pc(rf_hi)}  (train 93.6 - 95.9%)"],
            ["SVM", "RBF kernel, C = 1, probability = True", f"{pc(sv_lo)} - {pc(sv_hi)}  (served for diabetes)"]]
    K.table(s, 0.6, 1.5, [3.2, 4.6, 4.33], rows, font_size=15, row_h=0.78, first_bold=True)
    K.text(s, 0.6, 5.65, 12.13, 1.2, "The four algorithms are the ones named in the abstract. No hyper-parameter search was run: the settings are fixed. The linear model is the right family here because the synthetic labels come from a linear rule.", size=15, color=K.TXT, italic=True)
    notes[8] = "Four scikit-learn models per disease. Random Forest fits the training part to 93.6-95.9% but scores less on test, which shows over-fitting of the noisy labels."

    # ---- 9 training configuration
    s = S[8]
    set_title(s, "Training configuration")
    wipe(s)
    rows = [["Setting", "Value"],
            ["Split", F.SPLIT],
            ["Features", "scaled with StandardScaler (fit on train)"],
            ["Metrics", "accuracy, precision, recall, F1 (ROC-AUC added afterwards in run_experiments.py)"],
            ["Selection rule", "highest F1 on the test set; ties are not handled specially"],
            ["Saved files", "models/<disease>_model.pkl, _scaler.pkl, _features.pkl, training_report.json"],
            ["Pinned versions", "scikit-learn 1.8.0 (older versions fail on predict_proba), pandas 3.0.2, numpy 2.4.4"],
            ["Hardware / time", f"Intel Core i5-12500, CPU only; all 16 models fit in about {fit_total:.1f} s"]]
    K.table(s, 0.6, 1.45, [2.7, 9.43], rows, font_size=15, row_h=0.68, first_bold=True)
    notes[9] = "These are exactly the settings of train_models.py. The re-run in scripts/run_experiments.py reproduces all 16 stored results to four decimals."

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.6, 1.3, 4.2, 5.7)
    K.text(s, 5.1, 1.5, 7.6, 5.3, [
        ("**Choose the model.** For each disease the algorithm with the highest F1 on the hold-out set is saved: SVM for diabetes, Logistic Regression for heart, liver and kidney.", {"bullet": True, "space_after": 10}),
        ("**Predict.** Values are scaled, then model.predict gives the class; 1 = Positive (High Risk), 0 = Negative (Low Risk).", {"bullet": True, "space_after": 10}),
        ("**Confidence.** The percentage shown is max(predict_proba) of the predicted class (50-100%), not a calibrated risk.", {"bullet": True, "space_after": 10}),
        ("**Threshold.** The 0.5 rule is used as is; a sweep from 0.20 to 0.80 was run offline only.", {"bullet": True, "space_after": 10}),
        ("**Caution.** On ten different splits Logistic Regression had the best F1 every time, so the diabetes choice of SVM is within noise.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[10] = "The decision rule is simple and fixed. The model selection is not robust for diabetes: SVM beat Logistic Regression by 0.004 F1 on one split, inside the split-to-split spread."

    # ---- 11 data flow
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 6.4, 5.6)
    rows = [["Module", "Role"], ["generate_data.py", "writes the four synthetic CSVs"], ["train_models.py", "split, scale, train, evaluate, save"],
            ["app.py", "load bundles, field metadata, Flask pages"], ["backend/main.py", "7 JSON routes over app.py"],
            ["frontend/src/pages", "Landing, Overview, Predict, Models, Datasets, About"], ["scripts/", "smoke, API, experiment and screenshot scripts"]]
    K.table(s, 7.3, 1.45, [2.4, 3.03], rows, font_size=12, row_h=0.62)
    notes[11] = "Level-1 data flow. Level 0 and level 2 are in the report. The module table lists where each step of the flow lives."

    # ---- 12 data model and UML
    s = S[11]
    set_title(s, "Database and UML")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.6, 1.35, 6.0, 4.7)
    K.picture(s, F.diag("06_uml_class_diagram.png"), 6.85, 1.35, 5.9, 4.7)
    K.text(s, 0.6, 6.2, 12.13, 0.8, "There is no database: data live in CSV files, model bundles in pickle files and the report in JSON. The ER diagram is the logical data model; the code is function-oriented, so the UML shows modules and data holders rather than many classes.", size=13, color=K.TXT, align="c", italic=True)
    notes[12] = "State plainly that there is no database and no stored patient data. Predictions are computed and returned; nothing is kept."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    rows = [["Layer", "Technology"],
            ["Machine learning", "scikit-learn 1.8.0 (LR, DT, RF, SVC, StandardScaler), pandas 3.0.2, numpy 2.4.4, joblib 1.5.3"],
            ["Original pages", "Flask 3.1.3 with HTML / CSS templates (port 5000)"],
            ["API", "FastAPI 0.110.0, uvicorn 0.29.0, pydantic request model, CORS for localhost:5303"],
            ["Front end", "React 19, Vite 8, React Router 7, Recharts 3, framer-motion, three.js hero scene"],
            ["Testing and capture", "scripts/smoke_test.py, scripts/api_tests.py, Playwright for screenshots"],
            ["Run", "run.bat builds the venv and starts ports 8303 / 5303 (Python 3.11.9, Node 24, Windows 11)"]]
    K.table(s, 0.6, 1.5, [2.7, 9.43], rows, font_size=15, row_h=0.72, first_bold=True)
    notes[13] = "Everything is open source and runs on one machine; run.bat starts the stack. scikit-learn is pinned because the saved models need it."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Synthetic data generation, four CSVs", "Done", "generate_data.py, 8,000 rows"],
            ["Training and selection, 16 models", "Done", "train_models.py; reproduced to 4 d.p."],
            ["Flask pages and FastAPI layer (7 routes)", "Done", "smoke test 27 checks, API tests 41 / 41"],
            ["React interface, four forms, models, datasets", "Done", "20 real screenshots"],
            ["ROC-AUC, cross-validation, threshold analysis", "Offline only", "scripts/run_experiments.py"],
            ["Imputation, outlier handling, feature selection", "Not implemented", "no missing values in the synthetic data"],
            ["Per-prediction explanations, recommendations", "Not implemented", "objective 6 only partly met"]]
    K.table(s, 0.6, 1.5, [5.6, 2.1, 4.43], rows, font_size=14, row_h=0.66)
    notes[14] = "State of the code base. The gaps are imputation / outlier handling, explanations and recommendations; ROC-AUC and cross-validation exist only as offline scripts."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: 400-row test sets")
    wipe(s)
    sel = [(F.NAME[d], F.BEST[d], pc(best_acc[d]), f"{F.HO[d][F.BEST[d]]['f1']:.3f}") for d in D]
    for i, (v, l, c) in enumerate([(pc(best_acc["diabetes"]), "diabetes, SVM", K.BLUE), (pc(best_acc["heart"]), "heart, Logistic Regression", K.GREEN),
                                   (pc(best_acc["liver"]), "liver, Logistic Regression", K.ORANGE), (pc(best_acc["kidney"]), "kidney, Logistic Regression", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.08, 1.4, 2.88, 1.15, v, f"accuracy: {l}", c, vsize=26, lsize=11)
    K.bar_chart(s, 0.6, 2.75, 7.7, 4.1, [F.NAME[d] for d in D],
                [(F.SHORT[m], [F.HO[d][m]["f1"] for d in D]) for m in F.MODEL_ORDER],
                colors=[K.BLUE, K.GOLD, K.GREEN, K.PURPLE], number_format="0.00", font_size=11, max_value=1.0)
    K.card(s, 8.55, 2.8, 4.18, 4.05, "Reading", "F1 by algorithm. Logistic Regression has the best F1 on 10 of 10 random splits for all four diseases. Synthetic data: ceiling about 92%. On real public data the best algorithm scores 72-99% and the shipped models transfer only partly.", K.BLUE, "i", body_size=13)
    notes[15] = ("Hold-out accuracy of the served models: diabetes SVM 0.8950, heart LR 0.8850, liver LR 0.9175, kidney LR 0.8825. F1 values: " +
                 "; ".join(f"{a}: {b} F1 {c}" for a, b, _, c in sel) +
                 ". These are results on synthetic data. The supplementary real-data check gives 76.5% (Pima), 84.3% (Cleveland), 71.8% (ILPD), 98.8% (CKD) for the best algorithm.")

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    shots = [("05_predict-diabetes.png", "Diabetes: High Risk, SVM"), ("13_predict-diabetes-low-risk.png", "Diabetes: Low Risk"),
             ("17_validation-error.png", "Empty form: required fields"), ("09_models.png", "Model comparison"),
             ("10_datasets.png", "Datasets and synthetic warning"), ("20_predict-heart-false-positive.png", "Failure case: label 0, model says High Risk")]
    for i, (f, cap) in enumerate(shots):
        r, c = divmod(i, 3)
        x, y = 0.6 + c * 4.14, 1.3 + r * 2.85
        K.picture(s, F.shot(f), x, y, 3.95, 2.45, border=True)
        K.text(s, x, y + 2.48, 3.95, 0.3, cap, size=12, color=K.TXT, align="c")
    notes[16] = "Real screenshots of the running application with records drawn from the training CSVs. The last one is a genuine false positive: a label-0 heart record the model calls High Risk with 99.53% confidence."

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([("27", "smoke-test checks passed", K.GREEN), (f"{F.API_CHECKS[0]} / {F.API_CHECKS[1]}", "API checks: functional, consistency, validation", K.BLUE),
                                   (f"{lat_p50:.0f} ms", "median predict latency, 127.0.0.1", K.ORANGE), (f"{rps:.0f} / s", "requests per second, 4 clients", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.08, 1.4, 2.88, 1.2, v, l, c, vsize=26, lsize=11)
    rows = [["Observation", "Action"],
            ["API accepts impossible values (glucose 100,000 mg/dL gives High Risk 60.94%; age -5 is accepted)", "Only a UI warning exists; server range validation is a next step"],
            ["Missing, NaN, infinite, text and null inputs, unknown disease, malformed JSON", "Rejected with 422 / 404 / 405 as designed"],
            ["API prediction equals direct model.predict on 40 sampled rows", "Consistency with the Flask pages confirmed"]]
    K.table(s, 0.6, 3.0, [6.7, 5.43], rows, font_size=13, row_h=0.85)
    K.text(s, 0.6, 6.5, 12.13, 0.5, "No pytest suite, load test beyond 200 calls, security test or user study has been done. Timings were taken on a busy PC.", size=12, color=K.TXT, italic=True)
    notes[17] = ("Tests run against the live API. Latency: p50 about 13 ms, p95 about 25 ms; Windows timer resolution (about 15.6 ms) quantises single calls, "
                 "so treat the numbers as upper bounds.")

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["No public clinical data in the build environment", "Synthetic generator; a separate real-data check (Pima, UCI Cleveland, ILPD, CKD) was added for the documents"],
            ["Random Forest over-fits the noisy labels (train 93.6-95.9%, test 81.5-88.3%)", "Depth limited to 8; stated as a finding; cross-validation reported"],
            ["Model chosen on the test set with one split", "5-fold CV and 10 splits show Logistic Regression best; the shipped diabetes SVM is not changed"],
            ["Shipped models transfer partly to real records (liver 54.0% vs 71.4% majority baseline)", "Documented as a limitation; retraining on real data is future work"],
            ["Saved models need the same scikit-learn version", "requirements.txt pins scikit-learn 1.8.0"]]
    K.table(s, 0.6, 1.45, [6.0, 6.13], rows, font_size=13, row_h=0.8)
    notes[18] = "Each challenge was met either with a change in how results are reported or, where the code was not changed, with a clear statement of the limit."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Evaluate", "Report ROC curves, confusion matrices, cross-validation and threshold analysis in the final results.", K.BLUE, "1"),
        ("Explain", "Show the driving parameters per prediction (objective 6); compare with SHAP and counterfactuals of the base paper.", K.GREEN, "2"),
        ("Harden", "Range validation on the server, imputation and outlier handling, retraining on real public data.", K.ORANGE, "3"),
        ("Document", "Final report, IEEE paper and final deck with every figure regenerated from the JSON files.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.08, 1.6, 2.88, 4.6, t, b_, col, gl, body_size=16)
    notes[19] = "Plan to the final review. Real-data retraining and explanations are the two items that would most change the project's value."

    # ---- 20 thank you
    thank_you(S[19])
    notes[20] = "Likely questions: why synthetic data; why Logistic Regression is best here; why SVM is shipped for diabetes; why there are no explanations yet."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx")))
