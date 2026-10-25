"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.FIGS, "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                os.makedirs(F.FIGS, exist_ok=True)
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=22, bold=True, space_after=4)
    d.p(f"({F.SUBTITLE})", align="center", size=14, bold=True, space_after=8)
    d.p("A risk-screening web application: Intelligent Health Assistant", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: [PLACE]", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: [PLACE]", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the related work we studied, the maintainers of the public UCI Machine Learning Repository and of the Pima diabetes data used in our supplementary check, "
        "and the developers of the open-source software used in this project (Python, scikit-learn, pandas, NumPy, Flask, FastAPI, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    X, R = F.X, F.REAL
    acc = {k: F.pct(X[k]["holdout"][F.BEST[k]]["accuracy"]) for k in F.DISEASES}
    d.h1("ABSTRACT")
    d.p("Chronic diseases such as diabetes, heart disease, liver disease and kidney disease are often found late, and the tools that exist usually screen for one condition at a time. "
        "This project builds an Intelligent Health Assistant that screens four conditions from routine clinical values through one web interface. "
        "For each condition, Logistic Regression, a Decision Tree, a Random Forest and an RBF Support Vector Machine are trained on standardised inputs, compared on a stratified 80/20 hold-out split with accuracy, precision, recall and F1-score, "
        "and the model with the highest F1 is saved and served. The models are exposed through a FastAPI service with seven routes and a React interface with a prediction form for each disease, "
        "a model-comparison page and a dataset page; an earlier Flask version of the same models is kept.")
    d.p(f"The four training datasets (2,000 rows each) are synthetic: they were generated from a hand-written linear risk rule with 8% of the labels flipped, because the original build environment had no access to public data. "
        f"On the 400-row test sets the selected models reach accuracies of {acc['diabetes']} (diabetes, SVM), {acc['heart']} (heart, Logistic Regression), {acc['liver']} (liver, Logistic Regression) and {acc['kidney']} (kidney, Logistic Regression), "
        "close to the 92% ceiling that the label noise allows; they measure how well the pipeline recovers a known rule, not clinical performance. "
        "Cross-validation and ten random splits show that Logistic Regression is the best model for every disease, including diabetes where SVM was shipped, and that Random Forest over-fits the noisy labels.")
    d.p("A supplementary check on real public data (Pima diabetes, UCI Cleveland heart, UCI ILPD liver and UCI chronic kidney disease) gives accuracies of 77%, 84%, 72% and 99% for the best algorithm, and shows that the shipped models transfer only partly "
        "(for liver the accuracy of 54% is below the 71% majority-class baseline). The application does not show per-prediction explanations, does not impute missing values and has no clinical range validation; "
        f"the API passed {F.API_CHECKS[0]} of {F.API_CHECKS[1]} functional, validation and consistency checks. The system is a demonstration of a multi-disease screening pipeline and not a diagnostic tool.")
    d.p("**Keywords:** " + F.KEYWORDS + ".", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("ALT / AST", "Alanine / Aspartate Aminotransferase"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"),
    ("BMI", "Body Mass Index"), ("CKD", "Chronic Kidney Disease"), ("CORS", "Cross-Origin Resource Sharing"), ("CSV", "Comma-Separated Values"),
    ("CV", "Cross-Validation"), ("DFD", "Data Flow Diagram"), ("DT", "Decision Tree"), ("ECG", "Electrocardiogram"), ("ER", "Entity-Relationship"),
    ("FN / FP", "False Negative / False Positive"), ("HTTP", "Hypertext Transfer Protocol"), ("ILPD", "Indian Liver Patient Dataset"), ("JSON", "JavaScript Object Notation"),
    ("LR", "Logistic Regression"), ("ML", "Machine Learning"), ("RBF", "Radial Basis Function"), ("RF", "Random Forest"), ("REST", "Representational State Transfer"),
    ("ROC", "Receiver Operating Characteristic"), ("SHAP", "SHapley Additive exPlanations"), ("SMOTE", "Synthetic Minority Over-sampling Technique"),
    ("SPA", "Single-Page Application"), ("SVM", "Support Vector Machine"), ("TN / TP", "True Negative / True Positive"), ("UCI", "University of California, Irvine (Machine Learning Repository)"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("XAI", "Explainable Artificial Intelligence"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.6], size=11)
