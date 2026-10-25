"""
Layout kit for 19.IEEE_Paper.docx (two-column IEEE conference style, format per the documentation prompt).
Citations are written as {r:key1,key2} (keys of refs.NUM) and numbered by order of first appearance.
"""

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .kit import add_runs, set_font
from .refs import NUM, REFS

REF_BY_KEY = {k: REFS[n - 1] for k, n in NUM.items()}
FONT = "Times New Roman"
COL_W = 3.4  # inches


class Paper:
    def __init__(self):
        self.doc = Document()
        self.order = []          # reference keys in order of first citation
        self.fig_no = 0
        self.tab_no = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
        s.left_margin = s.right_margin = Inches(0.69)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0

    # ---------------------------------------------------------------- citations
    def cite(self, *keys):
        nums = []
        for k in keys:
            if k not in REF_BY_KEY:
                raise KeyError(k)
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        nums.sort()
        return ", ".join(f"[{n}]" for n in nums)

    def expand(self, text):
        text = text.replace("`", "")   # Times New Roman only: file names are set in the body font
        return re.sub(r"\{r:([a-z0-9,]+)\}", lambda m: self.cite(*m.group(1).split(",")), text)

    # ---------------------------------------------------------------- blocks
    def para(self, text, size=10, align="j", indent=0.2, bold=None, italic=None, after=0, before=0, keep_next=False):
        par = self.doc.add_paragraph()
        par.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        pf = par.paragraph_format
        pf.first_line_indent = Inches(indent) if indent else None
        pf.space_after, pf.space_before = Pt(after), Pt(before)
        pf.keep_with_next = keep_next
        add_runs(par, self.expand(text), size, bold, italic)
        return par

    def heading(self, roman, text):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.keep_with_next = True
        label = f"{roman}. {text.upper()}" if roman else text.upper()
        set_font(par.add_run(label), FONT, 10, True)

    def sub(self, letter, text):
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.keep_with_next = True
        set_font(par.add_run(f"{letter}. {text}"), FONT, 10, False, True)

    def figure(self, path, caption, width=COL_W):
        self.fig_no += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(caption) > 70 else WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(5)
        add_runs(cap, f"Fig. {self.fig_no}. {self.expand(caption)}", 8)

    def table(self, caption, header, rows, widths, size=7, align=None):
        self.tab_no += 1
        roman = ["I", "II", "III", "IV", "V", "VI"][self.tab_no - 1]
        c1 = self.doc.add_paragraph()
        c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c1.paragraph_format.space_before = Pt(4)
        c1.paragraph_format.keep_with_next = True
        add_runs(c1, f"TABLE {roman}", 8)
        c2 = self.doc.add_paragraph()
        c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c2.paragraph_format.space_after = Pt(2)
        c2.paragraph_format.keep_with_next = True
        add_runs(c2, self.expand(caption).upper(), 8)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, h, size, True)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (align is None or align[i] == "l") else WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, self.expand(str(v)), size)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(cs)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        set_font(sp.add_run(""), FONT, 4)

    def references(self):
        self.heading("", "References")
        for i, k in enumerate(self.order, 1):
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = par.paragraph_format
            pf.left_indent = Inches(0.3)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(1)
            add_runs(par, f"[{i}]  {REF_BY_KEY[k]}", 8)

    # ---------------------------------------------------------------- layout
    def begin_two_columns(self):
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.xpath("./w:cols")
        if cols:
            c = cols[0]
        else:
            c = OxmlElement("w:cols")
            sec._sectPr.append(c)
        c.set(qn("w:num"), "2")
        c.set(qn("w:space"), "360")  # 0.25 in

    def title_block(self, title, authors):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(8)
        set_font(par.add_run(title), FONT, 24, False)
        n = len(authors)
        t = self.doc.add_table(rows=1, cols=n)   # no style: borderless
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, a in enumerate(authors):
            c = t.rows[0].cells[i]
            c.width = Inches(7.12 / n)
            c.text = ""
            for j, line in enumerate(a):
                p = c.paragraphs[0] if j == 0 else c.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                add_runs(p, line, 9, j == 0)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def abstract(self, text, keywords):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.space_after = Pt(4)
        set_font(par.add_run("Abstract---"), FONT, 9, True, True)
        add_runs(par, text, 9)
        par2 = self.doc.add_paragraph()
        par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par2.paragraph_format.space_after = Pt(4)
        set_font(par2.add_run("Index Terms---"), FONT, 9, True, True)
        add_runs(par2, ", ".join(sorted(keywords, key=str.lower)) + ".", 9)

    def save(self, path):
        self.doc.save(path)
        return path
