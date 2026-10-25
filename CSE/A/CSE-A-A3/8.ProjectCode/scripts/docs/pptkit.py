"""
python-pptx helpers. The decks are built *from* the institution templates: every template slide is kept
(so slide count, order, layouts, master and banner are inherited) and its content shapes are replaced.
Removed pictures and charts also lose their relationships, so no template data stays in the package.
"""

import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1B, 0x2A, 0x41)
INK = RGBColor(0x0F, 0x17, 0x2A)
TXT = RGBColor(0x47, 0x55, 0x69)
CARD = RGBColor(0xF6, 0xF8, 0xFB)
LINE = RGBColor(0xD5, 0xDC, 0xE6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
GREEN = RGBColor(0x1B, 0xAF, 0x7A)
PURPLE = RGBColor(0x6B, 0x52, 0xD6)
GOLD = RGBColor(0xE8, 0xA3, 0x17)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FNAVY = RGBColor(0x1F, 0x38, 0x64)
FBLUE = RGBColor(0x2E, 0x75, 0xB6)
FHEAD = RGBColor(0xDE, 0xEA, 0xF6)
FZEBRA = RGBColor(0xF2, 0xF2, 0xF2)
DARK = RGBColor(0x22, 0x22, 0x22)

IN = 914400


def E(v):
    return Emu(int(v * IN))


# ------------------------------------------------------------------ cleanup
def drop_shape(slide, shape):
    el = shape._element
    for rid in el.xpath(".//@r:embed | .//@r:id | .//@r:link"):
        try:
            slide.part.drop_rel(rid)
        except KeyError:
            pass
    el.getparent().remove(el)


def clear(slide, keep):
    """Remove every shape for which keep(shape) is False."""
    for sh in list(slide.shapes):
        if not keep(sh):
            drop_shape(slide, sh)


def replace_text(shape, text, size=None):
    """Replace the text of a shape but keep the first run's formatting (font, colour, size)."""
    tf = shape.text_frame
    lines = text.split("\n")
    p0 = tf.paragraphs[0]
    proto = copy.deepcopy(p0._p)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    if p0.runs:
        p0.runs[0].text = lines[0]
        if size:
            p0.runs[0].font.size = Pt(size)
    else:
        p0.text = lines[0]
    for extra in lines[1:]:
        newp = copy.deepcopy(proto)
        tf._txBody.append(newp)
        from pptx.text.text import _Paragraph
        para = _Paragraph(newp, tf)
        runs = para.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        if runs:
            runs[0].text = extra
            if size:
                runs[0].font.size = Pt(size)


def set_notes(slide, text):
    ns = slide.notes_slide
    ns.notes_text_frame.text = text


# ------------------------------------------------------------------ primitives
def _font(run, size=None, bold=None, italic=None, color=None, name=None):
    f = run.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if name:
        f.name = name
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", name)


def box(slide, x, y, w, h, fill=CARD, line=LINE, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03):
    s = slide.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, paras, size=16, color=TXT, bold=False, align="l", anchor="t", font=None,
         bullets=False, space_after=4, line_spacing=None, italic=False):
    """paras: str | list of str | list of (str, dict(size,bold,color,italic,bullet))."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for item in paras:
        txt, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[opt.get("align", align)]
        p.space_after = Pt(opt.get("space_after", space_after))
        if line_spacing:
            p.line_spacing = line_spacing
        bl = opt.get("bullet", bullets)
        if bl:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(int(0.28 * IN)))
            ppr.set("indent", str(int(-0.22 * IN)))
            for tag in ("a:buNone", "a:buChar"):
                el = ppr.find(qn(tag))
                if el is not None:
                    ppr.remove(el)
            bu = etree.SubElement(ppr, qn("a:buChar"))
            bu.set("char", "•")
        # inline **bold** segments
        parts = []
        s_ = txt
        while "**" in s_:
            a, _, rest = s_.partition("**")
            b, _, s_ = rest.partition("**")
            parts.append((a, False))
            parts.append((b, True))
        parts.append((s_, False))
        for seg, bd in parts:
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            _font(r, opt.get("size", size), True if bd else opt.get("bold", bold), opt.get("italic", italic), opt.get("color", color), font)
    return tb


def circle(slide, x, y, d, color, glyph="", size=16, font=None):
    s = box(slide, x, y, d, d, fill=color, line=color, shape=MSO_SHAPE.OVAL)
    if glyph:
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = glyph
        _font(r, size, True, None, WHITE, font)
    return s


def card(slide, x, y, w, h, title, body, color=BLUE, glyph="", body_size=15, title_size=16, font=None, bullets=False):
    box(slide, x, y, w, h)
    circle(slide, x + 0.22, y + 0.2, 0.55, color, glyph, 16, font)
    text(slide, x + 0.92, y + 0.17, w - 1.1, 0.6, title, size=title_size, color=INK, bold=True, anchor="m", font=font)
    text(slide, x + 0.22, y + 0.95, w - 0.44, h - 1.1, body, size=body_size, color=TXT, bullets=bullets, font=font)


def tile(slide, x, y, w, h, value, label, color=BLUE, vsize=30, lsize=12, fill=CARD, font=None):
    box(slide, x, y, w, h, fill=fill)
    box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + 0.22, y + 0.1, w - 0.3, h * 0.55, value, size=vsize, color=INK, bold=True, anchor="m", font=font)
    text(slide, x + 0.22, y + h * 0.58, w - 0.3, h * 0.38, label, size=lsize, color=TXT, anchor="t", font=font)


def arrow_right(slide, x, y, w=0.26, h=0.26, color=RGBColor(0x64, 0x74, 0x8B)):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, E(x), E(y), E(w), E(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def flow(slide, x, y, w, h, steps, size=14, gap=0.28, font=None):
    """Row of coloured rounded boxes joined by arrows. steps: [(label, colour)]"""
    n = len(steps)
    bw = (w - gap * (n - 1)) / n
    for i, (label, col) in enumerate(steps):
        bx = x + i * (bw + gap)
        b = box(slide, bx, y, bw, h, fill=col, line=None, radius=0.08)
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        _font(r, size, True, None, WHITE, font)
        if i < n - 1:
            arrow_right(slide, bx + bw + (gap - 0.22) / 2, y + h / 2 - 0.13, 0.22, 0.26)


def picture(slide, path, x, y, w, h, border=False):
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(path, E(px), E(py), E(pw), E(ph))
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def table(slide, x, y, widths, rows, font_size=12, row_h=0.4, header_fill=NAVY, header_color=WHITE, zebra=CARD,
          body_color=INK, first_bold=False, font=None, align=None, header_size=None, border=LINE):
    nrows, ncols = len(rows), len(widths)
    gs = slide.shapes.add_table(nrows, ncols, E(x), E(y), E(sum(widths)), E(row_h * nrows))
    tbl = gs.table
    # strip the default style banding so our fills decide the look
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        if tblPr.get(a) is not None:
            tblPr.set(a, "0")
    for c, w in enumerate(widths):
        tbl.columns[c].width = E(w)
    for r in range(nrows):
        tbl.rows[r].height = E(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            val = str(rows[r][c])
            p = tf.paragraphs[0]
            a = "c" if r == 0 else (align[c] if align else "l")
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[a]
            run = p.add_run()
            run.text = val
            if r == 0:
                _font(run, header_size or font_size, True, None, header_color, font)
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            else:
                _font(run, font_size, first_bold and c == 0, None, body_color, font)
                cell.fill.solid()
                cell.fill.fore_color.rgb = zebra if (r % 2 == 0) else WHITE
            _cell_border(cell, border)
    return gs


def _cell_border(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.find(qn(tag))
        if ln is not None:
            tcPr.remove(ln)
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", "9525")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr"))
        clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    # schema order: borders must precede the fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill)
        tcPr.append(fill)


def bar_chart(slide, x, y, w, h, categories, series, colors=None, number_format='0.0%', font_size=12, legend=True, max_value=None, gap=60, stacked=False, label_size=None):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(x), E(y), E(w), E(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font_size)
    ch.has_title = False
    ch.has_legend = legend and len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    if max_value:
        va.maximum_scale = max_value
    va.minimum_scale = 0
    ch.category_axis.format.line.color.rgb = LINE
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size or font_size - 1)
    cols = colors or [BLUE, ORANGE, GREEN, PURPLE, GOLD]
    for s, c in zip(plot.series, cols):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c
    return gf


# ------------------------------------------------------------------ package hygiene
def strip_orphans_check(prs):
    """Names of media / chart parts reachable from the presentation (used by the verifier)."""
    names = set()
    for sl in prs.slides:
        for rel in sl.part.rels.values():
            names.add(rel.target_part.partname)
    return names
