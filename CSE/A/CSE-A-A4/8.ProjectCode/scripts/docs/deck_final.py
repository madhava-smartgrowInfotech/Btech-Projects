"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .lit import EXISTING_SYSTEM, GAPS, LIT_TABLE, NEED, OBJECTIVES, PROBLEM_STATEMENT, PROPOSED_APPROACH
from .refs import REFS, n as refno

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def cap(slide, x, y, w, h, text, size=12):
    return ftext(slide, x, y, w, h, text, size=size, color=K.FNAVY, italic=True)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.comments = ""
    cp.subject = f"{F.PRODUCT}, Batch {F.BATCH_NO} ({F.SECTION})"
    cp.keywords = "credit card fraud, restricted Boltzmann machine, class imbalance, machine learning"
    cp.last_modified_by = ""
    notes = {}
    sv = F.SAVED
    pcd, sas, qs, ens = sv["classical"], sv["sa"], sv["qa"], sv["ensemble"]
    bb = F.BASE_BAL
    srv = F.SERVING
    asv, scl = srv["as_served_no_scaling"]["ensemble"], srv["with_scaling"]["ensemble"]
    LAT = F.LATENCY
    lat1 = LAT["predict_single_fraud_example"]
    latb = LAT["predict_batch_20_rows"]
    n_test = F.SPLIT["test"][0]
    n_train = F.SPLIT["train"][0]
    ntest_fraud = F.SPLIT["test_fraud"]

    # footers on every slide except the title
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.FULL_TITLE, 30)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: SecurePay AI, a fraud-screening application built around the base paper's restricted Boltzmann machine classifier, and measured honestly."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Models\nDecision Workflow, Interaction & Storage\nImplementation Highlights\nResults: Shipped Models, Threshold, Baselines, Imbalance\nPerformance, Testing & the Serving Defect\nComparison with the Base Paper\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("Card fraud is rare, costly and changing: fixed rules are learned and bypassed, and a model trained naively on imbalanced data scores high accuracy while finding almost no fraud.", {"bullet": True, "space_after": 8}),
        ("SecurePay AI follows the base paper: a Gaussian-Bernoulli RBM classifier trained three ways (PCD, simulated-annealing sampling, quantum-inspired sampling), averaged into one score with approve / challenge / decline bands, behind a Flask API and a React interface.", {"bullet": True, "space_after": 8}),
        (f"Synthetic 39,000-transaction data, not the base paper's real data. On {n_test:,} held-out rows the ensemble reaches precision {F.pct(ens['precision'])}, recall {F.pct(ens['recall'])}, F1 {F.pct(ens['f1'])}, ROC-AUC {F.f3(ens['roc_auc'])}; the shipped PCD model flags nearly everything (accuracy {F.pct(pcd['accuracy'])}).", {"bullet": True, "space_after": 8}),
        (f"Logistic regression, random forest and XGBoost score F1 {F.pct(bb['Logistic Regression']['f1'])}, {F.pct(bb['Random Forest']['f1'])}, {F.pct(bb['XGBoost']['f1'])} on the same split: the simple baselines are better here.", {"bullet": True, "space_after": 8}),
        (f"Audit finding: the bulk route skips feature scaling, so accuracy on a fresh batch is {F.pct(asv['accuracy'])} as served against {F.pct(scl['accuracy'])} when scaled. A single check takes a median {lat1['median_ms']:.0f} ms.", {"bullet": True, "space_after": 8}),
        ("Not done: real transaction data, quantum hardware, drift, a user study.", {"bullet": True}),
    ], size=16)
    notes[3] = "Give the problem, the system, then the numbers with their conditions. End on the honest limits: synthetic data, baselines ahead, and the serving defect found by our own audit."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 0.75, F.BASE_PAPER, size=13, color=K.FNAVY)
    bt = F.BASE_TABLE
    rows = [["Aspect", "Neto, do Amaral and Temporao, IEEE Access 2026"],
            ["Method", "Gaussian RBM classifier (visible, hidden and label layers); energy written as a QUBO / Ising problem; three trainers: classical PCD, simulated-annealing sampling, D-Wave quantum sampling"],
            ["Data", "Real Stone fintech data: 145 million transactions, about 19 thousand fraud; manual undersampling to 39 thousand (50/50); 80/20 split; 50 epochs"],
            ["Table 1 (acc / prec / rec / F1, time)", "PCD 84 / 86 / 81 / 83 %, 2 min;  SAS 85 / 94 / 82 / 88 %, 4 h;  QS 85 / 91 / 79 / 84 %, 1.8 h"],
            ["Pre-processing", "One-hot, z-score, correlation filter, CatBoost importance; SMOTE and Tomek links performed poorly at this imbalance"],
            ["Stated limits", "Whether the advantage is general is an open question; SA metrics oscillate; hours of training; no interface or public baselines"]]
    ftable(s, 0.55, 2.6, [2.6, 9.6], rows, size=13, row_h=0.66)
    notes[4] = "Read from the PDF, including Table 1 as an image. We adopt the model family, the preprocessing recipe and the three trainers; we cannot reproduce the real data, the D-Wave runs or the hour-scale training."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Data / reported result", "Limitation"]]
    for r in LIT_TABLE:
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3], r[4]])
    ftable(s, 0.55, 1.75, [1.9, 0.6, 3.2, 3.5, 3.0], rows, size=10, row_h=0.58, align=["l", "c", "l", "l", "l"])
    notes[5] = "Seven verified works. Three further titles from the batch survey could not be verified and are flagged in the report's reference list, not cited."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in GAPS], size=20)
    notes[6] = "Six gaps from the survey. We address the compared pipeline, the recall-first evaluation, the open implementation with an interface, and the behaviour when fraud falls to 1 percent."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.5)
    ftext(s, 0.85, 1.9, 11.6, 2.4, PROBLEM_STATEMENT + [("**Need:** " + NEED, {"bullet": False})], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "A transaction (amount, merchant, channel, card present, country match, time, history signals) or a CSV of rows"),
                                 ("Output", "Fraud probability, label, risk band and per-model scores; bulk metrics when a Class column is present"),
                                 ("Record", "Downloadable predictions CSV, saved models and per-epoch histories")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.6, 3.9, 1.9, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.7, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.2, 3.5, 1.2, b_, size=14)
    notes[7] = "Four statements of the problem and the need, then what goes in and out of the system."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion states the status of each: 3 and 4 were met in the experiment script, not in the product."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 1.9, 12.2, 1.0, [("Transaction / CSV", K.FNAVY), ("Pre-process", K.FBLUE), ("3 RBM scorers", K.GREEN), ("Ensemble mean", K.PURPLE), ("Risk band", K.ORANGE)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 3.7, [(p, {"bullet": True, "space_after": 9}) for p in PROPOSED_APPROACH], size=16)
    notes[9] = "The proposed system as a line of boxes and the five design rules behind it."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.75, 6.6, 5.15)
    ftext(s, 7.4, 1.95, 5.4, 4.9, [
        ("**Presentation:** React and Vite on port 5304; every call goes through a Vite proxy to the API.", {"bullet": True, "space_after": 9}),
        ("**Application:** Flask on port 8304, 16 routes, JSON and multipart CSV, 500 responses carry an error message.", {"bullet": True, "space_after": 9}),
        ("**Model tier:** preprocessing, GRBMC, samplers, training loop, evaluation, synthetic data generator.", {"bullet": True, "space_after": 9}),
        ("**Data tier:** local files only (CSV, pickles, JSON); there is no SQL database.", {"bullet": True, "space_after": 9}),
        ("In-memory store holds the dataset, pipeline, models and histories after start-up.", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. The data tier is files, not a database; the store dictionary in app.py caches everything in memory."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.85, 12.2, 1.1, [("One-hot + z-score", K.FNAVY), ("Correlation filter", K.FBLUE), ("Undersample 50/50", K.GREEN), ("RF top-18", K.ORANGE), ("80/20 split", K.PURPLE), ("Train GRBMC", K.FBLUE), ("Evaluate", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        (f"Pre-processing as in the base paper: the backend rebuilds it from an 8,000-row sample; {n_train:,} training and {n_test:,} test rows; no column exceeded the 0.95 correlation limit.", {"bullet": True, "space_after": 8}),
        ("Training: 20 epochs; PCD with 200 hidden units and batch 512; SAS and QS with 65 hidden units, batch 32, 10 sweeps x 16 reads, 3,000 rows per epoch.", {"bullet": True, "space_after": 8}),
        ("QS is a classical simulation of a transverse-field tunnelling step, so it is quantum-inspired only; no quantum hardware was used.", {"bullet": True, "space_after": 8}),
        ("Added measurements: baselines (LR, RF, XGBoost), five seeds, threshold sweep, a fresh 2,000-row batch, a 1 % fraud variant with SMOTE, class weights and undersampling.", {"bullet": True, "space_after": 8}),
        ("Every number comes from scripts/run_experiments.py, written to experiments/metrics.json.", {"bullet": True}),
    ], size=15)
    notes[11] = "The pipeline and what we added around it: baselines, seeds, thresholds, imbalance. The history files choose a best epoch on the test split, so we report final-epoch figures as primary."

    # ---- 12 dataset, features, models
    s = S[11]
    set_title(s, "Dataset & Models")
    wipe(s)
    d = F.DATASET
    for i, (v, l, c) in enumerate([(f"{d['rows']:,}", "synthetic transactions", K.FBLUE), ("50 / 50", f"{d['fraud']:,} fraud, {d['legit']:,} legitimate", K.GREEN), ("18", "features kept by RF importance", K.ORANGE), (f"{n_train:,} / {n_test:,}", "train / test rows (80/20)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f01_dataset_overview.png"), 0.55, 3.0, 7.4, 1.9)
    rows = [["Trainer", "Hidden", "Batch", "Sampler"],
            ["PCD", str(pcd["n_hidden"]), str(pcd["batch_size"]), "Gibbs chain"],
            ["SAS", str(sas["n_hidden"]), str(sas["batch_size"]), "annealing sweeps"],
            ["QS", str(qs["n_hidden"]), str(qs["batch_size"]), "tunnelling flips"]]
    ftable(s, 8.2, 3.0, [1.0, 0.9, 0.9, 1.75], rows, size=12, row_h=0.45, align=["l", "c", "c", "l"])
    ftext(s, 0.55, 5.1, 12.2, 1.8, [
        ("The data is synthetic: generated by backend/synthetic_data.py to imitate the base paper's balanced 39 thousand set, so it is not comparable with real Stone data. Fraud is generated with larger amounts, online and card-not-present channels and shifted V3, V7, V10, V14, V17.", {"bullet": True, "space_after": 6}),
        (f"The backend works from an 8,000-row random sample of this file ({n_train:,} train, {n_test:,} test rows); every headline number uses that split unless stated.", {"bullet": True, "space_after": 6}),
        (f"Mean amount: fraud {d['amount_mean_fraud']:.0f}, legitimate {d['amount_mean_legit']:.0f}. Feature names V1 to V22 are generic; the UI's friendly labels for V3, V14, V17 are not real semantics.", {"bullet": True}),
    ], size=14)
    notes[12] = "Be frank: synthetic, balanced, and built so that amount, channel and a few V columns carry the signal. That is why simple baselines do so well."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("11_flowchart_decision.png"), 0.55, 1.75, 4.5, 5.15)
    for i, (v, l, c) in enumerate([("< 40 %", "LOW: approve", K.GREEN), ("40 - 70 %", "MEDIUM: challenge (3-D Secure)", K.ORANGE), ("> 70 %", "HIGH: decline and alert", K.RGBColor(0xB4, 0x36, 0x2E))]):
        ftile(s, 5.3 + i * 2.5, 1.85, 2.35, 1.05, v, l, c, 22)
    ftext(s, 5.3, 3.2, 7.45, 3.7, [
        ("Missing fields default to the dataset mean or mode; categorical fields are one-hot encoded; absent columns are filled with 0.", {"bullet": True, "space_after": 8}),
        ("The three models give p(fraud) from the free-energy softmax over the two label hypotheses; the ensemble is their mean, labelled FRAUD above 0.5.", {"bullet": True, "space_after": 8}),
        ("The risk band is chosen from the mean probability (0.4 and 0.7).", {"bullet": True, "space_after": 8}),
        ("The dashed branch is the bulk-route defect: the scaler call raises because it was fitted on 24 columns and only 13 are present, and a bare except lets the unscaled values through.", {"bullet": True}),
    ], size=15)
    notes[13] = "The scoring flow of predict_single and predict_batch, drawn as the code actually behaves, including the skipped scaling step."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_prediction.png"), 0.55, 1.75, 7.2, 5.15)
    K.picture(s, F.diag("08_use_case.png"), 7.95, 1.75, 4.85, 3.4)
    cap(s, 7.95, 5.3, 4.85, 1.6, "Left: one prediction from the browser through the Vite proxy and Flask to the three models. Right: the use cases that exist in the interface for a merchant or analyst and for the risk-operations team. There is no sign-in.")
    notes[14] = "Sequence of one check and the use cases. There is no authentication in the prototype."

    # ---- 15 storage design
    s = S[14]
    set_title(s, "Data Storage Design (no SQL database)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.75, 7.2, 4.9)
    rows = [["File", "Content"], ["creditcard_synthetic.csv", "39,000 rows, 29 columns"], ["pipeline.pkl", "scaler, columns, importances"], ["grbmc_*.pkl (3)", "trained models"],
            ["history_*.json (3)", "per-epoch metrics"], ["batch_predictions.csv", "last bulk upload result"]]
    ftable(s, 7.95, 1.85, [2.4, 2.45], rows, size=12, row_h=0.5)
    ftext(s, 7.95, 5.1, 4.85, 1.6, "pipeline.pkl was written with a different pandas version and does not load here, so the backend silently rebuilds the same pipeline at start-up (same sample, same seed).", size=12, color=K.FNAVY, italic=True)
    notes[15] = "Persistence is files plus an in-memory dictionary. The pickle issue is a real observation from this machine, handled by the backend's fallback."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("16", "API routes", K.FBLUE), ("7", "React pages", K.GREEN), (f"{F.BACKEND_LINES:,}", "backend lines (Python)", K.ORANGE), (f"{F.FRONTEND_LINES:,}", "frontend lines (React)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 26)
    ftext(s, 0.55, 3.05, 6.6, 3.9, [
        ("GRBMC is plain numpy: weights W (visible-hidden) and U (label-hidden), biases b, c, d, free-energy scoring and PCD gradients.", {"bullet": True, "space_after": 7}),
        ("SAS and QS replace the Gibbs negative phase with annealing-style sampling; they use the batch-mean visible vector, a simplification of the paper's procedure.", {"bullet": True, "space_after": 7}),
        ("Model lab: retrain from the browser with a choice of balancing, features, trainers and epochs; the models are saved on disk.", {"bullet": True, "space_after": 7}),
        ("Interface: light and dark themes, risk meter, toasts, offline banner, responsive layout.", {"bullet": True}),
    ], size=14)
    K.picture(s, F.diag("12_grbmc_structure.png"), 7.4, 3.05, 5.4, 3.5)
    notes[16] = "What was built and how big it is. The samplers are simplified approximations, which the report states plainly."

    # ---- 17 results: shipped models
    s = S[16]
    set_title(s, f"Results – Shipped Models (synthetic data, n = {n_test:,})")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.pct(ens["f1"]), "ensemble F1", K.FBLUE), (F.pct(sas["f1"]), "SAS F1", K.GREEN), (F.pct(qs["f1"]), "QS F1", K.ORANGE), (F.pct(pcd["accuracy"]), "PCD accuracy (flags almost all)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 26)
    K.picture(s, F.fig("f02_model_metrics.png"), 0.55, 3.0, 6.2, 3.5)
    K.picture(s, F.fig("f07_seed_study.png"), 6.95, 3.0, 5.85, 2.4)
    ftext(s, 6.95, 5.45, 5.85, 1.5, [
        (f"All four share ROC-AUC {F.f3(ens['roc_auc'])}: PCD ranks well but its 0.5 threshold is badly placed.", {"bullet": True, "space_after": 4}),
        (f"Five seeds, final epoch, F1: PCD {F.seed_cell('classical', 'f1')}, SAS {F.seed_cell('sa', 'f1')}, QS {F.seed_cell('qa', 'f1')}.", {"bullet": True}),
    ], size=12)
    cap(s, 0.55, 6.55, 6.2, 0.4, "Threshold 0.5. Best-epoch histories are chosen on the test split, so final-epoch figures are used.", 11)
    notes[17] = f"SAS and QS are around 94 percent F1; the ensemble 94.5. PCD at the 0.5 threshold has recall 100 percent and {pcd['fp']} false positives out of {pcd['tn'] + pcd['fp']} legitimate rows. Across seeds PCD is unstable."

    # ---- 18 confusion matrix & threshold
    s = S[17]
    set_title(s, "Confusion Matrices & Threshold Analysis")
    wipe(s)
    K.picture(s, F.fig("f03_confusion_matrices.png"), 0.55, 1.75, 11.0, 2.8)
    K.picture(s, F.fig("f06_threshold_sweep.png"), 0.55, 4.6, 6.5, 2.4)
    sw = {r["threshold"]: r for r in sv["ensemble_threshold_sweep"]}
    ftext(s, 7.3, 4.65, 5.5, 2.3, [
        (f"Ensemble: {ens['tn']} true negatives, {ens['fp']} false positives, {ens['fn']} false negatives, {ens['tp']} true positives.", {"bullet": True, "space_after": 6}),
        (f"Threshold 0.4: precision {F.pct(sw[0.4]['precision'])}, recall {F.pct(sw[0.4]['recall'])}; 0.5: {F.pct(sw[0.5]['precision'])} and {F.pct(sw[0.5]['recall'])}; 0.6: {F.pct(sw[0.6]['precision'])} and {F.pct(sw[0.6]['recall'])}.", {"bullet": True, "space_after": 6}),
        ("Above 0.7 nothing is flagged: no test row reaches the 'HIGH' band with the averaged probability.", {"bullet": True}),
    ], size=13)
    notes[18] = "Confusion matrices at 0.5 and the ensemble threshold sweep. The curve is steep: between 0.4 and 0.6 precision and recall trade places sharply, and the app's 0.7 band is never reached."

    # ---- 19 baselines and imbalance
    s = S[18]
    set_title(s, "Results – Baselines and a 1 % Fraud Variant")
    wipe(s)
    K.picture(s, F.fig("f08_baseline_comparison.png"), 0.55, 1.75, 6.2, 2.15)
    fu = F.FULL
    ftext(s, 7.0, 1.8, 5.8, 2.2, [
        (f"Same split: LR {F.pct(bb['Logistic Regression']['f1'])}, RF {F.pct(bb['Random Forest']['f1'])}, XGBoost {F.pct(bb['XGBoost']['f1'])} against RBM {F.pct(ens['f1'])}.", {"bullet": True, "space_after": 5}),
        (f"Whole 39,000-row file (n = 7,800): XGBoost {F.pct(fu['XGBoost']['f1'])}, RBM trainers {F.pct(fu['classical']['f1'])} to {F.pct(fu['qa']['f1'])}.", {"bullet": True, "space_after": 5}),
        ("Simple models win on this data; the RBM's only edge is that it is the base paper's method.", {"bullet": True}),
    ], size=13)
    K.picture(s, F.fig("f09_imbalance_study.png"), 0.55, 4.0, 9.3, 3.0)
    lr_n, lr_u = F.imb_row("Logistic Regression", "none"), F.imb_row("Logistic Regression", "undersample")
    xg_s = F.imb_row("XGBoost", "SMOTE")
    gr = F.imb_row("GRBMC (SA)", "undersample")
    ftext(s, 10.05, 4.05, 2.75, 2.95, [
        (f"Undersampling lifts recall to 0.96 - 0.98 but precision falls to about 0.2 (LR {lr_u['precision']['mean']:.2f}).", {"bullet": True, "space_after": 4}),
        (f"No resampling: LR precision {lr_n['precision']['mean']:.2f}, recall {lr_n['recall']['mean']:.2f}.", {"bullet": True, "space_after": 4}),
        (f"SMOTE with XGBoost: {xg_s['precision']['mean']:.2f} / {xg_s['recall']['mean']:.2f}.", {"bullet": True, "space_after": 4}),
        (f"RBM on undersampled data: {gr['precision']['mean']:.2f} / {gr['recall']['mean']:.2f}.", {"bullet": True}),
    ], size=11)
    notes[19] = "At 1 percent fraud (78 fraud among 7,800 test rows, five seeds) undersampling floods the analyst with false alarms."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance, Testing & the Serving Defect")
    wipe(s)
    sec = {k: F.SEEDS[k]["elapsed_sec"]["mean"] for k in ("classical", "sa", "qa")}
    for i, (v, l, c) in enumerate([(f"{lat1['median_ms']:.0f} ms", "median, single check (n = 100)", K.FBLUE), (f"{latb['median_ms']:.0f} ms", "median, 20-row CSV batch", K.GREEN),
                                   (f"{min(sec.values()):.0f} - {max(sec.values()):.0f} s", "20-epoch training, CPU", K.ORANGE), (f"{F.SMOKE_PASSED} / {F.SMOKE_TOTAL}", "smoke-test checks passed", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 26)
    rows = [["Test layer", "What it covers", "Result"],
            ["Smoke test", "running API on 8304: all read routes, predict, batch, errors, export", f"{F.SMOKE_PASSED} of {F.SMOKE_TOTAL}"],
            ["Latency", "127.0.0.1, Flask dev server", f"p95 {lat1['p95_ms']:.0f} ms"],
            ["Experiments", "baselines, seeds, thresholds, 1 % variant", "metrics.json"],
            ["Unit tests", "no pytest suite exists", "gap"],
            ["Not tested", "concurrency, real data, drift, user study", "gap"]]
    ftable(s, 0.55, 3.1, [1.7, 3.9, 1.5], rows, size=11, row_h=0.5)
    K.picture(s, F.fig("f10_serving_defect.png"), 7.9, 3.05, 4.9, 1.9)
    ftext(s, 7.9, 5.0, 4.9, 1.95, [
        (f"Fresh 2,000-row batch, ensemble: accuracy {F.pct(asv['accuracy'])} as served, {F.pct(scl['accuracy'])} when scaled (false positives {asv['fp']} against {scl['fp']}).", {"bullet": True, "space_after": 4}),
        (f"The bundled sample file flags {srv['sample_batch_as_served_flagged']} of {srv['sample_batch_rows']} rows as fraud; {srv['sample_batch_scaled_flagged']} when scaled (true fraud {srv['sample_batch_true_fraud']}).", {"bullet": True}),
    ], size=12)
    cap(s, 0.55, 6.5, 7.1, 0.45, "Intel Core i5-12500, CPU only. The defect is documented, not fixed.", 11)
    notes[20] = "Latency is far below a second. The defect was found by auditing the serving code and measured on a fresh batch generated with a different seed; the screenshot of the bulk page shows the same behaviour."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only: different data)")
    wipe(s)
    rows = [["", "Base paper (real data)", "This project (synthetic data, simulated samplers)"],
            ["F1: PCD / SAS / QS", f"{bt['PCD'][3]} / {bt['SAS'][3]} / {bt['QS'][3]} %", f"{F.pct(pcd['f1'])} / {F.pct(sas['f1'])} / {F.pct(qs['f1'])} (shipped models); five-seed means {F.seed_cell('classical', 'f1')} / {F.seed_cell('sa', 'f1')} / {F.seed_cell('qa', 'f1')}"],
            ["Precision: SAS", f"{bt['SAS'][1]} %", f"{F.pct(sas['precision'])}"],
            ["Training time", f"{bt['PCD'][4]} / {bt['SAS'][4]} / {bt['QS'][4]}", f"{sec['classical']:.0f} / {sec['sa']:.0f} / {sec['qa']:.0f} seconds (3,000-row subsample, 10 sweeps)"],
            ["Sampler hardware", "Julia QARBoM.jl; D-Wave simulators and quantum annealer", "numpy; classical simulation only"],
            ["Test set", "about 7.8 thousand rows (20 % of the 39 thousand balanced set)", f"{n_test:,} rows ({ntest_fraud} fraud) of generated data"]]
    ftable(s, 0.55, 1.8, [2.4, 4.4, 5.4], rows, size=12, row_h=0.5)
    K.picture(s, F.fig("f13_paper_comparison.png"), 0.55, 4.95, 6.4, 1.95)
    ftext(s, 7.2, 5.0, 5.6, 1.9, [
        ("The numbers sit side by side for context only: data, labels, test sets and samplers differ, so no superiority or inferiority is claimed.", {"bullet": True, "space_after": 5}),
        ("SAS is the steadiest of our three trainers and QS is within noise of it; the paper's SAS lead is therefore not contradicted, but our PCD result depends strongly on threshold and seed.", {"bullet": True}),
    ], size=12)
    notes[21] = "State plainly that the datasets differ: real Stone data against a generated set that is easier. The time column compares hours of simulator or cloud use with seconds of a simplified sampler."

    # ---- 22-24 output screens
    def screens(slide, title, a, b, ca, cb, note):
        set_title(slide, title)
        wipe(slide)
        K.picture(slide, F.shot(a), 0.55, 1.75, 6.0, 3.75, border=True)
        K.picture(slide, F.shot(b), 6.8, 1.75, 6.0, 3.75, border=True)
        cap(slide, 0.55, 5.6, 6.0, 1.3, ca, 12)
        cap(slide, 6.8, 5.6, 6.0, 1.3, cb, 12)

    screens(S[21], "Output Screens – Dashboard and Verify", "03_dashboard.png", "04_verify.png",
            "Dashboard: 39,000 transactions, 50 % fraud-labelled, best F1 94.8 % (simulated annealing, epoch 17, a best-epoch figure) and fraud rate by merchant category computed from the data.",
            "Verify: the fraud example (1,200, online, card not present, country mismatch) is flagged at 68.4 %: PCD 50.9 %, QS 59.5 %, SAS 94.7 %; recommended action 'challenge with 3-D Secure'.", "")
    notes[22] = "Real screenshots of the running application. The dashboard's best-F1 tile uses the best epoch chosen on the test split."
    screens(S[22], "Output Screens – Bulk Verification and Analytics", "05_bulk.png", "06_analytics-where.png",
            "Bulk: the 20-row sample file is flagged 20 of 20 (accuracy 50.0 %, recall 100 %): the unscaled bulk path in action (the single-check Verify route scales correctly). The page itself warns that every row was flagged.",
            "Analytics: patterns in the full dataset: Electronics and Online have the highest fraud rates, typical fraud amount 352 against a legitimate median of 80.", "")
    notes[23] = "The bulk screen is a faithful capture of the defect, not a mock-up. Analytics is computed from the CSV, not from a model."
    screens(S[23], "Output Screens – Model Comparison and Training Curves", "10_models-comparison.png", "11_models-curves.png",
            "Model comparison: metrics at each model's best epoch, and, in the table below the fold, the paper's training times next to the measured seconds.",
            "Training curves: per-epoch F1 of the three trainers; PCD jumps between 0 and 80 % while SAS settles near 94 %.", "")
    notes[24] = "The curves show the PCD instability that the five-seed study quantifies."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.1, [
        ("**Delivered:** a working fraud-screening application (Flask API with 16 routes, React interface, model lab) around the base paper's RBM classifier, with a 20/20 smoke test.", {"bullet": True, "space_after": 8}),
        (f"**Measured:** ensemble F1 {F.pct(ens['f1'])} and ROC-AUC {F.f3(ens['roc_auc'])} on synthetic data; LR, RF and XGBoost reach F1 {F.pct(bb['Random Forest']['f1'])} to {F.pct(bb['Logistic Regression']['f1'])}; PCD is unstable ({F.seed_cell('classical', 'f1')} F1 over five seeds).", {"bullet": True, "space_after": 8}),
        ("**Objectives:** 2 and 5 met; 3 and 4 met in the experiment script, not in the product; 1 partly (synthetic data, no public benchmark); 6 met, with the scaling defect.", {"bullet": True, "space_after": 8}),
        ("**Findings:** imbalance handling matters more than model family; undersampling at 1 % fraud costs precision; thresholds and calibration decide the PCD result.", {"bullet": True, "space_after": 8}),
        ("**Limits:** synthetic data, quantum-inspired simulation, test-split epoch selection, no authentication, no real-world validation.", {"bullet": True}),
    ], size=17)
    notes[25] = "Be explicit about which objectives were met and where. The honest result is that simple baselines beat the RBM on this data, and that our own audit found a serving defect."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 5.0, [
        ("Fix the scaling defect in the batch route and add a regression test that scores a known batch through the API.", {"bullet": True, "space_after": 10}),
        ("Choose the decision threshold on a validation split and calibrate the PCD model; select epochs on validation data, not the test split.", {"bullet": True, "space_after": 10}),
        ("Pilot on real, imbalanced transaction data with a temporal split; report precision at a fixed alert budget and a PR curve.", {"bullet": True, "space_after": 10}),
        ("Implement a faithful annealing sampler or use quantum cloud access; compare with a gradient-boosting champion model.", {"bullet": True, "space_after": 10}),
        ("Add authentication, HTTPS, a production server, drift monitoring and a database for scored transactions.", {"bullet": True}),
    ], size=20)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "rfpaper", "cfa", "shi", "bhatt", "dp18", "smote", "tomek", "hinton06", "cd", "pcd", "larochelle", "anneal", "sklearn"]

    def clean(k):
        return f"[{refno(k)}] " + REFS[refno(k) - 1].replace("*", "")

    ftext(s, 0.55, 1.8, 6.0, 5.1, [(clean(k), {"space_after": 7}) for k in keys[:7]], size=12)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(clean(k), {"space_after": 7}) for k in keys[7:]], size=12)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
