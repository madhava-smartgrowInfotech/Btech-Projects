"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os
import re
import sys
import time

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import (EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIMITATIONS_EXISTING, LIT_TABLE, NEED, OBJECTIVES, PROBLEM_STATEMENT,
                  PROPOSED_APPROACH, SUPPORT_EVIDENCE)
from .refs import KEYS, REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "credit card fraud, restricted Boltzmann machine, class imbalance, machine learning"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def ref_text(key):
    return f"[{KEYS[key]}] {REFS[KEYS[key] - 1]}".replace("*", "")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: SecurePay AI scores a card transaction with three restricted-Boltzmann-machine classifiers "
                "and returns a risk band; the work follows a 2026 IEEE Access base paper and is measured on synthetic data.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review covers why the problem matters and what is missing in current work; design, methodology and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: scoring a card payment")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Card fraud causes large losses. Fixed rules are learned and bypassed and raise false alarms. Genuine transactions vastly outnumber fraud, so a model can score high accuracy while finding almost no fraud.", K.ORANGE, "!"),
        ("Our idea", "Follow the base paper: a Gaussian restricted Boltzmann machine classifier trained three ways (PCD, simulated annealing, quantum-inspired sampling), on a balanced set, and judged by recall, precision and F1 rather than accuracy.", K.BLUE, "i"),
        ("The deliverable", "SecurePay AI: a Flask API and React interface for single checks, CSV bulk screening, analytics and model comparison, returning a risk score and an approve, challenge or decline action. Results are measured, with limits stated.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=20)
    notes[3] = ("Three points: fraud is rare and costly, the base paper's classifier is the starting idea, and the deliverable is a complete application. "
                "The quantum sampler is a classical simulation: no quantum hardware is used.")

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "A Gaussian-Bernoulli RBM classifier with a label layer, trained by persistent contrastive divergence, by simulated annealing and by D-Wave quantum sampling. One-hot encoding, z-score, correlation filter and CatBoost importance prepare the data.", K.BLUE, "M"),
        ("Data and reported results", "Real Stone data: 145 million transactions, about 19 thousand fraud, undersampled to 39 thousand (50/50), 80/20 split. F1: PCD 83 %, SAS 88 %, QS 84 %; training 2 minutes, 4 hours and 1.8 hours.", K.GREEN, "R"),
        ("Relevance to this project", "Same model family, preprocessing and balancing. Different data (synthetic), no quantum hardware, and a deployable interface with a risk band. The paper's numbers are never compared directly with ours.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=16)
    notes[4] = ("Read the citation. The base paper compares three ways of training the same RBM classifier on a private dataset. Table 1 of the paper gives F1 83, 88 and 84 percent. "
                "SMOTE and Tomek links did poorly there, so it used manual undersampling, which is what this project also does.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for r in LIT_TABLE[:6]:
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3] + "; limit: " + r[4][0].lower() + r[4][1:]])
    K.table(s, 0.6, 1.3, [2.2, 0.7, 4.4, 4.83], rows, font_size=11, row_h=0.74, header_size=12, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: classifier comparisons (random forest, XGBoost, ANN), fusion of tree models, imbalance handling, and the base paper's quantum-assisted RBM. "
                "Figures are quoted only where read from the paper or its publisher page; three titles in the batch survey could not be verified and are flagged in the report.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: accuracy is not enough")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 5.9, 5.2, ["Random forest\naccuracy [2]", "Random forest\nF1 [2]", "Fusion of trees\nAUC [3]", "Fusion of trees\nAUPRC [3]", "Fusion of trees\nF1 [3]"],
                     [("Reported value", [0.98, 0.98, 0.9405, 0.6699, 0.6373])], colors=[K.BLUE], number_format="0.00", font_size=11, max_value=1.0)
    for idx in (3, 4):
        pt = gf.chart.plots[0].series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 6.8, 1.45, 5.95, 4.9, [(e, {"bullet": True, "space_after": 9}) for e in SUPPORT_EVIDENCE], size=15, color=K.INK)
    K.text(s, 0.6, 6.65, 12.1, 0.45, "Different studies and datasets (94,682 transactions [2]; IEEE-CIS, 3.5 % fraud [3]); the bars are context, not a comparison.", size=12, color=K.TXT, italic=True)
    notes[6] = ("The left bars come from two papers on different data: a random forest on a small public set with accuracy and F1 near 0.98, and a fusion of tree models on IEEE-CIS where the rare-class measures (AUPRC 0.67, F1 0.64) are far lower. "
                "That is why this project reports precision, recall, F1 and ROC-AUC and tests a 1 percent fraud case.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Transaction arrives", K.NAVY), ("Rules fire", K.NAVY), ("Analyst reviews or auto-decline", K.NAVY), ("Customer approved or blocked", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Rule-based screening", "\n".join(EXISTING_SYSTEM[i] for i in (0, 1, 2)), K.BLUE, "1", body_size=15)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Machine-learning studies", "\n".join(EXISTING_SYSTEM[i] for i in (3, 4)), K.ORANGE, "2", body_size=15)
    notes[7] = "Today most screening starts with fixed rules and analyst review. Research studies use classifiers, but they compare models or sampling methods separately and rarely ship an interface."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Static rules", "Accuracy-led evaluation", "Resampling pitfalls", "Hard to reproduce"]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], LIMITATIONS_EXISTING[i].split(": ", 1)[1], C[i], str(i + 1), body_size=18)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also, no operational layer: most studies stop at a notebook, with no interface that flags a live transaction.", size=15, color=K.INK, anchor="m")
    notes[8] = "Five limitations from the survey. Each maps to a design decision later: undersampled 50/50 training, recall and F1 as the main measures, and a working interface."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 0.9
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.82, g, size=19, color=K.INK, anchor="m")
    notes[9] = ("The batch literature survey states the gap as a compared, tuned pipeline evaluated on recall rather than accuracy and deployable for live scoring. "
                "We address gaps 1, 2, 4, 5 and 6 and give an honest answer on gap 3: the base paper's model runs here only on synthetic data.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.2, fill=K.CARD)
    K.text(s, 0.95, 1.55, 11.45, 3.05, [(p, {"bullet": True, "space_after": 6}) for p in PROBLEM_STATEMENT] + [("**Need:** " + NEED, {"space_after": 0})], size=17, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "One transaction (amount, merchant category, channel, card present, country match, time and anonymised features) or a CSV of them"),
                                 ("Output", "A fraud probability, a risk band (low, medium, high) and an approve, challenge or decline action"),
                                 ("Record", "Per-model probabilities, a downloadable result CSV, and measured metrics for every claim")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.85, 3.85, 1.95)
        K.text(s, x + 0.25, 4.95, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.4, 3.4, 1.35, b_, size=14, color=K.TXT)
    notes[10] = "Read the four points slowly, then the need: handle severe imbalance and maximise fraud recall while keeping false alarms acceptable."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.4, [(f"{i}.  {o}", {"space_after": 20}) for i, o in enumerate(OBJECTIVES, 1)], size=24, color=K.INK)
    notes[11] = ("The six objectives from the proposal. Later reviews report which were met: the benchmark dataset is synthetic, and the Logistic Regression, Random Forest and XGBoost "
                 "comparison was run in an experiment script rather than shipped in the product.")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "One-hot encode merchant and channel, z-score the numbers, filter correlated columns, undersample to 50/50, keep the top 18 features, split 80/20.", K.BLUE, "1"),
        ("Model", "A Gaussian RBM classifier with visible, hidden and label layers, trained by PCD, by simulated annealing and by a quantum-inspired sampler.", K.GREEN, "2"),
        ("Decide", "Average the three fraud probabilities. Below 40 % approve, 40 to 70 % challenge, above 70 % decline and alert.", K.ORANGE, "3"),
        ("Serve and measure", "Flask API and React app; precision, recall, F1 and ROC-AUC against LR, RF and XGBoost, over seeds, thresholds and a 1 % fraud case.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=17)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "S", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "Fraud probability p = (p_PCD + p_SAS + p_QS) / 3;  fraud if p > 0.5", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.85, "Runs on a CPU with no paid service and no quantum hardware; the quantum sampler is a classical simulation. Everything is reproducible from scripts in the project.", size=14, color=K.TXT)
    notes[12] = " ".join(PROPOSED_APPROACH[:3]) + " The three trainers use 200 hidden units (PCD) and 65 hidden units (SAS, QS), as in the base paper."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    arch = F.diag("01_system_architecture.png")
    if os.path.exists(arch):
        K.picture(s, arch, 0.6, 1.3, 7.4, 5.5)
    else:
        K.box(s, 0.6, 1.3, 7.4, 5.5)
        K.text(s, 0.9, 3.8, 6.8, 0.6, "[system architecture diagram]", size=18, color=K.TXT, align="c")
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite web app (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: Flask API (port {F.API_PORT}) with {F.N_ENDPOINTS} routes.", {"bullet": True}),
        ("Processing: preprocessing, three RBM classifiers, ensemble and risk bands.", {"bullet": True}),
        ("Data: a CSV dataset, pickled models and training histories on the local disk; no database.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser calls the Flask API, which transforms the input, asks the three models and averages their probabilities. Everything is stored in files on the local disk."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working application", "Honest model comparison", "Imbalance evidence", "Measured, with limits", "Reproducible"]
    for i, (h, b_) in enumerate(zip(heads, EXPECTED_OUTCOMES)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=15)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=16)
    notes[14] = "Expected outcomes and contributions. The later reviews report measured results, including where the simple baselines do better than the RBM and what could not be measured."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "rfpaper", "cfa", "shi", "bhatt", "dp18", "smote", "pcd", "anneal", "sklearn"]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(ref_text(k), {"space_after": 8}) for k in keys], size=14, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report. Three survey titles that could not be verified are marked as such there and are not cited."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why a classical simulation instead of quantum hardware (no access; the quantum sampler here is quantum-inspired only); "
                 "why not simply use gradient boosting (it does score higher on this data and is reported); why the data is synthetic (the base paper's data is proprietary).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


def render(pptx_path, out_dir):
    """Export every slide to PNG with PowerPoint COM (needs Windows paths); serialised with the shared COM lock."""
    import win32com.client

    from .comlock import com_lock

    os.makedirs(out_dir, exist_ok=True)
    for attempt in range(4):
        try:
            with com_lock():
                app = win32com.client.DispatchEx("PowerPoint.Application")
                pres = app.Presentations.Open(os.path.abspath(pptx_path), True, False, False)
                pres.SaveAs(os.path.abspath(out_dir), 18)
                pres.Close()
                app.Quit()
            return out_dir
        except Exception as e:  # noqa: BLE001
            print("render retry", attempt + 1, e)
            time.sleep(8)
    return None


if __name__ == "__main__":
    out = build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx"))
    print(out)
    if "--render" in sys.argv:
        print(render(out, sys.argv[sys.argv.index("--render") + 1]))
