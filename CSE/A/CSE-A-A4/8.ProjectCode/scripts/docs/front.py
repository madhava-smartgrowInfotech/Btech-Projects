"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(F.TAGLINE, align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, whose work on quantum-assisted restricted Boltzmann machines this project builds on, and the maintainers of the open-source "
        "software used in it (Flask, React, Vite, NumPy, pandas, scikit-learn, imbalanced-learn, XGBoost and Matplotlib).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    sv, bb = F.SAVED, F.BASE_BAL
    ens, sas, qs, pcd = sv["ensemble"], sv["sa"], sv["qa"], sv["classical"]
    d.p("Credit card fraud causes large losses and erodes customer trust. Fixed rules are learned and bypassed by fraudsters and raise false alarms, and learning systems face class imbalance: "
        "genuine transactions vastly outnumber fraudulent ones, so a model can score a high accuracy while finding almost no fraud. This project builds SecurePay AI, a working fraud-screening application, and measures it honestly. "
        "It follows the base paper [1], which trains a Gaussian-Bernoulli restricted Boltzmann machine classifier three ways: persistent contrastive divergence (PCD), simulated-annealing sampling (SAS) and quantum sampling (QS). "
        "Quantum hardware was not available, so the quantum sampler is simulated classically and is only quantum-inspired. The pipeline one-hot encodes, z-scores, filters correlated columns, undersamples to 50/50, keeps 18 features and splits 80/20. "
        "A Flask API with 16 routes and a React interface offer single checks, CSV bulk screening, analytics, model comparison and a retraining lab; the three probabilities are averaged and mapped to approve, challenge or decline bands at 40 % and 70 %.")
    d.p(f"All results use a synthetic, balanced dataset of {F.DATASET['rows']:,} transactions, so they are not comparable with the base paper's real-data figures. "
        f"On the {F.SPLIT['test'][0]:,}-row held-out split the ensemble reaches {F.pct(ens['accuracy'])} accuracy, {F.pct(ens['recall'])} recall and F1 {F.pct(ens['f1'])} (ROC-AUC {F.f3(ens['roc_auc'])}); "
        f"SAS and QS reach F1 {F.pct(sas['f1'])} and {F.pct(qs['f1'])}, while the shipped PCD model flags almost every transaction (accuracy {F.pct(pcd['accuracy'])}). "
        f"Logistic regression, random forest and XGBoost score F1 {F.pct(bb['Logistic Regression']['f1'])}, {F.pct(bb['Random Forest']['f1'])} and {F.pct(bb['XGBoost']['f1'])} on the same split, so the simple baselines are better here. "
        f"Five retrainings of PCD give F1 {F.seed_cell('classical', 'f1')}, while SAS gives {F.seed_cell('sa', 'f1')}. "
        f"A code review found that the bulk-upload route silently skips feature scaling; on a fresh batch accuracy falls from {F.pct(F.SERVING['with_scaling']['ensemble']['accuracy'])} to {F.pct(F.SERVING['as_served_no_scaling']['ensemble']['accuracy'])}. "
        f"A single check takes a median {F.LATENCY['predict_single_fraud_example']['median_ms']:.0f} ms. No real transaction data, quantum hardware or user study was used.")
    d.p("**Keywords:** credit card fraud, machine learning, restricted Boltzmann machine, class imbalance, undersampling, SMOTE, simulated annealing, Flask, React.", align="left")


ABBR = [
    ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CD", "Contrastive Divergence"), ("CSV", "Comma-Separated Values"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"),
    ("FN", "False Negative"), ("FP", "False Positive"), ("GRBM", "Gaussian-Bernoulli Restricted Boltzmann Machine"), ("GRBMC", "GRBM Classifier"),
    ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"), ("LR", "Logistic Regression"), ("ML", "Machine Learning"),
    ("NISQ", "Noisy Intermediate-Scale Quantum"), ("PCD", "Persistent Contrastive Divergence"), ("PR", "Precision-Recall"), ("QA", "Quantum Annealing"),
    ("QS", "Quantum Sampling (quantum-inspired in this project)"), ("QUBO", "Quadratic Unconstrained Binary Optimization"), ("RBM", "Restricted Boltzmann Machine"),
    ("REST", "Representational State Transfer"), ("RF", "Random Forest"), ("ROC", "Receiver Operating Characteristic"), ("SA", "Simulated Annealing"),
    ("SAS", "Simulated-Annealing Sampling"), ("SMOTE", "Synthetic Minority Over-sampling Technique"), ("TN / TP", "True Negative / True Positive"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("XGB", "Extreme Gradient Boosting"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
