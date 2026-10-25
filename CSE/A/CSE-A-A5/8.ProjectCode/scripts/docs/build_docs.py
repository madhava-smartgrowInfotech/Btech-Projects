"""
Builds the six chapter documents and the full report (17.Final Document.docx) into the batch folder.

    python -m scripts.docs.build_docs            # from 8.ProjectCode; all documents
    python -m scripts.docs.build_docs --only 9   # one document
"""

import argparse
import os

from docx.shared import Inches, Pt

from . import appendix, ch1, ch2, ch3_4, ch5, ch6, ch7, ch8, ch9_10, diagram_index, front
from . import facts as F
from .blocks import FIG, H1, H2, P, render, renum
from .kit import Doc, add_runs
from .refs import REFS

NAMES = {
    9: "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx",
    10: "10.SCOPE & LIMITATIONS.docx",
    11: "11. BLOCK DIAGRAMS ALL.docx",
    12: "12.DESIGN &IMPLEMENTATION.docx",
    14: "14.RESULTS.docx",
    15: "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    17: "17.Final Document.docx",
}


def new_doc(doc_title, report=False):
    return Doc(doc_title, header_text=f"{F.PRODUCT}  |  {doc_title}", report=report)


def title_block(d, doc_title):
    d.footer_numbers()
    d.spacer(24)
    d.p(F.FULL_TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p(doc_title, align="center", size=14, bold=True, space_after=14)
    d.p(f"{F.DEPT}, {F.COLLEGE_SHORT}", align="center", size=12, space_after=2)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})  ·  Academic year {F.YEAR}", align="center", size=12, space_after=12)
    rows = [[str(i), n, r] for i, (n, r) in enumerate(F.STUDENTS, 1)]
    d.table("Project team", ["S.No.", "Student name", "Roll number"], rows, widths=[0.7, 3.4, 1.8], size=11, align=["c", "l", "c"])
    d.p(f"**Project guide:** {F.GUIDE}", align="center", size=12, space_after=10)
    d.p("Chapter and section numbers quoted in the text (for example Chapter 8 or Section 6.2) refer to the full project report, 17.Final Document. "
        "Figures and tables in this document are numbered from 1.", size=10.5, italic=True, space_after=6)


def save(d, name):
    path = d.save(os.path.join(F.BATCH, name))
    print("wrote", path)
    return path


def add_references(d):
    for i, r in enumerate(REFS, 1):
        par = d.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.45)
        par.paragraph_format.first_line_indent = Inches(-0.45)
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(5)
        add_runs(par, f"[{i}]  {r}", 10)


# ----------------------------------------------------------------- chapter documents
def doc9():
    d = new_doc("Methodology, Novelty and Key Contributions")
    title_block(d, "Methodology, Novelty and Key Contributions")
    d.h1("1. METHODOLOGY", page_break=False)
    d.p("This document describes how CipherGuard Shield works from end to end and what is new in it. Part 1 gives the pipeline, the data, the algorithms and the reasons they were chosen; "
        "Part 2 sets the work against the base paper and separates what is new from what is standard; Part 3 lists the contributions.")
    render(d, renum(ch6.methodology_blocks(), "1"))
    d.h1("2. NOVELTY AGAINST THE BASE PAPER")
    render(d, renum(ch9_10.novelty_blocks(), "2"))
    d.h1("3. KEY CONTRIBUTIONS")
    render(d, ch9_10.contribution_blocks())
    return save(d, NAMES[9])


def doc10():
    d = new_doc("Scope and Limitations")
    title_block(d, "Scope and Limitations")
    d.h1("1. SCOPE", page_break=False)
    render(d, renum(ch9_10.scope_blocks(), "1"))
    d.h1("2. ASSUMPTIONS")
    render(d, ch9_10.assumption_blocks())
    d.h1("3. LIMITATIONS")
    render(d, ch9_10.limitation_blocks())
    render(d, [P("Two figures show the limitation that matters most for reading the numbers. The first shows where the training data comes from decides everything: the shipped model is excellent on its own synthetic data and close to chance on real traffic. "
                 "The second shows which traffic categories the real-trained model handles well and where it does not."),
               FIG(F.fig("f07_synthetic_vs_real.png"), "Accuracy of the stacked ensemble by training data and test data: synthetic-trained models do not transfer to real UNSW-NB15 flows.", 5.6),
               FIG(F.fig("f05_per_category.png"), "Share of each traffic category handled correctly by the real-trained stacked model on the official test set.", 5.2)])
    d.h1("4. WHAT WAS NOT MEASURED")
    render(d, ch9_10.not_measured_blocks())
    return save(d, NAMES[10])


def doc11():
    d = new_doc("Block Diagrams")
    title_block(d, "Block Diagrams (all)")
    d.h1("1. INDEX OF DIAGRAMS", page_break=False)
    rows = [[str(i), t, f"`docs/diagrams/{f}`"] for i, (_, t, f, _, _) in enumerate(diagram_index.DIAGRAMS, 1)]
    d.p("Every diagram was generated from the real design by `scripts/docs_build/make_diagrams.py`, `diagrams_dfd.py` and `diagrams_more.py` (matplotlib, 150 dpi) and mirrors the code in `cipherguard/api/main.py`, `cipherguard/crypto/` and `cipherguard/ids/`. "
        "CipherGuard Shield keeps no database: alerts and keys live in process memory and the models and reports are files, so the entity-relationship diagram shows logical entities.")
    d.table("Diagrams in this document", ["No.", "Diagram", "Image file"], rows, widths=[0.5, 3.0, 2.8], size=10.5, align=["c", "l", "l"])
    d.h1("2. DIAGRAMS")
    for key, title, file, caption, text in diagram_index.DIAGRAMS:
        d.h2(title)
        d.figure(F.diag(file), caption, 5.0 if key in ("seq", "flow", "uc") else 5.8)
        d.p(text)
    return save(d, NAMES[11])


def doc12():
    d = new_doc("Design and Implementation")
    title_block(d, "Design and Implementation")
    d.h1("1. DESIGN", page_break=False)
    render(d, renum(ch5.blocks()[1:], "1"))
    d.h1("2. IMPLEMENTATION")
    render(d, renum(ch6.application_blocks(), "2"))
    d.h1("3. KEY CODE WALK-THROUGHS")
    render(d, renum(ch9_10.walkthrough_blocks(), "3"))
    d.h1("4. IMPLEMENTATION DECISIONS AND THEIR REASONS")
    render(d, ch9_10.decision_blocks())
    return save(d, NAMES[12])


def doc14():
    d = new_doc("Results")
    title_block(d, "Results")
    d.h1("1. RESULTS AND OUTPUT SCREENS", page_break=False)
    render(d, renum(ch8.blocks()[1:], "1"))
    return save(d, NAMES[14])


def doc15():
    d = new_doc("Conclusion, Future Scope and References")
    title_block(d, "Conclusion, Future Scope and References")
    d.h1("1. CONCLUSION", page_break=False)
    render(d, renum(ch9_10.conclusion_blocks(), "1"))
    render(d, [P("The results behind the key findings are shown below: how the stacked ensemble compares with its parts on real data, and how the measured accuracy sits next to the figures reported in the base paper."),
               FIG(F.fig("f02_model_comparison_official.png"), "Accuracy, F1 and false-positive rate of the stacked ensemble, its two parts and three baselines on the official UNSW-NB15 partition (5 seeds).", 6.2),
               FIG(F.fig("f08_base_paper_comparison.png"), "Mean accuracy reported in the base paper next to the values measured here under a comparable protocol (indicative only: different subsets).", 6.2)])
    d.h1("2. FUTURE SCOPE")
    render(d, renum(ch9_10.future_blocks(), "2"))
    d.h1("3. REFERENCES")
    add_references(d)
    return save(d, NAMES[15])


# ----------------------------------------------------------------- full report
def doc17():
    F.REPORT = True
    d = new_doc("Project Report", report=True)
    front.cover(d)
    d.new_section(fmt="lowerRoman", start=1)
    front.certificate(d)
    front.declaration(d)
    front.acknowledgement(d)
    front.abstract(d)
    d.h1("TABLE OF CONTENTS")
    d.toc_field('TOC \\o "1-2" \\h \\z \\u')
    d.h1("LIST OF FIGURES")
    d.toc_field('TOC \\h \\z \\c "Figure"')
    d.h1("LIST OF TABLES")
    d.toc_field('TOC \\h \\z \\c "Table"')
    d.h1("LIST OF ABBREVIATIONS")
    front.abbreviations(d)
    d.new_section(fmt="decimal", start=1)
    render(d, ch1.blocks(), h1_page_break=False)
    render(d, ch2.blocks())
    render(d, ch3_4.blocks())
    render(d, ch5.blocks())
    render(d, ch6.blocks())
    render(d, ch7.blocks())
    render(d, ch8.blocks())
    render(d, ch9_10.chapter9_blocks())
    render(d, ch9_10.chapter10_blocks())
    d.h1("REFERENCES")
    add_references(d)
    render(d, appendix.code_blocks())
    render(d, appendix.guide_blocks())
    render(d, appendix.screens_blocks())
    return save(d, NAMES[17])


BUILDERS = {9: doc9, 10: doc10, 11: doc11, 12: doc12, 14: doc14, 15: doc15, 17: doc17}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", type=int, nargs="*")
    a = ap.parse_args()
    for k, fn in BUILDERS.items():
        if not a.only or k in a.only:
            fn()


if __name__ == "__main__":
    main()
