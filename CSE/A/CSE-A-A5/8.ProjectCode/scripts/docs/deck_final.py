"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520

try:  # reference numbers are fixed (see the brief): 1 base, 2-5 batch survey, 7 UNSW-NB15, 9 XGBoost, 10 LeNet, 11 stacking, 12 GCM, 13 SP 800-56A, 14 HKDF, 20 Kadry, 23 Mbaya, 24 Rahman
    from .refs import REFS
except Exception:  # pragma: no cover - refs.py is written by another worker
    REFS = None


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "a5_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def cm_grid(slide, x, y, cm, title, w=3.9):
    """Native 2x2 confusion matrix: counts and share of each actual class."""
    ftext(slide, x, y, w, 0.5, title, size=13, color=K.FNAVY, bold=True)
    cw, ch = (w - 1.0) / 2, 0.85
    for j, lab in enumerate(["pred normal", "pred attack"]):
        ftext(slide, x + 1.0 + j * cw, y + 0.5, cw, 0.3, lab, size=11, color=K.DARK, align="c")
    for i, lab in enumerate(["actual normal", "actual attack"]):
        ftext(slide, x, y + 0.85 + i * ch, 1.0, ch, lab, size=11, color=K.DARK, anchor="m")
        tot = sum(cm[i])
        for j in range(2):
            good = i == j
            fill = K.RGBColor(0xC6, 0xE5, 0xD3) if good else K.RGBColor(0xFB, 0xDD, 0xC8)
            K.box(slide, x + 1.0 + j * cw + 0.03, y + 0.85 + i * ch + 0.03, cw - 0.06, ch - 0.06, fill=fill, line=None, radius=0.05)
            ftext(slide, x + 1.0 + j * cw, y + 0.85 + i * ch, cw, ch, [(f"{cm[i][j]:,}", {"bold": True, "size": 15}), (f"{cm[i][j] / tot * 100:.1f}% of row", {"size": 10})],
                  size=12, align="c", anchor="m", space_after=0)


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "cloud security, AES-256-GCM, ECDH, intrusion detection, XGBoost, 1D CNN, UNSW-NB15"
    notes = {}
    SH, OFS, PS, S42 = F.SHIPPED, F.OFFS, F.PAPS, F.S42
    st, xg = OFS["stack"], OFS["xgb"]
    pst = PS["stack"]
    bp = F.BASE_PAPER_RESULTS
    shr = F.SHIFT["shipped_on_real_official_test"]["stack"]
    rts = F.SHIFT["real_trained_on_synthetic_demo_csv"]["stack"]
    pc = F.OFF["per_category"]
    t5, t9 = F.thr_row(0.5), F.thr_row(0.9)
    sc = F.CRY["secure_channel_ecdh_aesgcm"]
    share = F.IP["share_of_transfer_median"]

    for sl in S:  # footers on every slide except the title
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh in S[0].shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.FULL_TITLE, 24)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and CipherGuard Shield: authenticated encryption plus machine-learning intrusion detection around one cloud data path."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Models\nScoring Workflow, Interaction & Data Design\nImplementation Highlights\nResults: Detection, Confusion Matrix, Real vs Synthetic\nPerformance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Limitations, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Data sent to a public cloud can be read, altered or attacked on the way. CipherGuard Shield puts two protections around one data path: authenticated encryption and a machine-learning intrusion detector.", {"bullet": True, "space_after": 8}),
        ("SecureChannel: a fresh P-256 key pair per message, ECDH + HKDF-SHA256 to a 256-bit key, AES-256-GCM. All 6,000 tamper trials behaved as expected; 100 KB is encrypted in 0.54 ms.", {"bullet": True, "space_after": 8}),
        ("ThreatSense Engine: XGBoost + 1D CNN stacked by logistic regression, with severity-graded alerts and a live React console.", {"bullet": True, "space_after": 8}),
        (f"Real UNSW-NB15: stacked ensemble {F.f3(st['accuracy']['mean'])} accuracy on the official 82,332-flow test set; XGBoost alone is better at {F.f3(xg['accuracy']['mean'])}. Base-paper protocol: {F.f3(pst['accuracy']['mean'])}.", {"bullet": True, "space_after": 8}),
        (f"The shipped default model is trained on synthetic data: its {F.pct(SH['accuracy'], 2)} is a property of the generator and drops to {F.f3(shr['accuracy'])} on real flows.", {"bullet": True, "space_after": 8}),
        (f"One protected 100 KB transfer costs {F.IP['protected_transfer_100kb']['median_ms']:.1f} ms; detection, not encryption, is {share['ids'] * 100:.0f}% of it.", {"bullet": True}),
    ], size=18)
    notes[3] = "Give the problem, the two protections, then the numbers with their protocol. Say clearly that 97.33% is on synthetic data."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: intrusion detection with fused quantum-inspired NN-LeNet and improved ECC"],
            ["Detector", "FLPF with EVD-QNN (quantum-inspired, simulated on classical hardware) + LeNet; scores fused by DSMMN score-level fusion"],
            ["Encryption", "Enhanced Elliptic Curve Cryptography (EECC): non-intrusive data encrypted before transfer; the paper states no ECDLP-based security proof is given"],
            ["Data", "UNSW-NB15, 9,000 balanced samples (4,500 per class), stratified 70/15/15, 10 runs averaged"],
            ["Reported", "Mean accuracy 0.9487 (0.9564 in the ablation), precision 0.97, FPR 0.03; EECC encrypts in 0.39 s and decrypts in 0.26 s"],
            ["Stated limits", "Balanced subset only (imbalanced data left to future work); no formal security proof for EECC; baselines re-implemented by the authors"]]
    ftable(s, 0.55, 2.7, [2.0, 10.2], rows, size=13, row_h=0.62)
    notes[4] = "Kalyanasundaram et al., IEEE Access 2026. We keep its idea of detection plus encryption on one path, but use standard, analysed primitives and a classical ensemble; neither the quantum-inspired network nor EECC is reproduced."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Finding / limitation"],
            ["Kalyanasundaram et al. [1] (base)", "2026", "Quantum-inspired NN + LeNet, DSMMN fusion, EECC", "UNSW-NB15, 9,000 flows; mean accuracy 0.9487; no EECC security proof"],
            ["Kadry et al. [20]", "2023", "Optimised QNN (WOA features) + ECC", "98.5% accuracy for real-time detection, as cited in [1]"],
            ["Mbaya et al. [23]", "2023", "SecFedIDM-V1: blockchain + BiLSTM-RNN on CIDDS", "Web app that labels packets normal or attack, as cited in [1]"],
            ["Rahman et al. [24]", "2024", "SYN-GAN: GAN-made synthetic data to train NIDS", "100% scores on BoT-IoT with synthetic training data, as cited in [1]"],
            ["Batch survey [2]", "2024", "Encryption + anomaly detection for cloud storage and transmission", "Title-level only; full text not reviewed"],
            ["Batch survey [3]", "2024", "Automated data deletion for multi-cloud with an IDS", "Title-level only; full text not reviewed"],
            ["Batch survey [4]", "2026", "Federated attack detection + improved Blowfish in cloud", "Title-level only; full text not reviewed"],
            ["Batch survey [5]", "2026", "Hybrid-encryption privacy framework for cloud transmission", "Title-level only; full text not reviewed"]]
    ftable(s, 0.55, 1.85, [3.0, 0.7, 4.6, 3.9], rows, size=12, row_h=0.54, align=["l", "c", "l", "l"])
    notes[5] = "Work found in the batch survey and in the base paper's related-work section. Rows described as 'as cited in [1]' are second-hand; rows for [2]-[5] are title-level only."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 15}) for g in [
        "The surveyed work treats cryptography and intrusion detection as separate concerns (batch literature survey).",
        "Where both appear, the combined cost of encrypting and detecting on the same path is not reported.",
        "The base paper's EECC has no formal security proof; standard ECDH + AES-GCM is analysed, standardised and available in common libraries.",
        "Base-paper results use a 9,000-flow balanced subset; the temporally separated official partition and class imbalance are left to future work.",
        "Synthetic or balanced evaluations hide domain shift: how a model trained on one data source behaves on another is rarely shown."]], size=20)
    notes[6] = "Five gaps. We address the first, second, third and fifth directly; the fourth we test with the official partition."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.6)
    ftext(s, 0.85, 1.9, 11.6, 2.5, [
        "Data placed in a public cloud is readable by the provider unless it is encrypted before upload.",
        "Symmetric encryption alone leaves the key exchange vulnerable to interception.",
        "Encryption cannot detect an attacker who already holds valid credentials.",
        "Most academic systems address either cryptography or intrusion detection, rarely both together.",
        "Need: encrypt with a hybrid scheme and monitor for intrusion with machine learning on the same path."], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Text (<= 1 MB), a file (<= 5 MB), one flow record or a CSV (<= 10,000 rows)"), ("Output", "Authenticated envelope, verified plaintext or a rejection, attack verdicts with severity"), ("Record", "Alert log, measured model metrics and benchmarks")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.7, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.8, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.3, 3.5, 1.0, b_, size=14)
    notes[7] = "Four problem points from the batch statement, then input, output and record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate([
        "Design a hybrid encryption scheme using AES for the payload and RSA for key exchange.",
        "Implement secure key generation, distribution and storage between sender and receiver.",
        "Build a transmission module that encrypts before upload and decrypts after download.",
        "Train a machine-learning classifier on a benchmark intrusion dataset to detect anomalous access.",
        "Integrate the detection layer so suspicious sessions are flagged and logged during transfer.",
        "Evaluate encryption and decryption time, throughput overhead and intrusion-detection accuracy."], 1)], size=22)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the real status of each, including that ECDH replaced RSA."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Lane A - SecureChannel", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Payload", K.FNAVY), ("Ephemeral P-256 key", K.FBLUE), ("ECDH + HKDF", K.GREEN), ("AES-256-GCM", K.ORANGE), ("Envelope", K.PURPLE), ("Cloud verify", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Lane B - ThreatSense Engine", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Flow record", K.FNAVY), ("Preprocess", K.FBLUE), ("XGBoost + 1D CNN", K.GREEN), ("LR stacking", K.ORANGE), ("Verdict", K.PURPLE), ("Alert log", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, [
        ("A new key for every message gives forward separation between messages; GCM authenticates ciphertext, nonce and tag, so any change is rejected.", {"bullet": True, "space_after": 7}),
        ("Detection scores flow records (one row or a CSV) and keeps an alert list with severity low / medium / high / critical.", {"bullet": True, "space_after": 7}),
        ("The two lanes run side by side in one console; the API does not block a transfer on a verdict - the operator decides.", {"bullet": True}),
    ], size=16)
    notes[9] = "The proposed system as two lanes. Be explicit that detection is not wired to block transfers."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.0, 5.1)
    ftext(s, 7.8, 1.95, 5.0, 4.9, [
        ("Console: React 19 + Vite on port 5305; legacy HTML dashboard", {"bullet": True, "space_after": 8}),
        ("API: FastAPI + Uvicorn on port 8305, 9 JSON routes plus /docs", {"bullet": True, "space_after": 8}),
        ("Core: cipherguard/crypto (SecureChannel) and cipherguard/ids (ThreatSense Engine)", {"bullet": True, "space_after": 8}),
        ("State: alert list and session-key cache in memory; no database", {"bullet": True, "space_after": 8}),
        ("Files: models/, reports/metrics.json, data/raw/synthetic_demo.csv", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. No database: alerts are an in-memory list capped at 500 entries."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("12_model_architecture.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, [
        ("Preprocess: MinMax on 39 numeric features, one-hot on proto, service, state.", {"bullet": True, "space_after": 8}),
        ("XGBoost: 200 trees, depth 6, learning rate 0.08.", {"bullet": True, "space_after": 8}),
        ("1D CNN: two conv + pool blocks, dense 64, dropout 0.3 (PyTorch, Adam, early stopping).", {"bullet": True, "space_after": 8}),
        ("Stacking: logistic regression on the two probabilities, fitted on the validation split; threshold 0.5.", {"bullet": True, "space_after": 8}),
        ("Measured on the official split (5 seeds), the base-paper protocol (10 runs) and a domain-shift test.", {"bullet": True}),
    ], size=14)
    notes[11] = "The model pipeline. The CNN loss is unweighted BCE in the PyTorch path that ran, although the README mentions balanced weighting."

    # ---- 12 dataset
    s = S[11]
    set_title(s, "Dataset & Models")
    wipe(s)
    ftile(s, 0.55, 1.85, 2.95, 1.05, "175,341", "official training flows (56,000 normal)", K.FBLUE, 26)
    ftile(s, 0.55, 3.05, 2.95, 1.05, "82,332", "official test flows (55.1% attack)", K.GREEN, 26)
    ftile(s, 0.55, 4.25, 2.95, 1.05, "42", "flow features -> 192 encoded inputs", K.ORANGE, 26)
    rows = [["Data", "Rows", "Used for"],
            ["Synthetic demo CSV (generated)", "12,000", "shipped model; 70/15/15 split; 2.5% label noise"],
            ["UNSW-NB15 official training set", "175,341", "fit 149,039 + validation 26,302"],
            ["UNSW-NB15 official testing set", "82,332", "held-out test, later capture period"],
            ["9,000 balanced flows (paper protocol)", "9,000", "6,300 / 1,350 / 1,350, ten repetitions"]]
    ftable(s, 3.75, 1.85, [3.4, 1.0, 4.65], rows, size=12, row_h=0.55)
    ftext(s, 3.75, 4.75, 9.0, 2.1, [
        ("UNSW-NB15 (Moustafa & Slay, UNSW Canberra, 2015): ten classes - Normal, Generic, Exploits, Fuzzers, DoS, Reconnaissance, Analysis, Backdoor, Shellcode, Worms.", {"bullet": True, "space_after": 6}),
        ("The shipped model never saw real traffic: its data is a generator with the UNSW-NB15 schema.", {"bullet": True, "space_after": 6}),
    ], size=14)
    ftext(s, 0.55, 5.55, 2.95, 1.3, "Licence: free for academic use; cite the original authors.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "Official train and test are separated in time, so the test is harder than a random split. The synthetic CSV is only a demo."

    # ---- 13 scoring workflow
    s = S[12]
    set_title(s, "Scoring Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_scoring.png"), 0.55, 1.8, 5.7, 5.1)
    ftext(s, 6.5, 1.95, 6.3, 4.9, [
        ("Missing features get defaults; column names are matched case-insensitively.", {"bullet": True, "space_after": 9}),
        ("Unknown protocols, services or states are ignored by the one-hot encoder, not crashed on.", {"bullet": True, "space_after": 9}),
        ("p = sigmoid(w1 * p_xgb + w2 * p_cnn + b); attack when p >= 0.5.", {"bullet": True, "space_after": 9}),
        ("Severity: p >= 0.85 critical, >= 0.70 high, otherwise medium; normal is low.", {"bullet": True, "space_after": 9}),
        (f"Shipped meta weights: {SH['meta_coef'][0][0]:.2f} for XGBoost and {SH['meta_coef'][0][1]:.2f} for the CNN.", {"bullet": True}),
    ], size=16)
    notes[13] = "The core algorithm of cipherguard/ids/inference.py as a flowchart."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_flow.png"), 0.55, 1.8, 6.4, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[14] = "Sequence of score, encrypt and decrypt requests, and the use cases of the operator, the developer and the mock cloud."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design (no database)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.7, 4.6)
    rows = [["Store", "Content"], ["alerts_log (memory)", "last 500 attack alerts"], ["_last_session_keys", "ephemeral key -> AES key (demo cache)"], ["cloud key pair", "one P-256 pair per process"],
            ["models/", "xgb, cnn, meta, preprocessor"], ["reports/metrics.json", "held-out test metrics"], ["data/raw/", "synthetic demo CSV"]]
    ftable(s, 7.5, 1.9, [2.2, 3.1], rows, size=11, row_h=0.5)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "No foreign keys and no persistence: alerts and keys are lost when the server restarts.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "CipherGuard Shield has no database. The ER diagram shows logical entities."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("9", "JSON routes (+ /docs)", K.FBLUE), (f"{F.PYTEST_TOTAL} + {F.SMOKE_TOTAL}", "pytest tests + smoke checks", K.GREEN), ("5", "React screens", K.ORANGE), ("93 B", "fixed envelope overhead", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        ("Input limits: text 1 MB, file 5 MB, CSV 10 MB and 10,000 rows; empty, non-CSV and malformed base64 return 400; engine not ready returns 503.", {"bullet": True, "space_after": 9}),
        ("Envelope = ciphertext + 12-byte random nonce + 16-byte tag + 65-byte ephemeral public key; HKDF-SHA256 with a fixed context string.", {"bullet": True, "space_after": 9}),
        ("The cloud derives the key by real ECDH with its private key when the demo cache misses (tested).", {"bullet": True, "space_after": 9}),
        ("Console: live scoring with CSV batch, secure transfer with tamper button, alerts, model performance; one command (run.bat) starts both servers.", {"bullet": True, "space_after": 9}),
        ("Not present: authentication, TLS, rate limiting, persistent keys.", {"bullet": True}),
    ], size=16)
    notes[16] = "Highlights and honest absences."

    # ---- 17 results: detection
    s = S[16]
    set_title(s, "Results – Detection on Real UNSW-NB15")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.f3(st["accuracy"]["mean"]), "stacked accuracy, official test (5 seeds)", K.FBLUE), (F.f3(st["f1"]["mean"]), "stacked F1", K.GREEN), (F.f3(st["roc_auc"]["mean"]), "stacked ROC-AUC", K.ORANGE), (F.f3(xg["accuracy"]["mean"]), "XGBoost alone: better", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.1, v, l, c, 26)
    models = [("stack", "Stacked"), ("xgb", "XGBoost"), ("cnn", "1D CNN"), ("mean", "Mean XGB+CNN"), ("rf", "Random forest"), ("dt", "Decision tree"), ("lr", "Logistic reg.")]
    K.bar_chart(s, 0.55, 3.0, 12.2, 3.2, [n for _, n in models],
                [("Accuracy", [OFS[k]["accuracy"]["mean"] for k, _ in models]), ("F1", [OFS[k]["f1"]["mean"] for k, _ in models]),
                 ("False-positive rate (lower is better)", [OFS[k]["fpr"]["mean"] for k, _ in models])],
                colors=[K.FBLUE, K.GREEN, K.ORANGE], number_format="0.00", font_size=11, max_value=1.0, legend=True, label_size=9)
    ftext(s, 0.55, 6.25, 12.2, 0.7, f"Five seeds, stacked accuracy std {st['accuracy']['std']:.4f}; stacking beats the plain mean by only 0.002 and loses to XGBoost. Base-paper protocol (9,000 balanced flows, 10 runs): stacked {F.f3(pst['accuracy']['mean'])}, XGBoost {F.f3(PS['xgb']['accuracy']['mean'])}, random forest {F.f3(PS['rf']['accuracy']['mean'])}: no gain from stacking. The test set is a later capture period than training.", size=12, color=K.FNAVY, italic=True)
    notes[17] = "The CNN lifts recall but doubles the false-alarm rate; XGBoost alone is the best single model on the temporally shifted test."

    # ---- 18 confusion matrix & threshold
    s = S[17]
    set_title(s, "Confusion Matrix & Threshold Analysis")
    wipe(s)
    cm_grid(s, 0.55, 1.75, S42["stack"]["confusion_matrix"], "Stacked, official test (82,332), seed 42")
    cm_grid(s, 4.7, 1.75, S42["xgb"]["confusion_matrix"], "XGBoost alone, same test set")
    cm_grid(s, 8.85, 1.75, SH["confusion_matrix"], "Shipped model, its synthetic split (1,800)")
    K.picture(s, F.fig("f04_threshold_sweep.png"), 0.55, 4.45, 4.3, 2.5)
    ftext(s, 5.2, 4.55, 7.6, 2.4, [
        (f"Stacked at 0.5: recall {F.pct(t5['recall'])}, FPR {F.pct(t5['fpr'])} ({S42['stack']['confusion_matrix'][0][1]:,} of 37,000 normal flows raise a false alarm).", {"bullet": True, "space_after": 6}),
        (f"Raising the threshold to 0.9: accuracy {F.f3(t9['accuracy'])}, precision {F.f3(t9['precision'])}, recall {F.f3(t9['recall'])}, FPR {F.pct(t9['fpr'])}.", {"bullet": True, "space_after": 6}),
        (f"Best-F1 point of the sweep is {max(F.OFF['thresholds'], key=lambda r: r['f1'])['threshold']:.2f}; the sweep was run on the test set, so it shows the trade-off rather than a tuned result.", {"bullet": True, "space_after": 6}),
        ("The shipped threshold stays 0.5.", {"bullet": True}),
    ], size=13)
    notes[18] = "Rows are the actual class. Most errors are normal traffic flagged as attack; the threshold trades missed attacks for fewer false alarms."

    # ---- 19 real vs synthetic and categories
    s = S[18]
    set_title(s, "Results – Real vs Synthetic Data, by Category")
    wipe(s)
    K.bar_chart(s, 0.55, 1.75, 6.4, 3.3, ["Shipped,\nits synthetic\ntest", "Shipped,\nfresh\nsynthetic", "Shipped,\nREAL test", "Real-trained,\nREAL test", "Real-trained,\nsynthetic CSV"],
                [("accuracy", [SH["accuracy"], F.SHIFT["shipped_on_fresh_synthetic_seed7"]["stack"]["accuracy"], shr["accuracy"], st["accuracy"]["mean"], rts["accuracy"]])],
                colors=[K.FBLUE], number_format="0.000", font_size=11, max_value=1.0, legend=False)
    K.picture(s, F.fig("f05_per_category.png"), 7.1, 1.8, 5.7, 3.2)
    ftext(s, 0.55, 5.1, 12.2, 1.8, [
        (f"The shipped {F.pct(SH['accuracy'], 2)} is on synthetic data; 2.5% injected label noise caps accuracy near 97.5% and XGBoost, CNN and stack all score the same. On the 82,332 real flows the same model scores {F.f3(shr['accuracy'])} (AUC {F.f3(shr['roc_auc'])}); real-trained models on the synthetic CSV score {F.f3(rts['accuracy'])}.", {"bullet": True, "space_after": 6}),
        (f"Detected: Generic {pc['Generic']['rate'] * 100:.1f}%, DoS {pc['DoS']['rate'] * 100:.1f}%, Exploits {pc['Exploits']['rate'] * 100:.1f}%; Fuzzers only {pc['Fuzzers']['rate'] * 100:.1f}%; normal traffic correctly passed {pc['Normal']['rate'] * 100:.1f}%. sttl dominates feature importance (a known dataset artefact).", {"bullet": True}),
    ], size=13)
    notes[19] = "The most important finding: where the training data comes from decides everything. Synthetic accuracy says nothing about real traffic."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{F.IP['ids_score_single']['median_ms']:.2f} ms", "score one flow (in process)", K.FBLUE), (f"{F.IP['encrypt_100kb']['median_ms']:.2f} ms", "encrypt 100 KB", K.GREEN),
                                   (f"{F.IP['decrypt_100kb']['median_ms']:.2f} ms", "decrypt 100 KB", K.ORANGE), (f"{F.IP['protected_transfer_100kb']['median_ms']:.2f} ms", f"protected transfer; IDS {share['ids'] * 100:.0f}%", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f09_crypto_latency.png"), 0.55, 2.9, 6.1, 2.35)
    K.picture(s, F.fig("f10_pipeline_latency.png"), 6.8, 2.9, 6.0, 2.35)
    rows = [["Check", "Result"],
            ["pytest (crypto 10, IDS engine 6, API 10)", f"{F.PYTEST_PASSED} of {F.PYTEST_TOTAL} passed in {F.PYTEST_SECONDS:.1f} s; smoke test {F.SMOKE_PASSED} of {F.SMOKE_TOTAL}"],
            ["Tamper tests: ciphertext, tag, nonce, truncation, wrong key", "6 cases x 1,000 trials: 6,000 of 6,000 as expected"],
            ["Nonce uniqueness", "20,000 encryptions of one text: 20,000 distinct nonces and ciphertexts"],
            ["Not tested", "load, concurrency, Docker, other browsers, adversarial evasion of the IDS"]]
    ftable(s, 0.55, 5.35, [4.7, 7.5], rows, size=11, row_h=0.3)
    notes[20] = "Medians on a shared i5-12500 (indicative). SecureChannel costs more than raw AES-GCM on large payloads because of base64 and ephemeral key work; RSA-OAEP decrypt costs 0.45 ms (2048-bit) versus 0.10 ms at 1 KB."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "This project"],
            ["Detector", "Quantum-inspired NN + LeNet, DSMMN fusion", "XGBoost + 1D CNN, LR stacking"],
            ["Encryption", "Enhanced ECC, no formal proof", "ECDH P-256 + HKDF + AES-256-GCM"],
            ["Data / protocol", "9,000 balanced UNSW-NB15 flows, 70/15/15, 10 runs", "same style of protocol, 10 runs; plus official 175,341 -> 82,332 split"],
            ["Accuracy", f"{bp['proposed_mean_accuracy']:.4f} mean ({bp['proposed_ablation_accuracy']:.4f} ablation)", f"{F.mstd(PS, 'stack', 'accuracy', 3)} (paper protocol); {F.mstd(OFS, 'stack', 'accuracy', 3)} (official)"],
            ["Precision / FPR", f"{bp['precision_at_90']:.2f} / {bp['fpr']:.2f}", f"{F.f3(pst['precision']['mean'])} / {F.f3(pst['fpr']['mean'])} (paper protocol)"],
            ["Encrypt / decrypt", f"{bp['enc_time_s']} s / {bp['dec_time_s']} s (payload not stated)", f"{sc['1 KB']['encrypt']['median_ms']:.2f} / {sc['1 KB']['decrypt']['median_ms']:.2f} ms (1 KB); {sc['100 KB']['encrypt']['median_ms']:.2f} / {sc['100 KB']['decrypt']['median_ms']:.2f} ms (100 KB)"]]
    ftable(s, 0.55, 1.85, [1.5, 2.9, 2.9], rows, size=11, row_h=0.62)
    K.picture(s, F.fig("f08_base_paper_comparison.png"), 8.1, 1.85, 4.75, 2.1)
    ftext(s, 8.1, 4.1, 4.75, 2.8, [
        (f"About one point below the paper's proposed model ({bp['proposed_mean_accuracy']:.3f}) and slightly above its best baseline (SqueezeNet {bp['baselines']['SqueezeNet']:.3f}).", {"bullet": True, "space_after": 6}),
        ("Subsets, hyper-parameters and test fraction are not published: indicative only.", {"bullet": True, "space_after": 6}),
        (f"EECC and the quantum-inspired network are not reproduced; the shipped {F.pct(SH['accuracy'], 2)} (synthetic) is excluded.", {"bullet": True}),
    ], size=12)
    notes[21] = "State plainly that datasets differ. No superiority over the base paper is claimed."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Live Scoring and Alerts")
    wipe(s)
    K.picture(s, crop_shot("03_live-scoring.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("08_alerts.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: the attack preset scored as ATTACK, severity HIGH, ensemble 0.84 (XGBoost 0.45, CNN 1.00). Right: alerts raised by the 120-row sample batch. The bundled sample is synthetic, drawn from the demo CSV.", size=13, color=K.FNAVY)
    notes[22] = "Real captures of the running console. Note the XGBoost and CNN disagree on this record; the stack decides."
    s = S[22]
    set_title(s, "Output – Secure Transfer and Tamper Test")
    wipe(s)
    K.picture(s, crop_shot("06_secure-transfer-verified.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("07_secure-transfer-tampered.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: text encrypted with a fresh ephemeral key, then decrypted and verified at the mock cloud (real ciphertext, nonce, tag, ephemeral public key). Right: one ciphertext byte flipped - the GCM tag check fails and the API answers 400.", size=13, color=K.FNAVY)
    notes[23] = "Encrypt, decrypt-and-verify, and the tamper button."
    s = S[23]
    set_title(s, "Output – Real UNSW-NB15 Flows: Shipped vs Real-Trained")
    wipe(s)
    K.picture(s, crop_shot("10_real-flows-shipped-model.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("12_real-flows-real-model.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, f"The same 120 real flows (60 normal, 60 attack, official test set). Left: shipped synthetic-trained model agrees with the labels on 78 of 120 (65.0%). Right: real-trained bundle agrees on 101 of 120 (84.2%); its Model-performance page shows {F.pct(S42['stack']['accuracy'], 2)} on 82,332 flows.", size=13, color=K.FNAVY)
    notes[24] = "The failure case and the fix side by side. The real-trained bundle is served with CIPHERGUARD_MODELS_DIR; it is not the default."

    # ---- 25 conclusion and objectives
    s = S[24]
    set_title(s, "Conclusion & Objectives vs Outcomes")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 1.3, [
        (f"A working platform: standard authenticated encryption (6,000 of 6,000 tamper trials as expected) and a detector reaching {F.f3(st['accuracy']['mean'])} stacked / {F.f3(xg['accuracy']['mean'])} XGBoost on real UNSW-NB15; {F.PYTEST_PASSED} tests and {F.SMOKE_PASSED} smoke checks pass.", {"bullet": True, "space_after": 4}),
        (f"Findings: detection, not encryption, is the overhead ({share['ids'] * 100:.0f}% of a protected transfer); stacking adds nothing over XGBoost; synthetic accuracy does not transfer.", {"bullet": True}),
    ], size=14)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Hybrid AES + RSA", "Partly: ECDH + AES-256-GCM implemented; RSA-OAEP only benchmarked"],
            ["2", "Key generation, distribution, storage", "Generation and distribution met; storage not (keys in memory)"],
            ["3", "Encrypt before upload, decrypt after download", "Met against a mock cloud; no real provider"],
            ["4", "Classifier on benchmark dataset", "Met in experiments (real UNSW-NB15); shipped default is synthetic"],
            ["5", "Flag and log suspicious sessions", "Partly: alert log in memory, not coupled to transfers"],
            ["6", "Evaluate time, overhead, accuracy", "Met: crypto latency, per-transfer overhead, detection metrics"]]
    ftable(s, 0.55, 3.4, [0.5, 4.2, 7.5], rows, size=12, row_h=0.5, align=["c", "l", "l"])
    notes[25] = "Be explicit about which objectives were met, partly met or not met."

    # ---- 26 limitations and future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 5.1, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("Default model trained on synthetic data", {"bullet": True, "space_after": 10}),
        ("22.6% false alarms on normal traffic at 0.5", {"bullet": True, "space_after": 10}),
        ("No live packet capture; flows are records or CSV rows", {"bullet": True, "space_after": 10}),
        ("Mock in-process cloud, no auth or TLS, keys in memory", {"bullet": True, "space_after": 10}),
        ("No adversarial or load testing", {"bullet": True, "space_after": 10}),
        ("EECC and quantum-inspired network not implemented", {"bullet": True, "space_after": 10}),
        ("Timings from one shared machine", {"bullet": True}),
    ], size=17)
    fbox(s, 6.8, 1.8, 6.0, 5.1, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Ship the real-trained bundle as default; retrain on newer data", {"bullet": True, "space_after": 10}),
        ("Calibrate the threshold on validation data; cut false alarms", {"bullet": True, "space_after": 10}),
        ("Capture live traffic and gate transfers on the verdict", {"bullet": True, "space_after": 10}),
        ("Real cloud storage, persistent key management, TLS, login", {"bullet": True, "space_after": 10}),
        ("Evaluate evasion attacks and the CICIDS / multi-class setting", {"bullet": True, "space_after": 10}),
        ("Explore standard post-quantum key exchange", {"bullet": True}),
    ], size=17)
    notes[26] = "Limitations first, then concrete next steps."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    if REFS:
        ids = [1, 2, 3, 4, 5, 7, 9, 10, 11, 12, 13, 14, 20, 24]
        short = [f"[{n}] {clean_ref(REFS[n - 1])}" for n in ids]
        short = [t if len(t) < 330 else t[:327] + "..." for t in short]
        ftext(s, 0.55, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[:7]], size=11)
        ftext(s, 6.8, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[7:]], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "CipherGuard")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
