"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "cloud security, authenticated encryption, AES-256-GCM, ECDH, intrusion detection, XGBoost, UNSW-NB15"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 32)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


# ---------------------------------------------------------------- slide text (batch documents + measured facts only)
PROBLEM_STATEMENT = [
    "Data placed in a public cloud is readable by the provider unless it is encrypted before upload",
    "Symmetric encryption alone leaves the key exchange itself open to interception",
    "Encryption cannot detect an attacker who already holds valid credentials",
    "Most academic systems address either cryptography or intrusion detection, rarely both together",
    "Need: a transmission framework that encrypts data and, on the same path, monitors traffic for intrusion with machine learning",
]
OBJECTIVES = [
    "Design a hybrid encryption scheme: a symmetric cipher for the payload and asymmetric key exchange",
    "Implement secure key generation, distribution and storage between sender and receiver",
    "Build a transmission module that encrypts before upload and decrypts after download",
    "Train a machine-learning classifier on a benchmark intrusion dataset to detect anomalous access",
    "Integrate the detection layer so suspicious sessions are flagged and logged during transfer",
    "Evaluate encryption and decryption time, throughput overhead and intrusion detection accuracy",
]


def build(out_path):
    from .lit import GAPS, LIT_TABLE
    from .refs import REFS

    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    bp = F.BASE_PAPER_RESULTS

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: CipherGuard Shield puts two protections on one cloud data path - authenticated encryption for confidentiality and integrity, "
                "and a machine-learning detector that flags attack traffic - and measures what both cost together.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design details and measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: protect the data and watch the traffic")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Data sent to a public cloud can be intercepted, altered or read by the provider. Encryption protects the content but cannot notice an attacker who already has valid access, and most academic systems treat encryption and intrusion detection separately.", K.ORANGE, "!"),
        ("Our idea", "Put both protections on one data path: authenticated encryption (ECDH P-256, HKDF-SHA256, AES-256-GCM) for confidentiality and integrity, and a machine-learning detector (XGBoost and a 1D CNN, stacked) that flags attack traffic and raises alerts.", K.BLUE, "i"),
        ("The deliverable", "CipherGuard Shield: a FastAPI service and a React console to encrypt text and files, verify at a mock cloud, test tampering, score flows and CSV batches, review alerts and see model performance, with the combined overhead measured.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=20)
    notes[3] = "Three points: the weakness of encryption-only and detection-only systems, the idea of putting both on one path, and what is delivered. Everything is local and runs without a GPU."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "A quantum-inspired network (FLPF with EVD-QNN) and LeNet classify UNSW-NB15 flows; their scores are fused (DSMMN-SLF). Data judged safe is then encrypted with an enhanced ECC scheme (EECC). The quantum part runs on classical hardware.", K.BLUE, "M"),
        ("Data and reported results", f"{bp['samples']:,} balanced samples, 70/15/15 split, {bp['runs']} runs. Mean accuracy {bp['proposed_mean_accuracy']:.4f}; best baseline SqueezeNet {bp['baselines']['SqueezeNet']:.4f}; precision {bp['precision_at_90']:.2f} at 90% training data; FPR {bp['fpr']:.2f}. EECC: {bp['enc_time_s']:.2f} s to encrypt, {bp['dec_time_s']:.2f} s to decrypt.", K.GREEN, "R"),
        ("Relevance to this project", "Same principle: detection and encryption in one framework. Different means: we use standard, vetted primitives (AES-GCM, ECDH, HKDF) and an XGBoost + CNN stack. The paper itself notes that EECC has no formal security proof.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=16)
    notes[4] = ("Read the citation. The paper improves both halves - a fused detector and a new elliptic-curve cipher - and reports its numbers on a 9,000-sample balanced subset. "
                "We keep the idea of one framework but replace the novel cipher by standard, analysed primitives; our numbers are never compared directly without stating the protocol.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for r in LIT_TABLE[:10]:
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.3, [2.6, 0.65, 4.6, 4.28], rows, font_size=11, row_h=0.5, header_size=12, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey (base paper plus four recent works) shows hybrid encryption, anomaly detection and federated detection, mostly studied separately. "
                "Figures are quoted only where they were read from the paper or its record; the complete verified reference list is in the report.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: data and protocol decide the number")
    wipe(s)
    vals = [F.SHIPPED["accuracy"], F.PAPS["stack"]["accuracy"]["mean"], F.OFFS["stack"]["accuracy"]["mean"], F.SHIFT["shipped_on_real_official_test"]["stack"]["accuracy"]]
    gf = K.bar_chart(s, 0.6, 1.4, 6.1, 5.2, ["Synthetic data\n(shipped model)", "Real flows,\nbalanced 9,000\n(paper-style)", "Real flows,\nofficial\ntrain/test split", "Shipped model\nscored on\nreal flows"],
                     [("Accuracy", vals)], colors=[K.BLUE], number_format="0.000", font_size=11, max_value=1.0)
    for idx in (0, 3):
        pt = gf.chart.plots[0].series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 7.0, 1.45, 5.75, 4.9, [
        ("**Pilot measurements of this project** (same stacked detector, accuracy):", {"space_after": 8}),
        ("Synthetic data flatters the model: 0.973, but only 0.660 on real UNSW-NB15 flows.", {"bullet": True, "space_after": 8}),
        ("A random balanced subset gives 0.938; the official, time-separated split gives 0.890.", {"bullet": True, "space_after": 8}),
        (f"The base paper's own baselines span 0.899 to {bp['baselines']['SqueezeNet']:.3f} on UNSW-NB15 [1].", {"bullet": True, "space_after": 8}),
        ("UNSW-NB15 offers 175,341 training and 82,332 testing flows in nine attack families [7].", {"bullet": True}),
    ], size=16, color=K.INK)
    K.text(s, 7.0, 6.4, 5.75, 0.5, "Context for the problem, not final results; full protocol in Review 2.", size=12, color=K.TXT, italic=True)
    notes[6] = ("This is early evidence from our own pilot runs, shown because it frames the problem: a detector's headline accuracy depends on where the data comes from and how it is split. "
                "The shipped demo model is trained on synthetic data and must not be read as a real-traffic result.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("User data", K.NAVY), ("Encrypt (AES, RSA, hybrid)", K.NAVY), ("Upload to cloud", K.NAVY), ("Separate IDS / firewall", K.NAVY)], size=15, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Encryption-only protection",
           "Symmetric ciphers such as AES are fast but need a secure channel for the key.\nRSA solves key exchange but is too slow for bulk data.\nHybrid schemes wrap an AES key with RSA, yet cannot see an attack that uses valid credentials.", K.BLUE, "1", body_size=19)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Detection-only protection",
           "Machine-learning intrusion detection reacts to attacks that encryption cannot prevent.\nIt is normally evaluated alone, on benchmark datasets, without the cost of the encryption layer on the same path.", K.ORANGE, "2", body_size=19)
    notes[7] = "Today the two protections are separate products or separate chapters of a paper: crypto protects the bytes, an IDS watches the traffic, and nobody measures the pair."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Key distribution", "Blind to valid access", "Separate concerns", "Hard to compare"]
    bodies = [
        "Symmetric encryption alone leaves the key exchange open to interception; RSA alone is too slow for bulk data.",
        "Encryption cannot detect an attacker who already holds valid credentials; a detector is needed as a second layer.",
        "Cryptography and intrusion detection are rarely built or measured together, so the combined overhead is unknown.",
        "Accuracy on UNSW-NB15 depends on the subset, balance and split, so published numbers cannot be lined up directly.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=17)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: new ciphers such as the base paper's EECC come without a formal security proof, so they are harder to trust than AES-GCM [1].", size=15, color=K.INK, anchor="m")
    notes[8] = "Four limitations from the survey and the batch problem statement. Each maps to a design choice: standard primitives, a detector on the same path, and an explicit evaluation protocol."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.35 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=18, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as: both operate together on a single transmission path and the combined overhead is measured. "
                "Gaps 4 and 5 come from our own reading of the base paper and from the pilot runs on real and synthetic data.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=19, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Text or a file (up to 5 MB) to send; network-flow records, one at a time or as a CSV (up to 10,000 rows, 42 features)"),
                                 ("Output", "Authenticated ciphertext envelope, verified plaintext or a rejection; an attack or normal verdict with confidence and severity"),
                                 ("Record", "Alert log of the latest 500 attacks, a model-performance report and measured timings")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=14, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a framework that encrypts the data and, on the same path, watches the traffic with machine learning."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Implementation note: key exchange uses ECDH P-256 instead of RSA; an RSA-OAEP + AES-GCM hybrid is built only as a benchmark baseline.", size=14, color=K.TXT, italic=True)
    notes[11] = ("The six objectives from the proposal. Later reviews report which were met. Objective 1 is met by substitution: ECDH P-256 gives a fresh key per message and a smaller key exchange than RSA; "
                 "the RSA hybrid is measured as the comparison. Key storage (objective 2) and automatic blocking (objective 5) will be only partly met.")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Encrypt", "A fresh P-256 key pair per message; ECDH with the cloud's public key; HKDF-SHA256 gives a 256-bit key; AES-256-GCM adds a 12-byte nonce and a 16-byte tag.", K.BLUE, "1"),
        ("Detect", "Flows are scaled and one-hot encoded; XGBoost and a 1D CNN each give an attack probability; a logistic-regression meta-model fuses them.", K.GREEN, "2"),
        ("Alert", "Attacks above 0.5 are labelled with a severity band and appended to a 500-entry alert log shown in the console.", K.ORANGE, "3"),
        ("Evaluate", "Encryption time by payload size against an RSA hybrid, detection on real UNSW-NB15, and the combined cost of one protected transfer.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "E", 18)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "Envelope = ciphertext + nonce (12 B) + tag (16 B) + ephemeral public key (65 B)", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "Delivered as a FastAPI service and a React console with no database and no GPU. A tampered ciphertext fails tag verification and is rejected.", size=14, color=K.TXT)
    notes[12] = ("Four stages. Encryption uses only standard, reviewed primitives from the cryptography library. The detector reuses a proven tabular pipeline. "
                 "The two lanes run side by side in the console; the API does not block a transfer on the verdict, which is stated as a scope limit.")

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.2, 7.5, 5.9)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite console (port {F.WEB_PORT}) and the original HTML dashboard.", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}) with size and format guards on every upload.", {"bullet": True}),
        ("Core: SecureChannel (ECDH, HKDF, AES-GCM) and the ThreatSense Engine (XGBoost, CNN, stacking).", {"bullet": True}),
        ("Data: model files, a metrics report and CSVs on local disk; alerts and keys stay in memory.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the console calls the API, the API calls the two core modules, and the modules read model files from disk. There is no database."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working platform", "Measured overhead", "Real-data detector", "Honest comparison", "Stated limits"]
    bodies = [
        "Encrypt, decrypt and verify text and files, reject tampering, score single flows and CSVs, keep alerts, show model performance.",
        "Encryption and decryption time by payload size against an RSA-OAEP hybrid, and the cost of detection plus encryption together.",
        "Detector trained and evaluated on real UNSW-NB15, with baselines, per-category and threshold analysis and repeated runs.",
        "Comparison with the base paper under a similar protocol, with every difference in data and method stated.",
        "Mock cloud, no live packet capture, synthetic-trained demo model: each limit documented, not hidden.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=17)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=18)
    notes[14] = "These are the outcomes we commit to; they are targets at this stage. Review 2 and the final review report what was actually measured, including what could not be."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [n for n in (1, 2, 3, 4, 5, 6, 7, 9, 12, 14) if n <= len(REFS)]
    K.text(s, 0.6, 1.3, 12.13, 5.7, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 4}) for n in idx], size=13, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why ECDH instead of RSA (fresh key per message, smaller key exchange; RSA measured as baseline); why standard AES-GCM instead of a new cipher (analysed, authenticated, no proof gap); "
                 "why synthetic data in the demo model (portable out of the box; real-data results are reported separately).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
