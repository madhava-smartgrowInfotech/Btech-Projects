"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the UNSW Canberra Cyber Range Lab for creating and sharing the UNSW-NB15 dataset, and the maintainers of the "
        "open-source software used in this project (FastAPI, React, Vite, the `cryptography` package, XGBoost, scikit-learn, PyTorch, pandas and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    sh, off, pap = F.SHIPPED, F.OFFS, F.PAPS
    real = F.SHIFT["shipped_on_real_official_test"]["stack"]
    tamper = F.CRY["tamper_tests"]
    n_tamper = sum(v["trials"] for k, v in tamper.items() if k != "untouched message accepted")
    ok_tamper = sum(v["as_expected"] for k, v in tamper.items() if k != "untouched message accepted")
    d.p("Data that travels to a cloud service crosses networks its owner does not control. Encryption protects the content but cannot notice an attacker who is already "
        "sending traffic, and intrusion detection watches the traffic but never sees inside the payload; most work treats the two separately. This project builds "
        f"{F.PRODUCT}, a platform that places both protections around one cloud data path. Its SecureChannel module encrypts every payload with a fresh ECDH P-256 key pair, "
        "HKDF-SHA256 and AES-256-GCM, so any change to the ciphertext is rejected at the cloud. Its ThreatSense Engine classifies network flows with XGBoost and a one-dimensional "
        "convolutional network fused by logistic-regression stacking; a FastAPI service and React console show verdicts and alerts.")
    d.p(f"The shipped demonstration model was trained on generated synthetic data and scores {F.pct(sh['accuracy'])} on it, but this figure does not transfer: on the "
        f"{F.OFF['test_rows']:,} real UNSW-NB15 test flows the same model reaches {F.pct(real['accuracy'])}. Trained on the official UNSW-NB15 training set, the stacked ensemble reaches "
        f"{F.mstd(off, 'stack', 'accuracy')} accuracy on the official test set (XGBoost alone {F.mstd(off, 'xgb', 'accuracy')}) and {F.f3(pap['stack']['accuracy']['mean'])} "
        f"under the base paper's 9,000-flow protocol, about one point below the paper's reported {F.f3(F.BASE_PAPER_RESULTS['proposed_mean_accuracy'])}. "
        f"All {ok_tamper:,} of {n_tamper:,} tampered messages were rejected, and one protected 100 KB transfer costs {F.f2(F.IP['protected_transfer_100kb']['median_ms'])} ms, "
        f"{F.pct(F.IP['share_of_transfer_median']['ids'], 0)} of it detection. The cloud is a mock, nothing is captured live, and no adversarial testing was done.")
    d.p("**Keywords:** cloud security, secure data transmission, AES-256-GCM, elliptic-curve Diffie-Hellman, intrusion detection, machine learning, XGBoost, 1D CNN, UNSW-NB15.")


ABBR = [
    ("AES", "Advanced Encryption Standard"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("BCE", "Binary Cross-Entropy"),
    ("CNN", "Convolutional Neural Network"), ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"),
    ("DFD", "Data Flow Diagram"), ("DoS", "Denial of Service"), ("ECC", "Elliptic Curve Cryptography"), ("ECDH", "Elliptic Curve Diffie-Hellman"),
    ("EECC", "Enhanced Elliptic Curve Cryptography"), ("ER", "Entity-Relationship"), ("FNR", "False-Negative Rate"), ("FPR", "False-Positive Rate"),
    ("GCM", "Galois/Counter Mode"), ("HKDF", "HMAC-based Key Derivation Function"), ("HTTP", "Hypertext Transfer Protocol"), ("IDS", "Intrusion Detection System"),
    ("JSON", "JavaScript Object Notation"), ("LR", "Logistic Regression"), ("MCC", "Matthews Correlation Coefficient"), ("ML", "Machine Learning"),
    ("NIDS", "Network Intrusion Detection System"), ("NIST", "National Institute of Standards and Technology"), ("OAEP", "Optimal Asymmetric Encryption Padding"),
    ("QNN", "Quantum Neural Network"), ("RF", "Random Forest"), ("ROC", "Receiver Operating Characteristic"), ("RSA", "Rivest-Shamir-Adleman"),
    ("TLS", "Transport Layer Security"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("UNSW-NB15", "University of New South Wales Network-Based 2015 dataset"),
    ("XGBoost", "Extreme Gradient Boosting"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
