"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from . import facts as F
from . import pptkit as K
from .ch1 import OBJECTIVES, OBJECTIVE_STATUS
from .ch2 import LIT_TABLE
from .ch8 import analysis
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520
FRAME = "POST frame (real face JPEG 640x480, monitoring on)"


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def cap(slide, x, y, w, text, size=12):
    ftext(slide, x, y, w, 0.5, text, size=size, color=K.FNAVY, align="c", italic=True)


def pc(x, d=1):
    return f"{x * 100:.{d}f}%"


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "eye aspect ratio, driver monitoring, impairment warning, MediaPipe, FastAPI"
    cp.last_modified_by = ""
    notes = {}

    a = analysis()
    cew, geo = F.CEW, F.GEO
    api = F.API_LAT
    fr = api[FRAME]
    rules = F.RULES

    def rob(substr):
        for c in F.ROB["conditions"]:
            if substr in c["condition"]:
                return c
        raise KeyError(substr)

    def sweep_row(closure, persist):
        for r in rules["parameter_sweep"]:
            if r["closure_seconds"] == closure and r["persistence_seconds"] == persist:
                return r
        raise KeyError((closure, persist))

    default_row, fast_row = sweep_row(1.5, 1.0), sweep_row(1.0, 0.5)
    gated, gdef = a["gated"], a["geo_default"]
    gbest = a["geo_best"]
    tests = F.TESTS

    # footers
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh in S[0].shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_.title()}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE}  ·  Academic Year {F.YEAR}")
    notes[1] = ("Introduce the team and the title. Be upfront: the title says alcohol, but what we built is an explainable eye-closure and head-tilt "
                "impairment-and-drowsiness warning for a rider. It cannot confirm alcohol, and we say so throughout.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap & Problem Statement\nObjectives and their status\n"
                               "Proposed System & Architecture\nMethodology, Dataset & Rules\nDecision Workflow, Interaction & Database\n"
                               "Implementation Highlights\nResults, Threshold Analysis, Robustness, Performance & Testing\nComparison with Base Paper\n"
                               "Output Screens\nConclusion, Limitations, Future Scope & References", size=18)
            for para in sh.text_frame.paragraphs:
                para.space_before = Pt(0)
                para.space_after = Pt(4)
    notes[2] = "The agenda follows the report: problem, design, method, results (with their limits), comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("Impaired riding is hard to police continuously: breath tests are occasional and need an officer; a rider's face, however, shows eye closure and head posture all the time.", {"bullet": True, "space_after": 9}),
        ("We built **Smart Rider**, a webcam system that measures the eye aspect ratio, closure time, closed-eye fraction and head roll, refuses to guess when the picture is poor, and shows one of three outputs with its reason: Normal, Possible impairment or drowsiness, Unable to assess.", {"bullet": True, "space_after": 9}),
        (f"Measured: on {cew['images_total']:,} public CEW faces the eye-aspect-ratio separates closed from open eyes with AUC **{a['geo_auc']:.3f}** ({pc(gdef['accuracy'])} accuracy at the default threshold); a closure must last about **{a['min_closure_that_warns']} s** to warn; one camera frame takes **{fr['p50_ms']:.0f} ms** through the web API; 20 unit tests and 34 API checks pass.", {"bullet": True, "space_after": 9}),
        ("**Honest scope:** it is not an alcohol detector. No deep-learning classifier was trained (no intoxication data existed), gaze is not measured, no ignition or contact action exists, and no study with riders was run.", {"bullet": True}),
    ], size=20)
    notes[3] = "Read the first two bullets, then the measured numbers, then stress the last bullet: what the batch objectives promised and what was and was not delivered."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper [1]"],
            ["Method", "Four-point EAR-4 and six-point EAR from MediaPipe; head pose; Whisper-tiny voice commands; fuzzy belief rule-based fusion; personal threshold from more than 300 frames"],
            ["Data", "NTHU-DDD: four subjects, 65,834 frames (29,804 open, 36,030 closed), 80/20 frame split; Random Forest and XGBoost"],
            ["Reported results", "EAR-4 recall 0.894 and F1 0.845 (Subject 2, Random Forest; six-point EAR 0.856 and 0.813); 25 FPS on a Jetson Orin Nano 8 GB"],
            ["Stated limits", "Targets drowsiness, fatigue and distraction, not alcohol; four subjects; subject-independent validation left as future work"]]
    ftable(s, 0.55, 2.9, [2.0, 10.2], rows, size=14, row_h=0.8)
    notes[4] = ("This is the batch's base paper. It is about drowsiness on embedded hardware. We borrow the idea of eye aspect ratio plus personal calibration; "
                "we do not reproduce EAR-4, the voice interface or the fuzzy belief engine.")

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    short = [
        ["Baroutian et al., 2025 [2]", "Landmark graph attention network + 3D ResNet, fused", "202 YouTube videos (101 sober, 101 intoxicated)", "Accuracy 95.82 %, precision 0.977, recall 0.97", "Scraped videos; GPU model; not a controlled study"],
        ["Deuber et al., 2026 [3]", "Smartwatch accelerometer + HRV; logistic regression and 1D CNN", "Test-track study, 54 participants", "AUROC 0.88 (any alcohol), 0.86 (above 0.05 g/dL)", "Needs a worn device; preprint under review"],
        ["Acharya, 2024 [5]", "Thermal images; NasNetMobile, MobileNet", "50 volunteers, four drinking levels", "NasNetMobile 85.10 %; MobileNet 74.07 % (sober vs drunk)", "Needs thermal camera; abstract only read"],
        ["Chai et al., 2025 [8]", "Thermal facial patterns; CNN features; MLP, SVM, KNN, RF", "Tufts Face Database + own thermal images", "MLP about 90 % accuracy", "Thermal camera; abstract only read"],
        ["Soukupova and Cech, 2016 [9]", "Eye aspect ratio from six landmarks", "Three standard blink datasets", "Comparable with the state of the art", "Blink detection, not impairment"],
        ["Song et al., 2014 [12] (CEW)", "Eye-closeness detection from still images", "2,423 faces: 1,192 closed, 1,231 open", "Used here as test data", "Closed and open from different sources"],
    ]
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"]] + short
    ftable(s, 0.55, 1.85, [2.1, 2.9, 2.6, 2.6, 2.0], rows, size=11, row_h=0.7)
    notes[5] = ("Intoxication from facial video, thermal images or a smartwatch has been studied; eye aspect ratio is the standard cue for blinks and drowsiness. "
                "Only results we could read are shown. The DIF dataset [4] is cited in the report; many of its source videos are no longer reachable.")

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 1.0, "Batch survey: the surveyed work concentrates on drowsiness and embedded driver-state assessment; the gap is an intoxication-specific vision model coupled to a prevention action rather than a warning alone.", size=15, color=K.FNAVY, italic=True)
    ftext(s, 0.55, 3.0, 12.2, 3.9, [
        ("**Cause is not separable:** eye-closure and head-posture monitors work for drowsiness, but the cues overlap with alcohol impairment, so such a monitor cannot tell the cause.", {"bullet": True, "space_after": 9}),
        ("**Data:** intoxication-specific video models need labelled intoxicated people; the best-known public set has partly disappeared and none suited this team.", {"bullet": True, "space_after": 9}),
        ("**Runnable path:** few works give a complete, explainable route from webcam frame to logged warning, with quality checks that refuse to guess.", {"bullet": True, "space_after": 9}),
        ("**We address** the third point and part of the first. We do **not** close the data gap or the prevention-action gap.", {"bullet": True}),
    ], size=20)
    notes[6] = "Four points from the survey. We built the complete explainable path and measured the eye-closure component; the data and prevention-action gaps remain open."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.9, 12.2, 2.1)
    ftext(s, 0.85, 1.95, 11.6, 2.0, "Given a live camera view of a rider's face, decide continuously whether the visible **eye and head behaviour** shows a sustained pattern that may indicate **impairment or drowsiness**, say **Unable to assess** when the picture cannot be trusted, give the reason for every output and keep a numeric record, through an interface a non-expert can run.", size=19, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Webcam frames, 640 x 480, from the browser or the desktop app; kept in memory only"),
                                 ("Output", "Normal, Possible impairment or drowsiness, or Unable to assess, with the reason"),
                                 ("Record", "Numeric event history in SQLite, CSV export, editable thresholds, calibration")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.3, 3.9, 2.0, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.4, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 4.9, 3.5, 1.3, b_, size=15)
    notes[7] = "Classify, explain, record. A camera cannot confirm alcohol, so the output says possible impairment or drowsiness."

    # ---- 8 objectives with status
    s = S[7]
    set_title(s, "Objectives and Their Status")
    wipe(s)
    ev = ["MediaPipe landmarks on 640 x 480 frames; 31 ms per frame through the API",
          "EAR, closure, closed fraction, blink rate, head roll; gaze not measured",
          "No intoxication data; no model trained; live decisions are rules",
          "Three outputs with a reason, warning and optional bell; no learned score",
          "Warning and bell only; no ignition control, no contact notification",
          "Public eye data, degraded images, scripted streams; no rider study"]
    short_obj = ["Capture facial video; detect face and landmarks in real time",
                 "Extract eye closure, blink, gaze and head cues",
                 "Train a deep-learning sober vs intoxicated classifier",
                 "Run inference continuously; alert above a threshold",
                 "Integrate a prevention action (ignition warning, contact)",
                 "Evaluate accuracy, false positives, low light, occlusion"]
    rows = [["#", "Objective (batch proposal)", "Status", "Evidence"]]
    for i in range(6):
        rows.append([str(i + 1), short_obj[i], OBJECTIVE_STATUS[i][0], ev[i]])
    ftable(s, 0.55, 1.85, [0.5, 4.3, 1.9, 5.5], rows, size=12.5, row_h=0.68, align=["c", "l", "l", "l"])
    notes[8] = ("Two objectives achieved (one as rules), two partly, two not achieved: the deep-learning classifier and the prevention action. "
                "The full wording of each objective and its evidence is in Chapter 1 and Chapter 8 of the report.")

    # ---- 9 proposed system & architecture
    s = S[8]
    set_title(s, "Proposed System & Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.7, 5.1)
    ftext(s, 8.5, 1.95, 4.3, 4.9, [
        ("**Proposed:** a webcam monitor that turns eye and head cues into one of three outputs, with the reason", {"bullet": True, "space_after": 8}),
        ("Browser app (React, port 5306) and the original Tk app share the same Python modules", {"bullet": True, "space_after": 8}),
        ("FastAPI (port 8306) holds one session per tab; frames stay in RAM", {"bullet": True, "space_after": 8}),
        ("core.py rules, vision.py MediaPipe, storage.py SQLite", {"bullet": True, "space_after": 8}),
        ("Quality gates always override the rules", {"bullet": True}),
    ], size=15)
    notes[9] = "The pipeline in one strip, then the four tiers. The decision logic is the same for the desktop and the web interface."

    # ---- 10 methodology
    s = S[9]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.9, 12.2, 1.1, [("Frame 640x480", K.FNAVY), ("Face landmarks", K.FBLUE), ("7 quality gates", K.PURPLE), ("EAR, roll, median", K.GREEN), ("Rules + persistence", K.ORANGE), ("Verdict + log", K.FNAVY)], size=13, gap=0.25, font=FONT)
    ftext(s, 0.55, 3.35, 12.2, 3.5, [
        ("Calibration: 8 to 15 s of open eyes; threshold = 0.72 x median EAR, clipped to 0.12 - 0.32.", {"bullet": True, "space_after": 9}),
        ("Verification is layered: 20 unit tests, a 34-check API smoke test, and measurements on public eye data, degraded images and scripted streams.", {"bullet": True, "space_after": 9}),
        ("The offline research pipeline (subject-disjoint, nearest-centroid) is built but has only run on a seeded synthetic file; it never drives the live decision.", {"bullet": True, "space_after": 9}),
        ("Every number comes from seeded scripts: measure_cew.py, measure_rules.py, measure_api.py, make_figures.py.", {"bullet": True}),
    ], size=20)
    notes[10] = "Methodology: pipeline, calibration, the layers of verification and the fact that the research classifier is not used live."

    # ---- 11 dataset, features, rules
    s = S[10]
    set_title(s, "Dataset, Features & Rules")
    wipe(s)
    ftile(s, 0.55, 1.85, 2.9, 1.05, f"{cew['images_total']:,}", "CEW faces: 1,192 closed, 1,231 open", K.FBLUE, 26)
    ftile(s, 0.55, 3.05, 2.9, 1.05, "478", "landmarks per face (pre-trained MediaPipe)", K.GREEN, 26)
    ftile(s, 0.55, 4.25, 2.9, 1.05, "0.21", "default EAR threshold (engineering heuristic)", K.ORANGE, 26)
    K.picture(s, F.fig("landmark_examples.png"), 3.7, 1.8, 4.9, 3.5)
    rows = [["Rule", "Default"], ["Continuous eye closure", ">= 1.5 s"], ["Closed fraction, last 10 s", ">= 0.55"], ["Head roll with low EAR", ">= 25 deg"], ["Persistence before warning", "1.0 s"], ["Alert cooldown", "15 s"]]
    ftable(s, 8.8, 1.9, [2.6, 1.4], rows, size=12, row_h=0.46)
    ftext(s, 0.55, 5.55, 12.2, 1.4, [
        ("No intoxication dataset was used. CEW (Internet closed-eye faces plus LFW open-eye faces) tests only the eye-open or closed component.", {"bullet": True, "space_after": 4}),
        ("The research CSV demo is **SYNTHETIC** (seeded, 6 subjects x 2 labels) and shows the pipeline runs, not performance.", {"bullet": True}),
    ], size=14)
    notes[11] = ("Dataset, features and rules. The landmark model is Google's, not trained by us. The rule values are demonstration choices, not tuned on data; "
                 "the CEW sweep later suggested 0.17 but the default was left at 0.21.")

    # ---- 12 decision workflow
    s = S[11]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 5.2, 5.1)
    ftext(s, 6.0, 1.95, 6.8, 4.9, [
        ("**Gate first:** a bad frame, a timing gap or an invalid value clears the evidence and gives Unable to assess.", {"bullet": True, "space_after": 10}),
        ("**Smooth:** median of the last three EAR values; closed means below the threshold.", {"bullet": True, "space_after": 10}),
        ("**Three rules** are checked every frame: closure time, closed fraction over 10 s, tilt with low EAR.", {"bullet": True, "space_after": 10}),
        ("**Persist:** a rule must hold 1 s more, with 2 s of reliable history, before the warning shows its reasons.", {"bullet": True, "space_after": 10}),
        ("**Alert:** the optional bell has a 15 s cooldown that survives signal loss.", {"bullet": True}),
    ], size=17)
    notes[12] = "The flowchart is drawn from RiskEngine.update in core.py. Blink rate is displayed but never warns on its own."

    # ---- 13 interaction
    s = S[12]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_diagram.png"), 0.55, 1.8, 7.9, 5.1)
    K.picture(s, F.diag("08_use_case_diagram.png"), 8.6, 1.8, 4.2, 5.1)
    notes[13] = "Left: one camera frame through the API. Right: what a rider or demonstrator and a researcher can do."

    # ---- 14 database
    s = S[13]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 5.0)
    rows = [["Column", "Type"], ["id", "INTEGER PK"], ["timestamp", "TEXT (UTC)"], ["source", "TEXT (LIVE CAMERA / SIMULATION)"], ["state, quality, reason", "TEXT"],
            ["ear, closure, blink_rate", "REAL"], ["closed_fraction, roll", "REAL"], ["alert", "INTEGER"]]
    ftable(s, 7.8, 1.9, [2.4, 2.6], rows, size=12, row_h=0.46)
    ftext(s, 7.8, 5.75, 5.0, 1.1, "One numeric events table plus settings.json; no video or face images are stored.", size=14, color=K.FNAVY, italic=True)
    notes[14] = "Only numbers are stored. Simulation rows are labelled so they can never be mistaken for rider data."

    # ---- 15 implementation
    s = S[14]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("20", "unit tests, all pass", K.GREEN), ("34", "API smoke checks, all pass", K.FBLUE), ("8", "web pages (React)", K.ORANGE), (f"{fr['p50_ms']:.0f} ms", "median camera frame via API", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.6, [
        ("HMAC-signed demo token (12 h), 401 on protected routes, validated thresholds (422), frames decoded in memory and never saved.", {"bullet": True, "space_after": 9}),
        ("One RiderSession per browser tab mirrors the desktop logic; a fresh SQLite connection per call keeps threads safe.", {"bullet": True, "space_after": 9}),
        ("Monitor screen: live landmarks, three outputs with reasons, calibration progress, alert sound, simulation scenarios for demonstrations without a camera.", {"bullet": True, "space_after": 9}),
        ("History with filters and CSV export, threshold editor, research screen, How-it-works page; light and dark themes; checked at 360, 768 and 1440 px.", {"bullet": True}),
    ], size=19)
    notes[15] = "Implementation highlights. The web layer adds no decision logic; it wraps the original modules."

    # ---- 16 results: eye-state classification on CEW
    s = S[15]
    set_title(s, "Results – Eye-state Classification on CEW")
    wipe(s)
    rows = [["Pipeline (threshold 0.21)", "Images assessed", "Accuracy", "Closed recall", "Open specificity", "AUC"],
            ["Shipped, quality gates on", f"{cew['assessed_good']:,} of {cew['images_total']:,} ({pc(a['coverage'])})", pc(gated["accuracy"]), pc(gated["recall"]), pc(gated["specificity"]), f"{a['gated_auc']:.3f}"],
            ["Landmarks only, no gates", f"{geo['face_found']:,} of {geo['images_total']:,}", pc(gdef["accuracy"]), pc(gdef["recall"]), pc(gdef["specificity"]), f"{a['geo_auc']:.3f}"]]
    ftable(s, 0.55, 1.85, [3.1, 3.0, 1.5, 1.6, 1.8, 1.2], rows, size=13, row_h=0.62, align=["l", "c", "c", "c", "c", "c"])
    cm = [["Landmarks only", "Predicted closed", "Predicted open"], ["Truly closed", str(gdef["tp"]), str(gdef["fn"])], ["Truly open", str(gdef["fp"]), str(gdef["tn"])]]
    ftable(s, 0.55, 4.0, [2.2, 1.9, 1.9], cm, size=13, row_h=0.55, align=["l", "c", "c"])
    for i, (v, l, c) in enumerate([(f"{cew['by_class']['closed']['ear_mean']:.3f}", "mean EAR, closed eyes", K.ORANGE), (f"{cew['by_class']['open']['ear_mean']:.3f}", "mean EAR, open eyes", K.GREEN)]):
        ftile(s, 6.9 + i * 3.0, 4.0, 2.85, 1.15, v, l, c, 26)
    ftext(s, 0.55, 5.8, 12.2, 1.1, [
        ("Gates pass only 44.1 % because 100-pixel crops are soft (sharpness gate). Closed images come from the Internet and open ones from LFW, so separation may be optimistic; still images only; coverage varies about +/- 2 points between runs.", {"bullet": True}),
    ], size=13)
    notes[16] = ("This is the measured result for one component: can the six-point EAR tell closed from open eyes. It is not an intoxication result. "
                 "The default threshold falsely calls about 14.5 % of open eyes closed on this data and misses under 1 % of closed ones.")

    # ---- 17 ROC and EAR distribution
    s = S[16]
    set_title(s, "ROC Curve & EAR Distribution (equivalent of a confusion matrix)")
    wipe(s)
    K.picture(s, F.fig("ear_histogram.png"), 0.55, 1.85, 6.1, 3.9)
    K.picture(s, F.fig("roc.png"), 6.75, 1.85, 6.1, 3.9)
    cap(s, 0.55, 5.8, 6.1, "EAR of the 1,068 images that passed the quality gates")
    cap(s, 6.75, 5.8, 6.1, f"ROC on the same images, AUC {a['gated_auc']:.3f}")
    ftext(s, 0.55, 6.3, 12.2, 0.6, "No sober-versus-intoxicated confusion matrix exists because no intoxication data was available.", size=14, color=K.FNAVY, bold=True, align="c")
    notes[17] = "Closed eyes cluster near 0.06 and open eyes near 0.29, with overlap; hence the high AUC and the threshold trade-off on the next slide."

    # ---- 18 threshold and parameter analysis
    s = S[17]
    set_title(s, "Threshold & Parameter Analysis")
    wipe(s)
    K.picture(s, F.fig("threshold_sweep.png"), 0.45, 1.8, 4.2, 2.8)
    K.picture(s, F.fig("closure_vs_warning.png"), 4.55, 1.8, 4.2, 2.8)
    K.picture(s, F.fig("parameter_tradeoff.png"), 8.65, 1.8, 4.2, 2.8)
    cap(s, 0.45, 4.6, 4.2, "EAR threshold sweep on CEW", 11)
    cap(s, 4.55, 4.6, 4.2, "Closure length vs warning (scripted)", 11)
    cap(s, 8.65, 4.6, 4.2, "Slow-blink alerts per hour (scripted)", 11)
    ftext(s, 0.55, 5.1, 12.2, 1.85, [
        (f"Best in-sample threshold is **0.17** ({pc(gbest['accuracy'])} accuracy) against **{pc(gdef['accuracy'])}** at the default 0.21. The default was not changed after this sweep; per-rider calibration sets a personal value.", {"bullet": True, "space_after": 9}),
        (f"On scripted streams closures up to {a['longest_closure_never_warns']} s never warn and closures of {a['min_closure_that_warns']} s or more always warn (30 of 30 runs); with 10 % slow blinks the default (1.5 s, 1.0 s) gives {default_row['alerts_per_hour_with_10pct_slow_blinks']:.1f} alerts per hour, while (1.0 s, 0.5 s) gives {fast_row['alerts_per_hour_with_10pct_slow_blinks']:.1f}.", {"bullet": True}),
    ], size=14)
    notes[18] = ("Left: threshold sweep on real CEW images. Middle and right: scripted streams, generated by seeded code, showing how the rule parameters trade detection against false alerts. "
                 "These show behaviour, not accuracy on people.")

    # ---- 19 robustness
    s = S[18]
    set_title(s, "Robustness: Quality Gating under Degraded Images")
    wipe(s)
    K.picture(s, F.fig("robustness.png"), 0.55, 1.8, 7.5, 4.9)
    sg, one, dark, tint, mask, noise = rob("sunglasses"), rob("one eye"), rob("x0.50"), rob("tinted"), rob("lower face"), rob("noise")
    ftext(s, 8.3, 1.9, 4.5, 4.9, [
        (f"Opaque sunglasses: **{pc(sg['good_rate'])}** assessed", {"bullet": True, "space_after": 6}),
        (f"One eye covered: **{pc(one['good_rate'])}** assessed", {"bullet": True, "space_after": 6}),
        (f"Brightness x0.5: **{pc(dark['good_rate'])}** assessed", {"bullet": True, "space_after": 6}),
        (f"Tinted lenses: {pc(tint['good_rate'])}; noise: {pc(noise['good_rate'])}; lower face masked: {pc(mask['good_rate'])}", {"bullet": True, "space_after": 6}),
        ("Gates **fail safe** (no verdict instead of a guess) but are **strict**: dim light already blocks assessment.", {"bullet": True, "space_after": 6}),
        (f"{F.ROB['n_images']} open-eye faces; low-resolution sources inflate the blur effect.", {"bullet": True}),
    ], size=15)
    notes[19] = "Objective 6 asked about low light and partial occlusion. The gates refuse to assess rather than guess, which is safe, but availability in dim light is poor."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{fr['p50_ms']:.1f} ms", f"camera frame via API, median (p95 {fr['p95_ms']:.1f})", K.FBLUE),
                                   (f"{F.CEW['latency_ms_vision_process_640x480']['p50']:.1f} ms", "landmarks + gates per frame, median", K.GREEN),
                                   (f"{rules['engine_cost']['microseconds_per_update']:.1f} us", "rule engine per observation", K.ORANGE),
                                   (f"{api['GET /api/health']['p50_ms']:.1f} ms", "health check, median", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.15, v, l, c, 24)
    rows = [["Test layer", "What it covers", "Result"],
            ["Unit tests", "EAR geometry, roll, rules, calibration, settings, SQLite, research pipeline", f"{tests['unittest']['passed']} of {tests['unittest']['tests']} passed"],
            ["API smoke test", "login, 401/409/422 guards, simulation, blank frame, history, CSV, thresholds, research", f"{tests['web_smoke']['passed']} of {tests['web_smoke']['checks']} checks passed"],
            ["Earlier desktop checks (23 Sep)", "Tk smoke test, camera diagnostic (81 real frames, no accuracy claim)", "passed"],
            ["User acceptance", "No user study was run", "not done"]]
    ftable(s, 0.55, 3.25, [3.0, 6.4, 2.8], rows, size=13, row_h=0.62)
    ftext(s, 0.55, 6.5, 12.2, 0.4, f"Measured on {F.HW['cpu']}, {F.HW['ram']} RAM, CPU only, 127.0.0.1.", size=12, color=K.FNAVY, italic=True)
    notes[20] = "All timings are on a laptop CPU. The API sustains about 30 frames per second; the browser sends about 9 per second."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different data, hardware, target)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "This project"],
            ["Target state", "Drowsiness, fatigue, distraction", "Possible impairment or drowsiness; not alcohol"],
            ["Eye feature", "EAR-4 (four upper-lid points) vs six-point EAR", "Conventional six-point EAR"],
            ["Decision", "Fuzzy belief rule-based fusion of EAR, head pose, voice", "Transparent duration rules: closure, closed fraction, tilt"],
            ["Data", "NTHU-DDD: 4 subjects, 65,834 frames; RF and XGBoost", "CEW: 2,423 still faces for the EAR component; no rider data"],
            ["Eye-state metric", "EAR-4 recall 0.894, F1 0.845 (Subject 2, RF); six-point EAR 0.856, 0.813", f"Six-point EAR at 0.21: recall {pc(gdef['recall'])}, specificity {pc(gdef['specificity'])}, F1 {gdef['f1']:.3f}"],
            ["Hardware / speed", "Jetson Orin Nano 8 GB, 25 FPS", f"Laptop CPU, about 30 frames/s via the API ({fr['p50_ms']:.0f} ms median)"]]
    ftable(s, 0.55, 1.85, [2.0, 5.2, 5.0], rows, size=12.5, row_h=0.62)
    ftext(s, 0.55, 6.3, 12.2, 0.6, "Not like-for-like: different datasets, subjects, hardware and target state. The numbers are context, not a ranking.", size=14, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that nothing here is directly comparable: different data, different hardware, and the base paper does not target alcohol either."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Live Monitor: Normal and Warning")
    wipe(s)
    K.picture(s, F.shot("06_monitor_camera_normal.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("04_monitor_simulation_warning.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    cap(s, 0.55, 5.7, 6.0, "Live camera path on a public-domain face: landmarks on the eyes, Normal, EAR 0.254, calibrated threshold 0.183")
    cap(s, 6.8, 5.7, 6.0, "Simulation 'Prolonged eye closure' raising Possible impairment or drowsiness, with the reason")
    notes[22] = "Left: real camera path (a public-domain portrait played as a fake webcam). Right: the scripted simulation scenario, labelled SIMULATION on screen."
    s = S[22]
    set_title(s, "Output – Calibration and Thresholds")
    wipe(s)
    K.picture(s, F.shot("05_monitor_calibrating.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("10_thresholds.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    cap(s, 0.55, 5.7, 6.0, "Calibration in progress (scrolled capture): EAR chart, 19 usable frames, default threshold 0.210")
    cap(s, 6.8, 5.7, 6.0, "Six validated thresholds with their bounds; saving stops monitoring")
    notes[23] = "Calibration sets a personal EAR threshold; the thresholds screen exposes the six rule values with validation."
    s = S[23]
    set_title(s, "Output – History and Research (synthetic data)")
    wipe(s)
    K.picture(s, F.shot("08_history.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("12_research_report.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    cap(s, 0.55, 5.7, 6.0, "Event history: 34 events at the time of capture (31 simulation, 3 live camera), EAR chart, CSV export")
    cap(s, 6.8, 5.7, 6.0, "SYNTHETIC data only: the 100 % figures prove the pipeline runs, not performance")
    ftext(s, 0.55, 6.35, 12.2, 0.5, "The research report on the right was built from a generated demonstration file and says nothing about real riders.", size=13, color=K.FNAVY, bold=True, align="c")
    notes[24] = "History is numeric only. The research screen carries a Synthetic data banner; do not read its 100 percent as a result."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 3.9, [
        ("**Delivered:** a complete, explainable path from webcam frame to a logged warning with a reason: landmarks, seven quality gates, calibration, a rule engine, a web and desktop interface, 20 unit tests and 34 API checks.", {"bullet": True, "space_after": 9}),
        (f"**Measured:** six-point EAR separates closed from open eyes on 2,423 public faces (AUC {a['geo_auc']:.3f}); the default 0.21 is conservative ({pc(a['default_false_closed_rate'])} of open eyes called closed); a closure must last about {a['min_closure_that_warns']} s to warn; {fr['p50_ms']:.0f} ms per frame.", {"bullet": True, "space_after": 9}),
        ("**Finding:** the quality gates fail safe but are strict; dim light, blur at low resolution, sunglasses or a covered eye give Unable to assess instead of a guess.", {"bullet": True, "space_after": 9}),
        ("**Not achieved:** alcohol detection, a deep-learning classifier (no intoxication data), gaze, and the prevention action. Not a medical device or a road-safety system.", {"bullet": True}),
    ], size=17)
    for i, (v, l, c) in enumerate([("2", "objectives achieved (1 and 4, the latter as rules)", K.GREEN), ("2", "partly achieved (2 and 6)", K.ORANGE), ("2", "not achieved (3 and 5)", K.PURPLE)]):
        ftile(s, 0.55 + i * 4.1, 5.75, 3.9, 1.1, v, l, c, 28)
    notes[25] = "Two of six objectives met, two partly, two not. We would rather state the gap than present a number we cannot support."

    # ---- 26 limitations & future scope
    s = S[25]
    set_title(s, "Limitations, What Was Not Measured & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.85, 6.0, 5.0, fill=K.FZEBRA)
    fbox(s, 6.8, 1.85, 6.0, 5.0, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.45, "Limitations and not measured", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("No alcohol detection; cues also arise from tiredness", {"bullet": True, "space_after": 9}),
        ("No intoxication data, no trained classifier, no gaze", {"bullet": True, "space_after": 9}),
        ("No prevention action (no ignition or contact link)", {"bullet": True, "space_after": 9}),
        ("No rider, helmet, visor, glasses, vibration or demographic study", {"bullet": True, "space_after": 9}),
        ("Default threshold conservative; gates strict in dim light", {"bullet": True, "space_after": 9}),
        ("Scripted streams show behaviour, not accuracy", {"bullet": True}),
    ], size=17)
    ftext(s, 7.0, 1.9, 5.6, 0.45, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Collect a consented, labelled sober vs intoxicated video dataset", {"bullet": True, "space_after": 9}),
        ("Gaze and saccade features; a temporal model checked on held-out subjects", {"bullet": True, "space_after": 9}),
        ("Tune thresholds on held-out riders; helmet-robust sensing", {"bullet": True, "space_after": 9}),
        ("Vehicle integration for a prevention action, with safety review", {"bullet": True, "space_after": 9}),
        ("Embedded deployment as in the base paper; formal user study", {"bullet": True}),
    ], size=17)
    notes[26] = "The limitations are the honest list; the future scope is what would turn this prototype into an intoxication-specific system."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    refs = [1, 2, 3, 4, 5, 6, 8, 9, 11, 12, 13, 14, 15, 16, 17, 24]
    clean = lambda n: REFS[n - 1].split(", doi")[0]
    ftext(s, 0.55, 1.85, 6.0, 5.0, [(f"[{n}] {clean(n)}", {"space_after": 4}) for n in refs[:8]], size=9)
    ftext(s, 6.8, 1.85, 6.0, 5.0, [(f"[{n}] {clean(n)}", {"space_after": 4}) for n in refs[8:]], size=9)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.scrub_theme(out_path, "SmartRider")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
