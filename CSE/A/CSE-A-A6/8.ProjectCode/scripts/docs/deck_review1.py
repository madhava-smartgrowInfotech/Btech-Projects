"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .ch1 import OBJECTIVES, OBJECTIVE_STATUS
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout/master footer text carries the template's deck title
    holders = list(prs.slide_layouts) + [prs.slide_master]
    for lay in holders:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.TITLE}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "driver monitoring, eye aspect ratio, MediaPipe, intoxication, drowsiness"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    geo = F.GEO

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, size=34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, f"Prototype: {F.PRODUCT}. Explainable eye and head cues from a webcam; a camera cannot confirm alcohol use")
        elif nm == "review label":
            K.replace_text(sh, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")
    notes[1] = ("Introduce the team and the project title. One honest sentence up front: the batch title says intoxication, but a camera cannot confirm alcohol; "
                "what we built is an explainable warning for possible drowsiness or impairment from eye and head cues.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: warn early, say so when unsure")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Impaired and drowsy riding is a major road-safety risk, and roadside breath checks are occasional. Existing driver monitoring mostly targets drowsiness, and intoxication research needs special data or sensors.", K.ORANGE, "!"),
        ("Our approach", "A webcam is read with MediaPipe face landmarks. Eye closure, closed-eye fraction and head tilt are judged by transparent rules, with a per-rider calibration and quality gates that answer 'Unable to assess' instead of guessing.", K.BLUE, "i"),
        ("What it is not", "A camera cannot confirm alcohol. The output is 'Possible impairment or drowsiness'. No intoxication dataset was available, so no learned alcohol classifier was trained, and no ignition or contact action exists yet.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=20)
    notes[3] = ("Three parts: problem, approach, limits. Say the limits aloud: the abstract in the batch file speaks of intoxication and a Safe or Intoxicated label; "
                "the delivered prototype is rule-based and stationary, and we state that plainly.")

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Eye aspect ratio from MediaPipe face landmarks (six-point EAR and a four-point EAR-4), head pose, Whisper-tiny voice commands, fused by a fuzzy belief rule-based engine; personal EAR threshold from more than 300 frames.", K.BLUE, "M"),
        ("Data and reported results", "NTHU-DDD, four subjects, 65,834 frames (29,804 open, 36,030 closed), 80/20 split. EAR-4 recall 0.894 against 0.856 for six-point EAR (Subject 2, Random Forest). 25 FPS on a Jetson Orin Nano.", K.GREEN, "R"),
        ("What we take, what differs", "We take the EAR idea and personal calibration. The base paper targets drowsiness and distraction, not alcohol, and its split is not subject-independent. We use six-point EAR, duration rules and no voice.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=17)
    notes[4] = ("Read the citation. The base paper is a drowsiness and distraction monitor, not an alcohol detector. Numbers are the paper's own, from its text and abstract; "
                "they were measured on different data and hardware than ours, so they are context, not a target.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result"],
            ["Feng and Li 2026 [1] (base)", "EAR-4 + head pose + voice, fuzzy belief rules", "NTHU-DDD, 4 subjects, 65,834 frames", "EAR-4 recall 0.894; 25 FPS (Jetson)"],
            ["Baroutian et al. 2025 [2]", "Landmark graph network fused with 3D ResNet", "202 videos (101 sober, 101 intoxicated), 3,542 clips", "Accuracy 95.82 %, recall 0.97"],
            ["Deuber et al. 2026 [3]", "Smartwatch motion + heart-rate variability, 1D CNN", "Test-track study, 54 participants", "AUROC 0.88 (any alcohol)"],
            ["Acharya 2024 [5]", "Thermal images, NasNetMobile / MobileNet", "50 volunteers, four drinking levels", "85.10 % (four levels); 74.07 % sober vs drunk"],
            ["Chai et al. 2025 [8]", "Thermal facial patterns, CNN features + MLP", "Tufts Face Database + own thermal images", "MLP about 90 % (abstract only)"],
            ["Soukupova and Cech 2016 [9]", "Eye aspect ratio from six landmarks for blinks", "Three standard blink datasets", "Comparable with the state of the art"],
            ["Song et al. 2014 [12] (CEW)", "Eye-closeness detection from still images", "2,423 faces: 1,192 closed, 1,231 open", "Dataset we use to test our EAR"]]
    K.table(s, 0.6, 1.4, [2.7, 3.5, 3.1, 2.83], rows, font_size=12, row_h=0.66, header_size=13)
    notes[5] = ("Pattern: video and thermal studies reach 85 to 96 percent but need special data or cameras; the wearable needs a device; EAR is the standard cheap eye measure. "
                "We read the PDFs of [1] to [3]; [5] and [8] are abstract only. The ICOEI 2026 taxi paper in the batch list could not be located, so it is not quoted.")

    # ---- 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence")
    wipe(s)
    K.bar_chart(s, 0.5, 1.4, 6.9, 4.7, ["Baroutian\n2025\naccuracy", "Chai\n2025\nMLP acc.", "Acharya\n2024\n4 levels", "Deuber\n2026\nAUROC", "Feng, Li\n2026\nrecall"],
                [("Headline result", [0.9582, 0.90, 0.851, 0.88, 0.894])], colors=[K.BLUE], number_format="0.0%", font_size=11, max_value=1.1)
    K.text(s, 0.6, 6.15, 6.8, 0.6, "Headline figures from the papers: different metrics, data and modalities, so not comparable with each other.", size=12, color=K.TXT, italic=True)
    K.picture(s, F.fig("ear_histogram.png"), 7.7, 1.45, 5.05, 3.0, border=True)
    g = geo["default_threshold_0.21"]
    K.text(s, 7.7, 4.55, 5.05, 2.2,
           [(f"**Our early check:** on {geo['face_found']:,} public CEW faces the six-point EAR separates closed from open eyes with AUC {geo['auc_ear_closed_vs_open']:.3f}.", {"space_after": 6}),
            (f"At the default threshold 0.21: accuracy {g['accuracy'] * 100:.1f} %, closed-eye recall {g['recall'] * 100:.1f} %, open-eye specificity {g['specificity'] * 100:.1f} %. Still images only; no alcohol data.", {})],
           size=13, color=K.INK)
    notes[6] = ("Left: what the literature reports. Right: our own first measurement, only to show the EAR signal is informative. Closed and open faces come from different sources, "
                "so the separation may be optimistic; it says nothing about alcohol.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Driver drinks", K.NAVY), ("Occasional roadside breath test", K.NAVY), ("Officer present, result", K.NAVY), ("Penalty after the fact", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Checks and sensors", "As stated in the project proposal: breath testing is occasional and cannot follow a rider through a journey, and hardware alcohol sensors add cost and can be bypassed by asking another person to blow.", K.BLUE, "1", body_size=20)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Research systems", "Driver-state monitors mostly target drowsiness [1]. Intoxication detection from video, thermal images or a smartwatch exists as research [2], [3], [5], [8], usually needing special data, a thermal camera or a worn device.", K.ORANGE, "2", body_size=20)
    notes[7] = "Describe today's path: a check that happens occasionally, a penalty afterwards. Research prototypes exist but are not everyday tools."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Intermittent checks", "A breath test samples one moment; it cannot watch the whole journey or act before a rider sets off.", K.ORANGE, "1"),
        ("Drowsiness-focused monitors", "The base paper [1] and most driver monitors detect tiredness and distraction, not alcohol impairment.", K.BLUE, "2"),
        ("Special data or hardware", "Thermal cameras [5], [8], a worn smartwatch [3] or GPU video models trained on scraped videos [2].", K.GREEN, "3"),
        ("Hard to reproduce", "Many source videos of the DIF intoxication dataset are no longer reachable, and numbers depend on the protocol [2].", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=19)
    notes[8] = "Four limitations. Each informs a design choice later: a continuous camera view, explainable rules, only a webcam, and a stated evaluation protocol."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Batch survey gap", "An intoxication-specific vision model coupled to a prevention action rather than a warning alone."),
            ("Data", "No public, controlled dataset of sober and intoxicated riders was available to us, so a learned classifier could not be trained honestly."),
            ("Transparency", "Our design target: every verdict states its reason, and poor light, blur or covered eyes give Unable to assess instead of a guess."),
            ("Degraded conditions", "The proposal asks for low light and partial occlusion to be evaluated; we measure them on test images."),
            ("Our scope", "We close a smaller gap: an explainable, quality-gated, calibrated eye and head monitor, measured honestly.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=17, color=K.INK, anchor="m")
    notes[9] = ("The first row is the gap stated in the batch literature survey. The other rows are what we found while trying to close it, and the last row is the part we actually address.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.55, fill=K.CARD)
    K.text(s, 0.95, 1.55, 11.45, 2.35, "Given a live webcam view of a rider's face, measure **eye closure** and **head tilt**, judge them with **explainable rules**, and warn of **possible impairment or drowsiness** when the signs persist, or say **Unable to assess** when the view is unreliable. A camera cannot confirm alcohol use.", size=22, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Webcam frames, 640 x 480, kept in memory only"), ("Output", "Normal, Possible impairment or drowsiness, or Unable to assess, with the reason"), ("Record", "Numeric event history, CSV export, adjustable thresholds")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.3, 3.85, 2.2)
        K.text(s, x + 0.25, 4.45, 3.4, 0.4, t, size=20, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.0, 3.4, 1.4, b_, size=18, color=K.TXT)
    notes[10] = "Read the statement slowly. Note the three outputs: the third one, Unable to assess, is deliberate. The system never guesses from a poor image."

    # ---- 11 objectives (with honest status)
    s = S[10]
    set_title(s, "Objectives and their status")
    wipe(s)
    short = ["Real-time webcam face and landmark detection", "Impairment cues: eye closure, blink rate, gaze, head movement",
             "Deep-learning sober vs intoxicated classifier", "Continuous inference with an alert above a threshold",
             "Prevention action: ignition warning or contact notification", "Evaluate accuracy, false alarms, low light, occlusion"]
    detail = ["Done: MediaPipe, about 31 ms per frame", "Gaze not measured; the rest are", "No data, no model; rules instead",
              "Done as rules, three outputs", "Not done: warning and bell only", "Measured on public data, no rider study"]
    rows = [["#", "Objective (from the project proposal)", "Status", "What exists"]]
    for i in range(6):
        rows.append([str(i + 1), short[i], OBJECTIVE_STATUS[i][0], detail[i]])
    K.table(s, 0.6, 1.4, [0.5, 5.2, 1.9, 4.53], rows, font_size=14, row_h=0.74, header_size=14)
    notes[11] = ("Six objectives from the project proposal, each with its honest status. Achieved: 1 and 4 (as rules). Partly: 2 and 6. Not achieved: 3 (no intoxication data, no deep-learning model) and 5 (no prevention action). "
                 "The full wording of the objectives is in the report, Chapter 1.")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Capture and landmarks", "Webcam frame, MediaPipe Face Landmarker, six points per eye, eye aspect ratio and head roll.", K.BLUE, "1"),
        ("Quality gates", "Seven checks (light, blur, motion, faces, eye size, eye contrast, geometry). A failed check means Unable to assess.", K.GREEN, "2"),
        ("Rules and calibration", "Closure of 1.5 s, closed-eye fraction, or tilt with low EAR, persisting 1 s. Threshold calibrated per rider.", K.ORANGE, "3"),
        ("Warn and log", "On-screen warning, optional bell, numeric event log, CSV export, adjustable validated thresholds.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=17)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "W", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Web application and future work", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "FastAPI service and React interface over the same Python modules, demo login. Future: a classifier trained on a consented sober and intoxicated dataset, and a real prevention action.", size=16, color=K.TXT)
    notes[12] = ("Four stages and the application around them. The classifier and the prevention action are future work: they need a consented labelled dataset and a vehicle interface that we did not have.")

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.5, 1.3, 7.3, 5.4, border=True)
    for i, (t, b_) in enumerate([("Client tier", "React app in the browser (webcam via getUserMedia) and the original Tkinter desktop app."),
                                 ("Application tier", "FastAPI service with a demo login and one monitoring session per browser tab."),
                                 ("Shared Python modules", "MediaPipe vision with quality gates, the rule engine, storage and an offline research tool."),
                                 ("Data tier", "Landmark model file, SQLite event log, JSON settings. Video frames stay in memory.")]):
        y = 1.4 + i * 1.32
        K.box(s, 8.1, y, 4.63, 1.2)
        K.text(s, 8.3, y + 0.1, 4.25, 0.35, t, size=15, color=K.INK, bold=True)
        K.text(s, 8.3, y + 0.48, 4.25, 0.7, b_, size=12.5, color=K.TXT)
    notes[13] = "Follow the diagram: front ends on the left, the API, the shared vision and rule modules, and the local data stores. Frames stay in memory and are never saved."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working monitor", "A real-time webcam monitor with three outputs and a stated reason for each.", K.BLUE, "1"),
            ("Safe failure", "No verdict when lighting, blur or occlusion make the eyes unreadable.", K.GREEN, "2"),
            ("Measured behaviour", "Eye-state accuracy on public data, degraded-image tests, rule timing and latency.", K.ORANGE, "3"),
            ("Usable application", "Web interface with login, history, thresholds and CSV export.", K.PURPLE, "4"),
            ("Honest scope", "Not an alcohol test; no trained classifier or prevention action in this version.", K.BLUE, "5"),
            ("Reproducible", "Seeded measurement scripts and automated tests for the core and the API.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=18)
    notes[14] = "These are expected outcomes, not results. We report whatever the measurements show, including weaknesses, and we state what is out of scope."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    refs = [(k, REFS[k - 1]) for k in (1, 2, 3, 4, 5, 8, 9, 11, 12, 17)]
    K.text(s, 0.6, 1.35, 12.13, 5.5, [(f"[{n_}] {t}", {"space_after": 7}) for n_, t in refs], size=14, color=K.INK)
    notes[15] = "Selected references; the full list is in the project report."

    # ---- 16 thank you
    s = S[15]
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")
    notes[16] = ("Likely questions: why rules and not deep learning (no intoxication data, explainability); can it detect alcohol (no, a camera cannot confirm it); "
                 "what about the ignition lock (future work).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)
    K.scrub_theme(out_path, "SmartRider")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
