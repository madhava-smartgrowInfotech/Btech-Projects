"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress. Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx.

Every number comes from docs/REPORT_FACTS.md via facts.py (experiments/**/metrics.json) or from chapter 8's analysis().
    python -m scripts.docs.deck_review2
"""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import ch8
from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"
KEEP = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY)
FRAME_KEY = "POST frame (real face JPEG 640x480, monitoring on)"


# ------------------------------------------------------------------ local helpers (no dependency on other decks)
def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "eye aspect ratio, driver monitoring, MediaPipe, explainable rules, FastAPI, React"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def geometry_sweep_chart():
    """Threshold sweep for the geometry-only CEW experiment (the numbers quoted on the results slide)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    sweep = F.GEO["threshold_sweep"]
    x = [c["threshold"] for c in sweep]
    fig, ax = plt.subplots(figsize=(6.4, 3.9), dpi=150)
    ax.plot(x, [c["accuracy"] for c in sweep], label="Accuracy", color="#264653")
    ax.plot(x, [c["recall"] for c in sweep], label="Closed-eye recall", color="#e76f51")
    ax.plot(x, [c["specificity"] for c in sweep], label="Open-eye specificity", color="#2a9d8f")
    ax.axvline(0.21, color="k", ls="--", lw=1.0)
    ax.text(0.2115, 0.07, "default 0.21", fontsize=8)
    ax.set(xlabel="EAR threshold", ylabel="Rate", ylim=(0, 1.02),
           title=f"Threshold sweep on CEW ({F.GEO['face_found']:,} faces, landmarks only)")
    ax.legend(frameon=False, fontsize=8, loc="lower left")
    fig.tight_layout()
    path = os.path.join(F.FIGS, "threshold_sweep_geometry.png")
    fig.savefig(path)
    plt.close(fig)
    return path


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}

    A = ch8.analysis()
    cew, geo, gate = F.CEW, F.GEO, A["gated"]
    gd, gb = A["geo_default"], A["geo_best"]
    fr = F.API_LAT[FRAME_KEY]
    vl = cew["latency_ms_vision_process_640x480"]
    ec = F.RULES["engine_cost"]
    ut, sm = F.TESTS["unittest"], F.TESTS["web_smoke"]
    pct = F.pct

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, f"{F.PRODUCT}: an explainable eye-closure and head-tilt warning with a FastAPI service and a React interface")
        elif nm == "review label":
            K.replace_text(sh, "REVIEW 2  ·  DESIGN, METHODOLOGY AND PROGRESS")
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")
    notes[1] = ("Introduce the team. This review shows how the system is designed and built, what has been implemented, the first measured results and the open items. "
                "Say at the start: a camera cannot confirm alcohol, so the product is a drowsiness / impairment warning and we state exactly what it does and does not do.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap and architecture"), ("02", "Methodology and dataset"), ("03", "Preprocessing and model"), ("04", "Training and decision logic"),
               ("05", "Data flow, database and UML"), ("06", "Technology stack"), ("07", "Progress and preliminary results"),
               ("08", "Output screens and testing"), ("09", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Drink-driving and fatigued riding are hard to police continuously: breath tests are occasional and sensors can be bypassed. The batch asked for a contactless, camera-based check.", K.ORANGE, "!"),
        ("What we built", "Webcam frames, MediaPipe eye landmarks, quality gates and explainable rules give Normal, Possible impairment or drowsiness, or Unable to assess, with the reason.", K.BLUE, "i"),
        ("What it is not", "A camera cannot confirm alcohol. There is no deep-learning classifier, no intoxication dataset and no ignition or contact alert in this version.", K.PURPLE, "×"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.45, W, 3.35, t, b_, col, gl, body_size=17)
    K.box(s, 0.6, 5.1, 12.13, 1.65)
    K.text(s, 0.85, 5.2, 11.6, 1.45, [
        ("**Base paper link.** Feng and Li, IEEE Access 2026: eye aspect ratio plus a personalised threshold for driver state. We reuse that idea with the "
         "conventional six-point EAR and transparent duration rules; we do not reproduce EAR-4, the voice interface or the fuzzy belief engine.", {})],
           size=17, color=K.INK, anchor="m")
    notes[3] = ("One-minute recap of Review 1. The batch title says alcohol; the abstract says visual signs of intoxication. What we can honestly deliver and measure is an eye and head "
                "based warning for possible drowsiness or impairment. The third card says plainly what is missing.")

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.6, 5.5)
    K.text(s, 8.5, 1.45, 4.25, 5.4, [
        ("**Four tiers**", {"space_after": 8}),
        ("Presentation: React app (port 5306) and the original Tkinter desktop app.", {"bullet": True}),
        ("Application: FastAPI service (port 8306), signed demo login, input validation.", {"bullet": True}),
        ("Core: `core.py` rules, `vision.py` landmarks, `storage.py`, `research.py`, shared by both front ends.".replace("`", ""), {"bullet": True}),
        ("Data: SQLite event table, settings file, landmark model file.", {"bullet": True}),
        ("Camera frames stay in RAM and are never saved.", {"bullet": True, "italic": True}),
    ], size=16, color=K.INK)
    notes[4] = ("Four tiers. The web service wraps the original desktop modules rather than rewriting them, so both front ends give the same verdict. "
                "The browser sends a 640 x 480 JPEG frame; the server runs MediaPipe and the rule engine and returns the state and an annotated preview.")

    # ---- 5 methodology
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.45, 12.13, 1.4, [("Camera frame 640 x 480", K.NAVY), ("Quality gates", K.PURPLE), ("Eye landmarks, EAR, roll", K.BLUE), ("Median + calibration", K.GREEN),
                                      ("Rule engine", K.ORANGE), ("Verdict + reason", K.NAVY), ("SQLite log, CSV", K.NAVY)], size=13, gap=0.22)
    for i, (t, b_) in enumerate([("Explainable by design", "Every warning states its reason: closure time, closed fraction or head tilt with low EAR. No hidden score."),
                                 ("Safe failure", "A failed quality gate always wins: the answer is Unable to assess and the evidence is cleared, never a guess."),
                                 ("One code path", "The desktop app and the web app use the same measurement, rule and storage modules.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.2, 3.85, 2.55)
        K.text(s, x + 0.25, 3.35, 3.4, 0.5, t, size=19, color=K.INK, bold=True)
        K.text(s, x + 0.25, 3.95, 3.4, 1.7, b_, size=17, color=K.TXT)
    K.text(s, 0.6, 6.0, 12.13, 0.8, "The same pipeline serves browser camera frames and four scripted simulation scenarios: Normal, Prolonged eye closure, Missing face and Poor lighting.",
           size=15, color=K.INK, italic=True, anchor="m")
    notes[5] = "Six stages from frame to log. The method is deliberately transparent, because there is no labelled intoxication data to train a classifier on."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Dataset and evidence used")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{cew['images_total']:,}", "public CEW faces: 1,192 closed, 1,231 open", K.BLUE), ("100 x 100", "warped crops, used for the EAR check only", K.GREEN),
                                   ("0", "intoxication recordings available", K.ORANGE), ("6 x 2 x 30 s", "synthetic CSV for the offline demo", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    K.picture(s, F.fig("landmark_examples.png"), 0.6, 2.95, 6.3, 3.9)
    K.text(s, 7.2, 3.0, 5.55, 3.85, [
        ("**Closed Eyes in the Wild** (Song et al., 2014): Internet closed-eye faces and LFW open-eye faces; not stored in the repository.", {"bullet": True, "space_after": 7}),
        ("**No intoxication data:** none was available, so no classifier was trained.", {"bullet": True, "space_after": 7}),
        ("**Synthetic CSV** for the offline pipeline is flagged Synthetic data on screen; its metrics are not performance.", {"bullet": True, "space_after": 7}),
        ("Live screenshots use one public-domain portrait as a fake webcam.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("CEW has 2,423 faces and is used only to test the eye open / closed measurement. It says nothing about alcohol. "
                "The synthetic CSV only proves the offline pipeline runs; the screen carries a Synthetic data banner.")

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Preprocessing: quality gates and measurements")
    wipe(s)
    K.text(s, 0.6, 1.45, 4.6, 5.3, [
        ("**1  Frame.** 640 x 480, grayscale copy.", {"space_after": 9}),
        ("**2  Gates.** Any failed gate gives Unable to assess.", {"space_after": 9}),
        ("**3  Landmarks.** Six points per eye and both outer eye corners.", {"space_after": 9}),
        ("**4  Measures.** EAR (mean of both eyes), head roll, 3-sample median.", {"space_after": 9}),
        ("Gates are strict on purpose: halving the brightness already blocks assessment.", {"italic": True, "size": 14, "color": K.TXT}),
    ], size=16, color=K.INK)
    K.box(s, 0.6, 4.55, 4.6, 2.2)
    K.text(s, 0.8, 4.65, 4.2, 2.05, [
        ("**Measured on 272 open-eye faces**", {"space_after": 5}),
        ("Sunglasses over both eyes: 270 of 272 answered Unable to assess.", {"bullet": True, "space_after": 3}),
        ("One eye covered: 272 of 272 Unable to assess.", {"bullet": True, "space_after": 3}),
        ("Brightness x0.5: only 1 of 272 still assessed.", {"bullet": True}),
    ], size=14, color=K.INK)
    rows = [["Quality gate", "Fails when"], ["Lighting", "mean gray < 40 or > 225"], ["Sharpness", "Laplacian variance < 25"], ["Motion", "frame difference > 35"],
            ["Face count", "no face or more than one"], ["Eye size / position", "eye width < 14 px or outside image"], ["Eye patch", "mean < 35 or > 230, or sd < 8"],
            ["Eye geometry", "left-right EAR difference > 0.12"], ["Timing", "gap > 0.5 s or time not increasing"]]
    K.table(s, 5.5, 1.45, [2.5, 4.73], rows, font_size=14, row_h=0.58, first_bold=True)
    notes[7] = ("Seven pixel and geometry gates plus a timing check run before any verdict. They are heuristics, not proof of occlusion. "
                "Our degraded-image test shows they fail safe but are strict.")

    # ---- 8 model
    s = S[7]
    set_title(s, "Model: pre-trained landmarks and a rule engine")
    wipe(s)
    rows = [["Component", "What it is", "Role"],
            ["Face Landmarker", "Google MediaPipe, 478 landmarks, pre-trained", "eye and head geometry"],
            ["EAR", "6 points per eye, mean of both eyes", "eye openness"],
            ["Head roll", "atan2 of outer eye corners, image plane", "tilt"],
            ["Median filter", "3 samples", "suppress outliers"],
            ["Rule engine", "state machine, 3 outputs", "live decision"],
            ["Offline classifier", "NumPy nearest centroid on 5 s windows", "research only: Model not trained"]]
    K.table(s, 0.6, 1.5, [2.3, 4.2, 2.6], rows, font_size=13, row_h=0.72, first_bold=True)
    K.tile(s, 9.9, 1.5, 2.83, 1.45, "478", "face landmarks, not trained by us", K.BLUE, vsize=28, lsize=12)
    K.tile(s, 9.9, 3.15, 2.83, 1.45, f"{ec['microseconds_per_update']:.0f} µs", "per rule-engine update", K.GREEN, vsize=26, lsize=12)
    K.tile(s, 9.9, 4.8, 2.83, 1.45, f"{vl['p50']:.0f} ms", "median landmarking, CPU only", K.ORANGE, vsize=26, lsize=12)
    notes[8] = ("There is no neural network trained by the team. MediaPipe supplies the landmarks; the decision is a small rule engine. "
                "The nearest-centroid classifier only runs offline on a consented labelled CSV, and none exists yet, so the screen says Model not trained.")

    # ---- 9 training configuration
    s = S[8]
    set_title(s, "Training configuration: calibration and defaults")
    wipe(s)
    rows = [["Setting", "Value"],
            ["Calibration", "8-15 s eyes open; at least 40 accepted frames and 70 % accepted; EAR sd / median <= 0.15"],
            ["Personal threshold", "clip(0.72 x median EAR, 0.12, 0.32); default 0.21, bounds 0.08-0.40"],
            ["Closure rule", "continuous closed eyes >= 1.5 s (bounds 0.5-5 s)"],
            ["Closed-fraction rule", ">= 0.55 of the last 10 s, with at least 8 s of history"],
            ["Tilt rule", "head roll >= 25 degrees together with a low EAR"],
            ["Persistence / cooldown", "1.0 s more evidence and 2 s reliable history; bell cooldown 15 s"],
            ["Tuning", "engineering defaults, not tuned on data; the CEW sweep later points to 0.17-0.18"]]
    K.table(s, 0.6, 1.5, [3.2, 8.93], rows, font_size=15, row_h=0.65, first_bold=True)
    notes[9] = ("There is no training loop. What is configured is the calibration and the rule constants. They are adjustable within validated bounds. "
                "After the CEW sweep we know 0.21 is conservative; we kept it as the documented default and rely on calibration per rider.")

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.6, 1.3, 5.2, 5.55)
    K.text(s, 6.2, 1.45, 6.55, 5.4, [
        ("**Quality first.** A failed gate or timing gap gives Unable to assess and clears the evidence; the alert cooldown is kept.", {"bullet": True, "space_after": 9}),
        ("**Three rules, every frame.** Closure >= 1.5 s; closed fraction >= 0.55 over 10 s; roll >= 25 degrees with low EAR.", {"bullet": True, "space_after": 9}),
        ("**Persistence.** The evidence must last 1 s more, with 2 s of reliable history, before the warning shows.", {"bullet": True, "space_after": 9}),
        ("**Blinks.** A closed interval of 0.08-0.8 s counts as a blink; blink rate is shown, never used alone.", {"bullet": True, "space_after": 9}),
        ("**Output.** Normal, Possible impairment or drowsiness, or Unable to assess, with the reason; optional bell every 15 s at most.", {"bullet": True}),
    ], size=17, color=K.INK)
    notes[10] = "The flowchart is drawn from RiskEngine.update. All three rules are evaluated every frame and their reasons are joined in the warning text."

    # ---- 11 data flow and modules
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 7.4, 5.55)
    rows = [["Module", "Role"], ["core.py", "rules, calibration, scenarios"], ["vision.py", "landmarks and quality gates"], ["storage.py", "SQLite, CSV, settings"],
            ["research.py", "offline evaluation"], ["backend/main.py", "22 API routes, login"], ["backend/session.py", "one session per tab"], ["frontend/src", "8 React pages"]]
    K.table(s, 8.2, 1.5, [2.1, 2.43], rows, font_size=12, row_h=0.55, first_bold=True)
    notes[11] = "Level-1 data flow: seven processes and three stores. Levels 0 and 2 are in the report."

    # ---- 12 database and UML
    s = S[11]
    set_title(s, "Database design and UML")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.6, 1.35, 6.0, 5.1)
    K.picture(s, F.diag("06_uml_class_diagram.png"), 6.8, 1.35, 5.9, 5.1)
    K.text(s, 0.6, 6.55, 12.13, 0.4, "One persisted table (events, 12 numeric columns), a settings file and the research CSV; sessions live in memory and no images are stored.",
           size=12, color=K.TXT, align="c", italic=True)
    notes[12] = "Left: what is persisted and what is not. Right: the main classes. Settings, Observation, Decision, RiskEngine and Calibration are pure logic; Vision, Store and RiderSession are the parts that touch the outside world."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    rows = [["Layer", "Technology"],
            ["Vision", "MediaPipe 1.0.1 (Face Landmarker, CPU), OpenCV contrib 5.0.0.93, NumPy 2.5.3"],
            ["Charts and offline tools", "Matplotlib 3.11.2, NumPy nearest centroid"],
            ["API", "FastAPI 0.162, Uvicorn 0.54, python-multipart; HMAC-signed demo token"],
            ["Desktop", "Tkinter (original prototype)"], ["Database", "SQLite and a JSON settings file"],
            ["Front end", "React 18, Vite 8, TypeScript, Tailwind 4, React Router 7, TanStack Query, Recharts"],
            ["Testing", "unittest, FastAPI TestClient, Playwright for screenshots"],
            ["Run", "Python 3.14.7 venv, run.bat (Windows), ports 8306 and 5306"]]
    K.table(s, 0.6, 1.45, [3.0, 9.13], rows, font_size=15, row_h=0.6, first_bold=True)
    notes[13] = "Everything is open source and runs on one laptop; run.bat rebuilds the environment and downloads the 3.7 MB landmark model if needed."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Eye and head measurement, 7 quality gates", "Done", "vision.py; measured on 2,423 CEW faces"],
            ["Rule engine, calibration, thresholds", "Done", f"core.py; {ut['tests']} unit tests"],
            ["FastAPI service and demo login", "Done", f"22 routes; {sm['checks']} smoke checks"],
            ["React interface, history, CSV export", "Done", "8 pages; 15 desktop and 4 phone screenshots"],
            ["Offline research pipeline", "Built, not trained", "synthetic demo only, flagged"],
            ["Deep-learning sober / intoxicated model", "Not done", "no labelled intoxication data"],
            ["Prevention action (ignition or contact alert)", "Not done", "on-screen warning and bell only"],
            ["Gaze deviation", "Not done", "not measured"]]
    K.table(s, 0.6, 1.45, [4.9, 2.1, 5.13], rows, font_size=14, row_h=0.58)
    notes[14] = "State of the code base. The three Not done rows are objectives from the batch file that this version does not meet; the report says so in the objectives table."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: eye state on CEW")
    wipe(s)
    for i, (v, l, c) in enumerate([(pct(gd["accuracy"]), "accuracy at threshold 0.21", K.BLUE), (pct(gd["recall"]), "closed-eye recall", K.GREEN),
                                   (pct(gd["specificity"]), "open-eye specificity", K.ORANGE), (f"{A['geo_auc']:.3f}", "ROC AUC of EAR", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.4, 2.9, 1.15, v, l, c, vsize=28, lsize=12)
    K.picture(s, geometry_sweep_chart(), 0.6, 2.75, 6.6, 4.1)
    K.tile(s, 7.5, 2.75, 2.55, 1.15, "about 2.5 s", "closure needed to warn", K.GREEN, vsize=22, lsize=11)
    K.tile(s, 10.2, 2.75, 2.53, 1.15, f"{fr['p50_ms']:.0f} ms", "API frame latency, median", K.BLUE, vsize=26, lsize=11)
    K.box(s, 7.5, 4.1, 5.23, 2.75)
    K.text(s, 7.7, 4.2, 4.85, 2.6, [
        (f"Geometry only, {geo['face_found']:,} CEW faces, still images.", {"bullet": True, "space_after": 4}),
        (f"Threshold 0.21 calls {pct(A['default_false_closed_rate'])} of open eyes closed; 0.17 does better ({pct(gb['accuracy'])}, chosen in-sample).", {"bullet": True, "space_after": 4}),
        (f"With the quality gates only {pct(A['coverage'])} of these 100 px images are assessed.", {"bullet": True, "space_after": 4}),
        ("Closed and open images come from different sources; scripted rule tests: closures up to 2.4 s never warn, 2.6 s always do.", {"bullet": True}),
    ], size=13, color=K.INK)
    notes[15] = (f"Numbers come from experiments/cew_ear and experiments/rule_engine. Accuracy {pct(gd['accuracy'])} at the default 0.21 on {geo['face_found']:,} faces; the default is conservative. "
                 "This tests only the eye open or closed measurement on still images, not intoxication. Closed faces come from the Internet and open faces from LFW, so separation may be optimistic. "
                 f"API frame latency: median {fr['p50_ms']:.1f} ms, 95th percentile {fr['p95_ms']:.1f} ms on this i5-12500.")

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    shots = [("06_monitor_camera_normal.png", "Live camera path: landmarks on a real face, Normal"),
             ("04_monitor_simulation_warning.png", "Scripted closure raises the warning, with its reason"),
             ("08_history.png", "Event history: 34 events, EAR chart, CSV export"),
             ("12_research_report.png", "Research report on SYNTHETIC data: not performance")]
    for i, (name, cap) in enumerate(shots):
        r, c = divmod(i, 2)
        x, y = 0.6 + c * 4.3, 1.35 + r * 2.9
        K.picture(s, F.shot(name), x, y, 4.1, 2.45, border=True)
        K.text(s, x, y + 2.48, 4.1, 0.35, cap, size=11, color=K.TXT, italic=True)
    K.text(s, 9.4, 1.45, 3.35, 5.3, [
        ("**What to look at**", {"space_after": 8}),
        ("Verdict and reason beside the preview; the measurements come straight from the rule engine.", {"bullet": True, "space_after": 7}),
        ("SIMULATION labels on scripted runs: they are not model predictions.", {"bullet": True, "space_after": 7}),
        ("The 100 % research metrics come from a generated file and say nothing about riders.", {"bullet": True, "space_after": 7}),
        ("Every screen is a real capture of the running app.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[16] = ("Real screenshots from the running application. Top left uses one public-domain portrait as a fake webcam; EAR 0.254 against a calibrated threshold of 0.183. "
                 "Bottom right is the offline evaluation screen on the generated demo file; it carries a Synthetic data banner.")

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{ut['passed']} / {ut['tests']}", "unit tests passed", K.GREEN), (f"{sm['passed']} / {sm['checks']}", "API smoke checks passed", K.BLUE),
                                   ("0 / 300", "scripted sessions with a false warning at normal blink rates", K.ORANGE), (f"{fr['p95_ms']:.0f} ms", "95th percentile frame latency", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.4, 2.9, 1.3, v, l, c, vsize=26, lsize=11)
    rows = [["Observation", "Action"],
            ["Halving the brightness blocks assessment: gates are strict", "Reported as a limitation; fails safe to Unable to assess"],
            [f"Default 0.21 calls {pct(A['default_false_closed_rate'])} of open CEW eyes closed", "Per-rider calibration; threshold sweep documented"],
            ["Only 44 % of 100 px CEW images pass the sharpness gate", "Reported; webcam frames are the intended input"],
            ["Closure + persistence above 3 s misses a 3 s closure", "Parameter trade-off measured and documented"]]
    K.table(s, 0.6, 3.0, [5.6, 6.53], rows, font_size=14, row_h=0.72)
    notes[17] = ("All suites use temporary storage and never touch the real history. The scripted blink streams are generated, not people. "
                 "The observations on the right drove the limitations section of the report.")

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["No labelled intoxication data exists for this project", "Explainable rules; offline pipeline kept for future consented data"],
            ["The batch title promises alcohol detection; a camera cannot confirm it", "Output says possible impairment or drowsiness, never alcohol"],
            ["Browser webcam needs localhost or https", "Frames sent from the browser; MediaPipe runs on the server"],
            ["Shipped virtual environment pointed at a missing Python", "run.bat rebuilds it and downloads the 3.7 MB landmark model"],
            ["Unrelated stills in MediaPipe video mode vary by about 2 points", "Stated as a measurement caveat"],
            ["The API sustains about 30 frames/s but the browser sends about 9", "Both figures reported"]]
    K.table(s, 0.6, 1.5, [5.8, 6.33], rows, font_size=14, row_h=0.72)
    notes[18] = "Each challenge led to a concrete change or a plainly stated limitation."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Document", "Final report, paper and deck, every number regenerated from the metrics files.", K.BLUE, "1"),
        ("Tune", "Choose thresholds on training-only subjects; the CEW sweep points to 0.17-0.18.", K.GREEN, "2"),
        ("Collect data", "If a consented labelled sober / intoxicated dataset is obtained, run the subject-disjoint offline pipeline.", K.ORANGE, "3"),
        ("Extend (future)", "Gaze features, helmet and visor robustness, vehicle integration for a prevention action, embedded deployment.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.1, 1.5, 2.9, 3.6, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.4, 12.13, 1.3)
    K.text(s, 0.85, 5.45, 11.6, 1.2, "Until a labelled intoxication dataset exists, Smart Rider stays a possible-impairment or drowsiness warning, not an alcohol test.",
           size=17, color=K.INK, bold=True, anchor="m")
    notes[19] = "Plan to the final review. The right-hand card lists future work that is outside this submission."

    # ---- 20 thank you
    s = S[19]
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")
    notes[20] = "Likely questions: why rules rather than a trained model; can it detect alcohol (no); what the measured limits are."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.scrub_theme(out_path, "SmartRider")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx")))
