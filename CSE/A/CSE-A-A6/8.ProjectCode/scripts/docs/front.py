"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=22, bold=True, space_after=8)
    d.p(f"Realised as the {F.PRODUCT}: an explainable, camera-based eye and head-pose warning", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p("Medbowli, Meerpet, Balapur, Hyderabad, Telangana - 500 097", align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the creators of the public Closed Eyes in the Wild dataset, and the developers of the open-source software "
        "used in this project (MediaPipe, OpenCV, NumPy, Matplotlib, FastAPI and React).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    g = F.GEO or {}
    d0 = g.get("default_threshold_0.21", {})
    c = F.CEW or {}
    lat = (F.API_LAT or {}).get("POST frame (real face JPEG 640x480, monitoring on)", {})
    rules = F.RULES or {}
    d.p("Drink-driving and drowsy riding are major causes of road deaths, and a rider on a two-wheeler is rarely checked while riding. Hardware alcohol sensors add cost and can be bypassed, "
        "and most driver-monitoring research targets drowsiness only. This project, titled Smart Alcohol Intoxication Detection and Prevention System Using AI, set out to watch a rider's face "
        "with an ordinary webcam and warn when the rider may be impaired.")
    d.p(f"What was built, and what was not. The delivered system, the {F.PRODUCT}, is an explainable rule-based warning. MediaPipe face landmarks give a six-point eye aspect ratio and the image-plane head roll; seven quality gates "
        "refuse to judge poor frames; a per-rider calibration sets the closed-eye threshold; and transparent duration rules output Normal, Possible impairment or drowsiness, or Unable to assess, with the reason shown. "
        "A FastAPI service and a React interface wrap the original Tkinter prototype, and every event is logged as numbers only. **A camera cannot confirm alcohol consumption, and this system does not claim to.** "
        "No deep-learning classifier was trained, because no sober-versus-intoxicated dataset was available, and the planned prevention action (ignition lock or contact alert) was not implemented.")
    if g and d0:
        d.p(f"The eye-state component was measured on the public Closed Eyes in the Wild images ({g['face_found']:,} faces found). At the default threshold of 0.21 the eye aspect ratio separates closed from open eyes with "
            f"{F.pct(d0['accuracy'])} accuracy, {F.pct(d0['recall'])} closed-eye recall and {F.pct(d0['specificity'])} open-eye specificity (area under the ROC curve {F.f3(g['auc_ear_closed_vs_open'])}); "
            f"a lower threshold of about 0.17 separates better. With the quality gates on, only {F.pct(c['coverage'])} of these low-resolution images are assessed; under darkness, blur, sunglasses or a covered eye the system declines to judge. "
            f"On scripted streams the rules never warn on closures up to 2.4 s and always warn from 2.6 s. One camera frame takes {lat.get('p50_ms', 0):.0f} ms through the API on a laptop CPU, and {F.TESTS['unittest']['passed']} unit tests and "
            f"{F.TESTS['web_smoke']['passed']} API checks pass.")
    d.p("The datasets, hardware and target state differ from those of the base paper, so its figures are given as context, not as a ranking. Gaze measurement, a rider study, helmet and visor robustness and any intoxication accuracy "
        "are listed as limitations and future work. The system is a stationary demonstration prototype and not a medical device.")
    d.p("**Keywords:** artificial intelligence, computer vision, driver monitoring, eye aspect ratio, facial landmarks, impairment warning, MediaPipe, traffic safety, FastAPI, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the ROC Curve"), ("BAC", "Blood Alcohol Concentration"),
    ("CEW", "Closed Eyes in the Wild (dataset)"), ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"),
    ("DFD", "Data Flow Diagram"), ("EAR", "Eye Aspect Ratio"), ("ER", "Entity-Relationship"), ("FN / FP", "False Negative / False Positive"), ("FPS", "Frames Per Second"),
    ("GUI", "Graphical User Interface"), ("HMAC", "Hash-based Message Authentication Code"), ("HTTP", "Hypertext Transfer Protocol"), ("IEEE", "Institute of Electrical and Electronics Engineers"),
    ("JPEG / PNG", "Image file formats"), ("JSON", "JavaScript Object Notation"), ("LFW", "Labeled Faces in the Wild (dataset)"), ("NTHU-DDD", "National Tsing Hua University Driver Drowsiness Detection dataset"),
    ("PERCLOS", "Percentage of Eyelid Closure"), ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("RAM", "Random Access Memory"),
    ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("TN / TP", "True Negative / True Positive"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
    ("WHO", "World Health Organization"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
