"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, whose study of diagnostic classification under incomplete clinical data shaped our evaluation, and the maintainers of the "
        "open-source software used in this project (FastAPI, React, Vite, XGBoost, SHAP, scikit-learn, RapidOCR, pypdf, ReportLab, SQLAlchemy, pandas and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    abl, comp = F.ABL, F.COMP
    st = F.STRESS
    d.p("Emergency departments must decide within minutes which patient is seen first. Triage varies with the person on duty, models built on structured fields cannot see what is written only in a "
        f"scanned report, and a prediction without reasons is hard to accept. This project builds {F.PRODUCT}, a decision-support tool that reads a PDF or scanned clinical report "
        "(text layer or local OCR), extracts symptoms, findings, history and six laboratory values with a lexicon that handles negation, and fuses them with eight vitals into "
        f"{F.N_FEATURES} inputs. An isotonic-calibrated XGBoost assigns a priority from 1 (critical) to 5 (non-urgent), a risk level and a confidence; per-patient SHAP shows which findings and vitals drove it. "
        "A clinician accepts or changes the priority with a stored reason; a live queue, a performance page and a summary PDF complete the FastAPI and React application.")
    d.p(f"The model is trained and evaluated on {F.ROWS:,} seeded synthetic cases, because no real patient records were available, so the figures measure recovery of a known rule, not agreement with clinicians. "
        f"On {F.N_TEST:,} held-out cases it reaches {F.pct(F.ACC)} exact accuracy and {F.pct(F.W1)} within one level (macro F1 {F.f3(F.MF1)}); label noise caps any model near {F.pct(F.CEILING, 0)}. "
        f"Fusing report and vitals gives {F.pct(abl['Fused (36)']['accuracy'], 2)} against {F.pct(abl['Vitals only (8)']['accuracy'], 2)} for vitals alone, but a logistic-regression baseline ({F.pct(comp['Logistic regression']['accuracy'])}) scores higher than the chosen model. "
        f"Priority-1 recall is only {F.pct(F.REC[0])}, so the sickest patients are under-triaged most often. The report reader is exact on generated samples but reaches {F.pct(st['recall'], 0)} recall on hand-written fragments. "
        f"A PDF report is read and triaged in a median {F.lat('flow_pdf_then_triage'):.0f} ms and a scanned image in about {F.lat('flow_scan_then_triage') / 1000:.1f} s. "
        "BioBERT, retrieval and multi-agent orchestration were not built.")
    d.p("**Keywords:** emergency triage, clinical decision support, optical character recognition, explainable AI, SHAP, XGBoost, probability calibration, healthcare AI.")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("ATHENA", "Autonomous Triage and Healthcare Electronic Neural Assistant"),
    ("AUC", "Area Under the Curve"), ("CDSS", "Clinical Decision Support System"), ("COPD", "Chronic Obstructive Pulmonary Disease"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CSE", "Computer Science and Engineering"), ("CSV", "Comma-Separated Values"),
    ("CT", "Computed Tomography"), ("DFD", "Data Flow Diagram"), ("ECE", "Expected Calibration Error"), ("ECG", "Electrocardiogram"), ("ED", "Emergency Department"),
    ("EHR", "Electronic Health Record"), ("ER", "Entity-Relationship"), ("ESI", "Emergency Severity Index"), ("FHIR", "Fast Healthcare Interoperability Resources"),
    ("GPU", "Graphics Processing Unit"), ("HL7", "Health Level Seven"), ("HTTP", "Hypertext Transfer Protocol"), ("ICD-10", "International Classification of Diseases, 10th Revision"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("LIME", "Local Interpretable Model-agnostic Explanations"), ("MLP", "Multi-Layer Perceptron"),
    ("NLP", "Natural Language Processing"), ("OCR", "Optical Character Recognition"), ("ONNX", "Open Neural Network Exchange"), ("ORM", "Object-Relational Mapping"),
    ("PDF", "Portable Document Format"), ("QWK", "Quadratic Weighted Kappa"), ("RAG", "Retrieval-Augmented Generation"), ("REST", "Representational State Transfer"),
    ("SHAP", "SHapley Additive exPlanations"), ("SPA", "Single-Page Application"), ("SpO2", "Peripheral Oxygen Saturation"), ("SQL", "Structured Query Language"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("XGBoost", "Extreme Gradient Boosting"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
