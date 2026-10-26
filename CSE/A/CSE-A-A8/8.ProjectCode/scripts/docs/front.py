"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=22, bold=True, space_after=8)
    d.p("LabLens: a web platform that reads lab reports, flags abnormal results, explains them in plain language and tracks them over time",
        align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p("Medbowli, Meerpet, Balapur, Hyderabad, Telangana - 500 097", align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper on retrieval-augmented analysis of blood test reports, whose work set the starting point of this project, and the developers of "
        "the open-source software it relies on (PyMuPDF, OpenCV, Tesseract, scikit-learn, SQLAlchemy, FastAPI, React, Vite and Recharts).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    ex = F.EXTRACTION["overall"]
    lat = F.LATENCY["end_to_end_pipeline_with_db"]["median_ms"]
    up = F.API_LAT["endpoints"]["POST /api/reports (PDF upload)"]["median_ms"]
    py = ((F.TESTS or {}).get("pytest") or {}).get("passed", 0)
    d.p("Laboratory reports are written for clinicians. Patients receive pages of test names, numbers and reference ranges that they cannot interpret, and they keep no history "
        "that would show whether a value is improving. This project builds LabLens, a web platform that reads a lab report supplied as a PDF or an image, extracts every test with its value, "
        "unit and reference range, flags each result as low, normal or high, explains each abnormal result in plain language with a medical disclaimer, and stores the report so "
        "that every test can be followed across reports on trend charts and in side-by-side comparisons.")
    d.p("The pipeline reads the text layer of digital PDFs with PyMuPDF and sends scans and photographs through OpenCV clean-up and Tesseract; a line parser recovers the fields and a "
        "TF-IDF character n-gram matcher snaps noisy test names to a 55-name dictionary; a rule classifies the result against its own range; rule-based text, with an optional "
        "language-model rewrite, produces the explanations; and SQLite, FastAPI and a React interface complete the application. The abstract's PP-OCR, CRF and retrieval-augmented "
        "components were replaced by these lighter methods, and the report says where that costs capability.")
    d.p(f"No labelled real reports were available, so the stages were measured on 120 generated reports with 1,680 known rows in eight layouts. Row-level precision is {F.f3(ex['row_precision'])}, "
        f"recall {F.f3(ex['row_recall'])} and F1 {F.f3(ex['row_f1'])}; status accuracy on the rows found is {F.f3(ex['status_acc'])}, unit accuracy {F.f3(ex['unit_acc'])}. Reports in which every table "
        f"cell is a separate text line yield no rows at all. A 14-row PDF passes the whole pipeline in a median of {lat:.1f} ms ({up:.0f} ms through the HTTP upload), and {py} automated tests pass. "
        "Scanned-image accuracy and the language-model rewrite were not measured. The system is an educational prototype and not a medical device.")
    d.p("**Keywords:** medical lab report analysis, OCR, information extraction, health insights, trend tracking, explainable health AI, FastAPI, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CRF", "Conditional Random Field"), ("CSS", "Cascading Style Sheets"), ("CSV", "Comma-Separated Values"), ("DFD", "Data Flow Diagram"),
    ("ER", "Entity-Relationship"), ("F1", "Harmonic mean of precision and recall"), ("HbA1c", "Glycated haemoglobin"), ("HDL / LDL", "High- / Low-Density Lipoprotein"),
    ("HMAC", "Hash-based Message Authentication Code"), ("HTTP", "Hypertext Transfer Protocol"), ("JPEG / PNG", "Image file formats"),
    ("JSON", "JavaScript Object Notation"), ("LLM", "Large Language Model"), ("NER", "Named Entity Recognition"), ("OCR", "Optical Character Recognition"),
    ("PDF", "Portable Document Format"), ("RAG", "Retrieval-Augmented Generation"), ("REST", "Representational State Transfer"), ("SPA", "Single-Page Application"),
    ("SQL", "Structured Query Language"), ("TF-IDF", "Term Frequency - Inverse Document Frequency"), ("TSH", "Thyroid-Stimulating Hormone"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("VLM", "Vision-Language Model"), ("WBC", "White Blood Cell"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
