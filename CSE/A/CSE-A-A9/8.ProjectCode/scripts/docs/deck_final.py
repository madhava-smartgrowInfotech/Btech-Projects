"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited).

    python -m scripts.docs.deck_final          # from 8.ProjectCode; writes the deck into the batch folder
"""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520

try:  # the full reference list is written by another worker; used only when it already holds the BotGraph list
    from .refs import REFS
except Exception:  # pragma: no cover
    REFS = None

# Verified references used when refs.py is not (yet) the BotGraph list. Authors/venues were checked against the PDFs in
# Documentary/ and archived_base_papers/, or are well-known classics.
LOCAL_REFS = [
    ("PruferNET", "A. Saha, N. K. Sarkar and S. Pratihar, \"PruferNET: An augmented graph neural network for influential node detection in social networks,\" IEEE Access, vol. 14, 2026, doi: 10.1109/ACCESS.2026.3695637."),
    ("RABot", "L. Zhang, X. Wang, H. Du, Y. Xu, Z. Liu and Y. Liu, \"RABot: Reinforcement-guided graph augmentation for imbalanced and noisy social bot detection,\" arXiv:2602.21749, 2026."),
    ("HW-GNN", "Z. Liu, J. Gao, Z. Ji and L. Zhao, \"HW-GNN: Homophily-aware Gaussian-window constrained graph spectral network for social network bot detection,\" arXiv:2511.22493, 2025."),
    ("LGB", "\"LGB: Language model and graph neural network-driven social bot detection,\" arXiv:2406.08762, 2024."),
    ("Semi-supervised", "T. N. Kipf and M. Welling, \"Semi-supervised classification with graph convolutional networks,\" in Proc. ICLR, 2017."),
    ("Graph attention", "P. Velickovic et al., \"Graph attention networks,\" in Proc. ICLR, 2018."),
    ("Inductive", "W. L. Hamilton, R. Ying and J. Leskovec, \"Inductive representation learning on large graphs,\" in Proc. NeurIPS, 2017, pp. 1024-1034."),
    ("Fast unfolding", "V. D. Blondel, J.-L. Guillaume, R. Lambiotte and E. Lefebvre, \"Fast unfolding of communities in large networks,\" J. Stat. Mech., 2008, P10008."),
    ("social circles", "J. McAuley and J. Leskovec, \"Learning to discover social circles in ego networks,\" in Proc. NIPS, 2012."),
    ("Adam", "D. P. Kingma and J. Ba, \"Adam: A method for stochastic optimization,\" in Proc. ICLR, 2015."),
    ("PyTorch", "A. Paszke et al., \"PyTorch: An imperative style, high-performance deep learning library,\" in Proc. NeurIPS, 2019."),
]


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "a9_deck_crops")


def crop_shot(name, box=(264, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def cm_grid(slide, x, y, c, title, w=3.9):
    """Native 2x2 confusion matrix (counts and share of each actual class). c = dict tn, fp, fn, tp."""
    cm = [[c["tn"], c["fp"]], [c["fn"], c["tp"]]]
    ftext(slide, x, y, w, 0.5, title, size=13, color=K.FNAVY, bold=True)
    cw, ch = (w - 1.0) / 2, 0.8
    for j, lab in enumerate(["pred genuine", "pred fake"]):
        ftext(slide, x + 1.0 + j * cw, y + 0.5, cw, 0.3, lab, size=11, color=K.DARK, align="c")
    for i, lab in enumerate(["actual genuine", "actual fake"]):
        ftext(slide, x, y + 0.85 + i * ch, 1.0, ch, lab, size=11, color=K.DARK, anchor="m")
        tot = sum(cm[i])
        for j in range(2):
            good = i == j
            fill = K.RGBColor(0xC6, 0xE5, 0xD3) if good else K.RGBColor(0xFB, 0xDD, 0xC8)
            K.box(slide, x + 1.0 + j * cw + 0.03, y + 0.85 + i * ch + 0.03, cw - 0.06, ch - 0.06, fill=fill, line=None, radius=0.05)
            ftext(slide, x + 1.0 + j * cw, y + 0.85 + i * ch, cw, ch, [(f"{cm[i][j]:,}", {"bold": True, "size": 15}), (f"{cm[i][j] / tot * 100:.1f}% of row", {"size": 10})],
                  size=12, align="c", anchor="m", space_after=0)


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def ref_lines():
    """[(number text, reference)] for the closing slide; real report numbers when refs.py holds the BotGraph list."""
    out = []
    if REFS and "Saha" in REFS[0]:
        for key, _ in LOCAL_REFS:
            for n, r in enumerate(REFS, 1):
                if key.lower() in r.lower():
                    out.append((n, clean_ref(r)))
                    break
        if len(out) >= 8:
            return sorted(out)
    return [(i, t) for i, (_, t) in enumerate(LOCAL_REFS, 1)]


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "graph neural networks, fake profile detection, botnet detection, GCN, GAT, GraphSAGE, Louvain"
    notes = {}

    mod, ex = F.MODELS, F.EXTRA
    bl, gc, ga, sg = mod["baseline"], mod["gcn"], mod["gat"], mod["sage"]
    sc1, sc8 = F.scal_row(1), F.scal_row(8)
    bp = F.BASE_PAPER_RESULTS
    thr = {r["threshold"]: r for r in ex["threshold_sweep"]["sage"]}
    thr_b = {r["threshold"]: r for r in ex["threshold_sweep"]["baseline"]}
    err = ex["error_analysis"]
    abl = ex["ablation"]

    for sl in S:  # footers on every slide except the title
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh in S[0].shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.FULL_TITLE, 28)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and BotGraph: fake accounts and coordinated botnets found from the shape of the social graph."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Models\nTraining, Clustering & Data Design\nImplementation Highlights\nResults: Model Comparison, Confusion Matrix, Threshold\nResults: Botnet Clusters & Feature Groups\nPerformance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Limitations, Future Scope & References", 15)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Fake profiles are created faster than moderators can remove them, and per-account classifiers miss coordinated networks whose single profiles look legitimate. BotGraph treats accounts as nodes and interactions as edges and learns from each account's neighbourhood.", {"bullet": True, "space_after": 8}),
        (f"Graph of {F.NODES:,} accounts and {F.EDGES:,} ties with {F.N_FEATURES} features; GCN, GAT and GraphSAGE compared with a no-graph baseline on one split (mean of 3 seeds).", {"bullet": True, "space_after": 8}),
        (f"GraphSAGE is best: F1 {F.f3(F.mean('sage', 'f1'))} against {F.f3(F.mean('baseline', 'f1'))} without the graph (+{F.lift('sage'):.1f} points); GCN adds only +{F.lift('gcn'):.1f}. Louvain over the embeddings recovers {F.REC['recovered']} of {F.REC['botnets']} injected botnets.", {"bullet": True, "space_after": 8}),
        (f"A moderator queue turns confirm/dismiss decisions into training labels; scalability measured up to {sc8['nodes']:,} accounts and {sc8['edges']:,} ties on CPU.", {"bullet": True, "space_after": 8}),
        ("Caveat: the botnets, the timing and interaction features and the handles are synthetic (seeded). GraphSAGE on node degree alone scores F1 0.965, above all 23 features: the injected structure is easy, so the headline F1 is not real-platform accuracy.", {"bullet": True}),
    ], size=18)
    notes[3] = "Problem, method, headline numbers with their protocol, then the caveat that botnets are injected."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "PruferNET: Prufer-sequence structural features with GraphSAGE"],
            ["Task", bp["task"] + " - an influence task, not bot detection"],
            ["Method", "Edge weight 1/(deg u + deg v) -> minimum spanning tree -> Prufer sequence -> frequency, entropy, unique ratio, transition influence, subtree depth; plus degree, closeness, PageRank, trust score"],
            ["Data", bp["data"]],
            ["Model and split", "GraphSAGE, 64 hidden units, dropout 0.5, AdamW lr 0.003, 300 epochs; 60/20/20 split"],
            ["Reported", "Test accuracy 99.63% with Prufer features against 96.54% with traditional features; 5-fold check on Facebook"],
            ["Not covered", "No fake or bot labels, no coordinated clusters, no moderator loop, no scalability measurement"]]
    ftable(s, 0.55, 2.6, [2.0, 10.2], rows, size=13, row_h=0.6)
    notes[4] = "Saha et al., IEEE Access 2026. BotGraph keeps the data family (SNAP Facebook), the GraphSAGE backbone idea and the 60/20/20 split; the Prufer features are listed as future work."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Finding / limitation"],
            ["Saha et al. [1] (base)", "2026", "GraphSAGE + Prufer-sequence structural features", "99.63% vs 96.54% accuracy on depth classes; no bot labels"],
            ["RABot (Zhang et al.)", "2026", "Neighbourhood-aware oversampling + RL edge filtering", "Targets class imbalance and topological noise; three benchmarks, four GNN backbones"],
            ["HW-GNN (Liu et al.)", "2025", "Homophily-aware Gaussian-window spectral network", "Average +4.3% F1 over existing methods, as reported by the authors"],
            ["LGB (arXiv 2406.08762)", "2024", "Language model + GNN for sparsely connected accounts", "Up to 10.95% over baselines on two datasets, as reported"],
            ["Batch survey: AdapTGN", "2026", "Adaptive temporal-relational GNN for bot detection", "Title only; not independently verified"],
            ["Batch survey: Arabic fake news", "2026", "GNN for fake-news identification in Arabic", "Title only; not independently verified"],
            ["Batch survey: metadata stacking", "2026", "Account metadata, feature selection, stacking", "Title only; nearest match found: Sharma & Borah, IJ-ICT 2026"]]
    ftable(s, 0.55, 1.85, [3.1, 0.7, 4.4, 4.0], rows, size=12, row_h=0.58, align=["l", "c", "l", "l"])
    notes[5] = "Verified from the PDFs (PruferNET, RABot, HW-GNN) or the arXiv record (LGB). The last three rows come from the batch survey with titles only; they are not cited for any number."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 15}) for g in [
        "The surveyed GNN work applies relational learning to influence and node classification; applying it to fake-profile and botnet detection at platform scale is the gap named in the batch survey.",
        "A per-account score leaves a moderator with thousands of items; the coordinated cluster is the unit they can act on.",
        "Whether the graph earns its place needs a no-graph baseline on the same split, features and training budget, with class imbalance (1 account in 8 is fake) and camouflaged profiles.",
        "Scalability has to be measured: how training and inference time grow with nodes and ties.",
        "Moderator decisions are valuable labels, but retraining on them must keep the test split untouched to stay comparable."]], size=20)
    notes[6] = "Five gaps. BotGraph addresses each with a measurement or a feature: baseline comparison, clusters, scalability sweep, moderator retraining."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.6)
    ftext(s, 0.85, 1.9, 11.6, 2.5, [
        "Fake profiles are created faster than manual moderation can remove them.",
        "Per-account classifiers miss coordinated networks whose individual profiles appear legitimate.",
        "Bot behaviour evolves quickly, so static rules and signatures become obsolete.",
        "Class imbalance and label noise degrade conventional supervised models.",
        "Need: a graph-based detector that classifies accounts from their relational context and exposes coordinated botnet clusters at scale."], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", f"Public datasets assembled into a graph of {F.NODES:,} accounts and {F.EDGES:,} ties, {F.N_FEATURES} features per account"),
                                 ("Output", "Fake-probability per account with reasons and neighbourhood, botnet clusters, a moderator queue"),
                                 ("Record", "Decision log, retraining runs with F1 before and after, measured metrics files")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.7, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.8, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.3, 3.5, 1.0, b_, size=14)
    notes[7] = "Four problem points from the batch statement and the stated need, then input, output and record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate([
        "Construct a graph representation of the social platform with accounts as nodes and interactions as edges.",
        "Engineer node features covering profile metadata, posting behaviour and interaction frequency.",
        "Implement and compare GCN, GAT and GraphSAGE architectures for account classification.",
        "Detect coordinated botnet clusters from the learned graph embeddings.",
        "Evaluate detection accuracy, precision and recall against labelled fake-account datasets.",
        "Assess scalability as the number of nodes and edges grows."], 1)], size=22)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Learn - from graph to scores", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Public datasets", K.FNAVY), ("Graph + 23 features", K.FBLUE), ("GCN / GAT / GraphSAGE", K.GREEN), ("Threshold on validation", K.ORANGE), ("Account scores", K.PURPLE)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Act - from scores to moderator decisions", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Embeddings", K.FNAVY), ("Louvain clusters", K.FBLUE), ("Moderator queue", K.GREEN), ("Confirm / dismiss", K.ORANGE), ("Retrain on decisions", K.PURPLE)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, [
        ("Every account is scored from its own features and its neighbourhood; a feature-only MLP on the same split shows what the graph adds.", {"bullet": True, "space_after": 7}),
        ("Accounts scoring 0.30 or more are grouped by community detection over the learned embeddings and real ties; dense groups are labelled coordinated, sparse ones loose.", {"bullet": True, "space_after": 7}),
        ("A decision on a cluster applies to all members; retraining never touches the test split, so before and after F1 are comparable.", {"bullet": True}),
    ], size=16)
    notes[9] = "The proposed system as two lanes: learn and act."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.0, 5.1)
    ftext(s, 7.8, 1.95, 5.0, 4.9, [
        (f"Web app: React 18 + Vite 5 on port {F.WEB_PORT}, V4 Soft Pastel Calm design kit", {"bullet": True, "space_after": 8}),
        (f"API: FastAPI on port {F.API_PORT}, 18 JSON endpoints, JWT roles (analyst, moderator)", {"bullet": True, "space_after": 8}),
        ("ml package: graph builder, GCN / GAT / GraphSAGE in plain PyTorch sparse ops, training, clustering, evaluation", {"bullet": True, "space_after": 8}),
        ("Storage: SQLite for users, decisions and retraining runs; graph, weights, scores and clusters as files loaded once at start-up", {"bullet": True, "space_after": 8}),
        ("No calls to outside services at run time", {"bullet": True}),
    ], size=15)
    notes[10] = "Web app, API, ML pipeline and storage. Everything runs on one Windows computer."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("12_model_architecture.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, [
        (f"{F.N_FEATURES} standardised features: profile 12, posting 4, temporal 6, interaction 1.", {"bullet": True, "space_after": 7}),
        ("Two message-passing layers, 64 hidden units, dropout 0.3; GCN, GraphSAGE and 4-head GAT, plus an MLP baseline.", {"bullet": True, "space_after": 7}),
        ("Adam, lr 0.01, weight decay 5e-4; class-weighted cross-entropy (weight sqrt(neg/pos)); up to 160 epochs, patience 40 on validation F1.", {"bullet": True, "space_after": 7}),
        ("Decision threshold tuned on the validation set; three seeds per model; best architecture chosen on validation F1.", {"bullet": True, "space_after": 7}),
        ("Scores explained by gradient x input attribution (first-order approximation).", {"bullet": True}),
    ], size=14)
    notes[11] = "Same split, features and training budget for every model, so differences come from the architecture."

    # ---- 12 dataset
    s = S[11]
    set_title(s, "Dataset & Graph")
    wipe(s)
    ftile(s, 0.55, 1.85, 2.95, 1.05, f"{F.NODES:,}", f"accounts ({F.REAL_NODES:,} real + {F.INJ_NODES} injected)", K.FBLUE, 26)
    ftile(s, 0.55, 3.05, 2.95, 1.05, f"{F.EDGES:,}", f"ties ({F.REAL_EDGES:,} real + {F.INJ_EDGES:,} injected)", K.GREEN, 26)
    ftile(s, 0.55, 4.25, 2.95, 1.05, f"{F.FAKE}", f"fake accounts ({F.FAKE_SHARE * 100:.1f}%): 420 in 8 botnets + {F.LONE_SPAMMERS} lone", K.ORANGE, 26)
    rows = [["Source", "Licence", "Used for"],
            ["SNAP Facebook social circles (ten ego networks)", "public research data", "real topology: 3,963 accounts, 88,156 ties"],
            ["Instagram fake / spammer / genuine accounts (696 rows)", "CC BY 3.0", "profile attributes, resampled"],
            ["Twitter bot detection dataset (50,000 rows)", "CC0", "posting behaviour (label carries almost no signal)"],
            ["Injected botnets (seed 7)", "generated", "8 botnets, 150 lone spammers, timing features"]]
    ftable(s, 3.75, 1.85, [3.9, 1.5, 3.65], rows, size=12, row_h=0.62)
    ftext(s, 3.75, 5.1, 9.0, 1.8, [
        (f"Stratified 60/20/20 split: {F.SPLITS['train']:,} train, {F.SPLITS['val']} validation, {F.SPLITS['test']} test ({F.TEST_FAKE} fake), the same for every model.", {"bullet": True, "space_after": 6}),
        (f"Camouflage by design: {F.CAMO_BOTS} bots get a genuine-looking profile and {F.SUSPICIOUS_GENUINE} genuine accounts a fake-looking one, so the profile alone cannot solve the task.", {"bullet": True, "space_after": 6}),
    ], size=14)
    ftext(s, 0.55, 5.55, 2.95, 1.3, "Nothing is scraped from a live platform.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "Topology is real; botnets, timing and interaction features and handles are synthetic and seeded so cluster detection can be scored against a known answer."

    # ---- 13 training and clustering workflow
    s = S[12]
    set_title(s, "Training & Clustering Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_pipeline.png"), 0.55, 1.8, 5.9, 4.7)
    K.picture(s, F.diag("11_flowchart_clustering.png"), 6.75, 1.8, 6.0, 4.7)
    ftext(s, 0.55, 6.55, 12.2, 0.5, "Left: build, 4 models x 3 seeds, threshold on validation, serve best. Right: candidates p >= 0.30, similarity graph + real ties, Louvain, groups < 5 dropped, coordinated if density >= 5%.", size=12, color=K.FNAVY)
    notes[13] = "The two core algorithms of the ml package as flowcharts (ml/train.py and ml/clustering.py)."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_flow.png"), 0.55, 1.8, 6.4, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[14] = "Sequence of sign-in, scoring, a moderator decision and retraining; and the use cases of the analyst, the moderator and the researcher."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.7, 4.6)
    rows = [["Store", "Content"], ["users", "email, name, role, bcrypt hash"], ["reviews", "decision, note, score, members, used_in_run"], ["training_runs", "F1 before / after, labels used, status"],
            ["data/processed/", "graph.npz, accounts.csv"], ["models/", "weights, scores.npz, clusters.json"], ["experiments/", "metrics and scalability JSON"]]
    ftable(s, 7.5, 1.9, [2.0, 3.3], rows, size=11, row_h=0.5)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "Three SQLite tables hold what moderators do; the graph and the model outputs are files, so the database stays small.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "reviews and training_runs have a foreign key to users; the graph, scores and clusters are logical entities stored as files."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("18", "API endpoints", K.FBLUE), (f"{F.SMOKE_PASSED} / {F.SMOKE_TOTAL}", "smoke-test checks pass", K.GREEN), ("10", "app screens (React)", K.ORANGE), ("~2,400", "lines of Python (ml, backend, scripts)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        ("Models are written against sparse tensors in plain PyTorch; torch-geometric was avoided because its Windows wheels often fail to build.", {"bullet": True, "space_after": 9}),
        ("Login returns a signed JWT (12 hours); passwords are bcrypt hashes; analysts get 403 on decisions and retraining; CORS allows only the web origin.", {"bullet": True, "space_after": 9}),
        (f"The API loads graph, scores and clusters once at start-up ({F.STARTUP_S} s); every screen has loading, empty and error states.", {"bullet": True, "space_after": 9}),
        ("Retraining runs as a background job: decisions become labels, the test split is never touched, F1 before and after is stored per model.", {"bullet": True, "space_after": 9}),
        ("Not present: password reset, MFA, rate limiting, live platform ingestion, a pytest suite.", {"bullet": True}),
    ], size=16)
    notes[16] = "Highlights and honest absences."

    # ---- 17 results: model comparison
    s = S[16]
    set_title(s, "Results – Model Comparison")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.f3(F.mean("sage", "f1")), "GraphSAGE F1 (served model)", K.FBLUE), (F.f3(F.mean("sage", "precision")), "GraphSAGE precision", K.GREEN), (F.f3(F.mean("sage", "recall")), "GraphSAGE recall", K.ORANGE), (F.f3(F.mean("sage", "roc_auc")), "GraphSAGE ROC-AUC", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.1, v, l, c, 26)
    names = [F.LABEL[m] if m != "baseline" else "No-graph baseline" for m in F.ORDER]
    K.bar_chart(s, 0.55, 3.0, 12.2, 3.2, names,
                [("Precision", [F.mean(m, "precision") for m in F.ORDER]), ("Recall", [F.mean(m, "recall") for m in F.ORDER]), ("F1", [F.mean(m, "f1") for m in F.ORDER])],
                colors=[K.FBLUE, K.GREEN, K.ORANGE], number_format="0.00", font_size=11, max_value=1.0, legend=True, label_size=10)
    ftext(s, 0.55, 6.25, 12.2, 0.7, f"Test split of {F.TEST_N} accounts ({F.TEST_FAKE} fake), mean of 3 seeds. F1 lift over the baseline: GCN +{F.lift('gcn'):.1f}, GAT +{F.lift('gat'):.1f}, GraphSAGE +{F.lift('sage'):.1f} points. GCN's precision is high but its recall ({F.f3(F.mean('gcn', 'recall'))}) is below the baseline's ({F.f3(F.mean('baseline', 'recall'))}). F1 std: GraphSAGE {F.std('sage', 'f1'):.3f}, GAT {F.std('gat', 'f1'):.3f}.", size=12, color=K.FNAVY, italic=True)
    notes[17] = "Accuracy is 0.959 for the baseline only because 87.4% of accounts are genuine; F1, precision and recall are the honest metrics here."

    # ---- 18 confusion matrix & threshold
    s = S[17]
    set_title(s, "Confusion Matrix & Threshold Analysis")
    wipe(s)
    t_s = F.MODELS["sage"]["threshold"]
    cm_grid(s, 0.55, 1.75, F.MODELS["baseline"]["confusion"], "No-graph baseline, test (908), seed 0")
    cm_grid(s, 4.7, 1.75, F.MODELS["gcn"]["confusion"], "GCN, same test split")
    cm_grid(s, 8.85, 1.75, F.MODELS["sage"]["confusion"], "GraphSAGE (served), same test split")
    K.picture(s, F.fig("f07_threshold_sweep.png"), 0.55, 4.35, 5.4, 2.65)
    t4, t5 = thr[0.4], thr[0.5]
    lo, hi = min(thr[t]["f1"] for t in (0.3, 0.35, 0.4, 0.45, 0.5, 0.55, 0.6)), max(thr[t]["f1"] for t in (0.3, 0.35, 0.4, 0.45, 0.5, 0.55, 0.6))
    ftext(s, 6.15, 4.4, 6.65, 2.6, [
        (f"GraphSAGE at the served threshold {t_s}: precision {F.f3(t5['precision'])}, recall {F.f3(t5['recall'])}; {t5['fp']} false alarms and {t5['fn']} missed fakes among {F.TEST_N} accounts.", {"bullet": True, "space_after": 6}),
        (f"F1 stays between {lo:.3f} and {hi:.3f} for every threshold from 0.30 to 0.60, so the result does not hinge on the tuned value; the sweep peaks at {t4['f1']:.3f} (0.40), on the test split itself, so it is analysis only.", {"bullet": True, "space_after": 6}),
        (f"Baseline at 0.5: {thr_b[0.5]['fp']} false alarms and {thr_b[0.5]['fn']} misses; to reach recall 0.91 it must accept precision 0.61 (threshold 0.25). GCN: {F.conf('gcn')[1]} false alarms, {F.conf('gcn')[2]} misses.", {"bullet": True, "space_after": 6}),
        ("The sweep uses the served scores (seed 0); the served threshold was tuned on validation.", {"bullet": True}),
    ], size=12)
    notes[18] = "Rows are the actual class; confusion matrices are from seed 0 (stored per model). The threshold sweep is descriptive, not a tuned result."

    # ---- 19 botnets and feature groups
    s = S[18]
    set_title(s, "Results – Botnet Clusters & Feature Groups")
    wipe(s)
    K.picture(s, F.fig("f06_cluster_recovery.png"), 0.55, 1.75, 6.0, 3.0)
    K.picture(s, F.fig("f08_ablation.png"), 6.8, 1.75, 6.0, 3.0)
    a = lambda k, m: abl[k][m]["f1_mean"]
    ftext(s, 0.55, 4.85, 6.0, 2.2, [
        (f"{F.REC['recovered']} of {F.REC['botnets']} injected botnets are single Louvain communities: mean recall {F.REC['mean_recall']:.3f}, mean precision {F.REC['mean_precision']:.3f}; one extra loose group of 160 lone spammers and look-alikes (tie density 0.0014).", {"bullet": True, "space_after": 5}),
        (f"Test split, served model: all {err['by_kind']['botnet member']['sage_caught']} botnet members caught (baseline {err['by_kind']['botnet member']['baseline_caught']}) but only {err['by_kind']['lone spammer']['sage_caught']} of {err['by_kind']['lone spammer']['n']} lone spammers (baseline {err['by_kind']['lone spammer']['baseline_caught']}): all {len(err['missed'])} misses are lone spammers of degree 4-8.", {"bullet": True, "space_after": 5}),
        ("Known-answer test on seeded botnets, not evidence about real ones.", {"bullet": True}),
    ], size=13)
    ftext(s, 6.8, 4.85, 6.0, 2.2, [
        (f"Ablation F1 (3 seeds, MLP / GraphSAGE): all 23 features {a('all 23 features', 'baseline'):.3f} / {a('all 23 features', 'sage'):.3f}; without the temporal group {a('without temporal pattern', 'baseline'):.3f} / {a('without temporal pattern', 'sage'):.3f} (largest drop, and those features are synthetic); without posting behaviour {a('without posting behaviour', 'baseline'):.3f} / {a('without posting behaviour', 'sage'):.3f} (adds nothing).", {"bullet": True, "space_after": 5}),
        (f"GraphSAGE on node degree alone scores {a('degree only (structure, no attributes)', 'sage'):.3f}, above all 23 features: the injected structure is easy, so headline F1 is not real-platform accuracy.", {"bullet": True, "bold": True}),
    ], size=13)
    notes[19] = "Botnet recovery per botnet and the feature-group ablation (scripts/run_extra_experiments.py). The ablation retrains in memory, so its all-features row can differ from the main table in the third decimal."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    med = [v[0] for v in F.LATENCY.values()]
    for i, (v, l, c) in enumerate([(f"{F.SMOKE_PASSED} / {F.SMOKE_TOTAL}", "smoke-test checks pass", K.FBLUE), (f"{min(med):.0f}-{max(med):.0f} ms", "median API latency, 10 endpoints", K.GREEN),
                                   (f"{F.scal_row(8)['models']['sage']['train_seconds_per_epoch']:.2f} s", f"GraphSAGE epoch, {sc8['nodes']:,} accounts", K.ORANGE), (f"{F.LOGIN_MS} ms", "sign-in (bcrypt check)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f05_scalability.png"), 0.55, 2.9, 6.3, 2.5)
    K.picture(s, F.fig("f10_api_latency.png"), 7.0, 2.9, 5.8, 2.5)
    rows = [["Check", "Result"],
            ["End-to-end smoke test (own API, scratch database)", f"{F.SMOKE_PASSED} of {F.SMOKE_TOTAL} passed: auth and roles, graph, models beat baseline, scores, clusters, decisions, retraining"],
            ["Scalability, disjoint copies, CPU", f"{sc1['nodes']:,} to {sc8['nodes']:,} accounts ({sc1['edges']:,} to {sc8['edges']:,} ties): GraphSAGE epoch {sc1['models']['sage']['train_seconds_per_epoch']:.2f} to {sc8['models']['sage']['train_seconds_per_epoch']:.2f} s, GAT {sc1['models']['gat']['train_seconds_per_epoch']:.2f} to {sc8['models']['gat']['train_seconds_per_epoch']:.2f} s"],
            ["Not tested", "load and concurrency, adversarial evasion, platform-scale graphs, formal user acceptance, a pytest suite"]]
    ftable(s, 0.55, 5.5, [3.7, 8.5], rows, size=11, row_h=0.36)
    notes[20] = "Latency: 20 requests per endpoint on 127.0.0.1, warm server, shared i5-12500 (upper bounds). Cost grows roughly in proportion to the number of ties; disjoint copies understate a well-mixed real graph."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper")
    wipe(s)
    rows = [["Aspect", "PruferNET [1]", "BotGraph"],
            ["Task", "3-class depth label in a spanning tree (influence)", "Binary fake-account label; botnet clusters"],
            ["Labels", "Derived from tree depth; no fake accounts", f"{F.FAKE} of {F.NODES:,} accounts fake ({F.FAKE_SHARE * 100:.1f}%); partly synthetic"],
            ["Data", "SNAP Facebook 4,039 nodes / 88,234 edges; Email-Eu-core", f"SNAP Facebook topology + {F.INJ_NODES} injected accounts; Instagram and Twitter rows"],
            ["Model", "GraphSAGE + Prufer features, dropout 0.5, AdamW 0.003, 300 epochs", "GCN, GAT, GraphSAGE vs MLP; dropout 0.3, Adam 0.01, <= 160 epochs"],
            ["Split", "60/20/20", "60/20/20 stratified, 3 seeds"],
            ["Reported", "Accuracy 99.63% with Prufer, 96.54% without", f"GraphSAGE F1 {F.f3(F.mean('sage', 'f1'))}, accuracy {F.f3(F.mean('sage', 'accuracy'))}; baseline F1 {F.f3(F.mean('baseline', 'f1'))}"],
            ["Beyond the paper", "-", "Clusters 8 of 8, moderator retraining, scalability sweep"]]
    ftable(s, 0.55, 1.8, [1.7, 5.2, 5.3], rows, size=12, row_h=0.56)
    fbox(s, 0.55, 6.4, 12.2, 0.6, fill=K.FZEBRA)
    ftext(s, 0.7, 6.4, 11.9, 0.6, "Not comparable: different task, labels and data, so 99.63% and 0.985 measure different things; no ranking is claimed.", size=13, color=K.FNAVY, bold=True, anchor="m")
    notes[21] = "State plainly that tasks, labels and datasets differ. What BotGraph takes from the paper: SNAP Facebook data, GraphSAGE, the 60/20/20 split, and structure-plus-attributes features."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Overview and Model Comparison")
    wipe(s)
    K.picture(s, crop_shot("04_overview.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("13_models_comparison.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, f"Left: trust and safety overview - {F.NODES:,} accounts, {F.FLAGGED} flagged (12.1%), 8 botnet clusters plus 1 loose group, F1 93.9% (+11.2 points vs no graph). Right: the four architectures on the same split with F1 ± std, ROC-AUC, PR-AUC, lift and training time.", size=13, color=K.FNAVY)
    notes[22] = "Real captures of the running app with the shipped model and committed graph."
    s = S[22]
    set_title(s, "Output – Botnet Clusters and Account Explanation")
    wipe(s)
    K.picture(s, crop_shot("08_clusters.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("07b_account_reasons_and_neighbourhood.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: Louvain groups with size, average fake probability, tie density and shared peak hour - eight coordinated clusters and one 160-account loose group (density 0%). Right: why one account scores high - account age, peak-hour share and burstiness push towards fake, and its 16 flagged neighbours add to the margin.", size=13, color=K.FNAVY)
    notes[23] = "Cluster list and the 'Why this score' view (gradient x input attribution, a first-order approximation)."
    s = S[23]
    set_title(s, "Output – Moderator Queue and Decisions")
    wipe(s)
    K.picture(s, crop_shot("11_queue_clusters.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("07_account_detail.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: queue of clusters to review with the retrain panel (demo database: 4 decisions confirmed fake, 1 dismissed, 38 labelled accounts ready for the next run). Right: an account page - fake probability 100% from all four models, member of a 56-account cluster, decision recorded as confirmed fake.", size=13, color=K.FNAVY)
    notes[24] = "Decisions are stored in SQLite and used as training labels at the next retraining run; the test split is never touched."

    # ---- 25 conclusion and objectives
    s = S[24]
    set_title(s, "Conclusion & Objectives vs Outcomes")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 1.4, [
        (f"A working end-to-end system: GraphSAGE reaches F1 {F.f3(F.mean('sage', 'f1'))} against {F.f3(F.mean('baseline', 'f1'))} without the graph; {F.REC['recovered']} of {F.REC['botnets']} injected botnets are recovered; {F.SMOKE_PASSED} smoke-test checks pass.", {"bullet": True, "space_after": 4}),
        (f"Findings: the graph helps, unequally (GCN +{F.lift('gcn'):.1f} points, GraphSAGE +{F.lift('sage'):.1f}); node degree alone nearly solves the injected benchmark, so F1 is not real-platform accuracy.", {"bullet": True}),
    ], size=14)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Graph of accounts and interactions", f"Met: {F.NODES:,} nodes, {F.EDGES:,} ties, documented schema"],
            ["2", "Node features", "Met: 23 features in four groups; no text features"],
            ["3", "Compare GCN, GAT, GraphSAGE", "Met: one split, 3 seeds, plus a no-graph baseline"],
            ["4", "Botnet clusters from embeddings", "Met on injected botnets (8 of 8); real botnets untested"],
            ["5", "Accuracy, precision, recall", "Met: six metrics, ROC and PR, confusion matrices; labels partly synthetic"],
            ["6", "Scalability", f"Met to {sc8['nodes']:,} accounts and {sc8['edges']:,} ties on CPU; platform scale untested"]]
    ftable(s, 0.55, 3.4, [0.5, 4.2, 7.5], rows, size=12, row_h=0.5, align=["c", "l", "l"])
    notes[25] = "Be explicit about which objectives were met and what 'met' covers."

    # ---- 26 limitations and future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 5.1, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("Botnets are injected and seeded", {"bullet": True, "space_after": 9}),
        ("Temporal and interaction features are synthetic; profile and posting rows are resampled", {"bullet": True, "space_after": 9}),
        ("No text or linguistic features; static snapshot graph", {"bullet": True, "space_after": 9}),
        ("One split, 3 seeds; GAT varies by ±0.01 between runs", {"bullet": True, "space_after": 9}),
        ("Scalability only to 36k accounts, on CPU", {"bullet": True, "space_after": 9}),
        ("No adversarial, load or user-acceptance testing; demo authentication only", {"bullet": True}),
    ], size=16)
    fbox(s, 6.8, 1.8, 6.0, 5.1, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Prufer / spanning-tree structural features from the base paper", {"bullet": True, "space_after": 9}),
        ("Text features from posts; temporal graphs that follow accounts over time", {"bullet": True, "space_after": 9}),
        ("Neighbour sampling and GPU training for platform-scale graphs", {"bullet": True, "space_after": 9}),
        ("Real labelled botnets; cross-validation and confidence intervals", {"bullet": True, "space_after": 9}),
        ("Robustness to bots that adapt; live ingestion with rate limits and MFA", {"bullet": True}),
    ], size=16)
    notes[26] = "Limitations first, then concrete next steps."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    lines = [f"[{n}] {t}" for n, t in ref_lines()]
    half = (len(lines) + 1) // 2
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in lines[:half]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in lines[half:]], size=11)
    notes[27] = "Selected references; the complete numbered list, including the batch-survey titles that could not be verified, is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "BotGraph")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
