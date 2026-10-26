"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the related work surveyed in Chapter 2, the Stanford Network Analysis Project for the social-circles data, the "
        "publishers of the Instagram and Twitter labelled datasets used in this project, and the maintainers of the open-source software used to build it (PyTorch, NetworkX, "
        "scikit-learn, FastAPI, SQLAlchemy, React, Vite and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    sage, base = "sage", "baseline"
    cm = F.conf(sage)
    sc1, sc8 = F.scal_row(1), F.scal_row(8)
    d.p("Fake profiles and coordinated botnets are created faster than moderators can remove them, and a classifier that reads one account at a time cannot see a network "
        "whose members each look genuine. This project builds "
        f"{F.PRODUCT}, which treats the platform as a graph in which accounts are nodes and interactions are edges, learns from each account's neighbourhood with graph neural "
        "networks, and groups flagged accounts into botnet clusters that a moderator can confirm or dismiss. The graph holds "
        f"{F.NODES:,} accounts and {F.EDGES:,} ties: the topology of ten Facebook ego networks, labelled profile and posting rows from public datasets, and eight seeded botnets "
        f"and {F.LONE_SPAMMERS} lone spammers added so that clusters can be scored against a known answer. Each account carries {F.N_FEATURES} features covering profile metadata, "
        "posting behaviour, timing and interaction frequency.")
    d.p(f"GCN, GAT and GraphSAGE were written in PyTorch and compared with a feature-only baseline on one stratified split ({F.TEST_N} test accounts, {F.TEST_FAKE} fake), over three seeds. "
        f"The baseline reaches F1 {F.f3(F.mean(base, 'f1'))}; GCN {F.f3(F.mean('gcn', 'f1'))}, GAT {F.f3(F.mean('gat', 'f1'))} and GraphSAGE {F.f3(F.mean(sage, 'f1'))}, "
        f"with GraphSAGE at precision {F.f3(F.mean(sage, 'precision'))}, recall {F.f3(F.mean(sage, 'recall'))} and ROC-AUC {F.f3(F.mean(sage, 'roc_auc'))} "
        f"(seed 0: {cm[1]} false positives and {cm[2]} missed fake accounts). GCN gains only {F.lift('gcn'):.1f} F1 point and loses recall. Louvain community detection over the learned "
        f"embeddings recovers all {F.REC['recovered']} injected botnets as single communities (mean precision {F.f3(F.REC['mean_precision'])}). A training epoch grows from "
        f"{sc1['models'][sage]['train_seconds_per_epoch']:.2f} s to {sc8['models'][sage]['train_seconds_per_epoch']:.2f} s as the graph grows from {sc1['nodes']:,} to {sc8['nodes']:,} accounts on a CPU. "
        f"A FastAPI service, a React moderator application and a {F.SMOKE_TOTAL}-check smoke test complete the system. The botnets are injected and several features are synthetic, so the "
        "results show what the method can do on this benchmark, not how it performs on a live platform; the base paper's figures measure a different task and are not compared.")
    d.p("**Keywords:** graph neural networks, fake profile detection, botnet detection, GCN, GAT, GraphSAGE, community detection, social network analysis.")


ABBR = [
    ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("bcrypt", "Adaptive password-hashing function"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("DFD", "Data Flow Diagram"),
    ("ELU", "Exponential Linear Unit"), ("ER", "Entity-Relationship"), ("F1", "Harmonic mean of precision and recall"), ("GAT", "Graph Attention Network"),
    ("GCN", "Graph Convolutional Network"), ("GNN", "Graph Neural Network"), ("GraphSAGE", "Graph SAmple and aggreGatE"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("kNN", "k Nearest Neighbours"), ("MLP", "Multi-Layer Perceptron"), ("MST", "Minimum Spanning Tree"),
    ("NLP", "Natural Language Processing"), ("ORM", "Object-Relational Mapping"), ("PR", "Precision-Recall"), ("REST", "Representational State Transfer"),
    ("ROC", "Receiver Operating Characteristic"), ("SNAP", "Stanford Network Analysis Project"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"),
    ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
