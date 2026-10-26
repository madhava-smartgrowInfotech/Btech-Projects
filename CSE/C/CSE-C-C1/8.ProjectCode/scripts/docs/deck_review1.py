"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import (EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIMITATIONS_EXISTING, LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT,
                  SUPPORT_EVIDENCE)
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout / master footer text carries the template's deck title
    parts = list(prs.slide_layouts) + list(prs.slide_masters)
    for lay in parts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "secure electronic voting, Aadhaar verification, blockchain, cryptographic hashing, RBAC, audit logging"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: AegisBallot lets a voter prove identity with a checked Aadhaar number, cast an encrypted ballot "
                "that is committed to a hash-chained ledger, and keep a receipt that an auditor can verify. It is a local demonstration, not a national-election system.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: a verifiable ballot from login to receipt")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Manual polling means queues, human error and slow counting. Ordinary e-voting keeps ballots in a database an administrator could change without trace, and without identity binding duplicate or ineligible votes cannot be excluded reliably.", K.ORANGE, "!"),
        ("Our idea", "Check a 12-digit Aadhaar number (Verhoeff checksum, salted-hash registry), issue a single-use token, encrypt the choice with AES-256-GCM and commit it to a hash-chained ledger approved by 2 of 3 validators.", K.BLUE, "i"),
        ("The deliverable", "AegisBallot: a FastAPI and React web app with administrator, voter and auditor roles, receipts, results and an audit log. A local demonstration with a synthetic registry, not a national-election system.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=20)
    notes[3] = "Three points: the integrity and identity problem, the combination of Aadhaar-style verification with an append-only ledger, and what is delivered. Say clearly that the registry is synthetic and UIDAI is not contacted."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Four tiers: MySQL, Ethereum nodes (Geth test chain), a C# web tier and client devices. Smart contracts register voters and candidates, cast and count votes; biometric authentication blocks duplicate votes in simulation.", K.BLUE, "1"),
        ("Data and reported results", "1000 synthetic voters; synchronous and asynchronous transactions in batches of 10, 20, 30 and 50. Times are plotted in figures (about 2 s per registration and 3-4 s per vote, read by eye); no table of values.", K.GREEN, "2"),
        ("Relevance to AegisBallot", "Same goal and a similar benchmark design. Different: an Aadhaar-style identity check, a local permissioned ledger instead of Ethereum, an admin and auditor workflow and tamper tests. Results are not directly comparable.", K.ORANGE, "3"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=17)
    notes[4] = ("Read the citation. The base paper builds and times a blockchain voting model for Oman; it concludes that pilots and public-trust studies are needed and names Layer-1 throughput and synthetic data as limits. "
                "AegisBallot keeps the idea of a ledger plus a registry but is a different, smaller system, so numbers are never compared directly.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for r in LIT_TABLE:
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.3, [2.5, 0.7, 4.6, 4.33], rows, font_size=12, row_h=0.6, header_size=13, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: Ethereum-based models with biometrics (base paper), Aadhaar and hash-chain concepts without performance data, surveys that name privacy and speed as the main issues, "
                "and an analysis of a deployed national-ID system. Items 3-5 and 7 were confirmed online from title, authors and abstract only; the complete list is in the report.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: why design and testing matter")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 5.7, 4.9, ["Ethereum", "Polygon", "Stellar", "BNB Chain", "Tron"],
                     [("tx per second", [119.1, 714.1, 1137, 2222, 2516])], colors=[K.BLUE], number_format="#,##0", font_size=12)
    pt = gf.chart.plots[0].series[0].points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 0.6, 6.3, 5.7, 0.6, "Theoretical throughput quoted by the base paper (Solana 65,000 not drawn). Context only.", size=12, color=K.TXT, italic=True)
    K.text(s, 6.6, 1.45, 6.1, 5.3, [(e, {"bullet": True, "space_after": 11}) for e in SUPPORT_EVIDENCE], size=18, color=K.INK)
    notes[6] = ("Two points. Real deployments of internet and blockchain-marketed voting have been broken by analysts (Estonia, Voatz). "
                "And the base paper itself quotes Ethereum at about 119 transactions per second, far below other chains, which is why a national-scale system would need more than Layer 1.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Register voter", K.NAVY), ("Check identity", K.NAVY), ("Cast ballot", K.NAVY), ("Store ballot", K.NAVY), ("Count and declare", K.NAVY)], size=16, gap=0.4)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Manual and database voting", "\n".join(EXISTING_SYSTEM[i] for i in (0, 1, 3, 4)), K.BLUE, "1", body_size=16)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Blockchain prototypes", EXISTING_SYSTEM[2] + "\nIdentity is often biometric or concept-level; few include an admin console, an auditor tool and measured tamper tests.", K.ORANGE, "2", body_size=16)
    notes[7] = "Today ballots are stored on paper, on machines or in one database that a single administrator controls. Blockchain prototypes improve integrity but are slower and rarely linked to a national-ID check."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Slow and error-prone", "Editable ballot box", "Weak identity binding", "Cost per transaction"]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], LIMITATIONS_EXISTING[i], C[i], str(i + 1), body_size=18)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: " + LIMITATIONS_EXISTING[4], size=15, color=K.INK, anchor="m")
    notes[8] = "Five limitations. Each maps to a design decision later: identity checked before a ballot, an append-only hash-chained ledger, a measured cost per vote, and an auditor who recomputes the result."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 0.9
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.82, g, size=19, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as combining a blockchain ledger with national identity verification and a full election administration workflow in one deployable system. "
                "We target gaps 2 to 5; gap 6 is a scalability limit we state rather than solve.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.9, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.65, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=20, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Voter credentials, a 12-digit Aadhaar number and a candidate choice; administrator settings"),
                                 ("Process", "Verify identity, issue a single-use token, encrypt and sign the ballot, commit to the ledger"),
                                 ("Output", "A receipt, a tally when the election closes, and a chain an auditor can recompute")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.65, 3.85, 1.9)
        K.text(s, x + 0.25, 4.77, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.25, 3.4, 1.25, b_, size=15, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a voting platform that verifies identity, records each ballot immutably and produces an independently auditable result."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 20}) for i, o in enumerate(OBJECTIVES, 1)], size=24, color=K.INK)
    notes[11] = "The six objectives from the proposal. Later reviews report which were met; objective 6 mentions TLS, which is not configured in the build, and that will be stated."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Verify", "12 digits, Verhoeff checksum, salted SHA-256 lookup in a registry; the number must belong to the signed-in voter and the voter must be eligible.", K.BLUE, "1"),
        ("Authorise", "A 256-bit single-use token, stored only as a hash, expires after 10 minutes and separates identity from the ballot.", K.GREEN, "2"),
        ("Record", "AES-256-GCM encrypts the choice; Ed25519 signs the transaction; 2 of 3 validators sign the block.", K.ORANGE, "3"),
        ("Audit", "Auditors recompute hashes and links, browse blocks and read the audit log; results publish on close.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "4", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "One voter, one vote: four independent layers", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.85, "Eligibility row, unique participation row, single-use token and ledger nonce / transaction-hash uniqueness. Built with FastAPI, SQLite and React; no external service is contacted.", size=14, color=K.TXT)
    notes[12] = "Four steps from identity to audit. The one-voter-one-vote rule does not rest on one check but on four independent ones, any of which stops a second ballot."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.5, 1.2, 7.7, 5.75)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite web app (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}) with session cookie, CSRF header, role checks and rate limits.", {"bullet": True}),
        ("Services: identity provider, voting service, crypto helpers, audit log and the local ledger with three validators.", {"bullet": True}),
        ("Data: one SQLite database of 13 tables and three key files; nothing leaves the machine.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser calls the JSON API, the services verify, encrypt and commit, and SQLite holds everything. The ledger is a local simulation of a permissioned chain."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working application", "One vote only", "Detectable edits", "Measured, with gaps", "Local demonstration"]
    for i, (h, b_) in enumerate(zip(heads, EXPECTED_OUTCOMES)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.55, h, b_, C[i % 4], str(i + 1), body_size=20)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=20)
    notes[14] = "Expected outcomes. The later reviews report measured results (benchmark, tests, tamper study) including what could not be measured."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 2, 3, 4, 5, 6, 7, 8, 10, 12, 15]
    K.text(s, 0.6, 1.25, 12.13, 5.7, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 6}) for n in idx], size=13, color=K.INK)
    notes[15] = "Selected references, base paper first; the complete numbered list of 31 entries is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: is Aadhaar really verified (no - a synthetic local registry, UIDAI is not contacted); is the ledger decentralised (no - three validators on one machine); "
                 "can an administrator change a vote (the auditor's check catches some edits, not all; the later reviews show the tamper study).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
