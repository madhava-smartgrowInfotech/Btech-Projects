"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p("AegisBallot: a local e-voting web application with Aadhaar verification, encrypted ballots and an auditable hash-chained ledger", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the maintainers of the open-source software used in this project (FastAPI, SQLAlchemy, the cryptography and argon2-cffi "
        "libraries, React, Vite, Tailwind CSS and Recharts), and the researchers whose published work on blockchain voting is listed in the references.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    s, h = F.BENCH_SYNC, F.LAT
    d.p("Elections depend on public confidence that every eligible person votes once, that no vote is altered, and that the result can be checked. Manual polling is slow and error-prone, while a conventional "
        "electronic voting system keeps ballots in a database that an administrator could change without trace. This project presents AegisBallot, a local web application that combines Aadhaar-style identity verification "
        "with an append-only, hash-chained ledger. A voter signs in, enters a 12-digit Aadhaar number that is checked for format, Verhoeff checksum and membership of a salted-hash registry (a synthetic demonstration registry; "
        "UIDAI is not contacted), and receives a single-use token that expires in ten minutes. The chosen candidate is encrypted with AES-256-GCM, the transaction is signed with Ed25519 and committed as a block "
        "approved by two of three validators; the voter keeps a receipt. Administrators manage elections and candidates, and auditors recompute the chain and read the audit log. The backend uses FastAPI and SQLite and the "
        "interface uses React.")
    d.p(f"Evaluation is functional and performance based because the system contains no learned model. All {F.PYTEST_TOTAL} automated tests passed and {F.MATRIX_PASSED} of {F.MATRIX_TOTAL} HTTP security and functional checks returned the expected status. "
        f"A benchmark of {F.N_VOTERS_BENCH} synthetic voters completed every vote without failure at a mean of {s.get('mean_ms', 0):.1f} ms per vote on a shared PC (about 16-17 ms on an idle PC), and a complete voter journey "
        f"over HTTP took {h['end_to_end']['mean_ms']:.0f} ms on average. A tamper study showed that the product's integrity check detects {F.N_TAMPER_PRODUCT} of {F.N_TAMPER} kinds of tampering; it misses edits to the table the tally is read from, "
        "edits made with every hash recomputed, forged validator signatures and removal of the newest block, which a stricter audit written for the study detects. The ledger is a single-machine simulation of a permissioned chain, "
        "ballot secrecy is partial, and TLS is not configured; the base paper's Ethereum results are not directly comparable.")
    d.p("**Keywords:** secure electronic voting, Aadhaar verification, blockchain, cryptographic hashing, role-based access control, audit logging, election integrity, digital democracy.")


ABBR = [
    ("AES", "Advanced Encryption Standard"), ("API", "Application Programming Interface"), ("CSRF", "Cross-Site Request Forgery"), ("DFD", "Data Flow Diagram"),
    ("ER", "Entity-Relationship"), ("GCM", "Galois/Counter Mode"), ("HTTP", "Hypertext Transfer Protocol"), ("HTTPS", "HTTP over TLS"), ("JSON", "JavaScript Object Notation"),
    ("KYC", "Know Your Customer"), ("OTP", "One-Time Password"), ("ORM", "Object-Relational Mapping"), ("PBFT", "Practical Byzantine Fault Tolerance"), ("RBAC", "Role-Based Access Control"),
    ("REST", "Representational State Transfer"), ("SHA", "Secure Hash Algorithm"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("TLS", "Transport Layer Security"),
    ("UIDAI", "Unique Identification Authority of India"), ("UML", "Unified Modeling Language"), ("UI", "User Interface"), ("UUID", "Universally Unique Identifier"), ("ZKP", "Zero-Knowledge Proof"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
