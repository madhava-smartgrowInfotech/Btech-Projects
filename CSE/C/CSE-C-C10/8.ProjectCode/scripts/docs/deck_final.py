"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os
import tempfile
import time

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .lit import GAP, OBJECTIVES, PROBLEM_POINTS, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520

NEEDED = ([F.diag(n) for n in ("01_system_architecture.png", "10_core_algorithm_flowchart.png", "07_sequence_diagram.png", "08_use_case_diagram.png", "09_er_diagram.png")]
          + [F.fig(n) for n in ("f01_model_comparison_real.png", "f02_confusion_matrix_rf.png", "f03_threshold_sweep.png", "f04_feature_importance.png", "f06_imbalance_strategies.png",
                                "f09_representation_study.png", "f13_class_balance.png")])


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def crop(src, box, name, tmp):
    out = os.path.join(tmp, name)
    with Image.open(src) as im:
        im.crop(box).save(out)
    return out


def build(out_path):
    tmp = tempfile.mkdtemp(prefix="vp_final_")
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "startup success prediction, machine learning, ensemble models, retrieval"
    cp.last_modified_by = ""
    notes = {}

    R, X, SA, LAT = F.R, F.X, F.S, F.LAT
    M = R["models"]
    rf, xgb, lr, dt, hgb = M["Random Forest"], M["XGBoost"], M["Logistic Regression"], M["Decision Tree"], M["HistGradientBoosting"]
    cm_rf = rf["confusion_matrix_mean"]
    cm_lr = lr["confusion_matrix_mean"]
    p70 = R["split_70_30"]
    app_lr = SA["models"]["Logistic Regression"]
    allf, nop = X["with_all_features"], X["without_posthoc_features"]
    rk = X["retrieval_only_structured"]
    ba_k = [rk[k]["balanced_accuracy"]["mean"] for k in ("10", "20", "30", "50")]
    auc_k = [rk[k]["roc_auc"]["mean"] for k in ("10", "20", "30", "50")]
    imb = R["imbalance"]
    p_pred = LAT["predict_k10"]
    p_pred50 = LAT["predict_k50"]

    # footers on every slide except the title
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE, 28)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: VenturePulse, a startup outcome predictor that returns a prediction, the comparable historical startups and the feature drivers."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Datasets & Model\nDecision Logic, Interaction & Data Design\nImplementation Highlights\nResults, Confusion Matrix & Threshold Analysis\nPerformance & Testing\nComparison with the Base Papers\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [
        ("Most startups fail and investors judge prospects largely by experience and intuition; a data-driven early signal that names its drivers is missing.", {"bullet": True, "space_after": 10}),
        ("VenturePulse scales funding, rounds, milestones, relationships, investor flags, state and industry, embeds an optional business description (MiniLM, 384-d) and fuses both. It returns SUCCESS or FAILURE with a confidence, the comparable historical startups and the top feature drivers.", {"bullet": True, "space_after": 10}),
        (f"On the real public table ({F.REAL_ROWS} start-ups, structured columns only, 5 x 80/20 splits) Random Forest and XGBoost reach balanced accuracy {F.pm(rf)[:5]} and ROC-AUC {rf['roc_auc']['mean']:.3f} / {xgb['roc_auc']['mean']:.3f}; the base paper's 85-92 % accuracy is not reproduced (best {xgb['accuracy']['mean'] * 100:.1f} %).", {"bullet": True, "space_after": 10}),
        (f"The shipped app runs on a synthetic fallback table ({F.SYN_ROWS:,} rows), whose scores (LR balanced accuracy {app_lr['balanced_accuracy']['mean']:.3f}) are a pipeline check, not evidence.", {"bullet": True, "space_after": 10}),
        (f"A FastAPI service and a React web app complete the system; {F.N_TESTS} automated checks pass.", {"bullet": True}),
    ], size=18)
    notes[3] = "Read the pipeline bullet, then the real-data headline with its protocol, then say plainly that the shipped app uses synthetic data and that the 85-92 % level was not reproduced."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper (Panjwani et al., IJRASET, 2026)"],
            ["Method", "Logistic Regression, Random Forest, SVM and Gradient Boosting on funding, team size, industry and location features; mean/mode imputation, one-hot or label encoding, standard scaling, ADASYN oversampling"],
            ["Data and protocol", "Public Crunchbase and Kaggle startup data (size not stated); 70:30 train / test split; hyper-parameter tuning; Flask back end with an HTML / CSS / JavaScript front end"],
            ["Reported results", "Ensembles (Random Forest, Gradient Boosting) best, accuracy 85-92 %; SVM slower; Logistic Regression the weaker baseline; no per-model table, no ROC-AUC values"],
            ["Stated limits", "Data quality and imbalance, no real-time data, external market conditions not modelled; future work: deep learning, real-time data, text and social media"]]
    ftable(s, 0.55, 2.9, [2.0, 10.2], rows, size=14, row_h=0.78)
    notes[4] = "A structured-data ensemble comparison with a Flask interface. We adopt its model set and its metrics; we add retrieved evidence, a React application and a measured, honest evaluation."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"],
            ["Panjwani et al. 2026 [1]", "LR, RF, SVM, GB; ADASYN; 70:30", "Crunchbase / Kaggle, size not stated", "Ensembles 85-92 % accuracy", "No table, no ROC-AUC"],
            ["Maarouf et al. 2026 [2]", "kNN in-context learning with an LLM vs LR, RF, XGBoost", "4,034 Crunchbase startups, 35.15 % success", "Balanced acc. 0.713 vs LR 0.631, RF 0.634, XGB 0.618", "Needs an LLM and text"],
            ["Mashhadi et al. 2026 [3]", "Interpretable ML, inverse weights, SMOTE-NC, SHAP, temporal splits", "Crunchbase firm-quarter panel 2010-2023 + USPTO", "AUROC 0.921 / 0.817 / 0.872", "Panel data not available here"],
            ["Bidgoli et al. 2024 [7] (via [1])", "RF, GB, SVM, LR compared", "As summarised by [1]", "Ensembles above 80 % accuracy", "Not re-read"],
            ["Zbikowski, Antosiuk 2021 [9] (via [1])", "Bias-free prediction from Crunchbase", "As summarised by [1]", "Look-ahead bias inflates accuracy", "Not re-read"],
            ["Gius 2025 [16]", "Judge disagreement as predictor", "67 venture competitions", "Disagreement predicts success", "Not a classifier"],
            ["Bhandari 2025 [17]", "Survey, correlation, regression", "120 respondents, Nepal", "Venture capital positively associated", "Small survey, association only"]]
    ftable(s, 0.55, 1.85, [2.3, 3.0, 2.5, 2.4, 2.0], rows, size=11, row_h=0.62)
    notes[5] = "Representative work; figures for [1]-[3], [16], [17] come from the papers or confirmed pages, entries marked via [1] are known only through the base paper's reference list."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, [
        (f"**Stated gap:** {GAP}", {"bullet": True, "space_after": 12}),
        ("**Evidence, not just a label:** most systems output a probability only; founders and investors need the comparable cases and the drivers behind it.", {"bullet": True, "space_after": 12}),
        ("**Leakage and imbalance are rarely tested:** fields recorded after the outcome window and class prevalence both change what a score means.", {"bullet": True, "space_after": 12}),
        ("**Fair baselines:** a trained-free retrieval baseline and a simple linear model are seldom reported next to the ensembles.", {"bullet": True, "space_after": 12}),
        ("**Usability:** few works package the predictor so that a non-expert can enter a startup profile and read the result.", {"bullet": True}),
    ], size=18)
    notes[6] = "Five gaps. We build for the second and fifth and test the third and fourth honestly; the stated gap about team data is only partly addressed because no team-experience field exists."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.95, 12.2, 2.1)
    ftext(s, 0.85, 2.0, 11.6, 2.0, PROBLEM_STATEMENT, size=16, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Funding, rounds, milestones, relationships, investor flags, state, industry, optional description"),
                                 ("Output", "SUCCESS / FAILURE, confidence, risk label, comparable startups, top feature drivers"),
                                 ("Evidence", "Held-out metrics on real and synthetic data, ablations, latency, tests")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.35, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.45, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 4.95, 3.5, 1.0, b_, size=14)
    ftext(s, 0.55, 6.15, 12.2, 0.5, "Need: " + PROBLEM_POINTS[0].rstrip(".") + "; the project needs " + "a machine learning model that predicts startup success early and names the factors behind each prediction.", size=12, color=K.FNAVY, italic=True)
    notes[7] = "The statement has three parts: predict, show the evidence and drivers, and report honestly how reliable the estimate is."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion reports which were met: objectives 1 and 2 only partly (no founder or team-experience data exist in either dataset)."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 2.0, 12.2, 1.1, [("Startup profile", K.FNAVY), ("Scale + one-hot (37)", K.FBLUE), ("MiniLM text (384)", K.GREEN), ("Fuse, alpha 0.5", K.ORANGE), ("Classifier + retrieval", K.PURPLE)], size=14, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Structured block: median / most-frequent imputation, StandardScaler, one-hot state and industry, log funding; text block: description cleaned, company name masked, embedded with all-MiniLM-L6-v2.", {"bullet": True, "space_after": 10}),
        ("Fused vector [0.5 s, 0.5 t], L2-normalised; cosine retrieval of k comparable startups from the TRAINING rows only, half SUCCESS and half FAILURE.", {"bullet": True, "space_after": 10}),
        ("Logistic Regression, Decision Tree, Random Forest and XGBoost are compared; the best by balanced accuracy serves predictions with a risk label.", {"bullet": True, "space_after": 10}),
        ("FastAPI service over the library; React app with Dashboard, Analyze, Model performance and Startup insights screens.", {"bullet": True}),
    ], size=17)
    notes[9] = "The pipeline as a row of boxes, then four properties that matter."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.85, 8.3, 5.0)
    ftext(s, 9.1, 2.0, 3.7, 4.7, [
        ("Presentation: React app (port 5410) plus the original Jinja pages", {"bullet": True, "space_after": 8}),
        ("Application: FastAPI (port 8410), 8 JSON endpoints, artefacts loaded once", {"bullet": True, "space_after": 8}),
        ("ML layer: dataset, preprocessing, embeddings, retrieval, models, evaluation", {"bullet": True, "space_after": 8}),
        ("Storage: joblib / npz / pkl files in models/; no database", {"bullet": True}),
    ], size=15)
    notes[10] = "Three tiers plus file storage; the services layer loads the artefacts into an in-memory state on first use and trains if they are missing."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Load table", K.FNAVY), ("Dedup + target map", K.FNAVY), ("Preprocess", K.PURPLE), ("5 x 80/20 splits", K.GREEN), ("Train 4-5 models", K.ORANGE), ("Metrics", K.FBLUE), ("Best model", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Target: acquired / IPO = SUCCESS, closed = FAILURE; duplicates removed; preprocessing is fitted on the training part only.", {"bullet": True, "space_after": 10}),
        ("Protocol: 5 repeated StratifiedShuffleSplit (80 / 20, seed 42); mean ± std; balanced accuracy is the primary metric, with accuracy, precision, recall, F1 and (real-data study) ROC-AUC.", {"bullet": True, "space_after": 10}),
        (f"Two tables: the {F.SYN_ROWS:,}-row synthetic fallback that the app serves, and the real public table ({F.REAL_ROWS} rows) used for all real measurements.", {"bullet": True, "space_after": 10}),
        ("Extra studies on the real table: class-imbalance strategies, threshold sweep, calibration, forest size, 70:30 protocol, synthetic-to-real transfer, leakage ablation, retrieval-only classifier.", {"bullet": True}),
    ], size=17)
    notes[11] = "The method and, as important, the studies that test whether the headline numbers can be trusted."

    # ---- 12 dataset, features, model
    s = S[11]
    set_title(s, "Dataset, Features & Model")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, f"{F.REAL_ROWS}", f"real table: {F.REAL_SUCCESS} success / {F.REAL_FAILURE} failure (64.6 % success)", K.FBLUE, 26)
    ftile(s, 0.55, 3.15, 2.9, 1.1, f"{F.SYN_ROWS:,}", f"synthetic fallback: {F.SYN_SUCCESS} / {F.SYN_FAILURE}", K.GREEN, 26)
    ftile(s, 0.55, 4.4, 2.9, 1.1, "37 + 384", "structured features + MiniLM text dimensions", K.ORANGE, 24)
    K.picture(s, F.fig("f13_class_balance.png"), 3.7, 1.85, 4.4, 2.75)
    K.picture(s, F.fig("f09_representation_study.png"), 8.3, 1.85, 4.5, 2.75)
    ftext(s, 3.7, 4.75, 9.1, 2.1, [
        ("Real table: US start-ups, structured columns only, no description text; success (acquired) is the MAJORITY class. Synthetic table: same schema plus template descriptions that differ by class, so the text channel leaks the label.", {"bullet": True, "space_after": 5}),
        (f"Models: Logistic Regression (balanced, C=1), Decision Tree (depth 8), Random Forest (300 trees), XGBoost (400 trees, depth 6, lr 0.05); HistGradientBoosting in the study.", {"bullet": True, "space_after": 5}),
        ("Kaggle needs a login, so the identical 923-row CSV came from a public GitHub copy.", {"bullet": True}),
    ], size=12)
    ftext(s, 0.55, 5.65, 3.0, 1.2, "Means over 5 splits; real test folds hold 185 rows.", size=11, color=K.FNAVY, italic=True)
    notes[12] = "Two datasets, never mixed up. The representation chart is from the synthetic table, where text carries the label by construction."

    # ---- 13 decision logic
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 5.0, 5.1)
    ftext(s, 5.9, 2.0, 6.9, 4.8, [
        ("The profile is validated (ranges, k in {10, 20, 30, 50}, description up to 2,000 characters), scaled and embedded, then fused.", {"bullet": True, "space_after": 9}),
        ("The served classifier returns P(success); outcome = the larger probability, confidence = max(P).", {"bullet": True, "space_after": 9}),
        ("Retrieval returns k training startups by cosine similarity with a half / half class quota; the weighted vote is sum(sim x label) / sum(sim).", {"bullet": True, "space_after": 9}),
        ("Risk label: Lower risk if SUCCESS with confidence at least 0.70, Moderate risk from 0.55, otherwise Higher uncertainty.", {"bullet": True, "space_after": 9}),
        (f"On the real table, retrieval alone (no trained model) scores balanced accuracy {min(ba_k):.3f}-{max(ba_k):.3f} and ROC-AUC {min(auc_k):.3f}-{max(auc_k):.3f}: equal to Logistic Regression, below Random Forest.", {"bullet": True}),
    ], size=15)
    notes[13] = "The core algorithm. The last bullet is a result from the extra studies: retrieval is competitive with the linear model without any training."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_diagram.png"), 0.55, 1.8, 6.6, 5.1)
    K.picture(s, F.diag("08_use_case_diagram.png"), 7.4, 1.8, 5.4, 5.1)
    notes[14] = "Sequence of one prediction request and the use cases. Sign-in is a demo account that only gates the React screens."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design (file based)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.5)
    rows = [["Store", "Content"], ["best_model.joblib", "served classifier"], ["preprocessor.joblib", "fitted scaler / encoder"], ["train_fused.npz", "fused vectors for retrieval"],
            ["text_embeddings.npz", "cached MiniLM vectors"], ["train_frame.pkl", "training records"], ["metadata.json, evaluation_results.json", "dataset flags, metrics"]]
    ftable(s, 7.8, 1.9, [2.6, 2.4], rows, size=11, row_h=0.5)
    ftext(s, 0.55, 6.35, 12.2, 0.5, "There is no SQL database: the diagram models the artefact files and the in-memory state. Predictions are not stored.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Honest point: the product stores files, not rows."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("8", "JSON endpoints (+ 5 original pages)", K.FBLUE), (str(F.N_TESTS), "automated checks", K.GREEN), ("6", "React screens", K.ORANGE), ("1", "command to start (run.bat)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Models, preprocessor and training frame are loaded once into an in-memory state; if an artefact is missing the training pipeline runs on first start.", {"bullet": True, "space_after": 9}),
        ("Validation: Pydantic ranges on every numeric field (for example funding 0 to 1e10, rounds 0 to 20), k whitelist, description length cap; invalid input returns 422.", {"bullet": True, "space_after": 9}),
        ("Demo sign-in issues an HMAC-signed token valid for 12 hours; a changed signature character is rejected (401).", {"bullet": True, "space_after": 9}),
        ("Each result shows outcome, probabilities, risk label, the top-5 signed feature terms and up to k comparable startups with similarity, funding and summary.", {"bullet": True}),
    ], size=17)
    notes[16] = "Highlights: lazy artefact loading, strict input validation, demo token and explained results."

    # ---- 17 results: model performance
    s = S[16]
    set_title(s, "Results – Model Performance (real table, 5 x 80/20)")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.f3(rf["balanced_accuracy"]["mean"]), f"balanced accuracy, RF and XGBoost ({F.f3(xgb['balanced_accuracy']['mean'])})", K.FBLUE),
                                   (F.f3(rf["roc_auc"]["mean"]), f"ROC-AUC, Random Forest (XGBoost {F.f3(xgb['roc_auc']['mean'])})", K.GREEN),
                                   (F.f3(xgb["accuracy"]["mean"]), f"best accuracy, XGBoost (RF {F.f3(rf['accuracy']['mean'])})", K.ORANGE),
                                   (F.f3(lr["balanced_accuracy"]["mean"]), f"balanced accuracy, Logistic Regression (AUC {F.f3(lr['roc_auc']['mean'])})", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26)
    K.picture(s, F.fig("f01_model_comparison_real.png"), 0.55, 3.1, 7.6, 3.75)
    ftext(s, 8.4, 3.15, 4.4, 3.7, [
        ("RF, XGBoost and HistGradientBoosting are within one standard deviation of each other; all beat Logistic Regression and the Decision Tree.", {"bullet": True, "space_after": 6}),
        (f"Without the four post-hoc fields (relationships, milestones, milestone ages) ROC-AUC falls from {allf['Random Forest']['roc_auc']['mean']:.2f} to {nop['Random Forest']['roc_auc']['mean']:.2f} (RF) and {nop['XGBoost']['roc_auc']['mean']:.2f} (XGBoost).", {"bullet": True, "space_after": 6}),
        (f"Shipped app on synthetic data: LR {app_lr['balanced_accuracy']['mean']:.3f}; not real-world evidence.", {"bullet": True}),
    ], size=13)
    notes[17] = "Headline: ensembles beat the linear model, balanced accuracy 0.746, ROC-AUC about 0.81. The leakage ablation says the real numbers are an upper bound for early-stage prediction."

    # ---- 18 confusion matrix and imbalance
    s = S[17]
    set_title(s, "Confusion Matrix & Class-Imbalance Study")
    wipe(s)
    K.picture(s, F.fig("f02_confusion_matrix_rf.png"), 0.55, 1.85, 4.8, 3.9)
    K.picture(s, F.fig("f06_imbalance_strategies.png"), 5.6, 1.85, 7.2, 3.9)
    ftext(s, 0.55, 5.85, 12.2, 1.05, [
        (f"Random Forest, mean per fold of 185 test rows: {cm_rf[0][0]} failures and {cm_rf[1][1]} successes right; {cm_rf[0][1]} failures called SUCCESS, {cm_rf[1][0]} successes called FAILURE. Logistic Regression: {cm_lr[0][0]} / {cm_lr[1][1]} right, {cm_lr[1][0]} successes missed.", {"bullet": True, "size": 13}),
        (f"Failure is the minority class here (35.4 %). Class weights and oversampling do not raise ROC-AUC (RF {imb['Random Forest']['none']['roc_auc']['mean']:.3f} none, {imb['Random Forest']['class_weight']['roc_auc']['mean']:.3f} weighted, {imb['Random Forest']['oversample']['roc_auc']['mean']:.3f} oversampled); they trade recall of the majority success class for precision.", {"bullet": True, "size": 13}),
    ], size=13)
    notes[18] = "Rows are the true class. The imbalance in this table is the reverse of the batch objective's assumption: success is the majority."

    # ---- 19 threshold + importance
    s = S[18]
    set_title(s, "Threshold Analysis & Feature Importance")
    wipe(s)
    K.picture(s, F.fig("f03_threshold_sweep.png"), 0.55, 1.85, 5.9, 3.7)
    K.picture(s, F.fig("f04_feature_importance.png"), 6.7, 1.85, 6.1, 3.7)
    ftext(s, 0.55, 5.65, 12.2, 1.3, [
        ("Logistic Regression threshold sweep: balanced accuracy is flat at about 0.71 between thresholds 0.30 and 0.50 (peak 0.714 at 0.45) while recall of success falls from 0.875 to 0.673 and precision rises from 0.781 to 0.830.", {"bullet": True, "size": 13}),
        ("Random Forest ranks relationships first (Gini 0.162, permutation 0.060 AUC drop), then milestones, last-milestone age and funding amount; Logistic Regression weights are led by state, top-500 and industry dummies, so the two rankings disagree.", {"bullet": True, "size": 13}),
    ], size=13)
    notes[19] = "Threshold and feature analysis. The default 0.5 is not special; the choice moves recall and precision, not balanced accuracy."

    # ---- 20 performance and testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{p_pred['median_ms']:.0f} ms", f"POST /predict, k=10, median (p95 {p_pred['p95_ms']:.0f} ms)", K.FBLUE), (f"{p_pred50['median_ms']:.0f} ms", f"POST /predict, k=50, median (p95 {p_pred50['p95_ms']:.0f} ms)", K.GREEN),
                                   (f"{LAT['health']['median_ms']:.1f} ms", f"GET /api/health median; insights {LAT['insights']['median_ms']:.1f} ms", K.ORANGE), (f"{R['runtime_seconds']:.0f} s", "whole real-data study on CPU", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26)
    rows = [["Test layer", "What it covers", "Result"],
            ["Unit tests (pytest)", "health, dataset, preprocessing, embeddings and fusion, retrieval, predict, pages", f"{F.N_UNIT} / {F.N_UNIT} passed"],
            ["Smoke test", "running server: endpoints, login, token, predict, invalid input, pages", f"{F.N_SMOKE} / {F.N_SMOKE} checks"],
            ["API edge tests", "limits, k whitelist, empty body, unknown categories, forms, auth, tampered token", f"{F.N_EDGE} / {F.N_EDGE} checks"],
            ["Total", "automated checks", f"{F.N_TESTS}"]]
    ftable(s, 0.55, 3.2, [2.4, 7.4, 2.4], rows, size=13, row_h=0.46)
    ftext(s, 0.55, 5.65, 12.2, 1.2, [
        ("Measured on 127.0.0.1 with a warm model on an Intel Core i5-12500 (CPU only); the first request pays the model-load cost.", {"bullet": True, "size": 13}),
        ("Not tested: load or stress, multi-user use, a security audit, a user study, React component tests.", {"bullet": True, "size": 13}),
    ], size=13)
    notes[20] = "Inference is fast enough for interactive use; the MiniLM embedding dominates the predict latency. The tests cover the library and the API, not the browser."

    # ---- 21 comparison with the base papers
    s = S[20]
    set_title(s, "Comparison with Base Papers (context only – different data)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "Maarouf et al. [2]", f"This project ({F.PRODUCT})"],
            ["Data", "Crunchbase / Kaggle, size not stated", "4,034 Crunchbase startups, 35.15 % success", f"Real table {F.REAL_ROWS} rows (64.6 % success); app on synthetic {F.SYN_ROWS:,}"],
            ["Protocol", "70:30 split, ADASYN", "5 x 80/20 stratified, balanced accuracy", "5 x 80/20 (and 5 x 70:30), balanced accuracy, ROC-AUC"],
            ["Models", "LR, RF, SVM, Gradient Boosting", "kNN-ICL (LLM) vs LR, RF, XGBoost", "LR, DT, RF, XGBoost, HistGB + balanced retrieval"],
            ["Reported", "Accuracy 85-92 % (best ensembles)", "Balanced acc. 0.713 vs 0.631 / 0.634 / 0.618", f"RF / XGB balanced acc. {F.f3(rf['balanced_accuracy']['mean'])}, accuracy {F.f3(rf['accuracy']['mean'])} / {F.f3(xgb['accuracy']['mean'])}; RF {F.f3(p70['Random Forest']['accuracy']['mean'])} on 70:30"],
            ["Compute", "Flask, local or cloud", "LLM prompts, no training", "CPU only, local; predict median 148-199 ms"]]
    ftable(s, 0.55, 1.85, [1.5, 3.1, 3.5, 4.1], rows, size=12, row_h=0.66)
    ftext(s, 0.55, 5.95, 12.2, 1.0, "The ensemble ranking of [1] is reproduced, its 85-92 % accuracy is not (best 77.5 % on 80/20), but its data and table are unspecified. Datasets differ from [2] (their descriptions are text; the real table here has none) and from the temporal Crunchbase panel of [3] (AUROC 0.872 for exits), so the figures are context, not a ranking.", size=12, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that the datasets and protocols differ; no superiority or inferiority claim is possible from these numbers."

    # ---- 22-24 output screens
    shots = F.SHOTS
    s = S[21]
    set_title(s, "Output – Dashboard & Analyze Form")
    wipe(s)
    K.picture(s, os.path.join(shots, "04_dashboard.png"), 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, os.path.join(shots, "06_analyze_form_filled.png"), 6.8, 1.85, 6.0, 4.2, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.8, f"Left: the dashboard of the running app after demo sign-in: {F.SYN_ROWS:,} historical startups ({F.SYN_SUCCESS} / {F.SYN_FAILURE}), the best model and a data note that the synthetic fallback is in use. Right: the Analyze form filled with a real submission (software startup, CA, USD 2.50 M over 2 rounds, 6 relationships, 3 milestones, VC and round A).", size=12, color=K.FNAVY)
    notes[22] = "Real screenshots from the running app. The data note on the dashboard tells the user that the synthetic fallback is in use."
    s = S[22]
    set_title(s, "Output – Prediction Result")
    wipe(s)
    top = crop(os.path.join(shots, "07_result.png"), (0, 0, 1440, 700), "result_top.png", tmp)
    bot = crop(os.path.join(shots, "07_result.png"), (0, 700, 1440, 1818), "result_bottom.png", tmp)
    K.picture(s, top, 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, bot, 6.8, 1.85, 6.0, 4.2, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.8, "Left (top of the page): SUCCESS at 83.3 % confidence from Logistic Regression, labelled Lower risk, the signed key features and the 25 / 25 evidence split of k=50. Right (rest of the page): the comparable startups (top similarity 0.580) and the submitted inputs. The prediction is an analytical estimate from a model trained on synthetic data, not a guarantee.", size=12, color=K.FNAVY)
    notes[23] = "One real prediction: outcome, drivers, evidence split and comparable startups. Say again that the model behind this screen was trained on the synthetic fallback."
    s = S[23]
    set_title(s, "Output – Model Performance & Startup Insights")
    wipe(s)
    perf = crop(os.path.join(shots, "08_performance.png"), (0, 0, 1440, 1250), "perf_top.png", tmp)
    ins = crop(os.path.join(shots, "09_insights.png"), (0, 0, 1440, 1250), "insights_top.png", tmp)
    K.picture(s, perf, 0.55, 1.85, 6.0, 4.2, border=True)
    K.picture(s, ins, 6.8, 1.85, 6.0, 4.2, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.8, "Left: the model comparison table, balanced-accuracy and F1 charts, precision-recall plot and confusion matrix of the served model (synthetic data: balanced accuracy 0.749 ± 0.025). Right: dataset insights, success rate by industry, median funding by outcome and correlations with success (top of each page).", size=12, color=K.FNAVY)
    notes[24] = "Analytics screens read their numbers from the stored evaluation of the served model, so they show synthetic-data statistics."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("**Delivered:** a working startup-outcome predictor: preprocessing, MiniLM text embedding, fused retrieval of comparable startups, four compared classifiers, explanations, a FastAPI service and a React app.", {"bullet": True, "space_after": 9}),
        (f"**Measured (real table, structured features):** RF / XGBoost balanced accuracy {F.f3(rf['balanced_accuracy']['mean'])}, ROC-AUC {F.f3(rf['roc_auc']['mean'])} / {F.f3(xgb['roc_auc']['mean'])}; ensembles beat Logistic Regression; {F.N_TESTS} automated checks pass; predict median 148-199 ms.", {"bullet": True, "space_after": 9}),
        ("**Objectives:** 3, 4, 5 and 6 met (6 with the caveat that ROC-AUC was not recorded for the shipped fused model); 1 and 2 only partly: no founder, team-experience or market data exist in either dataset.", {"bullet": True, "space_after": 9}),
        ("**Limits:** the shipped model is trained on synthetic data; the text channel is untested on real data; post-hoc fields inflate AUC by about 0.1; no temporal split; 922 rows from one US table; not investment advice.", {"bullet": True}),
    ], size=18)
    notes[25] = "Be explicit: the system works; the real-data numbers are an upper bound and the 85-92 % level of the base paper was not reproduced."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("Retrain the served model on the real table and re-run the study with a temporal split once dated records are available.", {"bullet": True, "space_after": 12}),
        ("Collect founder-experience, team and market-context fields and real business descriptions to test the text half of the fusion properly.", {"bullet": True, "space_after": 12}),
        ("Calibrate the probabilities, add SHAP-style explanations and report PR-AUC with confidence intervals and significance tests.", {"bullet": True, "space_after": 12}),
        ("Add HistGradientBoosting and the retrieval-only classifier to the app, with a user-selectable decision threshold.", {"bullet": True, "space_after": 12}),
        ("Replace the demo sign-in with real accounts, add load tests and a user study with investors or founders.", {"bullet": True}),
    ], size=19)
    notes[26] = "Concrete next steps; the first three follow directly from the findings."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [1, 2, 3, 7, 9, 16, 17, 18, 19, 21, 22, 25, 29, 32]
    ftext(s, 0.55, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 3}) for n in ids[:7]], size=10.5)
    ftext(s, 6.8, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 3}) for n in ids[7:]], size=10.5)
    notes[27] = "Selected references with their report numbers; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


def wait_for_inputs(timeout=1800):
    t0 = time.time()
    while True:
        missing = [p for p in NEEDED if not os.path.exists(p)]
        if not missing:
            return
        if time.time() - t0 > timeout:
            raise SystemExit("missing inputs: " + ", ".join(os.path.basename(m) for m in missing))
        time.sleep(15)


if __name__ == "__main__":
    wait_for_inputs()
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
