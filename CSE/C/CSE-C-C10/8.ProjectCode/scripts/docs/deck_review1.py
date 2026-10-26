"""7.Review_1_PPT.pptx: 16 slides, problem identification and research (VenturePulse). Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import NEED, OBJECTIVES, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "; ".join(F.KEYWORDS)
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 28)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 18)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}

    # ---- 1 title
    title_slide(S[0], "Structured funding signals and a short business description are fused to estimate startup outcome, with comparable historical startups and feature drivers shown as evidence",
                "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: VenturePulse reads a startup's funding, team and sector indicators plus a short description and returns SUCCESS or FAILURE "
                "with a confidence, the comparable historical startups and the features behind the estimate.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: an evidence-based early signal for startups")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Most startups fail, and investors and founders judge prospects largely on experience and intuition. Success is rare in the data, so naive models under-predict it, and a bare prediction gives a founder nothing to act on.", K.ORANGE, "!"),
        ("Our idea", "Learn from historical startups: structured signals (funding, rounds, milestones, relationships, sector, location) plus a text embedding of the business description. Compare Logistic Regression, Decision Tree, Random Forest and boosting, and retrieve comparable startups.", K.BLUE, "i"),
        ("The deliverable", "VenturePulse, a local web application: SUCCESS or FAILURE with confidence, the top features behind it, a table of comparable historical startups, and pages for model comparison and data insights.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: the problem, the idea of combining structured signals, text and retrieval of comparable startups, and the deliverable: a complete application, not only a model script."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Logistic Regression, Random Forest, SVM and Gradient Boosting on funding history, team size, industry and location. Mean/mode imputation, encoding, standard scaling and ADASYN oversampling; 70:30 split; Flask web application.", K.BLUE, "M"),
        ("Reported results", "Ensembles (Random Forest, Gradient Boosting) perform best, with overall accuracy between 85 and 92 %. SVM works well but is slower; Logistic Regression is the baseline. No dataset size or per-model table is given.", K.GREEN, "R"),
        ("What we take, what differs", "Taken: the model family, imbalance handling, several metrics, a web front end. Added: text embedding with comparable-startup retrieval, feature-level explanations, ROC-AUC, and a stated leakage check. The paper itself names look-ahead bias and missing text data as open problems.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The base paper reports that ensembles predict startup success best, at 85 to 92 percent accuracy, but gives no dataset size or table, so its figures cannot be reproduced exactly. "
                "Our project keeps its model family and adds retrieval-based evidence, explanations and a measured evaluation.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result / role"],
            ["Panjwani et al. 2026 [1] (base paper)", "LR, RF, SVM, Gradient Boosting; ADASYN; Flask app", "Crunchbase / Kaggle data, size not stated", "Ensembles best, accuracy 85-92 %"],
            ["Maarouf et al. 2026 [2]", "kNN retrieval of similar startups as examples for an LLM; compared with LR, RF, XGBoost", "4,034 Crunchbase startups, 35.15 % successful", "Balanced accuracy 0.713 vs 0.631 (LR), 0.634 (RF), 0.618 (XGBoost)"],
            ["Mashhadi et al. 2026 [3]", "Interpretable tree ensembles, inverse weights, SMOTE-NC, SHAP, temporal splits", "Crunchbase firm-quarter panel 2010-2023 + patents", "AUROC 0.921 (patents), 0.817 (funding), 0.872 (exit)"],
            ["Bidgoli et al. 2024 [7]", "RF, Gradient Boosting, SVM, LR compared (as summarised by [1])", "Startup data", "Ensembles above 80 % accuracy"],
            ["Zbikowski, Antosiuk 2021 [9]", "Bias-free success prediction from Crunchbase (as summarised by [1])", "Crunchbase", "Warns that look-ahead bias inflates accuracy"],
            ["Gius 2025 [16]", "Judge disagreement in venture competitions as a predictor", "67 venture competitions", "More disagreement, higher success chance"]]
    K.table(s, 0.6, 1.4, [2.7, 4.1, 2.8, 2.53], rows, font_size=12, row_h=0.76, header_size=13)
    notes[5] = ("The pattern: ensembles are the usual winners on structured startup data; the 2026 retrieval paper shows that giving a model similar past startups helps in small-data settings; "
                "the interpretable-ML paper insists on leakage-safe splits. Works [4]-[15] are known here only through the base paper's reference list, so no finding beyond that summary is claimed. "
                "The batch survey also lists two further items whose authors we could not verify; they appear in the report's reference list with that caveat.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: retrieval helps on small data")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 7.6, 5.3,
                     ["Logistic\nRegression", "XGBoost", "Random\nForest", "kNN-ICL\n(k = 50)", "Fused LLM\n(20,172 train)"],
                     [("Balanced accuracy", [0.631, 0.618, 0.634, 0.713, 0.743])], colors=[K.BLUE], number_format="0.000", font_size=12, max_value=0.9)
    for idx in (3, 4):
        pt = gf.chart.plots[0].series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = K.ORANGE
    K.card(s, 8.5, 1.45, 4.23, 2.6, "Takeaway", "On 4,034 Crunchbase startups with 5 stratified 80/20 splits, showing a model the most similar past startups (kNN-ICL) beats supervised baselines by about 8 points of balanced accuracy [2].", K.GREEN, "↑", body_size=13)
    K.card(s, 8.5, 4.2, 4.23, 2.5, "Caveat", "All bars are values reported in [2] on its own data. They are not our results, and the base paper's 85-92 % is accuracy on an unstated dataset, so no ranking against it is implied.", K.ORANGE, "!", body_size=13)
    notes[6] = ("Values are quoted from the 2026 retrieval paper: balanced accuracy 0.631 LR, 0.634 RF, 0.618 XGBoost, 0.713 kNN-ICL at 50 shots, 0.743 for the fused LLM trained on 20,172 startups. "
                "We reuse its protocol (five repeated stratified 80/20 splits, balanced accuracy) and its retrieval idea, without an LLM.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Founder pitch / data room", K.NAVY), ("Investor judgement", K.NAVY), ("Due diligence", K.NAVY), ("Funding decision", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Subjective assessment", "Investors and founders rely on experience, intuition and personal networks. The judgement is hard to repeat, differs between people and is open to bias.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Published ML predictors", "Classifiers on funding, team and sector features reach good scores, but are often reported without dataset sizes, leakage checks or reasons, and give a founder no comparable cases.", K.ORANGE, "2", body_size=17)
    notes[7] = "Today's path is judgement-led. Machine-learning predictors exist, but their outputs are often unexplained and their evaluation protocols are not always clear."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Subjective, inconsistent", "Startup success is judged subjectively, leading to inconsistent investment decisions.", K.ORANGE, "1"),
        ("Look-ahead bias", "Failure causes become obvious only after capital is committed; models trained on information recorded after the outcome look better than they are [1], [9].", K.BLUE, "2"),
        ("Rare success, little data", "Success is rare in many datasets, so naive models under-predict the positive class, and early-stage data is scarce.", K.GREEN, "3"),
        ("No stated drivers", "A prediction without drivers or comparable cases gives a founder nothing to act on.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=16)
    notes[8] = "Four limitations from the project's problem statement and the survey. Each maps to a design decision: balanced handling of the rare class, a stated leakage check, retrieval evidence and feature drivers."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Feature engineering", "Systematic features around funding velocity and team composition are missing from the surveyed work."),
            ("Rare success class", "Explicit, compared handling of class imbalance (weights, oversampling, thresholds) is not shown with numbers."),
            ("Evidence with answer", "Predictions are rarely shown together with comparable historical startups and the features that drive them."),
            ("Text and structure", "The base paper lists restricted use of text and poor integration of data sources as open gaps."),
            ("Honest evaluation", "Dataset size, leakage and protocol are often unstated; scores need baselines and stated limits.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=17, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as: systematic feature engineering around funding velocity and team composition, with explicit handling of the rare-success class. "
                "We add three more from reading the papers: evidence with the answer, text plus structure, and honest evaluation.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, PROBLEM_STATEMENT, size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Funding, rounds, milestones, relationships, investor flags, state, industry and an optional short description"),
                                 ("Output", "SUCCESS or FAILURE, probabilities, risk label, comparable startups, top feature drivers"),
                                 ("Record", "Model comparison, confusion matrix, feature importance, dataset insights")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.55)
        K.text(s, x + 0.25, 4.6, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.05, 3.4, 0.95, b_, size=14, color=K.TXT)
    K.text(s, 0.6, 6.25, 12.13, 0.5, "Need: " + NEED, size=14, color=K.TXT, italic=True)
    notes[10] = "Read the statement slowly. Input is structured indicators plus optional text, output is a two-way outcome with evidence, and the system keeps measured statistics that can be inspected."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "The six objectives from the proposal. The later reviews report which were met; collecting founder and market data (objectives 1 and 2) is the part limited by the public data available."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "De-duplicate, map acquired/IPO to SUCCESS and closed to FAILURE, impute, scale, one-hot encode and add log funding: 37 structured features.", K.BLUE, "1"),
        ("Represent", "Embed the description with all-MiniLM-L6-v2 (384 dimensions, company name masked) and fuse with the structured vector, weight 0.5 each.", K.GREEN, "2"),
        ("Predict", "Logistic Regression, Decision Tree, Random Forest and XGBoost with class weights; the best by balanced accuracy over 5 stratified 80/20 splits serves.", K.ORANGE, "3"),
        ("Explain", "Balanced top-k cosine retrieval of comparable startups (training rows only), top-5 feature drivers and a risk label.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "V", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Application", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "VenturePulse: React interface over a FastAPI service, running locally with no external prediction service. Analyse a startup, compare models, and browse data insights.", size=15, color=K.TXT)
    notes[12] = "Four stages from raw indicators to an explained prediction, and the application around them. The model comparison is kept so the choice of classifier is measured, not assumed."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 1.4, 1.3, 10.5, 5.3)
    K.text(s, 0.6, 6.62, 12.13, 0.3, "Browser UI, FastAPI service, the ML modules (preprocessing, embeddings, retrieval, models) and the saved model artefacts", size=12, color=K.TXT, align="c", italic=True)
    notes[13] = "Follow the diagram: the browser sends the startup profile to the FastAPI service, which calls the prediction and retrieval modules; the fitted artefacts are files on disk and there is no database."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working predictor", "A trained, compared set of four classifiers with the best one serving predictions.", K.BLUE, "1"),
            ("Measured comparison", "Accuracy, precision, recall, F1 and ROC-AUC with mean and spread over repeated splits.", K.GREEN, "2"),
            ("Imbalance handled", "Class weighting compared with oversampling and decision thresholds.", K.ORANGE, "3"),
            ("Drivers named", "Feature importance and per-prediction contributions.", K.PURPLE, "4"),
            ("Evidence shown", "Comparable historical startups with similarity, outcome and funding.", K.BLUE, "5"),
            ("Honest limits", "Data source, leakage-like fields and untested parts stated plainly.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=16)
    notes[14] = "These are targets, not results. We report whatever the held-out data shows, including a lower accuracy than the published 85-92 % if that is what the data gives."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 2, 3, 7, 9, 11, 16, 18, 22]
    K.text(s, 0.6, 1.35, 12.13, 5.5, [(f"[{n}] {REFS[n - 1]}", {"space_after": 5}) for n in idx], size=14, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = "Likely questions: why a classifier plus retrieval; why logistic regression may serve; where the data comes from; how leakage is checked; what success and failure mean in the data."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
