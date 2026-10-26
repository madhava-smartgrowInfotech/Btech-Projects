"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress (VenturePulse). Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


# ------------------------------------------------------------------ local helpers (self-contained)
def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  \u00b7  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "; ".join(F.KEYWORDS)
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 28)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 18)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS), 14)
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  \u00b7  Batch {F.BATCH_NO} ({F.SECTION})  \u00b7  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}
    M = F.R["models"]
    rf, xg, lr, dt, hg = M["Random Forest"], M["XGBoost"], M["Logistic Regression"], M["Decision Tree"], M["HistGradientBoosting"]
    ab = F.X["without_posthoc_features"]
    lat = F.LAT

    # ---- 1 title
    title_slide(S[0], "Design, methodology and progress: classifiers, comparable-startup retrieval and feature drivers behind a FastAPI service and a React interface",
                "REVIEW 2  \u00b7  DESIGN, METHODOLOGY AND PROGRESS")
    notes[1] = "This review shows how the system is designed and built, what is implemented, the first measured results on real data, and the open items."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap and architecture"), ("02", "Methodology and dataset"), ("03", "Preprocessing and model"), ("04", "Training and decision logic"),
               ("05", "Data flow, database and UML"), ("06", "Technology stack"), ("07", "Implementation progress and results"),
               ("08", "Output screens and testing"), ("09", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Startup success is judged subjectively; failure causes show up only after capital is committed; success is rare or imbalanced in data; and a bare prediction names no drivers.", K.ORANGE, "!"),
        ("Proposed system", "37 structured features plus a MiniLM text embedding are fused. Four classifiers are compared; the best serves a prediction, and a balanced retrieval returns comparable historical startups with the top feature drivers.", K.BLUE, "i"),
        ("Base paper link", "Panjwani et al. (2026) report that ensembles predict best (85-92 % accuracy, dataset not stated). We keep the model family and add retrieval, explanations and a measured, leakage-aware evaluation; their numbers are not reproducible from the paper.", K.GREEN, "B"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, 4.9, t, b_, col, gl, body_size=18)
    notes[3] = "One-minute recap of Review 1: the gap in feature engineering and rare-class handling, and how the base paper relates (model family and motivation only)."

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.6, 5.6)
    K.text(s, 8.5, 1.5, 4.2, 5.3, [
        ("**Three tiers**", {"space_after": 8}),
        ("Presentation: React 18 app (VenturePulse, port 5410) with landing, sign-in, dashboard, analyse, performance and insights screens.", {"bullet": True}),
        ("Application: FastAPI (port 8410), JSON endpoints plus the original Jinja pages; one model held in memory.", {"bullet": True}),
        ("Domain: app/ml (dataset, preprocessing, embeddings, retrieval, models, evaluation) and app/services (store, predictor, insights).", {"bullet": True}),
        ("Storage: artefact files in models/; no SQL database.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[4] = "The FastAPI layer calls the same predictor service for the JSON API and the original server-rendered pages, so both use one model and one retrieval index."

    # ---- 5 methodology workflow
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.2, [("Startup data", K.NAVY), ("Clean + encode", K.PURPLE), ("Embed text", K.PURPLE), ("Fuse + index", K.GREEN),
                                      ("Train 4 models", K.ORANGE), ("Select + save", K.BLUE), ("Predict + explain", K.NAVY)], size=13, gap=0.22)
    for i, (t, b_) in enumerate([("Leakage-safe", "Preprocessing is fitted on the training part of each split, and retrieval searches training rows only, so a test startup never retrieves itself."),
                                 ("Class weighting", "Balanced class weights in Logistic Regression, Decision Tree and Random Forest; imbalance strategies are compared on the real table."),
                                 ("Reproducible", "scripts/run_experiments.py and extra_experiments.py write JSON; every number in the documents is read from it.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.1, 3.85, 3.5)
        K.text(s, x + 0.25, 3.25, 3.4, 0.5, t, size=20, color=K.INK, bold=True)
        K.text(s, x + 0.25, 3.9, 3.4, 2.6, b_, size=17, color=K.TXT)
    notes[5] = "The same preprocessing objects are used for training and serving. The experiment scripts reuse the product's own preprocessing and model definitions."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Dataset: two tables, kept apart")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{F.SYN_ROWS:,}", "synthetic fallback rows (shipped app)", K.PURPLE), (f"{F.REAL_ROWS}", "real public rows after de-duplication", K.BLUE),
                                   (f"{F.REAL_SUCCESS} / {F.REAL_FAILURE}", "real: success / failure", K.GREEN), (f"{F.SYN_SUCCESS} / {F.SYN_FAILURE}", "synthetic: success / failure", K.ORANGE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    rows = [["", "Shipped synthetic table", "Real public table"],
            ["Source", "Generated with the same schema (public CSV was not available)", "Kaggle 'Startup Success Prediction' (public domain), US startups"],
            ["Outcome", "acquired = SUCCESS (42.5 %), closed = FAILURE", f"acquired = SUCCESS ({F.pct(F.R['dataset']['success_share'])}), closed = FAILURE"],
            ["Text", "Template descriptions that differ by class (they leak the label)", "No description column"],
            ["Used for", "Serving the app, screenshots", "All real measurements"]]
    K.table(s, 0.6, 2.95, [1.8, 5.5, 4.83], rows, font_size=14, row_h=0.62, first_bold=True)
    K.text(s, 0.6, 6.15, 12.13, 0.8, "In the real table SUCCESS is the majority class, the reverse of what the objectives assumed. Features that are recorded around or after the outcome (relationships, milestones) are flagged as leakage-like and tested later.", size=14, color=K.TXT, italic=True)
    notes[6] = ("The shipped app runs on a synthetic table of 1,400 rows because the public CSV needs a login; the identical 923-row public table was obtained separately and, after the app's own de-duplication, has 922 rows. "
                "All real results come from that table. Its lack of descriptions means the text half of the fusion cannot be tested on it.")

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Data preprocessing and feature extraction")
    wipe(s)
    rows = [["Step", "What the code does", "Result"],
            ["Clean rows", "Drop duplicates; map acquired/IPO to SUCCESS (1) and closed to FAILURE (0); coerce numbers", "Valid labelled rows"],
            ["Missing values", "Median imputation (numeric), most-frequent (flags), constant 'Unknown' (categorical)", "No NaN reaches a model"],
            ["Scale and encode", "StandardScaler on 9 numeric fields + log1p funding; one-hot state (10) and industry (10)", "37 structured features"],
            ["Text", "Clean description, mask company name, MiniLM embedding (384-d), L2-normalised", "Text vector"],
            ["Fuse", "x = [0.5 s, 0.5 t], then L2-normalise so cosine similarity is a dot product", "421-d fused vector"]]
    K.table(s, 0.6, 1.4, [2.0, 7.2, 2.93], rows, font_size=13, row_h=0.72, first_bold=True)
    K.text(s, 0.6, 5.9, 12.13, 1.0, [("**Seven flags and nine numeric fields:** has_VC, has_angel, has_roundA-D, is_top500; funding total and rounds, ages at first and last funding and milestone, relationships, milestones, average participants. "
                                      "Funding velocity is only approximated by the ages and the number of rounds; no team-experience field exists in either table.", {})], size=13, color=K.INK)
    notes[7] = "All steps live in app/ml/preprocessing.py, dataset.py and retrieval.py. The preprocessor is saved with the model so serving uses the same transforms."

    # ---- 8 model
    s = S[7]
    set_title(s, "Model: four classifiers on fused features")
    wipe(s)
    for i, (v, l, c) in enumerate([("4 + 1", "classifiers (+ HistGradientBoosting in the study)", K.BLUE), ("37", "structured features", K.GREEN),
                                   ("384", "text embedding dimensions", K.ORANGE), ("k = 50", "default comparable startups (10/20/30/50)", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=12)
    rows = [["Model", "Configuration (app/ml/models.py)"],
            ["Logistic Regression", "lbfgs, class_weight balanced, C = 1, max_iter 1000. Serves predictions in the app (best balanced accuracy on the shipped data)"],
            ["Decision Tree", "max_depth 8, min_samples_split 10, min_samples_leaf 5, balanced"],
            ["Random Forest", "300 trees, min_samples_leaf 2, balanced"],
            ["XGBoost", "400 trees, depth 6, learning rate 0.05, subsample 0.8, colsample 0.8"]]
    K.table(s, 0.6, 3.0, [2.6, 9.53], rows, font_size=14, row_h=0.66, first_bold=True)
    K.text(s, 0.6, 6.45, 12.13, 0.5, "Selection rule: highest mean balanced accuracy over five stratified 80/20 splits (seed 42).", size=14, color=K.TXT, italic=True)
    notes[8] = "The best model on the shipped synthetic table is Logistic Regression; on the real table the tree ensembles are ahead, which is reported in the results slide."

    # ---- 9 training configuration
    s = S[8]
    set_title(s, "Training and evaluation configuration")
    wipe(s)
    rows = [["Setting", "Value"],
            ["Protocol", "5 x StratifiedShuffleSplit, 80 % train / 20 % test, seed 42; preprocessing fitted on train only; mean \u00b1 std"],
            ["Metrics", "Accuracy, balanced accuracy (primary), precision, recall, F1, ROC-AUC, confusion matrix"],
            ["Studies", "Representation (structured / text / combined), retrieval k, imbalance, threshold, calibration, forest size, 70:30, leakage ablation"],
            ["Real-data run", f"{F.REAL_ROWS} rows, structured features, {F.R['runtime_seconds']:.0f} s on CPU for the full study"],
            ["Hardware", F.HARDWARE]]
    K.table(s, 0.6, 1.3, [2.2, 9.93], rows, font_size=12, row_h=0.43, first_bold=True)
    K.picture(s, F.diag("11_evaluation_protocol.png"), 0.6, 4.0, 7.6, 2.95)
    K.text(s, 8.5, 4.2, 4.2, 2.7,
           "The protocol follows the 2026 retrieval paper [2] (five repeated stratified 80/20 splits, balanced accuracy) and also runs the base paper's 70:30 split with five seeds, so both can be compared with the same data.", size=15, color=K.TXT)
    notes[9] = "Training is seconds, not hours: the whole real-data study of five models and all ablations takes about two minutes on the CPU."

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic: prediction, retrieval and risk label")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.6, 1.3, 5.2, 5.65)
    K.text(s, 6.1, 1.5, 6.6, 5.4, [
        ("**Predict.** The model returns P(success); the outcome is the larger of P(success) and P(failure), and the confidence is that maximum.", {"bullet": True, "space_after": 8}),
        ("**Retrieve.** Cosine similarity to training rows only; half the k slots go to successful and half to failed startups, so rare cases are not crowded out.", {"bullet": True, "space_after": 8}),
        ("**Explain.** Top five signed coefficient or importance terms, plus counts and average similarity of the retrieved cases.", {"bullet": True, "space_after": 8}),
        ("**Risk label.** 'Lower risk' if SUCCESS with confidence at least 0.70; 'Moderate risk' if confidence at least 0.55; else 'Higher uncertainty'.", {"bullet": True, "space_after": 8}),
        ("**Caveat.** Logistic Regression with balanced class weights under-states P(success) on the real table (calibration slide of the final review).", {"bullet": True}),
    ], size=14, color=K.INK)
    notes[10] = "Left: the core algorithm flowchart. The risk labels are rules on the displayed confidence; they are not validated probabilities."

    # ---- 11 data flow and modules
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 7.3, 5.6)
    rows = [["Module", "Role"], ["ml/dataset.py", "load CSV or synthetic fallback"], ["ml/preprocessing.py", "clean, impute, scale, encode"], ["ml/embeddings.py", "MiniLM text vectors"],
            ["ml/retrieval.py", "fusion, balanced top-k search"], ["ml/models.py", "four classifiers"], ["ml/evaluation.py", "5 x 80/20 study, charts"],
            ["services/predictor.py", "predict, explain, risk label"], ["main.py, api_extra.py", "routes, validation, demo sign-in"]]
    K.table(s, 8.2, 1.5, [2.3, 2.23], rows, font_size=11, row_h=0.52)
    notes[11] = "Level-1 data flow: predict, model statistics and insights, evaluation/training, with the artefact files as the store."

    # ---- 12 database and UML
    s = S[11]
    set_title(s, "Database design and UML")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.6, 1.35, 5.3, 4.9)
    K.picture(s, F.diag("06_uml_class_diagram.png"), 6.1, 1.3, 6.6, 5.1)
    K.text(s, 0.6, 6.5, 12.13, 0.45, "There is no SQL database: models and the training frame are files in models/, state is held in memory. The ER diagram models these artefacts.", size=12, color=K.TXT, align="c", italic=True)
    notes[12] = "The project stores files, not rows. The class diagram shows the real modules and the request model."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    rows = [["Layer", "Technology"], ["Machine learning", "Python 3.11, scikit-learn 1.5.2, XGBoost 2.1.3, pandas 2.2.3, NumPy, joblib"],
            ["Text embeddings", "sentence-transformers 3.3.1 (all-MiniLM-L6-v2), torch 2.5.1 CPU"], ["API", "FastAPI 0.115.6, uvicorn 0.34.0, Pydantic validation (port 8410)"],
            ["Front end", "React 18, Vite 6, Tailwind CSS 4, Recharts (port 5410); original Jinja pages kept"],
            ["Testing", "pytest, API smoke and edge scripts, Playwright for screenshots"], ["Hardware", F.HARDWARE]]
    K.table(s, 0.6, 1.5, [3.0, 9.13], rows, font_size=17, row_h=0.75, first_bold=True)
    notes[13] = "Everything is open source and starts with run.bat. No external prediction service or API key is used."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Preprocessing, embeddings, fusion, retrieval", "Done", "app/ml, 7 pytest tests pass"],
            ["Four classifiers + evaluation study", "Done", "5 x 80/20, metrics JSON, charts"],
            ["Real-data experiments (922 rows)", "Done", "run_experiments.py, extra_experiments.py"],
            ["FastAPI service and React interface", "Done", f"{F.N_SMOKE} smoke + {F.N_EDGE} API edge checks pass"],
            ["Founder / team-experience features", "Not done", "no such field in either table"],
            ["Temporal split, SHAP, calibration fix", "Not done", "table has no dates; listed as future work"],
            ["Text channel on real data, user study, load test", "Not done", "stated as limits"]]
    K.table(s, 0.6, 1.5, [4.9, 1.6, 5.63], rows, font_size=15, row_h=0.64)
    notes[14] = "State of the code base. The honest gaps are founder data, temporal validation and anything involving real descriptions."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: real table, 5 x 80/20 splits")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.f3(xg["balanced_accuracy"]["mean"]), "XGBoost balanced accuracy", K.BLUE), (F.f3(rf["roc_auc"]["mean"]), "Random Forest ROC-AUC", K.GREEN),
                                   (F.f3(rf["f1"]["mean"]), "Random Forest F1", K.ORANGE), (F.f3(lr["balanced_accuracy"]["mean"]), "Logistic Regression balanced accuracy", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=11)
    K.picture(s, F.fig("f01_model_comparison_real.png"), 0.6, 2.9, 7.9, 4.05)
    K.card(s, 8.7, 2.95, 4.03, 3.95, "Reading",
           f"Random Forest {F.pm(rf)}, XGBoost {F.pm(xg)}, Logistic Regression {F.pm(lr)} (balanced accuracy). Ensembles lead, within one std of each other. "
           f"Accuracy {F.f3(xg['accuracy']['mean'])} is below the base paper's 85-92 %. Without relationships and milestones, ROC-AUC drops to about {F.f3(ab['Random Forest']['roc_auc']['mean'])}.", K.BLUE, "i", body_size=14)
    notes[15] = ("Real 922-row table: ensembles reproduce the base paper's ranking but not its accuracy level. Dropping the four post-outcome-style fields lowers ROC-AUC from 0.82 to 0.72 (Random Forest), "
                 "so these numbers are an upper bound for early-stage prediction. The shipped app's own 0.749 is on synthetic data and is not a real-world figure.")

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    shots = [("07_result.png", "Result for a real submission: SUCCESS 83.3 %, Lower risk, 50 comparable startups", 0.6, 4.2),
             ("08_performance.png", "Model performance screen", 5.0, 3.4),
             ("09_insights.png", "Startup insights screen", 8.6, 4.13)]
    for fn, cap, x, w in shots:
        K.picture(s, F.shot(fn), x, 1.25, w, 4.85, border=True)
        K.text(s, x, 6.2, w, 0.7, cap, size=12, color=K.TXT, align="c")
    notes[16] = ("Real screenshots of the running app. The submission is a software startup in California with USD 2.50 M over two rounds; the screens are served by the model trained on the synthetic fallback table, "
                 "so the statistics shown on the performance screen are the synthetic ones.")

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([(str(F.N_TESTS), "automated checks passed", K.GREEN), (str(F.N_SMOKE), "API smoke checks", K.BLUE),
                                   (str(F.N_EDGE), "API edge and validation cases", K.ORANGE), (str(F.N_UNIT), "pytest unit tests", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    rows = [["Observation", "Action"],
            [f"POST /predict median {lat['predict_k10']['median_ms']:.0f} ms (k = 10) and {lat['predict_k50']['median_ms']:.0f} ms (k = 50); GET endpoints {lat['health']['median_ms']:.1f}-{lat['insights']['median_ms']:.1f} ms", "Measured on 127.0.0.1, CPU only"],
            ["Out-of-range funding, rounds, k outside {10,20,30,50} and over-long descriptions return 422", "Pydantic limits; 29 cases pass"],
            ["A demo token with one signature character changed is rejected (401)", "A first check that only appended text was corrected"]]
    K.table(s, 0.6, 3.1, [7.4, 4.73], rows, font_size=16, row_h=1.0)
    notes[17] = f"{F.N_UNIT} pytest tests, {F.N_SMOKE} smoke checks and {F.N_EDGE} API edge cases. No load, stress or penetration test was run."

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["Public CSV needs a login; the shipped model trained on synthetic data", "Obtained the identical public table, ran all real experiments on it, labelled app screens as synthetic-trained"],
            ["Synthetic text leaks the label, so text looks useful", "Reported text results as synthetic only; text untested on real data"],
            ["SUCCESS is the majority class, not the rare class assumed", "Compared weights, oversampling and thresholds; reported recall trade-offs"],
            ["Some fields look recorded after the outcome", "Ran a leakage ablation (ROC-AUC 0.82 to 0.72) and stated the limit"],
            ["Saved models failed to load with the pinned library versions", "Re-trained with the pinned versions; Logistic Regression unchanged"]]
    K.table(s, 0.6, 1.5, [5.6, 6.53], rows, font_size=16, row_h=0.95)
    notes[18] = "Each challenge changed either the experiment design or how results are reported."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Calibrate", "Correct the probability bias of the class-weighted model and re-check the risk label thresholds.", K.BLUE, "1"),
        ("Richer data", "Find a table with dates, descriptions and founder fields to test text and a temporal split.", K.GREEN, "2"),
        ("Explain better", "Add SHAP-style attributions and significance tests over more splits.", K.ORANGE, "3"),
        ("Finish and document", "Final deck, report and paper from the same measured numbers.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.1, 1.6, 2.9, 4.4, t, b_, col, gl, body_size=17)
    notes[19] = "Plan to the final review. None of these items is claimed as done."

    # ---- 20 thank you
    thank_you(S[19])
    notes[20] = "Likely questions: why the shipped model is not the best on real data; why text is untested on real data; how leakage was checked; what the base paper's numbers mean for us."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx")))
