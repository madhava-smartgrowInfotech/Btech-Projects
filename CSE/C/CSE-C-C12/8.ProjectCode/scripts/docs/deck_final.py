"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .lit import (BASE_PAPER_POINTS, EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT, PROPOSED_APPROACH,
                  SUPPORT_EVIDENCE)
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520
CROPS = os.path.join(tempfile.gettempdir(), "c12_deck_crops")


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def crop(name, box):
    """Crop a region (left, top, right, bottom in pixels) from a real screenshot into a temp PNG and return its path."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_{box[1]}.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "hybrid RAG, symptom monitoring, FAISS, BM25, cross-encoder, healthcare chatbot"
    notes = {}
    E, KB, G, SF, REL = F.ENV, F.KB, F.GEN, F.SAF, F.REL
    r = F.ret
    s1p, s1l, s2, s3, s4 = "S1_shipped_protocol_30", "S1_leave_one_out_30", "S2_verbatim_leave_one_out_176", "S3_two_symptoms_leave_one_out_176", "S4_lay_language_60"
    LAT, T5, FB = F.LAT, F.LAT_T5, F.LAT_FB
    shipped = F.SHIPPED_EVAL
    probe = "60 lay-language probes written by the team on a synthetic knowledge base"

    # footers on every slide except the title
    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.FULL_TITLE, 24)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: SymptoTrackAI, implemented as the CareLens AI web application, a local hybrid RAG assistant for symptoms."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Knowledge Base & Models\nRetrieval Workflow, Interaction & Storage\nImplementation Highlights\nResults: Retrieval, Confusion Matrix, Safety\nPerformance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [
        ("Delayed diagnosis and limited access push people to web search and general chatbots, which give ungrounded answers and keep no record of how symptoms evolve.", {"bullet": True, "space_after": 9}),
        (f"SymptoTrackAI, built as the CareLens AI web app, expands a symptom description, retrieves records with FAISS (MiniLM) and BM25, fuses them 0.65 / 0.35, reranks with a cross-encoder and builds the answer only from the retrieved records, with rule-based emergency warnings and a private consultation history.", {"bullet": True, "space_after": 9}),
        (f"On a synthetic {KB['records']}-record, {KB['conditions']}-condition base the shipped pipeline puts the right condition first for {F.pct(r('hybrid_rerank', s4))} of {F.RET['hybrid_rerank'][s4]['n']} team-written lay-language probes and {F.pct(r('hybrid_rerank', s3))} of two-symptom queries (leave-one-out). The shipped 100% figure is an upper bound.", {"bullet": True, "space_after": 9}),
        (f"Emergency rules flag {SF['by_group_pos']['literal']['detected']} of {SF['by_group_pos']['literal']['n']} literally worded cases but 0 of {SF['by_group_pos']['paraphrase']['n']} paraphrases; the local Flan-T5 generator echoes its prompt ({G['echo_answers']} of {G['n']}); relevance always shows 100%.", {"bullet": True, "space_after": 9}),
        (f"A full answer takes about {FB['chat_ms']['median'] / 1000:.1f} s (fallback answer) or {T5['chat_ms']['median'] / 1000:.1f} s (Flan-T5) on a CPU. {F.PYTEST_PASSED} of {F.PYTEST_TOTAL} tests and {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} smoke checks pass.", {"bullet": True, "space_after": 9}),
        ("Not done: trend analysis over time, a real medical knowledge base, clinician review, a user study.", {"bullet": True}),
    ], size=16)
    notes[3] = "Read the problem, then the system in one breath. Give each number with its protocol and end on the honest limits."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.8, F.BASE_PAPER.replace("*", ""), size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: a hybrid RAG framework for intelligent symptom monitoring and disease guidance"],
            ["Method", "Sequence-NQ transformer embeddings of a symptom-disease knowledge base, FAISS similarity search, RAG generation of guidance"],
            ["Data", "A symptom-disease dataset with advice, de-duplicated and indexed; size and source not stated"],
            ["Application", "Registration, login, chatbot symptom analysis and a stored consultation log (health log) for continued monitoring"],
            ["Reported results", "No accuracy, recall, latency or ablation figures; 'experimental observations' only, one example (symptoms leading to Typhoid)"],
            ["Stated limits", "None; future work lists a larger knowledge base, multilingual and multimodal support, cloud deployment"]]
    ftable(s, 0.55, 2.7, [2.0, 10.2], rows, size=13, row_h=0.62)
    notes[4] = "Madhury and Ramesh, IJERST 2026. We adopt the product idea and the web-app modules; the paper gives no numbers, so nothing can be compared quantitatively."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Data / finding / limitation"]]
    for row in LIT_TABLE:
        rows.append([row[0].replace(" (base paper)", " (base)"), row[1], row[2], row[3]])
    ftable(s, 0.55, 1.8, [2.5, 0.7, 4.7, 4.3], rows, size=10, row_h=0.55, align=["l", "c", "l", "l"])
    notes[5] = "Representative work: chatbots, RAG over books and hospital manuals, a guideline-grounded demo and HyDE. Three survey items could not be verified and are flagged in the report."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in GAPS], size=20)
    notes[6] = "Six gaps drawn from the survey. We address 1 to 4 and 6 directly; trend analysis (5) is only partly addressed."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.9, 12.2, 2.3)
    ftext(s, 0.85, 1.95, 11.6, 2.2, PROBLEM_STATEMENT, size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Symptoms in everyday language, 3 to 2,000 characters, from a signed-in user"), ("Output", "Evidence-backed answer, possible-condition cards, safety flag, relevance label"), ("Record", "Private consultation history: search, open, delete")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.5, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.6, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.1, 3.5, 1.0, b_, size=15)
    notes[7] = "The statement has three parts: ground the answer, warn on emergencies, and keep an inspectable private record."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion states the real status of each: 5 is not met beyond storage, 4 and 6 only partly."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.1, [("Symptom text", K.FNAVY), ("Expand", K.FBLUE), ("FAISS + BM25", K.GREEN), ("Fuse + rerank", K.ORANGE), ("Answer + safety", K.PURPLE)], size=15, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.35, 12.2, 3.5, [(p, {"bullet": True, "space_after": 10}) for p in PROPOSED_APPROACH] + [
        ("The assistant gives information, never a diagnosis: every screen carries the disclaimer.", {"bullet": True})], size=17)
    notes[9] = "The proposed system as one line of boxes and the design rules behind it."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 8.4, 5.1)
    ftext(s, 9.2, 2.0, 3.6, 4.8, [
        (f"Web app: React + Vite, port {F.WEB_PORT}", {"bullet": True, "space_after": 8}),
        (f"API: FastAPI, port {F.API_PORT}, 10 routes", {"bullet": True, "space_after": 8}),
        ("RAG: expansion, FAISS, BM25, fusion, rerank, generator, safety", {"bullet": True, "space_after": 8}),
        ("Storage: SQLite (users, consultations), index files", {"bullet": True, "space_after": 8}),
        ("Local first: no symptom text leaves the machine", {"bullet": True}),
    ], size=15)
    notes[10] = "Layers and storage; all analysis runs on the local machine, with no paid API."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.9, 12.2, 1.2, [("Normalise", K.FNAVY), ("HyDE template", K.FBLUE), ("FAISS + BM25", K.GREEN), ("Fusion", K.ORANGE), ("Cross-encoder", K.PURPLE), ("Answer", K.FBLUE), ("Safety + store", K.FNAVY)], size=13, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.4, 12.2, 3.4, [
        ("Nothing is trained: MiniLM, the MS-MARCO cross-encoder and Flan-T5 are used as published; the work is the pipeline, the safety layer and the measurement.", {"bullet": True, "space_after": 9}),
        (f"Retrieval measured on five query sets: shipped protocol (30), leave-one-out (30 and 176), two random symptoms (176) and {F.RET['hybrid_rerank'][s4]['n']} team-written lay-language probes.", {"bullet": True, "space_after": 9}),
        ("Seven configurations compared: BM25, dense, hybrid, hybrid + rerank (shipped), shipped without expansion, and two TF-IDF offline modes.", {"bullet": True, "space_after": 9}),
        (f"Safety layer probed with {SF['n_pos']} emergency and {SF['n_neg']} non-emergency texts; displayed relevance probed with {REL['out_kb_n']} out-of-base complaints.", {"bullet": True, "space_after": 9}),
        ("Every number is written by scripts/measure_metrics.py to docs/metrics/metrics.json.", {"bullet": True}),
    ], size=16)
    notes[11] = "Workflow and the measurements added on top of the shipped system."

    # ---- 12 knowledge base and models
    s = S[11]
    set_title(s, "Knowledge Base & Models")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.05, f"{KB['records']}", "records, synthetic (seed 42)", K.FBLUE, 26)
    ftile(s, 0.55, 3.1, 2.9, 1.05, f"{KB['conditions']}", "conditions, 4 to 6 records each", K.GREEN, 26)
    ftile(s, 0.55, 4.3, 2.9, 1.05, "384", "embedding dimensions (MiniLM)", K.ORANGE, 26)
    K.picture(s, F.fig("f07_kb_composition.png"), 3.7, 1.85, 5.2, 2.6)
    rows = [["Component", "Choice"], ["Dense", "all-MiniLM-L6-v2 + FAISS flat"], ["Lexical", "BM25Okapi"], ["Rerank", "ms-marco MiniLM cross-encoder"], ["Answer", "Flan-T5-small / fallback"]]
    ftable(s, 9.1, 1.9, [1.5, 2.3], rows, size=11, row_h=0.62, first_bold=True)
    ftext(s, 3.7, 4.55, 5.3, 2.3, [
        (f"Severity: {KB['severity']['moderate']} moderate, {KB['severity']['low']} low, {KB['severity']['high']} high, {KB['severity']['emergency']} emergency records.", {"bullet": True}),
        (f"All records of a condition share one description and guidance text ({KB['distinct_description_guidance_pairs']} distinct pairs), so retrieving the condition is easy.", {"bullet": True}),
        ("No public dataset was used; the base is team-written and not clinically reviewed.", {"bullet": True}),
    ], size=13)
    ftext(s, 0.55, 5.55, 3.0, 1.3, "Each record is searched as 'Disease: .. Symptoms: .. Description: .. Guidance: ..'.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "Be frank about the data: synthetic and written by the team. It makes the condition easy to find and the metrics optimistic."

    # ---- 13 retrieval workflow
    s = S[12]
    set_title(s, "Retrieval Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_retrieval.png"), 0.55, 1.8, 5.6, 5.1)
    ftext(s, 6.4, 1.95, 6.4, 4.9, [
        ("Query and template are embedded; FAISS returns the 15 nearest records.", {"bullet": True, "space_after": 10}),
        ("BM25 scores the same text lexically and returns 15 records.", {"bullet": True, "space_after": 10}),
        ("Each list is min-max scaled; fused = 0.65 x semantic + 0.35 x lexical.", {"bullet": True, "space_after": 10}),
        ("The cross-encoder scores the 15 query-record pairs; the top 5 are kept and the scores are min-max scaled for display.", {"bullet": True, "space_after": 10}),
        ("The answer is built from those records; rules check the query and filter the answer.", {"bullet": True}),
    ], size=16)
    notes[13] = "The core algorithm, in app/rag/pipeline.py."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_chat.png"), 0.55, 1.8, 7.2, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.9, 1.8, 4.9, 5.1)
    notes[14] = "Sequence of one chat request and the use cases of the user."

    # ---- 15 database design
    s = S[14]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.5)
    rows = [["Store", "Content"], ["table users", "id, unique email, bcrypt hash, created_at"], ["table consultations", "query, answer, top condition, evidence JSON, score, flag"], ["knowledge_base.json/csv", f"{KB['records']} synthetic records"],
            ["faiss.index", "384-d flat inner-product index"], ["bm25.json", "tokenised corpus"], ["records.json", "records aligned with the index"]]
    ftable(s, 7.8, 1.9, [2.4, 2.6], rows, size=11, row_h=0.55)
    ftext(s, 0.55, 6.4, 12.2, 0.5, "A consultation keeps its evidence as JSON text, so history stays readable even if the knowledge base is rebuilt; deleting a user cascades to the rows.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Two SQLite tables plus file-based artefacts."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("10", "REST routes", K.FBLUE), (f"{F.PYTEST_TOTAL} + {F.SMOKE_TOTAL}", "pytest tests + smoke checks", K.GREEN), ("8", "React pages", K.ORANGE), ("800", "lines of backend code", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Accounts: bcrypt hashing, HS256 token in an httponly cookie (1,440 minutes), per-user history isolation checked by a test (404 for another user's record).", {"bullet": True, "space_after": 9}),
        ("Validation: query 3 to 2,000 characters; Pydantic schemas; clear errors for 401, 400, 422, 503 and 500.", {"bullet": True, "space_after": 9}),
        ("Graceful degradation: MiniLM missing -> TF-IDF vectors; reranker missing -> hybrid order; generator missing -> evidence-based answer.", {"bullet": True, "space_after": 9}),
        ("History with search, detail and delete (with confirmation); light and dark themes; one command starts both servers (run.bat); Docker + nginx variant.", {"bullet": True}),
    ], size=17)
    notes[16] = "Highlights: accounts and isolation, validation, degradation paths and one-command start."

    # ---- 17 results: retrieval ablation
    s = S[16]
    set_title(s, "Results – Retrieval Configurations")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.pct(r('hybrid_rerank', s4)), f"shipped pipeline, lay language (n = 60)", K.FBLUE), (F.pct(r('hybrid_rerank', s3)), "two random symptoms (n = 176)", K.GREEN), (F.pct(r('bm25', s4)), "BM25 alone, lay language", K.ORANGE), (F.pct(r('hybrid_tfidf', s4)), "offline TF-IDF mode, lay language", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26)
    K.picture(s, F.fig("f02_ablation.png"), 0.55, 3.1, 7.4, 3.3)
    K.picture(s, F.fig("f09_hyde_effect.png"), 8.1, 3.1, 4.7, 2.5)
    ftext(s, 8.1, 5.6, 4.7, 1.3, f"Dropping the HyDE-style template scores {r('hybrid_rerank_nohyde', s4):.3f} on lay language against {r('hybrid_rerank', s4):.3f} shipped; Wilson 95% interval of the shipped result {F.ci('hybrid_rerank', s4)}.", size=12)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "Top-1 condition accuracy. The base is synthetic and the probes are team-written: these show how the components behave, not clinical accuracy.", size=12, color=K.FNAVY, italic=True)
    notes[17] = (f"Lay language: BM25 {r('bm25', s4):.3f}, dense {r('dense', s4):.3f}, hybrid {r('hybrid', s4):.3f}, hybrid plus cross-encoder {r('hybrid_rerank', s4):.3f}. "
                 f"Two random symptoms are harder: {r('hybrid_rerank', s3):.3f}. The reranker helps most; the expansion template does not help.")

    # ---- 18 confusion matrix and protocol comparison
    s = S[17]
    set_title(s, "Confusion Matrix & Evaluation Protocol")
    wipe(s)
    K.picture(s, F.fig("f01_lay_confusion.png"), 0.55, 1.85, 5.6, 4.6)
    K.picture(s, F.fig("f03_protocol_comparison.png"), 6.4, 1.85, 6.4, 3.3)
    miss = F.RET["hybrid_rerank"][s4]["misses"]
    ftext(s, 6.4, 5.25, 6.4, 1.6, [
        (f"Two lay-language misses: " + "; ".join(f"{m['truth']} read as {m['predicted']}" for m in miss) + ".", {"bullet": True, "size": 13}),
        (f"The shipped evaluation (n = {shipped.get('n_test', 30)}, top-1 = {shipped.get('top1_accuracy', 1.0):.1f}) keeps the held-out record in the index: an upper bound. Leave-one-out gives {r('hybrid_rerank', s1l):.3f}.", {"bullet": True, "size": 13}),
    ], size=13)
    ftext(s, 0.55, 6.5, 5.6, 0.5, "Rows: true condition; columns: predicted (60 lay-language probes).", size=12, color=K.FNAVY, italic=True)
    notes[18] = "The confusion matrix of the shipped pipeline on the lay-language probes: errors are common cold read as allergic rhinitis and typhoid read as appendicitis. The right figure shows how the protocol changes the number."

    # ---- 19 relevance display and safety
    s = S[18]
    set_title(s, "Results – Relevance Display & Safety Layer")
    wipe(s)
    K.picture(s, F.fig("f05_relevance_display.png"), 0.55, 1.8, 6.1, 3.2)
    K.picture(s, F.fig("f06_safety_rules.png"), 6.8, 1.8, 6.0, 3.2)
    ftext(s, 0.55, 5.15, 6.1, 1.8, [
        (f"All {REL['out_kb_n']} out-of-base complaints show 'Strong match' at 100%, like every in-base query.", {"bullet": True}),
        (f"The raw cross-encoder score separates them (AUROC {REL['auroc_in_vs_out_by_raw_rerank_score']:.1f}); min-max scaling hides it.", {"bullet": True}),
    ], size=13)
    hard = SF["by_group_neg"]["hard"]
    ftext(s, 6.8, 5.15, 6.0, 1.8, [
        (f"Emergency recall {SF['emergency_recall']:.3f}: literal wording {SF['by_group_pos']['literal']['detected']}/{SF['by_group_pos']['literal']['n']}, paraphrase {SF['by_group_pos']['paraphrase']['detected']}/{SF['by_group_pos']['paraphrase']['n']}.", {"bullet": True}),
        (f"Ordinary texts: 0 false alarms; hard negatives (negation, history): {len(hard['false_emergency'])} of {hard['n']} wrongly flagged.", {"bullet": True}),
        (f"Answer filter cleaned {SF['filter']['listed_removed']} of {SF['filter']['n_listed_phrases']} listed-phrase cases.", {"bullet": True}),
    ], size=13)
    notes[19] = "Two honest defects of the display and the safety layer, both measured with author-written probes."

    # ---- 20 performance and testing
    s = S[19]
    set_title(s, "Performance & Testing")
    st = LAT["stages_ms"]
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{st['embed']['median']:.0f} ms", "query embedding (median)", K.FBLUE), (f"{st['rerank']['mean']:.0f} ms", "cross-encoder, 15 pairs (mean)", K.GREEN),
                                   (f"{FB['chat_ms']['median'] / 1000:.2f} s", "/api/chat median, fallback answer", K.ORANGE), (f"{T5['chat_ms']['median'] / 1000:.2f} s", "/api/chat median, Flan-T5 on", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.15, v, l, c, 26)
    K.picture(s, F.fig("f04_latency_stages.png"), 0.55, 3.15, 6.0, 2.5)
    K.picture(s, F.fig("f08_api_latency.png"), 6.8, 3.15, 6.0, 2.5)
    ftext(s, 0.55, 5.8, 12.2, 1.1, [
        (f"Tests: pytest {F.PYTEST_PASSED} of {F.PYTEST_TOTAL} passed in {F.PYTEST_SECONDS:.0f} s; smoke test {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} checks; browser walk of 21 screens with 0 console errors.", {"size": 13, "color": K.FNAVY, "bold": True}),
        (f"{E['cpu']}, {E['ram_gb']} GB, CPU only, warm models; first start about 62 s to load three models. Not tested: concurrent users, a GPU, accessibility audit, a user study.", {"size": 12, "color": K.FNAVY, "italic": True}),
    ])
    notes[20] = "The cross-encoder is the dominant stage; Flan-T5 adds about 0.4 s per answer and, under transformers 5, returns the prompt."

    # ---- 21 comparison with the base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only – different data and scope)")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "This project (SymptoTrackAI)"],
            ["Retrieval", "Sequence-NQ embeddings + FAISS", "MiniLM + FAISS, BM25, min-max fusion, cross-encoder rerank"],
            ["Query handling", "Direct embedding of the symptom text", "Normalisation and a HyDE-style template (no measured gain)"],
            ["Safety", "None described", "10 emergency rules, caution keywords, answer filter, disclaimer"],
            ["Application", "Registration, login, chatbot, health log", "Same modules, plus evidence cards, history search / delete, themes"],
            ["Reported", "No numbers; one Typhoid example", f"Lay-language top-1 {r('hybrid_rerank', s4):.3f}; {probe}"],
            ["Trend tracking", "Claimed through the stored log", "Stored and searchable only; no trend analysis"]]
    ftable(s, 0.55, 1.85, [1.9, 4.6, 5.7], rows, size=12, row_h=0.62)
    ftext(s, 0.55, 6.3, 12.2, 0.6, "The base paper gives no quantitative results, so no numerical comparison is possible and no superiority is claimed.", size=13, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that the base paper has no numbers; the comparison is of what each system does, not of how well."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Symptom Assistant")
    wipe(s)
    K.picture(s, crop("07_assistant_answer.png", (248, 285, 1440, 900)), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("20_assistant_evidence.png", (248, 285, 1440, 900)), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: a fever query answered from the retrieved influenza records (evidence-based answer). Right: the matching evidence card for dengue, with typical symptoms, guidance, severity and 'seek urgent care if' red flags.", size=12, color=K.FNAVY)
    notes[22] = "Real screenshots from the running app. The answer is built from retrieved records; the evidence cards show what supports it."
    s = S[22]
    set_title(s, "Output – Safety Escalation and Local Generator")
    wipe(s)
    K.picture(s, crop("08_assistant_emergency.png", (248, 285, 1440, 900)), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("21_assistant_generator_echo.png", (248, 285, 1440, 900)), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: the emergency banner for chest pain with breathing difficulty. Right: the same app with the default Flan-T5 generator: the 'answer' is the prompt echoed back, the defect we measured.", size=12, color=K.FNAVY)
    notes[23] = "The right screenshot is a failure case shown on purpose: with transformers 5 the generator returns its own prompt."
    s = S[23]
    set_title(s, "Output – History and About Page")
    wipe(s)
    K.picture(s, crop("09_history.png", (248, 285, 1440, 900)), 0.55, 1.85, 6.0, 3.9, border=True)
    K.picture(s, crop("13_about_metrics.png", (248, 285, 1440, 900)), 6.8, 1.85, 6.0, 3.9, border=True)
    ftext(s, 0.55, 5.95, 12.2, 0.9, "Left: consultation history with safety flag and relevance (every row 100%). Right: the About page, which shows the shipped evaluation (100%, n = 30) and states that these scores are an upper bound on synthetic data.", size=12, color=K.FNAVY)
    notes[24] = "The product itself says the shipped metrics are an upper bound; the history shows the relevance display problem."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 5.0, [
        (f"**Delivered:** a working local application from symptom text to an evidence-backed answer, emergency warning and private history, with {F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks passing.", {"bullet": True, "space_after": 9}),
        (f"**Measured:** lay-language top-1 {r('hybrid_rerank', s4):.3f}, two-symptom {r('hybrid_rerank', s3):.3f}, leave-one-out verbatim {r('hybrid_rerank', s2):.3f}; BM25 alone {r('bm25', s4):.3f}; offline mode {r('hybrid_tfidf', s4):.3f}.", {"bullet": True, "space_after": 9}),
        ("**Objectives:** 1, 2, 3 met on a synthetic base; 4 partly (no per-statement citation); 5 not met beyond storage; 6 partly (groundedness not judged).", {"bullet": True, "space_after": 9}),
        ("**Defects:** prompt-echo generator, always-100% relevance, safety misses paraphrases, near-duplicate evidence cards.", {"bullet": True, "space_after": 9}),
        ("**Position:** information only; a clinician decides.", {"bullet": True}),
    ], size=19)
    notes[25] = "Be explicit about which objectives were met and which were not, and what was never measured."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 5.0, [
        ("Replace the synthetic base with a licensed medical source and have a clinician review conditions, guidance and red flags.", {"bullet": True, "space_after": 10}),
        ("Fix the generator for current transformers, add per-statement citations and a groundedness judge.", {"bullet": True, "space_after": 10}),
        ("Show a calibrated relevance from the raw cross-encoder score; return distinct conditions instead of near-duplicate records.", {"bullet": True, "space_after": 10}),
        ("Safety layer with synonyms, negation handling and a small classifier; evaluate with clinician-written cases.", {"bullet": True, "space_after": 10}),
        ("Symptom timeline and trend view over stored consultations; multilingual input; cache or shrink the reranker for speed.", {"bullet": True}),
    ], size=20)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [1, 5, 6, 8, 9, 11, 15, 16, 17, 19, 20, 21, 24, 25]
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(f"[{n}] {REFS[n - 1].replace(chr(42), chr(32)).replace(chr(32) * 2, chr(32))}", {"space_after": 5}) for n in ids[:7]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(f"[{n}] {REFS[n - 1].replace(chr(42), chr(32)).replace(chr(32) * 2, chr(32))}", {"space_after": 5}) for n in ids[7:]], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
