"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import (EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIMITATIONS_EXISTING, LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT,
                  PROPOSED_APPROACH, SUPPORT_EVIDENCE)
from .refs import CITE, REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "hybrid RAG, symptom monitoring, FAISS, BM25, cross-encoder, healthcare chatbot"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: SymptoTrackAI, implemented as the CareLens AI web application, turns symptoms written in everyday language "
                "into an answer built only from retrieved medical records, with evidence, a safety warning when needed and a private history.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: from symptoms to grounded guidance")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Healthcare suffers from delayed diagnosis and limited access. People search the web or ask general chatbots, which give frightening or ungrounded answers and keep no record of how symptoms evolve.", K.ORANGE, "!"),
        ("Our idea", "A hybrid RAG pipeline: expand the query, retrieve with FAISS semantic search and BM25, fuse the scores, rerank with a cross-encoder, and write the answer only from the retrieved records, with a safety layer on top.", K.BLUE, "i"),
        ("The deliverable", "A local web application: accounts, a symptom assistant with evidence cards and emergency warnings, a private searchable consultation history and an About page with measured metrics. Information only, never a diagnosis.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: the access and misinformation problem, the hybrid retrieval idea, and what is delivered: a complete application whose output is information, not diagnosis."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER.replace("*", ""), size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "SymptoTrackAI framework: a symptom-disease knowledge base is embedded with the Sequence-NQ transformer and indexed with FAISS; the nearest records are passed to a RAG step that writes the guidance. A web app adds registration, login and a stored consultation log.", K.BLUE, "M"),
        ("Data and reported results", "A symptom-disease dataset with advice text, de-duplicated and indexed. The paper reports no accuracy, recall or latency figures: only qualitative 'experimental observations' and one example (constipation, stomach ache, high fever, swollen lymph nodes: Typhoid).", K.GREEN, "R"),
        ("Relevance to this project", "Same product idea and same web-app modules. This project adds lexical BM25 retrieval, fusion, a cross-encoder reranker, a HyDE-style expansion, a safety layer and measured evaluation. The base paper has none of these.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The base paper is qualitative: it gives no numbers, so nothing can be compared quantitatively. "
                "It also claims long-term symptom tracking; what it actually describes is storing consultations, and this project does the same.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for i in (0, 1, 2, 4, 5, 6, 7):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.35, [2.5, 0.7, 4.5, 4.43], rows, font_size=13, row_h=0.66, header_size=13, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: chatbots built on machine learning or BERT, RAG over medical books or hospital manuals, a guideline-grounded demo (ClinicBot) and HyDE query expansion. "
                "Three more papers in the batch survey (items 2 to 4) could not be located, so they are listed in the report as not verified and not used for any claim.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: why retrieval and grounding")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 5.7, 5.2, ["Fine-tuned\nBGE-M3 [%d]" % CITE["emr"], "Multilingual\nE5-Large [%d]" % CITE["emr"]],
                [("Retrieval accuracy", [0.976, 0.897])], colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.2)
    K.text(s, 6.6, 1.45, 6.1, 5.0, [(e, {"bullet": True, "space_after": 11}) for e in SUPPORT_EVIDENCE], size=15, color=K.INK)
    K.text(s, 6.6, 6.35, 6.1, 0.5, "Hospital EMR-manual retrieval (5,931 pairs); a different task, shown as context, not as our result.", size=12, color=K.TXT, italic=True)
    notes[6] = ("Retrieval quality depends on the embedding model: one published EMR chatbot reports 97.6 and 89.7 percent for two embedders on its own data. "
                "We do not reuse those numbers; this project measures its own retrieval on its own knowledge base. The base paper offers no numbers at all.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Person has symptoms", K.NAVY), ("Web search or chatbot", K.NAVY), ("Generic answer", K.NAVY), ("Unrecorded decision", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Web search and general chatbots", "\n".join(EXISTING_SYSTEM[i] for i in (0, 1)) + "\nNeither keeps a private record of earlier episodes.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Retrieval chatbots", "\n".join(EXISTING_SYSTEM[i] for i in (2, 3, 4)), K.ORANGE, "2", body_size=17)
    notes[7] = "Today people search or chat. Retrieval-based chatbots exist but are rarely measured, rarely explain a match and rarely include emergency rules."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Generic results", "Ungrounded answers", "Unmeasured quality", "No emergency logic"]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], LIMITATIONS_EXISTING[i], C[i], str(i + 1), body_size=19)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: " + LIMITATIONS_EXISTING[4], size=15, color=K.INK, anchor="m")
    notes[8] = "Five limitations. Each maps to a design decision: retrieval grounding, hybrid retrieval with measurement, a rule-based safety layer and local processing."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 0.9
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.82, g, size=19, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as longitudinal symptom tracking and personalised insight on top of the retrieval pipeline. We address gaps 1 to 4 and 6; "
                "gap 5, trend analysis, is only partly addressed: consultations are stored and searchable, trends are not analysed.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.9, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.65, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=20, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Symptoms typed in everyday language (3 to 2,000 characters) by a signed-in user"), ("Output", "Evidence-backed answer, possible-condition cards, safety flag and relevance label"), ("Record", "A private, searchable consultation history that the user can delete")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.65, 3.85, 1.9)
        K.text(s, x + 0.25, 4.77, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.25, 3.4, 1.25, b_, size=16, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is an assistant that answers only from a curated knowledge base, shows its sources and keeps a history."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 20}) for i, o in enumerate(OBJECTIVES, 1)], size=24, color=K.INK)
    notes[11] = "The six objectives from the proposal. Later reviews report which were met; objective 5, trends over time, will be only partly met."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Expand", "Normalise the text and wrap it in a local HyDE-style template, with no external API.", K.BLUE, "1"),
        ("Retrieve", "MiniLM embeddings searched in FAISS, plus BM25 over the same record text: 15 candidates each.", K.GREEN, "2"),
        ("Fuse and rerank", "Min-max fusion 0.65 semantic + 0.35 lexical, then an MS-MARCO cross-encoder keeps the top 5 records.", K.ORANGE, "3"),
        ("Answer safely", "Answer built from the records, emergency rules, a medication filter and a permanent disclaimer.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "S", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "fused = 0.65 x semantic + 0.35 x lexical (each min-max normalised over the candidate pool)", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.85, "FastAPI service and React web app over SQLite, bcrypt and signed-cookie accounts, all models local and CPU-only. Offline fallbacks: TF-IDF embeddings and an evidence-based answer.", size=14, color=K.TXT)
    notes[12] = "Four stages from symptom text to a safe answer. The weights live in one configuration file; the fallbacks keep the app usable offline."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.4, 5.5)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"Presentation: React and Vite web app (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}) with auth, chat and history routes.", {"bullet": True}),
        ("RAG pipeline: expansion, FAISS and BM25 retrieval, fusion, rerank, generation, safety.", {"bullet": True}),
        ("Data: SQLite for users and consultations; knowledge base and indexes on local disk.", {"bullet": True}),
    ], size=17, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser sends the symptom text, the API runs the RAG pipeline, and the local files hold the knowledge base, the indexes and the users' consultations."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working application", "Measured retrieval", "Safety escalation", "Honest limits", "Information only"]
    for i, (h, b_) in enumerate(zip(heads, EXPECTED_OUTCOMES)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.55, h, b_, C[i % 4], str(i + 1), body_size=17)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=18)
    notes[14] = "Expected outcomes and contributions. The later reviews report measured results, including what could not be measured."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 5, 6, 8, 9, 11, 14, 17, 19, 20]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 4}) for n in idx], size=12, color=K.INK)
    notes[15] = f"Selected references; the complete numbered list of {len(REFS)} entries, with three marked as not verified, is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why BM25 next to dense search (exact symptom words); why a cross-encoder (precision on the top records); "
                 "what the About page numbers mean (an upper bound on synthetic data); and what stops the assistant from diagnosing (retrieval-only answers and a disclaimer).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
