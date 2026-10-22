"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                os.makedirs(os.path.dirname(BANNER), exist_ok=True)
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p("SymptoTrackAI, implemented as the CareLens AI web application: grounded medical information for symptoms described in everyday language, with safety escalation and a private consultation history",
        align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, whose framework gave this project its starting point, and the maintainers of the open-source software it depends on "
        "(FastAPI, React, Vite, SQLAlchemy, sentence-transformers, FAISS, rank-bm25, Hugging Face Transformers, PyTorch and scikit-learn), together with the "
        "authors of the all-MiniLM-L6-v2, ms-marco-MiniLM-L-6-v2 and Flan-T5 models that run inside the system.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    s2 = F.SHIP["S2_verbatim_leave_one_out_176"]
    s3 = F.SHIP["S3_two_symptoms_leave_one_out_176"]
    s4 = F.SHIP["S4_lay_language_60"]
    nh4 = F.RET["hybrid_rerank_nohyde"]["S4_lay_language_60"]
    tf4 = F.RET["hybrid_tfidf"]["S4_lay_language_60"]
    bm4 = F.RET["bm25"]["S4_lay_language_60"]
    sf = F.SAF
    g = F.GEN
    d.p("People with early, minor symptoms often delay seeking care, and a web search for those symptoms tends to return alarming or irrelevant results; a general chatbot, on the other hand, can state "
        "unsupported medical claims with confidence. This project presents SymptoTrackAI, implemented as the CareLens AI web application, a local symptom-monitoring assistant that answers only from a "
        "curated knowledge base. A symptom description in everyday language is normalised and expanded, then searched in two ways at once: a dense search (all-MiniLM-L6-v2 embeddings in a FAISS index) "
        f"and a lexical BM25 search. The two score lists are fused with weights {F.FUSION_WEIGHTS[0]} and {F.FUSION_WEIGHTS[1]}, the {F.CANDIDATE_K} best candidates are reranked by a cross-encoder, and the top "
        f"{F.TOP_K} records form an evidence-based answer that passes a rule-based safety layer. Every consultation is stored privately in SQLite and can be searched, reopened and deleted. The work follows the base paper in using retrieval-augmented generation with a symptom-disease knowledge base and extends it with BM25 fusion, neural reranking, "
        "safety escalation and a React interface.")
    d.p(f"The knowledge base is synthetic: {F.KB['records']} records covering {F.KB['conditions']} conditions, written by the team and not clinically validated. On {s4['n']} everyday-language probes written by "
        f"the team the shipped pipeline puts the correct condition first in {F.pct(s4['top1'])} of cases (BM25 alone {F.pct(bm4['top1'])}; the offline TF-IDF fallback {F.pct(tf4['top1'])}); with only two symptoms "
        f"given it reaches {F.pct(s3['top1'])}, and on leave-one-out verbatim queries {F.pct(s2['top1'])}. Removing the template query expansion raised the probe score to {F.pct(nh4['top1'])}, so that stage did not help. "
        f"The measurement also exposed defects, which are reported and not hidden: with the installed library version the local language model returns its own prompt ({g['echo_answers']} of {g['n']} answers), "
        f"the displayed relevance of the best match is always 100%, even for complaints outside the knowledge base, the emergency rules recognised {sf['true_emergency']} of {sf['n_pos']} test descriptions "
        f"(none of the {sf['by_group_pos']['paraphrase']['n']} paraphrased ones), and tracking of symptom trends over time was not built beyond the stored history. "
        f"A request takes about {F.LAT_FB['chat_ms']['mean'] / 1000:.1f} s on a CPU-only PC. The base paper gives no numbers, so no quantitative comparison with it is possible.")
    d.p("**Keywords:** hybrid retrieval-augmented generation, symptom monitoring, FAISS semantic search, BM25, neural reranking, query expansion, healthcare chatbot, safety escalation, consultation history, FastAPI, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUROC", "Area Under the Receiver Operating Characteristic curve"),
    ("BM25", "Best Matching 25 (probabilistic ranking function)"), ("CI", "Confidence Interval"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CPU", "Central Processing Unit"), ("CRUD", "Create, Read, Update, Delete"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"),
    ("FAISS", "Facebook AI Similarity Search"), ("HTTP", "Hypertext Transfer Protocol"), ("HyDE", "Hypothetical Document Embeddings"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("LLM", "Large Language Model"), ("LOO", "Leave-One-Out"),
    ("ML", "Machine Learning"), ("MRR", "Mean Reciprocal Rank"), ("NLP", "Natural Language Processing"), ("ORM", "Object-Relational Mapping"),
    ("RAG", "Retrieval-Augmented Generation"), ("REST", "Representational State Transfer"), ("SBERT", "Sentence-BERT"), ("SPA", "Single-Page Application"),
    ("SQL", "Structured Query Language"), ("TF-IDF", "Term Frequency - Inverse Document Frequency"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
