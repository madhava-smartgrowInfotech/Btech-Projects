"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx.

All measured numbers come from facts.py (docs/metrics/metrics.json, tests.json).
    python -m scripts.docs.deck_review1        # from 8.ProjectCode
"""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "client-side encryption, AES-256-GCM, RSA-OAEP, role-based access control, audit log"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 30)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 18)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}\nBatch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def trim(path, tag):
    """Copy of a diagram with its white margins removed (written to the temp folder, source untouched)."""
    from PIL import Image, ImageChops

    with Image.open(path) as im:
        im = im.convert("RGB")
        bbox = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).getbbox()
        if bbox:
            pad = 12
            bbox = (max(bbox[0] - pad, 0), max(bbox[1] - pad, 0), min(bbox[2] + pad, im.size[0]), min(bbox[3] + pad, im.size[1]))
            im = im.crop(bbox)
        out = os.path.join(tempfile.gettempdir(), f"ciphervault_r1_{tag}.png")
        im.save(out)
    return out


PROBLEM_STATEMENT = [
    "Files stored without client-side encryption are readable by the storage provider",
    "Without role-based control, any authenticated user may reach another user's files",
    "Tampering goes undetected when no integrity check is performed on retrieval",
    "Without audit logs, a security incident cannot be reconstructed afterwards",
    "Need: a storage system that encrypts before upload, enforces role-based access, verifies integrity and logs every action",
]
OBJECTIVES = [
    "Implement user registration and authentication with securely hashed credentials and JWT sessions",
    "Encrypt every uploaded file using AES-256-GCM before it reaches storage",
    "Implement role-based access control together with file-ownership verification",
    "Verify file integrity on download using cryptographic hashing",
    "Maintain structured audit logs of every upload, download, delete and access denial",
    "Evaluate encryption overhead, access-control correctness and integrity verification",
]

# [work, year, approach, reported result / limitation]  (checked against lit.ROWS; refs 3-5 are titles only)
LIT_ROWS = [
    [f"Nekouie et al. {N('base')} (base paper)", "2026", "Dual-attribute CP-ABE on an elliptic curve; ClinicalBERT attributes; Hyperledger Fabric ledger; health records",
     "Cost calculated from primitive timings (Raspberry Pi, HPE server); no product, no upload or audit latency"],
    [f"Sowmiya et al. {N('ijedr')}", "2026", "Client-side AES-256-GCM, RSA-2048 wrapping, 3-role RBAC, hash-chained audit log, MFA, 4 MB chunks",
     "+320 ms for 10 MB; 98.6% access accuracy; 50 of 50 tamper events caught; no hardware or method given"],
    [f"Thenmozhi et al. {N('geo')}", "2026", "Geocrypt: RSA file exchange gated by location and time; Ethereum smart contracts",
     "Abstract only; qualitative claim; location can be spoofed; no role model"],
    [f"Saginayev et al. {N('sagi')}", "2024", "Hybrid AES-256 + RSA cloud storage with SHA-256 integrity hash", "No benchmark in the abstract; no roles or audit trail"],
    [f"Backendal et al. {N('mega')}", "2023", "Attacks on MEGA's key hierarchy and file encryption", "Malleable key wrapping broke end-to-end storage; a design warning"],
    [f"Goh et al. {N('sirius')}; Kallahalla et al. {N('plutus')}", "2003", "SiRiUS and Plutus: per-file keys over untrusted network storage", "File-system clients, not web apps; no audit log in the abstracts"],
    [f"Goyal et al. {N('abe06')}; Bethencourt et al. {N('cpabe')}", "2006-07", "Key-policy and ciphertext-policy attribute-based encryption", "Pairing-based and heavier than symmetric keys"],
    [f"Ferraiolo et al. {N('rbac01')}", "2001", "NIST RBAC reference model: roles, permissions, constraints", "A model only; no key management or audit integrity"],
    [f"Schneier, Kelsey {N('audit99')}; Crosby, Wallach {N('tamper09')}", "1999-2009", "Hash-chained and Merkle-tree tamper-evident logs", "Need a separate verifier or published commitments"],
    [f"Batch survey {N('s2')}, {N('s3')}, {N('s4')}", "2026", "Three further titles (AES cloud storage, integrity and access control, chaotic-map image encryption)",
     "Not located online; titles only, nothing quoted"],
]

GAPS = [
    f"The base paper targets health records with attribute-based encryption, a ledger and ClinicalBERT; it has no file upload, role model or audit-log cost {N('base')}",
    f"The closest file-store paper reports +320 ms, 98.6% and 100% without hardware, per-operation timings or tamper methods {N('ijedr')}",
    "No surveyed work is a runnable general-purpose file store joining browser-side AES-256-GCM, per-recipient RSA key wrapping, server-enforced Owner/Editor/Viewer roles, a plaintext hash and a hash-chained audit log",
    f"Audit chains are proposed {N('audit99', 'tamper09')} but rarely tested against truncation, re-hashing or concurrent writers",
    "A measured study of every layer on one machine, reporting the weaknesses it finds, is missing",
]


def _ms(x):
    return f"{F.ms(x)} ms"


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    A = F.AES
    MB10, MB1, KB1 = "10485760", "1048576", "1024"
    e2e10, e2e1 = F.E2E[MB10], F.E2E[MB1]
    t_tail = F.TAMPER["audit_tamper"]["truncate_tail"]
    t_ok = [k for k, v in F.TAMPER["audit_tamper"].items() if v["detected"] == v["trials"]]
    t_all = len(F.TAMPER["audit_tamper"])

    # 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: CipherVault encrypts every file in the browser before upload, enforces Owner, Editor and Viewer roles on the server "
                "and keeps a tamper-evident audit log, and we have measured each of those layers.")

    # 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. The review covers why the problem matters and what is missing in current work, then what we built and what we have measured so far."

    # 3 abstract
    s = S[2]
    set_title(s, "Abstract: files the server can store but not read")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Organisations keep sensitive documents on infrastructure they do not control. The provider, insiders and network attackers are all exposure points, and without roles and logs a misuse cannot be limited or reconstructed.", K.ORANGE, "!"),
        ("Our idea", "Encrypt in the browser with AES-256-GCM, wrap each file key with the recipient's RSA-2048 public key, let the server enforce Owner, Editor and Viewer roles, store a SHA-256 of the plaintext, and chain every audit event with SHA-256.", K.BLUE, "i"),
        ("The deliverable", f"CipherVault: a FastAPI service with 24 endpoints and a React web app, with JWT and bcrypt sign-in. Ciphertext sits on disk, metadata in SQL. It is measured for encryption cost, access-control correctness and integrity.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: why readable storage is a risk, the client-side encryption plus roles plus audit design, and what is delivered and measured."

    # 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Elliptic-curve ciphertext-policy attribute-based encryption over user and data attributes; ClinicalBERT extracts data attributes from notes; a private Hyperledger Fabric ledger holds attribute public keys and policies; several key-generation centres avoid key escrow.", K.BLUE, "M"),
        ("Evaluation reported", "No dataset. Cost is calculated from operation counts and OpenSSL primitive timings on a Raspberry Pi 3 B+ and an HPE server: about 8,506 µs in total at 5 attributes on the Pi, against 8,812 to 71,323 µs for other schemes.", K.GREEN, "R"),
        ("Limits and our use", "It is not a file store: no upload, role model or audit-log cost. CipherVault adopts the trust model (provider sees ciphertext only), hash check, revocation and audit, and replaces ABE and the ledger with AES-GCM, RSA-OAEP and a hash chain.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The base paper is a different system, so we compare ideas, not scores: its timings are calculated in microseconds on other hardware, "
                "ours are measured end to end in milliseconds. The closest-architecture paper is Sowmiya et al., reference 2.")

    # 5 literature table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]] + LIT_ROWS
    K.table(s, 0.6, 1.1, [2.6, 0.8, 4.4, 4.33], rows, font_size=11, row_h=0.5, header_size=12, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey lists five items; items 3 to 5 in our numbering could not be located online, so only their titles appear. "
                "Other rows were checked against the papers or their abstract pages, and abstract-only sources are not credited with numbers.")

    # 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence: first measurements")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 6.6, 4.7, ["1 KiB", "1 MiB", "10 MiB"],
                [("Encrypted upload (browser)", [F.E2E[KB1]["encrypted_up"]["median"], e2e1["encrypted_up"]["median"], e2e10["encrypted_up"]["median"]]),
                 ("Plain upload, same endpoint", [F.E2E[KB1]["plain_up"]["median"], e2e1["plain_up"]["median"], e2e10["plain_up"]["median"]])],
                colors=[K.BLUE, K.ORANGE], number_format="0", font_size=11, max_value=260, legend=True)
    K.text(s, 0.6, 6.2, 6.6, 0.8, f"Median milliseconds per upload, Chrome to FastAPI over loopback ({F.ENV['cpu']}, shared with other jobs, so indicative).", size=11, color=K.TXT, italic=True)
    K.text(s, 7.5, 1.45, 5.25, 5.4, [
        ("**What we have measured**", {"space_after": 8}),
        (f"AES-256-GCM encrypts 10 MiB in {_ms(A[MB10]['encrypt']['median'])} (median); ciphertext is the plaintext plus 16 bytes.", {"bullet": True, "space_after": 8}),
        (f"Role matrix: {F.RBAC['matched']} of {F.RBAC['total']} operation cells behaved as specified; forged, expired and alg=none tokens all gave 401.", {"bullet": True, "space_after": 8}),
        (f"All {F.TAMPER['ciphertext_bitflip']['trials']} single-bit ciphertext flips were rejected by AES-GCM.", {"bullet": True, "space_after": 8}),
        (f"The closest paper reports +320 ms for 10 MB {N('ijedr')} but not its hardware, so no ranking is made.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = "These are our own early measurements, taken on one desktop; the full study follows in Review 2. They are not placed on one scale with published figures."

    # 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("User uploads the file", K.NAVY), ("Server stores it readable", K.NAVY), ("Provider-held key, if any", K.NAVY), ("Access by app checks only", K.NAVY), ("Editable plain log", K.NAVY)], size=14, gap=0.45)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Typical file storage",
           "The server sees the plaintext, so staff, backups and a breach all expose it.\nAny encryption at rest is keyed by the provider, not the user.\nA log that can be edited does not prove what happened.", K.BLUE, "1", body_size=18)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Research systems",
           f"Attribute-based encryption with a ledger {N('base')}, cryptographic storage prototypes {N('sirius', 'plutus')} and a client-side file-store design {N('ijedr')} each cover part of the need.\nNone is a runnable file store with its layers measured and its weak points reported.", K.ORANGE, "2", body_size=18)
    notes[7] = "Today most storage keeps files readable to the provider; research covers parts of the answer but not a measured, runnable whole."

    # 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Provider can read files", "No per-user control", "Tampering unnoticed", "No reconstruction"]
    bodies = [
        "Files stored without client-side encryption are readable by the storage provider.",
        "Without role-based control, any authenticated user may reach another user's files.",
        "Tampering goes undetected when no integrity check is performed on retrieval.",
        "Without audit logs, a security incident cannot be reconstructed afterwards.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=18)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"The base paper adds: attribute policies and a ledger that are heavy for a small deployment, and no file-level features {N('base')}.", size=14, color=K.INK, anchor="m")
    notes[8] = "The four limitations are the batch problem statement; the last line is from the base paper analysis."

    # 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=16, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as a general-purpose secure file store combining AES-256-GCM, RBAC, integrity verification and audit logging in one application. "
                "Gaps 2, 4 and 5 come from our reading of how the closest papers report their evaluation.")

    # 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", f"Files up to {F.UPLOAD_LIMIT_MB} MB, a password and, for sharing, a recipient and a role (Viewer or Editor)"),
                                 ("Output", "Ciphertext on disk, metadata in SQL, a decrypted and hash-verified download, and a verifiable audit chain"),
                                 ("Measures", "Encryption overhead, access-control correctness, tamper detection and verification time")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=15, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a store that encrypts before upload, enforces roles, verifies integrity and logs every action."

    # 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Slide 14 states what has been achieved and measured against these objectives, and where the system falls short.", size=13, color=K.TXT, italic=True)
    notes[11] = "The six objectives from the proposal. We say plainly later which are met and where the evidence has limits."

    # 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Register", f"The browser generates an RSA-2048 key pair and encrypts the private key under a PBKDF2 key ({F.PBKDF2_ITER:,} iterations). The server keeps a bcrypt password hash.", K.BLUE, "1"),
        ("Encrypt", "A fresh 256-bit AES key and 12-byte IV encrypt the file with AES-GCM; SHA-256 of the plaintext is computed; the key is wrapped with the owner's RSA public key.", K.GREEN, "2"),
        ("Share and enforce", "The owner re-wraps the file key for each recipient. The server checks Owner, Editor or Viewer on every request and logs denials.", K.ORANGE, "3"),
        ("Verify and audit", "On download the browser unwraps, decrypts and compares SHA-256. Every event joins a SHA-256 hash chain that anyone signed in can verify.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React web app", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, f"Dashboard, My files, Shared with me, file details with an access list, Audit log and a Security page. Backend on port {F.API_PORT}, frontend on {F.WEB_PORT}; SQLite by default, PostgreSQL through Docker.", size=14, color=K.TXT)
    notes[12] = "Four stages and the app that wraps them. No machine-learning model is involved: this is cryptography and access control."

    # 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, trim(F.diag("01_system_architecture.png"), "01"), 0.6, 1.2, 8.2, 5.9)
    K.text(s, 9.05, 1.5, 3.7, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"Browser: React app and Web Crypto API; the trusted crypto endpoint (port {F.WEB_PORT} in development).", {"bullet": True}),
        (f"FastAPI on port {F.API_PORT}: auth, users, files, sharing and audit routers.", {"bullet": True}),
        ("Services: role checks and the hash-chained audit log.", {"bullet": True}),
        ("SQL database with four tables; ciphertext only in storage/encrypted.", {"bullet": True}),
        ("The server never receives plaintext or private keys.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser encrypts, the API enforces roles and logs, the database holds metadata and the folder holds ciphertext only."

    # 14 outcomes
    s = S[13]
    set_title(s, "Expected outcomes and results so far")
    wipe(s)
    heads = ["Working platform", "Encryption cost", "Access control", "Integrity and audit", "Known gaps"]
    bodies = [
        f"24 endpoints; {F.PYTEST_PASSED} of {F.PYTEST_TOTAL} pytest cases, {F.E2E_PASSED} of {F.E2E_TOTAL} API checks and {F.UI_PASSED} of {F.UI_TOTAL} browser checks pass; uploads over {F.UPLOAD_LIMIT_MB} MB get HTTP 413.",
        f"10 MiB: AES-GCM {_ms(A[MB10]['encrypt']['median'])}, whole encrypted upload {_ms(e2e10['encrypted_up']['median'])} against {_ms(e2e10['plain_up']['median'])} unencrypted (medians).",
        f"{F.RBAC['matched']} of {F.RBAC['total']} role-matrix cells as specified; revoked users get HTTP {F.API['download_after_revoke_status']}; forged tokens get 401.",
        f"{len(t_ok)} of {t_all} audit-tamper types caught in every trial; {F.TAMPER['ciphertext_bitflip']['detected']} of {F.TAMPER['ciphertext_bitflip']['trials']} ciphertext bit flips rejected.",
        f"Tail truncation of the log went unseen ({t_tail['detected']} of {t_tail['trials']}); concurrent writers fork the chain; no MFA; revocation cannot recall a downloaded copy.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=14)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=16)
    notes[14] = ("Because the system is already built, this slide reports measured outcomes. Objectives 1 to 5 are implemented; objective 6 is measured, with the caveat "
                 "that timings come from a shared desktop and that two audit-log weaknesses were found and are documented, not fixed.")

    # 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = sorted(["base", "ijedr", "s2", "s3", "s4", "geo", "sagi", "mega", "sirius", "plutus", "abe06", "cpabe", "rbac01", "audit99", "tamper09"], key=lambda k: _IDX[k])

    def short(k):
        t = REFS[_IDX[k] - 1]
        for tail in (" (title as listed in the batch literature survey; not located online and not quoted)",):
            t = t.replace(tail, "")
        if k in ("s2", "s3", "s4"):
            t += " [title only; not located online]"
        return t

    half = (len(keys) + 1) // 2
    for col, ks in enumerate((keys[:half], keys[half:])):
        K.text(s, 0.6 + col * 6.15, 1.2, 5.95, 5.9, [(f"[{_IDX[k]}] {short(k)}", {"space_after": 4}) for k in ks], size=12, color=K.INK)
    notes[15] = "Selected references; the complete numbered list of 34 is in the project report."

    # 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why not attribute-based encryption as in the base paper (AES-GCM with per-recipient RSA wrapping is cheaper and simpler for a file store; the trade-off is no policy language); "
                 "why a hash chain and not a blockchain (single operator, no consensus needed); what the server can see (file name, size, MIME type, owner and timestamps).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
