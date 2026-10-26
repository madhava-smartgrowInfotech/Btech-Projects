"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        os.makedirs(os.path.dirname(BANNER), exist_ok=True)
        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, A. Nekouie, M. Vafaei Jahan, M. H. Moattar and R. Sheibani, whose work on access control for electronic health records "
        "framed the problem we set out to solve for general file storage, and G. Sowmiya, R. Nivas and E. Vikram, whose client-side-encryption architecture for cloud file storage is the closest design "
        "to ours. We are grateful to the W3C Web Cryptography working group and the browser vendors for the Web Crypto API on which the encryption in CipherVault runs, and to the maintainers of the open-source software "
        "used here: FastAPI, SQLAlchemy, the Python `cryptography`, bcrypt and PyJWT packages, React, Vite, Recharts, three.js and Playwright.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def _tamper_counts():
    """(classes, fully detected classes) for the audit-tamper study."""
    t = F.TAMPER["audit_tamper"]
    return len(t), sum(1 for v in t.values() if v["detected"] == v["trials"])


def abstract(d):
    aes10 = F.AES["10485760"]["encrypt"]["median"]
    up_enc, up_plain = F.E2E["10485760"]["encrypted_up"]["median"], F.E2E["10485760"]["plain_up"]["median"]
    flips = F.TAMPER["ciphertext_bitflip"]
    classes, caught = _tamper_counts()
    words = {8: "eight", 9: "nine"}
    d.p("Documents kept on storage the owner does not control can be read by the provider, its administrators and anyone who obtains a backup unless they are encrypted "
        f"before they leave the user. This project builds {F.PRODUCT}, a secure file-storage web application with a React front end and a FastAPI back end. The browser encrypts every file with a fresh "
        "AES-256-GCM key through the Web Crypto API, wraps that key with the RSA-2048-OAEP public key of the owner and of each recipient, and protects the user's private key with a key derived from the password "
        f"(PBKDF2-HMAC-SHA-256, {F.PBKDF2_ITER:,} iterations). The server holds only ciphertext, wrapped keys, bcrypt password hashes and metadata. It enforces Owner, Editor and Viewer roles on every request, "
        "the browser checks the AES-GCM tag and a stored SHA-256 digest on every download, and a hash-chained audit log records security events.")
    d.p(f"There is no machine-learning model and no dataset; the system was measured with fictional files on a shared Intel Core i5-12500, so timings are indicative. "
        f"In the browser, AES-256-GCM encryption of 10 MiB had a median of {F.f1(aes10)} ms, and a 10 MiB encrypted upload took {F.ms(up_enc)} ms against {F.ms(up_plain)} ms for a plain upload to the same endpoint. "
        f"An access matrix of 4 principals and 9 operations matched the rules in {F.RBAC['matched']} of {F.RBAC['total']} cells, all {flips['detected']:,} single-bit ciphertext corruptions were rejected, "
        f"and {F.PYTEST_PASSED} unit tests, {F.E2E_PASSED} end-to-end checks and an {F.UI_PASSED}-check browser walk passed.")
    d.p(f"The audit chain detected {words.get(caught, caught)} of {words.get(classes, classes)} tampering classes but not removal of the newest entries or a rewrite that recomputes later hashes; "
        "concurrent writers can fork it. Revocation cannot recall a key or ciphertext that was already downloaded, and files are limited to "
        f"{F.UPLOAD_LIMIT_MB} MB. The base paper concerns electronic health records, so its calculated timings are not compared with these measurements.")
    d.p("**Keywords:** secure file storage, client-side encryption, AES-256-GCM, RSA-OAEP, role-based access control, tamper-evident audit log, Web Crypto API, PBKDF2, SHA-256.")


ABBR = [
    ("ABE", "Attribute-Based Encryption"), ("AES", "Advanced Encryption Standard"), ("API", "Application Programming Interface"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CPU", "Central Processing Unit"), ("CRUD", "Create, Read, Update, Delete"), ("CSV", "Comma-Separated Values"), ("DFD", "Data Flow Diagram"),
    ("EHR", "Electronic Health Record"), ("ER", "Entity-Relationship"), ("FEK", "File Encryption Key"), ("FR", "Functional Requirement"),
    ("GCM", "Galois/Counter Mode"), ("HMAC", "Hash-based Message Authentication Code"), ("HTTP", "Hypertext Transfer Protocol"), ("IV", "Initialisation Vector"),
    ("JSON", "JavaScript Object Notation"), ("JWT", "JSON Web Token"), ("KEK", "Key-Encryption Key"), ("MFA", "Multi-Factor Authentication"),
    ("MIME", "Multipurpose Internet Mail Extensions"), ("NFR", "Non-Functional Requirement"), ("NIST", "National Institute of Standards and Technology"), ("OAEP", "Optimal Asymmetric Encryption Padding"),
    ("ORM", "Object-Relational Mapping"), ("PBKDF2", "Password-Based Key Derivation Function 2"), ("PKCS", "Public-Key Cryptography Standards"), ("RBAC", "Role-Based Access Control"),
    ("REST", "Representational State Transfer"), ("RSA", "Rivest-Shamir-Adleman"), ("SHA", "Secure Hash Algorithm"), ("SPA", "Single-Page Application"),
    ("SPKI", "Subject Public Key Info"), ("SQL", "Structured Query Language"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
    ("XSS", "Cross-Site Scripting"),
]


def abbreviations(d):
    """Two abbreviation/expansion pairs per row so the list takes half the space."""
    half = (len(ABBR) + 1) // 2
    rows = [list(ABBR[i]) + (list(ABBR[i + half]) if i + half < len(ABBR) else ["", ""]) for i in range(half)]
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion", "Abbreviation", "Expansion"], rows, widths=[0.95, 2.25, 0.95, 2.25], size=10)
