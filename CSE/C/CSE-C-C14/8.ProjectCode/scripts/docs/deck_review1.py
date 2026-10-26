"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import (BASE_PAPER_POINTS, EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIMITATIONS_EXISTING, LIT_TABLE, OBJECTIVES,
                  PROBLEM_STATEMENT, PROPOSED_APPROACH, SUPPORT_EVIDENCE)
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "AI-generated image detection, CNN ensemble, Grad-CAM, explainable AI, image forensics"
    cp.comments = ""
    cp.category = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 30)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: VeriVision AI is a local web application that says whether an image looks real or AI-generated, "
                "shows which regions drove the decision, and re-tests the verdict after compression, resizing and blur.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: detect, explain and re-test")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Generative models now produce images that people cannot reliably tell from photographs. Fakes drive misinformation and fraud, and a detector trained on one generator often fails on another.", K.ORANGE, "!"),
        ("Our idea", "Four ImageNet-pretrained CNN routes vote on every image. Their logits are fused with weights and two gating rules, Grad-CAM shows where the evidence lies, and a robustness check re-tests the image after JPEG, resize and blur.", K.BLUE, "i"),
        ("The deliverable", "VeriVision AI, a local web application: upload an image and read the verdict with confidence, per-route votes, heat maps and a consistency score. A decision aid, not proof of origin. Runs on a CPU.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=17)
    notes[3] = "Three points: the problem of realistic synthetic images, the idea of a heterogeneous ensemble with explanation and a robustness check, and what is delivered: a complete web application that runs locally."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "A one-class detector: it learns only from pristine faces plus synthesised pseudo-fakes, using hybrid-curvature representation learning on an EfficientNet-B4 backbone (19.3 M parameters).", K.BLUE, "M"),
        ("Data and reported results", "720 pristine FaceForensics++ videos; 200 epochs on two V100 GPUs (about 16 h). AUC 98.9% on FF++ and 77.0% on DFDC (lambda = 1 ablation); 74 frames per second on an RTX 3060.", K.GREEN, "R"),
        ("Relevance to VeriVision AI", "Shared aim: detection that holds up beyond the training source. Different scope: general images, supervised training, Grad-CAM, a robustness check and CPU-only use. The CTrue method itself is not implemented here.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=16)
    notes[4] = ("Read the citation. The base paper is a face-deepfake detector trained one-class; VeriVision AI is a supervised general-image detector, so numbers are never compared directly. "
                "The logit fusion and gating design follows HEDGE (arXiv 2604.03555), which is cited as a design reference.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for i in (0, 1, 2, 3, 4, 5, 6):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.35, [2.5, 0.7, 4.5, 4.43], rows, font_size=12, row_h=0.68, header_size=13, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: detection of generated images from GAN-era CNNs and frequency analysis, to CLIP-feature universal detectors, to heterogeneous ensembles, with a one-class face detector as the base paper. "
                "Figures are quoted only where they were read from the paper itself. The full list of references is in the report.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence from the literature")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 5.7, 5.0, ["CTrue AUC\nFF++ [1]", "CTrue AUC\nDFDC [1]", "HEDGE balanced\naccuracy, 5 sets [2]"],
                     [("Reported", [0.989, 0.770, 0.983])], colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.1)
    pt = gf.chart.plots[0].series[0].points[1]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 6.6, 1.45, 6.1, 4.9, [(e, {"bullet": True, "space_after": 9}) for e in SUPPORT_EVIDENCE], size=15, color=K.INK)
    K.text(s, 0.6, 6.45, 12.1, 0.5, "Different tasks, datasets and metrics (AUC vs balanced accuracy); the bars are context, not a comparison. The DFDC value is the lambda = 1 ablation.", size=12, color=K.TXT, italic=True)
    notes[6] = ("Two reasons for the design. A face detector that reaches AUC 98.9 percent in-domain falls to 77.0 percent on another dataset in the paper's own ablation, which is the generalisation problem. "
                "And HEDGE shows that heterogeneous routes with logit fusion and gating work on large GPU-trained backbones; we test the idea with small CNNs on a CPU.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Image circulates", K.NAVY), ("Human or tool inspects", K.NAVY), ("Verdict or score", K.NAVY), ("Decision / takedown", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Manual and metadata checks", "\n".join(EXISTING_SYSTEM[i] for i in (0, 3)) + "\nA mistaken verdict spreads misinformation or wrongly discredits a real image.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Automated detectors", "\n".join(EXISTING_SYSTEM[i] for i in (1, 2)), K.ORANGE, "2", body_size=17)
    notes[7] = "Today an analyst looks for visual artefacts or checks metadata; research detectors exist but usually return a bare score, and the strongest need GPUs or cloud services."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Human eyes fail", "Opaque score", "Generator shift", "Lost cues"]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], LIMITATIONS_EXISTING[i], C[i], str(i + 1), body_size=19)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: " + LIMITATIONS_EXISTING[4], size=15, color=K.INK, anchor="m")
    notes[8] = "Five limitations from the survey. Each maps to a design decision: voting routes with visible agreement, Grad-CAM, a robustness check, and local processing."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 0.9
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.82, g, size=20, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as pairing detection with Grad-CAM explanation and testing generalisation across unseen generators. "
                "We address gaps 2, 4, 5 and 6 directly and measure gap 3 (cross-domain behaviour) honestly rather than claim it is solved.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.9, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.65, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=21, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "An image: PNG, JPG, JPEG or WEBP, up to 10 MB"), ("Output", "Verdict, confidence, per-route votes, Grad-CAM heat maps and a robustness consistency score"), ("Record", "Measured test results with confidence intervals; every number traceable to a metrics file")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.65, 3.85, 1.9)
        K.text(s, x + 0.25, 4.77, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.25, 3.4, 1.25, b_, size=16, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a detector that identifies AI-generated images and explains its decision visually, with its limits stated."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.4, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=23, color=K.INK)
    notes[11] = "The six objectives of the proposal. The later reviews report which were met and which were not, including the cross-generator objective."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prepare", "Validate the upload, convert to RGB, resize per route (224 or 320 px) and normalise with ImageNet statistics.", K.BLUE, "1"),
        ("Classify", "Four routes: EfficientNet-B0 (two variants), ResNet50 with flip averaging, MobileNetV3-Large with focal loss.", K.GREEN, "2"),
        ("Fuse", "Weighted sum of logits (0.35 / 0.15 / 0.25 / 0.25), two deterministic gates, softmax, confidence wording.", K.ORANGE, "3"),
        ("Explain", "Grad-CAM heat maps on three routes and a re-test under JPEG, 2x resize and blur with a consistency score.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "S", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "Fused logits = 0.35 z(A1) + 0.15 z(A2) + 0.25 z(B) + 0.25 z(C), adjusted by Gate 1 and Gate 2", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.85, "Delivered as a Flask API (port 8414) and a React web app (port 5414); trained from ImageNet weights on a CIFAKE subset; no GPU needed; the verdict is a decision aid.", size=14, color=K.TXT)
    notes[12] = "Four stages from upload to explained verdict. " + PROPOSED_APPROACH[1] + ". The weights and gate thresholds live in one configuration file."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.4, 5.6)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite web app (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: Flask (port {F.API_PORT}); PNG, JPG, JPEG or WEBP up to {F.UPLOAD_LIMIT_MB} MB.", {"bullet": True}),
        ("Inference: PyTorch routes, logit fusion, Grad-CAM and robustness re-check.", {"bullet": True}),
        ("Data: model files, metadata and metric JSON on the local disk; no database, nothing leaves the machine.", {"bullet": True}),
    ], size=17, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser uploads the image, the Flask layer validates it, the inference tier runs four routes and the explanation, and local files hold models and metrics."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working application", "Measured performance", "Honest robustness", "Inspectable evidence", "Human judgement"]
    for i, (h, b_) in enumerate(zip(heads, EXPECTED_OUTCOMES)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.55, h, b_, C[i % 4], str(i + 1), body_size=17)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=18)
    notes[14] = "Expected outcomes and contributions. The later reviews report the measured results, including where the models fail."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 2, 3, 8, 12, 13, 19, 20, 21, 26]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 4}) for n in idx], size=14, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why four routes instead of one (disagreement is information, and HEDGE motivates heterogeneous routes); why a CPU-trained subset "
                 "(hardware); whether Grad-CAM proves why the model decided (it is a qualitative aid).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "VeriVision")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
