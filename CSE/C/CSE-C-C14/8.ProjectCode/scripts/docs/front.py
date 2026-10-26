"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations (VeriVision AI)."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs


def cover(d):
    d.spacer(14)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: a local web application that classifies an image as real or AI-generated with four CNN routes and explains the decision with Grad-CAM", align="center", size=12, italic=True, space_after=10)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=8)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=8)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the papers that shaped the design, the creators of the CIFAKE dataset, and the maintainers of the open-source software used in this project "
        "(PyTorch, torchvision, pytorch-grad-cam, scikit-learn, Flask, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    e, ci = F.ENS, F.CI["ensemble"]
    xd = F.XD["ensemble"]
    rob = F.ROB
    d.h1("ABSTRACT")
    d.p("Generative models now produce images that people cannot reliably tell from photographs, and detectors trained on one generator often fail on another. "
        f"This project presents {F.PRODUCT}, a local web application that classifies an uploaded image as real or AI-generated and explains the decision. Four ImageNet-pretrained CNN routes "
        "(two EfficientNet-B0 variants, ResNet50 and MobileNetV3-Large) are fused in logit space with configurable weights and two gating rules; Grad-CAM heat maps show where three of the routes looked, "
        "and a robustness check re-tests the image under JPEG compression, resizing and blur. The system is served by a Flask API and a React interface and runs on a CPU.")
    d.p(f"The routes were fine-tuned on a class-balanced subset of CIFAKE (6,000 training, 1,200 validation and {F.N_TEST:,} test images, 32 x 32 pixels, one generator). On the held-out test images the ensemble reaches "
        f"{F.pct(e['accuracy'], 2)} accuracy (95% interval {F.pct(ci[0])} to {F.pct(ci[1])}), F1 {F.f3(e['f1'])} and ROC-AUC {F.f4(e['roc_auc'])}; the best single route scores {F.pct(F.BR['a1']['accuracy'], 2)}, "
        "so the ensemble's gain in accuracy is not statistically distinguishable, and the two gates did not improve the result. The model leans toward the AI-generated class "
        f"(false-positive rate {F.pct(e['fpr'])}). Resizing and blur reduce accuracy to about 58-61%, and on 600 full-resolution images from another source the ensemble is at chance "
        f"({F.pct(xd['accuracy'])}, ROC-AUC {F.f3(xd['roc_auc'])}), so cross-generator generalisation was measured and not achieved. "
        f"A verdict takes a median of {F.LATENCY['predict_s']:.2f} s and a full analysis {F.LATENCY['analyze_s']:.1f} s on a CPU-only PC. "
        "Training used 6% of the available training images and a single run per route; the explanations were judged qualitatively and no user study was carried out. "
        "The base paper, a one-class face-deepfake detector, uses different data and metrics, so its results are not directly comparable.")
    d.p("**Keywords:** artificial intelligence, convolutional neural network, AI-generated image detection, explainable AI, Grad-CAM, image forensics, ensemble learning, robustness, Flask, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CE", "Cross-Entropy"),
    ("CI", "Confidence Interval"), ("CNN", "Convolutional Neural Network"), ("CPU", "Central Processing Unit"), ("DCT", "Discrete Cosine Transform"),
    ("DFD", "Data Flow Diagram"), ("ECE", "Expected Calibration Error"), ("ER", "Entity-Relationship"), ("FNR", "False-Negative Rate"),
    ("FPR", "False-Positive Rate"), ("GAN", "Generative Adversarial Network"), ("GPU", "Graphics Processing Unit"), ("Grad-CAM", "Gradient-weighted Class Activation Mapping"),
    ("HTTP", "Hypertext Transfer Protocol"), ("JPEG", "Joint Photographic Experts Group"), ("JSON", "JavaScript Object Notation"), ("ML", "Machine Learning"),
    ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("TTA", "Test-Time Augmentation"), ("UI", "User Interface"),
    ("UML", "Unified Modeling Language"), ("XAI", "Explainable Artificial Intelligence"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
