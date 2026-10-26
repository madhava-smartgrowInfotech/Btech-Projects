"""
19.IEEE_Paper.docx: a two-column IEEE-style conference paper built from docs/REPORT_FACTS.md-level facts
(experiments/*.json), the diagrams and the screenshots. Format per DOCS_PROMPT.md section 7.

    python -m scripts.docs.paper
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs, set_font
from .refs import REFS

FONT = "Times New Roman"
COL_W = 3.4  # inches


class Paper:
    def __init__(self):
        self.doc = Document()
        self.order = [1]         # global reference numbers in order of first citation
        self.fig_no = 0
        self.tab_no = 0
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin, s.bottom_margin = Inches(0.75), Inches(1.0)
        s.left_margin = s.right_margin = Inches(0.69)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0

    # ---------------------------------------------------------------- citations
    def cite(self, *globals_):
        nums = []
        for g in globals_:
            if g not in self.order:
                self.order.append(g)
            nums.append(self.order.index(g) + 1)
        nums.sort()
        return "[" + "], [".join(str(n) for n in nums) + "]"

    def expand(self, text):
        """Replace {r:1,5} by [n], [m] numbered by first appearance."""
        def sub(m):
            return self.cite(*[int(x) for x in m.group(1).split(",")])
        return re.sub(r"\{r:([\d,]+)\}", sub, text)

    # ---------------------------------------------------------------- blocks
    def para(self, text, size=10, align="j", indent=0.2, bold=None, italic=None, after=0, before=0, keep_next=False, space_line=1.0):
        par = self.doc.add_paragraph()
        par.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        pf = par.paragraph_format
        pf.first_line_indent = Inches(indent) if indent else None
        pf.space_after, pf.space_before = Pt(after), Pt(before)
        pf.keep_with_next = keep_next
        pf.line_spacing = space_line
        add_runs(par, self.expand(text), size, bold, italic)
        return par

    def eq(self, text, size=10):
        """Centred, numbered display equation in plain text (italic)."""
        self.eq_no = getattr(self, "eq_no", 0) + 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.keep_together = True
        import re as _re
        for part in _re.split(r"(_[A-Za-z0-9]+)", text):
            if not part:
                continue
            sub = part.startswith("_") and len(part) > 1
            r = par.add_run(part[1:] if sub else part)
            set_font(r, FONT, size, False, True)
            if sub:
                r.font.subscript = True
        r2 = par.add_run(f"        ({self.eq_no})")
        set_font(r2, FONT, size)

    def heading(self, roman, text):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.keep_with_next = True
        label = f"{roman}. {text.upper()}" if roman else text.upper()
        r = par.add_run(label)
        set_font(r, FONT, 10, True)

    def sub(self, letter, text):
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(f"{letter}. {text}")
        set_font(r, FONT, 10, False, True)

    def figure(self, path, caption, width=COL_W):
        self.fig_no += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(caption) > 70 else WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(5)
        add_runs(cap, f"Fig. {self.fig_no}. {self.expand(caption)}", 8)

    def table(self, caption, header, rows, widths, size=7.5, align=None):
        self.tab_no += 1
        roman = ["I", "II", "III", "IV", "V", "VI"][self.tab_no - 1]
        c1 = self.doc.add_paragraph()
        c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c1.paragraph_format.space_before = Pt(4)
        c1.paragraph_format.keep_with_next = True
        add_runs(c1, f"TABLE {roman}", 8)
        c2 = self.doc.add_paragraph()
        c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c2.paragraph_format.space_after = Pt(2)
        c2.paragraph_format.keep_with_next = True
        add_runs(c2, self.expand(caption).upper(), 8)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, h, size, True)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (align is None or align[i] == "l") else WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, self.expand(str(v)), size)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
            cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(cs)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        for r in sp.runs:
            pass
        set_font(sp.add_run(""), FONT, 4)

    def references(self):
        self.heading("", "References")
        for i, g in enumerate(self.order, 1):
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = par.paragraph_format
            pf.left_indent = Inches(0.3)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(1)
            add_runs(par, f"[{i}]  {REFS[g - 1]}", 8)

    # ---------------------------------------------------------------- layout
    def begin_two_columns(self):
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.xpath("./w:cols")
        if cols:
            c = cols[0]
        else:
            c = OxmlElement("w:cols")
            sec._sectPr.append(c)
        c.set(qn("w:num"), "2")
        c.set(qn("w:space"), "360")  # 0.25 in

    def title_block(self, title, authors):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(8)
        r = par.add_run(title)
        set_font(r, FONT, 24, False)
        n = len(authors)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, a in enumerate(authors):
            c = t.rows[0].cells[i]
            c.width = Inches(7.12 / n)
            c.text = ""
            first = True
            for j, line in enumerate(a):
                p = c.paragraphs[0] if first else c.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                add_runs(p, line, 9, j == 0)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def abstract(self, text, keywords):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.space_after = Pt(4)
        r = par.add_run("Abstract---")
        set_font(r, FONT, 9, True, True)
        add_runs(par, text, 9)
        par2 = self.doc.add_paragraph()
        par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par2.paragraph_format.space_after = Pt(4)
        r = par2.add_run("Index Terms---")
        set_font(r, FONT, 9, True, True)
        add_runs(par2, ", ".join(sorted(keywords, key=str.lower)) + ".", 9)

    def save(self, path):
        self.doc.save(path)
        return path
