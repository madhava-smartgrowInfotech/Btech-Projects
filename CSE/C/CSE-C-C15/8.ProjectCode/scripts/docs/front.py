"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the related work cited in this report, the Canadian Institute for Cybersecurity for the CICIDS2017 dataset used in the "
        "real-data evaluation, and the maintainers of the open-source software used in this project (Python, Flask, scikit-learn, pandas, NumPy, Scapy, psutil, SQLite, React, Vite, Recharts, "
        "matplotlib, pytest and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    k = F.K()
    d.h1("ABSTRACT")
    d.p("Distributed denial-of-service (DDoS) attacks flood a service with traffic from many sources and leave legitimate users without access. Detection that is not followed by action leaves "
        "the service down, while wrongly blocking legitimate sources is itself a denial of service. This project builds " + F.PRODUCT + ", a working real-time detection and adaptive mitigation "
        "system: Scapy (or a built-in simulator) feeds rolling traffic windows that are reduced to ten flow features and classified by a 200-tree Random Forest; an explainable rule layer scores how "
        "distributed the traffic is across sources and destinations (the carpet-bombing signature) and assigns a Low, Medium, High or Critical severity with its contributing factors; High and Critical "
        "windows quarantine the three busiest sources for a limited time, after which they recover automatically. A Flask service with twelve JSON endpoints, a SQLite history and a nine-page React "
        "console present the result. The work follows the resource-aware ensemble framework of Trigui et al. (IEEE Access, 2026) and adds the closed detection-to-mitigation loop that paper leaves as future work.")
    d.p(f"Measured honestly, the shipped model was trained on a labelled synthetic fallback because CICIDS2017 was unreachable when the product was built: it scores {k['ship_acc'] * 100:.2f} % accuracy on its own held-out split, "
        f"but only {k['shipreal_acc'] * 100:.1f} % accuracy ({k['shipreal_rec'] * 100:.1f} % recall, {k['shipreal_fpr'] * 100:.1f} % false-positive rate) on {k['real_rows']:,} real CICIDS2017 DDoS-day flows. Retrained on those real flows with the same features and "
        f"hyper-parameters, the forest reaches {k['rr_acc'][0] * 100:.2f} % accuracy over five seeds ({k['blk_acc'] * 100:.2f} % on a chronological block split), comparable with the 100 % reported by the base paper on different datasets. "
        f"The severity layer cuts false positives on benign windows from {k['w5_fpr_ml'] * 100:.1f} % to {k['w5_fpr_act'] * 100:.2f} % at 5-second windows, but window length matters: the model was calibrated on 10-second windows. "
        f"{F.PYTEST_PASSED} automated tests and {F.SMOKE_PASSED} smoke checks pass; live capture could not be run end to end here because Npcap is not installed, and mitigation is an application-level quarantine that never changes the operating-system firewall.")
    d.p("**Keywords:** DDoS detection, network security, Random Forest, CICIDS2017, Scapy packet capture, adaptive mitigation, carpet-bombing, Flask security dashboard.")


ABBR = [
    ("API", "Application Programming Interface"), ("AUC", "Area Under the ROC Curve"), ("CICIDS2017", "Canadian Institute for Cybersecurity Intrusion Detection System dataset 2017"),
    ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("DDoS", "Distributed Denial of Service"), ("DFD", "Data Flow Diagram"), ("DoS", "Denial of Service"),
    ("ER", "Entity-Relationship"), ("FN", "False Negative"), ("FP", "False Positive"), ("FPR", "False-Positive Rate"), ("HTTP", "Hypertext Transfer Protocol"),
    ("ICMP", "Internet Control Message Protocol"), ("IDS", "Intrusion Detection System"), ("IP", "Internet Protocol"), ("JSON", "JavaScript Object Notation"),
    ("kNN", "k-Nearest Neighbours"), ("LLM", "Large Language Model"), ("ML", "Machine Learning"), ("Npcap", "Packet-capture driver for Windows"), ("RAG", "Retrieval-Augmented Generation"),
    ("REST", "Representational State Transfer"), ("RF", "Random Forest"), ("ROC", "Receiver Operating Characteristic"), ("SDN", "Software-Defined Networking"),
    ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("SVM", "Support Vector Machine"), ("TCP", "Transmission Control Protocol"),
    ("TN", "True Negative"), ("TP", "True Positive"), ("UDP", "User Datagram Protocol"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
