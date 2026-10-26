"""Builds 19.IEEE_Paper.docx and 19.IEEE_Paper.pdf (two-column, 6 pages) for StoryVox.

    python -m scripts.docs.build_paper            # from 8.ProjectCode, system python (needs Word for the PDF)
    python -m scripts.docs.build_paper --no-pdf   # docx only
"""
import argparse
import os
import re
import sys

from . import facts as F
from . import paper_text
from .comlock import com_lock
from .paper_finish import export


def audit(docx_path):
    """Checks every citation is listed, every reference cited, and no template/donor leftovers."""
    from docx import Document

    d = Document(docx_path)
    body = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                body.append(c.text)
    text = "\n".join(body)
    refs_start = text.index("REFERENCES")
    main, refs = text[:refs_start], text[refs_start:]
    cited = set()
    for m in re.finditer(r"\[(\d+(?:\], \[\d+)*)\]", main):
        cited.update(int(x) for x in re.findall(r"\d+", m.group(1)))
    listed = {int(m.group(1)) for m in re.finditer(r"^\[(\d+)\]\s", refs, re.M)}
    print("citations:", len(cited), "references:", len(listed), "missing in list:", sorted(cited - listed), "uncited:", sorted(listed - cited))
    bad = [w for w in ("skincare", "Dayananda", "recommender system", "Fakeddit", "MCNN", "bank transaction", "social media manipulation", "[STUDENT NAME", "[ROLL NO",
                       "TrendCast", "H&M", "XGBoost", "FlowSense", "UNSW", "intrusion") if w.lower() in text.lower()]
    print("leftovers:", bad)
    words = len(re.findall(r"\S+", text))
    print("words (body+refs+captions+tables):", words)
    return cited, listed, bad, words


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    docx = os.path.join(F.BATCH, "19.IEEE_Paper.docx")
    pdf = os.path.join(F.BATCH, "19.IEEE_Paper.pdf")
    paper_text.build(docx)
    audit(docx)
    if not a.no_pdf:
        with com_lock():
            pages, words = export(docx, pdf)
        import pymupdf

        n = len(pymupdf.open(pdf))
        last = pymupdf.open(pdf)[n - 1]
        ytxt = max((b[3] for b in last.get_text("blocks")), default=0)
        print(f"Word pages {pages}, Word words {words}, PDF pages {n}, last page text ends at {ytxt:.0f} of {last.rect.height:.0f} pt")
    sys.exit(0)
