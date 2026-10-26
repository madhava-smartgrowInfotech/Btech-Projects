"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (letterhead, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from . import lit
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "storyvox_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def bl(items, space=8):
    return [(t, {"bullet": True, "space_after": space}) for t in items]


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "story generation, multilingual text-to-speech, MMS-TTS, Qwen3, Whisper"
    notes = {}

    LG = F.LANGS
    SS = F.STORY_SUMMARY
    A, Bc, C = SS["A_app_raw_prompt"], SS["B_chat_template"], SS["C_local_composer"]
    LS = F.LENGTH_SCALING
    EP = F.API["endpoints"]
    GEN = F.API["generate"]
    ASR = sorted(F.STORY_ASR["rows"], key=lambda r: r["story_id"])
    en_rows = [r for r in ASR if r["language"] == "en"]
    en_wer = sum(r["wer_norm"] for r in en_rows) / len(en_rows)
    en_cer = sum(r["cer_norm"] for r in en_rows) / len(en_rows)
    n_lm = lambda rows: sum(1 for r in rows if r["language_match"])  # noqa: E731
    condA, condB = F.STORY["conditions"]["A_app_raw_prompt"], F.STORY["conditions"]["B_chat_template"]
    nA, nB = n_lm(condA), n_lm(condB)
    nonen = [r for r in ASR if r["language"] != "en"]
    n_en_text = sum(1 for r in nonen if r.get("text_language") == "en")

    for sl in S:  # footers on every slide except the title
        for sh_ in sl.shapes:
            if sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 28)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and StoryVox: a local pipeline from a story idea to a story and its narration in 15 languages, measured honestly on a laptop CPU."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Data & Models\nDecision Workflow, Interaction & Data Design\nImplementation\nResults: Narration, Language Match, Parameter Analysis\nPerformance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Future Scope & References", 16)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, bl([
        "A learner or listener who wants a story in their own language, spoken aloud, has to chain a writing tool, a translator and a voice. StoryVox joins the steps in one local application: a story idea goes to a small language model (Qwen3-0.6B), the text goes to a per-language neural voice (MMS-TTS), and a recogniser (Whisper-small) checks the round trip.",
        f"Narration is the strong part: on 10 FLEURS sentences per language, the 12 languages narrated by MMS run at a mean real-time factor of {F.MMS_MEAN_RTF:.3f} on a CPU, with a mean normalised CER of {F.pct(F.MMS_MEAN_CER)} (best {F.LANG_NAMES[F.BEST_CER]} {F.pct(LG[F.BEST_CER]['cer_norm_corpus'])}, worst {F.LANG_NAMES[F.WORST_CER]} {F.pct(LG[F.WORST_CER]['cer_norm_corpus'])}).",
        f"Generation is the weak part: with the shipped prompt the 0.6B model answered in the requested language for only {nA} of {len(condA)} prompts (the chat-template experiment: {nB}); all {n_en_text} of {len(nonen)} non-English stories made through the API came out in English and were read by the foreign voice.",
        "Translation, voice or speed selection and user ratings, listed in the objectives, are not implemented. The app's own evaluation shows no WER/CER for story-length audio.",
        f"{F.PYTEST_PASSED} pytest tests pass; 11 real screenshots; every number comes from runs on the project's own code.",
    ], 9), size=16)
    notes[3] = "Give the product, then the honest result: narration works and is fast enough; the small story model mostly ignores the requested language; three objectives are not implemented."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: SAM Translator, cross-lingual text-to-speech translation"],
            ["Method", "mBART text-to-unit model with a sparse mixture of 8 experts (top-2 routing), self-paced learning (SPL) and data-augmented self-training (DAST); HuBERT units + Grad-TTS vocoder"],
            ["Data", "CVSS-C (21 source languages to English speech) and CoVoST-2; BLEU of the ASR transcript"],
            ["Cost", "611M-parameter baseline, 1.7B-parameter MoE; one A100 80 GB, about 1.5-2 days of training"],
            ["Reported", "SPL improved 18 of 21 languages; relative BLEU gains about +2-3% (high-resource) and +7-10% (low-resource); naturalness 4.12 vs 4.15 (10 listeners)"],
            ["Stated limits", "English is the only target; it translates but does not write stories or take style / length controls; DAST with 15K pairs lowered some languages"]]
    ftable(s, 0.55, 2.6, [1.7, 10.5], rows, size=12, row_h=0.7)
    notes[4] = "Tran and Sakti, IEEE Access 2026, the base paper named in the batch documents. StoryVox shares the goal of multilingual speech output and the ASR round trip as a check; it trains nothing."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    keys = ["omni", "multigen", "kahaani", "vits", "mms", "whisper", "qwen3", "tinystories"]
    short = lambda t, n: t if len(t) <= n else t[: n - 3].rsplit(" ", 1)[0] + "..."  # noqa: E731
    rows = [["Work", "Approach", "Limitation relevant here"]]
    for r in lit.ROWS:
        if r[0] in keys:
            rows.append([f"{r[1]} {N(r[0])}", short(r[2], 120), short(r[5], 120)])
    ftable(s, 0.55, 1.75, [2.5, 5.0, 4.7], rows, size=10.5, row_h=0.55)
    notes[5] = "Eight of the fifteen surveyed works. Speech models, an ASR judge and a small story model are all published and reusable; none joins them for story narration and measures the chain."

    # ---- 6 research gap + problem statement
    s = S[5]
    set_title(s, "Research Gap & Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 4.9, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Research gap", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.2, bl([
        "Speech-translation papers report BLEU on benchmark corpora; TTS papers report MOS or CER on read sentences.",
        "LLM + speech systems for children exist, but are not published as a locally runnable tool with a measured pipeline.",
        "Nothing tells a user whether a story requested in Hindi or Korean comes out in that language, is intelligible when spoken, and finishes in practical time on a laptop CPU.",
    ], 10), size=16)
    fbox(s, 6.8, 1.8, 6.0, 4.9, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Problem statement", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.2, bl([
        "Writing, translating and voicing a story are separate tools, each with its own account, cost and language coverage.",
        "A learner or listener needs one system: idea in, story in the chosen language out, spoken, saved and checkable.",
        "It must run without a cloud account, keep text and audio on the machine, and report honest quality numbers.",
    ], 10), size=16)
    notes[6] = "The gap is the join and the measurement, not any single model."

    # ---- 7 objectives and status
    s = S[6]
    set_title(s, "Objectives and Their Status")
    wipe(s)
    rows = [["#", "Objective (batch documents)", "Status", "What exists"],
            ["1", "Accept a user prompt", "Achieved", "Studio page, 2,000-character limit, validated"],
            ["2", "Story with length and reading-level controls", "Partly", "3 lengths and 8 styles; no reading-level control"],
            ["3", "Translate into target languages", "Not implemented", "No translation module; the language is requested in the prompt"],
            ["4", "Convert the story to speech", "Achieved (12 of 15)", "MMS-TTS for 12 languages; it, zh, ja use an OS voice"],
            ["5", "Playback with adjustable voice, speed, language", "Partly", "Browser player and WAV download; no voice or speed choice"],
            ["6", "Evaluate coherence, translation, naturalness via user rating", "Partly", "Automated WER/CER, similarity, ROUGE-L, RTF; no user rating"]]
    ftable(s, 0.55, 1.8, [0.5, 4.6, 2.3, 4.8], rows, size=12.5, row_h=0.62, align=["c", "l", "l", "l"])
    ftext(s, 0.55, 6.25, 12.2, 0.6, "Also not implemented: reference-audio upload / voice cloning. Status is stated against the code, not the abstract.", size=12.5, color=K.FNAVY, italic=True)
    notes[7] = "Be explicit: translation, voice/speed choice and user rating were objectives but are not in the product."

    # ---- 8 proposed system
    s = S[7]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Story pipeline (Studio)", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Idea + language, style, length", K.FNAVY), ("Prompt builder", K.FBLUE), ("Qwen3-0.6B", K.GREEN), ("Parse title / body", K.ORANGE), ("MMS-TTS narration", K.PURPLE), ("Save + play", K.FNAVY)], size=12, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Evaluation pipeline (Evaluation page)", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Fixed prompts", K.FNAVY), ("Story + narration", K.FBLUE), ("Whisper-small", K.GREEN), ("WER / CER", K.ORANGE), ("Similarity, ROUGE-L, RTF", K.PURPLE), ("Saved JSON + CSV", K.FNAVY)], size=12, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, bl([
        "Two fallbacks keep the app working: an offline composer replaces a model output under 20 words or a failed pipeline; an OS voice replaces a missing speech checkpoint.",
        "Text, audio and the SQLite library stay on the machine; models are downloaded once from Hugging Face.",
        "No model is trained or fine-tuned: StoryVox is a cascade of pretrained models behind a web application.",
    ], 7), size=15)
    notes[8] = "Two chains. The fallbacks are why a missing checkpoint produced speech in the wrong voice instead of an error."

    # ---- 9 architecture
    s = S[8]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.2, 5.1)
    ftext(s, 7.95, 1.95, 4.9, 4.9, bl([
        f"Web app: React 19 + Vite, port {F.WEB_PORT}; 5 pages (Landing, Studio, History, Story detail, Evaluation)",
        f"API: FastAPI + Uvicorn, port {F.API_PORT}; 11 routes plus a single-page-app fallback; no sign-in (single-user tool)",
        "Services: story generator, TTS providers, ASR, evaluation",
        "Storage: SQLite (stories, audio_outputs, evaluations); WAV files and evaluation JSON / CSV on disk",
        "Models run in-process on the CPU after a one-time download",
    ], 8), size=14)
    notes[9] = "Four tiers on one machine."

    # ---- 10 methodology
    s = S[9]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, bl([
        "Prompt: instruction text built from idea, language, style and length (220 / 420 / 700 new tokens); raw text, no chat template.",
        "Sampling: temperature 0.8, top-p 0.9.",
        "Narration: text split into chunks of up to 380 characters, one VITS call each, peak-normalised to 0.9, joined with 0.25 s pauses, 16 kHz.",
        "Check: Whisper-small transcript against the source text (WER, CER); RTF = synthesis time / audio duration.",
    ], 8), size=13.5)
    notes[10] = "Method in one picture. There is no training configuration: these are inference settings."

    # ---- 11 dataset, models
    s = S[10]
    set_title(s, "Data, Models and Tools")
    wipe(s)
    ftile(s, 0.55, 1.8, 2.1, 1.0, "15", "languages x 10 FLEURS sentences", K.FBLUE, 24, 11)
    ftile(s, 2.8, 1.8, 2.1, 1.0, "17 x 3", "story prompts x conditions", K.GREEN, 24, 11)
    ftile(s, 5.05, 1.8, 2.1, 1.0, "0", "models trained or tuned", K.ORANGE, 24, 11)
    rows = [["Item", "Detail"],
            ["Story model", "Qwen/Qwen3-0.6B (transformers pipeline) + offline composer"],
            ["Speech models", "facebook/mms-tts-* (VITS), 12 languages; OS voice for it, zh, ja"],
            ["Recogniser", "openai/whisper-small; jiwer WER / CER"],
            ["Similarity", "multilingual MiniLM (TF-IDF proxy offline); ROUGE-L"],
            ["Test data", "FLEURS test split (CC BY 4.0), first 10 length-filtered sentences per language"],
            ["Machine", "i5-12500, 12 threads, 15.7 GB RAM, CPU only; Python 3.11.9, torch 2.14.1"]]
    ftable(s, 0.55, 3.0, [1.6, 5.4], rows, size=11, row_h=0.5)
    K.picture(s, F.diag("12_model_architecture.png"), 7.95, 1.8, 4.9, 3.6)
    ftext(s, 7.95, 5.6, 4.9, 0.8, "No training data exist in this project; every model is used as published.", size=12, color=K.FNAVY, italic=True)
    notes[11] = "FLEURS is read speech of Wikipedia-style sentences, so it measures narration clarity, not story quality."

    # ---- 12 decision workflow
    s = S[11]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_algorithm.png"), 0.45, 1.75, 6.1, 4.7)
    K.picture(s, F.diag("11_flowchart_tts_evaluation.png"), 6.8, 1.75, 6.1, 4.7)
    ftext(s, 0.55, 6.5, 6.0, 0.5, "Story path: model output, parse, fall back to the composer if under 20 words.", size=12, color=K.FNAVY, bold=True)
    ftext(s, 6.9, 6.5, 6.0, 0.5, "Narration and evaluation path: MMS, then OS voice; Whisper scores the audio.", size=12, color=K.FNAVY, bold=True)
    notes[12] = "The two decisions that shape results: when the composer replaces the model, and when the OS voice replaces MMS."

    # ---- 13 interaction
    s = S[12]
    set_title(s, "User - System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_flow.png"), 0.55, 1.8, 6.4, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[13] = "Sequence of one generate-and-narrate request, and the use cases the product supports. There is no use case for translation, voice choice or rating."

    # ---- 14 database
    s = S[13]
    set_title(s, "Database Design")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.2, 4.2)
    ftext(s, 0.55, 6.1, 7.2, 0.8, "stories 1-to-many audio_outputs; evaluations stand alone. SQLite foreign keys are declared but not switched on, so deletion removes the audio rows in code.", size=12, color=K.FNAVY, italic=True)
    ftext(s, 8.0, 1.9, 4.8, 4.9, bl([
        "stories: prompt, title, text, language, style, length, word count, generation time",
        "audio_outputs: WAV path, duration, sample rate, TTS model, synthesis time, RTF",
        "evaluations: type, language, sample count, metrics JSON, created time",
        "Files: data/audio/*.wav and data/evaluations/*.json + CSV",
    ], 8), size=14)
    notes[14] = "Three tables. The ER diagram mirrors app/models."

    # ---- 15 implementation highlights
    s = S[14]
    set_title(s, "Implementation Highlights")
    wipe(s)
    ftile(s, 0.55, 1.8, 2.9, 1.0, "11 routes", "REST API + SPA fallback", K.FBLUE, 24, 12)
    ftile(s, 3.65, 1.8, 2.9, 1.0, f"{F.PYTEST_PASSED} tests", f"{F.PYTEST_API} API + {F.PYTEST_UNIT} unit, all pass", K.GREEN, 24, 12)
    ftile(s, 6.75, 1.8, 2.9, 1.0, "15 / 8 / 3", "languages / styles / lengths", K.ORANGE, 24, 12)
    ftile(s, 9.85, 1.8, 2.9, 1.0, "CPU only", "no GPU needed", K.PURPLE, 24, 12)
    ftext(s, 0.55, 3.1, 6.0, 3.8, bl([
        "TTSProvider abstract class: MMS (VITS) first, OS voice second; each can be replaced.",
        "chunk_text keeps sentences whole, splits on '.', '!', '?' and the Devanagari danda.",
        "Silence check: a narration that is silent is rejected, never saved.",
        "Models are loaded once and cached; the first request pays the load.",
    ], 7), size=14)
    ftext(s, 6.8, 3.1, 6.0, 3.8, bl([
        "Validation: language, style and length are checked; empty or over-long prompts return HTTP 422.",
        "Audio route resolves paths inside the audio folder only (escape attempts rejected).",
        "Evaluation reloads the similarity model on every call and reads audio with a plain Whisper call, which rejects clips over 30 s.",
        "ffmpeg and uroman are needed but not in requirements.txt.",
    ], 7), size=14)
    notes[15] = "Strengths on the left, defects found while measuring on the right."

    # ---- 16 results: narration
    s = S[15]
    set_title(s, "Results - Narration Speed and Intelligibility")
    wipe(s)
    ftile(s, 0.55, 1.75, 2.9, 0.95, f"{F.MMS_MEAN_RTF:.3f}", "mean RTF, 12 MMS languages", K.FBLUE, 24, 11)
    ftile(s, 3.65, 1.75, 2.9, 0.95, F.pct(F.MMS_MEAN_CER), "mean normalised CER", K.GREEN, 24, 11)
    ftile(s, 6.75, 1.75, 2.9, 0.95, f"{F.pct(LG['en']['wer_norm_corpus'])} / {F.pct(LG['en']['cer_norm_corpus'])}", "English WER / CER", K.ORANGE, 22, 11)
    ftile(s, 9.85, 1.75, 2.9, 0.95, f"{LG['ko']['mean_rtf']:.2f}", "Korean RTF (slowest MMS)", K.PURPLE, 24, 11)
    K.picture(s, F.fig("f01_rtf_by_language.png"), 0.45, 2.85, 6.2, 3.35)
    K.picture(s, F.fig("f02_wer_cer_by_language.png"), 6.75, 2.85, 6.2, 3.35)
    ftext(s, 0.55, 6.3, 12.2, 0.7, "RTF below 1 means faster than real time. WER/CER mix the speech model's error with Whisper-small's own, whose floor was not measured separately. Italian, Chinese and Japanese bars are the OS fallback voice.", size=12, color=K.FNAVY, italic=True)
    notes[16] = "Ten FLEURS sentences per language through the app's own services. 12 MMS languages are scored; the hatched bars are the fallback voice."

    # ---- 17 language match
    s = S[16]
    set_title(s, "Results - Did the Story Come Out in the Requested Language?")
    wipe(s)
    K.picture(s, F.fig("f07_language_match.png"), 0.55, 1.75, 12.2, 3.5)
    ftext(s, 0.55, 5.35, 12.2, 1.6, bl([
        f"Shipped path (A): the requested language in {nA} of {len(condA)} stories (English only). Chat template (B, experiment): {nB} of {len(condB)} - English, Spanish, German, Portuguese, Dutch, Chinese. Offline composer (C): English only.",
        f"Through the API, {n_en_text} of {len(nonen)} non-English stories (hi, fr, de, es) were English text, so the French, German and Spanish voices read English.",
        "Language detection is unreliable on very short text; chat-template output in Hindi, Arabic, Korean and others was English on inspection too.",
    ], 5), size=13)
    notes[17] = "This is the main weakness: the 0.6B model, given the shipped raw prompt, ignores the language request."

    # ---- 18 per-language table + story-length ASR
    s = S[17]
    set_title(s, "Results - Per-Language Outcome (Confusion-Matrix Equivalent)")
    wipe(s)
    rows = [["Lang", "Voice", "Spoken", "RTF", "WER", "CER", "Story A", "Story B"]]
    ma = {r["language"]: r["language_match"] for r in condA if r["style"] == "Adventure"}
    mb = {r["language"]: r["language_match"] for r in condB if r["style"] == "Adventure"}
    for c, e in LG.items():
        voice = "MMS" if c in F.MMS_LANGS else "OS"
        spoken = e.get("narrated", e["scored"] if e.get("models_used") else 0)
        rows.append([c, voice, f"{spoken}/{e['n']}", F.f2(e["mean_rtf"]) if e.get("mean_rtf") is not None else "-",
                     F.f2(e["wer_norm_corpus"]) if e.get("wer_norm_corpus") is not None else "-",
                     F.f2(e["cer_norm_corpus"]) if e.get("cer_norm_corpus") is not None else "-",
                     "yes" if ma.get(c) else "no", "yes" if mb.get(c) else "no"])
    ftable(s, 0.45, 1.7, [0.6, 0.7, 0.8, 0.7, 0.7, 0.7, 0.9, 0.9], rows, size=10, row_h=0.31, align=["c"] * 8)
    ftext(s, 7.35, 1.7, 5.5, 0.35, "Story-length audio (API stories, chunked Whisper)", size=13, color=K.FNAVY, bold=True)
    rows2 = [["Req.", "Text in", "Words", "Audio s", "WER", "CER"]]
    for r in ASR:
        rows2.append([r["language"], r.get("text_language") or "-", str(r["words"]), f"{r['audio_s']:.0f}", F.f2(r["wer_norm"]), F.f2(r["cer_norm"])])
    ftable(s, 7.35, 2.1, [0.7, 0.9, 0.85, 0.95, 0.8, 0.8], rows2, size=10.5, row_h=0.34, align=["c"] * 6)
    ftext(s, 7.35, 5.0, 5.5, 1.9, bl([
        f"English stories: mean WER {F.f2(en_wer)}, CER {F.f2(en_cer)}.",
        "Hindi voice made 4.4 s of audio from 185 English words; fr / de / es WER 0.6-1.3.",
        "The app's own evaluation reports no WER/CER for these clips (over 30 s).",
    ], 5), size=12)
    notes[18] = "Left: per language, was it spoken, how clear, and did the story match. Right: the stories made through the API, scored by a separate script because the app's evaluation fails on long audio."

    # ---- 19 parameter analysis
    s = S[18]
    set_title(s, "Parameter Analysis - Prompt Path and Story Length")
    wipe(s)
    K.picture(s, F.fig("f03_story_conditions.png"), 0.45, 1.75, 6.3, 2.6)
    K.picture(s, F.fig("f04_length_scaling.png"), 6.95, 1.75, 5.9, 3.3)
    ftext(s, 0.55, 4.45, 6.2, 2.5, bl([
        f"A (shipped): {A['mean_latency_s']:.1f} s, instruction echo {A['mean_instruction_echo']:.3f}, distinct-2 {A['mean_distinct_2']:.2f}.",
        f"B (chat template): {Bc['mean_latency_s']:.1f} s, echo {Bc['mean_instruction_echo']:.3f}, distinct-2 {Bc['mean_distinct_2']:.2f}.",
        f"C (composer): instant, {C['mean_words']:.0f} words; similarity {C['mean_similarity']:.2f} only because it quotes the idea.",
    ], 5), size=12.5)
    ftext(s, 6.95, 5.15, 5.9, 1.8, bl([
        "Length: " + "; ".join(f"{r['length']} {r['total_s']:.0f} s ({r['words']} words, {r['audio_s']:.0f} s audio)" for r in LS) + ".",
        "Chunk size (380 characters) and temperature were fixed, not swept.",
    ], 5), size=12.5)
    notes[19] = "The chat template removes instruction echo and raises language match but is not what the product ships. Time grows about linearly with the token budget."

    # ---- 20 performance and testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    K.picture(s, F.fig("f05_api_latency.png"), 0.45, 1.75, 6.4, 3.3)
    first = GEN[0]
    warm = [r for r in GEN if not r.get("cold")]
    ftile(s, 7.1, 1.75, 2.75, 0.95, f"{F.PYTEST_PASSED} / {F.PYTEST_TOTAL}", "pytest tests pass", K.GREEN, 24, 11)
    ftile(s, 10.0, 1.75, 2.75, 0.95, f"{first['request_s']:.0f} s", "first (cold) generate call", K.ORANGE, 24, 11)
    ftile(s, 7.1, 2.85, 2.75, 0.95, f"{min(r['request_s'] for r in warm):.0f}-{max(r['request_s'] for r in warm):.0f} s", "warm generate + narrate", K.FBLUE, 24, 11)
    ftile(s, 10.0, 2.85, 2.75, 0.95, f"{EP['GET /api/stories']['median_ms']:.1f} ms", "median list call", K.PURPLE, 24, 11)
    ftext(s, 0.55, 5.2, 12.2, 1.8, bl([
        f"Short story end to end {LS[0]['total_s']:.0f} s on a CPU (generation {LS[0]['gen_s']:.0f} s, narration {LS[0]['tts_s']:.0f} s); narration RTF about {LS[0]['rtf']:.2f}.",
        "Tests use stubbed models (validation, helpers, composer in 15 languages, API edge cases); the real models are exercised by the experiment scripts, not by pytest.",
        "Not done: load test, security audit, user study.",
    ], 5), size=13)
    notes[20] = "Cheap endpoints answer in milliseconds; the time is in the two models."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper (context only - nothing is score-comparable)")
    wipe(s)
    rows = [["Aspect", "SAM Translator (base paper)", "StoryVox"],
            ["Task", "Speech translation into English speech", "Story generation + narration in 15 languages"],
            ["Method", "mBART + SPL + 8-expert MoE + DAST (trained)", "Pretrained Qwen3-0.6B + MMS-TTS cascade (nothing trained)"],
            ["Scale / compute", "611M and 1.7B parameters; A100, 1.5-2 days", "0.6B story model + small VITS models; laptop CPU"],
            ["Data", "CVSS-C, CoVoST-2", "FLEURS test sentences (10 per language)"],
            ["Metric", "BLEU of the ASR transcript; MOS", "TTS-to-ASR WER / CER, RTF, language match"],
            ["Result", "SPL helped 18 of 21 languages", f"MMS: RTF {F.MMS_MEAN_RTF:.3f}, CER {F.pct(F.MMS_MEAN_CER)}; story language match {nA}/{len(condA)}"],
            ["Shared idea", "Transcribe the speech with an ASR model to check it", "Same round-trip check, with Whisper-small"]]
    ftable(s, 0.55, 1.8, [2.0, 5.0, 5.2], rows, size=12, row_h=0.52)
    ftext(s, 0.55, 6.1, 12.2, 0.85, "The two work on different tasks, data and metrics, so no score is compared. The comparison is of design: an end-to-end trained model against a cascade of pretrained models, and the cost each needs.", size=12.5, color=K.FNAVY, italic=True)
    notes[21] = "Say plainly: no score is comparable. SAM is trained, translates into English and reports BLEU; StoryVox trains nothing."

    # ---- 22-24 output screens
    def two_shots(slide, title, a, b, cap_a, cap_b):
        set_title(slide, title)
        wipe(slide)
        K.picture(slide, crop_shot(a), 0.45, 1.75, 6.0, 3.75, border=True)
        K.picture(slide, crop_shot(b), 6.85, 1.75, 6.0, 3.75, border=True)
        ftext(slide, 0.45, 5.65, 6.0, 1.2, cap_a, size=13, color=K.FNAVY, bold=True)
        ftext(slide, 6.85, 5.65, 6.0, 1.2, cap_b, size=13, color=K.FNAVY, bold=True)

    two_shots(S[21], "Output - Studio and Story Library", "03_studio_result.png", "04_history.png",
              "Studio: a Medium English adventure with audio player, WAV download and 'voice facebook/mms-tts-eng' (335 words, about 2 min). The title is 'Untitled Story'.",
              "Library: seven real stories (English, Spanish, German, French, Hindi) with play, download and delete.")
    notes[22] = "Real captures of the running app with the stories generated during measurement."
    two_shots(S[22], "Output - Story Pages", "05_story_detail.png", "10_story_detail_hindi.png",
              "Story detail for an English story with its narration.",
              "Hindi request: the text is English and the Hindi voice speaks 4.4 s of audio for 185 words - the language-mismatch failure, shown as it appeared.")
    notes[23] = "Success on the left, failure on the right."
    two_shots(S[23], "Output - Evaluation Page", "06_evaluation.png", "08_evaluation_samples.png",
              "Evaluation run of 5 samples: similarity 0.524; WER and CER show a dash because Whisper rejected the clips over 30 s.",
              "Per-sample rows: model, latency, RTF and the ASR error for each sample.")
    notes[24] = "The dash is the defect, not missing data: the app's evaluation cannot transcribe story-length audio."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.9, "StoryVox is a working, tested local application that joins story generation and multilingual narration; its measurements show narration is the reliable half and generation with a 0.6B model is the weak half.", size=16, color=K.FNAVY, bold=True)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Accept a user prompt", "Met"],
            ["2", "Length and reading-level controls", "Partly: 3 lengths, 8 styles; no reading level"],
            ["3", "Translate into target languages", "Not met: no translation module; model mostly answers in English"],
            ["4", "Convert to speech", f"Met for 12 languages (RTF {F.MMS_MEAN_RTF:.2f}, CER {F.pct(F.MMS_MEAN_CER)}); it, zh, ja use an OS voice"],
            ["5", "Adjustable voice, speed, language", "Partly: player and download only"],
            ["6", "Evaluate with user rating", "Partly: automated metrics only; WER/CER missing for long audio"]]
    ftable(s, 0.55, 2.75, [0.5, 4.2, 7.5], rows, size=12, row_h=0.5, align=["c", "l", "l"])
    ftext(s, 0.55, 6.35, 12.2, 0.6, "Key finding: language selection works in the speech step but not in the story step.", size=13, color=K.FNAVY, italic=True)
    notes[25] = "Be explicit about which objectives were met, partly met or not met."

    # ---- 26 limitations + future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 4.9, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.3, bl([
        "Qwen3-0.6B ignores the language request with the shipped prompt; titles often 'Untitled Story'",
        "No MMS checkpoint for it, zh, ja; OS voice gives 2/10 zh and 1/10 ja sentences",
        "Evaluation cannot score audio over 30 s; ffmpeg and uroman undeclared",
        "10 sentences per language; Whisper's own error floor not separated",
        "No human rating of story quality; metrics are proxies",
        "Tests use stubs; CPU only; broken project .venv on the test machine",
    ], 8), size=15)
    fbox(s, 6.8, 1.8, 6.0, 4.9, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.3, bl([
        "Use the chat template, a larger story model, or generate in English and translate",
        "Chunked ASR (chunk_length_s=30) in the evaluation; declare ffmpeg and uroman",
        "Voice and speed controls; other TTS engines for it, zh, ja",
        "User rating of coherence and naturalness; reading-level control",
        "Cache the similarity model; add a GPU path and a load test",
        "Reference-audio voice cloning; authentication and multi-user mode",
    ], 8), size=15)
    notes[26] = "Limitations first, then future work drawn from the measured defects."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "omni", "multigen", "kahaani", "vits", "mms", "whisper", "qwen3", "sbertm", "fleurs", "tinystories", "holtzman", "hubert", "mbart", "moe", "transformer"]
    short_ = [f"[{_IDX[k]}] {clean_ref(REFS[_IDX[k] - 1])}" for k in keys]
    short_ = [t if len(t) < 300 else t[:297] + "..." for t in short_]
    ftext(s, 0.55, 1.75, 6.0, 5.2, [(t, {"space_after": 4}) for t in short_[:8]], size=10)
    ftext(s, 6.8, 1.75, 6.0, 5.2, [(t, {"space_after": 4}) for t in short_[8:]] + [(f"The full numbered list of {len(REFS)} references, including the batch survey titles {N('s2')}-{N('s5')}, is in the report.", {"italic": True})], size=10)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh_ in S[27].shapes:
        if sh_.name == "Text 2":
            K.replace_text(sh_, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "StoryVox")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
