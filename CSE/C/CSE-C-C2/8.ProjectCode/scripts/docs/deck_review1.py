"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx.

All measured numbers come from facts.py (experiments/**/metrics.json).
    python -m scripts.docs.deck_review1        # from 8.ProjectCode
"""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "multilingual story generation, text-to-speech, MMS-TTS, Qwen3, Whisper, speech evaluation"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def trim(path, tag):
    """Copy of a diagram with its white margins removed (written to the temp folder, source untouched)."""
    from PIL import Image, ImageChops

    with Image.open(path) as im:
        im = im.convert("RGB")
        bbox = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).getbbox()
        if bbox:
            pad = 12
            bbox = (max(bbox[0] - pad, 0), max(bbox[1] - pad, 0), min(bbox[2] + pad, im.size[0]), min(bbox[3] + pad, im.size[1]))
            im = im.crop(bbox)
        out = os.path.join(tempfile.gettempdir(), f"storyvox_r1_{tag}.png")
        im.save(out)
    return out


PROBLEM_STATEMENT = [
    "Producing multilingual stories manually requires writers and narrators for every language",
    "Generic text-to-speech sounds mechanical and is unsuited to narrative delivery",
    "Low-resource Indian languages have limited synthesis support and poor voice quality",
    "Existing tools handle either text generation or speech synthesis, not both together",
    "Need: a single system that generates a story from a prompt and narrates it naturally in the user's chosen language",
]
OBJECTIVES = [
    "Accept a user prompt specifying theme, characters and target language",
    "Generate coherent narrative text using a language model, with controls for length and reading level",
    "Translate the generated story into the selected target languages while preserving meaning",
    "Convert the story text into natural speech using multilingual text-to-speech synthesis",
    "Provide playback with adjustable voice, speed and language selection",
    "Evaluate narrative coherence, translation quality and speech naturalness through user rating",
]

# [work, year, approach, reported result / limitation]  (checked against lit.ROWS; surveys 2, 4, 5 are titles only)
LIT_ROWS = [
    [f"Tran and Sakti {N('base')} (base paper)", "2026", "SAM Translator: mBART text-to-unit model with self-paced learning, sparse mixture-of-experts and self-training",
     "BLEU gains of +2-3% (high-resource) to +7-10% (low-resource) on CVSS-C; target side English only; 1.7B parameters, A100"],
    [f"Zhu et al. {N('omni')}", "2026", "OmniVoice: zero-shot TTS, masked-diffusion LM, 646 languages, 581k hours",
     "Average CER 4.00%, RTF 0.0319 on a GPU; speech only, no story text; 8 GPUs for 9.7 days of training"],
    [f"Gao et al. {N('multigen')}", "2025", "MultiGen: child-friendly multilingual speech from an LLM speech-token model",
     "CER Tamil 9.6 to 2.1, Malay 3.4 to 2.3 (%); three languages, private child-speech data"],
    [f"Park, Nakamura {N('emo')}", "2025", "EmoSSLSphere: multilingual emotional TTS with spherical emotion vectors", "Better intelligibility and expressiveness; English and Japanese only"],
    [f"Arif et al. {N('kahaani')}", "2024", "Kahaani: LLM co-creative storytelling with speech, music and video", "Interactive system for children; multilingual narration not reported"],
    [f"Kim et al. {N('vits')}", "2021", "VITS: end-to-end TTS with a conditional VAE, flow and adversarial training", "MOS close to recordings on LJ Speech; one language per checkpoint"],
    [f"Pratap et al. {N('mms')}", "2024", "MMS: wav2vec 2.0 pre-training plus ASR and VITS TTS for 1,107 languages", "One single-speaker TTS model per language; some scripts need romanisation"],
    [f"Radford et al. {N('whisper')}", "2023", "Whisper: weakly supervised multilingual ASR on 680,000 hours", "Robust zero-shot transcription; error differs by language; may hallucinate"],
    [f"Yang et al. {N('qwen3')}", "2025", "Qwen3: open LLM family, 0.6B to 235B, 119 languages", "The 0.6B model is far weaker than the flagship and needs its chat template"],
    [f"Eldan, Li {N('tinystories')}", "2023", "TinyStories: small models on simple synthetic stories write fluent English", "English only; restricted vocabulary and domain"],
    [f"Batch survey {N('s2')}, {N('s4')}, {N('s5')}", "2025-26", "Three further titles from the batch survey (multilingual story and speech, TTS with RAG, motion generation)",
     "Not located online; listed by title only, nothing quoted from them"],
]

GAPS = [
    f"No surveyed system joins controllable story generation (language, style, length) to multilingual narration in one application that runs locally {N('omni', 'multigen', 'kahaani')}",
    f"The base paper translates speech into English only and has no story content or control; it needs a 1.7B-parameter model and an A100 GPU {N('base')}",
    "Speech papers report MOS or CER on read sentences; none tells a user whether a story requested in Hindi or Korean really comes out in that language and is intelligible when spoken",
    f"Large systems rely on GPUs {N('base', 'omni')}; a laptop-CPU pipeline with a small story model {N('qwen3')} is rarely measured end to end",
    f"A speech recogniser can act as an objective judge of narration {N('whisper')}, but story pipelines seldom report round-trip WER/CER and timing together",
]


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    L = F.LANGS
    story_a, story_b = F.STORY["conditions"]["A_app_raw_prompt"], F.STORY["conditions"]["B_chat_template"]
    a_match = sum(1 for r in story_a if r["language_match"])
    b_match = sum(1 for r in story_b if r["language_match"])
    sasr = F.STORY_ASR["rows"]
    non_en = [r for r in sasr if r["language"] != "en"]
    non_en_english = sum(1 for r in non_en if r.get("text_language") == "en")
    hi_row = next(r for r in sasr if r["language"] == "hi")

    # 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: StoryVox turns a story idea into a story in a chosen language and reads it aloud, "
                "using pretrained models on a laptop CPU, and measures how well each step works.")

    # 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Outcomes measured so far"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. The review covers why the problem matters and what is missing in current work, then what we built and what we measured."

    # 3 abstract
    s = S[2]
    set_title(s, "Abstract: from a story idea to spoken stories")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Storytelling supports education and entertainment, but producing stories in several languages needs writers and narrators for each one. Generic speech sounds mechanical, and tools do either text or speech, not both.", K.ORANGE, "!"),
        ("Our idea", "One pipeline: a small instruction-tuned language model (Qwen3-0.6B) writes the story, a per-language VITS model from the MMS project narrates it, and a speech recogniser (Whisper-small) checks the result by transcribing it back.", K.BLUE, "i"),
        ("The deliverable", f"StoryVox: a FastAPI service (11 routes) and a 5-page React app with a SQLite library. It covers {len(F.LANGS)} languages, 8 styles and 3 lengths, runs on a CPU laptop, and is measured with real-time factor, WER and CER.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=19)
    notes[3] = "Three points: the manual effort of multilingual stories, the three-model cascade, and what is delivered and measured."

    # 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "mBART text-to-unit model with self-paced learning (easy samples first), sparse mixture-of-experts layers (8 experts, top-2) and self-training on filtered synthetic pairs; HuBERT units are turned into speech by Grad-TTS.", K.BLUE, "M"),
        ("Data and reported results", "CVSS-C (21 source languages to English speech) and CoVoST-2. BLEU of the transcript improves by 2-3% (high-resource) and 7-10% (low-resource) over mBART; self-training adds about 21.6% for Welsh.", K.GREEN, "R"),
        ("Limits and our use", "Target is English only, it translates but does not write stories, and it needs a 1.7B-parameter model on an A100 for 1.5-2 days. StoryVox keeps the goal (multilingual speech, judged by ASR) and none of the techniques.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=16.5)
    notes[4] = ("Read the citation. SAM Translator is a translation system; StoryVox is a generation-plus-narration pipeline. We compare ideas and compute cost, "
                "never scores: SAM reports BLEU, we report round-trip error rates.")

    # 5 literature table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]] + LIT_ROWS
    K.table(s, 0.6, 1.1, [2.3, 0.65, 4.7, 4.48], rows, font_size=10.5, row_h=0.47, header_size=11, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey lists five items; items 2, 4 and 5 could not be located online, so only their titles appear. "
                "All other rows were checked against the papers or their abstract pages.")

    # 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence: pretrained speech models work")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 6.6, 4.7, ["Tamil", "Malay", "Mandarin"],
                [("Before (as reported)", [9.6, 3.4, 10.7]), ("With MultiGen (as reported)", [2.1, 2.3, 4.0])], colors=[K.ORANGE, K.BLUE],
                number_format="0.0", font_size=11, max_value=13, legend=True)
    K.text(s, 0.6, 6.2, 6.6, 0.8, f"Character error rate in %, from Gao et al. {N('multigen')}; the paper's private child-speech data, not ours.", size=11, color=K.TXT, italic=True)
    K.text(s, 7.5, 1.45, 5.25, 5.4, [
        ("**What the literature shows**", {"space_after": 8}),
        (f"OmniVoice reaches an average CER of 4.00% against 5.11% for the ground-truth recordings, at RTF 0.0319 on a GPU {N('omni')}.", {"bullet": True, "space_after": 8}),
        (f"MMS offers one VITS model for each of 1,107 languages; its ASR halves Whisper's word error rate on 54 FLEURS languages {N('mms')}.", {"bullet": True, "space_after": 8}),
        (f"Tiny models write fluent short English stories when the domain is simple {N('tinystories')}.", {"bullet": True, "space_after": 8}),
        ("Caveat: all of this is on other data and hardware, so none of it is a baseline for our numbers.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = "These are published figures on other data; our own measurements are in the later slides and are not placed on one scale with them."

    # 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Writer drafts the story", K.NAVY), ("Translator for each language", K.NAVY), ("Separate TTS tool", K.NAVY), ("Narrator corrects audio", K.NAVY), ("One file per language", K.NAVY)], size=14, gap=0.45)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Manual and separate tools",
           "A writer and a narrator are needed for every language.\nText generators and speech synthesisers are different products, joined by hand.\nGeneric TTS sounds mechanical for storytelling.", K.BLUE, "1", body_size=19)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Research systems",
           f"Speech translation {N('base')}, massively multilingual TTS {N('omni', 'mms')} and LLM storytelling {N('kahaani')} each solve one part.\nNone is a locally runnable story-to-speech app with an honest end-to-end measurement.", K.ORANGE, "2", body_size=19)
    notes[7] = "Today the workflow is manual or stitched from separate tools; research covers the parts but not the joined, measured pipeline."

    # 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Effort per language", "Mechanical voices", "Low-resource languages", "Text or speech, not both"]
    bodies = [
        "Producing multilingual stories manually requires writers and narrators for every language.",
        "Generic text-to-speech sounds mechanical and is unsuited to narrative delivery.",
        "Indian and other low-resource languages have limited synthesis support and poor voice quality.",
        "Existing tools handle either text generation or speech synthesis, not both together.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=18)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"The base paper adds: English-only target speech, no story generation or content control, and a GPU-scale model {N('base')}.", size=14, color=K.INK, anchor="m")
    notes[8] = "The four limitations are the batch problem statement; the last line is from the base paper analysis."

    # 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=17, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as joining controllable story generation to multilingual narration in one pipeline aimed at education. "
                "Gaps 3 to 5 come from our reading of how speech papers report results.")

    # 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "A story idea (up to 2,000 characters), one of 15 languages, one of 8 styles and one of 3 lengths"),
                                 ("Output", "Story text with a title, a narration WAV, a saved library entry, and an evaluation report"),
                                 ("Measures", "Real-time factor, WER and CER of the round trip, prompt-story similarity, ROUGE-L, latency")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=15, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is one system that writes a story and narrates it in the chosen language."

    # 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Slide 14 states which objectives are met, partly met or not met in the system as built.", size=13, color=K.TXT, italic=True)
    notes[11] = "The six objectives from the proposal. We say plainly later which are met."

    # 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Prompt", "The idea, language, style and length become one instruction with a 'Title:' first-line rule and a paragraph target of 3, 5 or 8.", K.BLUE, "1"),
        ("Write", "Qwen3-0.6B generates 220, 420 or 700 tokens (temperature 0.8, top-p 0.9). If the output is under 20 words or the model fails, an offline composer answers.", K.GREEN, "2"),
        ("Narrate", "MMS-TTS (VITS) speaks the text in 380-character chunks at 16 kHz; an OS voice is the fallback. The WAV and a library row are saved.", K.ORANGE, "3"),
        ("Evaluate", "Whisper-small transcribes the audio back; WER, CER, real-time factor, similarity and ROUGE-L are computed and stored.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React app for the whole pipeline", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "Studio (write and narrate), History (library with playback), Story detail and Evaluation pages. Everything runs locally after the one-time model download; no GPU is required.", size=14, color=K.TXT)
    notes[12] = "Four stages and the app that wraps them. No model is trained: all three are pretrained and used as they are."

    # 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, trim(F.diag("01_system_architecture.png"), "01"), 0.6, 1.2, 8.2, 5.9)
    K.text(s, 9.05, 1.5, 3.7, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"React app on port {F.WEB_PORT} (the API also serves the built bundle).", {"bullet": True}),
        (f"FastAPI on port {F.API_PORT}: 11 routes plus an SPA fallback.", {"bullet": True}),
        ("Services: story generator, TTS provider chain, ASR and evaluation.", {"bullet": True}),
        ("SQLite with three tables (stories, audio_outputs, evaluations); WAV files on disk.", {"bullet": True}),
        ("Hugging Face models downloaded once, then used offline.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the app calls the API, the API calls the services, which call the three models and the database."

    # 14 outcomes
    s = S[13]
    set_title(s, "Outcomes measured so far")
    wipe(s)
    heads = ["Working platform", "Narration", "Story language", "Language mismatch", "Not implemented"]
    bodies = [
        f"11 routes, 5 pages, {len(F.LANGS)} languages, 8 styles, 3 lengths; {F.PYTEST_PASSED} of {F.PYTEST_TOTAL} tests pass ({F.PYTEST_API} original, {F.PYTEST_UNIT} added).",
        f"{len(F.MMS_LANGS)} languages use MMS: mean RTF {F.MMS_MEAN_RTF:.3f}, mean CER {F.pct(F.MMS_MEAN_CER)}. English WER {F.pct(L['en']['wer_norm_corpus'])}, CER {F.pct(L['en']['cer_norm_corpus'])} (10 FLEURS sentences).",
        f"The shipped prompt gives target-language text in only {a_match} of {len(story_a)} test cases; the chat template gives {b_match}. Titles often read 'Untitled Story'.",
        f"{non_en_english} of {len(non_en)} non-English API stories came out in English; the Hindi voice produced {hi_row['audio_s']:.1f} s of audio for {hi_row['words']} words.",
        "No translation module, voice or speed selection, user rating or reference-audio upload. Italian, Chinese and Japanese have no MMS checkpoint.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=14)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=16)
    notes[14] = ("Because the system is already built, this slide reports measured outcomes. Narration works well; story generation with the small model is the weak link. "
                 "Objective status: 1 achieved; 4 achieved for 12 languages; 2, 5 and 6 partly; 3 not as translation.")

    # 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = sorted(["base", "s2", "emo", "s4", "s5", "omni", "multigen", "kahaani", "vits", "mms", "whisper", "qwen3", "tinystories", "fleurs"], key=lambda k: _IDX[k])
    short = {"s2": "listed", "s4": "listed", "s5": "listed"}
    K.text(s, 0.6, 1.2, 12.13, 5.9, [(f"[{_IDX[k]}] {REFS[_IDX[k] - 1].replace(' (title as listed in the batch literature survey; not located online and not quoted)', ' (title only; not located online)').replace(' (title as listed in the batch literature survey; not located online, appears to concern motion synthesis, and is not quoted)', ' (title only; not located online)')}", {"space_after": 3}) for k in keys],
           size=12, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why a cascade and not an end-to-end model (no training budget; pretrained parts exist); why only 12 languages have MMS narration "
                 "(Italian, Chinese and Japanese checkpoints do not exist on the Hub); why the 0.6B story model (it runs on a CPU laptop).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
