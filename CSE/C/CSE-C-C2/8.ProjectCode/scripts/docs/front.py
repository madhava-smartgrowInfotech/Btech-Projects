"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        os.makedirs(os.path.dirname(BANNER), exist_ok=True)
        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, C. Tran and S. Sakti, whose work on cross-lingual text-to-speech translation framed this project; "
        "the Qwen team, the Meta MMS project and OpenAI, whose open Qwen3-0.6B, MMS-TTS and Whisper models StoryVox builds on; Google, for the FLEURS speech benchmark "
        "(released under CC BY 4.0) on which the speech round trip was measured; and the maintainers of the open-source software used here: "
        "PyTorch, Hugging Face Transformers, FastAPI, SQLAlchemy, React, Vite and Playwright.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    mms = len(F.MMS_LANGS)
    a, b, c = F.STORY_SUMMARY["A_app_raw_prompt"], F.STORY_SUMMARY["B_chat_template"], F.STORY_SUMMARY["C_local_composer"]
    en = F.LANGS["en"]
    d.p("Storytelling supports education, language learning and entertainment, but producing stories in several languages and reading them aloud normally takes writers and narrators for every language. "
        f"This project builds {F.PRODUCT}, a Multilingual Story Generation and Speech System that runs on a single computer: the user types a story idea and chooses a language, a style and a length; "
        "a small instruction-tuned language model (Qwen3-0.6B) writes the story; a per-language neural speech model (MMS-TTS, VITS) reads it aloud; and the story, the audio and the evaluation results are kept in a local database. "
        f"The system has a React web application with five pages, a FastAPI service with 11 routes, 15 supported languages, 8 story styles and 3 lengths, and {F.PYTEST_PASSED} passing automated tests. "
        "No model was trained or fine-tuned: StoryVox is a cascade of pretrained models, so it uses the base paper, SAM Translator, as the reference for multilingual speech output and does not reimplement its self-paced learning, mixture-of-experts or self-training.")
    d.p(f"We measured the pipeline on a CPU-only machine. Speech was checked by a round trip: 10 FLEURS sentences per language were synthesised and transcribed by Whisper-small. "
        f"Across the {mms} languages narrated by MMS the mean real-time factor is {F.MMS_MEAN_RTF:.3f} (faster than real time) and the mean normalised character error rate is {F.pct(F.MMS_MEAN_CER)}; "
        f"English scores a word error rate of {F.pct(en['wer_norm_corpus'])}. Italian, Chinese and Japanese have no MMS checkpoint on the model hub and fall back to an operating-system voice that has only English voices on the test machine: it read all 10 Italian sentences (word error rate {F.pct(F.LANGS['it']['wer_norm_corpus'])}) but narrated only {F.LANGS['zh']['narrated']} of 10 Chinese and {F.LANGS['ja']['narrated']} of 10 Japanese sentences. "
        f"The story model is the weak link: with the prompt the application sends, the story text matched the requested language in {F.pct(a['language_match_rate'], 0)} of 17 test cases, and all four non-English stories generated through the API came out in English, "
        f"which the foreign-language voice then read badly. A chat-template variant tried as an experiment raised the match to {F.pct(b['language_match_rate'], 0)}. "
        "The application's own evaluation page also reports no word or character error rate for story-length audio, because the recogniser call rejects audio longer than 30 seconds; the experiments use a separate chunked script.")
    d.p("The batch objectives are met in part: prompt entry, story generation with style and length control, speech output in 12 languages, playback and download, and automated evaluation are implemented; "
        "translation, reading-level control, voice and speed selection and user rating are not. We report these results and gaps as measured.")
    d.p("**Keywords:** multilingual story generation, text-to-speech, speech synthesis, natural language processing, storytelling system, MMS-TTS, Whisper, Qwen3, educational technology.")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("ASR", "Automatic Speech Recognition"), ("BLEU", "Bilingual Evaluation Understudy"),
    ("CER", "Character Error Rate"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("DAST", "Data-Augmented Self-Training"),
    ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"), ("FLEURS", "Few-shot Learning Evaluation of Universal Representations of Speech"), ("FR", "Functional Requirement"),
    ("GPU", "Graphics Processing Unit"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"), ("LLM", "Large Language Model"),
    ("MMS", "Massively Multilingual Speech"), ("MoE", "Mixture of Experts"), ("MOS", "Mean Opinion Score"), ("NFR", "Non-Functional Requirement"),
    ("NLP", "Natural Language Processing"), ("ORM", "Object-Relational Mapping"), ("REST", "Representational State Transfer"), ("RTF", "Real-Time Factor"),
    ("ROUGE-L", "Recall-Oriented Understudy for Gisting Evaluation, longest common subsequence"), ("SAM", "Self-paced learning and mixture-of-experts (SAM Translator)"),
    ("SPA", "Single-Page Application"), ("SPL", "Self-Paced Learning"), ("TTS", "Text-to-Speech"), ("UI", "User Interface"),
    ("UML", "Unified Modeling Language"), ("VITS", "Variational Inference with adversarial learning for end-to-end Text-to-Speech"), ("WAV", "Waveform Audio File Format"),
    ("WER", "Word Error Rate"),
]


def abbreviations(d):
    """Two abbreviation/expansion pairs per row so the list takes half the space."""
    half = (len(ABBR) + 1) // 2
    rows = [list(ABBR[i]) + (list(ABBR[i + half]) if i + half < len(ABBR) else ["", ""]) for i in range(half)]
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion", "Abbreviation", "Expansion"], rows, widths=[0.95, 2.25, 0.95, 2.25], size=10)
