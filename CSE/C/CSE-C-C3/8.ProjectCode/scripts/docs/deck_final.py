"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (letterhead, rule, title and footer chrome inherited)."""

import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import lit
from . import pptkit as K
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "c3_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def bl(items, space=8):
    return [(t, {"bullet": True, "space_after": space}) for t in items]


def pic(slide, path, x, y, w, h, border=False):
    """Picture if the file exists; otherwise a visible placeholder note (the diagram fork may still be drawing)."""
    if os.path.exists(path):
        return K.picture(slide, path, x, y, w, h, border=border)
    fbox(slide, x, y, w, h, fill=K.FZEBRA)
    ftext(slide, x + 0.2, y + 0.2, w - 0.4, 0.6, "[missing: " + os.path.basename(path) + "]", size=12)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.subject = F.PRODUCT
    cp.keywords = "smart parking, reservation, hold-back buffer, Trust Score, FastAPI, React"
    notes = {}

    SD = F.SEED
    sc = SD["status_counts"]
    pk = F.pop
    bd = F.bad
    probes = F.SEC

    for sl in S:  # footers on every slide except the title
        for sh_ in sl.shapes:
            if sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 30)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and ParkSure: an advance-reservation parking platform with a hold-back buffer, nearest-alternate re-allocation and a driver Trust Score."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Papers & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Data & Engine Rules\nInteraction & Data Design\nImplementation, Performance & Testing\nResults: Buffer Policies, Arrival Outcomes, Trust, Robustness\nComparison with the Base Papers\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, bl([
        "Drivers circle for parking and operators plan from paper registers. ParkSure is a web platform where drivers search live availability, reserve a bay in advance, pay online (Stripe test mode or demo checkout), check in and out, and administrators manage locations and see analytics. FastAPI + SQLite backend, React console, 22 API routes, 5 tables.",
        "A Smart Availability Engine admits a booking only if overlapping bookings + buffer < capacity, keeps a hold-back buffer per location, moves a driver to the nearest alternate (haversine) when a lot is full on arrival, and scores drivers with a 0-5 Trust Score that sets fee multipliers.",
        f"Verified by {F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks; a simulator matches the real engine on {F.PAR['total_cases']:,} decisions with {F.PAR['total_mismatches']} mismatches. Search takes {F.lat('GET /api/search'):.1f} ms, a reservation {F.lat('POST /api/reservations'):.1f} ms (127.0.0.1, 200 requests each).",
        f"Simulated buffer policies (80-110 vehicles, 30 runs): the shipped dynamic rule refuses the most reservations ({pk(100, 'dynamic_parksure', 'no_reservation'):.0f} of 1,000 at 100 vehicles against {pk(100, 'static_0', 'no_reservation'):.0f} with no buffer); 'no park' stays rare under every policy.",
        f"Honest limits: demo data are synthetic; probes found an over-booking race ({F.RACE['trials_with_overbooking']} of {F.RACE['trials']} trials), an unauthenticated payment-confirm route and no vehicle profile or receipt; no real-user study.",
    ], 8), size=16)
    notes[3] = "Product first, then the engine, the evidence and the honest limits."

    # ---- 4 base papers
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.75, F.BASE_PAPER, size=12, color=K.FNAVY)
    rows = [["Aspect", "Registered base paper (Jemmali et al.): passenger-aware allocation for medical cities"],
            ["Problem", "Allocate vehicles to parking spaces by the total number of passengers, not vehicle count, to balance the load between zones of a medical city"],
            ["Method", "Four-stage grouping procedure and six grouping-permutation heuristics; architecture with booking interface, IoT sensing, validation and a central allocation engine; QR-code booking"],
            ["Evaluation", "Nine instance classes, 2,790 generated test cases; best algorithm has the highest best-solution share (68.2 %); GVb matches or beats baseline MR on 1,547 instances (55.4 %), in sub-millisecond time"],
            ["Stated gap", "Existing systems optimise one vehicle or revenue; few use grouping criteria such as passenger count"],
            ["Not covered", "Payment, overstay handling, operator analytics, a working web product"],
            ["Also used", f"{F.TECH_SHORT}: dynamic buffer of non-reservable slots and a 0-5 star reward system; this is what ParkSure's engine implements {N('cabri')}"]]
    ftable(s, 0.55, 2.6, [1.7, 10.5], rows, size=12, row_h=0.62)
    notes[4] = "The registered base paper is Jemmali et al. (IEEE OJCS 2026). ParkSure does not implement its grouping heuristics. The reservation buffer and star ideas come from the Cabri et al. arXiv paper, which the report cites as the technical basis."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Approach", "Result / limitation (as stated in the base papers)"],
            [f"Jemmali et al. {N('base')} (base)", "Passenger-count zone balancing; six heuristics", "2,790 generated instances; GVb ties or beats baseline MR on 55.4 %; one facility, needs IoT"],
            [f"Cabri et al. {N('cabri')}", "Buffer of non-reservable slots + star reward", "SUMO, 80 spots / 8 areas: any buffer beats none, dynamic best trade-off; preliminary"],
            [f"Doulamis et al. {N('doulamis')}", "Interval scheduling of requests", "Accept or reject by time interval; no overstay handling described"],
            [f"Safitri et al. {N('safitri')}", "Mobile reservation, heuristic + haversine, occupancy display", "Mobile client; no buffer or reputation"],
            [f"Somani et al. {N('somani')}", "Part-payment confirms a booking", "Deters false bookings, not overstays"],
            [f"Mejri et al. {N('mejri')}; Kotb et al. {N('kotb')}", "Multi-objective ILP; iParker flexible times and pricing", "Optimisation / pricing models; no operating web platform described"],
            [f"Rizvi et al. {N('asap', 'mapark')}; Kokolaki {N('kokolaki')}", "Auction-based parking (agents, bidding)", "Suit competitive allocation, not a fixed-price booking service"],
            [f"Batch survey {N('s2')}-{N('s5')}", "Four reservation / smart-parking titles", "Titles only; not found online, no result quoted"]]
    ftable(s, 0.55, 1.8, [3.3, 4.2, 4.7], rows, size=11, row_h=0.52, align=["l", "l", "l"])
    notes[5] = "Only what the two base papers say about the related works is used. The four batch-survey titles could not be located online."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(g, {"bullet": True, "space_after": 14}) for g in [
        "The registered base paper optimises allocation inside a large managed facility; it has no booking-to-payment service for ordinary urban parking.",
        "The buffer and reward ideas (Cabri et al.) are shown only in a simulation of 80 spots; no operating product and no numeric table are reported.",
        "Surveyed systems describe reservation components one at a time: booking, part-payment, mobile access, pricing, auctions.",
        "Overstays are the known failure of reservation systems; few designs combine a buffer, a fallback alternate and a punctuality score.",
        "Operators lack a way to replay their own booking history under different buffer policies and see the effect.",
        "Gap: a general reservation platform with search, booking, payment, an overstay-aware engine and operator analytics, evaluated honestly on its own behaviour."]], size=17)
    notes[6] = "The gap statement from the batch literature survey, made concrete with the points the base papers leave open."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.6)
    ftext(s, 0.85, 1.9, 11.6, 2.5, [
        "Drivers circle repeatedly because there is no way to know which slots are free.",
        "Operators manage occupancy through manual registers with no live view.",
        "Without reservation, arriving at a facility gives no guarantee of a space.",
        "Cash collection is slow at entry and leaves no usable records.",
        "Need: a reservation platform that shows live availability, guarantees a slot in advance and handles payment digitally."], size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Registration, search window, reservation, payment, check-in and check-out; admin location data"),
                                 ("Output", "Availability per location, quote, booking status, payment record, Trust Score, admin overview and Strategy Insights"),
                                 ("Record", f"{F.PYTEST_PASSED} tests, simulator and latency metrics (JSON), 24 screenshots of the running app")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.7, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.8, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.3, 3.5, 1.0, b_, size=14)
    notes[7] = "The four problem points are taken from the batch problem statement."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate([
        "Build user registration, authentication and vehicle profile management.",
        "Maintain real-time parking slot availability across registered parking facilities.",
        "Implement slot search, advance reservation and booking cancellation.",
        "Integrate secure online payment and generate digital receipts.",
        "Provide an administrator dashboard for slots, bookings, occupancy and revenue.",
        "Evaluate booking success rate, allocation efficiency and system response time."], 1)], size=22)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Driver journey", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Search window", K.FNAVY), ("Reserve + quote", K.FBLUE), ("Pay (Stripe / demo)", K.GREEN), ("Check in", K.ORANGE), ("Park", K.PURPLE), ("Check out", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Smart Availability Engine (app/engine.py)", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("Admission rule", K.FNAVY), ("Hold-back buffer", K.FBLUE), ("Nearest alternate", K.GREEN), ("Trust Score", K.ORANGE), ("Fee multiplier", K.PURPLE), ("Strategy Insights", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, bl([
        "Admit only if overlapping bookings + buffer < capacity; refused requests are stored as 'declined' with the reason.",
        "If the lot is physically full on arrival: exactly one nearest alternate with forecast space, else 'failed' and the fee is credited automatically.",
        "On-time check-out (10 min grace) grows the buffer and the Trust Score; late check-out records a warning; the score changes the next fee.",
        "Administrators replay the stored history under four buffer policies (Strategy Insights).",
    ], 7), size=15)
    notes[9] = "The proposed system as two chains: what the driver does, and what the engine decides."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    pic(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.2, 5.1)
    ftext(s, 7.95, 1.95, 4.9, 4.9, bl([
        f"Console: React 19 + Vite, port {F.WEB_PORT}; landing, login, 3 driver pages, 5 admin pages",
        f"API: FastAPI, port {F.API_PORT}; 22 routes; JWT bearer, roles driver / admin",
        "Engine: app/engine.py (admission, buffer, haversine, trust, fees)",
        "Data: SQLAlchemy over one SQLite file, 5 tables; APScheduler resets buffers every 30 min",
        "Payments: Stripe Checkout in test mode, offline demo fallback",
    ], 8), size=15)
    notes[10] = "Four tiers; the engine is a plain module of the backend, not a separate service."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    pic(s, F.diag("02_high_level_workflow.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, bl([
        "Search counts bookable bays = capacity - buffer - overlapping bookings.",
        "Fee = max(0.25, hours) x hourly rate x trust multiplier (0.90 to 1.75); account credit is applied first.",
        "Buffer: start 1, +1 per on-time check-out, max 5, reset to 1 every 24 h.",
        "Trust: start 3; 3 benefits = +1; 2 warnings = -1; below 1.5 with a warning in 14 days = 24 h suspension.",
        "Evaluation: tests, a simulator parity-checked against the engine, API timing, probes.",
    ], 8), size=14)
    notes[11] = "Method in one picture and the four numeric rules that define the engine."

    # ---- 12 data and engine rules
    s = S[11]
    set_title(s, "Data and Engine Rules")
    wipe(s)
    ftile(s, 0.55, 1.8, 2.1, 1.0, f"{SD['reservations']}", "seeded reservations (synthetic)", K.FBLUE, 22, 11)
    ftile(s, 2.8, 1.8, 2.1, 1.0, f"{SD['locations']} + {SD['users']}", "locations + users seeded", K.GREEN, 22, 11)
    ftile(s, 5.05, 1.8, 2.1, 1.0, "60 days", "history + 8 Friday peaks", K.ORANGE, 22, 11)
    rows = [["Item", "Detail"],
            ["Source", "app/seed.py, random.Random(42): no external dataset, no ML model"],
            ["Statuses", f"{sc['completed']} completed, {sc['failed']} failed (unable to park), {sc['declined']} declined"],
            ["Locations", "5 lots, 8 to 25 bays, $2.00 to $5.00 per hour"],
            ["Fee multiplier", "score 0 / 1 / 2 / 3 / 4 / 5 = 1.75 / 1.50 / 1.25 / 1.00 / 0.95 / 0.90"],
            ["Payments", "Stripe test mode; demo checkout when keys are absent"],
            ["Caveat", "Synthetic data show the mechanism, not real demand"]]
    ftable(s, 0.55, 3.0, [1.5, 5.2], rows, size=11, row_h=0.46)
    pic(s, F.diag("12_trust_buffer_policy.png"), 7.5, 1.8, 5.35, 4.4)
    ftext(s, 7.5, 6.3, 5.35, 0.6, "Trust Score machine and buffer policy as implemented in app/engine.py and app/config.py.", size=12, color=K.FNAVY, italic=True)
    notes[12] = "There is no ML dataset. All demo data are generated by the product itself, so every result built on them is labelled synthetic."

    # ---- 13 core algorithm
    s = S[12]
    set_title(s, "Core Algorithm: Admission, Alternate and Trust")
    wipe(s)
    pic(s, F.diag("10_flowchart_algorithm.png"), 0.55, 1.75, 7.6, 5.2)
    ftext(s, 8.35, 1.9, 4.5, 5.0, bl([
        "Request time: suspended? refuse. overlapping + buffer < capacity? else decline. Quote and book.",
        "Arrival: parked >= capacity? search other lots by haversine distance (R = 6371 km); book the nearest with forecast space; else failed + credit.",
        "Departure: end + 10 min? late = warning, on time = buffer +1 and benefit.",
        f"Admission decision {F.MB['forecast_free (admission, 100 bookings at the location)']['mean_ms']:.2f} ms; nearest-alternate over 49 lots {F.MB['nearest_alternate over 49 other locations']['mean_ms']:.1f} ms (in-process).",
    ], 8), size=14)
    notes[13] = "The two decisions the engine makes, in code order. Implemented in app/engine.py and the reservation routes in app/main.py."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User - System Interaction")
    wipe(s)
    pic(s, F.diag("07_sequence_diagram.png"), 0.55, 1.8, 6.4, 5.1)
    pic(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[14] = "Sequence of one booking and the use cases of the driver, administrator and payment provider."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design (5 tables)")
    wipe(s)
    pic(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.7, 4.6)
    rows = [["Table", "Content"], ["users", "role, trust score, streak, warnings, credit, suspension"], ["parking_locations", "coordinates, capacity, buffer, price, last reset"],
            ["reservations", "window, status, fee, multiplier, check times, parent_id"], ["payments", "amount, provider, Stripe ids, status"], ["trust_events", "benefit / warning / reward / penalty log"]]
    ftable(s, 7.5, 1.9, [2.2, 3.1], rows, size=11, row_h=0.55)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "parent_id links an automatically created alternate booking to the booking it replaced; it is a plain integer, not a foreign key.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Five tables; one SQLite file; schema created by SQLAlchemy at start-up."

    # ---- 16 implementation, performance & testing
    s = S[15]
    set_title(s, "Implementation, Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([("22", "API routes, 5 tables", K.FBLUE), (f"{F.PYTEST_PASSED} + {F.SMOKE_PASSED}", "pytest tests + smoke checks, all pass", K.GREEN),
                                   (f"{F.PAR['total_cases']:,}", f"simulator vs engine decisions, {F.PAR['total_mismatches']} mismatches", K.ORANGE), (f"{F.lat('GET /api/search'):.1f} ms", "mean availability search", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 26)
    names = [("GET /api/search", "Search availability"), ("POST /api/reservations", "Reserve"), ("POST /api/reservations/{id}/checkin", "Check in"),
             ("POST /api/reservations/{id}/checkout", "Check out"), ("GET /api/admin/insights", "Admin insights (60-day replay)"),
             ("GET /api/reservations", "Own bookings (N+1 queries)"), ("GET /api/admin/bookings", "Admin bookings (N+1 queries)"), ("POST /api/auth/login", "Login (bcrypt cost 12)")]
    rows = [["Endpoint (127.0.0.1, 200 calls)", "mean ms", "p95 ms"]]
    for pre, lab in names:
        rows.append([lab, f"{F.lat(pre):.1f}", f"{F.lat(pre, 'p95_ms'):.1f}"])
    ftable(s, 0.55, 3.15, [3.9, 1.0, 1.0], rows, size=11, row_h=0.38, align=["l", "r", "r"])
    ftext(s, 6.7, 3.15, 6.1, 3.8, bl([
        f"{F.PYTEST_ENGINE} engine unit tests (haversine, fees, admission, alternate, trust, suspension, reset), {F.PYTEST_API} API tests, {F.PYTEST_PARITY} parity tests.",
        "27-check smoke test on a throw-away database; 8-thread search throughput 131 req/s.",
        "Measured on an i5-12500, 15.7 GB, Python 3.14.7, single uvicorn process, SQLite file.",
        "Not tested: real load, many users, Stripe live mode, accessibility audit, other browsers.",
    ], 7), size=13)
    notes[16] = "Implementation facts and the testing that was actually run, with its limits. Own-booking lists are slow because each row issues two extra queries."

    # ---- 17 results: buffer policy vs population
    s = S[16]
    set_title(s, "Results - Buffer Policies vs Population (25 % bad vehicles)")
    wipe(s)
    rows = [["Policy (100 vehicles, 1,000 requests)", "Refused", "Suspended", "No park", "Success"]]
    for st in F.STRATS:
        rows.append([F.STRAT_NAME[st], f"{pk(100, st, 'no_reservation'):.0f}", f"{pk(100, st, 'suspended'):.0f}", f"{pk(100, st, 'no_park'):.1f}", f"{pk(100, st, 'success_rate') * 100:.0f}%"])
    ftable(s, 0.55, 1.8, [3.1, 0.9, 1.1, 0.9, 0.9], rows, size=12, row_h=0.4, align=["l", "r", "r", "r", "r"])
    ftext(s, 7.7, 1.8, 5.1, 2.5, bl([
        "Refused = declined by the admission rule; Suspended = refused by the Trust lockout; means of 30 runs.",
        "The shipped dynamic rule (+1 per on-time check-out) climbs to the maximum buffer and refuses the most.",
        "The paper-literal rule (+1 per late check-out) refuses 165, close to static 3 (158).",
        "Simulator, not SUMO: layout of the base paper, assumptions listed in the report.",
    ], 4), size=11)
    pic(s, F.fig("f01_buffer_population.png"), 0.45, 4.3, 12.4, 2.75)
    notes[17] = "The central measurement. Say plainly that this is a simulator built from the paper's scenario, with ParkSure's own rules, and that the paper's claim that the dynamic buffer is best is not reproduced for the shipped rule."

    # ---- 18 arrival outcomes (confusion-matrix equivalent)
    s = S[17]
    set_title(s, "Arrival Outcomes at the Gate (50 % bad vehicles)")
    wipe(s)
    rows = [["Policy", "Parked directly", "Moved to alternate", "of which alt. full", "No park", "Refused (admission)", "Suspended"]]
    for st in F.STRATS:
        parked = bd(0.5, st, "parks_completed") - bd(0.5, st, "moved_to_alternate")
        rows.append([F.STRAT_NAME[st], f"{parked:.0f}", f"{bd(0.5, st, 'moved_to_alternate'):.1f}", f"{bd(0.5, st, 'alt_overflow'):.1f}", f"{bd(0.5, st, 'no_park'):.2f}", f"{bd(0.5, st, 'no_reservation'):.0f}", f"{bd(0.5, st, 'suspended'):.0f}"])
    ftable(s, 0.55, 1.8, [2.6, 1.6, 1.8, 1.7, 1.0, 1.9, 1.3], rows, size=12, row_h=0.42, align=["l", "r", "r", "r", "r", "r", "r"])
    pic(s, F.fig("f03_arrival_outcomes.png"), 0.55, 4.45, 5.4, 2.55)
    ftext(s, 6.2, 4.5, 6.6, 2.5, bl([
        f"Mean of 30 runs, population 80-110 ({bd(0.5, 'static_0', 'requests'):.0f} requests per run); this table plays the role of a confusion matrix.",
        "'No park' is at most about 1 per run: the nearest-alternate step turns almost every full arrival into a served one.",
        "The shipped code does not re-check the alternate's physical bays, so 'alt. full' counts drivers who would meet a full lot again (static 0: 2.9 per run).",
    ], 6), size=12)
    notes[18] = "Outcome table for the arrival stage. A driver who is moved to an alternate that is physically full is a hidden failure of the shipped design."

    # ---- 19 parameter analysis
    s = S[18]
    set_title(s, "Parameter Analysis - Bad-Vehicle Share and Trust Score")
    wipe(s)
    pic(s, F.fig("f02_buffer_badshare.png"), 0.45, 1.75, 12.4, 3.0)
    rows = [["Success rate at 50 % bad", "Trust on", "Trust off"]] + [[F.STRAT_NAME[st], f"{F.abl(0.5, True, st, 'success_rate') * 100:.0f}%", f"{F.abl(0.5, False, st, 'success_rate') * 100:.0f}%"] for st in F.STRATS]
    ftable(s, 0.55, 4.85, [2.9, 1.0, 1.0], rows, size=11, row_h=0.35, align=["l", "r", "r"])
    ftext(s, 6.2, 4.85, 6.6, 2.2, bl([
        f"25 % to 50 % bad vehicles: shipped-dynamic success {bd(0.25, 'dynamic_parksure', 'success_rate') * 100:.0f}% to {bd(0.5, 'dynamic_parksure', 'success_rate') * 100:.0f}%; Trust-lockout refusals {bd(0.25, 'static_0', 'suspended'):.0f} to {bd(0.5, 'static_0', 'suspended'):.0f} per run (static 0).",
        "With the Trust Score off, success is far higher: the lockout, not the buffer, is the largest source of refused requests.",
        "The lockout lasts at least 14 days after the last warning (base paper: one session).",
        "Buffer maximum, B = 3 and W = 2 were not varied.",
    ], 4), size=11)
    notes[19] = "Parameter analysis of the scenario variables. The Trust ablation shows how much of the refusal count comes from the lockout."

    # ---- 20 trust dynamics, seeded replay, probes
    s = S[19]
    set_title(s, "Trust Dynamics, Seeded Replay and Robustness Probes")
    wipe(s)
    pic(s, F.fig("f04_trust_dynamics.png"), 0.45, 1.75, 6.2, 2.7)
    pic(s, F.fig("f06_seeded_insights.png"), 6.75, 1.75, 6.1, 2.7)
    race = F.RACE
    ftext(s, 0.55, 4.55, 6.0, 2.5, bl([
        f"A driver who ends 30 % of sessions late finishes at trust {F.TD['0.3']['mean_final_trust']}, fee x{F.TD['0.3']['mean_fee_multiplier']}; {F.TD['0.3']['share_ever_suspended'] * 100:.0f}% are suspended at some point.",
        f"Seeded replay (synthetic): declined {SD['insights_60d']['totals']['fixed_0']['declined']} / {SD['insights_60d']['totals']['fixed_1']['declined']} / {SD['insights_60d']['totals']['fixed_3']['declined']} / {SD['insights_60d']['totals']['dynamic']['declined']} for no buffer / 1 / 3 / dynamic.",
    ], 6), size=12)
    ftext(s, 6.75, 4.55, 6.1, 2.5, bl([
        f"Race: {race['trials_with_overbooking']} of {race['trials']} trials accepted 2 bookings for 1 bookable bay (20 simultaneous requests).",
        "Unauthenticated GET /api/payments/demo-confirm marks any booking paid; a token forged with the default JWT secret is accepted; no login rate limit.",
        "Past-start bookings and 25-day-early check-ins are accepted; injection payloads returned 422; cross-user cancel blocked (404).",
    ], 5), size=11)
    notes[20] = "Trust dynamics, the synthetic replay and the defects found by the team's own probes. These are limitations, not achievements."

    # ---- 21 comparison with base papers
    s = S[20]
    set_title(s, "Comparison with the Base Papers (different problems and settings)")
    wipe(s)
    rows = [["Aspect", "Jemmali et al. (registered base)", "Cabri et al. (rules implemented)", "ParkSure (this work)"],
            ["Problem", "Zone load balancing by passenger count", "Overstay-proof reservations: buffer + star reward", "General reservation platform with the buffer, alternate and trust ideas"],
            ["Setting", "9 instance classes, 2,790 generated cases", "SUMO, 80 spots in 8 areas, 5 runs", "Own simulator (same layout, not SUMO), 30 runs; synthetic 60-day history"],
            ["Result", "GVb ties or beats MR on 55.4 %; sub-millisecond", "Any buffer beats none; dynamic best trade-off (figures only)", f"Shipped dynamic refuses most ({pk(100, 'dynamic_parksure', 'no_reservation'):.0f} vs {pk(100, 'static_0', 'no_reservation'):.0f} of 1,000); no park rare"],
            ["Product", "Architecture described (QR, IoT, engine)", "Concept and simulation only", f"Working app: 22 routes, {F.PYTEST_PASSED} tests, 24 screens"],
            ["Speed", "Sub-millisecond heuristics", "Not reported", f"Admission {F.MB['forecast_free (admission, 100 bookings at the location)']['mean_ms']:.2f} ms; reserve call {F.lat('POST /api/reservations'):.1f} ms"],
            ["Agreement", "-", f"Partly: static 3 lowers no park ({bd(0.5, 'static_0', 'no_park'):.2f} to {bd(0.5, 'static_3', 'no_park'):.2f} per run at 50 % bad); the shipped dynamic rule does not ({bd(0.5, 'dynamic_parksure', 'no_park'):.2f})", "Dynamic-best claim not reproduced for the shipped rule"]]
    ftable(s, 0.55, 1.8, [1.5, 3.4, 3.6, 3.7], rows, size=11, row_h=0.62)
    ftext(s, 0.55, 6.3, 12.2, 0.65, "Different problems, data and simulators: no figure of ParkSure is claimed to beat either paper. Jemmali's grouping heuristics are not implemented; Cabri's SUMO runs were not repeated.", size=12, color=K.FNAVY, italic=True)
    notes[21] = "Say plainly that the problems and settings differ. Only the qualitative buffer claim can be compared, and it is not reproduced for the shipped rule."

    # ---- 22-24 output screens
    def two_shots(slide, title, a, b, cap_a, cap_b, box_a=(256, 56, 1440, 900), box_b=(256, 56, 1440, 900)):
        set_title(slide, title)
        wipe(slide)
        K.picture(slide, crop_shot(a, box_a), 0.45, 1.75, 6.0, 3.75, border=True)
        K.picture(slide, crop_shot(b, box_b), 6.85, 1.75, 6.0, 3.75, border=True)
        ftext(slide, 0.45, 5.65, 6.0, 1.0, cap_a, size=13, color=K.FNAVY, bold=True)
        ftext(slide, 6.85, 5.65, 6.0, 1.0, cap_b, size=13, color=K.FNAVY, bold=True)

    two_shots(S[21], "Output - Search and Reserve", "07_search_full_lot.png", "08_reserve_dialog.png",
              "Search for a 1 h 20 min window: Kiosk Bay (1 bay, buffer 0) is Full, five other locations show free bays",
              "Reserve dialog: Airport Express Park, 3 h at $5.00/h, trust multiplier x0.95, estimated fee $14.25 (the server recalculates)",
              box_a=(256, 80, 1440, 800), box_b=(380, 200, 1060, 700))
    notes[22] = "Real screens from the running app on a fresh seeded database."
    two_shots(S[22], "Output - Full Lot on Arrival and Trust Score", "15_moved_to_alternate.png", "13_trust_score.png",
              "Second driver arrives at the full Kiosk Bay and is moved to Central Station Garage, 0.1 km away; booking #353 is marked Unable to park, #354 is parked",
              "Trust Score page: 4 / 5, fee multiplier x0.95, on-time streak 1 of 3, history with the new on-time check-out")
    notes[23] = "The alternate-move screen is produced by the real engine, not a mock-up."
    two_shots(S[23], "Output - Administrator", "17_admin_overview.png", "19_admin_insights.png",
              "Admin overview: 10 users, 6 locations, 354 reservations, $1604.00 paid, 40 unable to park, 29 declined",
              "Strategy insights: the same 60-day history replayed under four buffer policies; the dynamic policy declines the most")
    notes[24] = "Overview and Strategy Insights; the numbers come from the seeded synthetic history plus the demo run."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.8, "ParkSure is a working, tested reservation platform whose own measurements show where its engine helps and where it falls short.", size=17, color=K.FNAVY, bold=True)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Registration, authentication, vehicle profile", "Partly met: registration, JWT login, roles; no vehicle profile"],
            ["2", "Real-time slot availability", "Met as forecast bookable bays per window; no sensors"],
            ["3", "Search, advance reservation, cancellation", "Met: tested end to end"],
            ["4", "Secure online payment, digital receipts", "Partly met: Stripe test mode / demo; no receipt; demo-confirm route is open"],
            ["5", "Administrator dashboard", "Met: locations, users, bookings, revenue, Strategy Insights"],
            ["6", "Success rate, efficiency, response time", "Met by simulator, replay and timing; no real-user study"]]
    ftable(s, 0.55, 2.6, [0.5, 4.5, 7.2], rows, size=12, row_h=0.5, align=["c", "l", "l"])
    ftext(s, 0.55, 6.2, 12.2, 0.8, "Key findings: the shipped dynamic buffer is the most conservative policy; no park is rare because of the alternate step; the Trust lockout drives most refusals; concurrency and payment-confirm defects remain.", size=13, color=K.FNAVY, italic=True)
    notes[25] = "Be explicit about which objectives were met, partly met or not met."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 4.6, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, bl([
        "Demo data are synthetic; no real lot, no user study",
        "Simulator, not SUMO; arrival-process assumptions are ours",
        "Buffer grows on on-time check-outs (paper: on late departures)",
        "Admission check and insert are not atomic (over-booking race)",
        "Open demo-confirm route, default JWT secret, no rate limit",
        "No vehicle profile, no receipt, no hardware or sensors",
        "Own-booking lists are slow (N+1 queries)",
    ], 5), size=13)
    fbox(s, 6.8, 1.8, 6.0, 4.6, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, bl([
        "Lock or transaction around admission; remove the open payment route",
        "Test the paper-literal buffer rule and a short, session-based lockout",
        "Re-check the alternate's physical bays before moving a driver",
        "Vehicle profile and downloadable receipts; reject past windows",
        "Replace the N+1 queries with joined loads",
        "Run the scenario in SUMO and on real lot data; user study",
        "Entry barriers or sensors, mobile client, PostgreSQL",
    ], 5), size=13)
    notes[26] = "Limitations first, then concrete next steps drawn from the measured defects."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "cabri", "aljohani", "asap", "doulamis", "safitri", "somani", "mejri", "kotb", "chen", "sumo", "sinnott", "jwt", "bcrypt", "sqla", "stripe"]
    short_ = [f"[{_IDX[k]}] {clean_ref(REFS[_IDX[k] - 1])}" for k in keys]
    short_ = [t if len(t) < 300 else t[:297] + "..." for t in short_]
    ftext(s, 0.55, 1.75, 6.0, 5.2, [(t, {"space_after": 4}) for t in short_[:8]], size=10)
    ftext(s, 6.8, 1.75, 6.0, 5.2, [(t, {"space_after": 4}) for t in short_[8:]] + [(f"The full numbered list of {len(REFS)} references, including the batch survey titles {N('s2')}-{N('s5')} (not verified), is in the report.", {"italic": True})], size=10)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh_ in S[27].shapes:
        if sh_.name == "Text 2":
            K.replace_text(sh_, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "ParkSure")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
