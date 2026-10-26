"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\Abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import N, REFS, _IDX

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "parking reservation, slot booking, smart parking, web application, online payment, admin dashboard"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"Guide: {F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


# Problem statement and objectives are taken from 6.OBJECTIVES&PROBLEM STATEMENT.docx
PROBLEM_STATEMENT = [
    "Drivers circle repeatedly because there is no way to know which slots are free",
    "Parking operators manage occupancy through manual registers with no live view",
    "Without reservation, arriving at a facility gives no guarantee of a space",
    "Cash-based collection is slow at entry and provides no usable records",
    "Need: a reservation platform that shows live availability, guarantees a slot in advance, and handles payment digitally",
]
OBJECTIVES = [
    "Build user registration, authentication and vehicle profile management",
    "Maintain real-time parking slot availability across registered parking facilities",
    "Implement slot search, advance reservation and booking cancellation",
    "Integrate secure online payment and generate digital receipts",
    "Provide an administrator dashboard for slots, bookings, occupancy and revenue",
    "Evaluate booking success rate, allocation efficiency and system response time",
]

# [work, year, approach, reported result / limitation]  (only what the base papers state about each work)
LIT_ROWS = [
    [f"Jemmali et al. {N('base')} (base paper)", "2026", "Passenger-count-based zone balancing: four-stage grouping and six grouping-permutation heuristics",
     "2,790 generated instances; GVb matches or beats baseline MR on 55.4 %; one managed facility, no payment or overstay handling"],
    [f"Cabri et al. {N('cabri')}", "2026", "Buffer of non-reservable slots plus a 0-5 star reward system; one alternate area suggested",
     "SUMO, 80 spots, 8 areas: any buffer beats none, dynamic buffer best trade-off (figures only); simulation only"],
    [f"Doulamis et al. {N('doulamis')}", "2013", "Interval scheduling: requests are time intervals accepted or rejected", "No numbers in the base paper's summary; no overstay handling described"],
    [f"Safitri et al. {N('safitri')}", "2020", "Mobile reservation with occupancy display, heuristic and haversine formula", "Mobile client only; no buffer or reputation mechanism"],
    [f"Somani et al. {N('somani')}", "2018", "Cross-platform reservation; part of the fee paid to confirm a booking", "Deposit deters false bookings but not overstays"],
    [f"Mejri et al. {N('mejri')}", "2016", "Slot assignment as a multi-objective integer linear programme", "Optimisation model; no operating web platform described"],
    [f"Kotb et al. {N('kotb')}", "2016", "iParker: flexible reservation times, dynamic allocation, revenue and pricing", "Pricing by demand, not by the driver's punctuality"],
    [f"Rizvi et al. {N('asap', 'mapark')}", "2019-20", "Agent-assisted and multi-agent auction-based parking over IoT", "Auctions suit competitive allocation, not a fixed-price booking service"],
    [f"Yan et al. {N('yan')}", "2011", "SmartParking: on-the-fly spot reservation with real-time availability", "Security-focused architecture; no overstay mitigation described"],
    [f"Batch survey {N('s2')}-{N('s5')}", "2024-26", "Web and mobile reservation systems, slot prediction, face-recognition theft prevention (titles only)", "Full texts not located; no figures quoted"],
]

GAPS = [
    f"The registered base paper optimises allocation inside one large managed facility; it has no booking web service, payment or operator analytics {N('base')}",
    f"The buffer and star-reputation idea against overstays is evaluated only in simulation and is preliminary by its own words {N('cabri')}",
    f"Reservation components appear one at a time in the literature: booking, payment, mobile access, pricing {N('safitri', 'somani', 'kotb')}",
    "A general platform for ordinary urban parking that combines search, booking, payment, an overstay-aware availability engine and operator analytics is not reported",
    "Reported systems rarely measure their own response time, booking success rate and defects; this project measures them and states the failures",
]

ABSTRACT_POINTS = [
    ("The problem", "The growing number of vehicles in urban areas makes finding a parking space hard: more congestion, wasted time and poorly used facilities. Operators work from manual registers with no live view.", K.ORANGE, "!"),
    ("Our idea", "A web application where users search available slots, reserve in advance, pay online and manage bookings. It is designed to prevent double bookings through an automated reservation process and live availability (concurrency limits are measured later).", K.BLUE, "i"),
    ("The deliverable", "ParkSure: a secure, reliable and user-friendly platform with an admin panel for locations, slots, users, bookings and payment records.", K.GREEN, "✓"),
]


def pic_or_note(s, name, x, y, w, h):
    p = F.diag(name)
    if os.path.exists(p):
        K.picture(s, p, x, y, w, h)
        return True
    print("WARNING: missing", p)
    return False


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # 1 title
    title_slide(S[0], "Smart parking reservation: search, book, pay and check in", "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: ParkSure lets a driver search live availability, reserve a bay in advance, pay online and check in, "
                "and gives the operator a dashboard.")

    # 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. The review covers why the problem matters and what is missing in current work, then the approach, the architecture and what will be measured."

    # 3 abstract
    s = S[2]
    set_title(s, "Abstract: reserve parking before you arrive")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate(ABSTRACT_POINTS):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=20)
    K.text(s, 0.6, 6.6, 12.13, 0.5, "Keywords: online vehicle parking, reservation system, slot booking, online payment, user authentication, admin dashboard, smart parking.", size=12, color=K.TXT, italic=True)
    notes[3] = "Three points from the project abstract: the urban parking problem, the web reservation platform, and the admin panel with live availability."

    # 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.35, 12.13, 1.05, fill=K.CARD)
    K.text(s, 0.82, 1.4, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Parking spaces are allocated by the total number of passengers per vehicle. A four-stage grouping procedure and six grouping-permutation heuristics balance the number of people across parking zones. Booking interface, IoT sensing and a central allocation engine.", K.BLUE, "M"),
        ("Reported results", "Nine instance classes, 2,790 generated test cases. The best algorithm has the highest best-solution share (68.2 %). GVb matches or beats the baseline MR on 1,547 instances (55.4 %), in sub-millisecond time.", K.GREEN, "R"),
        ("What ParkSure takes and adds", "Takes the booking-plus-central-allocation architecture. Does not implement passenger grouping. Adds a general web reservation service with payment, an overstay-aware engine and operator analytics.", K.ORANGE, "A"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.6, cw, 3.15, t, b_, col, gl, body_size=14)
    K.box(s, 0.6, 5.95, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 5.97, 11.7, 1.05, f"Technical basis of the implemented rules: {F.TECH_PAPER} Reservation module with a buffer of non-reservable slots and a 0-5 star reward system; evaluated in SUMO with 80 spots in 8 areas {N('cabri')}.",
           size=12, color=K.INK, anchor="m")
    notes[4] = ("The registered base paper is the IEEE OJCS paper on medical-city parking. The rules that ParkSure actually implements for overstays, the buffer and the star reputation come from the "
                "second paper, so both are shown. ParkSure does not claim to improve the grouping algorithms.")

    # 5 literature table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]] + LIT_ROWS
    K.table(s, 0.6, 1.2, [2.55, 0.75, 4.55, 4.28], rows, font_size=10, row_h=0.5, header_size=11, align=["l", "c", "l", "l"])
    notes[5] = ("Details about related works are only what the two base papers state about them. The four batch-survey titles could not be located online, so no results are quoted from them.")

    # 6 supporting evidence
    s = S[5]
    set_title(s, "Supporting evidence from the base papers")
    wipe(s)
    cats = ["GVb better than MR", "GVb equal to MR", "MR better than GVb"]
    K.bar_chart(s, 0.6, 1.4, 6.6, 4.9, cats, [("Instances (of 2,790)", [57, 1490, 1243])], colors=[K.BLUE], number_format="#,##0", font_size=11, max_value=1700, legend=False)
    K.text(s, 0.6, 6.35, 6.6, 0.7, f"Base paper {N('base')}: instances out of 2,790 where the proposed GVb is better, equal or worse than baseline MR (2.0 %, 53.4 %, 44.6 %).",
           size=11, color=K.TXT, italic=True)
    K.text(s, 7.5, 1.45, 5.25, 5.0, [
        ("**What the literature shows**", {"space_after": 8}),
        (f"Allocation heuristics can run in sub-millisecond time and reach the best solution on 68.2 % of instances {N('base')}.", {"bullet": True, "space_after": 8}),
        (f"In the SUMO study with 80 spots in 8 areas, any buffer beats no buffer; the dynamic buffer gives the best trade-off between failed parks and rejected reservations {N('cabri')}.", {"bullet": True, "space_after": 8}),
        (f"Static buffers trade one failure for the other: larger buffers cut failed parks but refuse more reservations {N('cabri')}.", {"bullet": True, "space_after": 8}),
        ("Caveat: both studies use settings different from a deployed web service; their numbers are not put on one scale with ParkSure's own measurements.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("This is published evidence, on different data and settings from ours. The second paper reports its buffer results as figures only, so no numbers are quoted from it.")

    # 7 existing system
    s = S[6]
    set_title(s, "Existing system and current practice")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Driver circles to find a bay", K.NAVY), ("Arrives with no guarantee", K.NAVY), ("Cash paid at entry", K.NAVY), ("Manual register entry", K.NAVY), ("No live view for the operator", K.NAVY)], size=14, gap=0.45)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "For the driver",
           "There is no way to know which slots are free before arriving.\nArriving at a facility gives no guarantee of a space.\nCash collection is slow at the entry gate.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "For the operator",
           "Occupancy is kept in manual registers with no live view.\nCash payments leave no usable records.\nDemand cannot be seen, so capacity cannot be planned.", K.ORANGE, "2", body_size=17)
    notes[7] = "Current practice at ordinary urban facilities, taken from the project introduction and problem statement."

    # 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Repeated circling", "Manual, invisible occupancy", "No guarantee on arrival", "Cash and no records"]
    bodies = [
        "Drivers spend a significant share of each trip searching, which adds congestion, wasted fuel and avoidable emissions.",
        "Operators manage occupancy through registers and have no live view or demand data to plan capacity.",
        "Without advance reservation a driver cannot be sure of a space at the facility.",
        "Cash collection slows the entry and provides no usable digital record.",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=16)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"Reservation research adds: facility-level optimisation without a booking service {N('base')}, and a buffer idea evaluated only in simulation {N('cabri')}.", size=14, color=K.INK, anchor="m")
    notes[8] = "The four limitations are the batch problem statement; the last line is what the two base papers leave open."

    # 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=16, color=K.INK, anchor="m")
    notes[9] = ("The first gap is the survey conclusion of the batch: allocation inside large managed facilities, but no general reservation platform for ordinary urban parking. "
                "Gaps 2 to 5 come from reading the two base papers and the related work they cite.")

    # 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Registration details, a search window (arrival and departure), a chosen location, payment"),
                                 ("Output", "Live bookable bays per location, a confirmed reservation with fee, check-in and check-out, operator dashboard"),
                                 ("Measures", "Booking success rate, allocation efficiency and system response time")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=13, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a reservation platform with live availability, a guaranteed slot in advance and digital payment."

    # 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "Later reviews and the report state which objectives were met, partly met or not met, with measurements.",
           size=13, color=K.TXT, italic=True)
    notes[11] = "The six objectives from the proposal. The final presentation states plainly which are met."

    # 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Search and reserve", "Search a time window; each location shows the bays that can still be reserved. A booking is admitted only if overlapping bookings plus a hold-back buffer stay below capacity.", K.BLUE, "1"),
        ("Pay", "Fee = hours x rate x a multiplier from the driver's Trust Score. Stripe Checkout in test mode, or an offline demo checkout, with account credit applied.", K.GREEN, "2"),
        ("Arrive and park", "Check in at the lot. If it is physically full the driver is moved to the nearest alternate by haversine distance, or the fee is credited. Check out updates buffer and Trust Score.", K.ORANGE, "3"),
        ("Operate", "Admin panel for locations, users, bookings and revenue, and Strategy Insights: stored bookings replayed under four buffer policies.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React console + SQLite database", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "22 API routes, 5 tables, JWT login with driver and admin roles, a scheduler that resets buffers daily. The rules follow the buffer and star-reputation ideas of the second base paper; the approach is implemented in the project code.", size=14, color=K.TXT)
    notes[12] = "Four stages and the stack. The buffer and reputation rules follow the arXiv paper; search, booking, payment and the admin panel complete the platform."

    # 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    pic_or_note(s, "01_system_architecture.png", 0.6, 1.2, 7.6, 5.9)
    K.text(s, 8.4, 1.5, 4.35, 5.2, [
        ("**Layers**", {"space_after": 8}),
        (f"React console on port {F.WEB_PORT}; Vite proxies /api.", {"bullet": True}),
        (f"FastAPI on port {F.API_PORT}: 22 routes, JWT and role guards.", {"bullet": True}),
        ("Smart Availability Engine: admission, buffer, nearest alternate, Trust Score, fees.", {"bullet": True}),
        ("SQLite database with five tables.", {"bullet": True}),
        ("Scheduler job and Stripe test-mode payment as side services.", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the console calls the API, the API calls the engine and the database."

    # 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and what will be measured")
    wipe(s)
    heads = ["A working platform", "Reservation behaviour", "Buffer policies", "Response time", "Stated limits"]
    bodies = [
        "Registration and login, live search, reserve, cancel, pay, check in and out, and the admin dashboard, tested with an automated test suite.",
        "Booking success rate and allocation efficiency (bay utilisation), measured on a simulated scenario that follows the setting of the second base paper.",
        "Static and dynamic hold-back buffers compared on refused reservations and arrivals that find no bay. Results are not claimed in advance.",
        "Latency of the main API calls measured on one machine, and the cost of the engine decisions.",
        "Anything not achieved or not measured will be reported as a gap: for example no live sensors, no real-user study and no hardware gate.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=14)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=15)
    notes[14] = ("This is the problem-identification review, so the slide states what will be delivered and measured, not results. Results are in Review 2 and the final presentation.")

    # 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    keys = ["base", "cabri", "aljohani", "asap", "mapark", "doulamis", "kokolaki", "yan", "safitri", "somani", "mejri", "kotb", "s2", "s3", "s4", "s5"]
    keys = sorted(keys, key=lambda k: _IDX[k])
    items = [(f"[{_IDX[k]}] {REFS[_IDX[k] - 1]}", {"space_after": 4}) for k in keys]
    half = len(items) // 2
    K.text(s, 0.6, 1.25, 6.0, 5.8, items[:half], size=11.5, color=K.INK)
    K.text(s, 6.75, 1.25, 6.0, 5.8, items[half:], size=11.5, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why the registered base paper and the implemented rules come from two different papers (the batch registered the IEEE OJCS paper; the overstay buffer and star rules are in the arXiv paper); "
                 "what the hold-back buffer is; what happens when a lot is full on arrival.")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
