"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (letterhead, rule, title and footer chrome inherited)."""

import os
import tempfile

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import lit
from . import pptkit as K
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26, lsize=12):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=lsize, color=K.DARK, font=FONT)


CROPS = os.path.join(tempfile.gettempdir(), "c4_deck_crops")


def crop_shot(name, box=(256, 56, 1440, 900)):
    """Real screenshot with the sidebar cropped away (left, top, right, bottom in pixels) so the content is larger on the slide."""
    from PIL import Image

    os.makedirs(CROPS, exist_ok=True)
    out = os.path.join(CROPS, f"{os.path.splitext(name)[0]}_crop.png")
    with Image.open(F.shot(name)) as im:
        im.crop(box).save(out)
    return out


def clean_ref(s):
    return " ".join(s.replace("*", "").split())


def seedwise(tier, key):
    """Ten-ticker mean per seed, then mean over seeds."""
    S = len(F.RR["seeds"])
    return float(np.mean([np.mean([F.RR["tickers"][t]["models"][tier][s][key] for t in F.TICKERS]) for s in range(S)]))


def confusion(model="lstm_mil_attention"):
    """2x2 counts reconstructed from the seed-0 precision, recall and the test size (up = positive class)."""
    n = F.MIL["n_test_pooled"]
    run = F.MIL["models"][model]["runs"][0]
    pos = round(F.MIL["up_share_test"] * n)
    tp = round(run["recall"] * pos)
    pp = round(tp / run["precision"]) if run["precision"] > 0 else 0
    fp, fn = pp - tp, pos - tp
    tn = n - tp - fp - fn
    return tp, fp, fn, tn


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE"
    cp.comments = ""
    cp.last_modified_by = f"Batch {F.BATCH_NO}, Dept. of CSE"
    cp.subject = F.PRODUCT
    cp.keywords = "stock market prediction, LSTM, attention, FinBERT, federated averaging, multiple instance learning"
    notes = {}
    bp = F.BASE_PAPER_RESULTS
    T = F.TICKERS
    SUM, FWD, RR, MIL, API = F.SUM, F.FWD, F.RR, F.MIL, F.API
    ep = API["endpoints"]
    pers = float(np.mean([RR["tickers"][t]["persistence"]["mae"] for t in T]))
    mae_true = {k: seedwise(k, "mae_true") for k in F.TIER_LABEL}
    mae_ship = {k: seedwise(k, "mae") for k in F.TIER_LABEL}
    da_true, da_ship = seedwise("lstm_attention", "da_true"), seedwise("lstm_attention", "directional_accuracy")
    sh = SUM["models"]
    fs, fl = FWD["short"], FWD["long"]
    mm = MIL["models"]
    n_test = MIL["n_test_pooled"]
    always_up = MIL["always_up_accuracy_test"]
    pc = MIL["paired_vs_lstm_price"]
    rf_wins = SUM["best_mae_wins"].get("random_forest", 0)
    down5 = sum(1 for v in SUM["trend"].values() if v == "down")
    fc5 = float(np.mean(list(SUM["forecast_5d_change_pct"].values())))
    above = [RR["tickers"][t]["test_close_above_train_max_pct"] for t in T]
    n_above = sum(1 for a in above if a > 50)
    sent_train = sum(RR["tickers"][t]["sentiment_nonzero_in_train_part"] for t in T)

    for sl in S:  # footers on every slide except the title
        for sh_ in sl.shapes:
            if sh_.top == TOP_FOOT and sh_.has_text_frame and "Final Review" in sh_.text_frame.text:
                K.replace_text(sh_, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    for sh_ in S[0].shapes:
        if sh_.name == "Text 1":
            K.replace_text(sh_, F.FULL_TITLE, 28)
        elif sh_.name == "Text 2":
            K.replace_text(sh_, f"Final Review  ·  {F.PRODUCT}")
        elif sh_.name == "Text 3":
            K.replace_text(sh_, f"Team Details : Batch {F.BATCH_NO}")
        elif sh_.name == "Text 4":
            K.replace_text(sh_, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS_UPPER))
        elif sh_.name == "Text 5":
            K.replace_text(sh_, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh_.name == "Text 6":
            K.replace_text(sh_, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and QuantPulse: a working forecasting product that follows the base paper, plus an offline multiple-instance-learning study on real headlines."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh_ in s.shapes:
        if sh_.name == "Text 4":
            K.replace_text(sh_, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Data & Algorithm Workflow\nInteraction & Data Design\nImplementation Highlights\nResults: Forecasts, Scoring, MIL Study\nParameter Analysis, Performance & Testing\nComparison with Base Paper\nOutput Screens\nConclusion, Limitations, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 5.1, [
        ("Price-only forecasts ignore news and sentiment, and labelling every article is impractical. The project asks for multi-source prediction with Multiple Instance Learning (MIL). QuantPulse, the product, implements the base paper's pipeline: LSTM, multi-head attention, FinBERT sentiment and simulated FedAvg, as a FastAPI + React website for ten tickers.", {"bullet": True, "space_after": 8}),
        (f"Measured honestly, the forecasts did not hold up: a re-run on the same data scores MAE ${mae_true['lstm_attention']:.1f} against ${pers:.1f} for 'tomorrow = today'; the stored 5-day forecasts averaged {fs['mean_pred_change_pct']:+.1f}% against an actual {fs['mean_actual_change_pct']:+.1f}%, with the trend right for {fs['trend_hits']} of 10 tickers.", {"bullet": True, "space_after": 8}),
        (f"The app's directional score ({da_ship * 100:.0f}%) is inflated by clipped test targets; against real closes it is {da_true * 100:.0f}%. Its sentiment channel is zero in training and synthetic in test.", {"bullet": True, "space_after": 8}),
        (f"An offline MIL study on {MIL['instances']['n_headline_rows_used']:,} real FNSPID headlines (7 tickers, {n_test:,} test days) found no gain: the best MIL variant scored {mm['lstm_mil_attention']['summary']['accuracy']['mean'] * 100:.1f}% accuracy against {always_up * 100:.1f}% for always-up.", {"bullet": True, "space_after": 8}),
        (f"{F.PYTEST_PASSED} pytest tests and {F.SMOKE_PASSED} smoke checks pass. No social-media source, no MIL inside the product, no precision/recall/F1 in the app: stated on every results slide.", {"bullet": True}),
    ], size=16)
    notes[3] = "Give the product, then the honest headline results and the limits. The MIL study is offline; the product does not contain MIL."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Paper")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.8, F.BASE_PAPER, size=13, color=K.FNAVY)
    rows = [["Aspect", "Base paper: hybrid LSTM, Transformer attention, federated learning and sentiment analysis"],
            ["Models", "Baselines ARIMA, SARIMAX, Prophet, Random Forest, SVR vs LSTM (64 units + batch norm + dense); hybrid adds multi-head attention, FinBERT-tone sentiment and simulated FedAvg over ten clients"],
            ["Data", "Yahoo Finance for ten firms (Tesla, Apple, Airbnb, Amazon, FedEx, IBM, Microsoft, Nvidia, Oracle, Google); hybrid on a one-month window because NewsAPI returned 30 days of headlines"],
            ["Protocol", "Chronological split: 80/20 for baselines, 70/15/15 for the hybrid; 5- and 30-day horizons; R\u00b2, MSE, MAE and a directional measure"],
            ["Reported", f"Hybrid average R\u00b2 {bp['r2_avg']:.2f} (all companies above {bp['r2_min_company']:.2f}); TSLA R\u00b2 {bp['tsla_r2']:.3f}; baseline LSTM average trend accuracy {bp['lstm_trend_acc']:.2f}% (Airbnb 94%, Tesla 34.9%, IBM 18%)"],
            ["Stated limits", "Sentiment limited to 30 days; federated learning only simulated; performance varies across companies"]]
    ftable(s, 0.55, 2.6, [1.7, 10.5], rows, size=12, row_h=0.7)
    notes[4] = "Nejatbakhsh and Aliasgari, IEEE Access 2026. QuantPulse follows this pipeline. The paper's tables are images, so only figures stated in its text are quoted."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Finding / limitation"]] + [list(r) for r in lit.LIT_TABLE[:9]]
    ftable(s, 0.55, 1.75, [2.7, 0.8, 4.6, 4.1], rows, size=10, row_h=0.5, align=["l", "c", "l", "l"])
    notes[5] = "Rows marked as title-only could not be verified beyond bibliographic records, so no result is quoted for them. Numbers shown come from the PDFs we read."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 5.0, [(g, {"bullet": True, "space_after": 12}) for g in lit.GAPS], size=17)
    notes[6] = "The gap the batch asked us to address: an explicit multiple-instance treatment of a day's news, evaluated with direction metrics against a price-only baseline."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.85, 12.2, 2.7)
    ftext(s, 0.85, 1.9, 11.6, 2.6, lit.PROBLEM_STATEMENT, size=16, anchor="m", bullets=True)
    for i, (t, b_) in enumerate([("Input", "Daily OHLCV prices (Yahoo Finance) and financial headlines per ticker"), ("Output", "5- and 30-day forecasts with confidence bands, trend flag, model comparison and sentiment views"), ("Record", "metrics.json, forecasts.json, experiment JSON files and screenshots of the running website")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.8, 3.9, 1.6, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.88, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 5.35, 3.5, 1.0, b_, size=14)
    notes[7] = "The batch's problem statement, then the input, output and record of the product."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(lit.OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives from the batch document. The conclusion slide gives the status of each."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    ftext(s, 0.55, 1.75, 12.2, 0.4, "Data and features", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 2.15, 12.2, 0.85, [("Yahoo prices", K.FNAVY), ("20-day windows", K.FBLUE), ("Min-Max scaling", K.GREEN), ("Headlines", K.ORANGE), ("FinBERT-tone", K.PURPLE), ("Daily sentiment", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 3.2, 12.2, 0.4, "Models and serving", size=15, color=K.FNAVY, bold=True)
    K.flow(s, 0.55, 3.6, 12.2, 0.85, [("LSTM(64)", K.FNAVY), ("+ 4-head attention", K.FBLUE), ("+ sentiment channel", K.GREEN), ("FedAvg x10 clients", K.ORANGE), ("JSON artefacts", K.PURPLE), ("FastAPI + React", K.FNAVY)], size=13, gap=0.3, font=FONT)
    ftext(s, 0.55, 4.8, 12.2, 2.1, [
        ("Five classical baselines (ARIMA, SARIMAX, Prophet, Random Forest, SVR) are scored next to the four deep-model tiers on an 80/20 chronological split.", {"bullet": True, "space_after": 7}),
        ("Forecasts: 5- and 30-day recursive predictions with a 95% band, shown with actual closes on six screens.", {"bullet": True, "space_after": 7}),
        ("Multiple Instance Learning is not inside the product; it was studied offline on real headlines (bag = one day's news for one ticker).", {"bullet": True}),
    ], size=16)
    notes[9] = "The proposed system as two chains. Be explicit that MIL is an offline study and that the product's sentiment is one number per day."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.8, 7.2, 5.1)
    ftext(s, 7.95, 1.95, 4.9, 4.9, [
        (f"Website: React 19 + Vite, port {F.WEB_PORT}; six screens", {"bullet": True, "space_after": 8}),
        (f"API: FastAPI + Uvicorn, port {F.API_PORT}; 7 JSON endpoints and the static site", {"bullet": True, "space_after": 8}),
        ("Pipeline: data loader, sentiment, baselines, deep models, FedAvg, forecast modules", {"bullet": True, "space_after": 8}),
        ("State: JSON, CSV and Keras files; no database, no login", {"bullet": True, "space_after": 8}),
        ("The API serves stored results; the saved .keras models are not loaded at run time", {"bullet": True}),
    ], size=15)
    notes[10] = "Four tiers. No database and no authentication. The API serves metrics.json and forecasts.json."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.picture(s, F.diag("02_high_level_workflow.png"), 0.55, 1.8, 7.3, 4.9)
    ftext(s, 8.1, 1.9, 4.7, 5.0, [
        ("Window of 20 trading days of OHLCV (+ sentiment); Min-Max scalers fitted on the first 80% of rows.", {"bullet": True, "space_after": 8}),
        ("LSTM 64 units; attention tier adds 4 heads, residual, LayerNorm, pooling; batch 32, up to 50 epochs, early stopping.", {"bullet": True, "space_after": 8}),
        ("FedAvg: each ticker is a client; 2 rounds x 3 local epochs; weights averaged by sample count.", {"bullet": True, "space_after": 8}),
        ("Added here: five-seed re-run, persistence baseline, true-target scoring, forward test, MIL study.", {"bullet": True}),
    ], size=14)
    notes[11] = "Method in one picture, then the extra measurements that were added for the evaluation."

    # ---- 12 dataset
    s = S[11]
    set_title(s, "Dataset & Models")
    wipe(s)
    ftile(s, 0.55, 1.8, 3.9, 1.0, "10 x 501", "tickers x daily rows, 2024-09-09 to 2026-09-08", K.FBLUE, 24)
    ftile(s, 4.7, 1.8, 3.9, 1.0, f"{MIL['instances']['n_headline_rows_used']:,}", "real FNSPID headlines, 7 tickers (MIL study)", K.GREEN, 24)
    ftile(s, 8.85, 1.8, 3.9, 1.0, f"{sent_train}", "non-zero sentiment rows in the training part (shipped channel)", K.ORANGE, 24)
    rows = [["Data", "Size", "Used for"],
            ["Yahoo Finance OHLCV", "10 tickers", "all models; pinned snapshot for the experiments"],
            ["Synthetic template headlines", "166-191 per ticker", "sentiment channel (no NewsAPI key); last 90 days only"],
            ["FNSPID news [11]", "46.5k rows", "offline MIL study; dense only in 2022-2023"],
            ["Models", "9 + MIL", "5 baselines, 4 deep tiers; MIL variants offline"]]
    ftable(s, 0.55, 3.0, [2.1, 1.3, 3.5], rows, size=11, row_h=0.55)
    K.picture(s, F.fig("f06_sentiment_channel.png"), 7.8, 3.0, 5.0, 2.9)
    ftext(s, 0.55, 6.1, 12.2, 0.8, "The shipped sentiment channel is zero for every training row and carries synthetic headline scores only in the last 90 trading days, so it cannot show a sentiment effect (right).", size=12, color=K.FNAVY, italic=True)
    notes[12] = "No NewsAPI key was set, so the product uses labelled synthetic headlines. Real headlines are used only in the offline MIL study."

    # ---- 13 workflow
    s = S[12]
    set_title(s, "Forecast and Sentiment Workflow")
    wipe(s)
    K.picture(s, F.diag("10_flowchart_forecast.png"), 0.55, 1.8, 3.9, 5.1)
    K.picture(s, F.diag("11_flowchart_sentiment.png"), 4.6, 1.8, 4.6, 5.1)
    ftext(s, 9.4, 1.95, 3.5, 4.9, [
        ("Forecast: roll the predicted close forward, one step at a time.", {"bullet": True, "space_after": 8}),
        ("Band: 1.96 x residual std x sqrt(day).", {"bullet": True, "space_after": 8}),
        ("Sentiment: NewsAPI if a key exists, else synthetic templates; FinBERT-tone or lexicon; daily mean; fill forward, zero before the first headline.", {"bullet": True}),
    ], size=13)
    notes[13] = "The recursive forecast and the sentiment path as flowcharts."

    # ---- 14 interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_flow.png"), 0.55, 1.8, 6.4, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.1, 1.8, 5.7, 5.1)
    notes[14] = "Sequence of the Forecast screen flow and the use cases of the analyst and the evaluator."

    # ---- 15 data design
    s = S[14]
    set_title(s, "Data Design (no database)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 6.7, 4.6)
    rows = [["Store", "Content"], ["artifacts/metrics.json", "per-ticker metrics, histories, forecasts, headlines"], ["artifacts/forecasts.json", "5- and 30-day forecasts"], ["artifacts/serving/*.keras", "saved models (not loaded by the API)"],
            ["data/<T>.csv", "prices; rewritten on each live pull"], ["data/news_cache/", "NewsAPI cache (needs a key)"], ["experiments/**/metrics.json", "documentation experiments"]]
    ftable(s, 7.5, 1.9, [2.7, 2.6], rows, size=11, row_h=0.55)
    ftext(s, 0.55, 6.45, 12.2, 0.5, "No foreign keys and no persistence layer: the entities in the diagram are logical, files on disk.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "QuantPulse has no database. The ER diagram shows logical entities."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("7", "JSON endpoints", K.FBLUE), (f"{F.PYTEST_TOTAL} + {F.SMOKE_TOTAL}", "pytest tests + smoke checks", K.GREEN), ("6", "React screens", K.ORANGE), (f"{ep['forecast_short']['median_ms']:.0f} ms", "median forecast response", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.85, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.25, 12.2, 3.6, [
        ("Stored results are served from JSON; invalid symbols give 400, unknown 404, bad horizon 422, missing artefacts 503.", {"bullet": True, "space_after": 9}),
        ("The Forecast screen overlays the actual closes that followed the forecast date and warns that the models were trained up to 8 Sept 2026.", {"bullet": True, "space_after": 9}),
        ("/api/history downloads from Yahoo Finance on every request and rewrites the CSV; the first sentiment call loads FinBERT.", {"bullet": True, "space_after": 9}),
        ("run.bat sets up the virtual environment, trains if artefacts are missing and starts both servers; Docker files exist but were not run here.", {"bullet": True, "space_after": 9}),
        ("Not present: database, login, real federated deployment, social-media text, MIL inside the product.", {"bullet": True}),
    ], size=16)
    notes[16] = "Highlights and honest absences."

    # ---- 17 results: forecasts vs reality
    s = S[16]
    set_title(s, "Results – Forecasts versus Reality")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{fs['mean_pred_change_pct']:+.1f}%", f"mean stored 5-day forecast change (actual {fs['mean_actual_change_pct']:+.1f}%)", K.FBLUE),
                                   (f"{fs['trend_hits']} / 10", "5-day trend calls right (30-day: " + f"{fl['trend_hits']} / 10)", K.GREEN),
                                   (f"${fs['mean_mae']:.0f}", f"5-day MAE vs ${fs['mean_persistence_mae']:.0f} for 'no change'", K.ORANGE),
                                   (f"${pers:.1f}", f"persistence MAE vs ${mae_true['lstm_attention']:.1f} (LSTM + attention, 5 seeds)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.1, v, l, c, 26, lsize=11)
    K.picture(s, F.fig("f02_true_mae_vs_persistence.png"), 0.55, 3.0, 6.0, 3.1)
    K.picture(s, F.fig("f05_forecast_vs_actual_change.png"), 6.8, 3.0, 6.0, 3.1)
    ftext(s, 0.55, 6.2, 12.2, 0.8, f"Forward test: forecasts stored from the 2026-09-08 close against the 30 closes that followed. Shipped models: Random Forest has the lowest MAE on {rf_wins} of 10 tickers; no deep tier beats persistence on real closes. {down5} of 10 tickers were flagged 'down' (mean {fc5:+.1f}%).", size=12, color=K.FNAVY, italic=True)
    notes[17] = "The stored forecasts predicted a large fall that did not happen. Persistence beats every deep tier when scored against real closes."

    # ---- 18 classification-style results (MIL study)
    s = S[17]
    set_title(s, "Confusion Matrix & Precision / Recall / F1 (MIL Study)")
    wipe(s)
    names = [("majority_class", "Majority class"), ("logreg_price_C1.0", "Logistic regression, price"), ("lstm_price", "LSTM, price (same window)"), ("lstm_price_full_history", "LSTM, price (since 2016)"),
             ("lstm_flat_sentiment", "LSTM + flat sentiment"), ("lstm_mil_mean", "LSTM + MIL mean"), ("lstm_mil_max", "LSTM + MIL max"), ("lstm_mil_attention", "LSTM + MIL attention"), ("mil_attention_news_only", "MIL attention, news only")]
    rows = [["Model", "Acc.", "Prec.", "Recall", "F1"]]
    for k, nm in names:
        sm = mm[k]["summary"]
        rows.append([nm] + [f"{sm[m]['mean'] * 100:.1f}" for m in ("accuracy", "precision", "recall", "f1")])
    ftable(s, 0.55, 1.75, [2.9, 0.8, 0.8, 0.8, 0.8], rows, size=11, row_h=0.4, align=["l", "c", "c", "c", "c"])
    tp, fp, fn, tn = confusion()
    ftext(s, 7.4, 1.7, 5.4, 0.4, "LSTM + MIL attention, seed 0 (up = positive)", size=14, color=K.FNAVY, bold=True)
    ftable(s, 7.4, 2.15, [1.6, 1.8, 1.8], [["", "Predicted up", "Predicted down"], ["Actual up", f"{tp}", f"{fn}"], ["Actual down", f"{fp}", f"{tn}"]], size=13, row_h=0.5, align=["l", "c", "c"])
    ftext(s, 7.4, 3.8, 5.4, 0.9, f"Counts reconstructed from the seed-0 precision, recall and the {n_test:,} test days.", size=11, color=K.FNAVY, italic=True)
    ftext(s, 0.55, 5.9, 12.2, 1.1, [
        (f"Pooled test: {n_test:,} days from Q2-Q4 2023, 7 tickers, {MIL['up_share_test'] * 100:.1f}% up-days; five seeds. Always-up scores {always_up * 100:.1f}%, so no model is better than guessing 'up'.", {"bullet": True, "space_after": 3}),
        (f"MIL attention vs price-only LSTM: {pc['lstm_mil_attention']['acc_diff'] * 100:+.1f} points, 95% interval [{pc['lstm_mil_attention']['ci95'][0] * 100:+.1f}, {pc['lstm_mil_attention']['ci95'][1] * 100:+.1f}]; shuffled-bag placebo {mm['placebo_mil_attention_shuffled_bags']['summary']['accuracy']['mean'] * 100:.1f}%.", {"bullet": True}),
    ], size=12)
    notes[18] = "The classification-style evaluation the objectives ask for. The product itself reports price-level errors; this study is offline and shows no gain from news."

    # ---- 19 parameter analysis
    s = S[18]
    set_title(s, "Parameter Analysis – Seeds, Scoring and MIL Pooling")
    wipe(s)
    K.picture(s, F.fig("f07_seed_variation.png"), 0.55, 1.75, 6.0, 3.0)
    K.picture(s, F.fig("f10_mil_paired_ci.png"), 6.8, 1.75, 6.0, 3.0)
    rows = [["Tier (5 seeds, 10 tickers)", "MAE shipped scoring", "MAE real closes", "Direction shipped / real"]]
    for k in ("lstm", "lstm_attention", "lstm_attention_sentiment", "hybrid_federated"):
        rows.append([F.TIER_LABEL[k], f"${mae_ship[k]:.1f}", f"${mae_true[k]:.1f}", f"{seedwise(k, 'directional_accuracy') * 100:.0f}% / {seedwise(k, 'da_true') * 100:.0f}%"])
    rows.append(["Persistence (no change)", "-", f"${pers:.1f}", "-"])
    ftable(s, 0.55, 4.85, [2.7, 1.2, 1.2, 1.9], rows, size=11, row_h=0.34, align=["l", "c", "c", "c"])
    ftext(s, 7.9, 4.85, 4.9, 2.1, [
        (f"Test targets are clipped to the training range: {n_above} of 10 tickers have over half of the test closes above the training maximum.", {"bullet": True, "space_after": 4}),
        ("Adding sentiment or FedAvg raises error and seed-to-seed spread; attention helps most.", {"bullet": True, "space_after": 4}),
        ("MIL pooling (mean, max, attention) differ by under 1 point in accuracy, inside the intervals.", {"bullet": True}),
    ], size=12)
    notes[19] = "Parameter analysis: how much the result depends on the seed, on how the target is scored, and on the MIL pooling choice."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{ep['health']['median_ms']:.1f} ms", "health, median", K.FBLUE), (f"{ep['metrics_one']['median_ms']:.1f} ms", "one ticker's metrics, median", K.GREEN),
                                   (f"{ep['history_120d']['median_ms']:.0f} ms", "history (live Yahoo download)", K.ORANGE), (f"{ep['sentiment']['median_ms'] / 1000:.1f} s", "sentiment (FinBERT on CPU)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.8, 2.9, 1.0, v, l, c, 24)
    K.picture(s, F.fig("f11_api_latency.png"), 0.55, 2.95, 6.1, 2.6)
    K.picture(s, F.fig("f13_forward_tiers.png"), 6.8, 2.95, 6.0, 2.6)
    rows = [["Check", "Result"],
            ["pytest (pipeline units + API endpoints)", f"{F.PYTEST_PASSED} of {F.PYTEST_TOTAL} passed; smoke test {F.SMOKE_PASSED} of {F.SMOKE_TOTAL} checks on the running stack"],
            ["Not tested", "concurrent users, load, security testing (no login exists), Docker build, accessibility"]]
    ftable(s, 0.55, 5.7, [4.7, 7.5], rows, size=11, row_h=0.42)
    notes[20] = "Latencies on 127.0.0.1, i5-12500, CPU only, machine shared with other jobs. Right: forward-test error of the re-trained tiers against persistence."

    # ---- 21 comparison with base paper
    s = S[20]
    set_title(s, "Comparison with Base Paper")
    wipe(s)
    rows = [["Aspect", "Base paper [1]", "QuantPulse (this project)"],
            ["Data", "Ten firms; hybrid on a one-month (30-day) window; 5-day trend over three months", "Ten tickers, two years (501 rows each); 97 test windows per ticker"],
            ["Split", "70/15/15 hybrid, 80/20 baselines", "80/20 chronological (shipped); three expanding folds in the MIL study"],
            ["Fit", f"Hybrid average R\u00b2 {bp['r2_avg']:.2f}; TSLA R\u00b2 {bp['tsla_r2']:.3f}", f"Shipped hybrid mean R\u00b2 {sh['hybrid_federated']['r2']['mean']:.1f}; re-run real-close R\u00b2 {seedwise('lstm_attention', 'r2_true'):.1f} (LSTM + attention)"],
            ["Direction", f"Baseline LSTM trend accuracy {bp['lstm_trend_acc']:.2f}%", f"Shipped {da_ship * 100:.0f}% (inflated); real closes {da_true * 100:.0f}%; stored 5-day trend {fs['trend_hits']}/10"],
            ["Sentiment", "FinBERT-tone on NewsAPI headlines, 30 days", "Synthetic template headlines (no key); real FNSPID headlines offline: no gain"]]
    ftable(s, 0.55, 1.8, [1.2, 3.5, 3.7], rows, size=11, row_h=0.8)
    K.picture(s, F.fig("f12_base_paper_comparison.png"), 9.15, 1.8, 3.7, 2.5)
    ftext(s, 9.15, 4.4, 3.7, 2.5, [
        ("Datasets and protocols differ: indicative only, not a head-to-head.", {"bullet": True, "space_after": 6}),
        ("The paper's R\u00b2 0.91 was not reproduced.", {"bullet": True, "space_after": 6}),
        ("No superiority is claimed.", {"bullet": True}),
    ], size=12)
    notes[21] = "State plainly that the paper uses a 30-day window and different units. We did not reproduce its R-squared."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Dashboard and Forecast Overview")
    wipe(s)
    K.picture(s, crop_shot("03_dashboard.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("16_forecast_top.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: dashboard with the ten tickers, last close, trend flag and 5-day outlook from the stored forecasts. Right: NVDA 5-day outlook with the notice that the models were trained up to 8 Sept 2026.", size=13, color=K.FNAVY)
    notes[22] = "Real captures of the running website with live Yahoo data and the shipped artefacts."
    s = S[22]
    set_title(s, "Output – Forecast Chart and Table")
    wipe(s)
    K.picture(s, crop_shot("04_forecast_5day.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("14_forecast_table.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "TSLA: the dashed forecast drops while the actual close (solid line) stays near its level; the table shows the forecast, band, actual close and forecast miss for each day (about -14% to -16%).", size=13, color=K.FNAVY)
    notes[23] = "The Forecast screen shows the miss openly: the stored forecast against the closes that followed."
    s = S[23]
    set_title(s, "Output – Model Comparison and News Sentiment")
    wipe(s)
    K.picture(s, crop_shot("06_models_comparison.png"), 0.55, 1.8, 6.0, 4.25, border=True)
    K.picture(s, crop_shot("08_sentiment.png"), 6.8, 1.8, 6.0, 4.25, border=True)
    ftext(s, 0.55, 6.15, 12.2, 0.85, "Left: NVDA model comparison, four deep tiers and five baselines. Right: TSLA news sentiment; the source badge reads 'Built-in fallback' because no NewsAPI key was set, so the headlines are synthetic.", size=13, color=K.FNAVY)
    notes[24] = "Model comparison and sentiment screens. Say clearly that the sentiment shown is synthetic."

    # ---- 25 conclusion and objectives
    s = S[24]
    set_title(s, "Conclusion & Objectives vs Outcomes")
    wipe(s)
    ftext(s, 0.55, 1.8, 12.2, 1.45, [
        (f"A working, tested forecasting product ({F.PYTEST_PASSED} tests, {F.SMOKE_PASSED} smoke checks) that follows the base paper; Random Forest is the most accurate model on {rf_wins} of 10 tickers.", {"bullet": True, "space_after": 4}),
        (f"Measured against real closes the deep tiers do not beat 'no change'; stored 5-day trend right for {fs['trend_hits']} of 10; real headlines gave no measurable gain in the MIL study.", {"bullet": True}),
    ], size=13)
    rows = [["#", "Objective", "Outcome"],
            ["1", "Collect prices, news and social media", "Partly: prices yes; headlines synthetic unless a key is set; no social-media source"],
            ["2", "Align into daily instances per day", "Partly: one mean sentiment per day in the product; per-headline instances in the offline study"],
            ["3", "Bags of news under MIL", "Not in the product; done offline (FNSPID, 7 tickers, gated attention)"],
            ["4", "Hybrid LSTM + attention model", "Met as built; accuracy poor against real closes"],
            ["5", "Direction with accuracy, precision, recall, F1", "Offline only: accuracy ~50%, F1 0.36-0.59; the app reports MAE, R\u00b2 and a direction score"],
            ["6", "Multi-source vs price-only baseline", f"Done offline: MIL {mm['lstm_mil_attention']['summary']['accuracy']['mean'] * 100:.1f}% vs price-only {mm['lstm_price']['summary']['accuracy']['mean'] * 100:.1f}%, within noise"]]
    ftable(s, 0.55, 3.0, [0.5, 3.7, 8.0], rows, size=11, row_h=0.52, align=["c", "l", "l"])
    notes[25] = "Be explicit about which objectives were met, partly met or not met."

    # ---- 26 limitations and future scope
    s = S[25]
    set_title(s, "Limitations & Future Scope")
    wipe(s)
    fbox(s, 0.55, 1.8, 6.0, 5.1, fill=K.FZEBRA)
    ftext(s, 0.75, 1.9, 5.6, 0.4, "Known limitations", size=17, color=K.FNAVY, bold=True)
    ftext(s, 0.75, 2.4, 5.6, 4.4, [
        ("No MIL, no social media and no precision/recall/F1 in the product", {"bullet": True, "space_after": 7}),
        ("Test targets clipped to the training range inflate direction scores", {"bullet": True, "space_after": 7}),
        ("Sentiment channel zero in training; synthetic headlines in test", {"bullet": True, "space_after": 7}),
        ("Stored forecasts missed the realised moves; artefacts older than the CSVs", {"bullet": True, "space_after": 7}),
        ("FedAvg is simulated and mixes differently scaled clients", {"bullet": True, "space_after": 7}),
        (f"MIL study: {n_test:,} test days, 7 tickers, 2022-2023; differences under about 3 points are noise", {"bullet": True, "space_after": 7}),
        ("CPU only; no login, database or load testing", {"bullet": True}),
    ], size=14)
    fbox(s, 6.8, 1.8, 6.0, 5.1, fill=K.FHEAD)
    ftext(s, 7.0, 1.9, 5.6, 0.4, "Future scope", size=17, color=K.FNAVY, bold=True)
    ftext(s, 7.0, 2.4, 5.6, 4.4, [
        ("Score against real closes and add a persistence baseline in the product", {"bullet": True, "space_after": 7}),
        ("Move MIL into the product with a direction head and P/R/F1", {"bullet": True, "space_after": 7}),
        ("Real news archive and a social-media source for multi-year sentiment", {"bullet": True, "space_after": 7}),
        ("Walk-forward validation over many periods and more tickers", {"bullet": True, "space_after": 7}),
        ("Fix the sentiment channel; retrain on a schedule", {"bullet": True, "space_after": 7}),
        ("A real federated deployment with per-client scalers", {"bullet": True, "space_after": 7}),
        ("Login, persistence and monitoring for deployment", {"bullet": True}),
    ], size=14)
    notes[26] = "Limitations first, then concrete next steps."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = sorted(R[k] for k in ("base", "zhang", "liao", "deng", "dietterich", "ilse", "fnspid", "finbert_tone", "fedavg", "lstm", "transformer", "fischer", "sezer", "tetlock"))
    short = [f"[{n}] {clean_ref(REFS[n - 1])}" for n in ids]
    short = [t if len(t) < 330 else t[:327] + "..." for t in short]
    ftext(s, 0.55, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[:7]], size=11)
    ftext(s, 6.8, 1.8, 6.0, 5.1, [(t, {"space_after": 5}) for t in short[7:]] + [(f"The full numbered list of {len(REFS)} references is in the report.", {"italic": True})], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh_ in S[27].shapes:
        if sh_.name == "Text 2":
            K.replace_text(sh_, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "QuantPulse")
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "16.Final_PPT.pptx")))
