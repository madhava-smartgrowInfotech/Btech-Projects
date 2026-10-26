"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from . import lit
from .refs import CITE as R
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "stock market prediction, multiple instance learning, LSTM, attention, FinBERT, federated averaging"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


# ---------------------------------------------------------------- slide text (batch documents + verified facts only)
OBJECTIVES = lit.OBJECTIVES
PROBLEM_STATEMENT = lit.PROBLEM_STATEMENT
LIT_ROWS = lit.LIT_TABLE
GAPS = lit.GAPS


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]
    bp = F.BASE_PAPER_RESULTS

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: the batch title asks for stock-direction prediction from several sources with Multiple Instance Learning; "
                "the base paper is a hybrid LSTM, attention, FinBERT sentiment and federated-averaging forecaster, which QuantPulse implements as a working product.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design details and measured results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: prices alone are not enough")
    wipe(s)
    W, H = 3.85, 5.0
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Share prices respond to company results, market mood and news at once. Price-only models ignore the news, labelling every article is impractical and noisy, and volatile, non-stationary prices make forecasts degrade quickly.", K.ORANGE, "!"),
        ("Our idea", "Treat each trading day as a bag of news instances (Multiple Instance Learning) and fuse it with an LSTM over prices to predict up or down, judged by accuracy, precision, recall and F1 against a price-only baseline.", K.BLUE, "i"),
        ("The deliverable", "QuantPulse: the base paper's LSTM, attention, FinBERT sentiment and simulated federated averaging as a ten-ticker FastAPI + React product, plus an offline MIL study on real headlines.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = ("Three points: why price-only models fall short, the multiple-instance idea from the project title, and what was actually built. "
                "Be clear that the product forecasts price level and a trend flag; the multiple-instance model was studied offline.")

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "An LSTM (64 units) over a 20-step OHLCV window, a multi-head attention layer, FinBERT-tone sentiment of news headlines as an extra input, and a simulated federated-averaging stage in which each of ten companies is a client. Baselines: ARIMA, SARIMAX, Prophet, random forest, SVR.", K.BLUE, "M"),
        ("Reported results", f"Average R2 {bp['r2_avg']} over ten companies (all above {bp['r2_min_company']}); baseline LSTM trend accuracy {bp['lstm_trend_acc']} % (best: {bp['trend_best']}; worst: {bp['trend_worst']}). Horizons: 5 and 30 days.", K.GREEN, "R"),
        ("Relevance and limits", "Same core pipeline as the product. The paper states its own limits: sentiment covers only 30 days of NewsAPI headlines, federated learning is simulated in one process, and accuracy varies by company. It averages a day's news into one score.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("Read the citation. The paper is the pipeline QuantPulse implements. Its figures are quoted from its text; its tables are images and are not quoted. "
                "Its limitations section is the starting point for our research gap: the news is one averaged number and covers one month.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    rows += LIT_ROWS
    K.table(s, 0.6, 1.25, [2.6, 0.65, 4.75, 4.13], rows, font_size=10, row_h=0.5, header_size=11, align=["l", "c", "l", "l"])
    notes[5] = ("The batch survey is the base paper plus four recent works, extended here with the multiple-instance papers the project title comes from. "
                "Results are quoted only where we read them in the paper or its abstract; for the survey items without an accessible abstract only the title is used.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence from the literature")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 6.1, 4.9, ["Airbnb", "Google", "Average (10)", "Tesla", "IBM"],
                     [("LSTM trend accuracy, %", [94.0, 86.0, bp["lstm_trend_acc"], 34.9, 18.0])], colors=[K.BLUE], number_format="0.0",
                     font_size=12, max_value=100)
    pt = gf.chart.plots[0].series[0].points[2]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 0.6, 6.35, 6.1, 0.6, f"Base paper [{R['base']}]: price-only LSTM trend accuracy per company (values printed in its text); the spread shows how hard direction is.", size=12, color=K.TXT, italic=True)
    K.text(s, 7.0, 1.45, 5.75, 5.0, [
        ("**What the literature shows**", {"space_after": 8}),
        (f"A day's news can be one bag with a single label, so no article labels are needed: F1 up 6.9 % and 9.2 % over the nMIL baseline (2015, 2016) on the Shanghai index [{R['zhang']}].", {"bullet": True, "space_after": 8}),
        (f"Attentive pooling of all of a day's news lowered MAE from 0.1913 to 0.1777 (7.11 %) [{R['liao']}].", {"bullet": True, "space_after": 8}),
        (f"The day-as-bag idea has been applied to S&P 500 trend forecasting from financial news [{R['deng']}].", {"bullet": True, "space_after": 8}),
        (f"Price-only direction is hard and uneven across stocks [{R['base']}].", {"bullet": True}),
    ], size=15, color=K.INK)
    notes[6] = ("These numbers are the published evidence that motivates the project. They come from different markets, data and metrics, so they are not compared with each other. "
                "Our own measurements start in Review 2.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Price history (OHLCV)", K.NAVY), ("ARIMA / ML / LSTM forecaster", K.NAVY), ("One daily sentiment score", K.NAVY), ("Price-level forecast", K.NAVY)], size=15, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Price-only and classical models",
           f"ARIMA, SARIMAX, Prophet, random forest and SVR forecast from past prices alone [{R['box']}], [{R['prophet']}], [{R['rf']}], [{R['svr']}].\nLSTM models capture temporal dependence [{R['lstm']}], [{R['fischer']}] but see no news.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Hybrid and multi-source models",
           f"The base paper adds attention, FinBERT-tone sentiment and simulated federated averaging [{R['base']}].\nMultiple-instance models treat news events as instances of a day-level bag [{R['zhang']}], [{R['deng']}].", K.ORANGE, "2", body_size=17)
    notes[7] = "Most forecasters use prices alone. Where news is added it is usually one averaged sentiment score per day; the multiple-instance papers treat the day's news as a bag."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Ignore the news", "Short news window", "Simulated federation", "Narrow multi-source work"]
    bodies = [
        "Price-only forecasters ignore the news and opinion that move prices; volatile, non-stationary data make them degrade quickly.",
        f"The base paper's sentiment experiments cover only 30 days because the news API returns a month of history [{R['base']}].",
        f"Federated learning in the base paper is a single-process simulation, and its hybrid results rest on a one-month window [{R['base']}].",
        f"Multiple-instance stock models use event extraction for one Chinese index [{R['zhang']}] or news alone [{R['deng']}].",
    ]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], bodies[i], C[i], str(i + 1), body_size=16)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, f"Several surveyed systems are known by title only; their results could not be verified [{R['s_hybrid']}]-[{R['s_text']}]. Explanations for a prediction are rarely given.", size=14, color=K.INK, anchor="m")
    notes[8] = "The first four come from the batch problem statement and the base paper's own limitations; the last line is an honest note about the survey items."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.3 + i * 1.12
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 1.0, g, size=15, color=K.INK, anchor="m")
    notes[9] = ("Gap 1 is the batch survey's own statement: the explicit multiple-instance treatment of a day's news. Gaps 2 and 3 come from comparing the papers; "
                "gap 4 is the project's stated objective; gap 5 is our reading: a naive no-change forecast is rarely reported next to R2.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 3.1, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.85, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=18, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Daily OHLCV prices, financial headlines and social-media text for the target stocks, grouped per trading day"),
                                 ("Output", "Up or down for the next day; in QuantPulse a 5- and 30-day price forecast with confidence band and a trend flag"),
                                 ("Measures", "Accuracy, precision, recall and F1 against a price-only baseline; MAE, RMSE and R2 for price level")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.8, 3.85, 1.95)
        K.text(s, x + 0.25, 4.9, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.35, 3.4, 1.35, b_, size=13, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a model that fuses price history with news and sentiment under a multiple-instance formulation."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 4.9, [(f"{i}.  {o}", {"space_after": 14}) for i, o in enumerate(OBJECTIVES, 1)], size=21, color=K.INK)
    K.text(s, 0.6, 6.35, 12.13, 0.7, "How far each objective was met is reported in the final review; the product covers prices and headlines, the multiple-instance model was studied offline.", size=14, color=K.TXT, italic=True)
    notes[11] = ("The six objectives from the proposal. The final review reports honestly which were met: prices and headlines yes, social-media text no, "
                 "multiple-instance learning and precision, recall and F1 in an offline study rather than in the application.")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Collect", "Yahoo Finance prices for ten tickers; NewsAPI headlines (or a labelled synthetic fallback) scored by FinBERT-tone; real FNSPID headlines for the offline MIL study.", K.BLUE, "1"),
        ("Model", "LSTM (64 units), 4-head attention, a sentiment channel; for the day-as-bag formulation, gated-attention pooling over headline instances.", K.GREEN, "2"),
        ("Federate", "Each ticker is a client; weights are averaged proportionally to sample count (FedAvg), simulated in one process as in the base paper.", K.ORANGE, "3"),
        ("Evaluate", "Nine models on ten tickers (MAE, RMSE, R2) against a no-change baseline; for the MIL study accuracy, precision, recall and F1 over five seeds.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.7, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.35, 12.13, 1.4)
    K.circle(s, 0.85, 5.65, 0.7, K.NAVY, "UI", 16)
    K.text(s, 1.8, 5.5, 10.7, 0.45, "FastAPI service + React website: dashboard, forecast, model comparison, news sentiment and method screens", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.95, 10.7, 0.75, "Results are stored as JSON files; there is no database and no login. A research and education tool, not financial advice.", size=14, color=K.TXT)
    notes[12] = "Four stages and the website that wraps them. The application does not contain the multiple-instance model; that is studied in a separate offline script."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    if os.path.exists(F.diag("01_system_architecture.png")):
        K.picture(s, F.diag("01_system_architecture.png"), 0.5, 1.2, 8.4, 5.9)
    else:
        print("WARNING: diagram 01 missing")
    K.text(s, 9.1, 1.5, 3.7, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        (f"Presentation: React and Vite website (port {F.WEB_PORT}).", {"bullet": True}),
        (f"Application: FastAPI service (port {F.API_PORT}), seven JSON endpoints, input checks.", {"bullet": True}),
        ("Pipeline: data loader, FinBERT sentiment, baselines, LSTM and attention models, FedAvg, forecast.", {"bullet": True}),
        ("State: JSON, CSV and Keras files; no database, no authentication.", {"bullet": True}),
    ], size=14, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the website calls the API, which serves stored results and live prices; training writes the artefacts. There is no database."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working product", "Model comparison", "Honest evaluation", "MIL study", "Stated limits"]
    bodies = [
        "A website that shows ten tickers, 5- and 30-day forecasts with bands, model comparison and news sentiment, backed by a tested API.",
        "Five classical baselines and four deep tiers compared per ticker, with a no-change baseline next to them.",
        "Scoring against the real closes and a forward test of the stored forecasts on data that followed.",
        "A bag-of-headlines study on real news with accuracy, precision, recall and F1 against a price-only baseline.",
        "Synthetic fallback headlines, no social-media source, simulated federation: each limit documented, not hidden.",
    ]
    for i, (h, b_) in enumerate(zip(heads, bodies)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.65, h, b_, C[i % 4], str(i + 1), body_size=16)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.35, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=17)
    notes[14] = ("These are the outcomes we committed to. The measured results, including where they are weaker than the base paper reports, are presented in Review 2 and the final review.")

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [R[k] for k in ("base", "s_hybrid", "s_llm", "s_multimodal", "s_text", "zhang", "liao", "deng", "ilse", "fnspid")]
    K.text(s, 0.6, 1.25, 12.13, 5.8, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 3}) for n in idx], size=12, color=K.INK)
    notes[15] = "Selected references; the complete numbered list, with notes on how each entry was checked, is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why QuantPulse has no multiple-instance model (it implements the base paper; MIL was studied offline on real headlines); "
                 "why the sentiment is synthetic (no NewsAPI key); how the work differs from the base paper (two years of data, real-close scoring, a no-change baseline).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
