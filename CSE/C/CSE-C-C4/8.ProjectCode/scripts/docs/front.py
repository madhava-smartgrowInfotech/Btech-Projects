"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p(f"{F.PRODUCT}: {F.SUBTITLE}", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n.upper(), r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE.upper(), align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n.upper()}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n.upper()}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper and of the multiple-instance-learning papers that this work builds on, the creators of the FNSPID news dataset used in the offline study, "
        "and the maintainers of the open-source software used in this project (TensorFlow, PyTorch, scikit-learn, statsmodels, Prophet, Hugging Face Transformers, yfinance, FastAPI, React, Vite, Recharts, pandas and Playwright).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n.upper(), align="right", space_after=2)


def abstract(d):
    from .ch1 import K

    k = K()
    d.h1("ABSTRACT")
    d.p("Short-term stock forecasts are wanted by investors and analysts, but price history alone ignores the news that moves markets, and labelling every article is impractical. "
        f"This project builds {F.PRODUCT}, a working forecasting product that implements the pipeline of the base paper (Nejatbakhsh and Aliasgari, IEEE Access, 2026): an LSTM over twenty-day price windows, "
        "a four-head attention layer, a FinBERT-tone sentiment channel and a simulated federated-averaging stage, compared with ARIMA, SARIMAX, Prophet, random forest and SVR on ten tickers. "
        "A FastAPI service with seven endpoints and a six-screen React website present 5- and 30-day forecasts with confidence bands, model comparison and sentiment.")
    d.p(f"Measured honestly, the shipped models do not beat repeating yesterday's close: over five seeds the best deep tier has a mean absolute error of ${k['attn_true_mae'][0]:.2f} per share against ${k['persist']:.2f} for persistence, "
        f"and its directional accuracy on real closes is {k['attn_true_da'] * 100:.0f} %, not the {k['attn_ship_da'] * 100:.0f} % the application displays, because the application clips test targets to the training range. "
        f"A forward test of the stored 5-day forecasts against the closes that followed found a mean forecast change of {k['f5_pred']:+.1f} % against a realised {k['f5_act']:+.1f} % and a correct trend flag for {k['f5_hits']} of 10 tickers. "
        "The base paper's R2 of 0.91 was not reproduced; the datasets and protocols differ. The sentiment channel of the product rests on synthetic template headlines because no news key was set.")
    d.p(f"Multiple Instance Learning, the formulation named in the title, is not implemented in the product. It was studied offline on real FNSPID headlines, treating each ticker-day as a bag with a bag-level next-day direction label: "
        f"over {k['mil_n']:,} test days on {k['mil_tickers']} tickers, mean, max and gated-attention pooling reached {k['acc_lo'] * 100:.1f}-{k['acc_hi'] * 100:.1f} % accuracy, like the price-only LSTM, while always predicting up scores {k['always_up'] * 100:.1f} %. "
        f"No evidence of a gain from news was found at this sample size. {F.PYTEST_PASSED} automated tests and {F.SMOKE_PASSED} smoke checks pass.")
    d.p("**Keywords:** stock market prediction, multiple instance learning, LSTM, attention, FinBERT sentiment, federated averaging, financial news, forecast evaluation.")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("ARIMA", "Autoregressive Integrated Moving Average"), ("AUC", "Area Under the ROC Curve"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("DFD", "Data Flow Diagram"),
    ("EMH", "Efficient-Market Hypothesis"), ("ER", "Entity-Relationship"), ("FedAvg", "Federated Averaging"), ("FL", "Federated Learning"),
    ("FNSPID", "Financial News and Stock Price Integration Dataset"), ("GAP", "Global Average Pooling"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"),
    ("LSTM", "Long Short-Term Memory"), ("MAE", "Mean Absolute Error"), ("MAPE", "Mean Absolute Percentage Error"), ("MIL", "Multiple Instance Learning"),
    ("ML", "Machine Learning"), ("MSE", "Mean Squared Error"), ("OHLCV", "Open, High, Low, Close, Volume"), ("PCA", "Principal Component Analysis"),
    ("R2", "Coefficient of determination"), ("RF", "Random Forest"), ("RMSE", "Root Mean Squared Error"), ("SARIMAX", "Seasonal ARIMA with eXogenous variables"),
    ("SPA", "Single-Page Application"), ("SVR", "Support Vector Regression"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
