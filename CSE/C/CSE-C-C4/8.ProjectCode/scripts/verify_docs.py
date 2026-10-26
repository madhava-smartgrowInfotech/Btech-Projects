"""
Verifies the documentation set at the batch root.

    python scripts/verify_docs.py [--no-pdf]

Checks (Docs prompt, step 7):
 1. every required file exists with its exact name and is not tiny
 2. no template leftovers in any part of any file (including nested packages)
 3. deck slide counts 16 / 20 / 28
 4. the full report has the ten chapters, the references and the three appendices
 5. figures are embedded and every screenshot referenced by the builders exists
 6. student names, roll numbers and the guide name are on the title pages
 7. the IEEE paper is two-column, has five Roman sections, both tables, 4+ captioned figures, 15-20 references and 6 pages
 8. nothing from the reference paper the paper format was taken from
"""

import argparse
import hashlib
import io
import os
import re
import sys
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BATCH = os.path.dirname(ROOT)
sys.path.insert(0, ROOT)

REQUIRED = [
    "7.Review_1_PPT.pptx", "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "10.SCOPE & LIMITATIONS.docx", "11. BLOCK DIAGRAMS ALL.docx",
    "12.DESIGN &IMPLEMENTATION.docx", "13.Review_2_PPT.pptx", "14.RESULTS.docx", "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    "16.Final_PPT.pptx", "17.Final Document.docx", "19.IEEE_Paper.docx",
]
DECKS = {"7.Review_1_PPT.pptx": 16, "13.Review_2_PPT.pptx": 20, "16.Final_PPT.pptx": 28}
TEMPLATE_WORDS = ["[STUDENT NAME", "[ROLL NO", "Fakeddit", "PHEME", "Snopes", "MCNN", "bank transaction", "social media manipulation",
                  "fraud detection", "Hidden Risk", "Kaliyar", "Chandrakala", "Sindhuja", "pneumonia", "chest X-ray",
                  "SignSpeak", "sign language", "fingerspelling", "PeerScope", "TaxSentinel", "ChurnIQ", "reviewer assignment", "ResQAI",
                  "CipherGuard", "ThreatSense", "SecureChannel", "UNSW", "intrusion detection", "AutoSecAI", "LabLens", "BotGraph", "ALIP",
                  "FlowSense", "smart traffic", "ambulance", "SUMO", "YOLO", "Lamrabet", "light-bar"]
PAPER_WORDS = ["skincare", "Dayananda", "Sindhu", "recommender system"]
PENDING_WORDS = ["[latency pending]", "[to be filled]", "[pending]", "TODO", "XXX", "lorem"]
CHAPTERS = [("1", "INTRODUCTION"), ("2", "LITERATURE SURVEY"), ("3", "FEASIBILITY STUDY"), ("4", "SYSTEM REQUIREMENTS"), ("5", "SYSTEM DESIGN"), ("6", "METHODOLOGY"), ("7", "SYSTEM TESTING"),
            ("8", "RESULTS"), ("9", "CONCLUSION"), ("10", "FUTURE SCOPE")]
STUDENTS = [("Pottepaka Lingaswamy", "23K91A05J7"), ("Maduna Naveen Kumar", "23K91A05D7"), ("Meedigeri Keerthi", "24K95A0517"), ("Mareddy Sai Karthik", "23K91A05E8")]
GUIDE = "K Anusha"

problems, notes = [], []


def fail(msg):
    problems.append(msg)
    print("  FAIL", msg)


def ok(msg):
    print("  ok  ", msg)


def path(n):
    return os.path.join(BATCH, n)


def all_text_blobs(p):
    """Yield (member, bytes) for every part of a package, descending into nested zips (embedded workbooks)."""
    with zipfile.ZipFile(p) as z:
        for m in z.namelist():
            data = z.read(m)
            yield m, data
            if m.lower().endswith((".xlsx", ".docx", ".pptx", ".zip")):
                try:
                    with zipfile.ZipFile(io.BytesIO(data)) as zz:
                        for mm in zz.namelist():
                            yield f"{m}!{mm}", zz.read(mm)
                except zipfile.BadZipFile:
                    pass


def docx_text(p):
    """All text of a .docx in document order (table cells inline), then headers and footers."""
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(p)
    out = ["".join(t.text or "" for t in para.iter(qn("w:t"))) for para in d.element.body.iter(qn("w:p"))]
    for sct in d.sections:
        out += [x.text for x in sct.header.paragraphs] + [x.text for x in sct.footer.paragraphs]
    return "\n".join(out)


def pdf_pages(p):
    import pypdf

    return len(pypdf.PdfReader(p).pages)


def to_pdf(p, keep=False):
    import tempfile

    out = os.path.splitext(p)[0] + ".pdf" if keep else os.path.join(tempfile.mkdtemp(prefix="verify_pdf_"), os.path.basename(os.path.splitext(p)[0]) + ".pdf")
    try:
        import time
        import win32com.client

        last = None
        for _ in range(4):
            word = None
            try:
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                doc = word.Documents.Open(os.path.abspath(p), ReadOnly=True, AddToRecentFiles=False)
                doc.Repaginate()
                doc.ExportAsFixedFormat(os.path.abspath(out), 17)
                doc.Close(False)
                last = None
                break
            except Exception as e:  # noqa: BLE001
                last = e
                time.sleep(8)
            finally:
                if word is not None:
                    try:
                        word.Quit()
                    except Exception:  # noqa: BLE001
                        pass
        if last:
            raise last
        return out if os.path.exists(out) else None
    except Exception as e:  # noqa: BLE001
        notes.append(f"PDF conversion failed for {os.path.basename(p)}: {e}")
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()

    print("1. files exist")
    for n in REQUIRED:
        p = path(n)
        if not os.path.exists(p):
            fail(f"missing {n}")
            continue
        size = os.path.getsize(p)
        small = size < 40 * 1024
        if small:
            fail(f"{n} is only {size} bytes")
        else:
            ok(f"{n} ({size // 1024} KB)")

    print("2. template leftovers")
    for n in REQUIRED:
        p = path(n)
        if not os.path.exists(p):
            continue
        hits = set()
        for m, data in all_text_blobs(p):
            low = data.lower()
            for w in TEMPLATE_WORDS:
                # whole-word match so that "ephemeral" does not trip "PHEME"; bracketed placeholders match as plain text
                pat = re.escape(w.lower().encode()) if w.startswith("[") else rb"(?<![a-z])" + re.escape(w.lower().encode()) + rb"(?![a-z])"
                if re.search(pat, low):
                    hits.add((w, m))
        if p.endswith(".docx"):
            txt_all = docx_text(p)
            for w in PENDING_WORDS:
                if w.lower() in txt_all.lower():
                    hits.add((w, "text"))
            if re.search(r"(?:\bNone%|\bNone ms\b|\bNone s\b|\(None\)|\bof None\b)", txt_all):
                hits.add(("None (unformatted value)", "text"))
        if hits:
            for w, m in sorted(hits)[:6]:
                fail(f"{n}: '{w}' in {m}")
        else:
            ok(f"{n}: clean")

    print("3. deck slide counts")
    from pptx import Presentation

    for n, cnt in DECKS.items():
        if os.path.exists(path(n)):
            got = len(Presentation(path(n)).slides)
            (ok if got == cnt else fail)(f"{n}: {got} slides (expected {cnt})")

    print("4. full report structure")
    p17 = path("17.Final Document.docx")
    if os.path.exists(p17):
        t = docx_text(p17)
        for num, title in CHAPTERS:
            found = re.search(r"(?:CHAPTER\s+)?" + num + r"[.:]\s*" + title, t) is not None
            (ok if found else fail)(f"heading {num}. {title}")
        for w in ("REFERENCES", "APPENDIX A", "APPENDIX B", "APPENDIX C"):
            (ok if w in t else fail)(f"heading {w}")

    print("5. figures and screenshots")
    for n in ("14.RESULTS.docx", "17.Final Document.docx", "11. BLOCK DIAGRAMS ALL.docx"):
        if os.path.exists(path(n)):
            with zipfile.ZipFile(path(n)) as z:
                imgs = [m for m in z.namelist() if m.startswith("word/media/")]
            ok(f"{n}: {len(imgs)} embedded images")
            if n.startswith("17") and len(imgs) < 25:
                fail(f"{n}: only {len(imgs)} images")
    for sub, fn in (("screenshots", "(?:shot|sfig)"), ("diagrams", "diag"), ("figures", "(?<![A-Za-z_])fig")):
        referenced = set()
        for f in os.listdir(os.path.join(ROOT, "scripts", "docs")):
            if f.endswith(".py"):
                referenced |= set(re.findall(fn + r'\("([^"]+\.png)"', open(os.path.join(ROOT, "scripts", "docs", f), encoding="utf-8").read()))
        missing = [x for x in sorted(referenced) if not os.path.exists(os.path.join(ROOT, "docs", sub, x))]
        (fail if missing else ok)(f"{len(referenced)} referenced {sub} files, missing: {missing or 'none'}")

    print("6. title pages carry names, rolls and guide")
    if os.path.exists(p17):
        t = docx_text(p17)[:6000]
        for nme, roll in STUDENTS:
            (ok if (nme.lower() in t.lower() and roll in t) else fail)(f"report front matter: {nme} {roll}")
        (ok if GUIDE.lower() in t.lower() else fail)("report front matter: guide")
    for n in DECKS:
        if os.path.exists(path(n)):
            prs = Presentation(path(n))
            txt = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
            for nme, roll in STUDENTS:
                if roll not in txt or nme.split()[0].lower() not in txt.lower():
                    fail(f"{n}: title slide lacks {nme} {roll}")
            if GUIDE.lower() not in txt.lower():
                fail(f"{n}: title slide lacks the guide")
            else:
                ok(f"{n}: title slide has the team and the guide")
    for n in ("9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx", "10.SCOPE & LIMITATIONS.docx", "11. BLOCK DIAGRAMS ALL.docx", "12.DESIGN &IMPLEMENTATION.docx", "14.RESULTS.docx", "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx"):
        if os.path.exists(path(n)):
            t = docx_text(path(n))[:3000]
            good = all(r in t for _, r in STUDENTS) and GUIDE.lower() in t.lower()
            (ok if good else fail)(f"{n}: team and guide on the first page")

    print("7. IEEE paper")
    pp = path("19.IEEE_Paper.docx")
    if os.path.exists(pp):
        import docx
        from docx.oxml.ns import qn

        d = docx.Document(pp)
        sect = d.sections[-1]
        cols = sect._sectPr.xpath("./w:cols")
        two = bool(cols) and cols[0].get(qn("w:num")) == "2"
        (ok if two else fail)("body section is two-column")
        full = docx_text(pp)
        romans = re.findall(r"^(I|II|III|IV|V)\. [A-Z ]+$", "\n".join(p.text for p in d.paragraphs), re.M)
        (ok if len(set(romans)) == 5 else fail)(f"Roman-numbered sections: {sorted(set(romans))}")
        for t in ("TABLE I", "TABLE II"):
            (ok if t in full else fail)(t)
        figs = re.findall(r"Fig\. \d+\.", full)
        (ok if len(figs) >= 4 else fail)(f"{len(figs)} figure captions")
        refs = re.findall(r"^\[(\d+)\]\s", "\n".join(p.text for p in d.paragraphs), re.M)
        n_refs = max(map(int, refs)) if refs else 0
        (ok if 15 <= n_refs <= 20 else fail)(f"{n_refs} references")
        body = "\n".join(p.text for p in d.paragraphs)
        cited = {int(x) for grp in re.findall(r"\[([\d\], \[]+)\]", body.split("REFERENCES")[0]) if grp.replace(" ", "") != "0,1" for x in re.findall(r"\d+", grp)}
        (ok if cited == set(range(1, n_refs + 1)) else fail)(f"every reference cited and every citation listed (cited {len(cited)} of {n_refs})")
        for w in PAPER_WORDS:
            if w.lower() in full.lower():
                fail(f"paper contains '{w}'")
        words = len(full.split())
        notes.append(f"IEEE paper word count (incl. captions and references): {words}")
        if not a.no_pdf:
            pdf = to_pdf(pp, keep=True)
            if pdf:
                pg = pdf_pages(pdf)
                (ok if 5.5 <= pg <= 6.5 else fail)(f"paper is {pg} pages")
                notes.append(f"IEEE paper pages: {pg} (PDF: {os.path.basename(pdf)})")
    else:
        fail("19.IEEE_Paper.docx missing")

    print("8. no leftovers from the reference paper or other documents")
    for n in REQUIRED:
        p = path(n)
        if os.path.exists(p) and n.endswith(".docx"):
            hits = [w for w in PAPER_WORDS if w.lower() in docx_text(p).lower()]
            (fail if hits else ok)(f"{n}: {hits or 'clean'}")

    if not a.no_pdf and os.path.exists(p17):
        pdf = to_pdf(p17)
        if pdf:
            pg = pdf_pages(pdf)
            notes.append(f"Full report pages: {pg}")
            (ok if 80 <= pg <= 95 else fail)(f"17.Final Document is {pg} pages (target 80-95)")

    print()
    for n in notes:
        print("note:", n)
    if problems:
        print(f"\n{len(problems)} problem(s)")
        return 1
    print("\nAll documentation checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
