"""16.Final_PPT.pptx: 28 slides. Built from D:\\abstracts\\Templates\\Final_PPT.pptx (banner, rule, title and footer chrome inherited)."""

import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from . import facts as F
from . import pptkit as K
from .deck_review1 import scan_deck
from .lit import LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Final_PPT.pptx"
FONT = "Calibri"
TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT = 73152, 979985, 1098857, 6446520


def chrome(sh):
    return sh.top in (TOP_BANNER, TOP_RULE, TOP_TITLE, TOP_FOOT)


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.top == TOP_TITLE:
            K.replace_text(sh, text, size)
            return


def wipe(slide):
    K.clear(slide, chrome)


def fbox(slide, x, y, w, h, fill=K.FHEAD, line=None):
    return K.box(slide, x, y, w, h, fill=fill, line=line, radius=0.04)


def ftable(slide, x, y, widths, rows, size=13, row_h=0.45, align=None, first_bold=True):
    return K.table(slide, x, y, widths, rows, font_size=size, row_h=row_h, header_fill=K.FHEAD, header_color=K.FNAVY, zebra=K.FZEBRA,
                   body_color=K.DARK, first_bold=first_bold, font=FONT, align=align, border=K.RGBColor(0xBF, 0xBF, 0xBF))


def ftext(slide, x, y, w, h, paras, size=16, color=K.DARK, **kw):
    return K.text(slide, x, y, w, h, paras, size=size, color=color, font=FONT, **kw)


def ftile(slide, x, y, w, h, value, label, color=K.FBLUE, vsize=26):
    K.box(slide, x, y, w, h, fill=K.FHEAD, line=None, radius=0.05)
    K.box(slide, x, y, 0.08, h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    K.text(slide, x + 0.2, y + 0.06, w - 0.3, h * 0.55, value, size=vsize, color=K.FNAVY, bold=True, anchor="m", font=FONT)
    K.text(slide, x + 0.2, y + h * 0.58, w - 0.3, h * 0.4, label, size=12, color=K.DARK, font=FONT)


def build(out_path):
    prs = Presentation(TEMPLATE)
    S = list(prs.slides)
    assert len(S) == 28, len(S)
    cp = prs.core_properties
    cp.title = f"{F.TITLE} - Final Review"
    cp.author = f"Batch {F.BATCH_NO}, Dept. of CSE, TKRCET"
    cp.comments = ""
    cp.subject = F.PRODUCT
    cp.keywords = "campus placement, XGBoost, kNN, recommendation"
    cp.last_modified_by = ""
    notes = {}
    X, DS, PERF, REC, CONF = F.R, F.DS, F.PERF, F.REC, F.CONF

    for sl in S:
        for sh in sl.shapes:
            if sh.top == TOP_FOOT and sh.has_text_frame and "Final Review" in sh.text_frame.text:
                K.replace_text(sh, f"Dept. of CSE  |  Batch {F.BATCH_NO}  |  Final Review")

    # ---- 1 title
    s = S[0]
    for sh in s.shapes:
        if sh.name == "Text 1":
            K.replace_text(sh, F.TITLE, 28)
        elif sh.name == "Text 2":
            K.replace_text(sh, f"Final Review  ·  {F.PRODUCT}")
        elif sh.name == "Text 3":
            K.replace_text(sh, f"Team Details : Batch {F.BATCH_NO}")
        elif sh.name == "Text 4":
            K.replace_text(sh, "\n".join(f"{r_}   {n_}" for n_, r_ in F.STUDENTS))
        elif sh.name == "Text 5":
            K.replace_text(sh, f"Under the Esteemed Guidance of\n{F.GUIDE}\n{F.GUIDE_DESIGNATION}, Department of CSE")
        elif sh.name == "Text 6":
            K.replace_text(sh, f"DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING  ·  {F.COLLEGE_SHORT}  ·  Academic Year {F.YEAR}")
    notes[1] = "Introduce the team and the project: a placement-prediction and company-recommendation system with officer and recruiter tools."

    # ---- 2 table of contents
    s = S[1]
    set_title(s, "Table of Contents")
    for sh in s.shapes:
        if sh.name == "Text 4":
            K.replace_text(sh, "Abstract\nBase Paper & Literature Survey\nResearch Gap, Problem Statement & Objectives\nProposed System & Architecture\nMethodology, Dataset & Model\nDecision Workflow, Interaction & Storage\nImplementation Highlights\nResults, Confusion Matrices & Threshold Analysis\nPerformance & Testing\nComparison with Base Papers\nOutput Screens\nConclusion, Future Scope & References", 17)
    notes[2] = "The agenda follows the report: problem, design, method, results, comparison, screens, conclusion."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract")
    wipe(s)
    ftext(s, 0.55, 1.9, 12.2, 4.9, [
        ("Placement cells keep student records but do no predictive analysis: students learn their readiness late, shortlisting is manual and nobody says which skill or company to target.", {"bullet": True, "space_after": 10}),
        ("The Smart Campus Placement System estimates the probability that a student is placed at a chosen company and role (24 features, calibrated XGBoost, kNN refinement of borderline cases), gives skill-gap tips, recommends company-role pairs, and offers an officer dashboard and a recruiter shortlist.", {"bullet": True, "space_after": 10}),
        (f"On a **synthetic** 3,000-student dataset (450 test students) XGBoost + Platt reaches **{F.pct(X['xgb_calibrated']['accuracy'], 2)}** accuracy (AUC {X['xgb_calibrated']['roc_auc']:.4f}); the hybrid reaches {F.pct(X['hybrid_xgb_knn']['accuracy'], 1)} only when the true labels choose the refined rows, and {F.pct(X['hybrid_knn_labelfree']['accuracy'], 2)} without them.", {"bullet": True, "space_after": 10}),
        (f"A FastAPI service and React interface wrap the original code; {F.TEST_SUMMARY['passed']} scripted checks pass; a prediction takes about {PERF['POST /api/predict']['mean_ms']:.0f} ms, a recommendation about {PERF['POST /api/recommend']['mean_ms'] / 1000:.1f} s.", {"bullet": True, "space_after": 10}),
        ("Not done: real placement data, login, database storage and real-outcome validation of recommendations.", {"bullet": True}),
    ], size=17)
    notes[3] = "Read the first two bullets, then the headline numbers with their protocol and the honest caveat about the hybrid."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base Papers")
    wipe(s)
    ftext(s, 0.55, 1.85, 12.2, 0.95, REFS[0], size=12, color=K.FNAVY)
    rows = [["Aspect", "Declared base paper [1] (IJETMS 2026)", "Technique source [2] (arXiv:2603.29881)"],
            ["Task / data", "Placement prediction; synthetic records, 80:20 split", "Graduate admission prediction; 13,000 real GradCafe records, 70/15/15"],
            ["Method", "Logistic Regression, Random Forest, Gradient Boosting; Streamlit app, batch CSV, PDF report, e-mail", "XGBoost + Platt; kNN / SVM refinement of residual cases; university-program recommendation; SHAP"],
            ["Reported results", "None numeric (only the split and model names)", "XGBoost 0.7241; XGBoost + kNN hybrid 0.8723; Brier 0.1864; recommendation +65.1% mean acceptance probability"],
            ["Stated limits", "No company recommendation; no numbers", "Residual set uses true labels; US-centric self-reported data"]]
    ftable(s, 0.55, 2.9, [1.9, 4.9, 5.4], rows, size=13, row_h=0.78)
    notes[4] = "The batch's declared base paper reports no numbers, so numeric context comes from the admissions paper whose hybrid method the code adapts. Both are cited."

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature Survey")
    wipe(s)
    rows = [["Author / Year", "Methodology", "Dataset", "Key Results", "Limitations"]]
    for i in (0, 1, 2, 3, 4, 7):
        if i < len(LIT_TABLE):
            r = LIT_TABLE[i]
            rows.append([x.replace(" (declared base paper)", "").replace(" (technique base)", "") for x in r[:5]])
    ftable(s, 0.55, 1.85, [1.9, 2.9, 2.0, 2.9, 2.5], rows, size=9.5, row_h=0.62)
    notes[5] = "Representative work. Where a work's figures were not checked they are not quoted. Four further titles from the project survey could not be verified online and are listed in the report as unverified."

    # ---- 6 research gap
    s = S[5]
    set_title(s, "Research Gap")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.8, [
        ("**Prediction only:** placement predictors rarely recommend which company and role a student should target.", {"bullet": True, "space_after": 12}),
        ("**Actionable feedback:** few give the student the specific gaps against a company's requirements.", {"bullet": True, "space_after": 12}),
        ("**Calibrated probabilities:** a recommendation ranking needs reliable probabilities; calibration is seldom reported.", {"bullet": True, "space_after": 12}),
        ("**Honest evaluation:** results often lack several seeds, a confusion matrix, or a check that a refinement rule works without labels.", {"bullet": True, "space_after": 12}),
        ("**Real data:** placement records are private, so published work, including the base paper, uses synthetic data.", {"bullet": True}),
    ], size=19)
    notes[6] = "Five gaps. We address the first three directly, measure the fourth, and state the fifth as a limitation."

    # ---- 7 problem statement
    s = S[6]
    set_title(s, "Problem Statement")
    wipe(s)
    fbox(s, 0.55, 1.95, 12.2, 2.0)
    ftext(s, 0.85, 2.0, 11.6, 1.9, PROBLEM_STATEMENT, size=17, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Academics, aptitude and technical scores, projects, internships, backlogs, target company and role"), ("Output", "Probability, label, skill-gap tips and top-5 company-role pairs"), ("Users", "Students, placement officers, recruiters")]):
        x = 0.55 + i * 4.1
        fbox(s, x, 4.3, 3.9, 1.7, fill=K.FZEBRA)
        ftext(s, x + 0.2, 4.4, 3.5, 0.4, t, size=17, color=K.FNAVY, bold=True)
        ftext(s, x + 0.2, 4.9, 3.5, 1.0, b_, size=14)
    notes[7] = "The statement has three parts: estimate, explain and recommend, and serve three user groups."

    # ---- 8 objectives
    s = S[7]
    set_title(s, "Objectives")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=20)
    notes[8] = "The six objectives of the proposal. The conclusion reports for each whether it was met; objective 1 (collect real records) was met only with synthetic data."

    # ---- 9 proposed system
    s = S[8]
    set_title(s, "Proposed System")
    wipe(s)
    K.flow(s, 0.55, 2.0, 12.2, 1.1, [("Profile + company + role", K.FNAVY), ("24 features", K.FBLUE), ("XGBoost + Platt", K.GREEN), ("kNN on borderline", K.ORANGE), ("Probability, tips, top-5", K.PURPLE)], size=14, gap=0.35, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("Features include the gaps between the student's CGPA, aptitude and technical scores and the target company's cutoffs.", {"bullet": True, "space_after": 10}),
        ("Calibrated XGBoost; when the probability is between 0.4 and 0.6 it is blended with a kNN estimate; ready if p >= 0.5.", {"bullet": True, "space_after": 10}),
        ("Recommendation filters feasible companies (CGPA proximity, aptitude, backlogs, tier) and ranks company-role pairs.", {"bullet": True, "space_after": 10}),
        ("FastAPI service over the unchanged Python core; React interface and the original Flask pages.", {"bullet": True}),
    ], size=18)
    notes[9] = "The proposed system in one line of boxes, then four properties that matter."

    # ---- 10 architecture
    s = S[9]
    set_title(s, "System Architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.55, 1.85, 8.3, 5.0)
    ftext(s, 9.1, 2.0, 3.7, 4.7, [
        ("Presentation: React app (5405) and Flask pages (5000)", {"bullet": True, "space_after": 8}),
        ("Application: FastAPI (8405), 9 endpoints", {"bullet": True, "space_after": 8}),
        ("Core: features, training, prediction, recommendation", {"bullet": True, "space_after": 8}),
        ("Storage: CSV data and pickled models, no SQL database", {"bullet": True}),
    ], size=16)
    notes[10] = "Interfaces on the left, one shared core in the middle, files on the right."

    # ---- 11 methodology
    s = S[10]
    set_title(s, "Methodology")
    wipe(s)
    K.flow(s, 0.55, 1.95, 12.2, 1.2, [("Synthetic generator", K.FNAVY), ("24 features", K.FNAVY), ("70/15/15 split", K.GREEN), ("6 baselines", K.PURPLE), ("Platt calibration", K.ORANGE), ("Residual hybrid", K.FBLUE), ("metrics + plots", K.FNAVY)], size=12, gap=0.2, font=FONT)
    ftext(s, 0.55, 3.5, 12.2, 3.3, [
        ("One feature function for training, the API and the recommender: no train/serve mismatch.", {"bullet": True, "space_after": 10}),
        ("Stratified 70 / 15 / 15 split with seed 42; maps and scaler fitted on the training part only; 9-fold CV for XGBoost.", {"bullet": True, "space_after": 10}),
        ("Extra measurements: confusion matrices, label-free hybrid, borderline-band and threshold sweeps, five seeds, SHAP, recommender gains, latency.", {"bullet": True, "space_after": 10}),
        ("The shipped model files reproduce models/metrics.json exactly; every extra number comes from scripts/run_experiments.py.", {"bullet": True}),
    ], size=18)
    notes[11] = "Workflow and measurements. The shipped pickles reproduce the stored metrics; a fresh retrain drifts by at most 0.011."

    # ---- 12 dataset, features, model
    s = S[11]
    set_title(s, "Dataset, Features & Model")
    wipe(s)
    ftile(s, 0.55, 1.9, 2.9, 1.1, "3,000", "synthetic students, 19 columns", K.FBLUE, 26)
    ftile(s, 0.55, 3.15, 2.9, 1.1, "30", "invented companies (7 / 15 / 8 by tier)", K.GREEN, 26)
    ftile(s, 0.55, 4.4, 2.9, 1.1, "24", "features: 5 freq-encoded, 19 z-scored", K.ORANGE, 26)
    bt = DS["by_tier"]
    K.bar_chart(s, 3.7, 1.85, 4.8, 3.2, ["Tier 1", "Tier 2", "Tier 3"], [("Placed rate by target tier", [bt["1"]["mean"], bt["2"]["mean"], bt["3"]["mean"]])],
                colors=[K.FBLUE], number_format="0.0%", font_size=11, max_value=1.0)
    rows = [["Component", "Setting"], ["XGBoost", "100 trees, depth 8, lr 0.1"], ["Calibration", "Platt, cv = 3"], ["Local model", "kNN k = 9"], ["Split", "2,100 / 450 / 450"]]
    ftable(s, 8.7, 1.9, [1.6, 2.5], rows, size=11, row_h=0.5)
    ftext(s, 8.7, 4.6, 4.1, 2.2, [
        (f"Placed {F.pct(DS['placed_rate'])}; no missing values; no resampling.", {"bullet": True}),
        ("Labels come from a logistic rule plus noise, so linear models do well.", {"bullet": True}),
        ("The company table is invented, not recruiter data.", {"bullet": True}),
    ], size=13)
    ftext(s, 0.55, 5.7, 8.0, 1.0, f"Placed rate by tier (Tier 1 {DS['by_tier']['1']['mean'] * 100:.1f}%, Tier 2 {DS['by_tier']['2']['mean'] * 100:.1f}%, Tier 3 {DS['by_tier']['3']['mean'] * 100:.1f}%) follows how the generator was built; 9-fold CV accuracy of XGBoost is {F.CV9:.4f}.", size=12, color=K.FNAVY)
    notes[12] = "Dataset, features and model in one view. The dataset is synthetic; the tier effect is a property of the generator, not a finding about real placements."

    # ---- 13 decision workflow
    s = S[12]
    set_title(s, "Decision Workflow")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.55, 1.8, 5.0, 5.1)
    ftext(s, 5.9, 2.0, 6.9, 4.8, [
        ("Probability: calibrated XGBoost output (Platt scaling).", {"bullet": True, "space_after": 12}),
        ("Refine: if p is in [0.4, 0.6], p = 0.5 p + 0.5 kNN probability.", {"bullet": True, "space_after": 12}),
        ("Decide: placement ready if p >= 0.5, otherwise needs improvement plus tips.", {"bullet": True, "space_after": 12}),
        ("Recommend: the same prediction is called for every feasible company and role; top 5 by score = p + 0.05 proximity + location bonus.", {"bullet": True}),
    ], size=18)
    notes[13] = "The serving rule in src/recommend.py. The threshold behaviour is analysed on slide 19."

    # ---- 14 user-system interaction
    s = S[13]
    set_title(s, "User – System Interaction")
    wipe(s)
    K.picture(s, F.diag("07_sequence_prediction.png"), 0.55, 1.8, 6.6, 5.1)
    K.picture(s, F.diag("08_use_case.png"), 7.4, 1.8, 5.4, 5.1)
    notes[14] = "Sequence of one prediction request and the use cases of student, officer, recruiter and developer."

    # ---- 15 storage design
    s = S[14]
    set_title(s, "Database Design (file store)")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.55, 1.8, 7.0, 4.5)
    rows = [["Store", "Content"], ["data/placement_dataset.csv", "3,000 students"], ["data/companies.csv", "30 companies"], ["models/xgb_calibrated.pkl", "calibrated XGBoost"],
            ["models/knn_residual.pkl", "kNN (k = 9)"], ["models/svm_residual.pkl", "SVM (RBF)"], ["models/preprocess.pkl", "maps and scaler"], ["models/metrics.json", "9 models x 5 metrics"], ["static/*.png", "SHAP, PCA, histogram"]]
    ftable(s, 7.8, 1.9, [2.7, 2.3], rows, size=11, row_h=0.48)
    ftext(s, 0.55, 6.4, 12.2, 0.5, "There is no SQL database (the batch description mentions MySQL): the ER diagram models the CSV and pickle files.", size=13, color=K.FNAVY, italic=True)
    notes[15] = "Honest point: the project stores files, not rows."

    # ---- 16 implementation highlights
    s = S[15]
    set_title(s, "Implementation Highlights")
    wipe(s)
    for i, (v, l, c) in enumerate([("9", "REST endpoints (+ /docs)", K.FBLUE), (str(F.TEST_SUMMARY["total"]), "scripted checks, all pass", K.GREEN), ("6", "React pages", K.ORANGE), ("1", "command to start (run.bat)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.1, v, l, c, 28)
    ftext(s, 0.55, 3.3, 12.2, 3.5, [
        ("Pydantic validation: CGPA 0-10, scores 0-100, counts 0-20; unknown company gives 404, bad range 422, missing models 503.", {"bullet": True, "space_after": 9}),
        ("The FastAPI layer imports the original src.recommend functions, so the Flask and React interfaces give identical numbers.", {"bullet": True, "space_after": 9}),
        ("Skill-gap tips are cutoff rules (CGPA, aptitude, technical, internship); the officer dashboard shows metrics, SHAP, PCA and branch rates.", {"bullet": True, "space_after": 9}),
        ("CORS allows only the front-end origin; there is no authentication, rate limiting or database (documented limits).", {"bullet": True}),
    ], size=17)
    notes[16] = "Highlights: validation, reuse of the original code, rule-based tips and the known limits."

    # ---- 17 results: model performance
    s = S[16]
    set_title(s, "Results – Model Performance (450 test students, seed 42)")
    wipe(s)
    names = ["logreg", "rf", "svm", "knn", "xgb_calibrated", "hybrid_xgb_knn", "hybrid_knn_labelfree"]
    labels = {"logreg": "Logistic Regression", "rf": "Random Forest", "svm": "SVM (RBF)", "knn": "kNN (k = 9)", "xgb_calibrated": "XGBoost + Platt",
              "hybrid_xgb_knn": "Hybrid (labels used)", "hybrid_knn_labelfree": "Hybrid (label-free)"}
    ms_key = {"hybrid_knn_labelfree": "hybrid_knn_labelfree"}
    rows = [["Model", "Accuracy", "ROC-AUC", "F1", "Recall", "5-seed accuracy"]]
    for n in names:
        r = X[n]
        rows.append([labels[n], f"{r['accuracy']:.4f}", f"{r['roc_auc']:.4f}", f"{r['f1']:.4f}", f"{r['recall']:.4f}", F.ms(ms_key.get(n, n))])
    ftable(s, 0.55, 1.85, [2.6, 1.05, 1.05, 0.95, 0.95, 1.9], rows, size=12, row_h=0.45, align=["l", "c", "c", "c", "c", "c"])
    ftext(s, 0.55, 5.55, 8.5, 1.4, "Shipped models (models/metrics.json). The five-seed column is a retrain on five random splits (mean ± sd); one seed differs from the shipped model by at most 0.011. Differences of one to two points are within one standard deviation.", size=12, color=K.FNAVY)
    K.card(s, 9.4, 1.85, 3.4, 4.9, "Reading", f"Logistic Regression has the best AUC ({X['logreg']['roc_auc']:.4f}) because the labels come from a logistic rule. The label-based hybrid is the best accuracy ({X['hybrid_xgb_knn']['accuracy']:.4f}) but the label-free one equals XGBoost + Platt ({X['hybrid_knn_labelfree']['accuracy']:.4f} vs {X['xgb_calibrated']['accuracy']:.4f}).", K.FBLUE, "i", body_size=13)
    notes[17] = "The table is the headline result. Stress the reading: the data is synthetic, so Logistic Regression is strong, and the hybrid gain depends on using the labels."

    # ---- 18 confusion matrices
    s = S[17]
    set_title(s, "Confusion Matrices & Recommendation Gain")
    wipe(s)
    K.picture(s, F.fig("confusion_matrices.png"), 0.55, 1.85, 7.5, 2.6)
    c0, c1, c2 = CONF["xgb_calibrated"], CONF["hybrid_oracle"], CONF["hybrid_labelfree"]
    ftext(s, 0.55, 4.6, 7.5, 2.3, [
        (f"XGBoost + Platt: {c0['fn']} placed students missed, {c0['fp']} false alarms.", {"bullet": True, "size": 14}),
        (f"Hybrid with labels: {c1['fn']} missed (recall 0.9067): 72 of the 87 XGBoost errors are chosen for refinement only because the label shows they are wrong.", {"bullet": True, "size": 14}),
        (f"Hybrid serving rule (label-free): {CONF['hybrid_serving']['fn']} missed, {CONF['hybrid_serving']['fp']} false alarms; kNN is right on {F.pct(F.RES['knn_acc_on_borderline'])} of the 31 borderline rows against {F.pct(F.RES['xgb_acc_on_borderline'])} for XGBoost.", {"bullet": True, "size": 14}),
    ], size=14)
    K.picture(s, F.fig("reco_gain.png"), 8.3, 1.85, 4.5, 2.7)
    ftext(s, 8.3, 4.6, 4.5, 2.3, [
        (f"Recommendation on {REC['students']} test students rated below 0.5: {REC['with_at_least_one_better_alt']} get a better option; mean base {REC['mean_base_prob']:.3f}, top-1 {REC['mean_top1_prob']:.3f}. The histogram covers the 56 with at least one feasible company.", {"bullet": True, "size": 13}),
        (f"Top pick reaches 0.5 for {REC['top1_reaches_ready']} of {REC['students']}. Gains are model estimates, not realised placements; picks lean to Tier-3 companies.", {"bullet": True, "size": 13}),
    ], size=13)
    notes[18] = "Rows are the true class, columns the prediction. The hybrid's recall gain disappears without labels. The recommendation gains are estimates from the same model."

    # ---- 19 threshold / feature analysis
    s = S[18]
    set_title(s, "Threshold, Band & Feature Analysis")
    wipe(s)
    K.picture(s, F.fig("threshold_sweep.png"), 0.55, 1.85, 5.2, 3.2)
    K.picture(s, F.fig("shap_top10.png"), 5.9, 1.85, 6.9, 3.2)
    bands = F.BAND
    rows = [["Refinement band (kNN replaces p)"] + [f"{b['lo']:.2f}-{b['hi']:.2f}" if b['hi'] != b['lo'] else "none" for b in bands],
            ["Rows replaced of 450"] + [str(b["replaced"]) for b in bands],
            ["Accuracy"] + [f"{b['accuracy']:.4f}" for b in bands]]
    ftable(s, 0.55, 5.2, [3.0] + [1.5] * 6, rows, size=11, row_h=0.4, align=["l"] + ["c"] * 6)
    ftext(s, 0.55, 6.5, 12.2, 0.5, "Accuracy is flat (0.793-0.807) across thresholds 0.30-0.70 while precision rises 0.776 to 0.869 and recall falls 0.918 to 0.791; widening the band changes accuracy by at most 0.4 points. SHAP: CGPA gap, backlogs, technical gap, aptitude gap lead.", size=11, color=K.FNAVY)
    notes[19] = "The decision threshold trades precision for recall; accuracy hardly moves. The kNN refinement band makes no measurable difference. SHAP shows the gap features dominate, as the generator intended."

    # ---- 20 performance & testing
    s = S[19]
    set_title(s, "Performance & Testing")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{F.LAT_PREDICT['mean']:.0f} ms", "predict_for, in process (mean)", K.FBLUE), (f"{PERF['POST /api/predict']['mean_ms']:.0f} ms", "POST /api/predict (mean, HTTP)", K.GREEN),
                                   (f"{PERF['POST /api/recruiter/shortlist']['mean_ms']:.0f} ms", "POST shortlist (mean)", K.ORANGE), (f"{PERF['POST /api/recommend']['mean_ms'] / 1000:.1f} s", "POST /api/recommend (mean, n = 8)", K.PURPLE)]):
        ftile(s, 0.55 + i * 3.1, 1.9, 2.9, 1.15, v, l, c, 26)
    bg = F.TEST_SUMMARY["by_group"]
    rows = [["Test layer", "What it covers", "Result"],
            ["Unit", "tier deltas, 24-feature matrix, determinism, monotonicity, recommendation filters", f"{bg['unit'][0]} / {bg['unit'][1]}"],
            ["Integration and functional", "all 9 endpoints, predict, recommend, shortlist", f"{bg['integration'][0] + bg['functional'][0]} / {bg['integration'][1] + bg['functional'][1]}"],
            ["Validation", "out-of-range, missing and unknown values, boundary profiles", f"{bg['validation'][0]} / {bg['validation'][1]}"],
            ["Security", "injection-like input, path traversal, CORS, 100 kB input, undefined routes", f"{bg['security'][0]} / {bg['security'][1]}"],
            ["Smoke test", "running server, every endpoint the React app uses", f"{F.SMOKE_TOTAL} / {F.SMOKE_TOTAL}"]]
    ftable(s, 0.55, 3.3, [2.6, 7.2, 2.4], rows, size=13, row_h=0.5)
    ftext(s, 0.55, 6.4, 12.2, 0.5, "Sequential calls on 127.0.0.1, Intel Core i5-12500, CPU only. Recommendation is slow because four pickle files are reloaded per call. Not tested: real data, load, automated browser tests. No authentication exists.", size=12, color=K.FNAVY, italic=True)
    notes[20] = "46 scripted checks plus the 15-check smoke test. Recommendation takes 3 to 8 s; caching the models is the obvious fix and was not done."

    # ---- 21 comparison with the base papers
    s = S[20]
    set_title(s, "Comparison with Base Papers (context only – data differ)")
    wipe(s)
    rows = [["Aspect", "IJETMS paper [1]", "arXiv paper [2]", "This project"],
            ["Task", "Placement prediction", "Admission prediction + recommendation", "Placement prediction + company-role recommendation"],
            ["Data", "Synthetic, size not stated; 80:20", "13,000 real records; 70/15/15", "3,000 synthetic students; 70/15/15"],
            ["Models", "LR, RF, Gradient Boosting", "7 baselines; XGBoost + kNN / SVM", "6 baselines; calibrated XGBoost + kNN / SVM"],
            ["Accuracy", "not reported", f"XGBoost 0.7241; hybrid 0.8723", f"XGBoost {X['xgb']['accuracy']:.4f}; hybrid {X['hybrid_xgb_knn']['accuracy']:.4f} (labels) / {X['hybrid_knn_labelfree']['accuracy']:.4f} (label-free)"],
            ["Brier score", "not reported", "0.1864 (Platt)", f"{F.BRIER:.4f} (Platt)"],
            ["Explainability", "feature importance, skill gap", "SHAP", "SHAP ranking + cutoff-rule tips"],
            ["Deliverable", "Streamlit app, PDF, e-mail", "method only", "Flask + FastAPI + React, officer and recruiter tools"]]
    ftable(s, 0.55, 1.85, [1.7, 2.8, 3.2, 4.5], rows, size=12, row_h=0.56)
    ftext(s, 0.55, 6.5, 12.2, 0.45, "Different data and conditions: our data is synthetic and rule-generated (even Logistic Regression scores 0.8289), the paper's real data is noisier. No superiority is claimed.", size=12, color=K.FNAVY, italic=True)
    notes[21] = "State plainly that the datasets differ. The accuracy numbers are context, not a ranking."

    # ---- 22-24 output screens
    s = S[21]
    set_title(s, "Output – Eligibility and Recommendations")
    wipe(s)
    K.picture(s, F.shot("02_predict.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("03_recommend.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    ftext(s, 0.55, 5.8, 12.2, 1.0, "Left: Microsoft SDE for CGPA 8.4, aptitude 84, technical 79 gives 22.2%, needs improvement; the cutoff panel shows technical 79 below 82 and the tip. Right: for a CGPA 7.5 student targeting Google SDE (10.8%) the top picks are ValueLabs and Infosys at about 90%.", size=12, color=K.FNAVY)
    notes[22] = "Real screenshots of the running application. Point out that recommendations lean to easier Tier-3 companies because the ranking maximises probability."
    s = S[22]
    set_title(s, "Output – Officer Dashboard and SHAP")
    wipe(s)
    K.picture(s, F.shot("04_dashboard.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("05_dashboard_shap.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    ftext(s, 0.55, 5.8, 12.2, 1.0, "Left: 3,000 students, placed rate 59.7%, hybrid AUC 0.8968, 9-fold CV 0.8147, Brier 0.1352, model comparison and placed rate by branch. Right: mean |SHAP| of all 24 features; the gap features lead.", size=12, color=K.FNAVY)
    notes[23] = "The dashboard shows shipped metrics (the label-based hybrid); the label-free check is in this deck and the report."
    s = S[23]
    set_title(s, "Output – Recruiter Shortlist and Companies")
    wipe(s)
    K.picture(s, F.shot("06_recruiter.png"), 0.55, 1.85, 6.0, 3.75, border=True)
    K.picture(s, F.shot("07_companies.png"), 6.8, 1.85, 6.0, 3.75, border=True)
    ftext(s, 0.55, 5.8, 12.2, 1.0, "Left: minimum CGPA 7.5 and aptitude 65 match 1,394 students; the top 50 by CGPA are listed with their historical placed tag. Right: the 30-company table with tier, cutoffs, package, location and hiring rate (invented values).", size=12, color=K.FNAVY)
    notes[24] = "Recruiter shortlist and the company board. The placed tags are the synthetic labels."

    # ---- 25 conclusion
    s = S[24]
    set_title(s, "Conclusion")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("**Delivered:** a working placement system: 24-feature pipeline, calibrated XGBoost with kNN refinement, skill-gap tips, company-role recommendation, officer dashboard, recruiter shortlist, FastAPI and React interfaces, and 46 passing checks.", {"bullet": True, "space_after": 10}),
        (f"**Measured:** {F.pct(X['xgb_calibrated']['accuracy'], 2)} accuracy and AUC {X['xgb_calibrated']['roc_auc']:.4f} for XGBoost + Platt on 450 synthetic test students; the hybrid's {F.pct(X['hybrid_xgb_knn']['accuracy'], 1)} needs true labels, {F.pct(X['hybrid_knn_labelfree']['accuracy'], 2)} without; Logistic Regression has the best AUC.", {"bullet": True, "space_after": 10}),
        ("**Objectives:** 3, 4, 5 and 6 met; 2 met for encoding (no resampling needed); 1 only with synthetic data.", {"bullet": True, "space_after": 10}),
        ("**Limits:** synthetic data and invented companies, no real-outcome validation, 3-8 s recommendations, no login or database. Not a decision tool for real students.", {"bullet": True}),
    ], size=21)
    notes[25] = "Be explicit about which objectives were met and which were not."

    # ---- 26 future scope
    s = S[25]
    set_title(s, "Future Scope")
    wipe(s)
    ftext(s, 0.55, 1.95, 12.2, 4.9, [
        ("Train and evaluate on real, anonymised placement records with a time-based split.", {"bullet": True, "space_after": 12}),
        ("Cache the model files per process so a recommendation takes well under a second.", {"bullet": True, "space_after": 12}),
        ("Validate categories, add login and roles, and store data in a database as the batch plan describes.", {"bullet": True, "space_after": 12}),
        ("Batch CSV prediction, PDF reports and per-student SHAP explanations.", {"bullet": True, "space_after": 12}),
        ("Check recommendations against realised outcomes; balance probability with package and fit.", {"bullet": True}),
    ], size=22)
    notes[26] = "Concrete next steps, ordered by value."

    # ---- 27 references
    s = S[26]
    set_title(s, "References")
    wipe(s)
    ids = [n for n in (1, 2, 3, 4, 5, 6, 8, 9, 13, 15, 17, 18, 19, 21, 23) if n <= len(REFS)]
    half = (len(ids) + 1) // 2
    ftext(s, 0.55, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in ids[:half]], size=11)
    ftext(s, 6.8, 1.85, 6.0, 5.0, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in ids[half:]], size=11)
    notes[27] = "Selected references; the complete numbered list is in the report."

    # ---- 28 thank you
    for sh in S[27].shapes:
        if sh.name == "Text 2":
            K.replace_text(sh, "Questions & Suggestions")
    notes[28] = "Invite questions."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes.get(i, ""))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "Placement")
    return out_path


if __name__ == "__main__":
    p = build(os.path.join(F.BATCH, "16.Final_PPT.pptx"))
    print(p, scan_deck(p))
