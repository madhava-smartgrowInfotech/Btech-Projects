"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"

BANNED = ["social media", "fakeddit", "pheme", "snopes", "mcnn", "bank transaction", "hidden risk", "recommender system", "skincare",
          "dayananda", "sindhu", "[student name", "[roll no", "sign language", "signspeak", "churn", "fraud"]


def scan_deck(path):
    """Return banned strings found anywhere in the saved package (slides, notes, layouts, masters, theme, properties)."""
    hits = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                t = re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore")).lower()
                for b in BANNED:
                    if b in t:
                        hits.append((n, b))
    return hits


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "campus placement, XGBoost, kNN, recommendation, educational data mining"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 32)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def lit_rows(pick=(0, 1, 2, 3, 4, 5, 7), cols=(0, 1, 2, 3)):
    rows = []
    for i in pick:
        if i < len(LIT_TABLE):
            r = LIT_TABLE[i]
            rows.append([r[c].replace(" (declared base paper)", " (base)").replace(" (technique base)", " (technique)") for c in cols])
    return rows


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    X, R = F.R, F.RES

    # ---- 1 title
    title_slide(S[0], "Hybrid XGBoost + kNN placement prediction, skill-gap feedback and company-role recommendation for a campus placement cell",
                "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: the system estimates a student's probability of being placed at a chosen company, "
                "explains the gaps, recommends better company-role matches, and gives officers and recruiters dashboards.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: from a student profile to a placement decision")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Placement cells keep student records in spreadsheets and shortlist by hand. Students get no early signal of readiness, and recruiters get no consistent, explained ranking.", K.ORANGE, "!"),
        ("Our idea", "Estimate the probability of placement for a student, company and role with a calibrated XGBoost model refined by kNN on borderline cases, then recommend company-role pairs with a higher estimated chance.", K.BLUE, "i"),
        ("The deliverable", "A web application: eligibility with skill-gap tips, top-5 recommendations, an officer dashboard (metrics, SHAP, PCA) and a recruiter shortlist, evaluated honestly on a 3,000-student synthetic dataset.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=18)
    notes[3] = "Three points: the placement-cell problem, the hybrid-model idea adapted from admissions work, and a complete application. Stress that the data is synthetic."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, REFS[0] + "   Technique source: " + REFS[1], size=11, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Declared base paper [1]", "Predicts placement with Logistic Regression, Random Forest and Gradient Boosting on a synthetic dataset (80:20 split); Streamlit app with batch CSV, skill-gap analysis, PDF report and e-mail.", K.BLUE, "1"),
        ("What it reports", "No accuracy, F1 or AUC is reported: the results section states only the split and the three models. Its own limitation: prediction without company recommendation.", K.GREEN, "2"),
        ("Technique source [2]", "Admissions study on 13,000 real records: XGBoost 0.7241 accuracy, XGBoost + kNN hybrid 0.8723, Platt calibration, SHAP and a university-program recommender. We adapt it to placements.", K.ORANGE, "3"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=15)
    notes[4] = ("The batch's declared base paper is the IJETMS placement paper; it gives no numeric results, so numeric comparison is made with the arXiv admissions paper whose method the code adapts. "
                "Both are cited; the datasets differ.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported result"]] + lit_rows()
    n = len(rows)
    fs = 14 if n <= 3 else 12 if n <= 5 else 11 if n <= 7 else 10
    K.table(s, 0.6, 1.35, [2.2, 4.0, 2.6, 3.33], rows, font_size=fs, row_h=min(1.2 if n <= 5 else 0.8, 5.4 / n), header_size=12)
    notes[5] = ("Early placement predictors compare Logistic Regression, trees and boosting; the hybrid and recommendation ideas come from the admissions paper. "
                "Where a work's figures were not checked they are not quoted.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence: data and protocol set the number")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 7.6, 5.3,
                     ["[2] XGBoost\n(real data)", "[2] Hybrid\n(real data)", "Ours XGBoost\n+ Platt", "Ours Hybrid\n(labels used)", "Ours Hybrid\n(label-free)"],
                     [("Test accuracy", [0.7241, 0.8723, X["xgb_calibrated"]["accuracy"], X["hybrid_xgb_knn"]["accuracy"], X["hybrid_knn_labelfree"]["accuracy"]])],
                     colors=[K.BLUE], number_format="0.0%", font_size=12, max_value=1.1)
    for k in (2, 3, 4):
        pt = gf.chart.plots[0].series[0].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = K.ORANGE
    K.card(s, 8.5, 1.45, 4.23, 2.6, "Takeaway", "The same method scores 72% on real admissions data and 81% on our synthetic data. A number is meaningful only with its data and protocol.", K.GREEN, "↑", body_size=14)
    K.card(s, 8.5, 4.2, 4.23, 2.5, "Caveat", "The hybrid gains (+15 and +3 points) pick the refined cases with the true test labels; without labels ours is 80.9%, like plain XGBoost (80.7%).", K.ORANGE, "!", body_size=14)
    notes[6] = ("Blue bars are from the arXiv base paper [2] (real GradCafe data); orange bars are ours on a synthetic dataset, seed 42, 450 test students. "
                "The label-free bar is our own measurement: only borderline cases (p between 0.4 and 0.6) are refined, which is the only rule usable at serving time.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Student records in spreadsheets", K.NAVY), ("Manual eligibility check", K.NAVY), ("Recruiter shortlists by hand", K.NAVY), ("Result after the drive", K.NAVY)], size=15, gap=0.5)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Manual placement cell", "Records are stored, not analysed. Eligibility is judged against cutoffs by hand, shortlisting is slow and varies between recruiters, and students learn where they stand late.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Research predictors", "Studies classify placed / not placed from academics and skills. Many stop at a yes-or-no or an accuracy figure, with no company recommendation and no actionable gap feedback.", K.ORANGE, "2", body_size=17)
    notes[7] = "Today's path is manual. Published predictors exist, but few tell a student which company is realistic or which skill to improve."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("Data without analysis", "Placement cells store student data but perform no predictive analysis on it.", K.ORANGE, "1"),
        ("No early signal", "Students receive no indication of readiness while there is still time to improve.", K.BLUE, "2"),
        ("Manual shortlisting", "Shortlisting is slow and inconsistent across recruiters.", K.GREEN, "3"),
        ("No skill feedback", "Nothing says which specific skill a student should acquire next, or which company fits.", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=17)
    notes[8] = "The four problem points come from the batch problem statement. Each maps to a module: prediction, tips, recommendation, recruiter shortlist."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Prediction only", "Placement predictors rarely go on to recommend which company and role a student should target."),
            ("Actionable feedback", "Few give the student the specific gaps against a company's requirements."),
            ("Calibrated probabilities", "A recommendation ranking needs reliable probabilities, yet calibration is seldom reported."),
            ("Honest evaluation", "Results often lack the split, several seeds, a confusion matrix or a check that a refinement rule works without labels."),
            ("Real data", "Real placement records are private; published work, including the base paper, falls back to synthetic data.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=18, color=K.INK, anchor="m")
    notes[9] = ("The batch survey states the gap as extending prediction into company recommendation and actionable skill-gap feedback. "
                "We address the first three directly and measure the fourth; the fifth is a stated limitation.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, PROBLEM_STATEMENT, size=19, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Academic record, skill and aptitude scores, backlogs, target company and role"), ("Output", "Placement probability, label, skill-gap tips, top-5 company-role pairs"), ("Users", "Students, placement officers, recruiters")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.5)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 0.8, b_, size=14, color=K.TXT)
    notes[10] = "Read the statement slowly: the input is a profile plus a target, the output is a probability with reasons and alternatives, and three user groups are served."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 12}) for i, o in enumerate(OBJECTIVES, 1)], size=22, color=K.INK)
    notes[11] = "The six objectives from the proposal. Later reviews report which were met; real records (objective 1) could not be obtained, so a synthetic dataset was generated."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Data", "3,000 synthetic students and 30 companies; 24 features incl. gaps against each company's cutoffs; stratified 70/15/15 split.", K.BLUE, "1"),
        ("Predict", "Calibrated XGBoost (Platt) with kNN refinement for borderline probabilities 0.4 to 0.6; baselines for comparison.", K.GREEN, "2"),
        ("Recommend", "Feasibility filters on CGPA, aptitude, backlogs and tier; rank company-role pairs by probability plus proximity.", K.ORANGE, "3"),
        ("Explain", "Rule-based skill-gap tips, SHAP ranking, PCA of errors and a probability histogram for the officer.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.4, t, b_, col, gl, body_size=15)
    K.box(s, 0.6, 5.1, 12.13, 1.5)
    K.circle(s, 0.85, 5.4, 0.7, K.NAVY, "W", 18)
    K.text(s, 1.8, 5.3, 10.7, 0.5, "Application", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.75, "FastAPI service (9 endpoints) over the unchanged Python core, a React interface (eligibility, recommendations, officer dashboard, recruiter, companies) and the original Flask pages.", size=15, color=K.TXT)
    notes[12] = "Four stages from data to explanation, and the application around them. The hybrid idea is adapted from the admissions paper [2]."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 2.8, 1.3, 7.7, 5.3)
    K.text(s, 0.6, 6.62, 12.13, 0.3, "React and Flask front ends call the FastAPI layer and the shared core; models and data are files on disk (no SQL database)", size=12, color=K.TXT, align="c", italic=True)
    notes[13] = "Follow the diagram: interfaces on the left, the API, the core modules and the files on the right."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Placement predictor", "Calibrated probabilities with a transparent decision at 0.5.", K.BLUE, "1"),
            ("Skill-gap feedback", "Tips against the company's CGPA, aptitude and technical cutoffs.", K.GREEN, "2"),
            ("Company-role recommendation", "Top-5 alternatives with estimated gain and match rate.", K.ORANGE, "3"),
            ("Officer and recruiter tools", "Metrics, SHAP, PCA, branch rates and a filtered shortlist.", K.PURPLE, "4"),
            ("Honest evaluation", "Five seeds, confusion matrices and a label-free check of the hybrid rule.", K.BLUE, "5"),
            ("Reproducibility", "Scripts regenerate the data, models, metrics, figures and tests.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c = divmod(i, 3)
        K.card(s, 0.6 + c * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=16)
    notes[14] = "Expected outcomes. We report whatever the held-out data shows, including where the hybrid does not help."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [n for n in (1, 2, 3, 4, 5, 6, 15, 17, 18, 19, 21, 23) if n <= len(REFS)]
    half = (len(idx) + 1) // 2
    K.text(s, 0.6, 1.35, 5.95, 5.5, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in idx[:half]], size=10.5, color=K.INK)
    K.text(s, 6.78, 1.35, 5.95, 5.5, [(f"[{n}] {REFS[n - 1]}", {"space_after": 4}) for n in idx[half:]], size=10.5, color=K.INK)
    notes[15] = "Selected references; the complete numbered list is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = "Likely questions: why synthetic data; why the hybrid gain shrinks without labels; why Logistic Regression has the best AUC; how recommendations are validated (they are model estimates only)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, "Placement")
    return out_path


if __name__ == "__main__":
    p = build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx"))
    print(p, scan_deck(p))
