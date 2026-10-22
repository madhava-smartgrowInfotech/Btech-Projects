"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p("Placement-probability prediction, skill-gap feedback and company-role recommendation with a calibrated hybrid XGBoost and kNN model", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base papers and the developers of the open-source software used in this project "
        "(scikit-learn, XGBoost, SHAP, pandas, NumPy, Flask, FastAPI, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    r = F.R
    h, x, lf, sv = r["hybrid_xgb_knn"], r["xgb_calibrated"], r["hybrid_knn_labelfree"], r["hybrid_knn_serving"]
    d.p("Campus placement decides the early careers of engineering students, yet placement cells usually store student records in spreadsheets and shortlist by hand. Students learn where they stand "
        "only when a shortlist is published, and they receive no statement of which skill to improve. This project presents the Smart Campus Placement System, a web application that estimates the "
        "probability that a student is placed at a chosen company and role, explains the gaps against that company's requirements, recommends alternative company-role pairs, and gives placement "
        "officers and recruiters a dashboard and a shortlist. It adapts to placements the hybrid framework that a recent admission-prediction paper proposed: a Platt-calibrated XGBoost model, a "
        "k-nearest-neighbour model that refines borderline cases, SHAP explanations and feasibility-filtered recommendation. A FastAPI service with a React interface, and the original Flask "
        "application, serve the same Python core.")
    d.p(f"No real placement records were available, so the system is trained and evaluated on {F.DS['rows']:,} synthetic students and an invented table of {F.DS['companies']} companies; every result below "
        f"therefore describes this generated data and not real students. On a held-out test set of {F.TEST_N} students, calibrated XGBoost reaches {F.pct(x['accuracy'], 2)} accuracy (ROC-AUC {F.f4(x['roc_auc'])}, "
        f"Brier score {F.f4(F.BRIER)}), and the hybrid defined in the base paper reaches {F.pct(h['accuracy'], 2)}. That gain depends on the true test labels, which the paper's residual rule uses to choose the cases "
        f"to refine; the form that can run when a student is scored reaches {F.pct(lf['accuracy'], 2)} (borderline cases only) and {F.pct(sv['accuracy'], 2)} (blended serving rule), a difference of about one student "
        f"in {F.TEST_N}. Over five random splits the means are {F.ms('xgb_calibrated')} for calibrated XGBoost and {F.ms('hybrid_knn_serving')} for the serving rule, and Logistic Regression is as good ("
        f"{F.ms('logreg')}). For {F.REC['students']} test students rated below 0.5 for their own target, the recommender raises the estimated probability from {F.REC['mean_base_prob']:.2f} to {F.REC['mean_top1_prob']:.2f} "
        f"on average, an estimate that was not checked against real outcomes. A prediction takes about {F.PERF['POST /api/predict']['mean_ms']:.0f} ms, a recommendation about "
        f"{F.PERF['POST /api/recommend']['mean_ms'] / 1000:.1f} s. {F.TEST_SUMMARY['passed']} automated checks and a {F.SMOKE_TOTAL}-check smoke test pass. The system has no authentication or database, "
        "and a pilot with real placement data, access control and caching of the models are the main future steps.")
    d.p("**Keywords:** campus placement, placement prediction, educational data mining, XGBoost, probability calibration, k-nearest neighbours, SHAP, company recommendation, FastAPI, React.", align="left")


ABBR = [
    ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CGPA", "Cumulative Grade Point Average"), ("CORS", "Cross-Origin Resource Sharing"),
    ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"), ("CV", "Cross-Validation"), ("DFD", "Data Flow Diagram"), ("DSA", "Data Structures and Algorithms"),
    ("ER", "Entity-Relationship"), ("F1", "Harmonic mean of precision and recall"), ("FN / FP", "False Negative / False Positive"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JSON", "JavaScript Object Notation"), ("kNN", "k-Nearest Neighbours"), ("LPA", "Lakhs Per Annum"), ("ML", "Machine Learning"), ("PCA", "Principal Component Analysis"),
    ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("SHAP", "SHapley Additive exPlanations"), ("SPA", "Single-Page Application"),
    ("SVM", "Support Vector Machine"), ("TN / TP", "True Negative / True Positive"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("XGBoost", "Extreme Gradient Boosting"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
