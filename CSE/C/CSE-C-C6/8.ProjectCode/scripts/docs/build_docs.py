"""
Builds the six chapter documents and the full report (17.Final Document.docx) into the batch folder.

    python -m scripts.docs.build_docs            # from 8.ProjectCode; all documents
    python -m scripts.docs.build_docs --only 9   # one document
"""

import argparse
import os

from docx.shared import Inches, Pt

from . import appendix, ch1, ch2, ch3_4, ch5, ch6, ch7, ch8, ch9_10, front
from . import facts as F
from .blocks import FIG, H1, H2, P, render, renum
from .kit import Doc, add_runs
from .refs import REFS

NAMES = {
    9: "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx",
    10: "10.SCOPE & LIMITATIONS.docx",
    11: "11. BLOCK DIAGRAMS ALL.docx",
    12: "12.DESIGN &IMPLEMENTATION.docx",
    14: "14.RESULTS.docx",
    15: "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    17: "17.Final Document.docx",
}


def new_doc(doc_title):
    return Doc(doc_title, header_text=f"{F.PRODUCT}  |  {doc_title}")


def title_block(d, doc_title):
    d.footer_numbers()
    d.spacer(24)
    d.p(F.TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p(doc_title, align="center", size=14, bold=True, space_after=14)
    d.p(f"{F.DEPT}, {F.COLLEGE_SHORT}", align="center", size=12, space_after=2)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})  ·  Academic year {F.YEAR}", align="center", size=12, space_after=12)
    rows = [[str(i), n, r] for i, (n, r) in enumerate(F.STUDENTS, 1)]
    d.table("Project team", ["S.No.", "Student name", "Roll number"], rows, widths=[0.7, 3.4, 1.8], size=11, align=["c", "l", "c"])
    d.p(f"**Project guide:** {F.GUIDE}", align="center", size=12, space_after=10)
    d.p("Chapter and section numbers quoted in the text (for example Chapter 8 or Section 6.2) refer to the full project report, 17.Final Document. "
        "Figures and tables in this document are numbered from 1.", size=10.5, italic=True, space_after=6)


def save(d, name):
    path = d.save(os.path.join(F.BATCH, name))
    print("wrote", path)
    return path


def add_references(d):
    for i, r in enumerate(REFS, 1):
        par = d.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.45)
        par.paragraph_format.first_line_indent = Inches(-0.45)
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(5)
        add_runs(par, f"[{i}]  {r}", 10)


# ----------------------------------------------------------------- chapter documents
def doc9():
    d = new_doc("Methodology, Novelty and Key Contributions")
    title_block(d, "Methodology, Novelty and Key Contributions")
    d.h1("1. METHODOLOGY", page_break=False)
    d.p("This document describes how the fake news detection system (TruthGlow) works from end to end and what is new in it. Part 1 gives the pipeline, the data, the algorithms and the reasons they were chosen; "
        "Part 2 sets the work against the base paper and separates what is new from what is standard; Part 3 lists the contributions.")
    render(d, renum(ch6.methodology_blocks(), "1"))
    d.h1("2. NOVELTY AGAINST THE BASE PAPER")
    render(d, renum(ch9_10.novelty_blocks(), "2"))
    d.h1("3. KEY CONTRIBUTIONS")
    render(d, ch9_10.contribution_blocks())
    return save(d, NAMES[9])


def doc10():
    d = new_doc("Scope and Limitations")
    title_block(d, "Scope and Limitations")
    d.h1("1. SCOPE", page_break=False)
    render(d, renum(ch9_10.scope_blocks(), "1"))
    d.h1("2. ASSUMPTIONS")
    render(d, ch9_10.assumption_blocks())
    d.h1("3. LIMITATIONS")
    render(d, ch9_10.limitation_blocks())
    render(d, [P("Three figures illustrate the main limitations. The first shows how little of the feature variance the 300-component PCA cap keeps and how many components the Firefly search retains. "
                 "The second shows the accuracy of the verdicts on statements of different lengths: longer statements are harder. The third shows that the verdict and the class probabilities "
                 "shown in the application disagree on many statements."),
               FIG(F.fig("f08_feature_funnel.png"), "Feature counts after PCA and after the Firefly selection (the PCA cap of 300 components keeps only part of the variance).", 4.6),
               FIG(F.fig("f10_accuracy_by_length.png"), "Accuracy by statement length on the LIAR2 test split (seed 42).", 4.8),
               FIG(F.fig("f11_label_vs_probability.png"), "Verdict shown by the application against the class with the highest probability.", 4.8)])
    d.h1("4. WHAT WAS NOT MEASURED")
    render(d, ch9_10.not_measured_blocks())
    return save(d, NAMES[10])


def doc11():
    d = new_doc("Block Diagrams")
    title_block(d, "Block Diagrams (all)")
    d.h1("1. INDEX OF DIAGRAMS", page_break=False)
    rows = [[str(i), t, f"`docs/diagrams/{f}`"] for i, (_, t, f, _, _) in enumerate(ch5.DIAGRAMS, 1)]
    d.p("Every diagram was generated from the real design by `scripts/docs_build/make_diagrams.py` (matplotlib, 150 dpi) and mirrors the code in `backend/main.py` and `src/`. "
        "The application has no SQL database, so the entity-relationship diagram shows its file-based data model.")
    d.table("Diagrams in this document", ["No.", "Diagram", "Image file"], rows, widths=[0.5, 3.0, 2.8], size=10.5, align=["c", "l", "l"])
    d.h1("2. DIAGRAMS")
    for key, title, file, caption, text in ch5.DIAGRAMS:
        d.h2(title)
        d.figure(F.diag(file), caption, 5.0 if key in ("seq", "algo", "uc", "uml") else 5.8)
        d.p(text)
    return save(d, NAMES[11])


def doc12():
    d = new_doc("Design and Implementation")
    title_block(d, "Design and Implementation")
    d.h1("1. DESIGN", page_break=False)
    render(d, renum(ch5.blocks()[1:], "1"))
    d.h1("2. IMPLEMENTATION")
    render(d, renum(ch6.application_blocks(), "2"))
    d.h1("3. KEY CODE WALK-THROUGHS")
    render(d, renum(ch9_10.walkthrough_blocks(), "3"))
    d.h1("4. IMPLEMENTATION DECISIONS AND THEIR REASONS")
    render(d, ch9_10.decision_blocks())
    return save(d, NAMES[12])


def doc14():
    d = new_doc("Results")
    title_block(d, "Results")
    d.h1("1. RESULTS AND OUTPUT SCREENS", page_break=False)
    render(d, renum(ch8.blocks()[1:], "1"))
    return save(d, NAMES[14])


def doc15():
    d = new_doc("Conclusion, Future Scope and References")
    title_block(d, "Conclusion, Future Scope and References")
    d.h1("1. CONCLUSION", page_break=False)
    render(d, renum(ch9_10.conclusion_blocks(), "1"))
    render(d, [P("The comparison with the simpler models and the accuracy by confidence bin behind the key findings are shown below."),
               FIG(F.fig("f01_model_comparison.png"), "Test scores of the proposed pipeline and its comparisons on the LIAR2 test split.", 5.6),
               FIG(F.fig("f06_confidence_bins.png"), "Accuracy of the verdict by confidence bin (seed 42, test split).", 4.8)])
    d.h1("2. FUTURE SCOPE")
    render(d, renum(ch9_10.future_blocks(), "2"))
    d.h1("3. REFERENCES")
    add_references(d)
    return save(d, NAMES[15])


# ----------------------------------------------------------------- full report
def doc17():
    d = new_doc("Project Report")
    front.cover(d)
    d.new_section(fmt="lowerRoman", start=1)
    front.certificate(d)
    front.declaration(d)
    front.acknowledgement(d)
    front.abstract(d)
    d.h1("TABLE OF CONTENTS")
    d.toc_field('TOC \\o "1-2" \\h \\z \\u')
    d.h1("LIST OF FIGURES")
    d.toc_field('TOC \\h \\z \\c "Figure"')
    d.h1("LIST OF TABLES")
    d.toc_field('TOC \\h \\z \\c "Table"')
    d.h1("LIST OF ABBREVIATIONS")
    front.abbreviations(d)
    d.new_section(fmt="decimal", start=1)
    render(d, ch1.blocks(), h1_page_break=False)
    render(d, ch2.blocks())
    render(d, ch3_4.blocks())
    render(d, ch5.blocks())
    render(d, ch6.blocks())
    render(d, ch7.blocks())
    render(d, ch8.blocks())
    render(d, ch9_10.chapter9_blocks())
    render(d, ch9_10.chapter10_blocks())
    d.h1("REFERENCES")
    add_references(d)
    render(d, appendix.code_blocks())
    render(d, appendix.guide_blocks())
    render(d, appendix.screens_blocks())
    return save(d, NAMES[17])


BUILDERS = {9: doc9, 10: doc10, 11: doc11, 12: doc12, 14: doc14, 15: doc15, 17: doc17}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", type=int, nargs="*")
    a = ap.parse_args()
    for k, fn in BUILDERS.items():
        if not a.only or k in a.only:
            fn()


if __name__ == "__main__":
    main()
