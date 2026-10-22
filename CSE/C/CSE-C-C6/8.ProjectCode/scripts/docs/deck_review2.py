"""13.Review_2_PPT.pptx: 20 slides, design, methodology and progress. Built from D:\\abstracts\\Templates\\Review_2_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K

TEMPLATE = r"D:\abstracts\Templates\Review_2_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


# ------------------------------------------------------------------ local helpers (self-contained)
def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  \u00b7  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "fake news detection, PCA, Firefly algorithm, multi-class SVM, LIAR2"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 30)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 18)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS), 14)
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  \u00b7  Batch {F.BATCH_NO} ({F.SECTION})  \u00b7  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 2")
    S = list(prs.slides)
    assert len(S) == 20, len(S)
    notes = {}
    A, E = F.A, F.EA
    P, PD = A["proposed_test"], A["pca_svc_default_test"]
    RS, MO = A["random_search_test"], A["pca_svc_mask_only_test"]
    LR = A["ref_full_LogisticRegression_test"]
    npca, nsel = round(A["n_pca"]["mean"]), round(A["n_selected"]["mean"])
    var = A["pca_variance"]["mean"]
    lat = F.LAT[0]

    # ---- 1 title
    title_slide(S[0], "A light PCA + Firefly + multi-class SVM pipeline with a PARTIALLY TRUE class, behind a FastAPI service and a React interface",
                "REVIEW 2  \u00b7  DESIGN, METHODOLOGY AND PROGRESS")
    notes[1] = "This review shows how the system is designed and built, what is implemented, the first measured results on real data, and the open items."

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Recap and architecture"), ("02", "Methodology and dataset"), ("03", "Preprocessing and model"), ("04", "Training and decision logic"),
               ("05", "Data flow, database and UML"), ("06", "Technology stack"), ("07", "Implementation progress and results"),
               ("08", "Output screens and testing"), ("09", "Challenges and next steps")],
           [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda for the design and progress review."

    # ---- 3 recap
    s = S[2]
    set_title(s, "Recap: problem and proposed system")
    wipe(s)
    W = 3.85
    for i, (t, b_, col, gl) in enumerate([
        ("Problem", "Fabricated news spreads faster than fact-checks, and high-dimensional text features make classifiers slow and prone to overfitting. Default SVM parameters are rarely optimal, and most detectors have only two classes.", K.ORANGE, "!"),
        ("Proposed system", "TF-IDF and 25 stylometric features are reduced by PCA. The Firefly algorithm picks the PCA components and the SVM parameters. An RBF multi-class SVM returns REAL, PARTIALLY TRUE or FAKE with a confidence.", K.BLUE, "i"),
        ("Base paper link", "Saad et al. (2026) combine class weighting with transformer embeddings for Arabic, binary detection. We keep class weighting and baselines but use a cheap classical pipeline, English text, three classes and a CPU. Numbers are not comparable.", K.GREEN, "B"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, 4.3, t, b_, col, gl, body_size=17)
    notes[3] = "One-minute recap of Review 1: the gap between heavy neural detectors and a light, tunable pipeline, and how the base paper relates (motivation and class weighting only)."

    # ---- 4 architecture
    s = S[3]
    set_title(s, "System architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.6, 5.6)
    K.text(s, 8.5, 1.5, 4.2, 5.3, [
        ("**Three tiers**", {"space_after": 8}),
        ("Presentation: React 18 app (TruthGlow, port 5406) with six screens.", {"bullet": True}),
        ("Application: FastAPI (port 8406), 8 endpoints under /api, one model held in memory.", {"bullet": True}),
        ("Domain: the pure-Python library src/ (preprocessing, dataset, firefly, pipeline), shared with the CLI and the Streamlit app.", {"bullet": True}),
        ("Storage: one joblib model bundle; no SQL database.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[4] = "The FastAPI layer calls the same library functions as the Streamlit app and the command line, so all three front ends use one pipeline and one model file."

    # ---- 5 methodology workflow
    s = S[4]
    set_title(s, "Methodology: end-to-end workflow")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.2, [("Labelled text", K.NAVY), ("Clean + TF-IDF + style", K.PURPLE), ("Scale + PCA", K.PURPLE), ("80/20 split", K.GREEN),
                                      ("Firefly search", K.ORANGE), ("RBF MSVM", K.BLUE), ("Evaluate + save", K.NAVY)], size=13, gap=0.22)
    for i, (t, b_) in enumerate([("One fitted bundle", "Vectorizer, scaler, PCA, component mask, C, gamma and classifier are saved together, so training and serving use identical transforms."),
                                 ("Class weighting", "Class weights are balanced in the Firefly fitness and in the final SVM, as the base paper does for its imbalance."),
                                 ("Reproducible", "scripts/run_experiments.py runs the real-data study; its JSON feeds every figure and number in the documents.")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 3.1, 3.85, 2.7)
        K.text(s, x + 0.25, 3.25, 3.4, 0.5, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 3.85, 3.4, 2.4, b_, size=17, color=K.TXT)
    notes[5] = "The same workflow runs from train.py, the Streamlit app and the Retrain screen. The experiment script reuses the product's own pipeline class."

    # ---- 6 dataset
    s = S[5]
    set_title(s, "Dataset: LIAR2 and the demo corpus")
    wipe(s)
    for i, (v, l, c) in enumerate([(f"{F.NTRAIN_FULL:,}", "training statements", K.BLUE), (f"{F.NVALID:,}", "validation statements", K.GREEN),
                                   (f"{F.NTEST:,}", "test statements", K.ORANGE), (f"{F.NTRAIN:,}", "training subsample per run (3 seeds)", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=12)
    rows = [["LIAR2 label", "Project class", "Statements (whole dataset)"], ["true, mostly-true", "REAL", "2,585 + 3,429"], ["half-true", "PARTIALLY TRUE", "3,709"],
            ["barely-true, false, pants-on-fire", "FAKE", "3,603 + 6,605 + 3,031"]]
    K.table(s, 0.6, 3.0, [3.6, 2.2, 2.6], rows, font_size=14, row_h=0.6, align=["l", "l", "l"])
    tc = A["test_class_counts"]
    K.text(s, 9.3, 3.0, 3.45, 3.9, [
        (f"Test split: {tc[0]} REAL, {tc[1]} PARTIALLY TRUE, {tc[2]} FAKE. Always answering FAKE scores accuracy {F.f3(A['majority_test']['accuracy']['mean'])} but macro-F1 {F.f3(A['majority_test']['f1_macro']['mean'])}.", {"bullet": True}),
        ("Only the statement text is used; speaker, justification and counts are left out because they leak the label.", {"bullet": True}),
        ("Source: Hugging Face chengxuphd/liar2, Apache-2.0 [4].", {"bullet": True}),
    ], size=14, color=K.INK, space_after=8)
    K.text(s, 0.6, 5.55, 8.4, 1.4, [
        ("**Shipped demo corpus.** 750 synthetic template sentences (250 per class) train the model that ships with the project. Its 1.000 scores only check the pipeline: the same model scores accuracy "
         f"{F.f3(A['demo_model_on_liar2_test']['accuracy'])} and macro-F1 {F.f3(A['demo_model_on_liar2_test']['f1_macro'])} on real LIAR2 statements.", {}),
    ], size=14, color=K.INK)
    notes[6] = ("LIAR2 is 18,369 / 2,297 / 2,296 PolitiFact statements with six labels, mapped to the project's three classes. Training uses a stratified 3,000-statement subsample because the Firefly search needs thousands of SVM fits on a CPU. "
                "The demo corpus is synthetic, so its perfect scores are not evidence of accuracy on real news.")

    # ---- 7 preprocessing
    s = S[6]
    set_title(s, "Data preprocessing and feature extraction")
    wipe(s)
    K.picture(s, F.diag("11_feature_pipeline.png"), 0.6, 1.3, 12.13, 3.55)
    K.text(s, 0.6, 5.0, 6.0, 1.95, [
        ("**Cleaning**", {}),
        ("Lower-case, remove URLs, @mentions and #tags, keep letters only.", {"bullet": True}),
        ("Remove stop words and tokens of 2 characters or fewer.", {"bullet": True}),
        ("No stemming is implemented (objective 2 lists it).", {"bullet": True}),
    ], size=14, color=K.INK, space_after=3)
    K.text(s, 6.9, 5.0, 5.85, 1.95, [
        ("**Features**", {}),
        ("TF-IDF, 1-2 grams, sublinear tf, min_df 2, max_df 0.9, at most 4,000 columns.", {"bullet": True}),
        ("25 stylometric signals: capitals, ! and ?, digits, URLs, lexicon ratios, word lengths.", {"bullet": True}),
        ("StandardScaler, then PCA (95 % variance target, cap 300).", {"bullet": True}),
    ], size=14, color=K.INK, space_after=3)
    notes[7] = "The cleaning and the 25 stylometric features are in src/preprocessing.py. The scaler and PCA are fitted on the training subsample only and saved with the model."

    # ---- 8 model
    s = S[7]
    set_title(s, "Model: PCA, Firefly selection and MSVM")
    wipe(s)
    for i, (v, l, c) in enumerate([("4,025", "TF-IDF + stylometric features", K.BLUE), (f"{npca}", f"PCA components ({F.pct(var)} variance kept)", K.GREEN),
                                   (f"{nsel}", "components kept by Firefly (mean of 3 seeds)", K.ORANGE), ("3", "classes (one-vs-rest shape, RBF kernel)", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=12)
    rows = [["Stage", "What it does"], ["PCA", "Standardised features to at most 300 orthogonal components; the cap, not the 95 % target, is what binds on real text"],
            ["Firefly", "Chooses a component mask plus C in [0.1, 1000] and gamma in [1e-4, 10]; fitness is 3-fold CV macro-F1 minus 0.02 x selected share"],
            ["MSVM", "SVC with RBF kernel, balanced class weights, one-vs-rest decision shape, probability=True (Platt scaling)"]]
    K.table(s, 0.6, 3.0, [1.6, 10.53], rows, font_size=15, row_h=0.85, first_bold=True)
    K.text(s, 0.6, 6.5, 12.1, 0.5, "A finding: with a 300-component cap PCA keeps only about 29 % of the variance of the standardised real-text features.", size=14, color=K.TXT, italic=True)
    notes[8] = "The model is three fitted stages. The 25-feature stylometric block and the TF-IDF block are concatenated before scaling. Firefly's vector is 300 mask logits plus log10 C and log10 gamma."

    # ---- 9 training configuration
    s = S[8]
    set_title(s, "Training configuration")
    wipe(s)
    st = A["settings"]
    rows = [["Setting", "Value"],
            ["Firefly", f"{st['fireflies']} fireflies, {st['generations']} generations, alpha 0.2 (decay 0.97), beta0 1, gamma 1; about {A['n_firefly_evals']['mean']:.0f} fitness evaluations"],
            ["Fitness", "3-fold stratified CV macro-F1 minus 0.02 x (selected / total components)"],
            ["Split", f"80/20 stratified inside the {F.NTRAIN:,}-statement subsample; scored on the official validation and test splits"],
            ["Final SVM", "RBF, class_weight balanced, probability=True"],
            ["Seeds", f"{', '.join(str(x) for x in sorted(F.SEEDS))}; results are mean \u00b1 std"],
            ["Hardware", f"CPU only, Intel Core i5-12500 (6 cores, 12 threads); one full fit takes {A['train_time_s']['mean']:.0f} \u00b1 {A['train_time_s']['std']:.0f} s"]]
    K.table(s, 0.6, 1.35, [2.4, 9.73], rows, font_size=13, row_h=0.46, first_bold=True)
    K.picture(s, F.fig("f02_firefly_convergence.png"), 0.6, 4.75, 4.2, 2.3)
    K.text(s, 5.1, 4.95, 7.6, 2.0,
           f"Convergence: best fitness reaches {F.f3(A['firefly_fitness_best']['mean'])} \u00b1 {F.f3(A['firefly_fitness_best']['std'])} against {F.f3(A['random_fitness_best']['mean'])} \u00b1 {F.f3(A['random_fitness_best']['std'])} for random search with the same number of evaluations, "
           "most of the gain comes in the first generations, though one seed was still improving at generation 10. Firefly's hyper-parameters were not tuned.", size=14, color=K.TXT)
    notes[9] = "Settings come from src/firefly.py and src/pipeline.py. Training time is dominated by the Firefly search; random search with the same budget takes about as long."

    # ---- 10 decision logic
    s = S[9]
    set_title(s, "Decision logic: Firefly search and the verdict")
    wipe(s)
    K.picture(s, F.diag("10_core_algorithm_flowchart.png"), 0.6, 1.3, 5.2, 5.65)
    K.text(s, 6.1, 1.5, 6.6, 5.4, [
        ("**Move.** A dimmer firefly moves towards a brighter one with attraction exp(-gamma r\u00b2 / d) plus decaying random noise.", {"bullet": True, "space_after": 9}),
        ("**Decode.** Mask = sigmoid(logits) > 0.5; log10 C and log10 gamma are clipped to their ranges.", {"bullet": True, "space_after": 9}),
        ("**Verdict.** The label is SVC.predict; REAL / PARTIALLY TRUE / FAKE probabilities come from Platt scaling.", {"bullet": True, "space_after": 9}),
        (f"**Known mismatch.** On {F.pct(F.LV['disagree_rate'])} of test statements the verdict is not the highest-probability class, so the displayed confidence can belong to another class. Reported, not hidden.", {"bullet": True}),
    ], size=16, color=K.INK)
    notes[10] = "Left: the Firefly search and the final model. The mismatch arises because predict uses the class-weighted decision function while Platt scaling is fitted without that weighting; the demo model never showed it."

    # ---- 11 data flow and modules
    s = S[10]
    set_title(s, "Data flow and modules")
    wipe(s)
    K.picture(s, F.diag("04_dfd_level1.png"), 0.6, 1.3, 7.3, 5.6)
    rows = [["Module", "Role"], ["preprocessing.py", "clean text, 25 style features"], ["dataset.py", "load CSV, demo corpus"], ["firefly.py", "mask + C + gamma search"],
            ["pipeline.py", "fit, predict, save, load"], ["backend/main.py", "FastAPI routes, retrain job"], ["frontend/src", "React screens and charts"],
            ["train.py, predict.py", "command-line entry points"], ["app.py", "Streamlit front end"]]
    K.table(s, 8.2, 1.5, [2.2, 2.33], rows, font_size=12, row_h=0.52)
    notes[11] = "Level-1 data flow: detect, batch detect, train/retrain and statistics, with the model bundle and the temporary CSV as stores."

    # ---- 12 database and UML
    s = S[11]
    set_title(s, "Database design and UML")
    wipe(s)
    K.picture(s, F.diag("09_er_diagram.png"), 0.6, 1.35, 5.3, 4.9)
    K.picture(s, F.diag("06_uml_class_diagram.png"), 6.1, 1.3, 6.6, 5.1)
    K.text(s, 0.6, 6.5, 12.13, 0.45, "There is no SQL database: the model bundle is a file, retrain jobs live in memory, uploaded CSVs are temporary. The ER diagram models these stores.", size=12, color=K.TXT, align="c", italic=True)
    notes[12] = "The project stores files, not rows. The class diagram shows the real classes: the FastAPI module, the pipeline, the Firefly optimiser and the library modules."

    # ---- 13 technology stack
    s = S[12]
    set_title(s, "Technology stack")
    wipe(s)
    rows = [["Layer", "Technology"], ["Machine learning", "Python 3.11, scikit-learn 1.9 (TF-IDF, PCA, SVC), NumPy, SciPy, pandas, joblib"],
            ["Optimisation", "Firefly algorithm written for the project (src/firefly.py)"], ["API", "FastAPI, uvicorn, python-multipart (port 8406)"],
            ["Front end", "React 18, TypeScript, Vite 8, Tailwind 4, TanStack Query, Recharts, GSAP (port 5406)"],
            ["Other front ends", "Streamlit app and command line (train.py, predict.py), unchanged"], ["Testing", "Python check scripts (smoke and unit/edge), Playwright for screenshots"]]
    K.table(s, 0.6, 1.5, [3.0, 9.13], rows, font_size=16, row_h=0.68, first_bold=True)
    notes[13] = "Everything is open source and starts with run.bat. The ML library runs without the web layers."

    # ---- 14 implementation progress
    s = S[13]
    set_title(s, "Implementation progress")
    wipe(s)
    rows = [["Component", "Status", "Evidence"],
            ["Preprocessing, PCA, Firefly, MSVM pipeline", "Done", "src/, train.py, predict.py"],
            ["Real-data experiments on LIAR2", "Done", "3 seeds, experiments/liar2/*.json"],
            ["FastAPI service (8 endpoints)", "Done", f"{F.N_SMOKE} smoke + {F.N_UNIT_API} API edge checks pass"],
            ["React interface (6 screens)", "Done", "17 documentation screenshots, build passes"],
            ["Stemming (objective 2)", "Not done", "cleaning and stop words only"],
            ["Better probabilities / calibration", "Not done", "verdict and probability mismatch known"],
            ["Other datasets, user study, load test", "Not done", "stated as limits"]]
    K.table(s, 0.6, 1.5, [4.8, 1.6, 5.73], rows, font_size=14, row_h=0.62)
    notes[14] = "State of the code base. The honest gaps are stemming, calibration and evaluation beyond LIAR2."

    # ---- 15 preliminary results
    s = S[14]
    set_title(s, "Preliminary results: LIAR2 test split (3 seeds)")
    wipe(s)
    for i, (v, l, c) in enumerate([(F.f3(P["accuracy"]["mean"]), f"accuracy (n = {F.NTEST:,})", K.BLUE), (F.f3(P["f1_macro"]["mean"]), "macro-F1", K.GREEN),
                                   (F.f3(PD["f1_macro"]["mean"]), "PCA + default SVM macro-F1", K.ORANGE), (F.f3(LR["f1_macro"]["mean"]), "logistic regression, all data, macro-F1", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=28, lsize=11)
    K.picture(s, F.fig("f01_model_comparison.png"), 0.6, 2.9, 7.9, 4.05)
    K.card(s, 8.7, 2.95, 4.03, 3.95, "Reading",
           f"Proposed macro-F1 {F.pm(P)}; random search {F.pm(RS)}; Firefly mask only {F.pm(MO)}. The differences are within noise, so no gain from the optimiser is shown. "
           "PARTIALLY TRUE is the hardest class (F1 0.24), FAKE the easiest (0.63). Always answering FAKE gets accuracy 0.576 but macro-F1 0.244.", K.BLUE, "i", body_size=13)
    notes[15] = ("Headline: macro-F1 about 0.43 and accuracy about 0.49, mean of three seeds. The optimiser does not clearly beat equal-budget random search or the default SVM, and TF-IDF linear models are better and far cheaper. "
                 "The base paper's 95.5 % is for a different, binary Arabic task on a GPU and is not comparable.")

    # ---- 16 output screens
    s = S[15]
    set_title(s, "Output screens")
    wipe(s)
    K.picture(s, F.shot("11_liar2_detect_fake.png"), 0.6, 1.3, 7.4, 4.65, border=True)
    K.picture(s, F.shot("16_liar2_batch_results.png"), 8.3, 1.3, 4.43, 2.77, border=True)
    K.picture(s, F.shot("18_liar2_analytics_charts.png"), 8.3, 4.2, 4.43, 2.77, border=True)
    K.text(s, 0.6, 6.05, 7.4, 0.9, "Detect screen on a real LIAR2 test statement: verdict FAKE at 83.8 %, class probabilities and cue words. Right: batch check of 120 real statements, and the analytics screen (confusion matrix, Firefly curve, baseline table).", size=12, color=K.TXT)
    notes[16] = "Real screenshots of the running app with the model trained on LIAR2. The Detect example is a real test statement, not a mock-up."

    # ---- 17 testing
    s = S[16]
    set_title(s, "Initial testing and observations")
    wipe(s)
    for i, (v, l, c) in enumerate([(str(F.N_TESTS), "automated checks passed", K.GREEN), (str(F.N_SMOKE), "API smoke checks", K.BLUE),
                                   (f"{F.N_UNIT_LIB} + {F.N_UNIT_API}", "library + API edge checks", K.ORANGE), ("pass", "frontend type-check and build", K.PURPLE)]):
        K.tile(s, 0.6 + i * 3.1, 1.45, 2.9, 1.25, v, l, c, vsize=26, lsize=12)
    rows = [["Observation", "Action"],
            [f"Confidence is informative: accuracy {F.f3(E['by_confidence'][2]['accuracy'])} at 0.5-0.6, {F.f3(E['by_confidence'][3]['accuracy'])} at 0.6-0.7", "Shown on the Detect screen"],
            ["A label column with only true/false is parsed as booleans by pandas and mislabelled", "Recorded as a known issue"],
            [f"Server latency {lat['predict_server_ms']['mean']:.1f} ms per article; {lat['batch']['2296']['server_s']} s for {F.NTEST:,} rows", "Measured on 127.0.0.1, CPU idle"]]
    K.table(s, 0.6, 3.1, [6.0, 6.13], rows, font_size=15, row_h=0.85)
    notes[17] = f"{F.N_SMOKE} smoke checks plus {F.N_UNIT_LIB} library and {F.N_UNIT_API} API edge checks. The observations shaped how results and limitations are reported."

    # ---- 18 challenges
    s = S[17]
    set_title(s, "Challenges and corrective actions")
    wipe(s)
    rows = [["Challenge", "Corrective action"],
            ["No real labelled data shipped; the demo corpus gives a meaningless 1.000", "Ran experiments on LIAR2 and reported the demo score as a pipeline check only"],
            ["Firefly needs thousands of SVM fits on a CPU", "3,000-statement training subsample; about 8 minutes per full fit"],
            ["Optimiser shows no clear gain", "Added random-search, mask-only and linear baselines; reported honestly"],
            ["Verdict and probabilities disagree on real data", "Documented with a screenshot; calibration listed as future work"],
            ["PCA cap keeps about 29 % of the variance", "Reported; larger cap or sparse methods planned"]]
    K.table(s, 0.6, 1.5, [5.6, 6.53], rows, font_size=15, row_h=0.85)
    notes[18] = "Each challenge changed either the experiment design or how results are reported."

    # ---- 19 next steps
    s = S[18]
    set_title(s, "Next steps")
    wipe(s)
    for i, (t, b_, col, gl) in enumerate([
        ("Calibrate", "Make verdict and confidence agree, for example by calibrating on a held-out split with class-prior correction.", K.BLUE, "1"),
        ("Strengthen features", "Try a larger PCA cap or truncated SVD, and compare with embeddings; add stemming.", K.GREEN, "2"),
        ("Evaluate wider", "More seeds with a significance test, ISOT and other datasets, and cross-dataset checks.", K.ORANGE, "3"),
        ("Tune the search", "Tune Firefly settings and test other optimisers at equal budget.", K.PURPLE, "4")]):
        K.card(s, 0.6 + i * 3.1, 1.6, 2.9, 3.3, t, b_, col, gl, body_size=16)
    notes[19] = "Plan to the final review. None of these items is claimed as done."

    # ---- 20 thank you
    thank_you(S[19])
    notes[20] = "Likely questions: why the optimiser does not help here; why only three classes; how the verdict and probabilities can disagree; what the base paper's numbers mean for us (nothing directly)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "13.Review_2_PPT.pptx")))
