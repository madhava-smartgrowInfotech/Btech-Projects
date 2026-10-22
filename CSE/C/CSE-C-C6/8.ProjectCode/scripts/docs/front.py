"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.TITLE.upper(), align="center", size=18, bold=True, space_after=6)
    d.p("TruthGlow: a three-class screening aid built on PCA, the Firefly algorithm and a multi-class SVM", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the creators of the LIAR and LIAR2 datasets, and the developers of the open-source software used in this "
        "project (Python, scikit-learn, NumPy, SciPy, pandas, FastAPI, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    a = F.A
    prop = a["proposed_test"]
    d.p("Fake news spreads faster than fact-checkers can respond, and many detectors force a claim into two classes or need large neural models. This project builds TruthGlow, "
        "which gives one of three verdicts - REAL, PARTIALLY TRUE or FAKE - with a confidence. Text becomes TF-IDF and stylometric features, compressed by Principal Component "
        "Analysis; the Firefly algorithm selects components and sets the parameters of a multi-class RBF support vector machine. A FastAPI layer and a React web application "
        "provide single and batch detection, analytics and retraining.")
    d.p(f"The shipped model was trained on a synthetic corpus, so its perfect scores are only a pipeline check. On real fact-checked statements (LIAR2, {F.NTRAIN:,} training "
        f"statements, {len(F.SEEDS)} seeds) the pipeline reaches test macro-F1 {F.pm(prop)} and accuracy {F.pm(prop, 'accuracy')}. The Firefly optimisation gave no gain over a default "
        f"SVM (macro-F1 {F.f3(a['pca_svc_default_test']['f1_macro']['mean'])}) or an equal-budget random search, and logistic regression on plain TF-IDF scored "
        f"{F.f3(a['ref_full_LogisticRegression_test']['f1_macro']['mean'])}. Server time is about {F.LAT[0]['predict_server_ms']['mean']:.1f} ms per article and {F.N_TESTS} automated "
        "checks pass. Better accuracy at lower cost is therefore not demonstrated.")
    d.p("**Keywords:** " + ", ".join(k.lower() if i else k for i, k in enumerate(F.KEYWORDS)) + ".", align="left")


ABBR = [
    ("API", "Application Programming Interface"), ("CLI", "Command Line Interface"), ("CPU", "Central Processing Unit"), ("CSV", "Comma-Separated Values"),
    ("CV", "Cross-Validation"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"), ("GPU", "Graphics Processing Unit"), ("HTTP", "Hypertext Transfer Protocol"),
    ("JSON", "JavaScript Object Notation"), ("LIAR / LIAR2", "Fact-checked statement benchmarks from PolitiFact (original and extended)"), ("MSVM", "Multi-class Support Vector Machine"),
    ("NLP", "Natural Language Processing"), ("PCA", "Principal Component Analysis"), ("RBF", "Radial Basis Function"), ("REST", "Representational State Transfer"),
    ("SMOTE", "Synthetic Minority Over-sampling Technique"), ("SPA", "Single-Page Application"), ("SQL", "Structured Query Language"), ("SVM", "Support Vector Machine"),
    ("TF-IDF", "Term Frequency - Inverse Document Frequency"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"), ("OvR", "One-vs-Rest"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
