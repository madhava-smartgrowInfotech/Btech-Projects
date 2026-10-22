"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .refs import K as RK

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"

OBJECTIVES = [
    "Build learner profiles capturing interests, current skills, goals and completed courses.",
    "Construct a course catalogue with topic, level, prerequisite and outcome metadata.",
    "Implement collaborative filtering over learner interaction history.",
    "Implement content-based matching between learner profile and course metadata.",
    "Combine both into a hybrid recommendation engine and order the output into a coherent learning path.",
    "Evaluate using precision, recall and learner feedback on recommendation relevance.",
]


def c(*keys):
    return ", ".join(f"[{RK[k]}]" for k in keys)


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "course recommendation, hybrid ranking, learner success, uplift, offline policy evaluation"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors):
    cols = 3
    w, h, gx, gy = 3.85, 1.45, 0.29, 0.3
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c_ = divmod(i, cols)
        x, y = x0 + c_ * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, review_label, subtitle):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle)
        elif nm == "review label":
            K.replace_text(sh, review_label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s, body="Questions and suggestions are welcome."):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, body)


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    P = F.METRICS["policy_evaluation"]

    # ---- 1 title
    title_slide(S[0], "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH",
                f"{F.PRODUCT}: a hybrid course recommendation engine that ranks courses for each learner, estimates the chance of success and suggests the next best action")
    notes[1] = ("Introduce the team and the project title. One sentence: we build a web application that takes a learner profile and activity, ranks a course catalogue with an eight-term hybrid score "
                "and shows why each course was recommended.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey"), ("04", "Existing approaches and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; results come in later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: a next step chosen for the learner, with reasons")
    wipe(s)
    W, H = 3.85, 4.6
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Online platforms offer long course catalogues but little guidance. Popularity lists ignore a learner's skill level, goals and behaviour, new learners have no history, and most learners do not finish: the base paper notes completion rates are often below 10 %.", K.ORANGE, "!"),
        ("Our idea", "Build a learner profile, score every course with an eight-term hybrid score (text match, skill fit, topic, goal, rating, predicted success, behaviour, engagement) and estimate which learning action would help most. Every card carries a plain-language reason.", K.BLUE, "i"),
        ("The deliverable", "A web application: register, fill a profile, browse the catalogue, get ranked courses with reasons, follow a staged learning path, log activity and watch the ranking adapt, and inspect the model analytics.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=16)
    notes[3] = "Three parts: the problem, our idea, what we deliver. The completion-rate statement is quoted from the base paper's introduction."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Causal-AI pipeline on learner-week data: a predictive success model (XGBoost), uplift modelling of four actions (Quiz, Forum, Content, Homepage), an Actor-Critic reinforcement-learning agent with causal regularisation, off-policy evaluation (IPS, SNIPS, doubly robust) and a subgroup fairness audit.", K.BLUE, "M"),
        ("Data and reported results", "Open University Learning Analytics Dataset: 32,593 enrolments, about 1.7 M learner-week rows, 25,516 learners. XGBoost AUC 0.795, Brier 0.074; logged pass rate 0.589; Greedy-Uplift SNIPS 0.619 [0.613, 0.625].", K.GREEN, "R"),
        ("Limits stated by the authors", "Actions are the learner's own dominant activity, a proxy for system-recommended actions; no randomised trial; coarse action granularity; subgroup disparities by age and gender motivate fairness-aware extensions.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=14)
    notes[4] = ("Read the citation. The base paper prescribes learning actions; it does not recommend courses. We reuse its prediction, uplift, off-policy-evaluation and fairness-audit ideas for a course recommendation engine. "
                "We do not implement its Actor-Critic agent and we do not use real OULAD data.")

    # ---- 5 literature survey
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Approach", "Data", "Reported finding"],
            [f"Isiwu et al., 2026 {c('base')} (base)", "Uplift modelling + Actor-Critic RL, off-policy evaluation", "OULAD", "AUC 0.795; Greedy-Uplift SNIPS 0.619 vs logged 0.589"],
            [f"Khreis et al., 2026 {c('khreis')}", "BERT with contrastive learning and isotropy regularisation; Top-N by text query", "500+ engineering courses (uOttawa)", "More discriminative embeddings than vanilla BERT"],
            [f"Gharahighehi et al., 2025 {c('gharahighehi')}", "Collaborative filtering plus survival analysis (time to event)", "Three public MOOC datasets", "Better than CF on interactions alone (two measures)"],
            [f"Sujay et al., 2026 {c('eduguide')}", "TF-IDF + cosine similarity, rating-weighted hybrid score", "Coursera course data", "Hybrid of similarity and rating (metrics not recorded)"],
            [f"Survey item {c('s_ml')}*, 2025", "Course recommendation with machine learning", "Not recorded", "Not recorded"],
            [f"Survey item {c('s_lightgcn')}*, 2025", "LightGCN graph model on learning behaviour data", "Not recorded", "Not recorded"],
            [f"Survey item {c('s_review')}*, 2026", "Literature review of course recommendation for e-learning", "Review", "Not applicable"]]
    K.table(s, 0.6, 1.4, [2.9, 3.6, 2.5, 3.13], rows, font_size=10.5, row_h=0.6, header_size=12)
    K.text(s, 0.6, 6.5, 12.13, 0.4, "* Recorded in the batch literature survey by title and year only; authors and results were not recorded and are not guessed.", size=11, color=K.TXT, italic=True)
    notes[5] = ("The pattern: content similarity, graph and collaborative methods on one side, causal and prescriptive methods on the other. Rows 1 to 4 were read in their source; the three starred rows are "
                "known only by title and year from the batch survey, so no results are quoted for them.")

    # ---- 6 supporting evidence (chart)
    s = S[5]
    set_title(s, "Supporting evidence from the base paper")
    wipe(s)
    K.bar_chart(s, 0.6, 1.4, 7.4, 5.2, ["Logged pass rate\n(observed)", "Greedy-Uplift\npolicy (SNIPS)", "XGBoost AUC\n(prediction)"], [("Base paper, OULAD", [0.589, 0.619, 0.795])],
                colors=[K.BLUE], number_format="0.000", font_size=12, max_value=1.0)
    K.card(s, 8.3, 1.5, 4.43, 2.45, "Takeaway", "In the base paper, choosing the action with the highest estimated uplift moves the off-policy value from 0.589 to 0.619, with a good predictive base (AUC 0.795).", K.GREEN, "↑", body_size=14)
    K.card(s, 8.3, 4.15, 4.43, 2.45, "Caveat", "Its own Table 4 shows IPS at 9.956 against SNIPS 0.643: off-policy estimates can be unstable. We report ours with their effective sample size.", K.ORANGE, "!", body_size=14)
    notes[6] = "All values are quoted from the base paper (real OULAD). They are not our results; our own numbers, on synthetic data, appear in Review 2 and the final review."

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.55, 12.13, 1.1, [("Long course catalogue", K.NAVY), ("Search or popularity list", K.NAVY), ("Learner picks by title", K.NAVY), ("Enrols", K.NAVY), ("Progress not acted on", K.NAVY)], size=16, gap=0.4)
    K.card(s, 0.6, 3.1, 5.95, 3.55, "Search and popularity", "The learner filters by keyword or follows what is popular. Nothing in the list reflects the learner's current skill, goals or what they did last week, and a new learner gets the same list as everyone else.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 3.1, 5.95, 3.55, "Prediction without prescription", "Learning-analytics models can say who is likely to fail, but stop short of what to do about it " + c("base") + ". Course recommenders based on collaborative filtering struggle with cold start " + c("khreis") + ".", K.ORANGE, "2", body_size=17)
    notes[7] = "Describe today's typical path from catalogue to enrolment. Both statements on the right come from the base paper and the Khreis et al. introduction."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    items = [
        ("One list for everyone", "Popularity and keyword search ignore skill level, goals and behaviour, so the same courses are shown to a beginner and an expert.", K.ORANGE, "1"),
        ("Cold start", "Collaborative filtering needs interaction history; new learners and long-tail courses are poorly served " + c("khreis") + ".", K.BLUE, "2"),
        ("Prediction without action", "Success models flag at-risk learners but do not say which learning action to take next " + c("base") + ".", K.GREEN, "3"),
        ("Fragile evidence", "Effects are estimated from observational logs, not randomised trials, and off-policy estimates can be unstable " + c("base", "dudik") + ".", K.PURPLE, "4"),
    ]
    for i, (t, b_, col, gl) in enumerate(items):
        r, c_ = divmod(i, 2)
        K.card(s, 0.6 + c_ * 6.18, 1.5 + r * 2.7, 5.95, 2.5, t, b_, col, gl, body_size=16)
    notes[8] = "Four limitations our design responds to: a profile-aware hybrid score, content-first cold start, an estimated next action, and honest reporting of off-policy evidence."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    gaps = [("Cold start", "Content and profile matching should work from the first minute, before any history exists."),
            ("Explanation", "A recommendation should come with the reasons behind it, not only a rank."),
            ("Next action", "Few course recommendation tools also estimate which learning action (quiz, content, forum) would help the learner most."),
            ("Application", "Many works stop at a notebook; a usable web system with profiles, activity log and learning path is rarer."),
            ("Honest evaluation", "Offline metrics, fairness by group and a stated sample size should sit beside every claim.")]
    for i, (k, v) in enumerate(gaps):
        y = 1.45 + i * 1.08
        K.box(s, 0.6, y, 2.6, 0.9, fill=K.NAVY, line=None, radius=0.06)
        K.text(s, 0.6, y, 2.6, 0.9, k, size=16, color=K.WHITE, bold=True, align="c", anchor="m")
        K.text(s, 3.5, y, 9.2, 0.9, v, size=17, color=K.INK, anchor="m")
    notes[9] = "Each gap follows from the limitations. The batch problem statement phrases it as a need for a hybrid engine that matches courses to a profile and sequences them into a learning path."

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.75, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.5, "Learners face too many options and no reliable way to judge which suits them; platform recommendations optimise engagement rather than learning, new learners get poor suggestions, and courses are suggested one by one without order. "
           "The need is a **hybrid recommendation engine** that matches courses to a learner's profile and sequences them into a coherent **learning path**.", size=19, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "Learner profile (interests, skill, goals, weekly hours) and logged activity"), ("Output", "Ranked courses with reasons, a staged learning path, a suggested next action"), ("Record", "Interactions, recommendation log, model analytics")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.5, 3.85, 1.7)
        K.text(s, x + 0.25, 4.62, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.1, 3.4, 1.0, b_, size=14, color=K.TXT)
    notes[10] = "The statement and the need are taken from the batch objectives and problem-statement document."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 16}) for i, o in enumerate(OBJECTIVES, 1)], size=24, color=K.INK)
    notes[11] = ("Six objectives from the project proposal. Reviews 2 and 3 report their status honestly: collaborative filtering across learners (objective 3) was not implemented, "
                 "and objective 6 is only partly met (offline metrics and a persona check, no user study).")

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Profile and catalogue", "Learner profile (education, skill level, interests, topics, goals, weekly hours); 24-course catalogue with topic, level, prerequisites and outcomes.", K.BLUE, "1"),
        ("Match", "TF-IDF text similarity between profile and course, plus skill fit, topic overlap, goal match and rating.", K.GREEN, "2"),
        ("Predict and adapt", "Learner-success model on 8 activity features; one-vs-rest estimate of action benefit; behavioural boost from logged events.", K.ORANGE, "3"),
        ("Rank and sequence", "Eight-term hybrid score, top-N with reasons, learning path ordered by level then score.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=14)
    K.box(s, 0.6, 5.2, 12.13, 1.45)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "E", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.5, "Evaluation plan", size=17, color=K.INK, bold=True)
    K.text(s, 1.8, 5.85, 10.7, 0.75, "Success-model AUC, F1 and Brier on a learner-level held-out split; offline policy evaluation; fairness by group; precision@k on persona checks. Collaborative filtering across learners is not planned for lack of a real interaction log.", size=14, color=K.TXT)
    notes[12] = "Four stages and an evaluation plan. The training data are an OULAD-like synthetic generator because the real dataset was not used; the final review states this next to every number."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.7, 5.5)
    K.text(s, 8.6, 1.5, 4.15, 5.3, [
        ("**Four tiers**", {"space_after": 8}),
        ("Presentation: React 19 interface (port 5407) and the original Jinja pages (port 8407).", {"bullet": True}),
        ("Application: FastAPI routes, token login, validation.", {"bullet": True}),
        ("Recommendation: hybrid ranker, success model and action-benefit models.", {"bullet": True}),
        ("Data: SQLite with six tables, the 24-course catalogue, trained model files.", {"bullet": True}),
    ], size=16, color=K.INK, space_after=8)
    notes[13] = "Follow the diagram left to right: the user interface, the FastAPI routes, the recommendation modules, and the data stores."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    outs = [("Working engine", "Eight-term hybrid ranking of a 24-course catalogue for any profile.", K.BLUE, "1"),
            ("Visible reasoning", "Match score, success estimate and a plain-language reason on every card.", K.GREEN, "2"),
            ("Adapting ranking", "Logged events change the next ranking immediately, without retraining.", K.ORANGE, "3"),
            ("Measured models", "Success-model metrics, offline policy evaluation and fairness audit, with limits stated.", K.PURPLE, "4"),
            ("Usable system", "Web application with login, profile, learning path, insights and analytics.", K.BLUE, "5"),
            ("Reproducibility", "One script retrains the models; one script regenerates every figure from JSON.", K.GREEN, "6")]
    for i, (t, b_, col, gl) in enumerate(outs):
        r, c_ = divmod(i, 3)
        K.card(s, 0.6 + c_ * 4.14, 1.5 + r * 2.7, 3.85, 2.5, t, b_, col, gl, body_size=15)
    notes[14] = "These are expected outcomes, not results. We will report what the measurements show, including where the system falls short."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    from .refs import REFS

    keys = ["base", "khreis", "gharahighehi", "eduguide", "oulad", "dudik", "uplift", "lightgcn", "hrl"]
    K.text(s, 0.6, 1.35, 12.13, 5.5, [(f"[{RK[k]}] {REFS[RK[k] - 1].split(' (listing')[0]}", {"space_after": 9}) for k in sorted(keys, key=RK.get)], size=14, color=K.INK)
    notes[15] = "Full references, including the survey items known only by title and year, are in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = "Likely questions: why a hybrid score and not collaborative filtering (no real interaction log); why synthetic data (real OULAD not used); what the uplift numbers mean (estimated benefit, not proven causal effect)."

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.scrub_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
