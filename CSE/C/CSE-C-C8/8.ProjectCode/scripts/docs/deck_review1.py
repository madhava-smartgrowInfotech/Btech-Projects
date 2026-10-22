"""7.Review_1_PPT.pptx: 16 slides, problem identification and research. Built from D:\\abstracts\\Templates\\Review_1_PPT.pptx."""

import os

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import facts as F
from . import pptkit as K
from .lit import (EXISTING_SYSTEM, EXPECTED_OUTCOMES, GAPS, LIMITATIONS_EXISTING, LIT_TABLE, OBJECTIVES, PROBLEM_STATEMENT,
                  PROPOSED_APPROACH, SUPPORT_EVIDENCE)
from .refs import REFS

TEMPLATE = r"D:\abstracts\Templates\Review_1_PPT.pptx"
LAYOUT_FOOTER_OLD = "The Hidden Risk of Social Media"


def is_ph(sh, *types):
    return sh.is_placeholder and sh.placeholder_format.type in types


def prepare(prs, review_label):
    # layout footer text carries the template's deck title
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and LAYOUT_FOOTER_OLD in sh.text_frame.text:
                K.replace_text(sh, f"{F.PRODUCT}  ·  {review_label}")
    cp = prs.core_properties
    cp.title = f"{F.FULL_TITLE}: {review_label}"
    cp.subject = f"Batch {F.BATCH_NO} ({F.SECTION})"
    cp.author = ", ".join(n for n, _ in F.STUDENTS)
    cp.keywords = "research paper evaluation, reviewer recommendation, NLP, sentence embeddings, peer review"
    cp.comments = ""
    cp.last_modified_by = ""


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, text, size)
            return sh


def wipe(slide):
    K.clear(slide, lambda sh: sh.is_placeholder)


def agenda(slide, items, colors, cols=3):
    n = len(items)
    rows = (n + cols - 1) // cols
    gx, gy = 0.29, 0.3
    w = (12.13 - gx * (cols - 1)) / cols
    h = min(1.45, (5.3 - gy * (rows - 1)) / rows)
    x0, y0 = 0.6, 1.55
    for i, (n_, label) in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (w + gx), y0 + r * (h + gy)
        K.box(slide, x, y, w, h)
        K.text(slide, x + 0.25, y + 0.14, 1.0, 0.5, n_, size=26, color=colors[i % len(colors)], bold=True, font="Cambria")
        K.text(slide, x + 0.25, y + 0.72, w - 0.5, 0.65, label, size=16, color=K.INK, bold=True)


def title_slide(s, subtitle, label):
    for sh in s.shapes:
        nm = sh.name
        if is_ph(sh, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE):
            K.replace_text(sh, F.TITLE, 34)
        elif is_ph(sh, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY):
            K.replace_text(sh, subtitle, 20)
        elif nm == "review label":
            K.replace_text(sh, label)
        elif nm == "team":
            K.replace_text(sh, "\n".join(f"{n_}  {r_}" for n_, r_ in F.STUDENTS))
        elif nm == "Text 6":
            K.replace_text(sh, f"{F.GUIDE}, {F.GUIDE_DESIGNATION}")
        elif nm == "Text 7":
            K.replace_text(sh, f"{F.DEPT}\n{F.COLLEGE_SHORT}  ·  Batch {F.BATCH_NO} ({F.SECTION})  ·  {F.YEAR}")


def thank_you(s):
    for sh in s.shapes:
        if is_ph(sh, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            K.replace_text(sh, "Thank you")
        elif is_ph(sh, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE):
            K.replace_text(sh, "Questions and suggestions are welcome.")


def build(out_path):
    prs = Presentation(TEMPLATE)
    prepare(prs, "Review 1")
    S = list(prs.slides)
    assert len(S) == 16, len(S)
    notes = {}
    C = [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE]

    # ---- 1 title
    title_slide(S[0], F.SUBTITLE, "REVIEW 1  ·  PROBLEM IDENTIFICATION AND RESEARCH")
    notes[1] = ("Introduce the team and the title. One sentence: PeerScope reads a research-paper PDF, gives a preliminary and clearly caveated assessment, "
                "and ranks candidate reviewers with a score whose every term can be inspected.")

    # ---- 2 contents
    s = S[1]
    set_title(s, "Contents")
    wipe(s)
    agenda(s, [("01", "Abstract and overview"), ("02", "Base paper"), ("03", "Literature survey and evidence"), ("04", "Existing system and limitations"),
               ("05", "Research gap"), ("06", "Problem statement and objectives"), ("07", "Proposed approach and architecture"),
               ("08", "Expected outcomes"), ("09", "References")], [K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE, K.BLUE, K.GREEN, K.PURPLE, K.ORANGE])
    notes[2] = "Agenda. This review is about why the problem matters and what is missing in current work; design and results come in the later reviews."

    # ---- 3 abstract
    s = S[2]
    set_title(s, "Abstract: from a PDF to a reviewer ranking")
    wipe(s)
    W, H = 3.85, 4.9
    for i, (t, body, col, gl) in enumerate([
        ("The problem", "Journals and conferences receive more submissions than their reviewer pools can absorb. Reviewer choice is manual, keyword matching favours shared vocabulary over expertise, and first-pass screening is inconsistent.", K.ORANGE, "!"),
        ("Our idea", "Read the paper locally, describe it with structural signals, a trained decision signal and a local language-model critique, then rank reviewers with sentence embeddings, keyword and method overlap, availability and a conflict penalty.", K.BLUE, "i"),
        ("The deliverable", "PeerScope, a local web application: PDF upload, preliminary assessment, related literature, an auditable reviewer ranking with its score breakdown, a reviewer directory, history and export. Decision support only.", K.GREEN, "✓"),
    ]):
        K.card(s, 0.6 + i * (W + 0.29), 1.5, W, H, t, body, col, gl, body_size=19)
    notes[3] = "Three points: the editorial bottleneck, the hybrid idea (assessment plus auditable matching), and what is delivered: a complete application whose final judgement stays with a human."

    # ---- 4 base paper
    s = S[3]
    set_title(s, "Base paper")
    wipe(s)
    K.box(s, 0.6, 1.4, 12.13, 1.1, fill=K.CARD)
    K.text(s, 0.82, 1.48, 11.7, 0.95, F.BASE_PAPER, size=13, color=K.INK, anchor="m")
    cw = 3.85
    cards = [
        ("Method", "Language models only induce keywords: an ensemble of producer models proposes them, a judge model picks one list, then deterministic normalisation. Matchers compare TF-IDF, SBERT (all-MiniLM-L6-v2) and keyword-graph scores. An audit trail keeps every stage.", K.BLUE, "M"),
        ("Data and reported results", "663 papers from 85 conferences, 1,989 assignments, 524 reviewer profiles. Best MRR@3 0.7778 (titles) and 0.7843 (abstracts). Most text-encoder differences are not statistically significant.", K.GREEN, "R"),
        ("Relevance to PeerScope", "Same principle: an auditable matching score. Different scope: PeerScope adds paper assessment, availability and conflict terms inside the score, a directory and export, and runs locally. The base paper does not assess the paper itself.", K.ORANGE, "L"),
    ]
    for i, (t, b_, col, gl) in enumerate(cards):
        K.card(s, 0.6 + i * (cw + 0.29), 2.75, cw, 4.0, t, b_, col, gl, body_size=16)
    notes[4] = ("Read the citation. The base paper improves and audits the paper-reviewer affinity matrix; it does not assess submissions and left conflict of interest and capacity out of its experiments. "
                "PeerScope will be measured on a much smaller setup, so numbers are never compared directly.")

    # ---- 5 literature survey table
    s = S[4]
    set_title(s, "Literature survey")
    wipe(s)
    rows = [["Work", "Year", "Approach", "Reported result / limitation"]]
    for i in (0, 1, 2, 3, 4, 5, 6):
        r = LIT_TABLE[i]
        rows.append([r[0].replace(" (base paper)", " (base)"), r[1], r[2], r[3]])
    K.table(s, 0.6, 1.35, [2.5, 0.7, 4.5, 4.43], rows, font_size=13, row_h=0.66, header_size=13, align=["l", "c", "l", "l"])
    notes[5] = ("The pattern: reviewer-assignment research (learned suitability, expertise models, embeddings, an auditable LLM-keyword pipeline), decision prediction from review corpora, "
                "LLM-written reviews, and integrity studies. Figures are quoted only where they were read from the paper or its record. The complete list of 32 verified references is in the report.")

    # ---- 6 supporting evidence (chart + list)
    s = S[5]
    set_title(s, "Supporting evidence: why automation needs care")
    wipe(s)
    gf = K.bar_chart(s, 0.6, 1.4, 5.7, 5.2, ["Committee\ndisagreement\non outcomes [13]", "Hidden text\nsteers ChatGPT\nreview [2]", "Hidden text\nsteers Gemini\nreview [2]"],
                     [("Share of cases", [0.23, 0.78, 0.86])], colors=[K.BLUE], number_format="0%", font_size=12, max_value=1.0)
    for idx in (1, 2):
        pt = gf.chart.plots[0].series[0].points[idx]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = K.ORANGE
    K.text(s, 6.6, 1.45, 6.1, 5.3, [(e, {"bullet": True, "space_after": 11}) for e in SUPPORT_EVIDENCE], size=17, color=K.INK)
    K.text(s, 6.6, 6.35, 6.1, 0.5, "Different studies and protocols; the bars are context, not a comparison.", size=12, color=K.TXT, italic=True)
    notes[6] = ("Two reasons for the design. Human review is itself noisy (NeurIPS 2021: about 23 percent disagreement between independent committees). "
                "And LLM reviewers can be steered by hidden text (78 and 86 percent in one study), which is why PeerScope keeps a deterministic fallback, processes locally and frames its output as decision support.")

    # ---- 7 existing system
    s = S[6]
    set_title(s, "Existing system and current approaches")
    wipe(s)
    K.flow(s, 0.6, 1.5, 12.13, 1.1, [("Submission arrives", K.NAVY), ("Editor picks reviewers", K.NAVY), ("Reviewers write reviews", K.NAVY), ("Decision", K.NAVY)], size=16, gap=0.5)
    K.card(s, 0.6, 2.95, 5.95, 3.8, "Manual assignment", "\n".join(EXISTING_SYSTEM[i] for i in (0, 4)) + "\nA mismatched reviewer lowers review quality and delays the decision.", K.BLUE, "1", body_size=17)
    K.card(s, 6.78, 2.95, 5.95, 3.8, "Separate automated tools", "\n".join(EXISTING_SYSTEM[i] for i in (1, 2, 3)), K.ORANGE, "2", body_size=17)
    notes[7] = "Today an editor chooses reviewers by memory, bidding or keyword search. Research tools exist for matching, screening and review writing, but they are separate and rarely explain themselves."

    # ---- 8 limitations
    s = S[7]
    set_title(s, "Limitations of existing methods")
    wipe(s)
    heads = ["Weak reviewer fit", "Opaque scores", "Proxy labels", "Scarce real data"]
    for i in range(4):
        r, c = divmod(i, 2)
        K.card(s, 0.6 + c * 6.18, 1.45 + r * 2.35, 5.95, 2.15, heads[i], LIMITATIONS_EXISTING[i], C[i], str(i + 1), body_size=19)
    K.box(s, 0.6, 6.25, 12.13, 0.75, fill=K.CARD)
    K.text(s, 0.85, 6.25, 11.7, 0.75, "Also: " + LIMITATIONS_EXISTING[4], size=15, color=K.INK, anchor="m")
    notes[8] = "Five limitations from the literature survey. Each maps to a design decision later: visible score breakdown, availability and conflict terms in the score, local processing, and honest statements about data."

    # ---- 9 research gap
    s = S[8]
    set_title(s, "Research gap")
    wipe(s)
    for i, g in enumerate(GAPS):
        y = 1.4 + i * 0.9
        K.circle(s, 0.6, y + 0.1, 0.62, C[i % 4], str(i + 1), 18)
        K.text(s, 1.55, y, 11.15, 0.82, g, size=20, color=K.INK, anchor="m")
    notes[9] = ("The batch literature survey states the gap as combining auditable assignment with preliminary paper evaluation so both halves of the editorial workload are supported. "
                "We target gaps 1, 2, 3, 5 and 6; gap 4 shapes the responsible-use framing.")

    # ---- 10 problem statement
    s = S[9]
    set_title(s, "Problem statement")
    wipe(s)
    K.box(s, 0.6, 1.45, 12.13, 2.9, fill=K.CARD)
    K.text(s, 0.95, 1.6, 11.45, 2.65, [(p, {"bullet": True, "space_after": 8}) for p in PROBLEM_STATEMENT], size=21, color=K.INK, anchor="m")
    for i, (t, b_) in enumerate([("Input", "A research-paper PDF (up to 20 MB) and a directory of reviewer profiles"), ("Output", "Preliminary assessment, related literature and a ranked list of reviewers with score breakdowns"), ("Record", "Stored analyses with JSON and PDF export; every number traceable")]):
        x = 0.6 + i * 4.14
        K.box(s, x, 4.65, 3.85, 1.9)
        K.text(s, x + 0.25, 4.77, 3.4, 0.4, t, size=18, color=K.INK, bold=True)
        K.text(s, x + 0.25, 5.25, 3.4, 1.25, b_, size=16, color=K.TXT)
    notes[10] = "Read the statement slowly. The need is a framework that assesses submissions and assigns each to the most topically expert reviewer, as decision support for the editor."

    # ---- 11 objectives
    s = S[10]
    set_title(s, "Objectives")
    wipe(s)
    K.text(s, 0.6, 1.4, 12.13, 5.3, [(f"{i}.  {o}", {"space_after": 20}) for i, o in enumerate(OBJECTIVES, 1)], size=25, color=K.INK)
    notes[11] = "The six objectives from the proposal. Later reviews report which were met; the corpus of real publication profiles (1) and independent expert judgement (6) will be only partly met."

    # ---- 12 proposed approach
    s = S[11]
    set_title(s, "Proposed approach")
    wipe(s)
    cw = 2.88
    for i, (t, b_, col, gl) in enumerate([
        ("Read", "PyMuPDF extracts text, sections and references locally; keywords, methods and topics come from NLP features.", K.BLUE, "1"),
        ("Assess", "Structural signals, a TF-IDF and logistic-regression decision signal, a local language-model critique with a visible fallback, and related works from OpenAlex.", K.GREEN, "2"),
        ("Match", "Sentence-embedding similarity, keyword and method overlap, availability, minus a conflict penalty.", K.ORANGE, "3"),
        ("Explain", "Every recommendation shows its score terms, matching words and conflict check; results are stored and exportable.", K.PURPLE, "4"),
    ]):
        K.card(s, 0.6 + i * (cw + 0.2), 1.45, cw, 3.55, t, b_, col, gl, body_size=16)
    K.box(s, 0.6, 5.2, 12.13, 1.5)
    K.circle(s, 0.85, 5.5, 0.7, K.NAVY, "S", 18)
    K.text(s, 1.8, 5.35, 10.7, 0.45, "Match score = 0.55 expertise + 0.20 keyword + 0.15 method + 0.10 availability - 0.60 conflict", size=16, color=K.INK, bold=True)
    K.text(s, 1.8, 5.8, 10.7, 0.85, "Scaled to 0-100. Delivered as a FastAPI service and a React web application over SQLite, with no paid API and no GPU; the final judgement stays with a human.", size=14, color=K.TXT)
    notes[12] = "Four stages from PDF to explained recommendation. " + " ".join(PROPOSED_APPROACH[3:4]) + " The weights live in one configuration file."

    # ---- 13 architecture
    s = S[12]
    set_title(s, "High-level proposed architecture")
    wipe(s)
    K.picture(s, F.diag("01_system_architecture.png"), 0.6, 1.3, 7.4, 5.5)
    K.text(s, 8.3, 1.5, 4.4, 5.2, [
        ("**Four tiers**", {"space_after": 8}),
        ("Presentation: React and Vite web app (port 5408).", {"bullet": True}),
        ("Application: FastAPI service (port 8408), PDF-only upload with size and header checks.", {"bullet": True}),
        ("Services: extraction, features, decision signal, review, literature, embeddings, matching, export.", {"bullet": True}),
        ("Data: SQLite, model files and cached models on the local disk; only the OpenAlex query leaves the machine.", {"bullet": True}),
    ], size=17, color=K.INK)
    notes[13] = "Follow the diagram top to bottom: the browser uploads the PDF, the API runs the pipeline of services, and the local files hold the directory, the stored analyses and the trained model."

    # ---- 14 expected outcomes
    s = S[13]
    set_title(s, "Expected outcomes and contributions")
    wipe(s)
    heads = ["Working application", "Auditable ranking", "Screening aid", "Measured, with gaps", "Human in the loop"]
    for i, (h, b_) in enumerate(zip(heads, EXPECTED_OUTCOMES)):
        if i < 3:
            K.card(s, 0.6 + i * 4.14, 1.5, 3.85, 2.55, h, b_, C[i % 4], str(i + 1), body_size=18)
        else:
            K.card(s, 0.6 + (i - 3) * 6.18, 4.3, 5.95, 2.4, h, b_, C[i % 4], str(i + 1), body_size=18)
    notes[14] = "Expected outcomes and contributions. The later reviews report measured results for the decision model and the matcher, including what could not be measured."

    # ---- 15 references
    s = S[14]
    set_title(s, "References")
    wipe(s)
    idx = [1, 2, 6, 7, 10, 11, 13, 19, 22, 24]
    K.text(s, 0.6, 1.3, 12.13, 5.6, [(f"[{n}] {REFS[n - 1]}".replace("*", ""), {"space_after": 4}) for n in idx], size=13, color=K.INK)
    notes[15] = "Selected references; the complete numbered list of 32 verified entries is in the project report."

    # ---- 16 thank you
    thank_you(S[15])
    notes[16] = ("Likely questions: why a classical classifier for the decision signal (the data is synthetic and the signal is only one input); how conflicts of interest are found "
                 "(literal strings, stated limitation); why a local language model (confidentiality and no paid API).")

    for i, sl in enumerate(S, 1):
        K.set_notes(sl, notes[i])
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    K.rename_theme(out_path, F.PRODUCT)
    return out_path


if __name__ == "__main__":
    print(build(os.path.join(F.BATCH, "7.Review_1_PPT.pptx")))
