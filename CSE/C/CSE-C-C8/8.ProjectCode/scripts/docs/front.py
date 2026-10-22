"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(2)
    d.p(F.FULL_TITLE.upper(), align="center", size=17, bold=True, space_after=6)
    d.p("PeerScope: a local web application that turns a research-paper PDF into a preliminary assessment and an auditable reviewer ranking", align="center", size=12, italic=True, space_after=8)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=6)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=6)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=8)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.ADDRESS, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.FULL_TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, F.DEPT), ("Principal", F.PRINCIPAL, F.COLLEGE_SHORT)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.FULL_TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE_SHORT}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p(f"Place: {F.PLACE}", align="left", space_after=0)
    d.p(f"Date: {F.DATE}", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE_SHORT}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the maintainers of the open-source software used in this project (FastAPI, React, Vite, scikit-learn, "
        "sentence-transformers, PyMuPDF, Hugging Face Transformers, PyTorch and SQLAlchemy), the OpenAlex project for its open scholarly index, and the Qwen team for the open-weight language model.")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    sh, mm = F.SHIPPED, F.MINILM
    d.p("Journals and conferences receive more submissions than their reviewer pools can absorb. Choosing a suitable reviewer is slow and depends on an editor's personal knowledge of the field, "
        "keyword matching favours reviewers who share vocabulary rather than expertise, and first-pass screening of submissions is inconsistent. This project presents PeerScope, a local web application "
        "that turns a research-paper PDF into a preliminary, clearly caveated assessment and an auditable ranking of candidate reviewers. The paper is extracted with PyMuPDF, split into sections and "
        "described by structural signals; a TF-IDF and logistic-regression model gives a historical decision signal; a locally run instruction-tuned language model (Qwen2.5-1.5B) drafts a critique when it can be loaded; "
        "and related literature is retrieved from OpenAlex. Reviewers are ranked by a transparent score that combines sentence-embedding similarity (all-MiniLM-L6-v2), keyword and method overlap, availability and a "
        "conflict-of-interest penalty, and every recommendation shows its score breakdown. The work extends the auditable-matching idea of the base paper [1] with paper assessment, a reviewer directory, history and export.")
    d.p(f"On a held-out split of {F.DM['test']} synthetic labelled papers the decision model reaches {F.pct(sh['accuracy'])} accuracy (macro F1 {F.f3(sh['f1_macro'])}); because the labels were generated from the same phrases that appear in the text, "
        f"no text classifier can exceed about {F.pct(F.GEN['best_possible_accuracy_from_phrases'], 0)} on that data, so the model demonstrates the pipeline and is not evidence of predicting real peer-review decisions. "
        f"On 40 probe abstracts written by the team, the hybrid ranking puts an acceptable reviewer first in {F.pct(mm['full_hybrid_shipped_weights']['top1'], 0)} of cases with sentence embeddings and {F.pct(F.TFIDF['full_hybrid_shipped_weights']['top1'], 0)} with the offline TF-IDF fallback; "
        f"all {mm['conflict']['detected_by_kind']['conflict term']['n'] + mm['conflict']['detected_by_kind']['affiliation']['n'] + mm['conflict']['detected_by_kind']['full name']['n']} literal conflict links injected into papers were flagged, and none of the paraphrased links was. "
        f"The deterministic analysis takes seconds; with the local language model one analysis of a three-page paper takes {F.lat_demo_text()} on a CPU-only PC. "
        "The reviewer pool is a set of six demo profiles, the probe labels are the team's own, novelty is not assessed, and no real peer-review data or user study was used; the base paper's results, obtained on 663 papers and 524 reviewers, are therefore not comparable.")
    d.p("**Keywords:** artificial intelligence, research paper evaluation, reviewer recommendation, natural language processing, semantic similarity, sentence embeddings, large language model, peer review, conflict of interest, decision support, FastAPI, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("COI", "Conflict of Interest"),
    ("CORS", "Cross-Origin Resource Sharing"), ("CPU", "Central Processing Unit"), ("CRUD", "Create, Read, Update, Delete"), ("CSV", "Comma-Separated Values"),
    ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"), ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"),
    ("LLM", "Large Language Model"), ("MAP", "Mean Average Precision"), ("ML", "Machine Learning"), ("MRR", "Mean Reciprocal Rank"),
    ("NLP", "Natural Language Processing"), ("ORM", "Object-Relational Mapping"), ("PDF", "Portable Document Format"), ("RAP", "Reviewer Assignment Problem"),
    ("REST", "Representational State Transfer"), ("ROC", "Receiver Operating Characteristic"), ("SBERT", "Sentence-BERT"), ("SPA", "Single-Page Application"),
    ("SQL", "Structured Query Language"), ("TF-IDF", "Term Frequency - Inverse Document Frequency"), ("UI", "User Interface"), ("UML", "Unified Modeling Language"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
