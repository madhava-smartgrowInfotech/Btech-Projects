"""
Builds the six chapter documents and the full report (17.Final Document.docx) into the batch folder.

    python -m scripts.docs.build_docs            # from 8.ProjectCode
    python -m scripts.docs.build_docs --only 9   # one document
"""

import argparse
import os
import sys

from . import appendix, ch1, ch2, ch3_4, ch5, ch6, ch7, ch8, ch9_10, front, shared
from . import facts as F
from .blocks import H1, H2, H3, P, TAB, render
from .kit import Doc
from .refs import REFS

NAMES = {
    9: "9.METHODOLOGY&NOVELTY&KEY CONTRIBUTIONS.docx",
    10: "10.SCOPE & LIMITATIONS.docx",
    11: "11. BLOCK DIAGRAMS ALL.docx",
    12: "12.DESIGN &IMPLEMENTATION.docx",
    14: "14.RESULTS.docx",
    15: "15.CONCLUSION&FUTURE SCOPE&REFERENCES.docx",
    17: "17.Final Document.docx",
}


def new_doc(doc_title, compact=False):
    d = Doc(doc_title, header_text=f"{F.PRODUCT}  |  {doc_title}", compact=compact)
    return d


def title_block(d, doc_title, note=True):
    d.footer_numbers()
    d.spacer(24)
    d.p(F.TITLE.upper(), align="center", size=16, bold=True, space_after=6)
    d.p(doc_title, align="center", size=14, bold=True, space_after=14)
    d.p(f"Product: {F.PRODUCT}", align="center", size=12, space_after=2)
    d.p(f"{F.DEPT}, {F.COLLEGE}", align="center", size=12, space_after=2)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})  ·  Academic year {F.YEAR}", align="center", size=12, space_after=12)
    rows = [[str(i), n, r] for i, (n, r) in enumerate(F.STUDENTS, 1)]
    d.table("Project team", ["S.No.", "Student name", "Roll number"], rows, widths=[0.7, 3.4, 1.8], size=11, align=["c", "l", "c"])
    d.p(f"**Project guide:** {F.GUIDE}", align="center", size=12, space_after=10)
    if note:
        d.p("Chapter and section numbers quoted in the text (for example Section 6.6 or Chapter 8) refer to the full project report, 17.Final Document. "
            "Figures and tables in this document are numbered from 1.", size=10.5, italic=True, space_after=6)


def save(d, name):
    path = d.save(os.path.join(F.BATCH, name))
    print("wrote", path)
    return path


# ----------------------------------------------------------------- chapter documents
def doc9():
    d = new_doc("Methodology, Novelty and Key Contributions")
    title_block(d, "Methodology, Novelty and Key Contributions")
    d.h1("1. OVERVIEW OF THE METHODOLOGY", page_break=False)
    d.p("This document describes how the system works from end to end and what is new in it. Part 1 gives the pipeline, the data, the algorithms and the reasons they were chosen; "
        "Part 2 sets the work against the base paper and separates what is new from what is standard; Part 3 lists the contributions.")
    render(d, ch6.methodology_blocks())
    d.h1("2. NOVELTY AGAINST THE BASE PAPER")
    render(d, ch9_10.novelty_blocks())
    d.h1("3. KEY CONTRIBUTIONS")
    render(d, ch9_10.contribution_blocks())
    return save(d, NAMES[9])


def doc10():
    d = new_doc("Scope and Limitations")
    title_block(d, "Scope and Limitations")
    d.h1("1. SCOPE", page_break=False)
    render(d, ch9_10.scope_blocks())
    uc = [x for x in ch5.DIAGRAMS if x[0] == "uc"][0]
    d.figure(F.diag(uc[2]), "What the user can do within the scope of the system (use-case diagram)", 4.6)
    d.h1("2. ASSUMPTIONS")
    render(d, ch9_10.assumption_blocks())
    d.h1("3. LIMITATIONS")
    render(d, ch9_10.limitation_blocks())
    d.h1("4. WHAT WAS NOT MEASURED")
    render(d, ch9_10.not_measured_blocks())
    return save(d, NAMES[10])


def doc11():
    d = new_doc("Block Diagrams")
    title_block(d, "Block Diagrams (all)")
    d.h1("1. INDEX OF DIAGRAMS", page_break=False)
    rows = [[str(i), t, f"`docs/diagrams/{f}`"] for i, (_, t, f, _, _) in enumerate(ch5.DIAGRAMS, 1)]
    d.p("Every diagram is generated from the real design by `scripts/make_diagrams.py` (matplotlib, 150 dpi) and mirrors the code in `app/routes.py`, `app/api.py`, `app/services/`, `app/utils/validation.py` and `scripts/`.")
    d.table("Diagrams in this document", ["No.", "Diagram", "Image file"], rows, widths=[0.5, 3.0, 2.8], size=10.5, align=["c", "l", "l"])
    d.h1("2. DIAGRAMS")
    for key, title, file, caption, text in ch5.DIAGRAMS:
        d.h2(title)
        d.figure(F.diag(file), caption, 6.0 if key not in ("flowchart", "uc") else 5.0)
        d.p(text)
    return save(d, NAMES[11])


def doc12():
    d = new_doc("Design and Implementation")
    title_block(d, "Design and Implementation")
    d.h1("1. DESIGN", page_break=False)
    render(d, [b for b in ch5.blocks() if b[0] != "h1"])
    d.h1("2. IMPLEMENTATION")
    render(d, ch6.application_blocks())
    d.h1("3. KEY CODE WALK-THROUGHS")
    render(d, ch9_10.walkthrough_blocks())
    d.h1("4. IMPLEMENTATION DECISIONS AND THEIR REASONS")
    render(d, ch9_10.decision_blocks())
    return save(d, NAMES[12])


def doc14():
    d = new_doc("Results")
    title_block(d, "Results")
    render(d, ch8.blocks(), h1_page_break=False)
    return save(d, NAMES[14])


def doc15():
    d = new_doc("Conclusion, Future Scope and References")
    title_block(d, "Conclusion, Future Scope and References")
    render(d, ch9_10.conclusion_blocks(), h1_page_break=False)
    d.p(f"The figures below summarise the measurements behind the conclusions: the first chart compares the four classifiers on the same {F.N_TEST:,} held-out records and the second shows how "
        "precision and recall of the selected model trade off as the decision threshold moves.")
    d.figure(F.fig("model_comparison.png"), "The four classifiers compared on the held-out test split", 5.6)
    d.figure(F.fig("threshold_sweep.png"), "Threshold analysis of the selected XGBoost model", 5.0)
    render(d, ch9_10.future_blocks())
    d.h1("REFERENCES")
    add_references(d)
    return save(d, NAMES[15])


def add_references(d):
    from .kit import add_runs
    from docx.shared import Inches, Pt
    for i, r in enumerate(REFS, 1):
        par = d.doc.add_paragraph()
        par.paragraph_format.left_indent = Inches(0.45)
        par.paragraph_format.first_line_indent = Inches(-0.45)
        par.paragraph_format.line_spacing = 1.15
        par.paragraph_format.space_after = Pt(5)
        add_runs(par, f"[{i}]  {r}", 10)


# ----------------------------------------------------------------- full report
_SEEN_FIGS = set()
MAX_CODE_LINES = 22
SKIP_SECTIONS = ("5.9 ", "A.8 ", "A.9 ", "B.7 ")      # report-only: shown in documents 12 and 14 instead


def trim(blocks):
    """Report-only trimming: a figure already shown earlier is not repeated, long code listings are shortened."""
    out = []
    skipping = False
    for b in blocks:
        if b[0] in ("h1", "h2"):
            skipping = b[0] == "h2" and b[1].startswith(SKIP_SECTIONS)
        if skipping:
            continue
        if b[0] == "fig":
            key = os.path.basename(b[1])
            if key in _SEEN_FIGS:
                continue
            _SEEN_FIGS.add(key)
        elif b[0] == "code":
            lines = b[1].rstrip("\n").split("\n")
            if len(lines) > MAX_CODE_LINES:
                b = ("code", "\n".join(lines[:MAX_CODE_LINES]) + "\n    # ... (listing shortened; full source in the repository)", b[2])
        out.append(b)
    return out


def R(d, blocks, **kw):
    render(d, trim(blocks), **kw)


def doc17():
    _SEEN_FIGS.clear()
    d = new_doc("Project Report", compact=True)
    front.cover(d)
    d.new_section(fmt="lowerRoman", start=1)
    front.certificate(d)
    front.declaration(d)
    front.acknowledgement(d)
    front.abstract(d)
    d.h1("TABLE OF CONTENTS")
    d.toc_field('TOC \\o "1-2" \\h \\z \\u')
    d.h1("LIST OF FIGURES")
    d.toc_field('TOC \\h \\z \\c "Figure"')
    d.h1("LIST OF TABLES")
    d.toc_field('TOC \\h \\z \\c "Table"')
    d.h1("LIST OF ABBREVIATIONS")
    front.abbreviations(d)
    d.new_section(fmt="decimal", start=1)
    R(d, ch1.blocks(), h1_page_break=False)
    R(d, ch2.blocks())
    R(d, ch3_4.blocks())
    R(d, ch5.blocks())
    R(d, ch6.blocks())
    R(d, ch7.blocks())
    R(d, ch8.blocks())
    R(d, ch9_10.chapter9_blocks())
    R(d, ch9_10.chapter10_blocks())
    d.h1("REFERENCES")
    add_references(d)
    R(d, appendix.code_blocks())
    R(d, appendix.guide_blocks())
    R(d, appendix.screens_blocks())
    return save(d, NAMES[17])


BUILDERS = {9: doc9, 10: doc10, 11: doc11, 12: doc12, 14: doc14, 15: doc15, 17: doc17}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", type=int, nargs="*")
    a = ap.parse_args()
    for k, fn in BUILDERS.items():
        if not a.only or k in a.only:
            fn()


if __name__ == "__main__":
    main()
