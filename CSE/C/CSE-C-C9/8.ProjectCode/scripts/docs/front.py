"""Front matter of the full report: cover, certificate, declaration, acknowledgement, abstract, abbreviations."""

import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from . import facts as F
from .kit import add_runs

BANNER = os.path.join(F.DOCS, "figures", "college_banner.png")


def _banner():
    """The institution's letterhead image, taken from the final-review template (written once to docs/figures)."""
    if os.path.exists(BANNER):
        return BANNER
    try:
        from pptx import Presentation

        prs = Presentation(r"D:\abstracts\Templates\Final_PPT.pptx")
        for sh in prs.slides[0].shapes:
            if sh.shape_type == 13:
                with open(BANNER, "wb") as f:
                    f.write(sh.image.blob)
                return BANNER
    except Exception:
        pass
    return None


def _banner_par(d, width=5.9):
    b = _banner()
    if b:
        par = d.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(b, width=Inches(width))


def cover(d):
    _banner_par(d)
    d.spacer(8)
    d.p(F.TITLE.upper(), align="center", size=19, bold=True, space_after=8)
    d.p("An explainable machine-learning system for individual financial profiling and economic forecasting", align="center", size=13, italic=True, space_after=14)
    d.p("A project report submitted in partial fulfilment of the requirements for the award of the degree of", align="center", size=12, space_after=6)
    d.p("BACHELOR OF TECHNOLOGY", align="center", size=15, bold=True, space_after=2)
    d.p("in", align="center", size=12, space_after=2)
    d.p("COMPUTER SCIENCE AND ENGINEERING", align="center", size=15, bold=True, space_after=12)
    d.p("Submitted by", align="center", size=12, bold=True, space_after=4)
    rows = [[n, r] for n, r in F.STUDENTS]
    t = d.doc.add_table(rows=len(rows), cols=2)
    t.alignment = 1
    for i, (n, r) in enumerate(rows):
        for j, v in enumerate((n, r)):
            c = t.rows[i].cells[j]
            c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            par.paragraph_format.space_after = Pt(2)
            add_runs(par, v, 12, True)
            c.width = Inches(3.2 if j == 0 else 1.6)
    d.spacer(6)
    d.p(f"Batch {F.BATCH_NO} ({F.SECTION})", align="center", size=12, space_after=10)
    d.p("Under the esteemed guidance of", align="center", size=12, space_after=2)
    d.p(F.GUIDE, align="center", size=13, bold=True, space_after=0)
    d.p(F.GUIDE_DESIGNATION, align="center", size=12, space_after=12)
    d.p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align="center", size=13, bold=True, space_after=2)
    d.p(F.COLLEGE.upper(), align="center", size=12, bold=True, space_after=2)
    d.p(F.CITY, align="center", size=11, space_after=2)
    d.p(f"Academic year {F.YEAR}", align="center", size=12, space_after=0)


def certificate(d):
    d.h1("CERTIFICATE", page_break=False)
    _banner_par(d, 5.2)
    d.p(f"This is to certify that the project report entitled **\"{F.TITLE}\"** is a bonafide record of the work carried out by "
        + ", ".join(f"**{n}** ({r})" for n, r in F.STUDENTS)
        + f", in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, "
          f"during the academic year {F.YEAR}, under my supervision.")
    d.spacer(30)
    t = d.doc.add_table(rows=3, cols=3)
    cells = [("Project Guide", F.GUIDE, F.GUIDE_DESIGNATION), ("Head of the Department", F.HOD, f"{F.DEPT}"), ("Principal", F.PRINCIPAL, F.COLLEGE)]
    for j, (a, b, c) in enumerate(cells):
        for i, v in enumerate((a, b, c)):
            cell = t.rows[i].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.line_spacing = 1.15
            add_runs(par, v, 11, i == 0)
    d.spacer(30)
    d.p("External Examiner: [NAME AND SIGNATURE]", align="left")
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def declaration(d):
    d.h1("DECLARATION")
    d.p(f"We hereby declare that the project report entitled **\"{F.TITLE}\"**, submitted to the {F.DEPT}, {F.COLLEGE}, in partial fulfilment of the "
        f"requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering, is a record of original work carried out by us under the "
        f"guidance of {F.GUIDE}, {F.GUIDE_DESIGNATION}. The results embodied in this report have not been submitted to any other university or institution for the award of any degree or diploma.")
    d.spacer(24)
    for n, r in F.STUDENTS:
        d.p(f"{n}   {r}", align="right", space_after=2)
    d.spacer(18)
    d.p("Place: Hyderabad", align="left", space_after=0)
    d.p("Date: [DATE]", align="left")


def acknowledgement(d):
    d.h1("ACKNOWLEDGEMENT")
    d.p(f"We express our sincere gratitude to our project guide, {F.GUIDE}, {F.GUIDE_DESIGNATION}, for valuable guidance, encouragement and constructive suggestions "
        "throughout this project.")
    d.p(f"We thank {F.HOD}, Head of the {F.DEPT}, for providing the facilities and support needed to complete this work, and {F.PRINCIPAL}, Principal, {F.COLLEGE}, "
        "for the encouragement given to us.")
    d.p("We thank the project coordinators and all faculty and staff members of the department for their help, and our families and friends for their constant support.")
    d.p("We also acknowledge the authors of the base paper, the maintainers of the UCI Machine Learning Repository whose Adult (Census Income) data we used, and the developers of the "
        "open-source software used in this project (scikit-learn, XGBoost, SHAP, statsmodels, Flask, React and Vite).")
    d.spacer(14)
    for n, r in F.STUDENTS:
        d.p(n, align="right", space_after=2)


def abstract(d):
    d.h1("ABSTRACT")
    M = F.METRICS
    c = F.CM[F.BEST]
    py = (F.MEAS or {}).get("pytest", {}).get("passed", 0)
    d.p("Judging an individual's financial position is usually done with a narrow set of financial parameters, a static score and no explanation, so non-linear interactions between "
        "demographic and financial factors are missed and the decision cannot be defended to the person concerned. This project builds FinSight AI, an explainable web application that turns "
        "twelve demographic and financial inputs (age, education, occupation, marital status, household role, country, weekly hours, investment gains and losses and others) into a financial profile "
        "(Emerging, Developing or Established), a Low, Moderate or High risk band, a financial-strength score, a per-person explanation and a 12-month economic forecast.")
    d.p(f"Four classifiers (Logistic Regression, Random Forest, Gradient Boosting and XGBoost) were trained on the public UCI Adult / Census Income data ({F.N_ROWS:,} records) inside leakage-safe "
        f"scikit-learn pipelines and compared on a stratified 80/20 split with 5-fold cross-validation. XGBoost was selected: on {F.N_TEST:,} held-out records it reaches accuracy {F.pct(M['accuracy'])}, "
        f"precision {F.pct(M['precision'])}, recall {F.pct(M['recall'])}, F1 {F.f3(M['f1'])} and ROC-AUC {F.f3(M['roc_auc'])} (cross-validated AUC {F.f3(M['cv_roc_auc_mean'])} +/- {F.f3(M['cv_roc_auc_std'])}). "
        "SHAP supplies global importance and live, plain-language contributions for every assessment, and a what-if screen compares a baseline with a changed scenario. A Holt-Winters model produces the 12-month "
        "forecast with a 95% band, but on a documented synthetic index, because the census data have no time dimension.")
    d.p("The system is served by a Flask JSON API and a React interface with a sample login. "
        f"{py} automated tests and a 35-check smoke test pass, and a prediction takes a median of {F.MEAS['api_latency_ms_127_0_0_1']['POST /api/predict']['median']:.0f} ms on a desktop CPU. "
        "The target is income above 50K, not creditworthiness; sex, race and marital status are model inputs and no fairness audit was done; real macroeconomic indicators and transaction data are not integrated. "
        "The results are stated with these limits and the comparison with the base paper, which studies hospitals, is given as context and not as a ranking. FinSight AI is informational and is not financial advice.")
    d.p("**Keywords:** financial identification, demographic analytics, economic forecasting, machine learning, XGBoost, SHAP, explainable AI, Holt-Winters, Flask, React.", align="left")


ABBR = [
    ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"), ("AUC", "Area Under the Curve"), ("CSS", "Cascading Style Sheets"),
    ("CSV", "Comma-Separated Values"), ("CV", "Cross-Validation"), ("DFD", "Data Flow Diagram"), ("ER", "Entity-Relationship"),
    ("ETS", "Error, Trend, Seasonality (exponential smoothing)"), ("FN / FP", "False Negative / False Positive"), ("GB", "Gradient Boosting"),
    ("HTTP", "Hypertext Transfer Protocol"), ("JSON", "JavaScript Object Notation"), ("LR", "Logistic Regression"), ("MAE", "Mean Absolute Error"),
    ("MAPE", "Mean Absolute Percentage Error"), ("ML", "Machine Learning"), ("PR", "Precision-Recall"), ("REST", "Representational State Transfer"),
    ("RF", "Random Forest"), ("RMSE", "Root Mean Squared Error"), ("ROC", "Receiver Operating Characteristic"), ("SHAP", "SHapley Additive exPlanations"),
    ("SPA", "Single-Page Application"), ("TN / TP", "True Negative / True Positive"), ("UCI", "University of California, Irvine"), ("UI", "User Interface"),
    ("UML", "Unified Modeling Language"), ("XAI", "Explainable Artificial Intelligence"), ("XGB", "Extreme Gradient Boosting (XGBoost)"),
]


def abbreviations(d):
    d.table("Abbreviations used in this report", ["Abbreviation", "Expansion"], [list(x) for x in ABBR], widths=[1.4, 4.9], size=11)
